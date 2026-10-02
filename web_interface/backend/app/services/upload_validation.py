"""The pure upload validators: magic bytes, the zip-bomb bound, text extraction
(python-docx, and the PDF sandbox for PDFs) and the minimum-text rule.

Shared by /upload (app/api/upload.py re-imports every name) and the emailed-CV
intake (#1298). Lives in services/ so the worker can use it without importing
app.api or app.auth (see run_service.UPLOAD_DIR).
"""
import io
import logging
import os
import tempfile
import zipfile
import zlib

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError

from app.services.pdf_sandbox import extract_pdf_text

logger = logging.getLogger(__name__)
ZIP_MAGIC = b"PK\x03\x04"
PDF_MAGIC = b"%PDF-"
PDF_EXTENSION = ".pdf"

# Minimum extracted text (characters) for a document to be considered readable.
# A real CV runs into the thousands of characters; anything below this is almost
# certainly a scanned image, a password-protected file, or effectively blank.
MIN_EXTRACTED_CHARS = 500

# Expansion bound checked before python-docx parses (#793). zipfile stops
# inflating each entry at its declared file_size, so capping the declared
# totals caps what a parse can expand to. The 392 local corpus CVs top out at
# 37 entries and 7.2 MB uncompressed; a zip bomb declares gigabytes.
_DOCX_MAX_ENTRIES = 1000
_DOCX_MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024


def _validate_pdf_magic(content: bytes) -> bool:
    """Check if content starts with the PDF header."""
    return content[:len(PDF_MAGIC)] == PDF_MAGIC


def _validate_docx_magic(content: bytes) -> bool:
    """Check if content is a ZIP archive containing Word document structure,
    within the entry-count and uncompressed-size bounds above."""
    if content[:4] != ZIP_MAGIC:
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            entries = zf.infolist()
            if (len(entries) > _DOCX_MAX_ENTRIES
                    or sum(e.file_size for e in entries) > _DOCX_MAX_UNCOMPRESSED_BYTES):
                logger.warning("Rejected docx: %d entries, %d bytes uncompressed",
                               len(entries), sum(e.file_size for e in entries))
                return False
            return "word/document.xml" in zf.namelist()
    except (zipfile.BadZipFile, Exception):
        return False


# What python-docx raises on a zip-shaped upload it cannot read (measured on
# 1.2.0): a zip missing its parts -> KeyError; malformed part XML ->
# XMLSyntaxError; a truncated zip -> BadZipFile; not an OPC package at all ->
# PackageNotFoundError; the tempfile round-trip -> OSError. Anything else is a
# bug and must surface, not be swallowed (§5.4).
_DOCX_READ_ERRORS = (PackageNotFoundError, zipfile.BadZipFile, KeyError, XMLSyntaxError, OSError,
                     zlib.error)


def _extract_text(content: bytes, file_ext: str) -> str | None:
    """Best-effort text extraction for the empty-document guard.

    Returns the extracted text, an empty string when the file is readable but
    contains no text (scan/blank), or ``None`` when the document could not be
    read at all (a known python-docx read failure, see ``_DOCX_READ_ERRORS``).
    Callers treat ``None`` as "cannot determine" and skip the guard rather
    than block a possibly-valid upload.

    A PDF never fails open (#806): the run's conversion uses the same parser
    under the same limits, so a PDF this cannot read, the run cannot either.

    Raises:
        EncryptedPdfError: a password-protected PDF.
        PdfTooComplexError: a PDF over pdf_sandbox's page, memory or time limit.
        UnreadablePdfError: any other PDF parse failure.
        PdfBusyError: every PDF child slot is taken.
    """
    if file_ext == PDF_EXTENSION:
        return extract_pdf_text(content)
    if file_ext != ".docx":
        return None
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name
    try:
        doc = Document(tmp_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    except _DOCX_READ_ERRORS as e:
        logger.warning("Text extraction for empty-doc guard failed (%s): %s", file_ext, e, exc_info=True)
        return None
    finally:
        os.unlink(tmp_path)
