"""The finished document with the run doctor's findings flagged for review (#1388 C).

A reviewer sees each WARN or ERROR finding as a short Word comment on the text
it is about: the year that looks wrong, the second copy of a duplicate. A
finding with no one place in the document (a citation PubMed could not match,
an entry removed as a near-duplicate, a step that fell back) is an item in the
review-notes box closing the document, grouped by what to do about it. Where
the Appendix entries came from goes in stage 6's own Appendix note box. The
clean document is left as it is; this writes a copy beside it.
"""
import copy
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml.etree import XMLSyntaxError

from app.schemas import DoctorFindingInstance
from app.services.artifact_service import REVIEW_DOCX_SUFFIX
from app.services.run_quality_report import (
    LINT_COPY,
    TRUNCATION_MARK,
    _instance,
    _usable_findings,
)
from unified_pipeline.core.text_norm import squash  # noqa: E402
from unified_pipeline.doctor.lints.extraction import DEDUP_TEXT_CHARS  # noqa: E402
from unified_pipeline.stage4.schemas import TAXONOMY_LABELS  # noqa: E402
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402

#: What reading or writing a docx raises: quality_score's ``_load_docx`` set,
#: plus PackageNotFoundError, which python-docx raises for a file that is not a zip.
REVIEW_DOCX_ERRORS = (OSError, zipfile.BadZipFile, KeyError, ValueError, XMLSyntaxError,
                      PackageNotFoundError)
COMMENT_AUTHOR = "CViche check"
COMMENT_INITIALS = "CV"
COMMENTED_SEVERITIES = ("ERROR", "WARN")
#: Flags one lint may put in one document. A lint can fire hundreds of times.
MAX_FLAGS_PER_LINT = 25

#: What a reviewer reads, on the text it is about. Short, and an instruction
#: where there is one. Every lint the run page words (LINT_COPY) has one.
REVIEW_FLAGS = {
    "segmentation": "Text from the original CV may be missing or merged here.",
    "missed_headers": "This was a section heading in the original CV.",
    "bucket_status": "Grant is under the wrong funding heading.",
    "under_extraction": "Much of this entry was not read. Check it against the original CV.",
    "classified_unrendered": "Entries from the original CV are missing here.",
    "taxonomy_code_coverage": "Sent to the Appendix: move it to the right section.",
    "stage3b_fallback_ratio": "Some entries may be in the wrong section.",
    "output_hygiene": "Stray text: delete it.",
    "dead_sections": "This section is empty, but the original CV has entries for it.",
    "unrendered_records": "Part of this entry is missing.",
    "section_lost": "Entries from the original CV are missing from this section.",
    "enrichment_failures": "Not found in PubMed: check this citation.",
    "stage6_render_warnings": "Check this section.",
    "dedup_drops": "An entry was removed as a near-duplicate.",
    "pipe_leaks": "Citations run together: split them.",
    "table_shape": "One award split across rows: merge them.",
    "duplicate_passages": "Repeated: delete this copy.",
    "duplicate_records": "Duplicate: delete this copy.",
    "protected_data_in_output": "Protected personal data: delete it.",
    "invented_records": "Not in the original CV: check it or delete it.",
    "wrong_start_date": "Shows Present, but the original has an end date.",
    "table_lost": "A table from the original CV is mostly missing.",
    "date_only_lines": "A date on its own line: join it to its entry.",
    "stage3b_second_pass_error": "Some entries may be in the wrong section.",
    "offschema_fields": "Entries from the original CV are missing from this section.",
    "implausible_year": "Year looks wrong.",
    "stage4_group_failures": "Some entries may be missing details.",
    "python_repr_in_output": "Raw data code: replace it with plain text.",
    "llm_refusal_in_output": "Not CV text: delete it.",
    "llm_fallback_served": "A backup AI model wrote part of this document: check it.",
    "stage_failure_recorded": "A processing step failed: content may be missing.",
    "owner_missing_from_citation": "The CV owner's name is missing from these authors.",
    "etal_added": 'Authors cut to "et al.": restore the full list.',
    "multi_record_coverage": "Several entries merged into one: split them.",
    "year_not_in_source": "This year is not in the original CV.",
    "date_cell_shape": "Date looks wrong.",
    "junk_or_header_row": "A heading printed as an entry: delete it.",
    "teaching_postcheck": "Reworded wrongly: check it against the original CV.",
    "contact_slot_lost": "Office contact details are missing.",
    "pubmed_title_truncated": "Title cut short: complete it.",
    "enrichment_pubtype_mismatch": "Replaced by a correction notice: restore the original citation.",
    "section_consistency": "Probably in the wrong section.",
    "segmentation_collapse": "Sections of the original CV were not found: content may be misplaced.",
    "owner_contact_missing": "The CV owner's name was not found.",
    "pipeline_errors_present": "Processing error: content may be missing.",
    "no_output": "No document was produced.",
    "grant_boundary": "These details belong to the grant before or after.",
    "grant_bucket": "Grant is under the wrong funding heading.",
    "research_summary_call_failed": "The research summary is missing.",
    "span_count": "Separate years shown as one range.",
    "role_consistency": "Wrong principal investigator or role.",
    "fanout_cell_residue": "Leftover text: delete it.",
    "identical_rendered_rows": "Same as another row: delete one.",
    "split_child_unsourced": "Place or dates come from another entry.",
    "record_boundary": "This line belongs to the entry before.",
    "citation_field_dropped": 'Title, link or "..." missing from this citation.',
    "group_header_context": "Lost the heading it sat under.",
    "orphaned_fragments": "Text from the original CV may be missing from this entry.",
    "stage4_unplaced_items": "Records from the original CV are missing from this section.",
}
#: Review-notes group titles where the run page's title is internal wording.
NOTE_TITLES = {"output_hygiene": "Stray text to delete"}
PROTECTED_DATA_LINT = "protected_data_in_output"
#: protected_data_in_output names a category and a section, never the value:
#: "protected personal data (children / dependents) found in Appendix -- value withheld ...".
#: Also "a bare date found in the Personal Data block".
_PROTECTED_WHERE_RE = re.compile(r"(?:\((?P<category>[^)]+)\)|(?P<bare>a bare date)) found in "
                                 r"(?:the )?(?P<section>.+?)(?: block)?(?:\s+--|$)")
#: The lint stage 6's Appendix diversions arrive under (lint_stage6_warnings).
DIVERSION_LINT = "stage6_render_warnings"
#: Lints whose every quoted item is its own problem, each flagged where it is.
EACH_ITEM_LINTS = frozenset({
    "duplicate_records", "duplicate_passages", "pipe_leaks", "enrichment_failures",
    "stage6_render_warnings", "dedup_drops",
})
#: Lints that quote a record printed twice: the later copy is the one to delete.
REPEAT_LINTS = frozenset({"duplicate_records", "duplicate_passages"})
#: Lints whose problem is one year their detail names ("end_date=1912").
YEAR_LINTS = frozenset({"implausible_year", "year_not_in_source"})
_YEAR_RE = re.compile(r"\b(1[0-9]{3}|20[0-9]{2})\b")
#: dedup_drops evidence, "[entry N: ]K1 (jaccard=1.00, 89% covered by kept): dropped '<x>' vs
#: kept '<y>'" (doctor _drop_evidence); the run page may mark it cut with a trailing ellipsis.
_DEDUP_RE = re.compile(r"^(?:entry [^:]+: )?(?P<code>[A-Z][A-Z0-9]*) \(.*?\): "
                       r"dropped '(?P<dropped>.*)' vs kept '(?P<kept>.*)'\u2026?$")
#: stage 6's appendix_diversion messages, after the code prefix the run page
#: strips: "2 entries diverted to the Appendix ...", "1 entry ... recovered into the Appendix".
_DIVERTED_RE = re.compile(r"(?P<count>\d+) entr(?:y|ies)\b[^.]*?(?:diverted to|recovered into) the Appendix")
#: CViche's voice to the submitter: a one-cell light-gray table, in Arial,
#: titled "CViche ..." (the stage-6 Appendix note and the review notes).
REVIEW_NOTES_TITLE = "CViche review notes: delete this box before sending"
NOTES_FONT = "Arial"
BOX_FILL = "E7E6E6"  # these mirror stage6/formatting/docx.py CVICHE_BOX_* (#1552)
BOX_BORDER = "808080"
BOX_BORDER_SIZE = "6"  # eighths of a point: 0.75pt
BOX_FULL_WIDTH = "5000"  # fiftieths of a percent: the whole text column
BOX_PADDING = {"top": 120, "bottom": 120, "left": 160, "right": 160}  # twips: 6pt and 8pt
BOX_TITLE_PT = 9
BOX_TITLE_COLOR = RGBColor(0x59, 0x59, 0x59)
BOX_TEXT_PT = 10
BOX_OUTSIDE_PT = 6
BOX_LINE_AFTER_PT = 3
BOX_TITLE_AFTER_PT = 6
BOX_GROUP_BEFORE_PT = 6
BOX_GROUP_AFTER_PT = 2
NOTE_BULLET = "\u2022\t"
ITEM_INDENT = Inches(0.25)
BULLET_HANG = Inches(0.15)
#: "Removed:" / "Kept:" start at the first indent, the quote hangs at the
#: second: 0.75" apart, since "Removed:" in 10pt italic Arial is about 0.63".
PAIR_LABEL_INDENT = Inches(0.5)
PAIR_TEXT_INDENT = Inches(1.25)
DEDUP_TITLE = "Removed as near-duplicates"
DEDUP_INSTRUCTION = ("We kept one copy of each. If any of these is a separate entry, "
                     "add it back in the section shown.")

#: Squashed characters of a quote matched against the output. Long enough that
#: a hit is that passage, short enough to survive the doctor's own cuts.
QUOTE_PROBE_CHARS = 40
#: Shorter quotes ("PI", a year) would anchor on the first paragraph holding them.
MIN_QUOTE_CHARS = 12
#: Squashed length past which a paragraph is a record, not a heading.
MAX_HEADING_CHARS = 100
#: The doctor's locator before output text in a note: "block 337 repeats at 338: ".
_NOTE_LOCATOR_RE = re.compile(r"^(?:entry|row|blocks?|element_idx_start|stage) \d[^:]{0,40}:\s*")
#: pipe_leaks evidence names its section in brackets: "[bibliography] 93. ...".
_BRACKET_LOCATOR_RE = re.compile(r"^\[[^\]]{1,40}\]\s*")
_CODE_BY_LABEL = {label: code for code, label in TAXONOMY_LABELS.items()}
#: Top-level template headings for the codes stage 6's heading map leaves blank
#: (it routes no overflow there). Substrings, matched as stage 6 matches its own.
_TOP_LEVEL_HEADINGS = {
    "A": "PERSONAL DATA", "B": "EDUCATION", "C": "POSTDOCTORAL",
    "D": "PROFESSIONAL POSITIONS", "E": "EMPLOYMENT STATUS", "F": "LICENSURE",
    "G": "INSTITUTIONAL/HOSPITAL", "I": "PROFESSIONAL ORGANIZATIONS",
    "J": "PERCENT EFFORT", "T": "APPENDIX",
}
_APPENDIX_CODE = "T"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_QUOTES = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"'})


@dataclass(frozen=True)
class Flag:
    """One comment: the paragraph it sits on, the span of its text it covers
    (all of it when None) and what it says."""
    paragraph: Paragraph
    span: tuple[int, int] | None
    text: str


@dataclass(frozen=True)
class Note:
    """One item in the review-notes box, under the group its title and
    instruction name. ``removed`` and ``kept`` are a near-duplicate's two copies."""
    title: str
    instruction: str
    item: str | None = None
    removed: str | None = None
    kept: str | None = None


def review_docx_path(clean_docx: Path) -> Path:
    """Where the flagged copy of ``<uid>_wcm.docx`` is written."""
    return clean_docx.with_name(clean_docx.name.removesuffix("_wcm.docx") + REVIEW_DOCX_SUFFIX)


def _paragraph_text(p: Paragraph) -> str:
    # ponytail: w:t only, so a tracked insertion (enrichment) counts and a
    # tracked deletion (w:delText) does not; Paragraph.text skips w:ins runs.
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def _norm(text: str) -> str:
    """squash, with curly quotes straightened: a source "pharmacist’s" is a
    rendered "pharmacist's" once PubMed or stage 5d has rewritten the citation."""
    return squash(text).translate(_QUOTES)


def _probes(text: str) -> list[str]:
    """Squashed windows of a quote: its start, middle and end, since stage 6
    reformats the source text a quote was taken from (labels, date forms)."""
    text = _BRACKET_LOCATOR_RE.sub("", _NOTE_LOCATOR_RE.sub("", text))
    s = _norm(text.removesuffix(TRUNCATION_MARK))
    if len(s) < MIN_QUOTE_CHARS:
        return []
    starts = {0, max(0, len(s) // 2 - QUOTE_PROBE_CHARS // 2), max(0, len(s) - QUOTE_PROBE_CHARS)}
    return [s[i:i + QUOTE_PROBE_CHARS] for i in sorted(starts)]


def _find(text: str, surfaces: tuple[list[tuple[Paragraph, str]], ...],
          later_copy: bool = False) -> Paragraph | None:
    """The paragraph holding ``text``: searched paragraph by paragraph, then
    table row by table row (a record printed across cells). A quote's start
    is its own (a repeated record matches twice; ``later_copy`` takes the
    second); a window from its middle or end recurs across records, so it
    only counts when unique."""
    for surface in surfaces:
        for n, probe in enumerate(_probes(text)):
            hits = [p for p, body in surface if probe in body]
            if hits and (n == 0 or len(hits) == 1):
                return hits[1] if later_copy and len(hits) > 1 else hits[0]
    return None


def _rows(doc: Document) -> list[tuple[Paragraph, str]]:
    """Each table row as one squashed text, anchored on the row's first paragraph."""
    rows = []
    for tr in doc.element.body.iter(qn("w:tr")):
        first = next(tr.iter(qn("w:p")), None)
        if first is not None:
            rows.append((Paragraph(first, doc._body), _norm("".join(t.text or "" for t in tr.iter(qn("w:t"))))))
    return rows


def _heading(code: str, paragraphs: list[tuple[Paragraph, str]]) -> Paragraph | None:
    """The template heading stage 6 files ``code``'s records under, found by
    stage 6's own rule (_find_paragraph_with_text), kept to heading-length paragraphs."""
    text = squash(WCMTemplateGenerator._get_wcm_section_header(code)
                  or _TOP_LEVEL_HEADINGS.get(code[:1], ""))
    if not text:
        return None
    return next((p for p, body in paragraphs if text in body and len(body) <= MAX_HEADING_CHARS), None)


def _year_span(lint: str, detail: str, para: Paragraph) -> tuple[int, int] | None:
    """The span of the year a year lint's detail names, when the paragraph prints it."""
    year = _YEAR_RE.search(detail) if lint in YEAR_LINTS else None
    at = _paragraph_text(para).find(year.group(0)) if year else -1
    return (at, at + len(year.group(0))) if year and at >= 0 else None


def _item_flags(lint: str, inst: DoctorFindingInstance, label: str,
                surfaces: tuple[list[tuple[Paragraph, str]], ...]) -> list[Flag]:
    """Flags on the quoted text itself: every item for EACH_ITEM_LINTS, else the first found."""
    flags: list[Flag] = []
    for item in [*inst.quotes, *inst.notes]:
        para = _find(item, surfaces, later_copy=lint in REPEAT_LINTS)
        if para is not None and all(f.paragraph is not para for f in flags):
            flags.append(Flag(para, _year_span(lint, inst.detail, para), label))
            if lint not in EACH_ITEM_LINTS:
                break
    return flags


def _quoted(text: str) -> str:
    """A dedup text in quotes, marked cut when it is as long as the doctor quotes."""
    cut = TRUNCATION_MARK if len(text) >= DEDUP_TEXT_CHARS else ""
    return f'"{text.strip()}{cut}"'


def _dedup_notes(inst: DoctorFindingInstance) -> list[Note]:
    """One note per entry dedup removed: its section, what went and what stayed."""
    notes = []
    for item in [*inst.quotes, *inst.notes]:  # "entry N: ..." reads as a note
        m = _DEDUP_RE.search(item)
        if m is not None:
            notes.append(Note(DEDUP_TITLE, DEDUP_INSTRUCTION, TAXONOMY_LABELS.get(m["code"]),
                              _quoted(m["dropped"]), _quoted(m["kept"])))
    return notes or [Note(DEDUP_TITLE, DEDUP_INSTRUCTION)]


def _flags(finding: dict, surfaces: tuple[list[tuple[Paragraph, str]], ...]) -> tuple[list[Flag], list[Note]]:
    """Where one finding goes: a comment on its quoted text; else on its
    section's heading; else a review note saying what and where it was. A
    near-duplicate is always a note, since what it removed is not on the page."""
    lint, inst = finding["lint"], _instance(finding)
    if lint == "dedup_drops":
        return [], _dedup_notes(inst)
    label = REVIEW_FLAGS.get(lint, inst.detail)
    flags = _item_flags(lint, inst, label, surfaces)
    if flags:
        return flags, []
    heading = _heading(_CODE_BY_LABEL.get(inst.section or "", ""), surfaces[0])
    if heading is not None:
        return [Flag(heading, None, label)], []
    item = _note_item(lint, inst)
    if item is None:  # nothing to point the submitter at: it stays on the run page
        return [], []
    copy_ = LINT_COPY.get(lint)
    title = NOTE_TITLES.get(lint) or (copy_.title if copy_ else lint)
    return [], [Note(title, label, item)]


def _note_item(lint: str, inst: DoctorFindingInstance) -> str | None:
    """What a review note points at: its section and quoted text; for
    protected data, its section and category (the finding never quotes the value)."""
    if lint == PROTECTED_DATA_LINT:
        m = _PROTECTED_WHERE_RE.search(inst.detail)
        return f"{m['section']}: {m['category'] or m['bare']}" if m else None
    quote = f'"{inst.quotes[0]}"' if inst.quotes else None
    return ": ".join(x for x in (inst.section, quote) if x) or None


def _box(doc: Document, title: str) -> _Cell:
    """A one-cell, light-gray, bordered table the full width of the text
    column, closing the document: CViche's voice, deletable in one step.
    Returns the cell, holding the title (9pt bold gray small caps). Mirrors
    stage 6's add_cviche_box (#1552), which replaces this once merged."""
    table = doc.add_table(rows=1, cols=1)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = tbl_pr.makeelement(qn("w:tblW"), {})
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), BOX_FULL_WIDTH)
    margins = tbl_pr.makeelement(qn("w:tblCellMar"), {})
    for side, twips in BOX_PADDING.items():
        margins.append(margins.makeelement(qn(f"w:{side}"), {qn("w:w"): str(twips), qn("w:type"): "dxa"}))
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(margins)  # schema order: tblCellMar before tblLook
    else:
        tbl_pr.append(margins)
    cell = table.cell(0, 0)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is not None:
        tc_w.set(qn("w:type"), "pct")
        tc_w.set(qn("w:w"), BOX_FULL_WIDTH)
    tc_pr.append(tc_pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto",
                                                  qn("w:fill"): BOX_FILL}))
    borders = tc_pr.makeelement(qn("w:tcBorders"), {})
    for side in ("top", "left", "bottom", "right"):
        borders.append(borders.makeelement(qn(f"w:{side}"), {
            qn("w:val"): "single", qn("w:sz"): BOX_BORDER_SIZE, qn("w:color"): BOX_BORDER}))
    tc_pr.append(borders)
    title_para = cell.paragraphs[0]
    _spacing(title_para, after=BOX_TITLE_AFTER_PT)
    run = _run(title_para, title, size=BOX_TITLE_PT, bold=True)
    run.font.small_caps, run.font.color.rgb = True, BOX_TITLE_COLOR
    return cell


def _spacing(para: Paragraph, before: float = 0, after: float = BOX_LINE_AFTER_PT) -> None:
    fmt = para.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), 1.0


def _run(para: Paragraph, text: str, size: float = BOX_TEXT_PT, bold: bool = False,
         italic: bool = False) -> Run:
    run = para.add_run(text)
    run.font.name, run.font.size = NOTES_FONT, Pt(size)
    run.font.bold, run.font.italic = bold or None, italic or None
    return run


def _box_group(cell: _Cell, text: str, first: bool) -> None:
    para = cell.add_paragraph()
    _spacing(para, before=0 if first else BOX_GROUP_BEFORE_PT, after=BOX_GROUP_AFTER_PT)
    para.paragraph_format.keep_with_next = True
    _run(para, text, bold=True)


def _box_text(cell: _Cell, text: str) -> None:
    para = cell.add_paragraph()
    _spacing(para)
    _run(para, text)


def _box_item(cell: _Cell, text: str) -> None:
    """A bulleted item, its wrapped lines hanging under its text."""
    para = cell.add_paragraph()
    _spacing(para)
    para.paragraph_format.left_indent = ITEM_INDENT + BULLET_HANG
    para.paragraph_format.first_line_indent = -BULLET_HANG
    _run(para, f"{NOTE_BULLET}{text}")


def _box_pair(cell: _Cell, label: str, text: str, last: bool) -> None:
    """A "Removed:" / "Kept:" line: the label italic, the quote hanging under
    its opening mark; only the pair's last line has space after it."""
    para = cell.add_paragraph()
    _spacing(para, after=BOX_LINE_AFTER_PT if last else 0)
    fmt = para.paragraph_format
    fmt.left_indent, fmt.first_line_indent = PAIR_TEXT_INDENT, PAIR_LABEL_INDENT - PAIR_TEXT_INDENT
    fmt.tab_stops.add_tab_stop(PAIR_TEXT_INDENT)
    fmt.keep_with_next = not last or None
    _run(para, f"{label}\t", italic=True)
    _run(para, text)


def _split(r: BaseOxmlElement, t: BaseOxmlElement, lo: int, hi: int) -> BaseOxmlElement:
    """Split run ``r`` (one w:t) so ``t.text[lo:hi]`` is a run of its own; return that run."""
    text = t.text or ""
    middle = r
    for piece, is_middle in reversed(((text[:lo], False), (text[lo:hi], True), (text[hi:], False))):
        if not piece:
            continue
        clone = copy.deepcopy(r)
        clone_t = next(clone.iter(qn("w:t")))
        clone_t.text = piece
        clone_t.set(_XML_SPACE, "preserve")
        r.addnext(clone)
        if is_middle:
            middle = clone
    r.getparent().remove(r)
    return middle


def _runs(para: Paragraph, span: tuple[int, int] | None) -> list[Run]:
    """The first and last run covering ``span`` of the paragraph's text (all
    of it when None), splitting runs at its ends. Runs inside a tracked
    insertion count; a tracked deletion's do not."""
    runs = [r for r in para._p.iter(qn("w:r")) if r.getparent().tag != qn("w:del")]
    if not runs:
        return [para.add_run()]
    covered, at = [], 0
    for r in runs if span else []:
        ts = list(r.iter(qn("w:t")))
        length = sum(len(t.text or "") for t in ts)
        lo, hi = max(span[0] - at, 0), min(span[1] - at, length)
        if lo < hi:
            covered.append(_split(r, ts[0], lo, hi) if len(ts) == 1 else r)
        at += length
    chosen = covered or runs
    return [Run(chosen[0], para), Run(chosen[-1], para)]


def _add_review_notes(doc: Document, notes: list[Note]) -> None:
    """The review-notes box closing the document: a group per action, with
    its count and instruction, then an item per finding. The space above the
    box is set on the document's last paragraph, not a blank one."""
    if not notes:
        return
    groups: dict[tuple[str, str], list[Note]] = {}
    for note in notes:
        groups.setdefault((note.title, note.instruction), []).append(note)
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(BOX_OUTSIDE_PT)
    cell = _box(doc, REVIEW_NOTES_TITLE)
    for n, ((title, instruction), members) in enumerate(groups.items()):
        _box_group(cell, f"{title} ({len(members)})", first=n == 0)
        _box_text(cell, instruction)
        for note in dict.fromkeys(members):  # the same item once
            if note.item:
                _box_item(cell, note.item)
            if note.removed:
                _box_pair(cell, "Removed:", note.removed, last=not note.kept)
            if note.kept:
                _box_pair(cell, "Kept:", note.kept, last=True)


def write_review_docx(clean_docx: Path, doctor_payload: object) -> tuple[Path, int] | None:
    """Write the flagged copy of ``clean_docx``; return its path and how many
    flags (comments and review notes) it carries. None when there is nothing to flag."""
    if not isinstance(doctor_payload, dict):
        return None
    # Stage 6's Appendix diversions say which section CViche first tried, not
    # the heading the entry had in the CV, which the Appendix groups already
    # show: a second, different "came from" only contradicts them. Their
    # counts stay on the run page.
    findings = [f for f in _usable_findings(doctor_payload)[0]
                if f["severity"] in COMMENTED_SEVERITIES
                and not (f["lint"] == DIVERSION_LINT and _DIVERTED_RE.search(_instance(f).detail))]
    if not findings:
        return None
    doc = Document(str(clean_docx))
    paragraphs = [(p, _norm(_paragraph_text(p)))
                  for p in (Paragraph(el, doc._body) for el in doc.element.body.iter(qn("w:p")))]
    if not paragraphs:
        return None
    surfaces = (paragraphs, _rows(doc))
    flags: list[Flag] = []
    notes: list[Note] = []
    per_lint: dict[str, int] = {}
    for f in findings:
        found, noted = _flags(f, surfaces)
        for flag in found:
            per_lint[f["lint"]] = per_lint.get(f["lint"], 0) + 1
            if per_lint[f["lint"]] <= MAX_FLAGS_PER_LINT:
                flags.append(flag)
        notes.extend(noted)
    for flag in flags:
        doc.add_comment(_runs(flag.paragraph, flag.span), text=flag.text,
                        author=COMMENT_AUTHOR, initials=COMMENT_INITIALS)
    if not flags and not notes:
        return None
    _add_review_notes(doc, notes)
    out = review_docx_path(clean_docx)
    doc.save(str(out))
    return out, len(flags) + len(notes)
