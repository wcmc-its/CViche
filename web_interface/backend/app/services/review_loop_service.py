"""The reviewer's corrected copy of the delivered document (#1587).

The corrected copy is a CV. It is stored with the run's other artifacts under
`CORRECTED_PREFIX`, which the artifact downloads and the scorer never read:
those read `outputs/` only (artifact_service.resolve_artifact,
quality_score_service). The copy is served back by its own route only
(feedback_routes, GET /run/{id}/feedback/corrected-copies/{copy_id}): to the
run's owner, admins and staff, and behind the malware-scan gate the original
CV's download uses (services/malware_scan.py), since a reviewer uploaded it.
It is deleted with the run (storage.delete_run, and its corrected_copies row
with the run's rows). What is kept beside it is `doctor/docx_diff.py`'s typed
report: block positions, character counts and template section headings, no
CV text.

Once the copy is stored, `review_corrected_copy` adds two more (#1654):

- the verdicts: what happened at each comment of the run's review copy, read
  from the corrected copy (`doctor/comment_fate.py`), one per finding with its
  lint, shape, severity and entry index (the review copy's comment map,
  review_comments.CommentFinding). No CV text.
- the doctor re-run on the corrected copy, over the run's own stored
  artifacts, and its findings counted against the delivered run's report.

Each uploader keeps one copy per run (#1754): the keys carry the uploader's
user id, so an owner's and a staff member's copies never overwrite each other,
and a second upload by the same user replaces that user's copy. A copy stored
before that has the run's id alone in its keys (uploader None);
`stored_uploaders` lists both kinds. Each stored copy also has a
corrected_copies row (corrected_copy_service), which the runs list reads.
"""
import io
import json
import logging
import re
import tempfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from sqlalchemy.orm import Session

from app.errors import bad_request, conflict
from app.services import (
    corrected_copy_service,
    quality_score_service,  # noqa: F401  (puts src/ on sys.path for unified_pipeline)
)
from app.services.artifact_service import REVIEW_DOCX_SUFFIX, resolve_artifact
from app.services.quality_score_service import DOCTOR_SUFFIX
from app.services.review_comments import (
    COMMENT_AUTHOR,
    LEGACY_REVIEW_NOTES_TITLE,
    CommentFinding,
    comment_map_path,
    lint_of_flag,
)
from app.services.upload_validation import _DOCX_READ_ERRORS
from app.storage import get_storage
from unified_pipeline.doctor.comment_fate import comment_fates  # noqa: E402
from unified_pipeline.doctor.docx_diff import (  # noqa: E402  (path set by quality_score_service)
    diff_docx,
    to_report,
)
from unified_pipeline.doctor.precision import (  # noqa: E402
    REVIEW_VERDICTS,
    finding_shape,
)
from unified_pipeline.doctor.shared import STATUS_RAN  # noqa: E402
from unified_pipeline.stage6.formatting import is_cviche_box  # noqa: E402
from unified_pipeline.stage_errors import (  # noqa: E402
    STAGE_ERRORS_SUFFIX,
    stage_errors_path,
)

logger = logging.getLogger(__name__)

#: Where a reviewer's corrected copy and its diff live in the run's storage.
#: Never `outputs/`: the artifact downloads and the scorer read only that prefix.
CORRECTED_PREFIX = "corrected/"
#: The delivered document (orchestrator stage 6) and the stage-4 artifact
#: whose entries the diff maps changes to.
DELIVERED_DOCX_SUFFIX = "_wcm.docx"
STAGE4_FIELDS_SUFFIX = "_fields.json"


_DIFF_SUFFIX = "_diff.json"


def _corrected_key(run_id: str, uploader_id: int | None, suffix: str) -> str:
    """``corrected/<run>_u<user id><suffix>``; ``corrected/<run><suffix>`` for
    a copy stored before copies were kept per uploader (uploader None)."""
    owner = "" if uploader_id is None else f"_u{uploader_id}"
    return f"{CORRECTED_PREFIX}{run_id}{owner}{suffix}"


def corrected_docx_key(run_id: str, uploader_id: int | None) -> str:
    return _corrected_key(run_id, uploader_id, "_corrected.docx")


def corrected_diff_key(run_id: str, uploader_id: int | None) -> str:
    return _corrected_key(run_id, uploader_id, _DIFF_SUFFIX)


def corrected_verdicts_key(run_id: str, uploader_id: int | None) -> str:
    return _corrected_key(run_id, uploader_id, "_verdicts.json")


def corrected_doctor_key(run_id: str, uploader_id: int | None) -> str:
    return _corrected_key(run_id, uploader_id, "_doctor.json")


def corrected_doctor_compare_key(run_id: str, uploader_id: int | None) -> str:
    return _corrected_key(run_id, uploader_id, "_doctor_compare.json")


def stored_uploaders(run_id: str) -> list[int | None]:
    """Whose corrected copies of the run are stored, read from their diffs'
    keys: a user id per uploader, None for a copy stored under the run's id
    alone (before #1754). The pre-#1754 copy first, then by user id."""
    diff_key = re.compile(rf"{re.escape(CORRECTED_PREFIX + run_id)}(?:_u(\d+))?{re.escape(_DIFF_SUFFIX)}")
    found = [m for key in get_storage().list_files(run_id, CORRECTED_PREFIX) if (m := diff_key.fullmatch(key))]
    return sorted((None if m[1] is None else int(m[1]) for m in found), key=lambda u: -1 if u is None else u)


#: Where a verdict report's finding for each comment came from: the review
#: copy's comment map, or (a copy written before the map) the lint its
#: wording names, with no shape, severity or entry index.
FINDINGS_FROM_MAP = "comment_map"
#: run_doctor's artifact key for the stage-6 document, which the re-run reads
#: from the corrected copy instead.
DOCTOR_DOCX_ARTIFACT = "stage_6_docx"
FINDINGS_FROM_WORDING = "comment_wording"
#: Said in the verdict report when the comments' fate cannot be read.
NO_REVIEW_COPY_NOTE = "The run has no review copy, so there are no comments to follow."
NO_COMMENTS_NOTE = ("The uploaded copy holds no CViche comments: it is the clean document, or a "
                    "review copy with every comment deleted. A dismissed comment cannot be told "
                    "from one never there, so unchanged text is unknown, never not_a_problem.")


def _artifact_bytes(db: Session, run_id: str, name: str) -> bytes | None:
    """One of the run's artifacts, from the pod's disk or durable storage;
    None when neither has it."""
    resolved = resolve_artifact(db, run_id, name)
    if resolved.local_path is not None:
        return resolved.local_path.read_bytes()
    try:
        return get_storage().get_file(run_id, resolved.storage_key)
    except FileNotFoundError:
        return None


def _stage4_entries(db: Session, run_id: str) -> dict | None:
    """The run's stage-4 JSON, for entry indices; None when absent or
    unreadable (the diff then maps no entries, and says so per change)."""
    raw = _artifact_bytes(db, run_id, f"{run_id}{STAGE4_FIELDS_SUFFIX}")
    if raw is None:
        return None
    try:
        data = json.loads(raw)
    except (UnicodeDecodeError, json.JSONDecodeError):
        logger.warning("Stage-4 JSON of run %s is unreadable; diffing without entry indices", run_id)
        return None
    return data if isinstance(data, dict) else None


def record_corrected_docx(db: Session, run_id: str, uploader_id: int, corrected: bytes,
                          uploaded_at: datetime) -> int:
    """Diff a reviewer's corrected copy against the delivered document, store
    both the copy and the typed diff under CORRECTED_PREFIX as that uploader's
    (replacing any copy the same user uploaded before), record its
    corrected_copies row, and return how many changes the diff found. The
    caller has already checked access, size and file type. Nothing is stored
    when the diff cannot be made.

    Raises a 409 when the run has no delivered document, and a 400 when the
    corrected file cannot be read as a Word document.
    """
    delivered = _artifact_bytes(db, run_id, f"{run_id}{DELIVERED_DOCX_SUFFIX}")
    if delivered is None:
        raise conflict("This run has no finished document to compare with.")
    with tempfile.TemporaryDirectory(prefix=f"corrected_{run_id}_") as tmp:
        delivered_path, corrected_path = Path(tmp) / "delivered.docx", Path(tmp) / "corrected.docx"
        delivered_path.write_bytes(delivered)
        corrected_path.write_bytes(corrected)
        try:
            diff = diff_docx(delivered_path, corrected_path, _stage4_entries(db, run_id))
        # ValueError: python-docx's "not a Word file" (a zip whose main part
        # is some other content type), as scripts/docx_review_diff.py reads it.
        except (*_DOCX_READ_ERRORS, ValueError) as e:
            logger.warning("Corrected copy for run %s could not be diffed: %s", run_id, type(e).__name__)
            raise bad_request("We couldn't read this Word file. Re-save it in Word and try again.") from e
    storage = get_storage()
    docx_key = corrected_docx_key(run_id, uploader_id)
    storage.put_file(run_id, docx_key, corrected)
    report = json.dumps(to_report(run_id, diff), indent=1, sort_keys=True)
    storage.put_file(run_id, corrected_diff_key(run_id, uploader_id), report.encode("utf-8"))
    corrected_copy_service.record_copy(db, run_id, uploader_id, changes=len(diff.changes),
                                       storage_key=docx_key, uploaded_at=uploaded_at)
    return len(diff.changes)


#: What a corrected copy's download is named after the run's upload.
CORRECTED_DOWNLOAD_LABEL = "corrected"


def corrected_download_name(run_filename: str | None, uploader: str | None) -> str:
    """``<upload's name> - corrected by <uploader>.docx``; without the
    uploader when unknown (a copy stored before #1754)."""
    stem = Path((run_filename or "").strip()).stem or "CV"
    by = f" by {uploader}" if uploader else ""
    return f"{stem} - {CORRECTED_DOWNLOAD_LABEL}{by}.docx"


def changes_summary(changes: int) -> str:
    """The uploader's one-line confirmation."""
    return f"{changes} change{'' if changes == 1 else 's'} recorded"


# --- after the upload (#1654): verdicts and the doctor re-run ------------------

@dataclass(frozen=True)
class FindingVerdict:
    """One review-copy comment's finding and what the reviewer did there."""
    comment_id: int
    lint: str | None
    shape: str | None
    severity: str | None
    entry_index: int | None
    verdict: str


def review_corrected_copy(db: Session, run_id: str, uploader_id: int | None, corrected: bytes) -> None:
    """Store the verdicts and the doctor's re-run for this uploader's
    corrected copy of the run, once `record_corrected_docx` has stored it. Best-effort, each on its
    own: the copy and its diff are already stored, and a failure here is
    logged, never raised to the uploader."""
    for name, store in (("verdicts", store_verdicts), ("doctor re-run", store_corrected_doctor)):
        try:
            store(db, run_id, uploader_id, corrected)
        except Exception:  # noqa: BLE001 -- best-effort after a stored upload; logged with its traceback
            logger.exception("Corrected copy of run %s by user %s: %s failed", run_id, uploader_id, name)


def review_corrected_copy_in_background(run_id: str, uploader_id: int, corrected: bytes) -> None:
    """`review_corrected_copy` with a DB session of its own, for the upload
    route's background task: the request's session is closed by then."""
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        review_corrected_copy(db, run_id, uploader_id, corrected)
    finally:
        db.close()


def _comment_findings(db: Session, run_id: str, review: Path) -> tuple[dict[int, CommentFinding], str]:
    """The finding each review-copy comment marks, by comment id, and where
    that came from (FINDINGS_FROM_MAP or FINDINGS_FROM_WORDING)."""
    raw = _artifact_bytes(db, run_id, comment_map_path(Path(f"{run_id}{REVIEW_DOCX_SUFFIX}")).name)
    if raw is not None:
        rows = [CommentFinding(**row) for row in json.loads(raw)]
        return {row.comment_id: row for row in rows}, FINDINGS_FROM_MAP
    return ({c.comment_id: CommentFinding(c.comment_id, lint_of_flag(c.text), None, None, None)
             for c in Document(str(review)).comments if c.author == COMMENT_AUTHOR},
            FINDINGS_FROM_WORDING)


def comment_verdicts(db: Session, run_id: str, corrected: bytes) -> dict:
    """The verdict report: one FindingVerdict per review-copy comment, their
    counts, and a note when the comments' fate cannot be read."""
    report: dict = {"uid": run_id, "comments_tracked": False, "findings_from": None, "note": None,
                    "counts": dict.fromkeys(REVIEW_VERDICTS, 0), "findings": []}
    review = _artifact_bytes(db, run_id, f"{run_id}{REVIEW_DOCX_SUFFIX}")
    if review is None:
        return {**report, "note": NO_REVIEW_COPY_NOTE}
    with tempfile.TemporaryDirectory(prefix=f"verdicts_{run_id}_") as tmp:
        review_path, corrected_path = Path(tmp) / "review.docx", Path(tmp) / "corrected.docx"
        review_path.write_bytes(review)
        corrected_path.write_bytes(corrected)
        fates = comment_fates(review_path, corrected_path, COMMENT_AUTHOR)
        findings, source = _comment_findings(db, run_id, review_path)
    verdicts = []
    for fate in fates.fates:
        finding = findings.get(fate.comment_id) or CommentFinding(fate.comment_id, None, None, None, None)
        verdicts.append(FindingVerdict(fate.comment_id, finding.lint, finding.shape, finding.severity,
                                       finding.entry_index, fate.verdict))
    counts = Counter(v.verdict for v in verdicts)
    return {**report, "comments_tracked": fates.comments_tracked, "findings_from": source,
            "note": None if fates.comments_tracked else NO_COMMENTS_NOTE,
            "counts": {verdict: counts[verdict] for verdict in REVIEW_VERDICTS},
            "findings": [asdict(v) for v in verdicts]}


def store_verdicts(db: Session, run_id: str, uploader_id: int | None, corrected: bytes) -> dict:
    """Compute the verdict report and store it under CORRECTED_PREFIX, as the uploader's."""
    report = comment_verdicts(db, run_id, corrected)
    get_storage().put_file(run_id, corrected_verdicts_key(run_id, uploader_id),
                           json.dumps(report, indent=1, sort_keys=True).encode("utf-8"))
    return report


def _without_review_notes(docx: bytes) -> bytes:
    """The document without the review-notes box a review copy written
    before the notes became a comment closes with, which the delivered
    document never had; stage 6's own boxes stay, as the doctor read them on
    the delivered run. A later copy's notes are a comment, outside the body
    the doctor reads, so there is nothing to strip."""
    doc = Document(io.BytesIO(docx))
    for tbl in list(doc.element.body.iterchildren(qn("w:tbl"))):
        table = Table(tbl, doc)
        if is_cviche_box(table) and table.cell(0, 0).paragraphs[0].text.startswith(LEGACY_REVIEW_NOTES_TITLE):
            tbl.getparent().remove(tbl)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def doctor_on_corrected(db: Session, run_id: str, corrected: bytes) -> dict:
    """The doctor's report on the corrected copy, read with the run's own
    stored artifacts in place of what stage 6 wrote. An artifact the run no
    longer has leaves its lints skipped, as on any run."""
    return doctor_on_stored_outputs(db, run_id, _without_review_notes(corrected))


def doctor_on_stored_outputs(db: Session, run_id: str, document: bytes | None = None) -> dict:
    """The doctor's report on the run's stored artifacts, with ``document`` in
    place of the stage-6 document when given (a corrected copy). Pure file
    reads: no stage re-runs, no LLM call. The run's prompt logs stay on the
    pod that ran it, so llm_fallback_served sees none of their fallbacks, and
    a PDF upload's converted source is not stored, so the source lints skip."""
    from unified_pipeline.run_doctor import artifact_layout, run_doctor

    with tempfile.TemporaryDirectory(prefix=f"stored_doctor_{run_id}_") as tmp:
        root = Path(tmp)
        layout = artifact_layout(root, run_id)
        for key, path in layout.items():
            raw = (document if document is not None and key == DOCTOR_DOCX_ARTIFACT
                   else _artifact_bytes(db, run_id, path.name))
            if raw is not None:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(raw)
        errors = _artifact_bytes(db, run_id, f"{run_id}{STAGE_ERRORS_SUFFIX}")
        if errors is not None:
            stage_errors_path(root, run_id).parent.mkdir(parents=True, exist_ok=True)
            stage_errors_path(root, run_id).write_bytes(errors)
        source = _source_docx(run_id, root)
        return run_doctor(root, run_id, source=source)


def _source_docx(run_id: str, root: Path) -> Path | None:
    """The run's uploaded CV when it is a docx, written under ``root``; None
    for a PDF upload (its converted docx is not stored) or none stored."""
    try:
        raw = get_storage().get_file(run_id, f"input/{run_id}.docx")
    except FileNotFoundError:
        return None
    path = root / "uploads" / f"{run_id}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(raw)
    return path


def compare_doctor(delivered: dict, corrected: dict) -> dict:
    """The delivered run's findings against the corrected copy's, counted per
    lint and shape: how many cleared, persist (the same message in both) and
    are new. A lint either report did not run (skipped or unreadable), and a
    lint that reads no document (DOCUMENT_INDEPENDENT_LINTS: the re-run has no
    prompt logs, so llm_fallback_served's prompt-log findings would read as
    cleared), is listed apart, never counted as cleared."""
    from unified_pipeline.run_doctor import DOCUMENT_INDEPENDENT_LINTS

    findings = [p.get("findings") if isinstance(p.get("findings"), list) else []
                for p in (delivered, corrected)]
    not_run = sorted({f["lint"] for fs in findings for f in fs
                      if f.get("status", STATUS_RAN) != STATUS_RAN or f["lint"] in DOCUMENT_INDEPENDENT_LINTS})
    by_message: list[Counter[tuple[str, str | None, str]]] = []
    for fs in findings:
        by_message.append(Counter(
            (f["lint"], finding_shape(f["lint"], str(f.get("message") or "")), str(f.get("message") or ""))
            for f in fs if f.get("status", STATUS_RAN) == STATUS_RAN and f["lint"] not in not_run))
    before, after = by_message
    rows: dict[tuple[str, str | None], dict[str, int]] = {}
    for (lint, shape, message) in before.keys() | after.keys():
        row = rows.setdefault((lint, shape), {"delivered": 0, "corrected": 0, "persisting": 0})
        row["delivered"] += before[(lint, shape, message)]
        row["corrected"] += after[(lint, shape, message)]
        row["persisting"] += min(before[(lint, shape, message)], after[(lint, shape, message)])
    lints = [{"lint": lint, "shape": shape, **row, "cleared": row["delivered"] - row["persisting"],
              "new": row["corrected"] - row["persisting"]}
             for (lint, shape), row in sorted(rows.items(), key=lambda kv: (kv[0][0], kv[0][1] or ""))]
    totals = {k: sum(r[k] for r in lints) for k in ("delivered", "corrected", "persisting", "cleared", "new")}
    return {"not_compared": not_run, "totals": totals, "lints": lints}


def store_corrected_doctor(db: Session, run_id: str, uploader_id: int | None, corrected: bytes) -> dict | None:
    """Run the doctor on the corrected copy; store its report, and its
    comparison with the delivered run's report when that is stored. Returns
    the comparison, None without a delivered report."""
    storage = get_storage()
    report = doctor_on_corrected(db, run_id, corrected)
    storage.put_file(run_id, corrected_doctor_key(run_id, uploader_id), json.dumps(report, indent=1).encode("utf-8"))
    delivered = _artifact_bytes(db, run_id, f"{run_id}{DOCTOR_SUFFIX}")
    if delivered is None:
        logger.info("Run %s has no doctor report to compare its corrected copy with", run_id)
        return None
    comparison = {"uid": run_id, **compare_doctor(json.loads(delivered), report)}
    storage.put_file(run_id, corrected_doctor_compare_key(run_id, uploader_id),
                     json.dumps(comparison, indent=1, sort_keys=True).encode("utf-8"))
    return comparison
