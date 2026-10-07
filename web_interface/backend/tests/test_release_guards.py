"""Tests for the pre-release hardening guards.

User-facing failure modes hardened before the non-technical rollout:
  #5  restart_run enforces the per-user quota (so "restart" spam can't run up cost)
  #4  /upload rejects documents with no readable text (scan / encrypted / blank)
  #3  reconcile_stale_runs sweeps runs orphaned by a server restart
  #2  retry_step resumes the pipeline from the failed step (real retry)
  #1  per-pod admission control caps concurrent in-process pipelines (load)
"""
import io
import os

os.environ.setdefault("CVICHE_SESSION_SECRET", "test-secret-not-for-production")

import asyncio
import time
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
import redis
from docx import Document
from sqlalchemy.orm import object_session

from app.models import Log, Run, Step, User
from app.pipeline import concurrency
from app.pipeline.step_registry import STEP_REGISTRY
from app.services.run_service import (
    DEPLOY_INTERRUPT_MESSAGE,
    fail_runs_interrupted_by_shutdown,
    reconcile_stale_runs,
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# --- helpers ----------------------------------------------------------------

def _make_user(db, email="user@example.com", role="user", **kwargs):
    user = User(
        email=email, display_name="Test", role=role,
        auth_method="simple", consent_version="1.0", **kwargs,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def _auth_cookie(client, user):
    from app.auth import COOKIE_NAME, create_session_cookie
    # create_session_cookie reads the current epoch from a DB session;
    # `user` was just committed on the test's session, so borrow that one.
    client.cookies.set(COOKIE_NAME, create_session_cookie(user, object_session(user)))


def _docx_bytes(text: str) -> bytes:
    """A real .docx (valid magic bytes) whose body is ``text``."""
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- #3  reconcile_stale_runs ------------------------------------------------

class TestReconcileStaleRuns:
    def test_old_running_run_is_failed(self, db):
        run = Run(id="OLD001", filename="a.docx", file_type="docx", status="running",
                  started_at=datetime.now() - timedelta(minutes=120))
        db.add(run)
        db.add(Step(run_id="OLD001", step_number=3, step_name="x", status="running"))
        db.commit()

        swept = reconcile_stale_runs(db)

        assert swept == 1
        db.refresh(run)
        assert run.status == "failed"
        assert run.completed_at is not None
        assert "restart" in (run.error_message or "").lower()
        step = db.query(Step).filter(Step.run_id == "OLD001").first()
        assert step.status == "error"

    def test_recent_running_run_is_left_alone(self, db):
        run = Run(id="NEW001", filename="a.docx", file_type="docx", status="running",
                  started_at=datetime.now() - timedelta(minutes=5))
        db.add(run)
        db.commit()

        assert reconcile_stale_runs(db) == 0
        db.refresh(run)
        assert run.status == "running"

    def test_completed_run_is_ignored(self, db):
        run = Run(id="DONE01", filename="a.docx", file_type="docx", status="complete",
                  started_at=datetime.now() - timedelta(minutes=999))
        db.add(run)
        db.commit()

        assert reconcile_stale_runs(db) == 0
        db.refresh(run)
        assert run.status == "complete"

    def test_env_threshold_is_respected(self, db, monkeypatch):
        monkeypatch.setenv("CVICHE_STALE_RUN_MINUTES", "10")
        run = Run(id="ENV001", filename="a.docx", file_type="docx", status="running",
                  started_at=datetime.now() - timedelta(minutes=15))
        db.add(run)
        db.commit()

        assert reconcile_stale_runs(db) == 1
        db.refresh(run)
        assert run.status == "failed"


@pytest.fixture
def notified(monkeypatch):
    """Failure cards the shutdown path queues, as (run id, status, submitter)."""
    cards = []

    def record(run, score=None, submitter=None, doctor_report=None):
        cards.append((run.id, run.status, submitter))

    monkeypatch.setattr("app.services.notifications.notify_run_terminal", record)
    return cards


class TestFailRunsInterruptedByShutdown:
    def test_still_running_run_is_failed_with_the_deploy_message(self, db, notified):
        user = _make_user(db, email="Owner@Example.com")
        run = Run(id="DPLY01", filename="a.docx", file_type="docx", status="running",
                  started_at=datetime.now() - timedelta(minutes=30), user_id=user.id)
        db.add(run)
        db.add(Step(run_id="DPLY01", step_number=4, step_name="x", status="running"))
        db.commit()

        assert fail_runs_interrupted_by_shutdown(db, ["DPLY01"]) == 1

        db.refresh(run)
        assert run.status == "failed"
        assert run.error_message == DEPLOY_INTERRUPT_MESSAGE
        assert run.completed_at is not None
        assert db.query(Step).filter(Step.run_id == "DPLY01").one().status == "error"
        # #116: the run's own log says why it ended, and exactly one failure
        # card goes out -- the run's thread never gets to send one.
        assert [(log.level, log.message) for log in db.query(Log).filter(Log.run_id == "DPLY01")] \
            == [("ERROR", DEPLOY_INTERRUPT_MESSAGE)]
        assert notified == [("DPLY01", "failed", "owner@example.com")]

    def test_a_failed_submitter_lookup_still_sends_every_card(self, db, notified, monkeypatch):
        """The rows are committed before the cards go out; one run's lookup
        failing must not cost the other its card, nor raise to the drain."""
        from sqlalchemy.exc import OperationalError
        for run_id in ("DPLY04", "DPLY05"):
            db.add(Run(id=run_id, filename="a.docx", file_type="docx", status="running",
                       started_at=datetime.now() - timedelta(minutes=30)))
        db.commit()

        def lookup(db_, run):
            if run.id == "DPLY04":
                raise OperationalError("SELECT", {}, Exception("database gone"))
            return "owner@example.com"

        monkeypatch.setattr("app.services.run_service._submitter_email", lookup)

        assert fail_runs_interrupted_by_shutdown(db, ["DPLY04", "DPLY05"]) == 2
        assert sorted(notified) == [("DPLY04", "failed", None),
                                    ("DPLY05", "failed", "owner@example.com")]

    def test_run_that_finished_meanwhile_is_untouched(self, db, notified):
        finished_at = datetime.now() - timedelta(minutes=1)
        run = Run(id="DPLY02", filename="a.docx", file_type="docx", status="complete",
                  started_at=datetime.now() - timedelta(minutes=20), completed_at=finished_at)
        db.add(run)
        db.commit()

        assert fail_runs_interrupted_by_shutdown(db, ["DPLY02"]) == 0

        db.refresh(run)
        assert run.status == "complete"
        assert run.error_message is None
        assert run.completed_at == finished_at
        assert notified == []
        assert db.query(Log).filter(Log.run_id == "DPLY02").count() == 0

    def test_run_not_on_this_pod_is_untouched(self, db, notified):
        run = Run(id="DPLY03", filename="a.docx", file_type="docx", status="running",
                  started_at=datetime.now() - timedelta(minutes=5))
        db.add(run)
        db.commit()

        assert fail_runs_interrupted_by_shutdown(db, ["OTHER9"]) == 0

        db.refresh(run)
        assert run.status == "running"
        assert notified == []


# --- #5  restart quota enforcement ------------------------------------------

class TestRestartQuota:
    def test_restart_blocked_when_over_quota(self, client, db, seed_simple_mode):
        user = _make_user(db, daily_limit=0)
        _auth_cookie(client, user)
        run = Run(id="QUOTA1", filename="cv.docx", file_type="docx", status="failed",
                  user_id=user.id, started_at=datetime.now())
        db.add(run)
        db.commit()

        resp = client.post("/api/run/QUOTA1/restart")

        assert resp.status_code == 429
        assert resp.json()["detail"]["error"] == "rate_limited"

    def test_restart_passes_quota_check_when_under_limit(self, client, db, seed_simple_mode):
        # daily_limit None -> system default (10); one run is well under it, so the
        # quota check passes and we fall through to the missing-file 404. That 404
        # is the proof the quota gate did NOT false-block a legitimate restart.
        user = _make_user(db, email="ok@example.com")
        _auth_cookie(client, user)
        run = Run(id="QUOTA2", filename="cv.docx", file_type="docx", status="failed",
                  user_id=user.id, started_at=datetime.now())
        db.add(run)
        db.commit()

        resp = client.post("/api/run/QUOTA2/restart")

        assert resp.status_code == 404
        assert resp.json()["detail"]["error"] == "file_not_found"


# --- #4  empty / unreadable upload rejection --------------------------------

class TestEmptyUploadRejection:
    def test_extract_text_reads_docx(self):
        from app.api.upload import _extract_text
        text = _extract_text(_docx_bytes("Hello world\nSecond line"), ".docx")
        assert "Hello world" in text
        assert "Second line" in text

    def test_upload_rejects_a_corrupt_pdf(self, client, db, seed_simple_mode):
        # #806: unlike a docx read failure, a PDF that cannot be parsed never
        # fails open -- the run's conversion uses the same parser and would
        # fail too -- so it is a 400 up front, never a 500 or a doomed run.
        from app.services.pdf_sandbox import PDF_UNREADABLE_MESSAGE
        user = _make_user(db)
        _auth_cookie(client, user)
        resp = client.post(
            "/api/upload",
            files={"file": ("cv.pdf", b"%PDF-1.4 not really a pdf", "application/pdf")},
            data={"submission_type": "own_cv"},
        )
        assert resp.status_code == 400
        assert resp.json()["detail"]["message"] == PDF_UNREADABLE_MESSAGE
        assert db.query(Run).count() == 0

    def test_upload_rejects_image_only_pdf(self, client, db, seed_simple_mode):
        # #806: a fully scanned PDF is rejected before any LLM spend, and
        # (#1282) the message says it is scanned and how to fix it.
        from unified_pipeline.tests.test_pdf_to_docx import _make_pdf
        user = _make_user(db)
        _auth_cookie(client, user)
        resp = client.post(
            "/api/upload",
            files={"file": ("scan.pdf", _make_pdf([[]], image_pages=(0,)), "application/pdf")},
            data={"submission_type": "own_cv"},
        )
        assert resp.status_code == 400
        assert resp.json()["detail"]["message"].startswith("Page 1 of this PDF is a scanned image")
        assert db.query(Run).count() == 0

    def test_upload_rejects_unreadable_docx(self, client, db, seed_simple_mode):
        user = _make_user(db)
        _auth_cookie(client, user)
        content = _docx_bytes("Hi")  # far below the 500-char floor

        resp = client.post("/api/upload", files={"file": ("cv.docx", content, DOCX_MIME)}, data={"submission_type": "own_cv"})

        assert resp.status_code == 400
        assert "couldn't read" in resp.json()["detail"]["message"].lower()

    def test_upload_accepts_readable_docx(self, client, db, seed_simple_mode):
        user = _make_user(db, email="rich@example.com")
        _auth_cookie(client, user)
        content = _docx_bytes("Curriculum Vitae. Extensive academic record. " * 30)

        resp = client.post("/api/upload", files={"file": ("cv.docx", content, DOCX_MIME)}, data={"submission_type": "own_cv"})

        assert resp.status_code == 200
        run_id = resp.json()["run_id"]
        from app.api.upload import UPLOAD_DIR
        (UPLOAD_DIR / f"{run_id}.docx").unlink(missing_ok=True)


# --- #2  real per-step retry -------------------------------------------------

def _seed_failed_run(db, user, run_id="RETRY1", failed_at=6):
    run = Run(id=run_id, filename="cv.docx", file_type="docx", status="failed",
              user_id=user.id, started_at=datetime.now(), completed_at=datetime.now(),
              total_cost=0.5, error_message="boom")
    db.add(run)
    for sd in STEP_REGISTRY:
        if sd.number < failed_at:
            status, cost, err = "complete", 1.0, None
        elif sd.number == failed_at:
            status, cost, err = "error", None, "boom"
        else:
            status, cost, err = "pending", None, None
        db.add(Step(run_id=run_id, step_number=sd.number, stage_id=sd.stage_id,
                    step_name=sd.name, status=status, cost=cost, error_message=err))
    db.commit()
    return run


class TestRetryStep:
    def test_retry_resumes_from_failed_step(self, client, db, seed_simple_mode):
        user = _make_user(db)
        _auth_cookie(client, user)
        _seed_failed_run(db, user, run_id="RETRY1", failed_at=6)

        from app.api.upload import UPLOAD_DIR
        upload_file = UPLOAD_DIR / "RETRY1.docx"
        upload_file.write_bytes(b"dummy")
        try:
            with patch("app.api.runs.PipelineOrchestrator") as MockOrch:
                MockOrch.return_value.execute = AsyncMock(return_value=None)
                resp = client.post("/api/run/RETRY1/retry/6")

                assert resp.status_code == 200
                assert resp.json()["status"] == "running"
                MockOrch.return_value.execute.assert_awaited_once()
                _, kwargs = MockOrch.return_value.execute.call_args
                assert kwargs.get("start_step_number") == 6
            # TestClient runs the background task before returning; the retry
            # must have released the slot it took for this run (#116 drains by id).
            assert concurrency.active_run_ids() == []
        finally:
            upload_file.unlink(missing_ok=True)

        db.expire_all()
        run = db.query(Run).filter(Run.id == "RETRY1").first()
        assert run.status == "running"
        assert run.error_message is None
        assert run.completed_at is None

        steps = {s.step_number: s for s in db.query(Step).filter(Step.run_id == "RETRY1")}
        assert steps[5].status == "complete"        # before the failure: untouched
        assert steps[5].cost == 1.0
        assert steps[6].status == "pending"          # failed step: reset
        assert steps[6].error_message is None
        assert steps[12].status == "pending"         # downstream: reset

    def test_retry_restarts_started_at_so_the_reaper_skips_it(
        self, client, db, seed_simple_mode
    ):
        """#145: a retry of a run first started hours ago must not look stale.
        reconcile_stale_runs ages runs by started_at, and the retry used to keep
        the original one, so the next sweep failed the run mid-execution."""
        user = _make_user(db)
        _auth_cookie(client, user)
        run = _seed_failed_run(db, user, run_id="RETRY4", failed_at=6)
        run.started_at = datetime.now() - timedelta(hours=3)
        db.commit()

        from app.api.upload import UPLOAD_DIR
        upload_file = UPLOAD_DIR / "RETRY4.docx"
        upload_file.write_bytes(b"dummy")
        try:
            with patch("app.api.runs.PipelineOrchestrator") as MockOrch:
                MockOrch.return_value.execute = AsyncMock(return_value=None)
                assert client.post("/api/run/RETRY4/retry/6").status_code == 200
        finally:
            upload_file.unlink(missing_ok=True)

        db.expire_all()
        assert reconcile_stale_runs(db) == 0
        run = db.query(Run).filter(Run.id == "RETRY4").first()
        assert run.status == "running"
        assert run.started_at > datetime.now() - timedelta(minutes=1)

    def test_retry_rejects_non_failed_step(self, client, db, seed_simple_mode):
        user = _make_user(db)
        _auth_cookie(client, user)
        _seed_failed_run(db, user, run_id="RETRY2", failed_at=6)

        # step 5 is complete, not error -> 400
        resp = client.post("/api/run/RETRY2/retry/5")

        assert resp.status_code == 400
        assert resp.json()["detail"]["error"] == "bad_request"

    def test_retry_404_when_upload_file_missing(self, client, db, seed_simple_mode):
        user = _make_user(db)
        _auth_cookie(client, user)
        _seed_failed_run(db, user, run_id="RETRY3", failed_at=6)
        # no uploads/RETRY3.docx created

        resp = client.post("/api/run/RETRY3/retry/6")

        assert resp.status_code == 404


class TestRetryStepQueueModeEnqueueFailure:
    """B3: a failed enqueue during a queue-mode retry must leave the run AND
    its reset step retryable -- not the run reverted to "failed" while its
    step is stuck "pending" with no executor coming and retry_step's own
    `status != "error"` guard rejecting a second attempt outright."""

    def test_failed_enqueue_reverts_the_step_too_then_a_second_retry_succeeds(
        self, client, db, seed_simple_mode, monkeypatch,
    ):
        from app.pipeline import run_queue

        user = _make_user(db)
        _auth_cookie(client, user)
        _seed_failed_run(db, user, run_id="RETRYQ1", failed_at=6)
        monkeypatch.setenv("CVICHE_DISPATCH_MODE", "queue")
        monkeypatch.setenv("CVICHE_REDIS_URL", "redis://fake-valkey:6379/0")

        from app.api.upload import UPLOAD_DIR
        upload_file = UPLOAD_DIR / "RETRYQ1.docx"
        upload_file.write_bytes(b"dummy")

        def enqueue_fails(run_id, queue):
            raise redis.exceptions.ConnectionError("valkey unreachable")
        monkeypatch.setattr(run_queue, "enqueue", enqueue_fails)

        try:
            resp = client.post("/api/run/RETRYQ1/retry/6")

            assert resp.status_code == 503

            db.expire_all()
            run = db.query(Run).filter(Run.id == "RETRYQ1").one()
            assert run.status == "failed"
            step6 = db.query(Step).filter(Step.run_id == "RETRYQ1", Step.step_number == 6).one()
            assert step6.status == "error", "must stay retryable, not stranded pending with no executor"
            assert step6.error_message == "boom"
            step12 = db.query(Step).filter(Step.run_id == "RETRYQ1", Step.step_number == 12).one()
            assert step12.status == "pending", "unaffected downstream step: reset then reverted back to pending"

            # Second retry: enqueue now works.
            enqueued = []
            monkeypatch.setattr(run_queue, "enqueue", lambda run_id, queue: enqueued.append(run_id) or "1-0")

            resp2 = client.post("/api/run/RETRYQ1/retry/6")

            assert resp2.status_code == 202
            assert enqueued == ["RETRYQ1"]
            db.expire_all()
            run = db.query(Run).filter(Run.id == "RETRYQ1").one()
            assert run.status == "queued"
            step6 = db.query(Step).filter(Step.run_id == "RETRYQ1", Step.step_number == 6).one()
            assert step6.status == "pending"
        finally:
            upload_file.unlink(missing_ok=True)


# --- #2  orchestrator resume bookkeeping ------------------------------------

def test_prepare_resume_preserves_cost_and_loads_prior_outputs(db, tmp_path):
    import shutil

    from app.pipeline.orchestrator import PipelineOrchestrator

    run = Run(id="RSME01", filename="cv.docx", file_type="docx", status="failed",
              total_cost=1.5, started_at=datetime.now())
    db.add(run)
    db.commit()

    fake_file = tmp_path / "RSME01.docx"
    fake_file.write_bytes(b"x")
    orch = PipelineOrchestrator("RSME01", fake_file, db)
    try:
        # On-disk outputs exist for every stage before the failed step (#6 = stage '4').
        prior = {sid: tmp_path / f"{sid}.json" for sid in ["1a", "1b", "2", "3a", "3b"]}
        for p in prior.values():
            p.write_text("{}")
        all_paths = {**prior, "4": tmp_path / "missing_4.json"}

        with patch.object(orch, "_get_output_paths", return_value=all_paths):
            orch._prepare_resume(run, start_step_number=6)

        # Cost continues from the run's existing total rather than restarting at 0.
        assert orch.total_cost == 1.5
        # Only the completed earlier stages with on-disk output are registered.
        assert set(orch.stage_outputs) == {"1a", "1b", "2", "3a", "3b"}
    finally:
        shutil.rmtree(orch.web_output_dir, ignore_errors=True)


def test_resume_missing_upstream_output_gives_friendly_error(db, tmp_path):
    """When a resumed stage can't find an earlier stage's on-disk output (pod
    recycled), the failed banner must show a clear 'use Restart' message rather
    than a raw filesystem path + errno."""
    import asyncio
    import shutil

    from app.pipeline.orchestrator import PipelineOrchestrator

    run = Run(id="RSME02", filename="cv.docx", file_type="docx", status="running",
              total_cost=0.2, started_at=datetime.now())
    db.add(run)
    db.commit()

    fake_file = tmp_path / "RSME02.docx"
    fake_file.write_bytes(b"x")
    orch = PipelineOrchestrator("RSME02", fake_file, db)
    raw = "[Errno 2] No such file or directory: '/x/stage_4_field_extraction/RSME02_fields.json'"
    try:
        with patch.object(orch, "_copy_to_pipeline_input", return_value=str(fake_file)), \
             patch.object(orch, "execute_step", new=AsyncMock(side_effect=FileNotFoundError(raw))):
            with pytest.raises(FileNotFoundError):
                asyncio.run(orch.execute(start_step_number=8))

        db.refresh(run)
        assert run.status == "failed"
        msg = run.error_message
        assert "resume" in msg.lower() and "restart" in msg.lower()
        # The raw path / errno must NOT leak into the user-facing banner message.
        assert "Errno" not in msg and "stage_4" not in msg
    finally:
        shutil.rmtree(orch.web_output_dir, ignore_errors=True)


# --- #1  per-pod concurrency admission control ------------------------------

class TestConcurrencyModule:
    """The process-global slot counter that bounds concurrent in-process runs."""

    @pytest.fixture(autouse=True)
    def _reset(self):
        # The counter is module-global; isolate each test from slot leakage.
        concurrency._active_run_ids.clear()
        yield
        concurrency._active_run_ids.clear()

    def test_acquire_release_roundtrip(self, monkeypatch):
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "2")
        assert concurrency.try_acquire_slot("PLUM01") is True
        assert concurrency.try_acquire_slot("PEAR02") is True
        assert concurrency.active_count() == 2
        assert concurrency.try_acquire_slot("FIGS03") is False   # at cap -> rejected
        concurrency.release_slot("PLUM01")
        assert concurrency.active_count() == 1
        assert concurrency.active_run_ids() == ["PEAR02"]
        assert concurrency.try_acquire_slot("FIGS03") is True     # slot freed

    def test_env_cap_respected(self, monkeypatch):
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "1")
        assert concurrency.try_acquire_slot("SLOTRUN") is True
        assert concurrency.try_acquire_slot("SLOTRUN") is False

    def test_nonpositive_cap_falls_back_to_default(self, monkeypatch):
        # A 0/negative cap would wedge the pod; we treat it as the default.
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "0")
        assert concurrency.get_max_concurrent_runs() == concurrency.DEFAULT_MAX_CONCURRENT_RUNS

    def test_release_never_goes_negative(self):
        concurrency.release_slot("SLOTRUN")
        assert concurrency.active_count() == 0

    def test_losing_starter_release_keeps_winner_tracked(self, monkeypatch):
        # Two starters of the same run both hold a slot until claim_run_as_running
        # picks one; the loser's release must not untrack the winner's run.
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "3")
        assert concurrency.try_acquire_slot("TWIN01") is True
        assert concurrency.try_acquire_slot("TWIN01") is True
        concurrency.release_slot("TWIN01")
        assert concurrency.active_run_ids() == ["TWIN01"]

    def test_draining_refuses_new_slots(self, monkeypatch):
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "3")
        assert concurrency.try_acquire_slot("KEPT01") is True
        concurrency.begin_draining()
        assert concurrency.try_acquire_slot("LATE02") is False
        # The run admitted before the drain keeps its slot.
        assert concurrency.active_run_ids() == ["KEPT01"]

    def test_wait_for_drain_returns_once_the_active_run_finishes(self, monkeypatch):
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "3")
        assert concurrency.try_acquire_slot("SLOW01") is True

        async def finish_later():
            await asyncio.sleep(0.05)
            concurrency.release_slot("SLOW01")

        async def drain():
            finisher = asyncio.create_task(finish_later())
            started = time.monotonic()
            remaining = await concurrency.wait_for_drain(5, poll_seconds=0.01)
            await finisher
            return remaining, time.monotonic() - started

        remaining, waited = asyncio.run(drain())
        assert remaining == []
        assert 0.04 < waited < 1   # waited for the run, not for the 5s budget

    def test_wait_for_drain_returns_still_active_runs_at_budget(self, monkeypatch):
        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "3")
        assert concurrency.try_acquire_slot("STUCK1") is True
        remaining = asyncio.run(concurrency.wait_for_drain(0.05, poll_seconds=0.01))
        assert remaining == ["STUCK1"]


class TestConcurrencyAdmission:
    """The 429 gate on the endpoints that launch a pipeline."""

    @pytest.fixture(autouse=True)
    def _reset(self):
        concurrency._active_run_ids.clear()
        yield
        concurrency._active_run_ids.clear()

    def test_start_rejected_when_at_capacity(self, client, db, seed_simple_mode, monkeypatch):
        from app.api.upload import UPLOAD_DIR
        user = _make_user(db)
        _auth_cookie(client, user)
        run = Run(id="BUSY01", filename="cv.docx", file_type="docx", status="created",
                  user_id=user.id, started_at=datetime.now())
        db.add(run)
        db.commit()
        upload_file = UPLOAD_DIR / "BUSY01.docx"
        upload_file.write_bytes(b"dummy")

        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "1")
        assert concurrency.try_acquire_slot("SLOTRUN") is True   # fill the only slot
        try:
            resp = client.post("/api/run/BUSY01/start")

            assert resp.status_code == 429
            assert resp.json()["detail"]["error"] == "server_busy"
            assert resp.headers.get("Retry-After") == "30"
            # A rejected start must leave the run in its prior state, retryable.
            db.expire_all()
            assert db.query(Run).filter(Run.id == "BUSY01").first().status == "created"
        finally:
            upload_file.unlink(missing_ok=True)

    def test_start_proceeds_and_releases_under_capacity(self, client, db, seed_simple_mode):
        from app.api.upload import UPLOAD_DIR
        user = _make_user(db, email="free@example.com")
        _auth_cookie(client, user)
        run = Run(id="FREE01", filename="cv.docx", file_type="docx", status="created",
                  user_id=user.id, started_at=datetime.now())
        db.add(run)
        db.commit()
        upload_file = UPLOAD_DIR / "FREE01.docx"
        upload_file.write_bytes(b"dummy")
        try:
            with patch("app.api.runs.PipelineOrchestrator") as MockOrch:
                MockOrch.return_value.execute = AsyncMock(return_value=None)
                resp = client.post("/api/run/FREE01/start")

                assert resp.status_code == 200
                assert resp.json()["status"] == "running"
            # TestClient runs the background task before returning, so the slot
            # acquired for the run must have been released in the task's finally.
            assert concurrency.active_count() == 0
        finally:
            upload_file.unlink(missing_ok=True)

    def test_retry_rejected_when_at_capacity(self, client, db, seed_simple_mode, monkeypatch):
        from app.api.upload import UPLOAD_DIR
        user = _make_user(db, email="busy2@example.com")
        _auth_cookie(client, user)
        _seed_failed_run(db, user, run_id="BUSY02", failed_at=6)
        upload_file = UPLOAD_DIR / "BUSY02.docx"
        upload_file.write_bytes(b"dummy")

        monkeypatch.setenv("CVICHE_MAX_CONCURRENT_RUNS", "1")
        assert concurrency.try_acquire_slot("SLOTRUN") is True
        try:
            resp = client.post("/api/run/BUSY02/retry/6")

            assert resp.status_code == 429
            assert resp.json()["detail"]["error"] == "server_busy"
            # A rejected retry must not touch run or step state.
            db.expire_all()
            assert db.query(Run).filter(Run.id == "BUSY02").first().status == "failed"
            step6 = db.query(Step).filter(Step.run_id == "BUSY02", Step.step_number == 6).first()
            assert step6.status == "error"
        finally:
            upload_file.unlink(missing_ok=True)


# --- #799  atomic -> running transition -------------------------------------

def _stale_read_after_winner_commits(monkeypatch, run_id, stale_status):
    """Model the race: the loser reads the run, then the winner commits.

    Wraps ``check_run_access`` so that after the loser's read returns, another
    starter flips the row to "running" and commits; the loser's in-memory
    object still says ``stale_status``, so its status guard passes.
    """
    from sqlalchemy import update
    from sqlalchemy.orm.attributes import set_committed_value

    from app.api import runs as runs_api

    real = runs_api.check_run_access

    def racing(rid, user, db, **kw):
        run = real(rid, user, db, **kw)
        db.execute(update(Run).where(Run.id == run_id).values(status="running"))
        db.commit()
        set_committed_value(run, "status", stale_status)
        return run

    monkeypatch.setattr(runs_api, "check_run_access", racing)


class TestAtomicRunStart:
    @pytest.fixture(autouse=True)
    def _reset(self):
        concurrency._active_run_ids.clear()
        yield
        concurrency._active_run_ids.clear()

    @pytest.mark.parametrize("status", ["created", "paused"])
    def test_second_starter_loses_with_409_and_dispatches_nothing(
        self, client, db, seed_simple_mode, monkeypatch, status
    ):
        from app.api.upload import UPLOAD_DIR
        user = _make_user(db, email=f"race-{status}@example.com")
        _auth_cookie(client, user)
        run_id = f"RACE{status[:2].upper()}1"
        db.add(Run(id=run_id, filename="cv.docx", file_type="docx", status=status,
                   user_id=user.id, started_at=datetime.now()))
        db.commit()
        upload_file = UPLOAD_DIR / f"{run_id}.docx"
        upload_file.write_bytes(b"dummy")
        _stale_read_after_winner_commits(monkeypatch, run_id, status)
        try:
            with patch("app.api.runs.PipelineOrchestrator") as MockOrch:
                resp = client.post(f"/api/run/{run_id}/start")

                assert resp.status_code == 409
                assert resp.json()["detail"]["error"] == "conflict"
                MockOrch.assert_not_called()
            assert concurrency.active_count() == 0   # loser's slot returned
        finally:
            upload_file.unlink(missing_ok=True)

    def test_sequential_second_start_is_rejected_and_first_dispatches_once(
        self, client, db, seed_simple_mode
    ):
        from app.api.upload import UPLOAD_DIR
        user = _make_user(db, email="seq@example.com")
        _auth_cookie(client, user)
        db.add(Run(id="SEQ001", filename="cv.docx", file_type="docx", status="created",
                   user_id=user.id, started_at=datetime.now()))
        db.commit()
        upload_file = UPLOAD_DIR / "SEQ001.docx"
        upload_file.write_bytes(b"dummy")
        try:
            with patch("app.api.runs.PipelineOrchestrator") as MockOrch:
                MockOrch.return_value.execute = AsyncMock(return_value=None)
                first = client.post("/api/run/SEQ001/start")
                second = client.post("/api/run/SEQ001/start")

                assert first.status_code == 200
                assert second.status_code == 400
                MockOrch.return_value.execute.assert_awaited_once()
        finally:
            upload_file.unlink(missing_ok=True)

    def test_second_retry_loses_with_409_and_leaves_steps_untouched(
        self, client, db, seed_simple_mode, monkeypatch
    ):
        from app.api.upload import UPLOAD_DIR
        user = _make_user(db, email="raceretry@example.com")
        _auth_cookie(client, user)
        _seed_failed_run(db, user, run_id="RACER1", failed_at=6)
        upload_file = UPLOAD_DIR / "RACER1.docx"
        upload_file.write_bytes(b"dummy")
        _stale_read_after_winner_commits(monkeypatch, "RACER1", "failed")
        try:
            with patch("app.api.runs.PipelineOrchestrator") as MockOrch:
                resp = client.post("/api/run/RACER1/retry/6")

                assert resp.status_code == 409
                MockOrch.assert_not_called()
            assert concurrency.active_count() == 0
            db.expire_all()
            step6 = db.query(Step).filter(Step.run_id == "RACER1", Step.step_number == 6).first()
            assert step6.status == "error"   # the loser did not reset it
        finally:
            upload_file.unlink(missing_ok=True)


# --- #299  document_uid validation and per-run input copy --------------------

@pytest.mark.parametrize("stem", ["bad name", "a.b", "x" * 129, "\u00e9vil", "abc\n"])
def test_orchestrator_rejects_unsafe_document_uid(db, tmp_path, stem):
    from app.pipeline.orchestrator import PipelineOrchestrator

    with pytest.raises(ValueError, match="Invalid document_uid"):
        PipelineOrchestrator("REJ299", tmp_path / f"{stem}.docx", db)
    # A rejected constructor must not leave its output dir behind.
    assert not (Path(__file__).parent.parent.parent / "outputs" / "REJ299").exists()


@pytest.mark.parametrize("run_id", ["../evil", "a/b", ""])
def test_orchestrator_rejects_unsafe_run_id(db, tmp_path, run_id):
    from app.pipeline.orchestrator import PipelineOrchestrator

    with pytest.raises(ValueError, match="run_id"):
        PipelineOrchestrator(run_id, tmp_path / "cv.docx", db)


def test_orchestrator_rejects_document_uid_that_differs_from_run_id(db, tmp_path):
    from app.pipeline.orchestrator import PipelineOrchestrator

    with pytest.raises(ValueError, match="must equal run_id"):
        PipelineOrchestrator("MATCH1", tmp_path / "OTHER1.docx", db)
    with pytest.raises(ValueError, match="must equal run_id"):
        PipelineOrchestrator("ABC1", tmp_path / "abc1.docx", db)
    PipelineOrchestrator("MATCH2", tmp_path / "MATCH2.docx", db)


def test_copy_to_pipeline_input_is_per_run_and_overwrites(db, tmp_path, monkeypatch):
    import shutil

    from app.pipeline import orchestrator as orch

    monkeypatch.setattr(orch, "PARENT_DIR", tmp_path / "repo")
    src = tmp_path / "COPY01.docx"
    src.write_bytes(b"first")
    o = orch.PipelineOrchestrator("COPY01", src, db)
    try:
        dest = o._copy_to_pipeline_input()
        assert dest == str(tmp_path / "repo/data/sample_cvs/word/COPY01/COPY01.docx")
        assert Path(dest).read_bytes() == b"first"

        src.write_bytes(b"second")
        o._copy_to_pipeline_input()
        assert Path(dest).read_bytes() == b"second"
        assert o.pdf_conversion is None  # a docx is copied, never converted
    finally:
        shutil.rmtree(o.web_output_dir, ignore_errors=True)


def test_copy_to_pipeline_input_converts_a_pdf(db, tmp_path, monkeypatch, cv_pdf):
    """#806: a PDF upload is converted into the run's private docx path; the
    upload itself is left untouched as the original."""
    import shutil

    from app.pipeline import orchestrator as orch

    monkeypatch.setattr(orch, "PARENT_DIR", tmp_path / "repo")
    src = tmp_path / "CONV01.pdf"
    src.write_bytes(cv_pdf())
    o = orch.PipelineOrchestrator("CONV01", src, db)
    try:
        dest = o._copy_to_pipeline_input()
        assert dest == str(tmp_path / "repo/data/sample_cvs/word/CONV01/CONV01.docx")
        text = "\n".join(p.text for p in Document(dest).paragraphs)
        assert "EDUCATION" in text and "Doctor of Medicine, Example University" in text
        assert o.pdf_conversion.image_only_pages == []
        assert src.read_bytes() == cv_pdf()
    finally:
        shutil.rmtree(o.web_output_dir, ignore_errors=True)


def _execute_one_noop_step(db, monkeypatch, tmp_path, run_id, upload, attempts=(None,)):
    """execute() over the REAL _copy_to_pipeline_input and a one-step no-op
    registry, once per start_step_number in `attempts`; returns the run's
    Log rows."""
    import shutil
    from types import SimpleNamespace

    from app.pipeline import orchestrator as orch

    monkeypatch.setattr(orch, "PARENT_DIR", tmp_path / "repo")
    monkeypatch.setattr(orch, "event_emitter", AsyncMock())
    monkeypatch.setattr(orch, "STEP_REGISTRY", [SimpleNamespace(number=1, stage_id="1a", name="1a")])
    monkeypatch.setenv("CVICHE_RUN_DOCTOR", "0")
    db.add(Run(id=run_id, filename="cv.pdf", file_type="pdf", status="running",
               started_at=datetime.now()))
    db.commit()
    o = orch.PipelineOrchestrator(run_id, upload, db)
    monkeypatch.setattr(o, "execute_step", AsyncMock())
    try:
        for start in attempts:
            try:
                asyncio.run(o.execute(start_step_number=start))
            except Exception:
                pass  # a failed run is asserted on by the caller, via its row
    finally:
        shutil.rmtree(o.web_output_dir, ignore_errors=True)
    return db.query(Log).filter(Log.run_id == run_id).all()


def test_image_only_pdf_page_is_a_run_warning(db, tmp_path, monkeypatch, cv_pdf):
    """#536: a converted PDF's image-only page is named in a WARNING run-log
    row under the first stage, where the user sees it."""
    src = tmp_path / "SCAN01.pdf"
    src.write_bytes(cv_pdf(image_pages=(1,)))
    logs = _execute_one_noop_step(db, monkeypatch, tmp_path, "SCAN01", src)
    warnings = [log for log in logs if log.level == "WARNING"]
    assert len(warnings) == 1
    assert warnings[0].step_number == 1
    assert "page(s) 2 contain only images" in warnings[0].message


def test_image_only_warning_is_not_repeated_on_retry(db, tmp_path, monkeypatch, cv_pdf):
    """A retry/resume re-converts the PDF, but the step-1 log already
    carries the warning from the first attempt (Log rows are never deleted)."""
    src = tmp_path / "SCAN02.pdf"
    src.write_bytes(cv_pdf(image_pages=(1,)))
    logs = _execute_one_noop_step(db, monkeypatch, tmp_path, "SCAN02", src, attempts=(None, 1))
    assert len([log for log in logs if log.level == "WARNING"]) == 1


def test_pdf_over_a_parse_limit_fails_the_run_with_its_message(db, tmp_path, monkeypatch):
    """The run's conversion hits the same limits as the upload check and
    fails with the same user-facing message, not the generic one."""
    from app.services.pdf_sandbox import PDF_MAX_PAGES, PDF_TOO_COMPLEX_MESSAGE
    from unified_pipeline.tests.test_pdf_to_docx import _make_pdf
    src = tmp_path / "LONG01.pdf"
    src.write_bytes(_make_pdf([[]] * (PDF_MAX_PAGES + 1)))
    _execute_one_noop_step(db, monkeypatch, tmp_path, "LONG01", src)
    db.expire_all()
    run = db.get(Run, "LONG01")
    assert run.status == "failed"
    assert run.error_message == PDF_TOO_COMPLEX_MESSAGE


def test_run_waiting_too_long_for_a_pdf_slot_fails_with_its_message(db, tmp_path, monkeypatch, cv_pdf):
    """#806 review B1: the run's conversion waits a bounded time for a PDF
    child slot, then fails the run with a message that says why."""
    from app.services import pdf_sandbox
    src = tmp_path / "BUSY01.pdf"
    src.write_bytes(cv_pdf())
    monkeypatch.setattr(pdf_sandbox, "PDF_CONVERT_SLOT_WAIT_SECONDS", 0.2)
    for _ in range(pdf_sandbox.PDF_CHILD_SLOTS):
        assert pdf_sandbox._slots.acquire(blocking=False)
    try:
        _execute_one_noop_step(db, monkeypatch, tmp_path, "BUSY01", src)
    finally:
        for _ in range(pdf_sandbox.PDF_CHILD_SLOTS):
            pdf_sandbox._slots.release()
    db.expire_all()
    run = db.get(Run, "BUSY01")
    assert run.status == "failed"
    assert run.error_message == pdf_sandbox.PDF_BUSY_RUN_MESSAGE


def test_conversion_runs_off_the_event_loop_thread(db, tmp_path, monkeypatch):
    """#806 review N6: the conversion can wait for a slot and run for
    minutes; it must not block the run's event loop (cancel included)."""
    import threading

    from app.pipeline import orchestrator as orch
    seen = {}
    real = orch.PipelineOrchestrator._copy_to_pipeline_input

    def record(self):
        seen["thread"] = threading.current_thread()
        return real(self)
    monkeypatch.setattr(orch.PipelineOrchestrator, "_copy_to_pipeline_input", record)
    src = tmp_path / "LOOP01.docx"
    src.write_bytes(b"docx bytes")
    _execute_one_noop_step(db, monkeypatch, tmp_path, "LOOP01", src)
    assert seen["thread"] is not threading.main_thread()


def test_text_only_pdf_logs_no_warning(db, tmp_path, monkeypatch, cv_pdf):
    src = tmp_path / "TEXT01.pdf"
    src.write_bytes(cv_pdf())
    logs = _execute_one_noop_step(db, monkeypatch, tmp_path, "TEXT01", src)
    assert [log for log in logs if log.level == "WARNING"] == []
