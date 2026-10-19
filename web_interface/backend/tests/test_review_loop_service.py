"""app/services/review_loop_service.py (#1587): the corrected-copy diff. The
route over it is in test_feedback_notification.py. Every document here is
synthetic."""
import io
import json
import zipfile
from datetime import UTC, datetime

import pytest
from fastapi import HTTPException

from app.models import CorrectedCopy
from app.services import review_loop_service as svc

_RUN = "RVW001"
_USER = 7  # the uploader: every copy is stored as one user's (#1754)
_NOW = datetime(2026, 10, 13, 9, 30, tzinfo=UTC)
_LATER = datetime(2026, 10, 13, 11, 0, tzinfo=UTC)


def _docx(*paragraphs):
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def store(tmp_path, monkeypatch):
    from app.storage.local_storage import LocalRunStorage

    store = LocalRunStorage(str(tmp_path))
    monkeypatch.setattr(svc, "get_storage", lambda: store)
    return store


_ENTRIES = ("Fellowship in Quorvane studies, Northfield Institute, 2011-2013",
            "Brennic Foundation Award for Saltwick lattice modelling, 2015")


def test_record_corrected_docx_maps_changes_to_stage4_entries(db, store):
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _docx(*_ENTRIES))
    stage4 = {"entries": [{"element_idx_start": 12, "taxonomy_code": "H", "text": _ENTRIES[1]}]}
    store.put_file(_RUN, f"outputs/{_RUN}_fields.json", json.dumps(stage4).encode())

    assert svc.record_corrected_docx(db, _RUN, _USER, _docx(_ENTRIES[0]), _NOW) == svc.CopyCounts(1, 0)

    report = json.loads(store.get_file(_RUN, svc.corrected_diff_key(_RUN, _USER)))
    assert [(c["change_type"], c["element_idx"]) for c in report["changes"]] == [("deleted", 12)]


def test_each_uploader_keeps_one_copy_and_a_pre_1754_copy_is_still_listed(db, store):
    """Two users' copies sit side by side; the same user again replaces
    theirs; a copy stored under the run's id alone (before #1754) is listed
    as uploader None."""
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _docx(*_ENTRIES))
    store.put_file(_RUN, f"corrected/{_RUN}_diff.json", b"{}")
    svc.record_corrected_docx(db, _RUN, _USER, _docx(_ENTRIES[0]), _NOW)
    svc.record_corrected_docx(db, _RUN, 12, _docx(_ENTRIES[1]), _NOW)
    # One build, compared by identity of bytes: a second _docx() call can land
    # in the next 2-second zip timestamp and differ in its header bytes.
    replacement = _docx(*_ENTRIES)
    assert svc.record_corrected_docx(db, _RUN, _USER, replacement, _LATER) == svc.CopyCounts(0, 0)

    assert svc.stored_uploaders(_RUN) == [None, _USER, 12]
    assert store.get_file(_RUN, f"corrected/{_RUN}_u{_USER}_corrected.docx") == replacement
    assert json.loads(store.get_file(_RUN, svc.corrected_diff_key(_RUN, 12)))["by_type"]["deleted"] == 1
    assert svc.corrected_diff_key(_RUN, None) == f"corrected/{_RUN}_diff.json"
    # One corrected_copies row per uploader; the second upload replaced user 7's.
    rows = {r.user_id: r for r in db.query(CorrectedCopy).all()}
    assert set(rows) == {_USER, 12}
    assert (rows[_USER].changes, rows[_USER].uploaded_at) == (0, _LATER)
    assert rows[_USER].storage_key == svc.corrected_docx_key(_RUN, _USER)
    assert (rows[12].changes, rows[12].storage_key) == (1, svc.corrected_docx_key(_RUN, 12))


def test_record_corrected_docx_diffs_without_entries_when_stage4_is_unreadable(db, store):
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _docx(*_ENTRIES))
    store.put_file(_RUN, f"outputs/{_RUN}_fields.json", b"{not json")
    assert svc.record_corrected_docx(db, _RUN, _USER, _docx(_ENTRIES[0]), _NOW).changes == 1
    report = json.loads(store.get_file(_RUN, svc.corrected_diff_key(_RUN, _USER)))
    assert report["changes"][0]["element_idx"] is None


_REVIEWER = "Synthetic Reviewer"


def _commented(*authors):
    """The two entries, unedited, with one comment per author on the first."""
    from docx import Document

    doc = Document()
    first = doc.add_paragraph(_ENTRIES[0])
    doc.add_paragraph(_ENTRIES[1])
    for author in authors:
        doc.add_comment(first.runs, text="Synthetic note.", author=author)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_a_copy_with_only_the_reviewers_comments_counts_them_not_cviches(db, store):
    """#1937: a copy with the reviewer's comments and no text edit read "0
    changes". CViche's own comments, kept from the delivered document
    (stage 6's) and from the review copy, are not the reviewer's."""
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _commented("PubMed Enrichment"))
    store.put_file(_RUN, f"outputs/{_RUN}_wcm_review.docx", _commented("PubMed Enrichment", "CViche check"))
    corrected = _commented("PubMed Enrichment", "CViche check", _REVIEWER, _REVIEWER, "Second Reviewer")

    assert svc.record_corrected_docx(db, _RUN, _USER, corrected, _NOW) == svc.CopyCounts(0, 3)
    row = db.query(CorrectedCopy).one()
    assert (row.changes, row.comments) == (0, 3)


@pytest.mark.parametrize("corrected_authors, cviche_authors, reviewers", [
    ((), ("CViche check",), 0),                                    # no comments at all
    (("CViche check", "CViche"), ("CViche check", "CViche"), 0),   # CViche's only
    ((_REVIEWER,), (), 1),                                         # a run with no CViche comment
    ((_REVIEWER, "CViche check"), ("CViche check",), 1),
])
def test_count_reviewer_comments(corrected_authors, cviche_authors, reviewers):
    assert svc.count_reviewer_comments(_commented(*corrected_authors), frozenset(cviche_authors)) == reviewers


def test_without_a_review_copy_cviches_authors_come_from_the_delivered_document(db, store):
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _commented("PubMed Enrichment"))
    corrected = _commented("PubMed Enrichment", _REVIEWER)
    assert svc.record_corrected_docx(db, _RUN, _USER, corrected, _NOW) == svc.CopyCounts(0, 1)


def test_record_corrected_docx_stores_nothing_when_the_copy_cannot_be_read(db, store):
    """A zip that passes the magic check but whose document part is not XML."""
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _docx(*_ENTRIES))
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_docx("x"))) as src, zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            data = b"<not xml" if item.filename == "word/document.xml" else src.read(item)
            dst.writestr(item, data)
    with pytest.raises(HTTPException) as e:
        svc.record_corrected_docx(db, _RUN, _USER, buf.getvalue(), _NOW)
    assert e.value.status_code == 400
    assert store.list_files(_RUN, svc.CORRECTED_PREFIX) == []
    assert db.query(CorrectedCopy).count() == 0


@pytest.mark.parametrize("filename, uploader, name", [
    ("Synthetic Person CV.docx", "Pat Staff", "Synthetic Person CV - corrected by Pat Staff.docx"),
    ("Synthetic Person CV.pdf", None, "Synthetic Person CV - corrected.docx"),
    (None, "Pat Staff", "CV - corrected by Pat Staff.docx"),
])
def test_corrected_download_name(filename, uploader, name):
    assert svc.corrected_download_name(filename, uploader) == name


@pytest.mark.parametrize("changes, comments, line", [
    (0, 18, "0 text edits and 18 comments recorded"),
    (1, 1, "1 text edit and 1 comment recorded"),
    (7, 0, "7 text edits and 0 comments recorded"),
])
def test_changes_summary(changes, comments, line):
    assert svc.changes_summary(svc.CopyCounts(changes, comments)) == line


# --- #1654: verdicts from the review copy's comments, and the doctor re-run ---

_CITES = ("Quorvane T, Plesk M. Heliotropic drift in vexillary marmosets. J Synth Imag. 2019;12:34-56.",
          "Quorvane T, Abernoth R. Saltwick lattices under brennic loading. Synth Rep. 1912;3:7-19.",
          "Quorvane T. Tremulant gradients across fennish moorland. Moor Synth. 2016;8:101-110.")


def _finding(lint, message, evidence, severity="WARN"):
    return {"lint": lint, "severity": severity, "message": message, "evidence": evidence,
            "status": "ran", "reason": ""}


def _delivered(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("BIBLIOGRAPHY")
    for cite in _CITES:
        doc.add_paragraph(cite)
    path = tmp_path / f"{_RUN}_wcm.docx"
    doc.save(str(path))
    return path


def _with_review_copy(tmp_path, store, with_map=True):
    """The delivered document and its review copy, stored as the orchestrator
    stores them: comment 0 on citation A, 1 on B's year, 2 on C. Without
    ``with_map``, as a copy written before its comment map (#1654)."""
    from app.services.review_comments import comment_map_path, write_review_docx

    clean = _delivered(tmp_path)
    review, _ = write_review_docx(clean, {"findings": [
        _finding("enrichment_failures", "1 publication(s) failed PubMed enrichment", [_CITES[0]]),
        _finding("implausible_year", "entry 12 (S1): end_date=1912 -- before 1959", [_CITES[1]], "ERROR"),
        _finding("pipe_leaks", "1 numbered citation(s) fusing venue-date patterns", [_CITES[2]]),
    ]}, rows={})
    for path in (clean, review, *([comment_map_path(review)] if with_map else [])):
        store.put_file(_RUN, f"outputs/{path.name}", path.read_bytes())
    return review.read_bytes()


def _open(docx):
    from docx import Document

    return Document(io.BytesIO(docx))


def _saved(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _delete_comment(doc, comment_id):
    """What Word does on "Delete comment": the range, the reference and the comment go."""
    from docx.oxml.ns import qn

    for el in list(doc.element.body.iter(qn("w:commentRangeStart"), qn("w:commentRangeEnd"),
                                         qn("w:commentReference"))):
        if el.get(qn("w:id")) == str(comment_id):
            gone = el.getparent() if el.tag == qn("w:commentReference") else el  # the reference's run
            gone.getparent().remove(gone)
    comment = next(c for c in doc.comments if c.comment_id == comment_id)
    comment._comment_elm.getparent().remove(comment._comment_elm)


def _verdicts(report):
    return [(f["comment_id"], f["lint"], f["verdict"]) for f in report["findings"]]


def test_a_review_copy_uploaded_untouched_leaves_every_finding_unknown(db, store, tmp_path):
    review = _with_review_copy(tmp_path, store)
    report = svc.comment_verdicts(db, _RUN, review)
    assert report["comments_tracked"] and report["note"] is None
    assert report["findings_from"] == svc.FINDINGS_FROM_MAP
    assert _verdicts(report) == [(0, "enrichment_failures", "unknown"), (1, "implausible_year", "unknown"),
                                 (2, "pipe_leaks", "unknown")]
    assert report["counts"] == {"fixed": 0, "not_a_problem": 0, "unknown": 3}


def test_a_resolved_edit_is_fixed_and_a_dismissed_comment_not_a_problem(db, store, tmp_path):
    """Comment 1 resolved with B's year corrected; comment 2 deleted with C
    untouched; comment 0 kept. Stored per finding, with no CV text."""
    doc = _open(_with_review_copy(tmp_path, store))
    year = next(r for p in doc.paragraphs for r in p.runs if r.text == "1912")
    year.text = "2012"
    _delete_comment(doc, 1)
    _delete_comment(doc, 2)

    report = svc.store_verdicts(db, _RUN, _USER, _saved(doc))

    assert _verdicts(report) == [(0, "enrichment_failures", "unknown"), (1, "implausible_year", "fixed"),
                                 (2, "pipe_leaks", "not_a_problem")]
    assert report["findings"][1] == {"comment_id": 1, "lint": "implausible_year", "shape": None,
                                     "severity": "ERROR", "entry_index": 12, "verdict": "fixed"}
    stored = store.get_file(_RUN, svc.corrected_verdicts_key(_RUN, _USER)).decode()
    assert json.loads(stored) == report
    assert not any(cite[:30] in stored for cite in _CITES)


def test_the_clean_document_uploaded_has_no_comments_to_follow(db, store, tmp_path):
    _with_review_copy(tmp_path, store)
    report = svc.comment_verdicts(db, _RUN, (tmp_path / f"{_RUN}_wcm.docx").read_bytes())
    assert not report["comments_tracked"] and report["note"] == svc.NO_COMMENTS_NOTE
    assert report["counts"] == {"fixed": 0, "not_a_problem": 0, "unknown": 3}


def test_the_review_notes_comment_is_no_findings_comment(db, store, tmp_path):
    """The notes are one comment by their own author (id 3, after the
    findings'): it gets no verdict, and deleting it alone leaves the
    findings' comments followed, as Delete All Comments would not."""
    from app.services.review_comments import REVIEW_NOTES_AUTHOR

    doc = _open(_with_review_copy(tmp_path, store))
    assert [(c.comment_id, c.author) for c in doc.comments][-1] == (3, REVIEW_NOTES_AUTHOR)
    _delete_comment(doc, 3)
    report = svc.comment_verdicts(db, _RUN, _saved(doc))
    assert report["comments_tracked"]
    assert _verdicts(report) == [(0, "enrichment_failures", "unknown"), (1, "implausible_year", "unknown"),
                                 (2, "pipe_leaks", "unknown")]
    old_copy = svc.comment_verdicts(db, _RUN, _with_review_copy(tmp_path, store, with_map=False))
    assert [f["comment_id"] for f in old_copy["findings"]] == [0, 1, 2]


def test_a_review_copy_written_before_its_comment_map_reads_lints_from_the_wording(db, store, tmp_path):
    review = _with_review_copy(tmp_path, store, with_map=False)
    report = svc.comment_verdicts(db, _RUN, review)
    assert report["findings_from"] == svc.FINDINGS_FROM_WORDING
    assert [(f["lint"], f["severity"], f["entry_index"]) for f in report["findings"]] == [
        ("enrichment_failures", None, None), ("implausible_year", None, None), ("pipe_leaks", None, None)]


def test_a_run_with_no_review_copy_says_so(db, store, tmp_path):
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", _delivered(tmp_path).read_bytes())
    report = svc.comment_verdicts(db, _RUN, _delivered(tmp_path).read_bytes())
    assert report["note"] == svc.NO_REVIEW_COPY_NOTE and report["findings"] == []


_LEAK = "[M2A] Quorvane T. Brennic coupling in moorland lattices. Synth Rep. 2018;5:1-9."


@pytest.mark.parametrize("bad", ["../RVW001", "RVW/001", "RVW001\n", ""])
def test_the_stored_outputs_doctor_refuses_a_malformed_run_id(db, store, bad):
    """No temp path or storage key is built from a run id outside RUN_ID_RE."""
    with pytest.raises(ValueError, match="Invalid run id"):
        svc.doctor_on_stored_outputs(db, bad)


def test_an_artifact_path_outside_the_scratch_dir_is_refused(tmp_path):
    root = (tmp_path / "scratch").resolve()
    root.mkdir()
    assert svc._inside(root, root / "outputs" / "a.json") == root / "outputs" / "a.json"
    with pytest.raises(ValueError, match="escapes"):
        svc._inside(root, root / ".." / "elsewhere.json")


def test_the_doctor_reruns_on_the_corrected_copy_and_counts_what_cleared(db, store, tmp_path):
    """The delivered document leaks a taxonomy code; the reviewer removes it.
    Both reports are the real doctor's, over the same stored artifacts."""
    delivered = _docx("BIBLIOGRAPHY", _LEAK, _CITES[0])
    store.put_file(_RUN, f"outputs/{_RUN}_wcm.docx", delivered)
    store.put_file(_RUN, f"outputs/{_RUN}_doctor.json",
                   json.dumps(svc.doctor_on_corrected(db, _RUN, delivered)).encode())

    comparison = svc.store_corrected_doctor(db, _RUN, _USER, _docx("BIBLIOGRAPHY", _LEAK[6:], _CITES[0]))

    rerun = json.loads(store.get_file(_RUN, svc.corrected_doctor_key(_RUN, _USER)))
    assert rerun["artifacts"]["stage_6_docx"].endswith(f"{_RUN}_wcm.docx")
    assert "output_hygiene" not in {f["lint"] for f in rerun["findings"]}
    rows = {r["lint"]: r for r in comparison["lints"]}
    assert (rows["output_hygiene"]["cleared"], rows["output_hygiene"]["new"]) == (1, 0)
    assert "owner_contact_missing" not in rows and "owner_contact_missing" in comparison["not_compared"]
    assert json.loads(store.get_file(_RUN, svc.corrected_doctor_compare_key(_RUN, _USER))) == comparison


def test_the_rerun_drops_an_old_copys_review_notes_box_but_not_stage6s_own(tmp_path):
    """A review copy delivered before the notes became a comment closes with
    a box in the body; one uploaded later still has it stripped."""
    from docx import Document

    from app.services.review_comments import LEGACY_REVIEW_NOTES_TITLE
    from unified_pipeline.stage6.formatting import add_cviche_box

    doc = Document()
    doc.add_paragraph(_CITES[0])
    add_cviche_box(doc, "CViche note: delete this box before sending")
    add_cviche_box(doc, LEGACY_REVIEW_NOTES_TITLE)
    kept = _open(svc._without_review_notes(_saved(doc)))
    assert [t.cell(0, 0).paragraphs[0].text for t in kept.tables] == [
        "CViche note: delete this box before sending"]


def test_compare_counts_only_lints_that_ran_in_both_reports():
    ran = {"lint": "pipe_leaks", "severity": "WARN", "status": "ran", "message": "2 fused"}
    delivered = {"findings": [ran, {**ran, "lint": "segmentation", "message": "lost"}]}
    corrected = {"findings": [{**ran, "message": "1 fused"},
                              {"lint": "segmentation", "severity": "INFO", "status": "skipped",
                               "message": "skipped: missing source"}]}
    comparison = svc.compare_doctor(delivered, corrected)
    assert comparison["not_compared"] == ["segmentation"]
    assert comparison["lints"] == [{"lint": "pipe_leaks", "shape": None, "delivered": 1, "corrected": 1,
                                    "persisting": 0, "cleared": 1, "new": 1}]


def test_a_lint_that_reads_no_document_is_never_counted_cleared():
    """The re-run has no prompt logs: llm_fallback_served's prompt-log finding
    is absent from it whatever the reviewer did, and must not read as cleared."""
    served = {"lint": "llm_fallback_served", "severity": "WARN", "status": "ran",
              "message": "stage 3b served a fallback model on 1 call"}
    comparison = svc.compare_doctor({"findings": [served]}, {"findings": []})
    assert comparison["not_compared"] == ["llm_fallback_served"]
    assert comparison["lints"] == [] and comparison["totals"]["cleared"] == 0


def test_a_failing_step_is_logged_and_the_other_still_runs(db, store, monkeypatch, caplog):
    ran = []

    def boom(*_args):
        raise ValueError("synthetic")

    monkeypatch.setattr(svc, "store_verdicts", boom)
    monkeypatch.setattr(svc, "store_corrected_doctor", lambda *a: ran.append(a))
    svc.review_corrected_copy(db, _RUN, _USER, b"docx")
    assert len(ran) == 1
    assert "verdicts failed" in caplog.text
