"""Auto-run doctor integration (post-completion artifact lints).

Covers the orchestrator hook that runs unified_pipeline.run_doctor after a
successful run:

  * CVICHE_RUN_DOCTOR=0 disables it (no doctor call, no report file);
  * default-on: the report JSON is written under the pipeline outputs dir
    (stage_7_doctor/<uid>_doctor.json), registered on the final step's
    output_files, and summarized as a "Doctor" fact on the terminal Teams card;
  * a doctor crash is logged and swallowed -- the run still completes and the
    terminal notification still goes out (without the Doctor line).
"""
import asyncio
import json
import logging
import os
import sys
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

os.environ.setdefault("CVICHE_SESSION_SECRET", "test-secret-not-for-production")

# Make unified_pipeline.* importable directly (tests/ -> backend/ ->
# web_interface/ -> project_root/ -> src/), matching test_run_duration_coverage.
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent / "src"))

import unified_pipeline.run_doctor as run_doctor_mod
from app.services import notifications


def _doctor_payload():
    """A canned run_doctor report with one substantive (WARN) finding."""
    return {
        "document_uid": "DOC_ON",
        "root": "unused",
        "artifacts": {},
        "findings": [
            {"lint": "segmentation", "severity": "WARN",
             "message": "1 substantive source line(s) unaccounted for",
             "evidence": []},
        ],
        "counts": {"ERROR": 0, "WARN": 1, "INFO": 0},
        "worst_severity": "WARN",
    }


def _orchestrator(monkeypatch, tmp_path, db, run_id):
    """A run-to-completion orchestrator with the real pipeline stubbed out
    (mirrors test_run_duration_coverage) and pipeline outputs redirected to
    tmp_path so the doctor report never lands in the repo's outputs dir."""
    from app.models import Run, Step
    from app.pipeline import orchestrator as orch

    db.add(Run(
        id=run_id, filename="cv.docx", file_type="docx", status="running",
        started_at=datetime(2026, 7, 1, 12, 0, 0),
    ))
    # The final (stage 6) step row the doctor report gets attached to.
    db.add(Step(
        run_id=run_id, step_number=12, step_name="WCM Template Population",
        status="complete", output_files=json.dumps(["/x/cv_wcm.docx"]),
    ))
    db.commit()

    monkeypatch.setattr(orch, "event_emitter", AsyncMock())
    monkeypatch.setattr(
        orch, "STEP_REGISTRY",
        [SimpleNamespace(number=1, stage_id="1a", name="Hierarchy Extraction")],
    )

    o = orch.PipelineOrchestrator(run_id, tmp_path / f"{run_id}.docx", db)
    monkeypatch.setattr(o, "_copy_to_pipeline_input", lambda: str(tmp_path / "cv.docx"))
    monkeypatch.setattr(o, "execute_step", AsyncMock())
    monkeypatch.setattr(o, "pipeline_output_dir", tmp_path / "outputs")
    return o


def _capture_posts(monkeypatch):
    """Record every Teams webhook POST payload (started + terminal cards)."""
    posts = []

    def _ok_post(url, json=None, timeout=None):
        posts.append(json)
        return SimpleNamespace(status_code=200)

    monkeypatch.setattr(notifications._SESSION, "post", _ok_post)
    return posts


def _facts(payload):
    """{name: value} from the card's FactSet (as in test_notifications)."""
    card = payload["attachments"][0]["content"]
    factset = next(b for b in card["body"] if b["type"] == "FactSet")
    return {f["title"]: f["value"] for f in factset["facts"]}


def test_knob_off_skips_doctor(monkeypatch, tmp_path, db):
    """CVICHE_RUN_DOCTOR=0: run_doctor is never called and no report is written."""
    from app.models import Run

    monkeypatch.setenv("CVICHE_RUN_DOCTOR", "0")
    monkeypatch.delenv("CVICHE_TEAMS_WEBHOOK_URL", raising=False)

    called = []
    monkeypatch.setattr(
        run_doctor_mod, "run_doctor",
        lambda *a, **k: called.append(a) or _doctor_payload(),
    )

    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_OFF")
    asyncio.run(o.execute())

    db.expire_all()
    assert db.query(Run).filter(Run.id == "DOC_OFF").first().status == "complete"
    assert called == []
    assert not (tmp_path / "outputs" / "stage_7_doctor").exists()


def test_doctor_runs_by_default_and_publishes(monkeypatch, tmp_path, db):
    """Knob unset: the report is written where the stage-JSON viewer serves
    pipeline outputs, attached to the final step's output_files, and the
    terminal Teams card carries the Doctor summary line."""
    from app.models import Run, Step

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    monkeypatch.setenv("CVICHE_TEAMS_WEBHOOK_URL", "https://webhook.example/teams")
    monkeypatch.delenv("CVICHE_ALLOWED_ORIGINS", raising=False)

    payload = _doctor_payload()
    kwargs = []
    monkeypatch.setattr(run_doctor_mod, "run_doctor", lambda *a, **k: kwargs.append(k) or payload)
    posts = _capture_posts(monkeypatch)

    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_ON")
    asyncio.run(o.execute())

    # #1174: the doctor reads this run's prompt logs for fallback-served calls.
    from app.pipeline.orchestrator import PROMPT_LOGS_DIR
    assert kwargs[-1]["prompt_log_dir"] == PROMPT_LOGS_DIR / "DOC_ON"

    db.expire_all()
    assert db.query(Run).filter(Run.id == "DOC_ON").first().status == "complete"

    # Report written under the pipeline outputs dir (document_uid = run id).
    report = tmp_path / "outputs" / "stage_7_doctor" / "DOC_ON_doctor.json"
    assert json.loads(report.read_text()) == payload

    # Registered on the final step's output_files (what the viewer lists),
    # without clobbering the existing stage-6 output.
    step = db.query(Step).filter(Step.run_id == "DOC_ON").first()
    files = json.loads(step.output_files)
    assert files == ["/x/cv_wcm.docx", str(report)]

    # Two cards posted (started + terminal); the terminal one has the line.
    notifications.flush()
    assert len(posts) == 2
    assert _facts(posts[-1])["Doctor"] == "1 findings (top: segmentation)"


def _stage6_docx(tmp_path, uid, content=None):
    from docx import Document

    path = tmp_path / "outputs" / "stage_6_wcm_documents" / f"{uid}_wcm.docx"
    path.parent.mkdir(parents=True)
    if content is None:
        doc = Document()
        doc.add_paragraph("Example Medical College")
        doc.save(str(path))
    else:
        path.write_bytes(content)
    return path


def test_doctor_publishes_the_review_copy_beside_its_report(monkeypatch, tmp_path, db):
    """#1388: the doctor's findings as Word comments, attached after the
    report and mirrored to storage with it; the clean document is untouched."""
    from docx import Document

    from app.models import Step

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    payload = _doctor_payload()
    payload["findings"] = [{"lint": "output_hygiene", "severity": "WARN", "status": "ran",
                            "message": "1 boilerplate line(s)", "evidence": ["Insert dates here (MM/YYYY)"]}]
    monkeypatch.setattr(run_doctor_mod, "run_doctor", lambda *a, **k: payload)
    clean = _stage6_docx(tmp_path, "DOC_RV")
    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_RV")
    persisted = []
    monkeypatch.setattr(o, "_persist_outputs_to_storage", persisted.extend)

    asyncio.run(o._run_doctor())

    report = tmp_path / "outputs" / "stage_7_doctor" / "DOC_RV_doctor.json"
    review = clean.with_name("DOC_RV_wcm_review.docx")
    files = json.loads(db.query(Step).filter(Step.run_id == "DOC_RV").first().output_files)
    assert files == ["/x/cv_wcm.docx", str(report), str(review)]
    assert persisted == [str(report), str(review)]
    # The quoted text is not in the document, so it is a review note in the box closing the copy.
    notes = [p.text for p in Document(str(review)).tables[-1].cell(0, 0).paragraphs]
    assert notes[1:] == ["Stray text to delete (1)", "Stray text: delete it.", '\u2022\t"Insert dates here (MM/YYYY)"']
    assert len(list(Document(str(clean)).comments)) == 0


def test_the_review_copy_reads_the_runs_stage4_artifact_for_the_owner_role_fix(monkeypatch, tmp_path, db):
    """#1591: the orchestrator hands the run's stage-4 artifact to the review
    copy, so a grant table naming the owner as PI gets "PI" as a tracked
    insertion in its empty role cell."""
    from docx import Document
    from docx.oxml.ns import qn

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    payload = {**_doctor_payload(), "findings": [{
        "lint": "role_consistency", "severity": "WARN", "status": "ran", "evidence": ["Example Project"],
        "message": "entry 5: 'Name of Principal Investigator:' names the CV owner, and 'Your role:' "
                   "is empty (owner_pi_role_empty, #1403)"}]}
    monkeypatch.setattr(run_doctor_mod, "run_doctor", lambda *a, **k: payload)
    clean = _stage6_docx(tmp_path, "DOC_PI")
    doc = Document(str(clean))
    table = doc.add_table(rows=3, cols=2)
    for row, (label, value) in zip(table.rows, [("Project title:", "Example Project"),
                                                ("Name of Principal Investigator:", "Ada Testowner"),
                                                ("Your role:", "")], strict=True):
        row.cells[0].text, row.cells[1].text = label, value
    doc.save(str(clean))
    stage4 = tmp_path / "outputs" / "stage_4_field_extraction" / "DOC_PI_fields.json"
    stage4.parent.mkdir(parents=True)
    stage4.write_text(json.dumps({"cv_owner": {"first_name": "Ada", "last_name": "Testowner"}, "entries": [
        {"taxonomy_code": "M2B", "element_idx_start": 5, "text": "A grant",
         "extracted_fields": {"title": "Example Project", "pi_name": "Testowner"}}]}))
    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_PI")
    monkeypatch.setattr(o, "_persist_outputs_to_storage", lambda files: None)

    asyncio.run(o._run_doctor())

    review = Document(str(clean.with_name("DOC_PI_wcm_review.docx")))
    role_cell = review.tables[0].rows[2].cells[1]._tc
    assert ["".join(t.text for t in ins.iter(qn("w:t"))) for ins in role_cell.iter(qn("w:ins"))] == ["PI"]


def test_a_corrupt_stage4_artifact_leaves_the_review_copy_comments_only(monkeypatch, tmp_path, db, caplog):
    """#1591: a stage-4 JSON that does not parse costs only the tracked fix;
    the review copy is still written, with the finding as a comment."""
    from docx import Document

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    payload = {**_doctor_payload(), "findings": [{
        "lint": "role_consistency", "severity": "WARN", "status": "ran", "evidence": ["Example Project"],
        "message": "entry 5: 'Name of Principal Investigator:' names the CV owner, and 'Your role:' "
                   "is empty (owner_pi_role_empty, #1403)"}]}
    monkeypatch.setattr(run_doctor_mod, "run_doctor", lambda *a, **k: payload)
    clean = _stage6_docx(tmp_path, "DOC_BAD4")
    doc = Document(str(clean))
    doc.add_paragraph("Example Project")
    doc.save(str(clean))
    stage4 = tmp_path / "outputs" / "stage_4_field_extraction" / "DOC_BAD4_fields.json"
    stage4.parent.mkdir(parents=True)
    stage4.write_text('{"cv_owner": ')
    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_BAD4")
    monkeypatch.setattr(o, "_persist_outputs_to_storage", lambda files: None)

    with caplog.at_level(logging.WARNING):
        asyncio.run(o._run_doctor())

    review = Document(str(clean.with_name("DOC_BAD4_wcm_review.docx")))
    assert len(list(review.comments)) == 1
    assert any("Stage-4 JSON unreadable" in r.message for r in caplog.records)
    assert not any("Review-comment docx failed" in r.message for r in caplog.records)


def test_an_unreadable_document_still_publishes_the_report(monkeypatch, tmp_path, db, caplog):
    from app.models import Step

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    monkeypatch.setattr(run_doctor_mod, "run_doctor", lambda *a, **k: _doctor_payload())
    _stage6_docx(tmp_path, "DOC_BAD", content=b"not a zip")
    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_BAD")
    monkeypatch.setattr(o, "_persist_outputs_to_storage", lambda files: None)

    with caplog.at_level(logging.WARNING):
        asyncio.run(o._run_doctor())

    report = tmp_path / "outputs" / "stage_7_doctor" / "DOC_BAD_doctor.json"
    files = json.loads(db.query(Step).filter(Step.run_id == "DOC_BAD").first().output_files)
    assert files == ["/x/cv_wcm.docx", str(report)]
    assert any("Review-comment docx failed" in r.message for r in caplog.records)


def test_doctor_crash_never_fails_run(monkeypatch, tmp_path, db, caplog):
    """A doctor exception is logged and swallowed: the run stays complete and
    the terminal notification still goes out, just without a Doctor fact."""
    from app.models import Run

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    monkeypatch.setenv("CVICHE_TEAMS_WEBHOOK_URL", "https://webhook.example/teams")
    monkeypatch.delenv("CVICHE_ALLOWED_ORIGINS", raising=False)

    def _boom(*args, **kwargs):
        raise RuntimeError("doctor exploded")

    monkeypatch.setattr(run_doctor_mod, "run_doctor", _boom)
    posts = _capture_posts(monkeypatch)

    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_BOOM")
    with caplog.at_level(logging.WARNING):
        asyncio.run(o.execute())

    db.expire_all()
    assert db.query(Run).filter(Run.id == "DOC_BOOM").first().status == "complete"
    assert any("Run doctor failed" in r.message for r in caplog.records)
    notifications.flush()
    assert len(posts) == 2
    assert "Doctor" not in _facts(posts[-1])


def test_doctor_source_is_the_runs_private_input_copy(monkeypatch, tmp_path, db):
    """#299: the doctor reads the same per-run copy _copy_to_pipeline_input wrote."""
    from app.pipeline import orchestrator as orch

    monkeypatch.setattr(orch, "PARENT_DIR", tmp_path / "repo")
    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_SRC")
    copied = o._pipeline_input_path()
    copied.parent.mkdir(parents=True)
    copied.write_bytes(b"x")

    seen = {}
    monkeypatch.setattr(
        run_doctor_mod, "run_doctor",
        lambda out_dir, uid, source=None, prompt_log_dir=None: seen.update(source=source) or _doctor_payload(),
    )
    o._doctor_report()

    assert seen["source"] == copied


def test_doctor_report_records_the_executing_image_tag(monkeypatch, tmp_path, db):
    """#1239: the doctor json names the image that ran its lints."""
    from app.pipeline import orchestrator as orch

    monkeypatch.setattr(orch, "PARENT_DIR", tmp_path / "repo")
    monkeypatch.setenv("CVICHE_IMAGE_TAG", "dev-7.tag")
    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_TAG")
    monkeypatch.setattr(
        run_doctor_mod, "run_doctor",
        lambda out_dir, uid, source=None, prompt_log_dir=None: _doctor_payload(),
    )

    payload, out_path = o._doctor_report()

    assert payload["image_tag"] == "dev-7.tag"
    assert json.loads(out_path.read_text())["image_tag"] == "dev-7.tag"


def test_completed_run_gets_its_quality_columns(monkeypatch, tmp_path, db):
    """The run-end score is copied onto runs.quality_* on the orchestrator's thread."""
    from app.models import Run
    from app.services import quality_score_service

    monkeypatch.setenv("CVICHE_RUN_DOCTOR", "0")
    monkeypatch.delenv("CVICHE_TEAMS_WEBHOOK_URL", raising=False)
    score = {"totalScore": 62, "raw_score_before_caps": 62.0, "hard_fail_caps_applied": []}
    monkeypatch.setattr(quality_score_service, "compute_and_cache_score", lambda _rid: score)

    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_SCORE")
    asyncio.run(o.execute())

    db.expire_all()
    run = db.query(Run).filter(Run.id == "DOC_SCORE").first()
    assert (run.quality_score, run.quality_band, run.quality_cap) == (62, "YELLOW", None)


class _ScoreStorage:
    """Durable storage holding one scorable output; records what is put."""

    def __init__(self):
        self.files = {"outputs/DOC_fields.json": b"{}"}

    def list_files(self, run_id, prefix=""):
        return [k for k in self.files if k.startswith(prefix)]

    def get_file(self, run_id, key):
        return self.files[key]

    def put_file(self, run_id, key, data):
        self.files[key] = data


def _score_against(monkeypatch, storage):
    """Real compute_and_cache_score over ``storage``, with the scorer stubbed."""
    from app.services import quality_score_service

    monkeypatch.setattr(quality_score_service, "get_storage", lambda: storage)
    monkeypatch.setattr("unified_pipeline.quality_score.score_run",
                        lambda _d, _r: {"totalScore": 97, "data_complete": True})


def test_a_run_whose_doctor_raises_is_scored_not_checked(monkeypatch, tmp_path, db):
    """#1593 (IXJMKS): a run the doctor never checked must not read as a clean
    GREEN -- its cached score says not checked and incomplete, and so does the card."""
    from app.services import quality_score_service as svc

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    monkeypatch.setenv("CVICHE_TEAMS_WEBHOOK_URL", "https://webhook.example/teams")
    monkeypatch.delenv("CVICHE_ALLOWED_ORIGINS", raising=False)
    storage = _ScoreStorage()
    _score_against(monkeypatch, storage)
    posts = _capture_posts(monkeypatch)

    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_UNCK")
    monkeypatch.setattr(o, "_run_doctor", AsyncMock(side_effect=RuntimeError("doctor exploded")))
    asyncio.run(o.execute())

    cached = json.loads(storage.files[svc.CACHE_KEY])
    assert cached[svc.DOCTOR_STATUS_KEY] == svc.DOCTOR_NOT_CHECKED
    assert cached["data_complete"] is False
    notifications.flush()
    assert "not checked" in _facts(posts[-1])["Quality score"]


def test_the_score_sees_the_report_the_doctor_just_stored(monkeypatch, tmp_path, db):
    """#1593: the doctor runs before the score, so a run it checked is scored checked."""
    from app.services import quality_score_service as svc

    monkeypatch.delenv("CVICHE_RUN_DOCTOR", raising=False)
    monkeypatch.delenv("CVICHE_TEAMS_WEBHOOK_URL", raising=False)
    storage = _ScoreStorage()
    _score_against(monkeypatch, storage)

    async def _doctor_stores_its_report():
        storage.files[f"outputs/DOC{svc.DOCTOR_SUFFIX}"] = b"{}"
        return _doctor_payload()

    o = _orchestrator(monkeypatch, tmp_path, db, "DOC_CHKD")
    monkeypatch.setattr(o, "_run_doctor", _doctor_stores_its_report)
    asyncio.run(o.execute())

    cached = json.loads(storage.files[svc.CACHE_KEY])
    assert cached[svc.DOCTOR_STATUS_KEY] == svc.DOCTOR_CHECKED
    assert cached["data_complete"] is True
