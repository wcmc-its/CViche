"""POST /api/run/{run_id}/feedback also fires a best-effort Teams notification.

This covers the call-site wiring in feedback_routes.submit_feedback -- notify
is actually called with the right feedback/run/submitter, and a failure there
never affects the already-committed feedback or the 201 response. The
notifications.build_feedback_payload / notify_feedback_submitted helpers
themselves are covered in test_notifications.py; a helper test alone doesn't
prove the route actually calls them with the right arguments.
"""
import json
import os
from contextlib import contextmanager
from datetime import datetime
from unittest.mock import Mock

import pytest

os.environ.setdefault("CVICHE_SESSION_SECRET", "test-secret-not-for-production")

_RUN_ID = "_FDBK1"

_VALID_BODY = {
    "reviewer_role": "self",
    "overall_usefulness": 4,
    "manual_conversion_effort": "1-2 hours",
    "correction_effort": "1-2 hours",
    "likelihood_to_recommend": 5,
}


def _seed_run(db, email="reviewer@example.com"):
    """Insert a user + one of their runs, returning the committed user."""
    from app.models import Run, User

    user = User(email=email, display_name="Reviewer", role="user")
    db.add(user)
    db.flush()  # assign user.id without ending the transaction

    db.add(Run(
        id=_RUN_ID, filename="cv.docx", file_type="docx", status="complete",
        user_id=user.id, started_at=datetime(2026, 6, 16, 12, 0, 0),
        completed_at=datetime(2026, 6, 16, 12, 1, 0),
    ))
    db.commit()
    db.refresh(user)
    return user


@contextmanager
def _as_user(user):
    from app.auth import get_current_user
    from app.main import app

    app.dependency_overrides[get_current_user] = lambda: user
    try:
        yield
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_submit_feedback_notifies_teams_with_run_and_submitter(client, db, monkeypatch):
    user = _seed_run(db)
    notify = Mock()
    monkeypatch.setattr("app.services.notifications.notify_feedback_submitted", notify)

    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback", json=_VALID_BODY)

    assert resp.status_code == 201
    notify.assert_called_once()
    feedback_arg, run_arg, submitter_arg = notify.call_args[0]
    assert feedback_arg.run_id == _RUN_ID
    assert run_arg.id == _RUN_ID
    assert submitter_arg == "Reviewer"


def test_submit_feedback_succeeds_even_if_notification_raises(client, db, monkeypatch):
    """The Teams POST is best-effort: a failure there must not roll back the
    feedback or turn the 201 into a 500."""
    user = _seed_run(db)
    monkeypatch.setattr(
        "app.services.notifications.notify_feedback_submitted",
        Mock(side_effect=RuntimeError("webhook down")),
    )

    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback", json=_VALID_BODY)

    assert resp.status_code == 201

    from app.models import Feedback
    assert db.query(Feedback).filter(Feedback.run_id == _RUN_ID).count() == 1


def _staff(db):
    from app.models import User

    staff = User(email="staff@example.com", display_name="Staff User", role="staff")
    db.add(staff)
    db.commit()
    db.refresh(staff)
    return staff


def test_staff_can_read_but_not_submit_feedback_on_another_users_run(client, db, monkeypatch):
    """Staff is read-only: GET passes check_run_access(read_only=True), POST
    stays owner-or-admin and 403s with nothing written or notified."""
    _seed_run(db)
    staff = _staff(db)
    notify = Mock()
    monkeypatch.setattr("app.services.notifications.notify_feedback_submitted", notify)

    with _as_user(staff):
        read = client.get(f"/api/run/{_RUN_ID}/feedback")
        read_all = client.get(f"/api/run/{_RUN_ID}/feedback/all")
        write = client.post(f"/api/run/{_RUN_ID}/feedback", json=_VALID_BODY)

    assert read.status_code == 200
    assert read_all.status_code == 200
    assert write.status_code == 403
    notify.assert_not_called()
    from app.models import Feedback
    assert db.query(Feedback).filter(Feedback.run_id == _RUN_ID).count() == 0


def _add_feedback(db, user, at, **overrides):
    from app.models import Feedback

    fields = dict(
        run_id=_RUN_ID, user_id=user.id, reviewer_role="self", overall_usefulness=4,
        manual_conversion_effort="1-2 hours", correction_effort="1-2 hours",
        likelihood_to_recommend=5, submitted_at=at, issue_locations='["A", "M2"]',
        biggest_issue="Jane Testperson grants are missing",
    )
    fields.update(overrides)
    db.add(Feedback(**fields))
    db.commit()


def _other_user(db, email, role="user"):
    from app.models import User

    user = User(email=email, display_name="Other Reviewer", role=role)
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def test_feedback_all_lists_every_reviewer_newest_first(client, db):
    owner = _seed_run(db)
    other = _other_user(db, "other@example.com")
    _add_feedback(db, owner, datetime(2026, 9, 1))
    _add_feedback(db, other, datetime(2026, 9, 2), reviewer_role="colleague")

    with _as_user(owner):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/all")

    assert resp.status_code == 200
    rows = resp.json()
    assert [(r["display_name"], r["reviewer_role"]) for r in rows] == [
        ("Other Reviewer", "colleague"), ("Reviewer", "self")]
    submitted = datetime.fromisoformat(rows[0]["submitted_at"])
    assert submitted.tzinfo is not None  # carries a UTC offset (schemas.py rule)
    assert submitted.replace(tzinfo=None) == datetime(2026, 9, 2)
    assert rows[0]["issue_locations"] == ["A", "M2"]
    assert rows[0]["biggest_issue"] == "Jane Testperson grants are missing"
    assert rows[0]["overall_usefulness"] == 4
    assert rows[0]["summary_generated"] is None


def test_feedback_all_is_open_to_an_admin_and_closed_to_a_stranger(client, db):
    owner = _seed_run(db)
    _add_feedback(db, owner, datetime(2026, 9, 1))
    admin = _other_user(db, "root@example.com", role="admin")
    stranger = _other_user(db, "stranger@example.com")

    with _as_user(admin):
        assert client.get(f"/api/run/{_RUN_ID}/feedback/all").status_code == 200
    with _as_user(stranger):
        assert client.get(f"/api/run/{_RUN_ID}/feedback/all").status_code == 403
    with _as_user(admin):
        assert client.get("/api/run/NOSUCH/feedback/all").status_code == 404


def test_feedback_all_matches_the_single_feedback_serialisation(client, db):
    """One per-field serialisation: /feedback/all adds only display_name."""
    owner = _seed_run(db)
    _add_feedback(db, owner, datetime(2026, 9, 1))

    with _as_user(owner):
        single = client.get(f"/api/run/{_RUN_ID}/feedback").json()["feedback"]
        listed = client.get(f"/api/run/{_RUN_ID}/feedback/all").json()[0]

    listed.pop("display_name")
    assert listed == single


# --- section labels come from core/taxonomy_v7.json (#1782) --------------------


def test_section_labels_are_exactly_the_taxonomy_codes_and_labels():
    """Every code the feedback form can name, with its label, is taxonomy_v7's own."""
    from app.api.feedback_routes import _taxonomy_section_labels
    from unified_pipeline.stage3b.io import load_taxonomy

    taxonomy = {c["code"]: c["label"] for c in load_taxonomy()["codes"]}
    assert dict(_taxonomy_section_labels()) == taxonomy
    # The retired scheme the hand-kept map carried (S3 'Letters to Editor', S10-S15, T1-T7).
    assert _taxonomy_section_labels()["S3"] == "Books"
    assert {"S10", "S13", "S15", "T1", "T7"}.isdisjoint(_taxonomy_section_labels())


def _seed_populated_outputs(root, run_id, *section_codes):
    """A stage-4 metadata file plus one stage-2d enriched file per section code."""
    stage_4 = root / run_id / "stage_4_wcm_templates"
    stage_4.mkdir(parents=True)
    (stage_4 / "x_template_metadata.json").write_text(
        json.dumps({"records_processed": {"publications": 2}}), encoding="utf-8")
    enriched = root / run_id / "stage_2d_enriched"
    enriched.mkdir()
    for code in section_codes:
        (enriched / f"section_{code}_enriched.json").write_text("{}", encoding="utf-8")


def test_get_feedback_names_populated_sections_by_their_taxonomy_label(client, db, tmp_path, monkeypatch):
    from app.models import Step

    owner = _seed_run(db)
    db.add(Step(run_id=_RUN_ID, step_number=12, stage_id="6", step_name="WCM", status="complete"))
    db.commit()
    monkeypatch.setattr("app.api.feedback_routes.OUTPUTS_ROOT", tmp_path)
    _seed_populated_outputs(tmp_path, _RUN_ID, "S3", "S5", "S13")

    with _as_user(owner):
        sections = client.get(f"/api/run/{_RUN_ID}/feedback").json()["run_context"]["wcm_sections"]

    named = {s["section_id"]: s["section_name"] for s in sections if "section_id" in s}
    # S13 is not a taxonomy_v7 code, so it is not offered at all.
    assert named == {"S3": "Books", "S5": "Non-peer-reviewed Research Publications"}


# --- the review loop (#1587): the corrected copy -------------------------------


def _docx(*paragraphs):
    import io

    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Synthetic entries: nothing here is from a real CV.
_DELIVERED = ("Fellowship in Quorvane studies, Northfield Institute, 2011-2013",
              "Brennic Foundation Award for Saltwick lattice modelling, 2015",
              "Lecturer in Heliotropic Drift, Plesk College, 2016-2019")


def _copy_key(user, suffix):
    """Where the user's corrected copy of the run is stored: one per uploader (#1754)."""
    return f"corrected/{_RUN_ID}_u{user.id}_{suffix}"


@pytest.fixture
def run_storage(tmp_path, monkeypatch):
    from app.storage.local_storage import LocalRunStorage

    store = LocalRunStorage(str(tmp_path / "store"))
    monkeypatch.setattr("app.services.review_loop_service.get_storage", lambda: store)
    monkeypatch.setattr("app.api.feedback_routes.get_storage", lambda: store)
    return store


def test_corrected_docx_upload_stores_the_copy_and_its_diff_and_returns_one_line(client, db, run_storage):
    user = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    corrected = _docx(_DELIVERED[0], _DELIVERED[2])  # the reviewer deleted the award
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", corrected)})
    assert resp.status_code == 200
    assert resp.json() == {"changes": 1, "comments": 0, "summary": "1 text edit and 0 comments recorded"}

    import json
    assert run_storage.get_file(_RUN_ID, _copy_key(user, "corrected.docx")) == corrected
    report = json.loads(run_storage.get_file(_RUN_ID, _copy_key(user, "diff.json")))
    assert report["by_type"]["deleted"] == 1
    assert "Brennic" not in json.dumps(report)  # positions and counts only, no CV text
    # Never under outputs/: downloads and the scorer read only that prefix.
    assert run_storage.list_files(_RUN_ID, "outputs/") == [f"outputs/{_RUN_ID}_wcm.docx"]
    # #1654: the verdicts (this run has no review copy, and says so) and the doctor's re-run.
    verdicts = json.loads(run_storage.get_file(_RUN_ID, _copy_key(user, "verdicts.json")))
    assert verdicts["note"] and verdicts["findings"] == []
    rerun = json.loads(run_storage.get_file(_RUN_ID, _copy_key(user, "doctor.json")))
    assert rerun["document_uid"] == _RUN_ID and rerun["artifacts"]["stage_6_docx"]


def test_corrected_docx_upload_confirms_before_the_review_pass_runs(client, db, run_storage, monkeypatch):
    """#1654: the verdicts and the doctor's re-run are a background task, after the response."""
    from starlette.background import BackgroundTasks

    queued = []
    monkeypatch.setattr(BackgroundTasks, "add_task", lambda self, func, *args: queued.append((func, args)))
    user = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", _docx(_DELIVERED[0]))})
    assert resp.status_code == 200
    assert run_storage.list_files(_RUN_ID, "corrected/") == [
        _copy_key(user, "corrected.docx"), _copy_key(user, "diff.json")]
    (func, args), = queued
    func(*args)
    assert _copy_key(user, "verdicts.json") in run_storage.list_files(_RUN_ID, "corrected/")


def test_corrected_docx_upload_confirms_though_the_review_pass_fails(client, db, run_storage, monkeypatch):
    """#1654: the copy and its diff are stored; a failing verdict or re-run step is logged only."""
    def boom(*_args):
        raise ValueError("synthetic")

    monkeypatch.setattr("app.services.review_loop_service.store_verdicts", boom)
    monkeypatch.setattr("app.services.review_loop_service.store_corrected_doctor", boom)
    user = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", _docx(_DELIVERED[0]))})
    assert resp.status_code == 200
    assert run_storage.list_files(_RUN_ID, "corrected/") == [
        _copy_key(user, "corrected.docx"), _copy_key(user, "diff.json")]


@pytest.mark.parametrize("name, message", [
    ("corrected.pdf", "Word (.docx) file"),  # a valid Word file, but not named .docx
    ("corrected.docx", "does not match .docx format"),  # over the zip-bomb entry bound (patched to 1)
])
def test_corrected_docx_upload_refuses_what_the_cv_upload_refuses(
        client, db, run_storage, monkeypatch, name, message):
    if name.endswith(".docx"):
        monkeypatch.setattr("app.services.upload_validation._DOCX_MAX_ENTRIES", 1)
    user = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": (name, _docx(_DELIVERED[0]))})
    assert resp.status_code == 400
    assert message in resp.json()["detail"]["message"]
    assert run_storage.list_files(_RUN_ID, "corrected/") == []


def test_corrected_docx_upload_is_bounded_by_the_upload_size_cap(client, db, run_storage, monkeypatch):
    corrected = _docx(_DELIVERED[0])
    monkeypatch.setattr("app.api.feedback_routes.MAX_UPLOAD_SIZE", len(corrected) - 1)  # one byte over
    user = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", corrected)})
    assert resp.status_code == 400
    assert "too large" in resp.json()["detail"]["message"]
    assert run_storage.list_files(_RUN_ID, "corrected/") == []


def test_corrected_docx_upload_refuses_a_macro_carrying_docx(client, db, run_storage):
    import io
    import zipfile

    buf = io.BytesIO(_docx(*_DELIVERED))
    with zipfile.ZipFile(buf, "a") as zf:
        zf.writestr("word/vbaProject.bin", b"\x00")
    user = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", buf.getvalue())})
    assert resp.status_code == 400
    assert "macros" in resp.json()["detail"]["message"]
    assert run_storage.list_files(_RUN_ID, "corrected/") == []


def test_corrected_docx_upload_needs_a_delivered_document(client, db, run_storage):
    user = _seed_run(db)
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", _docx("x"))})
    assert resp.status_code == 409
    assert run_storage.list_files(_RUN_ID, "corrected/") == []


def test_corrected_docx_upload_is_open_to_staff_on_any_run(client, db, run_storage):
    """OFA staff may send a corrected copy of any run's CV, though they are
    read-only everywhere else (the survey stays closed to them -- see
    test_staff_can_read_but_not_submit_feedback_on_another_users_run)."""
    _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    corrected = _docx(_DELIVERED[0], _DELIVERED[2])
    staff = _staff(db)
    with _as_user(staff):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", corrected)})
    assert resp.status_code == 200
    assert resp.json() == {"changes": 1, "comments": 0, "summary": "1 text edit and 0 comments recorded"}
    assert run_storage.get_file(_RUN_ID, _copy_key(staff, "corrected.docx")) == corrected


def test_corrected_docx_upload_keeps_one_copy_per_uploader(client, db, run_storage):
    """#1754: the owner's and a staff member's copies of one run never
    overwrite each other; a second upload by the same user replaces theirs."""
    owner, staff = _seed_run(db), _staff(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    uploads = [(owner, _docx(_DELIVERED[0])), (staff, _docx(_DELIVERED[1])),
               (owner, _docx(_DELIVERED[0], _DELIVERED[2]))]
    for user, corrected in uploads:
        with _as_user(user):
            assert client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                               files={"file": ("corrected.docx", corrected)}).status_code == 200
    copies = [k for k in run_storage.list_files(_RUN_ID, "corrected/") if k.endswith("_corrected.docx")]
    assert sorted(copies) == sorted([_copy_key(owner, "corrected.docx"), _copy_key(staff, "corrected.docx")])
    assert run_storage.get_file(_RUN_ID, _copy_key(owner, "corrected.docx")) == uploads[2][1]
    assert run_storage.get_file(_RUN_ID, _copy_key(staff, "corrected.docx")) == uploads[1][1]


def test_corrected_docx_upload_is_closed_to_another_plain_user(client, db, run_storage):
    """The staff grant is for staff only: a plain user who does not own the
    run is still refused, with nothing stored."""
    _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    from app.models import User
    stranger = User(email="stranger@example.com", display_name="Stranger", role="user")
    db.add(stranger)
    db.commit()
    db.refresh(stranger)
    with _as_user(stranger):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", _docx(*_DELIVERED))})
    assert resp.status_code == 403
    assert run_storage.list_files(_RUN_ID, "corrected/") == []


# --- the corrected copies, listed and served back -----------------------------

_SCAN_FLAG = "CVICHE_REQUIRE_MALWARE_SCAN"


def _stranger(db):
    from app.models import User

    stranger = User(email="stranger@example.com", display_name="Stranger", role="user")
    db.add(stranger)
    db.commit()
    db.refresh(stranger)
    return stranger


def _upload(client, user, corrected):
    with _as_user(user):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", corrected)})
    assert resp.status_code == 200, resp.text


def _owner_and_staff_copies(client, db, run_storage):
    """The owner's and a staff member's copies of the run, staff's last."""
    owner, staff = _seed_run(db), _staff(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    _upload(client, owner, _docx(_DELIVERED[0]))
    _upload(client, staff, _docx(_DELIVERED[0], _DELIVERED[2]))
    return owner, staff


def _list(client, user):
    with _as_user(user):
        return client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies")


def test_the_run_owner_sees_staffs_corrected_copy_and_their_own(client, db, run_storage, monkeypatch):
    monkeypatch.delenv(_SCAN_FLAG, raising=False)
    owner, _ = _owner_and_staff_copies(client, db, run_storage)
    resp = _list(client, owner)
    assert resp.status_code == 200, resp.text
    rows = [(c["uploader"], c["role"], c["changes"], c["comments"], c["mine"], c["download"]) for c in resp.json()]
    assert rows == [("Staff User", "staff", 1, 0, False, "ready"), ("Reviewer", "owner", 2, 0, True, "ready")]


def test_a_copy_with_only_the_reviewers_comment_reports_and_lists_it(client, db, run_storage, monkeypatch):
    """#1937: comments and no text edit read "0 changes recorded" and looked lost."""
    import io

    from docx import Document

    monkeypatch.delenv(_SCAN_FLAG, raising=False)
    owner = _seed_run(db)
    run_storage.put_file(_RUN_ID, f"outputs/{_RUN_ID}_wcm.docx", _docx(*_DELIVERED))
    doc = Document(io.BytesIO(_docx(*_DELIVERED)))
    doc.add_comment(doc.paragraphs[0].runs, text="Synthetic note.", author="Synthetic Reviewer")
    corrected = io.BytesIO()
    doc.save(corrected)
    with _as_user(owner):
        resp = client.post(f"/api/run/{_RUN_ID}/feedback/corrected-docx",
                           files={"file": ("corrected.docx", corrected.getvalue())})
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"changes": 0, "comments": 1, "summary": "0 text edits and 1 comment recorded"}
    assert [(c["changes"], c["comments"]) for c in _list(client, owner).json()] == [(0, 1)]


def test_a_stranger_can_neither_list_nor_download_corrected_copies(client, db, run_storage):
    owner, _ = _owner_and_staff_copies(client, db, run_storage)
    copy_id = _list(client, owner).json()[0]["id"]
    stranger = _stranger(db)
    assert _list(client, stranger).status_code == 403
    with _as_user(stranger):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies/{copy_id}")
    assert resp.status_code == 403


def test_a_corrected_copy_downloads_under_a_quoted_name(client, db, run_storage, monkeypatch):
    """Local storage proxies the bytes; the name carries the upload's and the
    uploader's names, percent-encoded (RFC 5987) as the other downloads are."""
    from urllib.parse import quote

    monkeypatch.delenv(_SCAN_FLAG, raising=False)
    owner, staff = _owner_and_staff_copies(client, db, run_storage)
    staff_copy = _list(client, owner).json()[0]
    with _as_user(owner):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies/{staff_copy['id']}")
    assert resp.status_code == 200, resp.text
    assert resp.content == run_storage.get_file(_RUN_ID, _copy_key(staff, "corrected.docx"))
    assert resp.headers["content-disposition"] == (
        "attachment; filename*=utf-8''" + quote("cv - corrected by Staff User.docx", safe=""))


def test_a_corrected_copy_of_another_run_is_not_found(client, db, run_storage):
    owner, _ = _owner_and_staff_copies(client, db, run_storage)
    with _as_user(owner):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies/9999")
    assert resp.status_code == 404


def test_a_corrected_copy_still_being_scanned_is_listed_as_scanning_and_refused_409(
        client, db, run_storage, monkeypatch):
    """The same gate as the original CV: local storage carries no scan tag, so
    with the gate on every copy is still being scanned."""
    owner, _ = _owner_and_staff_copies(client, db, run_storage)
    monkeypatch.setenv(_SCAN_FLAG, "1")
    copies = _list(client, owner).json()
    assert [c["download"] for c in copies] == ["scanning", "scanning"]
    with _as_user(owner):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies/{copies[0]['id']}?as_url=true")
    assert resp.status_code == 409
    assert "still being scanned" in resp.json()["detail"]["message"]
    assert resp.headers["Retry-After"] == "60"


def test_a_corrected_copy_flagged_by_the_scan_is_never_served(client, db, run_storage, monkeypatch):
    owner, _ = _owner_and_staff_copies(client, db, run_storage)
    monkeypatch.setenv(_SCAN_FLAG, "1")
    monkeypatch.setattr(run_storage, "get_malware_scan_status", lambda run_id, key: "THREATS_FOUND")
    copies = _list(client, owner).json()
    assert [c["download"] for c in copies] == ["threats_found", "threats_found"]
    with _as_user(owner):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies/{copies[0]['id']}")
    assert resp.status_code == 403
    assert "flagged by the malware scan" in resp.json()["detail"]["message"]


def test_a_clean_corrected_copy_is_served_with_the_gate_on(client, db, run_storage, monkeypatch):
    owner, _ = _owner_and_staff_copies(client, db, run_storage)
    monkeypatch.setenv(_SCAN_FLAG, "1")
    monkeypatch.setattr(run_storage, "get_malware_scan_status", lambda run_id, key: "NO_THREATS_FOUND")
    copy_id = _list(client, owner).json()[0]["id"]
    with _as_user(owner):
        resp = client.get(f"/api/run/{_RUN_ID}/feedback/corrected-copies/{copy_id}?as_url=true")
    assert resp.status_code == 200
    assert resp.json() == {"url": None}  # local storage: the page then downloads from the route
    assert resp.headers["cache-control"] == "no-store"


def test_one_unreadable_scan_tag_marks_only_that_copy_unscannable(client, db, run_storage, monkeypatch):
    """A tag read that raises (object gone, no s3:GetObjectTagging) must not
    hide the other copies: the list still answers, that row is unscannable."""
    owner, staff = _owner_and_staff_copies(client, db, run_storage)
    monkeypatch.setenv(_SCAN_FLAG, "1")
    staff_key = _copy_key(staff, "corrected.docx")

    def tag(run_id, key):
        if key == staff_key:
            raise RuntimeError("AccessDenied")
        return "NO_THREATS_FOUND"

    monkeypatch.setattr(run_storage, "get_malware_scan_status", tag)
    resp = _list(client, owner)
    assert resp.status_code == 200, resp.text
    assert [c["download"] for c in resp.json()] == ["unscannable", "ready"]
