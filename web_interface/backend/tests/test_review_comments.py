"""The review copy: the run doctor's WARN/ERROR findings flagged in place (#1388 C).

Every document here is built in the test from invented text; no corpus CV.
"""
import copy
import json
import os
import sys
from pathlib import Path

os.environ.setdefault("CVICHE_SESSION_SECRET", "test-secret-not-for-production")
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent / "src"))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.services import review_comments as rc  # noqa: E402
from app.services.artifact_service import REVIEW_DOCX_SUFFIX  # noqa: E402
from app.services.run_quality_report import LINT_COPY  # noqa: E402
from unified_pipeline.doctor.blind_spots import blind_spots  # noqa: E402
from unified_pipeline.doctor.lints.completeness import lint_mentee_field_incomplete  # noqa: E402
from unified_pipeline.doctor.lints.extraction import (  # noqa: E402
    lint_grant_facts,
    lint_role_consistency,
)
from unified_pipeline.doctor.lints.render import lint_renderer_row_receipts  # noqa: E402
from unified_pipeline.doctor.precision import LintPrecision  # noqa: E402
from unified_pipeline.run_doctor import (  # noqa: E402
    read_docx_blocks,
    read_docx_table_rows,
)
from unified_pipeline.stage6.sections.mentoring import (  # noqa: E402
    MENTEE_CURRENT_POSITION_LABEL,
    MENTEE_TABLE_LABELS,
    PAST_MENTEES_HEADING,
)
from unified_pipeline.stage6.sections.research_support import (  # noqa: E402
    PI_NAME_LABEL,
    PROJECT_TITLE_LABEL,
    YOUR_ROLE_LABEL,
)

HEADER = "Example Medical College"
PAST_FUNDING_HEADING = "Past (Completed) Funding"
GRANT = "Zebrafish Fin Regrowth Study, Example Foundation, 1912-1914"
CITATION = "12. Quill AB, Marsh CD. Lanternfish vision under pressure. J Example Biol. 2019;4:1-9."
SECOND = "14. Quill AB, Reed EF. The squid's lenses and light. J Example Biol. 2021;6:3-8."
ENRICHED = "13. Quill AB. A tracked insertion only enrichment adds. J Example Biol. 2020;5:2-3."
TEACHING = "Tobacco Cessation Lecturer (PHRM 827) 4 hrs, 156 students"
#: source_line_coverage's group title in the box (#2002).
SLC_TITLE = "Lines from your CV that CViche couldn't place, often present in another form"
APPENDIX_LINES = ("1. Volunteer, Example Food Bank, 2015", "2. Marathon finisher, 2018")


def _clean_docx(tmp_path: Path, appendix: bool = True) -> Path:
    doc = Document()
    doc.add_paragraph(HEADER)
    doc.add_paragraph("RESEARCH")
    doc.add_paragraph(PAST_FUNDING_HEADING)
    doc.add_paragraph(GRANT)
    doc.add_paragraph("EDUCATIONAL CONTRIBUTIONS")
    doc.add_paragraph(TEACHING)
    doc.add_paragraph("BIBLIOGRAPHY")
    doc.add_paragraph(CITATION)
    doc.add_paragraph(CITATION)  # a repeated record
    doc.add_paragraph(SECOND)
    # Stage 5 enrichment writes a citation as a tracked insertion: no direct w:r.
    p = doc.add_paragraph()._p
    ins = p.makeelement(qn("w:ins"), {qn("w:id"): "90", qn("w:author"): "CViche"})
    run = doc.add_paragraph(ENRICHED).runs[0]._r
    ins.append(copy.deepcopy(run))
    p.append(ins)
    doc.element.body.remove(run.getparent())
    if appendix:
        doc.add_paragraph("T. APPENDIX")
        for line in APPENDIX_LINES:
            doc.add_paragraph(line)
    out = tmp_path / "DOC_wcm.docx"
    doc.save(str(out))
    return out


def _finding(lint, message, evidence=(), severity="WARN"):
    return {"lint": lint, "severity": severity, "message": message,
            "evidence": list(evidence), "status": "ran", "reason": ""}


def _report(*findings):
    return {"findings": list(findings)}


def _comments(path: Path, author: str = rc.COMMENT_AUTHOR) -> list[tuple[str, str]]:
    """(comment text, the exact text it is anchored to), in document order,
    of ``author``'s comments: the findings' by default, not the review notes."""
    doc = Document(str(path))
    texts = {c.comment_id: c.text for c in doc.comments if c.author == author}
    anchored: dict[int, str] = {}
    open_ids: list[int] = []
    for el in doc.element.body.iter():
        if el.tag == qn("w:commentRangeStart"):
            open_ids.append(int(el.get(qn("w:id"))))
            anchored.setdefault(open_ids[-1], "")
        elif el.tag == qn("w:commentRangeEnd"):
            open_ids.remove(int(el.get(qn("w:id"))))
        elif el.tag == qn("w:t"):
            for cid in open_ids:
                anchored[cid] += el.text or ""
    return [(texts[cid], body) for cid, body in anchored.items() if cid in texts]


def _notes_comment(path: Path):
    """The review-notes comment: the one comment by REVIEW_NOTES_AUTHOR."""
    (comment,) = [c for c in Document(str(path)).comments if c.author == rc.REVIEW_NOTES_AUTHOR]
    return comment


def _note_lines(path: Path) -> list[str]:
    """The review-notes comment's lines after its title."""
    lines = [p.text for p in _notes_comment(path).paragraphs]
    assert lines[0] == rc.REVIEW_NOTES_TITLE
    return lines[1:]


def _notes(path: Path) -> list[str]:
    """The review notes' lines under their title, up to the closing
    what-is-not-checked group."""
    lines = _note_lines(path)
    return lines[:lines.index(rc.NOT_CHECKED_TITLE)]


def _not_checked(path: Path) -> list[str]:
    """The review notes' closing group: what CViche does not check."""
    lines = _note_lines(path)
    return lines[lines.index(rc.NOT_CHECKED_TITLE):]


def _delete_all_comments(path: Path) -> Document:
    """What Word's Review > Delete All Comments does: every comment's range
    marks, its reference run and the comment itself go."""
    doc = Document(str(path))
    for el in list(doc.element.body.iter(qn("w:commentRangeStart"), qn("w:commentRangeEnd"),
                                         qn("w:commentReference"))):
        gone = el.getparent() if el.tag == qn("w:commentReference") else el  # the reference's run
        gone.getparent().remove(gone)
    for comment in list(doc.comments):
        comment._comment_elm.getparent().remove(comment._comment_elm)
    return doc


def _body_text(doc: Document) -> list[str]:
    """Every body paragraph's text in document order, table cells included,
    with tracked insertions (w:ins) read: what Paragraph.text skips."""
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


@pytest.fixture(autouse=True)
def _unmeasured_ledger(monkeypatch):
    """Every lint unmeasured unless a test passes its own ledger rows, so these
    tests read placement, not PRECISION.md's current figures."""
    monkeypatch.setattr(rc, "load_gate_ledger", dict)


def _with_appendix_note(tmp_path: Path) -> Path:
    """The clean document with stage 6's Appendix note box under T. APPENDIX."""
    doc = Document(str(_clean_docx(tmp_path)))
    heading = next(p for p in doc.paragraphs if p.text == "T. APPENDIX")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "CViche note: delete this box before sending"
    table.cell(0, 0).add_paragraph("These 5 entries from your original CV did not fit any section above.")
    heading._p.addnext(table._tbl)
    doc.save(str(tmp_path / "DOC_wcm.docx"))
    return tmp_path / "DOC_wcm.docx"


def _flag(lint):
    return rc.REVIEW_FLAGS[lint]


def test_every_lint_the_run_page_words_has_a_flag():
    assert set(LINT_COPY) <= set(rc.REVIEW_FLAGS)


def test_a_quoted_record_gets_its_short_flag(tmp_path):
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("enrichment_failures", "1 publication(s) failed PubMed enrichment", [SECOND])))
    assert out == tmp_path / f"DOC{REVIEW_DOCX_SUFFIX}" and n == 1
    assert _comments(out) == [(_flag("enrichment_failures"), SECOND)]
    assert [c.author for c in Document(str(out)).comments] == [rc.COMMENT_AUTHOR, rc.REVIEW_NOTES_AUTHOR]


def test_a_duplicate_is_flagged_on_its_later_copy(tmp_path):
    """The note names the first copy; the second is the one to delete."""
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("duplicate_records", "1 duplicated record(s)", [f"block 7 repeats at 8: {CITATION[:90]}"])))
    doc = Document(str(out))
    starts = [i for i, p in enumerate(doc.element.body.iter(qn("w:p")))
              if p.find(qn("w:commentRangeStart")) is not None]
    # paragraphs: 7 is the first CITATION, 8 the repeat; 0 carries the review notes
    assert starts == [0, 8]
    assert _comments(out) == [(_flag("duplicate_records"), CITATION)]


def test_each_record_a_finding_lists_gets_its_own_flag(tmp_path):
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("pipe_leaks", "2 numbered citation(s) fusing multiple venue-date patterns",
                 [f"[bibliography] {CITATION[:100]}", f"[bibliography] {SECOND}"])))
    assert n == 2
    assert [anchor for _, anchor in _comments(out)] == [CITATION, SECOND]


def test_a_wrong_year_is_flagged_on_the_year_alone(tmp_path):
    clean = _clean_docx(tmp_path)
    out, _ = rc.write_review_docx(clean, _report(
        _finding("implausible_year", "entry 3 (M2B): end_date=1912 -- before 1959",
                 [f"Grant Title: {GRANT}; Role: PI"])))
    assert _comments(out) == [(_flag("implausible_year"), "1912")]
    # Splitting the run to isolate the year leaves the text the doctor reads alone.
    assert read_docx_blocks(str(out)) == read_docx_blocks(str(clean))


def test_near_duplicates_are_one_group_naming_section_removed_and_kept(tmp_path):
    dropped = TEACHING.replace("4 hrs", "5 hrs")
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("dedup_drops", "2 dedup drop(s) poorly covered by the kept entry", [
            f"K1 (jaccard=1.00, 89% covered by kept): dropped '{dropped}' vs kept '{TEACHING}'",
            "S5 (jaccard=0.94, 88% covered by kept): dropped '(9) Quill AB. Squid optics.' vs kept '(11) Quill AB. Squid optics.'"])))
    assert n == 2 and _comments(out) == []
    assert _notes(out) == [
        "Removed as near-duplicates (2)", rc.DEDUP_INSTRUCTION,
        "\u2022 Didactic Teaching", f'Removed: "{dropped}"', f'Kept: "{TEACHING}"',
        "\u2022 Non-peer-reviewed Publications", 'Removed: "(9) Quill AB. Squid optics."',
        'Kept: "(11) Quill AB. Squid optics."',
    ]
    assert "jaccard" not in "".join(_notes(out))


def test_a_curly_apostrophe_in_the_quote_matches_a_straight_one_in_the_output(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("enrichment_failures", "1 failed", ["Quill AB, Reed EF. The squid\u2019s lenses"])))
    assert _comments(out) == [(_flag("enrichment_failures"), SECOND)]


def test_a_tracked_insertion_is_found_and_flagged(tmp_path):
    """Paragraph.text skips w:ins runs; the anchor search must not."""
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("enrichment_failures", "1 publication(s) failed PubMed enrichment", [ENRICHED])))
    assert _comments(out) == [(_flag("enrichment_failures"), ENRICHED)]


def test_a_window_that_recurs_across_records_does_not_anchor(tmp_path):
    """A middle window matching more than one paragraph is not this record:
    the finding becomes a review note, quoting what it was about."""
    other = "99. Other XY, Person Z. Lanternfish vision under pressure. J Example Biol. 2019;4:1-9."
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("enrichment_failures", "1 publication(s) failed PubMed enrichment", [other])))
    assert _comments(out) == []
    assert _notes(out) == ["PubMed lookup failed (1)", _flag("enrichment_failures"), f'\u2022 "{other}"']


def test_no_quote_falls_back_to_the_sections_template_heading(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("section_lost", "M2B: 1 entry absent from the RESEARCH section")))
    assert _comments(out) == [(_flag("section_lost"), PAST_FUNDING_HEADING)]


def test_a_lost_table_row_is_a_comment_on_its_sections_heading(tmp_path):
    """#534: renderer_row_missing names its taxonomy code the way the run
    page reads a section, so its flag sits on that section's heading rather
    than in a note that names no section. The lost row is not on the page,
    so its own text cannot carry the flag."""
    doc = Document(str(_clean_docx(tmp_path)))
    doc.add_paragraph("HONORS")
    doc.add_paragraph("2001 Imaginary Medal, Pretend Society")
    doc.save(str(tmp_path / "DOC_wcm.docx"))
    finding, = lint_renderer_row_receipts({"row_receipts": [{
        "section": "honors", "linted": True, "given": 2, "written": 1,
        "lost": [{"code": "H", "element_idx_start": 90,
                  "text": "2014 Fictional Award, Example Vascular Society"}]}]})
    out, n = rc.write_review_docx(tmp_path / "DOC_wcm.docx", _report(finding))
    assert n == 1 and _notes(out) == []
    assert _comments(out) == [(_flag("renderer_row_missing"), "HONORS")]


def test_findings_with_no_place_are_review_notes(tmp_path):
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("implausible_year", "entry 9 (K1): date=1905"),  # no Didactic Teaching heading
        _finding("llm_fallback_served", "stage 4 S8: the content filter blocked the primary model"),
        _finding("llm_fallback_served", "stage 4 S5: the content filter blocked the primary model")))
    assert n == 1 and _comments(out) == []
    # A note with nothing to point at (no quote, no section) stays on the run page.
    assert _notes(out) == ["Year probably wrong century (1)", _flag("implausible_year"), "\u2022 Didactic Teaching"]


_DIVERSIONS = (
    _finding("stage6_render_warnings", "stage 6 self-check: M2B: 2 entries diverted to the "
             "Appendix — declined by the research-support renderer as too sparse to table"),
    _finding("stage6_render_warnings", "stage 6 self-check: T: 3 entries diverted to the "
             "Appendix — no stage 6 section is routed to render this taxonomy code"),
)


def test_an_appendix_diversion_is_a_comment_on_its_sections_heading(tmp_path):
    """#1589: the count goes on the heading of the section the entries were
    meant for, not on any entry (the Appendix group headings name where each
    came from). The Appendix's own code is stage 6's Appendix note's to explain."""
    out, n = rc.write_review_docx(_with_appendix_note(tmp_path), _report(*_DIVERSIONS))
    assert n == 1 and _notes(out) == []
    assert _comments(out) == [
        ("2 entries meant for this section are in the Appendix: move any that belong here.",
         PAST_FUNDING_HEADING)]


def test_a_single_diverted_entry_reads_in_the_singular(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(_finding(
        "stage6_render_warnings", "stage 6 self-check: M2B: 1 entry diverted to the Appendix "
        "— declined by the research-support renderer as too sparse to table")))
    assert [text for text, _ in _comments(out)] == [
        "1 entry meant for this section is in the Appendix: move any that belong here."]


def test_a_diversion_with_no_heading_to_sit_on_is_a_note(tmp_path):
    """#1639, kept by #2002: a diversion whose section has no heading in this
    document, or that names no section, is listed in the box rather than vanishing."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("stage6_render_warnings", "stage 6 self-check: K1: 2 entries diverted to the "
                 "Appendix — declined by the renderer"),  # no Didactic Teaching heading
        _finding("stage6_render_warnings", "stage 6 self-check: 1 entry diverted to the "
                 "Appendix — no section named")))
    assert n == 2 and _comments(out) == []
    assert _notes(out) == [f"{rc.DIVERSION_TITLE} (1)", rc.DIVERSION_NOTE, "• Didactic Teaching",
                           f"{rc.DIVERSION_TITLE} (1)", rc.DIVERSION_UNPLACED_NOTE]


def test_a_builder_warning_no_comment_places_is_not_a_note(tmp_path):
    """#2002: a stage-6 builder warning whose quote is not on the page and
    whose section has no heading stays on the run page, even when its row
    is unmeasured (so the gate would put it on the text)."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(_finding(
        "stage6_render_warnings", "stage 6 self-check: K1 (Didactic Teaching): Content appears "
        "combined with semicolons instead of separate bullets", ["Squid optics outreach; marine biology open day"])))
    assert n == 0 and _comments(out) == [] and _notes(out) == []


def test_a_less_certain_finding_with_nothing_to_name_is_not_a_note(tmp_path):
    """#2002 (reversing #1639): off the text, it is the run page's alone."""
    rows = {("llm_fallback_served", None): LintPrecision("llm_fallback_served", 0, 3, "M1")}
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("llm_fallback_served", "stage 4 S8: the content filter blocked the primary model")),
        rows=rows)
    assert n == 0 and _comments(out) == [] and _notes(out) == []


_GATE_ROWS = {
    ("pipe_leaks", None): LintPrecision("pipe_leaks", 0, 2, "M1"),
    ("duplicate_records", None): LintPrecision("duplicate_records", 9, 10, "M1"),
    ("stage6_render_warnings", "appendix_grant_too_sparse"):
        LintPrecision("stage6_render_warnings", 1, 6, "M1", "appendix_grant_too_sparse"),
    ("citation_grounding", None): LintPrecision("citation_grounding", 1, 4, "YUY-CG"),
}


def test_a_finding_right_less_than_half_the_time_is_neither_a_comment_nor_a_note(tmp_path):
    """#1589's gate: below 50% hand-checked precision, nothing sits on the
    text. #2002: nor is it in the box; the run page's Diagnostics lists it."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("pipe_leaks", "1 numbered citation(s) fusing venue-date patterns", [f"[bibliography] {SECOND}"]),
        _finding("duplicate_records", "1 duplicated record(s)", [f"block 7 repeats at 8: {CITATION[:90]}"])),
        rows=_GATE_ROWS)
    assert n == 1
    assert _comments(out) == [(_flag("duplicate_records"), CITATION)]
    assert _notes(out) == []


def test_a_less_certain_diversion_is_neither_a_comment_nor_a_note(tmp_path):
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(_DIVERSIONS[0]), rows=_GATE_ROWS)
    assert n == 0 and _comments(out) == [] and _notes(out) == []


def test_the_box_lists_coverage_and_near_duplicates_but_not_less_certain_or_builder_notes(tmp_path):
    """#2002: of PPMTJH's and WZSMUO's box notes, only the source-line and
    near-duplicate groups held true findings. Less certain findings (here a
    date_cell_shape row at 1 of 4) and stage-6 builder warnings stay on the
    run page; the other two groups are listed as before."""
    rows = {("date_cell_shape", None): LintPrecision("date_cell_shape", 1, 4, "YUY-HO"),
            ("stage6_render_warnings", "semicolon_fused_bullets"):
                LintPrecision("stage6_render_warnings", 0, 2, "M1", "semicolon_fused_bullets")}
    lost = "Visiting lecturer in comparative squid anatomy, Example Polytechnic"
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("date_cell_shape", "entry 40 (K1): date cell reads '2014 2014'", ["Squid seminar, 2014 2014"]),
        _finding("stage6_render_warnings", "stage 6 self-check: K1 (Didactic Teaching): Content "
                 "appears combined with semicolons instead of separate bullets", ["Squid optics outreach; marine biology open day"]),
        _finding("source_line_coverage", "entry 31: 1 source line(s) with under 50% of their word pairs "
                 "anywhere in the output, Appendix included (lowest 0%)", [lost], severity="INFO"),
        _finding("dedup_drops", "1 dedup drop(s)", [
            "K1 (jaccard=1.00, 89% covered by kept): dropped 'Squid Optics Seminar' "
            "vs kept 'Squid Optics Seminar 2'"])),
        rows=rows)
    assert n == 2 and _comments(out) == []
    assert _notes(out) == [
        f"{SLC_TITLE} (1)",
        _flag("source_line_coverage"), f'\u2022 "{lost}"',
        f"{rc.DEDUP_TITLE} (1)", rc.DEDUP_INSTRUCTION, "\u2022 Didactic Teaching",
        'Removed: "Squid Optics Seminar"', 'Kept: "Squid Optics Seminar 2"']


def test_a_less_certain_appendix_code_diversion_is_not_a_note_either(tmp_path):
    """T's own entries are stage 6's Appendix note's to explain, gated or not."""
    rows = {**_GATE_ROWS, ("stage6_render_warnings", "appendix_no_route_T"):
            LintPrecision("stage6_render_warnings", 0, 9, "M1", "appendix_no_route_T")}
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(_DIVERSIONS[1]), rows=rows)
    assert n == 0 and _notes(out) == []


def test_a_less_certain_possibility_is_dropped_not_noted(tmp_path):
    """A REVIEW_COPY_ONLY_LINTS lint is only ever a comment on its citation."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("citation_grounding", "entry 14 (S1): author_2:initials_not_in_source", [SECOND[4:]],
                 severity="INFO")), rows=_GATE_ROWS)
    assert n == 0 and _comments(out) == [] and _notes(out) == []


def test_the_committed_ledger_is_read_when_no_rows_are_given(tmp_path, monkeypatch):
    """Without ``rows`` the gate reads PRECISION.md through precision.py."""
    monkeypatch.setattr(rc, "load_gate_ledger", lambda: _GATE_ROWS)
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("pipe_leaks", "1 numbered citation(s)", [f"[bibliography] {SECOND}"])))
    assert _comments(out) == []


def test_a_record_printed_across_table_cells_is_found_by_its_row(tmp_path):
    doc = Document(str(_clean_docx(tmp_path)))
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "Squid Optics Seminar Leader (BIOL 412)", "3 hrs, 40 students"
    doc.save(str(tmp_path / "DOC_wcm.docx"))
    out, _ = rc.write_review_docx(tmp_path / "DOC_wcm.docx", _report(
        _finding("enrichment_failures", "1 failed", ["Squid Optics Seminar Leader (BIOL 412) 3 hrs, 40 students"])))
    assert _comments(out) == [(_flag("enrichment_failures"), "Squid Optics Seminar Leader (BIOL 412)")]


def test_a_dedup_drop_led_by_its_entry_number_and_cut_at_the_doctors_length(tmp_path):
    """Current reports lead with "entry N: " (which the run page files as a
    note) and cut each text at DEDUP_TEXT_CHARS; the box marks the cut."""
    kept = ("Squid Optics Seminar Leader, Example School of Marine Biology, eight lectures yearly; " * 5)[:rc.DEDUP_TEXT_CHARS]
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(_finding("dedup_drops", "1 drop", [
        f"entry 12: K1 (jaccard=1.00, 89% covered by kept): dropped 'Squid Optics Seminar' vs kept '{kept}'"])))
    assert _notes(out)[2:] == ["\u2022 Didactic Teaching", 'Removed: "Squid Optics Seminar"', f'Kept: "{kept}\u2026"']


def test_a_dedup_quote_the_doctor_cut_still_reads(tmp_path):
    """The run page marks a quote cut at a doctor cap with an ellipsis."""
    evidence = ("K1 (jaccard=0.92, 86% covered by kept): dropped 'Core Squid Curriculum b Lecturer' "
                "vs kept 'Core Squid Curriculum b Coordinator, Lecturer, Workshop Facilitator, Squid Opti'")
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("dedup_drops", "1 dedup drop(s)", [evidence + "\u2026"])))
    assert _notes(out)[3] == 'Removed: "Core Squid Curriculum b Lecturer"'


def test_protected_data_flag_never_carries_the_finding_text(tmp_path):
    """The flag is the fixed wording alone: nothing from the finding, which
    names a category and a section, reaches the comment."""
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("protected_data_in_output", "protected personal data (children / dependents) found in "
                 "Appendix -- value withheld from this finding", severity="ERROR"),
        _finding("protected_data_in_output", "a bare date found in the Personal Data block", severity="ERROR")))
    assert _comments(out) == []
    assert _notes(out) == ["Protected personal data in document (2)", _flag("protected_data_in_output"),
                           "\u2022 Appendix: children / dependents", "\u2022 Personal Data: a bare date"]


def test_info_and_skipped_findings_get_no_comment(tmp_path):
    clean = _clean_docx(tmp_path)
    skipped = {**_finding("segmentation", "skipped"), "status": "skipped"}
    out, n = rc.write_review_docx(clean, _report(
        _finding("missed_headers", "1 header", severity="INFO"), skipped))
    assert n == 0 and _comments(out) == [] and _notes(out) == []


def test_a_citation_grounding_finding_is_a_comment_on_its_citation_though_info(tmp_path):
    """#1570 at 50% precision: never a run-page row, but a possibility
    commented on the citation it names, whatever its severity."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("citation_grounding", "entry 14 (S1): author_2:initials_not_in_source -- the stage 5d "
                 "citation names text its source line lacks", [SECOND[4:]], severity="INFO"),
        # Another lint's INFO finding quoting a line that IS in the document:
        # the exception is this lint's, so this one still gets no comment.
        _finding("owner_attribution", "entry 12 (S1): 1 publication(s) never name the owner",
                 [CITATION[4:]], severity="INFO")))
    assert n == 1
    assert _comments(out) == [(_flag("citation_grounding"), SECOND)]
    assert "may" in _flag("citation_grounding")


def test_a_citation_grounding_finding_not_in_the_document_flags_nothing(tmp_path):
    """No heading comment, no review note: a possibility only reads on the citation."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("citation_grounding", "entry 3 (S8): ordinal_not_in_source:10th -- the stage 5d "
                 "citation names text its source line lacks",
                 ["Quill A. A talk nobody printed. 10th Annual Meeting; 2004."], severity="INFO")))
    assert n == 0 and _comments(out) == [] and _notes(out) == []


def test_a_source_line_coverage_finding_is_one_review_note_per_line_though_info(tmp_path):
    """#1588: the quoted source lines are not on the page, so each one is a
    review note under its own title, never a comment, whatever the severity."""
    lost = ("Visiting lecturer in comparative squid anatomy, Example Polytechnic",
            "Organised the annual squid optics colloquium for graduate students")
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("source_line_coverage", "entry 31: 2 source line(s) with under 50% of their word pairs "
                 "anywhere in the output, Appendix included (lowest 0%)", lost, severity="INFO")))
    assert n == 2 and _comments(out) == []
    assert _notes(out) == [f"{SLC_TITLE} (2)", _flag("source_line_coverage"),
                           *(f'• "{line}"' for line in lost)]


def test_a_source_line_coverage_finding_below_the_bar_is_still_its_notes(tmp_path):
    """Its lines are notes already, never on the text, so the precision gate
    has nothing to move: a measured precision under 50% (PRECISION.md, YUY-SLC)
    must not drop them as it drops a less certain possibility."""
    lost = ("Visiting lecturer in comparative squid anatomy, Example Polytechnic",)
    rows = {("source_line_coverage", None): LintPrecision("source_line_coverage", 1, 4, "T")}
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("source_line_coverage", "entry 31: 1 source line(s) with under 50% of their word pairs "
                 "anywhere in the output, Appendix included (lowest 0%)", lost, severity="INFO")), rows=rows)
    assert n == 1 and _comments(out) == []
    assert _notes(out) == [f"{SLC_TITLE} (1)", _flag("source_line_coverage"),
                           f'• "{lost[0]}"']


def test_stray_text_is_titled_plainly_and_quotes_what_to_delete(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("output_hygiene", "1 boilerplate line(s) rendered in the appendix", ["Insert dates here (MM/YYYY)"])))
    assert _notes(out) == ["Stray text to delete (1)", _flag("output_hygiene"), '\u2022 "Insert dates here (MM/YYYY)"']


def test_review_notes_are_one_comment_on_the_first_paragraph_not_a_box(tmp_path):
    """The notes add nothing to the body: one comment by their own author on
    the document's first paragraph, its title and group labels bold."""
    clean = _clean_docx(tmp_path)
    out, _ = rc.write_review_docx(clean, _report(
        _finding("enrichment_failures", "1 failed", ["99. Nowhere AB. Unprinted paper. J Example. 2001."])))
    doc = Document(str(out))
    assert len(doc.tables) == len(Document(str(clean)).tables) == 0
    assert _comments(out, rc.REVIEW_NOTES_AUTHOR) == [(_notes_comment(out).text, HEADER)]
    title, group, instruction, item = _notes_comment(out).paragraphs[:4]
    assert title.text == rc.REVIEW_NOTES_TITLE and all(r.bold for r in title.runs if r.text)
    assert group.text == "PubMed lookup failed (1)" and group.runs[0].bold
    assert instruction.text == _flag("enrichment_failures") and not instruction.runs[0].bold
    assert item.text.startswith(rc.NOTE_BULLET)


def test_removed_and_kept_lead_with_italic_labels(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(_finding("dedup_drops", "1 drop", [
        "K1 (jaccard=1.00, 89% covered by kept): dropped 'Squid Lecture 5 hrs' vs kept 'Squid Lecture 4 hrs'"])))
    paragraphs = _notes_comment(out).paragraphs
    at = next(i for i, p in enumerate(paragraphs) if p.text.startswith("Removed:"))
    removed, kept = paragraphs[at:at + 2]
    for para, label, text in ((removed, "Removed: ", '"Squid Lecture 5 hrs"'),
                              (kept, "Kept: ", '"Squid Lecture 4 hrs"')):
        assert para.runs[0].text == label and para.runs[0].italic
        assert para.runs[1].text == text and not para.runs[1].italic


def test_delete_all_comments_leaves_the_clean_documents_text(tmp_path):
    """The copy is the one download: Word's Review > Delete All Comments
    must leave exactly the clean document's text, tracked insertions
    included, whatever the copy marked (a year split out of its run, a
    heading, a citation in a w:ins) or listed in its review notes."""
    clean = _with_appendix_note(tmp_path)
    out, n = rc.write_review_docx(clean, _report(
        _finding("implausible_year", "entry 3 (M2B): end_date=1912 -- before 1959", [f"Grant Title: {GRANT}"]),
        _finding("enrichment_failures", "1 publication(s) failed PubMed enrichment", [ENRICHED]),
        _finding("section_lost", "M2B: 1 entry absent from the RESEARCH section"),
        _finding("output_hygiene", "1 boilerplate line(s)", ["Insert dates here (MM/YYYY)"]),
        _finding("dedup_drops", "1 drop", ["K1 (jaccard=1.00): dropped 'Squid Lecture' vs kept 'Squid Talk'"])))
    assert n == 5 and len(_comments(out)) == 3 and _notes(out)  # every kind of mark is in play
    clean_doc = Document(str(clean))
    assert _body_text(Document(str(out))) == _body_text(clean_doc)  # comments sit outside the text
    deleted = _delete_all_comments(out)
    assert list(deleted.comments) == []
    assert _body_text(deleted) == _body_text(clean_doc)
    assert [el.tag for el in deleted.element.body] == [el.tag for el in clean_doc.element.body]
    assert not any(p.text.startswith(rc.REVIEW_NOTES_TITLE) for p in deleted.paragraphs)


def test_a_copy_with_nothing_to_flag_is_written_and_says_what_is_not_checked(tmp_path):
    """#1589: a quiet doctor is not a clean document, so every run gets its
    copy, closing with what CViche cannot see."""
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("llm_fallback_served", "stage 4 S8: the content filter blocked the primary model")))
    assert out == tmp_path / f"DOC{REVIEW_DOCX_SUFFIX}" and n == 0
    assert _comments(out) == [] and _notes(out) == []
    assert _not_checked(out) == [rc.NOT_CHECKED_TITLE, rc.NOT_CHECKED_INSTRUCTION,
                                 *(f"\u2022 {spot.sentence}" for spot in blind_spots())]
    assert len(_not_checked(out)) > 2  # COVERAGE.md lists at least one blind spot


def test_what_is_not_checked_closes_a_box_that_has_notes(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("output_hygiene", "1 boilerplate line(s)", ["Insert dates here (MM/YYYY)"])))
    lines = _note_lines(out)
    assert lines[0] == "Stray text to delete (1)"
    assert lines[-len(blind_spots()) - 2:] == _not_checked(out)


def test_one_lints_comments_are_capped(tmp_path):
    many = [_finding("output_hygiene", f"leak {i}", [CITATION]) for i in range(rc.MAX_FLAGS_PER_LINT + 5)]
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(*many, _finding("pipe_leaks", "one", [SECOND])))
    assert len(_comments(out)) == rc.MAX_FLAGS_PER_LINT + 1


def test_comments_leave_the_body_the_doctor_reads_unchanged(tmp_path):
    """Comments live in comments.xml: the doctor's view of the copy is the clean one's."""
    clean = _clean_docx(tmp_path)
    out, _ = rc.write_review_docx(clean, _report(
        _finding("pipe_leaks", "x", [CITATION]),
        _finding("section_lost", "M2B: 1 entry absent from the RESEARCH section")))
    assert _notes(out) == []
    assert read_docx_blocks(str(out)) == read_docx_blocks(str(clean))
    assert len(list(Document(str(clean)).comments)) == 0  # the clean document is untouched


def test_not_a_doctor_report_writes_nothing(tmp_path):
    assert rc.write_review_docx(_clean_docx(tmp_path), None) is None


def test_source_line_coverage_notes_are_capped_per_document(tmp_path):
    """Up to 3 lines a finding, up to 57 findings a run on YUYVIG: the notes
    stop at MAX_FLAGS_PER_LINT, as a comment-placing lint's flags do."""
    findings = [_finding("source_line_coverage", f"entry {i}: 3 source line(s) with under 50% of their "
                         "word pairs anywhere in the output, Appendix included (lowest 0%)",
                         [f"Squid optics seminar number {i} line {j} for graduate students" for j in range(3)])
                for i in range(20)]
    out, n = rc.write_review_docx(_clean_docx(tmp_path), _report(*findings))
    assert n == rc.MAX_FLAGS_PER_LINT
    assert _notes(out)[0] == f"{SLC_TITLE} ({rc.MAX_FLAGS_PER_LINT})"


# --- #1591: a certain fix as a tracked change ---------------------------------

OWNER = {"first_name": "Ada", "last_name": "Testowner"}
GRANT_TITLE = "Example Squid Optics Project"
OWNER_PI_ROLE_FINDING = _finding(
    "role_consistency", "entry 5: 'Name of Principal Investigator:' names the CV owner, and "
    "'Your role:' is empty (owner_pi_role_empty, #1403)", [GRANT_TITLE])


def _stage4(*pi_names):
    return {"cv_owner": OWNER, "entries": [
        {"taxonomy_code": "M2B", "element_idx_start": 5 + n, "text": "A grant",
         "extracted_fields": {"title": GRANT_TITLE, "pi_name": pi}} for n, pi in enumerate(pi_names)]}


def _with_grant_tables(tmp_path: Path, *tables: tuple[str, str]) -> Path:
    """The clean document with a rendered grant table per (PI cell, role
    cell); a role cell of None leaves the table without a "Your role:" row."""
    clean = _clean_docx(tmp_path)
    doc = Document(str(clean))
    for pi_name, role in tables:
        rows = [("Award Source:", "Example Foundation"), (PROJECT_TITLE_LABEL, GRANT_TITLE),
                (PI_NAME_LABEL, pi_name), *([(YOUR_ROLE_LABEL, role)] if role is not None else []),
                ("Your percent (%) effort:", "")]
        table = doc.add_table(rows=len(rows), cols=2)
        for row, (label, value) in zip(table.rows, rows, strict=True):
            row.cells[0].text, row.cells[1].text = label, value
    doc.save(str(clean))
    return clean


def test_a_grant_facts_finding_is_a_comment_on_its_grant_title_though_info(tmp_path):
    """#1588: a grant amount the CV entry states and the table lacks is a
    possibility, commented on the grant's title in its table, whatever its
    severity, and never a review note."""
    clean = _with_grant_tables(tmp_path, ("", "PI"))
    stage4 = {"entries": [{"element_idx_start": 7, "taxonomy_code": "M2B",
                           "text": f"{GRANT_TITLE}, Example Foundation, $40,000",
                           "extracted_fields": {"title": GRANT_TITLE}}]}
    findings = lint_grant_facts(stage4, read_docx_table_rows(str(clean)))
    out, n = rc.write_review_docx(clean, _report(*findings))
    assert [f["severity"] for f in findings] == ["INFO"] and n == 1
    assert _comments(out) == [(_flag("grant_facts"), GRANT_TITLE)]
    assert _notes(out) == []


def _with_past_mentees(tmp_path: Path, *current_positions: str) -> Path:
    """The clean document with a Past Mentees heading and a rendered mentee
    table per Current Position value; every other row filled."""
    clean = _clean_docx(tmp_path)
    doc = Document(str(clean))
    doc.add_paragraph(PAST_MENTEES_HEADING)
    for n, position in enumerate(current_positions):
        values = [f"Mentee Example {n}", "Example University - Fellow", "2010-2012",
                  "Squid lens optics", position, "Research"]
        table = doc.add_table(rows=len(MENTEE_TABLE_LABELS), cols=2)
        for row, label, value in zip(table.rows, MENTEE_TABLE_LABELS, values, strict=True):
            row.cells[0].text, row.cells[1].text = label, value
    doc.save(str(clean))
    return clean


def _commented_row_labels(path: Path) -> list[str]:
    """The first cell's text of the table row each finding comment starts in."""
    labels = []
    for start in Document(str(path)).element.body.iter(qn("w:commentRangeStart")):
        row = next((a for a in start.iterancestors() if a.tag == qn("w:tr")), None)
        if row is not None:
            labels.append("".join(t.text or "" for t in next(row.iter(qn("w:tc"))).iter(qn("w:t"))))
    return labels


def test_a_blank_mentee_field_is_one_comment_on_its_first_cell_though_info(tmp_path):
    """#1941: three past mentees with no Current Position ("N/A", blank, "none")
    are one comment, on the first one's cell, asking for the field; never a
    review note."""
    clean = _with_past_mentees(tmp_path, "Assistant Professor", "N/A", "", "none")
    findings = lint_mentee_field_incomplete(read_docx_table_rows(str(clean)),
                                            read_docx_blocks(str(clean)))
    out, n = rc.write_review_docx(clean, _report(*findings))
    assert [f["severity"] for f in findings] == ["INFO"] and n == 1
    assert _comments(out) == [(_flag("mentee_field_incomplete"), "N/A")]
    assert _commented_row_labels(out) == [MENTEE_CURRENT_POSITION_LABEL]
    assert _notes(out) == []
    assert json.loads(rc.comment_map_path(out).read_text())[0]["lint"] == "mentee_field_incomplete"


def test_a_blank_mentee_field_finding_naming_no_such_row_is_not_flagged(tmp_path):
    """A finding whose table or row the document lacks is no comment and no note."""
    clean = _with_past_mentees(tmp_path, "")
    for message in ("table 9, row 'Current Position:': blank or N/A in 1 of 1 Past Mentees tables",
                    "table 0, row 'Goals:': blank or N/A in 1 of 1 Past Mentees tables",
                    "Current Position blank"):
        out, n = rc.write_review_docx(clean, _report(_finding(
            "mentee_field_incomplete", message, severity="INFO")))
        assert n == 0 and _comments(out) == [] and _notes(out) == []


def _role_findings(clean: Path, stage4: dict) -> list[dict]:
    """What the doctor reports for ``clean``: role_consistency over its tables."""
    return lint_role_consistency(stage4, read_docx_table_rows(str(clean)))


def _role_cells(path: Path) -> list:
    """Each grant table's "Your role:" value cell element, in order."""
    return [row.cells[1]._tc for table in Document(str(path)).tables for row in table.rows
            if row.cells[0].text == YOUR_ROLE_LABEL]


def _reviewed(cell, accept: bool) -> str:
    """The cell's text once every tracked change is accepted (w:t, inserted
    runs included; a deletion's w:delText gone), or rejected (inserted runs
    gone; w:delText back)."""
    if accept:
        return "".join(t.text or "" for t in cell.iter(qn("w:t")))
    return "".join(el.text or "" for el in cell.iter(qn("w:t"), qn("w:delText"))
                   if el.tag == qn("w:delText")
                   or not any(a.tag == qn("w:ins") for a in el.iterancestors()))


def test_an_empty_role_beside_the_owner_as_pi_gets_pi_as_a_tracked_insertion(tmp_path):
    clean = _with_grant_tables(tmp_path, ("Ada Testowner", ""))
    out, n = rc.write_review_docx(clean, _report(OWNER_PI_ROLE_FINDING), _stage4("Testowner"))
    assert n == 1
    (cell,) = _role_cells(out)
    assert _reviewed(cell, accept=True) == rc.OWNER_PI_ROLE_FIX == "PI"
    assert _reviewed(cell, accept=False) == ""  # the delivered text
    (ins,) = cell.iter(qn("w:ins"))
    assert ins.get(qn("w:author")) == rc.COMMENT_AUTHOR and ins.get(qn("w:date")) is None
    # The ids follow the clean document's own tracked insertion (w:id 90).
    assert ins.get(qn("w:id")) == "91"
    # The fix replaces the comment; the clean document is untouched.
    assert _comments(out) == [] and _reviewed(_role_cells(clean)[0], accept=True) == ""


def test_only_the_tables_the_shape_reports_get_the_fix(tmp_path):
    """Another PI's table with an empty role, and the owner's with a role,
    are left as delivered."""
    clean = _with_grant_tables(tmp_path, ("Other Person", ""), ("Ada Testowner", "Co-PI"),
                               ("Ada Testowner", ""))
    stage4 = _stage4("Other Person", "Testowner", "Testowner")
    (finding,) = [f for f in _role_findings(clean, stage4) if "(owner_pi_role_empty," in f["message"]]
    out, n = rc.write_review_docx(clean, _report(finding), stage4)
    assert n == 1
    assert [_reviewed(c, accept=True) for c in _role_cells(out)] == ["", "Co-PI", "PI"]


def test_each_fix_is_its_own_revision(tmp_path):
    """Word accepts or rejects each change by its w:id: two fixes, two ids."""
    clean = _with_grant_tables(tmp_path, ("Ada Testowner", ""), ("Ada Testowner", ""))
    out, n = rc.write_review_docx(clean, _report(OWNER_PI_ROLE_FINDING), _stage4("Testowner", "Testowner"))
    ids = [ins.get(qn("w:id")) for cell in _role_cells(out) for ins in cell.iter(qn("w:ins"))]
    assert n == 2 and ids == ["91", "92"]


def test_a_reported_table_the_fix_cannot_reach_keeps_its_comment(tmp_path):
    """A table with no "Your role:" row shows the shape too, but has no cell
    for the fix: its finding stays a comment; the fixed table's goes."""
    clean = _with_grant_tables(tmp_path, ("Ada Testowner", ""), ("Ada Testowner", None))
    stage4 = _stage4("Testowner", "Testowner")
    findings = _role_findings(clean, stage4)
    assert [f["message"][:7] for f in findings] == ["entry 5", "entry 6"]
    out, n = rc.write_review_docx(clean, _report(*findings), stage4)
    assert n == 2
    assert [_reviewed(c, accept=True) for c in _role_cells(out)] == ["PI"]
    assert _comments(out) == [(_flag("role_consistency"), GRANT_TITLE)]


def test_without_the_stage4_artifact_the_finding_stays_a_comment(tmp_path):
    clean = _with_grant_tables(tmp_path, ("Ada Testowner", ""))
    out, n = rc.write_review_docx(clean, _report(OWNER_PI_ROLE_FINDING))
    assert n == 1
    assert _comments(out) == [(_flag("role_consistency"), GRANT_TITLE)]
    assert [_reviewed(c, accept=True) for c in _role_cells(out)] == [""]


def test_no_fix_without_the_doctors_finding(tmp_path):
    """The fix follows the doctor's report, not a fresh read of the tables."""
    clean = _with_grant_tables(tmp_path, ("Ada Testowner", ""))
    out, _ = rc.write_review_docx(clean, _report(
        _finding("enrichment_failures", "1 publication(s) failed PubMed enrichment", [SECOND])), _stage4("Testowner"))
    assert [_reviewed(c, accept=True) for c in _role_cells(out)] == [""]


def test_the_fix_leaves_every_other_cell_as_delivered(tmp_path):
    """Accepting the copy's changes differs from the clean document only in
    the fixed role cell (the doctor's own table reader, which reads w:ins)."""
    clean = _with_grant_tables(tmp_path, ("Ada Testowner", ""))
    out, _ = rc.write_review_docx(clean, _report(OWNER_PI_ROLE_FINDING), _stage4("Testowner"))
    before, after = read_docx_table_rows(str(clean)), read_docx_table_rows(str(out))
    assert len(after) == len(before)  # the review notes are a comment, not a table
    changed = [(t, r) for t, (a, b) in enumerate(zip(before, after, strict=True))
               for r, (x, y) in enumerate(zip(a, b, strict=True)) if x != y]
    assert changed == [(0, 3)]
    assert after[0][3] == [YOUR_ROLE_LABEL, "PI"]


def test_the_comment_map_names_each_comments_finding_by_its_id(tmp_path):
    """#1654: lint, message shape, severity and entry index per comment, no text."""
    reroute = ("stage 6 self-check: hierarchy-mismatch reroute K1->S8 refused fields do not fit: "
               "1 entry")
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report(
        _finding("implausible_year", "entry 3 (M2B): end_date=1912 -- before 1959",
                 [f"Grant Title: {GRANT}; Role: PI"], severity="ERROR"),
        _finding("stage6_render_warnings", reroute, [SECOND]),
        _finding("dedup_drops", "1 entry dropped", ["entry 9: K1 (jaccard=1.00): dropped 'x' vs kept 'y'"])))
    marked = json.loads(rc.comment_map_path(out).read_text(encoding="utf-8"))
    assert rc.comment_map_path(out).name == "DOC_wcm_review_comments.json"
    assert [m["comment_id"] for m in marked] == [c.comment_id for c in Document(str(out)).comments
                                                 if c.author == rc.COMMENT_AUTHOR]
    assert [(m["lint"], m["shape"], m["severity"], m["entry_index"]) for m in marked] == [
        ("implausible_year", None, "ERROR", 3),
        ("stage6_render_warnings", "reroute_refused", "WARN", None)]
    assert GRANT not in rc.comment_map_path(out).read_text(encoding="utf-8")


def test_a_copy_with_nothing_to_mark_has_an_empty_comment_map(tmp_path):
    out, _ = rc.write_review_docx(_clean_docx(tmp_path), _report())
    assert json.loads(rc.comment_map_path(out).read_text(encoding="utf-8")) == []


def test_a_comments_wording_names_its_lint_only_when_one_lint_words_it_so():
    """For a review copy written before its comment map."""
    assert rc.lint_of_flag(_flag("implausible_year")) == "implausible_year"
    assert _flag("bucket_status") == _flag("grant_bucket")
    assert rc.lint_of_flag(_flag("grant_bucket")) is None
    assert rc.lint_of_flag(rc.DIVERSION_FLAG.format(count="2 entries", verb="are")) == rc.DIVERSION_LINT
    assert rc.lint_of_flag("A finding's own message, with no flag of its own.") is None


def test_a_comment_map_that_cannot_be_written_leaves_no_review_copy(tmp_path):
    """The map is written before the copy, so its failure never strands a saved copy."""
    clean = _clean_docx(tmp_path)
    rc.comment_map_path(rc.review_docx_path(clean)).mkdir()
    with pytest.raises(rc.REVIEW_DOCX_ERRORS):
        rc.write_review_docx(clean, _report())
    assert not rc.review_docx_path(clean).exists()
