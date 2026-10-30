"""Template rules: the review-copy comments the WCM template's own asks imply (#1999).

Every document is the real WCM template with invented text written into it;
no corpus CV.
"""
import json
import os
import sys
from copy import deepcopy
from pathlib import Path

os.environ.setdefault("CVICHE_SESSION_SECRET", "test-secret-not-for-production")
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent / "src"))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from app.services import template_rules as tr  # noqa: E402
from unified_pipeline.quality_score import _TEMPLATE_DOCX_PATH  # noqa: E402
from unified_pipeline.stage6.empty_slots import NOT_APPLICABLE  # noqa: E402

CITATION = "Quill AB, Marsh CD. Lanternfish vision under pressure. J Example Biol. 2019;4:1-9."
SECTION_EMPTY_SLOTS = {r.slot for r in tr.load_rules() if r.condition == "section_empty"}


def _template():
    return Document(str(_TEMPLATE_DOCX_PATH))


def _paragraph(doc, start: str):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(start))


def _write_after(doc, start: str, lines: list[str]) -> None:
    """Paragraphs of ``lines`` right under the template paragraph starting ``start``."""
    at = _paragraph(doc, start)._p
    for line in reversed(lines):
        p = deepcopy(at)
        for child in [c for c in p if c.tag != qn("w:pPr")]:
            p.remove(child)
        at.addnext(p)
        Paragraph(p, doc._body).add_run(line)


def _cell(doc, label: str, row: int = 0, column: int = 1, exact: bool = False):
    """The cell ``row`` rows below and in ``column`` of the first table cell
    starting ``label`` (reading ``label`` whole when ``exact``)."""
    for table in doc.tables:
        for r, tr_ in enumerate(table.rows):
            for cell in tr_.cells:
                text = cell.text.strip()
                if text == label if exact else text.startswith(label):
                    return table.rows[r + row].cells[column]
    raise AssertionError(label)


def _fired(doc) -> dict[str, str]:
    """Each fired rule's shape -> the text of the paragraph its comment goes on."""
    return {rule.shape: para.text for para, rule in tr.template_rule_comments(doc)}


def _numbered(n: int) -> list[str]:
    return [f"{i}. {CITATION}" for i in range(1, n + 1)]


def test_every_rule_names_a_slot_the_template_ships():
    doc = _template()
    texts = [p.text.strip() for p in doc.paragraphs]
    texts += [c.text.strip() for t in doc.tables for row in t.rows for c in row.cells]
    for rule in tr.load_rules():
        assert any(t.startswith(rule.slot) for t in texts), rule.slot


def test_the_bare_template_fires_every_rule_but_the_caps_each_on_its_slot():
    fired = _fired(_template())
    expected = {r.shape for r in tr.load_rules() if not r.condition.startswith("max_")}
    assert set(fired) == expected
    for shape, anchor in fired.items():
        assert anchor.strip().startswith(shape.split(":", 1)[1]), shape


@pytest.mark.parametrize("slot,cap", [("Abstracts", 20), ("Case Reports", 10)])
def test_a_capped_subsection_fires_only_past_its_cap(slot, cap):
    at_cap, over = _template(), _template()
    _write_after(at_cap, slot, _numbered(cap))
    _write_after(over, slot, _numbered(cap + 1))
    assert f"max_entries:{slot}" not in _fired(at_cap)
    assert f"max_entries:{slot}" in _fired(over)


def test_a_capped_subsection_counts_numbered_entries_only_up_to_the_next_heading():
    doc = _template()
    _write_after(doc, "Abstracts", ["An unnumbered note."] * 30)
    _write_after(doc, "In review", _numbered(30))  # the slot above Abstracts
    assert "max_entries:Abstracts" not in _fired(doc)


def test_research_activities_fire_past_300_words():
    at_cap, over = _template(), _template()
    _write_after(at_cap, "Research Activities:", ["word " * 150, "word " * 150])
    _write_after(over, "Research Activities:", ["word " * 150, "word " * 151])
    assert "max_words:Research Activities:" not in _fired(at_cap)
    assert "max_words:Research Activities:" in _fired(over)


@pytest.mark.parametrize("label", ["Office address:", "Office telephone:", "Work email:",
                                   "Primary Hospital Affiliation:"])
def test_a_form_label_with_an_answer_does_not_fire(label):
    doc = _template()
    _cell(doc, label).text = "Example answer"
    fired = _fired(doc)
    assert f"blank:{label}" not in fired
    assert sum(s.startswith("blank:") for s in fired) == 5  # the other labels still do


def test_the_employer_line_with_a_value_after_its_tab_does_not_fire():
    doc = _template()
    _paragraph(doc, "Name of Current Employer(s):").add_run("\tExample Hospital")
    assert "blank:Name of Current Employer(s):" not in _fired(doc)


def test_percent_effort_fires_until_one_activity_has_an_effort():
    doc = _template()
    _cell(doc, "Percent Effort (%)", row=1, column=1).text = "40%"
    assert "blank:Percent Effort (%)" not in _fired(doc)


@pytest.mark.parametrize("slot", ["Academic Degree(s)", "Academic Appointments"])
def test_a_slot_holding_only_na_fires_and_one_with_an_entry_does_not(slot):
    na, filled = _template(), _template()
    header = {"Academic Degree(s)": "Degree, include field", "Academic Appointments": "Title"}[slot]
    _cell(na, header, row=1, column=0, exact=slot == "Academic Appointments").text = NOT_APPLICABLE
    _cell(filled, header, row=1, column=0, exact=slot == "Academic Appointments").text = "Example entry"
    assert f"not_applicable:{slot}" in _fired(na)
    assert f"not_applicable:{slot}" not in _fired(filled)


def test_appointments_rule_reads_its_own_table_not_the_hospital_one_below():
    doc = _template()
    _cell(doc, "Title", row=1, column=0, exact=True).text = NOT_APPLICABLE  # Academic Appointments
    hospital = [t for t in doc.tables if t.rows[0].cells[0].text == "Title"][1]
    hospital.rows[1].cells[0].text = "Example Attending"
    assert "not_applicable:Academic Appointments" in _fired(doc)


def test_a_section_with_one_entry_does_not_fire_and_the_others_still_do():
    doc = _template()
    _cell(doc, "Name of award", row=1, column=0).text = "Example Award"
    fired = {s for s in _fired(doc) if s.startswith("section_empty:")}
    assert fired == {f"section_empty:{s}" for s in SECTION_EMPTY_SLOTS - {"HONORS, AWARDS"}}


def test_a_section_answered_only_na_fires():
    doc = _template()
    _cell(doc, "Name of award", row=1, column=0).text = NOT_APPLICABLE
    assert "section_empty:HONORS, AWARDS" in _fired(doc)


def test_mentoring_with_one_mentee_does_not_fire():
    doc = _template()
    _cell(doc, "Name", row=0, column=1, exact=True).text = "Example Mentee"
    assert "section_empty:MENTORING" not in _fired(doc)


def test_nothing_after_the_appendix_counts():
    """The Appendix follows the template's last slot: its numbered lines are not that slot's."""
    doc = _template()
    doc.add_paragraph("T. APPENDIX")
    for line in _numbered(3):
        doc.add_paragraph(line)
    last = tr.TemplateRule("Other (media", "max_entries", "Please list two.", 2)
    assert tr.template_rule_comments(doc, (last,)) == []


def test_a_document_without_the_template_fires_nothing():
    doc = Document()
    doc.add_paragraph("RESEARCH")
    doc.add_paragraph("Abstracts are listed below.")
    assert tr.template_rule_comments(doc) == []


def test_rules_load_from_data_and_reject_what_no_code_reads(tmp_path):
    good = tmp_path / "good.json"
    good.write_text(json.dumps([{"slot": "Books:", "condition": "max_entries", "threshold": 2,
                                 "comment": "Please list two."}]))
    (rule,) = tr.load_rules(good)
    doc = _template()
    _write_after(doc, "Books:", _numbered(3))
    assert [(p.text, r.comment) for p, r in tr.template_rule_comments(doc, (rule,))] == [
        ("Books:", "Please list two.")]
    for bad in ([{"slot": "Books:", "condition": "max_lines", "comment": "x"}],
                [{"slot": "Books:", "condition": "max_entries", "comment": "x"}]):
        path = tmp_path / "bad.json"
        path.write_text(json.dumps(bad))
        tr.load_rules.cache_clear()
        with pytest.raises(ValueError):
            tr.load_rules(path)
