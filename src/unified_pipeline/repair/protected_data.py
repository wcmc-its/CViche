"""Remove protected personal data the rendered WCM document still carries (#1389).

`protected_data_in_output` is the last line of defence behind the pre-render
withhold (`stage6/pii_pass.py`): it reads the finished document and finds
what the withhold missed. Until now a hit meant a RED run and "remove it in
Word" for a person. This removes it before the document leaves stage 6.

The removal is outright, not a tracked deletion, for the reason the lint
scans tracked deletions at all (#1223): struck-through text is still in the
file and still readable. It follows the withhold's own presentation instead:
the value is cut, a body paragraph left with nothing in it carries
`PII_REDACTED_NOTICE`, and each touched paragraph gets one Word comment
naming the policy CATEGORY -- never the value, here or in the report.

Detection reuses the lint's own shapes and scopes, applied per paragraph
(body paragraphs and every table-cell paragraph) under the section the
paragraph's block falls in. Afterwards the real lint runs on the saved
document, and `ProtectedDataRepair.remaining` is what it still finds: the
repair never reports a fix the lint does not see.
"""
from __future__ import annotations

import json
import logging
import os
import re
import tempfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, NamedTuple

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from unified_pipeline.doctor.lints.protected_data import (
    _APPENDIX_SECTION,
    _BARE_DATE_RE,
    _DEA_LABEL_PRESENT_RE,
    _DEA_NUMBER_VALUE_RE,
    _DEA_VALUE_RE,
    _HOME_CONTACT_LABEL_ONLY_RE,
    _HOME_CONTACT_LABEL_RE,
    _INDEPENDENT_SHAPES,
    _PERSONAL_DATA_SECTION,
    _block_sections,
    _is_licensure_section,
    _scan_scope,
    _section_label,
    lint_protected_data_in_output,
)
from unified_pipeline.doctor.shared import docx_body_blocks
from unified_pipeline.stage6.normalization.pii import (
    CAT_CHILDREN,
    CAT_DEA,
    CAT_FAMILY,
    CAT_HOME_CONTACT,
    _pii_matches,
)
from unified_pipeline.stage6.normalization.records import _is_citation_shaped
from unified_pipeline.stage6.pii_pass import (
    PII_REDACTED_NOTICE,
    WITHHELD_COMMENT_AUTHOR,
)

if TYPE_CHECKING:
    from docx.document import Document as DocumentType

logger = logging.getLogger(__name__)

#: The switch both drivers read: "1" runs the repair at the end of stage 6,
#: anything else (the default) leaves the render as it is.
REPAIR_FLAG_ENV = "CVICHE_RUN_REPAIR"

#: The category a bare date in the Personal Data block is reported under. The
#: lint names no policy category for it ("a bare date found in ..."), and the
#: date may be a birth date or anything else, so the comment says only this.
BARE_DATE_CATEGORY = "date"

#: What the repair leaves on a paragraph it cut from. Categories only.
REPAIR_COMMENT_TEMPLATE = (
    "CViche removed protected personal data here after the document was built "
    "({categories}). The value is not reproduced in this comment or anywhere in "
    "this document. Review the source CV if this content is needed."
)

#: A paragraph with only these left after a cut held nothing but the leak.
_RESIDUE_RE = re.compile(r"[\s,;:|•·*–—-]*")
_MULTI_SPACE_RE = re.compile(r"[ \t]{2,}")

#: Words a paragraph must keep, once every independent-shape span is taken
#: out, before its citation shape counts as a title rather than a personal
#: line ("Born in <Place>, 1960." keeps ", 1960."). Over the #1392 review
#: renders the shortest such title kept 7 words; no real leak was citation-shaped.
TITLE_MIN_WORDS = 6
_WORD_RE = re.compile(r"[^\W\d_]{2,}")
#: A capitalised word directly before a capitalised label: "Foster Children -
#: <Title>", "Treating Children: <Subtitle>" -- a title-case phrase, not a
#: label opening its own line ("Children: ...", "Grandchildren – ...").
_TITLE_WORD_BEFORE_RE = re.compile(r"\b[A-Z][\w'’-]*[ \t]+$")
#: The categories whose label the review corpus found inside title-case talk
#: and paper titles; a spouse or birth label there was not seen, so those are
#: not gated on the phrase (only on `_is_title_text`).
_TITLE_LABEL_CATEGORIES = frozenset({CAT_CHILDREN, CAT_FAMILY})


class _Leak(NamedTuple):
    """A half-open span of one paragraph's text and its policy category."""
    start: int
    end: int
    category: str


@dataclass(frozen=True)
class Removal:
    """One paragraph the repair cut from: where (the lint's own section
    vocabulary) and the categories it removed. Never the value."""
    where: str
    categories: tuple[str, ...]


@dataclass(frozen=True)
class ProtectedDataRepair:
    """What the repair found, removed and left. `found` and `remaining` are
    lint finding counts on the document before and after."""
    found: int
    removals: list[Removal] = field(default_factory=list)
    remaining: int = 0

    def to_json(self) -> dict[str, object]:
        return {
            "repair": "protected_data",
            "found": self.found,
            "removed": [{"where": r.where, "categories": list(r.categories)} for r in self.removals],
            "remaining": self.remaining,
        }


def repairs_report_path(docx_path: str | Path) -> Path:
    """The report written beside a repaired `<uid>_wcm.docx`: `<uid>_repairs.json`."""
    path = Path(docx_path)
    return path.with_name(f"{path.stem.removesuffix('_wcm')}_repairs.json")


def _leaks(text: str, section: str | None, dea_label_in_block: bool) -> list[_Leak]:
    """Every span of one paragraph's accepted text the lint would flag in this
    section: the withhold's matcher, then the lint's own shapes, scoped as
    `lint_protected_data_in_output` scopes them."""
    leaks = [_Leak(m.start, m.end, m.category)
             for m in _pii_matches(text, _scan_scope(section),
                                   personal_data=section == _PERSONAL_DATA_SECTION)
             if not (m.category == CAT_HOME_CONTACT
                     and _HOME_CONTACT_LABEL_ONLY_RE.match(text[m.start:m.end]))]
    if section in (_PERSONAL_DATA_SECTION, _APPENDIX_SECTION):
        leaks += _independent_leaks(text, section)
    if section == _PERSONAL_DATA_SECTION:
        leaks += [_Leak(m.start(), m.end(), BARE_DATE_CATEGORY) for m in _BARE_DATE_RE.finditer(text)]
    if _is_licensure_section(section):
        # The lint gates the loose DEA shape on its BLOCK naming DEA (#1217);
        # in a table the label sits in another cell, so the gate is the block's.
        probe = _DEA_VALUE_RE if dea_label_in_block else _DEA_NUMBER_VALUE_RE
        leaks += [_Leak(m.start(), m.end(), CAT_DEA) for m in probe.finditer(text)]
    return leaks


def _is_title_text(text: str, shapes: list[_Leak]) -> bool:
    """Whether a paragraph reads as a citation or title once the shape spans
    are out of it: stage 6's own citation shape (`_is_citation_shaped`: a
    year and a closing period) over a residue of `TITLE_MIN_WORDS` words."""
    residue = "".join(ch for i, ch in enumerate(text) if not any(s.start <= i < s.end for s in shapes))
    return _is_citation_shaped(residue.strip()) and len(_WORD_RE.findall(residue)) >= TITLE_MIN_WORDS


def _in_title_case_phrase(text: str, leak: _Leak) -> bool:
    """Whether a child or family label, capitalised, continues a title-case phrase."""
    return (leak.category in _TITLE_LABEL_CATEGORIES and text[leak.start:leak.start + 1].isupper()
            and bool(_TITLE_WORD_BEFORE_RE.search(text[:leak.start])))


def _independent_leaks(text: str, section: str | None) -> list[_Leak]:
    """The lint's independent shapes in one paragraph, less those in title
    text (#1392 review). Those shapes false-match publication and talk titles
    ("...in Children: <Subtitle>", "...born in <Place>, 1989-1993") that the
    Appendix carries verbatim; a cut there would silently delete real CV
    content, so the hit stays a finding (`remaining`) for a person to read.
    A home phone or street address is never title text, so it is always cut.
    The Personal Data block holds no titles, so it is not gated."""
    shapes = [_Leak(m.start(), m.end(), category)
              for pattern, category in _INDEPENDENT_SHAPES for m in pattern.finditer(text)]
    if section != _APPENDIX_SECTION:
        return shapes
    title = _is_title_text(text, shapes)
    return [leak for leak in shapes
            if leak.category == CAT_HOME_CONTACT
            or not (title or _in_title_case_phrase(text, leak))]


def _deleted_leaks(text: str) -> list[_Leak]:
    """The lint's own shapes in a paragraph's tracked-deletion text (#1223)."""
    return [_Leak(m.start(), m.end(), category)
            for pattern, category in _INDEPENDENT_SHAPES for m in pattern.finditer(text)]


_SEPARATOR_AFTER_RE = re.compile(r"[ \t]*[;,][ \t]*")
_SEPARATOR_BEFORE_RE = re.compile(r"[ \t]*[;,][ \t]*$")


def _with_separator(text: str, leak: _Leak) -> _Leak:
    """The span plus one list separator beside it, so a cut out of "sailing;
    Husband: ..." or "A; Spouse: ...; B" leaves "sailing" and "A; B", not a
    dangling "; ". The separator after it first, else the one before."""
    after = _SEPARATOR_AFTER_RE.match(text, leak.end)
    if after:
        return leak._replace(end=after.end())
    before = _SEPARATOR_BEFORE_RE.search(text[:leak.start])
    return leak._replace(start=before.start()) if before else leak


def _merged(leaks: list[_Leak]) -> list[_Leak]:
    """Overlapping spans as one, in order; a merged span keeps every category
    joined so the comment names them all."""
    merged: list[_Leak] = []
    for leak in sorted(leaks):
        if merged and leak.start < merged[-1].end:
            last = merged[-1]
            categories = last.category if leak.category in last.category.split("; ") \
                else f"{last.category}; {leak.category}"
            merged[-1] = _Leak(last.start, max(last.end, leak.end), categories)
        else:
            merged.append(leak)
    return merged


def _text_nodes(p_el: BaseOxmlElement, tag: str) -> list[BaseOxmlElement]:
    return list(p_el.iter(qn(tag)))


def _paragraph_text(p_el: BaseOxmlElement) -> str:
    """The accepted view `doctor.shared._docx_text` reads: every `w:t`,
    inserted runs included, `w:delText` excluded."""
    return "".join(node.text or "" for node in _text_nodes(p_el, "w:t"))


def _set_text(node: BaseOxmlElement, text: str) -> None:
    node.text = text
    if text != text.strip():
        node.set(qn("xml:space"), "preserve")


def _cut(p_el: BaseOxmlElement, spans: list[_Leak]) -> None:
    """Remove the spans from a paragraph's `w:t` nodes, wherever the runs
    split them."""
    offset = 0
    cut = False
    for node in _text_nodes(p_el, "w:t"):
        text = node.text or ""
        node_start, node_end = offset, offset + len(text)
        offset = node_end
        kept = [ch for i, ch in enumerate(text, node_start)
                if not any(s.start <= i < s.end for s in spans)]
        if len(kept) != len(text):
            _set_text(node, _MULTI_SPACE_RE.sub(" ", "".join(kept)))
            cut = True
    if cut:
        _trim_trailing_space(p_el)


def _trim_trailing_space(p_el: BaseOxmlElement) -> None:
    """Drop the whitespace a cut at the end of the paragraph left behind
    ("(FAX) "), across however many trailing `w:t` nodes hold it."""
    for node in reversed(_text_nodes(p_el, "w:t")):
        text = node.text or ""
        _set_text(node, text.rstrip())
        if text.strip():
            return


def _drop_deleted(p_el: BaseOxmlElement, spans: list[_Leak]) -> None:
    """Remove the tracked deletions whose `w:delText` the spans touch: the
    struck-through value goes, the deletion around it with it."""
    offset = 0
    doomed = []
    for node in _text_nodes(p_el, "w:delText"):
        text = node.text or ""
        node_start, node_end = offset, offset + len(text)
        offset = node_end
        if any(s.start < node_end and node_start < s.end for s in spans):
            run = node.getparent()
            doomed.append(run.getparent() if run.getparent().tag == qn("w:del") else run)
    for element in dict.fromkeys(doomed):  # one w:del can hold several touched runs
        if element.getparent() is not None:
            element.getparent().remove(element)


def _anchor_runs(paragraph: Paragraph) -> list[Run]:
    """A comment anchors on runs that are the paragraph's own children (a run
    inside `w:ins` cannot carry the range markers); add an empty one if none."""
    runs = paragraph.runs
    return [runs[0], runs[-1]] if runs else [paragraph.add_run("")]


class _Leaf(NamedTuple):
    """One paragraph the repair scans: its element, the table cell it sits in
    (None for a body paragraph), and its block's section."""
    p_el: BaseOxmlElement
    cell: BaseOxmlElement | None
    section: str | None
    where: str
    dea_label_in_block: bool


def _leaves(doc: DocumentType) -> Iterator[_Leaf]:
    """Every body paragraph and table-cell paragraph, with the section and
    gates of the block the lint reads it in. A block holding the withheld
    notice is skipped, as the lint skips it."""
    # ponytail: the lint also reads each table row JOINED ("label | value",
    # `doctor.shared._table_lines`); a hit that exists only across that joint
    # is never seen here, and stays in `remaining` rather than being cut. Not
    # observed in the trial renders. If the corpus shows it, map joined-row
    # spans back to their cells instead of guessing a whole cell.
    blocks = docx_body_blocks(doc)
    sections = _block_sections(blocks)
    children = [c for c in doc.element.body.iterchildren() if c.tag in (qn("w:p"), qn("w:tbl"))]
    for child, (kind, text), section in zip(children, blocks, sections, strict=True):
        if PII_REDACTED_NOTICE in str(text):
            continue
        where = _section_label(kind, section)
        dea_label = bool(_DEA_LABEL_PRESENT_RE.search(str(text)))
        if kind == "p":
            yield _Leaf(child, None, section, where, dea_label)
            continue
        for p_el in child.iter(qn("w:p")):
            cell = next(a for a in p_el.iterancestors(qn("w:tc")))
            yield _Leaf(p_el, cell, section, where, dea_label)


def _home_value_leaf(leaves: list[_Leaf], i: int) -> _Leaf | None:
    """The paragraph a bare "Home address:" label at `leaves[i]` reads its
    value from, as `_home_contact_value_leaked` reads the next line: the next
    paragraph of its cell, or the first paragraph of the next non-empty cell."""
    here = leaves[i]
    for leaf in leaves[i + 1:]:
        if leaf.cell is here.cell:
            return leaf
        if _paragraph_text(leaf.p_el).strip() or any(
                _paragraph_text(p).strip() for p in leaf.cell.iter(qn("w:p"))):
            return leaf
    return None


def _plan(doc: DocumentType) -> dict[int, tuple[_Leaf, list[_Leak], list[_Leak]]]:
    """Per paragraph (by position among the leaves): its leaf, the accepted
    spans and the deleted spans to remove."""
    leaves = list(_leaves(doc))
    plan: dict[int, tuple[_Leaf, list[_Leak], list[_Leak]]] = {}

    def add(i: int, accepted: list[_Leak], deleted: list[_Leak]) -> None:
        _, a, d = plan.get(i, (leaves[i], [], []))
        plan[i] = (leaves[i], a + accepted, d + deleted)

    for i, leaf in enumerate(leaves):
        text = _paragraph_text(leaf.p_el)
        accepted = _leaks(text, leaf.section, leaf.dea_label_in_block)
        deleted = []
        if leaf.section in (_PERSONAL_DATA_SECTION, _APPENDIX_SECTION):
            deleted = _deleted_leaks("".join(n.text or "" for n in _text_nodes(leaf.p_el, "w:delText")))
        if accepted or deleted:
            add(i, accepted, deleted)
        if leaf.section == _PERSONAL_DATA_SECTION and leaf.cell is not None:
            for m in _HOME_CONTACT_LABEL_RE.finditer(text):
                if text[m.end():].strip():
                    continue  # a value on the label's own line is the matcher's
                value_leaf = _home_value_leaf(leaves, i)
                if value_leaf is not None:
                    value_text = _paragraph_text(value_leaf.p_el)
                    if re.search(r"\d", value_text):
                        add(leaves.index(value_leaf),
                            [_Leak(0, len(value_text), CAT_HOME_CONTACT)], [])
    return plan


def _apply(doc: DocumentType, plan: dict[int, tuple[_Leaf, list[_Leak], list[_Leak]]]) -> list[Removal]:
    removals = []
    for leaf, accepted, deleted in plan.values():
        text = _paragraph_text(leaf.p_el)
        accepted = _merged([_with_separator(text, leak) for leak in accepted])
        deleted = _merged(deleted)
        _cut(leaf.p_el, accepted)
        _drop_deleted(leaf.p_el, deleted)
        paragraph = Paragraph(leaf.p_el, doc._body)
        if leaf.cell is None and _RESIDUE_RE.fullmatch(_paragraph_text(leaf.p_el)):
            nodes = _text_nodes(leaf.p_el, "w:t")
            if nodes:
                _set_text(nodes[0], PII_REDACTED_NOTICE)
                for node in nodes[1:]:
                    _set_text(node, "")
            else:
                paragraph.add_run(PII_REDACTED_NOTICE)
        categories = tuple(dict.fromkeys(
            c for leak in accepted + deleted for c in leak.category.split("; ")))
        doc.add_comment(_anchor_runs(paragraph),
                        text=REPAIR_COMMENT_TEMPLATE.format(categories=", ".join(categories)),
                        author=WITHHELD_COMMENT_AUTHOR)
        removals.append(Removal(leaf.where, categories))
    return removals


def _findings(doc: DocumentType) -> int:
    return len(lint_protected_data_in_output(docx_body_blocks(doc),
                                             docx_body_blocks(doc, deleted=True)))


def _save_atomically(doc: DocumentType, path: Path) -> None:
    """Write beside the document, then replace it: a failed save never leaves
    a half-written docx where the driver will upload it."""
    fd, tmp = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    try:
        doc.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def repair_protected_data(docx_path: str | Path) -> ProtectedDataRepair:
    """Remove from the document at `docx_path`, in place, the protected
    personal data `protected_data_in_output` finds in it. A document with no
    finding, or none this can locate, is not rewritten."""
    path = Path(docx_path)
    doc = Document(str(path))
    found = _findings(doc)
    if not found:
        return ProtectedDataRepair(found=0)
    plan = _plan(doc)
    if not plan:
        return ProtectedDataRepair(found=found, remaining=found)
    removals = _apply(doc, plan)
    _save_atomically(doc, path)
    return ProtectedDataRepair(found=found, removals=removals,
                               remaining=_findings(Document(str(path))))


def repair_flag_on(value: object) -> bool:
    """The REPAIR_FLAG_ENV value as both drivers read it: on only for "1"."""
    return str(value or "").strip() == "1"


def repair_and_report(docx_path: str | Path) -> ProtectedDataRepair | None:
    """`repair_protected_data`, then its report beside the document. Never
    raises: a repair that fails leaves the render as stage 6 wrote it (the
    atomic save means it is never half-written) and the run goes on with the
    findings it had. Returns None in that case. A report that cannot be
    written is logged and the repair still returned: the document is already
    repaired, and the doctor's lint re-reads it rather than the report."""
    try:
        result = repair_protected_data(docx_path)
    except Exception as exc:  # noqa: BLE001 -- the render must survive its repair
        logger.warning("Protected-data repair failed, document left as rendered: %s",
                       type(exc).__name__)
        return None
    try:
        repairs_report_path(docx_path).write_text(json.dumps(result.to_json(), indent=2))
    except OSError:
        logger.warning("Protected-data repair report not written; the repaired document stands",
                       exc_info=True)
    if result.found:
        logger.info("Protected-data repair: %d finding(s), %d paragraph(s) cut, %d left",
                    result.found, len(result.removals), result.remaining)
    return result
