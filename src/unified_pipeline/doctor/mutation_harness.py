"""Inject one known defect into a good synthetic run and see which lint catches it (#1588).

`COVERAGE.md` measures the doctor's recall on the defects real autopsies
found, so a cell no autopsy ever filled stays "unmeasured". This harness
measures every cell it has a mutation for, the same way each time:

1. Build a small synthetic CV (`SYNTHETIC_CV`, invented names only) as a
   source docx plus the stage-2, stage-3b and stage-4 artifacts a clean run
   of it would produce, and render the WCM docx with the real stage 6.
2. For each `Mutation`, apply one defect to the entry list at the stage that
   makes that kind of defect, carry it into every later artifact, re-render,
   and run the doctor.
3. A mutation is caught when the doctor reports a finding the clean run did
   not. Each mutation sits in one `COVERAGE.md` cell; the catch rate per cell
   is written to `MUTATIONS.md`, which `COVERAGE.md` links.

`test_doctor_mutation_harness.py` runs this in CI. A cell whose catch rate
falls below the one `MUTATIONS.md` records fails it; a cell at 0% is reported,
not failed. After a doctor change moves a rate up, regenerate the table:

    PYTHONPATH=src python3 -m unified_pipeline.doctor.mutation_harness

No LLM call and no network: stage 6's `call_llm` is replaced by one that
raises for the length of each render, so a render that would reach a model
fails the harness instead of costing money.
"""
import argparse
import copy
import json
import logging
import re
import tempfile
from collections.abc import Callable
from pathlib import Path
from typing import Any, NamedTuple
from unittest import mock

from docx import Document

from unified_pipeline import stage_6_word_template
from unified_pipeline.run_doctor import STATUS_SKIPPED, read_docx_blocks, run_doctor

logger = logging.getLogger(__name__)

MUTATIONS_PATH = Path(__file__).with_name("MUTATIONS.md")
UID = "MUTSYN"

#: The stages a mutation can enter at. The defect is in that stage's artifact
#: and in every later one, the way a real run carries it forward.
STAGE_3B, STAGE_4, STAGE_6 = "stage_3b", "stage_4", "stage_6"

#: Stage-2 entries carry no classification or fields; these are its keys.
_STAGE_2_KEYS = ("element_idx_start", "element_idx_end", "element_type", "text", "hierarchy")
_FIELDS = "extracted_fields"

Entry = dict[str, Any]

OWNER = {"first_name": "Avery", "middle_name": "", "last_name": "Quillfeather", "suffix": "",
         "full_name": "Avery Quillfeather", "full_name_with_credentials": "Avery Quillfeather, MD"}


class CvLine(NamedTuple):
    key: str          # how a mutation names the line; never written to an artifact
    heading: str      # the source heading the line sits under
    code: str         # the taxonomy code a clean stage 3b gives it
    text: str         # the source line, which is also the entry's text
    fields: dict[str, str]  # what a clean stage 4 extracts


#: A one-page CV with every name, institution and journal made up. Each line
#: is one entry and one record, so a mutation names exactly what it breaks.
SYNTHETIC_CV: tuple[CvLine, ...] = (
    CvLine("name", "PERSONAL DATA", "A", "Avery Quillfeather, MD",
           {"name": "Avery Quillfeather, MD"}),
    CvLine("email", "PERSONAL DATA", "A", "avery.quillfeather@example.org",
           {"email": "avery.quillfeather@example.org"}),
    CvLine("degree", "EDUCATION", "B1", "MD, Northwind State University, 2004",
           {"degree": "MD", "institution": "Northwind State University", "year": "2004"}),
    CvLine("appt_now", "APPOINTMENTS", "D1",
           "Associate Professor of Medicine, Larkspur Medical College, 2015-present",
           {"title": "Associate Professor of Medicine", "institution": "Larkspur Medical College",
            "start_date": "2015", "end_date": "present"}),
    CvLine("appt_past", "APPOINTMENTS", "D1",
           "Assistant Professor of Medicine, Larkspur Medical College, 2009-2015",
           {"title": "Assistant Professor of Medicine", "institution": "Larkspur Medical College",
            "start_date": "2009", "end_date": "2015"}),
    CvLine("award_teaching", "HONORS", "H", "Brightwater Teaching Award, Larkspur Medical College, 2018",
           {"award_name": "Brightwater Teaching Award", "granting_body": "Larkspur Medical College",
            "date": "2018"}),
    CvLine("award_prize", "HONORS", "H",
           "Cobalt Young Investigator Prize, Society of Vascular Widgets, 2011",
           {"award_name": "Cobalt Young Investigator Prize",
            "granting_body": "Society of Vascular Widgets", "date": "2011"}),
    CvLine("grant_pi", "GRANTS", "M2A",
           "R01 HL900001, Mechanisms of glimmer cell repair, National Heart Institute, 2019-2024, PI",
           {"grant_number": "R01 HL900001", "title": "Mechanisms of glimmer cell repair",
            "pi_role": "PI", "agency": "National Heart Institute",
            "start_date": "2019", "end_date": "2024"}),
    CvLine("grant_coi", "GRANTS", "M2A",
           "U54 TR900002, Regional widget outcomes network, Clinical Translation Agency, "
           "2020-2025, Co-Investigator",
           {"grant_number": "U54 TR900002", "title": "Regional widget outcomes network",
            "pi_role": "Co-Investigator", "agency": "Clinical Translation Agency",
            "start_date": "2020", "end_date": "2025"}),
    CvLine("pub_first", "PUBLICATIONS", "S1",
           "Quillfeather A, Brambleton K, Okonkwo-Hale R. Glimmer cells in vascular repair. "
           "J Vasc Widgets. 2016;12(3):101-110.",
           {"authors": "Quillfeather A, Brambleton K, Okonkwo-Hale R", "year": "2016",
            "title": "Glimmer cells in vascular repair", "journal": "J Vasc Widgets",
            "volume": "12", "issue": "3", "pages": "101-110"}),
    CvLine("pub_middle", "PUBLICATIONS", "S1",
           "Marchetti L, Quillfeather A, Thistlewood P, Vanterpool S. Outcomes of widget therapy "
           "in older adults. Ann Widget Med. 2019;44(1):55-63.",
           {"authors": "Marchetti L, Quillfeather A, Thistlewood P, Vanterpool S", "year": "2019",
            "title": "Outcomes of widget therapy in older adults", "journal": "Ann Widget Med",
            "volume": "44", "issue": "1", "pages": "55-63"}),
    CvLine("pub_last", "PUBLICATIONS", "S1",
           "Quillfeather A, Delacroix-Byrne M. A registry of glimmer cell transplants. "
           "Widget Res Lett. 2021;7:200-208.",
           {"authors": "Quillfeather A, Delacroix-Byrne M", "year": "2021",
            "title": "A registry of glimmer cell transplants", "journal": "Widget Res Lett",
            "volume": "7", "pages": "200-208"}),
    # A trainee's paper the owner supervised, not wrote: a clean run files it
    # as a mentee output (N4). #1573's defect files it as the owner's own.
    CvLine("trainee_paper", "TRAINEE PUBLICATIONS", "N4",
           "Fenwright J, Osei-Barlow T. Lattice models of widget flow. "
           "Comput Widget Sci. 2020;3:14-22.",
           {"output_type": "publication", "mentee_name": "Fenwright J",
            "title": "Lattice models of widget flow", "date": "2020"}),
)


def _source_index() -> dict[str, int]:
    """Each line's paragraph index in the source docx: a heading paragraph
    opens each new heading, then one paragraph per line."""
    index, position, heading = {}, 0, None
    for line in SYNTHETIC_CV:
        if line.heading != heading:
            position, heading = position + 1, line.heading
        index[line.key] = position
        position += 1
    return index


_SOURCE_INDEX = _source_index()


def clean_entries() -> list[Entry]:
    """The stage-4 entries a clean run of `SYNTHETIC_CV` produces."""
    return [{"element_idx_start": _SOURCE_INDEX[line.key], "element_idx_end": _SOURCE_INDEX[line.key],
             "element_type": "paragraph", "text": line.text, "hierarchy": [line.heading],
             "taxonomy_code": line.code, "taxonomy_confidence": 0.95, "classification_source": "llm",
             _FIELDS: dict(line.fields), "extraction_success": True}
            for line in SYNTHETIC_CV]


def _entry(entries: list[Entry], key: str) -> Entry:
    """The entry for `SYNTHETIC_CV` line `key`, found by its source index."""
    return next(e for e in entries if e["element_idx_start"] == _SOURCE_INDEX[key])


def _without(entries: list[Entry], *keys: str) -> list[Entry]:
    starts = {_SOURCE_INDEX[key] for key in keys}
    return [e for e in entries if e["element_idx_start"] not in starts]


def _set_fields(entries: list[Entry], key: str, **fields: str) -> list[Entry]:
    _entry(entries, key)[_FIELDS].update(fields)
    return entries


# The mutations. Each takes a deep copy of the clean entry list and returns
# the defective one; `Mutation.stage` says which artifact first carries it.

def _drop_publication(entries: list[Entry]) -> list[Entry]:
    return _without(entries, "pub_middle")


def _drop_grant(entries: list[Entry]) -> list[Entry]:
    return _without(entries, "grant_coi")


def _drop_honors_section(entries: list[Entry]) -> list[Entry]:
    return _without(entries, "award_teaching", "award_prize")


def _drop_appointment_entry(entries: list[Entry]) -> list[Entry]:
    return _without(entries, "appt_past")


def _shift_publication_year(entries: list[Entry]) -> list[Entry]:
    return _set_fields(entries, "pub_first", year="2017")


def _shift_grant_end_year(entries: list[Entry]) -> list[Entry]:
    return _set_fields(entries, "grant_pi", end_date="2026")


def _swap_pi_and_coi(entries: list[Entry]) -> list[Entry]:
    _set_fields(entries, "grant_pi", pi_role="Co-Investigator")
    return _set_fields(entries, "grant_coi", pi_role="PI")


def _award_as_appointment(entries: list[Entry]) -> list[Entry]:
    """Stage 3b files a prize as an appointment; stage 4 then extracts the
    appointment schema from it."""
    entry = _entry(entries, "award_prize")
    entry["taxonomy_code"] = "D1"
    entry[_FIELDS] = {"title": "Cobalt Young Investigator Prize",
                      "institution": "Society of Vascular Widgets", "start_date": "2011"}
    return entries


def _truncate_authors_keeping_owner(entries: list[Entry]) -> list[Entry]:
    return _set_fields(entries, "pub_middle", authors="Marchetti L, Quillfeather A")


def _truncate_authors_dropping_owner(entries: list[Entry]) -> list[Entry]:
    return _set_fields(entries, "pub_middle", authors="Marchetti L")


def _copy_grant_values_across(entries: list[Entry]) -> list[Entry]:
    """#1575: one grant's dates and role come back on another grant."""
    coi = _entry(entries, "grant_coi")[_FIELDS]
    return _set_fields(entries, "grant_pi", start_date=coi["start_date"],
                       end_date=coi["end_date"], pi_role=coi["pi_role"])


def _trainee_paper_as_owners(entries: list[Entry]) -> list[Entry]:
    """#1573: a paper the owner did not write is filed as the owner's own
    peer-reviewed article."""
    entry = _entry(entries, "trainee_paper")
    entry["taxonomy_code"] = "S1"
    entry[_FIELDS] = {"authors": "Fenwright J, Osei-Barlow T", "year": "2020",
                      "title": "Lattice models of widget flow", "journal": "Comput Widget Sci",
                      "volume": "3", "pages": "14-22"}
    return entries


def _duplicate_title_in_field(entries: list[Entry]) -> list[Entry]:
    title = _entry(entries, "pub_last")[_FIELDS]["title"]
    return _set_fields(entries, "pub_last", title=f"{title}. {title}")


class Mutation(NamedTuple):
    name: str
    defect: str      # COVERAGE.md matrix row
    level: str       # COVERAGE.md matrix column
    stage: str       # STAGE_3B, STAGE_4 or STAGE_6
    apply: Callable[[list[Entry]], list[Entry]]
    note: str        # what the defect is, for the table


MUTATIONS: tuple[Mutation, ...] = (
    Mutation("drop_publication", "missing", "record", STAGE_6, _drop_publication,
             "stage 6 renders two of the three articles"),
    Mutation("drop_grant", "missing", "record", STAGE_6, _drop_grant,
             "stage 6 renders one of the two grants"),
    Mutation("drop_honors_section", "missing", "section", STAGE_6, _drop_honors_section,
             "stage 6 renders no Honors"),
    Mutation("drop_appointment_entry", "missing", "entry", STAGE_3B, _drop_appointment_entry,
             "stage 3b loses one appointment entry stage 2 found"),
    Mutation("shift_publication_year", "wrong value", "field", STAGE_4, _shift_publication_year,
             "an article's year is one later than the CV's"),
    Mutation("shift_grant_end_year", "wrong value", "field", STAGE_4, _shift_grant_end_year,
             "a grant's end year is two later than the CV's"),
    Mutation("swap_pi_and_coi", "wrong value", "field", STAGE_4, _swap_pi_and_coi,
             "the PI grant and the co-I grant swap roles"),
    Mutation("copy_grant_values_across", "wrong value", "field", STAGE_4, _copy_grant_values_across,
             "#1575: one grant's dates and role on another grant"),
    Mutation("award_as_appointment", "wrong place", "entry", STAGE_3B, _award_as_appointment,
             "a prize filed as an appointment"),
    Mutation("truncate_authors_keeping_owner", "missing", "field", STAGE_4,
             _truncate_authors_keeping_owner, "four authors cut to the first two, owner kept"),
    Mutation("truncate_authors_dropping_owner", "missing", "field", STAGE_4,
             _truncate_authors_dropping_owner, "four authors cut to the first, owner lost"),
    Mutation("trainee_paper_as_owners", "invented", "record", STAGE_3B, _trainee_paper_as_owners,
             "#1573: a trainee's paper rendered as the owner's article"),
    Mutation("duplicate_title_in_field", "duplicated", "field", STAGE_4, _duplicate_title_in_field,
             "an article's title printed twice in its title field"),
)


class Artifacts(NamedTuple):
    """The entry list each artifact is written from."""
    stage_3b: list[Entry]
    stage_4: list[Entry]
    render: list[Entry]


def artifacts_for(mutation: Mutation | None) -> Artifacts:
    """The clean entries, with `mutation` applied from its stage onward."""
    clean = clean_entries()
    if mutation is None:
        return Artifacts(clean, clean, clean)
    bad = mutation.apply(copy.deepcopy(clean))
    if mutation.stage == STAGE_3B:
        return Artifacts(bad, bad, bad)
    if mutation.stage == STAGE_4:
        return Artifacts(clean, bad, bad)
    return Artifacts(clean, clean, bad)


def _write_source(root: Path) -> None:
    doc = Document()
    heading = None
    for line in SYNTHETIC_CV:
        if line.heading != heading:
            doc.add_paragraph(line.heading)
            heading = line.heading
        doc.add_paragraph(line.text)
    doc.save(str(root / f"{UID}.docx"))


def _write_json(root: Path, stage_dir: str, suffix: str, payload: dict[str, Any]) -> Path:
    path = root / stage_dir / f"{UID}{suffix}"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"document_uid": UID, **payload}), encoding="utf-8")
    return path


def _no_llm(*_args: object, **_kwargs: object) -> None:
    raise RuntimeError("the mutation harness makes no LLM call; the synthetic CV reached one")


def _render(root: Path, entries: list[Entry]) -> None:
    """Render `entries` with the real stage 6 into the run's stage-6 dir. The
    render input sits outside the stage-4 dir, so a stage-6 mutation leaves
    the stage-4 artifact clean."""
    render_input = _write_json(root, "render_input", "_fields.json",
                               {"cv_owner": OWNER, "entries": entries})
    out_dir = root / "stage_6_wcm_documents"
    out_dir.mkdir(parents=True, exist_ok=True)
    generator = stage_6_word_template.WCMTemplateGenerator(verbose=False)
    with mock.patch.object(stage_6_word_template, "call_llm", _no_llm):
        generator.generate(str(render_input), str(out_dir / f"{UID}_wcm.docx"),
                           research_summary_path=None, discover_original_doc=False)
    # The sidecar lands next to the docx; the render input is not an artifact.
    render_input.unlink()
    render_input.parent.rmdir()


def write_run(root: Path, mutation: Mutation | None) -> Path:
    """Write the synthetic run, clean or with `mutation`, under `root`."""
    artifacts = artifacts_for(mutation)
    root.mkdir(parents=True, exist_ok=True)
    _write_source(root)
    _write_json(root, "stage_2_entry_extraction", "_entries.json",
                {"entries": [{k: e[k] for k in _STAGE_2_KEYS} for e in clean_entries()]})
    _write_json(root, "stage_3b_classified_entries", "_classified.json",
                {"entries": [{k: v for k, v in e.items() if k != _FIELDS} for e in artifacts.stage_3b]})
    _write_json(root, "stage_4_field_extraction", "_fields.json",
                {"cv_owner": OWNER, "entries": artifacts.stage_4})
    _render(root, artifacts.render)
    return root


class Finding(NamedTuple):
    lint: str
    severity: str
    message: str


def _findings(report: dict[str, Any], *, skipped: bool) -> list[dict[str, Any]]:
    return [f for f in report["findings"] if (f.get("status") == STATUS_SKIPPED) == skipped]


def doctor_findings(root: Path) -> frozenset[Finding]:
    """The doctor's findings on the run under `root`, minus the INFO each lint
    records when an artifact it reads is absent: the synthetic run has no
    stage 1a, 1b, 4.5, 5 or 5d."""
    return frozenset(Finding(f["lint"], f["severity"], f["message"])
                     for f in _findings(run_doctor(root, UID), skipped=False))


def skipped_lints(root: Path) -> tuple[str, ...]:
    """The lints the synthetic run cannot exercise, for the table's caveat."""
    return tuple(sorted({f["lint"] for f in _findings(run_doctor(root, UID), skipped=True)}))


def rendered_lines(root: Path) -> list[str]:
    return [text for _, text in read_docx_blocks(str(root / "stage_6_wcm_documents" / f"{UID}_wcm.docx"))]


class Outcome(NamedTuple):
    mutation: Mutation
    caught_by: tuple[Finding, ...]   # findings the clean run does not have

    @property
    def caught(self) -> bool:
        return bool(self.caught_by)


class Measurement(NamedTuple):
    clean: frozenset[Finding]
    skipped: tuple[str, ...]
    outcomes: tuple[Outcome, ...]


def measure(workdir: Path) -> Measurement:
    """Run the clean synthetic CV and every mutation under `workdir`."""
    clean_root = write_run(workdir / "clean", None)
    clean = doctor_findings(clean_root)
    outcomes = []
    for mutation in MUTATIONS:
        found = doctor_findings(write_run(workdir / mutation.name, mutation))
        outcomes.append(Outcome(mutation, tuple(sorted(found - clean))))
    return Measurement(clean, skipped_lints(clean_root), tuple(outcomes))


Cell = tuple[str, str]


def cell_counts(outcomes: tuple[Outcome, ...]) -> dict[Cell, tuple[int, int]]:
    """(caught, mutations) per COVERAGE.md cell, in first-seen order."""
    counts: dict[Cell, tuple[int, int]] = {}
    for outcome in outcomes:
        cell = (outcome.mutation.defect, outcome.mutation.level)
        caught, total = counts.get(cell, (0, 0))
        counts[cell] = (caught + outcome.caught, total + 1)
    return counts


_TABLE_HEAD = """# Doctor mutation catch rates

Generated by `doctor/mutation_harness.py` (#1588); do not edit by hand. Regenerate with
`PYTHONPATH=src python3 -m unified_pipeline.doctor.mutation_harness` after a doctor change moves a rate.
`test_doctor_mutation_harness.py` fails when a cell's rate falls below the one recorded here, and when
this file is stale. A cell at 0% is a blind spot reported, not a failure.

Each mutation injects one defect into a clean synthetic run (`SYNTHETIC_CV`: 13 one-record entries,
invented names) at the stage that makes that defect, carries it into the later artifacts, renders the
docx with the real stage 6 and runs the doctor. Caught means the doctor reports a finding, at any
severity, that the clean run does not. Cells are `COVERAGE.md`'s matrix cells.
"""


def _rate(caught: int, total: int) -> str:
    return f"{round(100 * caught / total)}%"


def render_table(measurement: Measurement) -> str:
    """MUTATIONS.md's text for `measurement`."""
    lines = [_TABLE_HEAD, "## Catch rate per cell", "",
             "| cell | caught / mutations | rate |", "|---|---|---|"]
    for (defect, level), (caught, total) in cell_counts(measurement.outcomes).items():
        lines.append(f"| {defect} / {level} | {caught} / {total} | {_rate(caught, total)} |")
    lines += ["", "## Mutations", "",
              "| mutation | cell | injected at | defect | caught by |", "|---|---|---|---|---|"]
    for outcome in measurement.outcomes:
        m = outcome.mutation
        by = ", ".join(sorted({f"`{f.lint}` ({f.severity})" for f in outcome.caught_by})) or "missed"
        lines.append(f"| `{m.name}` | {m.defect} / {m.level} | {m.stage} | {m.note} | {by} |")
    lines += ["", "## Not exercised", "",
              "The synthetic run has no stage 1a, 1b, 4.5, 5 or 5d artifact, so these lints skip on every "
              "mutation and cannot catch one here: "
              + ", ".join(f"`{lint}`" for lint in measurement.skipped) + "."]
    return "\n".join(lines) + "\n"


#: One row of the per-cell table: "| <defect> / <level> | <caught> / <total> | <rate> |".
_CELL_ROW_RE = re.compile(r"^\| (?P<defect>[a-z ]+) / (?P<level>[a-z]+) \| (?P<caught>\d+) / (?P<total>\d+) \|")


def recorded_counts(text: str) -> dict[Cell, tuple[int, int]]:
    """The per-cell (caught, mutations) MUTATIONS.md records."""
    return {(m["defect"], m["level"]): (int(m["caught"]), int(m["total"]))
            for m in map(_CELL_ROW_RE.match, text.splitlines()) if m}


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Regenerate doctor/MUTATIONS.md (#1588).")
    parser.add_argument("--out", type=Path, default=MUTATIONS_PATH)
    args = parser.parse_args(argv)
    with tempfile.TemporaryDirectory() as workdir:
        table = render_table(measure(Path(workdir)))
    args.out.write_text(table, encoding="utf-8")
    logger.warning("wrote %s", args.out)


if __name__ == "__main__":
    main()
