"""A typed diff between a delivered `<uid>_wcm.docx` and the copy a reviewer
corrected, turned into doctor labels (#1587).

Every edit a reviewer makes to the delivered document is a pipeline defect the
doctor either flagged or missed. This reads both files, lines their body blocks
up, and types each difference:

    deleted       a delivered block the reviewer removed (spurious, duplicate)
    added         a block the reviewer wrote in (content the pipeline lost)
    moved         a block the reviewer cut from one place and pasted elsewhere
    value_edited  a block the reviewer kept but changed

Tracked changes (the #1167 trap: python-docx's ``paragraph.text`` skips every
run inside ``<w:ins>``). Both files are read in the ACCEPTED view: text inside
``<w:ins>`` and ``<w:moveTo>`` counts as present, text inside ``<w:del>`` and
``<w:moveFrom>`` as gone (`w:delText` is never read). So a reviewer's edit counts the same whether they
left it as a pending revision or accepted it (or typed it with tracking off).
Stage 6's own tracked insertions (enriched citations, research summaries) sit
in both files and diff as unchanged; a reviewer who REJECTS one removes it from
the accepted view, which reads as `deleted` -- the reviewer's verdict that the
insertion was wrong. A deleted paragraph MARK (two paragraphs the reviewer
joined with tracking on) is not followed: the pair still reads as two blocks,
so the join diffs as an edit of the first and a deletion or edit of the
second rather than one merge. The pending revisions in each file are counted
(`Revisions`) but never attributed: a revision's author is a person's name.

PII: the corrected file is a CV. Nothing returned here carries its text -- a
change records block positions, character counts, the section heading (only
when it is a heading of the WCM template itself, `_section_names`) and the
mapped entry index. `to_label` reduces that further to what the label store may
hold (#1587): uid, entry index and change type.

The label is the `<uid>.json` schema `scripts/doctor_vs_autopsy.py` reads, so
a reviewed run scores the doctor the same way an autopsied one does (#1586).
"""
import difflib
import functools
import re
from collections.abc import Iterable, Mapping
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING

from unified_pipeline.core.text_norm import norm
from unified_pipeline.doctor.shared import (
    RENDER_TOKEN_MIN_COUNT,
    RENDER_TOKEN_OVERLAP,
    _fields_entries,
    _long_word_tokens,
    _output_section_header,
)

if TYPE_CHECKING:
    from docx.document import Document as DocumentType
    from docx.oxml.xmlchemy import BaseOxmlElement

CHANGE_DELETED = "deleted"
CHANGE_ADDED = "added"
CHANGE_MOVED = "moved"
CHANGE_VALUE_EDITED = "value_edited"
CHANGE_TYPES = (CHANGE_DELETED, CHANGE_ADDED, CHANGE_MOVED, CHANGE_VALUE_EDITED)

#: Two blocks inside one replaced run of the diff are the same block, edited,
#: when difflib's ratio reaches this; below it they are a deletion and an
#: addition. Set on the synthetic tests' edits (a date or a title word changed
#: in a citation scores 0.9+; two unrelated citations score under 0.4).
VALUE_EDIT_MIN_RATIO = 0.6
#: A deleted and an added block anywhere in the document are one moved block
#: at this ratio: near-identical, so a block moved AND lightly edited still
#: counts as one move rather than a loss plus an invention.
MOVE_MIN_RATIO = 0.9

#: Severity is the label schema's required cleanup-effort field
#: (scripts/score_vs_autopsy.py `SEVERITY_COST` rejects anything else). A diff
#: cannot judge severity, so this is a fixed proxy by change type: content the
#: pipeline lost is the costliest to notice and restore (whole_record_lost
#: autopsy findings are high); a wrong section or spurious block costs a cut;
#: a value edit is usually one field.
# ponytail: one severity per change type. A rewritten citation and a fixed
# year are both `value_edited`/low. If the precision report needs finer cost,
# grade value_edited by before/after length delta here.
REVIEW_SEVERITY: Mapping[str, str] = {
    CHANGE_ADDED: "high",
    CHANGE_DELETED: "medium",
    CHANGE_MOVED: "medium",
    CHANGE_VALUE_EDITED: "low",
}

#: `batch` of every review-derived label, so doctor_vs_autopsy.py's by_batch
#: recall separates review labels from autopsy batches.
REVIEW_LABEL_BATCH = "review"
#: Label `class` prefix: `review_deleted`, `review_added`, ...
REVIEW_CLASS_PREFIX = "review_"

#: The Appendix stage 6 adds after the template's own sections ("T. APPENDIX",
#: doctor/lints/render.py `_APPENDIX_HEADER`); the template has no such heading.
_APPENDIX_SECTION = "APPENDIX"
_LETTER_PREFIX_RE = re.compile(r"^[A-Z]\.\s+")

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_W_T, _W_P, _W_TBL = f"{_W_NS}t", f"{_W_NS}p", f"{_W_NS}tbl"
_W_TR, _W_TC = f"{_W_NS}tr", f"{_W_NS}tc"
_W_INS, _W_DEL = f"{_W_NS}ins", f"{_W_NS}del"
_W_MOVE_FROM, _W_MOVE_TO = f"{_W_NS}moveFrom", f"{_W_NS}moveTo"
#: Ancestors whose text is gone in the accepted view.
_REMOVED_WRAPPERS = frozenset({_W_DEL, _W_MOVE_FROM})
_CELL_JOINER = " | "


@dataclass(frozen=True)
class Block:
    """One body paragraph or table, in the accepted view."""
    kind: str             # "p" or "table"
    text: str             # whitespace-collapsed accepted text; never leaves this module
    section: str | None   # the WCM template heading it sits under, or None


@dataclass(frozen=True)
class Revisions:
    """Pending tracked changes in one file's body, by element count."""
    insertions: int = 0
    deletions: int = 0
    moves: int = 0


@dataclass(frozen=True)
class DocxChange:
    """One typed difference. Positions and counts only, no document text."""
    change_type: str
    section_before: str | None
    section_after: str | None
    block_before: int | None   # index into the delivered file's non-empty blocks
    block_after: int | None    # index into the corrected file's non-empty blocks
    element_idx: int | None    # the stage-4 entry the block renders, when one maps
    before_chars: int
    after_chars: int


@dataclass
class DocxDiff:
    """Every change between the two files, plus what was read to find them."""
    changes: list[DocxChange] = field(default_factory=list)
    delivered_blocks: int = 0
    corrected_blocks: int = 0
    delivered_revisions: Revisions = Revisions()
    corrected_revisions: Revisions = Revisions()


# ------------------------------------------------------------------ reading

@functools.cache
def _section_names() -> frozenset[str]:
    """Normalised WCM section headings: the template's own, plus the Appendix.
    Only these are ever reported as a section, so an all-caps line of CV text
    (a name, a lab) that `_output_section_header` accepts is never echoed."""
    from docx import Document

    from unified_pipeline.quality_score import _TEMPLATE_DOCX_PATH
    names = {_output_section_header(p.text) for p in Document(_TEMPLATE_DOCX_PATH).paragraphs}
    return frozenset(n for n in names if n) | {norm(_APPENDIX_SECTION)}


def _section_of(text: str) -> str | None:
    """The heading a paragraph is, upper-cased and without its letter prefix,
    when it is a WCM section heading; else None."""
    name = _LETTER_PREFIX_RE.sub("", text.strip())
    return name.upper() if norm(name) in _section_names() else None


def _accepted_text(element: BaseOxmlElement) -> str:
    """`w:t` text under `element` in document order, minus every run inside a
    `w:del` or `w:moveFrom` (gone once revisions are accepted). `w:ins` and
    `w:moveTo` text is kept. `w:delText` is never read."""
    parts = []
    for node in element.iter(_W_T):
        if not any(a.tag in _REMOVED_WRAPPERS for a in node.iterancestors()):
            parts.append(node.text or "")
    return "".join(parts)


def _table_text(tbl: BaseOxmlElement) -> str:
    """One line per row, its cells' accepted text joined by `_CELL_JOINER`.
    A nested table's text reads as part of its cell."""
    rows = []
    for tr in tbl.iterchildren(_W_TR):
        cells = (" ".join(_accepted_text(p) for p in tc.iter(_W_P)) for tc in tr.iterchildren(_W_TC))
        rows.append(_CELL_JOINER.join(c for c in cells if c.strip()))
    return "\n".join(r for r in rows if r)


def _collapse(text: str) -> str:
    """Whitespace-collapsed text: a re-wrapped line is not an edit, a changed
    character (case and punctuation included) is."""
    return " ".join(text.split())


def read_blocks(doc: DocumentType) -> list[Block]:
    """The non-empty body blocks of an open document, in the accepted view,
    each tagged with the section heading above it. CViche's own note boxes
    are skipped: the submitter is told to delete them, so removing one is not
    a correction."""
    from docx.table import Table

    from unified_pipeline.stage6.formatting import is_cviche_box

    blocks: list[Block] = []
    section: str | None = None
    for child in doc.element.body.iterchildren():
        if child.tag == _W_P:
            kind, text = "p", _collapse(_accepted_text(child))
            section = _section_of(text) or section
        elif child.tag == _W_TBL:
            if is_cviche_box(Table(child, doc)):
                continue
            kind, text = "table", _collapse(_table_text(child))
        else:
            continue
        if text:
            blocks.append(Block(kind, text, section))
    return blocks


def count_revisions(doc: DocumentType) -> Revisions:
    """Pending `w:ins`, `w:del` and `w:moveFrom` elements in the body."""
    body = doc.element.body
    return Revisions(insertions=sum(1 for _ in body.iter(_W_INS)),
                     deletions=sum(1 for _ in body.iter(_W_DEL)),
                     moves=sum(1 for _ in body.iter(_W_MOVE_FROM)))


# ------------------------------------------------------------ entry mapping

@dataclass(frozen=True)
class _EntryTokens:
    element_idx: int
    tokens: frozenset[str]


def entry_tokens(stage_json: Mapping) -> list[_EntryTokens]:
    """The long-word token set of each entry of a stage-4 (or later) artifact
    that carries an element_idx_start, truncated to its source element the way
    doctor_vs_autopsy.py `as_idx_set` reads one (474.1 -> 474)."""
    out = []
    for entry in _fields_entries(dict(stage_json)):
        try:
            idx = int(float(str(entry.element_idx)))
        except (TypeError, ValueError):
            continue
        tokens = frozenset(_long_word_tokens(entry.text))
        if len(tokens) >= RENDER_TOKEN_MIN_COUNT:
            out.append(_EntryTokens(idx, tokens))
    return out


def map_entry(text: str, entries: Iterable[_EntryTokens]) -> int | None:
    """The entry a block renders: of the entries whose tokens the block
    carries at `RENDER_TOKEN_OVERLAP` or more (the render check's own bar), the
    one with the best Jaccard overlap. None when no entry qualifies, or when
    two distinct entries tie for best (ambiguous is not a location)."""
    block = frozenset(_long_word_tokens(text))
    if not block:
        return None
    scored: dict[int, float] = {}
    for entry in entries:
        shared = len(entry.tokens & block)
        if shared / len(entry.tokens) >= RENDER_TOKEN_OVERLAP:
            jaccard = shared / len(entry.tokens | block)
            scored[entry.element_idx] = max(scored.get(entry.element_idx, 0.0), jaccard)
    if not scored:
        return None
    ranked = sorted(scored.items(), key=lambda kv: kv[1], reverse=True)
    if len(ranked) > 1 and ranked[0][1] == ranked[1][1]:
        return None
    return ranked[0][0]


# ------------------------------------------------------------------ diffing

def _ratio(a: str, b: str) -> float:
    matcher = difflib.SequenceMatcher(None, a, b, autojunk=False)
    return matcher.ratio() if matcher.quick_ratio() >= VALUE_EDIT_MIN_RATIO else 0.0


@dataclass
class _Pools:
    """Unpaired block indices left after the in-place pass, for the move pass."""
    deleted: list[int] = field(default_factory=list)
    added: list[int] = field(default_factory=list)
    edited: list[tuple[int, int]] = field(default_factory=list)


def _pair_replaced(before: list[Block], after: list[Block], i_range: range,
                   j_range: range, pools: _Pools) -> None:
    """Pair one replaced run in order: each delivered block takes the first
    later corrected block of the run it matches at `VALUE_EDIT_MIN_RATIO`."""
    # ponytail: greedy, order-preserving pairing inside one replaced run. A
    # run of hundreds of blocks (a whole section rewritten) is O(n*m) ratio
    # calls; fine at CV size. If that ever shows up in a profile, cap the
    # window per delivered block.
    j_next = j_range.start
    for i in i_range:
        match = next((j for j in range(j_next, j_range.stop)
                      if _ratio(before[i].text, after[j].text) >= VALUE_EDIT_MIN_RATIO), None)
        if match is None:
            pools.deleted.append(i)
            continue
        pools.added.extend(range(j_next, match))
        pools.edited.append((i, match))
        j_next = match + 1
    pools.added.extend(range(j_next, j_range.stop))


def _align(before: list[Block], after: list[Block]) -> _Pools:
    """difflib's block alignment, with each replaced run paired in place."""
    pools = _Pools()
    matcher = difflib.SequenceMatcher(None, [b.text for b in before],
                                      [b.text for b in after], autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "delete":
            pools.deleted.extend(range(i1, i2))
        elif tag == "insert":
            pools.added.extend(range(j1, j2))
        elif tag == "replace":
            _pair_replaced(before, after, range(i1, i2), range(j1, j2), pools)
    return pools


def _match_moves(before: list[Block], after: list[Block], pools: _Pools) -> list[tuple[int, int]]:
    """Pair each deleted block with the most similar added block at
    `MOVE_MIN_RATIO` or above; both leave their pools as one move."""
    moves = []
    for i in list(pools.deleted):
        best = max(((_ratio(before[i].text, after[j].text), j) for j in pools.added), default=None)
        if best is not None and best[0] >= MOVE_MIN_RATIO:
            moves.append((i, best[1]))
            pools.deleted.remove(i)
            pools.added.remove(best[1])
    return moves


def _change(kind: str, before: list[Block], after: list[Block], i: int | None,
            j: int | None, entries: list[_EntryTokens]) -> DocxChange:
    old = before[i] if i is not None else None
    new = after[j] if j is not None else None
    located = old if old is not None else new
    return DocxChange(
        change_type=kind,
        section_before=old.section if old else None,
        section_after=new.section if new else None,
        block_before=i, block_after=j,
        element_idx=map_entry(located.text, entries) if located else None,
        before_chars=len(old.text) if old else 0,
        after_chars=len(new.text) if new else 0)


def diff_blocks(before: list[Block], after: list[Block],
                entries: list[_EntryTokens] | None = None) -> list[DocxChange]:
    """Every change from `before` (delivered) to `after` (corrected), in
    delivered-then-corrected block order. `entries` maps each change to the
    entry it renders: the delivered block for a deletion, edit or move, the
    corrected block for an addition (a restored record's own words)."""
    entries = entries or []
    pools = _align(before, after)
    moves = _match_moves(before, after, pools)
    changes = [_change(CHANGE_DELETED, before, after, i, None, entries) for i in pools.deleted]
    changes += [_change(CHANGE_VALUE_EDITED, before, after, i, j, entries) for i, j in pools.edited]
    changes += [_change(CHANGE_MOVED, before, after, i, j, entries) for i, j in moves]
    changes += [_change(CHANGE_ADDED, before, after, None, j, entries) for j in pools.added]
    order = {kind: n for n, kind in enumerate(CHANGE_TYPES)}
    big = len(before) + len(after)
    changes.sort(key=lambda c: (c.block_before if c.block_before is not None else big,
                                c.block_after if c.block_after is not None else big,
                                order[c.change_type]))
    return changes


def diff_docx(delivered: Path, corrected: Path, stage_json: Mapping | None = None) -> DocxDiff:
    """Read both files and diff them. `stage_json` is the run's stage-4 (or
    later) artifact, for entry indices; without it every `element_idx` is None.
    An unreadable file raises (python-docx's own error): a diff of a file that
    could not be opened is not "no changes"."""
    from docx import Document

    delivered_doc, corrected_doc = Document(str(delivered)), Document(str(corrected))
    before, after = read_blocks(delivered_doc), read_blocks(corrected_doc)
    entries = entry_tokens(stage_json) if stage_json is not None else []
    return DocxDiff(changes=diff_blocks(before, after, entries),
                    delivered_blocks=len(before), corrected_blocks=len(after),
                    delivered_revisions=count_revisions(delivered_doc),
                    corrected_revisions=count_revisions(corrected_doc))


# ------------------------------------------------------------------ outputs

def to_report(uid: str, diff: DocxDiff) -> dict:
    """The typed diff stored with the run (positions, counts, template section
    headings; no document text)."""
    by_type = {kind: sum(c.change_type == kind for c in diff.changes) for kind in CHANGE_TYPES}
    return {"uid": uid, "delivered_blocks": diff.delivered_blocks,
            "corrected_blocks": diff.corrected_blocks,
            "delivered_revisions": asdict(diff.delivered_revisions),
            "corrected_revisions": asdict(diff.corrected_revisions),
            "by_type": by_type, "changes": [asdict(c) for c in diff.changes]}


def to_label(uid: str, diff: DocxDiff) -> dict:
    """The run's `<uid>.json` label in doctor_vs_autopsy.py's schema: one
    finding per change, carrying only its uid-scoped id, the change type (as
    `class` and its `REVIEW_SEVERITY`) and the entry index. No section, no
    positions, no text (#1587's PII rule for the label store)."""
    findings = [{"id": f"{uid}-R{n:02d}", "class": f"{REVIEW_CLASS_PREFIX}{c.change_type}",
                 # batch_class too: doctor_vs_autopsy.py groups recall by it, never by class
                 "class_ref": None, "batch_class": c.change_type, "stage": None,
                 "severity": REVIEW_SEVERITY[c.change_type],
                 "element_idx_start": c.element_idx, "records": None, "doctor_caught": None}
                for n, c in enumerate(diff.changes, start=1)]
    return {"uid": uid, "batch": REVIEW_LABEL_BATCH, "findings": findings, "doctor_review": []}
