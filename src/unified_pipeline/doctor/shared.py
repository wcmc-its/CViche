"""Primitives every lint needs, regardless of which stage it inspects (#493).

One responsibility: the vocabulary of a finding, and the text normalisation the
lints agree on. `_finding` is used by all sixteen; the rest are the matching
helpers that more than one domain needs, so they cannot live inside any single
lint module without one domain importing another.

Nothing here imports `run_doctor`. That is what keeps the dependency a straight
line -- run_doctor and every lint module depend on this, and it depends on
neither.
"""
import functools
import re
from typing import TYPE_CHECKING, NamedTuple

from unified_pipeline.core.text_norm import norm, squash

# The render-overlap/verbatim-piece constants and `_entry_pieces` used to be
# parallel COPIES of stage6/render_check.py's own (by-name, not by-import --
# "run_doctor is optional tooling and must not become a pipeline import").
# That rule is retired in this direction only (#825): the doctor may import
# the pipeline; no stage module imports the doctor (quality_score.py imports
# two doctor leaf modules -- #820, pre-#825 -- and is not a stage). Importing
# here instead of copying is what stops the two sets drifting the way they
# already had (#810 -- both missed the same 3b fallback). RENDER_PIECE_MIN_CHARS/WINDOW,
# RENDER_TOKEN_MIN_COUNT/OVERLAP and `_entry_pieces` are used only via
# re-export below (`lints/render.py`, `lints/extraction.py`, `run_doctor.py`
# import them from here unchanged); `_RENDER_TOKEN_RE` is also used directly
# by `_long_word_tokens` below.
from unified_pipeline.stage6.render_check import (  # noqa: F401
    RENDER_PIECE_MIN_CHARS,
    RENDER_PIECE_WINDOW,
    RENDER_TOKEN_MIN_COUNT,
    RENDER_TOKEN_OVERLAP,
    _RENDER_TOKEN_RE,
    _entry_pieces,
)

if TYPE_CHECKING:
    from docx.document import Document as DocumentType


# Separator joined between output lines in the containment haystack so a
# verbatim piece can't match across two unrelated lines (a control char that
# never occurs in real CV text).
_LINE_SENTINEL = "\x00"


def _long_word_tokens(text) -> set:
    """5+-letter token set for one string (lints 5/8 render-overlap checks)."""
    return set(_RENDER_TOKEN_RE.findall(norm(text)))


@functools.cache
def _template_haystack() -> str:
    """Squashed text of every paragraph and table cell of the pristine WCM
    template, sentinel-joined like the output haystack (#744). A piece found
    here is the template's own scaffolding, not faculty content, so finding it
    in a rendered document says nothing about THIS entry rendering. The path
    is quality_score's (one template, one place); imported lazily because
    quality_score imports this module.
    """
    from docx import Document
    from unified_pipeline.quality_score import _TEMPLATE_DOCX_PATH
    doc = Document(_TEMPLATE_DOCX_PATH)
    texts = [p.text for p in doc.paragraphs]
    texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return _LINE_SENTINEL.join(s for s in map(squash, texts) if s)


def _piece_in_template(piece: str) -> bool:
    """Whether a squashed entry piece occurs inside one template line."""
    return piece in _template_haystack()


# WCM output section headers come letter-prefixed ("T. APPENDIX") or as plain
# uppercase paragraphs ("RESEARCH", "MENTORING") — stage 6 emits both forms.
_SECTION_HEADER_RE = re.compile(r"^[A-Z]\.\s+\S")


def _output_section_header(text: str) -> str | None:
    """Normalized section name when a paragraph is a WCM output section
    header (either form above), else None."""
    stripped = str(text or "").strip()
    if _SECTION_HEADER_RE.match(stripped):
        return norm(re.sub(r"^[A-Z]\.\s+", "", stripped))
    if (3 <= len(stripped) <= 60 and stripped[0].isalpha()
            and stripped == stripped.upper()
            and not any(ch.isdigit() for ch in stripped)):
        return norm(stripped)
    return None


# Whether a lint ran, or why not (#750). Consumers read this field instead of
# parsing "skipped: missing ..." out of the human-readable message.
STATUS_RAN = "ran"
STATUS_SKIPPED = "skipped"          # an input artifact is absent
STATUS_UNREADABLE = "unreadable"    # an input artifact exists but won't load
FINDING_STATUSES = (STATUS_RAN, STATUS_SKIPPED, STATUS_UNREADABLE)


def _finding(lint: str, severity: str, message: str,
             evidence: list[str] | None = None, *,
             status: str = STATUS_RAN, reason: str = "") -> dict:
    """`reason` names the input artifact(s) behind a skip; empty when ran."""
    return {"lint": lint, "severity": severity, "message": message,
            "evidence": evidence or [], "status": status, "reason": reason}


def _magnitude_severity(observed: float, threshold: float) -> str:
    """WARN when this run is in the corpus's worst quartile, else INFO.

    Severity that encodes presence cannot rank anything: `echo_paragraphs=20`
    and `echo_paragraphs=1` are the same lint and used to be the same WARN."""
    return "WARN" if observed >= threshold else "INFO"


class Haystack(NamedTuple):
    text: str    # squashed containment haystack, _LINE_SENTINEL-joined
    tokens: set  # distinctive long-word token set

def _haystacks(blocks: list[tuple[str, str]]) -> Haystack:
    """Containment haystack + distinctive-token set for the output blocks.
    Read the result via its named fields (``h.text`` / ``h.tokens``), not
    positional unpacking."""
    pieces: list[str] = []
    tokens: set = set()
    for _, text in blocks:
        for line in str(text).split("\n"):
            if line.strip():
                pieces.append(squash(line))
                tokens.update(_long_word_tokens(line))
    return Haystack(_LINE_SENTINEL.join(pieces), tokens)


# ------------------------------------------------------------ docx text views
#
# Moved here from run_doctor.py (#820 round 2): the doctor's
# `protected_data_in_output` lint and `quality_score.score_protected_data`
# must scan the SAME blocks of the same rendered document (#825 -- two gate
# lists that drift), and quality_score cannot import run_doctor (a cycle
# through lints.enrichment). python-docx is imported lazily inside each
# function, as run_doctor's `_get_docx_document` does, so the JSON-only lints
# still import cleanly where the extra is absent.

def _docx_text(element, *, with_whitespace: bool = False) -> str:
    """All ``w:t`` text under a docx element in document order. Unlike
    python-docx's ``.text``, this INCLUDES text inside tracked-change ``<w:ins>``
    runs and EXCLUDES ``<w:delText>`` — the accepted-changes view a reader sees.
    Stage 6 inserts LLM-enriched content (research summaries, reformatted
    citations) as tracked INSERTIONS, so a reader that ignores ``<w:ins>``
    under-reports what actually rendered and false-flags content as 'unrendered'
    (issue #249: M1 summaries and reformatted citations read as dropped).

    ``with_whitespace=True`` also renders the run-level whitespace elements the
    way python-docx's ``.text`` does: ``<w:tab/>`` as ``"\\t"`` and a text-wrapping
    ``<w:br/>`` or ``<w:cr/>`` as ``"\\n"`` (#461: the scorer's raw-tab check and
    its template-paragraph match need them; the lints do not, so it is off by
    default). A ``<w:tab/>`` inside ``<w:pPr><w:tabs>`` is a tab STOP
    definition, not a character, and any element inside a ``<w:del>`` is
    deleted content; both are skipped."""
    from docx.oxml.ns import qn
    if not with_whitespace:
        return "".join(node.text or "" for node in element.iter(qn("w:t")))
    w_t, w_r, w_del = qn("w:t"), qn("w:r"), qn("w:del")
    whitespace = {qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n"}
    w_type = qn("w:type")
    parts = []
    for node in element.iter(w_t, *whitespace):
        if node.tag == w_t:
            parts.append(node.text or "")
            continue
        run = node.getparent()
        if run.tag != w_r or run.getparent().tag == w_del:
            continue
        if node.tag == qn("w:br") and node.get(w_type) not in (None, "textWrapping"):
            continue  # page / column break: python-docx renders no text
        parts.append(whitespace[node.tag])
    return "".join(parts)


def _cell_text(cell) -> str:
    """Track-change-aware equivalent of ``cell.text``: the cell's own paragraphs
    (nested tables excluded, matching python-docx), including ``<w:ins>`` text."""
    return "\n".join(_docx_text(p._p) for p in cell.paragraphs)


def _table_lines(tbl) -> list[str]:
    """Text lines of one Word table: each non-empty cell, nested tables
    recursed into (cell.text never surfaces them), and every row with more
    than one non-empty cell ALSO joined as one line — a record rendered as a
    structured row (label/value cells) keeps its tokens together the way one
    source line does only in the joined view. KEEP IN SYNC with the by-name
    mirror in stage_6_word_template.py's _rendered_output_lines() (the #221
    recovery pass, PR #225): both sides must agree on what counts as
    rendered. Extra lines only ever prove presence — strictly fewer false
    'absent' verdicts, never more."""
    lines: list[str] = []
    for row in tbl.rows:
        cell_texts = []
        for cell in row.cells:
            ctext = _cell_text(cell)
            if ctext.strip():
                cell_texts.append(ctext)
                lines.append(ctext)
            for nested in cell.tables:
                lines.extend(_table_lines(nested))
        if len(cell_texts) > 1:
            lines.append(" | ".join(" ".join(t.split()) for t in cell_texts))
    return lines


def docx_body_blocks(doc: DocumentType) -> list[tuple[str, str]]:
    """Body-order blocks of an OPEN python-docx Document: ("p", text) per
    paragraph, ("table", _table_lines joined by newlines) per table. Grants
    render as one Word table per grant, so any output check must read
    tables AND paragraphs. `run_doctor.read_docx_blocks` is the by-path
    form; `quality_score` calls this on the document it already loaded."""
    from docx.oxml.ns import qn
    from docx.table import Table

    blocks: list[tuple[str, str]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(("p", _docx_text(child)))
        elif child.tag == qn("w:tbl"):
            blocks.append(("table", "\n".join(_table_lines(Table(child, doc)))))
    return blocks
