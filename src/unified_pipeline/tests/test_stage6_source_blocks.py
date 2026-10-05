"""Unit tests for stage6/source_blocks.py (#1463): WCM-format detection,
block capture, the unfilled test and the copy cleaner. Synthetic text only.
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6 import source_blocks as sb  # noqa: E402
from unified_pipeline.stage6.sections.passthrough import _passthrough_letter  # noqa: E402

WCM_HEADINGS = ["PERSONAL DATA", "EDUCATION", "PROFESSIONAL POSITIONS & EMPLOYMENT", "EMPLOYMENT STATUS",
                "LICENSURE, BOARD CERTIFICATION", "INSTITUTIONAL/HOSPITAL AFFILIATION", "HONORS, AWARDS",
                "PERCENT EFFORT AND INSTITUTIONAL RESPONSIBILITIES", "BIBLIOGRAPHY"]


def _source(tmp_path, body_after: dict[str, list], headings=WCM_HEADINGS) -> str:
    """A docx with `headings` in order; body_after[h] is a list of str
    paragraphs or list-of-rows tables written under heading h."""
    doc = Document()
    for heading in headings:
        doc.add_paragraph(heading)
        for item in body_after.get(heading, []):
            if isinstance(item, str):
                doc.add_paragraph(item)
            else:
                table = doc.add_table(rows=len(item), cols=len(item[0]))
                for r, row in enumerate(item):
                    for c, text in enumerate(row):
                        table.cell(r, c).text = text
    path = tmp_path / "source.docx"
    doc.save(str(path))
    return str(path)


def test_wcm_format_needs_four_template_sections_in_order():
    assert sb.is_wcm_format(WCM_HEADINGS[:4])
    assert not sb.is_wcm_format(WCM_HEADINGS[:3])
    assert not sb.is_wcm_format(list(reversed(WCM_HEADINGS[:4])))


def test_a_label_line_naming_the_section_is_not_its_heading():
    assert sb.is_template_heading("EMPLOYMENT STATUS", _passthrough_letter)
    assert sb.is_template_heading("H. EMPLOYMENT STATUS", _passthrough_letter)
    assert not sb.is_template_heading(
        "Current Employment Status (Please choose one, list here, delete the others):", _passthrough_letter)


def test_capture_takes_each_block_up_to_the_next_heading(tmp_path):
    path = _source(tmp_path, {
        "EMPLOYMENT STATUS": ["Name of Current Employer(s): Example Medical Center",
                              "Current Employment Status (Please choose one, list here, delete the others):",
                              "Full-time salaried by Weill Cornell"],
        "PERCENT EFFORT AND INSTITUTIONAL RESPONSIBILITIES": [[["Activity", "Effort"], ["Teaching", "20%"]]],
    })
    blocks = sb.capture_source_blocks(path, _passthrough_letter)
    assert [sb.element_text(e) for e in blocks["E"]] == [
        "Name of Current Employer(s): Example Medical Center",
        "Current Employment Status (Please choose one, list here, delete the others):",
        "Full-time salaried by Weill Cornell"]
    assert [e.tag for e in blocks["J"]] == [qn("w:tbl")]
    assert blocks["G"] == []


def test_capture_skips_a_non_wcm_source_and_a_runaway_block(tmp_path):
    assert sb.capture_source_blocks(_source(tmp_path, {}, headings=["EDUCATION", "EMPLOYMENT STATUS"]),
                                    _passthrough_letter) == {}
    long_block = [f"Line {i}" for i in range(sb.MAX_BLOCK_ELEMENTS + 1)]
    blocks = sb.capture_source_blocks(_source(tmp_path, {"EMPLOYMENT STATUS": long_block}), _passthrough_letter)
    assert "E" not in blocks and "G" in blocks


def _block(paragraphs: list[str], table: list[list[str]] | None = None) -> list:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                t.cell(r, c).text = text
    return [el for el in doc.element.body.iterchildren() if el.tag in (qn("w:p"), qn("w:tbl"))]


TEMPLATE_E = ["Name of Current Employer(s):", "Current Employment Status (Please choose one, list here, delete the others):",
              "Full-time salaried by Weill Cornell", "Voluntary (self-employed or member of a P.C.)"]
TEMPLATE_G = [["Primary Hospital Affiliation:", ""], ["Other Hospital Affiliations:", ""],
              ["Other Institutional Affiliations:", ""]]


def test_unfilled_only_when_no_line_carries_content():
    template = _block(TEMPLATE_E)
    assert sb.is_unfilled(_block(TEMPLATE_E), template)
    assert sb.is_unfilled([], template)
    # A value typed after the label, in the same line or cell, is content.
    assert not sb.is_unfilled(_block(["Name of Current Employer(s): Example Medical Center"]), template)
    # Template lines only, but the option list cut down to the choice: filled.
    assert not sb.is_unfilled(_block(TEMPLATE_E[:3]), template)


def test_a_table_with_fewer_label_rows_is_not_a_choice():
    template = _block([], TEMPLATE_G)
    assert sb.is_unfilled(_block([], TEMPLATE_G[:1]), template)
    assert not sb.is_unfilled(_block([], [["Primary Hospital Affiliation: Example Hospital", ""]]), template)


def _paragraph(inner: str):
    return parse_xml(f'<w:p {nsdecls("w", "r")}>{inner}</w:p>')


def test_clean_copy_accepts_insertions_and_drops_source_references():
    p = _paragraph(
        '<w:pPr><w:pStyle w:val="SourceStyle"/><w:numPr><w:numId w:val="7"/></w:numPr></w:pPr>'
        '<w:commentRangeStart w:id="1"/>'
        '<w:r><w:t>kept </w:t></w:r>'
        '<w:ins w:id="2" w:author="a"><w:r><w:t>inserted </w:t></w:r></w:ins>'
        '<w:del w:id="3" w:author="a"><w:r><w:delText>deleted</w:delText></w:r></w:del>'
        '<w:hyperlink r:id="rId9"><w:r><w:t>linked</w:t></w:r></w:hyperlink>'
        '<w:r><w:drawing/></w:r>'
        '<w:commentRangeEnd w:id="1"/>')
    copy = sb.clean_copy(p)
    assert sb.element_text(copy) == "kept inserted linked"
    for tag in ("w:pStyle", "w:numPr", "w:ins", "w:del", "w:hyperlink", "w:drawing", "w:commentRangeStart"):
        assert next(copy.iter(qn(tag)), None) is None, tag
    assert not any(qn("r:id") in node.attrib for node in copy.iter())
    assert sb.element_text(p) == "kept inserted linked"  # the source is untouched
