"""Behaviour tests for stage_1b_hierarchy_mapper.py (#704 live-path coverage).

Pure-Python, no LLM: stage_1b never imports an LLM client, so nothing here
stubs call_llm. Covers is_header_match (strict/loose, punctuation/case/
whitespace variants), find_header_in_sequence (found / not found / document-
order enforcement / element-type filtering), map_hierarchy_node (nested
children, unmapped child, synthetic-node handling), get_first_child_element_idx,
has_mapped_children, compute_section_boundaries (sibling boundary, last
section to doc_length, nested bounds, preamble handling), and run_stage_1b
end to end against a synthetic docx built with python-docx in tmp_path plus
a hand-written hierarchy JSON. OutputManager is redirected to tmp_path via
monkeypatch (it otherwise writes into the repo's src/unified_pipeline/outputs/
by default) -- no real docx, corpus dir, or document uid is read or written.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage1b_hierarchy_mapper_behaviour.py -p no:cacheprovider
"""

import json
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import pytest  # noqa: E402
from docx import Document  # noqa: E402

from unified_pipeline import stage_1b_hierarchy_mapper as stage1b  # noqa: E402
from unified_pipeline.core.output_manager import OutputManager as _RealOutputManager  # noqa: E402
from unified_pipeline.stage_1b_hierarchy_mapper import (  # noqa: E402
    compute_section_boundaries,
    find_header_in_sequence,
    get_first_child_element_idx,
    has_mapped_children,
    is_header_match,
    map_hierarchy_node,
    normalize_text,
)


def _redirect_output_manager(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Force stage1b's OutputManager to write under tmp_path instead of the
    real src/unified_pipeline/outputs/ (OutputManager.__init__ creates and
    writes into that directory by default when base_output_dir is omitted)."""

    def factory(input_path, base_output_dir=None):
        return _RealOutputManager(input_path, base_output_dir=base_output_dir or tmp_path)

    monkeypatch.setattr(stage1b, "OutputManager", factory)


def _bold_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True


def _elem(unified_idx: int, text: str, elem_type: str = "paragraph") -> dict:
    return {"unified_idx": unified_idx, "type": elem_type, "text": text}


# ------------------------------------------------------------- normalize_text

def test_normalize_text_strips_punctuation_case_and_whitespace():
    # normalize_text strips only [:\.,-] -- '&' is deliberately left alone,
    # which is why this asserts "& training" survives rather than vanishing.
    assert normalize_text("  Education:  & Training,  ") == "education & training"
    assert normalize_text("EDUCATION AND TRAINING") == "education and training"


# ------------------------------------------------------------- is_header_match

def test_is_header_match_exact_match_always_accepted():
    assert is_header_match("education", "education", strict=True) is True


def test_is_header_match_short_paragraph_contained_in_longer_header():
    # para_text="education" is fully contained in the (longer) expected header
    # and well within the 2x length allowance.
    assert is_header_match("education and training", "education") is True


def test_is_header_match_word_boundary_prevents_false_positive():
    # The exact scenario named in the module's own docstring: "presentations"
    # must not match inside "representations" (no word boundary before the 'p').
    assert is_header_match("presentations", "representations", strict=True) is False


def test_is_header_match_paragraph_must_be_whole_words_inside_the_header():
    # #429: "national" is a substring of "international" but not a word of it,
    # so the sibling sub-label paragraph "National" must not satisfy the
    # expected header "International" (nor "regional" satisfy "subregional").
    assert is_header_match("international", "national", strict=True) is False
    assert is_header_match("international", "national") is False
    assert is_header_match("subregional", "regional", strict=True) is False
    # Prefix cases (the trailing \b): the paragraph is the start of a longer
    # word in the header ("ms" in "msc", "committee" in "committees").
    assert is_header_match("graduates msc advisees", "ms", strict=True) is False
    assert is_header_match("state and national committees", "committee", strict=True) is False
    # A whole word inside a multi-word header still matches.
    assert is_header_match("regional and national", "national", strict=True) is True
    # Regex metacharacters inside the paragraph are literal (re.escape): an
    # unescaped "(" or "+" would mis-match or raise re.error.
    assert is_header_match("awards (selected) list and honors", "awards (selected) list", strict=True) is True
    assert is_header_match("k+ channels research", "k+ channels", strict=True) is True
    # Paragraph ends in a non-word char: a \b boundary would reject these.
    assert is_header_match("honors (selected) and awards", "honors (selected)", strict=True) is True
    assert is_header_match("research & teaching", "research &", strict=True) is True
    assert is_header_match("k+ channels", "k+", strict=True) is True


def test_is_header_match_strict_mode_rejects_long_content_paragraph():
    # A content paragraph that happens to contain the header word, but is far
    # longer than 3x the header -- strict mode (sequence matching) must reject it
    # so content isn't mistaken for a header.
    long_para = normalize_text(
        "ASHA's continuing education program for the year requires many hours "
        "of documented professional development activities annually"
    )
    assert len(long_para) > len("education") * 3
    assert is_header_match("education", long_para, strict=True) is False


def test_is_header_match_non_strict_mode_allows_longer_paragraph():
    # Same paragraph, non-strict (single-header fallback) mode: allowed as long
    # as it's under 150 chars, even though strict mode rejects it.
    long_para = normalize_text(
        "ASHA's continuing education program for the year requires many hours "
        "of documented professional development activities annually"
    )
    assert len(long_para) < 150
    assert is_header_match("education", long_para, strict=False) is True


def test_is_header_match_strict_mode_accepts_short_header_plus_context():
    # Word-boundary branch (not exact, not "contained"): the header appears
    # as a whole word inside a short paragraph that stays within 3x length.
    assert is_header_match("awards", "awards and honors", strict=True) is True


def test_is_header_match_non_strict_mode_rejects_paragraph_at_150_chars_or_more():
    long_text = normalize_text("education " + ("filler word " * 20))
    assert len(long_text) >= 150
    assert is_header_match("education", long_text, strict=False) is False


def test_is_header_match_empty_paragraph_does_not_match():
    # Every caller skips empty text before calling is_header_match, which is
    # why this has never bitten a run -- but the function itself should say no.
    assert is_header_match("education", "") is False
    assert is_header_match("education", "", strict=True) is False


# ------------------------------------------------------------- find_header_in_sequence

def test_find_header_in_sequence_matches_in_document_order():
    elements = [_elem(0, "Education"), _elem(1, "PhD, Somewhere, 2001"), _elem(2, "Awards")]
    matches = find_header_in_sequence(elements, ["Education", "Awards"], 0)
    assert matches == [("Education", 0), ("Awards", 2)]


def test_find_header_in_sequence_returns_none_when_header_missing():
    elements = [_elem(0, "Education"), _elem(1, "PhD, Somewhere, 2001")]
    assert find_header_in_sequence(elements, ["Nonexistent Section"], 0) is None


def test_find_header_in_sequence_enforces_document_order():
    # "A" appears before "B" in the document, but the sequence asks for "A"
    # then "B" in the OPPOSITE order -- once "A" is found, the search moves
    # forward past it and can never find "B" earlier in the doc, so the whole
    # sequence must fail.
    elements = [_elem(0, "B"), _elem(1, "A")]
    assert find_header_in_sequence(elements, ["A", "B"], 0) is None
    # Reordered to match true document order succeeds.
    assert find_header_in_sequence(elements, ["B", "A"], 0) == [("B", 0), ("A", 1)]


def test_find_header_in_sequence_only_matches_paragraph_and_table_header_types():
    # A table_content element with matching text must be SKIPPED; only the
    # later table_header element (correct type) should be matched.
    elements = [
        _elem(0, "Publications", elem_type="table_content"),
        _elem(1, "Publications", elem_type="table_header"),
    ]
    assert find_header_in_sequence(elements, ["Publications"], 0) == [("Publications", 1)]


def test_find_header_in_sequence_empty_sequence_returns_empty_list():
    assert find_header_in_sequence([_elem(0, "Education")], [], 0) == []


def test_find_header_in_sequence_skips_paragraph_that_normalizes_to_empty():
    # A whitespace-only paragraph is non-empty text (passes the raw elem_text
    # guard) but normalize_text collapses it to "" -- the separate
    # `if not para_text: continue` check must still skip it.
    elements = [_elem(0, "   "), _elem(1, "Awards")]
    assert find_header_in_sequence(elements, ["Awards"], 0) == [("Awards", 1)]


def test_find_header_in_sequence_duplicate_header_advances_past_prior_match():
    # The module's own docstring says sequence matching exists "since headers
    # may not be unique" -- after matching the first "Publications" at idx 0,
    # the search must resume at idx 0 + 1, not idx 0 itself. Otherwise a
    # second occurrence of the same normalized header in the sequence would
    # be re-matched against the SAME element instead of advancing to the
    # later one at idx 1.
    elements = [_elem(0, "Publications"), _elem(1, "Publications")]
    matches = find_header_in_sequence(elements, ["Publications", "Publications"], 0)
    assert matches == [("Publications", 0), ("Publications", 1)]


# ------------------------------------------------------------- map_hierarchy_node

def test_map_hierarchy_node_maps_nested_children_with_correct_indices():
    elements = [
        _elem(0, "Education"),
        _elem(1, "Doctoral Degrees"),
        _elem(2, "PhD, Somewhere, 2001"),
        _elem(3, "Awards"),
    ]
    node = {
        "text": "Education",
        "level": "H1",
        "children": [{"text": "Doctoral Degrees", "level": "H2", "children": []}],
    }
    mapped, next_idx = map_hierarchy_node(node, elements, [], 0)

    assert mapped["element_idx"] == 0
    assert mapped["synthetic"] is False
    assert len(mapped["children"]) == 1
    assert mapped["children"][0]["text"] == "Doctoral Degrees"
    assert mapped["children"][0]["element_idx"] == 1
    # next_search_idx must move past the mapped child, not just the parent.
    assert next_idx == 2


def test_map_hierarchy_node_unmapped_child_gets_none_element_idx():
    elements = [_elem(0, "Education"), _elem(1, "PhD, Somewhere, 2001")]
    node = {
        "text": "Education",
        "level": "H1",
        "children": [{"text": "Subsection Not In Document", "level": "H2", "children": []}],
    }
    mapped, next_idx = map_hierarchy_node(node, elements, [], 0)

    assert mapped["element_idx"] == 0
    assert mapped["children"][0]["element_idx"] is None
    # Parent's own index (0) plus 1 still governs next_idx since the child
    # contributed no index of its own.
    assert next_idx == 1


def test_map_hierarchy_node_synthetic_node_never_searches_the_document():
    elements = [_elem(0, "Education"), _elem(1, "Awards")]
    node = {
        "text": "Synthetic Group",
        "level": "H1",
        "text_metadata": {"synthetic": True},
        "children": [{"text": "Awards", "level": "H2", "children": []}],
    }
    mapped, next_idx = map_hierarchy_node(node, elements, [], 0)

    assert mapped["synthetic"] is True
    assert mapped["element_idx"] is None
    # The child (real header "Awards") is still searched and found normally.
    assert mapped["children"][0]["element_idx"] == 1
    assert next_idx == 2


def test_map_hierarchy_node_synthetic_parent_children_all_search_from_the_same_start():
    # "Teaching" appears twice (idx 1 and 5); "Service" sits between at 3.
    # Under a synthetic parent every child restarts at start_search_idx, so
    # Teaching maps to its FIRST occurrence (1) even though Service (3) was
    # matched before it. A real parent would search sequentially and land
    # Teaching on 5 instead.
    elements = [
        _elem(0, "x"), _elem(1, "Teaching"), _elem(2, "y"),
        _elem(3, "Service"), _elem(4, "z"), _elem(5, "Teaching"),
    ]
    children = [
        {"text": "Service", "level": "H2", "children": []},
        {"text": "Teaching", "level": "H2", "children": []},
    ]
    synthetic = {"text": "Group", "level": "H1", "text_metadata": {"synthetic": True}, "children": children}
    real = {"text": "Group", "level": "H1", "children": children}

    mapped_syn, next_syn = map_hierarchy_node(synthetic, elements, [], 0)
    assert [(c["text"], c["element_idx"]) for c in mapped_syn["children"]] == [("Service", 3), ("Teaching", 1)]
    assert next_syn == 4

    mapped_real, next_real = map_hierarchy_node(real, elements, [], 0)
    assert [(c["text"], c["element_idx"]) for c in mapped_real["children"]] == [("Service", 3), ("Teaching", 5)]
    assert next_real == 6


def test_map_hierarchy_node_paragraph_index_negative_one_marks_synthetic():
    elements = [_elem(0, "Education")]
    node = {"text": "Whatever", "level": "H1", "paragraph_index": -1, "children": []}
    mapped, _ = map_hierarchy_node(node, elements, [], 0)
    assert mapped["synthetic"] is True
    assert mapped["element_idx"] is None


def test_map_hierarchy_node_defaults_parent_path_when_omitted():
    # parent_path defaults to None -> [] internally when the caller omits it
    # (map_hierarchy_node's own entry point, called without a parent context).
    elements = [_elem(0, "Awards")]
    node = {"text": "Awards", "level": "H1", "children": []}
    mapped, next_idx = map_hierarchy_node(node, elements, start_search_idx=0)
    assert mapped["element_idx"] == 0
    assert next_idx == 1


def test_map_hierarchy_node_fallback_wraps_around_to_find_header_before_start_idx():
    # The header actually sits at idx 0, but start_search_idx=2 is PAST it and
    # nothing from idx 2 onward matches -- find_header_in_sequence (forward
    # only) fails, and the plain forward search_for_header(start_search_idx,
    # len(elements)) fails too. The module comments this as handling "Stage
    # 1a hierarchy may not match document order": it must fall back to
    # search_for_header(0, start_search_idx) and find the EARLIER match
    # rather than giving up with element_idx=None.
    elements = [_elem(0, "Education"), _elem(1, "PhD, Somewhere, 2001"), _elem(2, "Filler Text")]
    node = {"text": "Education", "level": "H1", "children": []}
    mapped, next_idx = map_hierarchy_node(node, elements, [], start_search_idx=2)
    assert mapped["element_idx"] == 0
    assert next_idx == 1


def test_map_hierarchy_node_fallback_search_skips_wrong_type_and_empty_text():
    # search_sequence includes a parent header that isn't in the document at
    # all, so find_header_in_sequence fails and map_hierarchy_node falls back
    # to search_for_header -- which must skip a wrong-typed element AND an
    # empty-text paragraph before it reaches the real "Awards" match.
    elements = [
        _elem(0, "Awards", elem_type="table_content"),
        _elem(1, ""),
        _elem(2, "Awards"),
    ]
    node = {"text": "Awards", "level": "H1", "children": []}
    mapped, next_idx = map_hierarchy_node(node, elements, ["Nonexistent Parent"], 0)
    assert mapped["element_idx"] == 2
    assert next_idx == 3


# ------------------------------------------------------------- get_first_child_element_idx

def test_get_first_child_element_idx_picks_minimum_direct_child():
    children = [{"element_idx": 5}, {"element_idx": 2}, {"element_idx": 9}]
    assert get_first_child_element_idx(children) == 2


def test_get_first_child_element_idx_recurses_into_grandchildren():
    children = [
        {"text": "Unmapped Mid", "element_idx": None, "children": [{"element_idx": 7}]},
    ]
    assert get_first_child_element_idx(children) == 7


def test_get_first_child_element_idx_returns_none_when_no_indices_exist():
    children = [{"element_idx": None, "children": [{"element_idx": None}]}]
    assert get_first_child_element_idx(children) is None
    assert get_first_child_element_idx([]) is None


# ------------------------------------------------------------- has_mapped_children

def test_has_mapped_children_true_for_direct_child():
    assert has_mapped_children({"children": [{"element_idx": 3}]}) is True


def test_has_mapped_children_true_when_only_grandchild_is_mapped():
    node = {"children": [{"element_idx": None, "children": [{"element_idx": 4}]}]}
    assert has_mapped_children(node) is True


def test_has_mapped_children_false_when_no_children_have_indices():
    assert has_mapped_children({"children": [{"element_idx": None, "children": []}]}) is False
    assert has_mapped_children({"text": "leaf"}) is False


# ------------------------------------------------------------- compute_section_boundaries

def test_compute_section_boundaries_last_section_runs_to_doc_length():
    mapped = [
        {"text": "A", "level": "H1", "element_idx": 0, "children": []},
        {"text": "B", "level": "H1", "element_idx": 3, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=7)
    last = next(s for s in sections if s["hierarchy"] == ["B"])
    # No next sibling -> end runs all the way to doc_length - 1.
    assert last["element_idx_start"] == 3
    assert last["element_idx_end"] == 6


def test_compute_section_boundaries_sibling_boundary_is_disjoint():
    mapped = [
        {"text": "A", "level": "H1", "element_idx": 0, "children": []},
        {"text": "B", "level": "H1", "element_idx": 3, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=7)
    first = next(s for s in sections if s["hierarchy"] == ["A"])
    # First sibling's end is exactly one before the second sibling's start.
    assert first["element_idx_start"] == 0
    assert first["element_idx_end"] == 2


def test_compute_section_boundaries_nested_bounds_extend_parent_to_cover_child():
    mapped = [
        {
            "text": "Education",
            "level": "H1",
            "element_idx": 0,
            "children": [{"text": "Doctoral Degrees", "level": "H2", "element_idx": 1, "children": []}],
        },
        {"text": "Awards", "level": "H1", "element_idx": 4, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=6)
    parent = next(s for s in sections if s["hierarchy"] == ["Education"])
    child = next(s for s in sections if s["hierarchy"] == ["Education", "Doctoral Degrees"])
    assert child["element_idx_start"] == 1
    assert child["element_idx_end"] == 3
    # Parent must be extended to at least cover its child's end (docstring
    # invariant #3: "parent end >= last child end").
    assert parent["element_idx_start"] == 0
    assert parent["element_idx_end"] >= child["element_idx_end"]
    assert parent["has_children"] is True


def test_compute_section_boundaries_parent_end_extends_past_next_sibling_to_cover_a_late_child():
    # A's next sibling B sits at 3, so A's naive end is 2 -- but A's second
    # child A2 lives at 9, past B. The PARENT is still extended to cover its
    # full descendant subtree (A2's repaired 9..11), not stopped at 2
    # (docstring invariant #3, fixed by #851) -- parents keep their
    # subtree-covering range even though it overlaps B (#916 point 4).
    # A1 (a LEAF) is a different story: post-#916 it is clipped to end just
    # before the next mapped header in document order, which is B at 3, not
    # A2 at 9 -- so A1 now ends at 2, not 8.
    mapped = [
        {
            "text": "A",
            "level": "H1",
            "element_idx": 0,
            "children": [
                {"text": "A1", "level": "H2", "element_idx": 1, "children": []},
                {"text": "A2", "level": "H2", "element_idx": 9, "children": []},
            ],
        },
        {"text": "B", "level": "H1", "element_idx": 3, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=12)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert by_path[("A", "A1")]["element_idx_end"] == 2
    assert by_path[("A", "A2")]["element_idx_end"] == 11
    assert by_path[("A",)]["element_idx_end"] == 11
    assert by_path[("A",)]["element_idx_end"] > by_path[("B",)]["element_idx_start"]


def test_compute_section_boundaries_parent_covers_a_child_repaired_by_the_post_process_step():
    # A's child A1 sits at 8, but A's next sibling B is at 3, so A1's naive
    # end (2) is before its own start. The post-process step repairs A1 to
    # 8..11 (up to C at 12); the parent should then cover it too.
    mapped = [
        {
            "text": "A",
            "level": "H1",
            "element_idx": 2,
            "children": [{"text": "A1", "level": "H2", "element_idx": 8, "children": []}],
        },
        {"text": "B", "level": "H1", "element_idx": 3, "children": []},
        {"text": "C", "level": "H1", "element_idx": 12, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=20)
    child = next(s for s in sections if s["hierarchy"] == ["A", "A1"])
    parent = next(s for s in sections if s["hierarchy"] == ["A"])
    assert (child["element_idx_start"], child["element_idx_end"]) == (8, 11)
    assert parent["element_idx_end"] >= child["element_idx_end"]


def test_compute_section_boundaries_unmapped_sibling_bounds_by_its_first_child():
    # Sibling "B" itself has no element_idx (unmapped header), but its child
    # "B1" does -- "A"'s end must be bounded by B1's index (first mapped
    # descendant of the next sibling), not run past it.
    mapped = [
        {"text": "A", "level": "H1", "element_idx": 0, "children": []},
        {
            "text": "B",
            "level": "H1",
            "element_idx": None,
            "children": [{"text": "B1", "level": "H2", "element_idx": 5, "children": []}],
        },
    ]
    sections = compute_section_boundaries(mapped, doc_length=9)
    a = next(s for s in sections if s["hierarchy"] == ["A"])
    b1 = next(s for s in sections if s["hierarchy"] == ["B", "B1"])
    assert a["element_idx_start"] == 0
    assert a["element_idx_end"] == 4
    assert b1["element_idx_start"] == 5
    assert b1["element_idx_end"] == 8


def test_compute_section_boundaries_out_of_order_last_section_extends_to_doc_end():
    # Same out-of-order pathology as above, but this time the buggy child is
    # the LAST section in document order (nothing starts after it) -- the
    # POST-PROCESS repair has no next_starts, so it falls back to
    # doc_length - 1. Second #851 reproducer shape: the parent ("A") must
    # still be extended to cover the repaired child via this fallback path,
    # not just the next_starts path the other repair test covers.
    mapped = [
        {
            "text": "A",
            "level": "H1",
            "element_idx": 2,
            "children": [{"text": "A1", "level": "H2", "element_idx": 8, "children": []}],
        },
        {"text": "B", "level": "H1", "element_idx": 3, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=20)
    a1 = next(s for s in sections if s["hierarchy"] == ["A", "A1"])
    a = next(s for s in sections if s["hierarchy"] == ["A"])
    assert a1["element_idx_start"] == 8
    assert a1["element_idx_end"] == 19
    assert a["element_idx_end"] >= a1["element_idx_end"]


def test_compute_section_boundaries_duplicate_named_siblings_do_not_cross_extend():
    # Two DIFFERENT top-level nodes share the hierarchy name "X"; each has
    # its own child and neither child needs post-process repair. Matching
    # descendants by hierarchy-path prefix (rather than tree identity) would
    # make the first "X" absorb the second "X" (and the "Y" sibling between
    # them) because `other["hierarchy"][:1] == ["X"]` is true for both --
    # violating invariant #2 (siblings are disjoint) even though nothing was
    # repaired (#851 F1). Each "X" must only be extended by its OWN child.
    mapped = [
        {
            "text": "X",
            "level": "H1",
            "element_idx": 0,
            "children": [{"text": "X1", "level": "H2", "element_idx": 1, "children": []}],
        },
        {"text": "Y", "level": "H1", "element_idx": 5, "children": []},
        {
            "text": "X",
            "level": "H1",
            "element_idx": 10,
            "children": [{"text": "X2", "level": "H2", "element_idx": 11, "children": []}],
        },
    ]
    sections = compute_section_boundaries(mapped, doc_length=20)
    first_x = next(s for s in sections if s["hierarchy"] == ["X"] and s["element_idx_start"] == 0)
    second_x = next(s for s in sections if s["hierarchy"] == ["X"] and s["element_idx_start"] == 10)
    y = next(s for s in sections if s["hierarchy"] == ["Y"])
    assert first_x["element_idx_end"] == 4
    assert y["element_idx_start"] == 5
    assert y["element_idx_end"] == 9
    assert second_x["element_idx_start"] == 10
    assert second_x["element_idx_end"] == 19
    # The first "X" must not swallow "Y" or the second "X".
    assert first_x["element_idx_end"] < y["element_idx_start"]
    assert y["element_idx_end"] < second_x["element_idx_start"]


def test_compute_section_boundaries_grandchild_repair_cascades_through_multiple_levels():
    # A three-level chain A > A1 > A1a where the LEAF (A1a) is the one whose
    # naive end is invalid and needs post-process repair. Extending only
    # DIRECT children (by identity) still has to cascade the repaired value
    # up through A1 to A, one level per pass, because descendants are always
    # appended to the flat section list before their ancestor (#851 F4).
    mapped = [
        {
            "text": "A",
            "level": "H1",
            "element_idx": 2,
            "children": [
                {
                    "text": "A1",
                    "level": "H2",
                    "element_idx": 4,
                    "children": [{"text": "A1a", "level": "H3", "element_idx": 8, "children": []}],
                }
            ],
        },
        {"text": "B", "level": "H1", "element_idx": 3, "children": []},
        {"text": "C", "level": "H1", "element_idx": 12, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=20)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    a1a = by_path[("A", "A1", "A1a")]
    a1 = by_path[("A", "A1")]
    a = by_path[("A",)]
    assert (a1a["element_idx_start"], a1a["element_idx_end"]) == (8, 11)
    assert a1["element_idx_end"] >= a1a["element_idx_end"]
    assert a["element_idx_end"] >= a1["element_idx_end"]


def test_compute_section_boundaries_extend_never_shrinks_a_parent_below_its_own_naive_end():
    # Out-of-order top-level list [Z@10, P@2 (child C@8), W@0]: P's own naive
    # sibling-bound (from W@0) is invalid (-1), and its merge with C's own
    # (also invalid) naive end is still invalid, so BOTH fall back to the
    # top-level default_end (doc_length - 1 = 19) -- P's own pre-repair end
    # is 19, well past its child's eventual repaired end. C's post-process
    # repair is bounded tightly by the next later start (Z@10), giving C
    # 8..9 -- much SMALLER than P's already-settled 19. The extension step
    # must take max(parent's own end, child ends), never just the child
    # ends, or the parent wrongly SHRINKS from 19 down to 9 (#851 F3).
    mapped = [
        {"text": "Z", "level": "H1", "element_idx": 10, "children": []},
        {
            "text": "P",
            "level": "H1",
            "element_idx": 2,
            "children": [{"text": "C", "level": "H2", "element_idx": 8, "children": []}],
        },
        {"text": "W", "level": "H1", "element_idx": 0, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=20)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert by_path[("P", "C")]["element_idx_end"] == 9
    assert by_path[("P",)]["element_idx_end"] == 19


def test_compute_section_boundaries_preamble_extends_personal_data_parent_with_children():
    # "Personal Data" is itself a real top-level node with a mapped child,
    # and there's a gap before it -- both the section itself and the
    # has_children ancestor-update loop (which matches itself here) must
    # move its start back to 0.
    mapped = [
        {
            "text": "Personal Data",
            "level": "H1",
            "element_idx": 3,
            "children": [{"text": "Contact", "level": "H2", "element_idx": 4, "children": []}],
        },
        {"text": "Education", "level": "H1", "element_idx": 6, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=9)
    pd = next(s for s in sections if s["hierarchy"] == ["Personal Data"])
    assert pd["element_idx_start"] == 0
    assert pd["has_children"] is True


def test_compute_section_boundaries_skips_fully_unmapped_sibling_between_real_sections():
    # A sibling with neither an element_idx NOR children must be skipped
    # entirely by both the sibling-boundary scan and the top-level traversal
    # (it contributes no section and doesn't stop the search for a real
    # bound), and an unmapped-but-has-children node further down the list
    # still finds its own bound by scanning past it too.
    mapped = [
        {"text": "A", "level": "H1", "element_idx": 0, "children": []},
        {"text": "Empty", "level": "H1", "element_idx": None, "children": []},
        {
            "text": "B",
            "level": "H1",
            "element_idx": None,
            "children": [{"text": "B1", "level": "H2", "element_idx": 5, "children": []}],
        },
        {"text": "C", "level": "H1", "element_idx": 8, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=12)
    hierarchies = [s["hierarchy"] for s in sections]
    # "Empty" produced no section of its own.
    assert ["Empty"] not in hierarchies

    a = next(s for s in sections if s["hierarchy"] == ["A"])
    b1 = next(s for s in sections if s["hierarchy"] == ["B", "B1"])
    c = next(s for s in sections if s["hierarchy"] == ["C"])
    # A's bound skips past "Empty" to find B's first mapped child.
    assert a["element_idx_end"] == 4
    # B1's bound skips past nothing further and is capped by C.
    assert b1["element_idx_start"] == 5
    assert b1["element_idx_end"] == 7
    assert c["element_idx_start"] == 8
    assert c["element_idx_end"] == 11


def test_compute_section_boundaries_sibling_with_fully_unmapped_subtree_contributes_no_bound():
    # "B" has children but NONE of them (recursively) have an element_idx --
    # get_first_child_element_idx returns None for it, so it must NOT bound
    # "A"; "A" has to keep scanning past "B" to reach "C".
    mapped = [
        {"text": "A", "level": "H1", "element_idx": 0, "children": []},
        {
            "text": "B",
            "level": "H1",
            "element_idx": None,
            "children": [{"text": "B1", "level": "H2", "element_idx": None, "children": []}],
        },
        {"text": "C", "level": "H1", "element_idx": 6, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=10)
    hierarchies = [s["hierarchy"] for s in sections]
    assert hierarchies == [["A"], ["C"]]
    a = next(s for s in sections if s["hierarchy"] == ["A"])
    assert a["element_idx_end"] == 5


def test_compute_section_boundaries_preamble_creates_synthetic_personal_data():
    mapped = [{"text": "Education", "level": "H1", "element_idx": 3, "children": []}]
    sections = compute_section_boundaries(mapped, doc_length=6)
    preamble = sections[0]
    assert preamble["hierarchy"] == ["Personal Data"]
    assert preamble["element_idx_start"] == 0
    assert preamble["element_idx_end"] == 2
    assert preamble["synthetic"] is True


def test_compute_section_boundaries_preamble_extends_existing_personal_data():
    mapped = [
        {"text": "Personal Data", "level": "H1", "element_idx": 2, "children": []},
        {"text": "Education", "level": "H1", "element_idx": 5, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=8)
    pd = next(s for s in sections if s["hierarchy"] == ["Personal Data"])
    # No second synthetic section should be created -- the real one is
    # extended backwards to 0 instead.
    assert pd["element_idx_start"] == 0
    assert len([s for s in sections if s["hierarchy"] == ["Personal Data"]]) == 1


def test_compute_section_boundaries_preamble_extends_existing_contact_information():
    # "Contact Information" is one of the non-"Personal Data" aliases in
    # _PREAMBLE_SECTION_ALIASES -- it must be recognized the same way.
    mapped = [
        {"text": "Contact Information", "level": "H1", "element_idx": 2, "children": []},
        {"text": "Education", "level": "H1", "element_idx": 5, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=8)
    ci = next(s for s in sections if s["hierarchy"] == ["Contact Information"])
    assert ci["element_idx_start"] == 0
    assert len([s for s in sections if s["hierarchy"] == ["Contact Information"]]) == 1


def test_compute_section_boundaries_no_sections_creates_single_fallback_section():
    sections = compute_section_boundaries([], doc_length=4)
    assert len(sections) == 1
    assert sections[0]["hierarchy"] == ["Personal Data"]
    assert sections[0]["element_idx_start"] == 0
    assert sections[0]["element_idx_end"] == 3
    assert sections[0]["synthetic"] is True


def test_compute_section_boundaries_empty_hierarchy_and_zero_doc_length():
    assert compute_section_boundaries([], doc_length=0) == []


# --------------------------------------------------------- _next_start_lookup (#916)

def test_next_start_lookup_skips_ties_and_returns_none_past_the_last_start():
    # Shared by `_repair_out_of_order_section_bounds` and
    # `_clip_leaves_to_next_header` (#916 review point 1): `bisect_right`
    # must skip every section sharing a start (ties), matching the strict
    # `>` scan it replaces, and return None once nothing starts later.
    # `bisect_left` would wrongly return a tied start itself instead of
    # skipping past it.
    sections = [
        {"element_idx_start": 5, "element_idx_end": 9, "has_children": False},
        {"element_idx_start": 5, "element_idx_end": 9, "has_children": False},
        {"element_idx_start": 12, "element_idx_end": 20, "has_children": False},
    ]
    lookup = stage1b._next_start_lookup(sections)
    assert lookup(5) == 12
    assert lookup(0) == 5
    assert lookup(12) is None


# ------------------------------------------------- _clip_leaves_to_next_header (#916)

def test_out_of_order_sibling_leaf_ends_before_the_next_header():
    # B is listed after A but starts before it -- the OUT_OF_ORDER_FALLBACK
    # mechanism used to give both A and B an end of doc_length - 1
    # (overlapping). Each leaf must now end just before the next header in
    # document order instead.
    mapped = [
        {"text": "A", "level": "H1", "element_idx": 5, "children": []},
        {"text": "B", "level": "H1", "element_idx": 2, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=10)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert (by_path[("A",)]["element_idx_start"], by_path[("A",)]["element_idx_end"]) == (5, 9)
    assert (by_path[("B",)]["element_idx_start"], by_path[("B",)]["element_idx_end"]) == (2, 4)


def test_sibling_listed_out_of_document_order_does_not_stretch_its_predecessor():
    # Z is listed before Y under P but starts after it (SIBLING_DISORDER) --
    # X's naive end used to reach past Y's actual start into Z's range.
    # Each child now ends at the next header in document order regardless
    # of list position; the parent P still covers the whole subtree.
    mapped = [
        {
            "text": "P",
            "level": "H1",
            "element_idx": 0,
            "children": [
                {"text": "X", "level": "H2", "element_idx": 10, "children": []},
                {"text": "Z", "level": "H2", "element_idx": 30, "children": []},
                {"text": "Y", "level": "H2", "element_idx": 20, "children": []},
            ],
        },
    ]
    sections = compute_section_boundaries(mapped, doc_length=40)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert (by_path[("P", "X")]["element_idx_start"], by_path[("P", "X")]["element_idx_end"]) == (10, 19)
    assert (by_path[("P", "Y")]["element_idx_start"], by_path[("P", "Y")]["element_idx_end"]) == (20, 29)
    assert (by_path[("P", "Z")]["element_idx_start"], by_path[("P", "Z")]["element_idx_end"]) == (30, 39)
    assert (by_path[("P",)]["element_idx_start"], by_path[("P",)]["element_idx_end"]) == (0, 39)


def test_late_personal_data_section_is_not_extended_over_earlier_sections():
    # Personal Data starts AFTER SHORT BIO, so it must not be stretched back
    # to 0 (it would swallow SHORT BIO). A synthetic preamble section covers
    # the true gap before the first real section instead (#916).
    mapped = [
        {"text": "SHORT BIO", "level": "H1", "element_idx": 3, "children": []},
        {"text": "Personal Data", "level": "H1", "element_idx": 6, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=12)
    assert sections[0]["hierarchy"] == ["Personal Data"]
    assert sections[0]["synthetic"] is True
    assert (sections[0]["element_idx_start"], sections[0]["element_idx_end"]) == (0, 2)
    by_path = {tuple(s["hierarchy"]): s for s in sections if not s.get("synthetic")}
    assert (by_path[("SHORT BIO",)]["element_idx_start"], by_path[("SHORT BIO",)]["element_idx_end"]) == (3, 5)
    assert (by_path[("Personal Data",)]["element_idx_start"], by_path[("Personal Data",)]["element_idx_end"]) == (6, 11)


def test_first_personal_data_section_still_absorbs_the_preamble():
    # When Personal Data really is the first section in document order,
    # existing behaviour is unchanged: it absorbs the preamble by extending
    # backwards to 0, and no synthetic section is created.
    mapped = [
        {"text": "Personal Data", "level": "H1", "element_idx": 3, "children": []},
        {"text": "X", "level": "H1", "element_idx": 8, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=12)
    assert len(sections) == 2
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert (by_path[("Personal Data",)]["element_idx_start"], by_path[("Personal Data",)]["element_idx_end"]) == (0, 7)
    assert not any(s.get("synthetic") for s in sections)


def test_leaf_ranges_are_pairwise_disjoint_and_cover_the_document():
    # A 6-node tree mixing top-level out-of-order (C listed first but starts
    # late; A listed last but starts earliest), nested sibling disorder
    # (Z listed before Y under P but starts after it), and a preamble gap
    # before A -- every mechanism from #916 in one tree.
    mapped = [
        {"text": "C", "level": "H1", "element_idx": 25, "children": []},
        {
            "text": "P",
            "level": "H1",
            "element_idx": 5,
            "children": [
                {"text": "X", "level": "H2", "element_idx": 8, "children": []},
                {"text": "Z", "level": "H2", "element_idx": 20, "children": []},
                {"text": "Y", "level": "H2", "element_idx": 12, "children": []},
            ],
        },
        {"text": "A", "level": "H1", "element_idx": 3, "children": []},
    ]
    doc_length = 30
    sections = compute_section_boundaries(mapped, doc_length=doc_length)
    leaves = [s for s in sections if not s["has_children"]]
    parents = [s for s in sections if s["has_children"]]

    for i, a in enumerate(leaves):
        for b in leaves[i + 1:]:
            a_range = set(range(a["element_idx_start"], a["element_idx_end"] + 1))
            b_range = set(range(b["element_idx_start"], b["element_idx_end"] + 1))
            assert not (a_range & b_range), (a["hierarchy"], b["hierarchy"])

    covered = set()
    for leaf in leaves:
        covered.update(range(leaf["element_idx_start"], leaf["element_idx_end"] + 1))
    all_starts = [s["element_idx_start"] for s in sections]
    for parent in parents:
        later_starts = [s for s in all_starts if s > parent["element_idx_start"]]
        if later_starts:
            covered.update(range(parent["element_idx_start"], min(later_starts)))
    assert covered == set(range(doc_length))


def test_parent_ranges_still_cover_their_children_after_clipping():
    # P's first child C1 and a foreign top-level sibling Q sit between P's
    # start and its second child C2 -- if the clip step did not skip
    # has_children sections, P's own range would be clipped down to just
    # before C1, same as a leaf. Parents must keep covering every child
    # they have, clip or no clip (#916 point 4; invariant #3 from #851).
    mapped = [
        {
            "text": "P",
            "level": "H1",
            "element_idx": 0,
            "children": [
                {"text": "C1", "level": "H2", "element_idx": 2, "children": []},
                {"text": "C2", "level": "H2", "element_idx": 15, "children": []},
            ],
        },
        {"text": "Q", "level": "H1", "element_idx": 8, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=20)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert by_path[("P", "C1")]["element_idx_end"] == 7
    assert by_path[("P", "C2")]["element_idx_end"] == 19
    assert by_path[("P",)]["element_idx_end"] == 19
    assert by_path[("P",)]["has_children"] is True


def test_a_parent_header_bounds_the_preceding_leaf():
    # P (a PARENT, has_children=True) is listed FIRST but starts at 6; L (a
    # leaf) is listed SECOND but starts earlier, at 2 -- the out-of-order
    # fallback gives L a pre-clip end of doc_length - 1 (11), overlapping
    # P and C. The clip must bound L against the next start in document
    # order REGARDLESS of has_children -- P@6, not C@8 -- so L ends at 5,
    # not 7. A mutant that restricts the clip's bound to has_children==False
    # sections only (#916) would let C@8 bound L instead, giving [2-7].
    mapped = [
        {
            "text": "P",
            "level": "H1",
            "element_idx": 6,
            "children": [
                {"text": "C", "level": "H2", "element_idx": 8, "children": []},
            ],
        },
        {"text": "L", "level": "H1", "element_idx": 2, "children": []},
    ]
    sections = compute_section_boundaries(mapped, doc_length=12)
    by_path = {tuple(s["hierarchy"]): s for s in sections}
    assert (by_path[("L",)]["element_idx_start"], by_path[("L",)]["element_idx_end"]) == (2, 5)


# ------------------------------------------------------------- run_stage_1b (end to end)

def test_run_stage_1b_end_to_end_flat_hierarchy(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Researcher")
    doc.add_paragraph("123 Main St")
    _bold_paragraph(doc, "EDUCATION")
    doc.add_paragraph("PhD, Somewhere University, 2001")
    _bold_paragraph(doc, "AWARDS")
    doc.add_paragraph("Best Paper Award, 2010")
    docx_path = tmp_path / "flat_cv.docx"
    doc.save(docx_path)

    hierarchy = {
        "document_uid": "TESTUID",
        "hierarchy": [
            {"text": "Education", "level": "H1", "children": []},
            {"text": "Awards", "level": "H1", "children": []},
        ],
    }
    hierarchy_path = tmp_path / "flat_hierarchy.json"
    hierarchy_path.write_text(json.dumps(hierarchy))

    output_data, output_path = stage1b.run_stage_1b(str(docx_path), str(hierarchy_path))

    # Output must actually land under tmp_path, never the repo's real outputs dir.
    assert tmp_path in output_path.parents
    assert output_path.exists()

    assert output_data["document_uid"] == "TESTUID"
    assert output_data["document_length"] == 6

    edu_node = output_data["hierarchy_with_indices"][0]
    awards_node = output_data["hierarchy_with_indices"][1]
    assert edu_node["element_idx"] == 2
    assert awards_node["element_idx"] == 4

    boundaries = {tuple(s["hierarchy"]): s for s in output_data["section_boundaries"]}
    assert boundaries[("Personal Data",)]["element_idx_start"] == 0
    assert boundaries[("Personal Data",)]["element_idx_end"] == 1
    assert boundaries[("Education",)]["element_idx_start"] == 2
    assert boundaries[("Education",)]["element_idx_end"] == 3
    assert boundaries[("Awards",)]["element_idx_start"] == 4
    assert boundaries[("Awards",)]["element_idx_end"] == 5

    # Full coverage: every element index (0..5) is claimed by exactly one section.
    assert output_data["meta"]["coverage"]["percentage"] == 100.0
    assert output_data["meta"]["coverage"]["unmapped_count"] == 0
    assert output_data["meta"]["coverage"]["synthetic_sections_added"] == 1
    assert output_data["meta"]["total_sections"] == 3
    assert output_data["meta"]["leaf_sections"] == 3

    # What's on disk must match what was returned -- the write path is real.
    on_disk = json.loads(output_path.read_text())
    assert on_disk == output_data


def test_run_stage_1b_end_to_end_nested_hierarchy_marks_parent_non_leaf(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe")
    _bold_paragraph(doc, "EDUCATION")
    _bold_paragraph(doc, "DOCTORAL DEGREES")
    doc.add_paragraph("PhD, Somewhere, 2001")
    _bold_paragraph(doc, "AWARDS")
    doc.add_paragraph("Best Paper, 2010")
    docx_path = tmp_path / "nested_cv.docx"
    doc.save(docx_path)

    hierarchy = {
        "document_uid": "NESTEDUID",
        "hierarchy": [
            {
                "text": "Education",
                "level": "H1",
                "children": [{"text": "Doctoral Degrees", "level": "H2", "children": []}],
            },
            {"text": "Awards", "level": "H1", "children": []},
        ],
    }
    hierarchy_path = tmp_path / "nested_hierarchy.json"
    hierarchy_path.write_text(json.dumps(hierarchy))

    output_data, _ = stage1b.run_stage_1b(str(docx_path), str(hierarchy_path))

    boundaries = {tuple(s["hierarchy"]): s for s in output_data["section_boundaries"]}
    assert boundaries[("Education",)]["has_children"] is True
    assert boundaries[("Education", "Doctoral Degrees")]["has_children"] is False
    # 4 sections total (Personal Data, Education, Doctoral Degrees, Awards)
    # but only 3 are leaves -- Education itself is a parent.
    assert output_data["meta"]["total_sections"] == 4
    assert output_data["meta"]["leaf_sections"] == 3


def test_run_stage_1b_no_preamble_when_first_element_is_a_mapped_header(tmp_path, monkeypatch):
    # When the very first document element IS a mapped header, there's no
    # gap before it -- no synthetic "Personal Data" section should be added.
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    _bold_paragraph(doc, "EDUCATION")
    doc.add_paragraph("PhD, Somewhere, 2001")
    docx_path = tmp_path / "no_preamble_cv.docx"
    doc.save(docx_path)

    hierarchy = {"document_uid": "NOPRE", "hierarchy": [{"text": "Education", "level": "H1", "children": []}]}
    hierarchy_path = tmp_path / "no_preamble_hierarchy.json"
    hierarchy_path.write_text(json.dumps(hierarchy))

    output_data, _ = stage1b.run_stage_1b(str(docx_path), str(hierarchy_path))

    assert output_data["meta"]["coverage"]["synthetic_sections_added"] == 0
    assert [s["hierarchy"] for s in output_data["section_boundaries"]] == [["Education"]]


def test_run_stage_1b_autodetects_hierarchy_json_when_not_given(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe")
    _bold_paragraph(doc, "EDUCATION")
    doc.add_paragraph("PhD, Somewhere, 2001")
    docx_path = tmp_path / "zzz_autodetect_test_cv.docx"
    doc.save(docx_path)

    # Pre-place the hierarchy JSON exactly where OutputManager.get_stage1_json_path()
    # will look for it (om.file_handle is derived from the docx filename).
    om = _RealOutputManager(str(docx_path), base_output_dir=tmp_path)
    stage1_path = om.get_stage1_json_path()
    stage1_path.parent.mkdir(parents=True, exist_ok=True)
    stage1_path.write_text(
        json.dumps({"document_uid": "AUTO", "hierarchy": [{"text": "Education", "level": "H1", "children": []}]})
    )

    output_data, _ = stage1b.run_stage_1b(str(docx_path))

    assert output_data["document_uid"] == "AUTO"


def test_run_stage_1b_exits_when_no_hierarchy_json_can_be_found(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe, no hierarchy available anywhere")
    docx_path = tmp_path / "zzz_no_hierarchy_cv.docx"
    doc.save(docx_path)

    with pytest.raises(SystemExit) as exc_info:
        stage1b.run_stage_1b(str(docx_path))
    assert exc_info.value.code == 1


def test_map_hierarchy_node_international_does_not_bind_to_the_national_paragraph():
    # #429: the flat 1a list carries "International" AFTER an unrelated header,
    # so the forward search misses it and the wrap-around fallback runs. It
    # used to land on the earlier "National" paragraph (substring match),
    # giving both nodes the same element_idx.
    elements = [
        _elem(0, "National"),
        _elem(1, "Table body", "table_content"),
        _elem(2, "International"),
        _elem(3, "Publications"),
    ]
    national, nxt = map_hierarchy_node({"text": "National", "level": "H1"}, elements, [], 0)
    _pubs, nxt = map_hierarchy_node({"text": "Publications", "level": "H1"}, elements, [], nxt)
    international, _ = map_hierarchy_node({"text": "International", "level": "H1"}, elements, [], nxt)
    assert national["element_idx"] == 0
    assert international["element_idx"] == 2


def test_find_header_in_sequence_prefers_the_exact_paragraph_over_an_earlier_fragment():
    # #429 residual: "Doctoral:" is whole words inside the header "Doctoral
    # advisees", so it satisfied is_header_match and won as the first match
    # even though the verbatim header paragraph comes later.
    elements = [
        _elem(0, "Mentoring"),
        _elem(1, "Doctoral:"),
        _elem(2, "one entry line"),
        _elem(3, "Doctoral advisees"),
    ]
    assert find_header_in_sequence(elements, ["Mentoring", "Doctoral advisees"]) == [
        ("Mentoring", 0),
        ("Doctoral advisees", 3),
    ]


def test_find_header_in_sequence_keeps_the_first_fragment_when_no_paragraph_is_exact():
    elements = [_elem(0, "Doctoral:"), _elem(1, "one entry line"), _elem(2, "Advisees")]
    assert find_header_in_sequence(elements, ["Doctoral advisees"]) == [("Doctoral advisees", 0)]


def test_find_header_in_sequence_first_paragraph_longer_than_the_header_still_wins():
    # The header inside a longer paragraph (a table-of-contents line) keeps the
    # old first-match behaviour: only a SHORTER fragment yields to a later exact
    # paragraph, so this change cannot pull a header past a real paragraph.
    elements = [_elem(0, "Education and training"), _elem(1, "Education")]
    assert find_header_in_sequence(elements, ["Education"]) == [("Education", 0)]


def test_map_hierarchy_node_child_binds_to_the_exact_paragraph_not_an_earlier_fragment():
    elements = [
        _elem(0, "Mentoring"),
        _elem(1, "Doctoral:"),
        _elem(2, "one entry line"),
        _elem(3, "Doctoral advisees"),
    ]
    node = {
        "text": "Mentoring",
        "level": "H1",
        "children": [{"text": "Doctoral advisees", "level": "H2"}],
    }
    mapped, _ = map_hierarchy_node(node, elements, [], 0)
    assert mapped["element_idx"] == 0
    assert mapped["children"][0]["element_idx"] == 3


def test_map_hierarchy_node_fallback_search_prefers_the_exact_paragraph():
    # The parent is not in the document, so the sequence match fails and the
    # single-header fallback (search_for_header) runs.
    elements = [_elem(0, "Doctoral:"), _elem(1, "one entry line"), _elem(2, "Doctoral advisees")]
    node = {
        "text": "Absent parent",
        "level": "H1",
        "children": [{"text": "Doctoral advisees", "level": "H2"}],
    }
    mapped, _ = map_hierarchy_node(node, elements, [], 0)
    assert mapped["children"][0]["element_idx"] == 2


# ------------------------------------------- #916 / #1178: typographic fold,
# partial-only backward bound, exact-only sequence anchors

def test_normalize_text_folds_typographic_quotes_and_apostrophes():
    assert normalize_text("Example (CONT’D)") == normalize_text("Example (CONT'D)")
    assert normalize_text("“Example” heading") == normalize_text('"Example" heading')


def test_map_hierarchy_node_page_break_heading_matches_its_typographic_line():
    # Stage 1a writes "(CONT'D)" with an ASCII apostrophe; the source line has
    # a typographic one. The node must bind to that line, not its base heading.
    elements = [
        _elem(0, "Example Group"),
        _elem(1, "body line"),
        _elem(2, "Example Group (CONT’D)"),
        _elem(3, "body line"),
    ]
    mapped, _ = map_hierarchy_node({"text": "Example Group (CONT'D)", "level": "H2"}, elements, [], 3)
    assert mapped["element_idx"] == 2


def test_backward_fallback_rejects_a_partial_match_before_the_parent():
    # "Teaching Assistant" only contains the child's text and sits before the
    # parent's own heading, so it is not the child's line.
    elements = [_elem(0, "Teaching Assistant"), _elem(1, "Service"), _elem(2, "body")]
    child, _ = map_hierarchy_node(
        {"text": "Teaching", "level": "H2"}, elements, ["Service"], 2, parent_idx=1
    )
    assert child["element_idx"] is None


def test_backward_fallback_keeps_an_exact_match_before_the_parent():
    # The same bound must not touch an exact heading line (#1178 attempt 1).
    elements = [_elem(0, "Local"), _elem(1, "Service"), _elem(2, "body")]
    child, _ = map_hierarchy_node(
        {"text": "Local", "level": "H2"}, elements, ["Service"], 2, parent_idx=1
    )
    assert child["element_idx"] == 0


def test_backward_fallback_does_not_partial_match_a_top_level_heading():
    elements = [_elem(0, "Teaching"), _elem(1, "body")]
    node = {"text": "Teaching Assignments", "level": "H2"}
    barred, _ = map_hierarchy_node(
        node, elements, ["Service"], 1, top_level_headers=frozenset({"teaching"})
    )
    allowed, _ = map_hierarchy_node(node, elements, ["Service"], 1)
    assert barred["element_idx"] is None
    assert allowed["element_idx"] == 0


def test_find_header_in_sequence_needs_an_exact_parent_anchor():
    elements = [_elem(0, "Research Notes"), _elem(1, "Beta Grants")]
    assert find_header_in_sequence(elements, ["Research", "Beta Grants"], 0) is None
    elements[0] = _elem(0, "Research")
    assert find_header_in_sequence(elements, ["Research", "Beta Grants"], 0) == [
        ("Research", 0), ("Beta Grants", 1),
    ]


def test_map_hierarchy_node_child_is_not_sent_past_a_later_partial_parent_hit():
    # Beta's own line (2) sits right at the cursor. A later "Research Notes"
    # line must not become the parent anchor that drags Beta to the second
    # "Beta Grants" at 5 (#916 sequence-match path).
    elements = [
        _elem(0, "Research"),
        _elem(1, "Alpha Grants"),
        _elem(2, "Beta Grants"),
        _elem(3, "body line"),
        _elem(4, "Research Notes"),
        _elem(5, "Beta Grants"),
    ]
    node = {
        "text": "Research",
        "level": "H1",
        "children": [
            {"text": "Alpha Grants", "level": "H2"},
            {"text": "Beta Grants", "level": "H2"},
        ],
    }
    mapped, _ = map_hierarchy_node(node, elements, [], 0)
    assert [c["element_idx"] for c in mapped["children"]] == [1, 2]


def test_run_stage_1b_passes_the_top_level_headers_to_the_backward_fallback(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)
    doc = Document()
    doc.add_paragraph("Teaching")
    doc.add_paragraph("body line")
    docx_path = tmp_path / "toplevel_wire_cv.docx"
    doc.save(docx_path)
    hierarchy = {
        "document_uid": "WIRE",
        "hierarchy": [
            {"text": "Teaching", "level": "H1", "children": []},
            {"text": "Absent Parent", "level": "H1",
             "children": [{"text": "Teaching Assignments", "level": "H2"}]},
        ],
    }
    hierarchy_path = tmp_path / "toplevel_wire_hierarchy.json"
    hierarchy_path.write_text(json.dumps(hierarchy))

    output_data, _ = stage1b.run_stage_1b(str(docx_path), str(hierarchy_path))

    child = output_data["hierarchy_with_indices"][1]["children"][0]
    assert child["element_idx"] is None


def test_map_hierarchy_node_mapped_parent_is_not_re_searched_for_the_sequence():
    # The parent is mapped at 0, so Beta must not look for a second "Research"
    # line at 4 and take the "Beta Grants" after it (#916 sequence-match path).
    elements = [
        _elem(0, "Research"),
        _elem(1, "Alpha Grants"),
        _elem(2, "Beta Grants"),
        _elem(3, "body line"),
        _elem(4, "Research"),
        _elem(5, "Beta Grants"),
    ]
    node = {
        "text": "Research",
        "level": "H1",
        "children": [
            {"text": "Alpha Grants", "level": "H2"},
            {"text": "Beta Grants", "level": "H2"},
        ],
    }
    mapped, _ = map_hierarchy_node(node, elements, [], 0)
    assert [c["element_idx"] for c in mapped["children"]] == [1, 2]


def test_fallback_prefers_an_exact_line_behind_the_cursor_to_a_forward_fragment():
    # The bare "Research" ahead is only a fragment of the header; the exact
    # line behind the cursor is the node's own.
    elements = [
        _elem(0, "Pending Research Grants"),
        _elem(1, "body line"),
        _elem(2, "Research"),
        _elem(3, "body line"),
    ]
    mapped, _ = map_hierarchy_node({"text": "Pending Research Grants", "level": "H1"}, elements, [], 1)
    assert mapped["element_idx"] == 0


def test_fallback_keeps_a_forward_partial_that_is_more_than_a_fragment():
    # "Local talks" holds the whole header, so it still beats an
    # earlier exact "Local" (that out-of-order case is not widened).
    elements = [
        _elem(0, "Local"),
        _elem(1, "body line"),
        _elem(2, "Local talks"),
    ]
    mapped, _ = map_hierarchy_node({"text": "Local", "level": "H2"}, elements, [], 1)
    assert mapped["element_idx"] == 2


def test_map_hierarchy_node_child_floor_is_the_mapped_parent_line():
    # The parent maps to line 1; the child's only match is a partial line
    # before it. The recursion must hand the child the parent's line as the
    # partial floor, not 0 (#1178).
    elements = [_elem(0, "Teaching Assistant"), _elem(1, "Service"), _elem(2, "body line")]
    node = {"text": "Service", "level": "H1", "children": [{"text": "Teaching", "level": "H2"}]}
    mapped, _ = map_hierarchy_node(node, elements, [], 0)
    assert mapped["element_idx"] == 1
    assert mapped["children"][0]["element_idx"] is None


def test_map_hierarchy_node_top_level_node_may_partial_match_a_top_level_heading():
    # The top-level exclusion protects only children; a top-level node has no
    # parent path, so it keeps the partial match.
    elements = [_elem(0, "Teaching"), _elem(1, "body line")]
    mapped, _ = map_hierarchy_node(
        {"text": "Teaching Assignments", "level": "H1"}, elements, [], 1,
        top_level_headers=frozenset({"teaching"}),
    )
    assert mapped["element_idx"] == 0


def test_map_hierarchy_node_child_is_not_bound_to_an_earlier_top_level_section():
    # #1178's own shape: a child heading must not bind onto an earlier
    # top-level section whose heading text it merely contains; the child
    # stays unmapped and its parent stays a leaf with its body.
    elements = [
        _elem(0, "Presentations"),
        _elem(1, "body line"),
        _elem(2, "Publications"),
        _elem(3, "PosterPresentations:"),
        _elem(4, "body line"),
    ]
    pubs = {
        "text": "Publications",
        "level": "H1",
        "children": [{"text": "Poster Presentations", "level": "H2"}],
    }
    tops = frozenset({"presentations", "publications"})
    first, nxt = map_hierarchy_node({"text": "Presentations", "level": "H1"}, elements, [], 0, top_level_headers=tops)
    second, _ = map_hierarchy_node(pubs, elements, [], nxt, top_level_headers=tops)
    assert first["element_idx"] == 0
    assert second["element_idx"] == 2
    assert second["children"][0]["element_idx"] is None
    sections = compute_section_boundaries([first, second], doc_length=len(elements))
    pubs_section = next(s for s in sections if s["hierarchy"] == ["Publications"])
    assert not pubs_section["has_children"]
    assert (pubs_section["element_idx_start"], pubs_section["element_idx_end"]) == (2, 4)


def test_fallback_takes_a_forward_fragment_when_no_exact_line_is_behind_the_cursor():
    # No exact "Pending Grants" line exists anywhere and no bounded partial
    # exists behind the cursor, so the forward fragment "Grants" is the node's
    # only line and must be kept.
    elements = [
        _elem(0, "Intro"),
        _elem(1, "body line"),
        _elem(2, "Research"),
        _elem(3, "Grants"),
        _elem(4, "body line"),
    ]
    node = {
        "text": "Research",
        "level": "H1",
        "children": [{"text": "Pending Grants", "level": "H2"}],
    }
    mapped, _ = map_hierarchy_node(node, elements, [], 2)
    assert mapped["element_idx"] == 2
    assert mapped["children"][0]["element_idx"] == 3
