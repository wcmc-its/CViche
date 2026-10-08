"""Wire tests for the two #575 fixes, driven through the real render paths,
plus the licensure half of #658 (bare substring label matching).

1. `_fill_licensure` formats F1 dates through `format_date_for_section`. A
   month-only issue date ("March 2019") must land in the Licensure table as
   "03/2019" — not the fabricated "03/01/2019" the old fallback produced — and
   a full-precision date must keep its day.

2. `_fill_research_support` rebuckets a grant by its own status string via
   `grant_status_rebucket_target`. An M2A grant whose status reads "In review"
   must render under Pending Funding, not Current Research Funding — the same
   symptom #210 fixed for the literal "Under review". The move's note is read
   off the rendered comment rather than off the caller's entry, because M2
   classifies on its own copies of the records and leaves the input alone.

3. `_classify_licensure_entry`'s `license_type` label check, and
   `_fill_dea_npi`'s template-row label check, used bare `'npi' in label` /
   `'dea' in label` substring tests -- so a label containing "dea" or "npi"
   as a run of characters inside an unrelated word (e.g. "Idea Number")
   misclassified. Bounded the same way #573 already bounded the raw-text
   check for "Dean".

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_licensure_render.py -p no:cacheprovider

Self-contained: no DB, no LLM calls. Loads the bundled WCM template like
test_stage6_funding_appendix.py does.
"""

import logging
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.stage6.sections.licensure import (  # noqa: E402
    KIND_DEA,
    KIND_LICENSE,
    KIND_NPI,
    LicenseRecord,
    _classify_licensure_entry,
    _resolve_licensure,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402

_LICENSE_ENTRY = {
    "text": "New York State Medical License 123456, issued March 2019",
    "extracted_fields": {
        "state": "New York",
        "license_number": "123456",
        "issue_date": "March 2019",       # month-only: no day stated
        "expiration_date": "2021-06-30",  # full precision: day must survive
    },
}


def _rendered_license_row(doc):
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            if len(cells) >= 4 and cells[0] == "New York" and cells[1] == "123456":
                return cells
    return None


def test_licensure_month_only_issue_date_is_not_given_a_day():
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)

    gen._fill_licensure([_LICENSE_ENTRY])

    cells = _rendered_license_row(gen.doc)
    assert cells is not None, "license row was not rendered into any table"
    assert cells[2] == "03/2019", f"issue date fabricated a day: {cells[2]!r}"
    assert cells[3] == "06/30/2021", f"stated day was lost: {cells[3]!r}"


_IN_REVIEW_GRANT = {
    "text": "Metacognitive Reflection Training | Role: PI | Status: In review",
    "taxonomy_code": "M2A",
    "extracted_fields": {
        "title": "Metacognitive Reflection Training Study",
        "agency": "SDRME Foundation",
        "pi_role": "Principal Investigator",
        "status": "In review",
        "start_date": "07/2025",
        "end_date": "06/2030",  # future end date: date inference must not move it
    },
}


def _body_index(doc, element):
    return list(doc.element.body).index(element)


def test_in_review_grant_renders_under_pending_funding():
    gen = WCMTemplateGenerator(verbose=False, emit_comments=True)
    gen.doc = Document(gen.template_path)

    entry = {**_IN_REVIEW_GRANT, "extracted_fields": dict(_IN_REVIEW_GRANT["extracted_fields"])}
    submitted = {**entry, "extracted_fields": dict(entry["extracted_fields"])}
    gen._fill_research_support({"M2A": [entry]})

    notes = [c["text"] for c in gen._comments if c["author"] == "Reclassification"]
    assert notes, "status rebucket left no reclassification comment"
    assert any("Pending (M2C)" in note for note in notes), notes
    assert entry == submitted, "rendering wrote back into the caller's own record"

    grant_table = None
    for table in gen.doc.tables:
        if "Metacognitive Reflection Training Study" in table._element.xml:
            grant_table = table
            break
    assert grant_table is not None, "grant table was not rendered"

    # The template orders the headers Current -> Past -> Pending, so rendering
    # after the Pending Funding paragraph proves the grant left Current.
    pending_idx = gen._find_paragraph_with_text("Pending Funding")
    assert pending_idx is not None
    pending_pos = _body_index(gen.doc, gen.doc.paragraphs[pending_idx]._element)
    table_pos = _body_index(gen.doc, grant_table._tbl)
    assert table_pos > pending_pos, (
        "grant with status 'In review' rendered before the Pending Funding "
        "header — it stayed in Current Research Funding"
    )


# ---------------------------------------------------------------------------
# 3. #658: bounded label matching in licensure classification


def test_license_type_label_embedded_dea_substring_does_not_misclassify():
    """A `license_type` value with "dea" embedded in an unrelated word must
    not route to the DEA slot. Bare `'dea' in label` matched "dea" inside
    "Idea" the same way the pre-#573 raw-text check matched it inside
    "Dean"."""
    kind = _classify_licensure_entry(
        state="New York", license_number="123456",
        license_type="Idea for renewal", original_text="")
    assert kind == KIND_LICENSE


def test_license_type_label_embedded_npi_substring_does_not_misclassify():
    kind = _classify_licensure_entry(
        state="New York", license_number="123456",
        license_type="Alnpine board license", original_text="")
    assert kind == KIND_LICENSE


def test_license_type_label_still_classifies_a_real_npi():
    """Control: a genuine whole-word "NPI" label still routes to the NPI
    slot -- the fix bounds the match, it does not remove it."""
    kind = _classify_licensure_entry(
        state="", license_number="1234567890",
        license_type="NPI", original_text="")
    assert kind == KIND_NPI


def test_license_type_label_still_classifies_a_real_dea():
    kind = _classify_licensure_entry(
        state="", license_number="AB1234567",
        license_type="DEA registration", original_text="")
    assert kind == KIND_DEA


def _dea_npi_table(gen):
    gen.doc.add_paragraph("Identifiers")
    table = gen.doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "DEA number: (optional)"
    # A row whose label contains "dea" only as a substring of an unrelated
    # word -- the finder above still locates this table via row 0's own
    # "DEA number" text, but this row must not itself be treated as the
    # DEA row.
    table.rows[1].cells[0].text = "Idea Reference Number"
    table.rows[2].cells[0].text = "NPI number: (optional)"
    return table


def test_fill_dea_npi_does_not_write_into_an_embedded_substring_row():
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    table = _dea_npi_table(gen)

    gen._fill_dea_npi("AB1234567", "1234567890")

    assert table.rows[0].cells[1].text == "AB1234567"
    assert table.rows[1].cells[1].text == "", (
        "the 'Idea Reference Number' row was written to -- 'dea' matched as "
        "a bare substring of 'Idea'"
    )
    assert table.rows[2].cells[1].text == "1234567890"


# ---------------------------------------------------------------------------
# Round 4 (#658 round 4 review): T5 (7 items, 1 and 5 are the same
# `_write_license_row` narrow-table point) and T4/T6 (label-source and
# number-shape coverage).


def _one_row_table(gen, cols):
    return gen.doc.add_table(rows=1, cols=cols)


_A_RECORD = LicenseRecord(state="NY", number="123456", issue_date="01/2020",
                          last_registration_date="01/2022")


def test_write_license_row_leaves_a_0_column_table_alone():
    """T5.1/T5.5: the column check must run BEFORE `table.add_row()`, not
    after -- otherwise a 0/1-column table gets an empty row appended even
    though the function returns False and the docstring says the table is
    'left alone'."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    table = _one_row_table(gen, cols=0)

    result = gen._write_license_row(table, _A_RECORD)

    assert result is False
    assert len(table.rows) == 1, "a row was added to a 0-column table"


def test_write_license_row_leaves_a_1_column_table_alone():
    """T5.1/T5.5, the 1-column case."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    table = _one_row_table(gen, cols=1)

    result = gen._write_license_row(table, _A_RECORD)

    assert result is False
    assert len(table.rows) == 1, "a row was added to a 1-column table"


def test_write_license_row_control_2_column_table_gets_the_row():
    """Control: a 2-column table still gets the row, proving the narrowed
    check above did not also block the case that must keep working."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    table = _one_row_table(gen, cols=2)

    result = gen._write_license_row(table, _A_RECORD)

    assert result is True
    assert len(table.rows) == 2
    assert [c.text for c in table.rows[-1].cells] == ["NY", "123456"]


# ---------------------------------------------------------------------------
# T4/T5.2: both label sources -- `license_type` (already covered above) and
# raw `original_text` -- and genuine raw-text DEA/NPI positive controls.


def test_raw_text_dean_of_students_does_not_misclassify():
    """A raw-text "Dean" mention must not misclassify via `original_text`,
    the same way `test_license_type_label_embedded_dea_substring_does_not_misclassify`
    already proves for the `license_type` field -- #573's original fix
    target, restated for the field this round's review flagged as
    untested."""
    kind = _classify_licensure_entry(
        state="NY", license_number="12345", license_type="",
        original_text="Dean of Students license 12345")
    assert kind == KIND_LICENSE


def test_raw_text_alnpine_medical_board_does_not_misclassify():
    kind = _classify_licensure_entry(
        state="", license_number="", license_type="",
        original_text="Alnpine Medical Board")
    assert kind == KIND_LICENSE


def test_raw_text_genuine_dea_still_classifies():
    """Positive control through `_classify_licensure_entry` directly."""
    kind = _classify_licensure_entry(
        state="", license_number="AB1234567", license_type="",
        original_text="DEA AB1234567")
    assert kind == KIND_DEA


def test_raw_text_genuine_npi_still_classifies():
    kind = _classify_licensure_entry(
        state="", license_number="1234567890", license_type="",
        original_text="NPI 1234567890")
    assert kind == KIND_NPI


def test_resolve_licensure_raw_text_dea_is_withheld_npi_reaches_identifiers():
    """Same positive controls, once more through `_resolve_licensure` with
    dict entries -- the label sources feed classification the same way
    whether reached from a unit-level string or a full raw entry dict. The
    NPI (public) still reaches `identifiers.npi`; the DEA (#821: withheld)
    never reaches `identifiers.dea` -- `dea_withheld` records it instead."""
    entries = [
        {"text": "DEA AB1234567",
         "extracted_fields": {"license_number": "AB1234567", "date": "2019"}},
        {"text": "NPI 1234567890",
         "extracted_fields": {"license_number": "1234567890", "date": "2020"}},
    ]
    result = _resolve_licensure(entries)
    assert result.identifiers.dea is None
    assert result.dea_withheld is True
    assert result.identifiers.npi == "1234567890"
    assert result.licenses == ()


# ---------------------------------------------------------------------------
# #1217: the agency's full name is the DEA label too. A row whose first
# column spells out "Drug Enforcement Administration" carries no "DEA" token,
# and the shape tiebreak is off because that column lands in `state`.


def test_agency_full_name_in_the_state_field_classifies_as_dea():
    kind = _classify_licensure_entry(
        state="Drug Enforcement Administration", license_number="AB1234567",
        license_type="", original_text="")
    assert kind == KIND_DEA


def test_agency_full_name_in_raw_text_classifies_as_dea_with_a_state_present():
    """The raw-text path alone, with the state field naming a jurisdiction
    (so neither the state test nor the shape tiebreak can answer)."""
    kind = _classify_licensure_entry(
        state="NY", license_number="AB1234567", license_type="",
        original_text="Drug Enforcement Administration  AB1234567  2011 - Present")
    assert kind == KIND_DEA


def test_agency_name_variants_classify_as_dea():
    for name in ("Drug Enforcement Agency", "U.S. DRUG ENFORCEMENT ADMINISTRATION",
                 "Drug  Enforcement\nAdministration"):
        kind = _classify_licensure_entry(
            state=name, license_number="", license_type="", original_text=name)
        assert kind == KIND_DEA, name


def test_a_line_that_only_mentions_drug_enforcement_is_not_dea():
    """Negative control: the label is the agency's name, not any two of its
    words -- a licence line about drug-enforcement coursework stays a licence."""
    kind = _classify_licensure_entry(
        state="NY", license_number="12345", license_type="",
        original_text="Drug Enforcement Policy course, NY licence 12345")
    assert kind == KIND_LICENSE


# ---------------------------------------------------------------------------
# ND3: a raw-text DEA label makes a DEA record only when it heads the entry,
# the entry names no other jurisdiction, or a DEA-like number is present. A
# state licence whose text merely MENTIONS a DEA registration stays a licence.


def test_a_state_licence_that_mentions_a_dea_registration_is_a_licence():
    kind = _classify_licensure_entry(
        state="NY", license_number="", license_type="",
        original_text="NY Medical Board 2001\tActive NY medical license and DEA license")
    assert kind == KIND_LICENSE


def test_a_state_licence_mentioning_dea_renders_its_row_and_withholds_nothing():
    entries = [{"text": "NY Medical Board 2001\tActive NY license and DEA license",
                "extracted_fields": {"state_country": "NY", "issue_date": "2001"}}]
    result = _resolve_licensure(entries)
    assert [r.state for r in result.licenses] == ["NY"]
    assert result.dea_withheld is False


def test_a_dea_headed_entry_with_a_state_is_still_dea():
    for text in ("DEA registration, NY", "dea registration, NY",
                 "\u2022 2005 - DEA registration, NY",
                 "Drug Enforcement Administration registration (NY)"):
        kind = _classify_licensure_entry(
            state="NY", license_number="", license_type="", original_text=text)
        assert kind == KIND_DEA, text


def test_a_dea_mention_with_a_dea_shaped_number_cell_is_still_dea():
    for number in ("AB1234567", "12345, AB12345O7"):
        kind = _classify_licensure_entry(
            state="NY", license_number=number, license_type="",
            original_text="NY license and DEA registration")
        assert kind == KIND_DEA, number


def test_a_dea_mention_with_a_dea_number_in_the_text_is_still_dea():
    for text in ("NY license 12345; DEA AB1234567",
                 "NY license 12345; DEA # AB-1234567",
                 "NY license 12345; DEA AB 123456O"):
        kind = _classify_licensure_entry(
            state="NY", license_number="12345", license_type="",
            original_text=text)
        assert kind == KIND_DEA, text


def test_a_nine_letter_word_beside_a_dea_mention_is_not_a_dea_number():
    kind = _classify_licensure_entry(
        state="NY", license_number="12345", license_type="",
        original_text="NY license 12345, Certified, and DEA registered")
    assert kind == KIND_LICENSE


def test_a_longer_token_holding_a_dea_shape_is_not_a_dea_number():
    # Both ends of the in-text shape are bounded: a token with letters glued
    # before the two-letter prefix, or with an eighth character after the
    # seven, is a longer identifier (a state licence number), not a DEA one.
    for text in ("NY license MED1234567 and DEA registered",
                 "NY license AB12345678 and DEA registered"):
        kind = _classify_licensure_entry(
            state="NY", license_number="", license_type="",
            original_text=text)
        assert kind == KIND_LICENSE, text


def test_a_dea_mention_naming_no_state_is_still_dea():
    kind = _classify_licensure_entry(
        state="", license_number="", license_type="",
        original_text="Active medical license and DEA license")
    assert kind == KIND_DEA


# ---------------------------------------------------------------------------
# T5.3: the number-shape fallback, uncovered before this round -- only
# reached when the entry carries no label and (for the NPI/DEA branch) no
# stated jurisdiction.


def test_shape_fallback_npi_digit_count():
    kind = _classify_licensure_entry(
        state="", license_number="1234567890", license_type="",
        original_text="")
    assert kind == KIND_NPI


def test_shape_fallback_dea_two_letters_seven_alphanumeric():
    kind = _classify_licensure_entry(
        state="", license_number="AB1234567", license_type="",
        original_text="")
    assert kind == KIND_DEA


def test_shape_fallback_does_not_win_when_a_state_is_present():
    """The important tiebreak guard: an NPI-shaped 10-digit number with a
    stated jurisdiction is a state licence number, not an NPI -- shape
    alone cannot distinguish them, so a stated state disables the
    shape tiebreak entirely (#573)."""
    kind = _classify_licensure_entry(
        state="NY", license_number="1234567890", license_type="",
        original_text="")
    assert kind == KIND_LICENSE


def test_shape_fallback_no_number_is_a_license():
    kind = _classify_licensure_entry(
        state="NY", license_number="", license_type="", original_text="")
    assert kind == KIND_LICENSE


# ---------------------------------------------------------------------------
# T5.4: the fused "NPI1234567890" contract -- a non-obvious regex detail
# that must not silently regress to a stricter `\bNPI\b`.


def test_fused_npi_label_is_recognized():
    """`_NPI_LABEL_RE` intentionally permits digits immediately after the
    acronym because corpus extraction can produce a fused value
    ("NPI15180546000" on corpus CV HU4DXA, per the module docstring). A
    future simplification to `\\bNPI\\b` would silently reintroduce that
    corpus failure -- pinned here so it cannot.

    `license_number` is deliberately NOT NPI/DEA-shaped (6 digits, not
    10-11): the shape fallback only runs when the label tests find nothing,
    so an NPI-shaped number here would make this test pass under
    `\\bNPI\\b` too, via the shape tiebreak, and never actually exercise the
    label regex this test exists to pin."""
    kind = _classify_licensure_entry(
        state="", license_number="123456", license_type="",
        original_text="NPI1234567890")
    assert kind == KIND_NPI


# ---------------------------------------------------------------------------
# T5.6: the single-valued identifier/duplicate policy, through
# `_resolve_licensure` (and caplog for the warning).


def test_resolve_licensure_second_different_npi_is_ignored_with_a_warning(caplog):
    entries = [
        {"text": "NPI 1111111111",
         "extracted_fields": {"license_type": "NPI", "license_number": "1111111111",
                               "date": "2020"}},
        {"text": "NPI 2222222222",
         "extracted_fields": {"license_type": "NPI", "license_number": "2222222222",
                               "date": "2018"}},
    ]
    with caplog.at_level(logging.WARNING):
        result = _resolve_licensure(entries)

    assert result.identifiers.npi == "1111111111", (
        "the later (more recent, sorted first) candidate should win, not "
        "be silently overwritten by the second"
    )
    assert any("second NPI candidate" in r.message for r in caplog.records)


def test_resolve_licensure_second_same_npi_is_silent(caplog):
    """Control: a second candidate carrying the SAME number is not a
    conflict and must not warn."""
    entries = [
        {"text": "NPI 1234567890",
         "extracted_fields": {"license_type": "NPI", "license_number": "1234567890",
                               "date": "2020"}},
        {"text": "NPI 1234567890",
         "extracted_fields": {"license_type": "NPI", "license_number": "1234567890",
                               "date": "2018"}},
    ]
    with caplog.at_level(logging.WARNING):
        result = _resolve_licensure(entries)

    assert result.identifiers.npi == "1234567890"
    assert not any("second NPI candidate" in r.message for r in caplog.records)


# ---------------------------------------------------------------------------
# T5.7: DEA/NPI entries are consumed -- they reach `identifiers`, never
# `licenses`.


def test_resolve_licensure_dea_npi_and_state_license_are_routed_correctly():
    entries = [
        {"text": "DEA AB1234567",
         "extracted_fields": {"license_type": "DEA", "license_number": "AB1234567",
                               "date": "2019"}},
        {"text": "NPI 1234567890",
         "extracted_fields": {"license_type": "NPI", "license_number": "1234567890",
                               "date": "2020"}},
        {"text": "New York State Medical License",
         "extracted_fields": {"state": "New York", "license_number": "987654",
                               "date": "2021"}},
    ]
    result = _resolve_licensure(entries)

    assert len(result.licenses) == 1
    assert result.licenses[0].state == "New York"
    assert result.licenses[0].number == "987654"
    assert result.identifiers.dea is None
    assert result.dea_withheld is True
    assert result.identifiers.npi == "1234567890"


# ---------------------------------------------------------------------------
# 4. #862: zero entries still clears the placeholder row


def _licensure_table(doc):
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "State":
            return table
    return None


def test_no_entries_clears_the_placeholder_row():
    """#862: `_fill_licensure([])` now clears the WCM template's own blank
    placeholder data row instead of returning before `_clear_table_data`
    runs, so a CV with zero F1 entries no longer ships it. The DEA/NPI
    table (a separate template table `_fill_dea_npi` writes) is untouched
    on this path, same as before."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)

    before = _licensure_table(gen.doc)
    assert before is not None
    assert len(before.rows) == 2, "template no longer ships a placeholder row"

    gen._fill_licensure([])

    after = _licensure_table(gen.doc)
    assert len(after.rows) == 1
    assert gen.stats['tables_populated'] == 0


def test_no_entries_with_a_foreign_table_is_left_alone():
    """Positive shape guard (#862): on the zero-entry path, a table located
    after the Licensure heading whose header row is NOT the licensure
    table's own is left untouched -- a template variant could put an
    unrelated table right after the same heading, and there is no
    data-driven signal to catch that when there are no entries."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen.doc.add_paragraph("F1. LICENSURE")
    table = gen.doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Organization"
    table.rows[0].cells[1].text = "Date (yyyy-yyyy)"
    table.rows[1].cells[0].text = "foreign stale row"

    gen._fill_licensure([])

    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    assert rows == [["foreign stale row", ""]]
    assert gen.stats['tables_populated'] == 0


# --- EBYSBC E14: the licence type and status have no column of their own ------

def _one_license(**fields):
    base = {"state_country": "Wrenshire", "license_number": "77001",
            "expiration_date": "2015-06-30"}
    base.update(fields)
    return [{"text": "Wrenshire licence 77001", "extracted_fields": base}]


def test_status_follows_the_last_registration_date():
    """SEKQUI-07: a relinquished licence read as one that simply lapsed."""
    (record,) = _resolve_licensure(_one_license(status="voluntarily surrendered")).licenses
    assert record.last_registration_date == "06/30/2015 (voluntarily surrendered)"
    assert record.state == "Wrenshire"


def test_status_alone_when_there_is_no_date():
    (record,) = _resolve_licensure(_one_license(status="Inactive", expiration_date="")).licenses
    assert record.last_registration_date == "Inactive"


def test_license_type_follows_the_state():
    """MYNQRA 56: a row with a country and a year said nothing about what
    the licence was."""
    (record,) = _resolve_licensure(_one_license(license_type="Pilot certificate")).licenses
    assert record.state == "Wrenshire (Pilot certificate)"
    assert record.number == "77001"


def test_a_type_the_row_already_states_is_not_repeated():
    (record,) = _resolve_licensure(
        _one_license(state_country="Wrenshire Medical License", license_type="medical license")).licenses
    assert record.state == "Wrenshire Medical License"


def test_a_blank_or_non_scalar_status_or_type_adds_nothing():
    for blank in ("", "   ", None, [], {}, ["Active"], True):
        (record,) = _resolve_licensure(_one_license(status=blank, license_type=blank)).licenses
        assert (record.state, record.last_registration_date) == ("Wrenshire", "06/30/2015"), blank


def test_status_and_type_reach_the_document(tmp_path):
    import json

    entry = {"taxonomy_code": "F1", "element_idx_start": 0, "text": "Wrenshire licence 77001",
             "extracted_fields": {"state_country": "Wrenshire", "license_number": "77001",
                                  "expiration_date": "2015-06-30", "status": "Lapsed by request",
                                  "license_type": "Pilot certificate"}}
    source, target = tmp_path / "in.json", tmp_path / "out.docx"
    source.write_text(json.dumps({"document_uid": "TESTAA", "entries": [entry]}))
    WCMTemplateGenerator(verbose=False).generate(str(source), str(target), research_summary_path=None)
    rows = [[c.text for c in row.cells] for table in Document(str(target)).tables
            for row in table.rows if row.cells and "77001" in row.cells[min(1, len(row.cells) - 1)].text]
    assert rows and rows[0][0] == "Wrenshire (Pilot certificate)"
    assert rows[0][-1] == "06/30/2015 (Lapsed by request)"


# --- #1576: the raw-text fallback carries the whole text, never a slice ------

# Over 100 characters, with a tail that the old `[:100]` cut dropped mid-word.
_LONG_UNSTRUCTURED = ("Wrenshire Board certificate in tidal-turbine maintenance for low, "
                      "medium and high flow sites, later extended to the graded-zone survey protocol")


def test_unstructured_entry_tail_reaches_the_document(tmp_path):
    """HXBPCT 145-155: an F1 entry with no state and no number printed its
    raw text in the State cell cut silently at 100 characters mid-word."""
    import json

    assert len(_LONG_UNSTRUCTURED) > 100
    entry = {"taxonomy_code": "F1", "element_idx_start": 0, "text": _LONG_UNSTRUCTURED,
             "extracted_fields": {"state_country": None, "license_number": None}}
    source, target = tmp_path / "in.json", tmp_path / "out.docx"
    source.write_text(json.dumps({"document_uid": "TESTAA", "entries": [entry]}))
    WCMTemplateGenerator(verbose=False).generate(str(source), str(target), research_summary_path=None)
    (table,) = [t for t in Document(str(target)).tables
                if any("Wrenshire Board" in c.text for r in t.rows for c in r.cells)]
    # Read every w:t, including text inside tracked insertions (w:ins).
    cells = ["".join(node.text or "" for node in cell._tc.iter()
                     if node.tag.endswith("}t"))
             for row in table.rows for cell in row.cells]
    assert _LONG_UNSTRUCTURED in cells
