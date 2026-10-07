"""The "Organization" fallback table guard must not index blindly (#544).

`_fill_memberships` falls back to searching for a table whose first cell reads
"Organization" when no table follows the section header. That fallback can
match a table belonging to a different section, so it double-checks the second
column header says "Date" before accepting it -- but the check itself used to
read `table.rows[0].cells[1]` unconditionally. A matched table with zero rows,
or exactly one column, raised IndexError instead of falling through to "no
table found".

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_memberships_table_guard.py -p no:cacheprovider
"""

import sys
from pathlib import Path

import docx

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6.sections.memberships import (
    MembershipsSection,  # noqa: E402
)


class _StubGenerator(MembershipsSection):
    """Just enough of WCMTemplateGenerator to drive the fallback branch.

    `_find_table_after_paragraph` always misses so `_fill_memberships` falls
    through to the "Organization" header search, and `_find_table_with_cell_text`
    hands back whatever table the test is probing.
    """

    def __init__(self, fallback_table):
        self.verbose = False
        self.stats = {"tables_populated": 0, "entries_inserted": 0}
        self._fallback_table = fallback_table

    def _find_header_paragraph(self, search_text):
        return 0

    def _find_table_after_paragraph(self, para_idx):
        return None

    def _find_table_with_cell_text(self, search_text):
        return self._fallback_table

    def _add_entry_comments(self, para, entry):
        pass


def _entries():
    return [{"text": "Member, Some Society", "extracted_fields": {}}]


def test_zero_row_fallback_table_is_rejected_without_indexerror():
    table = docx.Document().add_table(rows=0, cols=2)
    generator = _StubGenerator(table)

    generator._fill_memberships(_entries())  # must not raise IndexError

    assert generator.stats["tables_populated"] == 0


def test_single_column_fallback_table_is_rejected_without_indexerror():
    table = docx.Document().add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Organization"
    generator = _StubGenerator(table)

    generator._fill_memberships(_entries())  # must not raise IndexError

    assert generator.stats["tables_populated"] == 0


def test_matching_two_column_fallback_table_is_still_accepted():
    table = docx.Document().add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Organization"
    table.rows[0].cells[1].text = "Date"
    generator = _StubGenerator(table)

    generator._fill_memberships(_entries())

    assert generator.stats["tables_populated"] == 1
    assert len(table.rows) == 2  # header row plus the one entry written


if __name__ == "__main__":
    test_zero_row_fallback_table_is_rejected_without_indexerror()
    test_single_column_fallback_table_is_rejected_without_indexerror()
    test_matching_two_column_fallback_table_is_still_accepted()
    print("OK")
