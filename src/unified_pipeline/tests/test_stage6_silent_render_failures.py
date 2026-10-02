"""Regression guard for issue #547 (the silent-failure half) and the render
contracts of the three section writers #739 reworked: a mispositioned patent
table and a skipped postdoc section both used to report success, and the
mentoring/other-education/patents writers had no test pinning where their
content lands.

`patents.py:157-164` (pre-fix line numbers) caught a repositioning failure
with a bare ``except (ValueError, IndexError): pass`` and still incremented
both counters -- so a table stranded at the end of the document, instead of
under its heading, was invisible three times over: nothing logged, the stats
claimed a normal render, and the content is present so a text-only render
gate sees no change either. `postdoc_training.py:309-317` (pre-fix line
numbers) returned silently -- twice -- when the POSTDOCTORAL/TRAINING
heading or its table could not be found in the template, dropping every
postdoc entry with no record of any kind.

Both now emit `logger.warning` naming the section and the entry count,
mirroring `service.py`'s `_fill_journal_reviewing` pattern (:691-698, added
by 993642bf). Every body splice now goes through one helper, `_insert_after`
(stage6/formatting/docx.py), anchored on an element rather than a paragraph
index; a detached anchor raises `DetachedAnchorError`.

Judgement call (disclosed in the PR body): on a reposition failure, patents.py
increments a `tables_misplaced` counter instead of `tables_populated`, so the
latter keeps meaning "landed where it should have"; `entries_inserted` still
increments either way because the content is present in the document. The
cursor is advanced only when the table was placed.

Tests marked "real template" open `key_files/wcm_cv_template_faculty_october_2022_final.docx`
the way `generate()` does and assert on body order around the section
headings, so the synthetic fixtures cannot drift from the shipped structure.
Fixtures are fictional.

    python3 -m pytest src/unified_pipeline/tests/test_stage6_silent_render_failures.py -p no:cacheprovider
"""

import json
import logging
import sys
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from unified_pipeline.stage6.formatting import DetachedAnchorError, _insert_after  # noqa: E402
from unified_pipeline.stage6.sections.mentoring import (  # noqa: E402
    MenteeRecord,
    N1_HEADING,
    N2_HEADING,
    N2_INSTRUCTION,
    _looks_like_training_grant_table,
    _program_leadership_line,
    _training_grant_is_sparse,
    _training_grant_rows,
)
import unified_pipeline.stage_6_word_template as s6  # noqa: E402
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    GEO_SCOPE_FAILURE_STAT,
    RECLASSIFY_FAILURE_STAT,
    WCMTemplateGenerator,
)

MENTORING_LOGGER = 'unified_pipeline.stage6.sections.mentoring'
PATENTS_LOGGER = 'unified_pipeline.stage6.sections.patents'
POSTDOC_LOGGER = 'unified_pipeline.stage6.sections.postdoc_training'


def _new_generator() -> WCMTemplateGenerator:
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    return gen


def _template_generator() -> WCMTemplateGenerator:
    """A generator over the real WCM template, the way `generate()` opens it."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    return gen


def _warnings(caplog, logger_name: str) -> list[logging.LogRecord]:
    return [r for r in caplog.records
            if r.levelno == logging.WARNING and r.name == logger_name]


def _body_after(doc, heading_text: str, count: int) -> list[tuple[str, str]]:
    """The `count` body elements following the paragraph whose stripped text
    equals `heading_text` (case-insensitive), as ('p', text) or
    ('tbl', text of row 0 cell 1) -- a mentee's name in a mentee table, a
    patent's first value in a patent table."""
    body = list(doc.element.body)
    for para in doc.paragraphs:
        if para.text.strip().lower() == heading_text.lower():
            start = body.index(para._element)
            break
    else:
        raise AssertionError(f"no paragraph {heading_text!r} in document")
    shape = []
    for element in body[start + 1:start + 1 + count]:
        if element.tag == qn('w:sectPr'):
            break
        if element.tag == qn('w:tbl'):
            table = Table(element, doc)
            shape.append(('tbl', table.rows[0].cells[1].text))
        else:
            shape.append(('p', Paragraph(element, doc).text))
    return shape


def _mentee(code: str, name: str, **fields) -> dict:
    return {'taxonomy_code': code, 'text': f"{name} mentee entry",
            'extracted_fields': {'mentee_name': name, **fields}}


# --- shared insertion helper ---------------------------------------------------

def test_insert_after_moves_the_element_to_follow_the_anchor():
    """`add_paragraph`/`add_table` append at the end of the body; the helper
    MOVES the new element (one copy, no duplicate left behind)."""
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    doc.add_paragraph("between")
    moved = doc.add_paragraph("moved")

    _insert_after(anchor._p, moved._p)

    assert [p.text for p in doc.paragraphs] == ["anchor", "moved", "between"]


def test_insert_after_rejects_a_detached_anchor():
    """An anchor removed from the body has nowhere for a sibling to go;
    the helper says so by name rather than lxml's root-element TypeError."""
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    anchor._p.getparent().remove(anchor._p)

    with pytest.raises(DetachedAnchorError):
        _insert_after(anchor._p, OxmlElement('w:p'))


# --- mentoring: ordering, anchors, stats ----------------------------------------

def test_mentoring_tables_render_in_order_with_spacing_between():
    """Reverse insertion is the fragile part: three current mentees must come
    out as A -> spacing -> B -> spacing -> C under "Current Mentees:", each
    table carrying its own mentee's name."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    gen.doc.add_paragraph("Current Mentees:")
    gen.doc.add_paragraph("Past Mentees:")
    entries_by_code = {'N3A': [_mentee('N3A', 'Alice A'),
                               _mentee('N3A', 'Bob B'),
                               _mentee('N3A', 'Carol C')]}

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, "Current Mentees:", 7) == [
        ('tbl', 'Alice A'), ('p', ''),
        ('tbl', 'Bob B'), ('p', ''),
        ('tbl', 'Carol C'), ('p', ''),
        ('p', 'Past Mentees:'),
    ]


def test_mentoring_replaces_the_template_placeholder_tables():
    """The template ships one placeholder table under each heading; each is
    removed exactly once and only the rendered tables remain."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    gen.doc.add_paragraph("Current Mentees:")
    gen.doc.add_table(rows=6, cols=2).rows[0].cells[0].text = "Name"
    gen.doc.add_paragraph("Past Mentees:")
    gen.doc.add_table(rows=6, cols=2).rows[0].cells[0].text = "Name"

    gen._fill_mentoring({'N3A': [_mentee('N3A', 'Alice A')],
                         'N3B': [_mentee('N3B', 'Bob B', end_date='2014')]})

    assert [t.rows[0].cells[1].text for t in gen.doc.tables] == ['Alice A', 'Bob B']
    assert gen.stats['tables_populated'] == 2


def test_mentoring_current_only_removes_the_past_placeholder_too():
    """#845: mentees under ONLY "Current Mentees:" still get "Past
    Mentees:"'s own placeholder removed -- removal no longer depends on
    that heading having content of its own."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    gen.doc.add_paragraph("Current Mentees:")
    gen.doc.add_table(rows=6, cols=2).rows[0].cells[0].text = "Name"
    gen.doc.add_paragraph("Past Mentees:")
    gen.doc.add_table(rows=6, cols=2).rows[0].cells[0].text = "Name"

    gen._fill_mentoring({'N3A': [_mentee('N3A', 'Alice A')]})

    assert [t.rows[0].cells[1].text for t in gen.doc.tables] == ['Alice A']
    assert gen.stats['tables_populated'] == 1


def test_mentoring_no_entries_does_not_remove_a_non_placeholder_table():
    """Guard negative path, with no mentees at all: a table that is not a
    pristine "Name" placeholder survives the now-unconditional removal
    (#845) -- the guard, not the content check, is what protects it."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    gen.doc.add_paragraph("Current Mentees:")
    foreign = gen.doc.add_table(rows=1, cols=3)
    foreign.rows[0].cells[0].text = "Name of Committee"
    gen.doc.add_paragraph("Past Mentees:")

    gen._fill_mentoring({})

    assert len(gen.doc.tables) == 1
    assert gen.doc.tables[0].rows[0].cells[0].text == "Name of Committee"


def test_mentoring_real_template_no_entries_removes_both_mentee_placeholders():
    """#845, real template: with zero N1-N4 entries, N2's own placeholder
    is removed by `_fill_training_grants` (pre-existing #529 behavior) and
    both mentee placeholders are now ALSO removed on purpose -- three
    tables gone, none rebuilt -- while N1's own template slot and its
    neighbours are untouched."""
    gen = _template_generator()
    template_tables = len(gen.doc.tables)
    n1_neighbours_before = _body_after(gen.doc, N1_HEADING, 2)

    gen._fill_mentoring({})

    assert len(gen.doc.tables) == template_tables - 3
    assert not any(t.rows[0].cells[0].text.strip() == 'Name' for t in gen.doc.tables)
    assert _body_after(gen.doc, N1_HEADING, 2) == n1_neighbours_before
    assert gen.stats['tables_populated'] == 0
    assert gen.stats['entries_inserted'] == 0


def test_mentoring_missing_every_heading_logs_warning_with_count(caplog):
    gen = _new_generator()
    gen.doc.add_paragraph("Some Unrelated Section")
    entries_by_code = {'N3A': [_mentee('N3A', 'Alice A')],
                       'N4': [{'taxonomy_code': 'N4', 'text': 'An outcome.',
                               'extracted_fields': {}}]}

    with caplog.at_level(logging.WARNING, logger=MENTORING_LOGGER):
        gen._fill_mentoring(entries_by_code)

    assert [w.getMessage() for w in _warnings(caplog, MENTORING_LOGGER)] == [
        "Mentoring: none of 'Current Mentees:', 'Past Mentees:' or MENTORING "
        "found in template; 2 entries not rendered"]
    assert len(gen.doc.tables) == 0


def test_mentoring_missing_one_heading_logs_warning_and_renders_the_other(caplog):
    """Only "Current Mentees:" exists: the past group is reported with its
    count, the current group still renders."""
    gen = _new_generator()
    gen.doc.add_paragraph("Current Mentees:")
    entries_by_code = {'N3A': [_mentee('N3A', 'Alice A')],
                       'N3B': [_mentee('N3B', 'Bob B', end_date='2014'),
                               {'taxonomy_code': 'N3B', 'text': 'Completed: 27',
                                'extracted_fields': {}}]}

    with caplog.at_level(logging.WARNING, logger=MENTORING_LOGGER):
        gen._fill_mentoring(entries_by_code)

    assert [w.getMessage() for w in _warnings(caplog, MENTORING_LOGGER)] == [
        "Mentoring: 'Past Mentees:' heading not found in template; "
        "2 entries not rendered"]
    assert _body_after(gen.doc, "Current Mentees:", 2) == [('tbl', 'Alice A'), ('p', '')]


def test_mentoring_missing_mentoring_header_logs_warning_for_outcomes(caplog):
    gen = _new_generator()
    gen.doc.add_paragraph("Current Mentees:")
    gen.doc.add_paragraph("Past Mentees:")
    entries_by_code = {'N4': [{'taxonomy_code': 'N4', 'text': 'An outcome.',
                               'extracted_fields': {}}]}

    with caplog.at_level(logging.WARNING, logger=MENTORING_LOGGER):
        gen._fill_mentoring(entries_by_code)

    assert [w.getMessage() for w in _warnings(caplog, MENTORING_LOGGER)] == [
        "Mentoring: MENTORING heading not found in template; "
        "1 outcome lines not rendered"]


def test_mentoring_single_heading_fallback_renders_both_groups():
    """A template with only a MENTORING header (neither mentee heading):
    both groups land under it, current above past, nothing dropped."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    entries_by_code = {'N3A': [_mentee('N3A', 'Alice A')],
                       'N3B': [_mentee('N3B', 'Bob B', end_date='2014')]}

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, "MENTORING", 4) == [
        ('tbl', 'Alice A'), ('p', ''), ('tbl', 'Bob B'), ('p', '')]


def test_mentoring_stats_count_only_tables_that_were_built():
    """A record with no name renders nothing and moves no counter; a named
    one moves both. The line writer counts entries_inserted on its own."""
    gen = _new_generator()
    anchor = gen.doc.add_paragraph("Current Mentees:")._element

    assert gen._create_mentee_table_with_spacing(MenteeRecord(name=''), anchor) is None
    assert gen.stats['tables_populated'] == 0
    assert gen.stats['entries_inserted'] == 0

    assert gen._create_mentee_table_with_spacing(MenteeRecord(name='Alice A'), anchor) is not None
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['entries_inserted'] == 1

    gen._insert_mentoring_line("Completed: 27", anchor)
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['entries_inserted'] == 2


def test_mentoring_detached_anchor_raises_instead_of_stranding_content():
    """The old `except (ValueError, IndexError): pass` left the paragraph at
    the end of the document and still counted it. A cursor that no longer
    points into the body is now a loud failure, and nothing is counted."""
    gen = _new_generator()
    anchor = gen.doc.add_paragraph("Current Mentees:")._element
    anchor.getparent().remove(anchor)

    with pytest.raises(DetachedAnchorError):
        gen._insert_mentoring_line("Completed: 27", anchor)
    with pytest.raises(DetachedAnchorError):
        gen._create_mentee_table_with_spacing(MenteeRecord(name='Alice A'), anchor)
    assert gen.stats['entries_inserted'] == 0
    assert gen.stats['tables_populated'] == 0


def test_mentoring_real_template_body_order():
    """Real template: summary line, then the mentee tables with spacing,
    directly under each heading and above the template's own "Duplicate
    table below as needed" instruction; the two placeholder tables are gone;
    N4 sits under MENTORING."""
    gen = _template_generator()
    template_tables = len(gen.doc.tables)
    entries_by_code = {
        'N3A': [_mentee('N3A', 'Alice A', start_date='2022'),
                _mentee('N3A', 'Bob B', start_date='2023'),
                {'taxonomy_code': 'N3A', 'text': 'Current Ph.D. Students: 2',
                 'extracted_fields': {}}],
        'N3B': [_mentee('N3B', 'Carol C', start_date='2015', end_date='2019')],
        'N4': [{'taxonomy_code': 'N4', 'text': 'One mentee now leads a lab.',
                'extracted_fields': {}}],
    }

    gen._fill_mentoring(entries_by_code)

    current = _body_after(gen.doc, "Current Mentees:", 6)
    assert current[:5] == [('p', 'Current Ph.D. Students: 2'),
                           ('tbl', 'Alice A'), ('p', ''),
                           ('tbl', 'Bob B'), ('p', '')]
    assert current[5][0] == 'p' and current[5][1].startswith('Duplicate table below')
    past = _body_after(gen.doc, "Past Mentees:", 3)
    assert past[:2] == [('tbl', 'Carol C'), ('p', '')]
    assert past[2][1].startswith('Duplicate table below')
    assert _body_after(gen.doc, "MENTORING", 1) == [('p', 'One mentee now leads a lab.')]
    # Two mentee placeholders out, three mentee tables in; the N2 placeholder
    # also disappears as a side effect -- `_fill_mentoring` calls
    # `_fill_training_grants([])` first, and its removal is now
    # unconditional-but-guarded (#529 round 2, F3), so it fires even though
    # this test has no N2 entries.
    assert len(gen.doc.tables) == template_tables - 3 + 3
    assert gen.stats['tables_populated'] == 3
    assert gen.stats['entries_inserted'] == 5


# --- N1/N2: pure builders (#529) -------------------------------------------------

def test_program_leadership_line_uses_text_verbatim_when_present():
    """(a) The line equals the entry `text` exactly (stripped) -- no field
    assembly at all (#529 round 3)."""
    fields = {'role': 'Director', 'program_name': 'Scholars Program',
              'institution': 'Test University', 'start_date': '2019',
              'end_date': '2022'}
    assert _program_leadership_line(fields, '  Directed the Scholars Program.  ') == \
        'Directed the Scholars Program.'


def test_program_leadership_line_ignores_fields_when_text_present():
    """A role-only entry keeps its FULL text rather than shrinking to just
    the field -- the round 1-2 defect this round fixes (#529 round 3)."""
    assert _program_leadership_line(
        {'role': 'Director'}, 'Directed a mentoring initiative in 2019.') == \
        'Directed a mentoring initiative in 2019.'


def test_program_leadership_line_empty_text_falls_back_to_joined_fields():
    """(b) Empty `text` with fields present -> the non-empty fields joined
    with ', ' -- the only fallback (#529 round 3)."""
    fields = {'role': 'Director', 'program_name': 'Scholars Program',
              'institution': 'Test University', 'start_date': '2019',
              'end_date': '2022'}
    assert _program_leadership_line(fields, '') == \
        'Director, Scholars Program, Test University, 2019-2022'


def test_program_leadership_line_empty_text_falls_back_to_partial_fields():
    assert _program_leadership_line({'role': 'Director'}, '') == 'Director'
    assert _program_leadership_line(
        {'program_name': 'Scholars Program', 'institution': 'Test University'}, '') == \
        'Scholars Program, Test University'


def test_program_leadership_line_empty_text_and_fields_is_empty_string():
    """Never a blank paragraph: `_fill_program_leadership` checks for a
    non-empty line before inserting, so this is the one input that produces
    no line at all rather than an inserted blank one (#529 round 3)."""
    assert _program_leadership_line({}, '') == ''
    assert _program_leadership_line({}, '   ') == ''


def test_program_leadership_line_none_text_is_safe():
    """(d) `None`/missing `text` is safe -- falls back to fields exactly
    like an empty string does (#529 round 3)."""
    assert _program_leadership_line({'role': 'Director'}, None) == 'Director'
    assert _program_leadership_line({}, None) == ''


def test_program_leadership_line_none_field_values_are_safe():
    fields = {'role': None, 'program_name': None, 'institution': None,
              'start_date': None, 'end_date': None}
    assert _program_leadership_line(fields, None) == ''
    assert _program_leadership_line(fields, '') == ''


def test_training_grant_rows_full_fields():
    fields = {'agency': 'National Test Institute', 'grant_number': 'T32-999',
              'role': 'Mentor', 'grant_title': 'Test Training Program',
              'start_date': '2018', 'end_date': '2021'}
    assert _training_grant_rows(fields) == [
        ('Award Source (funding agency, type of grant):',
         'National Test Institute (T32-999) (Mentor)'),
        ('Project title:', 'Test Training Program'),
        ('Duration of support (mm/yyyy-mm/yyyy):', '2018-2021'),
    ]


def test_training_grant_rows_grant_number_already_in_agency_is_not_repeated():
    fields = {'agency': 'National Test Institute T32-999', 'grant_number': 'T32-999'}
    assert _training_grant_rows(fields)[0] == (
        'Award Source (funding agency, type of grant):',
        'National Test Institute T32-999')


def test_training_grant_rows_title_precedence_grant_title_then_title():
    assert _training_grant_rows({'grant_title': 'A', 'title': 'B'})[1] == \
        ('Project title:', 'A')
    assert _training_grant_rows({'title': 'B'})[1] == ('Project title:', 'B')


def test_training_grant_rows_none_values_are_safe():
    fields = {'agency': None, 'grant_number': None, 'role': None,
              'grant_title': None, 'title': None, 'start_date': None, 'end_date': None}
    assert _training_grant_rows(fields) == [
        ('Award Source (funding agency, type of grant):', ''),
        ('Project title:', ''),
        ('Duration of support (mm/yyyy-mm/yyyy):', ''),
    ]


def test_training_grant_is_sparse_true_when_no_identifying_field():
    assert _training_grant_is_sparse({}) is True
    assert _training_grant_is_sparse({'role': 'Mentor'}) is True


def test_training_grant_is_sparse_false_when_any_identifying_field_present():
    assert _training_grant_is_sparse({'grant_number': 'T32-1'}) is False
    assert _training_grant_is_sparse({'agency': 'A'}) is False
    assert _training_grant_is_sparse({'grant_title': 'A'}) is False


# --- N1/N2: rendering into their template slots (#529) ---------------------------

def _n1(**fields) -> dict:
    text = fields.pop('text', 'N1 entry')
    return {'taxonomy_code': 'N1', 'text': text, 'extracted_fields': fields}


def _n2(**fields) -> dict:
    text = fields.pop('text', 'N2 entry')
    return {'taxonomy_code': 'N2', 'text': text, 'extracted_fields': fields}


def test_n1_missing_anchor_falls_back_to_mentoring_header(caplog):
    """Blank Document() with only MENTORING: warning emitted, the entry
    still renders (as a fallback line), nothing raised (#529)."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")

    with caplog.at_level(logging.WARNING, logger=MENTORING_LOGGER):
        gen._fill_mentoring({'N1': [_n1()]})

    assert [w.getMessage() for w in _warnings(caplog, MENTORING_LOGGER)] == [
        "Mentoring: 'Leadership and mentoring in programs (Describe "
        "activity; include dates)' heading not found; 1 entries rendered "
        "under MENTORING instead"]
    assert _body_after(gen.doc, "MENTORING", 1) == [('p', 'N1 entry')]


def test_n2_missing_anchor_falls_back_to_mentoring_header(caplog):
    """Same fallback shape as N1, for the table-shaped code (#529)."""
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")

    with caplog.at_level(logging.WARNING, logger=MENTORING_LOGGER):
        gen._fill_mentoring({'N2': [_n2(agency='National Test Institute')]})

    assert [w.getMessage() for w in _warnings(caplog, MENTORING_LOGGER)] == [
        "Mentoring: 'Institutional Training Grants and Mentored Trainee "
        "Grants' heading not found; 1 entries rendered under MENTORING "
        "instead"]
    assert len(gen.doc.tables) == 1
    assert gen.doc.tables[0].rows[0].cells[1].text == 'National Test Institute'


def test_n1_real_template_three_lines_in_order_text_verbatim_and_fallback():
    """(c) Real template: three N1 entries land after "Leadership and
    mentoring in programs..." in input order (#529). The first two render
    their own `text` verbatim -- unshortened, even though the second's
    fields alone would describe less -- and the third has no `text` at all
    so it falls back to its non-empty fields joined with ', ' (#529
    round 3)."""
    gen = _template_generator()
    entries_by_code = {
        'N1': [
            _n1(role='Director', text='Directed the Scholars Program from 2015-2020.'),
            _n1(program_name='Scholars Program', institution='Test University',
                text='Served as a mentor to junior faculty in the program.'),
            _n1(role='Advisor', institution='Test University', text=''),
        ],
    }

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, N1_HEADING, 3) == [
        ('p', 'Directed the Scholars Program from 2015-2020.'),
        ('p', 'Served as a mentor to junior faculty in the program.'),
        ('p', 'Advisor, Test University'),
    ]
    assert gen.stats['entries_inserted'] == 3


def test_n1_entry_with_empty_text_and_no_fields_inserts_no_blank_line():
    """An N1 entry with nothing to say -- empty `text`, no populated fields --
    must not become a blank paragraph under the heading (#529 round 3). The
    `if line:` guard in `_fill_program_leadership` is what keeps it out; the
    two real entries around it still render, in order, and the counter only
    counts what was written."""
    gen = _template_generator()
    entries_by_code = {
        'N1': [
            _n1(text='First real line.'),
            _n1(text='', role='', program_name=None),
            _n1(text='Second real line.'),
        ],
    }

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, N1_HEADING, 2) == [
        ('p', 'First real line.'),
        ('p', 'Second real line.'),
    ]
    assert gen.stats['entries_inserted'] == 2


def test_n2_real_template_tables_in_order_placeholder_removed_sparse_as_line():
    """Real template: two N2 entries -> placeholder table gone, two 3-row
    tables land AFTER the instruction paragraph ("Duplicate table below
    as needed...") in input order with a spacer between -- heading ->
    instruction -> tables, the template's own order (#529 round 2, F1;
    was between the heading and the instruction before this fix). A
    third, sparse entry renders as one plain line and adds no table.
    No N3A/N3B/N4 entries here, so both mentee placeholders are also
    removed on purpose (#845)."""
    gen = _template_generator()
    template_tables = len(gen.doc.tables)
    entries_by_code = {
        'N2': [
            _n2(grant_title='Test Training Program A', agency='National Test Institute',
                grant_number='T32-100', role='Mentor', start_date='2018', end_date='2021'),
            _n2(title='Test Training Program B', agency='Regional Test Foundation',
                start_date='2020', end_date='2023'),
            _n2(text='A sparse training-grant line with no identifying field.'),
        ],
    }

    gen._fill_mentoring(entries_by_code)

    n2_region = _body_after(gen.doc, N2_HEADING, 6)
    assert n2_region[0] == ('p', N2_INSTRUCTION)
    assert n2_region[1] == ('tbl', 'National Test Institute (T32-100) (Mentor)')
    assert n2_region[2] == ('p', '')
    assert n2_region[3] == ('tbl', 'Regional Test Foundation')
    assert n2_region[4] == ('p', '')
    assert n2_region[5] == (
        'p', 'A sparse training-grant line with no identifying field.')
    body = list(gen.doc.element.body)
    instruction_idx = next(
        i for i, el in enumerate(body) if el.tag == qn('w:p')
        and Paragraph(el, gen.doc).text.strip() == N2_INSTRUCTION)
    table_a = Table(body[instruction_idx + 1], gen.doc)
    assert [row.cells[1].text for row in table_a.rows] == [
        'National Test Institute (T32-100) (Mentor)', 'Test Training Program A', '2018-2021']
    # N2's own placeholder out, two grant tables in; the sparse entry built
    # none; both mentee placeholders also out (no mentees at all, #845).
    assert len(gen.doc.tables) == template_tables - 1 - 2 + 2
    assert not any(t.rows[0].cells[0].text.strip() == 'Name' for t in gen.doc.tables)
    assert gen.stats['tables_populated'] == 2
    assert gen.stats['entries_inserted'] == 3


def test_n2_real_template_tables_after_heading_when_instruction_paragraph_missing():
    """F1 fallback (#529 round 2): with the instruction paragraph absent
    from the template, N2's tables anchor on the heading itself instead of
    raising or landing under the wrong content."""
    gen = _template_generator()
    instruction_idx = next(
        i for i, p in enumerate(gen.doc.paragraphs)
        if p.text.strip() == N2_INSTRUCTION)
    instruction_element = gen.doc.paragraphs[instruction_idx]._element
    instruction_element.getparent().remove(instruction_element)
    assert gen._find_paragraph_exact(N2_INSTRUCTION) is None

    gen._fill_mentoring({'N2': [_n2(agency='National Test Institute')]})

    assert _body_after(gen.doc, N2_HEADING, 1) == [('tbl', 'National Test Institute')]


def test_n2_foreign_table_after_heading_survives_the_shape_guard():
    """Shape guard, negative path (#529 round 2, F3): a table that does not
    look like N2's own placeholder is never removed, even though it is the
    first (and only) table `_first_table_after` finds; the new N2 table is
    inserted ahead of it, after the heading."""
    gen = _new_generator()
    gen.doc.add_paragraph(N2_HEADING)
    foreign = gen.doc.add_table(rows=5, cols=2)
    foreign.rows[0].cells[0].text = 'Something else:'
    assert not _looks_like_training_grant_table(foreign)

    gen._fill_mentoring({'N2': [_n2(agency='National Test Institute')]})

    assert len(gen.doc.tables) == 2
    # New N2 table, its spacer paragraph, then the foreign table -- still
    # there, unharmed.
    region = _body_after(gen.doc, N2_HEADING, 3)
    assert region[0] == ('tbl', 'National Test Institute')
    assert region[1] == ('p', '')
    assert region[2] == ('tbl', '')
    tables_by_label = {t.rows[0].cells[0].text: t for t in gen.doc.tables}
    assert 'Something else:' in tables_by_label
    assert len(tables_by_label['Something else:'].rows) == 5


def test_n2_real_template_no_entries_removes_placeholder_region_otherwise_unchanged():
    """F3's removal policy: unconditional-but-guarded. With NO N2 entries at
    all, the placeholder table is still removed (#836's real-pipeline
    cascade already does this on every render, N2 content or not -- see the
    PR description), but every other element in the MENTORING -> Mentees
    body region is untouched at the XML level (#529 round 2, F3)."""
    from lxml import etree

    def region_elements(doc):
        body = list(doc.element.body)
        start = next(i for i, el in enumerate(body) if el.tag == qn('w:p')
                     and Paragraph(el, doc).text.strip() == 'MENTORING')
        end = next(i for i, el in enumerate(body) if el.tag == qn('w:p')
                   and Paragraph(el, doc).text.strip() == 'Mentees')
        return body[start:end]

    gen = _template_generator()
    before = region_elements(gen.doc)
    assert sum(1 for el in before if el.tag == qn('w:tbl')) == 1
    before_non_tables = [etree.tostring(el) for el in before if el.tag != qn('w:tbl')]

    gen._fill_mentoring({})

    after = region_elements(gen.doc)
    after_non_tables = [etree.tostring(el) for el in after if el.tag != qn('w:tbl')]
    assert after_non_tables == before_non_tables
    assert not any(el.tag == qn('w:tbl') for el in after)


def test_n1_n2_only_entries_still_render_when_no_n3_n4_content():
    """N1/N2 must not depend on the N3/N4 partition being non-empty -- a CV
    with ONLY N1/N2 content must not hit `_fill_mentoring`'s early return
    for an empty mentee/outcome partition (#529)."""
    gen = _template_generator()
    gen._fill_mentoring({'N1': [_n1()],
                         'N2': [_n2(agency='National Test Institute')]})

    assert _body_after(gen.doc, N1_HEADING, 1) == [('p', 'N1 entry')]
    n2_region = _body_after(gen.doc, N2_HEADING, 2)
    assert n2_region[0] == ('p', N2_INSTRUCTION)
    assert n2_region[1] == ('tbl', 'National Test Institute')


# --- other education: template structure, three-column rendering -----------------

OTHER_EDUCATION_LOGGER = 'unified_pipeline.stage6.sections.other_education'


def _b2(program: str, **fields) -> dict:
    return {'taxonomy_code': 'B2', 'text': f"{program} entry",
            'extracted_fields': {'program_name': program, **fields}}


def test_other_education_missing_heading_logs_warning_with_count(caplog):
    gen = _new_generator()
    gen.doc.add_paragraph("Some Unrelated Section")

    with caplog.at_level(logging.WARNING, logger=OTHER_EDUCATION_LOGGER):
        gen._fill_other_education([_b2('Certificate A'), _b2('Certificate B')])

    assert [w.getMessage() for w in _warnings(caplog, OTHER_EDUCATION_LOGGER)] == [
        "Other Educational Experiences: section heading not found in template; "
        "2 entries not rendered"]
    assert gen.stats['tables_populated'] == 0
    assert gen.stats['entries_inserted'] == 0


def test_other_education_missing_table_logs_warning_with_count(caplog):
    """The heading exists but no table follows it: template structure,
    not empty data, and reported as such."""
    gen = _new_generator()
    gen.doc.add_paragraph("OTHER EDUCATIONAL")
    gen.doc.add_paragraph("No table follows this heading.")

    with caplog.at_level(logging.WARNING, logger=OTHER_EDUCATION_LOGGER):
        gen._fill_other_education([_b2('Certificate A')])

    assert [w.getMessage() for w in _warnings(caplog, OTHER_EDUCATION_LOGGER)] == [
        "Other Educational Experiences: table not found after section heading; "
        "1 entries not rendered"]
    assert gen.stats['tables_populated'] == 0


@pytest.mark.parametrize("columns", [2, 4])
def test_other_education_wrong_column_count_logs_warning_and_writes_nothing(caplog, columns):
    """The writer always builds program | institution | dates; a grid of any
    other width would land content in the wrong column, so it is refused."""
    gen = _new_generator()
    gen.doc.add_paragraph("OTHER EDUCATIONAL")
    table = gen.doc.add_table(rows=1, cols=columns)

    with caplog.at_level(logging.WARNING, logger=OTHER_EDUCATION_LOGGER):
        gen._fill_other_education([_b2('Certificate A')])

    assert [w.getMessage() for w in _warnings(caplog, OTHER_EDUCATION_LOGGER)] == [
        f"Other Educational Experiences: expected a 3-column table, found "
        f"{columns} columns; 1 entries not rendered"]
    assert len(table.rows) == 1
    assert gen.stats['tables_populated'] == 0
    assert gen.stats['entries_inserted'] == 0


def test_other_education_real_template_three_column_rows():
    """Real template: the B2 table has three columns; the header row survives,
    the placeholder row is cleared, entries land most-recent-first with
    program | institution, location | dates, and a sparse entry (no program,
    no institution) is skipped and not counted."""
    gen = _template_generator()
    entries = [
        _b2('Certificate in Epidemiology', institution='Some College',
            location='Boston, MA', start_date='2017-08', end_date='2018-05'),
        _b2('Workshop on Grant Writing', institution='Another Institute', year='2021'),
        {'taxonomy_code': 'B2', 'text': 'stray line', 'extracted_fields': {'year': '2020'}},
    ]

    gen._fill_other_education(entries)

    table = gen._find_table_after_paragraph(
        gen._find_paragraph_with_text("OTHER EDUCATIONAL"))
    assert len(table.columns) == 3
    assert [c.text for c in table.rows[0].cells] == [
        'Description', 'Institution, city and state',
        'Dates attended (mm/yy – mm/yy)']
    assert [[c.text for c in row.cells] for row in table.rows[1:]] == [
        ['Workshop on Grant Writing', 'Another Institute', '2021'],
        ['Certificate in Epidemiology', 'Some College, Boston, MA', '08/17-05/18'],
    ]
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['entries_inserted'] == 2


# --- patents: reposition failure ------------------------------------------------

def _patent(text: str, **fields) -> dict:
    return {'taxonomy_code': 'M2D', 'text': text, 'extracted_fields': fields}


def test_patents_detached_cursor_logs_warning_counts_misplaced_and_holds_cursor(caplog):
    """The real failure mode: `_add_spacing_paragraph` (shared, on the
    generator) removes its paragraph from the body and re-inserts it with
    `body.insert`; when that insert fails it swallows the error and returns
    the now-DETACHED element, which the section had been adopting as its
    cursor. Forcing the FIRST `body.insert` call to raise (and only that
    one) reproduces exactly that -- and, unlike an always-raising mock,
    lets a second table's own insert succeed if the code ever hands it a
    chance to. It must not get that chance: once the cursor is detached,
    `_place_patent_table` must keep failing (and keep reporting) for every
    later table instead of quietly recovering on the next placement it
    happens to try.

    The first table is placed (before any spacing exists). The one forced
    `body.insert` failure detaches the spacing paragraph after it, so every
    later table finds a detached cursor: each is logged with its ordinal
    and the caught exception, counted in tables_misplaced, and left at the
    document end -- immediately after the previous one, with no spacing
    paragraph between them, because the cursor never became attached again.

    This pins the regression a cursor-advances-on-failure bug would cause:
    if the guard were removed, the second table's own spacing attempt would
    reach the now-unmocked `body.insert`, succeed, and the third table
    would land correctly spaced after the second -- one warning and
    `tables_populated == 2` instead of two warnings and `tables_populated
    == 1`, with a spacing paragraph between Widget B and Widget C that must
    not exist here."""
    gen = _new_generator()
    gen.doc.add_paragraph().add_run("Patents & Inventions").bold = True

    real_insert = gen.doc.element.body.insert
    calls = {'n': 0}

    def _raise_once(*args, **kwargs):
        calls['n'] += 1
        if calls['n'] == 1:
            raise ValueError("forced for test")
        return real_insert(*args, **kwargs)

    gen.doc.element.body.insert = _raise_once

    entries = [_patent('newest', title='Widget A', year='2021'),
               _patent('middle', title='Widget B', year='2020'),
               _patent('oldest', title='Widget C', year='2019')]

    with caplog.at_level(logging.WARNING, logger=PATENTS_LOGGER):
        gen._fill_patents(entries)  # must not raise

    warnings = _warnings(caplog, PATENTS_LOGGER)
    assert [w.getMessage() for w in warnings] == [
        "Patents & Inventions: table reposition failed for entry 2 of 3; "
        "table left at document end instead of under the section heading",
        "Patents & Inventions: table reposition failed for entry 3 of 3; "
        "table left at document end instead of under the section heading",
    ]
    # exc_info=True on the warning call -- the caught DetachedAnchorError
    # must reach the log, not just its message text.
    assert all(w.exc_info is not None and isinstance(w.exc_info[1], DetachedAnchorError)
               for w in warnings)

    # Content not lost: all three tables exist, the first under the heading.
    assert [t.rows[0].cells[1].text for t in gen.doc.tables] == [
        'Widget A', 'Widget B', 'Widget C']
    assert _body_after(gen.doc, "Patents & Inventions", 1) == [('tbl', 'Widget A')]
    # Widget C lands directly after Widget B with no spacing between them --
    # the cursor never recovered, so no spacing insert was ever attempted
    # for either. A guard that advanced the cursor on failure would let the
    # forced insert fail once and then heal, spacing C correctly after B.
    assert _body_after(gen.doc, "Patents & Inventions", 3) == [
        ('tbl', 'Widget A'), ('tbl', 'Widget B'), ('tbl', 'Widget C')]
    # Failure stats: placed vs misplaced are distinct; entries_inserted counts both.
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['tables_misplaced'] == 2
    assert gen.stats['entries_inserted'] == 3


def test_patents_success_stats_have_no_misplaced_count():
    gen = _new_generator()
    gen.doc.add_paragraph().add_run("Patents & Inventions").bold = True

    gen._fill_patents([_patent('a', title='Widget A', year='2021'),
                       _patent('b', title='Widget B', year='2020')])

    assert gen.stats['tables_populated'] == 2
    assert gen.stats['entries_inserted'] == 2
    assert gen.stats.get('tables_misplaced', 0) == 0


def test_patents_missing_heading_logs_warning_with_count(caplog):
    gen = _new_generator()
    gen.doc.add_paragraph("Some Unrelated Section")

    with caplog.at_level(logging.WARNING, logger=PATENTS_LOGGER):
        gen._fill_patents([_patent('a', title='Widget A'), _patent('b', title='Widget B')])

    assert [w.getMessage() for w in _warnings(caplog, PATENTS_LOGGER)] == [
        "Patents & Inventions: section heading not found in template; "
        "2 entries not rendered"]
    assert len(gen.doc.tables) == 0


def test_patents_render_twice_replaces_the_first_render():
    """Rendering the section again into the same document used to append a
    second set of tables; the previously rendered tables (and the spacing
    between them) are now cleared first, so the body shape is identical."""
    gen = _new_generator()
    gen.doc.add_paragraph().add_run("Patents & Inventions").bold = True
    gen.doc.add_paragraph("Please include inventors, title of invention and patent number.")
    gen.doc.add_paragraph("MENTORING")
    entries = [_patent('a', title='Widget A', year='2021'),
               _patent('b', title='Widget B', year='2020')]

    gen._fill_patents(entries)
    first = _body_after(gen.doc, "Patents & Inventions", 6)
    gen._fill_patents(entries)

    assert first == [('tbl', 'Widget A'), ('p', ''), ('tbl', 'Widget B'),
                     ('p', ''), ('p', 'MENTORING')]
    assert _body_after(gen.doc, "Patents & Inventions", 6) == first
    assert len(gen.doc.tables) == 2


def test_patents_real_template_body_order():
    """Real template: the tables sit directly under the heading, most recent
    first, one spacing paragraph between them, the blanked instruction line
    and the template's two blank paragraphs below, then MENTORING."""
    gen = _template_generator()
    template_tables = len(gen.doc.tables)

    gen._fill_patents([_patent('older', title='Older Widget', patent_number='US1',
                               issue_date='2019-05-01', year='2019'),
                       _patent('newer', title='Newer Widget', patent_number='US2',
                               issue_date='2021-11-01', year='2021')])

    assert _body_after(gen.doc, "Patents & Inventions", 7) == [
        ('tbl', 'Newer Widget'), ('p', ''), ('tbl', 'Older Widget'),
        ('p', ''), ('p', ''), ('p', ''), ('p', 'MENTORING')]
    assert len(gen.doc.tables) == template_tables + 2
    assert gen.stats['tables_populated'] == 2
    assert gen.stats.get('tables_misplaced', 0) == 0


# --- postdoc training: missing template structure --------------------------------

def test_postdoc_missing_heading_logs_warning(caplog):
    """A template with neither "POSTDOCTORAL" nor "TRAINING" anywhere used to
    return silently, dropping every entry with no diagnostic (postdoc_training.py,
    first bare return)."""
    gen = _new_generator()
    gen.doc.add_paragraph("Some Unrelated Section")

    entries_by_code = {
        'C1': [{'taxonomy_code': 'C1', 'text': 'A postdoc research entry',
                'extracted_fields': {'institution': 'Some Hospital'}}],
    }

    with caplog.at_level(logging.WARNING, logger=POSTDOC_LOGGER):
        gen._fill_postdoc_training(entries_by_code)  # must not raise

    warnings = _warnings(caplog, POSTDOC_LOGGER)
    assert len(warnings) == 1, "expected exactly one warning, got: %r" % (
        [r.message for r in warnings],)
    message = warnings[0].getMessage()
    assert 'Postdoctoral Training' in message
    assert '1 entries not rendered' in message
    assert len(gen.doc.tables) == 0


def test_postdoc_missing_table_logs_warning(caplog):
    """A template whose POSTDOCTORAL heading exists but has no table after
    it used to return silently (postdoc_training.py, second bare return)."""
    gen = _new_generator()
    gen.doc.add_paragraph("POSTDOCTORAL TRAINING")
    gen.doc.add_paragraph("No table follows this heading.")

    entries_by_code = {
        'C1': [{'taxonomy_code': 'C1', 'text': 'A postdoc research entry',
                'extracted_fields': {'institution': 'Some Hospital'}}],
    }

    with caplog.at_level(logging.WARNING, logger=POSTDOC_LOGGER):
        gen._fill_postdoc_training(entries_by_code)  # must not raise

    warnings = _warnings(caplog, POSTDOC_LOGGER)
    assert len(warnings) == 1, "expected exactly one warning, got: %r" % (
        [r.message for r in warnings],)
    message = warnings[0].getMessage()
    assert 'Postdoctoral Training' in message
    assert '1 entries not rendered' in message


# ---- geographic-scope classification failure (#547, instance 3) ----

S6_LOGGER = 'unified_pipeline.stage_6_word_template'

_OWNER_LOCATION = {
    'inference_success': True,
    'primary_location': {'institution': 'Example University Medical Center',
                         'city': 'Springfield', 'state': 'XX'},
    'locations': [],
    'metro_area': 'Springfield',
}
_PRESENTATION = {'text': 'Talk at Sample Institute', 'taxonomy_code': 'R',
                 'extracted_fields': {'organization': 'Sample Institute'}}


def _raising_llm(*args, **kwargs):
    raise RuntimeError('simulated LLM outage')


def _scope_generator() -> WCMTemplateGenerator:
    gen = WCMTemplateGenerator(verbose=False)
    gen.cv_owner_location = _OWNER_LOCATION
    return gen


def test_geo_scope_failure_warns_counts_and_defaults_national_when_not_verbose(
        monkeypatch, caplog):
    monkeypatch.setattr(s6, 'call_llm', _raising_llm)
    gen = _scope_generator()
    with caplog.at_level(logging.WARNING, logger=S6_LOGGER):
        assert gen._classify_geographic_scope(_PRESENTATION) == 'National'
    warnings = _warnings(caplog, S6_LOGGER)
    assert len(warnings) == 1
    assert 'Geographic scope classification failed' in warnings[0].getMessage()
    assert warnings[0].exc_info and warnings[0].exc_info[0] is RuntimeError
    assert gen.stats[GEO_SCOPE_FAILURE_STAT] == 1


def test_geo_scope_success_does_not_count_a_failure(monkeypatch, caplog):
    monkeypatch.setattr(
        s6, 'call_llm', lambda *a, **k: {'content': '{"scope": "Regional"}'})
    gen = _scope_generator()
    with caplog.at_level(logging.WARNING, logger=S6_LOGGER):
        assert gen._classify_geographic_scope(_PRESENTATION) == 'Regional'
    assert not _warnings(caplog, S6_LOGGER)
    assert gen.stats[GEO_SCOPE_FAILURE_STAT] == 0
    assert gen._geo_scope_failure_warnings() == []


def test_geo_scope_failure_reaches_the_render_warnings_sidecar(
        monkeypatch, tmp_path):
    monkeypatch.setattr(s6, 'call_llm', _raising_llm)
    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: None
    input_path = tmp_path / 'in.json'
    input_path.write_text(json.dumps({
        'document_uid': 'T547G',
        'cv_owner_location': _OWNER_LOCATION,
        'entries': [
            {'text': 'Name: Jane Q. Public, MD', 'taxonomy_code': 'A',
             'extracted_fields': {}, 'element_idx_start': 0},
            dict(_PRESENTATION, element_idx_start=1),
        ]}))
    gen.generate(str(input_path), str(tmp_path / 'out.docx'))
    sidecar = json.loads((tmp_path / 'T547G_render_warnings.json').read_text())
    found = [w for w in sidecar['warnings']
             if w.get('check') == GEO_SCOPE_FAILURE_STAT]
    assert len(found) == 1
    assert found[0]['severity'] == 'WARN'
    assert found[0]['evidence'] == [f'{GEO_SCOPE_FAILURE_STAT}=1']


# ---- segment-reclassification failure (#652) ----

_RECLASSIFY_TEXT = 'Visiting Lecturer, Example Institute, 2010'


def test_reclassify_failure_warns_counts_and_returns_none_when_verbose(
        monkeypatch, caplog):
    """The `except Exception` arm of `_reclassify_entry_segments` returns
    None, logs, and counts the failure."""
    monkeypatch.setattr(s6, 'call_llm', _raising_llm)
    gen = WCMTemplateGenerator(verbose=True)
    with caplog.at_level(logging.WARNING, logger=S6_LOGGER):
        assert gen._reclassify_entry_segments(_RECLASSIFY_TEXT, 'P') is None
    warnings = _warnings(caplog, S6_LOGGER)
    assert len(warnings) == 1
    assert warnings[0].getMessage() == (
        'LLM reclassification failed; entry stays in the appendix unsplit')
    assert warnings[0].exc_info and warnings[0].exc_info[0] is RuntimeError
    assert gen.stats[RECLASSIFY_FAILURE_STAT] == 1


def test_reclassify_failure_warns_and_counts_when_not_verbose(
        monkeypatch, caplog):
    """#652: production runs are not verbose, so the fallback must be visible
    without it (CODING_STANDARDS 5.3): a warning with the traceback and a
    stat. The None return is unchanged."""
    monkeypatch.setattr(s6, 'call_llm', _raising_llm)
    gen = WCMTemplateGenerator(verbose=False)
    with caplog.at_level(logging.WARNING, logger=S6_LOGGER):
        assert gen._reclassify_entry_segments(_RECLASSIFY_TEXT, 'P') is None
    warnings = _warnings(caplog, S6_LOGGER)
    assert len(warnings) == 1
    assert warnings[0].getMessage() == (
        'LLM reclassification failed; entry stays in the appendix unsplit')
    assert warnings[0].exc_info and warnings[0].exc_info[0] is RuntimeError
    assert gen.stats[RECLASSIFY_FAILURE_STAT] == 1


def test_reclassify_success_does_not_count_a_failure(monkeypatch, caplog):
    monkeypatch.setattr(
        s6, 'call_llm', lambda *a, **k: {'content': f'P: {_RECLASSIFY_TEXT}'})
    gen = WCMTemplateGenerator(verbose=False)
    with caplog.at_level(logging.WARNING, logger=S6_LOGGER):
        assert gen._reclassify_entry_segments(_RECLASSIFY_TEXT, 'P')
    assert not _warnings(caplog, S6_LOGGER)
    assert gen.stats[RECLASSIFY_FAILURE_STAT] == 0
    assert gen._reclassify_failure_warnings() == []


# ---- reclassify reply must cover its source (#1230) ----

_FUSED_WORDS = ['zeolite', 'marigold', 'quartzite', 'nasturtium', 'basalt',
                'chrysanthemum', 'obsidian', 'hyacinth', 'feldspar',
                'delphinium', 'granite', 'snapdragon']
_FUSED_LINES = [
    f'Doe J, Roe K. Synthetic study of {word} outcomes. '
    f'Journal of Invented {word.title()} Findings; volume {i}.'
    for i, word in enumerate(_FUSED_WORDS)
]
_FUSED_TEXT = '\n'.join(_FUSED_LINES)


def _reply_llm(content):
    return lambda *a, **k: {'content': content}


def test_reclassify_summary_reply_is_rejected_and_counted(monkeypatch, caplog):
    """A reply that stands one bracketed line in for the whole list covers a
    sliver of the source: refused (None, so the caller keeps the entry whole)
    and counted with the failures."""
    summary = 'KEEP: [All 12 publication entries follow, first Doe J0 last Doe J11]'
    monkeypatch.setattr(s6, 'call_llm', _reply_llm(summary))
    gen = WCMTemplateGenerator(verbose=False)
    with caplog.at_level(logging.WARNING, logger=S6_LOGGER):
        assert gen._reclassify_entry_segments(_FUSED_TEXT, 'T') is None
    assert gen.stats[RECLASSIFY_FAILURE_STAT] == 1
    assert any('covered under' in w.getMessage()
               for w in _warnings(caplog, S6_LOGGER))


def test_reclassify_faithful_reply_is_accepted(monkeypatch):
    reply = '\n'.join(f'KEEP: {line}' for line in _FUSED_LINES)
    monkeypatch.setattr(s6, 'call_llm', _reply_llm(reply))
    gen = WCMTemplateGenerator(verbose=False)
    segments = gen._reclassify_entry_segments(_FUSED_TEXT, 'T')
    assert [text for text, _ in segments] == _FUSED_LINES
    assert gen.stats[RECLASSIFY_FAILURE_STAT] == 0


def test_reconsider_keeps_the_whole_entry_when_the_reply_is_a_summary(
        monkeypatch):
    """End to end through `_reconsider_appendix_entries`: the summary reply is
    refused, so every source line lands in the appendix, none vanishes."""
    summary = 'KEEP: [All 12 publication entries follow, first Doe J0 last Doe J11]'
    monkeypatch.setattr(s6, 'call_llm', _reply_llm(summary))
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen._appendix_pending = [({'text': _FUSED_TEXT, 'taxonomy_code': 'T',
                               'extracted_fields': {}}, 5.0)]
    gen._reconsider_appendix_entries()
    body = '\n'.join(p.text for p in gen.doc.paragraphs)
    assert all(line in body for line in _FUSED_LINES)
    assert 'publication entries follow' not in body


def test_reclassify_failure_reaches_the_render_warnings_sidecar(
        monkeypatch, tmp_path):
    monkeypatch.setattr(s6, 'call_llm', _raising_llm)
    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    real_reconsider = gen._reconsider_appendix_entries

    def seed_then_reconsider():
        # Queue one overflow entry, then run the REAL reconsideration so the
        # real _reclassify_entry_segments except-arm fires inside generate().
        gen._appendix_pending = [({'text': _RECLASSIFY_TEXT,
                                   'taxonomy_code': 'P',
                                   'extracted_fields': {}}, 40)]
        return real_reconsider()

    gen._reconsider_appendix_entries = seed_then_reconsider
    input_path = tmp_path / 'in.json'
    input_path.write_text(json.dumps({
        'document_uid': 'T652G',
        'entries': [{'text': 'Name: Jane Q. Public, MD', 'taxonomy_code': 'A',
                     'extracted_fields': {}, 'element_idx_start': 0}]}))
    gen.generate(str(input_path), str(tmp_path / 'out.docx'))
    sidecar = json.loads((tmp_path / 'T652G_render_warnings.json').read_text())
    found = [w for w in sidecar['warnings']
             if w.get('check') == RECLASSIFY_FAILURE_STAT]
    assert len(found) == 1
    assert found[0]['severity'] == 'WARN'
    # Literals, not the constant: lint_stage6_warnings copies evidence and
    # message verbatim into the doctor report and onto the Teams card.
    assert found[0]['evidence'] == ['segment_reclassification_failures=1']
    assert found[0]['section'] == 'appendix'
    assert found[0]['message'] == (
        '1 appendix entry reclassification(s) failed or came back incomplete; those entries stayed in '
        'the appendix whole instead of being split and routed to their sections')


def test_reconsider_sends_entry_to_appendix_when_reclassify_fails(monkeypatch):
    """The None fallback reaches the caller: the entry goes to the appendix
    with its original text and code instead of being dropped."""
    monkeypatch.setattr(s6, 'call_llm', _raising_llm)
    gen = WCMTemplateGenerator(verbose=False)
    sent = []
    gen._add_remaining_to_appendix = lambda items: sent.extend(items) or []
    entry = {'text': _RECLASSIFY_TEXT, 'taxonomy_code': 'P',
             'extracted_fields': {}}
    gen._appendix_pending = [(entry, 40)]
    gen._reconsider_appendix_entries()
    assert sent == [(_RECLASSIFY_TEXT, 'P', 40)]
