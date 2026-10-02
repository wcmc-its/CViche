"""Regression tests for the code-'A' gap: unconsumed Personal Data entries were
silently deleted, and one class of them was silently rendered.

`mapped_codes` lists 'A' with the comment "unused A entries go to appendix".
They did not: the exclusion covered the whole code, on the stated assumption
that "A entries are all used in Personal Data section". `_fill_personal_data`
reads only phone/address/email, so an entry carrying "Citizenship: US" or
"Fax: ..." is consumed by nothing and was then excluded from the appendix too.
146 such orphans exist across 80 of the 100 corpus CVs.

The opposite failure also exists and is worse: web07's Office address row
renders "Cincinnati, Ohio" today, lifted straight out of
"PLACE OF BIRTH: Cincinnati, Ohio" by the address catch-all.

The two paths need DIFFERENT granularities, which is the whole point of these
tests:

- consumption path -> deny by value PROVENANCE. Dropping the whole entry breaks
  three corpus CVs whose contact block also carries a birth date, costing them
  real office addresses, an office phone and a work email.
- appendix path -> deny the whole ENTRY. It renders nothing, so discarding it
  is free, and fragment-level filtering would keep the birth date and drop only
  its label.

Text is read by walking w:t nodes, NOT Paragraph.text, which is blind to runs
inside <w:ins> tracked changes and under-reports by 11-19% (#461).

    python3 -m pytest src/unified_pipeline/tests/test_stage6_personal_data_recovery.py -p no:cacheprovider
"""

import json
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6.normalization.pii import (  # noqa: E402
    SCOPE_ALL_CODES,
    _PII_FIELD_KEY_RE,
)
from unified_pipeline.stage6.pii_pass import (  # noqa: E402
    WITHHELD_COMMENT_AUTHOR,
    WITHHELD_COMMENT_HEADER,
    run_pii_pass,
)
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    PII_REDACTED_NOTICE,
    RENDER_ROUTED_CODES,
    TAXONOMY_TO_SECTION,
    WCMTemplateGenerator,
    _pii_fragments,
)

W_T = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"


def _all_text(docx_path) -> str:
    doc = Document(str(docx_path))
    return "\n".join(n.text or "" for n in doc.element.body.iter(W_T))


def _list_item_texts(docx_path) -> set[str]:
    """Every paragraph's text (raw w:t join, matching `_paragraph_texts`)
    for paragraphs that carry `w:numPr` -- i.e. render as a real Word list
    item rather than a literal bullet glyph. `_add_remaining_to_appendix`
    switched from `f"• {text}"` to `_apply_list_bullet` (#864), so an
    appendix residual's paragraph text no longer starts with the glyph; the
    list membership is asserted through the numbering property instead."""
    doc = Document(str(docx_path))
    out = set()
    for para in doc.paragraphs:
        pPr = para._p.pPr
        numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
        if numPr is not None:
            out.add("".join(n.text or "" for n in para._p.iter(W_T)))
    return out


def _render(tmp_path, entries, cv_owner=None, **generator_kwargs) -> str:
    gen = WCMTemplateGenerator(verbose=False, **generator_kwargs)
    gen._reconsider_appendix_entries = lambda: None
    ip, op = tmp_path / "in.json", tmp_path / "out.docx"
    payload = {"document_uid": "TESTPD", "entries": entries}
    if cv_owner:
        payload["cv_owner"] = cv_owner
    ip.write_text(json.dumps(payload))
    gen.generate(str(ip), str(op), research_summary_path=None)
    return _all_text(op)


def _entry(text, code, fields=None, idx=0):
    return {"text": text, "taxonomy_code": code,
            "extracted_fields": fields or {}, "element_idx_start": idx}


def _comments(docx_path) -> list[tuple[str, str]]:
    """(author, text) of every comment in the docx's comments part, text
    lines joined with newlines; [] when the part does not exist."""
    with zipfile.ZipFile(docx_path) as z:
        if "word/comments.xml" not in z.namelist():
            return []
        root = ET.fromstring(z.read("word/comments.xml"))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    out = []
    for c in root.iter(f"{W}comment"):
        lines = ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in c.iter(f"{W}p")]
        out.append((c.get(f"{W}author"), "\n".join(lines)))
    return out


def _paragraph_texts(docx_path) -> list[str]:
    """Every paragraph's text, read from raw w:t nodes like `_all_text` --
    one string per PARAGRAPH, which `_all_text` cannot give (it joins node
    by node), and which an assertion about what a line renders AS rather
    than what it CONTAINS needs."""
    doc = Document(str(docx_path))
    return ["".join(n.text or "" for n in para._p.iter(W_T))
            for para in doc.paragraphs]


def _a(text, fields=None, idx=0):
    return {"text": text, "taxonomy_code": "A",
            "extracted_fields": fields or {}, "element_idx_start": idx}


# --------------------------------------------------------------------------
# consumption path
# --------------------------------------------------------------------------

def test_birthplace_is_not_rendered_as_an_address(tmp_path):
    """web07's shape: stage 4 extracted the birthplace into `address`, and the
    address catch-all put it in the Office address row."""
    text = _render(tmp_path, [
        _a("PLACE OF BIRTH: Cincinnati, Ohio", {"address": "Cincinnati, Ohio"}),
    ])
    assert "Cincinnati, Ohio" not in text, "birthplace rendered as an address"


def test_real_contact_data_survives_an_entry_that_also_carries_pii(tmp_path):
    """web086's shape, and the reason this denies by value rather than by entry.

    Dropping the whole entry -- the obvious implementation -- costs this CV its
    office address and office phone. This test fails if someone makes that
    change.
    """
    text = _render(tmp_path, [
        _a("Born: December 13, 1965\nOffice: 8200 East Belleview Ave, Greenwood Village, CO\n"
           "Phone: (303) 694-3333",
           {"address": "8200 East Belleview Ave, Greenwood Village, CO",
            "phone": "(303) 694-3333"}),
    ])
    assert "8200 East Belleview Ave" in text, "real office address was dropped with the PII"
    assert "(303) 694-3333" in text, "real office phone was dropped with the PII"
    assert "December 13, 1965" not in text, "birth date reached the document"


def test_email_regex_fallback_does_not_harvest_from_a_pii_fragment(tmp_path):
    text = _render(tmp_path, [
        _a("Spouse's name: Pat Doe (pat.doe@example.com)"),
    ])
    assert "pat.doe@example.com" not in text


def test_personal_data_path_keeps_phone_and_email_beside_a_birth_date(tmp_path: Path) -> None:
    """#735 review item 3: the same entry carries a birth date AND a real
    office phone AND a work email; the consumption path denies by value, so
    both live contact values survive while only the birth date is denied.
    Extends `test_real_contact_data_survives_an_entry_that_also_carries_pii`
    (address + phone) with the one field it does not cover, email.

    The email is deliberately absent from the raw text (unlike address and
    phone, which the entry's raw text also carries): `_fill_personal_data`
    has a second, regex-based email fallback that scans raw text
    independently of `extracted_fields['email']`, and it would mask a
    regression in the field-based classification this test targets if the
    email string were also sitting in the text for it to find."""
    text = _render(tmp_path, [
        _a("Born: April 1, 1958\n"
           "Office: 1300 York Avenue\n"
           "Phone: (212) 555-0100",
           {"address": "1300 York Avenue",
            "phone": "(212) 555-0100",
            "email": "jane.roe@med.example.edu"}),
    ])
    assert "(212) 555-0100" in text, "real office phone was dropped with the PII"
    assert "jane.roe@med.example.edu" in text, "real work email was dropped with the PII"
    assert "April 1, 1958" not in text, "birth date reached the document"


# --------------------------------------------------------------------------
# appendix recovery path
# --------------------------------------------------------------------------

def test_unconsumed_entry_reaches_the_appendix(tmp_path):
    """The bug: an A entry supplying no phone/address/email rendered nowhere."""
    text = _render(tmp_path, [
        _a("Citizenship: United States of America"),
        _a("Foreign Languages: Spanish, Portuguese"),
    ])
    assert "United States of America" in text, "citizenship was silently deleted"
    assert "Spanish, Portuguese" in text, "languages were silently deleted"


def test_consumed_entry_is_not_duplicated_into_the_appendix(tmp_path):
    """An entry that DID fill a slot must not also be appended."""
    text = _render(tmp_path, [
        _a("Office: 1300 York Avenue, New York, NY",
           {"address": "1300 York Avenue, New York, NY"}),
    ])
    assert text.count("1300 York Avenue") == 1, "consumed address duplicated"


def test_name_banner_already_on_the_page_is_not_appended_again(tmp_path):
    """The redundancy filter, and the reason recovery runs post-render.

    The commonest orphan by far is the faculty member's own name/title banner:
    76 of the corpus's 146, of which 75 already render -- from `cv_owner`, not
    from the A entry, so the entry itself looks unconsumed. Without this filter
    every one of them is appended a second time.
    """
    owner = "Jane Q. Public, M.D., Ph.D."
    text = _render(tmp_path, [_a(owner)], cv_owner={"full_name_with_credentials": owner})
    assert text.count(owner) == 1, "name banner duplicated into the appendix"


def test_pii_orphan_is_redacted_with_a_visible_notice(tmp_path):
    """Withheld, but not silently: the reader is told something was removed."""
    text = _render(tmp_path, [
        _a("Date of Birth: 12/13/1947"),
        _a("Marital Status: Married, spouse Jane Roe"),
        _a("Citizenship: Canada"),
    ])
    assert "12/13/1947" not in text, "date of birth reached the document"
    assert "Jane Roe" not in text, "spouse name reached the document"
    assert PII_REDACTED_NOTICE in text, "content was withheld with no indication"
    assert "Canada" in text, "non-PII entry withheld along with the PII"


def test_home_phone_only_orphan_is_withheld_not_appended(tmp_path):
    """#821: the template has no home-phone row, so an entry supplying only
    a home phone was ALREADY unconsumed by `_fill_personal_data` (nothing
    excludes `home_phone` from the appendix-orphan check the way a
    classified value does) -- before the #821 policy row, that raw entry
    text reached the Appendix unfiltered as a safety-net "orphan"
    (test_unconsumed_entry_reaches_the_appendix, above, is exactly that
    mechanism). Now it is withheld with notice like everything else."""
    text = _render(tmp_path, [
        _a("Home phone: 555-111-2222", {"phone": "555-111-2222"}),
    ])
    assert "555-111-2222" not in text, "a home phone reached the document"
    assert PII_REDACTED_NOTICE in text


def test_one_redaction_notice_per_document(tmp_path):
    text = _render(tmp_path, [
        _a("Date of Birth: 12/13/1947"),
        _a("Place of Birth: Rochester, NY"),
        _a("Marital Status: Married"),
    ])
    assert text.count(PII_REDACTED_NOTICE) == 1


def test_pii_named_only_by_field_key_is_caught(tmp_path):
    """6 of the corpus PII entries carry no recognisable text label; stage 4
    names the field instead."""
    text = _render(tmp_path, [
        _a("Additional information", {"marital_status_children": "Ann, Bob"}),
    ])
    assert "Ann, Bob" not in text
    assert PII_REDACTED_NOTICE in text


def test_appendix_removal_catches_raw_text_and_field_key_pii_alike(tmp_path: Path) -> None:
    """#735 review item 1, end-to-end through the real appendix removal site
    (`_unconsumed_personal_data_batch`, `stage_6_word_template.py:2390`): an
    A-coded entry whose RAW TEXT carries a PII label, and a sibling entry
    whose text carries no label but whose `extracted_fields` alone carries a
    PII key, are both absent from the generated document with one shared
    withheld notice -- and a third, non-PII sibling in the same appendix
    batch still renders. `test_pii_orphan_is_redacted_with_a_visible_notice`
    and `test_pii_named_only_by_field_key_is_caught` pin the two halves
    separately; this pins them together in one batch, which is what item 1
    actually asks for.

    #821 R2 F3 (issue #834) narrowed the raw-text-cut half of this to a
    per-VALUE deny (see the test right below) -- this test's middle entry
    stays a whole-ENTRY deny on purpose: its PII lives ONLY in
    `extracted_fields`, so `run_pii_pass` never touches `entry['text']` at
    all ("Additional information" is never cut), and there is no residual
    fragment for the #834 fix to have anything to recover."""
    text = _render(tmp_path, [
        _a("Date of Birth: 04/01/1958"),
        _a("Additional information", {"marital_status_spouse": "Pat Roe"}),
        _a("Foreign Languages: French"),
    ])
    assert "04/01/1958" not in text, "raw-text-labelled PII reached the document"
    assert "Additional information" not in text, (
        "field-key-only PII entry reached the document"
    )
    assert "Pat Roe" not in text, "field-key-only PII value reached the document"
    assert PII_REDACTED_NOTICE in text, "content was withheld with no indication"
    assert "French" in text, "a non-PII sibling in the same batch was dropped too"


def test_appendix_residual_survives_a_fused_withheld_and_kept_fragment(tmp_path: Path) -> None:
    """#821 R2 F3 (issue #834): an A-coded orphan whose raw text FUSES a
    withheld fragment with unrelated, non-PII content used to lose the
    WHOLE entry -- `_unconsumed_personal_data_batch` denied by ENTRY
    (`entry['_pii_withheld']`), the same granularity the appendix path
    deliberately uses for an entry that renders nothing on its own, but
    this entry is not "nothing": `run_pii_pass` (`pii_pass.py`'s
    `_cut_spans`) already cut only the withheld SPAN out of `entry['text']`,
    so the kept fragment survives in the entry's own residual text and
    should have reached the Appendix all along (the corpus shape: a fused
    "Home Phone" + "Citizenship" orphan). The residual renders as an
    ordinary Appendix bullet with its dangling separator punctuation
    trimmed (`_strip_dangling_separators`) and its normal classification
    comment attached, same as any other appendix line."""
    text = _render(tmp_path, [
        _a("Home Phone: 555-123-4567; Citizenship: US"),
    ], emit_comments=True)
    assert "Citizenship: US" in text, (
        "the kept fragment was dropped along with its withheld sibling"
    )
    assert "555-123-4567" not in text, "the withheld home phone reached the document"
    assert PII_REDACTED_NOTICE in text
    comments = _comments(tmp_path / "out.docx")
    assert any("Originally classified" in body for _, body in comments), (
        "the recovered residual line lost its usual classification comment"
    )


def test_withheld_comment_names_appendix_when_the_residual_rendered_there(tmp_path: Path) -> None:
    """#848: the withheld-item Word comment names the section the item's
    entry actually rendered in. A fused orphan whose residual reached the
    Appendix is "Appendix" -- not the A code's nominal "Personal Data" --
    while an item that never had a residual (a bare Date of Birth, nothing
    left to render) keeps its nominal label."""
    _render(tmp_path, [
        _a("Home Phone: 555-123-4567; Citizenship: US"),
        _a("Date of Birth: 04/01/1958"),
    ])
    bodies = [b for _, b in _comments(tmp_path / "out.docx") if " \u2022 " in b]
    assert len(bodies) == 1
    lines = [ln for ln in bodies[0].splitlines() if ln.startswith(" \u2022 ")]
    # Per category, not "any": an off-by-one in the entry index would swap
    # the two labels and still satisfy an any()-shaped check.
    phone = [ln for ln in lines if "home address / phone" in ln]
    dob = [ln for ln in lines if "date of birth" in ln]
    assert len(phone) == 1 and len(dob) == 1, lines
    assert phone[0].endswith("Appendix)"), phone
    assert dob[0].endswith("Personal Data)"), (
        "the bare Date of Birth has no residual and must keep its label")


def test_appendix_residual_is_refused_when_the_value_sits_past_a_hard_delimiter(
        tmp_path: Path) -> None:
    """#821 R2 F3 safety check, corpus-observed: a "Label: |
    value" shape -- a pipe, the pipeline's own inline-cell separator,
    sitting directly after the label's colon. `pii.py`'s label span runs
    "to the next hard delimiter", and `|` IS one, so an UNEXTENDED cut
    would remove only "Home telephone:" and leave the phone number
    completely uncut in `entry['text']` -- indistinguishable, at that
    point, from a safe residual (`_clean_inline_tabs` then drops the
    now-empty label cell entirely, leaving a bare, unlabelled phone
    number). `pii_pass.py`'s `_extend_bare_label_span` closes this at the
    pass itself (the phone number is now PART of the cut fragment, so this
    entry has nothing left at all -- see the sibling test right below for
    the case where a value trails the same shape); `_unconsumed_personal_
    data_batch`'s own `_pii_cut_left_a_bare_label` is the second layer,
    for any shape that extension does not catch."""
    text = _render(tmp_path, [
        _a("Home telephone: | 212 555 1234"),
    ])
    assert "212 555 1234" not in text, (
        "a home phone number separated from its label by a hard delimiter "
        "reached the document unlabelled"
    )
    assert "Home telephone" not in text
    assert PII_REDACTED_NOTICE in text


def test_appendix_residual_recovers_a_sibling_field_past_a_wide_gap(
        tmp_path: Path) -> None:
    """#821 R2 F3 (issue #834), the corpus shape that needed
    `_extend_bare_label_span`: a fused "Home Phone" + "Citizenship"
    A-coded orphan column-aligns its fields with a wide run of spaces
    instead of a colon-adjacent value ("Home Phone:        <phone>
    Citizenship: ..."). 3+ spaces is ALSO one of `pii.py`'s hard
    delimiters, so the UNEXTENDED label match stops at the colon, same as
    the pipe shape above -- but here a sibling field follows, not the end
    of the entry, so simply discarding the whole entry (as F3's original,
    unextended fix did) cost the citizenship line. The extension pulls the
    orphaned phone value into the SAME cut, stopping before the next hard
    delimiter, so the residual is exactly the sibling field with nothing
    of the phone left in it."""
    text = _render(tmp_path, [
        _a("Home Phone:        212 555 1234"
           "                                               Citizenship:  US"),
    ], emit_comments=True)
    assert "Citizenship" in text and "US" in text, (
        "the sibling field past the wide gap was dropped along with the phone"
    )
    assert "212 555 1234" not in text, "the withheld home phone reached the document"
    assert "Home Phone" not in text
    assert PII_REDACTED_NOTICE in text


def test_appendix_residual_renders_exactly_the_kept_fragment(tmp_path: Path) -> None:
    """#821 R3 F-E: the residual is asserted by EQUALITY on the rendered
    paragraph, not by substring. The cut stops at the hard delimiter that
    separated the withheld fragment from its neighbour, so the delimiter
    itself is left dangling on the front of the residual (`"; Citizenship:
    US"`); `_strip_dangling_separators` trims it. A substring assertion
    ("Citizenship: US" in text) passes either way and left that trim
    untested on the wire -- this one fails the moment the residual renders
    with its leading separator. The appendix bullet is a real Word list
    item (#864), not a literal glyph, so the equality check is against
    `_list_item_texts`, not a "\u2022 " prefix."""
    _render(tmp_path, [
        _a("Home Phone: 555-123-4567; Citizenship: US"),
    ])
    list_items = _list_item_texts(tmp_path / "out.docx")
    assert "Citizenship: US" in list_items, (
        "the residual did not render as a clean appendix list item: "
        f"{[x for x in _paragraph_texts(tmp_path / 'out.docx') if 'Citizenship' in x]}"
    )


def test_appendix_residual_renders_when_a_bare_label_had_no_value_after_it(
        tmp_path: Path) -> None:
    """#821 R3 F-D: a protected label with NOTHING after it -- a fielded
    template's leftover empty row -- is not a leak. The R2 refusal keyed on
    "the cut fragment ends in a colon", which is equally true of a label
    whose value survived the cut and of one that never had a value at all,
    so this entry lost its citizenship line (a #821 "render" item) to
    protect nothing. `run_pii_pass` now records which of the two happened
    (`_pii_orphaned_value`) and only the first refuses the residual."""
    text = _render(tmp_path, [
        _a("Citizenship: US\nHome Address:"),
    ])
    assert "Citizenship: US" in text, (
        "a bare protected label with no value after it withheld its sibling field"
    )
    assert "Home Address" not in text, "the protected label itself rendered"
    assert PII_REDACTED_NOTICE in text


def test_appendix_residual_is_refused_when_the_value_sits_on_the_next_line(
        tmp_path: Path) -> None:
    """The negative half of the case above, and the reason the refusal
    still exists: the same bare label, but with an unlabelled value on the
    line below it. The pass will not extend a cut across a newline (a
    newline is the source document's own field separator -- extending
    across one swallowed the NEXT field, #821 R3 F-B), so the address is
    still in the residual, unlabelled, indistinguishable from safe content.
    The whole entry is denied instead -- including its other, harmless
    field, which is the price of not leaking the address."""
    text = _render(tmp_path, [
        _a("Citizenship: US\nHome Address:\n12 Example Street"),
    ])
    assert "12 Example Street" not in text, (
        "an unlabelled home address reached the Appendix as a safe residual"
    )
    assert "Citizenship" not in text, (
        "the residual rendered even though the protected value was still in it"
    )
    assert PII_REDACTED_NOTICE in text


def test_personal_data_redacted_counts_every_entry_the_policy_cut(
        tmp_path: Path) -> None:
    """#821 R3 F-H: `stats['personal_data_redacted']` means "A-coded
    orphans the policy removed something from", so it counts the entry
    whose residual rendered as well as the one denied whole -- before the
    #834 residual fix those were the same population and the counter sat on
    the deny branch alone, which silently turned it into "entries dropped".
    `personal_data_recovered` is the count of what actually reached the
    Appendix, so the two still separate the cases: 2 cut, 2 rendered (the
    residual and the clean sibling)."""
    gen = WCMTemplateGenerator(verbose=False)
    gen._reconsider_appendix_entries = lambda: None
    ip, op = tmp_path / "in.json", tmp_path / "out.docx"
    ip.write_text(json.dumps({"document_uid": "TESTPD", "entries": [
        _a("Home Phone: 555-123-4567; Citizenship: US"),
        _a("Date of Birth: 04/01/1958", idx=1),
        _a("Foreign Languages: French", idx=2),
    ]}))
    gen.generate(str(ip), str(op), research_summary_path=None)
    assert gen.stats["personal_data_redacted"] == 2
    assert gen.stats["personal_data_recovered"] == 2


def test_pii_in_a_tab_separated_cell_is_still_caught(tmp_path):
    """The scan must read the RAW entry text. _clean_inline_tabs rewrites the
    first '\\t' to ': ' and ' | ' to ' — ', destroying the fragment boundaries
    the scan splits on -- scan the cleaned text instead and the PII cell merges
    into its neighbour, no longer starts a fragment, and renders."""
    text = _render(tmp_path, [
        _a("Office: 1300 York Ave\tDate of Birth: 12/13/1947"),
        _a("Home | Marital Status: Married"),
    ])
    assert "12/13/1947" not in text
    assert "Married" not in text
    assert PII_REDACTED_NOTICE in text


def test_web057_shape_is_withheld_from_a_t_coded_appendix_line(tmp_path):
    """#820's local instance, end to end: a T-coded line with the spouse
    label buried after another label (`Personal Information:: Husband:`)
    -- the unanchored-matching case, through the real Appendix path."""
    text = _render(tmp_path, [
        _entry("Personal Information:: Husband: Pat Example, MD", "T"),
        _entry("Foreign Languages: French", "T"),
    ])
    assert "Pat Example" not in text, "spouse name reached the Appendix"
    assert PII_REDACTED_NOTICE in text
    assert "French" in text


def test_notice_is_keyed_on_something_withheld_not_on_appendix_redaction(tmp_path):
    """M20: the notice used to be keyed on the A-orphan redaction count.
    PII withheld from a RENDERED section (an F1 entry's text) with no
    A-coded orphan at all must still produce the notice."""
    text = _render(tmp_path, [
        _entry("Medical license, Example State; SSN: 123-45-6789", "F1",
               {"license_number": "X1", "state": "Example State"}),
    ])
    assert "123-45-6789" not in text
    assert PII_REDACTED_NOTICE in text


def test_no_notice_when_the_only_label_is_out_of_scope(tmp_path):
    """The round-1 false notice: an ambiguous label on a routed content
    code is not withheld, so nothing is recorded and the document carries
    neither the notice nor a comment (the render gate's CHANGED 0).

    An F1 DEA entry used to ride alongside this as a second "out of scope"
    example; #821 flipped that decision (see
    test_dea_in_licensure_is_withheld_with_notice_and_comment below), so it
    is no longer a valid negative control here and was removed rather than
    updated in place -- the S4 title is this test's whole subject."""
    text = _render(tmp_path, [
        _entry("Children: Research, Practice and Policy. Example Press, 2001.", "S4",
               {"title": "Children: Research, Practice and Policy",
                "authors": "Roe J", "year": "2001", "publisher": "Example Press"}),
    ])
    assert "Children: Research" in text, "a content-code title was withheld"
    assert PII_REDACTED_NOTICE not in text
    assert _comments(tmp_path / "out.docx") == []


def test_dea_in_licensure_is_withheld_with_notice_and_comment(tmp_path):
    """#821: a licensure DEA entry is withheld with notice through the full
    render path -- the slot renders no value, the document-wide notice
    paragraph appears, and the Word comment names "DEA number" under
    "Licensure" (the same vocabulary/mechanism every other withheld
    category uses -- see WITHHELD_COMMENT_HEADER)."""
    text = _render(tmp_path, [
        _entry("DEA registration AB1234567", "F1", {"license_number": "AB1234567"}),
    ])
    assert "AB1234567" not in text
    assert PII_REDACTED_NOTICE in text
    comments = _comments(tmp_path / "out.docx")
    assert len(comments) == 1
    author, comment_text = comments[0]
    assert author == WITHHELD_COMMENT_AUTHOR
    assert comment_text.startswith(WITHHELD_COMMENT_HEADER)
    assert "DEA number" in comment_text
    assert "Licensure" in comment_text


def test_email_field_fallback_does_not_harvest_from_a_pii_fragment_on_any_code(tmp_path):
    """M23: the all-entries email fallback (`personal_data.py`, the
    `_pii_fragments` read at its second site) must consult the pass's
    stored fragments -- a T-coded emergency-contact entry whose
    `extracted_fields['email']` was lifted from the fragment must not
    supply the Work email slot."""
    text = _render(tmp_path, [
        _entry("Emergency Contact: Pat Example, pat.example@example.com", "T",
               {"email": "pat.example@example.com"}),
    ])
    assert "pat.example@example.com" not in text


def test_pass_runs_before_dedup_so_a_dropped_near_duplicate_is_still_scanned(tmp_path):
    """`_recover_unrendered_records` re-scans the PRE-dedup entries, so an
    entry dedup drops can still put a record line on the page. The pass
    therefore runs before the pre-dedup snapshot: the dropped
    near-duplicate's SSN is recorded (notice emitted) rather than skipped.
    With the pass after dedup this document carried no notice."""
    kept = ("2001 - 2002\tExample Award for Distinguished Example Work, Example Society of Examples\n"
            "2003 - 2004\tExample Prize for Outstanding Example Research, Example Foundation of Examples")
    dropped = ("2001 - 2002\tExample Award for Distinguished Example Work, Example Society\n"
               "2003 - 2004\tExample Prize for Outstanding Example Research; SSN: 123-45-6789")
    text = _render(tmp_path, [_entry(kept, "H"), _entry(dropped, "H", idx=1)])
    assert "123-45-6789" not in text
    assert PII_REDACTED_NOTICE in text
    assert "Example Award for Distinguished Example Work" in text


# --------------------------------------------------------------------------
# the Word comment on the notice (A-820 addendum)
# --------------------------------------------------------------------------

def test_withheld_notice_carries_one_word_comment_listing_categories(tmp_path):
    """Round trip: save, reopen the package, read the comment part back.
    Exactly one comment, on the notice, author "CViche", listing both
    categories with counts and sections and NOT the values -- and present
    with `emit_comments=False` (it is not a classification comment, #153)."""
    text = _render(tmp_path, [
        _entry("Date of Birth: 01/02/1970", "A"),
        _entry("Visa Status: O-1", "T"),
    ], emit_comments=False)
    assert "01/02/1970" not in text and "O-1" not in text
    assert PII_REDACTED_NOTICE in text
    comments = _comments(tmp_path / "out.docx")
    assert len(comments) == 1, comments
    author, body = comments[0]
    assert author == WITHHELD_COMMENT_AUTHOR
    lines = body.split("\n")
    assert lines[0] == WITHHELD_COMMENT_HEADER
    assert " • date of birth (1 item, Personal Data)" in lines
    assert " • visa / immigration status (1 item, Appendix)" in lines
    assert "01/02/1970" not in body and "O-1" not in body, "a withheld value re-leaked into the comment"


def test_the_comment_element_splits_one_w_p_per_line(tmp_path):
    """`_create_comments_xml` must emit one <w:p> per line of the comment
    text, not join every line into a single paragraph's <w:t> -- Word
    collapses an embedded newline inside one <w:t> to a space, running the
    header, every bullet and the footer together on one line (#820 R3
    finding 2). The round-trip helper `_comments()` re-groups text by
    whatever <w:p> boundaries already exist, so it reads a single paragraph
    holding embedded '\\n' characters identically to one <w:p> per line and
    cannot tell them apart -- this test reads the raw <w:p> COUNT instead."""
    text = _render(tmp_path, [
        _entry("Date of Birth: 01/02/1970", "A"),
        _entry("Visa Status: O-1", "T"),
    ], emit_comments=False)
    assert PII_REDACTED_NOTICE in text
    with zipfile.ZipFile(tmp_path / "out.docx") as z:
        root = ET.fromstring(z.read("word/comments.xml"))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    comments = root.findall(f"{W}comment")
    assert len(comments) == 1, comments
    paragraphs = comments[0].findall(f"{W}p")
    # header line + one bullet per category (date of birth, visa) + footer.
    expected = 1 + 2 + 1
    assert len(paragraphs) == expected, [
        "".join(t.text or "" for t in p.iter(f"{W}t")) for p in paragraphs]
    # And every <w:p> holds exactly one line's worth of text -- no line was
    # merged into its neighbor.
    for p in paragraphs:
        line_text = "".join(t.text or "" for t in p.iter(f"{W}t"))
        assert "\n" not in line_text


def test_withheld_comment_is_anchored_on_the_notice_paragraph(tmp_path):
    _render(tmp_path, [_entry("Date of Birth: 01/02/1970", "A")], emit_comments=False)
    doc = Document(str(tmp_path / "out.docx"))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    anchored = [p for p in doc.paragraphs
                if p._p.find(f"{W}commentRangeStart") is not None]
    assert len(anchored) == 1
    assert PII_REDACTED_NOTICE in anchored[0].text


def test_withheld_comment_is_the_only_comment_even_with_classification_comments_on(tmp_path):
    """With emit_comments=True the Appendix's other bullets get their
    classification comments; the notice paragraph gets ONLY the withheld
    summary, not a second "Originally classified A" comment."""
    _render(tmp_path, [
        _entry("Date of Birth: 01/02/1970", "A"),
        _entry("Foreign Languages: French", "A"),
    ], emit_comments=True)
    comments = _comments(tmp_path / "out.docx")
    assert len([c for c in comments if c[0] == WITHHELD_COMMENT_AUTHOR]) == 1
    assert any("Originally classified" in body for _, body in comments), (
        "the classification comments were not emitted with emit_comments=True")
    doc = Document(str(tmp_path / "out.docx"))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    notice = [p for p in doc.paragraphs if PII_REDACTED_NOTICE in p.text]
    assert len(notice) == 1
    assert len(notice[0]._p.findall(f"{W}commentRangeStart")) == 1


def test_notice_and_comment_still_emitted_with_recovery_disabled(tmp_path):
    """#820 R3 finding 3: `recover_unrendered_records=False` must disable
    ONLY the #221 record-line recovery -- never the withheld notice and its
    Word comment. The pass has already stripped the PII from the entry text
    either way; before this fix both were built inside
    `_recover_unrendered_records`, gated behind
    `if not self.recover_unrendered_records: return`, so setting the flag
    False silently suppressed the reader's only indication that something
    was withheld -- a second, silent loss stacked on top of the first."""
    text = _render(tmp_path, [
        _entry("SSN: 123-45-6789", "A"),
    ], recover_unrendered_records=False)
    assert "123-45-6789" not in text, "an SSN reached the rendered document"
    assert PII_REDACTED_NOTICE in text, (
        "the withheld notice was suppressed by an unrelated recovery flag")
    comments = _comments(tmp_path / "out.docx")
    assert len(comments) == 1, comments
    author, body = comments[0]
    assert author == WITHHELD_COMMENT_AUTHOR
    assert "social security number" in body.lower()
    assert "123-45-6789" not in body


# --------------------------------------------------------------------------
# the deny predicate itself
# --------------------------------------------------------------------------

def test_deny_predicate_requires_a_colon_terminator():
    """The colon terminator is load-bearing in two separate ways, and each
    relaxation deletes real CV content.

    Allowing '-' or an en-dash as a terminator swallows publication and award
    titles that merely begin with a listed word. Allowing slop before the colon
    ("children" + anything + ':') swallows any title containing one. All of
    these strings, or their shapes, occur in the corpus.
    """
    for keeper in [
        # would match if '-' / en-dash were accepted as a terminator
        "Born - Digital: A Study of Youth Media Practices",
        "Children - Oncology Group Consortium, Emeritus",
        "Spouse – A Documentary Film Review",
        # would match if slop were allowed before the colon
        "Children and Fire: Research on Burn Prevention",
        "Children of the Pandemic: A Longitudinal Cohort",
        "Born in the USA: Birth Cohort Methods",
        # ordinary corpus content that must never match
        "Children's Oncology Group - Emeritus",
        "Children's Hospital Colorado - Pillar Award",
        "Religion and Healing in America",
    ]:
        assert not _pii_fragments(keeper), f"false positive on {keeper!r}"


def test_deny_predicate_catches_the_real_labels():
    for pii in [
        "Date of Birth: 12/13/1947",
        "PLACE OF BIRTH: Cincinnati, Ohio",
        "Born: December 13, 1965",
        "Marital Status: Married",
        "Spouse's name: Jane Roe",
        "Children: Ann, Bob",
        "D.O.B.: 1/1/1970",
    ]:
        assert _pii_fragments(pii), f"missed {pii!r}"


def test_person_field_keys_accept_a_suffix():
    """Stage 4 names these keys freely, so every person stem takes a suffix --
    'wife_name' must be caught wherever 'spouse_name' is. Anchoring is what
    keeps ordinary keys out, so the suffix costs no precision."""
    for key in ["spouse_name", "wife_name", "husband_name", "children_names",
                "dependents_names", "marital_status_children",
                "personal_wife", "birthplace", "dob", "ssn"]:
        assert _PII_FIELD_KEY_RE.match(key), f"missed {key!r}"
    for keeper in ["institutional_email", "office_address", "housewife"]:
        assert not _PII_FIELD_KEY_RE.match(keeper), f"false positive {keeper!r}"


def test_research_vocabulary_still_off_the_list_is_not_denied():
    """sex/race are still deliberately NOT on the list -- ordinary research
    vocabulary that occurs as publication titles, unchanged by #820.

    `gender`/`ethnicity`/`religion` ARE on the list since #820, but only at
    the Personal Data / Appendix scope (`SCOPE_PERSONAL_AND_APPENDIX`): the
    #473 control "Gender: A Review of the Literature" is restored for every
    routed content code, where the ambiguous rows never apply -- see
    `test_protected_class_labels_deny_only_the_colon_form` for the two
    scopes side by side."""
    for keeper in [
        "Sex: differences in galanin expression",
        "Race: reporting practices in clinical trials",
    ]:
        assert not _pii_fragments(keeper), f"false positive on {keeper!r}"
    for content_title in [
        "Gender: A Review of the Literature",
        "Children: Research, Practice and Policy",
        "Health: A Journal Title",
    ]:
        assert not _pii_fragments(content_title, SCOPE_ALL_CODES), (
            f"content-code false positive on {content_title!r}")


def test_protected_class_labels_deny_only_the_colon_form():
    """#820 comment: religion/ethnicity/gender/veteran status/disability/
    health/blood type join the vocabulary as colon-terminated labels --
    every one of these leaked on a real CV before this ticket.

    The colon terminator is the same guard #473 built for
    marital-status/spouse/children, applied to this class for the first
    time: a bare occurrence of the word, with no colon directly after it,
    must still pass through untouched -- these are the ticket's own named
    negative controls (`docs` A-820, MUST NOT list). A *colon-terminated*
    title starting with one of these words ("Gender: A Review of the
    Literature") is denied ONLY where the entry is A-coded or Appendix-bound
    (#820 round 2's scope split); on a routed content code it renders.
    """
    for pii in [
        "Religion: Catholic",
        "Ethnicity: Hispanic",
        "Gender: Female",
        "Veteran Status: Yes",
        "Disability: None",
        "Health: Good",
        "Blood Type: O+",
    ]:
        assert _pii_fragments(pii), f"missed {pii!r}"
    for keeper in [
        "Health Sciences",
        "Public Health",
        "Gender Medicine",
        "Veterans Affairs Medical Center",
        "Age-related macular degeneration",
    ]:
        assert not _pii_fragments(keeper), f"false positive on {keeper!r}"


# --------------------------------------------------------------------------
# #821 R4 F-1: a bare-label cut may stop only at a KNOWN field label
#
# The round-3 stop searched the run for "1-4 word tokens then a colon", a
# WORD SHAPE, and the protected value's own trailing words satisfy it: a
# home address followed on the same run by a sibling field had its
# city/state -- or, for an address with no digits, nearly all of it --
# left in the residual and rendered into the Appendix. The stop is now a
# VOCABULARY (`normalization/pii.py::_KNOWN_FIELD_LABEL_RE`, built from
# `WITHHOLD_POLICY`'s label rows plus `_RENDER_SET_FIELD_LABELS`), so it
# can only stop where a field is actually NAMED, and a run with no known
# label in it is cut whole.
#
# Every case is asserted twice: on the pass's own residual by EQUALITY
# (`run_pii_pass` -- the leak-proof statement: the residual contains no
# token of the value), and on the rendered document (`_render` -- the
# wire: the survivor reached it, every forbidden token did not). All
# values synthetic.
# --------------------------------------------------------------------------

_KNOWN_LABEL_CASES = [
    # (id, text, residual, survives_in_document, forbidden_in_document)
    ("value_tail_is_words_then_a_known_label",
     "Home Address:        12 Elm St New York NY Citizenship: US",
     "Citizenship: US", "• Citizenship: US", ["Elm", "New York"]),
    ("value_is_all_words_then_a_known_label",
     "Home Address:        Elm House Oak Lane Citizenship: US",
     "Citizenship: US", "• Citizenship: US", ["Elm", "Oak Lane", "House"]),
    ("pipe_value_tail_is_words_then_a_known_label",
     "Home Address: | 12 Elm St Apt 4 Anytown Citizenship: US",
     "Citizenship: US", "• Citizenship: US", ["Elm", "Anytown", "Apt"]),
    ("value_carries_its_own_colon",
     # "Apt:" is colon-terminated and would stop a shape-based search
     # dead, leaving "4 Anytown" in the residual. It is not a field this
     # codebase renders, so it is not in the vocabulary and is cut with
     # the rest of the value.
     "Home Address: | 12 Elm St Apt: 4 Anytown Citizenship: US",
     "Citizenship: US", "• Citizenship: US", ["Elm", "Anytown", "Apt"]),
    ("gapped_phone_then_a_known_label",
     # The corpus shape the extension was built for: a column-aligned
     # field pair whose label is padded to a fixed width with spaces, a
     # phone value, then a sibling field one plain space later.
     "Home Phone:        555-0100 Citizenship: US",
     "Citizenship: US", "• Citizenship: US", ["555-0100"]),
    ("gapped_phone_then_a_known_render_set_label",
     # The survivor here is a render-set field rather than a policy
     # "render" default, and it reaches the PERSONAL DATA table's Work
     # email cell rather than the Appendix -- which is why the document
     # assertion is a substring and the residual assertion is equality.
     "Home Phone:        555-0100 Work Email: someone@example.org",
     "Work Email: someone@example.org", "someone@example.org", ["555-0100"]),
    ("no_known_label_in_the_run_cuts_the_whole_run",
     # The stated cost of a vocabulary: "Foo:" names no field this
     # pipeline renders, so the run is cut whole. The sibling is lost;
     # nothing of the address is leaked.
     "Home Address:        12 Elm St Anytown Foo: bar",
     "", None, ["Elm", "Anytown", "Foo", "bar"]),
    ("a_known_label_glued_into_a_word_is_not_a_label",
     # The word-start anchor. Without it the search stops on the
     # "citizenship" INSIDE "Noncitizenship" and everything before it --
     # the whole address -- renders.
     "Home Address: | 12 Elm St Noncitizenship: none",
     "", None, ["Elm", "Noncitizenship", "citizenship"]),
    ("newline_before_a_known_label_is_untouched",
     "Home Address:\nCitizenship: US",
     "Citizenship: US", "• Citizenship: US", ["Home Address"]),
    ("newline_before_a_numbered_known_label_is_untouched",
     "2. Home Address:\n3. Work Email: someone@example.org",
     "2. \n3. Work Email: someone@example.org", "someone@example.org",
     ["Home Address"]),
    ("delimiter_then_newline_refuses_the_residual",
     # #821 R3 F-2 / the verifier's g13b mutant: the `text[value_start]
     # == "\n"` guard in `_extend_bare_label_span`. The delimiter is a
     # pipe, so the extension starts, but the value is on the NEXT line --
     # the one place a cut may not reach. The value stays in the residual
     # uncut, the pass flags the entry (`_pii_orphaned_value`), and the
     # Appendix refuses it whole. Without the guard the address renders.
     "Home Address: | \n12 Elm St",
     "| \n12 Elm St", None, ["Elm"]),
]


@pytest.mark.parametrize(
    "text,residual,survives,forbidden",
    [(text, residual, survives, forbidden)
     for _, text, residual, survives, forbidden in _KNOWN_LABEL_CASES],
    ids=[case_id for case_id, *_ in _KNOWN_LABEL_CASES],
)
def test_a_bare_label_cut_stops_only_at_a_known_field_label(
        tmp_path: Path, text: str, residual: str, survives: str | None,
        forbidden: list[str]) -> None:
    """#821 R4 F-1, both halves in one case list.

    `residual` is `entry['text']` after the real `run_pii_pass`, asserted
    by equality: that is the statement "no token of the protected value is
    left behind", which a substring assertion cannot make. `survives` is
    what must reach the rendered document (None: nothing from this entry
    may), and `forbidden` is every value token that must not, read from
    the whole document rather than from one paragraph. A leading "• "
    on `survives` is a case-list sentinel meaning "must render as the
    appendix's Word list item", not a literal glyph (#864) -- checked via
    `_list_item_texts`, not a text-prefix match."""
    entry = {"text": text, "taxonomy_code": "A", "extracted_fields": {}}
    run_pii_pass({"A": [entry]}, routed_codes=RENDER_ROUTED_CODES,
                 section_names=TAXONOMY_TO_SECTION)
    assert entry["text"] == residual, (
        "the pass's residual is not exactly the sibling field")

    rendered = _render(tmp_path, [_a(text)])
    for token in forbidden:
        assert token not in rendered, f"{token!r} reached the document"
    if survives is None:
        assert "Citizenship: US" not in rendered
    elif survives.startswith("•"):
        expected_item = survives[len("• "):]
        assert expected_item in _list_item_texts(tmp_path / "out.docx"), (
            "the residual did not render as exactly that appendix list item: "
            f"{[p for p in _paragraph_texts(tmp_path / 'out.docx') if p.strip()][-6:]}")
    else:
        assert survives in rendered, "the sibling field did not reach the document"
    assert PII_REDACTED_NOTICE in rendered


def test_redaction_notice_is_not_reported_as_a_recovered_appendix_entry(tmp_path):
    """#531 x #820: `_recover_unrendered_records` reports the code of every
    appendix bullet it recovered so the render_warnings sidecar can name it
    as an `appendix_diversion`. The withheld notice is a bullet in that
    appendix too, but it is a notice, not a recovered entry -- it must not
    show up in the sidecar as an A-coded `recovered_unrendered` diversion."""
    _render(tmp_path, [
        _a("Date of Birth: 12/13/1947"),
        _a("Marital Status: Married, spouse Jane Roe"),
    ])
    sidecar = json.loads((tmp_path / "TESTPD_render_warnings.json").read_text())
    recovered = [w for w in sidecar["warnings"]
                 if w.get("check") == "appendix_diversion"
                 and w.get("reason") == "recovered_unrendered"]
    assert recovered == [], recovered


# --------------------------------------------------------------------------
# the template's own visa rows (#897)
# --------------------------------------------------------------------------

def _personal_data_rows(docx_path) -> dict[str, str]:
    """label -> value for the PERSONAL DATA table, located by its "Work
    email:" cell the way `_write_personal_data_table_cells` locates it."""
    doc = Document(str(docx_path))
    for table in doc.tables:
        cells = {c.text.strip() for r in table.rows for c in r.cells}
        if "Work email:" in cells:
            return {r.cells[0].text.strip(): r.cells[1].text.strip() for r in table.rows}
    raise AssertionError("no PERSONAL DATA table in the render")


def test_the_visa_answers_fill_the_template_slots_not_the_placeholder(tmp_path):
    """A CV already in the WCM template answers the two visa rows as
    label|value entries. Before #897 nothing consumed them: the eligibility
    slot rendered the template's own "Yes/No" placeholder, so the faculty's
    "No" read back as unanswered, and the row was not in the Appendix
    either (stage 2 had dropped it -- the other half of #897)."""
    gen = WCMTemplateGenerator(verbose=False)
    gen._reconsider_appendix_entries = lambda: None
    ip, op = tmp_path / "in.json", tmp_path / "out.docx"
    ip.write_text(json.dumps({"document_uid": "TESTPD", "entries": [
        _a("Is your eligibility to work in the U.S. based on an employment visa?: | Yes"),
        _a("If yes, please provide Visa type (Examples: J-1, H-1B, E-3, TN, etc.): | H-1B"),
    ]}))
    gen.generate(str(ip), str(op), research_summary_path=None)

    rows = _personal_data_rows(op)
    assert rows["Is your eligibility to work in the U.S. based on an employment visa?:"] == "Yes"
    assert rows["If yes, please provide Visa type (Examples: J-1, H-1B, E-3, TN, etc.):"] == "H-1B"
    # consumed into the slots, so neither row is appended as an
    # "<label> — <value>" Appendix residual
    text = _all_text(op)
    assert "— Yes" not in text and "— H-1B" not in text


def test_an_unanswered_visa_row_leaves_the_placeholder_alone(tmp_path):
    gen = WCMTemplateGenerator(verbose=False)
    gen._reconsider_appendix_entries = lambda: None
    ip, op = tmp_path / "in.json", tmp_path / "out.docx"
    ip.write_text(json.dumps({"document_uid": "TESTPD", "entries": [
        _a("Is your eligibility to work in the U.S. based on an employment visa?: | Yes/No"),
    ]}))
    gen.generate(str(ip), str(op), research_summary_path=None)

    rows = _personal_data_rows(op)
    assert rows["Is your eligibility to work in the U.S. based on an employment visa?:"] == "Yes/No"


# --------------------------------------------------------------------------
# slot ranking and phone-type keywords (#1222)
# --------------------------------------------------------------------------

def _contact_rows(tmp_path, entries) -> dict[str, str]:
    """The Personal Data table as {label cell: value cell}, read from raw
    w:t nodes so tracked insertions count."""
    _render(tmp_path, entries)
    rows = {}
    for table in Document(str(tmp_path / "out.docx")).tables:
        for row in table.rows:
            cells = ["".join(n.text or "" for n in c._tc.iter(W_T)).strip()
                     for c in row.cells]
            if len(cells) >= 2 and cells[0].endswith(":"):
                rows.setdefault(cells[0], cells[1])
    return rows


def test_labelled_work_address_outranks_an_earlier_banner_address(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Example University", {"address": "Example University"}, idx=0),
        _a("Business Address: 1 Sample Way, Exampleton, ZZ 00000\nPhone: 555-0100",
           {"address": "1 Sample Way, Exampleton, ZZ 00000",
            "phone": "555-0100"}, idx=1),
    ])
    assert "1 Sample Way" in rows["Office address:"]
    assert "Example University" not in rows["Office address:"]


def test_first_labelled_work_address_is_not_displaced_by_a_later_one(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Office Address: 1 Sample Way, Exampleton, ZZ 00000",
           {"address": "1 Sample Way, Exampleton, ZZ 00000"}, idx=0),
        _a("Work Address: 2 Other Road, Exampleton, ZZ 00000",
           {"address": "2 Other Road, Exampleton, ZZ 00000"}, idx=1),
    ])
    assert "1 Sample Way" in rows["Office address:"]
    assert "2 Other Road" not in rows["Office address:"]


def test_unlabelled_address_still_fills_an_empty_slot(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("3 Plain Street, Exampleton, ZZ 00000",
           {"address": "3 Plain Street, Exampleton, ZZ 00000"}),
    ])
    assert "3 Plain Street" in rows["Office address:"]


def test_labelled_work_address_never_surfaces_a_home_address(tmp_path):
    text = _render(tmp_path, [
        _a("Home Address: 9 Private Lane, Exampleton, ZZ 00000",
           {"address": "9 Private Lane, Exampleton, ZZ 00000"}, idx=0),
        _a("Business Address: 1 Sample Way, Exampleton, ZZ 00000",
           {"address": "1 Sample Way, Exampleton, ZZ 00000"}, idx=1),
    ])
    assert "1 Sample Way" in text
    assert "9 Private Lane" not in text


def test_cellular_in_a_department_name_does_not_make_the_phone_a_cell(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Department of Cellular Example Studies\nTel: 555-0100",
           {"phone": "555-0100"}),
    ])
    assert rows["Office telephone:"] == "555-0100"
    assert rows["Cell phone:"] == ""


def test_a_cell_label_still_routes_to_cell_phone(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Cell: 555-0100", {"phone": "555-0100"}),
    ])
    assert rows["Cell phone:"] == "555-0100"


@pytest.mark.parametrize("text", [
    "Mobile telephone: 555-0100",
    "Mobile Phone Number: 555-0100",
    "555-0100 cell\tdoe@example.com",
    "Cellphone: 555-0100",
    "Cell - 555-0100",
    "Cell/Text: 555-0100",
    "Mobile Ph: 555-0100",
    "Cell (preferred): 555-0100",
    "Cell phone 555-0100",
    "Name (cell) 555-0100",
])
def test_every_written_cell_label_shape_still_routes_to_cell_phone(tmp_path, text):
    rows = _contact_rows(tmp_path, [_a(text, {"phone": "555-0100"})])
    assert rows["Cell phone:"] == "555-0100"
    assert rows["Office telephone:"] == ""


def test_labelled_work_phone_outranks_an_earlier_unlabelled_phone(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Dr Example Doe\n1 Plain Street, Exampleton, ZZ 00000\n555-0101",
           {"phone": "555-0101"}, idx=0),
        _a("Primary Work Address:\t1 Sample Way\tExampleton, ZZ 00000\tPhone 555-0102",
           {"phone": "555-0102", "address": "1 Sample Way, Exampleton, ZZ 00000"}, idx=1),
    ])
    assert rows["Office telephone:"] == "555-0102"
    assert "1 Sample Way" in rows["Office address:"]


def test_labelled_home_phone_never_fills_the_office_row(tmp_path):
    text = _render(tmp_path, [
        _a("Home Phone: 555-0199", {"phone": "555-0199"}, idx=0),
        _a("Office Phone: 555-0102", {"phone": "555-0102"}, idx=1),
    ])
    assert "555-0102" in text
    assert "555-0199" not in text


def test_cell_as_the_first_word_of_a_program_name_is_not_a_label(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Cell Biology Program, Tel. 555-0100", {"phone": "555-0100"}),
    ])
    assert rows["Office telephone:"] == "555-0100"
    assert rows["Cell phone:"] == ""


def test_a_number_followed_by_a_cell_subject_is_not_a_cell_label(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Tel 555-0100 cell biology laboratory", {"phone": "555-0100"}),
    ])
    assert rows["Office telephone:"] == "555-0100"
    assert rows["Cell phone:"] == ""


@pytest.mark.parametrize("label", [
    "Office Phone", "Work Telephone", "Work Tel",
])
def test_every_work_phone_noun_outranks_a_banner_phone(tmp_path, label):
    rows = _contact_rows(tmp_path, [
        _a("Dr Example Doe\n555-0101", {"phone": "555-0101"}, idx=0),
        _a(f"{label}: 555-0102", {"phone": "555-0102"}, idx=1),
    ])
    assert rows["Office telephone:"] == "555-0102"


def test_a_residence_line_does_not_rank_an_address_for_the_office_slot(tmp_path):
    text = _render(tmp_path, [
        _a("1 Banner Street, Exampleton, ZZ 00000",
           {"address": "1 Banner Street, Exampleton, ZZ 00000"}, idx=0),
        _a("Residence: 9 Private Lane, Exampleton, ZZ 00000; Office phone: 555-0102",
           {"address": "9 Private Lane, Exampleton, ZZ 00000",
            "phone": "555-0102"}, idx=1),
    ])
    assert "1 Banner Street" in text
    assert "9 Private Lane" not in text


def test_a_school_name_before_tel_does_not_displace_a_street_address(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("1 Banner Street, Exampleton, ZZ 00000",
           {"address": "1 Banner Street, Exampleton, ZZ 00000"}, idx=0),
        _a("Example Business School, Tel. 555-0102",
           {"address": "Example Business School", "phone": "555-0102"}, idx=1),
    ])
    assert "1 Banner Street" in rows["Office address:"]


def test_a_fax_number_does_not_displace_an_office_phone(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Tel: 555-0101", {"phone": "555-0101"}, idx=0),
        _a("Office Fax Number: 555-0199", {"phone": "555-0199"}, idx=1),
    ])
    assert rows["Office telephone:"] == "555-0101"


def test_a_displaced_banner_address_is_recovered_not_lost(tmp_path):
    text = _render(tmp_path, [
        _a("Example Center for Sample Research",
           {"address": "Example Center for Sample Research"}, idx=0),
        _a("Business Address: 1 Sample Way, Exampleton, ZZ 00000",
           {"address": "1 Sample Way, Exampleton, ZZ 00000"}, idx=1),
    ])
    assert "1 Sample Way" in text
    assert "Example Center for Sample Research" in text


def test_a_value_from_a_structured_address_is_not_displaced(tmp_path):
    rows = _contact_rows(tmp_path, [
        _a("Contact", {"address": {"office_address": "5 Dict Street, Exampleton",
                                   "home_address": "6 Hidden Lane, Exampleton"}}, idx=0),
        _a("Business Address: 1 Sample Way, Exampleton, ZZ 00000",
           {"address": "1 Sample Way, Exampleton, ZZ 00000"}, idx=1),
    ])
    assert "5 Dict Street" in rows["Office address:"]


@pytest.mark.parametrize("home_line", ["(h) 555-0199", "H: 555-0199"])
def test_a_displaced_banner_entry_never_surfaces_a_home_lettered_number(
        tmp_path, home_line):
    """The pii pass cuts only home WORDS; a displaced entry recovered whole
    must not print a number labelled "(h)" or "H:" (#1222)."""
    text = _render(tmp_path, [
        _a(f"1 Banner Street, Exampleton, ZZ 00000\n{home_line}",
           {"address": "1 Banner Street, Exampleton, ZZ 00000",
            "phone": "555-0199"}, idx=0),
        _a("Business Address: 1 Sample Way, Exampleton, ZZ 00000",
           {"address": "1 Sample Way, Exampleton, ZZ 00000"}, idx=1),
    ])
    assert "1 Sample Way" in text
    assert "1 Banner Street" in text, "the displaced address was lost"
    assert "555-0199" not in text, "a home-lettered number was rendered"


def test_a_displaced_phone_entry_never_surfaces_a_home_lettered_number(tmp_path):
    text = _render(tmp_path, [
        _a("Tel: 555-0101\n(h) 555-0199",
           {"phone": "555-0101, 555-0199"}, idx=0),
        _a("Office Phone: 555-0102", {"phone": "555-0102"}, idx=1),
    ])
    assert "555-0102" in text
    assert "555-0199" not in text, "a home-lettered number was rendered"


def test_a_cell_label_written_straight_before_a_number_without_punctuation(tmp_path):
    """The label-end alternative `Cell 555-...` with no colon or dash; the two
    numbers do not pair with the extracted value, so only the block-level
    regex can route it."""
    rows = _contact_rows(tmp_path, [
        _a("Cell 555-0100, Office 555-0101",
           {"phone": "555-0100, 555-0101"}),
    ])
    assert rows["Cell phone:"] != ""
