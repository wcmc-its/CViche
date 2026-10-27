"""Regression guards for the #221 unrendered-record recovery pass.

Grounded in run KFGXBW: a fused D1 entry held five date-prefixed employment
records as text lines, extracted_fields held only the first, and the positions
table rendered exactly one row — the other four records appeared nowhere in
the output docx. The same class hits the F1/H/P/Q1/R and grant render paths.
All fixture strings here are invented; shapes mirror the affected corpus
entries (flat singular fields, tab rows with column headers, pipe-delimited
split table_rows with string element_idx values).

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_unrendered_recovery.py -p no:cacheprovider

Self-contained: no DB, no LLM calls. Loads the bundled WCM template like
test_stage6_funding_appendix.py does.
"""

import ast
import copy
import inspect
import json
import sys
import textwrap
import unicodedata
from pathlib import Path

import pytest

# Ensure the repo's ``src`` directory is importable regardless of cwd/rootdir.
_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from lxml import etree  # noqa: E402

from unified_pipeline.stage4.schemas import STAGE4_RECORDS_KEY  # noqa: E402
from unified_pipeline.stage6.fan_out import (  # noqa: E402
    FANNED_OUT_FROM,
    LAST_STAGE4_RECORD,
)
from unified_pipeline.stage6.formatting import plain_entry_left_twips  # noqa: E402
from unified_pipeline.stage6.normalization import (  # noqa: E402
    _clean_inline_tabs,
    _squash,
)
from unified_pipeline.stage6.parsing.text import (  # noqa: E402
    DATE_LED_SOURCE_TEXT,
    split_date_led_records,
)
from unified_pipeline.stage6.render_check import (  # noqa: E402
    _RENDER_TOKEN_RE,
    _entry_pieces,
    _norm,
    _record_rendered,
    _record_tokens,
    _whole_record_rendered,
    record_chunks,
    record_pieces,
    record_siblings,
    record_twins,
    segment_held_by_fields,
    segments_cover_source,
    short_tab_part_rendered,
    short_tab_part_words,
    sibling_record_rendered,
    tab_aware_record_lines,
)
from unified_pipeline.stage6.sections.appendix import recovery_candidates  # noqa: E402
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    WCMTemplateGenerator,
    _record_lines,
    run_stage6,
    segment_already_rendered,
)


def _generator(**kwargs):
    gen = WCMTemplateGenerator(verbose=False, **kwargs)
    gen.doc = Document(gen.template_path)
    return gen


def _num_level(para):
    """The paragraph's w:ilvl, or None when it carries no list markup.

    Mirrors test_stage6_k_list_bullets.py's helper of the same name.
    """
    pPr = para._p.pPr
    numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
    if numPr is None:
        return None
    return numPr.find(qn("w:ilvl")).get(qn("w:val"))


def _bulleted_texts(paragraphs):
    """Text of every real Word list paragraph among the given paragraphs.

    Both recovery emitters -- `_insert_reconsidered_segment` and
    `_add_remaining_to_appendix` -- write real list paragraphs as of #483 R2
    (previously `_insert_reconsidered_segment` prefixed a literal "•", which
    is what the old `startswith("•")` scans across this file matched on).
    """
    return [p.text for p in paragraphs if _num_level(p) is not None]


# Mirrors the KFGXBW entry-21 shape: date-range comma lines, flat singular
# extracted_fields holding only the FIRST record.
def _d1_entry():
    return {
        "element_idx_start": 21,
        "element_type": "break",
        "hierarchy": ["ACADEMIC APPOINTMENTS"],
        "taxonomy_code": "D1",
        "taxonomy_confidence": 0.85,
        "text": "\n".join([
            "Jan 2024-Present, Vice Chair of Interdimensional Affairs",
            "Aug 2019-Dec 2023, Associate Director of Quantitative "
            "Basketweaving, Institute of Speculative Metrics, Norvale "
            "University, Crab Hollow, ZQ",
            "Mar 2015-Jul 2019, Program Coordinator for Wombat Logistics, "
            "Office of Curious Ventures, Norvale University, Crab Hollow, ZQ",
        ]),
        "extraction_coverage": {"extraction_coverage_percent": 11.8},
        "extracted_fields": {
            "title": "Vice Chair of Interdimensional Affairs",
            "institution": None,
            "start_date": "Jan 2024",
            "end_date": "Present",
        },
    }


# ------------------------------------------------------- record-line heuristic

def test_record_lines_heuristics():
    text = "\n".join([
        "Aug 2019-Dec 2023, Associate Director of Quantitative Basketweaving",
        "Aug, 2026, Judge, Regional Pie Symposium",  # comma month: no heuristic
        "short",
        "Chair | Committee on Extremely Long Titles | Guild of Meandering "
        "Auditors | 2011-2014",
    ])
    assert _record_lines(text) == [
        "Aug 2019-Dec 2023, Associate Director of Quantitative Basketweaving",
        "Chair | Committee on Extremely Long Titles | Guild of Meandering "
        "Auditors | 2011-2014",
    ]


def test_tab_aware_record_lines_split_a_tab_fused_line_into_its_dated_records():
    # #1243: one line, three posts joined by tabs. `_record_lines` counts the
    # line once, so the entry never reached the recovery pass.
    fused = ("Executive Committee of the Quokka Society 2013-present\t"
             "President 5/2017-10/2020\tPresident Elect 5/2015-5/2017")
    assert _record_lines(fused) == [fused]
    assert tab_aware_record_lines(fused) == [
        "Executive Committee of the Quokka Society 2013-present",
        "President 5/2017-10/2020",
        "President Elect 5/2015-5/2017",
    ]
    # A tab row with one dated part is one record line, as before.
    row = "Chair\tCommittee on Extremely Long Titles\tGuild of Meandering Auditors\t2011-2014"
    assert tab_aware_record_lines(row) == [row] == _record_lines(row)


def test_tab_fused_entry_recovers_the_record_the_render_dropped():
    gen = _generator()
    entry = {
        "taxonomy_code": "Q1",
        "element_idx_start": 281,
        "text": ("Steering Committee of the Quokka Guild 2013-present\t"
                 "Chairperson of Marsupial Outreach 5/2017-10/2020\t"
                 "Secretary of Wallaby Affairs 5/2013-5/2015"),
        "extracted_fields": {"organization": "Quokka Guild", "role": "Steering Committee"},
    }
    gen.doc.add_paragraph("Steering Committee of the Quokka Guild 2013-present")

    gen._recover_unrendered_records({"Q1": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Chairperson of Marsupial Outreach 5/2017-10/2020" in t for t in bullets)
    assert any("Secretary of Wallaby Affairs 5/2013-5/2015" in t for t in bullets)
    # The part that did render is never written a second time.
    assert not any("Steering Committee of the Quokka Guild" in t for t in bullets)


def test_tab_part_the_extracted_fields_hold_is_not_recovered():
    # #1243: a mentee's name and level, then his fellowship, each dated. The
    # row renders them cell by cell, so no output line carries the first part
    # whole, but stage 4 holds every word of it.
    gen = _generator()
    text = ("2007-2011 Quimby Farrow, MSPH, PhD (Epidemiology)\t"
            "Marsupial Heart Association Fellowship (7/2009-6/2011)")
    entry = {
        "taxonomy_code": "N3B", "text": text,
        "extracted_fields": {"mentee_name": "Quimby Farrow", "mentee_level": "MSPH, PhD (Epidemiology)",
                             "funding_source": "Marsupial Heart Association Fellowship (7/2009-6/2011)"},
    }
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records({"N3B": [entry]})

    assert len(gen.doc.paragraphs) == before
    assert segment_held_by_fields("2007-2011 Quimby Farrow, MSPH, PhD (Epidemiology)", entry["extracted_fields"])
    assert not segment_held_by_fields("Chairperson of Marsupial Outreach 5/2017-10/2020",
                                      entry["extracted_fields"])


def test_tab_parts_of_an_entry_stage_4_split_into_records_are_not_recovered():
    # The records render as children of their own, reformatted; the parent's
    # last child keeps the whole fused text.
    gen = _generator()
    entry = {
        "taxonomy_code": "Q1",
        "text": ("Regional Quokka Advisory Council (Appointed by Governor) 1995 - 2001\t"
                 "Committee on Quokka Facilities, Chairman, 2002 - 2004"),
        "extracted_fields": {"role": "Chairman", "organization": "Regional Quokka Advisory Council"},
        "stage4_last_record": {"key": "stage4_records", "index": 1, "count": 2, "text": None},
    }
    entry.update({LAST_STAGE4_RECORD: entry.pop("stage4_last_record")})
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records({"Q1": [entry]})

    assert len(gen.doc.paragraphs) == before


def test_tab_part_held_only_by_a_list_field_is_still_recovered():
    # #1243: `additional_roles` items are not always printed (fan-out declined,
    # or a field the section does not name), so a part whose words only a list
    # holds is not the extracted record.
    gen = _generator()
    entry = {
        "taxonomy_code": "I",
        "text": ("Member of the Quokka Society 2009-present\t"
                 "Chairman, Credentials Committee, 2009-2011\t"
                 "Secretary of the Wombat Panel 2012-2014"),
        "extracted_fields": {"organization": "Quokka Society", "role": "Member",
                             "additional_roles": ["Chairman, Credentials Committee",
                                                  "Secretary of the Wombat Panel"]},
    }
    gen.doc.add_paragraph("Member of the Quokka Society 2009-present")

    gen._recover_unrendered_records({"I": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Chairman, Credentials Committee, 2009-2011" in t for t in bullets)
    assert any("Secretary of the Wombat Panel 2012-2014" in t for t in bullets)
    assert not segment_held_by_fields("Chairman, Credentials Committee, 2009-2011",
                                      entry["extracted_fields"])


def test_recovered_tab_part_does_not_vouch_for_its_sibling():
    # #1243: siblings of one society share their vocabulary and differ in the
    # dates. The first one written must not make the next read as rendered.
    gen = _generator()
    entry = {
        "taxonomy_code": "Q1",
        "text": ("Member of the Quokka Society 1989-present\t"
                 "Chair, Program Committee of the Quokka Society 2015-2017\t"
                 "Vice Chair, Program Committee of the Quokka Society 2012-2015"),
        "extracted_fields": {"organization": "Quokka Society", "role": "Member"},
    }
    gen.doc.add_paragraph("Member of the Quokka Society 1989-present")

    gen._recover_unrendered_records({"Q1": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Chair, Program Committee of the Quokka Society 2015-2017" in t for t in bullets)
    assert any("Vice Chair, Program Committee of the Quokka Society 2012-2015" in t for t in bullets)


def _quokka_entry(*dated_parts):
    return {
        "taxonomy_code": "Q1",
        "text": "\t".join(("Member of the Quokka Society 1989-present",) + dated_parts),
        "extracted_fields": {"organization": "Quokka Society", "role": "Member"},
    }


def test_a_tab_part_recovered_from_one_entry_vouches_for_its_duplicate_in_another():
    # #1243: dedup keeps near-duplicate entries apart until later, so the same
    # fused entry can reach the pass twice; each part is written once.
    gen = _generator()
    parts = ("Chair, Program Committee of the Quokka Society 2015-2017",
             "Treasurer of the Wombat Panel 2012-2014")
    gen.doc.add_paragraph("Member of the Quokka Society 1989-present")

    gen._recover_unrendered_records({"Q1": [_quokka_entry(*parts), _quokka_entry(*parts)]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    for part in parts:
        assert sum(part in t for t in bullets) == 1
    assert gen.stats["unrendered_records_recovered"] == 2


@pytest.mark.parametrize("tab_first", [True, False])
def test_a_record_recovered_as_a_tab_part_is_not_rewritten_from_a_newline_entry(tab_first):
    # #1243: the same records, fused by tabs in one entry and one per line in
    # another; each is written once whichever entry the pass reads first.
    gen = _generator()
    parts = ("2015-2017 Chair, Program Committee of the Quokka Society",
             "2012-2014 Treasurer of the Wombat Panel")
    gen.doc.add_paragraph("Member of the Quokka Society 1989-present")
    newline_entry = {"taxonomy_code": "Q1", "text": "\n".join(parts), "extracted_fields": {}}
    entries = [_quokka_entry(*parts), newline_entry]

    gen._recover_unrendered_records({"Q1": entries if tab_first else entries[::-1]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    for part in parts:
        assert sum(part in t for t in bullets) == 1
    assert gen.stats["unrendered_records_recovered"] == 2


def test_a_recovered_tab_part_does_not_vouch_for_a_later_entry_that_differs_in_its_years():
    # #1243: two entries of one society share their vocabulary; the second
    # one's part carries other years, so it is its own record.
    gen = _generator()
    first = "Chair, Program Committee of the Quokka Society 2015-2017"
    second = "Chair, Program Committee of the Quokka Society 2009-2011"
    gen.doc.add_paragraph("Member of the Quokka Society 1989-present")

    gen._recover_unrendered_records({"Q1": [
        _quokka_entry(first, "Treasurer of the Wombat Panel 2012-2014"),
        _quokka_entry(second, "Auditor of the Numbat Panel 2005-2007"),
    ]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any(first in t for t in bullets)
    assert any(second in t for t in bullets)


def test_record_line_held_by_fields_is_still_recovered_when_it_is_not_a_tab_part():
    # The held-by-fields skip is for tab parts only: an ordinary record line
    # whose words the fields hold, and that the render dropped, is lost.
    gen = _generator()
    entry = {
        "taxonomy_code": "Q1",
        "text": ("2009-2011 Member of the Quokka Society in good standing\n"
                 "2012-2014 Chair of the Wombat Panel for three years"),
        "extracted_fields": {"notes": "2009-2011 Member of the Quokka Society in good standing "
                                      "2012-2014 Chair of the Wombat Panel for three years"},
    }
    gen.doc.add_paragraph("Unrelated line about marsupial husbandry in the highlands")

    gen._recover_unrendered_records({"Q1": [entry]})

    texts = [p.text for p in gen.doc.paragraphs]
    assert "2012-2014 Chair of the Wombat Panel for three years" in texts
    assert gen.stats['unrendered_records_recovered'] == 2


def test_short_officer_tab_parts_the_render_dropped_are_recovered():
    # #1243: an officer post under its society's line has one distinctive word
    # and its years, too few for the overlap test, so the pass skipped it. Its
    # word and its years identify it: a line naming the same office with other
    # years does not vouch for it.
    gen = _generator()
    entry = {
        "taxonomy_code": "Q1",
        "text": ("American Board of Quokka Surgery, term: 2007-2015\t"
                 "• Chairman, Credentials Committee, 2009-2011\t"
                 "• Secretary, 2009-2011\t• Vice President, 2011-2012"),
        "extracted_fields": {"role": "Member", "organization": "American Board of Quokka Surgery",
                             "additional_roles": [{"role": "Secretary"}, {"role": "Vice President"}]},
    }
    gen.doc.add_paragraph("Member, American Board of Quokka Surgery 2007-2015")
    gen.doc.add_paragraph("Chairman, Credentials Committee 2009-2011")
    gen.doc.add_paragraph("Secretary of the Wombat Panel, 2005-2007")

    gen._recover_unrendered_records({"Q1": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Secretary, 2009-2011" in t for t in bullets)
    assert any("Vice President, 2011-2012" in t for t in bullets)
    assert not any("Chairman" in t for t in bullets)
    assert gen.stats["unrendered_records_recovered"] == 2


def test_short_tab_part_rendered_elsewhere_or_held_by_fields_is_not_recovered():
    # #1243: the short part that did render, and a mentee line whose name a
    # scalar field holds (its "Dates:" label and years render in rows of their
    # own), are never written a second time.
    gen = _generator()
    officers = {
        "taxonomy_code": "Q1",
        "text": ("Member of the Quokka Society 1989-present\t"
                 "Treasurer, 2004-2005\tSecond Vice President 2006-2007"),
        "extracted_fields": {"organization": "Quokka Society", "role": "Member"},
    }
    mentee = {
        "taxonomy_code": "N3B",
        "text": ("Dr. Lisa Quokkason, Ph.D.    Dates: 1989-1990\t"
                 "Assistant Professor of Wombat Surgery (1991-1996)\tNorvale University"),
        "extracted_fields": {"mentee_name": "Lisa Quokkason", "mentee_level": "Ph.D.",
                             "current_position": "Assistant Professor of Wombat Surgery (1991-1996)"},
    }
    gen.doc.add_paragraph("Member of the Quokka Society 1989-present")
    gen.doc.add_paragraph("Treasurer 2004-2005")
    gen.doc.add_paragraph("Second Vice President 2006-2007")
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records({"Q1": [officers], "N3B": [mentee]})

    assert len(gen.doc.paragraphs) == before
    assert gen.stats["unrendered_records_recovered"] == 0


def test_short_record_line_outside_tabs_stays_unverifiable():
    # The words-and-years check reads tab parts only: a short line of a
    # newline entry is still too short to verify, and is never written.
    gen = _generator()
    entry = {"taxonomy_code": "Q1", "extracted_fields": {},
             "text": "2009-2011 Secretary of it\n2012-2014 Treasurer of it"}
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records({"Q1": [entry]})

    assert len(gen.doc.paragraphs) == before


@pytest.mark.parametrize("part, words", [
    ("• Secretary, 2009-2011", {"secretary"}),
    ("Secretary, July, 1993 - October 1994", {"secretary"}),
    ("President Elect 5/2015-5/2017", {"president", "elect"}),
    ("1996-97 Vice-President, Acting President", {"president", "acting"}),
    ("2008-present Africa Liaison", {"africa", "liaison"}),
    ("Dr. Lisa Quokkason, Ph.D.    Dates: 1989-1990", {"quokkason"}),
    # Nothing to check it by: a heading with a year, a head count, a bare date.
    ("COMMITTEES 1993-1994", None),
    ("students 2013-2014: 12", None),
    ("July 1992 - October 1994", None),
])
def test_short_tab_part_words(part, words):
    assert short_tab_part_words(part) == words


def test_short_tab_part_rendered_needs_its_years_on_the_line():
    sets = [{"secretary", "wombat", "panel", "2005", "2007"}]
    assert short_tab_part_rendered("Secretary, 2009-2011", {}, sets) is False
    assert short_tab_part_rendered("Secretary, 2005-2007", {}, sets) is True
    assert short_tab_part_rendered("Secretary, 2009-2011", {"role": "Secretary"}, []) is True
    assert short_tab_part_rendered("COMMITTEES 1993-1994", {}, []) is None
    # #1243: a part dated by one year cannot be told from the row whose range
    # cell holds that year ("Fellow, Quokka Society" beside "2015-Present").
    assert short_tab_part_rendered("Elected fellow 2015.", {}, []) is None
    assert short_tab_part_rendered("Elected fellow 2015-2016.", {}, []) is False


def test_tab_parts_dated_by_one_year_the_render_dropped_are_recovered():
    # #1243 (BFSUMA 98's shape): a society line whose posts mostly carry one
    # year. The posts no line names are recovered; the one the render wrote
    # into another line is not, and neither is a short one-year part.
    gen = _generator()
    entry = {
        "taxonomy_code": "I",
        "text": ("Quokka College of Pharmacy, 1979-present\tFounding Member, 1979\t"
                 "Elected fellow 2015.\tResearch Affairs Committee, 1985\t"
                 "Chairperson, Awards Committee, 1987-1988"),
        "extracted_fields": {"organization": "Quokka College of Pharmacy",
                             "membership_type": "Founding Member, 1979"},
    }
    gen.doc.add_paragraph("Founding Member, 1979; Fellow, 1985, Quokka College of Pharmacy 1979-present")
    gen.doc.add_paragraph("Fellow, Quokka College of Pharmacy")

    gen._recover_unrendered_records({"I": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Research Affairs Committee, 1985" in t for t in bullets)
    assert any("Chairperson, Awards Committee, 1987-1988" in t for t in bullets)
    assert not any("Founding Member" in t or "Elected fellow" in t for t in bullets)
    assert gen.stats["unrendered_records_recovered"] == 2


# ------------------------------------------------ D1 positions path (KFGXBW)

def test_d1_remainder_recovered_in_positions_section():
    gen = _generator()
    entry = _d1_entry()

    # The real render path: one table row from extracted_fields, remainder
    # lines dropped (#221).
    gen._fill_positions({"D1": [entry]})
    gen._recover_unrendered_records({"D1": [entry]})

    paras = gen.doc.paragraphs
    texts = [p.text for p in paras]
    bullets = _bulleted_texts(paras)
    assert any(t.startswith("Aug 2019-Dec 2023, Associate Director of "
                            "Quantitative Basketweaving") for t in bullets)
    assert any(t.startswith("Mar 2015-Jul 2019, Program Coordinator for "
                            "Wombat Logistics") for t in bullets)
    # The extracted record rendered as a table row — never duplicated as a
    # bullet.
    assert not any("Interdimensional Affairs" in t for t in bullets)
    # Recovered into the entry's own section, not the appendix.
    header_idx = next(i for i, t in enumerate(texts)
                      if "Academic Appointments" in t)
    bullet_idx = next(i for i, (p, t) in enumerate(zip(paras, texts))
                      if _num_level(p) is not None and t.startswith("Aug 2019"))
    assert bullet_idx > header_idx
    assert not any("T. APPENDIX" in t for t in texts)


def test_reformatted_rendered_line_not_recovered():
    gen = _generator()
    # Simulate a stage-6 reformat of the first record: different date format
    # and field order, but its distinctive tokens stay on one line.
    gen.doc.add_paragraph(
        "Curator of Improbable Archives, Norvale University (2019 - 2023)")
    entry = {
        "element_idx_start": 7,
        "taxonomy_code": "D1",
        "text": "\n".join([
            "Aug 2019-Dec 2023, Curator of Improbable Archives, Norvale "
            "University",
            "Feb 2011-Jul 2015, Keeper of Ceremonial Spreadsheets, Bureau of "
            "Ornamental Data, Quexley College",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"D1": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert not any("Improbable Archives" in t for t in bullets)
    assert any(t.startswith("Feb 2011-Jul 2015, Keeper of Ceremonial "
                            "Spreadsheets") for t in bullets)


def test_short_title_extracted_record_not_double_rendered():
    """#221 review (HIGH): a short extracted title ('Professor') whose source
    line carries department tokens the positions table omits pushes the
    per-line token overlap below 0.7, so the already-rendered guard — not the
    token check — must keep the record from being re-emitted as a bullet next
    to its own table row."""
    gen = _generator()
    entry = {
        "element_idx_start": 40,
        "taxonomy_code": "D1",
        "text": "\n".join([
            "Jul 2014-Present, Professor, Department of Speculative "
            "Rehabilitation Metrics, Norvale University Medical College",
            "Aug 2009-Jun 2014, Associate Professor, Bureau of Ornamental "
            "Data, Quexley College",
        ]),
        "extracted_fields": {
            "title": "Professor",
            "institution": "Norvale University Medical College",
            "start_date": "Jul 2014",
            "end_date": "Present",
        },
    }

    gen._fill_positions({"D1": [entry]})
    gen._recover_unrendered_records({"D1": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    # The extracted record rendered as a table row — never also as a bullet.
    assert not any("Jul 2014-Present" in t for t in bullets)
    # The dropped sibling is still recovered.
    assert any(t.startswith("Aug 2009-Jun 2014, Associate Professor")
               for t in bullets)


def test_short_title_guard_requires_both_date_anchors():
    """The short-title conjunction must not mark a fused sibling as already
    rendered when it only shares a boundary date and a title suffix
    ('Associate Professor' contains 'Professor')."""
    fields = {
        "title": "Professor",
        "institution": "Norvale University Medical College",
        "start_date": "Jul 2014",
        "end_date": "Present",
    }
    line = ("Jul 2014-Present, Professor, Department of Speculative "
            "Rehabilitation Metrics, Norvale University Medical College")
    sibling = ("Aug 2009-Jul 2014, Associate Professor, Norvale University "
               "Medical College")
    assert segment_already_rendered(line, fields)
    assert not segment_already_rendered(sibling, fields)


def test_unverifiable_short_line_never_recovered():
    gen = _generator()
    entry = {
        "taxonomy_code": "D1",
        "text": "\n".join([
            # <3 distinctive tokens: absent verbatim, but unverifiable (None)
            "Jun 2011-Jun 2014, Dean, ZQ Camp",
            "Feb 2001-Jul 2005, Cartographer of Forgotten Stairwells, "
            "Quexley College",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"D1": [entry]})

    texts = [p.text for p in gen.doc.paragraphs]
    assert not any("ZQ Camp" in t for t in texts)
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any(t.startswith("Feb 2001-Jul 2005, Cartographer of Forgotten "
                            "Stairwells") for t in bullets)


# --------------------------------------------------- F1 licensure path (2071)

def _f1_entry():
    return {
        "element_idx_start": 20,
        "element_idx_end": 21,
        "element_type": "table",
        "hierarchy": ["Certification and Licensure"],
        "taxonomy_code": "F1",
        "text": "\n".join([
            "Licensure",
            # Column-header row: tabbed and >60 chars (a record line under the
            # tab heuristic) but with no digit payload — must never become a
            # recovered bullet.
            "State/Country\t\tLicense Number\t\tStatus\t\t\tDate of Issue\t\t"
            "Date of Expiration",
            "Commonwealth of Greater Zephyria\t\t4471\t\t\tActive\t03/2011\t\t"
            "06/30/2027",
            "Board of Speculative Osteopathy\t83412\t\t\tCert 09/2015\t\t"
            "09/22/2028",
        ]),
        "extracted_fields": {
            "state_country": "Commonwealth of Greater Zephyria",
            "license_number": "4471",
            "issue_date": "2011-03",
            "expiration_date": "2027-06-30",
        },
    }


_CJK_RECORD = ("2015年4月 東京大学医学部附属病院 循環器内科 准教授として心不全の臨床研究に従事 "
               "한국어로된논문제목입니다 中华人民共和国国家自然科学基金资助项目")


def test_cjk_never_forms_a_render_token_so_cjk_lines_are_unverifiable():
    """#722: CJK has no word boundaries, so the 5-letter floor is not
    meaningful for it; CJK is excluded, not measured. A CJK-only record line
    is None (not verifiable), never False, and Latin/Cyrillic words in a
    mixed-script line still count. Invented text."""
    assert _RENDER_TOKEN_RE.findall(_norm(_CJK_RECORD)) == []
    mixed = _norm("東京大学医学部 Cardiology Иванов 心不全の臨床研究 Blorvane")
    assert set(_RENDER_TOKEN_RE.findall(mixed)) == {
        "cardiology", "иванов", "blorvane"}
    other = [set(_RENDER_TOKEN_RE.findall(_norm("Completely unrelated "
                                                "Cardiology Blorvane line")))]
    assert _record_rendered(_CJK_RECORD, "", other) is None
    latin = "Jun 2011, Quexley Cartographer Stairwells Forgotten College"
    assert _record_rendered(latin, "", other) is False


@pytest.mark.parametrize("run", [
    "ひらがなのことばです",          # Hiragana
    "カタカナノコトバデス",          # Katakana
    "ｶﾀｶﾅｶﾅｶﾅｶﾅ",                    # Halfwidth Katakana (raw, pre-NFKD)
    "한국어한ᄀ",  # Hangul Jamo (NFKD form)
    "".join(chr(0x3400 + i * 16) for i in range(12)),        # CJK extension A
    "".join(chr(0x20000 + i * 16) for i in range(12)),       # CJK extension B
    "ㄱㄲㄴㄷㄹㅁ",                  # Hangul compatibility Jamo
    "ㇰㇱㇲㇳㇴㇵ",                  # Katakana phonetic extensions
    "ꥠꥡꥢꥣꥤꥥ",                  # Hangul Jamo extended-A
    "".join(chr(0xF900 + i) for i in range(6)),              # CJK compatibility ideographs (NFC-unstable literal)
    "ힰힱힲힳힴힵ",                  # Hangul Jamo extended-B
])
def test_every_cjk_block_is_excluded_from_the_token_regex(run):
    """#722: each range in `_CJK_CLASS` must keep a 5+ run out of the token
    set on its own (the sentence-level test above hits only Han/Hangul)."""
    assert _RENDER_TOKEN_RE.findall(run) == []


_CJK_NAME_PREFIXES = ("CJK ", "HIRAGANA", "KATAKANA", "HENTAIGANA", "HANGUL",
                      "HALFWIDTH KATAKANA", "HALFWIDTH HANGUL")


def test_cjk_exclusion_matches_unicode_names_across_every_code_point():
    """#722: pins every `_CJK_CLASS` endpoint at once. A letter whose Unicode
    name is CJK/kana/Hangul never forms a token, and every other letter does
    (Bopomofo is the named, deliberate gap: its names start 'BOPOMOFO')."""
    escaped, over_excluded = [], []
    for cp in range(sys.maxunicode + 1):
        ch = chr(cp)
        if not unicodedata.category(ch).startswith("L"):
            continue
        is_cjk = unicodedata.name(ch, "").startswith(_CJK_NAME_PREFIXES)
        forms_token = _RENDER_TOKEN_RE.fullmatch(ch * 5) is not None
        if is_cjk and forms_token:
            escaped.append(hex(cp))
        elif not is_cjk and not forms_token:
            over_excluded.append(hex(cp))
    assert escaped == [], escaped[:20]
    assert over_excluded == [], over_excluded[:20]


def test_f1_dropped_license_recovered_column_header_not():
    gen = _generator()
    entry = _f1_entry()

    gen._fill_licensure([entry])
    gen._recover_unrendered_records({"F1": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("83412" in b for b in bullets)          # dropped second license
    assert not any("License Number" in b for b in bullets)  # column header
    assert not any("4471" in b for b in bullets)       # rendered license


def test_two_cell_no_digit_record_recovered():
    """The 2071 D2 residual: a real dateless two-cell record (title \\t
    institution) must not be mistaken for a column-header row — only
    full-width no-digit rows (>=2 separators) are header furniture."""
    gen = _generator()
    entry = {
        "element_idx_start": 31,
        "taxonomy_code": "D2",
        "text": "\n".join([
            "Jul 2014-Present, Attending Wrangler of Nocturnal Metrics, "
            "Quexley College Medical Pavilion",
            "Visiting Lecturer of Ornamental Calculus\tNorvale Institute of "
            "Applied Whimsy",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"D2": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Ornamental Calculus" in t for t in bullets)


def test_dateless_committee_row_recovered():
    """#221 review: a real committee row with no digits and >=2 separators
    ('Member | Committee ... | Organization | description') is a record, not
    column-header furniture — only majority-label rows are skipped (see the
    F1 column-header test above)."""
    gen = _generator()
    entry = {
        "element_idx_start": 77,
        "taxonomy_code": "P",
        "text": "\n".join([
            "Member | Committee on Subterranean Balloon Safety Standards | "
            "Guild of Meandering Auditors | reviews annual certification "
            "checklists",
            "Chair | Panel of Improbable Weights and Measures 2013-2016 | "
            "Norvale Metrology Circle",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"P": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Subterranean Balloon" in t for t in bullets)


def test_near_duplicate_variant_across_entries_recovered_once():
    """#221 review: dedup drops entries precisely because they NEAR-duplicate
    a kept one, so the same unrendered record can reach recovery twice with
    punctuation variance. The first insert must count as rendered for the
    second variant."""
    gen = _generator()
    base = ("Aug 2019-Dec 2023, Associate Director of Quantitative "
            "Basketweaving, Institute of Speculative Metrics, Norvale "
            "University")
    entry_a = {
        "element_idx_start": 50,
        "taxonomy_code": "D1",
        "text": "\n".join([
            base,
            "Mar 2015-Jul 2019, Program Coordinator for Wombat Logistics, "
            "Office of Curious Ventures, Norvale University",
        ]),
        "extracted_fields": {},
    }
    entry_b = {
        "element_idx_start": 54,
        "taxonomy_code": "D1",
        "text": "\n".join([
            base.replace("Aug 2019-Dec 2023", "Aug 2019 - Dec 2023") + ".",
            "Feb 2011-Jul 2015, Keeper of Ceremonial Spreadsheets, Bureau of "
            "Ornamental Data, Quexley College",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"D1": [entry_a, entry_b]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    hits = [t for t in bullets if "Quantitative Basketweaving" in t]
    assert len(hits) == 1


# ------------------------------------------------ S bibliography anchor (#221)

def test_s_code_recovery_lands_in_bibliography_section():
    """#221 review: BIBLIOGRAPHY has no bold subsection header, and a plain
    substring fallback anchored S-code recoveries on the MENTORING section's
    '**Optional: List publications...' instruction paragraph. The recovered
    bullet must land after the real BIBLIOGRAPHY header."""
    gen = _generator()
    entry = {
        "element_idx_start": 90,
        "taxonomy_code": "S1",
        "text": "\n".join([
            "Fennwick, R., Marblegate, Q.\tOrnamental calculus for the "
            "reluctant curator\tJournal of Improbable Results\t2017",
            "Marblegate, Q., Fennwick, R.\tSubaquatic origami and its "
            "discontents\tAnnals of Speculative Metrics\t2019",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"S1": [entry]})

    paras = gen.doc.paragraphs
    texts = [p.text for p in paras]
    biblio_idx = next(i for i, t in enumerate(texts)
                      if t.strip() == "BIBLIOGRAPHY")
    optional_idx = next(i for i, t in enumerate(texts)
                        if t.strip().startswith("**Optional"))
    assert optional_idx < biblio_idx  # template invariant the bug relied on
    bullet_idxs = [i for i, (p, t) in enumerate(zip(paras, texts))
                   if _num_level(p) is not None and "Ornamental calculus" in t]
    assert bullet_idxs, "citation line was not recovered"
    assert all(i > biblio_idx for i in bullet_idxs)


# -------------------------------------------------- M2B grant table_row (2054)

def _m2b_entry():
    return {
        # Split table_rows carry STRING element_idx values plus parent keys.
        "element_idx_start": "66.16",
        "element_idx_end": "66.16",
        "parent_idx": 66,
        "table_index": 14,
        "row_index": 16,
        "element_type": "table_row",
        "hierarchy": ["PROFESSIONAL ACTIVITIES", "2. Research and Training"],
        "taxonomy_code": "M2B",
        "text": "\n".join([
            "01/01/2011-05/31/2014 | 01/01/2011-05/31/2014 | Cartography of "
            "Subterranean Cloud Formations, 5R01ZZ049917-04",
            '"Cartography of Subterranean Cloud Formations"',
            "07/01/2016-06/30/2019 | 07/01/2016-06/30/2019 | Panuvian Wetland "
            "Acoustics Survey, 3R21QX041776-02",
        ]),
        "extracted_fields": {
            "grant_number": "5R01ZZ049917-04",
            "title": "Cartography of Subterranean Cloud Formations",
            "agency": "NIOF",
            "total_funding": "$46,176",
        },
    }


def test_m2b_sibling_routed_home_extracted_grant_guarded():
    gen = _generator()
    entry = _m2b_entry()

    # Nothing rendered at all (the 2054 case: even the extracted grant's
    # tokens end up split across grant-table label/value rows). The
    # identifying-fields guard — not the token check — must keep the
    # extracted grant from being re-inserted.
    gen._recover_unrendered_records({"M2B": [entry]})

    texts = [p.text for p in gen.doc.paragraphs]
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert any("Panuvian Wetland Acoustics Survey" in b for b in bullets)
    assert not any("Cartography of Subterranean" in b for b in bullets)
    # Routed to the funding subsection (the #214 precedent), not the appendix.
    header_idx = next(i for i, t in enumerate(texts)
                      if "Past (Completed) Funding" in t)
    bullet_idx = next(i for i, t in enumerate(texts) if "Panuvian" in t)
    assert bullet_idx > header_idx
    assert not any("T. APPENDIX" in t for t in texts)


def test_m2b_datelike_field_values_do_not_vouch_sibling_recovered():
    """#221 post-review (corpus CV 2054, entry 66.16): sibling records of a
    fused grant entry share their leading date ranges, so a date-like
    identifying-field value ('07/01/2016-06/30/2019', 21 chars — over the
    guard's length bar) textually present in a sibling line must not mark
    that sibling as already rendered. Same for the extracted grant's bare
    grant number, which carries no alphabetic word at all. Before the fix
    the guard vouched on these and the recovery pass skipped a genuinely
    absent grant record line."""
    gen = _generator()
    entry = {
        "element_idx_start": "66.16",
        "element_idx_end": "66.16",
        "parent_idx": 66,
        "element_type": "table_row",
        "hierarchy": ["PROFESSIONAL ACTIVITIES", "2. Research and Training"],
        "taxonomy_code": "M2B",
        "text": "\n".join([
            "07/01/2016-06/30/2019 | 07/01/2016-06/30/2019 | Cartography of "
            "Subterranean Cloud Formations, 5R01ZZ049917-04",
            "07/01/2016-06/30/2019 | 07/01/2016-06/30/2019 | Meandering "
            "Auditors Guild Junior Fellowship, 3R21QX041776-02",
        ]),
        "extracted_fields": {
            "grant_number": "5R01ZZ049917-04",
            "title": "Cartography of Subterranean Cloud Formations",
            "agency": "NIOF",
            # Extraction noise mirroring the corpus shape: a shared date
            # range landed in an identifying field.
            "award_source": "07/01/2016-06/30/2019",
        },
    }

    # Neither record rendered anywhere (the 2054 case). Only prose values may
    # block recovery: the extracted grant's own line is vouched by its title;
    # the sibling shares nothing with the fields but dates + punctuation.
    assert not segment_already_rendered(
        "07/01/2016-06/30/2019 | 07/01/2016-06/30/2019 | Meandering "
        "Auditors Guild Junior Fellowship, 3R21QX041776-02",
        entry["extracted_fields"])

    gen._recover_unrendered_records({"M2B": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    hits = [b for b in bullets if "Meandering Auditors Guild Junior" in b]
    assert len(hits) == 1  # sibling recovered, exactly once
    # The extracted grant's line still blocked by its (prose) title.
    assert not any("Cartography of Subterranean" in b for b in bullets)


def test_prose_identifying_values_still_vouch_and_block_recovery():
    """Regression for the date-like tightening: an extracted record whose
    title and organization DO appear in its record line is still treated as
    rendered — the guard must keep vouching on prose values."""
    gen = _generator()
    entry = {
        "element_idx_start": 12,
        "taxonomy_code": "P",
        "text": "\n".join([
            "Chair | Panel of Improbable Weights and Measures | Norvale "
            "Metrology Circle | 2013-2016",
            "Member | Committee on Subterranean Balloon Safety Standards | "
            "Guild of Meandering Auditors | 2017-2020",
        ]),
        "extracted_fields": {
            "title": "Panel of Improbable Weights and Measures",
            "organization": "Norvale Metrology Circle",
        },
    }

    gen._recover_unrendered_records({"P": [entry]})

    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert not any("Improbable Weights" in b for b in bullets)
    assert any("Subterranean Balloon" in b for b in bullets)


def test_recovery_not_swallowed_by_appendix_group_head():
    """The KFGXBW re-render regression: _fill_appendix leaves bold
    'From "ACADEMIC APPOINTMENTS":' group heads at document end, and the
    subsection search must not anchor recovered D1 records there instead of
    the real positions section."""
    gen = _generator()
    entry = _d1_entry()
    gen._fill_positions({"D1": [entry]})

    # Simulate an earlier _fill_appendix run whose group head echoes the
    # source section name.
    appendix_head = gen.doc.add_paragraph()
    run = appendix_head.add_run("T. APPENDIX")
    run.bold = True
    group_head = gen.doc.add_paragraph()
    run = group_head.add_run('From "ACADEMIC APPOINTMENTS":')
    run.bold = True

    gen._recover_unrendered_records({"D1": [entry]})

    paras = gen.doc.paragraphs
    texts = [p.text for p in paras]
    header_idx = next(i for i, t in enumerate(texts)
                      if "Academic Appointments" in t)
    appendix_idx = next(i for i, t in enumerate(texts) if "T. APPENDIX" in t)
    bullet_idx = next(i for i, (p, t) in enumerate(zip(paras, texts))
                      if _num_level(p) is not None and t.startswith("Aug 2019"))
    assert header_idx < bullet_idx < appendix_idx


# ------------------------------------------------- appendix fallback + option

def test_unmapped_code_falls_back_to_appendix_untruncated():
    gen = _generator()
    entry = {
        "taxonomy_code": "F2",  # no subsection header mapping
        "text": "\n".join([
            "Jan 2012-Dec 2016, Fellowship of Ornithological Cryptography, "
            "Guild of Meandering Auditors, Crab Hollow, ZQ",
            "Feb 2017-Dec 2021, Fellowship of Subaquatic Origami, Guild of "
            "Meandering Auditors, Crab Hollow, ZQ",
        ]),
        "extracted_fields": {},
    }

    gen._recover_unrendered_records({"F2": [entry]})

    texts = [p.text for p in gen.doc.paragraphs]
    assert any("T. APPENDIX" in t for t in texts)
    # Untruncated verbatim, as a real Word list paragraph (#483) -- the
    # 200-char _fill_appendix truncation is part of what masked the KFGXBW
    # loss.
    matches = [p for p in gen.doc.paragraphs
               if p.text == "Jan 2012-Dec 2016, Fellowship of Ornithological "
                             "Cryptography, Guild of Meandering Auditors, "
                             "Crab Hollow, ZQ"]
    assert len(matches) == 1
    assert _num_level(matches[0]) == "0"
    # No other paragraph duplicates it (the pre-#483 glyph path could
    # theoretically double-insert without the exact-text match above
    # catching it, since a glyph prefix changes the string).
    assert len([t for t in texts if "Ornithological Cryptography" in t]) == 1


def test_reconsider_insert_failure_falls_back_to_appendix(monkeypatch):
    """A reconsidered segment whose code has no usable section anchor must
    land in the appendix, not vanish (#221 hardening of the #214 path)."""
    gen = _generator()
    entry = {"text": "blob", "taxonomy_code": "D9", "extracted_fields": {}}
    gen._appendix_pending = [(entry, 5.0)]
    monkeypatch.setattr(
        gen, "_reclassify_entry_segments",
        lambda text, code: [("Custodian of Improbable Ledgers, Guild of "
                             "Meandering Auditors, 2011-2014", "D9")])

    gen._reconsider_appendix_entries()

    texts = [p.text for p in gen.doc.paragraphs]
    assert any("Custodian of Improbable Ledgers" in t for t in texts)


# ------------------------------------- K3/K4/K5 teaching anchors (#1225)

# The template's own K subsection headings, written out here rather than read
# back from the map under test: `_get_wcm_section_header` once returned strings
# that were substrings of no heading ("Curriculum Development", "Other
# Teaching"), so K4/K5 segments fell to the Appendix, and K3's "Mentoring"
# matched the MENTORING section instead of "Administrative teaching".
_K_SUBSECTION_HEADINGS = {
    "K1": "Didactic teaching",
    "K2": "Clinical teaching",
    "K3": "Administrative teaching",
    "K4": "Continuing education and professional education",
    "K5": "Other education/outreach activities",
}


def _heading_above(gen, text):
    """Text of the nearest bold paragraph above the first paragraph holding
    `text` -- the subsection heading that bullet sits under."""
    paras = gen.doc.paragraphs
    idx = next(i for i, p in enumerate(paras) if text in p.text)
    for para in reversed(paras[:idx]):
        if para.text.strip() and para.runs and para.runs[0].bold:
            return para.text.strip()
    return None


def _section_header_map_codes():
    """Every code `_get_wcm_section_header` maps explicitly. The map is a local
    of that method, so read its keys from the source instead of copying them
    into the test, where the copy could drift from the map."""
    tree = ast.parse(textwrap.dedent(
        inspect.getsource(WCMTemplateGenerator._get_wcm_section_header)))
    for node in ast.walk(tree):
        if (isinstance(node, ast.Assign)
                and any(isinstance(t, ast.Name) and t.id == "subsection_map"
                        for t in node.targets)):
            return list(ast.literal_eval(node.value))
    raise AssertionError("subsection_map not found in _get_wcm_section_header")


@pytest.mark.parametrize("code", sorted(_K_SUBSECTION_HEADINGS))
def test_reconsidered_k_segment_lands_under_its_own_template_subsection(code):
    gen = _generator()
    segment = f"Synthetic {code} teaching segment, Example University, 2015"

    assert gen._insert_reconsidered_segment(segment, code) is True

    assert _heading_above(gen, segment).startswith(_K_SUBSECTION_HEADINGS[code])


def _plain_left(para):
    """The left indent, in twips, of a paragraph with no list markup."""
    assert _num_level(para) is None, "expected no bullet (#1943)"
    return para.paragraph_format.left_indent.twips


@pytest.mark.parametrize(("code", "bulleted"), [("K1", False), ("L1", False), ("D1", True)])
def test_a_reconsidered_k_or_l_segment_carries_no_bullet(code, bulleted):
    # #1943: K and L entries render without bullets; a segment recovered into
    # those sections matches them. Other sections keep the list bullet.
    gen = _generator()
    segment = f"Synthetic {code} recovered segment, Example University, 2015"

    assert gen._insert_reconsidered_segment(segment, code) is True

    para = next(p for p in gen.doc.paragraphs if p.text == segment)
    if bulleted:
        assert _num_level(para) == "0"
    else:
        assert _plain_left(para) == plain_entry_left_twips(0)


def test_every_section_header_map_code_anchors_in_the_blank_template():
    codes = _section_header_map_codes()
    assert {"K3", "K4", "K5"} <= set(codes), "map harvest is vacuous"
    gen = _generator()

    unanchored = [c for c in codes
                  if not gen._insert_reconsidered_segment(
                      f"Synthetic {c} segment, Example University", c)]

    assert unanchored == []


@pytest.mark.parametrize("code", ["K3", "K4", "K5"])
def test_table_cell_overflow_for_k3_k4_k5_stays_in_the_teaching_section(code):
    """The third `_get_wcm_section_header` caller: a K entry whose paragraph
    sits in a table cell is inserted at the end of the teaching section, and
    queued for the Appendix only when no header resolves."""
    gen = _generator()
    cell_para = gen.doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    text = f"Synthetic {code} overflow prose, Example University, 2015"
    gen._overflow_entries = [(
        {"taxonomy_code": code, "text": text,
         "extraction_coverage": {"extraction_coverage_percent": 12}},
        cell_para, code)]

    gen._route_overflow_entries()

    assert gen._appendix_pending == []
    assert gen.stats["overflow_bullets_added"] == 1
    texts = [p.text for p in gen.doc.paragraphs]
    start = next(i for i, t in enumerate(texts) if t.strip() == "EDUCATIONAL CONTRIBUTIONS")
    end = next(i for i, t in enumerate(texts) if t.startswith("CLINICAL PRACTICE, INNOVATION"))
    assert start < texts.index(text) < end
    # #1943: the section's entries carry no bullet, and neither does its overflow.
    assert _plain_left(gen.doc.paragraphs[texts.index(text)]) == plain_entry_left_twips(0)


@pytest.mark.parametrize(("code", "bulleted"), [("K1", False), ("L1", False), ("D1", True)])
def test_body_paragraph_overflow_is_unbulleted_for_k_and_l(code, bulleted):
    """The other overflow route: the entry's own line is a body paragraph,
    replaced by the full text at the end of its section (#1943)."""
    gen = _generator()
    idx = gen._find_paragraph_with_text("EDUCATIONAL CONTRIBUTIONS")
    short = gen._insert_bulleted_entry(idx + 1, "Synthetic short line", None, list_level=0)
    text = f"Synthetic {code} overflow prose, Example University, 2015"
    gen._overflow_entries = [(
        {"taxonomy_code": code, "text": text,
         "extraction_coverage": {"extraction_coverage_percent": 12}},
        short, code)]

    gen._route_overflow_entries()

    assert gen.stats["overflow_bullets_added"] == 1
    para = next(p for p in gen.doc.paragraphs if p.text == text)
    if bulleted:
        assert _num_level(para) == "0"
    else:
        assert _plain_left(para) == plain_entry_left_twips(0)


@pytest.mark.parametrize("bulleted", [True, False])
def test_overflow_levels_survive_with_or_without_bullets(bulleted):
    """A title, a context segment and a label's content: levels 0, 1, 2, as
    ilvl when bulleted and as indent steps when not (#1943)."""
    gen = _generator()
    idx = gen._find_paragraph_with_text("EDUCATIONAL CONTRIBUTIONS")
    gen._insert_overflow_bullet(idx + 1, "Synthetic title\tScope:\tSynthetic scope prose",
                                bulleted=bulleted)
    paras = {p.text: p for p in gen.doc.paragraphs}
    expected = {"Synthetic title": 0, "Scope:": 1, "Synthetic scope prose": 2}
    for text, level in expected.items():
        if bulleted:
            assert _num_level(paras[text]) == str(level)
        else:
            assert _plain_left(paras[text]) == plain_entry_left_twips(level)


def test_reconsider_prompt_names_k3_k4_k5_by_their_taxonomy_labels(monkeypatch):
    """The prompt's code hint was stale in the same way as the anchors (K3
    "Mentoring/Advising", K4 "Curriculum Development"), so the model coded
    segments to meanings the taxonomy does not give those codes (#1225)."""
    prompts = []

    def fake_llm(*args, **kwargs):
        prompts.append(kwargs["messages"][-1]["content"])
        return {**_PRICED_LLM_RESULT, "content": "D1: Synthetic position"}

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm", fake_llm)
    taxonomy = json.loads(
        (_SRC / "unified_pipeline" / "core" / "taxonomy_v7.json").read_text())
    labels = {c["code"]: c["label"] for c in taxonomy["codes"]}

    _generator()._reclassify_entry_segments("Synthetic position, Example University", "D1")

    hint_lines = prompts[0].splitlines()
    for code in ("K3", "K4", "K5"):
        assert any(line.startswith(f"{code}: {labels[code]}") for line in hint_lines)


def test_option_off_restores_old_behavior():
    gen = _generator(recover_unrendered_records=False)
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records({"D1": [_d1_entry()]})

    assert len(gen.doc.paragraphs) == before


def test_recovery_defaults_on():
    assert WCMTemplateGenerator(verbose=False).recover_unrendered_records is True


def test_single_record_entry_not_touched():
    gen = _generator()
    entry = {
        "taxonomy_code": "D1",
        "text": "Aug 2019-Dec 2023, Wrangler of Nocturnal Metrics, Quexley "
                "College",
        "extracted_fields": {},
    }
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records({"D1": [entry]})

    assert len(gen.doc.paragraphs) == before


# ------------------------------------------------------------ generate() wiring

def test_generate_recovers_remainder_by_default(tmp_path):
    payload = {
        "document_uid": "TESTUID221",
        "entries": [_d1_entry()],
        "cv_owner": {"full_name": "Quenby Marblegate"},
    }
    input_path = tmp_path / "TESTUID221_fields.json"
    input_path.write_text(json.dumps(payload))

    out = run_stage6(str(input_path), str(tmp_path / "out.docx"), verbose=False)

    bullets = _bulleted_texts(Document(out).paragraphs)
    assert any("Quantitative Basketweaving" in t for t in bullets)


def test_deduped_entry_unique_record_still_recovered(tmp_path, monkeypatch):
    """The 2071 entry-168 residual: dedup keeps the longer near-duplicate and
    drops a fused sibling whose SECOND record exists nowhere else. Recovery
    must scan the pre-dedup entries so that record still surfaces."""
    monkeypatch.setattr(
        "unified_pipeline.stage_6_word_template.call_llm",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            RuntimeError("no LLM in tests")))
    kept_entry = {
        "element_idx_start": 168,
        "taxonomy_code": "P",
        "text": ("Member | Committee on Subterranean Balloon Safety Standards"
                 " | Guild of Meandering Auditors | reviews annual protocols "
                 "and certification checklists for subterranean balloon "
                 "safety inspections across member lodges, 2013-2016"),
        "extracted_fields": {
            "role": "Member",
            "committee": "Committee on Subterranean Balloon Safety Standards",
        },
    }
    dropped_entry = {
        "element_idx_start": 171,
        "taxonomy_code": "P",
        "text": ("Member | Committee on Subterranean Balloon Safety Standards"
                 " | Guild of Meandering Auditors\n"
                 "Chair | Panel of Improbable Weights and Measures 2013-2016 "
                 "| Norvale Metrology Circle"),
        "extracted_fields": {
            "role": "Member",
            "committee": "Committee on Subterranean Balloon Safety Standards",
        },
    }
    # Preflight: the fixture must actually trigger the dedup drop (title
    # containment) so this test exercises the pre-dedup snapshot. The kept
    # text states the dropped record's 2013-2016 range: a dropped entry whose
    # date range the kept one lacks is kept, not dropped (#666).
    from unified_pipeline.stage_6_word_template import deduplicate_entries
    assert deduplicate_entries([kept_entry, dropped_entry],
                               verbose=False) == [kept_entry]

    payload = {
        "document_uid": "TESTUID221P",
        "entries": [kept_entry, dropped_entry],
        "cv_owner": {"full_name": "Quenby Marblegate"},
    }
    input_path = tmp_path / "TESTUID221P_fields.json"
    input_path.write_text(json.dumps(payload))

    out = run_stage6(str(input_path), str(tmp_path / "dedup.docx"),
                     verbose=False)

    texts = [p.text for p in Document(out).paragraphs]
    hits = [t for t in texts if "Improbable Weights and Measures" in t]
    assert len(hits) == 1  # recovered exactly once


def test_generate_flag_off_drops_remainder(tmp_path):
    payload = {
        "document_uid": "TESTUID221",
        "entries": [_d1_entry()],
        "cv_owner": {"full_name": "Quenby Marblegate"},
    }
    input_path = tmp_path / "TESTUID221_fields.json"
    input_path.write_text(json.dumps(payload))

    out = run_stage6(str(input_path), str(tmp_path / "off.docx"),
                     verbose=False, recover_unrendered_records=False)

    texts = [p.text for p in Document(out).paragraphs]
    # The flag being off means the record is never touched at all -- checking
    # for a plain substring, not a glyph-prefixed one, since #483 R2 removed
    # the glyph representation the recovery path could have produced.
    assert not any("Quantitative Basketweaving" in t for t in texts)


@pytest.mark.parametrize("call", ["geographic_scope", "reclassify"])
def test_stage6_llm_outage_propagates_but_other_errors_default(monkeypatch, call):
    """A provider outage past the budget fails the run (#810); any other
    call_llm error still takes the method's default."""
    from unified_pipeline.llm.retry import LLMOutageError

    gen = _generator()
    gen.cv_owner_location = {"primary_location": {"institution": "Example Medical College",
                                                  "city": "Springfield", "state": "ZZ"}}
    entry = {"text": "Visiting Lecturer", "extracted_fields": {"organization": "Example Institute"}}

    def run():
        if call == "geographic_scope":
            return gen._classify_geographic_scope(entry)
        return gen._reclassify_entry_segments("Visiting Lecturer, Example Institute, 2010", "P")

    def outage(*args, **kwargs):
        raise LLMOutageError("provider down", seconds_waited=1800.0)

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm", outage)
    with pytest.raises(LLMOutageError):
        run()

    def blip(*args, **kwargs):
        raise RuntimeError("connection reset")

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm", blip)
    assert run() == ("National" if call == "geographic_scope" else None)


_PRICED_LLM_RESULT = {"content": '{"scope": "National"}', "cost": 0.125,
                      "prompt_tokens": 11, "completion_tokens": 5,
                      "cache_read_tokens": 2, "cache_write_tokens": 1}


@pytest.mark.parametrize("call", ["geographic_scope", "reclassify"])
def test_stage6_llm_calls_are_added_to_the_generators_usage(monkeypatch, call):
    """#1177: stage 6 reported no cost at all; every one of its call_llm sites
    now feeds the generator's LlmUsage. The scope decision makes two calls when
    the activity names no reach of its own (#1579): the named-reach step, then
    the distance step."""
    gen = _generator()
    gen.cv_owner_location = {"primary_location": {"institution": "Example Medical College",
                                                  "city": "Springfield", "state": "ZZ"}}
    entry = {"text": "Visiting Lecturer", "extracted_fields": {"organization": "Example Institute"}}
    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm",
                        lambda *args, **kwargs: dict(_PRICED_LLM_RESULT))

    if call == "geographic_scope":
        gen._classify_geographic_scope(entry)
    else:
        gen._reclassify_entry_segments("Visiting Lecturer, Example Institute, 2010", "P")

    calls = 2 if call == "geographic_scope" else 1
    assert gen.llm_usage.cost == pytest.approx(0.125 * calls)
    assert gen.llm_usage.prompt_tokens == 11 * calls
    assert gen.llm_usage.completion_tokens == 5 * calls
    assert gen.llm_usage.cache_read_tokens == 2 * calls
    assert gen.llm_usage.cache_write_tokens == 1 * calls


def test_run_stage6_hands_the_callers_usage_to_the_generator(monkeypatch):
    from unified_pipeline import stage_6_word_template as stage6
    from unified_pipeline.llm_client import LlmUsage

    seen = {}

    class _FakeGenerator:
        def __init__(self, **kwargs):
            seen.update(kwargs)

        def generate(self, *args, **kwargs):
            return "unused.docx"

    monkeypatch.setattr(stage6, "WCMTemplateGenerator", _FakeGenerator)
    usage = LlmUsage()

    stage6.run_stage6("unused.json", llm_usage=usage)

    assert seen["llm_usage"] is usage


def test_remaining_appendix_opens_with_the_shared_cviche_note_box():
    # #534/#1388: the bullet writer creates T. APPENDIX with the same CViche
    # note box as _fill_appendix -- never claiming the entries appear nowhere
    # else (stage 6 cannot prove that) -- and generate() restates its count.
    from unified_pipeline.stage6.formatting import is_cviche_box
    from unified_pipeline.stage6.sections.appendix import (
        APPENDIX_NOTE_TITLE,
        appendix_note_text,
    )

    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen._add_remaining_to_appendix([("Example Leftover Committee, 2019-2021", "T", 0.0)])
    [box] = [t for t in gen.doc.tables if is_cviche_box(t)]
    assert [p.text for p in box.cell(0, 0).paragraphs] == [APPENDIX_NOTE_TITLE, appendix_note_text(1)]
    # The box sits right after the heading, before the bullet.
    body = list(gen.doc.element.body)
    heading = next(i for i, el in enumerate(body) if el.tag.endswith("}p") and
                   "".join(t.text or "" for t in el.iter(qn("w:t"))) == "T. APPENDIX")
    assert body[heading + 1] is box._tbl

    # No blank paragraph around the box: the heading carries the space above
    # it and the first line after it the space below (#1388).
    assert body[heading].find(qn("w:pPr")).find(qn("w:spacing")).get(qn("w:after")) == "120"
    after = body[heading + 2]
    assert "".join(t.text or "" for t in after.iter(qn("w:t"))) == "Example Leftover Committee, 2019-2021"
    assert after.find(qn("w:pPr")).find(qn("w:spacing")).get(qn("w:before")) == "120"

    gen._set_appendix_note_count(7)
    assert box.cell(0, 0).paragraphs[1].text == appendix_note_text(7)


def test_the_appendix_note_counts_numbered_lines_and_every_recovered_bullet():
    """#1388: the box is written by the first writer; the reconsider pass and
    the unrendered-record recovery add bullets after it, and the box must say
    the total (the sidecar's appendix_diversion counts sum to the same)."""
    from unified_pipeline.stage6.formatting import is_cviche_box
    from unified_pipeline.stage6.sections.appendix import (
        RecoveredLine,
        appendix_note_text,
    )

    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen._write_appendix_header(1)
    gen._reconsider_appendix_entries = lambda: [RecoveredLine("T", "a"), RecoveredLine("T", "b")]
    gen._recover_unrendered_records = lambda entries_by_code, cv_owner: [RecoveredLine("D1", "c")]
    recovered = gen._recover_appendix_lines([{"text": "x", "taxonomy_code": "T"}], {}, None)
    assert len(recovered) == 3
    [box] = [t for t in gen.doc.tables if is_cviche_box(t)]
    assert box.cell(0, 0).paragraphs[1].text == appendix_note_text(4)


def _withheld_generator():
    from unified_pipeline.stage6.normalization.pii import WithheldItem

    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen._pii_result.withheld = [WithheldItem("DEA registration number", "Personal Data", 0)]
    gen._reconsider_appendix_entries = lambda: []
    return gen


def _appendix_note(gen):
    from unified_pipeline.stage6.formatting import is_cviche_box

    [box] = [t for t in gen.doc.tables if is_cviche_box(t)]
    return box.cell(0, 0).paragraphs[1].text


def test_an_appendix_holding_only_the_withheld_notice_states_no_entry_count():
    """#1582: with nothing written or recovered, the withheld notice alone
    opens the Appendix; it is not an entry, so the box restated a count of 0
    as "These 0 entries ..." above no entries. It now says what is there."""
    from unified_pipeline.stage6.pii_pass import PII_REDACTED_NOTICE
    from unified_pipeline.stage6.sections.appendix import (
        APPENDIX_NOTE_WITHHELD_ONLY_TEXT,
    )

    gen = _withheld_generator()
    recovered = gen._recover_appendix_lines([], {}, None)

    assert recovered == []
    note = _appendix_note(gen)
    assert note == APPENDIX_NOTE_WITHHELD_ONLY_TEXT
    assert "0 entries" not in note
    assert PII_REDACTED_NOTICE in [p.text for p in gen.doc.paragraphs]


def test_the_withheld_notice_is_not_counted_beside_a_real_appendix_entry():
    """#1582: the notice beside one numbered line keeps the one-entry wording."""
    from unified_pipeline.stage6.sections.appendix import appendix_note_text

    gen = _withheld_generator()
    gen._write_appendix_header(1)
    gen._recover_appendix_lines([{"text": "x", "taxonomy_code": "T"}], {}, None)

    assert _appendix_note(gen) == appendix_note_text(1)
    assert _appendix_note(gen).startswith("This entry from your original CV")


# ------------------------------------ #530: foreign scaffolding never recovered

def test_add_remaining_to_appendix_drops_foreign_template_scaffolding():
    """The bullet writer's own filter (the recovery path, not `_fill_appendix`)
    drops another institution's instruction line and a "label: N/A" placeholder
    and still writes the genuine segment."""
    gen = _generator()
    written = gen._add_remaining_to_appendix([
        ("C. Sample Appointments (include institution, title and dates)", "T", 0.0),
        ("1. Sample Leave: N/A", "T", 0.0),
        ("Served on the sample review panel for the Example Society", "T", 0.0),
    ])
    assert [(line.code, line.text) for line in written] == [
        ("T", "Served on the sample review panel for the Example Society")]
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert "Served on the sample review panel for the Example Society" in bullets
    assert not any("Sample Appointments" in t or "Sample Leave" in t for t in bullets)


# ------------------------------------------- #1230: T-coded fused entries

_T_ROWS = [
    f"Example Grant Title {title} | Example Funder {funder} | Role {n} | {n}7% | ${n}1,500 | 19{n}1-19{n}2"
    for n, (title, funder) in enumerate((("Zeolite Catalysis", "Marigold Foundation"),
                                         ("Quartzite Weathering", "Nasturtium Institute"),
                                         ("Obsidian Fracturing", "Hyacinth Society")), start=3)
]


def _t_entry():
    return {
        "element_idx_start": 7,
        "element_type": "table_row",
        "hierarchy": ["COMPLETED GRANTS"],
        "taxonomy_code": "T",
        "text": "\n".join(_T_ROWS),
        "extraction_coverage": {"extraction_coverage_percent": 3.0},
        "extracted_fields": {},
    }


def _appendix_count(gen, text):
    """How many paragraphs of the document carry *text*."""
    return sum(text in p.text for p in gen.doc.paragraphs)


def test_t_fused_entry_rows_past_the_cap_render_once_in_the_appendix():
    """`_fill_appendix` renders a T body whole (#1230), so every row reaches
    the page and the recovery pass finds nothing to re-add."""
    gen = _generator()
    entry = _t_entry()
    assert len(entry["text"]) > 200
    gen._fill_appendix([entry])
    gen._recover_unrendered_records({"T": [entry]})
    assert gen.stats["unrendered_records_recovered"] == 0
    assert all(_appendix_count(gen, row.split(" | ")[0]) == 1 for row in _T_ROWS)


def test_t_single_line_over_the_cap_is_not_printed_twice():
    """EBYSBC XELRLZ shape: a 302-character T footnote printed cut as a
    numbered item, then again whole as a recovered bullet (#1230)."""
    footnote = ("Synthetic footnote 2004: the zeolite appointment above was held "
                "jointly with the Marigold Institute, Example City, and the role "
                "carried teaching duties in quartzite petrology, field courses in "
                "obsidian sites, and a review panel seat for the Example Society.")
    assert len(footnote) > 200
    entry = {**_t_entry(), "text": footnote, "element_type": "paragraph"}
    gen = _generator()
    gen._fill_appendix([entry])
    gen._recover_unrendered_records({"T": [entry]})
    assert _appendix_count(gen, footnote) == 1
    assert gen.stats["unrendered_records_recovered"] == 0


def test_t_undated_lines_and_near_duplicate_rows_past_the_cap_are_kept():
    """Undated non-record lines (a duty bullet) and sibling rows that differ
    only in a short word or a digit reach the page whole (#1230)."""
    lines = ["Synthetic duty: directed the zeolite workgroup and its quarterly reviews",
             "Member | Marigold Panel A | Example Society",
             "Member | Marigold Panel B | Example Society",
             "Member | Marigold Panel 7 | Example Society",
             "Synthetic note: obsidian outreach was shared with the nasturtium team"]
    entry = {**_t_entry(), "text": "\n".join(lines)}
    gen = _generator()
    gen._fill_appendix([entry])
    body = "\n".join(p.text for p in gen.doc.paragraphs)
    assert all(_clean_inline_tabs(line) in body for line in lines)


def test_t_row_already_on_the_page_is_not_duplicated():
    gen = _generator()
    entry = _t_entry()
    gen.doc.add_paragraph(_T_ROWS[0])  # rendered whole, e.g. by another section
    gen._recover_unrendered_records({"T": [entry]})
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert _T_ROWS[0] not in bullets
    assert _T_ROWS[1] in bullets and _T_ROWS[2] in bullets


def test_t_row_whose_title_cell_alone_renders_is_still_recovered():
    """The per-cell check treats a rendered title as the record rendered; for T
    the whole line must be accounted for, so funder and amount are not lost."""
    gen = _generator()
    entry = _t_entry()
    gen.doc.add_paragraph("Example Grant Title Quartzite Weathering")
    gen._recover_unrendered_records({"T": [entry]})
    assert _T_ROWS[1] in _bulleted_texts(gen.doc.paragraphs)


def test_whole_record_rendered_verdicts():
    row = _T_ROWS[0]
    sets = [set(_RENDER_TOKEN_RE.findall(_norm(row)))]
    assert _whole_record_rendered(row, _squash_text(row), sets) is True
    assert _whole_record_rendered(_T_ROWS[1], "", sets) is False
    assert _whole_record_rendered("A | B | 1", "", []) is None


def test_whole_record_rendered_cut_line_vouches_only_for_a_record_it_carries_whole():
    from unified_pipeline.stage6.render_check import _record_tokens
    carried = "Example Grant Title Zeolite Catalysis | Example Funder Marigold Foundation | 1987"
    cut = [_record_tokens(carried + " | extra words here ...")]
    assert _whole_record_rendered(carried, "", [], cut) is True
    longer = carried + " | with a tail the cut line lost entirely"
    assert _whole_record_rendered(longer, "", [], cut) is False


# #829 (YDSVDB): the revised template's mentee rows are "<label> | <answer>".
# One stray "Goals/expected Outcomes — ..." line on the page vouched, through
# its label, for every mentee's goal line. Labels are the template's own text;
# the answers are invented.
_GOALS_LABEL = "Goals/expected Outcomes"
_SUPERVISION_LABEL = "Type of Supervision (Research, clinical, teaching, leadership)"
_STRAY_GOAL = f"{_GOALS_LABEL} — Catalogued the marigold quartzite herbarium"


def _output_line_tokens(*lines):
    return [set(_RENDER_TOKEN_RE.findall(_norm(line))) for line in lines]


def test_a_template_label_cell_is_no_verbatim_piece():
    line = f"{_GOALS_LABEL} | Shorter zeolite clinic waits"
    assert record_pieces(line) == [_squash_text(" Shorter zeolite clinic waits")]
    assert len(_entry_pieces(line)) == 2  # whole-entry checks keep the label piece (#744)


def test_a_template_label_cell_does_not_vouch_for_an_unrendered_record():
    line = f"{_GOALS_LABEL} | Shorter zeolite clinic waits for nasturtium patients"
    assert _record_rendered(line, _squash_text(_STRAY_GOAL),
                            _output_line_tokens(_STRAY_GOAL)) is False
    rendered = "Goals: Shorter zeolite clinic waits for nasturtium patients"
    assert _record_rendered(line, _squash_text(rendered),
                            _output_line_tokens(rendered)) is True


def test_a_short_answer_under_a_template_label_is_unverifiable():
    """The label's 5 tokens no longer outvote the answer's 1, so the line
    matches no stray copy of the label and is None, not rendered."""
    line = f"{_SUPERVISION_LABEL} | Leadership"
    stray = f"{_SUPERVISION_LABEL} — Lead"
    assert _record_rendered(line, _squash_text(stray), _output_line_tokens(stray)) is None


def test_record_chunks_drop_only_template_label_cells():
    plain = "Basalt Obsidian Feldspar | Hyacinth Delphinium"
    assert record_chunks(plain) == [plain, "Basalt Obsidian Feldspar ", " Hyacinth Delphinium"]
    labelled = f"{_GOALS_LABEL} | Hyacinth Delphinium"
    assert record_chunks(labelled) == [" Hyacinth Delphinium"]
    assert record_chunks(f"{_GOALS_LABEL} | {_SUPERVISION_LABEL}") == []


def test_answers_to_different_labels_need_not_share_one_output_line():
    """A5IZ6Q shape: a grant form row carries the PI name and the role, each
    under its own label; stage 6 renders them as two table rows. Once the
    labels stop vouching, the answers are judged one by one (each too thin to
    verify), not as one joined line that no single output row holds."""
    line = ("Name of Principal Investigator: | Quillon Marbury\tYour role*\toversight"
            "\tYour percent (%) effort:\t1%")
    page = ["Name of Principal Investigator: | Quillon Marbury", "Your role: | oversight",
            "Your percent (%) effort: | 1%"]
    haystack = _squash_text("\n".join(page))
    assert _record_rendered(line, haystack, _output_line_tokens(*page)) is None


def _squash_text(text):
    from unified_pipeline.stage6.normalization import _squash
    return _squash(text)


def test_segments_cover_source_threshold():
    words = ["zeolite", "marigold", "quartzite", "nasturtium", "basalt",
             "chrysanthemum", "obsidian", "hyacinth", "feldspar", "delphinium"] * 2
    source = "\n".join(f"Synthetic Record {w}{'x' * i} Entry" for i, w in enumerate(words))
    faithful = source.split("\n")
    assert segments_cover_source(faithful, source) is True
    assert segments_cover_source(["[all 20 records follow]"], source) is False
    assert segments_cover_source(faithful[:10], source) is False
    assert segments_cover_source([], "tiny") is True  # too few tokens to judge


def test_t_header_rows_the_appendix_would_drop_are_not_recovered():
    """A T entry stage 3b confirmed as a table header row is scaffolding the
    Appendix drops; recovery must not resurrect its lines."""
    header_rows = [
        "Years Inclusive | Grant Number and Title | Source | Annual Direct Costs",
        "Years Inclusive | Years Inclusive | Grant Number and Title | Source",
    ] * 4  # past APPENDIX_MAX_CHARS, so the entry reaches the candidate scan
    entry = {
        "element_idx_start": 3, "element_type": "table_row",
        "hierarchy": ["GRANTS"], "taxonomy_code": "T",
        "text": "\n".join(header_rows),
        "classification_reasoning":
            "[T-validation confirmed] Column header row - structural artifact",
        "extraction_coverage": {"extraction_coverage_percent": 0.0},
        "extracted_fields": {},
    }
    gen = _generator()
    gen._recover_unrendered_records({"T": [entry]})
    assert not any("Years Inclusive" in t
                   for t in _bulleted_texts(gen.doc.paragraphs))


def test_t_sibling_row_differing_only_in_amount_and_years_is_recovered():
    """Two rows share title, funder and role words and differ in amount and
    years. The one on the page must not vouch for the other (#1230)."""
    base = "Example Grant Title Zeolite Catalysis | Example Funder Marigold Foundation | Site Lead | "
    on_page = base + "$1,500 | 1991-1992"
    sibling = base + "$8,200 | 1987-1988"
    entry = _t_entry()
    entry["text"] = "\n".join([on_page, sibling, _T_ROWS[1]])
    gen = _generator()
    gen.doc.add_paragraph(on_page)
    gen._recover_unrendered_records({"T": [entry]})
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert sibling in bullets
    assert on_page not in bullets


def test_t_prose_lines_past_the_cap_are_recovered():
    """Lines of a capped T entry that are not record-shaped (no pipe, tab or
    dated prefix) are recovered too: the cap cut them just the same."""
    lines = [f"Doe J. Synthetic {w} lecture series, Example University, 200{i}."
             for i, w in enumerate(("zeolite", "marigold", "quartzite", "nasturtium",
                                    "basalt", "obsidian"))]
    entry = _t_entry()
    entry["text"] = "\n".join(lines)
    gen = _generator()
    gen._fill_appendix([entry])
    gen._recover_unrendered_records({"T": [entry]})
    body = "\n".join(p.text for p in gen.doc.paragraphs)
    assert all(line in body for line in lines)


def test_t_single_record_entry_cut_by_the_cap_is_recovered():
    long_line = ("Doe J. 2003. Synthetic invited lecture on zeolite catalysis at the Example "
                 "Symposium, hosted by the Marigold Institute, Example City, with a "
                 "long tail of detail that lies well past the two hundred character cap.")
    entry = _t_entry()
    entry["text"] = long_line
    gen = _generator()
    gen._fill_appendix([entry])
    gen._recover_unrendered_records({"T": [entry]})
    assert long_line in "\n".join(p.text for p in gen.doc.paragraphs)


def test_t_undated_prose_and_undated_merged_cell_rows_are_not_recovered():
    """A line with no year and no row shape (template text, an objective) is
    not a record, and neither is a merged-cell row that collapses to one such
    cell: recovery leaves both to the Appendix pointer."""
    prose = ("Example instruction: list every activity in reverse order and "
             "describe each briefly, with the role held and the venue used. " * 2)
    merged = " | ".join(["Example label for a category of activities"] * 3)
    entry = _t_entry()
    entry["text"] = "\n".join([prose, merged])
    gen = _generator()
    gen._recover_unrendered_records({"T": [entry]})
    assert gen.stats["unrendered_records_recovered"] == 0
    body = "\n".join(p.text for p in gen.doc.paragraphs)
    assert "Example instruction" not in body and "Example label for" not in body


def test_t_merged_cell_description_after_a_record_is_recovered_once():
    """A description spanning six grid columns is flattened to six copies; the
    recovered line carries it once."""
    description = ("Funding for a synthetic pilot study of zeolite weathering "
                   "in Example County, used to collect first-round samples.")
    entry = _t_entry()
    entry["text"] = "\n".join([_T_ROWS[0], " | ".join([description] * 6), _T_ROWS[1]])
    gen = _generator()
    gen._fill_appendix([entry])
    gen._recover_unrendered_records({"T": [entry]})
    body = "\n".join(p.text for p in gen.doc.paragraphs)
    assert description in body
    assert description + " | " + description not in body


def test_t_description_row_after_a_record_is_recovered():
    """An undated description row directly after a grant row travels with it."""
    description = ("Funding for a synthetic pilot study of zeolite weathering "
                   "in Example County, used to collect first-round samples.")
    entry = _t_entry()
    entry["text"] = "\n".join([_T_ROWS[0], description, _T_ROWS[1]])
    gen = _generator()
    gen._fill_appendix([entry])
    gen._recover_unrendered_records({"T": [entry]})
    assert description in "\n".join(p.text for p in gen.doc.paragraphs)


def test_t_entry_under_the_cap_is_left_alone():
    entry = _t_entry()
    entry["text"] = "Doe J. Short synthetic note, Example University."
    gen = _generator()
    gen._recover_unrendered_records({"T": [entry]})
    assert gen.stats["unrendered_records_recovered"] == 0


def test_segments_cover_source_rejects_a_reply_that_summarises_a_few_lines():
    """17 of 20 lines kept verbatim plus one bracketed summary: pooled coverage
    is high, but three source lines are carried by no segment (#1230)."""
    words = ["zeolite", "marigold", "quartzite", "nasturtium", "basalt",
             "chrysanthemum", "obsidian", "hyacinth", "feldspar", "delphinium",
             "granite", "snapdragon", "tourmaline", "gardenia", "pumice",
             "lavender", "dolomite", "begonia", "gypsum", "camellia"]
    lines = [f"Doe J. Synthetic {w} outcomes. Journal of Invented {w.title()} stuff {i}."
             for i, w in enumerate(words)]
    source = "\n".join(lines)
    reply = lines[:17] + ["[3 further publications follow]"]
    assert segments_cover_source(lines, source) is True
    assert segments_cover_source(reply, source) is False


def test_segments_cover_source_ignores_a_repeated_page_header_the_reply_drops():
    """A page-break header repeats in the source and a faithful reply rightly
    drops it; that must not reject the reply (#1230)."""
    pairs = (("zeolite", "geyser"), ("marigold", "estuary"), ("quartzite", "lagoon"),
             ("nasturtium", "tundra"), ("basalt", "savanna"), ("obsidian", "fjord"),
             ("hyacinth", "canyon"), ("feldspar", "glacier"))
    lines = [f"Doe J. Synthetic {a} and {b} outcomes {i}." for i, (a, b) in enumerate(pairs)]
    header = "Example Header Continued"
    source = "\n".join(lines[:4] + [header] + lines[4:] + [header])
    assert segments_cover_source(lines, source) is True
    assert segments_cover_source(lines[:-1], source) is False


def test_segments_cover_source_rejects_a_reply_that_drops_a_near_duplicate_sibling():
    """Six committee rows told apart by one word: a reply that keeps five of
    them carries 92% of the sixth row's tokens, so the pooled per-line share
    forgave the loss. A row with a sibling needs every one of its tokens in
    the reply (#1230)."""
    names = ("Zeolite", "Marigold", "Quartzite", "Nasturtium", "Obsidian", "Hyacinth")
    rows = [f"Member | {n} Committee | Example Department of Synthetic Geology | "
            f"Example College of Sciences | Example University | 1987-1992" for n in names]
    source = "\n".join(rows)
    assert segments_cover_source(rows, source) is True
    assert segments_cover_source(rows[:-1], source) is False
    assert segments_cover_source(rows[1:], source) is False
    # Two sibling rows merged into one segment still carry both whole.
    assert segments_cover_source(["; ".join(rows[:2])] + rows[2:], source) is True


def test_segments_cover_source_accepts_a_sibling_row_split_across_segments():
    """The prompt asks for one segment per responsibility, so a reply that
    keeps every token of a sibling row but splits it at its separators is a
    faithful reply and must not be rejected (#1230). Only a lost token fails."""
    names = ("Zeolite", "Marigold", "Quartzite", "Nasturtium", "Obsidian", "Hyacinth")
    rows = [f"Director, {n} Laboratory, Example Hospital, 2004-2010; supervised residents; "
            f"chaired the quality committee" for n in names]
    source = "\n".join(rows)
    split = [part for row in rows for part in row.split("; ")]
    assert segments_cover_source(split, source) is True
    # A split reply that loses the one word telling a sibling apart is rejected.
    lost = [part.replace("Hyacinth", "") for part in split]
    assert segments_cover_source(lost, source) is False


def test_segments_cover_source_short_line_counts_as_carried():
    """A line too short to verify (under RENDER_TOKEN_MIN_COUNT tokens) cannot
    be held against the reply, though the long lines are all kept (#1230)."""
    pairs = (("zeolite", "geyser"), ("marigold", "estuary"), ("quartzite", "lagoon"),
             ("nasturtium", "tundra"), ("basalt", "savanna"), ("obsidian", "fjord"))
    lines = [f"Doe J. Synthetic {a} and {b} outcomes {i}." for i, (a, b) in enumerate(pairs)]
    source = "\n".join(lines + ["Page 2"])
    assert segments_cover_source(lines, source) is True


def test_segments_cover_source_pools_tokens_for_a_line_without_a_sibling():
    """The all-tokens rule is for sibling rows only: a line with no near
    duplicate keeps the pooled share, so a reply that splits it across two
    segments, or loses one of its five tokens (0.8 carried), is still
    believed (#1230)."""
    triples = (("zeolite", "geyser", "basalt"), ("marigold", "estuary", "savanna"),
               ("quartzite", "lagoon", "granite"), ("nasturtium", "tundra", "prairie"),
               ("obsidian", "fjord", "meadow"), ("hyacinth", "glacier", "plateau"))
    lines = [f"Synthetic {a} {b} {c} outcomes." for a, b, c in triples]
    source = "\n".join(lines)
    halves = [half for line in lines for half in line.split(" ", 2)]
    assert segments_cover_source(halves, source) is True
    # One token of a non-sibling line lost: 4 of 5 still carried, accepted.
    # A strict all-tokens rule on every line would reject this reply.
    lossy = [line.replace("geyser ", "") for line in lines]
    assert segments_cover_source(lossy, source) is True
    # Two tokens lost from the same line: 3 of 5, rejected.
    assert segments_cover_source([line.replace("geyser ", "").replace("basalt ", "")
                                  for line in lines[:1]] + lines[1:], source) is False


def _two_by_two_grid():
    """Two committees, each served in the same two terms: every token of each
    row also sits in another row (#1230). Each row has 9 record tokens, as a
    real row does; the two terms of one committee share 7 of them, under
    `RECLASSIFY_MIN_TOKEN_COVERAGE`."""
    return [f"Member | {name} Committee | Example State Medical College | {term}"
            for name in ("Admissions", "Curriculum") for term in ("2010-2012", "2014-2016")]


def test_segments_cover_source_rejects_a_dropped_row_of_a_two_by_two_grid():
    """Dropping any one row of the grid leaves each of its tokens in the
    rows the reply kept, so the all-tokens rule accepted it. Each sibling
    row needs its own copy of the tokens it shares (#1230)."""
    rows = _two_by_two_grid()
    source = "\n".join(rows)
    assert segments_cover_source(rows, source) is True
    for dropped in rows:
        assert segments_cover_source([row for row in rows if row != dropped], source) is False


def test_segments_cover_source_counts_split_and_merged_grid_rows():
    """Rows split at their separators, or two rows written whole into one
    segment, keep every row's tokens and are accepted. A segment that writes
    two terms as one row carries the shared words once and is rejected,
    which keeps the entry whole (#1230)."""
    rows = _two_by_two_grid()
    source = "\n".join(rows)
    split = [cell for row in rows for cell in row.split(" | ")]
    assert segments_cover_source(split, source) is True
    assert segments_cover_source(["; ".join(rows[:2])] + rows[2:], source) is True
    merged = ("Member | Admissions Committee | Example State Medical College | "
              "2010-2012 and 2014-2016")
    assert segments_cover_source([merged] + rows[2:], source) is False


def test_segments_cover_source_unrelated_lines_do_not_pay_for_a_dropped_grid_row():
    """Unrelated lines that between them hold every word of a dropped grid row
    do not count for it: only segments that are a piece of a sibling row, or
    hold one whole, are counted (#1230)."""
    rows = _two_by_two_grid()
    unrelated = ["Member of the Faculty Senate, Riverbend Example University, 2014",
                 "Admissions Prize Committee juror, 2016",
                 "State Medical College seminar organiser"]
    source = "\n".join(rows + unrelated)
    assert segments_cover_source(rows + unrelated, source) is True
    assert segments_cover_source([row for row in rows if row != rows[1]] + unrelated,
                                 source) is False


def test_segments_cover_source_a_partner_line_may_lose_a_word_no_sibling_row_holds():
    """A line that holds a sibling row whole (its partner) is counted only for
    the sibling row's words: trimming a word of its own keeps the pooled
    share that any line without a sibling gets (#1230)."""
    rows = _two_by_two_grid()
    partner = rows[0] + " | organised the annual workshop retreat"
    source = "\n".join(rows + [partner])
    trimmed = partner.replace(" retreat", "")
    assert segments_cover_source(rows + [trimmed], source) is True
    assert segments_cover_source(rows[1:] + [partner], source) is False


def test_record_siblings_skips_a_sibling_that_holds_the_whole_line():
    """A line every token of which sits in a sibling is part of that sibling,
    so the sibling's render carries it and it needs no token of its own
    (#1230). The grid rows keep each other as siblings."""
    rows = _two_by_two_grid()
    short = rows[0].replace("Example State Medical College | ", "")
    siblings = record_siblings([short] + rows)
    assert _record_tokens(short) <= _record_tokens(rows[0])
    assert _record_tokens(rows[0]) not in siblings[short]
    assert siblings[short] == [_record_tokens(rows[2])]
    assert set(rows) <= set(siblings)


def test_non_t_grid_row_rendered_only_through_its_siblings_is_recovered():
    """Three rows of the grid on the page vouched for the fourth: every one
    of its words sits in a sibling's row. A sibling row is rendered only by
    an output line that holds a token each sibling lacks (#1230)."""
    rows = _two_by_two_grid()
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    for row in rows[:3]:
        gen.doc.add_paragraph(row.replace(" | ", ", "))
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]
    assert bullets == [rows[3].replace(" | ", " \u2014 ")]


def test_non_t_grid_rows_recovered_one_after_another_each_land_once():
    """One grid row on the page: the other three are recovered in turn, and
    each recovered row then counts as on the page for the rows after it. The
    lines it may vouch through grow as rows are recovered (web207 crashed on
    a list that did not grow with them, #1230)."""
    rows = _two_by_two_grid()
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    gen.doc.add_paragraph(rows[0].replace(" | ", ", "))
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]
    assert sorted(bullets) == sorted(row.replace(" | ", " \u2014 ") for row in rows[1:])


def test_non_t_grid_rows_all_rendered_are_not_recovered():
    """Every row on the page, reformatted: nothing is re-inserted (#1230)."""
    rows = _two_by_two_grid()
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    for row in rows:
        gen.doc.add_paragraph(row.replace(" | ", ", ").replace("-", " - "))
    gen._recover_unrendered_records({"P": [entry]})
    assert not [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]


def test_record_siblings_holds_the_same_row_in_another_term():
    """The two terms of one committee share 7 of 9 tokens, under the sibling
    share, yet differ only in their numbers: they are siblings (#1230). Rows
    that differ only in a month name are not: a render may write the month as
    a number."""
    rows = _two_by_two_grid()
    siblings = record_siblings(rows)
    assert _record_tokens(rows[1]) in siblings[rows[0]]
    assert _record_tokens(rows[0]) in siblings[rows[1]]
    months = [f"Visiting Lecturer | Example State Medical College | {month} 2015"
              for month in ("January", "March")]
    assert record_siblings(months) == {}


def test_non_t_one_committee_in_two_terms_recovers_the_missing_term():
    """One committee in two terms, one term on the page: its row vouched for
    the other term, which shares every word. A row whose own numbers are on
    no rendered line with its words is recovered (#1230)."""
    rows = _two_by_two_grid()[:2]
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    gen.doc.add_paragraph(rows[0].replace(" | ", ", "))
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]
    assert bullets == [rows[1].replace(" | ", " \u2014 ")]


def _twin_rows():
    """Two rows told apart only by a word under five letters: `_record_tokens`
    gives them one token set (#1956)."""
    return [f"Member | Graduate Medical Education Committee | {site} | 2015-2018"
            for site in ("NYP", "WCM")]


def test_segments_cover_source_rejects_a_reply_that_drops_a_twin_row():
    """Twin rows share one token set, so neither was a sibling of the other
    and a reply that kept one carried every token of both. Each twin needs
    its own copy of the tokens (#1956)."""
    rows = _twin_rows()
    source = "\n".join(rows)
    assert _record_tokens(rows[0]) == _record_tokens(rows[1])
    assert segments_cover_source(rows, source) is True
    assert segments_cover_source(rows[:1], source) is False
    assert segments_cover_source(rows[1:], source) is False


def test_record_twins_maps_each_twin_row_to_the_other_squashed():
    """`record_siblings` cannot hold twins: no token tells them apart. Their
    squashed text does (#1956). A row with other words is no twin, and a row
    repeated with other spacing is the same record, not a twin."""
    rows = _twin_rows()
    other = "Member | Graduate Medical Education Committee | Example Hospital | 2015-2018"
    assert record_siblings(rows) == {}
    assert record_twins(rows + [other]) == {rows[0]: [_squash(rows[1])],
                                            rows[1]: [_squash(rows[0])]}
    assert record_twins([rows[0], rows[0].replace(" | ", " |  ")]) == {}


def test_non_t_twin_row_vouched_for_only_by_its_twin_is_recovered():
    """One twin on the page vouched for the other through their shared words.
    A row whose only vouching line is its twin's row verbatim is recovered
    (#1956)."""
    rows = _twin_rows()
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    gen.doc.add_paragraph(rows[1])
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]
    assert bullets == [rows[0].replace(" | ", " \u2014 ")]


def test_non_t_twin_rows_neither_on_the_page_are_both_recovered():
    """Neither twin on the page: the first is recovered, and then must not
    vouch for the second (#1956)."""
    rows = _twin_rows()
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]
    assert sorted(bullets) == [row.replace(" | ", " \u2014 ") for row in rows]


def test_non_t_twin_rendered_reformatted_beside_its_twin_verbatim_is_not_recovered():
    """One twin on the page reformatted past every verbatim piece, the other
    verbatim: the reformatted line vouches by its words and holds no twin, so
    it may be the row itself and the verdict stands (#1956). Both rows on the
    page verbatim re-insert nothing either."""
    rows = _twin_rows()
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    reformatted = "Graduate Medical Education (GME) Committee, NYP: Member, 2015 - 2018"
    assert not any(piece in _squash(reformatted) for piece in record_pieces(rows[0]))
    for page in ([reformatted, rows[1]], rows):
        gen = _generator()
        for line in page:
            gen.doc.add_paragraph(line)
        gen._recover_unrendered_records({"P": [entry]})
        assert not [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]


def test_sibling_record_rendered_keeps_a_twin_verdict_no_twin_line_overrules():
    """A twin's verdict turns False only when output lines vouch for it and
    every one is a twin's row verbatim (#1956). A line that holds one of its
    verbatim pieces without most of its words still vouches; with no
    haystack line vouching (a recovered tab part did), the verdict stands."""
    head = "Member, Graduate Medical Education Committee"
    rows = [f"{head} and Residency Curriculum Oversight Subcommittee, {site}, 2015-2018"
            for site in ("NYP", "WCM")]
    twins = record_twins(rows)[rows[0]]
    for page in ([rows[1]], [head, rows[1]]):
        haystack = "\x00".join(_squash(line) for line in page)
        token_sets = [set(_RENDER_TOKEN_RE.findall(_norm(line))) for line in page]
        expected = len(page) > 1
        assert sibling_record_rendered(rows[0], None, haystack, token_sets, [],
                                       twins) is expected
    token_sets = [set(_RENDER_TOKEN_RE.findall(_norm(rows[0])))]
    assert sibling_record_rendered(rows[0], None, "", token_sets, [], twins) is True


def test_non_t_date_led_grid_row_rendered_only_through_its_siblings_is_recovered():
    """The grid written date first, three rows on the page: the fourth is
    recovered (#1230)."""
    rows = [f"{term} Member, {name} Committee, Example State Medical College"
            for name in ("Admissions", "Curriculum") for term in ("2010-2012", "2014-2016")]
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    for row in rows[:3]:
        gen.doc.add_paragraph(row.replace("-", " - "))
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if "Member" in t]
    assert bullets == [rows[3]]


def test_non_t_rows_differing_only_in_a_month_name_rendered_as_numbers_are_kept():
    """Two rows told apart only by a month name, both on the page with the
    month as a number: neither is re-inserted (#1230)."""
    rows = [f"Visiting Lecturer | Example State Medical College | Boston | {month} 2015"
            for month in ("January", "March")]
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    for number in ("01", "03"):
        gen.doc.add_paragraph(
            f"{number}/2015 Visiting Lecturer, Example State Medical College, Boston")
    gen._recover_unrendered_records({"P": [entry]})
    assert not [t for t in _bulleted_texts(gen.doc.paragraphs) if "Lecturer" in t]


def test_non_t_sibling_rows_sharing_a_month_name_are_still_told_apart():
    """Two committees in the same month: the month is no row's own token, so
    the rendered row, which spells it, cannot vouch for the other (#1230)."""
    rows = [f"Member | {name} Committee | Example State Medical College | January 2015"
            for name in ("Admissions", "Curriculum")]
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    gen.doc.add_paragraph(rows[0].replace(" | ", ", "))
    gen._recover_unrendered_records({"P": [entry]})
    bullets = [t for t in _bulleted_texts(gen.doc.paragraphs) if t.startswith("Member")]
    assert bullets == [rows[1].replace(" | ", " \u2014 ")]


def _recovered_terms(rows, out):
    """The bullets the recovery pass re-inserts for P `rows` with `out` on the
    page, one paragraph per item."""
    entry = {"element_idx_start": 88, "taxonomy_code": "P",
             "text": "\n".join(rows), "extracted_fields": {}}
    gen = _generator()
    for text in out:
        gen.doc.add_paragraph(text)
    gen._recover_unrendered_records({"P": [entry]})
    return [t for t in _bulleted_texts(gen.doc.paragraphs) if "Committee" in t]


def _one_committee_two_terms(place="Weill Cornell Medicine"):
    return [f"Member | Admissions Committee | {place} | {term}"
            for term in ("2010-2012", "2014-2016")]


def test_non_t_two_terms_rendered_as_table_rows_without_the_place_are_kept():
    """Both terms on the page in the committee table's form, the place left
    out: each line holds half the row's words, so no line held the share of
    the whole row, and both rows were re-inserted. The line holding the
    row's committee cell and its own years vouches for it (#1230)."""
    rows = _one_committee_two_terms()
    assert not _recovered_terms(rows, [f"Admissions Committee | Member | {term}"
                                       for term in ("2010-2012", "2014-2016")])
    assert _recovered_terms(rows, ["Admissions Committee | Member | 2010-2012"]) == [
        rows[1].replace(" | ", " \u2014 ")]


def test_non_t_two_terms_under_a_place_heading_are_kept():
    """Both terms on the page under the place written once as a heading:
    neither row is re-inserted (#1230)."""
    rows = _one_committee_two_terms("Weill Cornell Medical College")
    assert not _recovered_terms(rows, ["Weill Cornell Medical College",
                                       "Member, Admissions Committee, 2010-2012",
                                       "Member, Admissions Committee, 2014-2016"])


def test_non_t_two_terms_with_the_years_dropped_are_kept():
    """Both terms on the page without their years: no line is either term's
    row, so nothing tells them apart and the verdict stands (#1230)."""
    rows = _one_committee_two_terms()
    assert not _recovered_terms(rows, ["Member, Admissions Committee, Weill Cornell Medicine"] * 2)


@pytest.mark.parametrize("years, recovered", [
    (("2010-12", "2014-16"), True),
    (("'10-'12", "'14-'16"), False),
])
def test_non_t_two_terms_with_shortened_years_are_kept(years, recovered):
    """Both terms on the page with each range's end written in two digits: a
    year counts by its last two digits standing alone, so neither row is
    re-inserted. With one term on the page, the other is recovered only when
    the line spells a year of its sibling whole, as "2010-12" does; "'10-'12"
    is no line positively the sibling's, and keeps the verdict (#1230)."""
    rows = _one_committee_two_terms()
    line = "Member, Admissions Committee, Weill Cornell Medicine, {}"
    assert not _recovered_terms(rows, [line.format(y) for y in years])
    assert _recovered_terms(rows, [line.format(years[0])]) == (
        [rows[1].replace(" | ", " \u2014 ")] if recovered else [])


@pytest.mark.parametrize("open_term", [
    "2014-Present", "2014-current", "2014-ongoing", "2014-now", "2014 to date"])
def test_non_t_a_term_and_an_open_term_are_told_apart(open_term):
    """One committee 2010-2012 and again in an open term ("2014-Present"):
    the open-term word counts like a year, so with either row on the page the
    other is recovered, and with both on the page neither is (#1230)."""
    rows = [f"Member | Admissions Committee | Weill Cornell Medical College | {term}"
            for term in ("2010-2012", open_term)]
    line = "Member, Admissions Committee, Weill Cornell Medical College, {}"
    terms = ("2010-2012", open_term)
    assert not _recovered_terms(rows, [line.format(term) for term in terms])
    assert _recovered_terms(rows, [line.format(terms[0])]) == [rows[1].replace(" | ", " \u2014 ")]
    assert _recovered_terms(rows, [line.format(terms[1])]) == [rows[0].replace(" | ", " \u2014 ")]


@pytest.mark.parametrize("line", [
    "Member, Admissions Committee, Example State Medical College, {}",
    "Committee on Admissions, Member, Example State Med. College, {}",
])
def test_non_t_consecutive_terms_with_shortened_years_are_kept(line):
    """Terms 2010-2012 and 2012-2014 on the page as "2010-12" and "2012-14",
    with a verbatim cell or reworded: the second term's own year is on its
    line only as "14", and the first term's line spells its sibling's 2010
    whole, so without the two-digit form it was re-inserted (#1230)."""
    rows = [f"Member | Admissions Committee | Example State Medical College | {term}"
            for term in ("2010-2012", "2012-2014")]
    assert not _recovered_terms(rows, [line.format(y) for y in ("2010-12", "2012-14")])


def test_non_t_a_year_inside_another_number_does_not_tell_terms_apart():
    """A missing term's year as two digits inside a longer number ("10014")
    or its sibling's year inside one ("20101") is not that year (#1230)."""
    rows = _one_committee_two_terms()
    place = "Member, Admissions Committee, Weill Cornell Medicine, New York, NY 10014"
    assert _recovered_terms(rows, [f"{place}, 2010-2012"]) == [rows[1].replace(" | ", " \u2014 ")]
    room = "Member, Admissions Committee, Weill Cornell Medicine, room 20101"
    assert not _recovered_terms(rows, [room, room])


def test_non_t_a_line_sharing_one_cell_word_does_not_vouch_for_a_term():
    """An unrelated line holding the missing term's years and one word of its
    row ("Member") does not vouch for it: a chunk too short for the overlap
    test is skipped, as `_record_rendered` skips it (#1230)."""
    rows = _one_committee_two_terms()
    assert _recovered_terms(rows, ["Member, Admissions Committee, Weill Cornell Medicine, 2010-2012",
                                   "Member, Faculty Council, 2014-2016"]) == [
        rows[1].replace(" | ", " \u2014 ")]


def test_non_t_row_whose_years_hold_its_siblings_is_kept_with_a_short_range():
    """'2010' and '2010-2012' rows, both on the page, the range written
    '2010-12': the range's row is not re-inserted (#1230)."""
    rows = [f"Member | Admissions Committee | Weill Cornell Medicine | {term}"
            for term in ("2010", "2010-2012")]
    assert not _recovered_terms(rows, [f"Member, Admissions Committee, Weill Cornell Medicine, {y}"
                                       for y in ("2010", "2010-12")])


def test_sibling_record_rendered_without_siblings_is_record_rendered():
    """A line with no sibling rows keeps `_record_rendered`'s verdict, and a
    sibling row on the page verbatim is rendered (#1230)."""
    rows = _two_by_two_grid()
    out = [rows[0].replace("2010-2012", "2014-2016")]
    letters = [set(_RENDER_TOKEN_RE.findall(_norm(line))) for line in out]
    records = [_record_tokens(line) for line in out]
    haystack = "\x00".join(_squash(line) for line in out)
    siblings = record_siblings(rows)
    assert sibling_record_rendered(rows[0], None, haystack, letters, records) is True
    assert sibling_record_rendered(rows[0], siblings[rows[0]], haystack, letters,
                                   records) is False
    assert sibling_record_rendered(rows[1], siblings[rows[1]], haystack, letters,
                                   records) is True


def test_t_sibling_rows_differing_only_in_one_name_are_all_recovered():
    """Committee rows with the same years and place, told apart by one word:
    the complete row on the page must not vouch for its siblings (#1230)."""
    names = ("Zeolite", "Marigold", "Quartzite", "Nasturtium", "Obsidian", "Hyacinth")
    rows = [f"Member | {n} Committee | Example Department of Synthetic Geology | Example College of Sciences | Example University | 1987-1992"
            for n in names]
    entry = _t_entry()
    entry["text"] = "\n".join(rows)
    gen = _generator()
    gen.doc.add_paragraph(rows[0])
    gen._recover_unrendered_records({"T": [entry]})
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert all(row in bullets for row in rows[1:])
    assert rows[0] not in bullets


def test_t_row_that_is_its_own_non_t_entry_is_not_recovered():
    """A T entry that repeats a table whose rows are entries of their own: the
    row renders from its entry, reformatted, so recovering it duplicates it."""
    row = _T_ROWS[1]
    own = dict(_t_entry(), taxonomy_code="N", text=row)
    entry = _t_entry()
    entry["text"] = "\n".join(_T_ROWS)
    gen = _generator()
    gen.doc.add_paragraph("Reformatted Zeolite")  # the entry rendered, differently
    gen._recover_unrendered_records({"T": [entry], "N": [own]})
    bullets = _bulleted_texts(gen.doc.paragraphs)
    assert row not in bullets
    assert _T_ROWS[2] in bullets


# --- #1205: supplementary prose as tracked-deleted sub-points -----------------
# The generator glue for stage6/supplementary.py: the flag, the pass's place
# before the overflow routing and this recovery, and both skipping what a
# sub-point carries. The sub-point module itself is test_stage6_supplementary.py.

_DUTY_PROSE = ("Responsible for coordinating the departmental curriculum committee "
               "and supervising graduate trainees in the example programme")


def _duty_appointment():
    """A D1 row whose source text ends in duty prose no field holds, split
    over two record-shaped lines the way SQMWHM 29's was (YUYVIG)."""
    return {
        "element_idx_start": 29,
        "taxonomy_code": "D1",
        "hierarchy": ["PROFESSIONAL EXPERIENCE"],
        "text": ("July 2010 - present\tAssociate Professor and Program Director, Department of "
                 f"Example Studies, Norvale University\nMain Annex, ZQ.\t{_DUTY_PROSE}"),
        "extraction_coverage": {"extraction_coverage_percent": 21.0},
        "extracted_fields": {
            "title": "Associate Professor and Program Director",
            "institution": "Department of Example Studies, Norvale University",
            "start_date": "2010-07",
            "end_date": "present",
        },
    }


def _subpoint_rows(gen):
    author = "CViche: source text with no template field"
    return [tr for tr in gen.doc.element.body.iter(qn("w:tr"))
            if (d := tr.find(f"{qn('w:trPr')}/{qn('w:del')}")) is not None
            and d.get(qn("w:author")) == author]


def _render_duty_appointment(**kwargs):
    gen = _generator(**kwargs)
    entry = _duty_appointment()
    gen._fill_positions({"D1": [entry]})
    gen._add_supplementary_subpoints({"D1": [entry]})
    return gen, entry


def test_supplementary_subpoints_default_off():
    assert WCMTemplateGenerator(verbose=False).supplementary_subpoints is False


def test_run_stage6_passes_the_subpoint_flag_to_the_generator(monkeypatch):
    from unified_pipeline import stage_6_word_template as stage6

    seen = {}

    class _FakeGenerator:
        def __init__(self, **kwargs):
            seen.update(kwargs)

        def generate(self, *args, **kwargs):
            return "unused.docx"

    monkeypatch.setattr(stage6, "WCMTemplateGenerator", _FakeGenerator)
    stage6.run_stage6("unused.json")
    assert seen["supplementary_subpoints"] is True  # on by default since 2026-10-09
    stage6.run_stage6("unused.json", supplementary_subpoints=False)
    assert seen["supplementary_subpoints"] is False


def test_duty_prose_is_a_deleted_row_under_its_appointment():
    gen, _entry = _render_duty_appointment(supplementary_subpoints=True)

    rows = _subpoint_rows(gen)
    assert len(rows) == 1
    above = rows[0].getprevious()
    assert "Associate Professor and Program Director" in "".join(
        t.text for t in above.iter(qn("w:t")))
    assert "".join(t.text for t in rows[0].iter(qn("w:delText"))) == _DUTY_PROSE
    assert gen.stats["supplementary_subpoints"] == 1
    # The recovery reads the prose, not the sub-point's line: it stands
    # for a line, of any entry, that is no record of its own and that it
    # is part of.
    assert gen._subpoint_prose == [_DUTY_PROSE] and gen._subpoint_lines == []


@pytest.mark.parametrize("kwargs", [{}, {"supplementary_subpoints": True, "emit_track_changes": False}])
def test_no_sub_point_with_the_flag_off_or_track_changes_off(kwargs):
    gen, _entry = _render_duty_appointment(**kwargs)
    assert _subpoint_rows(gen) == []
    assert "supplementary_subpoints" not in gen.stats
    assert gen._subpoint_lines == [] and gen._subpoint_entry_ids == set()


@pytest.mark.parametrize("flag, queued", [(True, 0), (False, 1)])
def test_the_overflow_pass_skips_an_entry_whose_prose_is_a_sub_point(flag, queued):
    gen, entry = _render_duty_appointment(supplementary_subpoints=flag)
    cell_para = gen.doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    gen._overflow_entries = [(entry, cell_para, "D1")]

    gen._route_overflow_entries()

    assert len(gen._appendix_pending) == queued


@pytest.mark.parametrize("flag, recovered", [(True, 0), (False, 1)])
def test_the_recovery_reads_a_sub_point_as_its_entry_rendered(flag, recovered):
    gen, entry = _render_duty_appointment(supplementary_subpoints=flag)

    gen._recover_unrendered_records({"D1": [entry]})

    assert gen.stats["unrendered_records_recovered"] == recovered


def test_the_recovery_reads_no_sub_point_as_a_record_of_its_own():
    """#1205 review: a record of its own beside the entry's
    (`supplementary.is_sub_record`) is re-emitted though a written sub-point
    is part of it. The planner leaves such a line to the recovery by its
    verdict before the overflow pass, which the recovery's own can differ
    from; read as rendered by the sub-point, the record would be in neither
    document."""
    entry = _duty_appointment_with(_DUTY_SUB_RECORD)
    recovered = []
    for prose in ((), ("Chair, Departmental Curriculum Committee",)):
        gen = _generator(supplementary_subpoints=True)
        gen._fill_positions({"D1": [entry]})
        gen._subpoint_prose = list(prose)
        gen._recover_unrendered_records({"D1": [entry]})
        recovered.append(gen.stats["unrendered_records_recovered"])
    assert recovered[0] == recovered[1] == 2


_TITLE_ROLE = "Visiting Fellow in Example Science"
_TITLE_DUTIES = "to coordinate the example curriculum committee"
#: The source states the duties at more length than stage 4 kept in `title`,
#: so the text's duty paragraph is not mostly field values: only the title
#: row's sub-point in the pass's haystack keeps it from being planned again.
_TEXT_DUTIES = (f"{_TITLE_DUTIES}, supervise graduate trainees, organise regional "
                "seminars, maintain programme records and evaluate outcomes")


def _title_duty_appointment():
    """A D3 row whose `title` holds duties (#1641's split), and whose text
    holds them too, as DYLJXC 661/668's did (YUYVIG)."""
    return {
        "element_idx_start": 661,
        "taxonomy_code": "D3",
        "hierarchy": ["PROFESSIONAL EXPERIENCE"],
        "text": f"2013-2015\t{_TITLE_ROLE}, Example Policy Office\t{_TEXT_DUTIES}",
        "extracted_fields": {
            "title": f"{_TITLE_ROLE}, Example Policy Office, {_TITLE_DUTIES}",
            "organization": "Example Institute for Science",
            "start_date": "2013", "end_date": "2015",
        },
    }


def test_title_duties_are_offered_once_and_read_as_rendered():
    """#1205: the duties the positions writer made a sub-point join the pass's
    haystack, so they are not planned again from the entry's text, and the
    #221 recovery reads them as the entry rendered."""
    gen = _generator(supplementary_subpoints=True)
    entry = _title_duty_appointment()
    gen._fill_positions({"D3": [entry]})
    gen._add_supplementary_subpoints({"D3": [entry]})

    rows = _subpoint_rows(gen)
    texts = ["".join(t.text for t in row.iter(qn("w:delText"))) for row in rows]
    assert texts[0].endswith(_TITLE_DUTIES)
    # No second sub-point repeats the duties the title row already offers.
    assert [t for t in texts[1:] if _TITLE_DUTIES in t] == []
    assert gen.stats["supplementary_subpoints"] == len(rows)
    gen._recover_unrendered_records({"D3": [entry]})
    assert gen.stats["unrendered_records_recovered"] == 0


def test_the_recovery_reads_title_duties_as_a_row_of_them_rendered():
    """#1205 review: the recovery reads the title duties' sub-point
    (`_subpoint_lines`) with the document, so a row of the entry's text
    stating the duties is read as rendered and not re-emitted. Coverage for
    the recovery's view of the document, which this round moved into
    `_recovery_view`."""
    gen = _generator(supplementary_subpoints=True)
    entry = _title_duty_appointment()
    entry["text"] += f"\nSeconded to the regional annex.\t{_TITLE_DUTIES}, as the role states"
    assert len(recovery_candidates("D3", entry)) == 2
    gen._fill_positions({"D3": [entry]})
    gen._add_supplementary_subpoints({"D3": [entry]})

    gen._recover_unrendered_records({"D3": [entry]})

    assert gen.stats["unrendered_records_recovered"] == 0


def test_sub_point_anchors_stop_at_the_appendix_heading(monkeypatch):
    """The Appendix repeats entry text in bullets: no place for a sub-point."""
    from unified_pipeline.stage6 import supplementary

    gen = _generator(supplementary_subpoints=True)
    gen._write_appendix_header(1)
    appendix = gen.doc.paragraphs[gen._find_header_paragraph("T. APPENDIX")]._p
    seen = []
    real = supplementary.anchor_candidates

    def spy(doc, stop_before=None):
        seen.append(stop_before)
        return real(doc, stop_before)

    monkeypatch.setattr(supplementary, "anchor_candidates", spy)
    gen._add_supplementary_subpoints({"D1": [_duty_appointment()]})
    assert seen == [appendix]


# --- #1205: Q1/Q2 extramural service, the next code family -------------------
# Both render as table rows in Q's tables, so the prose is a deleted row under
# the entry's own row, spanning that table's columns.

_SERVICE_PROSE = ("Advised the society on regional workforce planning and drafted "
                  "the annual training standards for member programmes")


def _board_service():
    """A Q2 row whose source line goes on past its committee, role and
    organization into the committee's remit, as YUYVIG ECXGAT 2129's did."""
    return {
        "element_idx_start": 2129,
        "taxonomy_code": "Q2",
        "text": f"2015-2018\tMember, Example Advisory Committee, Example Society\t{_SERVICE_PROSE}",
        "extracted_fields": {"committee_name": "Example Advisory Committee", "role": "Member",
                             "organization": "Example Society",
                             "start_date": "2015", "end_date": "2018"},
    }


def _extramural_leadership():
    """A Q1 row whose source holds the organization's mission, as EBYSBC
    RGUNJV 3057's did."""
    return {
        "element_idx_start": 3057,
        "taxonomy_code": "Q1",
        "text": f"2009-2014\tExample Foundation, President\t{_SERVICE_PROSE}",
        "extracted_fields": {"organization": "Example Foundation", "role": "President",
                             "start_date": "2009", "end_date": "2014"},
    }


def _texts_of(body):
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in body.iter(qn("w:p"))]


def _render_service(code, entry, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_service({code: [entry]})
    gen._add_supplementary_subpoints({code: [entry]})
    return gen


@pytest.mark.parametrize("code, entry, anchor_text, columns", [
    ("Q2", _board_service, "Example Advisory Committee", 4),
    ("Q1", _extramural_leadership, "Example Foundation", 3),
])
def test_service_prose_is_a_deleted_row_under_its_own_row(code, entry, anchor_text, columns):
    gen = _render_service(code, entry(), supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert anchor_text in "".join(t.text for t in above.iter(qn("w:t")))
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _SERVICE_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == str(columns)
    assert gen.stats["supplementary_subpoints"] == 1


@pytest.mark.parametrize("code, entry", [("Q2", _board_service), ("Q1", _extramural_leadership)])
def test_service_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it(code, entry):
    off = _render_service(code, entry(), supplementary_subpoints=False)
    on = _render_service(code, entry(), supplementary_subpoints=True)

    # Reject All: the deleted text is the reader's again.
    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _SERVICE_PROSE in rejected
    assert _SERVICE_PROSE not in "".join(_texts_of(off.doc.element.body))
    # Accept All: the deleted row goes, and the document is the flag-off one,
    # empty paragraphs included.
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)


def test_service_prose_already_in_its_cell_is_not_offered_again():
    """A Q2 row with no committee field shows the entry's whole text in its
    cell (`_fill_service_boards`' fallback): nothing is left to offer."""
    from unified_pipeline.stage6 import supplementary

    entry = _board_service()
    entry["extracted_fields"] = {"start_date": "2015", "end_date": "2018"}
    gen = _render_service("Q2", entry, supplementary_subpoints=True)
    assert _subpoint_rows(gen) == []
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry, "Q2", haystack) == []


def test_journal_reviewing_q4d_stays_excluded():
    entry = dict(_board_service(), taxonomy_code="Q4D")
    entry["extracted_fields"] = {"journal_name": "Example Journal of Training", "role": "Reviewer"}
    gen = _render_service("Q4D", entry, supplementary_subpoints=True)
    assert _subpoint_rows(gen) == []


def _journal_reviewer_coded_q2():
    """A Q2 entry `_route_q2_entries` reroutes to Journal Reviewing: its
    rendered line is a Q4D row, but its own code stays Q2."""
    return {
        "element_idx_start": 1514,
        "taxonomy_code": "Q2",
        "text": f"2015-2018\tReviewer, Example Journal of Training\t{_SERVICE_PROSE}",
        "extracted_fields": {"role": "Reviewer", "organization": "Example Journal of Training",
                             "start_date": "2015", "end_date": "2018"},
    }


def test_a_q2_entry_rendered_as_journal_reviewing_gets_no_sub_point():
    """The 2026-10-02 decision excludes journal reviewing. A Q2 entry the
    section writes into the Journal Reviewing table is that, whatever its
    code says, so it is excluded by where it rendered."""
    from unified_pipeline.stage6.sections.service import _route_q2_entries

    entry = _journal_reviewer_coded_q2()
    assert len(_route_q2_entries([entry])[0]) == 1
    gen = _render_service("Q2", entry, supplementary_subpoints=True)
    assert "Example Journal of Training" in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == []
    assert _SERVICE_PROSE not in "".join(
        t.text or "" for t in gen.doc.element.body.iter(qn("w:delText")))


# --- #1205: H honors, the next code family -----------------------------------
# The honors section writes every H entry as a row of its three-column table
# (`_add_honors_row`), so the prose is a deleted row under the award's row.

_HONOR_PROSE = ("Awarded annually to an early career investigator whose published "
                "research advanced clinical understanding of respiratory infections")


def _honor():
    """An H row whose source line goes on past the award, its granting body
    and year into the award's criteria, as EBYSBC QITQWH 39-53's did."""
    return {
        "element_idx_start": 39,
        "taxonomy_code": "H",
        "text": f"2019\tYoung Investigator Award, Example Thoracic Society\t{_HONOR_PROSE}",
        "extracted_fields": {"award_name": "Young Investigator Award",
                             "granting_body": "Example Thoracic Society", "date": "2019"},
    }


def _render_honors(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_honors(entries)
    gen._add_supplementary_subpoints({"H": entries})
    return gen


def test_honor_prose_is_a_deleted_row_under_the_award_row():
    gen = _render_honors([_honor()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = "".join(t.text for t in row.getprevious().iter(qn("w:t")))
    assert "Young Investigator Award" in above
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _HONOR_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "3"
    assert gen.stats["supplementary_subpoints"] == 1


def test_honor_prose_goes_under_its_own_award_among_several():
    other = {"element_idx_start": 41, "taxonomy_code": "H",
             "text": "2015\tOutstanding Mentor Award, Example Medical College",
             "extracted_fields": {"award_name": "Outstanding Mentor Award",
                                  "granting_body": "Example Medical College", "date": "2015"}}
    gen = _render_honors([other, _honor()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = "".join(t.text for t in row.getprevious().iter(qn("w:t")))
    assert "Young Investigator Award" in above and "Mentor" not in above
    below = row.getnext()
    assert below is not None and "Outstanding Mentor Award" in "".join(
        t.text for t in below.iter(qn("w:t")))


def test_honor_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_honors([_honor()], supplementary_subpoints=False)
    on = _render_honors([_honor()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _HONOR_PROSE in rejected
    assert _HONOR_PROSE not in "".join(_texts_of(off.doc.element.body))
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)


def test_honor_prose_already_in_its_award_cell_is_not_offered_again():
    """With no award field, the award cell shows the entry's whole text
    (`parse_honor_entry`'s fallback): nothing is left to offer."""
    from unified_pipeline.stage6 import supplementary

    entry = _honor()
    entry["extracted_fields"] = {"date": "2019"}
    gen = _render_honors([entry], supplementary_subpoints=True)
    assert _HONOR_PROSE in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == []
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry, "H", haystack) == []


def _rerouted(entry, expected_code, heading):
    """`entry` flagged by stage 3b's hierarchy check, at low confidence, under
    a heading that expects `expected_code`: `_correct_mismatch_if_needed`
    moves it across families."""
    return dict(entry, taxonomy_confidence=0.5, hierarchy_mismatch_flag=True,
                hierarchy_mismatch_detail={"expected_codes": [expected_code],
                                           "hierarchy": [heading]})


def _render_grouped(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_honors(grouped.get("H", []))
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def test_an_honor_rerouted_out_of_h_follows_its_destination_and_gets_no_sub_point():
    """An H entry rendered as another code is that code's record: N1 is
    excluded, so neither its own row nor another award's carries its prose."""
    other = {"element_idx_start": 41, "taxonomy_code": "H",
             "text": "2015\tYoung Investigator Award, Example Thoracic Society",
             "extracted_fields": {"award_name": "Young Investigator Award",
                                  "granting_body": "Example Thoracic Society", "date": "2015"}}
    moved = _rerouted(_honor(), "N1", "MENTORING PROGRAMS")
    gen, grouped = _render_grouped([other, moved], supplementary_subpoints=True)

    assert grouped["N1"] == [moved] and grouped["H"] == [other]
    assert _subpoint_rows(gen) == []
    assert _HONOR_PROSE not in "".join(
        t.text or "" for t in gen.doc.element.body.iter(qn("w:delText")))


def test_a_record_rerouted_into_h_follows_the_honors_rule():
    """A society fellowship stage 3b coded I under an honors heading renders
    as an award: its prose is a deleted row under that award's row."""
    fellowship = dict(_honor(), taxonomy_code="I")
    gen, grouped = _render_grouped([_rerouted(fellowship, "H", "HONORS")],
                                   supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["H"]] == [39]
    (row,) = _subpoint_rows(gen)
    assert "Young Investigator Award" in "".join(t.text for t in row.getprevious().iter(qn("w:t")))
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _HONOR_PROSE


def test_an_honor_whose_award_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass dropped the award's name, so the honors writer
    renders no row for what is left. Its prose is not offered either, under
    another award of the same granting body or anywhere else."""
    other = {"element_idx_start": 41, "taxonomy_code": "H",
             "text": "2015\tOutstanding Mentor Award, Example Thoracic Society",
             "extracted_fields": {"award_name": "Outstanding Mentor Award",
                                  "granting_body": "Example Thoracic Society", "date": "2015"}}
    withheld = _honor()
    del withheld["extracted_fields"]["award_name"]
    withheld["_pii_dropped_fields"] = ["award_name"]
    withheld["_pii_withheld"] = True  # the pass sets both
    gen = _render_honors([other, withheld], supplementary_subpoints=True)

    assert "Young Investigator Award" not in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == []
    assert _HONOR_PROSE not in "".join(
        t.text or "" for t in gen.doc.element.body.iter(qn("w:delText")))


# --- #1205: N3A/N3B mentees, the next code family ----------------------------
# The mentoring section writes each named mentee as a six-row "Label:" table of
# its own, so the prose is a deleted row at the end of that mentee's table.

_MENTEE_PROSE = ("Received the departmental travel award and presented the thesis findings "
                 "at the regional pharmacology meeting")


def _mentee(code="N3B", idx=1477, name="Ana Li", end="2017", prose=_MENTEE_PROSE):
    """A mentee whose source line goes on past the name, level and years into
    what the mentee won and presented, as EBYSBC MUHLLD 226's did. The name
    has no five-letter word, so only the level's words can anchor it: another
    postdoc's Site/Position row holds all of them (YUYVIG XNWSZN 188)."""
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": f"{name} (Postdoctoral Associate, 2014-{end})" + (f"\t{prose}" if prose else ""),
        "extracted_fields": {"mentee_name": name, "mentee_level": "Postdoctoral Associate",
                             "start_date": "2014", "end_date": end},
    }


def _other_postdoc(code="N3B", end="2017"):
    """A postdoc rendered before `_mentee()`'s, with no prose of its own."""
    return _mentee(code, idx=1470, name="Casey Sample", end=end, prose="")


def _render_mentoring(entries_by_code, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_mentoring(entries_by_code)
    gen._add_supplementary_subpoints(entries_by_code)
    return gen


def _mentee_named_by(row):
    """The Name: value of the mentee table `row` is a row of."""
    first_row_cells = row.getparent().find(qn("w:tr")).findall(qn("w:tc"))
    return "".join(t.text or "" for t in first_row_cells[1].iter(qn("w:t")))


def _mentee_heading_above(row):
    """The mentee heading the table holding `row` sits under."""
    for element in row.getparent().itersiblings(preceding=True):
        text = "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()
        if text in ("Current Mentees:", "Past Mentees:"):
            return text
    return None


def _mentee_names(gen):
    """The Name: value of every mentee table, in document order."""
    return [_mentee_named_by(tbl.find(qn("w:tr"))) for tbl in gen.doc.element.body.iter(qn("w:tbl"))
            if "".join(t.text or "" for t in tbl.find(qn("w:tr")).find(qn("w:tc")).iter(qn("w:t")))
            == "Name:"]


def _row_texts(body):
    return [["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.findall(qn("w:tc"))]
            for tr in body.iter(qn("w:tr"))]


def _deleted_texts(gen):
    return "".join(t.text or "" for t in gen.doc.element.body.iter(qn("w:delText")))


@pytest.mark.parametrize("code, end, heading", [("N3B", "2017", "Past Mentees:"),
                                                ("N3A", "present", "Current Mentees:")])
def test_mentee_prose_is_a_deleted_row_at_the_end_of_its_table(code, end, heading):
    gen = _render_mentoring({code: [_mentee(code, end=end)]}, supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert row.getnext() is None
    assert _mentee_named_by(row) == "Ana Li"
    assert _mentee_heading_above(row) == heading
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _MENTEE_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "2"
    assert gen.stats["supplementary_subpoints"] == 1


def test_mentee_prose_goes_under_its_own_table_among_several():
    """The postdoc before it shares its Site/Position row word for word; the
    prose still goes to the end of the table whose Name: row is its mentee."""
    gen = _render_mentoring({"N3B": [_other_postdoc(), _mentee()]}, supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _mentee_named_by(row) == "Ana Li" and row.getnext() is None
    assert _mentee_names(gen) == ["Casey Sample", "Ana Li"]


@pytest.mark.parametrize("code, end, other_end, heading", [
    ("N3B", "present", "2017", "Current Mentees:"),   # the open range makes it current
    ("N3A", "2017", "present", "Past Mentees:"),      # an ended period makes it past
])
def test_a_mentee_the_section_files_under_the_other_heading_keeps_its_sub_point(
        code, end, other_end, heading):
    gen = _render_mentoring({code: [_other_postdoc(code, other_end), _mentee(code, end=end)]},
                            supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _mentee_named_by(row) == "Ana Li"
    assert _mentee_heading_above(row) == heading
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _MENTEE_PROSE


def test_mentee_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_mentoring({"N3B": [_other_postdoc(), _mentee()]}, supplementary_subpoints=False)
    on = _render_mentoring({"N3B": [_other_postdoc(), _mentee()]}, supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _MENTEE_PROSE in rejected
    assert _MENTEE_PROSE not in "".join(_texts_of(off.doc.element.body))
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def test_mentee_prose_already_in_its_table_is_not_offered_again():
    """Stage 4 put the prose in `awards`, which the Project/Accomplishments row
    prints, though no N3B rendered field (`_RENDERED_FIELDS`) names it: the
    document holds it, so nothing is left to offer."""
    from unified_pipeline.stage6 import supplementary

    entry = _mentee()
    entry["extracted_fields"]["awards"] = _MENTEE_PROSE
    gen = _render_mentoring({"N3B": [entry]}, supplementary_subpoints=True)
    assert _MENTEE_PROSE in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == []
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry, "N3B", haystack) == []


def test_a_mentee_naming_no_one_renders_its_whole_line_and_gets_no_sub_point():
    """An aggregate N3B line ("<n> postdocs ...") is a line, not a table."""
    entry = _mentee()
    del entry["extracted_fields"]["mentee_name"]
    gen = _render_mentoring({"N3B": [_other_postdoc(), entry]}, supplementary_subpoints=True)

    assert any(_MENTEE_PROSE in text for text in _texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == [] and _MENTEE_PROSE not in _deleted_texts(gen)


def _render_grouped_mentoring(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_mentoring(grouped)
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def test_a_mentee_rerouted_to_n1_follows_its_destination_and_gets_no_sub_point():
    """N1 is excluded: the entry renders as an N1 line of its text, and no
    mentee table carries its prose, not even the postdoc's that shares its
    level."""
    moved = _rerouted(_mentee(), "N1", "Leadership and mentoring in programs")
    gen, grouped = _render_grouped_mentoring([_other_postdoc(), moved],
                                             supplementary_subpoints=True)

    assert grouped["N1"] == [moved] and [e["element_idx_start"] for e in grouped["N3B"]] == [1470]
    assert any(_MENTEE_PROSE in text for text in _texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == [] and _MENTEE_PROSE not in _deleted_texts(gen)


def test_a_record_rerouted_into_n3b_follows_the_mentee_rule():
    """A trainee stage 3b coded K2 under a past-mentees heading renders as a
    mentee table: its prose is a deleted row at that table's end."""
    trainee = dict(_mentee(), taxonomy_code="K2")
    gen, grouped = _render_grouped_mentoring([_rerouted(trainee, "N3B", "PAST MENTEES")],
                                             supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["N3B"]] == [1477]
    (row,) = _subpoint_rows(gen)
    assert _mentee_named_by(row) == "Ana Li" and row.getnext() is None
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _MENTEE_PROSE


def test_a_rerouted_entrys_placement_records_the_code_it_rendered_as():
    """The sidecar's record of the sub-point names the mentee code the K2
    trainee rendered as, and the entry's own index (the doctor's key)."""
    from unified_pipeline.stage6.supplementary import SubpointPlacement

    trainee = dict(_mentee(), taxonomy_code="K2")
    gen, _ = _render_grouped_mentoring([_rerouted(trainee, "N3B", "PAST MENTEES")],
                                       supplementary_subpoints=True)

    assert gen._subpoint_placements == [SubpointPlacement(1477, "N3B", (_MENTEE_PROSE,))]


def test_the_render_sidecar_records_each_sub_point_with_its_entry_and_code(tmp_path, monkeypatch):
    """Through `run_stage6`: one placement per sub-point the docx holds, under
    the entry it was written for -- a mentee's planned row and the duties the
    positions writer split off a title -- so the doctor reads each as its own
    entry's only (EBYSBC VGHNZD: one mentee's sub-point silenced another
    mentee's WARN). The postdoc with no prose has none."""
    from unified_pipeline.doctor.shared import docx_subpoint_lines
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm",
                        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    payload = {"document_uid": "T1205P", "cv_owner": {"full_name": "Quenby Marblegate"},
               "entries": [_other_postdoc(), _mentee(), _title_duty_appointment()]}
    input_path = tmp_path / "T1205P_fields.json"
    input_path.write_text(json.dumps(payload))

    out = run_stage6(str(input_path), str(tmp_path / "out.docx"), verbose=False)

    placements = json.loads((tmp_path / "T1205P_render_warnings.json").read_text())[
        SUBPOINT_PLACEMENTS_KEY]
    assert sorted((p["element_idx_start"], p["code"]) for p in placements) == [
        (661, "D3"), (1477, "N3B")]
    by_entry = {p["element_idx_start"]: p["paragraphs"] for p in placements}
    assert by_entry == {661: [_TITLE_DUTIES], 1477: [_MENTEE_PROSE]}
    assert sorted(text for p in placements for text in p["paragraphs"]) == sorted(
        docx_subpoint_lines(Document(out)))


def test_a_generator_rendering_twice_records_only_this_renders_placements(tmp_path, monkeypatch):
    """The placements reset per render, as the section failures do: a second
    CV with no sub-point has none in its sidecar."""
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm",
                        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    gen = WCMTemplateGenerator(verbose=False, supplementary_subpoints=True)
    placements = []
    for uid, entries in (("T1205R1", [_mentee()]), ("T1205R2", [_other_postdoc()])):
        input_path = tmp_path / f"{uid}_fields.json"
        input_path.write_text(json.dumps({"document_uid": uid, "entries": entries}))
        gen.generate(str(input_path), str(tmp_path / f"{uid}.docx"), discover_original_doc=False)
        sidecar = json.loads((tmp_path / f"{uid}_render_warnings.json").read_text())
        placements.append([p["element_idx_start"] for p in sidecar[SUBPOINT_PLACEMENTS_KEY]])
    assert placements == [[1477], []]


def test_a_mentee_whose_field_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass dropped one of the mentee's fields. What is left
    still describes what was withheld, so none of it is offered, under its
    own table or the postdoc's."""
    withheld = _mentee()
    withheld["_pii_dropped_fields"] = ["site_position"]
    withheld["_pii_withheld"] = True  # the pass sets both
    gen = _render_mentoring({"N3B": [_other_postdoc(), withheld]}, supplementary_subpoints=True)

    assert "Ana Li" in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == [] and _MENTEE_PROSE not in _deleted_texts(gen)


def test_a_mentee_named_only_inside_an_award_title_gets_no_sub_point():
    """#1702: the section refuses a mentee whose name is an award's and lists
    the entry in the Appendix. No table names that person, so its prose goes
    under no table, not even the postdoc's that shares its level."""
    refused = _mentee(idx=1480, name="Jordan Example")
    refused["text"] = f"Recipient of the Jordan Example Award, Postdoctoral Associate\t{_MENTEE_PROSE}"
    gen = _render_mentoring({"N3B": [_other_postdoc(), refused]}, supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in gen._declined_grant_entries] == [1480]
    assert "Jordan Example" not in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == [] and _MENTEE_PROSE not in _deleted_texts(gen)


def test_a_one_segment_mentee_line_is_offered_whole_its_identity_included():
    """A fragment is a whole tab segment, so a line giving the mentee's name,
    level and years before its prose is offered as the CV wrote it: Reject
    All prints those three a second time, beside the table's own rows."""
    line = f"Ana Li, Postdoctoral Associate, 2014-2017, {_MENTEE_PROSE}"
    gen = _render_mentoring({"N3B": [dict(_mentee(), text=line)]}, supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _mentee_named_by(row) == "Ana Li"
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == line


def _mentee_line(text, *periods):
    """One source line stage 4 returned as a record per (name, start, end),
    the entry's own fields being the last record's. Fan-out gives every
    earlier record a line of its own and the last one the whole line."""
    records = [{"mentee_name": name, "mentee_level": "Postdoctoral Associate",
                "start_date": start, "end_date": end} for name, start, end in periods]
    return {"element_idx_start": 1490, "taxonomy_code": "N3B", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_LISTED_MENTEES = (("Ana Li", "2014", "2017"), ("Casey Sample", "2012", "2015"),
                   ("Jordan Example", "2010", "2013"))


def test_a_line_listing_several_mentees_is_offered_under_none_of_their_tables():
    """YUYVIG XNWSZN 188: twelve postdocs on one line. Each renders as a table
    of its own, so the line is no one mentee's prose: under the last one's
    table, Reject All would print every other name and period again."""
    line = "Postdoctoral associates mentored in the laboratory: " + ", ".join(
        f"{name} ({start}-{end})" for name, start, end in _LISTED_MENTEES)
    gen, grouped = _render_grouped_mentoring([_mentee_line(line, *_LISTED_MENTEES)],
                                             supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["N3B"][-1] and grouped["N3B"][-1]["text"] == line
    assert sorted(_mentee_names(gen)) == sorted(name for name, _, _ in _LISTED_MENTEES)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


_ANA = "Ana Li, Postdoctoral Associate, 2014-2017"
_CASEY = "Casey Sample, Postdoctoral Associate, 2012-2015"


@pytest.mark.parametrize("segments, owner", [
    # The segments pair with the records: each prose goes under its own mentee.
    ([f"{_ANA}, {_MENTEE_PROSE}", _CASEY], "Ana Li"),
    ([_ANA, f"{_CASEY}, {_MENTEE_PROSE}"], "Casey Sample"),
    # They do not: whose prose it is cannot be told, so it goes under neither.
    ([_ANA, _MENTEE_PROSE, _CASEY], None),
])
def test_prose_on_a_line_of_several_mentees_goes_under_its_own_or_none(segments, owner):
    entry = _mentee_line("\t".join(segments), ("Ana Li", "2014", "2017"),
                         ("Casey Sample", "2012", "2015"))
    gen, _ = _render_grouped_mentoring([entry], supplementary_subpoints=True)

    assert sorted(_mentee_names(gen)) == ["Ana Li", "Casey Sample"]
    assert [_mentee_named_by(row) for row in _subpoint_rows(gen)] == ([owner] if owner else [])
    assert _deleted_texts(gen).count(_MENTEE_PROSE) == (1 if owner else 0)


def test_one_mentees_two_periods_on_one_line_keep_its_prose():
    """YUYVIG ECXGAT 1479: stage 4 made one record per period of the same
    mentee, so the whole line is that mentee's, and its prose goes once,
    under the first of the two tables. Another line's mentee is no record
    of this one."""
    line = f"Ana Li, Postdoctoral Associate, 2008-2010 and 2012-2014, {_MENTEE_PROSE}"
    entry = _mentee_line(line, ("Ana Li", "2008", "2010"), ("Ana Li", "2012", "2014"))
    gen, _ = _render_grouped_mentoring([_other_postdoc(), entry], supplementary_subpoints=True)

    assert sorted(_mentee_names(gen)) == ["Ana Li", "Ana Li", "Casey Sample"]
    (row,) = _subpoint_rows(gen)
    assert _mentee_named_by(row) == "Ana Li" and _deleted_texts(gen) == line


# --- #1205: B1 degrees, the next code family ----------------------------------
# The education section writes every B1 entry as a row of its four-column
# table, so the prose is a deleted row under the degree's own row: the row the
# education writer recorded for that entry (the row-receipt ledger).

_THESIS = ("Thesis: Seasonal variation of nitrogen fixation in coastal wetland soils "
           "under experimental warming")
_DEGREE_FIELD = "Example Marine Studies"


def _degree(idx=19, degree="M.S.", year="2001", prose=_THESIS, code="B1", *,
            discipline=_DEGREE_FIELD, institution="Norvale University", **fields):
    """A degree whose source line goes on past the degree, field, institution
    and year into its thesis, as YUYVIG FMIGLR 19's did. "M.S." and "Ph.D."
    hold no five-letter word, so another degree of the same field and
    institution shares every distinctive word of its row (`_doctorate`). A
    value passed as None is in neither the line nor the fields."""
    named = {"degree": degree, "discipline": discipline, "institution": institution, "year": year}
    line = ", ".join(value for value in named.values() if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _doctorate():
    """A later degree of the same field and institution, with no prose: it
    renders first, reverse chronologically."""
    return _degree(idx=17, degree="Ph.D.", year="2005", prose="")


def _render_education(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_education(entries)
    gen._add_supplementary_subpoints({"B1": entries})
    return gen


def _row_text(tr):
    return "".join(t.text or "" for t in tr.iter(qn("w:t")))


def _degree_table_row():
    """The same degree read from a source table: stage 2 joins the row's
    cells with pipes, and its position names the row ("96.3")."""
    entry = _degree()
    entry["element_idx_start"] = "96.3"
    entry["text"] = entry["text"].replace(", ", " | ").replace("\t", " | ")
    return entry


@pytest.mark.parametrize("entry", [_degree, _degree_table_row])
def test_degree_prose_is_a_deleted_row_under_the_degree_row(entry):
    gen = _render_education([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == f"M.S., {_DEGREE_FIELD}Norvale University2001"
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _THESIS
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "4"
    assert gen.stats["supplementary_subpoints"] == 1


def test_degree_prose_goes_under_its_own_row_among_degrees_of_one_institution():
    """The Ph.D. row renders first and holds every distinctive word of the
    M.S. row, so words alone put the M.S. thesis under the Ph.D. (YUYVIG
    FMIGLR 19). Only the row the writer recorded for the M.S. entry takes it."""
    gen = _render_education([_doctorate(), _degree()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert _row_text(above).startswith("M.S.,")
    assert _row_text(above.getprevious()).startswith("Ph.D.,")
    assert row.getnext() is None


def test_degree_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_education([_doctorate(), _degree()], supplementary_subpoints=False)
    on = _render_education([_doctorate(), _degree()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _THESIS in rejected
    assert _THESIS not in "".join(_texts_of(off.doc.element.body))
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


@pytest.mark.parametrize("field, value, line", [
    # The degree's honors print in parentheses in the degree cell.
    ("honors", "summa cum laude, with highest departmental honors",
     "Graduated summa cum laude with highest departmental honors"),
    # Stage 4 put the thesis title in `major`, which the degree cell prints.
    ("major", _THESIS.removeprefix("Thesis: "), _THESIS),
])
def test_degree_prose_already_in_its_row_is_not_offered_again(field, value, line):
    """The degree cell prints these, and the document holds most of each
    line's words (`_record_rendered`), so neither prints twice. A line with
    more words of its own than that is judged by its share of the printed
    field values (`test_a_degree_line_mostly_restating_its_cell_is_not_offered`),
    and one with more new words still is offered whole
    (`test_degree_prose_restating_part_of_its_cell_is_offered_as_written`)."""
    from unified_pipeline.stage6 import supplementary

    entry = _degree(prose=line, **{field: value})
    gen = _render_education([entry], supplementary_subpoints=True)
    assert value in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_rows(gen) == []
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry, "B1", haystack) == []


def test_every_degree_cell_field_the_pass_counts_is_printed_in_the_degree_cell():
    """`supplementary._DEGREE_CELL_FIELDS` restates the education writer's
    degree cell as literals (`_RENDERED_FIELDS` marks schema fields only):
    each one, alone on an entry, reaches that cell."""
    from unified_pipeline.stage6 import supplementary

    for field in supplementary._DEGREE_CELL_FIELDS:
        marker = f"Zq{field.replace('_', 'x')}Z"
        gen = _render_education([_degree(prose="", **{field: marker})], supplementary_subpoints=True)
        (row,) = [tr for tr in gen.doc.element.body.iter(qn("w:tr"))
                  if _row_text(tr).startswith("M.S.,")]
        assert marker in _row_text(row.find(qn("w:tc"))), field


_ADVISOR = "Orrin Halvard Vellacott"


@pytest.mark.parametrize("field, value, line", [
    ("major", "Estuarine Sediment Geochemistry",
     "Concentration: Estuarine Sediment Geochemistry and Coastal Processes"),
    ("field_of_study", "Estuarine Sediment Geochemistry",
     "Concentration: Estuarine Sediment Geochemistry and Coastal Processes"),
    ("honors", "summa cum laude with distinction",
     "Awarded summa cum laude with distinction in research"),
    ("major_professor", _ADVISOR, f"Research supervised by {_ADVISOR}"),
    ("mentor", _ADVISOR, f"Research supervised by {_ADVISOR}"),
    # The minor is new, but half the line's words are the printed major: the
    # share rule every family uses (`SUBPOINT_FIELD_TOKEN_MAX_SHARE`) reads it
    # as the row, so it is not offered.
    ("major", "Estuarine Sediment Geochemistry",
     "Major in Estuarine Sediment Geochemistry with minor in Chemistry"),
])
def test_a_degree_line_mostly_restating_its_cell_is_not_offered(field, value, line):
    """The degree cell prints `major` or `field_of_study`, the honors and an
    advisor under any of its names, though no B1 schema field is any of them
    (`_DEGREE_CELL_FIELDS`). A line at least half of whose words are theirs
    restates the row; the document holds too few of its words for
    `_record_rendered` to see that."""
    from unified_pipeline.stage6 import supplementary

    entry = _degree(prose=line, **{field: value})
    gen = _render_education([entry], supplementary_subpoints=True)
    assert value in _row_text(next(tr for tr in gen.doc.element.body.iter(qn("w:tr"))
                                   if _row_text(tr).startswith("M.S.,")))
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary._record_rendered(line, haystack.text, haystack.line_tokens) is False
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("field, value, line", [
    ("advisor", _ADVISOR, f"Dissertation advisor: Professor {_ADVISOR}, Department of Oceanography"),
    ("honors", "magna cum laude", "Graduated magna cum laude with departmental distinction honors"),
])
def test_degree_prose_restating_part_of_its_cell_is_offered_as_written(field, value, line):
    """More of the line's words are its own than the cell's, so it is
    offered, and as the CV wrote it: Reject All prints the advisor or the
    honors again, under the row that holds them."""
    gen = _render_education([_degree(prose=line, **{field: value})], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert value in _row_text(row.getprevious())
    assert _deleted_texts(gen) == line


def _short_degree(kind):
    """A degree whose row holds too few distinctive (5+ letter) words for the
    word anchor: "PhD", "MD", a year and a four-letter name hold none, and a
    stage-5b name the row prints may share none with stage 4's field."""
    if kind == "one word":
        return _degree(degree="PhD", discipline=None, institution="Kell University", year="2010")
    if kind == "no word":
        return _degree(degree="MD", discipline=None, institution="Vaux", year=None)
    entry = _degree(degree="Ph.D.", discipline=None,
                    institution="Norvale Graduate School of Biomedical Sciences", year="2008")
    entry["institution_enrichment"] = {"cleaned_name": "Halbrook Medical College"}
    return entry


@pytest.mark.parametrize("kind", ["one word", "no word", "renamed"])
def test_a_short_degrees_prose_goes_under_its_own_row(kind):
    """The writer's record says which row is the entry's, so the row is
    taken as it is (`_pick_anchor`), however few words it shares with the
    fields. Before, the word test found no anchor, and the prose reached
    the document in neither arm. Accept All is still the flag-off render."""
    off = _render_education([_short_degree(kind)], supplementary_subpoints=False)
    on = _render_education([_short_degree(kind)], supplementary_subpoints=True)

    (row,) = _subpoint_rows(on)
    assert _row_text(row.getprevious()).startswith(_short_degree(kind)["extracted_fields"]["degree"])
    assert _deleted_texts(on) == _THESIS
    assert _THESIS not in "".join(_texts_of(off.doc.element.body))
    row.getparent().remove(row)
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def test_a_thesis_citation_on_a_row_naming_no_degree_gets_no_sub_point():
    """Stage 4 coded a thesis citation B1 with a field and an institution but
    no degree (EBYSBC VGHNZD 1861). The section writes it a row of its own,
    the field alone in the degree cell, beside the master's degree the
    thesis was for; the prose is of no degree on that row, so it gets none."""
    citation = _degree(idx=23, degree=None, year="2001", prose="",
                       author="Example AB", title=_THESIS.removeprefix("Thesis: "))
    citation["text"] = f"Example AB. {_THESIS.removeprefix('Thesis: ')}. Norvale University, 2001."
    gen = _render_education([_degree(prose=""), citation], supplementary_subpoints=True)

    assert any(_row_text(tr).startswith(f"{_DEGREE_FIELD}Norvale University")
               for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_degree_the_section_skips_gets_no_sub_point_under_another_row():
    """An entry naming no degree and no institution is not written (`_fill_
    education`), so it has no row of its own; its field's words would put
    its prose under the other degree's row."""
    skipped = _degree(idx=21, degree=None, institution=None,
                      prose="Completed coursework in estuarine ecology and coastal sediment chemistry")
    gen = _render_education([_degree(prose=""), skipped], supplementary_subpoints=True)

    assert len([tr for tr in gen.doc.element.body.iter(qn("w:tr"))
                if _row_text(tr).startswith("M.S.,")]) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _render_grouped_education(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_education(grouped.get("B1", []))
    gen._fill_other_education(grouped.get("B2", []))
    gen._fill_postdoc_training(grouped, [e for group in grouped.values() for e in group])
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def test_a_degree_rerouted_out_of_b1_follows_its_destinations_rule():
    """B2 is excluded: a school with no degree renders under Other
    Educational Experiences (`_route_precollegiate_education`), and neither
    the degree row of the same institution nor any other carries its prose.
    A training row stage 3b coded B1 under a postdoctoral heading renders as
    C: its prose is a deleted row under its own training row."""
    school = _degree(idx=32, degree=None, discipline=None, institution="Example Preparatory School",
                     prose="Completed the advanced laboratory sequence in marine field methods")
    training = _rerouted(_degree(idx=31, degree=None, year=None, discipline=None,
                                 prose="Trained in quantitative methods for estuarine monitoring "
                                       "programmes", dates_attended="2002-2004"),
                         "C", "POSTDOCTORAL TRAINING")
    gen, grouped = _render_grouped_education([_degree(prose=""), school, training],
                                             supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["B2"]] == [32]
    assert [e["element_idx_start"] for e in grouped["C"]] == [31]
    assert "Example Preparatory School" in "".join(_texts_of(gen.doc.element.body))
    (row,) = _subpoint_rows(gen)
    assert _table_header(row).startswith("Title, include area of training")
    assert _deleted_texts(gen) == "Trained in quantitative methods for estuarine monitoring programmes"
    assert [(p.element_idx_start, p.code) for p in gen._subpoint_placements] == [(31, "C")]


def test_a_record_rerouted_into_b1_follows_the_degree_rule():
    """A degree programme stage 3b coded C under an education heading renders
    as a degree row: its thesis is a deleted row under that row."""
    gen, grouped = _render_grouped_education(
        [_rerouted(_degree(idx=30, code="C"), "B1", "EDUCATION")], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["B1"]] == [30]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()).startswith("M.S.,")
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _THESIS


def test_a_degree_whose_field_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass dropped one of the degree's fields. What is left
    still describes what was withheld, so none of it is offered."""
    withheld = _degree()
    withheld["_pii_dropped_fields"] = ["advisor"]
    withheld["_pii_withheld"] = True  # the pass sets both
    gen = _render_education([_doctorate(), withheld], supplementary_subpoints=True)

    assert any(_row_text(tr).startswith("M.S.,") for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _degree_line(text, *degrees):
    """One source line stage 4 returned as a record per (degree, year), the
    entry's own fields being the last record's. Fan-out gives every earlier
    record a line of its own and the last one the whole line."""
    records = [{"degree": degree, "discipline": _DEGREE_FIELD, "institution": "Norvale University",
                "year": year} for degree, year in degrees]
    return {"element_idx_start": 25, "taxonomy_code": "B1", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_MS_LINE = f"M.S., {_DEGREE_FIELD}, Norvale University, 2001"
_PHD_LINE = f"Ph.D., {_DEGREE_FIELD}, Norvale University, 2005"


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own degree.
    ([f"{_MS_LINE}, {_THESIS}", _PHD_LINE], [("M.S.", f"{_MS_LINE}, {_THESIS}")]),
    ([_MS_LINE, f"{_PHD_LINE}, {_THESIS}"], [("Ph.D.", f"{_PHD_LINE}, {_THESIS}")]),
    # They do not: whose thesis it is cannot be told, so it goes under neither.
    ([_MS_LINE, _THESIS, _PHD_LINE], []),
])
def test_prose_on_a_line_of_several_degrees_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the degree's line as well as its thesis brings the degree, field,
    institution and year along, and Reject All prints them a second time,
    under the row that holds them (`test_a_thesis_on_its_degrees_own_line_
    brings_the_line_with_it`)."""
    entry = _degree_line("\t".join(segments), ("M.S.", "2001"), ("Ph.D.", "2005"))
    gen, grouped = _render_grouped_education([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["B1"][-1] and len(grouped["B1"]) == 2
    assert [(_row_text(row.getprevious()).split(",")[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def test_a_thesis_on_its_degrees_own_line_brings_the_line_with_it():
    """A fragment splits only on tabs, line breaks and pipes, so a thesis in
    the same sentence run as its degree is offered with the degree's line,
    the shape every family has (#1673's honors, #1935's mentees): Reject All
    prints the degree, field, institution and year again, under the row
    that holds them. The CV's line as written is the shape the owner
    approved (#1205, 2026-10-09)."""
    line = f"{_PHD_LINE}. {_THESIS}"
    entry = _degree(idx=17, degree="Ph.D.", year="2005", prose="")
    entry["text"] = line
    gen = _render_education([entry], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == f"Ph.D., {_DEGREE_FIELD}Norvale University2005"
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == line


def test_the_doctor_matches_a_degree_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the M.S. entry's sub-point
    under its own index and B1, and the doctor's per-entry match counts it for
    that entry, not for the Ph.D. that shares its words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm",
                        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    payload = {"document_uid": "T1205B", "cv_owner": {"full_name": "Quenby Marblegate"},
               "entries": [_doctorate(), _degree()]}
    input_path = tmp_path / "T1205B_fields.json"
    input_path.write_text(json.dumps(payload))

    out = run_stage6(str(input_path), str(tmp_path / "out.docx"), verbose=False)

    sidecar = json.loads((tmp_path / "T1205B_render_warnings.json").read_text())
    assert [(p["element_idx_start"], p["code"], p["paragraphs"])
            for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(19, "B1", [_THESIS])]
    subpoints = subpoint_text(docx_subpoint_lines(Document(out)), subpoint_placements(sidecar))
    assert offered_as_subpoint(_THESIS, subpoints, element_idx=19)
    assert not offered_as_subpoint(_THESIS, subpoints, element_idx=17)


# --- #1205: O institutional leadership, the next code family -----------------
# The leadership section writes every O entry as a row of its three-column
# table, so the prose is a deleted row under the role's own row: the row the
# leadership writer recorded for that entry (the row-receipt ledger).

_ROLE_PROSE = ("Oversaw an annual operating budget and twelve clinical staff, and launched "
               "the regional outreach programme for community clinics")
_ROLE_TITLE = "Director, Example Outreach Programme"
_ROLE_INSTITUTION = "Norvale University"


def _role(idx=57, title=_ROLE_TITLE, start="2008", end="2012", prose=_ROLE_PROSE, code="O", *,
          institution=_ROLE_INSTITUTION, **fields):
    """A leadership role whose source line goes on past its dates, title and
    institution into what it oversaw, as EBYSBC SEKQUI 75's did. Another
    role of the same title and institution shares every distinctive word of
    its row, and so does a later title holding its words ("Assistant
    Director, ..."): only the dates, which are digits, tell them apart. A
    value passed as None is in neither the line nor the fields."""
    named = {"leadership_role": title, "institution": institution, "start_date": start,
             "end_date": end}
    dates = "-".join(value for value in (start, end) if value)
    line = "\t".join(part for part in (dates, ", ".join(v for v in (title, institution) if v)) if part)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_role(title=_ROLE_TITLE):
    """A later role of the same institution, with no prose: it renders first,
    reverse chronologically."""
    return _role(idx=55, title=title, start="2016", end="2020", prose="")


def _render_leadership(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_leadership(entries)
    gen._add_supplementary_subpoints({"O": entries})
    return gen


def _role_table_row():
    """The same role read from a source table: stage 2 joins the row's cells
    with pipes, and its position names the row ("41.2")."""
    entry = _role()
    entry["element_idx_start"] = "41.2"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


_ROLE_ROW = f"{_ROLE_TITLE}{_ROLE_INSTITUTION}2008-2012"


@pytest.mark.parametrize("entry", [_role, _role_table_row])
def test_leadership_prose_is_a_deleted_row_under_the_role_row(entry):
    gen = _render_leadership([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _ROLE_ROW
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _ROLE_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "3"
    assert gen.stats["supplementary_subpoints"] == 1


def test_a_roles_several_prose_lines_are_one_row_of_paragraphs_in_source_order():
    lines = (_ROLE_PROSE, "Chaired the programme steering group and its annual quality review")
    gen = _render_leadership([_role(prose="\t".join(lines))], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _ROLE_ROW
    assert ["".join(t.text for t in p.iter(qn("w:delText"))) for p in row.iter(qn("w:p"))] \
        == list(lines)


@pytest.mark.parametrize("later", [
    # One title held twice at one institution: the rows differ in their dates only.
    _ROLE_TITLE,
    # A later title at the same institution holding every word of this one.
    f"Assistant {_ROLE_TITLE}",
])
def test_leadership_prose_goes_under_its_own_row_among_roles_of_one_institution(later):
    """The later role renders first and holds every distinctive word of the
    earlier one's row, so words alone put the earlier role's prose under the
    later row. Only the row the writer recorded for the earlier entry takes it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_leadership([_later_role(later), _role()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert _row_text(above) == _ROLE_ROW
    assert _row_text(above.getprevious()) == f"{later}{_ROLE_INSTITUTION}2016-2020"
    assert row.getnext() is None
    # Words alone would have taken the later role's row.
    lines = list(supplementary.anchor_candidates(_render_leadership(
        [_later_role(later), _role()], supplementary_subpoints=False).doc))
    rows = [line for line in lines if _ROLE_INSTITUTION in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_role(), "O"), rows) is rows[0]
    assert rows[0].text.startswith(later)


def test_leadership_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_leadership([_later_role(), _role()], supplementary_subpoints=False)
    on = _render_leadership([_later_role(), _role()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _ROLE_PROSE in rejected
    assert _ROLE_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_rows(on)) == 1
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def _role_in_division():
    """Stage 4 put the programme's description in `division_department`,
    which the institution cell prints."""
    return _role(division_department=_ROLE_PROSE)


def _role_with_no_role_or_institution():
    """No role and no institution: the role cell prints the entry's text."""
    return _role(title=None, institution=None)


def _wrapped_role_row():
    """One source table row whose cells wrap (#987): the rejoined text is the row."""
    entry = _role(title=None, institution=None, start=None, end=None, prose="")
    entry.update(element_type="table_row", element_idx_start=41, element_idx_end=41,
                 text=f"{_ROLE_TITLE} | {_ROLE_INSTITUTION}\n{_ROLE_PROSE} | 2008 -\n2012")
    return entry


@pytest.mark.parametrize("entry", [_role_in_division, _role_with_no_role_or_institution,
                                   _wrapped_role_row])
def test_leadership_prose_already_in_its_row_is_not_offered_again(entry):
    """The row prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_leadership([entry()], supplementary_subpoints=True)
    assert _ROLE_PROSE in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "O", haystack) == []


def test_every_leadership_row_field_the_pass_counts_is_printed_in_the_role_row():
    """`supplementary._LEADERSHIP_ROW_FIELDS`, the off-schema names among the
    fields the leadership writer reads its cells from (`_RENDERED_FIELDS`
    marks schema fields only): each one, alone on an entry with the schema's
    field empty, reaches the row."""
    from unified_pipeline.stage6 import supplementary

    for field in supplementary._LEADERSHIP_ROW_FIELDS:
        marker = f"Zq{field}Z"
        entry = _role(title=None, institution=None, prose="", **{field: marker})
        gen = _render_leadership([entry], supplementary_subpoints=True)
        assert any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr"))), field


def test_every_field_the_role_row_prints_is_one_the_pass_counts():
    """The other direction: a field the row prints that `_value_tokens` does
    not count would offer a fragment restating the row. The names are the
    role and organization fields of the codes an entry is rerouted into O
    from (D1-D3, L3, P, Q1, Q2), each alone on an entry."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    counted = _RENDERED_FIELDS["O"] | set(supplementary._LEADERSHIP_ROW_FIELDS)
    names = set().union(*(_RENDERED_FIELDS[code] for code in ("D1", "D2", "D3", "L3", "P", "Q1", "Q2")))
    for field in sorted(names - {"start_date", "end_date"}):
        marker = f"Zq{field}Z"
        entry = {"element_idx_start": 57, "taxonomy_code": "O", "text": "2008-2012",
                 "extracted_fields": {field: marker, "start_date": "2008", "end_date": "2012"}}
        gen = _render_leadership([entry], supplementary_subpoints=True)
        printed = any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
        assert not printed or field in counted, field


@pytest.mark.parametrize("field", ["role", "position", "organization"])
def test_a_leadership_line_mostly_restating_its_row_is_not_offered(field):
    """Stage 4 named the role or the organization under an off-schema key,
    which the row prints when the schema's is empty. A line at least half
    of whose words restate it is the row (`SUBPOINT_FIELD_TOKEN_MAX_SHARE`);
    the document holds too few of its words for `_record_rendered` to see that."""
    from unified_pipeline.stage6 import supplementary

    value = "Regional Outreach Clinics Director"
    line = f"{value}: expanded statewide pediatric services"
    emptied = {"institution": None} if field == "organization" else {"title": None}
    entry = _role(prose=line, **emptied, **{field: value})
    gen = _render_leadership([entry], supplementary_subpoints=True)
    assert any(value in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary._record_rendered(line, haystack.text, haystack.line_tokens) is False
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_flattened_leadership_table_prints_every_line_as_a_row_and_offers_nothing():
    """Three or more lines are a source table flattened into one entry: the
    writer prints each line as a row of its own (`_add_multiline_leadership_
    rows`), a prose line included, so nothing is left to offer. Every row is
    the entry's own (the ledger), and none of them gets a sub-point."""
    block = (f"Chair, Example Admissions Committee | {_ROLE_INSTITUTION} | 2001-2004\n"
             f"Director, Example Outreach Programme | {_ROLE_INSTITUTION} | 2005-2010\n"
             f"{_ROLE_PROSE}")
    entry = {"element_idx_start": 60, "taxonomy_code": "O", "text": block,
             "extracted_fields": {"leadership_role": "Chair, Example Admissions Committee",
                                  "institution": _ROLE_INSTITUTION,
                                  "start_date": "2001", "end_date": "2004"}}
    gen = _render_leadership([entry], supplementary_subpoints=True)

    rows = gen._row_receipts.elements_for(entry)
    assert len(rows) == 3 and _ROLE_PROSE in _row_text(rows[2])
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_leadership_entry_the_section_cannot_place_gets_no_sub_point():
    """The leadership writer binds only to its own table (#625, #664). With
    none, it writes no row, so no row is the entry's own: its prose goes
    under no other line of the document, however many of its words that
    line holds."""
    gen = _generator(supplementary_subpoints=True)
    table = gen._find_table_after_paragraph(gen._find_paragraph_exact(
        "INSTITUTIONAL LEADERSHIP ACTIVITIES"))
    # No longer the leadership table's header, and a line of the section
    # holding every distinctive word of the role's fields.
    table.rows[0].cells[0].text = f"{_ROLE_TITLE}, {_ROLE_INSTITUTION}"
    role = _role()
    gen._fill_leadership([role])
    gen._add_supplementary_subpoints({"O": [role]})

    assert gen._row_receipts.elements_for(role) == []
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _render_grouped_leadership(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_leadership(grouped.get("O", []))
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _role_line(text, *roles):
    """One source line stage 4 returned as a record per (role, start, end),
    the entry's own fields being the last record's. Fan-out gives every
    earlier record a line of its own and the last one the whole line."""
    records = [{"leadership_role": role, "institution": _ROLE_INSTITUTION, "start_date": start,
                "end_date": end} for role, start, end in roles]
    return {"element_idx_start": 63, "taxonomy_code": "O", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_CHAIR = ("Chair, Example Admissions Committee", "2001", "2004")
_CHAIR_LINE = f"Chair, Example Admissions Committee, {_ROLE_INSTITUTION}, 2001-2004"
_DIRECTOR = (_ROLE_TITLE, "2005", "2010")
_DIRECTOR_LINE = f"{_ROLE_TITLE}, {_ROLE_INSTITUTION}, 2005-2010"


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own role.
    ([f"{_CHAIR_LINE}. {_ROLE_PROSE}", _DIRECTOR_LINE],
     [("Chair", f"{_CHAIR_LINE}. {_ROLE_PROSE}")]),
    ([_CHAIR_LINE, f"{_DIRECTOR_LINE}. {_ROLE_PROSE}"],
     [("Director", f"{_DIRECTOR_LINE}. {_ROLE_PROSE}")]),
    # They do not: which role the prose is about cannot be told, so neither takes it.
    ([_CHAIR_LINE, _ROLE_PROSE, _DIRECTOR_LINE], []),
])
def test_prose_on_a_line_of_several_roles_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the role's line as well as its prose brings the role, institution
    and dates along, and Reject All prints them a second time, under the
    row that holds them (the shape every family has; #1205, 2026-10-09)."""
    entry = _role_line("\t".join(segments), _CHAIR, _DIRECTOR)
    gen, grouped = _render_grouped_leadership([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["O"][-1] and len(grouped["O"]) == 2
    assert [(_row_text(row.getprevious()).split(",")[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def test_generic_list_fan_out_children_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, the entry prints its whole
    text in the role cell, and nothing is left to offer either."""
    records = [{"leadership_role": "Chair, Example Admissions Committee",
                "institution": _ROLE_INSTITUTION, "start_date": "2001", "end_date": "2004"},
               {"leadership_role": _ROLE_TITLE, "institution": _ROLE_INSTITUTION,
                "start_date": "2005", "end_date": "2010"}]
    for text, children in ((f"{_CHAIR_LINE}\t{_DIRECTOR_LINE}", 2),
                           (f"{_CHAIR_LINE}\t{_DIRECTOR_LINE}\t{_ROLE_PROSE}", 1)):
        entry = {"element_idx_start": 64, "taxonomy_code": "O", "text": text,
                 "extracted_fields": {"leadership_roles": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_leadership([entry], supplementary_subpoints=True)
        assert len(grouped["O"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["O"])
        assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_leadership_role_rerouted_out_of_o_follows_its_destination_and_gets_no_sub_point():
    """N1 is excluded: a role under a mentoring-programmes heading renders as
    N1. A record holding only an institution and a start date under a
    training heading renders as C, whose prose goes only under its own C row
    (none is written here). Neither the role row of the same institution
    nor any other carries their prose."""
    programme = _rerouted(_role(idx=70), "N1", "MENTORING PROGRAMS")
    training = _rerouted(_role(idx=71, title=None, end=None,
                               prose="Trained in quantitative methods for community outreach "
                                     "programme evaluation"), "C", "POSTDOCTORAL TRAINING")
    gen, grouped = _render_grouped_leadership([_later_role(), programme, training],
                                              supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["N1"]] == [70]
    assert [e["element_idx_start"] for e in grouped["C"]] == [71]
    assert [e["element_idx_start"] for e in grouped["O"]] == [55]
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_refused_reroute_keeps_the_leadership_rule():
    """A role rerouted to I keeps its O code: I's renderer writes none of its
    fields (`_cross_family_refusal`). It renders as an O row and its prose is
    a deleted row under that row."""
    gen, grouped = _render_grouped_leadership([_rerouted(_role(), "I", "MEMBERSHIPS")],
                                              supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["O"]] == [57]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _ROLE_ROW


def test_a_record_rerouted_into_o_follows_the_leadership_rule():
    """A clinical leadership role stage 3b coded L3 under the institutional
    leadership heading renders as an O row: its prose is a deleted row under
    that row."""
    gen, grouped = _render_grouped_leadership(
        [_rerouted(_role(code="L3"), "O", "INSTITUTIONAL LEADERSHIP")], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["O"]] == [57]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _ROLE_ROW
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _ROLE_PROSE


def test_a_leadership_role_whose_field_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass dropped one of the role's fields. What is left
    still describes what was withheld, so none of it is offered."""
    withheld = _role()
    withheld["_pii_dropped_fields"] = ["division_department"]
    withheld["_pii_withheld"] = True  # the pass sets both
    gen = _render_leadership([_later_role(), withheld], supplementary_subpoints=True)

    assert any(_row_text(tr) == _ROLE_ROW for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_leadership_role_dedup_drops_offers_its_prose_once_under_the_kept_row():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no row, so its
    prose goes only under the row of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_role(idx=57), _role(idx=58)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["O"]) == 2 and len(kept["O"]) == 1
    gen._fill_leadership(kept["O"])
    gen._add_supplementary_subpoints(before)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _ROLE_ROW
    assert _deleted_texts(gen) == _ROLE_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["O"][0]["element_idx_start"]


def test_leadership_and_appointment_prose_each_stay_under_their_own_section_row():
    """An appointment and a leadership role can share a title and an
    institution: each one's prose goes under its own section's row, never
    under the other's."""
    appointment_prose = ("Supervised the example clinic's residents and its weekly "
                         "interdisciplinary case conference")
    appointment = {"element_idx_start": 29, "taxonomy_code": "D1",
                   "text": f"2008-2012\t{_ROLE_TITLE}, {_ROLE_INSTITUTION}\t{appointment_prose}",
                   "extracted_fields": {"title": _ROLE_TITLE, "institution": _ROLE_INSTITUTION,
                                        "start_date": "2008", "end_date": "2012"}}
    gen = _generator(supplementary_subpoints=True)
    role = _role()
    gen._fill_positions({"D1": [appointment]})
    gen._fill_leadership([role])
    gen._add_supplementary_subpoints({"D1": [appointment], "O": [role]})

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {appointment_prose, _ROLE_PROSE}
    leadership_header = "Role(s)/Position"
    assert _row_text(by_text[_ROLE_PROSE].getparent().find(qn("w:tr"))).startswith(leadership_header)
    assert not _row_text(by_text[appointment_prose].getparent().find(qn("w:tr"))).startswith(
        leadership_header)
    assert by_text[_ROLE_PROSE].getparent() is not by_text[appointment_prose].getparent()


def test_the_doctor_matches_a_leadership_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier role's sub-point
    under its own index and O, and the doctor's per-entry match counts it for
    that entry, not for the later role that shares its words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm",
                        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    payload = {"document_uid": "T1205O", "cv_owner": {"full_name": "Quenby Marblegate"},
               "entries": [_later_role(), _role()]}
    input_path = tmp_path / "T1205O_fields.json"
    input_path.write_text(json.dumps(payload))

    out = run_stage6(str(input_path), str(tmp_path / "out.docx"), verbose=False)

    sidecar = json.loads((tmp_path / "T1205O_render_warnings.json").read_text())
    assert [(p["element_idx_start"], p["code"], p["paragraphs"])
            for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(57, "O", [_ROLE_PROSE])]
    subpoints = subpoint_text(docx_subpoint_lines(Document(out)), subpoint_placements(sidecar))
    assert offered_as_subpoint(_ROLE_PROSE, subpoints, element_idx=57)
    assert not offered_as_subpoint(_ROLE_PROSE, subpoints, element_idx=55)


def _stage6_render(tmp_path, monkeypatch, entries, uid="T1205P", **kwargs):
    """Through `run_stage6`, so the PII pass and the date-led split (#983)
    run as in a real render: the document and its sidecar's placements.
    `kwargs` go to `run_stage6` (`supplementary_subpoints`)."""
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    monkeypatch.setattr("unified_pipeline.stage_6_word_template.call_llm",
                        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    tmp_path.mkdir(parents=True, exist_ok=True)
    input_path = tmp_path / f"{uid}_fields.json"
    input_path.write_text(json.dumps({"document_uid": uid,
                                      "cv_owner": {"full_name": "Quenby Marblegate"},
                                      "entries": entries}))
    out = run_stage6(str(input_path), str(tmp_path / "out.docx"), verbose=False, **kwargs)
    sidecar = json.loads((tmp_path / f"{uid}_render_warnings.json").read_text())
    return Document(out), [p["paragraphs"] for p in sidecar[SUBPOINT_PLACEMENTS_KEY]]


_SPOUSE = "Zelphine Qorvath"
_HOME_PHONE = "212-555-0147"


def _naming_the_spouse(entry, prose):
    """The spouse under a field named for it, and in the prose with no label."""
    entry["extracted_fields"]["spouse_name"] = _SPOUSE
    entry["text"] = entry["text"].replace(prose, f"{prose}, with my spouse {_SPOUSE}")
    return entry


def _with_a_home_phone(entry, prose):
    """A home phone number in the middle of the prose's line."""
    entry["text"] = entry["text"].replace(prose, f"{prose}. Home phone: {_HOME_PHONE}")
    return entry


@pytest.mark.parametrize("make, prose", [(_role, _ROLE_PROSE), (_honor, _HONOR_PROSE)],
                         ids=["O", "H"])
@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, make, prose, protect, value):
    """#1205 review, probes P7 and P9. The PII pass drops a field named for
    protected data and marks the entry withheld, but lists no
    `_pii_dropped_fields`; and no scope withholds a phone number in the
    middle of a line. Both printed in the entry's sub-point on Reject All,
    while the Appendix said the spouse was withheld. The same entry without
    them still has its sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [make()])
    assert placements == [[prose]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch, [protect(make(), prose)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


_CHAIR_PROSE = ("Established the departmental mentoring scheme and doubled faculty retention "
                "across affiliated hospitals")


def test_a_line_of_two_dated_roles_offers_neither_ones_prose(tmp_path, monkeypatch):
    """#1205 review, probe P12. One source line holds two roles, each opening
    with its dates and followed by its prose, and stage 4 returned the second.
    Both prose lines went under the second role's row on Reject All, the
    first role's achievement with them. Which role a line is about cannot be
    told, so neither is offered; the second role alone on its line still is."""
    first = f"2001-2004 Chair, Example Admissions Committee, {_ROLE_INSTITUTION}"
    second = f"2005-2010 {_ROLE_TITLE}, {_ROLE_INSTITUTION}"

    def role(text):
        return {"element_idx_start": 90, "taxonomy_code": "O", "text": text,
                "extracted_fields": {"leadership_role": _ROLE_TITLE, "institution": _ROLE_INSTITUTION,
                                     "start_date": "2005", "end_date": "2010"}}

    _doc, placements = _stage6_render(tmp_path / "one", monkeypatch, [role(f"{second}\t{_ROLE_PROSE}")])
    assert placements == [[_ROLE_PROSE]]

    doc, placements = _stage6_render(
        tmp_path / "two", monkeypatch, [role(f"{first}\t{_CHAIR_PROSE}\t{second}\t{_ROLE_PROSE}")])
    assert placements == []
    deleted = "".join(t.text or "" for t in doc.element.body.iter(qn("w:delText")))
    assert _CHAIR_PROSE not in deleted and _ROLE_PROSE not in deleted


# --- #1205: P institutional committees, the next code family -----------------
# The administrative-activities section writes every P entry as a row of its
# three-column table, so the prose is a deleted row under the committee's own
# row: the one row the writer recorded for that entry (the row-receipt
# ledger). An entry holding several records offers none.

_COMMITTEE_PROSE = ("Reviewed every proposed change to the core clerkship curriculum and drafted "
                    "the annual assessment report for the faculty council")
_COMMITTEE_NAME = "Example Curriculum Steering Committee"
_COMMITTEE_INSTITUTION = "Norvale University"


def _committee(idx=81, name=_COMMITTEE_NAME, role="Member", start="2008", end="2012",
               prose=_COMMITTEE_PROSE, code="P", *, institution=_COMMITTEE_INSTITUTION, **fields):
    """A committee whose source line goes on past its name, institution, role
    and years into what it did, as EBYSBC MQSUIC 615's did. Another term of
    the same committee shares every distinctive word of its row, and so does
    a later committee holding its words ("Executive ..."): only the years,
    which are digits, tell them apart. A value passed as None is in neither
    the line nor the fields."""
    named = {"committee_name": name, "role": role, "institution": institution,
             "start_date": start, "end_date": end}
    dates = "-".join(value for value in (start, end) if value)
    line = ", ".join(value for value in (name, institution, role, dates) if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_committee(name=_COMMITTEE_NAME):
    """A later term at the same institution, with no prose: it renders first,
    reverse chronologically."""
    return _committee(idx=79, name=name, start="2016", end="2020", prose="")


def _render_committees(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_administrative_activities(entries)
    gen._add_supplementary_subpoints({"P": entries})
    return gen


def _committee_table_row():
    """The same committee read from a source table: stage 2 joins the row's
    cells with pipes, and its position names the row ("41.2")."""
    entry = _committee()
    entry["element_idx_start"] = "41.2"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


_COMMITTEE_ROW = f"{_COMMITTEE_NAME}, {_COMMITTEE_INSTITUTION}Member2008-2012"
_COMMITTEE_HEADER = "Name of Committee"


def _table_header(row):
    return _row_text(row.getparent().find(qn("w:tr")))


@pytest.mark.parametrize("entry", [_committee, _committee_table_row])
def test_committee_prose_is_a_deleted_row_under_the_committee_row(entry):
    gen = _render_committees([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _COMMITTEE_ROW
    assert _table_header(row).startswith(_COMMITTEE_HEADER)
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _COMMITTEE_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "3"
    assert gen.stats["supplementary_subpoints"] == 1


def test_a_committees_several_prose_lines_are_one_row_of_paragraphs_in_source_order():
    lines = (_COMMITTEE_PROSE, "Chaired its assessment sub-committee and the annual review of electives")
    gen = _render_committees([_committee(prose="\t".join(lines))], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _COMMITTEE_ROW
    assert ["".join(t.text for t in p.iter(qn("w:delText"))) for p in row.iter(qn("w:p"))] \
        == list(lines)


@pytest.mark.parametrize("later", [
    # One committee in two terms at one institution: the rows differ in their years only.
    _COMMITTEE_NAME,
    # A later committee at the same institution holding every word of this one.
    f"Executive {_COMMITTEE_NAME}",
])
def test_committee_prose_goes_under_its_own_row_among_committees_of_one_institution(later):
    """The later committee renders first and holds every distinctive word of
    the earlier one's row, so words alone put the earlier one's prose under
    the later row. Only the row the writer recorded for the earlier entry
    takes it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_committees([_later_committee(later), _committee()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert _row_text(above) == _COMMITTEE_ROW
    assert _row_text(above.getprevious()) == f"{later}, {_COMMITTEE_INSTITUTION}Member2016-2020"
    assert row.getnext() is None
    # Words alone would have taken the later committee's row.
    lines = list(supplementary.anchor_candidates(_render_committees(
        [_later_committee(later), _committee()], supplementary_subpoints=False).doc))
    rows = [line for line in lines if _COMMITTEE_INSTITUTION in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_committee(), "P"), rows) is rows[0]
    assert rows[0].text.startswith(later)


def test_committee_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_committees([_later_committee(), _committee()], supplementary_subpoints=False)
    on = _render_committees([_later_committee(), _committee()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _COMMITTEE_PROSE in rejected
    assert _COMMITTEE_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_rows(on)) == 1
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def test_a_short_committees_prose_goes_under_its_own_row():
    """An acronym, an acronym's institution and a short role hold one
    distinctive word, too few for the word test: the writer's record says
    the one row is the entry's own (`_pick_anchor`)."""
    gen = _render_committees([_committee(name="IRB", institution="NVU", role="Chair")],
                             supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == "IRB, NVUChair2008-2012"
    assert _deleted_texts(gen) == _COMMITTEE_PROSE


def _committee_with_no_name():
    """No committee name: the name cell prints the entry's line, its role
    and years taken out (#1346)."""
    return _committee(name=None, institution=None)


def _committee_with_parenthetical_years():
    """Nothing extracted: the writer reads the role and years out of the
    line's parenthesis and prints the rest of the line (#660)."""
    entry = _committee(name=None, role=None, start=None, end=None, institution=None)
    entry["text"] = f"{_COMMITTEE_NAME} (Chair 2008-2012). {_COMMITTEE_PROSE}"
    return entry


def _committee_on_two_lines_with_nothing_extracted():
    """Two lines, nothing extracted: the writer re-parses them, one row each."""
    return {"element_idx_start": 81, "taxonomy_code": "P", "extracted_fields": {},
            "text": f"{_COMMITTEE_NAME}, {_COMMITTEE_INSTITUTION}, 2008-2012\n{_COMMITTEE_PROSE}"}


@pytest.mark.parametrize("entry", [_committee_with_no_name, _committee_with_parenthetical_years,
                                   _committee_on_two_lines_with_nothing_extracted])
def test_committee_prose_already_in_its_row_is_not_offered_again(entry):
    """The row prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_committees([entry()], supplementary_subpoints=True)
    assert _COMMITTEE_PROSE in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "P", haystack) == []


def test_every_committee_row_field_the_pass_counts_is_printed_in_the_committee_row():
    """`supplementary._ADMINISTRATIVE_ROW_FIELDS`, the off-schema names among
    the fields the writer reads its name cell from (`_RENDERED_FIELDS` marks
    schema fields only): each one, alone on an entry with the schema's field
    empty, reaches the row."""
    from unified_pipeline.stage6 import supplementary

    assert supplementary._ADMINISTRATIVE_ROW_FIELDS
    for field in supplementary._ADMINISTRATIVE_ROW_FIELDS:
        marker = f"Zq{field}Z"
        entry = _committee(name=None, prose="", **{field: marker})
        gen = _render_committees([entry], supplementary_subpoints=True)
        assert any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr"))), field


def test_every_field_the_committee_row_prints_is_one_the_pass_counts():
    """The other direction: a field the row prints that `_value_tokens` does
    not count would offer a fragment restating the row. The names are the
    fields of the codes an entry is rerouted into P from (K3, O, Q1, Q2, and
    the positions D1-D3), each alone on an entry."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    counted = _RENDERED_FIELDS["P"] | set(supplementary._ADMINISTRATIVE_ROW_FIELDS)
    names = set().union(*(_RENDERED_FIELDS[code] for code in ("D1", "D2", "D3", "O", "Q1", "Q2")),
                        supplementary.ADMINISTRATIVE_ACTIVITY_FIELDS, {"name", "title", "position"})
    for field in sorted(names - {"start_date", "end_date"}):
        marker = f"Zq{field}Z"
        entry = {"element_idx_start": 81, "taxonomy_code": "P", "text": "2008-2012",
                 "extracted_fields": {field: marker, "start_date": "2008", "end_date": "2012"}}
        gen = _render_committees([entry], supplementary_subpoints=True)
        printed = any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
        assert not printed or field in counted, field


@pytest.mark.parametrize("field", ["committee", "activity"])
def test_a_committee_line_mostly_restating_its_row_is_not_offered(field):
    """Stage 4 named the committee under an off-schema key, which the row
    prints when the schema's is empty. A line at least half of whose words
    restate it is the row (`SUBPOINT_FIELD_TOKEN_MAX_SHARE`); the document
    holds too few of its words for `_record_rendered` to see that."""
    from unified_pipeline.stage6 import supplementary

    value = "Regional Pharmacy Formulary Review Panel"
    line = f"{value}: approved statewide pediatric dosing"
    entry = _committee(name=None, prose=line, **{field: value})
    gen = _render_committees([entry], supplementary_subpoints=True)
    assert any(value in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary._record_rendered(line, haystack.text, haystack.line_tokens) is False
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


# The role first, as many CVs write it. A line opening with the name cell's
# own text is read as that row, prose and all (`_record_rendered`'s verbatim
# piece), and is not offered.
_ADMISSIONS_LINE = f"Chair, Example Admissions Committee, {_COMMITTEE_INSTITUTION}, 2001-2004"
_ADMISSIONS_PROSE = ("Interviewed every applicant shortlisted for the rural medicine track "
                     "and chaired the appeals hearings")


def test_a_list_of_committee_records_prints_a_row_each_and_offers_nothing():
    """Stage 4 packed two committees as a list under the committee name
    (#208/#248): the writer prints a row for each, both the entry's own (the
    ledger), and which of them a prose line is about cannot be shown. The
    line names no years, so the writer's rows alone say it holds two."""
    entry = {"element_idx_start": 60, "taxonomy_code": "P",
             "text": f"Chair, Example Admissions Committee; Member, {_COMMITTEE_NAME}"
                     f"\t{_ADMISSIONS_PROSE}",
             "extracted_fields": {"committee_name": [
                 {"committee_name": "Example Admissions Committee", "role": "Chair",
                  "start_date": "2001", "end_date": "2004"},
                 {"committee_name": _COMMITTEE_NAME, "role": "Member",
                  "start_date": "2008", "end_date": "2012"}]}}
    gen = _render_committees([entry], supplementary_subpoints=True)

    assert len(gen._row_receipts.elements_for(entry)) == 2
    assert _ADMISSIONS_PROSE not in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_block_of_dated_committees_prints_a_row_each_and_offers_nothing():
    """Two or more dated lines are a source table flattened into one entry
    (#660): the writer prints each line as a row of its own
    (`_multiline_committee_rows`), a prose line included, so nothing is left
    to offer. Every row is the entry's own."""
    block = (f"Example Admissions Committee | Chair | 2001-2004\n"
             f"{_COMMITTEE_NAME} | Member | 2008-2012\n"
             f"{_COMMITTEE_PROSE}")
    entry = {"element_idx_start": 60, "taxonomy_code": "P", "text": block,
             "extracted_fields": {"committee_name": "Example Admissions Committee",
                                  "role": "Chair", "start_date": "2001", "end_date": "2004"}}
    gen = _render_committees([entry], supplementary_subpoints=True)

    rows = gen._row_receipts.elements_for(entry)
    assert len(rows) == 3 and _COMMITTEE_PROSE in _row_text(rows[2])
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _fused_committees():
    """Three committees each followed by its years, which stage 2 joined with
    tabs (EBYSBC YOXXOH 74's shape): one line, so the writer prints one row,
    of the record stage 4 extracted, the last; the other two are its records,
    not its prose."""
    parts = ["Alpha Questionnaire Working Group, Kestrel Medical Society, 1979-1980",
             "Bravo Strategic Planning Group, Osprey Regional Council, 1980-1982",
             f"{_COMMITTEE_NAME}, {_COMMITTEE_INSTITUTION}, Steering Group, 1986"]
    return {"element_idx_start": 74, "taxonomy_code": "P", "text": "\t".join(parts),
            "extracted_fields": {"committee_name": _COMMITTEE_NAME, "role": "Steering Group",
                                 "institution": _COMMITTEE_INSTITUTION, "start_date": "1986"}}


@pytest.mark.parametrize("flag", [True, False])
def test_a_line_of_several_dated_committees_offers_none_and_its_records_are_recovered(flag):
    """#1205 comment 6275741663: the #221 recovery reads every sub-point
    line as rendered, so a committee record offered as another's prose would
    be in the accepted document nowhere and still not be recovered. None is
    offered, and the recovery re-emits the same records with sub-points on
    as off."""
    entry = _fused_committees()
    gen = _render_committees([entry], supplementary_subpoints=flag)
    assert len(gen._row_receipts.elements_for(entry)) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""

    gen._recover_unrendered_records({"P": [entry]})

    assert gen.stats["unrendered_records_recovered"] == 2
    page = "".join(_texts_of(gen.doc.element.body))
    assert "Alpha Questionnaire Working Group" in page and "Bravo Strategic Planning Group" in page


def test_a_committee_the_section_writes_no_row_for_gets_no_sub_point():
    """A record naming no committee is skipped (#256), so no row is the
    entry's own: its prose goes under no other line of the document, though
    the other committee's row holds every distinctive word of its line."""
    nameless = _committee(name=None)
    nameless["extracted_fields"] = {"committee_name": [{"role": "Member", "start_date": "2008"}]}
    later = _later_committee()
    gen = _render_committees([later, nameless], supplementary_subpoints=True)

    assert gen._row_receipts.elements_for(nameless) == []
    assert gen._row_receipts.elements_for(later)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


_KESTREL_LINE = "Member, Example Admissions Appeals Committee, Kestrel Medical Society"


def _committee_after_an_undated_committee():
    """Another committee, undated, then the extracted one with its years: one
    dated part, so the writer's #660 test reads one record."""
    entry = _committee()
    entry["text"] = f"{_KESTREL_LINE}\t{entry['text']}"
    return entry


def _committee_list_with_a_nameless_record():
    """Stage 4 packed a record naming no committee and the committee in a
    list (#208/#248): the writer prints no row for the nameless one (#256)."""
    entry = _committee(prose=None)
    entry["text"] = (f"Member, 2001-2004, {_ADMISSIONS_PROSE}\t"
                     f"{_COMMITTEE_NAME}, {_COMMITTEE_INSTITUTION}, Member")
    entry["extracted_fields"] = {"committee_name": [
        {"role": "Member", "start_date": "2001", "end_date": "2004"},
        {"committee_name": _COMMITTEE_NAME, "institution": _COMMITTEE_INSTITUTION,
         "role": "Member"}]}
    return entry


@pytest.mark.parametrize("entry", [_committee_after_an_undated_committee,
                                   _committee_list_with_a_nameless_record])
def test_a_committee_row_holding_another_record_gets_no_sub_point(entry):
    """The writer prints one row, the entry's own, though the line holds
    another record: that record's text would be offered as the committee's
    prose, under its row."""
    entry = entry()
    gen = _render_committees([entry], supplementary_subpoints=True)

    assert len(gen._row_receipts.elements_for(entry)) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _render_grouped_committees(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_administrative_activities(grouped.get("P", []))
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _committee_line(text, *committees):
    """One source line stage 4 returned as a record per (name, role, start,
    end), the entry's own fields being the last record's. Fan-out gives every
    earlier record a line of its own and the last one the whole line."""
    records = [{"committee_name": name, "role": role, "institution": _COMMITTEE_INSTITUTION,
                "start_date": start, "end_date": end} for name, role, start, end in committees]
    return {"element_idx_start": 63, "taxonomy_code": "P", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_ADMISSIONS = ("Example Admissions Committee", "Chair", "2001", "2004")
_STEERING = (_COMMITTEE_NAME, "Member", "2005", "2010")
_STEERING_LINE = f"Member, {_COMMITTEE_NAME}, {_COMMITTEE_INSTITUTION}, 2005-2010"


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own committee.
    ([f"{_ADMISSIONS_LINE}. {_ADMISSIONS_PROSE}", _STEERING_LINE],
     [("Example Admissions Committee", f"{_ADMISSIONS_LINE}. {_ADMISSIONS_PROSE}")]),
    ([_ADMISSIONS_LINE, f"{_STEERING_LINE}. {_COMMITTEE_PROSE}"],
     [(_COMMITTEE_NAME, f"{_STEERING_LINE}. {_COMMITTEE_PROSE}")]),
    # They do not: which committee the prose is about cannot be told, so neither takes it.
    ([_ADMISSIONS_LINE, _COMMITTEE_PROSE, _STEERING_LINE], []),
])
def test_prose_on_a_line_of_several_committees_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the committee's line as well as its prose brings the name,
    institution, role and years along, and Reject All prints them a second
    time, under the row that holds them (the shape every family has; #1205,
    2026-10-09)."""
    entry = _committee_line("\t".join(segments), _ADMISSIONS, _STEERING)
    gen, grouped = _render_grouped_committees([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["P"][-1] and len(grouped["P"]) == 2
    assert [(_row_text(row.getprevious()).split(",")[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def test_a_committee_line_opening_with_its_name_cell_is_read_as_its_row():
    """The name cell folds the institution in ("Committee, Institution"), so
    a line that opens with that text holds the row's own text verbatim, and
    `_record_rendered` reads the whole line as rendered: none of it is
    offered, its prose included. A line opening with the role is offered
    whole, as the CV wrote it."""
    name_first = f"{_COMMITTEE_NAME}, {_COMMITTEE_INSTITUTION}, Member, 2005-2010. {_COMMITTEE_PROSE}"
    role_first = f"{_STEERING_LINE}. {_COMMITTEE_PROSE}"
    for line, offered in ((name_first, ""), (role_first, role_first)):
        entry = _committee(start="2005", end="2010", prose=None)
        entry["text"] = line
        gen = _render_committees([entry], supplementary_subpoints=True)
        assert _deleted_texts(gen) == offered


def test_generic_list_fan_out_children_of_committees_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, the entry prints its whole
    line in the name cell, and nothing is left to offer either."""
    records = [{"committee_name": "Example Admissions Committee", "role": "Chair",
                "institution": _COMMITTEE_INSTITUTION, "start_date": "2001", "end_date": "2004"},
               {"committee_name": _COMMITTEE_NAME, "role": "Member",
                "institution": _COMMITTEE_INSTITUTION, "start_date": "2005", "end_date": "2010"}]
    for text, children in ((f"{_ADMISSIONS_LINE}\t{_STEERING_LINE}", 2),
                           (f"{_ADMISSIONS_LINE}\t{_STEERING_LINE}\t{_COMMITTEE_PROSE}", 1)):
        entry = {"element_idx_start": 64, "taxonomy_code": "P", "text": text,
                 "extracted_fields": {"committees": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_committees([entry], supplementary_subpoints=True)
        assert len(grouped["P"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["P"])
        assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_committee_rerouted_out_of_p_follows_its_destination_and_gets_no_sub_point():
    """N1 is excluded: a committee under a mentoring-programmes heading
    renders as N1. A record holding only an institution and a start year
    under a training heading renders as C, whose prose goes only under its
    own C row (none is written here). Neither the committee row of the same
    institution nor any other carries their prose."""
    programme = _rerouted(_committee(idx=70), "N1", "MENTORING PROGRAMS")
    training = _rerouted(_committee(idx=71, name=None, role=None, end=None,
                                    prose="Trained in quantitative methods for clerkship "
                                          "curriculum evaluation"), "C", "POSTDOCTORAL TRAINING")
    gen, grouped = _render_grouped_committees([_later_committee(), programme, training],
                                              supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["N1"]] == [70]
    assert [e["element_idx_start"] for e in grouped["C"]] == [71]
    assert [e["element_idx_start"] for e in grouped["P"]] == [79]
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_refused_reroute_keeps_the_committee_rule():
    """A committee rerouted to I keeps its P code: I's renderer would drop
    its committee name (`_fields_lost_in_reroute`). It renders as a P row and
    its prose is a deleted row under that row."""
    gen, grouped = _render_grouped_committees([_rerouted(_committee(), "I", "MEMBERSHIPS")],
                                              supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["P"]] == [81]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _COMMITTEE_ROW


@pytest.mark.parametrize("code", ["Q2", "K3"])
def test_a_record_rerouted_into_p_follows_the_committee_rule(code):
    """A board seat stage 3b coded Q2, or an administrative teaching role it
    coded K3, under the institutional administrative heading renders as a P
    row: its prose is a deleted row under that row."""
    gen, grouped = _render_grouped_committees(
        [_rerouted(_committee(code=code), "P", "INSTITUTIONAL ADMINISTRATIVE ACTIVITIES")],
        supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["P"]] == [81]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _COMMITTEE_ROW
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _COMMITTEE_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.code == "P"


def test_a_committee_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _committee()
    withheld["_pii_withheld"] = True
    gen = _render_committees([_later_committee(), withheld], supplementary_subpoints=True)

    assert any(_row_text(tr) == _COMMITTEE_ROW for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_committee_dedup_drops_offers_its_prose_once_under_the_kept_row():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no row, so its
    prose goes only under the row of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_committee(idx=81), _committee(idx=82)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["P"]) == 2 and len(kept["P"]) == 1
    gen._fill_administrative_activities(kept["P"])
    gen._add_supplementary_subpoints(before)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _COMMITTEE_ROW
    assert _deleted_texts(gen) == _COMMITTEE_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["P"][0]["element_idx_start"]


def test_committee_leadership_and_appointment_prose_each_stay_under_their_own_section_row():
    """A committee, the leadership role of chairing it and an appointment can
    share a title and an institution: each one's prose goes under its own
    section's row, never under another's."""
    title = f"Chair, {_COMMITTEE_NAME}"
    appointment_prose = ("Supervised the example clinic's residents and its weekly "
                         "interdisciplinary case conference")
    role_prose = ("Oversaw an annual operating budget and twelve clinical staff, and launched "
                  "the regional outreach programme for community clinics")
    appointment = {"element_idx_start": 29, "taxonomy_code": "D1",
                   "text": f"2008-2012\t{title}, {_COMMITTEE_INSTITUTION}\t{appointment_prose}",
                   "extracted_fields": {"title": title, "institution": _COMMITTEE_INSTITUTION,
                                        "start_date": "2008", "end_date": "2012"}}
    role = {"element_idx_start": 57, "taxonomy_code": "O",
            "text": f"2008-2012\t{title}, {_COMMITTEE_INSTITUTION}\t{role_prose}",
            "extracted_fields": {"leadership_role": title, "institution": _COMMITTEE_INSTITUTION,
                                 "start_date": "2008", "end_date": "2012"}}
    committee = _committee(role="Chair")
    gen = _generator(supplementary_subpoints=True)
    gen._fill_positions({"D1": [appointment]})
    gen._fill_leadership([role])
    gen._fill_administrative_activities([committee])
    gen._add_supplementary_subpoints({"D1": [appointment], "O": [role], "P": [committee]})

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {appointment_prose, role_prose, _COMMITTEE_PROSE}
    assert _table_header(by_text[_COMMITTEE_PROSE]).startswith(_COMMITTEE_HEADER)
    assert _table_header(by_text[role_prose]).startswith("Role(s)/Position")
    assert not _table_header(by_text[appointment_prose]).startswith(
        (_COMMITTEE_HEADER, "Role(s)/Position"))
    assert len({id(row.getparent()) for row in by_text.values()}) == 3


def test_the_doctor_matches_a_committee_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier term's sub-point
    under its own index and P, and the doctor's per-entry match counts it for
    that entry, not for the later term that shares its words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_committee(), _committee()],
                                     uid="T1205C")
    assert placements == [[_COMMITTEE_PROSE]]
    sidecar = json.loads((tmp_path / "T1205C_render_warnings.json").read_text())
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(81, "P")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_COMMITTEE_PROSE, subpoints, element_idx=81)
    assert not offered_as_subpoint(_COMMITTEE_PROSE, subpoints, element_idx=79)


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_committee_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line (#1205 review of O). The same committee without them still has its
    sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_committee()])
    assert placements == [[_COMMITTEE_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_committee(), _COMMITTEE_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


# --- #1205: R invited presentations, the next code family --------------------
# The presentations section writes every R entry as one row of its Regional,
# National or International table, so the prose is a deleted row under the
# talk's own row: the one row the writer recorded for that entry (the
# row-receipt ledger). An entry holding several talks offers none.

_TALK_PROSE = ("Delivered to an audience of about two hundred residents and nursing staff, "
               "followed by a moderated question session")
_TALK_TITLE = "Example Advances in Coastal Wetland Hydrology"
_TALK_EVENT = "Example Grand Rounds"
_TALK_VENUE = "Norvale University, Crab Hollow, ZQ"


def _talk(idx=601, title=_TALK_TITLE, event=_TALK_EVENT, venue=_TALK_VENUE, date="2016",
          prose=_TALK_PROSE, code="R", **fields):
    """A talk whose source line goes on past its title, meeting, venue and
    year into who it was for, as YUYVIG SIJYJZ 621's did. Another talk of the
    same title at the same venue shares every distinctive word of its row:
    only the year, a digit, tells them apart. A value passed as None is in
    neither the line nor the fields."""
    named = {"title": title, "event_name": event, "location": venue, "date": date}
    line = ", ".join(value for value in (title, event, venue, date) if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_talk(title=_TALK_TITLE):
    """The same talk given again, at the same venue, with no prose: it renders
    first, reverse chronologically."""
    return _talk(idx=599, title=title, date="2019", prose="")


def _render_talks(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_presentations(entries)
    gen._add_supplementary_subpoints({"R": entries})
    return gen


def _talk_table_row():
    """The same talk read from a source table: stage 2 joins the row's cells
    with pipes, and its position names the row ("41.2")."""
    entry = _talk()
    entry["element_idx_start"] = "41.2"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


_TALK_ROW = f"{_TALK_TITLE}{_TALK_EVENT}, {_TALK_VENUE}2016"
_TALK_HEADER = "TitleInstitution/LocationDates (yyyy)"


@pytest.mark.parametrize("entry", [_talk, _talk_table_row])
def test_talk_prose_is_a_deleted_row_under_the_talk_row(entry):
    gen = _render_talks([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TALK_ROW
    assert _table_header(row) == _TALK_HEADER
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _TALK_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "3"
    assert gen.stats["supplementary_subpoints"] == 1


def test_a_talks_several_prose_lines_are_one_row_of_paragraphs_in_source_order():
    lines = (_TALK_PROSE, "Co-presented with two hospital pharmacists from the regional formulary service")
    gen = _render_talks([_talk(prose="\t".join(lines))], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TALK_ROW
    assert ["".join(t.text for t in p.iter(qn("w:delText"))) for p in row.iter(qn("w:p"))] \
        == list(lines)


@pytest.mark.parametrize("later", [
    # One talk given twice at one venue: the rows differ in their years only.
    _TALK_TITLE,
    # A later talk at the same venue holding every word of this one's title.
    f"Revisiting {_TALK_TITLE}",
])
def test_talk_prose_goes_under_its_own_row_among_talks_of_one_venue(later):
    """The later talk renders first and holds every distinctive word of the
    earlier one's row, so words alone put the earlier one's prose under the
    later row. Only the row the writer recorded for the earlier entry takes
    it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_talks([_later_talk(later), _talk()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert _row_text(above) == _TALK_ROW
    assert _row_text(above.getprevious()) == f"{later}{_TALK_EVENT}, {_TALK_VENUE}2019"
    assert row.getnext() is None
    # Words alone would have taken the later talk's row.
    lines = list(supplementary.anchor_candidates(_render_talks(
        [_later_talk(later), _talk()], supplementary_subpoints=False).doc))
    rows = [line for line in lines if _TALK_VENUE in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_talk(), "R"), rows) is rows[0]
    assert rows[0].text.startswith(later)


def test_talks_filed_under_two_scope_tables_each_take_their_own_prose():
    """A talk the CV files under its International heading renders in that
    table, the other in National: each prose goes under its own talk's row,
    in its own table."""
    abroad_prose = "Keynote for the closing plenary, followed by a workshop for early career faculty"
    abroad = _talk(idx=603, venue="Kestrel Institute, Osprey Bay", prose=abroad_prose)
    abroad["hierarchy"] = ["INVITATIONS TO SPEAK/PRESENT", "International"]
    gen = _render_talks([abroad, _talk()], supplementary_subpoints=True)

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {abroad_prose, _TALK_PROSE}
    assert _row_text(by_text[_TALK_PROSE].getprevious()) == _TALK_ROW
    assert _row_text(by_text[abroad_prose].getprevious()).startswith(_TALK_TITLE)
    assert "Kestrel Institute" in _row_text(by_text[abroad_prose].getprevious())
    assert by_text[_TALK_PROSE].getparent() is not by_text[abroad_prose].getparent()


def test_talk_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_talks([_later_talk(), _talk()], supplementary_subpoints=False)
    on = _render_talks([_later_talk(), _talk()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _TALK_PROSE in rejected
    assert _TALK_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_rows(on)) == 1
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def test_a_short_talks_prose_goes_under_its_own_row():
    """An acronym title at an acronym's venue holds no distinctive word, too
    few for the word test: the writer's record says the one row is the
    entry's own (`_pick_anchor`)."""
    gen = _render_talks([_talk(title="IBD", event=None, venue="NVU")], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == "IBDNVU2016"
    assert _deleted_texts(gen) == _TALK_PROSE


def _untitled_talk():
    """No title: the Title cell is the entry's line without the meeting,
    venue and year the other cells show (`_untitled_talk_title`), so it
    prints the prose too."""
    return _talk(title=None)


def _talk_titled_by_its_prose():
    """The prose is the title stage 4 extracted."""
    entry = _talk(prose=None)
    entry["extracted_fields"]["title"] = _TALK_PROSE
    entry["text"] = f"{_TALK_PROSE}, {_TALK_EVENT}, {_TALK_VENUE}, 2016"
    return entry


@pytest.mark.parametrize("entry", [_untitled_talk, _talk_titled_by_its_prose])
def test_talk_prose_already_in_its_row_is_not_offered_again(entry):
    """The row prints it (`_record_rendered`), so it does not print twice: a
    title derived from the line counts as rendered."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_talks([entry()], supplementary_subpoints=True)
    assert _TALK_PROSE in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "R", haystack) == []


# A date-shaped marker for the date fields, which the Dates cell prints
# only as a date; an end date only after a start.
_DATE_MARKERS = {"year": "1987", "start_date": "1987", "end_date": "1989"}


def _talk_with_only(field, marker):
    entry = _talk(title=None, event=None, venue=None, date=None, prose="")
    entry["extracted_fields"] = {field: marker}
    if field == "end_date":
        entry["extracted_fields"]["start_date"] = "1985"
    entry["text"] = "Example invited lecture"
    return entry


def test_every_talk_row_field_the_pass_counts_is_printed_in_the_talk_row():
    """`supplementary._PRESENTATION_ROW_FIELDS`, the off-schema names among
    the fields the writer reads its cells from (`_RENDERED_FIELDS` marks
    schema fields only): each one, alone on an entry, reaches the row."""
    from unified_pipeline.stage6 import supplementary

    assert supplementary._PRESENTATION_ROW_FIELDS
    for field in supplementary._PRESENTATION_ROW_FIELDS:
        marker = _DATE_MARKERS.get(field, f"Zq{field}Z")
        gen = _render_talks([_talk_with_only(field, marker)], supplementary_subpoints=True)
        assert any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr"))), field


def test_every_field_the_talk_row_prints_is_one_the_pass_counts():
    """The other direction: a field the row prints that `_value_tokens` does
    not count would offer a fragment restating the row. The names are the
    fields of the codes an entry is rerouted into R from (S8, S9, K1-K3,
    Q2, H, D1), each alone on an entry."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    counted = _RENDERED_FIELDS["R"] | set(supplementary._PRESENTATION_ROW_FIELDS)
    names = set().union(*(_RENDERED_FIELDS[code] for code in
                          ("S8", "S9", "K1", "K2", "K3", "Q2", "H", "D1") if code in _RENDERED_FIELDS),
                        supplementary.PRESENTATION_TITLE_FIELDS, supplementary.PRESENTATION_VENUE_FIELDS,
                        supplementary.PRESENTATION_EVENT_FIELDS, {"name", "organization", "position"})
    for field in sorted(names):
        marker = f"Zq{field}Z"
        gen = _render_talks([_talk_with_only(field, marker)], supplementary_subpoints=True)
        printed = any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
        assert not printed or field in counted, field


@pytest.mark.parametrize("field", ["presentation_title", "institution", "venue", "conference_name"])
def test_a_talk_line_mostly_restating_its_row_is_not_offered(field):
    """Stage 4 named the talk's title, venue or meeting under an off-schema
    key, which the row prints when the schema's is empty. A line at least
    half of whose words restate it is the row
    (`SUBPOINT_FIELD_TOKEN_MAX_SHARE`); the document holds too few of its
    words for `_record_rendered` to see that."""
    from unified_pipeline.stage6 import supplementary

    value = "Regional Pharmacy Formulary Review Symposium"
    line = "Formulary Review Symposium for Regional Pharmacy: approved statewide pediatric dosing"
    schema_name = {"presentation_title": "title", "institution": "location",
                   "venue": "location", "conference_name": "event_name"}[field]
    entry = _talk(prose=line)
    del entry["extracted_fields"][schema_name]
    entry["extracted_fields"][field] = value
    gen = _render_talks([entry], supplementary_subpoints=True)
    assert any(value in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary._record_rendered(line, haystack.text, haystack.line_tokens) is False
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


_OTHER_TALK = ("Kestrel Lecture on Estuary Sediment Transport and Marsh Restoration Outcomes, "
               "Osprey Regional Council, Bramble Harbour")


def _talks_in_one_paragraph():
    """Two talks each opening with its year, which the CV wrote in one
    paragraph (EBYSBC BMAMWE 1440's shape): one line, so the writer prints
    one row, of the talk stage 4 extracted; the other is its record, not its
    prose."""
    entry = _talk(prose=None)
    entry["text"] = (f"2016 {_TALK_VENUE}: {_TALK_TITLE}, {_TALK_EVENT} "
                     f"2016 {_OTHER_TALK}, Crab Hollow, ZQ")
    return entry


def _talks_each_ending_in_its_year():
    """Two talks each ending in its year, which stage 2 joined with a tab:
    the writer-side #660 test reads two dated parts."""
    entry = _talk(prose=None)
    entry["text"] = f"{_TALK_TITLE}, {_TALK_EVENT}, {_TALK_VENUE}, 2016\t{_OTHER_TALK}, 2014"
    return entry


def _talks_each_ending_in_its_bracketed_year():
    """The same, each year in brackets, which ends no part for the #660 test."""
    entry = _talk(prose=None)
    entry["text"] = (f"{_TALK_TITLE}, {_TALK_EVENT}, {_TALK_VENUE} (2016)\t{_OTHER_TALK} (2014)"
                     f"\t{_TALK_PROSE}")
    return entry


@pytest.mark.parametrize("make", [_talks_in_one_paragraph, _talks_each_ending_in_its_year,
                                  _talks_each_ending_in_its_bracketed_year])
@pytest.mark.parametrize("flag", [True, False])
def test_a_line_of_several_dated_talks_offers_none_and_its_records_are_recovered(make, flag):
    """#1205 comment 6275741663: the #221 recovery reads every sub-point
    line as rendered, so a talk offered as another's prose would be in the
    accepted document nowhere and still not be recovered. None is offered,
    and the recovery re-emits the same lines with sub-points on as off."""
    entry = make()
    gen = _render_talks([entry], supplementary_subpoints=flag)
    assert len(gen._row_receipts.elements_for(entry)) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""

    gen._recover_unrendered_records({"R": [entry]})
    off = _render_talks([make()], supplementary_subpoints=False)
    off._recover_unrendered_records({"R": [make()]})

    assert gen.stats["unrendered_records_recovered"] == off.stats["unrendered_records_recovered"]
    assert _texts_of(gen.doc.element.body) == _texts_of(off.doc.element.body)


@pytest.mark.parametrize("dates", [
    # Stage 4 kept both dates under the talk's own key ...
    {"date": "2014; 2016"},
    # ... or one under it and the other as a further span (#1245).
    {"date": "2016", "additional_dates": ["2014"]},
])
def test_one_talk_given_on_two_dates_is_one_record_and_offers_its_prose(dates):
    """The same title at two venues, each opening with its year: the Dates
    cell prints both years, so the two dated parts are this talk's own dates,
    not two talks (`_talk_dates`). Its prose goes under its row."""
    entry = _talk(prose=None, **dates)
    entry["text"] = (f"2014 {_TALK_TITLE}, {_TALK_EVENT}, Kestrel Institute\t"
                     f"2016 {_TALK_TITLE}, {_TALK_EVENT}, {_TALK_VENUE}\t{_TALK_PROSE}")
    gen = _render_talks([entry], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert set(_row_texts_of(row.getprevious())[-1].split(", ")) == {"2014", "2016"}
    assert _deleted_texts(gen) == _TALK_PROSE


@pytest.mark.parametrize("dates", [
    {},
    # A further span repeating the talk's own date, which the cell drops.
    {"additional_dates": ["2016"]},
    # Its one date only under a further span, the string "None" its own.
    {"date": None, "year": "None", "additional_dates": ["2016"]},
], ids=["one_date", "span_repeats_its_date", "none_string"])
def test_the_same_line_with_one_date_kept_is_two_talks_and_offers_none(dates):
    """With one date kept, the row prints one occasion: the other dated part
    is a talk the row does not print."""
    entry = _talk(prose=None, **dates)
    entry["text"] = (f"2014 {_TALK_TITLE}, {_TALK_EVENT}, Kestrel Institute\t"
                     f"2016 {_TALK_TITLE}, {_TALK_EVENT}, {_TALK_VENUE}\t{_TALK_PROSE}")
    gen = _render_talks([entry], supplementary_subpoints=True)
    (row,) = gen._row_receipts.elements_for(entry)
    assert _row_texts_of(row)[-1] == "2016"
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _talk_quoting_another_title():
    """Two quoted titles on one undated line: the other talk's title shares
    no word with this one's."""
    entry = _talk(prose=None)
    entry["text"] = (f"“{_TALK_TITLE}”, {_TALK_EVENT}, {_TALK_VENUE}. “Estuary Sediment "
                     f"Transport in Tidal Marshes”, Osprey Regional Council. {_TALK_PROSE}")
    return entry


def _talk_list_kept_whole():
    """Stage 4 listed a second talk under an off-schema key, and fan-out kept
    the list whole because the text holds prose no record does (#983): the
    writer prints the entry's own fields only."""
    entry = _talk()
    entry["extracted_fields"]["additional_presentations"] = [
        {"title": "Estuary Sediment Transport in Tidal Marshes", "location": "Osprey Bay",
         "date": "2014"}]
    return entry


@pytest.mark.parametrize("entry", [_talk_quoting_another_title, _talk_list_kept_whole])
def test_a_talk_row_holding_another_talk_gets_no_sub_point(entry):
    """The writer prints one row, the entry's own, though the line holds
    another talk: that talk's text would be offered as this one's prose,
    under its row."""
    entry = entry()
    gen = _render_talks([entry], supplementary_subpoints=True)

    assert len(gen._row_receipts.elements_for(entry)) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_talk_quoting_its_own_title_differently_punctuated_offers_its_prose():
    entry = _talk(prose=None)
    entry["text"] = f"“{_TALK_TITLE.lower()}.” {_TALK_EVENT}, {_TALK_VENUE}, 2016\t{_TALK_PROSE}"
    gen = _render_talks([entry], supplementary_subpoints=True)
    assert _deleted_texts(gen) == _TALK_PROSE


def test_a_talk_the_section_writes_no_row_for_gets_no_sub_point():
    """An entry no row was written for -- one the section skipped -- has no
    row of its own: its prose goes under no other line of the document,
    though the other talk's row holds every distinctive word of its line."""
    skipped, later = _talk(), _later_talk()
    gen = _generator(supplementary_subpoints=True)
    gen._fill_presentations([later])
    gen._add_supplementary_subpoints({"R": [later, skipped]})

    assert gen._row_receipts.elements_for(skipped) == []
    assert gen._row_receipts.elements_for(later)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _render_grouped_talks(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_presentations(grouped.get("R", []))
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _talk_line(text, *talks):
    """One source line stage 4 returned as a record per (title, venue, date),
    the entry's own fields being the last record's. Fan-out gives every
    earlier record a line of its own and the last one the whole line."""
    records = [{"title": title, "location": venue, "date": date} for title, venue, date in talks]
    return {"element_idx_start": 663, "taxonomy_code": "R", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_ESTUARY = ("Estuary Sediment Transport in Tidal Marshes", "Kestrel Institute", "2014")
_WETLAND = (_TALK_TITLE, _TALK_VENUE, "2016")
_ESTUARY_LINE = "Estuary Sediment Transport in Tidal Marshes, Kestrel Institute, 2014"
_WETLAND_LINE = f"{_TALK_TITLE}, {_TALK_VENUE}, 2016"
_ESTUARY_PROSE = ("Opened the morning session on sediment budgets and led a discussion with "
                  "three regional water managers")


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own talk.
    ([f"{_ESTUARY_LINE}. {_ESTUARY_PROSE}", _WETLAND_LINE],
     [("Estuary Sediment Transport in Tidal Marshes", f"{_ESTUARY_LINE}. {_ESTUARY_PROSE}")]),
    ([_ESTUARY_LINE, f"Invited speaker, {_WETLAND_LINE}. {_TALK_PROSE}"],
     [(_TALK_TITLE, f"Invited speaker, {_WETLAND_LINE}. {_TALK_PROSE}")]),
    # They do not: which talk the prose is about cannot be told, so neither takes it.
    ([_ESTUARY_LINE, _TALK_PROSE, _WETLAND_LINE], []),
])
def test_prose_on_a_line_of_several_talks_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the talk's line as well as its prose brings the title, venue and
    year along, and Reject All prints them a second time, under the row that
    holds them (the shape every family has; #1205, 2026-10-09)."""
    entry = _talk_line("\t".join(segments), _ESTUARY, _WETLAND)
    gen, grouped = _render_grouped_talks([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["R"][-1] and len(grouped["R"]) == 2
    assert [(_row_texts_of(row.getprevious())[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def _row_texts_of(tr):
    return ["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.findall(qn("w:tc"))]


def test_a_talk_line_opening_with_its_title_cell_is_read_as_its_row():
    """A line that opens with the Title cell's whole text holds the row's own
    text verbatim (`_record_rendered`'s first `RENDER_PIECE_WINDOW`
    characters), and the whole line reads as rendered: none of it is offered,
    its prose included. A line opening with the role is offered whole, as the
    CV wrote it."""
    title_first = f"{_TALK_TITLE}, {_TALK_EVENT}, {_TALK_VENUE}, 2016. {_TALK_PROSE}"
    role_first = f"Invited speaker, {title_first}"
    for line, offered in ((title_first, ""), (role_first, role_first)):
        entry = _talk(prose=None)
        entry["text"] = line
        gen = _render_talks([entry], supplementary_subpoints=True)
        assert _deleted_texts(gen) == offered


def test_generic_list_fan_out_children_of_talks_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, the entry renders its own
    fields, and the list is talks the row does not print (`_lists_talks`)."""
    records = [{"title": _ESTUARY[0], "location": _ESTUARY[1], "date": _ESTUARY[2]},
               {"title": _WETLAND[0], "location": _WETLAND[1], "date": _WETLAND[2]}]
    for text, children in ((f"{_ESTUARY_LINE}\t{_WETLAND_LINE}", 2),
                           (f"{_ESTUARY_LINE}\t{_WETLAND_LINE}\t{_TALK_PROSE}", 1)):
        entry = {"element_idx_start": 664, "taxonomy_code": "R", "text": text,
                 "extracted_fields": {"presentations": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_talks([entry], supplementary_subpoints=True)
        assert len(grouped["R"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["R"])
        assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("N1", "MENTORING PROGRAMS"), ("S0", "BIBLIOGRAPHY")])
def test_a_talk_rerouted_out_of_r_to_an_excluded_code_gets_no_sub_point(code, heading):
    """N1 and S0 are excluded: a talk rendered as either is that code's
    record, so neither the talk row of the same title nor any other carries
    its prose."""
    gen, grouped = _render_grouped_talks([_later_talk(), _rerouted(_talk(), code, heading)],
                                         supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [601]
    assert [e["element_idx_start"] for e in grouped["R"]] == [599]
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("C", "POSTDOCTORAL TRAINING"), ("I", "MEMBERSHIPS")])
def test_a_refused_reroute_keeps_the_talk_rule(code, heading):
    """A talk rerouted to C or I keeps its R code: their renderers would drop
    its title (`_fields_lost_in_reroute`). It renders as an R row and its
    prose is a deleted row under that row."""
    gen, grouped = _render_grouped_talks([_rerouted(_talk(), code, heading)],
                                         supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["R"]] == [601]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TALK_ROW


def _abstract_under_invited_talks():
    """An abstract stage 3b coded S8, under the invited-talks heading: its
    meeting is S8's `conference_name`, which the row reads (#1570)."""
    return _talk(code="S8", event=None, conference_name=_TALK_EVENT, authors="Marblegate Q")


@pytest.mark.parametrize("make", [_abstract_under_invited_talks, lambda: _talk(code="K3")])
def test_a_record_rerouted_into_r_follows_the_talk_rule(make):
    """An abstract coded S8, or a lecture coded K3, under the invited-talks
    heading renders as an R row: its prose is a deleted row under that row."""
    gen, grouped = _render_grouped_talks(
        [_rerouted(make(), "R", "INVITATIONS TO SPEAK/PRESENT")], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["R"]] == [601]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TALK_ROW
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _TALK_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.code == "R"


def test_a_talk_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _talk()
    withheld["_pii_withheld"] = True
    gen = _render_talks([_later_talk(), withheld], supplementary_subpoints=True)

    assert any(_row_text(tr) == _TALK_ROW for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_talk_dedup_drops_offers_its_prose_once_under_the_kept_row():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no row, so its
    prose goes only under the row of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_talk(idx=601), _talk(idx=602)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["R"]) == 2 and len(kept["R"]) == 1
    gen._fill_presentations(kept["R"])
    gen._add_supplementary_subpoints(before)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TALK_ROW
    assert _deleted_texts(gen) == _TALK_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["R"][0]["element_idx_start"]


def test_talk_and_honor_prose_each_stay_under_their_own_section_row():
    """An award for a talk can share its title and venue: each one's prose
    goes under its own section's row, never under the other's."""
    honor_prose = ("Awarded each year to the best lecture given by an early career speaker "
                   "at the regional meeting")
    honor = {"element_idx_start": 39, "taxonomy_code": "H",
             "text": f"2016\tBest Lecture Award for {_TALK_TITLE}, {_TALK_EVENT}\t{honor_prose}",
             "extracted_fields": {"award_name": f"Best Lecture Award for {_TALK_TITLE}",
                                  "granting_body": _TALK_EVENT, "date": "2016"}}
    talk = _talk()
    gen = _generator(supplementary_subpoints=True)
    gen._fill_honors([honor])
    gen._fill_presentations([talk])
    gen._add_supplementary_subpoints({"H": [honor], "R": [talk]})

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {honor_prose, _TALK_PROSE}
    assert _table_header(by_text[_TALK_PROSE]) == _TALK_HEADER
    assert _table_header(by_text[honor_prose]) != _TALK_HEADER
    assert by_text[_TALK_PROSE].getparent() is not by_text[honor_prose].getparent()


def test_the_doctor_matches_a_talk_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier talk's sub-point
    under its own index and R, and the doctor's per-entry match counts it for
    that entry, not for the later talk that shares its words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_talk(), _talk()], uid="T1205R")
    assert placements == [[_TALK_PROSE]]
    sidecar = json.loads((tmp_path / "T1205R_render_warnings.json").read_text())
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(601, "R")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_TALK_PROSE, subpoints, element_idx=601)
    assert not offered_as_subpoint(_TALK_PROSE, subpoints, element_idx=599)


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_talk_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line (#1205 review of O). The same talk without them still has its
    sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_talk()])
    assert placements == [[_TALK_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_talk(), _TALK_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


# --- #1205: C postdoctoral training, the next code family --------------------
# The postdoctoral-training section writes every C entry as one row of its
# three-column table, so the prose is a deleted row under the position's own
# row: the one row the writer recorded for that entry (the row-receipt
# ledger). A line holding several dated records offers none.

_TRAINING_PROSE = ("Program director: Doctor Bramwell Fenn; rotations in paediatric surgery, "
                   "neonatal intensive care and transport medicine")
_TRAINING_TYPE = "Example Pharmacy Residency"
_TRAINING_SITE = "Norvale General Hospital"


def _training(idx=113, ttype=_TRAINING_TYPE, site=_TRAINING_SITE, start="2009-07", end="2010-06",
              prose=_TRAINING_PROSE, code="C", **fields):
    """A residency whose source line goes on past its type, hospital and
    dates into who directed it, as EBYSBC TXTATQ 113's did. The next year's
    residency at the same hospital shares every distinctive word of its row:
    only the dates, digits, tell them apart. A value passed as None is in
    neither the line nor the fields."""
    named = {"training_type": ttype, "institution": site, "start_date": start, "end_date": end}
    dates = "-".join(value for value in (start, end) if value)
    line = "\t".join(value for value in (ttype, site, dates) if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_training(ttype=_TRAINING_TYPE):
    """The next year's residency at the same hospital, with no prose: it
    renders first, reverse chronologically."""
    return _training(idx=108, ttype=ttype, start="2010-07", end="2011-06", prose="")


def _render_training(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_postdoc_training({"C": entries}, entries)
    gen._add_supplementary_subpoints({"C": entries})
    return gen


def _training_table_row():
    """The same residency read from a source table: stage 2 joins the row's
    cells with pipes, and its position names the row ("29.2")."""
    entry = _training()
    entry["element_idx_start"] = "29.2"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


_TRAINING_ROW = f"{_TRAINING_TYPE}{_TRAINING_SITE}07/09-06/10"
_LATER_TRAINING_ROW = f"{_TRAINING_TYPE}{_TRAINING_SITE}07/10-06/11"
_TRAINING_HEADER = "Title, include area of trainingInstitution, city and stateDates (mm/yy - mm/yy)"


@pytest.mark.parametrize("entry", [_training, _training_table_row])
def test_training_prose_is_a_deleted_row_under_the_position_row(entry):
    gen = _render_training([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TRAINING_ROW
    assert _table_header(row) == _TRAINING_HEADER
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _TRAINING_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "3"
    assert gen.stats["supplementary_subpoints"] == 1
    (placement,) = gen._subpoint_placements
    assert (placement.element_idx_start, placement.code) == (entry()["element_idx_start"], "C")


def test_a_positions_several_prose_lines_are_one_row_of_paragraphs_in_source_order():
    lines = (_TRAINING_PROSE, "Research topic: nutrient cycling in tidal marsh sediments and their microbes")
    gen = _render_training([_training(prose="\t".join(lines))], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TRAINING_ROW
    assert ["".join(t.text for t in p.iter(qn("w:delText"))) for p in row.iter(qn("w:p"))] \
        == list(lines)


@pytest.mark.parametrize("later", [
    # One title held twice at one hospital: the rows differ in their dates only.
    _TRAINING_TYPE,
    # A later combined position holding every word of this one's type
    # (EBYSBC MOFVYH 43, a residency under a residency-and-fellowship row).
    f"{_TRAINING_TYPE} and Example Critical Care Fellowship",
])
def test_training_prose_goes_under_its_own_row_among_positions_at_one_hospital(later):
    """The later position renders first and holds every distinctive word of
    the earlier one's row, so words alone put the earlier one's prose under
    the later row. Only the row the writer recorded for the earlier entry
    takes it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_training([_later_training(later), _training()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert _row_text(above) == _TRAINING_ROW
    assert _row_text(above.getprevious()) == f"{later}{_TRAINING_SITE}07/10-06/11"
    # Words alone would have taken the later position's row.
    lines = list(supplementary.anchor_candidates(_render_training(
        [_later_training(later), _training()], supplementary_subpoints=False).doc))
    rows = [line for line in lines if _TRAINING_SITE in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_training(), "C"), rows) is rows[0]
    assert rows[0].text.startswith(later)


def test_positions_at_one_hospital_each_take_their_own_prose():
    """Two residencies at one hospital, each with its program director
    (EBYSBC TXTATQ 108 and 113): each prose goes under its own row."""
    later_prose = ("Program director: Doctor Ottoline Marsh; rotations in critical care "
                   "pharmacotherapy and toxicology")
    later = _training(idx=108, start="2010-07", end="2011-06", prose=later_prose)
    gen = _render_training([later, _training()], supplementary_subpoints=True)

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {later_prose, _TRAINING_PROSE}
    assert _row_text(by_text[later_prose].getprevious()) == _LATER_TRAINING_ROW
    assert _row_text(by_text[_TRAINING_PROSE].getprevious()) == _TRAINING_ROW
    assert [(p.element_idx_start, p.paragraphs) for p in gen._subpoint_placements] \
        == [(108, (later_prose,)), (113, (_TRAINING_PROSE,))]


def test_training_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_training([_later_training(), _training()], supplementary_subpoints=False)
    on = _render_training([_later_training(), _training()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _TRAINING_PROSE in rejected
    assert _TRAINING_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_rows(on)) == 1
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def test_a_short_positions_prose_goes_under_its_own_row():
    """A type and a hospital of acronyms hold no distinctive word, too few
    for the word test: the writer's record says the one row is the entry's
    own (`_pick_anchor`)."""
    gen = _render_training([_training(ttype="PGY-1", site="NGH")], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == "PGY-1NGH07/09-06/10"
    assert _deleted_texts(gen) == _TRAINING_PROSE


def _position_typed_by_its_prose():
    """The prose is the type stage 4 extracted."""
    return _training(ttype=_TRAINING_PROSE, prose=None)


def _position_with_its_prose_as_role():
    """The prose is the role stage 4 extracted, which the type cell folds in."""
    return _training(role=_TRAINING_PROSE)


@pytest.mark.parametrize("entry", [_position_typed_by_its_prose, _position_with_its_prose_as_role])
def test_training_prose_already_in_its_row_is_not_offered_again(entry):
    """The row prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_training([entry()], supplementary_subpoints=True)
    assert _TRAINING_PROSE in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "C", haystack) == []


def _position_with_only(field, marker):
    entry = _training(ttype=None, site=None, start=None, end=None, prose="")
    entry["extracted_fields"] = {field: marker}
    entry["text"] = "Example training position"
    return entry


def test_every_training_row_field_the_pass_counts_is_printed_in_the_training_row():
    """The off-schema names among the fields `_value_tokens` counts for a
    training position (`_RENDERED_FIELDS` marks schema fields only): each
    one, alone on an entry, is counted and reaches the row."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    off_schema = [name for name in (*supplementary.POSTDOC_TYPE_FIELDS,
                                    *supplementary.POSTDOC_FIELD_OF_STUDY_FIELDS,
                                    *supplementary._POSTDOC_PLAIN_FIELDS)
                  if name not in _RENDERED_FIELDS["C"]]
    assert off_schema == ["title", "field_of_study", "location"]
    for field in off_schema:
        marker = f"Zq{field}Z"
        entry = _position_with_only(field, marker)
        assert supplementary._tokens(marker) <= supplementary._value_tokens(entry, "C"), field
        gen = _render_training([entry], supplementary_subpoints=True)
        assert any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr"))), field


def test_every_field_the_training_row_prints_is_one_the_pass_counts():
    """The other direction: a field the row prints that `_value_tokens` does
    not count would offer a fragment restating the row. The names are the
    fields of the codes an entry is rerouted into C from (B1, O), of the
    codes stage 3b confuses with it (D1, D2, K2, N3B), and the writer's
    own, each alone on an entry."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    names = set().union(*(_RENDERED_FIELDS[code] for code in
                          ("B1", "O", "D1", "D2", "K2", "N3B") if code in _RENDERED_FIELDS),
                        supplementary.POSTDOC_TYPE_FIELDS, supplementary.POSTDOC_FIELD_OF_STUDY_FIELDS,
                        {"location", "department", "mentor", "organization", "position", "name"})
    for field in sorted(names):
        marker = f"Zq{field}Z"
        entry = _position_with_only(field, marker)
        gen = _render_training([entry], supplementary_subpoints=True)
        printed = any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
        counted = supplementary._tokens(marker) <= supplementary._value_tokens(entry, "C")
        assert not printed or counted, field


_SHADOWED_ALIAS = "Saltmarsh restoration and estuarine sediment track"


@pytest.mark.parametrize("printed, shadowed", [("training_type", "title"),
                                               ("field_of_study", "specialty")])
def test_a_line_restating_the_alias_its_row_does_not_print_is_offered(printed, shadowed):
    """A cell prints the first of its aliases holding a value
    (`first_field_value`). The other one's words are in no line of the
    document, so a line restating it is the position's prose, not its row."""
    entry = _training(prose=_SHADOWED_ALIAS, **{printed: "Example Geochemistry",
                                                 shadowed: _SHADOWED_ALIAS})
    gen = _render_training([entry], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert "Example Geochemistry" in _row_text(row.getprevious())
    assert _SHADOWED_ALIAS not in _row_text(row.getprevious())
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _SHADOWED_ALIAS


def _positions_each_opening_with_their_dates():
    """A residency and the fellowship after it on one line, each opening
    with its dates and followed by its prose; stage 4 returned the
    fellowship. Which one a prose line is about cannot be told."""
    return _training(start="2010-07", end="2012-06", prose=None) | {
        "text": (f"2007-2010 {_TRAINING_TYPE}, {_TRAINING_SITE}\t{_TRAINING_PROSE}\t"
                 "2010-2012 Example Critical Care Fellowship, Norvale General Hospital\t"
                 "Research topic: nutrient cycling in tidal marsh sediments and their microbes")}


def _a_position_then_the_next_dated_at_its_end():
    """A residency, then on the same line the fellowship after it, its
    dates ending its part; stage 4 returned the residency (#1205 review)."""
    return _training(prose=None) | {
        "text": (f"{_TRAINING_TYPE}, {_TRAINING_SITE}, 2009-2010\t"
                 "Example Critical Care Fellowship, Kestrel Institute, 2010-2013")}


def _a_position_then_the_next_dated_in_months():
    """The same, the fellowship's dates in months."""
    return _training(prose=None) | {
        "text": (f"{_TRAINING_TYPE}, {_TRAINING_SITE}, 2009-2010\t"
                 "Example Critical Care Fellowship, Kestrel Institute, July 2010 - June 2013")}


@pytest.mark.parametrize("positions", [_positions_each_opening_with_their_dates,
                                       _a_position_then_the_next_dated_at_its_end,
                                       _a_position_then_the_next_dated_in_months])
@pytest.mark.parametrize("flag", [True, False])
def test_a_line_of_several_dated_positions_offers_none_and_its_records_are_recovered(flag, positions):
    """#1205 comment 6275741663: the #221 recovery reads every sub-point
    line as rendered, so a position offered as another's prose would be in
    the accepted document nowhere and still not be recovered. None is
    offered, and the recovery re-emits the same lines with sub-points on as
    off, wherever the other position's part puts its dates."""
    entry = positions()
    gen = _render_training([entry], supplementary_subpoints=flag)
    assert len(gen._row_receipts.elements_for(entry)) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""

    gen._recover_unrendered_records({"C": [entry]})
    off = _render_training([positions()], supplementary_subpoints=False)
    off._recover_unrendered_records({"C": [positions()]})

    assert gen.stats["unrendered_records_recovered"] == off.stats["unrendered_records_recovered"]
    assert _texts_of(gen.doc.element.body) == _texts_of(off.doc.element.body)


def test_a_position_the_section_writes_no_row_for_gets_no_sub_point():
    """An entry no row was written for -- one the section skipped -- has no
    row of its own: its prose goes under no other line of the document,
    though the other position's row holds every distinctive word of its
    line."""
    skipped, later = _training(), _later_training()
    gen = _generator(supplementary_subpoints=True)
    gen._fill_postdoc_training({"C": [later]}, [later])
    gen._add_supplementary_subpoints({"C": [later, skipped]})

    assert gen._row_receipts.elements_for(skipped) == []
    assert gen._row_receipts.elements_for(later)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_the_training_span_starts_at_the_heading_its_writer_looks_up():
    """The span starts at the heading the writer looks its table up by
    (`POSTDOC_TRAINING_HEADING`), and holds that table. The other callers of
    the code's heading lookup still name none for C: the reconsider pass and
    the #221 recovery keep a C record in the Appendix, as before C's
    sub-points."""
    gen = _render_training([_training()], supplementary_subpoints=True)
    span = gen._subpoint_section_spans()["C"]
    heading = gen.doc.element.body[span.start]
    assert "".join(t.text or "" for t in heading.iter(qn("w:t"))).startswith("POSTDOCTORAL TRAINING")
    (row,) = _subpoint_rows(gen)
    assert span.start < list(gen.doc.element.body).index(row.getparent()) < span.end
    assert WCMTemplateGenerator._get_wcm_section_header("C") == ""
    assert gen._insert_reconsidered_segment("Example Pharmacy Residency, 2009-2010", "C") is False


def test_the_training_span_is_not_an_earlier_heading_naming_training():
    """A heading that names training before the postdoctoral one (a training
    grants table) is neither the writer's table nor the span's start."""
    gen = _generator(supplementary_subpoints=True)
    gen.doc = Document()
    gen.doc.add_paragraph("INSTITUTIONAL TRAINING GRANTS")
    gen.doc.add_table(rows=1, cols=3)
    heading = gen.doc.add_paragraph("C. POSTDOCTORAL TRAINING")
    table = gen.doc.add_table(rows=1, cols=3)
    entry = _training()
    gen._fill_postdoc_training({"C": [entry]}, [entry])
    gen._add_supplementary_subpoints({"C": [entry]})

    body = list(gen.doc.element.body)
    assert body[gen._subpoint_section_spans()["C"].start] is heading._p
    (row,) = _subpoint_rows(gen)
    assert row.getparent() is table._tbl


def test_a_template_with_no_training_heading_places_no_training_prose():
    """Without the heading neither the writer nor the span finds the
    section: no row is written, and the prose goes under no other line."""
    gen = _generator(supplementary_subpoints=True)
    gen.doc = Document()
    gen.doc.add_paragraph("A. PERSONAL DATA")
    entry = _training()
    gen._fill_postdoc_training({"C": [entry]}, [entry])
    gen._add_supplementary_subpoints({"C": [entry]})

    assert "C" not in gen._subpoint_section_spans()
    assert _deleted_texts(gen) == ""


def _render_grouped_training(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_postdoc_training(grouped, [e for group in grouped.values() for e in group])
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _training_line(text, *positions):
    """One source line stage 4 returned as a record per (type, hospital,
    start), the entry's own fields being the last record's. Fan-out gives
    every earlier record a line of its own and the last one the whole line."""
    records = [{"training_type": ttype, "institution": site, "start_date": start}
               for ttype, site, start in positions]
    return {"element_idx_start": 663, "taxonomy_code": "C", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_INTERNSHIP = ("Example Surgical Internship", "Kestrel Institute", "2007")
_RESIDENCY = (_TRAINING_TYPE, _TRAINING_SITE, "2009")
_INTERNSHIP_LINE = "Example Surgical Internship, Kestrel Institute, 2007"
_RESIDENCY_LINE = f"{_TRAINING_TYPE}, {_TRAINING_SITE}, 2009"
_INTERNSHIP_PROSE = ("Rotated through the trauma, burns and vascular services under the "
                     "chief of surgery")


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own position.
    ([f"{_INTERNSHIP_LINE}. {_INTERNSHIP_PROSE}", _RESIDENCY_LINE],
     [("Example Surgical Internship", f"{_INTERNSHIP_LINE}. {_INTERNSHIP_PROSE}")]),
    ([_INTERNSHIP_LINE, f"Chief resident, {_RESIDENCY_LINE}. {_TRAINING_PROSE}"],
     [(_TRAINING_TYPE, f"Chief resident, {_RESIDENCY_LINE}. {_TRAINING_PROSE}")]),
    # They do not: which position the prose is about cannot be told, so neither takes it.
    ([_INTERNSHIP_LINE, _TRAINING_PROSE, _RESIDENCY_LINE], []),
])
def test_prose_on_a_line_of_several_positions_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the position's line as well as its prose brings the type, hospital
    and year along, and Reject All prints them a second time, under the row
    that holds them (the shape every family has; #1205, 2026-10-09)."""
    entry = _training_line("\t".join(segments), _INTERNSHIP, _RESIDENCY)
    gen, grouped = _render_grouped_training([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["C"][-1] and len(grouped["C"]) == 2
    assert [(_row_texts_of(row.getprevious())[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def test_generic_list_fan_out_children_of_positions_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, the entry renders its own
    fields, and the list is positions the row does not print
    (`_own_training_line`): their lines would be offered as its prose."""
    records = [{"training_type": _INTERNSHIP[0], "institution": _INTERNSHIP[1],
                "start_date": _INTERNSHIP[2]},
               {"training_type": _RESIDENCY[0], "institution": _RESIDENCY[1],
                "start_date": _RESIDENCY[2]}]
    for text, children in ((f"{_INTERNSHIP_LINE}\t{_RESIDENCY_LINE}", 2),
                           (f"{_INTERNSHIP_LINE}\t{_RESIDENCY_LINE}\t{_TRAINING_PROSE}", 1)):
        entry = {"element_idx_start": 664, "taxonomy_code": "C", "text": text,
                 "extracted_fields": {"positions": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_training([entry], supplementary_subpoints=True)
        assert len(grouped["C"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["C"])
        assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("N1", "MENTORING PROGRAMS"), ("S0", "BIBLIOGRAPHY")])
def test_a_position_rerouted_out_of_c_to_an_excluded_code_gets_no_sub_point(code, heading):
    """N1 and S0 are excluded: a position rendered as either is that code's
    record, so neither the training row of the same hospital nor any other
    carries its prose."""
    gen, grouped = _render_grouped_training([_later_training(), _rerouted(_training(), code, heading)],
                                            supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [113]
    assert [e["element_idx_start"] for e in grouped["C"]] == [108]
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("I", "PROFESSIONAL ORGANIZATIONS"),
                                           ("F2", "Board Certification"), ("D1", "ACADEMIC APPOINTMENTS")])
def test_a_refused_reroute_keeps_the_training_rule(code, heading):
    """A position rerouted to I, to F2, whose span the heading lookup does
    not find, or to D1 keeps its C code: their renderers would drop its type
    (`_fields_lost_in_reroute`). It renders as a C row and its prose is a
    deleted row under that row."""
    gen, grouped = _render_grouped_training([_rerouted(_training(), code, heading)],
                                            supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["C"]] == [113]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TRAINING_ROW


def _residency_coded_as_a_role():
    """A residency stage 3b coded O, naming only its hospital and start."""
    return _training(code="O", ttype=None, end=None)


def _residency_coded_as_a_degree():
    """A residency stage 3b coded B1, naming no degree."""
    return _training(code="B1", ttype=None, end=None)


@pytest.mark.parametrize("make", [_residency_coded_as_a_role, _residency_coded_as_a_degree])
def test_a_record_rerouted_into_c_follows_the_training_rule(make):
    """A residency under the postdoctoral-training heading that stage 3b
    coded O or B1 renders as a C row: its prose is a deleted row under that
    row, recorded as C."""
    gen, grouped = _render_grouped_training(
        [_rerouted(make(), "C", "POSTDOCTORAL TRAINING")], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["C"]] == [113]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()).startswith("Postdoctoral")
    assert _TRAINING_SITE in _row_text(row.getprevious())
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _TRAINING_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.code == "C"


def test_a_position_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _training()
    withheld["_pii_withheld"] = True
    gen = _render_training([_later_training(), withheld], supplementary_subpoints=True)

    assert any(_row_text(tr) == _TRAINING_ROW for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_position_dedup_drops_offers_its_prose_once_under_the_kept_row():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no row, so its
    prose goes only under the row of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_training(idx=113), _training(idx=114)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["C"]) == 2 and len(kept["C"]) == 1
    gen._fill_postdoc_training(kept, kept["C"])
    gen._add_supplementary_subpoints(before)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _TRAINING_ROW
    assert _deleted_texts(gen) == _TRAINING_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["C"][0]["element_idx_start"]


def test_a_low_coverage_position_with_a_sub_point_is_not_also_sent_whole_to_the_appendix():
    """The content-overflow pass skips an entry whose prose is a sub-point.
    It reads the entry the row's comment hook recorded, so the writer hands
    that hook the entry stage 6 grouped, not a copy of it."""
    long_prose = "\t".join(f"{_TRAINING_PROSE}, block {n} of the rotation schedule" for n in range(9))
    entry = _training(prose=long_prose)
    entry["extraction_coverage"] = {"extraction_coverage_percent": 10.0}
    assert len(entry["text"]) > 1000
    gen = _generator(supplementary_subpoints=True)
    gen._fill_postdoc_training({"C": [entry]}, [entry])
    assert [queued for queued, _para, _code in gen._overflow_entries] == [entry]
    assert gen._overflow_entries[0][0] is entry
    gen._place_unrendered_prose({"C": [entry]})

    assert len(_subpoint_rows(gen)) == 1
    assert gen._appendix_pending == []


def test_the_doctor_matches_a_training_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier residency's
    sub-point under its own index and C, and the doctor's per-entry match
    counts it for that entry, not for the later residency that shares its
    words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_training(), _training()],
                                     uid="T1205C")
    assert placements == [[_TRAINING_PROSE]]
    sidecar = json.loads((tmp_path / "T1205C_render_warnings.json").read_text())
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(113, "C")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_TRAINING_PROSE, subpoints, element_idx=113)
    assert not offered_as_subpoint(_TRAINING_PROSE, subpoints, element_idx=108)


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_training_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line (#1205 review of O). The same position without them still has its
    sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_training()])
    assert placements == [[_TRAINING_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_training(), _TRAINING_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


# --- #1205: I professional memberships, the next code family -----------------
# The memberships section writes every I entry as one row of its two-column
# table, so the prose is a deleted row under the membership's own row: the one
# row the writer recorded for that entry (the row-receipt ledger). A line
# holding another membership offers none, and a dated office the #221
# recovery reads as one of the entry's records is left to that recovery.

_MEMBERSHIP_PROSE = ("Chaired the scientific session on estuarine sediment budgets at the "
                     "annual meeting and edited its proceedings")
_SOCIETY_NAME = "Example Society of Estuarine Ecology"


def _membership(idx=135, org=_SOCIETY_NAME, mtype="Member", start="2009", end="2014",
                prose=_MEMBERSHIP_PROSE, code="I", **fields):
    """A membership whose source line goes on past its type, society and
    years into a session the member chaired, as EBYSBC BFSUMA 98's did. A
    later membership of the same society and type shares every distinctive
    word of its row: only the years, digits, tell them apart. A value passed
    as None is in neither the line nor the fields."""
    named = {"organization": org, "membership_type": mtype, "start_date": start, "end_date": end}
    dates = "-".join(value for value in (start, end) if value)
    line = ", ".join(value for value in (mtype, org, dates) if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_membership(mtype="Member"):
    """The same society's membership from 2015, with no prose: it renders
    first, reverse chronologically."""
    return _membership(idx=122, mtype=mtype, start="2015", end="present", prose="")


def _render_memberships(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_memberships(entries)
    gen._add_supplementary_subpoints({"I": entries})
    return gen


def _membership_table_row():
    """The same membership read from a source table: stage 2 joins the row's
    cells with pipes, and its position names the row ("88.2")."""
    entry = _membership()
    entry["element_idx_start"] = "88.2"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


_MEMBERSHIP_ROW = f"Member, {_SOCIETY_NAME}2009-2014"
_LATER_MEMBERSHIP_ROW = f"Member, {_SOCIETY_NAME}2015-Present"
_MEMBERSHIP_HEADER = "OrganizationDate (yyyy-yyyy)"


@pytest.mark.parametrize("entry", [_membership, _membership_table_row])
def test_membership_prose_is_a_deleted_row_under_the_membership_row(entry):
    gen = _render_memberships([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _MEMBERSHIP_ROW
    assert _table_header(row) == _MEMBERSHIP_HEADER
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _MEMBERSHIP_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "2"
    assert gen.stats["supplementary_subpoints"] == 1
    (placement,) = gen._subpoint_placements
    assert (placement.element_idx_start, placement.code) == (entry()["element_idx_start"], "I")


def test_a_memberships_several_prose_lines_are_one_row_of_paragraphs_in_source_order():
    lines = (_MEMBERSHIP_PROSE, "Delegate for the northern region to the annual assembly of members")
    gen = _render_memberships([_membership(prose="\t".join(lines))], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _MEMBERSHIP_ROW
    assert ["".join(t.text for t in p.iter(qn("w:delText"))) for p in row.iter(qn("w:p"))] \
        == list(lines)


@pytest.mark.parametrize("later", [
    # One type held twice in one society: the rows differ in their years only.
    "Member",
    # A later fellowship of the same society, whose row holds every word of this one's.
    "Elected Member",
])
def test_membership_prose_goes_under_its_own_row_among_memberships_of_one_society(later):
    """The later membership renders first and holds every distinctive word
    of the earlier one's row, so words alone put the earlier one's prose
    under the later row. Only the row the writer recorded for the earlier
    entry takes it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_memberships([_later_membership(later), _membership()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    above = row.getprevious()
    assert _row_text(above) == _MEMBERSHIP_ROW
    assert _row_text(above.getprevious()) == f"{later}, {_SOCIETY_NAME}2015-Present"
    # Words alone would have taken the later membership's row.
    lines = list(supplementary.anchor_candidates(_render_memberships(
        [_later_membership(later), _membership()], supplementary_subpoints=False).doc))
    rows = [line for line in lines if _SOCIETY_NAME in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_membership(), "I"), rows) is rows[0]
    assert rows[0].text.startswith(later)


def test_memberships_of_one_society_each_take_their_own_prose():
    later_prose = "Organized the society's first regional workshop on tidal marsh restoration methods"
    later = _membership(idx=122, start="2015", end="present", prose=later_prose)
    gen = _render_memberships([later, _membership()], supplementary_subpoints=True)

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {later_prose, _MEMBERSHIP_PROSE}
    assert _row_text(by_text[later_prose].getprevious()) == _LATER_MEMBERSHIP_ROW
    assert _row_text(by_text[_MEMBERSHIP_PROSE].getprevious()) == _MEMBERSHIP_ROW
    assert [(p.element_idx_start, p.paragraphs) for p in gen._subpoint_placements] \
        == [(122, (later_prose,)), (135, (_MEMBERSHIP_PROSE,))]


def test_membership_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_memberships([_later_membership(), _membership()], supplementary_subpoints=False)
    on = _render_memberships([_later_membership(), _membership()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _MEMBERSHIP_PROSE in rejected
    assert _MEMBERSHIP_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_rows(on)) == 1
    for row in _subpoint_rows(on):
        row.getparent().remove(row)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


def test_a_short_memberships_prose_goes_under_its_own_row():
    """A society of an acronym holds no distinctive word, too few for the
    word test: the writer's record says the one row is the entry's own
    (`_pick_anchor`)."""
    gen = _render_memberships([_membership(org="ESEE")], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == "Member, ESEE2009-2014"
    assert _deleted_texts(gen) == _MEMBERSHIP_PROSE


def _membership_named_by_its_prose():
    """The prose is the organization stage 4 extracted."""
    return _membership(org=_MEMBERSHIP_PROSE, prose=None)


def _membership_typed_by_its_prose():
    """The prose is the membership type stage 4 extracted, which the
    Organization cell prefixes."""
    return _membership(mtype=_MEMBERSHIP_PROSE)


@pytest.mark.parametrize("entry", [_membership_named_by_its_prose, _membership_typed_by_its_prose])
def test_membership_prose_already_in_its_row_is_not_offered_again(entry):
    """The row prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_memberships([entry()], supplementary_subpoints=True)
    assert _MEMBERSHIP_PROSE in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "I", haystack) == []


def _membership_with_only(field, marker, text="Kestrel Ornithology Union"):
    """An entry holding `field` alone. Its text names an organization the
    writer's fallback scan reads, so a type has an organization to prefix;
    an organization name is its whole text, as the fallback prints a name
    only when the entry's words all reach the row (`_recovery_covers_text`)."""
    entry = _membership(org=None, mtype=None, start=None, end=None, prose="")
    entry["extracted_fields"] = {field: marker}
    entry["text"] = text
    return entry


def test_every_membership_row_field_the_pass_counts_is_printed_in_the_membership_row():
    """The off-schema names among the fields `_value_tokens` counts for a
    membership (`_RENDERED_FIELDS` marks schema fields only): each one, alone
    on an entry, is counted and reaches the row."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    off_schema = [name for name in (*supplementary.MEMBERSHIP_ORGANIZATION_FIELDS,
                                    *supplementary.MEMBERSHIP_TYPE_FIELDS)
                  if name not in _RENDERED_FIELDS["I"]]
    assert off_schema == ["institution", "organization_name", "society", "fellowship_designation", "role"]
    for field in off_schema:
        marker = f"Zq{field}Z"
        names_organization = field in supplementary.MEMBERSHIP_ORGANIZATION_FIELDS
        entry = _membership_with_only(field, marker, *([marker] if names_organization else []))
        assert supplementary._tokens(marker) <= supplementary._value_tokens(entry, "I"), field
        gen = _render_memberships([entry], supplementary_subpoints=True)
        assert any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr"))), field


def test_every_field_the_membership_row_prints_is_one_the_pass_counts():
    """The other direction: a field the row prints that `_value_tokens` does
    not count would offer a fragment restating the row. The names are the
    fields of the codes an entry is rerouted into I from (Q1, Q2), of the
    codes stage 3b confuses with it (H, P, O), and the writer's own, each
    alone on an entry, and beside the schema's organization."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    names = set().union(*(_RENDERED_FIELDS[code] for code in ("Q1", "Q2", "H", "P", "O")
                          if code in _RENDERED_FIELDS),
                        supplementary.MEMBERSHIP_ORGANIZATION_FIELDS, supplementary.MEMBERSHIP_TYPE_FIELDS,
                        {"title", "position", "name", "department", "committee"})
    for field in sorted(names):
        marker = f"Zq{field}Z"
        for entry in (_membership_with_only(field, marker),
                      _membership(prose="", **{field: marker})):
            gen = _render_memberships([entry], supplementary_subpoints=True)
            printed = any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
            counted = supplementary._tokens(marker) <= supplementary._value_tokens(entry, "I")
            assert not printed or counted, field


@pytest.mark.parametrize("alias, mtype", [("institution", "Member"), ("society", "Member"),
                                          ("role", "Member"), ("role", None)])
def test_a_line_restating_a_name_its_row_does_not_print_is_offered(alias, mtype):
    """Beside the schema's organization the row prints no off-schema name
    (`membership_cell_values`), not even a role when the schema names no
    type. Its words are in no line of the document, so a line restating it
    is the membership's prose, not its row."""
    restated = "Coastal sediment monitoring network steering working party"
    entry = _membership(mtype=mtype, prose=restated, **{alias: restated})
    gen = _render_memberships([entry], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert restated not in _row_text(row.getprevious())
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == restated


def test_a_membership_the_writer_parses_into_several_rows_gets_no_sub_point():
    """A line holding several memberships, one per source table row, which
    the writer prints a row each (#476): which of them a prose line is about
    cannot be shown."""
    entry = _membership(prose=None)
    entry["text"] = (f"Member | {_SOCIETY_NAME} | 2009-2014\n"
                     "Fellow | Kestrel Academy of Tidal Studies | 2012-2016\n"
                     f"{_MEMBERSHIP_PROSE}")
    gen = _render_memberships([entry], supplementary_subpoints=True)

    assert len(gen._row_receipts.elements_for(entry)) > 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _memberships_in_one_sentence():
    """One sentence listing the society among others (EBYSBC NJIKGI 65's
    shape): the writer prints the one stage 4 extracted."""
    entry = _membership(prose=None)
    entry["text"] = (f"Member of the {_SOCIETY_NAME}, the Kestrel Statistical Association, and the "
                     "Norvale Academy of Tidal Studies, contributing to their annual meetings")
    return entry


def _a_chapter_fused_with_another_college():
    """A chapter's line fused with another college's fellowship (YUYVIG
    FCAAUV 30's shape)."""
    entry = _membership(org="Theta Chapter, Example Honor Guild", start="1986", end=None, prose=None)
    entry["text"] = ("Fellow Kestrel College of Physicians, Theta Chapter, Example Honor Guild 1986\t"
                     f"{_MEMBERSHIP_PROSE}")
    return entry


@pytest.mark.parametrize("make", [_memberships_in_one_sentence, _a_chapter_fused_with_another_college])
@pytest.mark.parametrize("flag", [True, False])
def test_a_membership_line_naming_another_society_offers_none_and_its_records_are_recovered(make, flag):
    """#1205 comment 6275741663: the #221 recovery reads every sub-point
    line as rendered, so another membership offered as this one's prose
    would be in the accepted document nowhere. None is offered, and the
    recovery re-emits the same lines with sub-points on as off."""
    entry = make()
    gen = _render_memberships([entry], supplementary_subpoints=flag)
    assert len(gen._row_receipts.elements_for(entry)) == 1
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""

    gen._recover_unrendered_records({"I": [entry]})
    off = _render_memberships([make()], supplementary_subpoints=False)
    off._recover_unrendered_records({"I": [make()]})

    assert gen.stats["unrendered_records_recovered"] == off.stats["unrendered_records_recovered"]
    assert _texts_of(gen.doc.element.body) == _texts_of(off.doc.element.body)


_OFFICES = ("Member, Estuarine Sediment Working Group, 2003-2008",
            "Chair, Marsh Restoration Task Force for Coastal Wetlands, 2010-2012")


def _membership_with_dated_offices():
    """A society line, then two offices held in it, each ending in its years,
    then prose (EBYSBC VQFSDI 92's shape). The recovery reads the dated parts
    as the entry's records (`recovery_candidates`); stage 4 extracted the
    membership."""
    entry = _membership(start="2009", end="present", prose=None)
    entry["text"] = "\t".join((f"{_SOCIETY_NAME}, 2009-present", *_OFFICES, _MEMBERSHIP_PROSE))
    return entry


def test_a_dated_office_the_recovery_reads_as_a_record_is_left_to_it():
    """#1205 comment 6317001032, 'Counts that do not add up' item 2: dev's
    #221 recovery re-emits each office to the Appendix, with sub-points on
    and off. Offered as a sub-point, the office's line read as rendered, so
    the recovery left it out, and after Accept All it was in the document
    nowhere. It stays the recovery's; the prose beside it is the sub-point,
    and the document after Accept All is the flag-off one."""
    entry = _membership_with_dated_offices()
    assert recovery_candidates("I", entry)[1:] == list(_OFFICES)

    def recovered(flag):
        gen = _render_memberships([_membership_with_dated_offices()], supplementary_subpoints=flag)
        gen._recover_unrendered_records({"I": [_membership_with_dated_offices()]})
        return gen

    on, off = recovered(True), recovered(False)
    (row,) = _subpoint_rows(on)
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _MEMBERSHIP_PROSE
    assert on.stats["unrendered_records_recovered"] == off.stats["unrendered_records_recovered"] == 2
    body_text = "".join(_texts_of(on.doc.element.body))
    assert all(office in body_text for office in _OFFICES)
    row.getparent().remove(row)
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


_OFFICES_BESIDE_THEIR_PROSE = (
    ("Chair, Education Committee, 2015-2018",
     "As chair of the education committee, developed the society's continuing education curriculum"),
    ("Chair, Annual Meeting Program Committee, 2012-2013",
     "Chaired the annual meeting program committee sessions and selected the plenary speakers"),
)


def _membership_with_an_office_beside_its_prose(office, prose):
    """A society line, a dated office held in it, then prose naming the
    office's committee: the prose holds every distinctive word of the
    office's line (#1205 review of I). Stage 4 extracted the membership."""
    entry = _membership(start="2009", end="present", prose=None)
    entry["text"] = "\t".join((f"Member, {_SOCIETY_NAME}, 2009-present", office, prose))
    return entry


@pytest.mark.parametrize("office, prose", _OFFICES_BESIDE_THEIR_PROSE, ids=["education", "program"])
def test_a_dated_office_beside_prose_repeating_its_words_stays_in_the_accepted_document(
        tmp_path, monkeypatch, office, prose):
    """#1205 review of I, through `run_stage6`. The prose is the sub-point,
    and the office, a record the #221 recovery reads, is not part of it. The
    sub-point's line held every word of the office's line, so the recovery
    read the office as rendered (`_record_rendered`'s token overlap) and left
    it out: it was in neither the accepted nor the rejected document. The
    recovery reads no membership sub-point (`RECOVERY_RECORD_CODES`) and
    re-emits the office as with sub-points off."""
    entry = _membership_with_an_office_beside_its_prose(office, prose)
    assert recovery_candidates("I", entry)[1:] == [office]

    def texts(flag):
        doc, placements = _stage6_render(tmp_path / str(flag), monkeypatch, [copy.deepcopy(entry)],
                                         uid="T1205J", supplementary_subpoints=flag)
        body = doc.element.body
        return ("".join(t.text or "" for t in body.iter(qn("w:t"))),
                "".join(t.text or "" for t in body.iter(qn("w:delText"))), placements)

    accepted_off, deleted_off, placements_off = texts(False)
    accepted, deleted, placements = texts(True)
    assert (deleted_off, placements_off) == ("", [])
    assert placements == [[prose]] and deleted == prose
    assert accepted.count(office) == accepted_off.count(office) == 1


@pytest.mark.parametrize("office, prose", _OFFICES_BESIDE_THEIR_PROSE, ids=["education", "program"])
def test_accept_all_of_an_office_beside_its_prose_is_the_flag_off_document(office, prose):
    """The same through the section and the recovery: with the sub-point's
    row removed, the document is the flag-off one."""
    def recovered(flag):
        entry = _membership_with_an_office_beside_its_prose(office, prose)
        gen = _render_memberships([entry], supplementary_subpoints=flag)
        gen._recover_unrendered_records({"I": [entry]})
        return gen

    on, off = recovered(True), recovered(False)
    (row,) = _subpoint_rows(on)
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == prose
    assert on.stats["unrendered_records_recovered"] == off.stats["unrendered_records_recovered"] == 1
    row.getparent().remove(row)
    assert _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _row_texts(on.doc.element.body) == _row_texts(off.doc.element.body)


#: #1205, every sub-point family: an entry whose line holds its own record,
#: then a sub-record of the entry with its own years (inside the span its row
#: prints), then prose repeating the sub-record's words. Invented names.
#: `own` and `sub` end in their years; `own_led` and `sub_led` open with them.
_SUB_RECORD_FAMILIES = {
    "K1": dict(
        code="K1", idx=40,
        fields={"course_title": "Introduction to Estuarine Biology", "role": "Course Director",
                "institution": "Example University", "course_code": "BIOL 210",
                "start_date": "2010", "end_date": "2018"},
        own="Course Director, Introduction to Estuarine Biology (BIOL 210), Example University, 2010-2018",
        own_led=("2010-2018 Course Director, Introduction to Estuarine Biology (BIOL 210), "
                 "Example University"),
        sub="Module Leader, Sediment Transport Laboratory Sessions, 2012-2014",
        sub_led="2012-2014 Module Leader, Sediment Transport Laboratory Sessions",
        prose=("As module leader for the sediment transport laboratory sessions, "
               "redesigned the field sampling exercises")),
    "M2A": dict(
        code="M2A", idx=52,
        fields={"title": ("Sediment Budgets and Nutrient Retention of Restored Tidal Marshes "
                          "in Northern Estuaries"),
                "agency": "Example Foundation", "pi_role": "Principal Investigator",
                "start_date": "2018", "end_date": "2026"},
        own=("Sediment Budgets and Nutrient Retention of Restored Tidal Marshes in Northern "
             "Estuaries, Example Foundation, Principal Investigator, 2018-2026"),
        own_led=("2018-2026 Sediment Budgets and Nutrient Retention of Restored Tidal Marshes in "
                 "Northern Estuaries, Example Foundation, Principal Investigator"),
        sub="Chair, Restoration Substudy Steering Committee, 2019-2021",
        sub_led="2019-2021 Chair, Restoration Substudy Steering Committee",
        prose=("As chair of the restoration substudy steering committee, coordinated the field "
               "teams across three marshes")),
    "D1": dict(
        code="D1", idx=30,
        fields={"title": "Associate Professor", "department": "Department of Estuarine Science",
                "institution": "Example University", "start_date": "2010", "end_date": "2020"},
        own="Associate Professor, Department of Estuarine Science, Example University, 2010-2020",
        own_led="2010-2020 Associate Professor, Department of Estuarine Science, Example University",
        sub="Director, Graduate Admissions Committee, 2014-2017",
        sub_led="2014-2017 Director, Graduate Admissions Committee",
        prose=("As director of the graduate admissions committee, redesigned the interview "
               "process for applicants")),
    "Q1": dict(
        code="Q1", idx=60,
        fields={"role": "President", "organization": "Example Society of Estuarine Ecology",
                "start_date": "2015", "end_date": "2017"},
        own="President, Example Society of Estuarine Ecology, 2015-2017",
        own_led="2015-2017 President, Example Society of Estuarine Ecology",
        sub="Chair, Annual Meeting Program Committee, 2016-2017",
        sub_led="2016-2017 Chair, Annual Meeting Program Committee",
        prose=("As chair of the annual meeting program committee, organized the plenary "
               "sessions on marsh restoration")),
    "Q2": dict(
        code="Q2", idx=61,
        fields={"role": "Member", "committee_name": "Scientific Advisory Board",
                "organization": "Example Estuarine Foundation", "start_date": "2012", "end_date": "2018"},
        own="Member, Scientific Advisory Board, Example Estuarine Foundation, 2012-2018",
        own_led="2012-2018 Member, Scientific Advisory Board, Example Estuarine Foundation",
        sub="Chair, Grants Review Subcommittee, 2014-2016",
        sub_led="2014-2016 Chair, Grants Review Subcommittee",
        prose="As chair of the grants review subcommittee, led the evaluation of seed grant proposals"),
    "H": dict(
        code="H", idx=70,
        fields={"award_name": "Early Career Fellowship",
                "granting_body": "Example Foundation for Estuarine Science", "date": "2011-2014"},
        own="Early Career Fellowship, Example Foundation for Estuarine Science, 2011-2014",
        own_led="2011-2014 Early Career Fellowship, Example Foundation for Estuarine Science",
        sub="Cohort Representative, Fellows Advisory Council, 2012-2013",
        sub_led="2012-2013 Cohort Representative, Fellows Advisory Council",
        prose=("As cohort representative on the fellows advisory council, organized the annual "
               "fellows retreat")),
    "N3A": dict(
        code="N3A", idx=82,
        fields={"mentee_name": "Orrin Feldspar", "mentee_level": "Postdoctoral Fellow",
                "research_focus": ("sediment transport modelling in restored tidal marsh channels "
                                   "and creeks"),
                "start_date": "2019"},
        own=("Orrin Feldspar, Postdoctoral Fellow, sediment transport modelling in restored tidal "
             "marsh channels and creeks, 2019-present"),
        own_led=("2019-present Orrin Feldspar, Postdoctoral Fellow, sediment transport modelling in "
                 "restored tidal marsh channels and creeks"),
        sub="Laboratory Safety Officer, Estuarine Ecology Laboratory, 2020-2021",
        sub_led="2020-2021 Laboratory Safety Officer, Estuarine Ecology Laboratory",
        prose=("As laboratory safety officer for the estuarine ecology laboratory, trained incoming "
               "students on field safety")),
    "N3B": dict(
        code="N3B", idx=83,
        fields={"mentee_name": "Pella Quartzmoor", "mentee_level": "PhD Student",
                "start_date": "2012", "end_date": "2017",
                "thesis_title": ("Nutrient Cycling and Denitrification Rates in Brackish Coastal "
                                 "Lagoons under Seasonal Hypoxia")},
        own=("Pella Quartzmoor, PhD Student, Nutrient Cycling and Denitrification Rates in Brackish "
             "Coastal Lagoons under Seasonal Hypoxia, 2012-2017"),
        own_led=("2012-2017 Pella Quartzmoor, PhD Student, Nutrient Cycling and Denitrification Rates "
                 "in Brackish Coastal Lagoons under Seasonal Hypoxia"),
        sub="Teaching Assistant Coordinator, Marine Biology Program, 2014-2015",
        sub_led="2014-2015 Teaching Assistant Coordinator, Marine Biology Program",
        prose=("As teaching assistant coordinator for the marine biology program, scheduled the "
               "laboratory sections")),
    "B1": dict(
        code="B1", idx=10,
        fields={"degree": "PhD", "discipline": "Estuarine Ecology", "institution": "Example University",
                "year": "2006"},
        own="PhD, Estuarine Ecology, Example University, 2001-2006",
        own_led="2001-2006 PhD, Estuarine Ecology, Example University",
        sub="Graduate Student Representative, Faculty Search Committee, 2004-2005",
        sub_led="2004-2005 Graduate Student Representative, Faculty Search Committee",
        prose=("As graduate student representative on the faculty search committee, interviewed "
               "candidates")),
    "O": dict(
        code="O", idx=90,
        fields={"leadership_role": "Director", "institution": "Example Estuarine Research Center",
                "start_date": "2006", "end_date": "2012"},
        own="Director, Example Estuarine Research Center, 2006-2012",
        own_led="2006-2012 Director, Example Estuarine Research Center",
        sub="Chair, Field Station Steering Committee, 2008-2010",
        sub_led="2008-2010 Chair, Field Station Steering Committee",
        prose="As chair of the field station steering committee, set the station's research priorities"),
    "R": dict(
        code="R", idx=101,
        fields={"title": "Sediment Budgets of Restored Tidal Marshes",
                "event_name": "Example Estuarine Institute Lectures", "location": "Example City",
                "role": "Invited Lecturer", "date": "2014; 2016"},
        own=("Sediment Budgets of Restored Tidal Marshes, Example Estuarine Institute Lectures, "
             "Example City, 2014-2016"),
        own_led=("2014-2016 Sediment Budgets of Restored Tidal Marshes, Example Estuarine Institute "
                 "Lectures, Example City"),
        sub="Series Organizer, Visiting Scholars Committee, 2015-2016",
        sub_led="2015-2016 Series Organizer, Visiting Scholars Committee",
        prose=("As series organizer for the visiting scholars committee, recruited speakers for "
               "the lectures")),
    "C": dict(
        code="C", idx=20,
        fields={"training_type": "Postdoctoral Fellow", "institution": "Example University Hospital",
                "start_date": "2006", "end_date": "2009"},
        own="Postdoctoral Fellow, Example University Hospital, 2006-2009",
        own_led="2006-2009 Postdoctoral Fellow, Example University Hospital",
        sub="Chief Resident, Education Committee, 2007-2008",
        sub_led="2007-2008 Chief Resident, Education Committee",
        prose=("As chief resident on the education committee, developed the residency's teaching "
               "curriculum")),
    "L3": dict(
        code="L3", idx=1482,
        fields={"leadership_role": "Director, Example Wound Care Program",
                "institution": "Example Estuarine Hospital", "start_date": "2010", "end_date": "2018"},
        own="Director, Example Wound Care Program, Example Estuarine Hospital, 2010-2018",
        own_led="2010-2018 Director, Example Wound Care Program, Example Estuarine Hospital",
        sub="Chair, Limb Salvage Quality Committee, 2012-2014",
        sub_led="2012-2014 Chair, Limb Salvage Quality Committee",
        prose=("As chair of the limb salvage quality committee, set the clinic's amputation "
               "review standards")),
    "Q4B": dict(
        code="Q4B", idx=41,
        fields={"role": "Associate Editor", "journal_name": "Journal of Estuarine Sediment Studies",
                "start_date": "2012", "end_date": "2018"},
        own="Associate Editor, Journal of Estuarine Sediment Studies, 2012-2018",
        own_led="2012-2018 Associate Editor, Journal of Estuarine Sediment Studies",
        sub="Guest Editor, Special Issue on Marsh Hydrology, 2014-2015",
        sub_led="2014-2015 Guest Editor, Special Issue on Marsh Hydrology",
        prose=("As guest editor of the special issue on marsh hydrology, recruited its twelve "
               "invited reviews")),
    "Q4C": dict(
        code="Q4C", idx=42,
        fields={"journal_name": "Kestrel Review of Tidal Ecology", "start_date": "2010", "end_date": "2020"},
        own="Kestrel Review of Tidal Ecology, 2010-2020",
        own_led="2010-2020 Kestrel Review of Tidal Ecology",
        sub="Chair, Early Career Reviewers Committee, 2015-2017",
        sub_led="2015-2017 Chair, Early Career Reviewers Committee",
        prose=("As chair of the early career reviewers committee, mentored the journal's new "
               "referees")),
    "S1": dict(
        code="S1", idx=1200,
        fields={"authors": "Quill AM, Rook BT", "title": "Sediment Budgets of Restored Tidal Marshes",
                "journal": "Journal of Estuarine Sediment Studies", "year": "2013-2014"},
        own=("Quill AM, Rook BT. Sediment Budgets of Restored Tidal Marshes. Journal of Estuarine "
             "Sediment Studies, 2013-2014"),
        own_led=("2013-2014 Quill AM, Rook BT. Sediment Budgets of Restored Tidal Marshes. Journal of "
                 "Estuarine Sediment Studies"),
        sub="Featured Article, Example Coastal Society Newsletter, 2016-2017",
        sub_led="2016-2017 Featured Article, Example Coastal Society Newsletter",
        prose=("As the coastal society newsletter's featured article, summarised the findings for "
               "marsh managers")),
    "S2": dict(
        code="S2", idx=1210,
        fields={"authors": "Quill AM", "title": "Restoring Tidal Marshes: A Review of Sediment Budgets",
                "journal": "Kestrel Review of Tidal Ecology", "year": "2015-2016"},
        own=("Quill AM. Restoring Tidal Marshes: A Review of Sediment Budgets. Kestrel Review of Tidal "
             "Ecology, 2015-2016"),
        own_led=("2015-2016 Quill AM. Restoring Tidal Marshes: A Review of Sediment Budgets. Kestrel "
                 "Review of Tidal Ecology"),
        sub="Invited Commentary, Example Estuarine Society Bulletin, 2018-2019",
        sub_led="2018-2019 Invited Commentary, Example Estuarine Society Bulletin",
        prose=("As the estuarine society bulletin's invited commentary, answered the review's "
               "critics")),
}
#: The two line shapes in which sub-points on lost the sub-record: tab parts
#: ending in their years, where the sub-point offered the sub-record itself;
#: and date-led tab parts, where the prose repeating its words read as the
#: sub-record rendered. B1, C, H and O lost it in the first shape only: their
#: own rules (`_own_record_line`, the honors writer's rows) offer nothing in
#: the second. P lost it in neither (`_own_committee_line`), I since #1991.
#: L3, Q4B, Q4C, S1 and S2, which joined after #2010, are pinned in the first
#: shape: their own rule (`_own_record_line`) offers nothing in the second.
_TRAILING_YEARS, _DATE_LED = "tab_trailing_years", "tab_date_led"
_SUB_RECORD_CASES = [*((name, _TRAILING_YEARS) for name in _SUB_RECORD_FAMILIES),
                     *((name, _DATE_LED) for name in ("K1", "M2A", "D1", "Q1", "Q2", "N3A", "N3B", "R"))]


def _sub_record_entry(name, shape):
    """The family's entry in `shape`, its sub-record and its prose."""
    family = _SUB_RECORD_FAMILIES[name]
    own, sub = (family["own"], family["sub"]) if shape == _TRAILING_YEARS else (
        family["own_led"], family["sub_led"])
    entry = {"element_idx_start": family["idx"], "taxonomy_code": family["code"],
             "text": "\t".join((own, sub, family["prose"])), "extracted_fields": dict(family["fields"])}
    return entry, sub, family["prose"]


def _accepted_and_deleted(tmp_path, monkeypatch, entry, flag, uid, *others):
    """Through `run_stage6`: the document's accepted text (`w:t`), its
    deleted text (`w:delText`, what Reject All adds) and its sidecar's
    sub-point placements. `others` are entries listed before `entry`."""
    entries = [copy.deepcopy(e) for e in (*others, entry)]
    doc, placements = _stage6_render(tmp_path / str(flag), monkeypatch, entries,
                                     uid=uid, supplementary_subpoints=flag)
    body = doc.element.body
    return ("".join(t.text or "" for t in body.iter(qn("w:t"))),
            "".join(t.text or "" for t in body.iter(qn("w:delText"))), placements)


@pytest.mark.parametrize("name, shape", _SUB_RECORD_CASES,
                         ids=[f"{name}-{shape}" for name, shape in _SUB_RECORD_CASES])
def test_accept_all_keeps_a_dated_sub_record_in_every_family(tmp_path, monkeypatch, name, shape):
    """#1205, every family: the #221 recovery re-emits the sub-record with
    sub-points off. With them on, the sub-point offered it (its own years
    at its end) or its prose repeated its words, and the recovery read it as
    rendered: after Accept All it was in the document nowhere. It is now the
    recovery's (`_left_to_recovery`, `held_by_subpoint`): Accept All is the
    flag-off document, Reject All prints it once, and the prose is the
    sub-point."""
    entry, sub, prose = _sub_record_entry(name, shape)
    accepted_off, _deleted_off, _ = _accepted_and_deleted(tmp_path, monkeypatch, entry, False, "T1205F")
    accepted, deleted, placements = _accepted_and_deleted(tmp_path, monkeypatch, entry, True, "T1205F")

    assert accepted_off.count(sub) == 1
    assert accepted == accepted_off
    assert sub not in deleted
    assert placements == [[prose]]


#: A dated sub-record of the appointment whose words its duties repeat
#: (`_DUTY_PROSE`: "departmental curriculum committee").
_DUTY_SUB_RECORD = "2014-2016 Chair, Departmental Curriculum Committee"


def _duty_appointment_with(sub_record):
    """`_duty_appointment` (SQMWHM 29's shape: the duties a tab row of their
    own, with no date), with `sub_record` as a line before the duties."""
    entry = _duty_appointment()
    if sub_record:
        head, duties = entry["text"].split("\n")
        entry["text"] = "\n".join((head, sub_record, duties))
    return entry


@pytest.mark.parametrize("sub_record", [None, _DUTY_SUB_RECORD], ids=["duties", "and_a_dated_sub_record"])
def test_appointment_duties_stay_its_sub_point_beside_a_dated_sub_record(tmp_path, monkeypatch,
                                                                         sub_record):
    """Owner decision 2026-10-08: an appointment's duties are its
    tracked-deleted sub-point, not a recovered line. The #221 recovery reads
    their row (a tab line with no date of its own) as rendered by the
    sub-point (`held_by_subpoint`), so Accept All drops the copy it re-emits
    with sub-points off, as decided (SQMWHM 29), and Reject All prints the
    duties once. A dated sub-record beside them is the recovery's in both,
    though the duties repeat its words."""
    entry = _duty_appointment_with(sub_record)
    accepted_off, deleted_off, _ = _accepted_and_deleted(tmp_path, monkeypatch, entry, False, "T1205D")
    accepted, deleted, placements = _accepted_and_deleted(tmp_path, monkeypatch, entry, True, "T1205D")

    assert placements == [[_DUTY_PROSE]] and deleted == _DUTY_PROSE
    assert deleted_off == "" and accepted_off.count(_DUTY_PROSE) == 1
    assert _DUTY_PROSE not in accepted
    if sub_record:
        assert accepted.count(sub_record) == accepted_off.count(sub_record) == 1


_TIDAL_DUTIES = ("Coordinated the tidal flats sampling programme and trained visiting field technicians",
                 "Maintained the shared sediment core archive and catalogued every seasonal survey")


def test_every_paragraph_of_every_sub_point_stands_for_its_duty_row(tmp_path, monkeypatch):
    """#1205: the recovery reads each paragraph of each written sub-point
    (`_subpoint_prose`, `held_by_subpoint`). Two appointments in SQMWHM
    29's shape, the second with two duty rows: every duty row is read as
    rendered by its paragraph, so each duty prints once with sub-points off,
    not after Accept All, and once on Reject All."""
    second = {"element_idx_start": 31, "taxonomy_code": "D1", "hierarchy": ["PROFESSIONAL EXPERIENCE"],
              "text": ("July 2012 - present\tAdjunct Professor, Department of Tidal Example Science, "
                       f"Example College\nNorth Annex, ZQ.\t{_TIDAL_DUTIES[0]}\nSouth Annex, ZQ.\t"
                       f"{_TIDAL_DUTIES[1]}"),
              "extracted_fields": {"title": "Adjunct Professor",
                                   "institution": "Department of Tidal Example Science, Example College",
                                   "start_date": "2012-07", "end_date": "present"}}
    first = _duty_appointment()
    accepted_off, deleted_off, _ = _accepted_and_deleted(tmp_path, monkeypatch, second, False, "T1205E",
                                                         first)
    accepted, deleted, placements = _accepted_and_deleted(tmp_path, monkeypatch, second, True, "T1205E",
                                                          first)

    assert placements == [[_DUTY_PROSE], list(_TIDAL_DUTIES)] and deleted_off == ""
    for duty in (_DUTY_PROSE, *_TIDAL_DUTIES):
        assert accepted_off.count(duty) == 1 and duty not in accepted and deleted.count(duty) == 1


def test_title_duties_stay_a_sub_point_through_run_stage6(tmp_path, monkeypatch):
    """#1641's title duties, which the positions writer makes the
    appointment's sub-point (owner decision 2026-10-09: DYLJXC 661/668),
    are unchanged: Accept All is the flag-off document without its
    continuation row of the duties, Reject All prints them once, and the
    text's own statement of them is not offered again."""
    entry = _title_duty_appointment()
    accepted_off, deleted_off, _ = _accepted_and_deleted(tmp_path, monkeypatch, entry, False, "T1205T")
    accepted, deleted, placements = _accepted_and_deleted(tmp_path, monkeypatch, entry, True, "T1205T")

    assert deleted_off == "" and len(placements) == 1
    (duties,) = placements[0]
    assert duties.endswith(_TITLE_DUTIES) and deleted == duties
    assert accepted == accepted_off.replace(duties, "", 1) != accepted_off


def _appointment_entry(text, title="Associate Professor", start="2010", end="present",
                       institution="Example University"):
    """An invented D1 entry with `text` and its row's fields."""
    return {"element_idx_start": 30, "taxonomy_code": "D1", "text": text,
            "extracted_fields": {"title": title, "department": "Department of Estuarine Science",
                                 "institution": institution, "start_date": start,
                                 "end_date": end}}


_SURVEY = "A survey of estuarine habitats, sediment transport and nutrient cycling for second year students"
_PRACTICUM = "A practicum in coring, sampling and laboratory analysis held at the example field station"
_ACCRETION = "This study will measure accretion and nutrient retention across restored marshes"
_HYPOXIA = "This study mapped invertebrate communities under seasonal hypoxia"
_MANAGE = ("Plan the estuary survey cruises each season and keep the shared sampling archive. "
           "Train new field staff in coring.")
_SERVE = ("Oversee the joint tidal flats sampling station; advise visiting graduate students on "
          "their sediment budgets.")
_SUPERVISED = "Supervised the department's graduate teaching assistants and their laboratory sections"
#: #1205 review: an entry of two date-led lines, each a record with its prose
#: in a tab part. The first is the entry's own (its fields' years), which its
#: row renders, so the #221 recovery reads it as rendered and its prose
#: reaches the document only as the sub-point. The second is another record
#: with years of its own (`restored`), which the recovery re-emits with
#: sub-points off, or reads as rendered too (`restored` None). Invented text.
_TWO_RECORD_LINES = {
    "D1_another_appointment": dict(
        entry=_appointment_entry(
            f"July 2010 - present\tProfessor and Field Studies Coordinator, Department of Estuarine "
            f"Science, Example University, Example City.\t{_MANAGE}\nMarch 2008 - present\tAdjunct "
            f"Professor, Department of Tidal Geology, Example Coastal College\t{_SERVE}",
            title="Professor and Field Studies Coordinator", start="July 2010"),
        kept=(_MANAGE,), restored="Adjunct Professor, Department of Tidal Geology"),
    "D1_both_rendered": dict(
        entry=_appointment_entry(
            f"2015-2020\tAssociate Professor, Department of Estuarine Science, Example University\t"
            f"{_DUTY_PROSE}\n2010-2015\tAssistant Professor, Department of Estuarine Science, Example "
            f"University\t{_SUPERVISED}", start="2015", end="2020"),
        kept=(_DUTY_PROSE, _SUPERVISED), restored=None),
    "K1": dict(
        entry={"element_idx_start": 40, "taxonomy_code": "K1",
               "text": (f"2010-2018\tCourse Director, Introduction to Estuarine Biology (BIOL 210)\t{_SURVEY}"
                        f"\n2012-2016\tCourse Director, Field Methods in Marsh Ecology (BIOL 340)\t{_PRACTICUM}"),
               "extracted_fields": {"course_title": "Introduction to Estuarine Biology", "role": "Course Director",
                                    "institution": "Example University", "course_code": "BIOL 210",
                                    "start_date": "2010", "end_date": "2018"}},
        kept=(_SURVEY,), restored="Field Methods in Marsh Ecology"),
    "M2A": dict(
        entry={"element_idx_start": 52, "taxonomy_code": "M2A",
               "text": (f"2018-2026\tSediment Budgets of Restored Tidal Marshes, Example Foundation, Principal "
                        f"Investigator\t{_ACCRETION}\n2014-2017\tBenthic Invertebrates of Brackish Lagoons, "
                        f"Example Trust, Co-Investigator\t{_HYPOXIA}"),
               "extracted_fields": {"title": "Sediment Budgets of Restored Tidal Marshes",
                                    "agency": "Example Foundation", "pi_role": "Principal Investigator",
                                    "start_date": "2018", "end_date": "2026"}},
        kept=(_ACCRETION,), restored="Benthic Invertebrates of Brackish Lagoons"),
}


@pytest.mark.parametrize("name", list(_TWO_RECORD_LINES))
def test_reject_all_keeps_the_prose_of_a_record_line_the_recovery_reads_as_rendered(
        tmp_path, monkeypatch, name):
    """#1205 review (FINSIS 511): a fragment merely part of a dated record
    line was left to the #221 recovery whether or not it re-emits the line.
    It reads the entry's own line as rendered (its row prints it), so the
    prose was then in neither document. Now only a record of its own that
    the recovery re-emits keeps its prose (`RecordLine.lost`): Accept All is
    the flag-off document, which re-emits the other record once, and Reject
    All adds the own line's prose, once."""
    case = _TWO_RECORD_LINES[name]
    accepted_off, _, _ = _accepted_and_deleted(tmp_path, monkeypatch, case["entry"], False, "T1205R")
    accepted, deleted, _ = _accepted_and_deleted(tmp_path, monkeypatch, case["entry"], True, "T1205R")

    assert accepted == accepted_off
    for prose in case["kept"]:
        assert prose not in accepted_off and deleted.count(prose) == 1
    if case["restored"]:
        assert accepted.count(case["restored"]) == 1 and case["restored"] not in deleted


_CORING = "Direct the sediment laboratory's coring program."
_TWELVE = "Supervise twelve doctoral researchers in marsh hydrology."


@pytest.mark.parametrize("entry, duties", [
    # FINSIS 511's shape: the institution and title, then the appointment's
    # own dates and its duties as tab parts.
    (_appointment_entry(f"Example University Medical College / Example Hospital\tAssociate Professor of "
                        f"Estuarine Science\n2010\u2013Present\t{_CORING}\t{_TWELVE}",
                        title="Associate Professor of Estuarine Science",
                        institution="Example University Medical College / Example Hospital"),
     (_CORING, _TWELVE)),
    # SQMWHM 29's shape, the duties' row ending on a year or a range.
    (dict(_duty_appointment(), text=_duty_appointment()["text"] + " since 2015"),
     (f"{_DUTY_PROSE} since 2015",)),
    (dict(_duty_appointment(), text=_duty_appointment()["text"] + " (2012-2014)"),
     (f"{_DUTY_PROSE} (2012-2014)",)),
], ids=["beside_its_own_dates", "row_ends_on_a_year", "row_ends_on_a_range"])
def test_duties_beside_their_appointments_dates_stay_its_sub_point(tmp_path, monkeypatch, entry, duties):
    """#1205 review, owner decision 2026-10-08: duties the recovery re-emits
    with sub-points off, beside the appointment's own dates or on a row of
    their own that ends on a year, are no record of their own
    (`is_sub_record`): they are the appointment's sub-point, and Accept All
    drops the flag-off copy, as for SQMWHM 29."""
    accepted_off, deleted_off, _ = _accepted_and_deleted(tmp_path, monkeypatch, entry, False, "T1205S")
    accepted, deleted, placements = _accepted_and_deleted(tmp_path, monkeypatch, entry, True, "T1205S")

    assert deleted_off == "" and placements == [list(duties)]
    for duty in duties:
        assert accepted_off.count(duty) == 1 and duty not in accepted and deleted.count(duty) == 1


#: Another entry's row holding `_DUTY_PROSE` among cells of its own, and a
#: second row (an invented reviewing record).
_ROW_HOLDING_THE_DUTIES = {
    "element_idx_start": 12, "taxonomy_code": "Q4D",
    "text": (f"Example Sediment Science Advisory Panel Steering Group | {_DUTY_PROSE} | Example "
             f"Regional Water Authority\nExample Coastal Monitoring Network Data Quality Board | "
             f"Reviewed the network's monitoring protocols | Example State Agency"),
    "extracted_fields": {"journal_name": "Example Coastal Monitoring Network Data Quality Board"}}


@pytest.mark.parametrize("first", [True, False], ids=["read_first", "read_last"])
def test_another_entrys_row_holding_the_duties_leaves_them_once_on_reject_all(tmp_path, monkeypatch,
                                                                              first):
    """#1205: the recovery reads a written sub-point for a line of any entry
    that is no record of its own and that the prose is part of
    (`held_by_subpoint`). Another entry's row holding the duties, read
    before the appointment or after it, is read so, and the duties are the
    owner decision's (2026-10-08): flag off prints them once, Accept All
    drops them, and Reject All prints them exactly once. Read only for the
    sub-point's own entry, the row read first was re-emitted with them, and
    Reject All printed them twice."""
    other, appointment = _ROW_HOLDING_THE_DUTIES, _duty_appointment()
    entry, before = (appointment, other) if first else (other, appointment)
    accepted_off, deleted_off, _ = _accepted_and_deleted(tmp_path, monkeypatch, entry, False, "T1205X",
                                                         before)
    accepted, deleted, placements = _accepted_and_deleted(tmp_path, monkeypatch, entry, True, "T1205X",
                                                          before)
    rejected = accepted + deleted

    assert placements == [[_DUTY_PROSE]] and deleted == _DUTY_PROSE
    assert deleted_off == "" and accepted_off.count(_DUTY_PROSE) == 1
    assert _DUTY_PROSE not in accepted and rejected.count(_DUTY_PROSE) == 1
    # Nothing of the row prints twice on Reject All. Its other cells are not
    # pinned: read first, Accept All lacks them as on dev, a pre-existing
    # loss (a row that holds the offered prose plus cells the page prints
    # nowhere), filed separately.
    for cell in ("Advisory Panel Steering Group", "Regional Water Authority"):
        assert rejected.count(cell) <= 1


def test_the_membership_span_starts_at_the_heading_its_writer_looks_up():
    """The span starts at the heading the writer looks its table up by
    (`MEMBERSHIPS_HEADING`), and holds that table. The other callers of the
    code's heading lookup still name none for I: the reconsider pass and the
    #221 recovery keep an I record in the Appendix, as before I's
    sub-points."""
    gen = _render_memberships([_membership()], supplementary_subpoints=True)
    span = gen._subpoint_section_spans()["I"]
    heading = gen.doc.element.body[span.start]
    assert "".join(t.text or "" for t in heading.iter(qn("w:t"))).startswith("PROFESSIONAL ORGANIZATIONS")
    (row,) = _subpoint_rows(gen)
    assert span.start < list(gen.doc.element.body).index(row.getparent()) < span.end
    assert WCMTemplateGenerator._get_wcm_section_header("I") == ""
    assert gen._insert_reconsidered_segment(_OFFICES[0], "I") is False


def test_a_template_with_no_membership_heading_places_no_membership_prose():
    """Without the heading neither the writer nor the span finds the
    section: no row is written, and the prose goes under no other line."""
    gen = _generator(supplementary_subpoints=True)
    gen.doc = Document()
    gen.doc.add_paragraph("A. PERSONAL DATA")
    entry = _membership()
    gen._fill_memberships([entry])
    gen._add_supplementary_subpoints({"I": [entry]})

    assert "I" not in gen._subpoint_section_spans()
    assert _deleted_texts(gen) == ""


def _render_grouped_memberships(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_memberships(grouped.get("I", []))
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _membership_line(text, *memberships):
    """One source line stage 4 returned as a record per (type, society,
    start), the entry's own fields being the last record's. Fan-out gives
    every earlier record a line of its own and the last one the whole line."""
    records = [{"membership_type": mtype, "organization": org, "start_date": start}
               for mtype, org, start in memberships]
    return {"element_idx_start": 663, "taxonomy_code": "I", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_FELLOWSHIP = ("Fellow", "Kestrel Academy of Tidal Studies", "2012")
_MEMBER = ("Member", _SOCIETY_NAME, "2009")
_FELLOWSHIP_LINE = "Fellow, Kestrel Academy of Tidal Studies, 2012"
_MEMBER_LINE = f"Member, {_SOCIETY_NAME}, 2009"
_FELLOWSHIP_PROSE = "Reviewed the academy's annual awards and judged its student poster competition"


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own membership.
    ([f"{_FELLOWSHIP_LINE}. {_FELLOWSHIP_PROSE}", _MEMBER_LINE],
     [("Fellow, Kestrel Academy of Tidal Studies", f"{_FELLOWSHIP_LINE}. {_FELLOWSHIP_PROSE}")]),
    ([_FELLOWSHIP_LINE, f"{_MEMBER_LINE}. {_MEMBERSHIP_PROSE}"],
     [(f"Member, {_SOCIETY_NAME}", f"{_MEMBER_LINE}. {_MEMBERSHIP_PROSE}")]),
    # They do not: which membership the prose is about cannot be told, so neither takes it.
    ([_FELLOWSHIP_LINE, _MEMBERSHIP_PROSE, _MEMBER_LINE], []),
])
def test_prose_on_a_line_of_several_memberships_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the membership's line as well as its prose brings the type,
    society and year along, and Reject All prints them a second time, under
    the row that holds them (the shape every family has; #1205, 2026-10-09)."""
    entry = _membership_line("\t".join(segments), _FELLOWSHIP, _MEMBER)
    gen, grouped = _render_grouped_memberships([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["I"][-1] and len(grouped["I"]) == 2
    assert [(_row_texts_of(row.getprevious())[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def test_generic_list_fan_out_children_of_memberships_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, the entry renders its own
    fields, and the list is memberships the row does not print
    (`_own_membership_line`): their lines would be offered as its prose."""
    records = [{"membership_type": mtype, "organization": org, "start_date": start}
               for mtype, org, start in (_FELLOWSHIP, _MEMBER)]
    for text, children in ((f"{_FELLOWSHIP_LINE}\t{_MEMBER_LINE}", 2),
                           (f"{_FELLOWSHIP_LINE}\t{_MEMBER_LINE}\t{_MEMBERSHIP_PROSE}", 1)):
        entry = {"element_idx_start": 664, "taxonomy_code": "I", "text": text,
                 "extracted_fields": {"memberships": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_memberships([entry], supplementary_subpoints=True)
        assert len(grouped["I"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["I"])
        assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("N1", "MENTORING PROGRAMS"), ("S0", "BIBLIOGRAPHY")])
def test_a_membership_rerouted_out_of_i_to_an_excluded_code_gets_no_sub_point(code, heading):
    """N1 and S0 are excluded: a membership rendered as either is that
    code's record, so neither the membership row of the same society nor any
    other carries its prose."""
    gen, grouped = _render_grouped_memberships([_later_membership(), _rerouted(_membership(), code, heading)],
                                               supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [135]
    assert [e["element_idx_start"] for e in grouped["I"]] == [122]
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("F2", "Board Certification"),
                                           ("C", "POSTDOCTORAL TRAINING"),
                                           ("Q2", "SERVICE ON BOARDS")])
def test_a_refused_reroute_keeps_the_membership_rule(code, heading):
    """A membership rerouted to F2, whose span the heading lookup does not
    find, or to C or Q2, keeps its I code: their renderers write none of its
    fields or would drop its type (`_cross_family_refusal`). It renders as an
    I row and its prose is a deleted row under that row."""
    gen, grouped = _render_grouped_memberships([_rerouted(_membership(), code, heading)],
                                               supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["I"]] == [135]
    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _MEMBERSHIP_ROW


def test_a_record_rerouted_into_i_follows_the_membership_rule():
    """A board seat stage 3b coded Q2 under the memberships heading, whose
    role is a plain "Member" the memberships row states anyway, renders as
    an I row: its prose is a deleted row under that row, recorded as I."""
    seat = _membership(code="Q2", mtype=None, role="Member")
    gen, grouped = _render_grouped_memberships(
        [_rerouted(seat, "I", "PROFESSIONAL ORGANIZATIONS AND SOCIETY MEMBERSHIPS")],
        supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["I"]] == [135]
    (row,) = _subpoint_rows(gen)
    assert _SOCIETY_NAME in _row_text(row.getprevious())
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _MEMBERSHIP_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.code == "I"


def test_a_membership_the_section_writes_no_row_for_gets_no_sub_point():
    """An entry no row was written for -- one the section skipped -- has no
    row of its own: its prose goes under no other line of the document,
    though the other membership's row holds every distinctive word of its
    line."""
    skipped, later = _membership(), _later_membership()
    gen = _generator(supplementary_subpoints=True)
    gen._fill_memberships([later])
    gen._add_supplementary_subpoints({"I": [later, skipped]})

    assert gen._row_receipts.elements_for(skipped) == []
    assert gen._row_receipts.elements_for(later)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_membership_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _membership()
    withheld["_pii_withheld"] = True
    gen = _render_memberships([_later_membership(), withheld], supplementary_subpoints=True)

    assert any(_row_text(tr) == _MEMBERSHIP_ROW for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_membership_dedup_drops_offers_its_prose_once_under_the_kept_row():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no row, so its
    prose goes only under the row of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_membership(idx=135), _membership(idx=136)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["I"]) == 2 and len(kept["I"]) == 1
    gen._fill_memberships(kept["I"])
    gen._add_supplementary_subpoints(before)

    (row,) = _subpoint_rows(gen)
    assert _row_text(row.getprevious()) == _MEMBERSHIP_ROW
    assert _deleted_texts(gen) == _MEMBERSHIP_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["I"][0]["element_idx_start"]


def test_the_doctor_matches_a_membership_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier membership's
    sub-point under its own index and I, and the doctor's per-entry match
    counts it for that entry, not for the later membership that shares its
    words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_membership(), _membership()],
                                     uid="T1205I")
    assert placements == [[_MEMBERSHIP_PROSE]]
    sidecar = json.loads((tmp_path / "T1205I_render_warnings.json").read_text())
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(135, "I")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_MEMBERSHIP_PROSE, subpoints, element_idx=135)
    assert not offered_as_subpoint(_MEMBERSHIP_PROSE, subpoints, element_idx=122)


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_membership_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line (#1205 review of O). The same membership without them still has
    its sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_membership()])
    assert placements == [[_MEMBERSHIP_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_membership(), _MEMBERSHIP_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


# --- #1205: L3 clinical leadership, the next code family ---------------------
# The clinical-practice section writes every L3 entry as a line of its own,
# "Role, Institution, Dates", with no bullet: the WCM template ships no table
# under the heading. So the prose is a deleted paragraph one indent step
# under the role's own line, the one paragraph the writer recorded for that
# entry (the row-receipt ledger). A template shipping a clinical table there
# gets a deleted row under the role's own row.

_CLINICAL_PROSE = ("Oversaw a twelve-bed inpatient unit and its outpatient clinic, and built "
                   "the regional referral network for community physicians")
_CLINICAL_ROLE = "Director, Example Wound Care Program"
_CLINICAL_SITE = "Norvale General Hospital"


def _clinical_role(idx=1482, name=_CLINICAL_ROLE, site=_CLINICAL_SITE, start="2010", end="2015",
                   prose=_CLINICAL_PROSE, code="L3", **fields):
    """A clinical leadership role whose source line goes on past its title,
    hospital and years into what the role oversaw, as YUYVIG SQMWHM 1482's
    did. A later role of the same title at the same hospital shares every
    distinctive word of its line: only the years, digits, tell them apart.
    A value passed as None is in neither the line nor the fields."""
    named = {"leadership_role": name, "institution": site, "start_date": start, "end_date": end}
    dates = "-".join(value for value in (start, end) if value)
    line = ", ".join(value for value in (name, site, dates) if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_clinical_role(name=_CLINICAL_ROLE):
    """The same hospital's role from 2016, with no prose: it renders first,
    reverse chronologically."""
    return _clinical_role(idx=1478, name=name, start="2016", end="present", prose="")


def _render_clinical_roles(entries, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_clinical_practice({"L3": entries})
    gen._add_supplementary_subpoints({"L3": entries})
    return gen


def _clinical_role_table_row():
    """The same role read from a source table: stage 2 joins the row's cells
    with pipes, and its position names the row ("41.2")."""
    entry = _clinical_role()
    entry["element_idx_start"] = "41.2"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


# The role's line is an entry card (#2019): the source's own words, the
# years on a right tab stop.
_CLINICAL_LINE = f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}\t2010-2015"
_LATER_CLINICAL_LINE = f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}\t2016-present"


def _subpoint_paragraphs(gen):
    """The body paragraphs the sub-point pass wrote, in document order."""
    from unified_pipeline.stage6 import supplementary

    return [p for p in gen.doc.element.body.iterchildren(qn("w:p")) if supplementary._is_subpoint(p)]


def _paragraph_text(p):
    """Its run text, a run's tab (an entry card's date, #2019) as '\t'; a
    tab STOP (`w:pPr/w:tabs/w:tab`) is no text."""
    return "".join(("\t" if t.getparent().tag == qn("w:r") else "") if t.tag == qn("w:tab")
                   else t.text or ""
                   for t in p.iter(qn("w:t"), qn("w:delText"), qn("w:tab")))


def _left_indent(p):
    return p.find(f"{qn('w:pPr')}/{qn('w:ind')}").get(qn("w:left"))


def _accept_all(gen):
    """Word's Accept All on the sub-points: each deleted paragraph (its mark
    and run deleted) goes, and each deleted row goes."""
    for element in [*_subpoint_paragraphs(gen), *_subpoint_rows(gen)]:
        element.getparent().remove(element)
    return gen.doc.element.body


@pytest.mark.parametrize("entry", [_clinical_role, _clinical_role_table_row])
def test_clinical_leadership_prose_is_a_deleted_paragraph_under_the_role_line(entry):
    gen = _render_clinical_roles([entry()], supplementary_subpoints=True)

    (para,) = _subpoint_paragraphs(gen)
    above = para.getprevious()
    assert _paragraph_text(above) == _CLINICAL_LINE
    assert "".join(t.text for t in para.iter(qn("w:delText"))) == _CLINICAL_PROSE
    assert list(para.iter(qn("w:t"))) == []
    assert para.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None
    # Unbulleted, one indent step under the role's line (#1943).
    assert para.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None
    assert _left_indent(above) == str(plain_entry_left_twips(0))
    assert _left_indent(para) == str(plain_entry_left_twips(1))
    span = gen._subpoint_section_spans()["L3"]
    assert span.start < list(gen.doc.element.body).index(para) < span.end
    assert gen.stats["supplementary_subpoints"] == 1
    (placement,) = gen._subpoint_placements
    assert (placement.element_idx_start, placement.code) == (entry()["element_idx_start"], "L3")


def test_a_roles_several_prose_lines_are_paragraphs_in_source_order():
    lines = (_CLINICAL_PROSE, "Recruited four wound care nurses and launched the limb preservation clinic")
    gen = _render_clinical_roles([_clinical_role(prose="\t".join(lines))], supplementary_subpoints=True)

    paragraphs = _subpoint_paragraphs(gen)
    assert [_paragraph_text(p) for p in paragraphs] == list(lines)
    assert _paragraph_text(paragraphs[0].getprevious()) == _CLINICAL_LINE
    assert paragraphs[1].getprevious() is paragraphs[0]


@pytest.mark.parametrize("later", [
    # One title held twice at one hospital: the lines differ in their years only.
    _CLINICAL_ROLE,
    # A later role holding every word of this one's title.
    f"{_CLINICAL_ROLE} and Example Limb Preservation Service",
])
def test_clinical_leadership_prose_goes_under_its_own_line_among_roles_of_one_hospital(later):
    """The later role renders first and holds every distinctive word of the
    earlier one's line, so words alone put the earlier one's prose under the
    later line. Only the line the writer recorded for the earlier entry
    takes it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_clinical_roles([_later_clinical_role(later), _clinical_role()], supplementary_subpoints=True)

    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para.getprevious()) == _CLINICAL_LINE
    assert _paragraph_text(para.getprevious().getprevious()) == f"{later}, {_CLINICAL_SITE}\t2016-present"
    # Words alone would have taken the later role's line.
    lines = list(supplementary.anchor_candidates(_render_clinical_roles(
        [_later_clinical_role(later), _clinical_role()], supplementary_subpoints=False).doc))
    roles = [line for line in lines if line.text.endswith(("2015", "present"))]
    assert supplementary.best_anchor(supplementary._value_tokens(_clinical_role(), "L3"), roles) is roles[0]
    assert roles[0].text.startswith(later) and roles[0].text.endswith("present")


def test_roles_at_one_hospital_each_take_their_own_prose():
    later_prose = ("Ran the weekly limb salvage conference and the hospital's hyperbaric "
                   "oxygen therapy service")
    later = _clinical_role(idx=1478, start="2016", end="present", prose=later_prose)
    gen = _render_clinical_roles([later, _clinical_role()], supplementary_subpoints=True)

    by_text = {_paragraph_text(p): p for p in _subpoint_paragraphs(gen)}
    assert set(by_text) == {later_prose, _CLINICAL_PROSE}
    assert _paragraph_text(by_text[later_prose].getprevious()) == _LATER_CLINICAL_LINE
    assert _paragraph_text(by_text[_CLINICAL_PROSE].getprevious()) == _CLINICAL_LINE
    assert [(p.element_idx_start, p.paragraphs) for p in gen._subpoint_placements] \
        == [(1478, (later_prose,)), (1482, (_CLINICAL_PROSE,))]


def test_clinical_leadership_prose_goes_under_its_own_line_not_a_clinical_practice_line():
    """The section's L1 lines sit above the L3 ones: an attending post at the
    same programme holds every distinctive word of the role's line, and
    words alone take it. The role's own line takes the prose."""
    practice = {"element_idx_start": 1401, "taxonomy_code": "L1",
                "text": f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}, attending wound care sessions",
                "extracted_fields": {}}
    grouped = {"L1": [practice], "L3": [_clinical_role()]}
    gen = _generator(supplementary_subpoints=True)
    gen._fill_clinical_practice(grouped)
    gen._add_supplementary_subpoints(grouped)

    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para.getprevious()) == _CLINICAL_LINE
    # Words alone would have taken the practice line, which renders first.
    from unified_pipeline.stage6 import supplementary

    lines = [line for line in supplementary.anchor_candidates(gen.doc) if _CLINICAL_SITE in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_clinical_role(), "L3"), lines) is lines[0]
    assert lines[0].text.endswith("sessions")


def test_clinical_leadership_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off = _render_clinical_roles([_later_clinical_role(), _clinical_role()], supplementary_subpoints=False)
    on = _render_clinical_roles([_later_clinical_role(), _clinical_role()], supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _CLINICAL_PROSE in rejected
    assert _CLINICAL_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_paragraphs(on)) == 1 and _subpoint_rows(on) == []
    accepted = _accept_all(on)
    # Empty paragraphs and every table row, cell by cell, included; and the
    # markup itself, so no paragraph's properties moved.
    assert _texts_of(accepted) == _texts_of(off.doc.element.body)
    assert _row_texts(accepted) == _row_texts(off.doc.element.body)
    assert etree.tostring(accepted, method="c14n") == etree.tostring(off.doc.element.body, method="c14n")


def test_a_short_roles_prose_goes_under_its_own_line():
    """A role and a hospital of acronyms hold no distinctive word, too few
    for the word test: the writer's record says the one line is the entry's
    own (`_pick_anchor`)."""
    gen = _render_clinical_roles([_clinical_role(name="MD", site="NGH")], supplementary_subpoints=True)

    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para.getprevious()) == "MD, NGH\t2010-2015"
    assert _deleted_texts(gen) == _CLINICAL_PROSE


def _role_titled_by_its_prose():
    """The prose is the role stage 4 extracted."""
    return _clinical_role(name=_CLINICAL_PROSE, prose=None)


def _role_with_no_role():
    """No role: the line is titled from the entry's text, and the parts after
    the first print as its detail lines."""
    return _clinical_role(name=None)


@pytest.mark.parametrize("entry", [_role_titled_by_its_prose, _role_with_no_role])
def test_clinical_leadership_prose_already_in_its_line_is_not_offered_again(entry):
    """The line prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_clinical_roles([entry()], supplementary_subpoints=True)
    assert _CLINICAL_PROSE in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "L3", haystack) == []


def _role_with_only(field, marker):
    entry = _clinical_role(name=None, site=None, start=None, end=None, prose="")
    entry["extracted_fields"] = {field: marker}
    entry["text"] = "Example clinical leadership role"
    return entry


def test_every_clinical_leadership_line_field_the_pass_counts_is_printed_in_the_line():
    """The off-schema names among the fields `_value_tokens` counts for a
    clinical leadership role (`_RENDERED_FIELDS` marks schema fields only):
    each one, alone on an entry, is counted and reaches the line."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    off_schema = [name for name in (*supplementary.CLINICAL_LEADERSHIP_ROLE_FIELDS,
                                    *supplementary.CLINICAL_LEADERSHIP_INSTITUTION_FIELDS,
                                    *supplementary._CLINICAL_LEADERSHIP_PLAIN_FIELDS)
                  if name not in _RENDERED_FIELDS["L3"]]
    assert off_schema == ["role", "title", "organization"]
    for field in off_schema:
        marker = f"Zq{field}Z"
        entry = _role_with_only(field, marker)
        assert supplementary._tokens(marker) <= supplementary._value_tokens(entry, "L3"), field
        gen = _render_clinical_roles([entry], supplementary_subpoints=True)
        assert any(marker in text for text in _texts_of(gen.doc.element.body)), field


def test_every_field_the_clinical_leadership_line_prints_is_one_the_pass_counts():
    """The other direction: a field the line prints that `_value_tokens`
    does not count would offer a fragment restating the line. The names are
    the fields of the code an entry is rerouted into L3 from (D2), of the
    codes stage 3b confuses with it (O, D1, D3, Q1, K3), the schema's
    optional ones and the writer's own, each alone on an entry."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    names = set().union(*(_RENDERED_FIELDS[code] for code in ("D2", "O", "D1", "D3", "Q1", "K3")
                          if code in _RENDERED_FIELDS),
                        supplementary.CLINICAL_LEADERSHIP_ROLE_FIELDS,
                        supplementary.CLINICAL_LEADERSHIP_INSTITUTION_FIELDS,
                        {"description", "program", "scope", "narrative", "position", "department"})
    for field in sorted(names):
        marker = f"Zq{field}Z"
        entry = _role_with_only(field, marker)
        gen = _render_clinical_roles([entry], supplementary_subpoints=True)
        printed = any(marker in text for text in _texts_of(gen.doc.element.body))
        counted = supplementary._tokens(marker) <= supplementary._value_tokens(entry, "L3")
        assert not printed or counted, field


_SHADOWED_NAME = "Saltmarsh restoration and estuarine sediment track"


@pytest.mark.parametrize("printed, shadowed", [("role", "leadership_role"),
                                               ("institution", "organization")])
def test_a_line_restating_the_name_its_line_does_not_print_is_offered(printed, shadowed):
    """The line prints the first of its aliases holding a value
    (`clinical_leadership_cell_values`). The other one's words are in no
    line of the document, so a line restating it is the role's prose."""
    entry = _clinical_role(prose=_SHADOWED_NAME, **{printed: "Example Geochemistry",
                                                    shadowed: _SHADOWED_NAME})
    gen = _render_clinical_roles([entry], supplementary_subpoints=True)

    (para,) = _subpoint_paragraphs(gen)
    assert "Example Geochemistry" in _paragraph_text(para.getprevious())
    assert _SHADOWED_NAME not in _paragraph_text(para.getprevious())
    assert _paragraph_text(para) == _SHADOWED_NAME


def _clinical_table_generator(**kwargs):
    """A template that ships a "Title | Institution/Location | Dates" table
    under the heading: the section writes each role as a row of it
    (`_clinical_header_match`)."""
    gen = _generator(**kwargs)
    gen.doc = Document()
    gen.doc.add_paragraph().add_run("Clinical Leadership").bold = True
    table = gen.doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Title", "Institution/Location", "Dates (yyyy)"),
                          strict=True):
        cell.text = text
    gen.doc.add_paragraph("M. RESEARCH")
    return gen, table


def _render_clinical_table(entries, **kwargs):
    gen, table = _clinical_table_generator(**kwargs)
    gen._fill_clinical_practice({"L3": entries})
    gen._add_supplementary_subpoints({"L3": entries})
    return gen, table


def test_clinical_leadership_prose_in_a_clinical_table_is_a_deleted_row_under_the_roles_own_row():
    """The table form: the writer records each role's row, so the prose is a
    deleted row under the role's own row, not under the later role's row of
    the same words, and the row spans the table."""
    gen, table = _render_clinical_table([_later_clinical_role(), _clinical_role()],
                                        supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert row.getparent() is table._tbl
    assert _row_texts_of(row.getprevious()) == ["2010-2015", _CLINICAL_ROLE, _CLINICAL_SITE]
    assert _row_texts_of(row.getprevious().getprevious())[0] == "2016-Present"
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _CLINICAL_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "3"
    (placement,) = gen._subpoint_placements
    assert (placement.element_idx_start, placement.code) == (1482, "L3")


def test_clinical_table_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it():
    off, _ = _render_clinical_table([_later_clinical_role(), _clinical_role()], supplementary_subpoints=False)
    on, _ = _render_clinical_table([_later_clinical_role(), _clinical_role()], supplementary_subpoints=True)

    assert _CLINICAL_PROSE in _deleted_texts(on)
    assert len(_subpoint_rows(on)) == 1 and _subpoint_paragraphs(on) == []
    accepted = _accept_all(on)
    assert _texts_of(accepted) == _texts_of(off.doc.element.body)
    assert _row_texts(accepted) == _row_texts(off.doc.element.body)
    assert etree.tostring(accepted, method="c14n") == etree.tostring(off.doc.element.body, method="c14n")


@pytest.mark.parametrize("field", ["description", "program"])
def test_a_description_the_table_prints_is_not_offered_and_the_line_form_offers_it(field):
    """The table form prints a role's description or programme in its third
    cell, so a line restating it is in the document. The line form prints
    neither: the same line is the role's prose, under its line."""
    gen, _table = _render_clinical_table([_clinical_role(**{field: _CLINICAL_PROSE})],
                                         supplementary_subpoints=True)
    assert any(_CLINICAL_PROSE in "".join(cells) for cells in _row_texts(gen.doc.element.body))
    assert _deleted_texts(gen) == ""

    gen = _render_clinical_roles([_clinical_role(**{field: _CLINICAL_PROSE})], supplementary_subpoints=True)
    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para) == _CLINICAL_PROSE


def _roles_each_opening_with_their_dates():
    """A directorship and the chief's post after it on one line, each opening
    with its years and followed by its prose; stage 4 returned the first.
    Which one a prose line is about cannot be told."""
    return _clinical_role(prose=None) | {
        "text": (f"2010-2015 {_CLINICAL_ROLE}, {_CLINICAL_SITE}\t{_CLINICAL_PROSE}\t"
                 "2016-2020 Chief, Example Limb Preservation Service, Kestrel Institute\t"
                 "Ran the weekly limb salvage conference and the hyperbaric oxygen service")}


def _a_role_then_the_next_dated_at_its_end():
    """A directorship, then on the same line the chief's post after it, its
    years ending its part; stage 4 returned the directorship."""
    return _clinical_role(prose=None) | {
        "text": (f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}, 2010-2015\t"
                 "Chief, Example Limb Preservation Service, Kestrel Institute, 2016-2020")}


def _a_role_then_the_next_dated_in_months():
    """The same, the chief's post dated in months."""
    return _clinical_role(prose=None) | {
        "text": (f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}, 2010-2015\t"
                 "Chief, Example Limb Preservation Service, Kestrel Institute, July 2016 - June 2020")}


@pytest.mark.parametrize("roles", [_roles_each_opening_with_their_dates,
                                   _a_role_then_the_next_dated_at_its_end,
                                   _a_role_then_the_next_dated_in_months])
def test_a_line_of_several_dated_roles_keeps_accept_all_the_flag_off_document(roles):
    """#1205 comment 6275741663 and #2010: another role on the role's line is
    never left in neither document. Opening with its years, the line offers
    nothing (`_own_record_line`); dated as a tab part the #221 recovery
    re-emits, it is the recovery's (`_left_to_recovery`). Either way the
    recovery re-emits with sub-points on what it re-emits with them off, and
    Accept All is the flag-off document."""
    def recovered(flag):
        entry = roles()
        gen = _render_clinical_roles([entry], supplementary_subpoints=flag)
        assert len(gen._row_receipts.elements_for(entry)) == 1
        gen._recover_unrendered_records({"L3": [entry]})
        return gen

    on, off = recovered(True), recovered(False)
    assert on.stats["unrendered_records_recovered"] == off.stats["unrendered_records_recovered"] == 1
    assert on._subpoint_placements == [] and "Limb Preservation" not in _deleted_texts(on)
    accepted = _accept_all(on)
    assert _texts_of(accepted) == _texts_of(off.doc.element.body)


def test_a_role_the_section_writes_no_line_for_gets_no_sub_point():
    """An entry no line was written for -- one the section skipped -- has no
    line of its own: its prose goes under no other line of the document,
    though the other role's line holds every distinctive word of its line."""
    skipped, later = _clinical_role(), _later_clinical_role()
    gen = _generator(supplementary_subpoints=True)
    gen._fill_clinical_practice({"L3": [later]})
    gen._add_supplementary_subpoints({"L3": [later, skipped]})

    assert gen._row_receipts.elements_for(skipped) == []
    assert gen._row_receipts.elements_for(later)
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


def test_a_structural_label_the_section_skips_gets_no_sub_point():
    """The writer skips a source heading stage 3b coded L3
    (`_is_structural_label`): it writes no line, so none of its words go
    under the role line that shares them."""
    label = {"element_idx_start": 1470, "taxonomy_code": "L3",
             "extracted_fields": {"leadership_role": None, "institution": None},
             "text": "CLINICAL LEADERSHIP POSITIONS HELD AT NORVALE GENERAL HOSPITAL"}
    gen = _render_clinical_roles([label, _later_clinical_role()], supplementary_subpoints=True)

    assert gen._row_receipts.elements_for(label) == []
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


def test_the_clinical_leadership_span_starts_at_the_heading_its_writer_looks_up():
    """The span starts at the heading the writer looks the section up by,
    the one the code's heading lookup names, and holds the role's line and
    its sub-point."""
    gen = _render_clinical_roles([_clinical_role()], supplementary_subpoints=True)
    span = gen._subpoint_section_spans()["L3"]
    heading = gen.doc.element.body[span.start]
    assert _paragraph_text(heading).strip() == WCMTemplateGenerator._get_wcm_section_header("L3") \
        == "Clinical Leadership"
    (para,) = _subpoint_paragraphs(gen)
    assert span.start < list(gen.doc.element.body).index(para) < span.end


def test_a_template_with_no_clinical_leadership_heading_places_no_clinical_leadership_prose():
    """Without the heading neither the writer nor the span finds the
    section: no line is written, and the prose goes under no other line."""
    gen = _generator(supplementary_subpoints=True)
    gen.doc = Document()
    gen.doc.add_paragraph("A. PERSONAL DATA")
    entry = _clinical_role()
    gen._fill_clinical_practice({"L3": [entry]})
    gen._add_supplementary_subpoints({"L3": [entry]})

    assert "L3" not in gen._subpoint_section_spans()
    assert _deleted_texts(gen) == ""


def _render_grouped_clinical_roles(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_clinical_practice(grouped)
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _clinical_role_line(text, *roles):
    """One source line stage 4 returned as a record per (role, hospital,
    start), the entry's own fields being the last record's. Fan-out gives
    every earlier record a line of its own and the last one the whole line."""
    records = [{"leadership_role": role, "institution": site, "start_date": start}
               for role, site, start in roles]
    return {"element_idx_start": 663, "taxonomy_code": "L3", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_CLINICAL_CHIEF = ("Chief, Example Limb Preservation Service", "Kestrel Institute", "2007")
_CLINICAL_DIRECTOR = (_CLINICAL_ROLE, _CLINICAL_SITE, "2010")
_CLINICAL_CHIEF_LINE = "Chief, Example Limb Preservation Service, Kestrel Institute, 2007"
_CLINICAL_DIRECTOR_LINE = f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}, 2010"
_CLINICAL_CHIEF_PROSE = ("Ran the weekly limb salvage conference, the hyperbaric oxygen service and its "
                "wound outcomes registry")


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own role.
    ([f"Founding {_CLINICAL_CHIEF_LINE}. {_CLINICAL_CHIEF_PROSE}", _CLINICAL_DIRECTOR_LINE],
     [(_CLINICAL_CHIEF_LINE, f"Founding {_CLINICAL_CHIEF_LINE}. {_CLINICAL_CHIEF_PROSE}")]),
    ([_CLINICAL_CHIEF_LINE, f"{_CLINICAL_PROSE} (interim, {_CLINICAL_DIRECTOR_LINE})"],
     [(_CLINICAL_DIRECTOR_LINE, f"{_CLINICAL_PROSE} (interim, {_CLINICAL_DIRECTOR_LINE})")]),
    # They do not: which role the prose is about cannot be told, so neither takes it.
    ([_CLINICAL_CHIEF_LINE, _CLINICAL_PROSE, _CLINICAL_DIRECTOR_LINE], []),
])
def test_prose_on_a_line_of_several_clinical_roles_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the role's line as well as its prose brings the role, hospital and
    year along, and Reject All prints them a second time, under the line
    that holds them (the shape every family has; #1205, 2026-10-09)."""
    entry = _clinical_role_line("\t".join(segments), _CLINICAL_CHIEF, _CLINICAL_DIRECTOR)
    gen, grouped = _render_grouped_clinical_roles([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["L3"][-1] and len(grouped["L3"]) == 2
    assert [(_paragraph_text(p.getprevious()), _paragraph_text(p)) for p in _subpoint_paragraphs(gen)] \
        == offered


def test_prose_opening_with_the_roles_line_as_it_prints_reads_as_rendered():
    """A residual, every family's `_record_rendered` rule: a fragment whose
    opening stretch prints verbatim reads as rendered. The line form prints
    a role as "Role, Institution, Dates", the order a CV often writes it in,
    so prose sharing its segment with the line as it prints is not offered:
    it is in neither document, with sub-points on as off (dev)."""
    entry = _clinical_role(prose=None)
    entry["text"] = f"{_CLINICAL_LINE}. {_CLINICAL_PROSE}"
    on = _render_clinical_roles([entry], supplementary_subpoints=True)
    off = _render_clinical_roles([copy.deepcopy(entry)], supplementary_subpoints=False)

    assert _deleted_texts(on) == "" and _texts_of(on.doc.element.body) == _texts_of(off.doc.element.body)
    assert _CLINICAL_PROSE not in "".join(_texts_of(off.doc.element.body))


def test_generic_list_fan_out_children_of_roles_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, and the entry renders its
    own fields."""
    records = [{"leadership_role": role, "institution": site, "start_date": start}
               for role, site, start in (_CLINICAL_CHIEF, _CLINICAL_DIRECTOR)]
    for text, children in ((f"{_CLINICAL_CHIEF_LINE}\t{_CLINICAL_DIRECTOR_LINE}", 2),
                           (f"{_CLINICAL_CHIEF_LINE}\t{_CLINICAL_DIRECTOR_LINE}\t{_CLINICAL_PROSE}", 1)):
        entry = {"element_idx_start": 664, "taxonomy_code": "L3", "text": text,
                 "extracted_fields": {"roles": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_clinical_roles([entry], supplementary_subpoints=True)
        assert len(grouped["L3"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["L3"])
        if children == 2:
            assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("N1", "MENTORING PROGRAMS"), ("S0", "BIBLIOGRAPHY")])
def test_a_role_rerouted_out_of_l3_to_an_excluded_code_gets_no_sub_point(code, heading):
    """N1 and S0 are excluded: a role rendered as either is that code's
    record, so neither the role line of the same hospital nor any other
    carries its prose."""
    gen, grouped = _render_grouped_clinical_roles(
        [_later_clinical_role(), _rerouted(_clinical_role(), code, heading)], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [1482]
    assert [e["element_idx_start"] for e in grouped["L3"]] == [1478]
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("F2", "Board Certification"),
                                           ("C", "POSTDOCTORAL TRAINING")])
def test_a_refused_reroute_keeps_the_clinical_leadership_rule(code, heading):
    """A role rerouted to F2, whose span the heading lookup does not find,
    or to C keeps its L3 code: their renderers would drop its role
    (`_cross_family_refusal`). It renders as an L3 line and its prose is a
    deleted paragraph under that line."""
    gen, grouped = _render_grouped_clinical_roles([_rerouted(_clinical_role(), code, heading)],
                                                  supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["L3"]] == [1482]
    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para.getprevious()) == _CLINICAL_LINE


def _hospital_paragraph_with_a_directorship():
    """A hospital-employment paragraph stage 3b coded D2, which stage 4 split
    into an attending post and a directorship, the directorship's segment
    going on into what it oversaw (#1731, TJDPJE's shape)."""
    attending = {"title": "Attending Physician", "institution": _CLINICAL_SITE, "start_date": "2008"}
    director = {"title": _CLINICAL_ROLE, "institution": _CLINICAL_SITE, "start_date": "2010"}
    return {"element_idx_start": 46, "taxonomy_code": "D2",
            "text": (f"Attending Physician, {_CLINICAL_SITE}, 2008\t"
                     f"{_CLINICAL_PROSE} ({_CLINICAL_ROLE}, {_CLINICAL_SITE}, 2010)"),
            "extracted_fields": {**director, STAGE4_RECORDS_KEY: [attending, director]}}


def test_a_record_rerouted_into_l3_follows_the_clinical_leadership_rule():
    """A split D2 record whose title opens with a leadership title renders
    under Clinical Leadership (#1731): its prose is a deleted paragraph under
    its L3 line, recorded as L3. The attending post stays D2."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_hospital_paragraph_with_a_directorship()])
    gen._fill_positions(grouped)
    gen._fill_clinical_practice(grouped)
    gen._add_supplementary_subpoints(grouped)

    (role,) = grouped["L3"]
    assert role["taxonomy_code_original"] == "D2" and [e["extracted_fields"]["title"]
                                                       for e in grouped["D2"]] == ["Attending Physician"]
    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para.getprevious()).startswith(f"{_CLINICAL_ROLE}, {_CLINICAL_SITE}, 2010")
    assert _CLINICAL_PROSE in _paragraph_text(para)
    (placement,) = gen._subpoint_placements
    assert placement.code == "L3"


def test_a_role_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _clinical_role()
    withheld["_pii_withheld"] = True
    gen = _render_clinical_roles([_later_clinical_role(), withheld], supplementary_subpoints=True)

    assert _CLINICAL_LINE in [_paragraph_text(p) for p in gen.doc.element.body.iter(qn("w:p"))]
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


def test_a_role_dedup_drops_offers_its_prose_once_under_the_kept_line():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no line, so its
    prose goes only under the line of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_clinical_role(idx=1482), _clinical_role(idx=1483)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["L3"]) == 2 and len(kept["L3"]) == 1
    gen._fill_clinical_practice(kept)
    gen._add_supplementary_subpoints(before)

    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para.getprevious()) == _CLINICAL_LINE
    assert _deleted_texts(gen) == _CLINICAL_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["L3"][0]["element_idx_start"]


@pytest.mark.parametrize("twin_prose", [_CLINICAL_PROSE,
                                        "Chaired the hospital's limb salvage quality review panel"],
                         ids=["same_prose", "other_prose"])
def test_a_role_the_record_dedup_drops_as_an_affiliates_twin_gets_no_second_sub_point(twin_prose):
    """`record_dedup._same_leadership` (EQGGRB-07's shape): one role and unit
    in one span at a hospital and its university is one record, and the
    sparser copy is dropped. It wrote no line, so it gets no sub-point of
    its own: the kept copy's prose is written once, under the kept line."""
    unit = "Example Wound Care Program"
    kept_role = _clinical_role(unit_program=unit)
    twin = _clinical_role(idx=1490, site="Norvale University", prose=twin_prose)
    twin["extracted_fields"]["unit_program"] = unit
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([kept_role, twin])
    kept, before, decisions = gen._dedup_grouped_entries(grouped)
    assert [e["element_idx_start"] for e in kept["L3"]] == [1482]
    assert any(d["metric"].endswith("leadership_affiliate") for d in decisions)
    gen._fill_clinical_practice(kept)
    gen._add_supplementary_subpoints(before)

    assert gen._row_receipts.elements_for(before["L3"][1]) == []
    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para) == _CLINICAL_PROSE
    assert [p.element_idx_start for p in gen._subpoint_placements] == [1482]


def _low_coverage_role(*parts):
    """A role whose line stage 4 covered at 10%: the content overflow takes
    it, and with sub-points off writes its whole text as blue text in the
    section, in place of the role's line."""
    entry = _clinical_role(prose="\t".join(parts))
    entry["extraction_coverage"] = {"extraction_coverage_percent": 10.0}
    return entry


def _render_low_coverage_role(entry, **kwargs):
    gen = _generator(**kwargs)
    gen._fill_clinical_practice({"L3": [entry]})
    # The entry the line's comment hook queued is the one stage 6 grouped.
    assert [queued for queued, _para, _code in gen._overflow_entries] == [entry]
    assert gen._overflow_entries[0][0] is entry
    gen._place_unrendered_prose({"L3": [entry]})
    return gen


def test_a_low_coverage_role_whose_sub_point_holds_all_its_prose_keeps_its_line():
    """The sub-point and the role's line hold every part of the role's text,
    so the overflow pass skips the entry (`yields_to_overflow`): the line
    stays, and the prose is the sub-point only, not also blue overflow text
    in the line's place."""
    blocks = [f"{_CLINICAL_PROSE}, block {n} of the service plan" for n in range(9)]
    gen = _render_low_coverage_role(_low_coverage_role(*blocks), supplementary_subpoints=True)

    paragraphs = _subpoint_paragraphs(gen)
    assert [_paragraph_text(p) for p in paragraphs] == blocks
    assert _paragraph_text(paragraphs[0].getprevious()) == _CLINICAL_LINE
    assert gen.stats["overflow_bullets_added"] == 0
    assert not any(_CLINICAL_PROSE in text for text in _texts_of(gen.doc.element.body))
    assert gen._appendix_pending == []


@pytest.mark.parametrize("parts", [
    ("Budget: $4M",),                            # too short to be offered
    ("Staff:", "12 nurses"),                     # a label and a count
    ("2013: opened the limb salvage clinic",),   # a date-led line
    ("Department of Surgery, Norvale General Hospital",),  # a unit the line does not print
], ids=["short", "label", "date_led", "unprinted_unit"])
def test_a_low_coverage_role_with_a_part_no_sub_point_holds_keeps_its_overflow(parts):
    """#1205 review: the planner offers the long prose but not these parts.
    With the overflow skipped for the sub-point, they were in neither Accept
    All nor Reject All. The role gets no sub-point, and the render is the
    one with sub-points off: the overflow writes every part in the section."""
    prose = (_CLINICAL_PROSE, "Recruited four wound care nurses and launched the limb "
             "preservation clinic for veterans with the affiliated medical center")
    entry = _low_coverage_role(prose[0], *parts, prose[1])
    on = _render_low_coverage_role(entry, supplementary_subpoints=True)
    off = _render_low_coverage_role(copy.deepcopy(entry), supplementary_subpoints=False)

    assert _subpoint_paragraphs(on) == [] and on._subpoint_placements == []
    assert on.stats["overflow_bullets_added"] == 1
    texts = _texts_of(on.doc.element.body)
    assert all(part in texts for part in (*prose, *parts))
    assert etree.tostring(on.doc.element.body, method="c14n") \
        == etree.tostring(off.doc.element.body, method="c14n")


def test_the_doctor_matches_a_clinical_leadership_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier role's sub-point
    under its own index and L3, and the doctor's per-entry match counts it
    for that entry, not for the later role that shares its words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_clinical_role(), _clinical_role()],
                                     uid="T1205L")
    assert placements == [[_CLINICAL_PROSE]]
    sidecar = json.loads((tmp_path / "T1205L_render_warnings.json").read_text())
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(1482, "L3")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_CLINICAL_PROSE, subpoints, element_idx=1482)
    assert not offered_as_subpoint(_CLINICAL_PROSE, subpoints, element_idx=1478)


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_clinical_leadership_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line (#1205 review of O). The same role without them still has its
    sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_clinical_role()])
    assert placements == [[_CLINICAL_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_clinical_role(), _CLINICAL_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


# --- #1205: Q4B/Q4C editorial board memberships, the next code family --------
# The service section's Other Service writer writes every Q4B and Q4C entry
# as a row of the Editorial Board Membership table, "Board / Organization Name
# | Dates", the role and the journal joined in its first cell. So the prose is
# a deleted row under the seat's own row, the one the writer recorded for
# that entry (the row-receipt ledger).

_EDITORIAL_PROSE = ("Edited the special issue on sediment budgets of restored tidal marshes "
                    "and recruited its invited reviews")
_EDITORIAL_JOURNAL = "Journal of Estuarine Sediment Studies"
_EDITORIAL_SEAT = "Associate Editor"


def _board_seat(idx=104, code="Q4B", role=_EDITORIAL_SEAT, journal=_EDITORIAL_JOURNAL, start="2012",
                end="2016", prose=_EDITORIAL_PROSE, **fields):
    """A seat whose source line goes on past its role, journal and years into
    the special issue it edited, as EBYSBC QITQWH 196's did. A later seat on
    the same journal shares every distinctive word of its row: only the
    years, digits, tell them apart. A value passed as None is in neither the
    line nor the fields."""
    named = {"journal_name": journal, "role": role, "start_date": start, "end_date": end}
    dates = "-".join(value for value in (start, end) if value)
    line = ", ".join(value for value in (role, journal, dates) if value)
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": line + (f"\t{prose}" if prose else ""),
        "extracted_fields": {key: value for key, value in {**named, **fields}.items()
                             if value is not None},
    }


def _later_board_seat(code="Q4B", role=_EDITORIAL_SEAT):
    """The same journal's seat from 2017, with no prose: it renders first,
    reverse chronologically."""
    return _board_seat(idx=98, code=code, role=role, start="2017", end="present", prose="")


def _board_member(idx=104, **kwargs):
    """A Q4C seat: its schema names the journal and the years, no role."""
    return _board_seat(idx=idx, code="Q4C", role=None, **kwargs)


def _render_editorial_boards(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = {}
    for entry in entries:
        grouped.setdefault(entry["taxonomy_code"], []).append(entry)
    gen._fill_service(grouped)
    gen._add_supplementary_subpoints(grouped)
    return gen


def _board_seat_table_row():
    """The same seat read from a source table: stage 2 joins the row's cells
    with pipes, and its position names the row ("41.3")."""
    entry = _board_seat()
    entry["element_idx_start"] = "41.3"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


_SEAT_ROW = [f"{_EDITORIAL_SEAT}, {_EDITORIAL_JOURNAL}", "2012-2016"]
_MEMBER_ROW = [_EDITORIAL_JOURNAL, "2012-2016"]
_EDITORIAL_HEADER = "Board / Organization NameDates (yyyy-yyyy)"


@pytest.mark.parametrize("entry, cells, code", [(_board_seat, _SEAT_ROW, "Q4B"),
                                                (_board_seat_table_row, _SEAT_ROW, "Q4B"),
                                                (_board_member, _MEMBER_ROW, "Q4C")],
                         ids=["Q4B", "Q4B_table_row", "Q4C"])
def test_editorial_board_prose_is_a_deleted_row_under_the_seats_row(entry, cells, code):
    gen = _render_editorial_boards([entry()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == cells
    assert _table_header(row) == _EDITORIAL_HEADER
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == _EDITORIAL_PROSE
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == "2"
    assert gen.stats["supplementary_subpoints"] == 1
    (placement,) = gen._subpoint_placements
    assert (placement.element_idx_start, placement.code) == (entry()["element_idx_start"], code)


def test_a_seats_several_prose_lines_are_one_row_of_paragraphs_in_source_order():
    lines = (_EDITORIAL_PROSE, "Handled the manuscripts on coastal hydrology for the journal's methods section")
    gen = _render_editorial_boards([_board_seat(prose="\t".join(lines))], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == _SEAT_ROW
    assert ["".join(t.text for t in p.iter(qn("w:delText"))) for p in row.iter(qn("w:p"))] == list(lines)


@pytest.mark.parametrize("make", [_board_seat, _board_member], ids=["Q4B", "Q4C"])
def test_editorial_board_prose_goes_under_its_own_row_among_seats_on_one_journal(make):
    """The later seat renders first and holds every distinctive word of the
    earlier one's row, so words alone put the earlier one's prose under the
    later row. Only the row the writer recorded for the earlier entry takes
    it."""
    from unified_pipeline.stage6 import supplementary

    code = make()["taxonomy_code"]
    later = make(idx=98, start="2017", end="present", prose="")
    gen = _render_editorial_boards([later, make()], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious())[1] == "2012-2016"
    assert _row_texts_of(row.getprevious().getprevious())[1] == "2017-Present"
    # Words alone would have taken the later seat's row.
    lines = list(supplementary.anchor_candidates(_render_editorial_boards(
        [make(idx=98, start="2017", end="present", prose=""), make()], supplementary_subpoints=False).doc))
    rows = [line for line in lines if _EDITORIAL_JOURNAL in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(make(), code), rows) is rows[0]
    assert rows[0].text.endswith("2017-Present")


def test_seats_on_one_journal_each_take_their_own_prose():
    later_prose = "Chaired the journal's reviewer training workshops for early career scientists"
    later = _board_seat(idx=98, start="2017", end="present", prose=later_prose)
    gen = _render_editorial_boards([later, _board_seat()], supplementary_subpoints=True)

    by_text = {"".join(t.text for t in row.iter(qn("w:delText"))): row for row in _subpoint_rows(gen)}
    assert set(by_text) == {later_prose, _EDITORIAL_PROSE}
    assert _row_texts_of(by_text[later_prose].getprevious())[1] == "2017-Present"
    assert _row_texts_of(by_text[_EDITORIAL_PROSE].getprevious()) == _SEAT_ROW
    assert [(p.element_idx_start, p.paragraphs) for p in gen._subpoint_placements] \
        == [(98, (later_prose,)), (104, (_EDITORIAL_PROSE,))]


@pytest.mark.parametrize("entries", [lambda: [_later_board_seat(), _board_seat()],
                                     lambda: [_board_member(idx=98, start="2017", end="present", prose=""),
                                              _board_member(),
                                              _board_seat(idx=110, start="2009", end="2011", prose=_GUEST_PROSE)]],
                         ids=["Q4B", "Q4C_and_Q4B"])
def test_editorial_board_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it(entries):
    off = _render_editorial_boards(entries(), supplementary_subpoints=False)
    on = _render_editorial_boards(entries(), supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert _EDITORIAL_PROSE in rejected
    assert _EDITORIAL_PROSE not in "".join(_texts_of(off.doc.element.body))
    assert len(_subpoint_rows(on)) == len(entries()) - 1
    accepted = _accept_all(on)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(accepted) == _texts_of(off.doc.element.body)
    assert _row_texts(accepted) == _row_texts(off.doc.element.body)
    assert etree.tostring(accepted, method="c14n") == etree.tostring(off.doc.element.body, method="c14n")


def test_a_short_seats_prose_goes_under_its_own_row():
    """An acronym journal and a short role hold one distinctive word, too few
    for the word test: the writer's record says the one row is the entry's
    own (`_pick_anchor`)."""
    gen = _render_editorial_boards([_board_seat(role="Editor", journal="JESS")], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == ["Editor, JESS", "2012-2016"]
    assert _deleted_texts(gen) == _EDITORIAL_PROSE


def _seat_named_by_its_prose():
    """The prose is the journal stage 4 extracted."""
    return _board_seat(journal=_EDITORIAL_PROSE, prose=None)


def _seat_titled_by_its_prose():
    """The prose is the role stage 4 extracted, which the cell prefixes."""
    return _board_seat(role=_EDITORIAL_PROSE)


def _seat_whose_section_is_its_prose():
    """The prose is the Q4B section stage 4 extracted, which the cell adds
    after the journal."""
    return _board_seat(section=_EDITORIAL_PROSE)


@pytest.mark.parametrize("entry", [_seat_named_by_its_prose, _seat_titled_by_its_prose,
                                   _seat_whose_section_is_its_prose])
def test_editorial_board_prose_already_in_its_row_is_not_offered_again(entry):
    """The row prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_editorial_boards([entry()], supplementary_subpoints=True)
    assert _EDITORIAL_PROSE in "".join(_row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "Q4B", haystack) == []


def _seat_with_only(code, field, marker):
    """An entry holding `field` alone, its text a line no field holds."""
    entry = _board_seat(code=code, role=None, journal=None, start=None, end=None, prose="")
    entry["extracted_fields"] = {field: marker}
    entry["text"] = "Example editorial line"
    return entry


@pytest.mark.parametrize("code", ["Q4B", "Q4C"])
def test_every_editorial_board_row_field_the_pass_counts_is_printed_in_the_row(code):
    """The off-schema names among the fields `_value_tokens` counts for an
    editorial board row (`_RENDERED_FIELDS` marks schema fields only): each
    one, alone on an entry, is counted and reaches the row."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    off_schema = [name for name in ("role", *supplementary.OTHER_SERVICE_ORGANIZATION_FIELDS)
                  if name not in _RENDERED_FIELDS[code]]
    assert off_schema == (["organization", "committee_name", "agency"] if code == "Q4B"
                          else ["role", "organization", "committee_name", "agency"])
    for field in off_schema:
        marker = f"Zq{field}Z"
        entry = _seat_with_only(code, field, marker)
        assert supplementary._tokens(marker) <= supplementary._value_tokens(entry, code), field
        gen = _render_editorial_boards([entry], supplementary_subpoints=True)
        assert any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr"))), field


@pytest.mark.parametrize("code", ["Q4B", "Q4C"])
def test_every_field_an_editorial_board_row_prints_is_one_the_pass_counts(code):
    """The other direction: a field the row prints that `_value_tokens` does
    not count would offer a fragment restating the row. The names are the
    fields of the codes the section writes beside them (Q1-Q4D), of a
    membership, and the writer's own, each alone on an entry, and beside the
    schema's journal and role."""
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.fan_out import _RENDERED_FIELDS

    names = set().union(*(_RENDERED_FIELDS[c] for c in ("Q1", "Q2", "Q3", "Q4", "Q4A", "Q4B", "Q4C", "Q4D", "I")),
                        supplementary.OTHER_SERVICE_ORGANIZATION_FIELDS,
                        {"title", "name", "publisher", "organization_name", "institution", "panel_name"})
    for field in sorted(names):
        marker = f"Zq{field}Z"
        for entry in (_seat_with_only(code, field, marker),
                      _board_seat(code=code, prose="", **{field: marker})):
            gen = _render_editorial_boards([entry], supplementary_subpoints=True)
            printed = any(marker in _row_text(tr) for tr in gen.doc.element.body.iter(qn("w:tr")))
            counted = supplementary._tokens(marker) <= supplementary._value_tokens(entry, code)
            assert not printed or counted, field


def test_a_line_restating_the_row_with_a_role_it_does_not_print_is_offered():
    """A Q4C row prints no role when stage 4 names none and the writer reads
    none from the line. The role says what kind of board it was and the
    seat's standing, which no cell holds: the line is offered as the CV
    wrote it. A role the row prints is the row's: a line restating it with
    that role is not offered."""
    line = f"Honorary Member, Editorial Advisory Board, {_EDITORIAL_JOURNAL}, 2012-2016"
    entry = _board_member(prose=None)
    entry["text"] = line
    gen = _render_editorial_boards([entry], supplementary_subpoints=True)

    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == _MEMBER_ROW
    assert "".join(t.text for t in row.iter(qn("w:delText"))) == line
    shown = _board_seat(code="Q4C", role="Honorary Member, Editorial Advisory Board", prose=None)
    assert shown["text"] == line
    gen = _render_editorial_boards([shown], supplementary_subpoints=True)
    assert [f"Honorary Member, Editorial Advisory Board, {_EDITORIAL_JOURNAL}", "2012-2016"] \
        in _row_texts(gen.doc.element.body)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_line_restating_the_row_with_only_its_headings_words_is_not_offered():
    """A generic seat ("Editorial Board Member", "Member, Editorial Board")
    names only the words of the Editorial Board Membership heading its row
    sits under. A Q4C row printing no role there says the same, so the line
    is not offered: Reject All would print the journal and the years twice
    and add nothing. An acronym journal, or one of two words, holds too few
    of the line's words for the share test. A seat with a word of its own
    beyond the heading's is still offered."""
    lines = {201: "Editorial Board Member, JESS, 2015-present",
             205: "Member, Editorial Board, Marsh Hydrology, 2010-2014",
             203: "Editorial Advisory Board, Kestrel Ecology, 2021-"}
    entries = [_board_member(idx=201, journal="JESS", start="2015", end="present", prose=None),
               _board_member(idx=205, journal="Marsh Hydrology", start="2010", end="2014", prose=None),
               _board_member(idx=203, journal="Kestrel Ecology", start="2021", end=None, prose=None)]
    for entry in entries:
        entry["text"] = lines[entry["element_idx_start"]]
    gen = _render_editorial_boards(entries, supplementary_subpoints=True)

    rows = _row_texts(gen.doc.element.body)
    assert ["JESS", "2015-Present"] in rows and ["Marsh Hydrology", "2010-2014"] in rows
    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == ["Kestrel Ecology", "2021-Present"]
    assert _deleted_texts(gen) == lines[203]
    assert [p.element_idx_start for p in gen._subpoint_placements] == [203]


def _editorial_template_generator(columns, **kwargs):
    """A template whose Editorial Board Membership table has `columns`
    cells: three or more print the role, the journal and the dates each in
    its own; one prints "<journal> (<dates>)"."""
    doc = Document()
    doc.add_paragraph("EXTRAMURAL PROFESSIONAL RESPONSIBILITIES")
    doc.add_paragraph("Editorial Board Membership")
    table = doc.add_table(rows=1, cols=columns)
    for cell, text in zip(table.rows[0].cells, ("Role", "Journal", "Dates", "Notes")[:columns], strict=True):
        cell.text = text
    run = doc.add_paragraph().add_run("INVITATIONS TO SPEAK/PRESENT")
    run.bold = run.underline = True
    gen = _generator(**kwargs)
    gen.doc = doc  # its paragraphs are the template's (`_find_template_label`)
    return gen, table


@pytest.mark.parametrize("columns, cells", [
    (3, [_EDITORIAL_SEAT, _EDITORIAL_JOURNAL, "2012-2016"]),
    (4, [_EDITORIAL_SEAT, _EDITORIAL_JOURNAL, "2012-2016", ""]),
    (1, [f"{_EDITORIAL_JOURNAL} (2012-2016)"]),
], ids=["three_columns", "four_columns", "one_column"])
def test_editorial_board_prose_goes_under_its_own_row_whatever_the_tables_columns(columns, cells):
    """The writer prints the row's cells by the table's width; the prose is a
    deleted row spanning the table under the seat's own row, and Accept All
    is the flag-off render."""
    def render(flag):
        gen, table = _editorial_template_generator(columns, supplementary_subpoints=flag)
        entries = [_later_board_seat(), _board_seat()]
        gen._fill_service({"Q4B": entries})
        gen._add_supplementary_subpoints({"Q4B": entries})
        return gen, table

    on, table = render(True)
    (row,) = _subpoint_rows(on)
    assert row.getparent() is table._tbl
    assert _row_texts_of(row.getprevious()) == cells
    span = row.find(f"{qn('w:tc')}/{qn('w:tcPr')}/{qn('w:gridSpan')}")
    assert span.get(qn("w:val")) == str(columns)
    off, _ = render(False)
    accepted = _accept_all(on)
    assert _row_texts(accepted) == _row_texts(off.doc.element.body)
    assert etree.tostring(accepted, method="c14n") == etree.tostring(off.doc.element.body, method="c14n")


def test_an_editorial_board_seat_the_section_writes_no_row_for_gets_no_sub_point():
    """An entry no row was written for has no row of its own: its prose goes
    under no other line of the document, though the other seat's row holds
    every distinctive word of its line."""
    skipped, later = _board_seat(), _later_board_seat()
    gen = _generator(supplementary_subpoints=True)
    gen._fill_service({"Q4B": [later]})
    gen._add_supplementary_subpoints({"Q4B": [later, skipped]})

    assert gen._row_receipts.elements_for(skipped) == []
    assert gen._row_receipts.elements_for(later)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_template_with_no_editorial_board_label_places_no_editorial_board_prose():
    """Without the label the writer finds no table and writes no row; the
    EXTRAMURAL span is there, and the prose goes under no other line."""
    gen, _table = _editorial_template_generator(2, supplementary_subpoints=True)
    for paragraph in gen.doc.paragraphs:
        if paragraph.text == "Editorial Board Membership":
            paragraph.text = "Other Roles"
    entries = [_board_seat()]
    gen._fill_service({"Q4B": entries})
    gen._add_supplementary_subpoints({"Q4B": entries})

    assert "Q4B" in gen._subpoint_section_spans()
    assert gen._row_receipts.elements_for(entries[0]) == []
    assert _deleted_texts(gen) == ""


def test_the_editorial_board_span_is_the_extramural_section():
    """The code's heading lookup names the EXTRAMURAL section ('Q' prefix),
    which holds the Editorial Board Membership table, and the lookup is
    unchanged for the reconsider pass and the #221 recovery."""
    gen = _render_editorial_boards([_board_seat()], supplementary_subpoints=True)
    span = gen._subpoint_section_spans()["Q4B"]
    heading = gen.doc.element.body[span.start]
    assert "".join(t.text or "" for t in heading.iter(qn("w:t"))).startswith("EXTRAMURAL PROFESSIONAL")
    (row,) = _subpoint_rows(gen)
    assert span.start < list(gen.doc.element.body).index(row.getparent()) < span.end
    assert gen._subpoint_section_spans()["Q4C"] == span
    assert WCMTemplateGenerator._get_wcm_section_header("Q4B") == "EXTRAMURAL PROFESSIONAL"


def _journal_reviewing_template_generator(**kwargs):
    """A template whose Editorial Board Membership label has no table of its
    own: the next table after it is Journal Reviewing's."""
    doc = Document()
    doc.add_paragraph("EXTRAMURAL PROFESSIONAL RESPONSIBILITIES")
    doc.add_paragraph("Editorial Board Membership")
    doc.add_paragraph("Journal Reviewing/Ad hoc Reviewing")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Journal / Organization Name", "Dates"
    gen = _generator(**kwargs)
    gen.doc = doc
    return gen, table


def test_an_editorial_board_row_in_the_journal_reviewing_table_gets_no_sub_point():
    """Q4D is excluded, and so is every row of the table its writer fills,
    whatever code the entry written there carries: a seat the writer
    appends to that table gets none, and its prose goes under no other
    line."""
    gen, table = _journal_reviewing_template_generator(supplementary_subpoints=True)
    reviewing = {"element_idx_start": 120, "taxonomy_code": "Q4D", "text": "Kestrel Tidal Review, 2015",
                 "extracted_fields": {"journal_name": "Kestrel Tidal Review", "start_date": "2015"}}
    seat = _board_member()
    grouped = {"Q4D": [reviewing], "Q4C": [seat]}
    gen._fill_service(grouped)
    gen._add_supplementary_subpoints(grouped)

    (own,) = gen._row_receipts.elements_for(seat)
    assert own.getparent() is table._tbl and table._tbl in gen._subpoint_excluded_tables
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def _render_grouped_editorial_boards(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_service(grouped)
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


def _board_line(text, *seats):
    """One source line stage 4 returned as a record per (role, journal,
    start), the entry's own fields being the last record's. Fan-out gives
    every earlier record a line of its own and the last one the whole line."""
    records = [{"role": role, "journal_name": journal, "start_date": start} for role, journal, start in seats]
    return {"element_idx_start": 663, "taxonomy_code": "Q4B", "text": text,
            "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}


_GUEST_SEAT = ("Guest Editor", "Kestrel Tidal Review", "2014")
_ASSOCIATE_SEAT = (_EDITORIAL_SEAT, "Kestrel Marsh Letters", "2012")
_GUEST_LINE = "Guest Editor, Kestrel Tidal Review, 2014"
_ASSOCIATE_LINE = f"{_EDITORIAL_SEAT}, Kestrel Marsh Letters, 2012"
_GUEST_PROSE = "Ran the review's special issue on mangrove carbon storage and wrote its editorial"


@pytest.mark.parametrize("segments, offered", [
    # The segments pair with the records: the prose goes under its own seat.
    ([f"{_GUEST_LINE}. {_GUEST_PROSE}", _ASSOCIATE_LINE],
     [("Guest Editor, Kestrel Tidal Review", f"{_GUEST_LINE}. {_GUEST_PROSE}")]),
    ([_GUEST_LINE, f"{_ASSOCIATE_LINE}. {_EDITORIAL_PROSE}"],
     [(f"{_EDITORIAL_SEAT}, Kestrel Marsh Letters", f"{_ASSOCIATE_LINE}. {_EDITORIAL_PROSE}")]),
    # They do not: which seat the prose is about cannot be told, so neither takes it.
    ([_GUEST_LINE, _EDITORIAL_PROSE, _ASSOCIATE_LINE], []),
])
def test_prose_on_a_line_of_several_seats_goes_under_its_own_or_none(segments, offered):
    """A sub-point is the whole tab segment, as the CV wrote it: one that
    holds the seat's line as well as its prose brings the role, journal and
    year along, and Reject All prints them a second time, under the row that
    holds them (the shape every family has; #1205, 2026-10-09)."""
    entry = _board_line("\t".join(segments), _GUEST_SEAT, _ASSOCIATE_SEAT)
    gen, grouped = _render_grouped_editorial_boards([entry], supplementary_subpoints=True)

    assert LAST_STAGE4_RECORD in grouped["Q4B"][-1] and len(grouped["Q4B"]) == 2
    assert [(_row_texts_of(row.getprevious())[0],
             "".join(t.text for t in row.iter(qn("w:delText")))) for row in _subpoint_rows(gen)] \
        == offered


def test_a_seats_line_opening_with_its_rows_first_cell_is_read_as_its_row():
    """A line that opens with the first cell's whole text, role and journal,
    holds the row's own text verbatim (`_record_rendered`'s first
    `RENDER_PIECE_WINDOW` characters), and the whole line reads as rendered:
    none of it is offered, its prose included, as for a talk. A line opening
    otherwise is offered whole, as the CV wrote it."""
    cell_first = f"{_EDITORIAL_SEAT}, {_EDITORIAL_JOURNAL}, 2012-2016. {_EDITORIAL_PROSE}"
    society_first = f"Appointed by the society, {cell_first}"
    for line, offered in ((cell_first, ""), (society_first, society_first)):
        entry = _board_seat(prose=None)
        entry["text"] = line
        gen = _render_editorial_boards([entry], supplementary_subpoints=True)
        assert _deleted_texts(gen) == offered


def test_generic_list_fan_out_children_of_seats_offer_nothing_their_fields_do_not_hold():
    """Fan-out splits a record list stage 4 filed under a key the schema does
    not define only when the fields hold every word of the text, and each
    child's text is then the line built from its fields: nothing is prose.
    With prose in the text the list is not split, and the entry renders its
    own fields."""
    records = [{"role": role, "journal_name": journal, "start_date": start}
               for role, journal, start in (_GUEST_SEAT, _ASSOCIATE_SEAT)]
    for text, children in ((f"{_GUEST_LINE}\t{_ASSOCIATE_LINE}", 2),
                           (f"{_GUEST_LINE}\t{_ASSOCIATE_LINE}\t{_EDITORIAL_PROSE}", 1)):
        entry = {"element_idx_start": 664, "taxonomy_code": "Q4B", "text": text,
                 "extracted_fields": {"seats": copy.deepcopy(records)}}
        gen, grouped = _render_grouped_editorial_boards([entry], supplementary_subpoints=True)
        assert len(grouped["Q4B"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["Q4B"])
        if children == 2:
            assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("N1", "MENTORING PROGRAMS"), ("S0", "BIBLIOGRAPHY")])
def test_a_seat_rerouted_out_of_q4c_to_an_excluded_code_gets_no_sub_point(code, heading):
    """N1 and S0 are excluded: a seat rendered as either is that code's
    record, so neither the seat row of the same journal nor any other
    carries its prose."""
    later = _board_member(idx=98, start="2017", end="present", prose="")
    gen, grouped = _render_grouped_editorial_boards([later, _rerouted(_board_member(), code, heading)],
                                                    supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [104]
    assert [e["element_idx_start"] for e in grouped["Q4C"]] == [98]
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("F2", "Board Certification"),
                                           ("I", "PROFESSIONAL ORGANIZATIONS AND SOCIETY MEMBERSHIPS"),
                                           ("C", "POSTDOCTORAL TRAINING")])
def test_a_refused_reroute_keeps_the_editorial_board_rule(code, heading):
    """A seat rerouted to F2, whose span the heading lookup does not find,
    or to I or C keeps its Q4C code: their renderers write no journal
    (`_cross_family_refusal`). It renders as a Q4C row and its prose is a
    deleted row under that row."""
    gen, grouped = _render_grouped_editorial_boards([_rerouted(_board_member(), code, heading)],
                                                    supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["Q4C"]] == [104]
    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == _MEMBER_ROW


@pytest.mark.parametrize("code", ["K5", "P"])
def test_no_reroute_leads_into_an_editorial_board_code(code):
    """Stage 6 moves no entry into Q4B or Q4C: `TAXONOMY_TO_SECTION` names no
    section for Q4A-Q4D, so a hierarchy-mismatch reroute never targets one,
    not even a record whose fields fit it (K5 renders its text, so nothing
    would be lost), and records no decision; the service section reroutes Q2
    only to Q3 and Q4D. The entry keeps its code and its rule, and no
    editorial board row is written for it."""
    from unified_pipeline.stage_6_word_template import TAXONOMY_TO_SECTION

    assert not {"Q4A", "Q4B", "Q4C", "Q4D"} & set(TAXONOMY_TO_SECTION)
    entry = _rerouted(_board_seat(code=code, role=None), "Q4C", "Editorial Board Membership")
    gen, grouped = _render_grouped_editorial_boards([entry], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [104] and "Q4C" not in grouped
    assert "taxonomy_code_original" not in entry
    assert not [w for w in gen._section_failures if w.get("check") == "hierarchy_mismatch_reroute"]
    assert _row_texts(gen.doc.element.body).count(_MEMBER_ROW) == 0
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_seat_rerouted_out_to_board_certification_gets_no_sub_point():
    """F2 is not offered (its heading lookup finds the title block). A Q4C
    entry whose only field stage 4 filled is a certifying board, under the
    Board Certification heading, renders as F2: it is that code's record,
    and no editorial board row carries its prose."""
    entry = _board_seat(code="Q4C", role=None, journal=None, start=None, end=None)
    entry["extracted_fields"] = {"certifying_board": "Example Board of Estuarine Medicine"}
    gen, grouped = _render_grouped_editorial_boards(
        [_rerouted(entry, "F2", "Board Certification")], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["F2"]] == [104] and "Q4C" not in grouped
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_a_seat_stage_3b_pinned_to_journal_reviewing_gets_no_sub_point():
    """Stage 3b's reviewer pin (#1413) recodes an unsure, role-less Q4C after
    a reviewer lead-in to Q4D, keeping the model's code in `pre_pin_code`:
    it renders as a Journal Reviewing row, which is excluded."""
    entry = dict(_board_member(), taxonomy_code="Q4D", pre_pin_code="Q4C")
    gen = _render_editorial_boards([entry], supplementary_subpoints=True)

    (own,) = gen._row_receipts.elements_for(entry)
    assert any(own.getparent() is table for table in gen._subpoint_excluded_tables)
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_an_editorial_board_seat_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _board_seat()
    withheld["_pii_withheld"] = True
    gen = _render_editorial_boards([_later_board_seat(), withheld], supplementary_subpoints=True)

    assert any(_row_texts_of(tr) == _SEAT_ROW for tr in gen.doc.element.body.iter(qn("w:tr")))
    assert _subpoint_rows(gen) == [] and _deleted_texts(gen) == ""


def test_an_editorial_board_seat_dedup_drops_offers_its_prose_once_under_the_kept_row():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no row, so its
    prose goes only under the row of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_board_seat(idx=104), _board_seat(idx=105)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["Q4B"]) == 2 and len(kept["Q4B"]) == 1
    gen._fill_service(kept)
    gen._add_supplementary_subpoints(before)

    (row,) = _subpoint_rows(gen)
    assert _row_texts_of(row.getprevious()) == _SEAT_ROW
    assert _deleted_texts(gen) == _EDITORIAL_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["Q4B"][0]["element_idx_start"]


def test_a_low_coverage_seat_never_reaches_the_content_overflow():
    """The content overflow collects an entry through the comment hook,
    which the Other Service writer does not call: a seat whose coverage and
    length qualify (`is_overflow_candidate`) is not queued, so its sub-point
    stands for nothing the overflow would write, and Accept All is the
    flag-off render."""
    from unified_pipeline.stage_6_word_template import is_overflow_candidate

    blocks = "\t".join(f"{_EDITORIAL_PROSE}, part {n} of the issue's plan" for n in range(9))

    def render(flag):
        entry = _board_seat(prose=blocks)
        entry["extraction_coverage"] = {"extraction_coverage_percent": 5.0}
        gen = _generator(supplementary_subpoints=flag)
        gen._fill_service({"Q4B": [entry]})
        assert is_overflow_candidate(entry) and gen._overflow_entries == []
        gen._place_unrendered_prose({"Q4B": [entry]})
        return gen

    on, off = render(True), render(False)
    (row,) = _subpoint_rows(on)
    assert len(list(row.iter(qn("w:p")))) == 9
    assert on.stats["overflow_bullets_added"] == off.stats["overflow_bullets_added"] == 0
    accepted = _accept_all(on)
    assert etree.tostring(accepted, method="c14n") == etree.tostring(off.doc.element.body, method="c14n")


def test_the_doctor_matches_an_editorial_board_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier seat's sub-point
    under its own index and Q4B, and the doctor's per-entry match counts it
    for that entry, not for the later seat that shares its words."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_board_seat(), _board_seat()],
                                     uid="T1205Q")
    assert placements == [[_EDITORIAL_PROSE]]
    sidecar = json.loads((tmp_path / "T1205Q_render_warnings.json").read_text())
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(104, "Q4B")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_EDITORIAL_PROSE, subpoints, element_idx=104)
    assert not offered_as_subpoint(_EDITORIAL_PROSE, subpoints, element_idx=98)


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_editorial_board_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line (#1205 review of O). The same seat without them still has its
    sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_board_seat()])
    assert placements == [[_EDITORIAL_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_board_seat(), _EDITORIAL_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


# --- #1205: S1/S2 publications, the next code family -------------------------
# The bibliography writer writes every S1 and S2 entry as one citation
# paragraph under its subsection label, and records that paragraph against
# the entry. So the prose is a deleted paragraph under the paper's own
# citation, the one the writer recorded for that entry (the row-receipt
# ledger), unless the citation is itself a tracked change.

_PAPER_AUTHORS = "Quill AM, Rook BT, Marblegate Q"
_PAPER_TITLE = "Sediment budgets of restored tidal marshes"
_PAPER_JOURNAL = "Journal of Estuarine Sediment Studies"
_PAPER_PROSE = "Designed the sampling study, analysed the sediment cores and wrote the manuscript"
_LATER_PAPER_PROSE = "Led the field campaign across the three restored marshes and drafted the figures"


def _publication(idx=1200, code="S1", title=_PAPER_TITLE, year="2014", prose=_PAPER_PROSE, **fields):
    """A paper whose source line goes on past its citation into what the
    owner did for it, on a line of its own (EBYSBC SJWASY's shape). A later
    paper of the same group in the same journal shares every distinctive
    word of its citation but the title's last ones."""
    citation = f"{_PAPER_AUTHORS}. {title}. {_PAPER_JOURNAL}. {year};12(3):45-52."
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "text": citation + (f"\t{prose}" if prose else ""),
        "extracted_fields": {"authors": _PAPER_AUTHORS, "title": title, "journal": _PAPER_JOURNAL,
                             "year": year, "volume": "12", "issue": "3", "pages": "45-52", **fields},
    }


def _later_publication(code="S1", prose=""):
    """The group's 2018 paper: it renders first, reverse chronologically."""
    return _publication(idx=1190, code=code, title=f"{_PAPER_TITLE} after a decade", year="2018",
                        prose=prose)


def _enriched(entry):
    """`entry` as stage 5 leaves an enriched citation: the writer deletes the
    CV's whole line and inserts PubMed's citation beside it."""
    return dict(entry, enrichment_status="enriched", enrichment_source="pubmed")


def _render_publications(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = {}
    for entry in entries:
        grouped.setdefault(entry["taxonomy_code"], []).append(entry)
    gen._fill_bibliography(grouped, cv_owner={"last_name": "Marblegate"})
    gen._add_supplementary_subpoints(grouped)
    return gen


def _citation_text(p):
    """A citation paragraph's text in Accept All: its plain and inserted runs."""
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _without_revision_dates(body):
    """`body` with every `w:date` dropped: an enriched citation's revisions
    are stamped with the clock, which differs between two renders."""
    body = copy.deepcopy(body)
    for element in body.iter():
        element.attrib.pop(qn("w:date"), None)
    return etree.tostring(body, method="c14n")


def _publication_table_row(code="S1"):
    """The same paper read from a source table: stage 2 joins the row's
    cells with pipes, and its position names the row ("41.3")."""
    entry = _publication(code=code)
    entry["element_idx_start"] = "41.3"
    entry["text"] = entry["text"].replace("\t", " | ")
    return entry


@pytest.mark.parametrize("code, make", [("S1", _publication), ("S2", _publication),
                                        ("S1", _publication_table_row)], ids=["S1", "S2", "S1_table_row"])
def test_publication_prose_is_a_deleted_paragraph_under_its_own_citation(code, make):
    entry = make(code=code)
    gen = _render_publications([entry], supplementary_subpoints=True)

    (para,) = _subpoint_paragraphs(gen)
    above = para.getprevious()
    assert _citation_text(above).startswith(f"1. {_PAPER_AUTHORS}. {_PAPER_TITLE}.")
    assert "".join(t.text for t in para.iter(qn("w:delText"))) == _PAPER_PROSE
    assert list(para.iter(qn("w:t"))) == []
    assert para.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None
    # No bullet and the citation's indent (the writer gives a citation
    # none); a paragraph follows, so the mark stays deleted and Accept All
    # leaves no empty line.
    assert para.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None
    assert para.find(f"{qn('w:pPr')}/{qn('w:ind')}") is None
    assert para.getnext().tag == qn("w:p")
    span = gen._subpoint_section_spans()[code]
    assert span.start < list(gen.doc.element.body).index(para)
    assert gen.stats["supplementary_subpoints"] == 1
    (placement,) = gen._subpoint_placements
    assert (placement.element_idx_start, placement.code) == (entry["element_idx_start"], code)


def test_a_papers_several_prose_lines_are_paragraphs_in_source_order():
    lines = (_PAPER_PROSE, "Presented the findings at the Example Coastal Ecology Congress plenary session")
    gen = _render_publications([_publication(prose="\t".join(lines))], supplementary_subpoints=True)

    paragraphs = _subpoint_paragraphs(gen)
    assert [_paragraph_text(p) for p in paragraphs] == list(lines)
    assert _citation_text(paragraphs[0].getprevious()).startswith("1. ")
    assert paragraphs[1].getprevious() is paragraphs[0]


@pytest.mark.parametrize("code", ["S1", "S2"])
def test_publication_prose_goes_under_its_own_citation_among_papers_of_one_group(code):
    """The later paper renders first and holds every distinctive word of the
    earlier one's citation, so words alone put the earlier one's prose under
    the later citation. Only the paragraph the writer recorded for the
    earlier entry takes it."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_publications([_later_publication(code), _publication(code=code)],
                               supplementary_subpoints=True)

    (para,) = _subpoint_paragraphs(gen)
    assert _citation_text(para.getprevious()).startswith(f"2. {_PAPER_AUTHORS}. {_PAPER_TITLE}.")
    assert _citation_text(para.getprevious().getprevious()).startswith("1. ")
    # Words alone would have taken the later paper's citation.
    off = _render_publications([_later_publication(code), _publication(code=code)],
                               supplementary_subpoints=False)
    lines = [line for line in supplementary.anchor_candidates(off.doc) if _PAPER_AUTHORS in line.text]
    assert supplementary.best_anchor(supplementary._value_tokens(_publication(code=code), code),
                                     lines) is lines[0]
    assert lines[0].text.startswith("1. ")


def test_papers_of_one_group_each_take_their_own_prose():
    gen = _render_publications([_later_publication(prose=_LATER_PAPER_PROSE), _publication()],
                               supplementary_subpoints=True)

    by_text = {_paragraph_text(p): p for p in _subpoint_paragraphs(gen)}
    assert set(by_text) == {_LATER_PAPER_PROSE, _PAPER_PROSE}
    assert _citation_text(by_text[_LATER_PAPER_PROSE].getprevious()).startswith("1. ")
    assert _citation_text(by_text[_PAPER_PROSE].getprevious()).startswith("2. ")
    assert [(p.element_idx_start, p.paragraphs) for p in gen._subpoint_placements] \
        == [(1190, (_LATER_PAPER_PROSE,)), (1200, (_PAPER_PROSE,))]


@pytest.mark.parametrize("entries", [
    lambda: [_later_publication(), _publication()],
    # An enriched paper beside them, and a review: neither changes.
    lambda: [_enriched(_later_publication(prose=_LATER_PAPER_PROSE)), _publication(),
             _publication(idx=1210, code="S2", title="Restoring tidal marshes: a review", year="2016")],
], ids=["S1", "S1_enriched_and_S2"])
def test_publication_sub_point_accept_all_is_the_flag_off_render_and_reject_restores_it(entries):
    off = _render_publications(entries(), supplementary_subpoints=False)
    on = _render_publications(entries(), supplementary_subpoints=True)

    rejected = "".join(t.text or "" for t in on.doc.element.body.iter(qn("w:t"), qn("w:delText")))
    assert rejected.count(_PAPER_PROSE) == len(on._subpoint_placements)
    assert _PAPER_PROSE not in "".join(_texts_of(off.doc.element.body))
    accepted = _accept_all(on)
    # Empty paragraphs and every table row, cell by cell, included.
    assert _texts_of(accepted) == _texts_of(off.doc.element.body)
    assert _row_texts(accepted) == _row_texts(off.doc.element.body)
    assert _without_revision_dates(accepted) == _without_revision_dates(off.doc.element.body)


def test_an_enriched_citation_gets_no_sub_point_and_reject_all_restores_its_prose_once():
    """An enriched citation deletes the CV's whole line beside PubMed's
    inserted one. Reject All already gives the line back, prose and all, so
    a sub-point would print the prose twice; it gets none, and the document
    is the flag-off one. The re-rank's shape: the citation is inside w:ins,
    no word of it visible to the word test, which took the paper before it
    (EBYSBC, 21 of 64 placements)."""
    from unified_pipeline.stage6 import supplementary

    enriched = _enriched(_publication())
    on = _render_publications([_later_publication(), enriched], supplementary_subpoints=True)
    off = _render_publications([_later_publication(), _enriched(_publication())],
                               supplementary_subpoints=False)

    (own,) = on._row_receipts.elements_for(enriched)
    assert supplementary._holds_tracked_change(own) and _citation_text(own).startswith("2. ")
    assert _subpoint_paragraphs(on) == [] and on._subpoint_placements == []
    assert _deleted_texts(on).count(_PAPER_PROSE) == 1
    assert _without_revision_dates(on.doc.element.body) == _without_revision_dates(off.doc.element.body)


def _paper_titled_by_its_prose():
    """The prose is the title stage 4 extracted, which the citation prints."""
    return _publication(title=_PAPER_PROSE)


def _paper_whose_journal_is_its_prose():
    """The prose is the journal stage 4 extracted."""
    return _publication(journal=_PAPER_PROSE)


@pytest.mark.parametrize("entry", [_paper_titled_by_its_prose, _paper_whose_journal_is_its_prose])
def test_publication_prose_already_in_its_citation_is_not_offered_again(entry):
    """The citation prints it (`_record_rendered`), so it does not print twice."""
    from unified_pipeline.stage6 import supplementary

    gen = _render_publications([entry()], supplementary_subpoints=True)
    assert _PAPER_PROSE in "".join(_texts_of(gen.doc.element.body))
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""
    haystack = supplementary.rendered_haystack(gen._rendered_output_lines(), gen.doc)
    assert supplementary.unrendered_prose(entry(), "S1", haystack) == []


def _declined_paper():
    """No title and no year: the writer declines it as no citation (#446)
    and writes no paragraph for it."""
    entry = _publication()
    entry["extracted_fields"] = {"authors": _PAPER_AUTHORS, "journal": _PAPER_JOURNAL}
    return entry


def _superseded_paper():
    """An "in press" copy stage 5 found published: written as one tracked
    deletion of its line, with no citation number and no receipt."""
    return dict(_publication(), in_press_superseded=True,
                in_press_note="Published since; listed again below.")


def _fused_papers():
    """Two papers stage 5d formatted as one block (#208): the writer splits
    it into two citations of its own copies, neither of them the entry."""
    entry = _publication()
    second = f"{_PAPER_AUTHORS}. {_PAPER_TITLE} after a decade. {_PAPER_JOURNAL}. 2018;16(1):1-9."
    entry["extracted_fields"].update(
        formatted_citation=f"{_PAPER_AUTHORS}. {_PAPER_TITLE}. {_PAPER_JOURNAL}. 2014;12(3):45-52.\n{second}",
        formatting_source="stage_5d_llm")
    return entry


@pytest.mark.parametrize("make", [_declined_paper, _superseded_paper, _fused_papers],
                         ids=["declined", "in_press_superseded", "split_by_the_writer"])
def test_a_paper_the_bibliography_writes_no_citation_of_its_own_for_gets_no_sub_point(make):
    entry = make()
    gen = _render_publications([_later_publication(), entry], supplementary_subpoints=True)

    assert gen._row_receipts.elements_for(entry) == []
    assert _subpoint_paragraphs(gen) == [] and gen._subpoint_placements == []


def test_the_publication_span_is_the_bibliography():
    """`_get_wcm_section_header` falls back on the 'S' prefix to the whole
    Bibliography, which holds every citation the writer writes; the sub-point
    pass needs no heading of its own for S1 or S2."""
    gen = _render_publications([_publication(), _publication(idx=1210, code="S2", title="Restoring tidal marshes",
                                                             prose=_LATER_PAPER_PROSE)],
                               supplementary_subpoints=True)
    body = list(gen.doc.element.body)
    for code in ("S1", "S2"):
        assert WCMTemplateGenerator._get_wcm_section_header(code) == "BIBLIOGRAPHY"
        span = gen._subpoint_section_spans()[code]
        assert "".join(t.text or "" for t in body[span.start].iter(qn("w:t"))).startswith("BIBLIOGRAPHY")
    assert len(_subpoint_paragraphs(gen)) == 2


def _render_grouped_publications(entries, **kwargs):
    gen = _generator(**kwargs)
    grouped = gen._group_entries_by_code(entries)
    gen._fill_bibliography(grouped, cv_owner={"last_name": "Marblegate"})
    gen._add_supplementary_subpoints(grouped)
    return gen, grouped


@pytest.mark.parametrize("code, heading", [("S0", "RESEARCHER PROFILES"), ("N1", "MENTORING PROGRAMS")])
def test_a_paper_rerouted_out_of_s1_to_an_excluded_code_gets_no_sub_point(code, heading):
    """S0 and N1 are excluded: a paper rendered as either is that code's
    record, so neither its citation of the same words nor any other carries
    its prose."""
    gen, grouped = _render_grouped_publications(
        [_later_publication(), _rerouted(_publication(), code, heading)], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[code]] == [1200]
    assert [e["element_idx_start"] for e in grouped["S1"]] == [1190]
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("code, heading", [("S8", "ABSTRACTS"), ("F2", "BOARD CERTIFICATION")])
def test_a_refused_reroute_keeps_the_publication_rule(code, heading):
    """A peer-reviewed article under an abstracts heading keeps its S1 code
    (S1 -> S8 is a confusion the taxonomy warns of,
    `_TAXONOMY_WARNED_CONFUSIONS`), and one under the board certification
    heading too, the code the span gap blocks (its renderer would drop the
    title, `_fields_lost_in_reroute`). It renders as an S1 citation and its
    prose is a deleted paragraph under it."""
    gen, grouped = _render_grouped_publications([_rerouted(_publication(), code, heading)],
                                                supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped["S1"]] == [1200]
    (para,) = _subpoint_paragraphs(gen)
    assert _citation_text(para.getprevious()).startswith(f"1. {_PAPER_AUTHORS}.")
    assert _paragraph_text(para) == _PAPER_PROSE


@pytest.mark.parametrize("assigned, target, heading", [
    ("S8", "S1", "Peer-Reviewed Research Articles"),
    ("S3", "S2", "Reviews and Editorials"),
])
def test_a_record_rerouted_into_a_publication_code_follows_the_publication_rule(assigned, target, heading):
    """An abstract under a peer-reviewed heading, or a book under the reviews
    heading, renders as that heading's citation: its prose is a deleted
    paragraph under it, recorded under the code it rendered as."""
    gen, grouped = _render_grouped_publications(
        [_rerouted(_publication(code=assigned), target, heading)], supplementary_subpoints=True)

    assert [e["element_idx_start"] for e in grouped[target]] == [1200]
    (para,) = _subpoint_paragraphs(gen)
    assert _paragraph_text(para) == _PAPER_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.code == target


def test_a_paper_whose_entry_the_pii_pass_withheld_gets_no_sub_point():
    """#892: the PII pass touched the entry. What is left can still describe
    what was withheld, so none of it is offered."""
    withheld = _publication()
    withheld["_pii_withheld"] = True
    gen = _render_publications([_later_publication(), withheld], supplementary_subpoints=True)

    assert gen._row_receipts.elements_for(withheld)
    assert _subpoint_paragraphs(gen) == [] and _deleted_texts(gen) == ""


@pytest.mark.parametrize("protect, value", [(_naming_the_spouse, _SPOUSE),
                                            (_with_a_home_phone, _HOME_PHONE)],
                         ids=["spouse", "home_phone"])
def test_publication_prose_holding_protected_data_is_not_offered_through_the_pii_pass(
        tmp_path, monkeypatch, protect, value):
    """The PII pass drops a field named for protected data and marks the
    entry withheld; no scope withholds a phone number in the middle of a
    line. The same paper without them still has its sub-point."""
    _doc, placements = _stage6_render(tmp_path / "plain", monkeypatch, [_publication()])
    assert placements == [[_PAPER_PROSE]]

    doc, placements = _stage6_render(tmp_path / "protected", monkeypatch,
                                     [protect(_publication(), _PAPER_PROSE)])
    assert placements == []
    assert value not in "".join(t.text or "" for t in doc.element.body.iter(qn("w:t"), qn("w:delText")))


def test_a_paper_dedup_drops_offers_its_prose_once_under_the_kept_citation():
    """Stage 6 dedups before the sections render and hands the sub-point
    pass the entries from before it: the dropped copy wrote no citation, so
    its prose goes only under the citation of the copy dedup kept, once."""
    gen = _generator(supplementary_subpoints=True)
    grouped = gen._group_entries_by_code([_publication(idx=1200), _publication(idx=1201)])
    kept, before, _decisions = gen._dedup_grouped_entries(grouped)
    assert len(before["S1"]) == 2 and len(kept["S1"]) == 1
    gen._fill_bibliography(kept, cv_owner={"last_name": "Marblegate"})
    gen._add_supplementary_subpoints(before)

    (para,) = _subpoint_paragraphs(gen)
    assert _citation_text(para.getprevious()).startswith(f"1. {_PAPER_AUTHORS}.")
    assert _deleted_texts(gen) == _PAPER_PROSE
    (placement,) = gen._subpoint_placements
    assert placement.element_idx_start == kept["S1"][0]["element_idx_start"]


def test_a_low_coverage_paper_never_reaches_the_content_overflow():
    """`is_overflow_candidate` turns every S code away, so a paper whose
    coverage and length would qualify another code is not queued: its
    sub-point stands for nothing the overflow would write (#2013's
    neither-document gap cannot arise), and Accept All is the flag-off
    render."""
    from unified_pipeline.stage_6_word_template import is_overflow_candidate

    blocks = "\t".join(f"{_PAPER_PROSE}, part {n} of the study's methods" for n in range(12))

    def render(flag):
        entry = _publication(prose=blocks)
        entry["extraction_coverage"] = {"extraction_coverage_percent": 5.0}
        gen = _generator(supplementary_subpoints=flag)
        gen._fill_bibliography({"S1": [entry]}, cv_owner={"last_name": "Marblegate"})
        assert not is_overflow_candidate(entry) and gen._overflow_entries == []
        assert is_overflow_candidate(dict(entry, taxonomy_code="D1"))
        gen._place_unrendered_prose({"S1": [entry]})
        return gen

    on, off = render(True), render(False)
    assert len(_subpoint_paragraphs(on)) == 12
    assert on.stats["overflow_bullets_added"] == off.stats["overflow_bullets_added"] == 0
    accepted = _accept_all(on)
    assert etree.tostring(accepted, method="c14n") == etree.tostring(off.doc.element.body, method="c14n")


def test_fan_out_children_of_a_papers_line_offer_nothing_their_own_segment_does_not_hold():
    """Stage 4 returned two papers from one line: fan-out gives the earlier
    one its own segment and the last one the whole line, which reads only
    its own segment when the segments pair with the records
    (`_own_record_line`). Prose on a segment of its own pairs with no
    record, so neither child offers it, as which paper it is about cannot
    be told. The same line kept whole, with no records, offers it: the
    later paper, by the same group in the same journal, restates the
    entry's fields (`_restates_row`) and holds its title, so it reads as the
    entry's own citation, not another paper's (`_cites_another_paper`)."""
    later = f"{_PAPER_AUTHORS}. {_PAPER_TITLE} after a decade. {_PAPER_JOURNAL}. 2018;16(1):1-9."
    own = f"{_PAPER_AUTHORS}. {_PAPER_TITLE}. {_PAPER_JOURNAL}. 2014;12(3):45-52."
    records = [{"authors": _PAPER_AUTHORS, "title": f"{_PAPER_TITLE} after a decade", "journal": _PAPER_JOURNAL,
                "year": "2018"},
               {"authors": _PAPER_AUTHORS, "title": _PAPER_TITLE, "journal": _PAPER_JOURNAL, "year": "2014"}]
    text = f"{later}\t{own}\t{_PAPER_PROSE}"
    for listed, children, offered in ((True, 2, []), (False, 1, [(_PAPER_PROSE,)])):
        fields = {**records[-1], STAGE4_RECORDS_KEY: copy.deepcopy(records)} if listed else dict(records[-1])
        entry = {"element_idx_start": 1300, "taxonomy_code": "S1", "text": text, "extracted_fields": fields}
        gen, grouped = _render_grouped_publications([entry], supplementary_subpoints=True)
        assert len(grouped["S1"]) == children
        assert all(gen._row_receipts.elements_for(e) for e in grouped["S1"])
        assert [p.paragraphs for p in gen._subpoint_placements] == offered


def test_the_doctor_matches_a_publication_sub_point_to_its_own_entry_only(tmp_path, monkeypatch):
    """Through `run_stage6`: the sidecar records the earlier paper's
    sub-point under its own index and S1, and the doctor's per-entry match
    counts it for that entry, not for the later paper that shares its words.
    The sub-point sits directly under its citation at the citation's own
    indent, none: Reject All shows it as an unnumbered paragraph level with
    the citations."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )
    from unified_pipeline.stage6 import supplementary
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    doc, placements = _stage6_render(tmp_path, monkeypatch, [_later_publication(), _publication()],
                                     uid="T1205S")
    assert placements == [[_PAPER_PROSE]]
    sidecar = json.loads((tmp_path / "T1205S_render_warnings.json").read_text())
    assert [(p["element_idx_start"], p["code"]) for p in sidecar[SUBPOINT_PLACEMENTS_KEY]] == [(1200, "S1")]
    subpoints = subpoint_text(docx_subpoint_lines(doc), subpoint_placements(sidecar))
    assert offered_as_subpoint(_PAPER_PROSE, subpoints, element_idx=1200)
    assert not offered_as_subpoint(_PAPER_PROSE, subpoints, element_idx=1190)
    (para,) = [p for p in doc.element.body.iterchildren(qn("w:p")) if supplementary._is_subpoint(p)]
    assert _citation_text(para.getprevious()).startswith("2. ")
    assert para.find(f"{qn('w:pPr')}/{qn('w:ind')}") is None
    assert para.getprevious().find(f"{qn('w:pPr')}/{qn('w:ind')}") is None


_FUSED_PAPER = "Fenn CD, Quill AM. Saltmarsh accretion under rising tides. Coastal Studies Review. 2011;9:101-12."


@pytest.mark.parametrize("second, placed", [(_FUSED_PAPER, []), (_PAPER_PROSE, [[_PAPER_PROSE]])],
                         ids=["another_paper", "own_prose"])
def test_the_doctor_still_reports_another_paper_fused_onto_a_publications_line(tmp_path, monkeypatch,
                                                                               second, placed):
    """Through `run_stage6` (#1205 review): the CV's next paper joined to
    the entry's line with no list number is no sub-point, so the doctor's
    `source_line_coverage` reports it missing from the entry with sub-points
    on as with them off. The paper's own prose in its place is the
    sub-point, which the doctor reads as there."""
    from unified_pipeline.doctor.lints.coverage import lint_source_line_coverage
    from unified_pipeline.doctor.shared import docx_body_blocks

    entry = _publication(prose=second)
    citation = entry["text"].split("\t")[0]

    def missing(flag):
        doc, placements = _stage6_render(tmp_path / str(flag), monkeypatch, [copy.deepcopy(entry)],
                                         supplementary_subpoints=flag)
        findings = lint_source_line_coverage([citation, second], docx_body_blocks(doc),
                                             docx_body_blocks(doc, deleted=True),
                                             stage3b={"entries": [copy.deepcopy(entry)]})
        return placements, [(f["message"].split(":")[0], f["evidence"]) for f in findings]

    assert missing(False) == ([], [("entry 1200", [second])])
    assert missing(True) == (placed, [] if placed else [("entry 1200", [second])])


# ------------------------------------------- date-led sibling records (#983)
# Stage 2 joins an entry's paragraphs with tabs, so three society records that
# each open with their own date reach stage 6 as one line. Stage 4 extracted
# the last one only (the web228/web26 I-entry shape, synthetic names).

_SOCIETY_RECORDS = [
    "2001-    Guild of Meandering Auditors (GMA)",
    "2004-2008  Executive Board Member, Treasurer.",
    "2007-2009 President, Organizing Committee, 44th Annual Meeting, GMA, Crab Hollow, ZQ",
]


def _society_entry(code="I", **fields):
    return {
        "element_idx_start": 60,
        "element_idx_end": 62,
        "element_type": "paragraph",
        "hierarchy": ["PROFESSIONAL MEMBERSHIPS"],
        "taxonomy_code": code,
        "taxonomy_confidence": 0.9,
        "text": "\t".join(_SOCIETY_RECORDS),
        "extracted_fields": fields or {
            "organization": "Guild of Meandering Auditors (GMA)",
            "start_date": "2007",
            "end_date": "2009",
            "membership_type": "President, Organizing Committee, 44th Annual Meeting",
        },
        "extraction_coverage": {"extraction_coverage_percent": 40.0},
    }


def _fieldless_society_entry(code="I"):
    # Stage 4 extracted nothing, so dev's renderer shows the text.
    entry = _society_entry(code=code)
    entry["extracted_fields"] = {}
    return entry


def _all_paragraph_texts(path):
    """Every paragraph's w:t text, table cells and tracked insertions included."""
    body = Document(path).element.body
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in body.iter(qn("w:p"))]


def test_date_led_sibling_records_render_one_line_each(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "unified_pipeline.stage_6_word_template.call_llm",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    payload = {"document_uid": "TESTUID983", "entries": [_fieldless_society_entry()],
               "cv_owner": {"full_name": "Quenby Marblegate"}}
    input_path = tmp_path / "TESTUID983_fields.json"
    input_path.write_text(json.dumps(payload))

    texts = _all_paragraph_texts(run_stage6(str(input_path), str(tmp_path / "out.docx"),
                                            verbose=False))

    # Each record is its own paragraph, once, where dev rendered the three
    # fused into one line; none is recovered a second time beside it.
    for record in _SOCIETY_RECORDS:
        assert texts.count(record) == 1, record
    assert not [t for t in texts if "Treasurer" in t and "President" in t]


@pytest.mark.parametrize("code", ["I", "O", "P"])
def test_date_led_siblings_render_one_line_each_when_one_records_fields_render(
        tmp_path, monkeypatch, code):
    # The web228/web26 (I) and web181 (O) shape: stage 4 extracted the last
    # record only. On dev the I renderer wrote that record's fields and lost
    # the 2001 membership; O and P wrote the records out of order or fused
    # into one paragraph. Now each source record is its own paragraph, once.
    monkeypatch.setattr(
        "unified_pipeline.stage_6_word_template.call_llm",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    payload = {"document_uid": "TESTUID983F", "entries": [_society_entry(code=code)],
               "cv_owner": {"full_name": "Quenby Marblegate"}}
    input_path = tmp_path / "TESTUID983F_fields.json"
    input_path.write_text(json.dumps(payload))

    texts = _all_paragraph_texts(run_stage6(str(input_path), str(tmp_path / "out.docx"),
                                            verbose=False))

    assert [t for t in texts if t in _SOCIETY_RECORDS] == _SOCIETY_RECORDS


@pytest.mark.parametrize("tail", [
    "Liaison, Quimby Farrow, (212) 555-0147",
    "Liaison, Quimby Farrow, qfarrow@example.org",
    "Member No. X1234567",
], ids=["phone", "email", "ID number"])
def test_an_entry_holding_contact_details_or_an_id_keeps_its_tabs(tail):
    # #983 review round 3: the O/P line parser read the last four digits of a
    # phone or an ID as the date column. Such an entry renders as on dev.
    entry = _society_entry()
    entry["text"] = "\t".join(_SOCIETY_RECORDS + [f"2010-2012 {tail}"])
    entry = _grouped_text(entry)
    assert "\n" not in entry["text"]
    assert DATE_LED_SOURCE_TEXT not in entry


def _grouped_text(entry):
    gen = _generator()
    grouped = gen._group_and_cut_entries([entry])
    (only,) = [e for entries in grouped.values() for e in entries]
    return only


def test_grouping_puts_date_led_siblings_on_lines_and_keeps_stage2_text():
    entry = _grouped_text(_fieldless_society_entry())
    assert entry["text"] == "\n".join(_SOCIETY_RECORDS)
    assert entry[DATE_LED_SOURCE_TEXT] == "\t".join(_SOCIETY_RECORDS)


@pytest.mark.parametrize("fields, splits", [
    ({}, True),
    (None, True),
    ({"organization": "", "start_date": None, "offices": []}, True),  # nothing extracted
    # One record's fields: the renderer writes the lines instead, so the
    # sibling records those fields leave out reach the page (#983, after #1859).
    ({"organization": "Guild of Meandering Auditors"}, True),
    ({"end_date": "2009"}, True),
    # Records stage 4 separated: their renderer owns them, not the lines.
    ({"offices": [{"office": "Treasurer"}, {"office": "President"}]}, False),
])
def test_an_entry_is_split_unless_stage4_separated_its_records(fields, splits):
    entry = _society_entry()
    entry["extracted_fields"] = fields
    entry = _grouped_text(entry)
    assert ("\n" in entry["text"]) is splits
    assert (DATE_LED_SOURCE_TEXT in entry) is splits


@pytest.mark.parametrize("code", ["M2A", "M2B", "M2C", "M2D"])
def test_a_grant_keeps_its_tabs(code):
    entry = _grouped_text(_fieldless_society_entry(code=code))
    assert entry["text"] == "\t".join(_SOCIETY_RECORDS)
    assert DATE_LED_SOURCE_TEXT not in entry


@pytest.mark.parametrize("code", ["K1", "K2", "K4"])
def test_a_teaching_entry_keeps_its_tabs(code):
    # The teaching renderer splits the text on its tabs itself.
    entry = _grouped_text(_fieldless_society_entry(code=code))
    assert entry["text"] == "\t".join(_SOCIETY_RECORDS)


def test_records_stage4_separated_keep_their_tabs():
    # A list of records under a key the I schema does not define: fan-out
    # declines it (no schema key), and its renderer, not the text, owns them.
    entry = _grouped_text(_society_entry(
        organization="Guild of Meandering Auditors (GMA)",
        offices=[{"office": "Treasurer", "start_date": "2004"},
                 {"office": "President", "start_date": "2007"}]))
    assert entry["text"] == "\t".join(_SOCIETY_RECORDS)


@pytest.mark.parametrize("marker", [FANNED_OUT_FROM, LAST_STAGE4_RECORD])
def test_a_record_fan_out_made_keeps_its_tabs(marker):
    # PR #1850 review: no test held the fan-out exclusion. A record fan-out
    # made is one of stage 4's records; dedup reads its marker, not its lines.
    entry = _fieldless_society_entry()
    entry[marker] = {"key": "records", "index": 0, "count": 2}
    entry = _grouped_text(entry)
    assert entry["text"] == "\t".join(_SOCIETY_RECORDS)
    assert DATE_LED_SOURCE_TEXT not in entry


# The web181 Q1 shape: a long lead-in, then dated offices, most of them
# single years. #1243's tab-part rules read it as one record (fewer than half
# its parts give a date range, and the lead-in's colon makes the single years
# details); its date-led lines would read as two.
_FORUM_PARTS = [
    "Example Rehabilitation Forum. A not-for-profit consortium of academic departments:",
    "2005-present Founding President.",
    "2009 Conference Director: First World Conference, Example City.",
    "2011 Conference Director: World Congress, Example City",
    "2012 Conference co-director, World Congress, Example City",
    "2016-present Example Fellowship Director",
]


def test_recovery_reads_stage2_text_not_the_date_led_lines():
    entry = _grouped_text({
        "taxonomy_code": "Q1",
        "text": "\t".join(_FORUM_PARTS),
        "extracted_fields": {},
    })
    assert "\n" in entry["text"]  # the lines render ...
    # ... but the recovery pass reads records as it did before the split.
    assert recovery_candidates("Q1", entry) == []
    lines_only = {**entry}
    del lines_only[DATE_LED_SOURCE_TEXT]
    assert recovery_candidates("Q1", lines_only) != []


def test_a_split_entry_whose_fields_render_recovers_nothing():
    # The mentee shape of `test_tab_part_the_extracted_fields_hold_is_not_recovered`,
    # with both parts date-led. Grouping splits it, but recovery reads stage
    # 2's text, so it adds nothing beside the mentee table, as on dev.
    gen = _generator()
    entry = {
        "taxonomy_code": "N3B",
        "text": ("2007-2011 Quimby Farrow, MSPH, PhD (Epidemiology)\t"
                 "2009-2011 Marsupial Heart Association Fellowship"),
        "extracted_fields": {"mentee_name": "Quimby Farrow", "mentee_level": "MSPH, PhD (Epidemiology)",
                             "start_date": "2007", "end_date": "2011",
                             "funding_source": "Marsupial Heart Association Fellowship (2009-2011)"},
    }
    grouped = gen._group_and_cut_entries([entry])
    assert "\n" in grouped["N3B"][0]["text"]
    assert DATE_LED_SOURCE_TEXT in grouped["N3B"][0]
    before = len(gen.doc.paragraphs)

    gen._recover_unrendered_records(grouped)

    assert len(gen.doc.paragraphs) == before


@pytest.mark.parametrize("code", ["Q1", "D1"])
@pytest.mark.parametrize("pii", ["SSN: 123-45-6789", "Date of Birth: 01/02/1960"])
def test_a_withheld_fragment_in_a_date_led_record_is_not_recovered(tmp_path, monkeypatch,
                                                                   code, pii):
    # PR #1850 review: the PII pass cut `text`, but the recovery pass read the
    # pre-split stage-2 text and put the withheld line back verbatim.
    monkeypatch.setattr(
        "unified_pipeline.stage_6_word_template.call_llm",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("no LLM in tests")))
    entry = _fieldless_society_entry(code=code)
    entry["text"] = "\t".join([_SOCIETY_RECORDS[0], f"{_SOCIETY_RECORDS[1]} {pii}",
                               _SOCIETY_RECORDS[2]])
    payload = {"document_uid": "TESTUID983P", "entries": [entry],
               "cv_owner": {"full_name": "Quenby Marblegate"}}
    input_path = tmp_path / "TESTUID983P_fields.json"
    input_path.write_text(json.dumps(payload))

    texts = _all_paragraph_texts(run_stage6(str(input_path), str(tmp_path / "out.docx"),
                                            verbose=False))

    value = pii.split(": ", 1)[1]
    assert not [t for t in texts if value in t]


# ------------------- #983 review round 2: the PII pass sees stage 2's tabs
# The pass's bare-label extension and its continuation-cell absorption run on
# over a tab and stop at a line break. A split that ran first turned the tab
# after "Children:" or "Date of Birth:" into a line break, and the value in
# the next cell rendered. Each case renders twice: as the branch does, and with
# the split switched off, which is dev's grouping and PII pass unchanged
# (stage6/pii_pass.py is dev's file). The pass must get the same input, cut the
# same spans, and leave every protected value in the docx the same number of
# times. Synthetic names and values throughout.
#
# Round 3 adds the render side. For every case, each of stage 2's tab parts
# that reaches the branch's docx reached dev's too, and no part reaches it cut
# across two paragraphs unless dev's cut it too (a trailing ID or phone number
# read as the date column). Every shape here is protected data, contact
# details or an ID, so an entry whose fields render is not split and renders
# exactly as on dev.

_PII_DATED_LEAD = ["2001-2005 Member, Guild of Meandering Auditors",
                   "2006-2008 Treasurer, Guild of Meandering Auditors"]
_PII_UNDATED_LEAD = ["Member, Guild of Meandering Auditors",
                     "Treasurer, Guild of Meandering Auditors"]
# name -> (the parts after the lead, the protected values, whether those parts
# alone hold two date-led records and so split behind an undated lead too)
_PII_SHAPES = {
    "bare label, value in the next cell": (
        ["Date of Birth:", "01/02/1960, Quillhaven, Ohio"], ["01/02/1960"], False),
    "bare birth label, place in the next cell": (
        ["Born:", "1960 Quillhaven, Ohio"], ["1960 Quillhaven"], False),
    "children, year-led continuation cells": (
        ["Children:", "2001 Ann Marblegate", "2004 Bob Marblegate"],
        ["Ann Marblegate", "Bob Marblegate"], True),
    "children, continuation cells": (
        ["Children: Ann Marblegate (b. 2001)", "Bob Marblegate (b. 2004)"],
        ["Ann Marblegate", "Bob Marblegate"], False),
    "home address, continuation cell": (
        ["Home Address: 12 Elm Street", "2019 Apartment 4, Quillhaven"],
        ["12 Elm Street", "Apartment 4"], False),
    "date of birth, birthplace cell": (
        ["Date of Birth: 01/02/1960", "1960 Quillhaven General Hospital"],
        ["01/02/1960", "Quillhaven General Hospital"], False),
    "in-line SSN": (["President, Organizing Committee SSN: 123-45-6789"], ["123-45-6789"], False),
    "in-line DOB": (["Auditor Date of Birth: 01/02/1960"], ["01/02/1960"], False),
    # A value shape, not a label: withheld only in an Appendix-bound entry,
    # and the full-scope `_pii_fragments` does not see it, so only the order
    # keeps the split's copy of stage 2's text from holding it uncut.
    "third-party email": (["2009-2011 Liaison, Quimby Farrow, qfarrow@example.org"],
                          ["qfarrow@example.org"], False),
    "third-party phone": (["2009-2011 Liaison, Quimby Farrow, (212) 555-0147"],
                          ["555-0147"], False),
    # A protected label after a leading date: withheld since #1859, so the
    # entry is not split, whether or not its fields render.
    "spouse after a date": (["2005 Spouse: Jane Quillfeather"], ["Jane Quillfeather"], False),
    "salary after a date": (["2015 Salary: $250,000"], ["$250,000"], False),
    "passport number after a date": (["2015 Passport No. X1234567"], ["X1234567"], False),
}
_PII_CODES = ["T", "A", "I", "D1", "Q1", "O", "P"]
# Codes whose entries get the full Personal Data / Appendix cut; a routed code
# withholds only the shapes in `_PII_ROUTED_SHAPES`.
_PII_FULL_SCOPE_CODES = {"T", "A"}
_PII_ROUTED_SHAPES = {"bare label, value in the next cell", "date of birth, birthplace cell",
                      "in-line SSN", "in-line DOB", "spouse after a date",
                      "salary after a date", "passport number after a date"}
# Withheld in an Appendix-bound (T) entry only: an A entry's contacts are the
# owner's, and a routed code keeps a contact.
_PII_THIRD_PARTY_SHAPES = {"third-party email", "third-party phone"}


def _pii_case_withholds(code, shape):
    if shape in _PII_THIRD_PARTY_SHAPES:
        return code == "T"
    return code in _PII_FULL_SCOPE_CODES or shape in _PII_ROUTED_SHAPES


def _docx_part_texts(path):
    """Each paragraph's w:t + w:delText text in every XML part of the docx
    (body, comments, headers): a protected value must be nowhere in it."""
    import zipfile

    from lxml import etree

    texts = []
    with zipfile.ZipFile(path) as docx:
        for name in docx.namelist():
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(docx.read(name))
            for para in root.iter(qn("w:p")):
                texts.append("".join(node.text or "" for node in para.iter(qn("w:t"), qn("w:delText"))))
    return texts


def _render_watching_pii_pass(tmp_path, monkeypatch, entries, *, dev_split_order):
    """Render `entries`; return (what the PII pass got and gave, paragraph texts).
    `dev_split_order` switches the date-led split off: dev's path."""
    import copy

    import unified_pipeline.stage_6_word_template as stage6

    seen = {}
    real_pass = stage6.run_pii_pass

    def watched_pass(entries_by_code, **kwargs):
        seen["input"] = {code: [e.get("text") for e in group]
                         for code, group in entries_by_code.items()}
        result = real_pass(entries_by_code, **kwargs)
        seen["output"] = {code: [(e.get("text"), copy.deepcopy(e.get("_pii_fragments")),
                                  e.get("_pii_orphaned_value")) for e in group]
                          for code, group in entries_by_code.items()}
        seen["withheld"] = list(result.withheld)
        return result

    arm = "dev" if dev_split_order else "branch"
    with monkeypatch.context() as patch:
        patch.setattr(stage6, "call_llm", lambda *args, **kwargs: (_ for _ in ()).throw(
            RuntimeError("no LLM in tests")))
        patch.setattr(stage6, "run_pii_pass", watched_pass)
        if dev_split_order:
            patch.setattr(stage6, "_split_date_led_groups", lambda entries_by_code: None)
        input_path = tmp_path / f"TESTUID983{arm}_fields.json"
        input_path.write_text(json.dumps({"document_uid": f"TESTUID983{arm}", "entries": entries,
                                          "cv_owner": {"full_name": "Quenby Marblegate"}}))
        out = run_stage6(str(input_path), str(tmp_path / f"{arm}.docx"), verbose=False)
    return seen, _docx_part_texts(out)


def _pii_case_entry(code, lead, shape, *, fields=False):
    entry = _society_entry(code=code) if fields else _fieldless_society_entry(code=code)
    entry["text"] = "\t".join(lead + _PII_SHAPES[shape][0])
    return entry


def _squashed(text):
    return "".join(text.split())


def _parts_on_page(texts, parts):
    """The `parts` whose text is in the docx: in one paragraph, or run on
    across paragraphs (a number cut into two cells still reached the page)."""
    page = _squashed("".join(texts))
    return {part for part in parts if _squashed(part) in page}


def _parts_cut_across_paragraphs(texts, parts):
    """The `parts` on the page that no single paragraph holds whole."""
    return {part for part in _parts_on_page(texts, parts)
            if not any(_squashed(part) in _squashed(text) for text in texts)}


# (whether stage 4 extracted fields, the lead). An entry whose fields render
# is the round-3 case; its undated lead adds nothing the dated one does not.
_PII_VARIANTS = [(False, "dated"), (False, "undated"), (True, "dated")]


@pytest.mark.parametrize("code", _PII_CODES)
@pytest.mark.parametrize("fields,lead", _PII_VARIANTS)
@pytest.mark.parametrize("shape", list(_PII_SHAPES))
def test_the_pii_pass_cuts_what_dev_cuts_with_or_without_the_split(tmp_path, monkeypatch,
                                                                   code, fields, lead, shape):
    _parts, protected, parts_split_alone = _PII_SHAPES[shape]
    entry = _pii_case_entry(code, _PII_DATED_LEAD if lead == "dated" else _PII_UNDATED_LEAD,
                            shape, fields=fields)
    # The case is what it says: stage 2's text splits behind a dated lead, and
    # behind an undated one only when the PII parts alone hold two records.
    assert (split_date_led_records(entry["text"]) != entry["text"]) == (
        lead == "dated" or parts_split_alone)

    dev_seen, dev_texts = _render_watching_pii_pass(
        tmp_path, monkeypatch, [copy.deepcopy(entry)], dev_split_order=True)
    seen, texts = _render_watching_pii_pass(
        tmp_path, monkeypatch, [copy.deepcopy(entry)], dev_split_order=False)

    assert seen["input"] == dev_seen["input"]  # stage 2's tabs, as on dev
    assert seen["output"] == dev_seen["output"]
    assert seen["withheld"] == dev_seen["withheld"]
    withholds = _pii_case_withholds(code, shape)
    assert bool(seen["withheld"]) == withholds
    # Round 3: the split puts no source line on the page that dev left off it,
    # and cuts none across two paragraphs.
    source_parts = entry["text"].split("\t")
    assert _parts_on_page(texts, source_parts) <= _parts_on_page(dev_texts, source_parts)
    assert (_parts_cut_across_paragraphs(texts, source_parts)
            <= _parts_cut_across_paragraphs(dev_texts, source_parts))
    for value in protected:
        assert (sum(t.count(value) for t in texts)
                == sum(t.count(value) for t in dev_texts)), value
        if withholds and code in _PII_FULL_SCOPE_CODES:
            assert not [t for t in texts if value in t], value
    if fields:
        assert texts == dev_texts  # its fields render: not split, so dev's docx


def test_an_entry_the_pii_pass_cut_is_not_split():
    # Its text and the recovery pass's input stay dev's cut text, tabs and all.
    # Field-less, so only the pass's mark keeps its dated lead unsplit.
    entry = _grouped_text(_pii_case_entry("T", _PII_DATED_LEAD, "children, year-led continuation cells"))
    assert entry["_pii_withheld"] is True
    assert entry["text"] == "\t".join(_PII_DATED_LEAD)
    assert DATE_LED_SOURCE_TEXT not in entry


@pytest.mark.parametrize("fields", [{"date_of_birth": "01/02/1960"}, {"notes": "SSN 123-45-6789"}])
def test_an_entry_the_pii_pass_touched_through_its_fields_is_not_split(fields):
    # No fragment in the text: the pass's own mark, set before the split runs,
    # and the fields it touched each keep the entry as dev renders it.
    entry = _society_entry(code="Q1", organization="Guild of Meandering Auditors", **fields)
    entry = _grouped_text(entry)
    assert entry["_pii_withheld"] is True
    assert entry["text"] == "\t".join(_SOCIETY_RECORDS)


def test_a_routed_entry_holding_protected_data_is_not_split():
    # An I entry gets the routed scope, which does not withhold a children
    # list, so the pass leaves it alone. Field-less, so only the full-scope
    # fragment test keeps its lines as dev renders them.
    entry = _grouped_text(_pii_case_entry("I", _PII_DATED_LEAD, "children, year-led continuation cells"))
    assert "_pii_withheld" not in entry
    assert "\n" not in entry["text"]
    assert DATE_LED_SOURCE_TEXT not in entry


# Judgement call 6: the web205/web210 shape. Date-led mentee lines whose
# mentee table already renders them, now with a withheld fragment in one line.
# Recovery must not write the mentee lines back beside the table.
_MENTEE_PII_ENTRY = {
    "element_idx_start": 80,
    "element_idx_end": 81,
    "element_type": "paragraph",
    "hierarchy": ["MENTORING"],
    "taxonomy_code": "N3B",
    "taxonomy_confidence": 0.9,
    "text": ("2007-2011 Quimby Farrow, MSPH, PhD (Epidemiology), SSN: 123-45-6789\t"
             "2009-2011 Marsupial Heart Association Fellowship"),
    "extracted_fields": {"mentee_name": "Quimby Farrow", "mentee_level": "MSPH, PhD (Epidemiology)",
                         "start_date": "2007", "end_date": "2011",
                         "funding_source": "Marsupial Heart Association Fellowship (2009-2011)"},
}


def test_a_pii_touched_date_led_mentee_entry_recovers_no_duplicate_lines(tmp_path, monkeypatch):
    without_pii = _MENTEE_PII_ENTRY["text"].replace(", SSN: 123-45-6789", "")
    assert split_date_led_records(without_pii) != without_pii  # date-led siblings
    dev_seen, dev_texts = _render_watching_pii_pass(
        tmp_path, monkeypatch, [copy.deepcopy(_MENTEE_PII_ENTRY)], dev_split_order=True)
    seen, texts = _render_watching_pii_pass(
        tmp_path, monkeypatch, [copy.deepcopy(_MENTEE_PII_ENTRY)], dev_split_order=False)

    assert seen["withheld"] == dev_seen["withheld"] and seen["withheld"]
    assert not [t for t in texts if "123-45-6789" in t]
    for token in ("Quimby Farrow", "Marsupial Heart Association"):
        # Once, in the mentee table, on both arms: never recovered again.
        assert sum(token in t for t in texts) == sum(token in t for t in dev_texts) == 1, token
