"""Stage-6 render-report units (#227/#228): deduplicate_entries records its
drop decisions, and _validate_output returns structured dicts whose messages
match the VALIDATION WARNINGS banner lines byte-for-byte — both feed the
<uid>_render_warnings.json sidecar the run doctor re-emits.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_render_report.py -p no:cacheprovider

Self-contained: no DB, no LLM, no template — documents are built in-test.
"""

import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.stage_6_word_template import (  # noqa: E402
    WCMTemplateGenerator,
    deduplicate_entries,
)

# ------------------------------------------------------- dedup decision trail

# The 2Q1_ZQ B1 drop shape: distinct degrees whose signature words all overlap
# because the degree tokens are single characters after tokenization (#227).
_PHD = {"text": "PhD: Instructional Systems Technology, Indiana University"}
_MS = {"text": "M.S: Instructional Systems Technology, Indiana University"}


def test_dedup_records_drop_decisions():
    decisions = []
    kept = deduplicate_entries([_PHD, _MS], decisions=decisions)
    assert len(kept) == 1
    assert len(decisions) == 1
    d = decisions[0]
    assert d["dropped_text"] == _MS["text"]
    assert d["kept_text"] == _PHD["text"]
    assert d["metric"].startswith("jaccard=")
    assert 0 <= d["jaccard"] <= 1
    assert set(d) >= {"metric", "jaccard", "containment", "title_containment",
                      "dropped_text", "kept_text"}


def test_dedup_decisions_default_is_backward_compatible():
    # decisions omitted: same return value, no crash (the pre-#227 call shape).
    assert len(deduplicate_entries([_PHD, _MS])) == 1
    assert deduplicate_entries([_PHD]) == [_PHD]


def test_dedup_no_decisions_recorded_when_nothing_dropped():
    decisions = []
    distinct = [{"text": "Editorial Board Member, BMC Medical Education"},
                {"text": "Invited talk on assessment design at AMEE 2024 conference"}]
    assert deduplicate_entries(distinct, decisions=decisions) == distinct
    assert decisions == []


# ------------------------------------------------------- _validate_output dicts

def _generator_for(doc: Document) -> WCMTemplateGenerator:
    """_validate_output only reads self.doc; skip the heavyweight ctor."""
    gen = WCMTemplateGenerator.__new__(WCMTemplateGenerator)
    gen.doc = doc
    return gen


def test_validate_output_returns_structured_semicolon_warning():
    doc = Document()
    doc.add_paragraph("Administrative teaching")
    fused = ("• Course Director (Modules: Introduction; Ethics; Practice; "
             "Assessment; Feedback)")
    doc.add_paragraph(fused)
    warnings = _generator_for(doc)._validate_output()
    assert len(warnings) == 1
    w = warnings[0]
    assert w["check"] == "semicolon_fused_bullets"
    assert w["code"] == "K3"
    assert w["section"] == "Administrative teaching"
    # The banner prints w["message"]; keep it byte-identical to the pre-#228
    # string so pod-log greps keep working.
    assert w["message"] == ("K3 (Administrative teaching): Content appears "
                            "combined with semicolons instead of separate bullets")
    assert w["evidence"] == [fused[:200]]


def test_validate_output_clean_doc_returns_empty_list():
    doc = Document()
    doc.add_paragraph("Administrative teaching")
    doc.add_paragraph("• Course Director, AI in Medicine")
    assert _generator_for(doc)._validate_output() == []


def _empty_teaching_doc() -> Document:
    doc = Document()
    doc.add_paragraph("EDUCATIONAL CONTRIBUTIONS")
    doc.add_paragraph("CLINICAL PRACTICE")
    return doc


def test_empty_teaching_section_warns_only_when_a_teaching_code_was_given():
    """#1221: a CV with no K entry has nothing to show under EDUCATIONAL
    CONTRIBUTIONS, so the empty-section check is not a finding there."""
    def checks(entries_by_code):
        warnings = _generator_for(_empty_teaching_doc())._validate_output(entries_by_code)
        return [w["check"] for w in warnings]

    assert checks({"K2": [{}]}) == ["no_visible_teaching_content"]
    assert checks({"K5": [{}], "H": [{}]}) == ["no_visible_teaching_content"]
    assert checks({"H": [{}], "N3B": [{}]}) == []
    assert checks({}) == []
    assert checks(None) == ["no_visible_teaching_content"]  # no entries to hand: as before


# ------------------------------------------------- warnings go to the logger

def test_comment_failure_logs_via_project_logger_not_print(caplog, capsys):
    """The monolith's except-branch advisories reach the run log through the
    project logger (dictConfig's _DynamicStdout), not a bare print(). One
    representative branch stands in for all twelve: a None paragraph makes
    _add_word_comment's `para._p` raise, which is the fallback path."""
    import logging

    gen = WCMTemplateGenerator.__new__(WCMTemplateGenerator)
    gen.emit_comments = True
    gen.verbose = True
    gen._comment_id = 0
    gen._comments = []
    gen.stats = {"comments_added": 0}

    with caplog.at_level(logging.WARNING, logger="unified_pipeline.stage_6_word_template"):
        gen._add_word_comment(None, "note")

    assert any("Could not add comment" in r.message for r in caplog.records)
    assert capsys.readouterr().out == ""
