"""Tests for doctor/comment_fate.py (#1654): a verdict per review-copy comment,
read from the reviewer's corrected copy.

Every document is built here from synthetic text: a review copy with one
comment per record, then the corrected copy a reviewer could upload. The four
fixtures the issue names: a comment kept, a comment resolved with its text
edited, a comment deleted with its text untouched, and the clean document
uploaded instead of the review copy.

    python3 -m pytest src/unified_pipeline/tests/test_doctor_comment_fate.py -p no:cacheprovider
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.doctor import comment_fate as cf  # noqa: E402
from unified_pipeline.stage6.formatting import add_cviche_box  # noqa: E402

AUTHOR = "CViche check"
CITE_A = "Quorvane T, Plesk M. Heliotropic drift in vexillary marmosets. J Synth Imag. 2019;12:34-56."
CITE_B = "Quorvane T, Abernoth R. Saltwick lattices under brennic loading. Synth Rep. 1912;3:7-19."
CITE_C = "Quorvane T. Tremulant gradients across fennish moorland. Moor Synth. 2016;8:101-110."
CITE_D = "Quorvane T, Lisk V. Brennic coupling in moorland lattices. Synth Rep. 2018;5:1-9."
FLAG = "Check this citation."
YEAR_FLAG = "Year looks wrong."


def _document(*, comments: bool, box: bool = False) -> Document:
    """BIBLIOGRAPHY and four citations; with ``comments``, one review-copy
    comment on each (ids 0..3), B's on its year alone."""
    doc = Document()
    doc.add_paragraph("BIBLIOGRAPHY")
    paragraphs = [doc.add_paragraph(text) for text in (CITE_A, CITE_B, CITE_C, CITE_D)]
    if comments:
        for p in paragraphs:
            if p.text == CITE_B:
                before, year, after = p.text.partition("1912")
                p.runs[0].text = before
                p.add_run(year)
                p.add_run(after)
                doc.add_comment(p.runs[1], text=YEAR_FLAG, author=AUTHOR)
            else:
                doc.add_comment(p.runs, text=FLAG, author=AUTHOR)
    if box:
        add_cviche_box(doc, "CViche review notes: delete this box before sending")
    return doc


def _save(doc: Document, path: Path) -> Path:
    doc.save(str(path))
    return path


def _paragraph(doc: Document, starts: str):
    return next(p for p in doc.paragraphs if p.text.startswith(starts))


def _delete_comment(doc: Document, comment_id: int) -> None:
    """What Word does on "Delete comment": the range, the reference and the comment go."""
    for el in list(doc.element.body.iter(qn("w:commentRangeStart"), qn("w:commentRangeEnd"),
                                         qn("w:commentReference"))):
        if el.get(qn("w:id")) == str(comment_id):
            gone = el.getparent() if el.tag == qn("w:commentReference") else el  # the reference's run
            gone.getparent().remove(gone)
    comment = next(c for c in doc.comments if c.comment_id == comment_id)
    comment._comment_elm.getparent().remove(comment._comment_elm)


def _fates(tmp_path: Path, corrected: Document) -> cf.FateReport:
    review = _save(_document(comments=True, box=True), tmp_path / "review.docx")
    return cf.comment_fates(review, _save(corrected, tmp_path / "corrected.docx"), AUTHOR)


def _verdicts(report: cf.FateReport) -> dict[int, str]:
    return {f.comment_id: f.verdict for f in report.fates}


def test_a_kept_comment_on_untouched_text_is_unknown(tmp_path):
    report = _fates(tmp_path, _document(comments=True, box=True))
    assert report.comments_tracked
    assert _verdicts(report) == dict.fromkeys(range(4), cf.REVIEW_UNKNOWN)


def test_a_comment_resolved_with_its_text_edited_is_fixed(tmp_path):
    """The reviewer corrects B's year and deletes the comment; the box goes too."""
    doc = _document(comments=True)
    _paragraph(doc, "Quorvane T, Abernoth").runs[1].text = "2012"
    _delete_comment(doc, 1)
    assert _verdicts(_fates(tmp_path, doc)) == {0: cf.REVIEW_UNKNOWN, 1: cf.REVIEW_FIXED,
                                                2: cf.REVIEW_UNKNOWN, 3: cf.REVIEW_UNKNOWN}


def test_an_edit_outside_a_comments_span_leaves_its_text_unchanged(tmp_path):
    """B's comment is on its year; the reviewer fixes the page range and keeps the year."""
    doc = _document(comments=True)
    _paragraph(doc, "Quorvane T, Abernoth").runs[2].text = ";3:7-21."
    assert _verdicts(_fates(tmp_path, doc))[1] == cf.REVIEW_UNKNOWN


def test_a_comment_deleted_with_its_text_untouched_is_not_a_problem(tmp_path):
    doc = _document(comments=True)
    _delete_comment(doc, 2)
    assert _verdicts(_fates(tmp_path, doc)) == {0: cf.REVIEW_UNKNOWN, 1: cf.REVIEW_UNKNOWN,
                                                2: cf.REVIEW_NOT_A_PROBLEM, 3: cf.REVIEW_UNKNOWN}


def test_a_record_deleted_with_its_comment_is_fixed(tmp_path):
    doc = _document(comments=True)
    p = _paragraph(doc, "Quorvane T, Lisk")._p
    p.getparent().remove(p)
    _delete_comment(doc, 3)
    assert _verdicts(_fates(tmp_path, doc))[3] == cf.REVIEW_FIXED


def test_a_tracked_deletion_reads_as_the_reviewer_deleting_the_text(tmp_path):
    """Accepted view: a pending w:del is gone, the comment kept beside it."""
    doc = _document(comments=True)
    run = _paragraph(doc, "Quorvane T. Tremulant").runs[0]._r
    wrapper = run.makeelement(qn("w:del"), {qn("w:id"): "901", qn("w:author"): "Synthetic Reviewer"})
    run.addprevious(wrapper)
    wrapper.append(run)
    for t in run.iter(qn("w:t")):
        t.tag = qn("w:delText")
    assert _verdicts(_fates(tmp_path, doc))[2] == cf.REVIEW_FIXED


def test_a_pending_deletion_in_the_review_copy_is_not_part_of_the_anchored_text(tmp_path):
    """Stage 6 writes tracked deletions; a commented paragraph holding one
    (here a w:t inside w:del) is unchanged when the reviewer leaves it alone."""
    def with_deletion():
        doc = _document(comments=True)
        run = _paragraph(doc, "Quorvane T. Tremulant").runs[0]._r
        gone = run.makeelement(qn("w:del"), {qn("w:id"): "901", qn("w:author"): "CViche"})
        deleted = run.makeelement(qn("w:r"), {})
        t = deleted.makeelement(qn("w:t"), {})
        t.text = " Retracted."
        deleted.append(t)
        gone.append(deleted)
        run.addnext(gone)
        return doc

    review = _save(with_deletion(), tmp_path / "review.docx")
    report = cf.comment_fates(review, _save(with_deletion(), tmp_path / "corrected.docx"), AUTHOR)
    assert _verdicts(report)[2] == cf.REVIEW_UNKNOWN


def test_the_clean_document_uploaded_has_no_comments_to_follow(tmp_path):
    """Unchanged text is unknown, never not_a_problem; an edit is still fixed."""
    doc = _document(comments=False)
    p = _paragraph(doc, "Quorvane T, Plesk")
    p.runs[0].text = p.text.replace("2019", "2020")
    report = _fates(tmp_path, doc)
    assert not report.comments_tracked
    assert _verdicts(report) == {0: cf.REVIEW_FIXED, 1: cf.REVIEW_UNKNOWN,
                                 2: cf.REVIEW_UNKNOWN, 3: cf.REVIEW_UNKNOWN}


def test_another_authors_comment_is_not_followed(tmp_path):
    """A reviewer's own comment neither counts as a CViche comment kept nor gets a verdict."""
    doc = _document(comments=False)
    doc.add_comment(_paragraph(doc, "Quorvane T, Plesk").runs, text=FLAG, author="Synthetic Reviewer")
    report = _fates(tmp_path, doc)
    assert not report.comments_tracked
    assert [f.comment_id for f in report.fates] == [0, 1, 2, 3]


def test_two_comments_of_one_wording_on_one_block_each_need_their_own_copy(tmp_path):
    """Both comments sit on A; the reviewer deletes one of them."""
    review = _document(comments=True)
    doc = _document(comments=True)
    for d in (review, doc):
        d.add_comment(_paragraph(d, "Quorvane T, Plesk").runs, text=FLAG, author=AUTHOR)
    _delete_comment(doc, 4)
    report = cf.comment_fates(_save(review, tmp_path / "r.docx"), _save(doc, tmp_path / "c.docx"), AUTHOR)
    assert sorted(f.verdict for f in report.fates if f.comment_id in (0, 4)) == [
        cf.REVIEW_NOT_A_PROBLEM, cf.REVIEW_UNKNOWN]


def test_doctor_review_writes_fixed_as_tp_and_dismissed_as_fp_and_drops_unknown():
    findings = [
        {"lint": "implausible_year", "shape": None, "severity": "WARN", "entry_index": 12, "verdict": "fixed"},
        {"lint": "pipe_leaks", "shape": None, "severity": "WARN", "entry_index": None,
         "verdict": "not_a_problem"},
        {"lint": "pipe_leaks", "shape": None, "severity": "WARN", "entry_index": 3, "verdict": "unknown"},
        {"lint": None, "shape": None, "severity": None, "entry_index": None, "verdict": "fixed"},
    ]
    assert cf.doctor_review(findings) == [
        {"lint": "implausible_year", "shape": None, "severity": "WARN", "verdict": "TP",
         "element_idx_start": 12},
        {"lint": "pipe_leaks", "shape": None, "severity": "WARN", "verdict": "FP", "element_idx_start": None},
    ]
