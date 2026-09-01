"""Regression test for the `_add_table_row` mixin collision (hotfix, no
issue number yet -- introduced by #625 / commit aeea7a7).

Reproduces the exact crash: `ClinicalPracticeSection._add_table_row` (4
positional args, no `entry`) shadowed `MembershipsSection._add_table_row`
(4th arg `entry`, called with `entry=entry` at `memberships.py:124` and
`:147`) because `ClinicalPracticeSection` precedes `MembershipsSection` in
`WCMTemplateGenerator`'s base list and Python's MRO resolves `self.<name>` to
the first definition it finds. The result was:

    TypeError: ClinicalPracticeSection._add_table_row() got an unexpected
    keyword argument 'entry'

on 52 of 66 corpus CVs, with the 1287-test pipeline suite green throughout --
none of those tests reach `_add_table_row` through the COMPOSED
`WCMTemplateGenerator` class from the memberships call site; they either
construct `MembershipsSection` in isolation or exercise `ClinicalPracticeSection`
through the composed class instead.

This drives the memberships path itself: build a real `WCMTemplateGenerator`
(so the method resolves through the actual 23-mixin MRO, not a hand-picked
subset), hand it a python-docx table and a memberships-shaped entry, and call
`self._add_table_row(table, data, entry=entry)` exactly as
`memberships.py:147` does. Before the fix this raises the `TypeError` above;
after, `ClinicalPracticeSection._add_table_row` is gone (renamed to
`_add_clinical_table_row`) and the call resolves to
`MembershipsSection._add_table_row`, which accepts `entry` and writes the row.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_memberships_table_row_regression.py -p no:cacheprovider
"""

import sys
from pathlib import Path

import pytest
from docx import Document

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


def _generator_with_table(header_cells):
    """A real WCMTemplateGenerator over a minimal hand-built table -- same
    shape as the fixture in test_stage6_clinical_practice_round2.py's
    `_fake_doc_with_table`, so this test resolves methods through the
    identical composed class the production crash did.
    """
    gen = WCMTemplateGenerator(verbose=False)
    doc = Document()
    table = doc.add_table(rows=1, cols=len(header_cells))
    for cell, text in zip(table.rows[0].cells, header_cells):
        cell.text = text
    gen.doc = doc
    return gen, table


class TestMembershipsAddTableRowThroughComposedClass:
    def test_entry_kwarg_does_not_raise_and_row_text_lands(self):
        gen, table = _generator_with_table(["Organization", "Dates"])

        entry = {
            "text": "Fellow, American College of Surgeons | 2015-Present",
            "extracted_fields": {
                "organization": "American College of Surgeons",
                "membership_type": "Fellow",
                "start_date": "2015",
                "end_date": "Present",
            },
        }

        # This is memberships.py:147's own call, unmodified:
        #     self._add_table_row(table, [org_text, date_str], entry=entry)
        # Pre-fix this raises TypeError because the name resolves to
        # ClinicalPracticeSection's 4-positional-arg version instead.
        gen._add_table_row(
            table,
            ["Fellow, American College of Surgeons", "2015-Present"],
            entry=entry,
        )

        assert len(table.rows) == 2, "no row was added"
        added_row = table.rows[1]
        assert added_row.cells[0].text == "Fellow, American College of Surgeons"
        assert added_row.cells[1].text == "2015-Present"

    def test_multi_membership_shape_also_survives(self):
        """The other memberships.py call site (line 124, inside the
        multi-membership loop) -- same signature, different caller context.
        """
        gen, table = _generator_with_table(["Organization", "Dates"])

        entry = {"text": "Member\nSociety A\nSociety B | date1\ndate2"}

        gen._add_table_row(table, ["Society A", "date1"], entry=entry)
        gen._add_table_row(table, ["Society B", "date2"], entry=entry)

        assert len(table.rows) == 3
        assert table.rows[1].cells[0].text == "Society A"
        assert table.rows[2].cells[0].text == "Society B"


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
