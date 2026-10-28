"""#985: an institution sub-heading fills the empty cells of the rows under it.

`stage6/institution_headings.py` is pure, so every case below is entries in,
entries and groups out, with stage 4's own weak heading-shape test passed in
as stage 6 passes it. The wiring (that `_group_entries_by_code` runs the pass
on the routed entries) is pinned in `test_stage6_render_drops.py`, next to the
other tests of that method; the last test here renders a whole document.

Every value is invented. The shapes are the ones the #985 census found in the
stored runs: an institution line over committees (WZSMUO 331), a society line
over its offices (FLYBMX 298-341), a mentee list's institution line (WZSMUO
166-198), an unextracted name line (NJIKGI 212), and the near-misses the pass
must leave alone.

    python3 -m pytest src/unified_pipeline/tests/test_stage6_institution_headings.py -p no:cacheprovider
"""
import copy
import json
from collections import defaultdict

import pytest
from docx import Document

from unified_pipeline import stage_6_word_template as stage6
from unified_pipeline.stage4.context_headings import _is_heading_shaped
from unified_pipeline.stage6 import institution_headings as ih

SERVICE = ('Institutional Service',)


def _entry(code: str, text: str, el: int, hierarchy: tuple[str, ...] = SERVICE, **fields) -> dict:
    return {'taxonomy_code': code, 'text': text, 'element_idx_start': el, 'element_idx_end': el,
            'hierarchy': list(hierarchy), 'extracted_fields': dict(fields)}


def _committee(name: str, el: int, code: str = 'P', **fields) -> dict:
    """A dated committee row, as a CV writes one: never heading-shaped."""
    return _entry(code, f'{name}, 2001-2003', el, committee_name=name, start_date='2001', **fields)


def _at(entry: dict, hierarchy: tuple[str, ...]) -> dict:
    return {**entry, 'hierarchy': list(hierarchy)}


def _coded(*entries: dict) -> list[tuple[dict, str]]:
    return [(entry, entry['taxonomy_code']) for entry in entries]


def _consume(*entries: dict) -> tuple[dict[str, list[dict]], list[dict]]:
    coded = _coded(*entries)
    groups: dict[str, list[dict]] = defaultdict(list)
    for entry, code in coded:
        groups[code].append(entry)
    warnings: list[dict] = []
    ih.consume_institution_headings(coded, groups, _is_heading_shaped, warnings=warnings)
    return groups, warnings


def _filled(entry: dict, key: str) -> object:
    return entry['extracted_fields'].get(key)


# ----------------------------------------------------------------- the shapes

def test_an_institution_line_fills_the_committees_under_it_and_is_consumed() -> None:
    heading = _entry('P', 'Example University', 10, institution='Example University')
    first, second = _committee('Admissions Committee', 11), _committee('Space Task Force', 12)
    groups, _ = _consume(heading, first, second)
    assert _filled(first, 'institution') == _filled(second, 'institution') == 'Example University'
    assert groups['P'] == [first, second]


def test_a_society_line_fills_the_organization_of_its_offices_and_memberships() -> None:
    heading = _entry('I', 'Northfield Society of Invented Pharmacy', 20,
                     organization='Northfield Society of Invented Pharmacy')
    office = _entry('Q1', 'President, 2007-2009', 21, role='President', start_date='2007')
    member = _entry('I', 'Member, 2007-2018', 22, membership_type='Member', start_date='2007')
    groups, _ = _consume(heading, office, member)
    assert _filled(office, 'organization') == _filled(member, 'organization') == heading['text']
    assert groups['I'] == [member]


def test_a_mentee_list_institution_line_fills_the_mentee_institution() -> None:
    hierarchy = ('Graduate Student Mentorship', 'Thesis Advisor')
    heading = _entry('N3B', 'Example University', 30, hierarchy, site_position='Example University')
    past = _entry('N3B', 'A. Student, Neuroscience Program, Ph.D. 2002', 31, hierarchy,
                  mentee_name='A. Student', site_position='Neuroscience Program')
    current = _entry('N3A', 'B. Student, Neuroscience Program', 32, hierarchy,
                     mentee_name='B. Student', site_position='Neuroscience Program')
    groups, _ = _consume(heading, past, current)
    assert _filled(past, 'institution') == _filled(current, 'institution') == 'Example University'
    assert groups['N3B'] == [past]


def test_a_trailing_colon_is_not_part_of_the_name() -> None:
    heading = _entry('K2', 'Example University:', 40, institution='Example University')
    rows = [_entry('N3B', f'Fellow {n}, MD', 41 + n, mentee_name=f'Fellow {n}') for n in range(2)]
    groups, _ = _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Example University'] * 2
    assert 'K2' not in groups


def test_rows_on_a_sub_level_of_the_heading_are_its_rows() -> None:
    heading = _entry('P', 'Example University', 50, ('Administrative Duties',),
                     institution='Example University')
    programs = ('Administrative Duties', 'Example University Programs')
    committees = ('Administrative Duties', 'Example University Committees')
    rows = [_entry('O', 'Director, Combined Program, 1997-2001', 52, programs, leadership_role='Director'),
            _at(_committee('Steering Committee', 53), programs),
            _at(_committee('Search Committee', 56), committees),
            _at(_committee('Planning Committee', 57), committees)]
    _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Example University'] * 4


def test_an_unextracted_name_line_is_a_heading_when_a_row_under_it_names_it() -> None:
    heading = _entry('Q2', 'Exampleton', 60)
    rows = [_committee('Provost committee on work life', 62),
            _committee('Exampleton advisory board', 64, institution='Exampleton'),
            _committee('Faculty welfare subcommittee', 66)]
    groups, _ = _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Exampleton'] * 3
    assert 'Q2' not in groups


def test_a_row_naming_the_heading_institution_or_a_unit_of_it_keeps_it_and_the_run_goes_on() -> None:
    heading = _entry('Q2', 'Exampleton', 70, organization='Exampleton')
    unit = _committee('Exampleton Cancer Center search committee', 72, institution='Exampleton Cancer Center')
    rows = [_committee('Judicial Panels', 71), unit, _committee('Focus Group', 73)]
    _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Exampleton', 'Exampleton Cancer Center', 'Exampleton']
    society = _entry('I', 'Example Society of Imaginary Medicine (ESIM)', 80,
                     organization='Example Society of Imaginary Medicine (ESIM)')
    rows = [_committee('Program Committee', 81, code='Q2'),
            _committee('Other ESIM Roles', 82, code='Q2', organization='ESIM'),
            _committee('Audit Committee', 83, code='Q2')]
    _consume(society, *rows)
    assert [_filled(row, 'organization') for row in rows] == [society['text'], 'ESIM', society['text']]


def test_an_alias_and_reordered_connectives_still_read_as_the_name() -> None:
    """WZSMUO 354's shape: stage 4 reorders a name with its parenthesised
    alias, and drops or adds a connective; the text is still only the name."""
    heading = _entry('P', 'Examplefield (University of Examplefield, UEF)', 10,
                     institution='Examplefield University (UEF)')
    rows = [_committee('Admissions Committee', 11), _committee('Search Committee', 12)]
    _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == [heading['text']] * 2


def test_restating_and_blank_fields_do_not_make_a_heading_a_record() -> None:
    """A heading stage 5c formatted (WZSMUO 132's `formatted_text`), or with a
    blank field stage 4 left, is still only a name."""
    heading = _entry('P', 'Example University', 10, institution='Example University',
                     formatted_text='Example University', formatting_source='stage_5c_llm',
                     role='  ', notes=[], extra={})
    rows = [_committee('Admissions Committee', 11), _committee('Search Committee', 12)]
    _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Example University'] * 2


def test_fanned_out_records_of_one_element_are_one_run() -> None:
    """WYMVGU 930: stage 6's fan-out gives each record of an entry its own
    child, all on the entry's element; a row may also span several elements."""
    heading = _entry('I', 'Example Section of Imaginary Research', 930,
                     organization='Example Section of Imaginary Research')
    rows = [_committee('Session co-chairman', 931, code='Q1'), _committee('Abstract reviewer', 931, code='Q1'),
            {**_committee('Vice-President', 932, code='Q1'), 'element_idx_end': 934},
            _committee('President', 935, code='Q1')]
    _consume(heading, *rows)
    assert [_filled(row, 'organization') for row in rows] == [heading['text']] * 4


def test_rows_of_one_element_on_a_sub_level_set_no_negative_spacing() -> None:
    """The first step measured can be between two fanned-out records of one
    element; it counts as adjacent, so the next adjacent row stays in the run."""
    heading = _entry('P', 'Example University', 50, ('Duties',), institution='Example University')
    level = ('Duties', 'Example University Committees')
    rows = [_at(_committee('Admissions Committee', 52), level), _at(_committee('Search Committee', 52), level),
            _at(_committee('Planning Committee', 53), level)]
    _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Example University'] * 3


def test_a_mentee_list_run_ends_at_a_committee_row() -> None:
    hierarchy = ('Mentoring',)
    heading = _entry('N3B', 'Example University', 30, hierarchy, site_position='Example University')
    mentees = [_entry('N3B', f'Student {n}, Neuroscience Program', 31 + n, hierarchy, mentee_name=f'Student {n}')
               for n in range(2)]
    after = [_at(_committee('Admissions Committee', 33), hierarchy), _at(_committee('Search Committee', 34), hierarchy)]
    _consume(heading, *mentees, *after)
    assert [_filled(row, 'institution') for row in (*mentees, *after)] == ['Example University'] * 2 + [None] * 2


def test_the_heading_bound_is_stage_4s() -> None:
    from unified_pipeline.stage4 import context_headings
    assert ih.MAX_HEADING_CHARS == context_headings.MAX_HEADING_CHARS


def test_an_unextracted_dated_line_is_not_a_heading_even_when_a_row_names_it() -> None:
    _unchanged(_entry('Q2', 'Exampleton 2005', 60), _committee('Provost committee', 62),
               _committee('Exampleton advisory board', 64, institution='Exampleton'),
               _committee('Faculty welfare subcommittee', 66))


def test_entries_without_integer_element_indexes_are_still_read() -> None:
    for index in ('table_3', None):
        heading = {**_entry('P', 'Example University', 10, institution='Example University'),
                   'element_idx_start': index, 'element_idx_end': index}
        rows = [{**_committee('Admissions Committee', 0), 'element_idx_start': None, 'element_idx_end': None},
                {**_committee('Search Committee', 0), 'element_idx_start': 'table_4', 'element_idx_end': 'table_4'}]
        _consume(heading, *rows)
        assert [_filled(row, 'institution') for row in rows] == ['Example University'] * 2


def test_a_leadership_row_naming_its_organization_keeps_it() -> None:
    """Leadership prints `institution` or `organization`, so an O row holding
    either names its own; a blank one is not a name."""
    heading = _entry('P', 'Example University', 10, institution='Example University')
    rows = [_committee('Admissions Committee', 11),
            _entry('O', 'Chair, Example University Senate, 2001', 12, leadership_role='Chair',
                   organization='Example University Senate'),
            _committee('Search Committee', 13, institution='   ')]
    _consume(heading, *rows)
    assert [_filled(row, 'institution') for row in rows] == ['Example University', None, 'Example University']
    assert _filled(rows[1], 'organization') == 'Example University Senate'


def test_a_long_name_up_to_the_bound_is_a_heading_and_past_it_is_not() -> None:
    def outcome(name: str) -> list[object]:
        rows = [_committee('Admissions Committee', 11), _committee('Search Committee', 12)]
        _consume(_entry('P', name, 10, institution=name), *rows)
        return [_filled(row, 'institution') for row in rows]

    at_bound = 'Example ' + 'x' * (ih.MAX_HEADING_CHARS - len('Example '))
    assert outcome(at_bound) == [at_bound] * 2
    assert outcome(at_bound + 'y') == [None] * 2


def test_consuming_without_a_warnings_list_still_fills() -> None:
    heading = _entry('P', 'Example University', 10, institution='Example University')
    rows = [_committee('Admissions Committee', 11), _committee('Search Committee', 12)]
    groups = {'P': [heading, *rows]}
    ih.consume_institution_headings(_coded(heading, *rows), groups, _is_heading_shaped)
    assert groups == {'P': rows} and _filled(rows[0], 'institution') == 'Example University'


def test_the_warning_names_element_indexes_and_no_text() -> None:
    heading = _entry('P', 'Example University', 10, institution='Example University')
    _, warnings = _consume(heading, _committee('Admissions Committee', 11), _committee('Space Task Force', 12))
    assert warnings == [{
        'check': ih.CONSUMED_HEADING_CHECK, 'code': None, 'section': None,
        'message': 'institution sub-heading moved into the rows under it: 1 heading, 2 rows filled',
        'evidence': ['element_idx_start 10 filled 11, 12'], 'severity': 'INFO',
        ih.CONSUMED_HEADINGS_KEY: [[10, [11, 12]]]}]
    assert ih.heading_fills(json.loads(json.dumps(warnings))) == {10: [11, 12]}


@pytest.mark.parametrize('warnings', [
    [],
    [{'check': 'another_check', ih.CONSUMED_HEADINGS_KEY: [[10, [11]]]}],
    [{'check': ih.CONSUMED_HEADING_CHECK, ih.CONSUMED_HEADINGS_KEY: [[10, 11], [12], 'x']}],
    ['not a record', {'check': ih.CONSUMED_HEADING_CHECK}],
])
def test_heading_fills_reads_nothing_from_other_or_malformed_records(warnings: list) -> None:
    assert ih.heading_fills(warnings) == {}


def test_a_filled_row_gets_its_own_fields_dict() -> None:
    shared = {'committee_name': 'Admissions Committee'}
    first = {**_committee('Admissions Committee', 11), 'extracted_fields': shared}
    second = {**_committee('Admissions Committee', 12), 'extracted_fields': shared}
    _consume(_entry('P', 'Example University', 10, institution='Example University'), first, second)
    assert shared == {'committee_name': 'Admissions Committee'}
    assert first['extracted_fields'] is not second['extracted_fields']


def test_a_heading_that_is_its_codes_only_entry_leaves_no_empty_group() -> None:
    groups, _ = _consume(_entry('G', 'Example Medical Group', 10, organization='Example Medical Group'),
                         _committee('Board Committee', 11), _committee('Strategy Committee', 12))
    assert sorted(groups) == ['P']


# ------------------------------------------------------------- the near-misses

def _unchanged(*entries: dict) -> None:
    before = copy.deepcopy(list(entries))
    groups, warnings = _consume(*entries)
    assert list(entries) == before and warnings == []
    assert sum(len(group) for group in groups.values()) == len(entries)


def test_a_one_line_record_with_an_institution_and_a_date_is_not_a_heading() -> None:
    _unchanged(_entry('I', 'Example Society of Imaginary Surgery 2002-2018', 10,
                      organization='Example Society of Imaginary Surgery', start_date='2002'),
               _entry('Q1', 'President, 2014-2015', 11, role='President', start_date='2014'),
               _entry('Q1', 'Treasurer, 2010-2012', 12, role='Treasurer', start_date='2010'))


def test_rows_that_each_name_their_own_institution_are_left_alone() -> None:
    _unchanged(_entry('P', 'Example University', 10, institution='Example University'),
               _committee('Cancer Committee, Westbrook Hospital', 11, institution='Westbrook Hospital'),
               _committee('IRB, Eastfield Clinic', 12, institution='Eastfield Clinic'))


def test_a_heading_over_one_row_is_left_alone() -> None:
    _unchanged(_entry('P', 'Example University', 10, institution='Example University'),
               _committee('Admissions Committee', 11))


@pytest.mark.parametrize('flag', ['is_fragment', 'is_duplicate'])
@pytest.mark.parametrize('line', [
    _entry('P', 'Other Example Hospital', 123, institution='Other Example Hospital'),
    _committee('Other Example Hospital guidelines', 123),
])
def test_a_line_stage_3b_dropped_ends_the_run_before_it(line: dict, flag: str) -> None:
    """VQFSDI 126: an institution line 3b flagged as a fragment of the row
    above. Stage 4 skips it, so the run must not reach the rows below it."""
    heading = _entry('P', 'Example University', 120, institution='Example University')
    first, second = _committee('Admissions Committee', 121), _committee('Search Committee', 122)
    fragment = {**copy.deepcopy(line), flag: True}
    after = [_committee('Guidelines Committee', 124), _committee('Fundraising Committee', 125)]
    _consume(heading, first, second, fragment, *after)
    assert [_filled(row, 'institution') for row in (first, second, *after)] == ['Example University'] * 2 + [None] * 2


def test_a_line_stage_4_dropped_ends_the_run_at_the_gap_it_leaves() -> None:
    """VQFSDI 126: the fragment is not in stage 4's output at all, only the
    element index it held is missing between two rows that were adjacent."""
    heading = _entry('P', 'Example University', 120, institution='Example University')
    first, second = _committee('Admissions Committee', 121), _committee('Search Committee', 122)
    after = [_committee('Guidelines Committee', 124), _committee('Fundraising Committee', 125)]
    _consume(heading, first, second, *after)
    assert [_filled(row, 'institution') for row in (first, second, *after)] == ['Example University'] * 2 + [None] * 2


def test_rows_a_blank_line_apart_are_one_run() -> None:
    heading = _entry('Q2', 'Exampleton', 172, organization='Exampleton')
    rows = [_committee(f'Committee {n}', 174 + 2 * n, code='Q2') for n in range(3)]
    _consume(heading, *rows)
    assert [_filled(row, 'organization') for row in rows] == ['Exampleton'] * 3


def test_a_heading_whose_next_entries_sit_under_another_leaf_is_left_alone() -> None:
    _unchanged(_entry('P', 'Example University', 10, ('Administrative Duties', 'Programs'),
                      institution='Example University'),
               _committee('Admissions Committee', 12) | {'hierarchy': ['Administrative Duties', 'Committees']},
               _committee('Search Committee', 13) | {'hierarchy': ['Administrative Duties', 'Committees']})


@pytest.mark.parametrize('code, fields', [
    ('K1', {'course_title': 'Neurobiology', 'role': 'Lecturer'}),
    ('R', {'title': 'Invited talk', 'date': '2001'}),
    ('D1', {'title': 'Professor', 'start_date': '2001'}),
])
def test_a_row_whose_renderer_prints_no_filled_institution_gets_nothing(code: str, fields: dict) -> None:
    _unchanged(_entry('P', 'Example University', 10, institution='Example University'),
               _entry(code, 'A row of another kind, 2001', 11, **fields),
               _committee('Admissions Committee', 12), _committee('Search Committee', 13))


@pytest.mark.parametrize('name', ['Medical Center', 'University', 'School of Medicine', 'Board of Directors',
                                  'Department of Medicine', 'College of Pharmacy', 'Health System'])
def test_a_name_of_a_kind_of_institution_is_not_a_heading(name: str) -> None:
    _unchanged(_entry('P', name, 10, institution=name),
               _committee('Admissions Committee', 11), _committee('Search Committee', 12))


def test_an_appointment_line_is_never_a_heading() -> None:
    _unchanged(_entry('D3', 'Example University', 10, organization='Example University'),
               _entry('O', 'Director, Example Program', 11, leadership_role='Director'),
               _entry('O', 'Chair, Example Council', 12, leadership_role='Chair'))


def test_a_row_naming_another_institution_ends_the_run() -> None:
    heading = _entry('I', 'Example Society of Imaginary Medicine', 60,
                     organization='Example Society of Imaginary Medicine')
    rows = [_committee('Research Committee', 61, code='Q2'), _committee('Program Committee', 62, code='Q2'),
            _committee('Liaison Committee, Westbrook Academy', 63, code='Q2', organization='Westbrook Academy'),
            _committee('Finance Committee', 64, code='Q2')]
    _consume(heading, *rows)
    assert [_filled(row, 'organization') for row in rows] == [heading['text']] * 2 + ['Westbrook Academy', None]


@pytest.mark.parametrize('ender', [
    _entry('Q2', 'Council on Professional Affairs', 63, committee_name='Council on Professional Affairs'),
    _entry('T', 'Meeting Location', 63, description='Meeting Location'),
    _committee('Hospital Committee', 63),
    _entry('Q2', 'A row stage 4 extracted nothing from', 63),
])
def test_a_sub_heading_another_family_t_or_an_unextracted_row_ends_the_run(ender: dict) -> None:
    heading = _entry('I', 'Example Society of Imaginary Medicine', 60,
                     organization='Example Society of Imaginary Medicine')
    rows = [_committee('Research Committee', 61, code='Q2'), _committee('Program Committee', 62, code='Q2')]
    after = [_committee('Finance Committee', 64, code='Q2'), _committee('Audit Committee', 65, code='Q2')]
    _consume(heading, *rows, ender, *after)
    assert [_filled(row, 'organization') for row in (*rows, *after)] == [heading['text']] * 2 + [None] * 2


def test_the_next_heading_ends_the_run_and_starts_its_own() -> None:
    """The next heading here is not heading-shaped (a comma), so only being a
    heading ends the first run."""
    first = _entry('I', 'Example Society of Imaginary Medicine', 60,
                   organization='Example Society of Imaginary Medicine')
    second = _entry('Q2', 'Westbrook Academy, Eastfield', 63, organization='Westbrook Academy, Eastfield')
    rows = [_committee('Research Committee', 61, code='Q2'), _committee('Program Committee', 62, code='Q2')]
    after = [_committee('Finance Committee', 64, code='Q2'), _committee('Audit Committee', 65, code='Q2')]
    groups, warnings = _consume(first, *rows, second, *after)
    assert [_filled(row, 'organization') for row in (*rows, *after)] == [first['text']] * 2 + [second['text']] * 2
    assert groups['Q2'] == [*rows, *after]
    assert warnings[0]['evidence'] == ['element_idx_start 60 filled 61, 62', 'element_idx_start 63 filled 64, 65']
    assert warnings[0]['message'].endswith(': 2 headings, 4 rows filled')


def test_a_unit_heading_under_a_heading_starts_its_own_run() -> None:
    """A sub-heading naming a unit of the heading's institution is a heading:
    the first run ends there, and the rows under it take the unit's name."""
    first = _entry('Q2', 'Exampleton', 60, organization='Exampleton')
    unit = _entry('P', 'Exampleton Cancer Center, Main Campus', 63,
                  institution='Exampleton Cancer Center, Main Campus')
    rows = [_committee('Judicial panels', 61), _committee('Focus group', 62)]
    after = [_committee('Search committee', 64), _committee('Program committee', 65)]
    _, warnings = _consume(first, *rows, unit, *after)
    assert [_filled(row, 'institution') for row in (*rows, *after)] == ['Exampleton'] * 2 + [unit['text']] * 2
    assert warnings[0]['evidence'] == ['element_idx_start 60 filled 61, 62', 'element_idx_start 63 filled 64, 65']


def test_a_row_naming_a_numbered_part_of_the_heading_ends_the_run() -> None:
    """A row's number is a word of its name: "Example Society 2" is not
    "Example Society of Imaginary Medicine", so the run stops there."""
    heading = _entry('I', 'Example Society of Imaginary Medicine', 60,
                     organization='Example Society of Imaginary Medicine')
    rows = [_committee('Research committee', 61, code='Q2'), _committee('Program committee', 62, code='Q2'),
            _committee('Liaison committee', 63, code='Q2', organization='Example Society 2'),
            _committee('Finance committee', 64, code='Q2')]
    _consume(heading, *rows)
    assert [_filled(row, 'organization') for row in rows] == [heading['text']] * 2 + ['Example Society 2', None]


def test_an_unextracted_name_line_no_row_names_is_left_alone() -> None:
    _unchanged(_entry('Q2', 'Exampleton', 60), _committee('Provost committee', 62),
               _committee('Faculty welfare subcommittee', 64))


def test_a_cell_of_its_own_row_is_not_a_heading() -> None:
    """HVLDGS 1878: stage 2 split a table row's first cell off as an entry."""
    _unchanged(_entry('Q2', 'Celebrate Example Science', 10, organization='Celebrate Example Science'),
               _committee('Celebrate Example Science\tDepartment Contact', 10),
               _committee('Department Contact and Organizer', 11))


def test_a_heading_with_more_than_its_name_is_not_a_heading() -> None:
    """Even with a row under it naming the same institution, which would
    corroborate a line stage 4 extracted nothing from."""
    _unchanged(_entry('P', 'Example University', 10, institution='Example University', role='Member'),
               _committee('Admissions Committee', 11), _committee('Search Committee', 12),
               _committee('Example University Senate', 13, institution='Example University'))


@pytest.mark.parametrize('text', [
    'Example University\tMain Campus', 'Example University | Main Campus',
    'Example University\nMain Campus', 'Example University (Main Campus',
])
def test_a_line_with_a_tab_pipe_break_or_open_parenthesis_is_not_a_name(text: str) -> None:
    _unchanged(_entry('P', text, 10, institution=text),
               _committee('Admissions Committee', 11), _committee('Search Committee', 12))


@pytest.mark.parametrize('code', ['T', 'D3'])
def test_a_t_or_appointment_line_is_not_a_heading_even_when_a_row_names_it(code: str) -> None:
    """T is stage 4's stamp's (#1047); an appointment is positions'. Neither
    with the name extracted nor without it."""
    for fields in ({'institution': 'Exampleton'}, {}):
        _unchanged(_entry(code, 'Exampleton', 10, **fields),
                   _committee('Admissions Committee', 11), _committee('Search Committee', 12),
                   _committee('Exampleton advisory board', 13, institution='Exampleton'))


def test_an_unextracted_name_of_a_kind_of_institution_is_not_a_heading() -> None:
    _unchanged(_entry('Q2', 'Medical Center', 10), _committee('Admissions Committee', 11),
               _committee('Search Committee', 12), _committee('Medical Center board', 13, institution='Medical Center'))


def test_a_name_field_holding_a_list_is_not_a_name() -> None:
    _unchanged(_entry('P', 'Example University', 10, institution=['Example University']),
               _committee('Admissions Committee', 11), _committee('Search Committee', 12))


def test_a_row_with_no_fields_dict_or_a_placeholder_name_ends_the_run() -> None:
    heading = _entry('P', 'Example University', 10, institution='Example University')
    rows = [_committee('Admissions Committee', 11), _committee('Search Committee', 12)]
    for ender in ({**_committee('Space Committee', 13), 'extracted_fields': None},
                  _committee('Space Committee', 13, institution='-')):
        after = [_committee('Planning Committee', 14), _committee('Budget Committee', 15)]
        _consume(copy.deepcopy(heading), *copy.deepcopy(rows), ender, *after)
        assert [_filled(row, 'institution') for row in after] == [None, None]


def test_a_text_with_words_its_name_fields_lack_is_not_a_heading() -> None:
    _unchanged(_entry('Q1', 'Other Example Leadership Roles', 10, organization='Example'),
               _entry('Q1', 'Chair, 2010-2011', 11, role='Chair'),
               _entry('Q1', 'Vice Chair, 2009-2010', 12, role='Vice Chair'))


# ------------------------------------------------------------ a rendered page

def test_the_rendered_committee_rows_carry_the_institution_and_no_heading_row(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(stage6, 'call_llm', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('no LLM in tests')))
    entries = [
        {'taxonomy_code': 'A', 'element_idx_start': 0, 'text': 'Name: Pat Example, MD', 'extracted_fields': {}},
        _entry('P', 'Quillmoor University', 10, institution='Quillmoor University'),
        _committee('Fernhollow Admissions Committee', 11),
        _committee('Fernhollow Search Committee', 12),
    ]
    source = tmp_path / 'TESTIH_fields.json'
    source.write_text(json.dumps({'document_uid': 'TESTIH', 'entries': entries}))
    out = stage6.run_stage6(str(source), str(tmp_path / 'out.docx'), verbose=False)
    document = Document(out)
    rows = [' | '.join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
    lines = [p.text.strip() for p in document.paragraphs] + rows
    assert sum('Fernhollow Admissions Committee, Quillmoor University' in row for row in rows) == 1
    assert sum('Fernhollow Search Committee, Quillmoor University' in row for row in rows) == 1
    assert not [line for line in lines if 'Quillmoor' in line and 'Fernhollow' not in line]
    sidecar = json.loads((tmp_path / 'TESTIH_render_warnings.json').read_text())
    assert [w['check'] for w in sidecar['warnings'] if w['check'] == ih.CONSUMED_HEADING_CHECK] == [
        ih.CONSUMED_HEADING_CHECK]
