"""Tests for the gold-set segmentation regression harness (#208 precondition).

Pure-function tests only — no LLM, no gold corpus needed. The one docx used
is built in-test to prove the coverage metric sees INSIDE table cells (the
89HQVQ blind spot: whole sections living in 1×1 layout-table cells).

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_segmentation_regression.py -p no:cacheprovider
"""

import sys
from collections import Counter
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.segmentation_regression import (  # noqa: E402
    _COUNT_KEYS,
    _has_compact_window,
    BODY_BLOCK,
    COVERAGE_DROP_TOLERANCE_PTS,
    COVERAGE_WINDOW_SLACK,
    MEGA_ENTRY_MIN_RECORDS,
    SUBSTANTIVE_LINE_CHARS,
    compare_metrics,
    compute_metrics,
    find_lost_blocks,
    iter_source_block_lines,
    iter_source_lines,
    lint_metrics,
)


# ------------------------------------------------------------- source lines

def test_iter_source_lines_sees_into_table_cells(tmp_path):
    doc = Document()
    doc.add_paragraph("EDUCATION")
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.paragraphs[0].text = "Grants Awarded"
    cell.add_paragraph("FSMB Foundation Grant | Role: Co-PI | Status: Awarded 2026")
    path = tmp_path / "cv.docx"
    doc.save(path)

    lines = iter_source_lines(str(path))
    assert "EDUCATION" in lines
    assert "Grants Awarded" in lines
    assert any("FSMB" in l for l in lines)


def test_iter_source_lines_reads_through_tracked_insertion_in_cell(tmp_path):
    """#557: a run nested inside w:ins (tracked-change insertion) must not be
    dropped mid-word. Reproduces "Down Syndrome" -> "Down yndrome": a table
    cell paragraph built as 'Down ' + <w:ins>S</w:ins> + 'yndrome' must read
    back whole, not with the tracked-change character silently skipped."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    para = cell.paragraphs[0]
    para.add_run("Down ")
    tail_run = para.add_run("yndrome")
    ins = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t>S</w:t></w:r></w:ins>'
    )
    tail_run._r.addprevious(ins)

    path = tmp_path / "cv.docx"
    doc.save(path)

    lines = iter_source_lines(str(path))
    assert "Down Syndrome" in lines
    assert "Down yndrome" not in lines


def test_iter_source_lines_recurses_into_nested_table(tmp_path):
    """iter_source_lines's docstring claims recursive table support
    ("INCLUDING paragraphs inside table cells (recursively)") -- prove a
    table nested inside a cell is actually walked, not just a single level
    of cells."""
    doc = Document()
    outer_table = doc.add_table(rows=1, cols=1)
    outer_cell = outer_table.rows[0].cells[0]
    outer_cell.paragraphs[0].text = "Outer cell text"
    nested_table = outer_cell.add_table(rows=1, cols=1)
    nested_table.rows[0].cells[0].paragraphs[0].text = "Nested table text"

    path = tmp_path / "cv.docx"
    doc.save(path)

    lines = iter_source_lines(str(path))
    assert "Outer cell text" in lines
    assert "Nested table text" in lines


# ------------------------------------------------------------------ metrics

_GRANT_A = "FSMB Foundation Grant | Shapiro, M. (PI) | Improving Access to Healthcare | Role: Co-PI"
_GRANT_B = "John Templeton Foundation OFI | Jung, E. (PI) | Intellectual Humility | Role: PI"
_GRANT_C = "NBME Stemmler Education Grant | Jung, E. (PI) | Use of AI | Role: PI | Status: Not Funded"

_STAGE1A = {"hierarchy": [
    {"text": "GRANTS", "children": [
        {"text": "Grants Awarded", "children": []},
        {"text": "Grants Under Review & Submitted", "children": []},
    ]},
]}


def _entry(text, etype="paragraph", start=0, hierarchy=None):
    return {"text": text, "element_type": etype, "element_idx_start": start,
            "hierarchy": hierarchy or ["GRANTS"]}


def test_metrics_on_clean_extraction():
    source = ["GRANTS", _GRANT_A, _GRANT_B, _GRANT_C]
    stage2 = {"entries": [_entry(_GRANT_A, start=1), _entry(_GRANT_B, start=2),
                          _entry(_GRANT_C, start=3)]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["text_coverage_pct"] == 100.0
    assert m["lost_lines"] == []
    assert m["mega_entries"] == 0
    assert m["headers_detected"] == 3
    assert m["per_h1_content_counts"] == {"grants": 3}


def test_metrics_detect_mega_entry_and_lost_text():
    # The 89HQVQ shape: grants fused into one entry, one grant dropped entirely.
    source = ["GRANTS", _GRANT_A, _GRANT_B, _GRANT_C]
    mega = _entry(_GRANT_A + "\n" + _GRANT_B + "\n" + _GRANT_A + " extra", start=1)
    stage2 = {"entries": [mega]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["mega_entries"] == 1
    assert any("NBME" in l for l in m["lost_lines"])
    assert m["text_coverage_pct"] < 100.0


def test_metrics_count_dups_and_empties():
    stage2 = {"entries": [
        _entry(_GRANT_A, etype="table", start=30.0),
        _entry(_GRANT_A, etype="table", start=30.0),  # duplicate
        _entry("", start=5),                           # empty content
        _entry("", etype="break", start=6),            # break: NOT empty-content
    ]}
    m = compute_metrics([], _STAGE1A, stage2)
    assert m["duplicate_entries"] == 1
    assert m["empty_content"] == 1


def test_coverage_is_whitespace_insensitive():
    """Stage 2 fuses text across in-paragraph breaks with NO whitespace
    ('Present position:' + break + 'Attending' -> 'position:Attending') and
    drops tabs from label/value lines. Neither may count as lost text."""
    source = [
        "Present position:\tAttending Pathologist of the Department",
        "Urine cytology: Pitfall due to a\n\"remnant\" lesion. Cytojournal. 2015",
    ]
    stage2 = {"entries": [
        _entry("Present position:Attending Pathologist of the Department", start=1),
        _entry('Urine cytology: Pitfall due to a"remnant" lesion. Cytojournal. 2015', start=2),
    ]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["lost_lines"] == []
    assert m["text_coverage_pct"] == 100.0


def test_coverage_line_cannot_match_across_entry_boundary():
    # First half in one entry, second half in the next: that IS a lost line
    # (the sentinel between entries must prevent a spanning match).
    source = ["alpha beta gamma delta epsilon zeta"]
    stage2 = {"entries": [_entry("alpha beta gamma", start=1),
                          _entry("delta epsilon zeta", start=2)]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["lost_lines"] == ["alpha beta gamma delta epsilon zeta"]


def test_coverage_words_scattered_through_unrelated_entry_are_lost():
    """#610: an entry that merely CONTAINS the line's words, far apart and out
    of order inside an unrelated sentence, does not cover the line."""
    source = ["Diabetes Research Grant Foundation Award"]
    stage2 = {"entries": [_entry(
        "The foundation gave an award to the committee chair after a long "
        "review of her research on diabetes, and a further grant was pending "
        "in the spring of the following year", start=1)]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["lost_lines"] == ["Diabetes Research Grant Foundation Award"]
    assert m["text_coverage_pct"] == 0.0


def test_coverage_reformatted_line_with_spliced_phrase_is_still_covered():
    """The mid-line merge the token fallback exists for (web053/web057):
    stage 2 splices a phrase into the line, so it is no contiguous substring,
    but every token is present in order and close together."""
    source = ["\t\t1984-1989\t\t\t\tB.S.\t (Biology)"]
    stage2 = {"entries": [_entry("1984-1989    B.S. University of Utah (Biology)", start=1)]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["lost_lines"] == []
    assert m["text_coverage_pct"] == 100.0


def test_coverage_reordered_but_compact_line_is_still_covered():
    """Stage 2 can reorder a line's parts inside one cell; nothing is lost."""
    source = ["Maple Ridge Medical School, Springfield, Freedonia - Visiting Faculty"]
    stage2 = {"entries": [_entry(
        "Visiting Faculty Northwind Maple Ridge Medical School, Springfield, "
        "Freedonia  June 2006", start=1)]}
    m = compute_metrics(source, _STAGE1A, stage2)
    assert m["lost_lines"] == []


def test_compact_window_slack_boundary():
    need = Counter(["a", "b"])
    filler = ["x"] * COVERAGE_WINDOW_SLACK
    span = 2 + COVERAGE_WINDOW_SLACK
    assert _has_compact_window(need, ["a"] + filler + ["b"], span)
    assert not _has_compact_window(need, ["a"] + filler + ["x", "b"], span)
    # Multiplicity counts: two "a" tokens need two occurrences in the window.
    assert not _has_compact_window(Counter(["a", "a"]), ["a", "b", "c"], 9)
    # The window slides: a later compact stretch qualifies.
    assert _has_compact_window(need, ["a"] + filler + ["x", "a", "b"], span)

    # A surplus repeat of a needed token must not count toward the missing total.
    assert not _has_compact_window(Counter("ab"), list("aaxxx"), 5)
    # The left edge advances one token at a time: the window [a, y, b] (span 3)
    # is reachable only by not skipping the "a" that precedes it.
    assert _has_compact_window(Counter("ab"), list("xayb"), 3)
    assert not _has_compact_window(Counter("ab"), list("xayb"), 2)


def test_coverage_long_reformatted_line_is_covered_by_its_own_length():
    """#610: the span budget scales with the line's token count, so a long
    line (16 tokens) spliced with 2 extra words (window 18, over the bare
    slack) is still covered."""
    words = [f"term{c}" for c in "abcdefghijklmnop"]
    source = [" ".join(words)]
    entry = " ".join(words[:8] + ["inserted", "words"] + words[8:])
    m = compute_metrics(source, _STAGE1A, {"entries": [_entry(entry, start=1)]})
    assert m["lost_lines"] == []


def test_coverage_span_budget_counts_repeated_tokens():
    """#610: the budget is the line's token count WITH multiplicity. A 14-token
    line over 2 distinct words, spliced with 5 extra tokens (window 19), fits
    14 + slack but not distinct-count + slack (2 + 12)."""
    source = [" ".join(["alpha", "beta"] * 7)]
    entry = " ".join(["alpha", "beta"] * 4 + ["z1", "z2", "z3", "z4", "z5"]
                     + ["alpha", "beta"] * 3)
    m = compute_metrics(source, _STAGE1A, {"entries": [_entry(entry, start=1)]})
    assert m["lost_lines"] == []


def test_coverage_substantive_line_without_tokens_is_lost():
    """A long punctuation-only line has no word tokens to match on, so unless
    it appears verbatim it is lost (not vacuously covered)."""
    source = ["----------------------------------------"]
    m = compute_metrics(source, _STAGE1A,
                        {"entries": [_entry("Some unrelated paragraph text", start=1)]})
    assert m["lost_lines"] == ["----------------------------------------"]


def test_template_scaffolding_is_neither_covered_nor_lost():
    """#815 (976WPY): a CV written on the WCM template keeps the template's
    prompts, which stage 2 rightly drops. They must not count as lost."""
    verbatim = "Other Educational Experiences (i.e., certificates, etc)"
    reworded = "Academic Degree(s) (Bachelor's and higher)"  # near-match only
    source = [verbatim, reworded, _GRANT_A]
    m = compute_metrics(source, _STAGE1A, {"entries": [_entry(_GRANT_A, start=1)]})
    assert m["lost_lines"] == []
    assert m["text_coverage_pct"] == 100.0


def test_short_template_text_and_real_content_still_count_as_lost():
    """Short template strings double as real values ("Full-time salaried by
    Weill Cornell" is the Employment Status answer a CV kept), so below the
    length floor they still count; so does absent real content."""
    source = ["Full-time salaried by Weill Cornell", "2. Principal Investigator",
              _GRANT_B, _GRANT_A]
    m = compute_metrics(source, _STAGE1A, {"entries": [_entry(_GRANT_A, start=1)]})
    assert m["lost_lines"] == source[:3]
    assert m["text_coverage_pct"] == 25.0


def test_substantive_line_chars_boundary():
    """A source line of exactly SUBSTANTIVE_LINE_CHARS is substantive (and here
    lost); one char shorter is ignored by the coverage metric."""
    at_floor = "a" * SUBSTANTIVE_LINE_CHARS
    below_floor = "a" * (SUBSTANTIVE_LINE_CHARS - 1)
    m_at = compute_metrics([at_floor], _STAGE1A, {"entries": []})
    assert m_at["substantive_lines"] == 1
    assert m_at["lost_lines"] == [at_floor]
    assert m_at["text_coverage_pct"] == 0.0
    m_below = compute_metrics([below_floor], _STAGE1A, {"entries": []})
    assert m_below["source_lines"] == 1
    assert m_below["substantive_lines"] == 0
    assert m_below["lost_lines"] == []
    assert m_below["text_coverage_pct"] == 100.0


def _records(n):
    return "\n".join([_GRANT_A, _GRANT_B, _GRANT_C][:n])


def test_mega_entry_threshold_boundary():
    """MEGA_ENTRY_MIN_RECORDS - 1 record-like lines is NOT a mega-entry; the
    threshold count is."""
    below = MEGA_ENTRY_MIN_RECORDS - 1
    m_below = compute_metrics([], _STAGE1A, {"entries": [_entry(_records(below))]})
    assert m_below["mega_entries"] == 0
    m_at = compute_metrics([], _STAGE1A, {"entries": [_entry(_records(MEGA_ENTRY_MIN_RECORDS))]})
    assert m_at["mega_entries"] == 1


def test_same_text_different_start_is_not_a_duplicate():
    """Duplicate identity is (type, text, start): identical text at a different
    element_idx_start (a genuinely repeated line) must not be flagged."""
    stage2 = {"entries": [_entry(_GRANT_A, start=1), _entry(_GRANT_A, start=2)]}
    assert compute_metrics([], _STAGE1A, stage2)["duplicate_entries"] == 0


def test_header_and_break_entries_excluded_from_content_metrics():
    stage2 = {"entries": [
        _entry(_records(MEGA_ENTRY_MIN_RECORDS), etype="header", start=1),
        _entry(_records(MEGA_ENTRY_MIN_RECORDS), etype="break", start=2),
        _entry(_GRANT_A, start=3),
    ]}
    m = compute_metrics([], _STAGE1A, stage2)
    assert m["entries_total"] == 3
    assert m["entries_content"] == 1
    assert m["mega_entries"] == 0
    assert m["per_h1_content_counts"] == {"grants": 1}


def test_empty_source_reports_full_coverage_and_no_lost_lines():
    m = compute_metrics([], _STAGE1A, {"entries": [_entry(_GRANT_A, start=1)]})
    assert m["source_lines"] == 0
    assert m["text_coverage_pct"] == 100.0
    assert m["lost_lines"] == []


def test_header_titles_exact_and_partial_shapes_tolerated():
    m = compute_metrics([], _STAGE1A, {"entries": []})
    assert m["header_titles"] == ["grants", "grants awarded",
                                  "grants under review & submitted"]
    # Missing hierarchy / untitled node / entry with no keys: defensive paths.
    sparse = {"hierarchy": [{"children": [{"text": "Kept"}]}]}
    m = compute_metrics([], sparse, {"entries": [{}]})
    assert m["header_titles"] == ["kept"]
    assert m["empty_content"] == 1
    assert m["per_h1_content_counts"] == {"(none)": 1}
    m = compute_metrics([], {}, {})
    assert m["headers_detected"] == 0
    assert m["entries_total"] == 0


def test_iter_source_block_lines_tags_each_table_and_matches_flat_lines(tmp_path):
    doc = Document()
    doc.add_paragraph("A body paragraph outside every table")
    first = doc.add_table(rows=2, cols=1)
    first.cell(0, 0).text = "first table row one"
    first.cell(1, 0).text = "first table row two"
    second = doc.add_table(rows=1, cols=1)
    second.cell(0, 0).text = "second table cell"
    second.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "nested table cell"
    path = tmp_path / "blocks.docx"
    doc.save(path)

    pairs = iter_source_block_lines(str(path))
    assert pairs == [
        (BODY_BLOCK, "A body paragraph outside every table"),
        (0, "first table row one"),
        (0, "first table row two"),
        (1, "second table cell"),
        (2, "nested table cell"),
    ]
    assert iter_source_lines(str(path)) == [line for _, line in pairs]


def _covered_line(i):
    return f"covered record line number {i} alpha"


def _lost_line(i):
    return f"lost personal data line {i} zebra"


def test_small_table_lost_whole_is_found_despite_high_document_coverage():
    """#815 (web207): a small table lost whole is a rounding error against
    the rest of the document, so the document-wide lint stays quiet."""
    covered = [_covered_line(i) for i in range(200)]
    lost = [_lost_line(i) for i in range(5)]
    block_lines = [(0, l) for l in covered] + [(1, l) for l in lost]
    stage2 = {"entries": [_entry(l, start=i) for i, l in enumerate(covered)]}

    metrics = compute_metrics([l for _, l in block_lines], _STAGE1A, stage2)
    assert metrics["text_coverage_pct"] >= 97.0
    assert not any(f.startswith("coverage") for f in lint_metrics(metrics))

    assert find_lost_blocks(block_lines, stage2) == [
        {"block": 1, "substantive_lines": 5, "lost_lines": lost}]
    # Exactly half of a table lost is already a lost block.
    half = [(2, covered[0]), (2, covered[1]), (2, lost[0]), (2, lost[1])]
    assert find_lost_blocks(half, stage2) == [
        {"block": 2, "substantive_lines": 4, "lost_lines": lost[:2]}]


def test_lost_lines_scattered_across_tables_or_in_body_are_not_a_lost_block():
    covered = [_covered_line(i) for i in range(50)]
    lost = [_lost_line(i) for i in range(5)]
    stage2 = {"entries": [_entry(l, start=i) for i, l in enumerate(covered)]}
    # One lost line in each of five 10-line tables: 10% per table.
    scattered = [(i // 10, l) for i, l in enumerate(covered)] + [
        (i, l) for i, l in enumerate(lost)]
    assert find_lost_blocks(scattered, stage2) == []
    # A whole table lost below the size floor, and lost body text.
    # (Short lines are not substantive, so they do not lift it over the floor.)
    small = [(0, l) for l in covered] + [
        (1, lost[0]), (1, lost[1]), (1, "2016"), (1, "PhD")]
    body = [(0, l) for l in covered] + [(BODY_BLOCK, l) for l in lost]
    assert find_lost_blocks(small, stage2) == []
    assert find_lost_blocks(body, stage2) == []


# ------------------------------------------------------------------ compare

_BASE = {
    "text_coverage_pct": 99.0, "lost_lines": ["old lost line"],
    "headers_detected": 10, "header_titles": ["grants", "education"],
    "mega_entries": 1, "duplicate_entries": 0, "empty_content": 0,
}


def _cand(**over):
    c = dict(_BASE)
    c.update(over)
    return c


def test_compare_flags_coverage_drop_beyond_tolerance():
    verdict, reasons = compare_metrics(_BASE, _cand(text_coverage_pct=97.0))
    assert verdict == "REGRESSION"
    assert any("coverage" in r for r in reasons)


def test_compare_tolerates_llm_wobble():
    verdict, _ = compare_metrics(_BASE, _cand(text_coverage_pct=98.5))
    assert verdict == "OK"


def test_compare_flags_newly_lost_lines_even_within_tolerance():
    verdict, reasons = compare_metrics(
        _BASE, _cand(lost_lines=["old lost line", "a brand new lost line"]))
    assert verdict == "REGRESSION"
    assert any("newly lost" in r for r in reasons)


def test_compare_flags_disappeared_headers_by_name():
    verdict, reasons = compare_metrics(
        _BASE, _cand(headers_detected=9, header_titles=["education"]))
    assert verdict == "REGRESSION"
    assert any("grants" in r for r in reasons)


def test_compare_reports_improvement():
    verdict, reasons = compare_metrics(_BASE, _cand(mega_entries=0))
    assert verdict == "IMPROVED"
    assert any("mega_entries" in r for r in reasons)


def test_compare_coverage_drop_exact_tolerance_boundary():
    """A drop EQUAL to the tolerance is wobble (OK); just above is REGRESSION."""
    base = _cand(text_coverage_pct=99.0)
    at = round(99.0 - COVERAGE_DROP_TOLERANCE_PTS, 1)
    assert compare_metrics(base, _cand(text_coverage_pct=at)) == ("OK", [])
    above = round(at - 0.1, 1)
    verdict, reasons = compare_metrics(base, _cand(text_coverage_pct=above))
    assert verdict == "REGRESSION"
    assert any("coverage" in r for r in reasons)


@pytest.mark.parametrize("key", _COUNT_KEYS)
def test_compare_count_key_regression_and_improvement(key):
    base = _cand(**{key: 1})
    verdict, reasons = compare_metrics(base, _cand(**{key: 2}))
    assert verdict == "REGRESSION"
    assert reasons == [f"{key} 1 -> 2"]
    verdict, reasons = compare_metrics(base, _cand(**{key: 0}))
    assert verdict == "IMPROVED"
    assert reasons == [f"{key} 1 -> 0"]


def test_compare_reports_every_simultaneous_regression_reason():
    verdict, reasons = compare_metrics(_BASE, _cand(
        text_coverage_pct=90.0, lost_lines=["old lost line", "brand new loss"],
        headers_detected=9, header_titles=["education"],
        mega_entries=2, duplicate_entries=1, empty_content=1))
    assert verdict == "REGRESSION"
    assert len(reasons) == 3 + len(_COUNT_KEYS)  # coverage, lost lines, headers


# --------------------------------------------------------------------- lint

def test_lint_flags_and_clean():
    dirty = _cand(text_coverage_pct=90.0, mega_entries=2)
    assert lint_metrics(dirty)
    clean = _cand(text_coverage_pct=99.5, mega_entries=0,
                  duplicate_entries=0, empty_content=0)
    assert lint_metrics(clean) == []
