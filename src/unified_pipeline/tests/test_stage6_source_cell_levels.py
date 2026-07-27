"""Sub-bullets inside a source table cell must render as level-1 children (#423).

The reader flattens a table cell to newline-joined text
(`docx_structure_extractor.py:340`), so the per-paragraph `w:numPr` that marks
which lines are children of the entry's title never reaches any artifact --
`list_level` appears in zero persisted stage JSON. Branch A of
`_insert_teaching_entry` then emits one same-level bullet per line, which is the
flat dump #423 reports. The cell still has the markup and the entry still has
the table/row coordinates, so stage 6 re-reads it from the source.

Measured on the issue's own CV (C0ZGFW): 64 of its branch-A K entries convert.
The 100-CV web-harvest corpus converts 0 -- it holds 8 title-then-bullets cells
against production's 223 -- so the corpus is the no-collateral arm of the gate,
never the evidence the fix works.

These tests guard the WIRE, not just the helper. The branch-A loop runs over
`reversed(original_lines)`, so a fix that indexes with `j` instead of
`len - 1 - j` inverts the hierarchy -- the title becomes a child of its own
children -- while every helper-level assertion still passes.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_source_cell_levels.py -p no:cacheprovider
"""

import ast
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402

_MODULE = Path(WCMTemplateGenerator.__module__.replace(".", "/"))
_SOURCE = (_SRC / _MODULE).with_suffix(".py")

TITLE = "Weill Bugango Medical Center POCUS Program"
CHILDREN = ["Develop curriculum for medicine residents",
            "Coordinate with the Global Health program"]


def _bullet(para):
    """Give a fixture paragraph real list markup, without calling the code under test."""
    pPr = para._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "1")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)


def _source_docx(tmp_path, lines, bulleted):
    """A one-row table whose first cell holds `lines`, bulleting where flagged."""
    doc = Document()
    cell = doc.add_table(rows=1, cols=2).rows[0].cells[0]
    for i, text in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.add_run(text)
        if bulleted[i]:
            _bullet(para)
    path = tmp_path / "source.docx"
    doc.save(str(path))
    return str(path)


def _entry(lines):
    return {
        "taxonomy_code": "K3",
        "element_type": "table_row",
        "table_index": 0,
        "row_index": 0,
        "text": "\n".join(lines),
        "extracted_fields": {"formatted_text": " | ".join(lines)},
    }


def _render(source_docx, lines):
    gen = WCMTemplateGenerator(verbose=False, source_docx=source_docx)
    gen.doc = Document(gen.template_path)
    gen._fill_teaching({"K3": [_entry(lines)]})
    return gen


def _level(para):
    """The paragraph's w:ilvl, or None when it carries no list markup."""
    pPr = para._p.pPr
    numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
    if numPr is None:
        return None
    return numPr.find(qn("w:ilvl")).get(qn("w:val"))


def _levels(gen, lines):
    """Rendered ilvl per source line, in source order."""
    out = []
    for line in lines:
        hits = [p for p in gen.doc.paragraphs if line in p.text]
        assert len(hits) == 1, f"expected one paragraph for {line!r}, got {len(hits)}"
        out.append(_level(hits[0]))
    return out


# --- the wire ----------------------------------------------------------------

def test_source_sub_bullets_render_as_level_1_children(tmp_path):
    lines = [TITLE] + CHILDREN
    src = _source_docx(tmp_path, lines, [False, True, True])
    assert _levels(_render(src, lines), lines) == ["0", "1", "1"]


def test_the_title_is_never_demoted_below_its_children(tmp_path):
    # Indexing the reversed loop with j instead of len-1-j returns exactly
    # ["1", "1", "0"] here, and every helper-level assertion still passes.
    lines = [TITLE] + CHILDREN
    src = _source_docx(tmp_path, lines, [False, True, True])
    levels = _levels(_render(src, lines), lines)
    assert levels[0] == "0", "the title must stay at the top level"
    assert all(lv == "1" for lv in levels[1:]), "children must sit below the title"


# --- and every guard returns None, leaving today's flat render ---------------

def test_without_a_source_docx_every_line_stays_level_0(tmp_path):
    lines = [TITLE] + CHILDREN
    _source_docx(tmp_path, lines, [False, True, True])  # exists, but not passed
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen._fill_teaching({"K3": [_entry(lines)]})
    assert _levels(gen, lines) == ["0", "0", "0"]


def test_a_flat_source_cell_stays_flat(tmp_path):
    lines = [TITLE] + CHILDREN
    src = _source_docx(tmp_path, lines, [False, False, False])
    assert _levels(_render(src, lines), lines) == ["0", "0", "0"]


def test_a_cell_of_peer_bullets_stays_flat(tmp_path):
    # No title, so there is no hierarchy to restore. Without the
    # title-then-bullets test the whole run is promoted to level 1 and a plain
    # bulleted list silently gains an indent step -- and a prior measurement put
    # 104 of 217 source sub-bullet clusters in exactly this shape, where
    # flattening is the correct WCM rendering.
    lines = [TITLE] + CHILDREN
    src = _source_docx(tmp_path, lines, [True, True, True])
    assert _levels(_render(src, lines), lines) == ["0", "0", "0"]


def test_a_row_that_does_not_line_up_is_rejected(tmp_path):
    # Stage 2 rewrites some rows, and a paragraph count that disagrees with the
    # line count means the lookup found the wrong content. Guessing an alignment
    # would weld one line's hierarchy onto another.
    lines = [TITLE] + CHILDREN
    src = _source_docx(tmp_path, lines + ["A fourth source paragraph"],
                       [False, True, True, True])
    assert _levels(_render(src, lines), lines) == ["0", "0", "0"]


def test_a_row_whose_text_disagrees_is_rejected(tmp_path):
    # Measured on 6 real entries: same paragraph count, different content. The
    # count guard alone would have accepted all six.
    lines = [TITLE] + CHILDREN
    src = _source_docx(tmp_path, ["An unrelated row"] + CHILDREN, [False, True, True])
    assert _levels(_render(src, lines), lines) == ["0", "0", "0"]


def test_a_missing_source_file_is_not_fatal(tmp_path):
    lines = [TITLE] + CHILDREN
    assert _levels(_render(str(tmp_path / "gone.docx"), lines), lines) == ["0", "0", "0"]


# --- the call site actually asks for a computed level ------------------------

def test_branch_a_passes_a_computed_level_not_a_constant():
    # Guards the wire: `list_level=0` here is what shipped the flat dump, and it
    # passes every helper-level test in this file.
    tree = ast.parse(_SOURCE.read_text())
    fn = next(n for n in ast.walk(tree)
              if isinstance(n, ast.FunctionDef) and n.name == "_insert_teaching_entry")
    calls = [c for c in ast.walk(fn)
             if isinstance(c, ast.Call) and isinstance(c.func, ast.Attribute)
             and c.func.attr == "_insert_bulleted_entry"]
    levels = [kw.value for c in calls for kw in c.keywords if kw.arg == "list_level"]
    assert levels, "no call site passes list_level"
    assert any(not isinstance(v, ast.Constant) for v in levels), \
        "every list_level is a constant -- the source-cell levels are unwired"


def test_both_pipelines_hand_stage_6_the_source_docx():
    # The CLI and the backend orchestrator are parallel reimplementations that
    # share no code, so a correct helper behind one unwired caller is invisible:
    # run_stage6 defaults source_docx to None and the render silently stays flat.
    root = _SRC.parent
    for rel in ("run_full_pipeline.py",
                "web_interface/backend/app/pipeline/orchestrator.py"):
        tree = ast.parse((root / rel).read_text())
        calls = [c for c in ast.walk(tree) if isinstance(c, ast.Call)
                 and (getattr(c.func, "id", None) == "run_stage6"
                      or any(getattr(a, "id", None) == "run_stage6" for a in c.args))]
        assert calls, f"{rel}: no run_stage6 call found"
        assert any(kw.arg == "source_docx" for c in calls for kw in c.keywords), \
            f"{rel} calls run_stage6 without source_docx -- #423 is unwired here"
