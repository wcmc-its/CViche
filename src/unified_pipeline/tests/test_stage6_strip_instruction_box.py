"""Regression guard for Stage 6 gray "instruction box" removal.

The WCM template ships a leading gray table on page 1 that instructs the author
how to fill in the CV ("When preparing the WCM CV template ... delete this
instruction box"). That box is template scaffolding baked into the .docx, not
extracted CV content, so the Stage 2 entry filter never sees it — the box has
to be dropped from the generated document itself. ``strip_template_instructions``
(the Run column, default ON) gates the removal.

These tests load the real bundled template into the generator, run the removal
directly, and assert the box (and only the box) is gone. Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_strip_instruction_box.py -p no:cacheprovider

Self-contained: no DB, no FastAPI app. Requires only ``python-docx`` and the
bundled WCM template (auto-located by the generator).
"""

import sys
from pathlib import Path

# Ensure the repo's ``src`` directory is importable regardless of cwd/rootdir.
_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402

_SIGNATURE = "when preparing the wcm cv template"
_CONTENT_TABLE_MARKER = "work email:"  # a real personal-data table, must survive


def _table_texts(doc):
    return [
        " ".join(c.text for row in t.rows for c in row.cells).lower()
        for t in doc.tables
    ]


def test_default_flag_matches_db_column():
    """Constructor default mirrors the Run column default (strip ON)."""
    gen = WCMTemplateGenerator(verbose=False)
    assert gen.strip_template_instructions is True


def test_template_actually_contains_the_box():
    """Guard the fixture: the bundled template must still have the box, else
    the removal tests would pass vacuously."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    texts = _table_texts(gen.doc)
    assert any(_SIGNATURE in t for t in texts), "template lost its instruction box"
    assert any(_CONTENT_TABLE_MARKER in t for t in texts), "template lost its personal-data table"


def test_removal_drops_exactly_the_box():
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    before = len(gen.doc.tables)

    gen._remove_instruction_box()

    after_texts = _table_texts(gen.doc)
    # The box is gone...
    assert not any(_SIGNATURE in t for t in after_texts), "instruction box survived removal"
    # ...exactly one table was removed...
    assert len(gen.doc.tables) == before - 1
    # ...and real content tables are untouched.
    assert any(_CONTENT_TABLE_MARKER in t for t in after_texts), "removal nuked real content"


def test_flag_off_keeps_the_box():
    """With the flag off, generate()'s gate must leave the box in place. We
    assert on the gate condition the generator uses."""
    gen = WCMTemplateGenerator(verbose=False, strip_template_instructions=False)
    gen.doc = Document(gen.template_path)
    before = len(gen.doc.tables)

    # Mirror generate(): removal only runs when the flag is set.
    if gen.strip_template_instructions:
        gen._remove_instruction_box()

    assert len(gen.doc.tables) == before
    assert any(_SIGNATURE in t for t in _table_texts(gen.doc)), "box wrongly removed with flag off"


def test_table_styling_leaves_personal_data_row0_unshaded():
    """_apply_table_styling_to_all_tables shades row 0 of every table as a
    header row. The PERSONAL DATA table has no header row -- its row 0 is
    "Office address:" -- so it must be skipped (faculty feedback 2026-09-15)
    while every other table's row 0 is still shaded."""
    from docx.oxml.ns import qn

    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen._apply_table_styling_to_all_tables()

    def row0_fills(table):
        return [
            (shd.get(qn("w:fill")) if (shd := c._tc.find(".//" + qn("w:shd"))) is not None else None)
            for c in table.rows[0].cells
        ]

    tables = list(gen.doc.tables)  # python-docx builds fresh proxies per access; snapshot once
    texts = _table_texts(gen.doc)
    personal = [t for t, txt in zip(tables, texts) if _CONTENT_TABLE_MARKER in txt]
    assert len(personal) == 1, "expected exactly one personal-data table"
    assert row0_fills(personal[0]) == [None, None], "personal-data row 0 must stay unshaded"

    others = [t for t, txt in zip(tables, texts) if _CONTENT_TABLE_MARKER not in txt and t.rows]
    assert others, "no other tables to check shading on"
    for t in others:
        assert set(row0_fills(t)) == {"D9D9D9"}, "section-table header row lost its shading"
