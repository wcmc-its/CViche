"""Behaviour tests for stage_2_entry_extraction.py (#704 round 2, packet 3).

Covers ``detect_entries_for_section`` (the LLM-driven per-section delimiter
detection: response-shape fallbacks, table-row and pseudo-row extraction,
the table_row recovery backstop, multi-batch splitting, exception handling
per batch, subset-delimiter collapsing, and malformed/out-of-bounds
delimiter rejection) and ``run_stage_2`` end to end (synthetic docx +
hand-built stage-1b hierarchy JSON, ``OutputManager`` redirected to
``tmp_path``, ``call_llm`` stubbed and routed by the exact
``**CV Section Header:** \\`...\\``` prompt line -- never by a raw substring
search, since a section's own body text can innocently contain another
section's name and cause a stub to answer the wrong call). Also covers
``strip_template_instructions`` on/off and ``cancel_check`` raising mid-run.

A sibling packet covers the pure helpers below detect_entries_for_section
(get_hierarchy_path, build_element_index_map, get_element_text,
split_merged_row_into_pseudo_rows, extract_leaf_sections_with_boundaries,
collect_header_indices/info, remove_subset_delimiters,
recover_unclaimed_table_rows, _dedup_idx_key, filter_extraction_noise) --
none of those are re-tested here in isolation, only as they are exercised
through detect_entries_for_section/run_stage_2's own control flow.

Untestable / effectively dead within this scope (see the module notes
inline below): the ``element_idx_start.startswith("table_")`` delimiter
branch (~lines 901-914) can never be populated via any doc_elements shape
build_element_index_map/detect_entries_for_section's own element-building
step produces -- the outer ``isinstance(idx, int)`` gate (~line 534)
requires an integer unified_idx before any element enters
``section_elements`` at all, so a batch's ``idx`` key is never a
"table_N" string by the time delimiters are validated. The same reasoning
makes sort_key's "table_N"/bare-numeric-string elif branches (~1286-1288)
and each ``doc.paragraphs[idx]`` fallback used when an index is absent
from ``element_index_map`` (~1118, 1226, 1245) dead in practice --
element_index_map is always built from the same doc_structure that
produced every index run_stage_2 looks up. The parent-header gap/break
loop at ~1148-1163 requires a 3-level-deep hierarchy where an
intermediate has_children node contributes no leaf of its own (so
extract_leaf_sections_with_boundaries's own gap-fill never fires for the
outer parent) -- reachable only by a contrived fixture, not by any
shape stage 1b's compute_section_boundaries is known to emit; left
uncovered rather than pinned to a shape with no traced provenance.
``__main__`` (argparse/CLI entry, ~1372-1382) is not exercised -- no
corpus, no real docx.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage2_detect_and_run_behaviour.py -p no:cacheprovider
"""

import inspect
import io
import json
import logging
import sys
import threading
import time
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import pytest  # noqa: E402
from docx import Document  # noqa: E402

from unified_pipeline import stage_2_entry_extraction as stage2  # noqa: E402
from unified_pipeline.core.output_manager import OutputManager as _RealOutputManager  # noqa: E402


# --------------------------------------------------------------- shared helpers

def _llm_result(content_obj, cost=0.001, total_tokens=10, prompt_tokens=8, completion_tokens=2):
    """Build a call_llm-shaped return dict; content_obj is JSON-serialized."""
    return {
        "content": json.dumps(content_obj),
        "cost": cost,
        "total_tokens": total_tokens,
        "prompt_tokens": prompt_tokens,
        "completion_tokens": completion_tokens,
    }


def _para(idx, text):
    return {"unified_idx": idx, "type": "paragraph", "text": text}


def _idx_map(elements):
    return stage2.build_element_index_map({"elements": elements})


def _redirect_output_manager(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Force run_stage_2's OutputManager to write under tmp_path instead of the
    real src/unified_pipeline/outputs/."""

    def factory(input_path, base_output_dir=None):
        return _RealOutputManager(input_path, base_output_dir=base_output_dir or tmp_path)

    monkeypatch.setattr(stage2, "OutputManager", factory)


def _route_call_llm(monkeypatch: pytest.MonkeyPatch, routes: dict, calls: list | None = None) -> None:
    """Stub stage2.call_llm, routed by the EXACT `CV Section Header` prompt
    line rather than a raw substring search of the whole prompt -- a
    section's own raw-data body can innocently contain another section's
    name (e.g. "Intro paragraph before Books subsection" contains "Books"),
    which would misroute a naive `"Books" in prompt` stub."""

    def fake(**kwargs):
        prompt = kwargs["messages"][1]["content"]
        header_line = next(line for line in prompt.splitlines() if "CV Section Header" in line)
        if calls is not None:
            calls.append(header_line)
        for key, delimiters in routes.items():
            if f"`{key}`" in header_line:
                return _llm_result({"delimiters": delimiters})
        return _llm_result({"delimiters": []})

    monkeypatch.setattr(stage2, "call_llm", fake)


def _write_hierarchy(tmp_path: Path, name: str, document_uid: str, hierarchy_with_indices: list, section_boundaries: list) -> Path:
    path = tmp_path / name
    path.write_text(json.dumps({
        "document_uid": document_uid,
        "hierarchy_with_indices": hierarchy_with_indices,
        "section_boundaries": section_boundaries,
    }))
    return path


# =============================================================== _element_text_or_fallback

def test_element_text_or_fallback_uses_the_map_silently_when_index_is_present(caplog):
    elements = [_para(0, "Jane Doe")]
    text = stage2._element_text_or_fallback(0, _idx_map(elements), Document())
    assert text == "Jane Doe"
    assert caplog.records == []  # positive case: no warning when the index IS in the map


def test_element_text_or_fallback_warns_and_falls_back_to_doc_paragraphs_when_index_missing(caplog):
    # #915 item 2: measured on the 111-uid local corpus, 2/106 matched uids
    # had a stage-1b hierarchy referencing indices past a fresh
    # extraction's doc_length -- a real, if rare, mismatch, not dead code.
    # Negative case: index absent from element_index_map falls back to
    # doc.paragraphs AND logs, rather than silently returning "".
    doc = Document()
    doc.add_paragraph("Fallback Text")
    with caplog.at_level(logging.WARNING, logger="unified_pipeline.stage_2_entry_extraction"):
        text = stage2._element_text_or_fallback(0, {}, doc)
    assert text == "Fallback Text"
    assert len(caplog.records) == 1
    assert "missing from element_index_map" in caplog.records[0].message


def test_element_text_or_fallback_returns_empty_string_past_doc_paragraphs_bounds(caplog):
    with caplog.at_level(logging.WARNING, logger="unified_pipeline.stage_2_entry_extraction"):
        text = stage2._element_text_or_fallback(5, {}, Document())
    assert text == ""
    assert len(caplog.records) == 1


# =============================================================== detect_entries_for_section

def test_delimiters_key_merges_multiline_entry_and_uses_given_confidence(monkeypatch):
    elements = [_para(10, "Award A"), _para(11, "continuation detail")]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": 10, "element_idx_end": 11, "element_type": "paragraph", "confidence": 0.9}]}
    ))
    entries, cost_info = stage2.detect_entries_for_section(
        ["Awards"], elements, 10, 11, element_index_map=_idx_map(elements)
    )
    assert entries == [{
        "element_idx_start": 10, "element_idx_end": 11, "element_type": "paragraph",
        "confidence": 0.9, "text": "Award A\tcontinuation detail",
    }]
    assert cost_info == {"cost": 0.001, "tokens": 10, "prompt_tokens": 8, "completion_tokens": 2}


def test_entries_key_used_as_fallback_when_delimiters_key_absent(monkeypatch):
    elements = [_para(1, "Solo paragraph")]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"entries": [{"element_idx_start": 1, "element_idx_end": 1, "element_type": "paragraph", "confidence": 0.7}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["S"], elements, 1, 1, element_index_map=_idx_map(elements))
    assert [e["text"] for e in entries] == ["Solo paragraph"]


def test_result_key_used_as_fallback(monkeypatch):
    elements = [_para(1, "X")]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"result": [{"element_idx_start": 1, "element_idx_end": 1, "element_type": "paragraph"}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["S"], elements, 1, 1, element_index_map=_idx_map(elements))
    assert [e["text"] for e in entries] == ["X"]


def test_data_key_used_as_fallback(monkeypatch):
    elements = [_para(1, "Y")]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"data": [{"element_idx_start": 1, "element_idx_end": 1, "element_type": "paragraph"}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["S"], elements, 1, 1, element_index_map=_idx_map(elements))
    assert [e["text"] for e in entries] == ["Y"]


def test_bare_list_response_and_missing_confidence_defaults_to_one(monkeypatch):
    elements = [_para(20, "Solo entry")]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        [{"element_idx_start": 20, "element_idx_end": 20, "element_type": "paragraph"}]
    ))
    entries, _ = stage2.detect_entries_for_section(["S"], elements, 20, 20, element_index_map=_idx_map(elements))
    assert entries[0]["confidence"] == 1.0
    assert entries[0]["text"] == "Solo entry"


def test_unrecognized_response_key_yields_no_entries(monkeypatch):
    elements = [_para(21, "Whatever")]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"unexpected_key": [{"element_idx_start": 21, "element_idx_end": 21}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["S"], elements, 21, 21, element_index_map=_idx_map(elements))
    assert entries == []


def test_empty_section_returns_early_without_calling_llm(monkeypatch):
    calls = {"n": 0}
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: calls.__setitem__("n", calls["n"] + 1) or _llm_result({"delimiters": []}))
    entries, cost_info = stage2.detect_entries_for_section(["S"], [], 5, 5)
    assert entries == []
    assert cost_info == {"cost": 0, "tokens": 0}
    assert calls["n"] == 0


def test_all_elements_excluded_by_header_indices_returns_early_without_calling_llm(monkeypatch):
    # The section's only element is itself a section header -- it never
    # becomes an LLM candidate, so section_elements ends up empty and the
    # function short-circuits exactly like a truly empty section.
    elements = [_para(5, "SECTION HEADER TEXT")]
    calls = {"n": 0}
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: calls.__setitem__("n", calls["n"] + 1) or _llm_result({"delimiters": []}))
    entries, cost_info = stage2.detect_entries_for_section(
        ["X"], elements, 5, 5, header_indices={5}, element_index_map=_idx_map(elements)
    )
    assert entries == []
    assert cost_info == {"cost": 0, "tokens": 0}
    assert calls["n"] == 0


def test_pseudo_row_split_claims_one_segment_and_recovers_its_sibling(monkeypatch):
    # A merged table row (both cells share the same \n\n segment count) is
    # split into pseudo-rows "5.0.0" / "5.0.1" before the LLM ever sees it.
    row = [{"text": "SGIM workshop\n\nPOCUS mentor"}, {"text": "2020\n\n2021"}]
    elements = [{"unified_idx": 5, "type": "table_content", "table_index": 7, "data": [row], "rows": 1, "cols": 2}]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": "5.0.0", "element_idx_end": "5.0.0", "element_type": "table_row", "confidence": 0.77}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["Teaching"], elements, 5, 5, element_index_map=_idx_map(elements))
    by_idx = {e["element_idx_start"]: e for e in entries}
    # The LLM-claimed pseudo-row: model-attested confidence, correct text.
    assert by_idx["5.0.0"]["text"] == "SGIM workshop | 2020"
    assert by_idx["5.0.0"]["confidence"] == 0.77
    assert "recovered_row" not in by_idx["5.0.0"]
    # Its sibling pseudo-row was never claimed -- the #420 backstop recovers
    # it structurally, distinguishable by recovered_row + lower confidence.
    assert by_idx["5.0.1"]["text"] == "POCUS mentor | 2021"
    assert by_idx["5.0.1"]["recovered_row"] is True
    assert by_idx["5.0.1"]["confidence"] < 0.77


def test_normal_table_row_claimed_and_unclaimed_sibling_row_recovered(monkeypatch):
    rows = [[{"text": "Year"}, {"text": "Award"}], [{"text": "2020"}, {"text": "Best Paper"}]]
    elements = [{"unified_idx": 9, "type": "table_content", "table_index": 2, "data": rows, "rows": 2, "cols": 2}]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"entries": [{"element_idx_start": "9.1", "element_idx_end": "9.1", "element_type": "table_row", "confidence": 0.6}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["Awards"], elements, 9, 9, element_index_map=_idx_map(elements))
    by_idx = {e["element_idx_start"]: e for e in entries}
    assert by_idx["9.1"]["text"] == "2020 | Best Paper"
    assert by_idx["9.1"]["table_index"] == 2
    assert by_idx["9.1"]["row_index"] == 1
    assert by_idx["9.0"]["text"] == "Year | Award"
    assert by_idx["9.0"]["recovered_row"] is True


def test_table_content_with_no_row_data_falls_back_to_single_integer_element(monkeypatch):
    elements = [{"unified_idx": 20, "type": "table_content", "text": "Fallback summary text", "data": [], "rows": 0, "table_index": 3}]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": 20, "element_idx_end": 20, "element_type": "table", "confidence": 0.5}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["Grants"], elements, 20, 20, element_index_map=_idx_map(elements))
    assert entries == [{
        "element_idx_start": 20, "element_idx_end": 20, "element_type": "table",
        "confidence": 0.5, "text": "Fallback summary text",
    }]


def test_legacy_table_multi_row_claimed_and_sibling_row_recovered(monkeypatch):
    elements = [{
        "unified_idx": 30, "type": "table", "table_index": 2,
        "data": [[{"text": "R0C0"}, {"text": "R0C1"}], [{"text": "R1C0"}, {"text": "R1C1"}]],
        "rows": 2, "cols": 2,
    }]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": "30.1", "element_idx_end": "30.1", "element_type": "table_row", "confidence": 0.7}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["Committees"], elements, 30, 30, element_index_map=_idx_map(elements))
    by_idx = {e["element_idx_start"]: e for e in entries}
    assert by_idx["30.1"]["text"] == "R1C0 | R1C1"
    assert by_idx["30.0"]["text"] == "R0C0 | R0C1"
    assert by_idx["30.0"]["recovered_row"] is True


def test_legacy_table_single_row_falls_back_to_whole_table_element(monkeypatch):
    elements = [{
        "unified_idx": 40, "type": "table", "table_index": 5,
        "data": [[{"text": "Solo Row A"}, {"text": "Solo Row B"}]], "rows": 1, "cols": 2,
    }]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": 40, "element_idx_end": 40, "element_type": "table", "confidence": 0.4}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["Misc"], elements, 40, 40, element_index_map=_idx_map(elements))
    # Single-row legacy tables never get split into table_row sub-indices --
    # they stay one integer-idx "table" element, cells tab-joined.
    assert entries == [{
        "element_idx_start": 40, "element_idx_end": 40, "element_type": "table",
        "confidence": 0.4, "text": "Solo Row A\tSolo Row B",
    }]


def test_non_integer_element_idx_is_silently_skipped(monkeypatch):
    # Defensive branch for an element whose idx is neither an int nor a
    # unified_idx-produced value -- it must be dropped without affecting
    # any other, well-formed element in the same doc_elements list.
    elements = [
        {"unified_idx": "not-an-int", "type": "paragraph", "text": "Should be skipped"},
        {"unified_idx": 7, "type": "paragraph", "text": "Real entry"},
    ]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": 7, "element_idx_end": 7, "element_type": "paragraph", "confidence": 0.9}]}
    ))
    entries, _ = stage2.detect_entries_for_section(["Sec"], elements, 7, 7, element_index_map=_idx_map(elements))
    assert entries == [{
        "element_idx_start": 7, "element_idx_end": 7, "element_type": "paragraph",
        "confidence": 0.9, "text": "Real entry",
    }]


def test_table_header_not_excluded_by_header_indices_is_included_and_claimable(monkeypatch):
    # A table_header element NOT listed in header_indices (the "rare case"
    # the source comments call out) is added to section_elements/the prompt
    # like any other candidate, and can be claimed by an integer delimiter
    # exactly like a paragraph.
    elements = [{"unified_idx": 3, "type": "table_header", "text": "SUB HEADER"}]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result(
        {"delimiters": [{"element_idx_start": 3, "element_idx_end": 3, "element_type": "table", "confidence": 0.6}]}
    ))
    entries, _ = stage2.detect_entries_for_section(
        ["Sec"], elements, 3, 3, header_indices=set(), element_index_map=_idx_map(elements)
    )
    assert entries == [{
        "element_idx_start": 3, "element_idx_end": 3, "element_type": "table",
        "confidence": 0.6, "text": "SUB HEADER",
    }]


def test_table_content_dict_and_string_row_shapes_are_recovered(monkeypatch):
    # extract_unified_elements always hands table_content a list-of-cell-dicts
    # row, but the function also defends a dict-shaped row and a bare string
    # row (row_data straight from JSON, or a pre-flattened value). Leave both
    # unclaimed so the #420 recovery backstop proves they were extracted at
    # all -- recovered_row entries only exist for rows the code actually saw.
    elements = [{
        "unified_idx": 9, "type": "table_content", "table_index": 4,
        "data": [{"text": "Dict row text"}, "Plain string row"], "rows": 2, "cols": 1,
    }]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result({"delimiters": []}))
    entries, _ = stage2.detect_entries_for_section(["Sec"], elements, 9, 9, element_index_map=_idx_map(elements))
    by_idx = {e["element_idx_start"]: e for e in entries}
    assert by_idx["9.0"]["text"] == "Dict row text"
    assert by_idx["9.0"]["recovered_row"] is True
    assert by_idx["9.1"]["text"] == "Plain string row"
    assert by_idx["9.1"]["recovered_row"] is True


def test_legacy_table_pseudo_row_dict_row_and_string_row_all_recovered(monkeypatch):
    # Same dict/string row-shape defenses, but in the LEGACY "table" branch,
    # combined with a merged (pseudo-row-splittable) first row -- exercises
    # all three legacy sub-branches in one section.
    merged_row = [{"text": "A1\n\nA2"}, {"text": "B1\n\nB2"}]
    elements = [{
        "unified_idx": 12, "type": "table", "table_index": 6,
        "data": [merged_row, {"text": "Dict legacy row"}, "Plain legacy row"], "rows": 3, "cols": 2,
    }]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result({"delimiters": []}))
    entries, _ = stage2.detect_entries_for_section(["Sec"], elements, 12, 12, element_index_map=_idx_map(elements))
    by_idx = {e["element_idx_start"]: e for e in entries}
    assert by_idx["12.0.0"]["text"] == "A1 | B1"
    assert by_idx["12.0.1"]["text"] == "A2 | B2"
    assert by_idx["12.1"]["text"] == "Dict legacy row"
    assert by_idx["12.2"]["text"] == "Plain legacy row"
    assert all(e["recovered_row"] is True for e in by_idx.values())


def test_batch_exception_is_swallowed_and_later_batch_still_succeeds(monkeypatch):
    # 51 elements force a second 50-element-cap batch (BATCH_SIZE=50). The
    # first batch's call_llm raises; the loop must not abort the section --
    # it logs and continues, and the second batch's entry still lands.
    elements = [_para(100 + i, f"Line {i}") for i in range(51)]
    state = {"n": 0}

    def flaky(**kw):
        state["n"] += 1
        if state["n"] == 1:
            raise RuntimeError("simulated LLM outage")
        return _llm_result({"delimiters": [{"element_idx_start": 150, "element_idx_end": 150, "element_type": "paragraph", "confidence": 0.6}]})

    monkeypatch.setattr(stage2, "call_llm", flaky)
    entries, cost_info = stage2.detect_entries_for_section(
        ["Big"], elements, 100, 150, element_index_map=_idx_map(elements)
    )
    assert state["n"] == 2
    assert entries == [{
        "element_idx_start": 150, "element_idx_end": 150, "element_type": "paragraph",
        "confidence": 0.6, "text": "Line 50",
    }]
    # Cost/tokens only reflect the surviving batch -- the raised batch never
    # reached the accumulation lines.
    assert cost_info["cost"] == 0.001
    assert cost_info["tokens"] == 10


def test_remove_subset_delimiters_collapses_subset_keeps_larger_span(monkeypatch):
    elements = [_para(10 + i, f"P{i}") for i in range(6)]  # idx 10..15
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result({"delimiters": [
        {"element_idx_start": 10, "element_idx_end": 12, "element_type": "paragraph", "confidence": 0.5},
        {"element_idx_start": 10, "element_idx_end": 15, "element_type": "paragraph", "confidence": 0.9},
    ]}))
    entries, _ = stage2.detect_entries_for_section(["Sec"], elements, 10, 15, element_index_map=_idx_map(elements))
    assert len(entries) == 1
    assert entries[0]["element_idx_start"] == 10
    assert entries[0]["element_idx_end"] == 15
    assert entries[0]["text"] == "P0\tP1\tP2\tP3\tP4\tP5"


def test_malformed_and_out_of_bounds_delimiters_are_silently_ignored(monkeypatch):
    # The malformed and out-of-bounds delimiters must be individually SKIPPED
    # (the isinstance/bounds guards `continue` past them), not treated as a
    # batch-fatal error -- a THIRD, well-formed delimiter in the same LLM
    # response must still land. Asserting only `entries == []` here would
    # also pass if the malformed delimiter instead raised and the batch's
    # outer `except Exception` swallowed it (a different, wrong route that
    # happens to yield the same empty result for THIS shape) -- see #6.1:
    # a fallback assertion must distinguish the intended route from the
    # error-swallow route, so the valid delimiter's presence is the real
    # proof the loop kept going rather than aborting.
    elements = [_para(10 + i, f"P{i}") for i in range(6)]  # idx 10..15
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result({"delimiters": [
        {"element_idx_start": "weird", "element_idx_end": 12},  # not int, no "." -> skipped
        {"element_idx_start": 999, "element_idx_end": 999},      # in-bounds type, but out of section range
        {"element_idx_start": 11, "element_idx_end": 11, "element_type": "paragraph", "confidence": 0.8},
    ]}))
    entries, cost_info = stage2.detect_entries_for_section(["Sec"], elements, 10, 15, element_index_map=_idx_map(elements))
    assert entries == [{
        "element_idx_start": 11, "element_idx_end": 11, "element_type": "paragraph",
        "confidence": 0.8, "text": "P1",
    }]
    assert cost_info["cost"] == 0.001


def test_pre_llm_scrub_wire_dob_and_ssn_never_reach_call_llm_messages(monkeypatch, tmp_path):
    """#847 wire test: `extract_unified_elements`'s pre-LLM scrub survives
    into the EXACT `messages` stage 2 hands to `call_llm` -- driven from a
    real synthetic docx through the real reader, not a hand-built element
    dict, so nothing between the reader and the prompt can silently
    reintroduce a raw value."""
    from unified_pipeline.core.docx_structure_extractor import extract_unified_elements

    doc = Document()
    doc.add_paragraph("Date of Birth: 01/02/1970")
    doc.add_paragraph("SSN: 123-45-6789")
    docx_path = tmp_path / "wire_test.docx"
    doc.save(str(docx_path))
    elements = extract_unified_elements(str(docx_path))["elements"]

    captured = {}

    def fake(**kwargs):
        captured["messages"] = kwargs["messages"]
        return _llm_result({"delimiters": []})

    monkeypatch.setattr(stage2, "call_llm", fake)
    stage2.detect_entries_for_section(
        ["S"], elements,
        elements[0]["unified_idx"], elements[-1]["unified_idx"],
        element_index_map=_idx_map(elements),
    )

    prompt_text = "\n".join(m["content"] for m in captured["messages"])
    assert "01/02/1970" not in prompt_text
    assert "123-45-6789" not in prompt_text
    assert "[withheld]" in prompt_text


def test_pre_llm_scrub_wire_dob_in_table_cells_never_reach_call_llm_messages(monkeypatch, tmp_path):
    """#847 round 2 wire test: a DOB label in one table CELL and its value
    in the NEXT cell of the same row -- stage 2 builds its table-row prompt
    text by joining `elem["data"]` cells directly
    (`" | ".join(cell.get("text") ...)`, stage_2_entry_extraction.py), so
    the scrub must have already run per-cell, not just on the element's own
    joined `text` field."""
    from unified_pipeline.core.docx_structure_extractor import extract_unified_elements

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Date of Birth:"
    table.cell(0, 1).text = "01/02/1970"
    docx_path = tmp_path / "wire_test_table.docx"
    doc.save(str(docx_path))
    elements = extract_unified_elements(str(docx_path))["elements"]

    captured = {}

    def fake(**kwargs):
        captured["messages"] = kwargs["messages"]
        return _llm_result({"delimiters": []})

    monkeypatch.setattr(stage2, "call_llm", fake)
    stage2.detect_entries_for_section(
        ["S"], elements,
        elements[0]["unified_idx"], elements[-1]["unified_idx"],
        element_index_map=_idx_map(elements),
    )

    prompt_text = "\n".join(m["content"] for m in captured["messages"])
    assert "01/02/1970" not in prompt_text
    assert "[withheld]" in prompt_text


# =============================================================== run_stage_2 (end to end)

def test_run_stage_2_flat_hierarchy_sorts_headers_content_and_breaks(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    for text in ["Jane Researcher", "EDUCATION", "PhD, Somewhere University, 2001", "AWARDS", "", "Best Paper Award, 2010"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "flat.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "flat_h.json", "TESTUID2",
        hierarchy_with_indices=[
            {"text": "Education", "level": "H1", "element_idx": 1, "children": []},
            {"text": "Awards", "level": "H1", "element_idx": 3, "children": []},
        ],
        section_boundaries=[
            {"hierarchy": ["Education"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False},
            {"hierarchy": ["Awards"], "element_idx_start": 3, "element_idx_end": 5, "has_children": False},
        ],
    )
    calls = []
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Education": [{"element_idx_start": 2, "element_idx_end": 2, "element_type": "paragraph", "confidence": 0.95}],
        "Awards": [{"element_idx_start": 5, "element_idx_end": 5, "element_type": "paragraph", "confidence": 0.85}],
    }, calls)

    output_data, output_path = stage2.run_stage_2(str(docx_path), str(hpath))

    assert tmp_path in output_path.parents
    assert len(calls) == 3
    assert output_data["document_uid"] == "TESTUID2"
    assert output_data["document_length"] == 6
    assert output_data["total_cost"] == pytest.approx(0.003)
    assert output_data["total_tokens"] == 30
    assert output_data["coverage"]["coverage_percentage"] == 100.0
    assert output_data["coverage"]["unaccounted_indices"] == []

    entries = output_data["entries"]
    # Sort order follows element_idx_start ascending regardless of which
    # section/entry-type produced the record.
    assert [(e["element_idx_start"], e["element_type"], e["text"]) for e in entries] == [
        (0, "paragraph", "Jane Researcher"),
        (1, "header", "EDUCATION"),
        (2, "paragraph", "PhD, Somewhere University, 2001"),
        (3, "header", "AWARDS"),
        (4, "break", ""),
        (5, "paragraph", "Best Paper Award, 2010"),
    ]
    assert entries[0]["hierarchy"] == ["Personal Data"]
    assert entries[2]["hierarchy"] == ["Education"]
    assert entries[4]["hierarchy"] == ["Awards"]
    assert entries[5]["hierarchy"] == ["Awards"]

    # What's on disk must match what was returned.
    on_disk = json.loads(output_path.read_text())
    assert on_disk == output_data


def test_run_stage_2_strip_template_instructions_true_drops_instruction_entry(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    for text in ["Jane Doe", "EDUCATION", "Faculty Curriculum Vitae Template"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "tmpl.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "tmpl_h.json", "TMPL1",
        hierarchy_with_indices=[{"text": "Education", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[{"hierarchy": ["Education"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Education": [{"element_idx_start": 2, "element_idx_end": 2, "element_type": "paragraph", "confidence": 0.85}],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))  # default strip=True

    texts = [e["text"] for e in output_data["entries"]]
    assert "Faculty Curriculum Vitae Template" not in texts
    assert texts == ["Jane Doe", "EDUCATION"]
    assert output_data["total_entries"] == 2
    # Coverage is computed from index-assignment tracking BEFORE the
    # template filter runs, so the filtered-out index still counts as
    # "covered" even though its entry vanished from the final list.
    assert output_data["coverage"]["coverage_percentage"] == 100.0


def test_run_stage_2_strip_template_instructions_false_keeps_instruction_entry(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    for text in ["Jane Doe", "EDUCATION", "Faculty Curriculum Vitae Template"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "tmpl2.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "tmpl2_h.json", "TMPL2",
        hierarchy_with_indices=[{"text": "Education", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[{"hierarchy": ["Education"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Education": [{"element_idx_start": 2, "element_idx_end": 2, "element_type": "paragraph", "confidence": 0.85}],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath), strip_template_instructions=False)

    texts = [e["text"] for e in output_data["entries"]]
    assert texts == ["Jane Doe", "EDUCATION", "Faculty Curriculum Vitae Template"]
    assert output_data["total_entries"] == 3


def test_run_stage_2_drops_an_older_template_revision_s_reworded_instruction(tmp_path, monkeypatch):
    """#829: a faculty copy of another template revision words the licensure
    note "...to the NYP Hospital staff..." where the tracked 2020 template
    says "...to the Hospital staff...". Kept, it was classified F1 and stage
    4 read "New York State" out of it as a license the faculty member does
    not hold. The near-match drops it; the real licence row beside it stays."""
    _redirect_output_manager(monkeypatch, tmp_path)

    note = ("Licensure: Every physician appointed to the NYP Hospital staff, except "
            "interns, and aliens in the US via non-immigrant visas, must have a New "
            "York State license or a temporary certificate in lieu of the license.")
    doc = Document()
    for text in ["Jane Doe", "LICENSURE", note, "Fictional State Medical License 12345"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "lic.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "lic_h.json", "LIC1",
        hierarchy_with_indices=[{"text": "Licensure", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[{"hierarchy": ["Licensure"], "element_idx_start": 1, "element_idx_end": 3, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Licensure": [
            {"element_idx_start": 2, "element_idx_end": 2, "element_type": "paragraph", "confidence": 0.85},
            {"element_idx_start": 3, "element_idx_end": 3, "element_type": "paragraph", "confidence": 0.85},
        ],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))

    texts = [e["text"] for e in output_data["entries"]]
    assert note not in texts
    assert "Fictional State Medical License 12345" in texts


def test_run_stage_2_cancel_check_raises_mid_run_and_no_output_is_written(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    for text in ["Jane Researcher", "EDUCATION", "PhD content", "AWARDS", "", "Best Paper Award, 2010"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "cancel.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "cancel_h.json", "CANCEL1",
        hierarchy_with_indices=[
            {"text": "Education", "level": "H1", "element_idx": 1, "children": []},
            {"text": "Awards", "level": "H1", "element_idx": 3, "children": []},
        ],
        section_boundaries=[
            {"hierarchy": ["Education"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False},
            {"hierarchy": ["Awards"], "element_idx_start": 3, "element_idx_end": 5, "has_children": False},
        ],
    )
    calls = []
    _route_call_llm(monkeypatch, {}, calls)

    class _Cancelled(Exception):
        pass

    state = {"n": 0}

    def cancel_check():
        state["n"] += 1
        if state["n"] == 2:  # allow the first section (Personal Data) through
            raise _Cancelled("stop")

    # #881 made run_stage_2's default `workers` > 1 (sections run on a pool);
    # this test pins the exact pre-#881 serial call count, so it needs
    # workers=1 explicitly -- the concurrent cancel behavior itself is
    # covered by test_cancel_raised_in_one_section_starts_no_further_llm_call
    # below.
    with pytest.raises(_Cancelled):
        stage2.run_stage_2(str(docx_path), str(hpath), cancel_check=cancel_check, workers=1)

    assert state["n"] == 2
    assert len(calls) == 1  # only Personal Data's section actually ran
    om = _RealOutputManager(str(docx_path), base_output_dir=tmp_path)
    assert not om.get_stage2_path().exists()


def test_run_stage_2_parent_header_and_dedup_of_duplicate_header_record(tmp_path, monkeypatch):
    # A has_children parent whose own header sits before its subsection:
    # the header gets emitted once via the parent_header_indices pre-loop
    # AND once via the enclosing "Personal Data" preamble's own per-section
    # header loop (both cover element_idx_start=0) -- filter_extraction_noise
    # must collapse the exact duplicate down to a single header entry.
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    for text in ["PUBLICATIONS", "Intro paragraph before Books subsection", "BOOKS", "Book Title Two Thousand Twenty"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "nested.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "nested_h.json", "NESTED2",
        hierarchy_with_indices=[{
            "text": "Publications", "level": "H1", "element_idx": 0,
            "children": [{"text": "Books", "level": "H2", "element_idx": 2, "children": []}],
        }],
        section_boundaries=[
            {"hierarchy": ["Publications"], "element_idx_start": 0, "element_idx_end": 3, "has_children": True},
            {"hierarchy": ["Publications", "Books"], "element_idx_start": 2, "element_idx_end": 3, "has_children": False},
        ],
    )
    calls = []
    _route_call_llm(monkeypatch, {
        "Publications > Books": [{"element_idx_start": 3, "element_idx_end": 3, "element_type": "paragraph", "confidence": 0.85}],
        "Publications": [{"element_idx_start": 1, "element_idx_end": 1, "element_type": "paragraph", "confidence": 0.9}],
    }, calls)

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))

    assert len(calls) == 3  # Personal Data (no content claimed there), Publications, Publications > Books
    entries = output_data["entries"]
    assert len(entries) == 4
    headers_at_zero = [e for e in entries if e["element_idx_start"] == 0]
    assert len(headers_at_zero) == 1
    assert headers_at_zero[0]["element_type"] == "header"
    assert headers_at_zero[0]["text"] == "PUBLICATIONS"
    assert headers_at_zero[0]["hierarchy"] == ["Publications"]

    assert [(e["element_idx_start"], e["element_type"], e["text"], e["hierarchy"]) for e in entries] == [
        (0, "header", "PUBLICATIONS", ["Publications"]),
        (1, "paragraph", "Intro paragraph before Books subsection", ["Publications"]),
        (2, "header", "BOOKS", ["Publications", "Books"]),
        (3, "paragraph", "Book Title Two Thousand Twenty", ["Publications", "Books"]),
    ]
    assert output_data["coverage"]["coverage_percentage"] == 100.0


def _build_table_gap_fixture(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict:
    """Shared fixture for the two table/doc_length coverage tests below: a
    table sandwiched BETWEEN two paragraphs (the realistic shape, not last
    in the doc), one row LLM-claimed and its sibling recovered by the #420
    backstop. With the table mid-document its own unified index (2) falls
    inside range(doc_length), which is what lets the parent-index-tracking
    code (and the coverage_percentage bug below) actually be exercised --
    with the table last, index 2 sits outside range(doc_length) and every
    assertion here would be trivially satisfied regardless."""
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("AWARDS")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "2020"
    table.cell(0, 1).text = "Best Paper"
    table.cell(1, 0).text = "2021"
    table.cell(1, 1).text = "Rising Star"
    doc.add_paragraph("Committee reviewer, NIH study section, 2022")
    docx_path = tmp_path / "tbl.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "tbl_h.json", "TBL1",
        hierarchy_with_indices=[{"text": "Awards", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[{"hierarchy": ["Awards"], "element_idx_start": 1, "element_idx_end": 3, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Awards": [
            {"element_idx_start": "2.1", "element_idx_end": "2.1", "element_type": "table_row", "confidence": 0.8},
            {"element_idx_start": 3, "element_idx_end": 3, "element_type": "paragraph", "confidence": 0.9},
        ],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))
    return output_data


def test_run_stage_2_sorts_table_row_string_indices_and_counts_table_element_in_doc_length(tmp_path, monkeypatch):
    output_data = _build_table_gap_fixture(tmp_path, monkeypatch)

    # A whole table occupies exactly ONE unified index. Since #870,
    # document_length is the unified element count (paragraphs + tables),
    # not num_paragraphs + num_empty, so it includes the table's own index
    # (0, 1, 2, 3 -- four elements) even though the table's rows are
    # addressed by string sub-indices within `entries`.
    assert output_data["document_length"] == 4

    entries = output_data["entries"]
    assert [e["element_idx_start"] for e in entries] == [0, 1, "2.0", "2.1", 3]
    row0, row1 = entries[2], entries[3]
    assert row0["text"] == "2020 | Best Paper" and row0["recovered_row"] is True
    assert row1["text"] == "2021 | Rising Star" and row1.get("recovered_row") is not True
    assert row1["confidence"] == 0.8
    assert entries[4]["text"] == "Committee reviewer, NIH study section, 2022"

    # The table's own unified index (2) must be tracked as assigned via its
    # claimed/recovered rows' parent link -- not left to surface as a
    # spurious "break"/unaccounted gap between the header and the trailing
    # paragraph.
    coverage = output_data["coverage"]
    assert coverage["unaccounted_indices"] == []
    assert not any(e["element_type"] == "break" for e in entries)


def test_run_stage_2_coverage_percentage_never_exceeds_100_percent(tmp_path, monkeypatch):
    # Same fixture as the sort-order test above. Every doc index is
    # accounted for there (coverage.unaccounted_indices == []), so a correct
    # coverage metric must read exactly 100% -- never a figure inflated by
    # triple-counting one table row across its string sub-index, its
    # sibling's string sub-index, and their shared artificial int parent
    # marker on top of the real integer positions.
    output_data = _build_table_gap_fixture(tmp_path, monkeypatch)
    coverage = output_data["coverage"]
    assert coverage["coverage_percentage"] <= 100.0
    assert coverage["coverage_percentage"] == pytest.approx(100.0)
    assert coverage["assigned_indices"] <= output_data["document_length"]


def test_run_stage_2_coverage_percentage_bounded_when_table_is_last_section(tmp_path, monkeypatch):
    # Distinct shape from the fixture above: with NO paragraph after the
    # table. Before #870, the table's own unified index (2) sat OUTSIDE
    # range(document_length) == {0, 1} (document_length counted only the 2
    # real paragraphs), so an unbounded integer_assigned would have read
    # len({0, 1, 2}) / 2 == 150% -- only intersecting with
    # range(document_length) (_bound_coverage_indices) kept that at 100%.
    # Since #870, document_length is the unified element count (3: two
    # paragraphs + the table), so index 2 is naturally inside the range and
    # this configuration can no longer overflow by construction -- this is
    # now a plain regression guard that coverage still reads 100% here.
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("AWARDS")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "2020"
    table.cell(0, 1).text = "Best Paper"
    table.cell(1, 0).text = "2021"
    table.cell(1, 1).text = "Rising Star"
    docx_path = tmp_path / "tbl_last.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "tbl_last_h.json", "TBL2",
        hierarchy_with_indices=[{"text": "Awards", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[{"hierarchy": ["Awards"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Awards": [
            {"element_idx_start": "2.1", "element_idx_end": "2.1", "element_type": "table_row", "confidence": 0.8},
        ],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))

    assert output_data["document_length"] == 3
    coverage = output_data["coverage"]
    assert coverage["unaccounted_indices"] == []
    assert coverage["assigned_indices"] <= output_data["document_length"]
    assert coverage["coverage_percentage"] <= 100.0
    assert coverage["coverage_percentage"] == pytest.approx(100.0)


def test_run_stage_2_no_personal_data_preamble_when_first_element_is_mapped_header(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("PhD content here")
    docx_path = tmp_path / "nopre.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "nopre_h.json", "NOPRE1",
        hierarchy_with_indices=[{"text": "Education", "level": "H1", "element_idx": 0, "children": []}],
        section_boundaries=[{"hierarchy": ["Education"], "element_idx_start": 0, "element_idx_end": 1, "has_children": False}],
    )
    calls = []
    _route_call_llm(monkeypatch, {
        "Education": [{"element_idx_start": 1, "element_idx_end": 1, "element_type": "paragraph", "confidence": 0.9}],
    }, calls)

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))

    assert len(calls) == 1  # a single section call -- no separate Personal Data section
    assert not any(e["hierarchy"] == ["Personal Data"] for e in output_data["entries"])
    assert [e["element_type"] for e in output_data["entries"]] == ["header", "paragraph"]


def test_run_stage_2_reports_unaccounted_indices_for_gap_between_sibling_sections(tmp_path, monkeypatch):
    # extract_leaf_sections_with_boundaries only fills the gap between a
    # PARENT header and its first child -- a gap between two unrelated
    # top-level siblings is claimed by neither section's range and is never
    # visited by any header/content/break assignment loop at all.
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    for text in ["Jane Doe", "SECTION A", "A content", "orphan line 1", "orphan line 2", "SECTION B", "B content"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "gap.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "gap_h.json", "GAP1",
        hierarchy_with_indices=[
            {"text": "Section A", "level": "H1", "element_idx": 1, "children": []},
            {"text": "Section B", "level": "H1", "element_idx": 5, "children": []},
        ],
        section_boundaries=[
            {"hierarchy": ["Section A"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False},
            {"hierarchy": ["Section B"], "element_idx_start": 5, "element_idx_end": 6, "has_children": False},
        ],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Section A": [{"element_idx_start": 2, "element_idx_end": 2, "element_type": "paragraph", "confidence": 0.9}],
        "Section B": [{"element_idx_start": 6, "element_idx_end": 6, "element_type": "paragraph", "confidence": 0.9}],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))

    assert output_data["coverage"]["unaccounted_indices"] == [3, 4]
    assert output_data["coverage"]["coverage_percentage"] == pytest.approx(71.4)
    assert not any(3 <= e["element_idx_start"] <= 4 for e in output_data["entries"] if isinstance(e["element_idx_start"], int))


def test_run_stage_2_autodetects_hierarchy_json_when_not_given(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("PhD content here")
    docx_path = tmp_path / "auto.docx"
    doc.save(docx_path)

    # Pre-place the hierarchy JSON exactly where om.get_stage1b_path() looks.
    om = _RealOutputManager(str(docx_path), base_output_dir=tmp_path)
    hpath = om.get_stage1b_path()
    hpath.parent.mkdir(parents=True, exist_ok=True)
    hpath.write_text(json.dumps({
        "document_uid": "AUTO1",
        "hierarchy_with_indices": [{"text": "Education", "level": "H1", "element_idx": 0, "children": []}],
        "section_boundaries": [{"hierarchy": ["Education"], "element_idx_start": 0, "element_idx_end": 1, "has_children": False}],
    }))
    _route_call_llm(monkeypatch, {
        "Education": [{"element_idx_start": 1, "element_idx_end": 1, "element_type": "paragraph", "confidence": 0.9}],
    })

    output_data, _ = stage2.run_stage_2(str(docx_path))

    assert output_data["document_uid"] == "AUTO1"


def test_run_stage_2_doc_length_is_unified_element_count_and_flags_trailing_paragraph(tmp_path, monkeypatch):
    # #870: doc_length used to be num_paragraphs + num_empty (2 real
    # paragraphs + the trailing one = 3 here), one less than the 4-element
    # unified stream (paragraph, header, table, trailing paragraph) because
    # that count ignores the table. A trailing paragraph placed AFTER the
    # last section boundary (element_idx_end=2, stopping at the table) is
    # never visited by any header/content/break assignment loop, so it is
    # never in all_assigned_indices either way -- but under the old
    # doc_length (3) it also fell OUTSIDE range(doc_length), so
    # _bound_coverage_indices could never flag it: range(3) == {0, 1, 2}
    # already reads "fully covered" without index 3 ever being checked.
    # Under the fixed doc_length (4, the real element count), range(4)
    # includes index 3 and the loss becomes visible in unaccounted_indices.
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("AWARDS")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "2020"
    table.cell(0, 1).text = "Best Paper"
    doc.add_paragraph("Committee reviewer, NIH study section, 2022")
    docx_path = tmp_path / "trailing.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "trailing_h.json", "TRAIL1",
        hierarchy_with_indices=[{"text": "Awards", "level": "H1", "element_idx": 1, "children": []}],
        # Stops at the table (element_idx_end=2) -- the trailing paragraph
        # (index 3) is outside every section's range, mirroring a stage 1b
        # hierarchy that never mapped it (the real #870 shape).
        section_boundaries=[{"hierarchy": ["Awards"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Awards": [],  # table row left unclaimed -- recovered as a break, not a content entry
    })

    output_data, _ = stage2.run_stage_2(str(docx_path), str(hpath))

    assert output_data["document_length"] == 4
    coverage = output_data["coverage"]
    assert coverage["unaccounted_indices"] == [3]
    # The pre-#870 denominator (3) would have bounded range() to {0, 1, 2},
    # masking index 3 entirely rather than flagging it.
    old_doc_length = 3
    assert 3 not in set(range(old_doc_length))
    assert not any(e["element_idx_start"] == 3 for e in output_data["entries"])


# ================================================ run_stage_2: section pool (#881)
#
# run_stage_2 now extracts each section's entries on a thread pool
# (map_in_order, core/batch_pool.py) instead of one after another -- same
# shape as stage 3b's group pool (_classify_group / _group_progress_printer /
# run_stage_3b's map_in_order call). A synthetic doc with several sibling
# top-level sections (plus the "Personal Data" preamble) gives map_in_order
# real width to fan out over.

_MANY_SECTIONS = [f"SECTION {i:02d}" for i in range(8)]


def _build_many_sections_docx(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Jane Researcher")
    for name in _MANY_SECTIONS:
        doc.add_paragraph(name)
        doc.add_paragraph(f"Entry text for {name}")
    docx_path = tmp_path / "many_sections.docx"
    doc.save(docx_path)
    return docx_path


def _build_many_sections_hierarchy(tmp_path: Path, document_uid: str) -> Path:
    hierarchy_with_indices = []
    section_boundaries = []
    for i, name in enumerate(_MANY_SECTIONS):
        header_idx = 1 + 2 * i
        hierarchy_with_indices.append(
            {"text": name, "level": "H1", "element_idx": header_idx, "children": []}
        )
        section_boundaries.append({
            "hierarchy": [name],
            "element_idx_start": header_idx,
            "element_idx_end": header_idx + 1,
            "has_children": False,
        })
    return _write_hierarchy(
        tmp_path, "many_sections_h.json", document_uid, hierarchy_with_indices, section_boundaries
    )


_MANY_SECTION_KEYS = ["Personal Data", *_MANY_SECTIONS]


def _section_key_from_prompt(prompt: str, keys: list[str]) -> str:
    """Same routing rule as _route_call_llm: the EXACT `CV Section Header`
    prompt line, never a raw substring search of the whole prompt."""
    header_line = next(line for line in prompt.splitlines() if "CV Section Header" in line)
    return next(k for k in keys if f"`{k}`" in header_line)


def _many_section_delimiters(key: str) -> list[dict]:
    if key == "Personal Data":
        return [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}]
    i = _MANY_SECTIONS.index(key)
    header_idx = 1 + 2 * i
    return [{
        "element_idx_start": header_idx + 1, "element_idx_end": header_idx + 1,
        "element_type": "paragraph", "confidence": 0.9,
    }]


def test_parallel_sections_write_the_same_artifact_as_the_serial_loop(tmp_path, monkeypatch):
    """Sections finish in reverse order under the pool; the artifact must not."""
    in_flight = {"now": 0, "peak": 0}
    lock = threading.Lock()
    # Distinct per-section cost: belt-and-braces on top of the full-artifact
    # equality below -- total_cost is a float sum, so an accumulation-order
    # mutant (e.g. completion order) COULD shift the last bit or two of the
    # sum, but in practice reversed(results) still passes this assertion
    # (float addition is associative enough here). The real guard against an
    # accumulation-order mutant is test_tied_start_indices_keep_section_
    # submission_order; this assertion only pins that the totals are
    # order-independent in value, not that ordering is enforced.
    costs = {key: round(0.1 * (i + 1), 4) for i, key in enumerate(_MANY_SECTION_KEYS)}

    def slow_early_sections(**kwargs):
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, _MANY_SECTION_KEYS)
        idx = _MANY_SECTION_KEYS.index(key)
        with lock:
            in_flight["now"] += 1
            in_flight["peak"] = max(in_flight["peak"], in_flight["now"])
        # Later sections answer first: SECTION 07 sleeps least.
        time.sleep((len(_MANY_SECTION_KEYS) - idx) * 0.005)
        with lock:
            in_flight["now"] -= 1
        return _llm_result({"delimiters": _many_section_delimiters(key)}, cost=costs[key])

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY1")
    monkeypatch.setattr(stage2, "call_llm", slow_early_sections)

    serial, _ = stage2.run_stage_2(str(docx_path), str(hpath), workers=1)
    assert in_flight["peak"] == 1  # workers=1 really is the serial loop

    parallel, _ = stage2.run_stage_2(str(docx_path), str(hpath), workers=4)
    assert in_flight["peak"] > 1  # and workers=4 really overlapped

    # No generated_at/datetime.now field exists in stage 2's output to strip
    # before comparing -- pin that so a future field addition is noticed here.
    src = inspect.getsource(stage2)
    assert "generated_at" not in src
    assert "datetime.now" not in src

    assert serial == parallel
    # Exact float equality: sum order (not just the final entries list) must
    # match the serial loop too.
    assert serial["total_cost"] == parallel["total_cost"]


def test_section_progress_lines_are_monotonic_and_unspliced(tmp_path, monkeypatch, capsys):
    def slow_early_sections(**kwargs):
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, _MANY_SECTION_KEYS)
        idx = _MANY_SECTION_KEYS.index(key)
        time.sleep((len(_MANY_SECTION_KEYS) - idx) * 0.005)
        return _llm_result({"delimiters": _many_section_delimiters(key)})

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY2")
    monkeypatch.setattr(stage2, "call_llm", slow_early_sections)

    stage2.run_stage_2(str(docx_path), str(hpath), workers=4)

    out = capsys.readouterr().out.splitlines()
    total = len(_MANY_SECTION_KEYS)
    marker = f"/{total}]"
    progress = [line for line in out if line.startswith("[") and marker in line]
    # orchestrator.py's PROGRESS_PATTERNS read "[N/M]"; N must never go backwards.
    assert [int(line[1:line.index("/")]) for line in progress] == list(range(1, total + 1))
    # Each section's block is one atomic print: Processing -> Elements -> summary -> blank.
    for i, line in enumerate(out):
        if line.startswith("[") and marker in line:
            assert out[i + 1].strip().startswith("Elements:"), out[i:i + 4]
            assert out[i + 2].strip().startswith(("✓ Extracted", "- No content")), out[i:i + 4]
            assert out[i + 3] == "", out[i:i + 4]


def test_sections_run_inside_the_callers_run_id_context(tmp_path, monkeypatch):
    from unified_pipeline.core import prompt_logger

    seen = set()

    def record(**kwargs):
        seen.add(prompt_logger._current_run_id.get())
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, _MANY_SECTION_KEYS)
        return _llm_result({"delimiters": _many_section_delimiters(key)})

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY3")
    monkeypatch.setattr(stage2, "call_llm", record)

    token = prompt_logger.set_current_run_id("run-stage2")
    try:
        stage2.run_stage_2(str(docx_path), str(hpath), workers=4)
    finally:
        prompt_logger.reset_current_run_id(token)

    assert seen == {"run-stage2"}


class _FakeRoutedStdout:
    """Minimal stand-in for orchestrator.py's ``_RoutedStdout`` -- a dict
    keyed by ``threading.get_ident()``. A registered thread's write lands in
    its own capture; any other thread's write lands in ``leak`` instead. The
    real ``_RoutedStdout`` would not actually lose an unregistered thread's
    write (since #883 it falls back to the ``_capture_var`` ContextVar that
    ``map_in_order`` copies onto every pool thread) -- this fake omits that
    fallback on purpose, to isolate the property this test actually checks:
    output must come from the calling thread's ``on_result`` callback, not
    from a print inside ``_extract_section`` itself.
    """

    def __init__(self, leak):
        self._leak = leak
        self._captures = {}

    def register(self, ident, capture):
        self._captures[ident] = capture

    def write(self, text):
        return self._captures.get(threading.get_ident(), self._leak).write(text)

    def flush(self):
        pass


def test_section_progress_prints_only_from_the_calling_thread(tmp_path, monkeypatch):
    """LEAD defect shape: a print from inside _extract_section would run on
    map_in_order's pool thread instead of the calling thread's on_result
    callback. That would not literally lose the output -- since #883 the
    real _RoutedStdout falls back to a ContextVar map_in_order copies onto
    each pool thread -- but it would break two other guarantees: the
    "[N/M]" number is assigned by completion order on the calling thread
    (on_result), and each section's block (Processing/Elements/summary/
    blank) must go out as one atomic print so concurrent sections' blocks
    cannot splice together. Install a fake routed stdout that only
    registers the main (calling) thread and run a real thread pool
    (workers=4): every "[N/M]" line must land in the registered capture,
    and none may leak.
    """
    def slow_early_sections(**kwargs):
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, _MANY_SECTION_KEYS)
        idx = _MANY_SECTION_KEYS.index(key)
        time.sleep((len(_MANY_SECTION_KEYS) - idx) * 0.005)
        return _llm_result({"delimiters": _many_section_delimiters(key)})

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY4")
    monkeypatch.setattr(stage2, "call_llm", slow_early_sections)

    registered = io.StringIO()
    leak = io.StringIO()
    fake_stdout = _FakeRoutedStdout(leak)
    fake_stdout.register(threading.get_ident(), registered)

    real_stdout = sys.stdout
    monkeypatch.setattr(sys, "stdout", fake_stdout)
    try:
        stage2.run_stage_2(str(docx_path), str(hpath), workers=4)
    finally:
        monkeypatch.setattr(sys, "stdout", real_stdout)

    total = len(_MANY_SECTION_KEYS)
    marker = f"/{total}]"

    def progress_lines(text):
        return [line for line in text.splitlines() if line.startswith("[") and marker in line]

    registered_progress = progress_lines(registered.getvalue())
    leak_progress = progress_lines(leak.getvalue())
    assert [int(line[1:line.index("/")]) for line in registered_progress] == list(range(1, total + 1))
    assert leak_progress == []
    # Not just the "[N/M]" shape: nothing at all -- a mutant that prints the
    # section body (not just the progress line) from the pool thread must
    # fail this test too, not just the narrower progress-line check above.
    assert leak.getvalue() == ""


def test_workers_one_never_overlaps_and_workers_four_does(tmp_path, monkeypatch):
    in_flight = {"now": 0, "peak": 0}
    lock = threading.Lock()

    def track(**kwargs):
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, _MANY_SECTION_KEYS)
        with lock:
            in_flight["now"] += 1
            in_flight["peak"] = max(in_flight["peak"], in_flight["now"])
        time.sleep(0.01)
        with lock:
            in_flight["now"] -= 1
        return _llm_result({"delimiters": _many_section_delimiters(key)})

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY5")
    monkeypatch.setattr(stage2, "call_llm", track)

    stage2.run_stage_2(str(docx_path), str(hpath), workers=1)
    assert in_flight["peak"] == 1

    in_flight["peak"] = 0
    stage2.run_stage_2(str(docx_path), str(hpath), workers=4)
    assert in_flight["peak"] >= 2


def test_cancel_raised_in_one_section_starts_no_further_llm_call(tmp_path, monkeypatch):
    workers, cancel_after = 4, 3
    checks = 0
    cancelled = threading.Event()
    started_after_cancel = 0
    lock = threading.Lock()

    class _Cancelled(Exception):
        pass

    def cancel_check():
        nonlocal checks
        with lock:
            checks += 1
            if checks > cancel_after:
                cancelled.set()
                raise _Cancelled("stop")

    def llm(**kwargs):
        nonlocal started_after_cancel
        # cancel_check gates every task before its LLM call, so once it has
        # raised no further section reaches call_llm. In-flight ones finish.
        if cancelled.is_set():
            with lock:
                started_after_cancel += 1
        time.sleep(0.01)
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, _MANY_SECTION_KEYS)
        return _llm_result({"delimiters": _many_section_delimiters(key)})

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY6")
    monkeypatch.setattr(stage2, "call_llm", llm)

    with pytest.raises(_Cancelled):
        stage2.run_stage_2(str(docx_path), str(hpath), cancel_check=cancel_check, workers=workers)

    assert started_after_cancel == 0


def test_detect_entries_for_section_raising_propagates_and_writes_no_output(tmp_path, monkeypatch):
    """#915 review item 5: detect_entries_for_section normally swallows an
    LLM/JSON/network error per batch (logged, next batch continues), but a
    bug outside that guard would raise out of it. Pin what run_stage_2 does
    then: the exception propagates (map_in_order's own contract -- first
    raise cancels the queue and re-raises) and, since that happens before
    the json.dump that writes stage 2's output, no artifact exists."""
    target_section = "SECTION 03"

    def flaky_detect(hierarchy_path, *args, **kwargs):
        if hierarchy_path == [target_section]:
            raise RuntimeError("synthetic detect_entries_for_section failure")
        return [], {"cost": 0, "tokens": 0}

    _redirect_output_manager(monkeypatch, tmp_path)
    docx_path = _build_many_sections_docx(tmp_path)
    hpath = _build_many_sections_hierarchy(tmp_path, "MANY7")
    monkeypatch.setattr(stage2, "detect_entries_for_section", flaky_detect)

    with pytest.raises(RuntimeError, match="synthetic detect_entries_for_section failure"):
        stage2.run_stage_2(str(docx_path), str(hpath), workers=4)

    om = _RealOutputManager(str(docx_path), base_output_dir=tmp_path)
    assert not om.get_stage2_path().exists()


def test_stage2_workers_config_knob_is_read_from_env(monkeypatch):
    # STAGE2_SECTION_WORKERS itself is bound once, at import time, so it
    # can't observe an env var set by a test -- this pins the reader it's
    # built from instead: workers_from_config("CVICHE_STAGE2_SECTION_WORKERS").
    monkeypatch.setenv("CVICHE_STAGE2_SECTION_WORKERS", "7")
    assert stage2.workers_from_config("CVICHE_STAGE2_SECTION_WORKERS") == 7


def test_run_stage_2_defaults_to_the_config_knob():
    # Pins the default itself, not just the reader: a hardcoded literal
    # default (e.g. `workers: int = 1`) would pass every other test here
    # (they all pass workers= explicitly) while silently dropping the
    # CVICHE_STAGE2_SECTION_WORKERS knob in production.
    default = inspect.signature(stage2.run_stage_2).parameters["workers"].default
    assert default is stage2.STAGE2_SECTION_WORKERS
    assert default == stage2.workers_from_config("CVICHE_STAGE2_SECTION_WORKERS")


def test_assigned_indices_union_matches_the_serial_loop(tmp_path, monkeypatch):
    # Table row fixture (parent-table int path): a claimed "N.M" string
    # sub-index plus its recovered sibling, so section_assigned's string +
    # parent-int bookkeeping is exercised, not just plain integer ranges.
    _redirect_output_manager(monkeypatch, tmp_path)

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("AWARDS")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "2020"
    table.cell(0, 1).text = "Best Paper"
    table.cell(1, 0).text = "2021"
    table.cell(1, 1).text = "Rising Star"
    doc.add_paragraph("Committee reviewer, NIH study section, 2022")
    docx_path = tmp_path / "tbl_union.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "tbl_union_h.json", "TBLUNION",
        hierarchy_with_indices=[{"text": "Awards", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[{"hierarchy": ["Awards"], "element_idx_start": 1, "element_idx_end": 3, "has_children": False}],
    )
    _route_call_llm(monkeypatch, {
        "Personal Data": [{"element_idx_start": 0, "element_idx_end": 0, "element_type": "paragraph", "confidence": 0.9}],
        "Awards": [
            {"element_idx_start": "2.1", "element_idx_end": "2.1", "element_type": "table_row", "confidence": 0.8},
            {"element_idx_start": 3, "element_idx_end": 3, "element_type": "paragraph", "confidence": 0.9},
        ],
    })

    serial, _ = stage2.run_stage_2(str(docx_path), str(hpath), workers=1)
    parallel, _ = stage2.run_stage_2(str(docx_path), str(hpath), workers=4)

    assert serial["coverage"] == parallel["coverage"]
    assert serial["coverage"]["unaccounted_indices"] == []
    assert serial["coverage"]["coverage_percentage"] == pytest.approx(100.0)


def test_string_row_index_never_lands_in_assigned_only_its_parent_int_does(monkeypatch):
    """#915 review item 1: ``section_assigned.add(str(start_idx_entry))``
    was inert -- ``section_assigned`` is later diffed against
    ``section_range = set(range(...))``, which holds only ints, and
    ``_bound_coverage_indices`` (#856) filters ``isinstance(idx, int)`` by
    construction too, so the string form is never read by any consumer.
    Pin the INVARIANT the string-add was removed for rather than the
    deleted line itself: a section whose sole entry carries a "N.M"
    table-row string index still ends up with an int-only ``assigned`` set
    (never the "1.0" string) and the parent table index alone still keeps
    that element out of ``unassigned_in_section`` / the break entries.
    """
    doc_elements = [
        _para(0, "AWARDS"),
        {"unified_idx": 1, "type": "table_content", "data": [["2020", "Best Paper"]]},
    ]
    element_index_map = _idx_map(doc_elements)
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result({"delimiters": [
        {"element_idx_start": "1.0", "element_idx_end": "1.0", "element_type": "table_row", "confidence": 0.8},
    ]}))

    context = stage2._DocumentContext(
        doc_elements=doc_elements,
        header_indices={0},
        element_index_map=element_index_map,
        header_info={0: ["Awards"]},
        doc=Document(),
    )
    result = stage2._extract_section(
        context, ["Awards"], 0, 1,
        document_uid="TEST",
        cancel_check=None,
    )

    assert result.assigned == {0, 1}  # header idx 0 + parent table idx 1, ints only
    assert "1.0" not in result.assigned
    break_starts = {e["element_idx_start"] for e in result.entries if e["element_type"] == "break"}
    assert 1 not in break_starts  # the parent int alone kept it out of break_entries


def test_extract_section_signature_bundles_the_five_constants_into_one_context_param():
    """#915 review item 6: doc_elements, header_indices, element_index_map,
    header_info, doc are constant across every call in a run_stage_2
    invocation; pin the bundled signature shape (context first and
    positional so map_in_order's per-call (hierarchy_path, start_idx,
    end_idx) tuple still binds correctly via partial) rather than just
    trusting the corpus A/B to catch a drift back to five loose params."""
    params = list(inspect.signature(stage2._extract_section).parameters.values())
    names = [p.name for p in params]
    assert names == ["context", "hierarchy_path", "start_idx", "end_idx", "document_uid", "cancel_check"]
    by_name = {p.name: p for p in params}
    assert by_name["context"].kind == inspect.Parameter.POSITIONAL_OR_KEYWORD
    assert by_name["hierarchy_path"].kind == inspect.Parameter.POSITIONAL_OR_KEYWORD
    for kw_only in ("document_uid", "cancel_check"):
        assert by_name[kw_only].kind == inspect.Parameter.KEYWORD_ONLY
    assert stage2._DocumentContext._fields == (
        "doc_elements", "header_indices", "element_index_map", "header_info", "doc",
    )


def test_tied_start_indices_keep_section_submission_order(tmp_path, monkeypatch):
    """A stage-1b leaf section can legitimately enclose another leaf
    section's range (corpus A/B, 21/107 real CVs: a wide catch-all leaf
    alongside a narrower nested leaf sharing an element index). Both
    sections then independently call the LLM over the shared index and can
    each independently claim it as a same-type content entry -- the text
    comes from the shared element_index_map, so it is identical either way
    -- and filter_extraction_noise's (element_type, text, idx) dedup then
    keeps only the FIRST of the two survivors in `all_entries`. Because the
    later `all_entries.sort(key=sort_key)` is stable, "first" means
    accumulation order, not LLM-completion order -- this pins that
    accumulation stays SUBMISSION order (the earlier `reversed(results)`
    mutant flips which section's hierarchy the tied entry carries) even
    when the later-submitted section's call finishes first.
    """
    doc = Document()
    for text in ["Jane Researcher", "CURRENT POSITION", "Professor of Chemistry"]:
        doc.add_paragraph(text)
    docx_path = tmp_path / "tie.docx"
    doc.save(docx_path)

    hpath = _write_hierarchy(
        tmp_path, "tie_h.json", "TIE1",
        hierarchy_with_indices=[{"text": "Current Position", "level": "H1", "element_idx": 1, "children": []}],
        section_boundaries=[
            # Wide leaf enclosing the narrower leaf below -- element 2 is in
            # BOTH ranges. Order here fixes submission order: Personal Data
            # is section index 0, Current Position is section index 1.
            {"hierarchy": ["Personal Data"], "element_idx_start": 0, "element_idx_end": 2, "has_children": False},
            {"hierarchy": ["Current Position"], "element_idx_start": 1, "element_idx_end": 2, "has_children": False},
        ],
    )

    submission_order = {"Personal Data": 0, "Current Position": 1}

    def slow_first_submitted_finishes_last(**kwargs):
        prompt = kwargs["messages"][1]["content"]
        key = _section_key_from_prompt(prompt, list(submission_order))
        # Personal Data (submitted first) sleeps longest, so under
        # workers=4 Current Position (submitted second) completes first --
        # completion order is the exact reverse of submission order.
        time.sleep((len(submission_order) - submission_order[key]) * 0.02)
        return _llm_result({"delimiters": [
            {"element_idx_start": 2, "element_idx_end": 2, "element_type": "paragraph", "confidence": 0.9},
        ]})

    _redirect_output_manager(monkeypatch, tmp_path)
    monkeypatch.setattr(stage2, "call_llm", slow_first_submitted_finishes_last)

    serial, _ = stage2.run_stage_2(str(docx_path), str(hpath), workers=1)
    parallel, _ = stage2.run_stage_2(str(docx_path), str(hpath), workers=4)

    assert serial == parallel

    tied = [e for e in parallel["entries"] if e["element_idx_start"] == 2]
    assert len(tied) == 1  # filter_extraction_noise collapsed the duplicate claim
    assert tied[0]["hierarchy"] == ["Personal Data"]  # first-SUBMITTED section wins


def test_batch_progress_goes_to_the_logger_not_stdout(monkeypatch, caplog, capsys):
    # 120 elements / BATCH_SIZE=50 -> 3 batches, so "Processing batch" fires.
    elements = [_para(i, f"Line {i}") for i in range(120)]
    monkeypatch.setattr(stage2, "call_llm", lambda **kw: _llm_result({"delimiters": []}))

    with caplog.at_level(logging.INFO, logger=stage2.logger.name):
        stage2.detect_entries_for_section(
            ["Big Section"], elements, 0, 120,
            document_uid="BATCHLOG", header_indices=set(), element_index_map=_idx_map(elements),
        )

    messages = [r.message for r in caplog.records]
    assert any("Processing batch 1/3" in m for m in messages)
    assert any("Processing batch 2/3" in m for m in messages)
    assert any("Processing batch 3/3" in m for m in messages)

    out = capsys.readouterr().out
    assert "Processing batch" not in out


# =============================================================== detail-line fold (#986)

def _detail_fold_docx(tmp_path):
    doc = Document()
    for text in ["MENTORING", "2016 Ana Cruz, Graduate Student", "Project: Study of tides",
                 "2015 Lee Park, Undergraduate", "Project: Study of kelp", "",
                 "ADVISING", "2014 Kim Ortiz, Fellow", "Role: Primary advisor"]:
        doc.add_paragraph(text)
    path = tmp_path / "fold.docx"
    doc.save(path)
    hpath = _write_hierarchy(
        tmp_path, "fold_h.json", "FOLDUID",
        hierarchy_with_indices=[
            {"text": "Mentoring", "level": "H1", "element_idx": 0, "children": []},
            {"text": "Advising", "level": "H1", "element_idx": 6, "children": []},
        ],
        section_boundaries=[
            {"hierarchy": ["Mentoring"], "element_idx_start": 0, "element_idx_end": 5, "has_children": False},
            {"hierarchy": ["Advising"], "element_idx_start": 6, "element_idx_end": 8, "has_children": False},
        ],
    )
    return path, hpath


def test_run_stage_2_folds_detail_lines_whether_detector_missed_or_split_them(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)
    path, hpath = _detail_fold_docx(tmp_path)
    # Mentoring: the detector returns nothing, so every line is an unassigned
    # "break". Advising: it splits the Role line off as its own entry.
    _route_call_llm(monkeypatch, {
        "Advising": [
            {"element_idx_start": 7, "element_idx_end": 7, "element_type": "paragraph", "confidence": 0.9},
            {"element_idx_start": 8, "element_idx_end": 8, "element_type": "paragraph", "confidence": 0.9},
        ],
    })

    output_data, _ = stage2.run_stage_2(str(path), str(hpath))

    rows = [(e["element_idx_start"], e["element_idx_end"], e["element_type"], e["text"])
            for e in output_data["entries"] if e["element_type"] != "header"]
    assert rows == [
        (1, 2, "break", "2016 Ana Cruz, Graduate Student\tProject: Study of tides"),
        (3, 4, "break", "2015 Lee Park, Undergraduate\tProject: Study of kelp"),
        (5, 5, "break", ""),
        (7, 8, "paragraph", "2014 Kim Ortiz, Fellow\tRole: Primary advisor"),
    ]
    assert output_data["coverage"]["coverage_percentage"] == 100.0
    assert output_data["coverage"]["unaccounted_indices"] == []
    assert output_data["coverage"]["content_entries"] == 1  # only the paragraph; folded breaks are 3b-visible breaks


def test_run_stage_2_fold_gives_the_same_artifact_serial_and_parallel(tmp_path, monkeypatch):
    _redirect_output_manager(monkeypatch, tmp_path)
    path, hpath = _detail_fold_docx(tmp_path)
    _route_call_llm(monkeypatch, {})
    serial, _ = stage2.run_stage_2(str(path), str(hpath), workers=1)
    parallel, _ = stage2.run_stage_2(str(path), str(hpath), workers=4)
    assert serial["entries"] == parallel["entries"]
