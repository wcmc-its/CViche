"""Regression guard for issue #659 (five of ten sites): entry.get('extracted_fields', {})
crashes on an explicit None -- plus the behaviour each touched section
writer is expected to produce once it no longer crashes (#739 review).

`dict.get(key, default)` only substitutes `default` when the key is absent --
an entry that carries `extracted_fields: None` explicitly still gets `None`
back, and the very next `.get()` call on it raises `AttributeError: 'NoneType'
object has no attribute 'get'`. `licensure.py` and `postdoc_training.py`
already use the defended form `entry.get('extracted_fields') or {}`; this
fixes the same idiom in mentoring.py (the N3B ongoing-mentorship rule and
the mentee normalizer, which the two table-fill loops now go through),
other_education.py and patents.py.

research_support.py (4 sites) and formatting/values.py (1 site) are the
remaining five of the ten sites #659 names; they land in sibling PRs (#659
closes only once all three land -- see PR body).

Every test drives the real `_fill_*` section-writer entrypoint (or the
module-level classifier it delegates to) on a minimal document and asserts
where the content landed, not merely that the call returned. Fixtures are
fictional.

    python3 -m pytest src/unified_pipeline/tests/test_stage6_extracted_fields_none_safety.py -p no:cacheprovider
"""

import sys
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from unified_pipeline.stage6.sections.mentoring import (  # noqa: E402
    _normalize_mentee,
    _partition_mentoring_entries,
)
from unified_pipeline.stage6.sections.other_education import (  # noqa: E402
    _normalize_other_education_entry,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


def _new_generator() -> WCMTemplateGenerator:
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    return gen


def _body_after(doc, heading_text: str, count: int) -> list[tuple[str, str]]:
    """The `count` body elements following the paragraph whose stripped text
    equals `heading_text`, as ('p', text) or ('tbl', text of row 0 cell 1) --
    the cell that carries a mentee's name in a mentee table."""
    body = list(doc.element.body)
    for para in doc.paragraphs:
        if para.text.strip().lower() == heading_text.lower():
            start = body.index(para._element)
            break
    else:
        raise AssertionError(f"no paragraph {heading_text!r} in document")
    shape = []
    for element in body[start + 1:start + 1 + count]:
        if element.tag == qn('w:sectPr'):
            break
        if element.tag == qn('w:tbl'):
            table = Table(element, doc)
            shape.append(('tbl', table.rows[0].cells[1].text))
        else:
            shape.append(('p', Paragraph(element, doc).text))
    return shape


def _mentoring_doc() -> WCMTemplateGenerator:
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    gen.doc.add_paragraph("Current Mentees:")
    gen.doc.add_paragraph("Past Mentees:")
    return gen


def _mentee(code: str, name: str, **fields) -> dict:
    return {'taxonomy_code': code, 'text': f"{name} mentee entry",
            'extracted_fields': {'mentee_name': name, **fields}}


#: The year the partition judges an N3A period against
#: (`_n3a_entry_has_ended`), fixed so no test depends on the clock.
_CURRENT_YEAR = 2026


# --- mentoring: extracted_fields=None at every #659 site ------------------------

@pytest.mark.parametrize("code, heading", [("N3A", "Current Mentees:"),
                                           ("N3B", "Past Mentees:")])
def test_mentoring_none_extracted_fields_renders_as_a_summary_line(code, heading):
    """An N3A or N3B entry with extracted_fields=None used to crash -- N3B in
    the ongoing-mentorship rule (mentoring.py:90 pre-fix), both codes in the
    table-fill loops (:161/:182 pre-fix). With no fields it names no mentee,
    so it is an aggregate summary: rendered as a plain line directly under
    its heading, with its text intact."""
    gen = _mentoring_doc()
    entry = {'taxonomy_code': code, 'text': 'Ph.D. Graduated: 38',
             'extracted_fields': None}

    gen._fill_mentoring({code: [entry]})

    assert _body_after(gen.doc, heading, 1) == [('p', 'Ph.D. Graduated: 38')]
    assert len(gen.doc.tables) == 0
    assert gen.stats['entries_inserted'] == 1


def test_normalize_mentee_tolerates_none_extracted_fields():
    """The table-fill loops read fields through `_normalize_mentee`; an
    explicit None yields an empty record (no name), never an AttributeError."""
    record = _normalize_mentee({'taxonomy_code': 'N3A', 'extracted_fields': None}, ongoing=True)
    assert record.name == ''
    assert record.site_position == ''
    assert record.mentoring_period == ''


def test_partition_tolerates_none_extracted_fields_in_ongoing_rule():
    """The N3B ongoing-mentorship rule is the first #659 site reached; a None
    entry is neither moved nor dropped -- it partitions as a past summary."""
    entry = {'taxonomy_code': 'N3B', 'text': 'Completed: 27', 'extracted_fields': None}
    partition = _partition_mentoring_entries({'N3B': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.past_summaries == (entry,)
    assert partition.moved_to_current == 0
    assert partition.current == () and partition.past == ()


# --- N1/N2: extracted_fields=None (#529) -----------------------------------------

def _n1_n2_doc() -> WCMTemplateGenerator:
    gen = _new_generator()
    gen.doc.add_paragraph("MENTORING")
    gen.doc.add_paragraph(
        "Leadership and mentoring in programs (Describe activity; include dates)")
    gen.doc.add_paragraph(
        "Institutional Training Grants and Mentored Trainee Grants")
    return gen


def test_n1_none_extracted_fields_renders_as_its_raw_text():
    """`_fill_mentoring` reads `entry.get('extracted_fields') or {}`, the
    #659-defended idiom, before handing fields to `_program_leadership_line`;
    an explicit None degrades to the entry's text, same as an entry with no
    fields at all -- never an AttributeError."""
    gen = _n1_n2_doc()
    entry = {'taxonomy_code': 'N1', 'text': 'Directed a mentoring initiative.',
             'extracted_fields': None}

    gen._fill_mentoring({'N1': [entry]})

    assert _body_after(gen.doc, "Leadership and mentoring in programs "
                       "(Describe activity; include dates)", 1) == [
        ('p', 'Directed a mentoring initiative.')]
    assert gen.stats['entries_inserted'] == 1


def test_n2_none_extracted_fields_renders_as_its_raw_text():
    """Same idiom for N2: a None `extracted_fields` reads as sparse (no
    title/agency/grant_number), so it degrades to a plain line, not a
    crash and not an empty table."""
    gen = _n1_n2_doc()
    entry = {'taxonomy_code': 'N2', 'text': 'A training grant with no fields.',
             'extracted_fields': None}

    gen._fill_mentoring({'N2': [entry]})

    assert _body_after(gen.doc, "Institutional Training Grants and "
                       "Mentored Trainee Grants", 1) == [
        ('p', 'A training grant with no fields.')]
    assert len(gen.doc.tables) == 0
    assert gen.stats['entries_inserted'] == 1


# --- mentoring: Past -> Current migration ---------------------------------------

@pytest.mark.parametrize("fields", [
    {'start_date': '2019', 'end_date': 'present'},
    {'start_date': '2019', 'end_date': 'Present'},
    {'start_date': '2019', 'end_date': '2019-present'},
    {'start_date': '2019', 'end_date': 'ongoing'},
    {'start_date': '2019', 'end_date': 'current'},
    {'start_date': '2019', 'end_date': 'now'},
], ids=lambda f: repr(f.get('end_date', '<absent>')))
def test_partition_moves_ongoing_past_mentee_to_current(fields):
    entry = _mentee('N3B', 'Ada Lovelace', **fields)
    partition = _partition_mentoring_entries({'N3B': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.current == (entry,)
    assert partition.past == ()
    assert partition.moved_to_current == 1


@pytest.mark.parametrize("text", [
    "Ada Lovelace, PhD student, 2019-",
    "Ada Lovelace (2019- ) PhD student",
    "Ada Lovelace, PhD student, 2019 - present",
    "Ada Lovelace, PhD student, 2019 \u2013 Current",
    "Ada Lovelace, PhD student, 2019 to date",
    "Ada Lovelace, PhD student, 09/2019-Present",
])
@pytest.mark.parametrize("end_date", ['', None, '<absent>'])
def test_partition_moves_start_only_mentee_when_source_leaves_range_open(text, end_date):
    """A start date with no end date is current only when the entry's own
    text writes the range as open (class 1, 2026-10-02 s7ab autopsy)."""
    fields = {'start_date': '2019'}
    if end_date != '<absent>':
        fields['end_date'] = end_date
    entry = _mentee('N3B', 'Ada Lovelace', **fields)
    entry['text'] = text
    partition = _partition_mentoring_entries({'N3B': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.current == (entry,)
    assert partition.past == ()
    assert partition.moved_to_current == 1


@pytest.mark.parametrize("text", [
    "Ada Lovelace, PhD 2019",
    "2019 Ada Lovelace, MD thesis; now Assistant Professor",
    "Ada Lovelace, Summer student 2019. Current position: Resident",
    "Ada Lovelace, 2019 - Excellence Award winner",
    "Ada Lovelace, MD 2019; Fellow - now Assistant Professor",
    "Ada Lovelace, class of 2019, present address withheld",
])
def test_partition_keeps_start_only_mentee_with_a_lone_year_in_past(text):
    """A lone completion, visit or class year stored as start_date stays a
    past mentee. "now"/"Current"/"present" that describe the mentee, not the
    range, and a dash used as a column separator, do not make it ongoing."""
    entry = _mentee('N3B', 'Ada Lovelace', start_date='2019', end_date=None)
    entry['text'] = text
    partition = _partition_mentoring_entries({'N3B': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.past == (entry,)
    assert partition.current == ()
    assert partition.moved_to_current == 0


@pytest.mark.parametrize("fields", [
    {'start_date': '2015', 'end_date': '2019'},
    {'end_date': '2019'},
    {'start_date': '2019'},                    # lone year, text not open
    {},                                        # no dates at all: not ongoing
], ids=lambda f: repr(f))
def test_partition_keeps_ended_or_undated_past_mentee_in_past(fields):
    entry = _mentee('N3B', 'Ada Lovelace', **fields)
    partition = _partition_mentoring_entries({'N3B': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.past == (entry,)
    assert partition.current == ()
    assert partition.moved_to_current == 0


def test_ongoing_past_mentee_renders_under_current_mentees():
    """The migration is visible in the document: the N3B entry's table sits
    under "Current Mentees:" and nothing sits under "Past Mentees:"."""
    gen = _mentoring_doc()
    entries_by_code = {
        'N3B': [_mentee('N3B', 'Ada Lovelace', start_date='2019', end_date='present'),
                _mentee('N3B', 'Grace Hopper', start_date='2010', end_date='2014')],
    }

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, "Current Mentees:", 2) == [
        ('tbl', 'Ada Lovelace'), ('p', '')]
    assert _body_after(gen.doc, "Past Mentees:", 2) == [
        ('tbl', 'Grace Hopper'), ('p', '')]


def _mentoring_period_cell(doc, name: str) -> str:
    """The Mentoring Period value of the mentee table headed `name`."""
    for table in doc.tables:
        if table.rows[0].cells[1].text == name:
            return table.rows[2].cells[1].text
    raise AssertionError(f"no mentee table for {name!r}")


def test_start_only_past_mentee_renders_bare_year_under_past_mentees():
    """The class-1 render: a past mentee with one stored year (month added
    by stage 4) renders that year under "Past Mentees:", not "-present"
    under "Current Mentees:". A current mentee keeps "-present"."""
    gen = _mentoring_doc()
    past = _mentee('N3B', 'Ada Lovelace', start_date='2007-05', end_date=None)
    past['text'] = "Ada Lovelace, MS 2007"
    current = _mentee('N3A', 'Grace Hopper', start_date='2021', end_date=None)
    current['text'] = "Grace Hopper, PhD student, 2021-"

    gen._fill_mentoring({'N3A': [current], 'N3B': [past]}, current_year=_CURRENT_YEAR)

    assert _body_after(gen.doc, "Past Mentees:", 1) == [('tbl', 'Ada Lovelace')]
    assert _body_after(gen.doc, "Current Mentees:", 1) == [('tbl', 'Grace Hopper')]
    assert _mentoring_period_cell(gen.doc, 'Ada Lovelace') == '2007'
    assert _mentoring_period_cell(gen.doc, 'Grace Hopper') == '2021-present'


def test_end_only_past_mentee_renders_bare_year():
    """Class 2: an end date with no start renders that year, not an empty
    Mentoring Period cell."""
    gen = _mentoring_doc()
    entry = _mentee('N3B', 'Ada Lovelace', start_date=None, end_date='2004-12')
    entry['text'] = "Ada Lovelace, MPH 2004"

    gen._fill_mentoring({'N3B': [entry]})

    assert _body_after(gen.doc, "Past Mentees:", 1) == [('tbl', 'Ada Lovelace')]
    assert _mentoring_period_cell(gen.doc, 'Ada Lovelace') == '2004'


def test_end_only_period_written_in_source_renders_as_written():
    """The renderer hands the entry's text to the period formatter: a CV's
    own "1999-02" range, stored as a year-month end date, renders as the
    author wrote it, not as the single year 1999."""
    gen = _mentoring_doc()
    entry = _mentee('N3B', 'Ada Lovelace', start_date=None, end_date='1999-02')
    entry['text'] = "Ada Lovelace, MS 1999-02"

    gen._fill_mentoring({'N3B': [entry]})

    assert _mentoring_period_cell(gen.doc, 'Ada Lovelace') == '1999-02'


def test_start_only_mentee_moved_to_current_renders_open_range():
    """An N3B entry the ongoing rule moves renders "-present" under
    "Current Mentees:"."""
    gen = _mentoring_doc()
    entry = _mentee('N3B', 'Ada Lovelace', start_date='2019', end_date=None)
    entry['text'] = "Ada Lovelace, PhD student, 2019-"

    gen._fill_mentoring({'N3B': [entry]})

    assert _body_after(gen.doc, "Current Mentees:", 1) == [('tbl', 'Ada Lovelace')]
    assert _mentoring_period_cell(gen.doc, 'Ada Lovelace') == '2019-present'


# --- mentoring: Current -> Past migration (class E9, EBYSBC autopsy) -------------

@pytest.mark.parametrize("fields", [
    {'start_date': '2015'},                                  # a lone past year
    {'start_date': '2015', 'end_date': None},
    {'start_date': '2023-09'},
    {'start_date': '2023', 'end_date': '2024'},              # a closed past period
    {'start_date': '2024', 'end_date': str(_CURRENT_YEAR - 1)},
    {'end_date': '2020'},
], ids=lambda f: repr(f))
def test_partition_moves_an_ended_current_mentee_to_past(fields):
    """Stage 3b codes N3A for a lone past year or a closed period; neither
    is still running, so the entry renders under Past Mentees."""
    entry = _mentee('N3A', 'Ada Lovelace', **fields)
    entry['text'] = "Ada Lovelace, PhD dissertation"
    partition = _partition_mentoring_entries({'N3A': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.past == (entry,)
    assert partition.current == ()
    assert partition.moved_to_past == 1
    assert partition.moved_to_current == 0


@pytest.mark.parametrize("text", [
    "2015 Ada Lovelace, MD thesis; now Assistant Professor",
    "Ada Lovelace, summer student 2015. Current position: Resident",
    "Ada Lovelace, MD 2015; Fellow - now Assistant Professor",
    "2015 \u2013 2016 Ada Lovelace, postdoctoral fellow",
])
def test_partition_moves_an_ended_current_mentee_whatever_its_text_says_about_the_mentee(text):
    """"now ..." and "Current position" describe the mentee, not the
    mentoring, and a closed range is not an open marker."""
    entry = _mentee('N3A', 'Ada Lovelace', start_date='2015')
    entry['text'] = text
    partition = _partition_mentoring_entries({'N3A': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.past == (entry,)
    assert partition.moved_to_past == 1


@pytest.mark.parametrize("fields", [
    {'start_date': str(_CURRENT_YEAR)},                     # this year: may still run
    {'start_date': '2025', 'end_date': str(_CURRENT_YEAR)},
    {'start_date': '2029'},                                  # an expected completion
    {'start_date': '2026', 'end_date': '2028'},
    {'start_date': '2029 (anticipated)'},                    # no readable year
    {'start_date': '2016-17'},
    {'start_date': '2015', 'end_date': 'present'},           # says it is running
    {},                                                      # no dates at all
], ids=lambda f: repr(f))
def test_partition_keeps_a_current_mentee_that_has_not_ended(fields):
    entry = _mentee('N3A', 'Ada Lovelace', **fields)
    entry['text'] = "Ada Lovelace, PhD dissertation"
    partition = _partition_mentoring_entries({'N3A': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.current == (entry,)
    assert partition.past == ()
    assert partition.moved_to_past == 0


@pytest.mark.parametrize("start_date, text", [
    ('2015', "Ada Lovelace, PhD student, 2015-"),
    ('2015', "Ada Lovelace, PhD student, 2015 - present"),
    # The entry opens with the year and a dash: "since 2015", with the tab
    # after the dash lost by the reader.
    ('2015', "2015 \u2013 Ada Lovelace, PhD student"),
    ('2015-03', "2015.03 \u2013 Ada Lovelace, PhD student"),
    # Open-range markers the strict ongoing check does not read: a
    # box-drawing dash, a two-digit year, "pres", a dash before stage 2's
    # cell separator, and a later period left open.
    ('2015-03', "Ada Lovelace, PhD student 03/2015 \u2500 current"),
    ('2004', "Faculty advisor for the chief resident 04-present (Ada Lovelace)"),
    ('2015', "2015-2016 Ada Lovelace, intern\t2016-pres PhD student"),
    ('2015', "Ada Lovelace | PhD student | 2015 \u2013 | Fictional Institute"),
])
def test_partition_keeps_a_current_mentee_whose_source_leaves_the_year_open(start_date, text):
    entry = _mentee('N3A', 'Ada Lovelace', start_date=start_date)
    entry['text'] = text
    partition = _partition_mentoring_entries({'N3A': [entry]}, current_year=_CURRENT_YEAR)
    assert partition.current == (entry,)
    assert partition.moved_to_past == 0


def test_partition_judges_an_n3a_period_by_the_year_it_is_given():
    """The boundary, both sides: a period ending the year before
    `current_year` is past, one ending in `current_year` is not."""
    entry = _mentee('N3A', 'Ada Lovelace', start_date='2019', end_date='2021')
    assert _partition_mentoring_entries(
        {'N3A': [entry]}, current_year=2022).past == (entry,)
    assert _partition_mentoring_entries(
        {'N3A': [entry]}, current_year=2021).current == (entry,)


def test_ended_current_mentee_renders_its_year_under_past_mentees():
    """The render: an N3A lone past year sits under "Past Mentees:" as the
    bare year, not under "Current Mentees:" as "<year>-present"."""
    gen = _mentoring_doc()
    ended = _mentee('N3A', 'Ada Lovelace', start_date='2015', end_date=None)
    ended['text'] = "2015 Ada Lovelace, PhD dissertation"
    running = _mentee('N3A', 'Grace Hopper', start_date='2021', end_date=None)
    running['text'] = "2021 \u2013 Grace Hopper, PhD student"

    gen._fill_mentoring({'N3A': [ended, running]}, current_year=_CURRENT_YEAR)

    assert _body_after(gen.doc, "Past Mentees:", 1) == [('tbl', 'Ada Lovelace')]
    assert _body_after(gen.doc, "Current Mentees:", 1) == [('tbl', 'Grace Hopper')]
    assert _mentoring_period_cell(gen.doc, 'Ada Lovelace') == '2015'
    assert _mentoring_period_cell(gen.doc, 'Grace Hopper') == '2021-present'


# --- mentoring: summary and N4 outcome lines ------------------------------------

def test_mentoring_summary_lines_sit_above_the_tables():
    """Aggregate N3A/N3B lines land directly under their heading, above the
    mentee tables, in input order."""
    gen = _mentoring_doc()
    entries_by_code = {
        'N3A': [_mentee('N3A', 'Ada Lovelace'),
                {'taxonomy_code': 'N3A', 'text': 'Current Ph.D. Students: 12',
                 'extracted_fields': {'mentee_name': None}},
                {'taxonomy_code': 'N3A', 'text': 'Current Postdocs: 3',
                 'extracted_fields': {}}],
    }

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, "Current Mentees:", 4) == [
        ('p', 'Current Ph.D. Students: 12'),
        ('p', 'Current Postdocs: 3'),
        ('tbl', 'Ada Lovelace'),
        ('p', ''),
    ]
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['entries_inserted'] == 3


def test_n4_outcomes_render_under_the_mentoring_header():
    """N4 lines -- routed as N4, or rewritten to N3A by the mismatch corrector
    with the original stashed -- go under MENTORING, not under either mentee
    heading, in input order (N4 first, then the reclaimed N3A)."""
    gen = _mentoring_doc()
    entries_by_code = {
        'N4': [{'taxonomy_code': 'N4', 'text': 'Three mentees now hold faculty posts.',
                'extracted_fields': {}}],
        'N3A': [{'taxonomy_code': 'N3A', 'taxonomy_code_original': 'N4',
                 'text': 'Two mentees won K awards.', 'extracted_fields': {}}],
    }

    gen._fill_mentoring(entries_by_code)

    assert _body_after(gen.doc, "MENTORING", 2) == [
        ('p', 'Three mentees now hold faculty posts.'),
        ('p', 'Two mentees won K awards.'),
    ]
    assert _body_after(gen.doc, "Current Mentees:", 1) == [('p', 'Past Mentees:')]
    assert len(gen.doc.tables) == 0


# --- other education / patents: extracted_fields=None -------------------------

def test_other_education_none_extracted_fields_does_not_raise():
    """A B2 entry with extracted_fields=None used to crash the fill loop at
    other_education.py:85. With no fields it has no program and no
    institution, so it is skipped and the table keeps only its header row."""
    gen = _new_generator()
    gen.doc.add_paragraph("OTHER EDUCATIONAL")
    gen.doc.add_table(rows=1, cols=3)

    entries = [{'taxonomy_code': 'B2', 'text': 'Some training program',
                'extracted_fields': None}]

    gen._fill_other_education(entries)  # must not raise AttributeError

    assert len(gen.doc.tables[0].rows) == 1


def test_patents_none_extracted_fields_does_not_raise():
    """An M2D entry with extracted_fields=None used to crash the fill loop
    at patents.py:80. With no title and no patent number it is sparse, so
    no table is built."""
    gen = _new_generator()
    gen.doc.add_paragraph().add_run("Patents & Inventions").bold = True

    entries = [{'taxonomy_code': 'M2D', 'text': 'A patent with no fields',
                'extracted_fields': None}]

    gen._fill_patents(entries)  # must not raise AttributeError

    assert len(gen.doc.tables) == 0


# --- other education: date precedence and attribution ----------------------------

def _b2_doc() -> WCMTemplateGenerator:
    gen = _new_generator()
    gen.doc.add_paragraph("OTHER EDUCATIONAL")
    header = gen.doc.add_table(rows=1, cols=3).rows[0].cells
    header[0].text, header[1].text, header[2].text = "Description", "Institution", "Dates"
    return gen


def _insertions(cell) -> list[tuple[str, str]]:
    """(author, text) of every tracked insertion in a cell, in order."""
    return [(ins.get(qn('w:author')),
             ''.join(t.text or '' for t in ins.iter(qn('w:t'))))
            for ins in cell._tc.iter(qn('w:ins'))]


@pytest.mark.parametrize("fields, text, dates, from_text", [
    ({'start_date': '2017-08', 'end_date': '2021-07', 'year': '2019'},
     'Certificate, August 2020', '08/17-07/21', False),
    ({'start_date': '2017-08', 'year': '2019'}, 'Certificate, August 2020',
     '08/17', False),
    ({'end_date': '2021-07', 'year': '2019'}, 'Certificate, August 2020',
     '07/21', False),
    ({'year': '2019'}, 'Certificate, August 2020', '2019', False),
    ({'year_awarded': '2018'}, 'Certificate, August 2020', '2018', False),
    ({}, 'Certificate in Epidemiology, August 2020', '2020', True),
    ({}, 'Certificate 2019-2021', '2021', True),
    ({}, 'Certificate in Epidemiology', '', False),
    ({'start_date': '2020'}, 'Workshop, Example Institute, 2020', '2020', False),
    ({'start_date': '2020'}, 'Workshop, Example Institute, 2020-', '2020-Present', False),
    ({'start_date': '2020', 'end_date': 'present'}, 'Workshop, 2020', '2020-Present', False),
], ids=['range-beats-year-and-text', 'start-only-beats-year', 'end-only-beats-year',
        'year-beats-text', 'year_awarded-beats-text', 'text-month-year',
        'text-range-takes-end-year', 'nothing', 'start-only-is-one-occasion',
        'start-only-open-dash-in-source', 'end-says-present'])
def test_b2_date_precedence(fields, text, dates, from_text):
    """start/end -> year (year, then year_awarded) -> raw text, each step
    winning outright; only the raw-text step is marked as inferred."""
    record = _normalize_other_education_entry(
        {'taxonomy_code': 'B2', 'text': text,
         'extracted_fields': {'program_name': 'Certificate', **fields}})
    assert (record.dates, record.dates_from_text) == (dates, from_text)


def test_b2_start_only_renders_the_bare_date_in_the_table():
    """#1220: one date in the source is the date attended, not an open range."""
    gen = _b2_doc()
    gen._fill_other_education([
        {'taxonomy_code': 'B2', 'text': 'Workshop, Example Institute, 2020',
         'extracted_fields': {'program_name': 'Workshop', 'start_date': '2020'}},
        {'taxonomy_code': 'B2', 'text': 'Course, Example Institute, 06/2019-',
         'extracted_fields': {'program_name': 'Course', 'start_date': '2019-06'}}])

    dates = {row.cells[0].text: row.cells[2].text for row in gen.doc.tables[0].rows[1:]}
    assert dates == {'Workshop': '2020', 'Course': '06/19-Present'}


def test_b2_text_extracted_year_renders_as_a_tracked_insertion():
    gen = _b2_doc()
    gen._fill_other_education([
        {'taxonomy_code': 'B2', 'text': 'Certificate in Epidemiology, August 2020',
         'extracted_fields': {'program_name': 'Certificate in Epidemiology'}}])

    row = gen.doc.tables[0].rows[1]
    assert row.cells[0].text == 'Certificate in Epidemiology'
    assert _insertions(row.cells[2]) == [('Text Extraction', '2020')]
    assert _insertions(row.cells[0]) == []


def test_b2_field_year_renders_as_plain_text():
    gen = _b2_doc()
    gen._fill_other_education([
        {'taxonomy_code': 'B2', 'text': 'Certificate in Epidemiology, August 2020',
         'extracted_fields': {'program_name': 'Certificate in Epidemiology',
                              'year': '2019'}}])

    row = gen.doc.tables[0].rows[1]
    assert row.cells[2].text == '2019'
    assert _insertions(row.cells[2]) == []


def test_b2_enriched_location_is_attributed_inside_the_institution_cell():
    """Stage-5b enrichment supplies the cleaned name (plain text) and the
    city/state, which is appended as a tracked insertion after it."""
    gen = _b2_doc()
    gen._fill_other_education([
        {'taxonomy_code': 'B2', 'text': 'Certificate in Biostatistics, Duke University, 2019',
         'extracted_fields': {'program_name': 'Certificate in Biostatistics',
                              'institution': 'Duke University, Durham', 'year': '2019'},
         'institution_enrichment': {'cleaned_name': 'Duke University', 'city': 'Durham',
                                    'state': 'North Carolina', 'country_code': 'US'}}])

    cell = gen.doc.tables[0].rows[1].cells[1]
    plain_runs = ''.join(r.text for r in cell.paragraphs[0].runs)
    assert plain_runs == 'Duke University'
    assert _insertions(cell) == [('Institution Enrichment', ', Durham, NC')]


def test_b2_extracted_location_is_plain_text():
    gen = _b2_doc()
    gen._fill_other_education([
        {'taxonomy_code': 'B2', 'text': 'Certificate, Some College',
         'extracted_fields': {'program_name': 'Certificate', 'institution': 'Some College',
                              'location': 'Boston, MA'}}])

    cell = gen.doc.tables[0].rows[1].cells[1]
    assert cell.text == 'Some College, Boston, MA'
    assert _insertions(cell) == []


def test_b2_institution_none_sentinel_is_blank():
    gen = _b2_doc()
    gen._fill_other_education([
        {'taxonomy_code': 'B2', 'text': 'Certificate',
         'extracted_fields': {'program_name': 'Certificate', 'institution': 'none'}}])

    assert gen.doc.tables[0].rows[1].cells[1].text == ''


# --- patents: row counts, ordering, instruction removal --------------------------

def _patent(text: str, **fields) -> dict:
    return {'taxonomy_code': 'M2D', 'text': text, 'extracted_fields': fields}


def _patents_doc() -> WCMTemplateGenerator:
    gen = _new_generator()
    gen.doc.add_paragraph().add_run("Patents & Inventions").bold = True
    gen.doc.add_paragraph("MENTORING")
    return gen


def _labels(table) -> list[str]:
    return [row.cells[0].text for row in table.rows]


def test_patents_all_eight_attributes_render_exactly_eight_rows_in_order():
    gen = _patents_doc()
    gen._fill_patents([_patent(
        'full patent', title='Widget for Doing Things', patent_number='US1234567',
        inventors='A. Inventor, B. Inventor', filing_date='2019-03-01',
        issue_date='2021-07-15', status='Issued', assignee='Some University',
        narrative='A device that does the things a widget should do.')])

    table = gen.doc.tables[0]
    assert _labels(table) == [
        'Title of invention:', 'Patent number:', 'Inventors:', 'Status:',
        'Filing date:', 'Issue date:', 'Assignee:', 'Description:']
    assert [row.cells[1].text for row in table.rows] == [
        'Widget for Doing Things', 'US1234567', 'A. Inventor, B. Inventor', 'Issued',
        '03/2019', '07/2021', 'Some University',
        'A device that does the things a widget should do.']


@pytest.mark.parametrize("fields, labels", [
    ({'title': 'Widget'}, ['Title of invention:']),
    ({'patent_number': 'US1', 'inventors': 'A. Inventor'},
     ['Patent number:', 'Inventors:']),
    ({'title': 'Widget', 'narrative': 'too short'}, ['Title of invention:']),
    ({'title': 'Widget', 'narrative': 'abcdefghij'}, ['Title of invention:']),
    ({'title': 'Widget', 'narrative': 'abcdefghijk'},
     ['Title of invention:', 'Description:']),
    ({'title': 'Widget', 'status': 'Pending', 'assignee': 'Some University'},
     ['Title of invention:', 'Status:', 'Assignee:']),
], ids=['title-only', 'number-and-inventors', 'narrative-9-chars-dropped',
        'narrative-10-chars-dropped', 'narrative-11-chars-kept', 'title-status-assignee'])
def test_patents_sparse_records_render_only_the_rows_they_have(fields, labels):
    gen = _patents_doc()
    gen._fill_patents([_patent('sparse patent', **fields)])

    assert len(gen.doc.tables) == 1
    assert _labels(gen.doc.tables[0]) == labels


@pytest.mark.parametrize("fields", [
    {'inventors': 'A. Inventor'},
    {'status': 'Pending', 'filing_date': '2019-03-01'},
    {},
], ids=['inventors-only', 'status-and-filing-only', 'no-fields'])
def test_patents_without_title_or_number_render_no_table(fields):
    gen = _patents_doc()
    gen._fill_patents([_patent('sparse patent', **fields)])

    assert len(gen.doc.tables) == 0
    assert gen.stats['entries_inserted'] == 0


def test_patents_multiple_records_keep_order_and_add_no_trailing_spacing():
    """Three entries, the middle one sparse: the two rendered tables come out
    most recent first with one spacing paragraph between them and none after
    the last -- the spacing is counted against rendered patents, not input."""
    gen = _patents_doc()
    gen._fill_patents([
        _patent('older', title='Older Widget', year='2019'),
        _patent('sparse', inventors='Nobody Named', year='2020'),
        _patent('newer', title='Newer Widget', year='2021'),
    ])

    assert _body_after(gen.doc, "Patents & Inventions", 4) == [
        ('tbl', 'Newer Widget'), ('p', ''), ('tbl', 'Older Widget'), ('p', 'MENTORING')]
    assert gen.stats['tables_populated'] == 2
    assert gen.stats['entries_inserted'] == 2


def test_patents_instruction_paragraph_is_blanked_and_left_below_the_tables():
    gen = _new_generator()
    gen.doc.add_paragraph().add_run("Patents & Inventions").bold = True
    instruction = gen.doc.add_paragraph(
        "Please include inventors, title of invention and patent number.")
    gen.doc.add_paragraph("MENTORING")

    gen._fill_patents([_patent('one', title='Widget')])

    assert instruction.text == ''
    body = list(gen.doc.element.body)
    assert body.index(instruction._element) == body.index(gen.doc.tables[0]._tbl) + 1
    assert _body_after(gen.doc, "Patents & Inventions", 3) == [
        ('tbl', 'Widget'), ('p', ''), ('p', 'MENTORING')]
