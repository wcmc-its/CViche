"""Regression guard for issue #442: the whole document is lost when the
stage-4 ``address`` field is a dict.

Stage 4 stores raw LLM JSON and ``coerce_field_value_types`` deliberately leaves
dicts intact, so a CV whose contact block is a two-column Home/Office table
arrives at stage 6 with ``address`` as a dict. ``home_address.replace(...)``
then raised ``AttributeError: 'dict' object has no attribute 'replace'`` and
aborted the render — two of 96 runs in the 2026-07-25 corpus batch (web094,
web147) produced no deliverable at all, having succeeded through stage 5d.

``_address_cell_text`` normalises at the read boundary and, because the dict is
the richer form, fills the home and office slots from it separately instead of
forcing the whole dict into one.

    python3 -m pytest src/unified_pipeline/tests/test_stage6_dict_address.py -p no:cacheprovider

Self-contained: no DB, no network, no PII (the two address shapes below are the
observed key vocabulary with invented values).
"""

import json
import sys
from pathlib import Path

from docx import Document

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6.normalization.pii import CAT_HOME_CONTACT  # noqa: E402
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    WCMTemplateGenerator,
    _address_cell_text,
    _labels_its_own_address_slots,
)

SURVIVES = "DISTINCTIVE_SURVIVING_TOKEN"


def test_string_address_is_returned_untouched():
    """The 101 string addresses in the corpus must render byte-identically."""
    assert _address_cell_text("1300 York Ave, New York, NY 10065", 'office') == \
        "1300 York Ave, New York, NY 10065"
    assert _address_cell_text("12 Elm St\tRye, NY 10580", 'home') == \
        "12 Elm St\tRye, NY 10580"
    assert _address_cell_text(None, 'home') == ""
    assert _address_cell_text("", 'office') == ""


def test_dict_address_splits_into_the_matching_slot():
    # web094's observed key pair
    value = {"home_address": "508 Howe Road, Merion, PA 19066",
             "office_address": "1701 N. 13th Street, Philadelphia PA 19122"}
    assert _address_cell_text(value, 'home') == "508 Howe Road, Merion, PA 19066"
    assert _address_cell_text(value, 'office') == "1701 N. 13th Street, Philadelphia PA 19122"

    # web147's observed key pair — 'business_address' rather than 'office_address'
    value = {"business_address": "3010 Old Clinic Bldg, Chapel Hill, NC 27599",
             "home_address": "111 Simerville Road, Chapel Hill, NC 27517"}
    assert _address_cell_text(value, 'office') == "3010 Old Clinic Bldg, Chapel Hill, NC 27599"
    assert _address_cell_text(value, 'home') == "111 Simerville Road, Chapel Hill, NC 27517"


def test_dict_with_only_one_slot_leaves_the_other_empty():
    value = {"office_address": "1300 York Ave, New York, NY 10065"}
    assert _address_cell_text(value, 'office') == "1300 York Ave, New York, NY 10065"
    assert _address_cell_text(value, 'home') == ""


def test_unlabelled_dict_is_one_address_joined_whole():
    """No slot key at all: this is a single address split into parts, so it is
    joined for whichever slot the caller asks for and routed by entry text."""
    value = {"street": "1300 York Ave", "city": "New York", "state": "NY", "zip": ""}
    assert _address_cell_text(value, 'office') == "1300 York Ave; New York; NY"
    assert _address_cell_text(value, 'home') == "1300 York Ave; New York; NY"
    assert _labels_its_own_address_slots(value) is False


def test_partly_labelled_dict_does_not_drop_the_unmatched_key():
    """One recognised key must not disable the fallback for the rest of the
    dict. The key vocabulary is unbounded -- stage 4 extracts with no schema."""
    value = {"home_address": "508 Howe Road, Merion, PA 19066",
             "mailing_address": "1701 N. 13th Street, Philadelphia PA 19122"}
    assert _address_cell_text(value, 'home') == "508 Howe Road, Merion, PA 19066"
    assert _address_cell_text(value, 'office') == "1701 N. 13th Street, Philadelphia PA 19122"


def test_nested_value_under_a_slot_key_is_recursed_into_not_dropped():
    """A recognised key whose value is itself structured must not vanish, and
    must not render as a Python repr."""
    value = {"home_address": {"street": "12 Elm St", "city": "Rye, NY 10580"},
             "office_address": ["1300 York Ave", "New York, NY 10065"]}
    assert _address_cell_text(value, 'home') == "12 Elm St; Rye, NY 10580"
    assert _address_cell_text(value, 'office') == "1300 York Ave; New York, NY 10065"
    for slot in ('home', 'office'):
        rendered = _address_cell_text(value, slot)
        assert "{" not in rendered and "[" not in rendered, "dict/list repr leaked"


def test_empty_dict_yields_empty_text():
    assert _address_cell_text({}, 'home') == ""
    assert _address_cell_text({}, 'office') == ""


def test_list_and_other_types_do_not_raise():
    assert _address_cell_text([{"home_address": "12 Elm St"}, {"home_address": "9 Oak Rd"}],
                              'home') == "12 Elm St; 9 Oak Rd"
    assert _address_cell_text(12345, 'office') == "12345"


def test_every_return_is_a_string_so_the_render_sites_cannot_raise():
    """The render sites call .replace() on whatever comes back (#442)."""
    shapes = [None, "", "  ", "a street", 12345, 3.5, True, [], {}, ["a", None],
              {"home_address": None}, {"office": {"x": ["y", {"z": 1}]}},
              {"street": {"a": 1}}, [{"home": []}], {"home_address": 7}]
    for value in shapes:
        for slot in ('home', 'office'):
            out = _address_cell_text(value, slot)
            assert isinstance(out, str), f"{value!r} -> {out!r}"
            out.replace('\t', '\n')  # the exact call that crashed


def _render(tmp_path, entries):
    """Render and return (full text, {contact label: cell value}, generator).

    The per-label mapping matters: asserting only that an address appears
    *somewhere* leaves the home/office wiring untested, and a swap of the two
    is an objectively wrong document that no blob assertion can see. The
    generator is returned too so a #821 home-address test can check
    `gen._pii_result.withheld` -- the home slot renders empty either way it
    is routed wrong, so the withheld list is what still tells "correctly
    routed to home, then withheld" apart from "never routed at all"."""
    gen = WCMTemplateGenerator(verbose=False)
    # Neutralize the LLM-driven appendix-reconsider pass: keeps the render
    # deterministic and credential-free.
    gen._reconsider_appendix_entries = lambda: None
    ip = tmp_path / "in.json"
    op = tmp_path / "out.docx"
    ip.write_text(json.dumps({"document_uid": "TESTAA", "entries": entries}))
    gen.generate(str(ip), str(op), research_summary_path=None)
    doc = Document(str(op))
    parts = [p.text for p in doc.paragraphs]
    contact = {}
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            parts += cells
            if len(cells) >= 2 and cells[0].strip().endswith(':'):
                contact.setdefault(cells[0].strip(), cells[1].strip())
    return "\n".join(parts), contact, gen


def test_dict_address_still_produces_a_document(tmp_path):
    """The actual #442 symptom: the render aborted and nothing was written.

    #821: the home half is withheld with notice, so it no longer lands in
    the home CELL -- `gen._pii_result.withheld` is what still proves it was
    routed there correctly (not dropped, not misrouted to office) before
    being withheld."""
    entries = [
        {"text": "Home:  Office Address:\t508 Howe Road\t1701 N. 13th Street",
         "taxonomy_code": "A", "element_idx_start": 0,
         "extracted_fields": {
             "name": "Jane Q. Public, MD",
             "address": {"home_address": "508 Howe Road, Merion, PA 19066",
                         "office_address": "1701 N. 13th Street, Philadelphia PA 19122"},
         }},
        {"text": f"{SURVIVES} Professor of Medicine, 2010-2014",
         "taxonomy_code": "C", "element_idx_start": 5,
         "extracted_fields": {"title": f"{SURVIVES} Professor of Medicine"}},
    ]
    text, contact, gen = _render(tmp_path, entries)
    assert SURVIVES in text, "whole document was lost — the #442 crash"
    # The office half still lands in its own cell, untouched by #821; the
    # home half is withheld rather than landing in ITS cell or, worse,
    # bleeding into the office one.
    assert contact.get("Home address:", "") == ""
    assert contact.get("Office address:") == "1701 N. 13th Street, Philadelphia PA 19122"
    assert "508 Howe Road" not in text, "a withheld home address reached the document"
    assert "home_address" not in text, "the dict repr leaked into the document"
    assert any(item.category == CAT_HOME_CONTACT for item in gen._pii_result.withheld)


def test_unlabelled_dict_is_routed_by_the_entry_text_not_forced_to_office(tmp_path):
    """A dict that names no slot must obey the entry's own Home/Office label.
    Putting a home address in the office row is worse than the old drop.

    #821: the routed-to-home value is now withheld rather than rendered, so
    the home cell is empty either way it was routed -- `gen._pii_result.
    withheld` is what still proves it reached the home slot (and was
    withheld there) rather than never being routed, or landing in office."""
    entries = [
        {"text": "Home Address:\t12 Elm St, Rye, NY 10580", "taxonomy_code": "A",
         "element_idx_start": 0,
         "extracted_fields": {"name": "Jane Q. Public, MD",
                              "address": {"street": "12 Elm St", "city": "Rye",
                                          "state": "NY 10580"}}},
    ]
    _, contact, gen = _render(tmp_path, entries)
    assert contact.get("Home address:", "") == ""
    assert not contact.get("Office address:"), \
        "a home-labelled address was rendered into the office row"
    assert any(item.category == CAT_HOME_CONTACT for item in gen._pii_result.withheld)


def test_string_address_lands_in_the_same_cell_as_before(tmp_path):
    entries = [
        {"text": "Office Address: 1300 York Ave", "taxonomy_code": "A",
         "element_idx_start": 0,
         "extracted_fields": {"name": "Jane Q. Public, MD",
                              "address": "1300 York Ave, New York, NY 10065"}},
    ]
    _, contact, _gen = _render(tmp_path, entries)
    assert contact.get("Office address:") == "1300 York Ave, New York, NY 10065"
    assert not contact.get("Home address:")


def test_home_labelled_string_still_lands_in_the_home_cell(tmp_path):
    """#821: a home-labelled string address is still routed to the home
    slot (not office), it is just withheld once it gets there rather than
    rendered -- `gen._pii_result.withheld` is what distinguishes that from
    never being routed at all."""
    entries = [
        {"text": "Home Address: 12 Elm St", "taxonomy_code": "A",
         "element_idx_start": 0,
         "extracted_fields": {"name": "Jane Q. Public, MD",
                              "address": "12 Elm St, Rye, NY 10580"}},
    ]
    _, contact, gen = _render(tmp_path, entries)
    assert contact.get("Home address:", "") == ""
    assert not contact.get("Office address:")
    assert any(item.category == CAT_HOME_CONTACT for item in gen._pii_result.withheld)


if __name__ == "__main__":
    import tempfile
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            if "tmp_path" in fn.__code__.co_varnames[:fn.__code__.co_argcount]:
                with tempfile.TemporaryDirectory() as d:
                    fn(Path(d))
            else:
                fn()
    print("OK")
