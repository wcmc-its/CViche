"""Honors-table rendering fixes (#229): the multi-award fallback used to put
citation blobs in the name cell, leak state abbreviations/cities into the
Organization column, leave the year column empty, and duplicate the org
inside the name. Fixtures mirror the 2Q1_ZQ fused H entry with fictional
content.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_honors_render.py -p no:cacheprovider
"""

import sys
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


def _render_honors(entries):
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen.doc.add_paragraph("H. HONORS AND AWARDS")
    table = gen.doc.add_table(rows=1, cols=3)
    for i, header in enumerate(["Name of award", "Organization",
                                "Date awarded (yyyy)"]):
        table.rows[0].cells[i].text = header
    gen._fill_honors(entries)
    return [[c.text for c in r.cells] for r in table.rows[1:]]


_FUSED_LINES = [
    # 2Q1_ZQ row-1 shape: citation blob with venue/city/state/month-year tail;
    # stage 4 extracted clean fields for this first award only.
    "Groundbreaking Widget Design Award – Runner-up Presentation. "
    "Alice Smith, Bob Jones. Groundbreaking Widget Design. "
    "Fictional Services University Education Day, Bethesda, MD, August 2025.",
    # leading year + org duplicated as the trailing comma segment
    "2020 AICT Outstanding Article Award, Association for Imaginary "
    "Communication and Technology (AICT)",
    # leading year RANGE
    "2015-2017 Imaginary Resources National Scholar Fellowship, "
    "The Fictional Benefactor Foundation",
    # org followed by a bare city segment
    "2013 Jerrold Q. Fake Fellowship, Made-up University, Bloomington",
]

_FUSED_ENTRY = {
    "taxonomy_code": "H",
    "text": "\n".join(_FUSED_LINES),
    "extracted_fields": {
        "award_name": "Groundbreaking Widget Design Award – Runner-up Presentation",
        "granting_body": "Fictional Services University Education Day",
        "date": "2025",
    },
}


def test_fused_entry_renders_clean_rows():
    rows = _render_honors([_FUSED_ENTRY])
    assert len(rows) == 4

    # Row 1: stage-4 fields used verbatim — no citation blob, no 'MD' org
    assert rows[0] == ["Groundbreaking Widget Design Award – Runner-up Presentation",
                       "Fictional Services University Education Day", "2025"]

    # Row 2: leading year extracted, org not duplicated in the name cell
    assert rows[1][0] == "AICT Outstanding Article Award"
    assert rows[1][1] == ("Association for Imaginary Communication and "
                          "Technology (AICT)")
    assert rows[1][2] == "2020"

    # Row 3: leading year range extracted
    assert rows[2][0] == "Imaginary Resources National Scholar Fellowship"
    assert rows[2][1] == "The Fictional Benefactor Foundation"
    assert rows[2][2].startswith("2015")

    # Row 4: org is the university, not the city; city tail stripped from name
    assert rows[3][0] == "Jerrold Q. Fake Fellowship"
    assert rows[3][1] == "Made-up University"
    assert rows[3][2] == "2013"


def test_single_award_trailing_year_with_period():
    entry = {"taxonomy_code": "H",
             "text": "Best Poster Award, Imaginary Society of Things, 2021.",
             "extracted_fields": {}}
    rows = _render_honors([entry])
    assert rows == [["Best Poster Award", "Imaginary Society of Things", "2021"]]


def test_org_extraction_never_returns_state_abbrev_or_city():
    gen = WCMTemplateGenerator(verbose=False)
    # an institutional-keyword segment wins over later city/state segments
    # (the old single-pass returned 'MD' / 'Bloomington' first, #229)
    assert gen._extract_organization_from_award(
        "Best Talk Award, Made-up University, Bloomington, MD") == \
        "Made-up University"
    # no IKW anywhere: state abbrevs and digit segments rejected, proper-noun
    # fallback still works
    assert gen._extract_organization_from_award(
        "Some Prize, Weill Cornell, NY") == "Weill Cornell"


def test_split_award_year_shapes():
    gen = WCMTemplateGenerator(verbose=False)
    assert gen._split_award_year("2020 Great Award, Some Org") == (
        "Great Award, Some Org", "2020")
    assert gen._split_award_year("2015-2017 Long Fellowship") == (
        "Long Fellowship", "2015-2017")
    assert gen._split_award_year("Venue Day, Bethesda, August 2025.") == (
        "Venue Day, Bethesda", "2025")
    assert gen._split_award_year("No year here at all") == (
        "No year here at all", "")


def test_892_a_pii_withheld_award_leaves_no_row_for_the_entry():
    """The PII pass dropped award_name (an O-1 visa); the cut-text remnant,
    the USCIS org and the year must not render as an honors row."""
    from unified_pipeline.stage6.pii_pass import run_pii_pass
    from unified_pipeline.stage_6_word_template import (
        RENDER_ROUTED_CODES,
        TAXONOMY_TO_SECTION,
    )
    entry = {"taxonomy_code": "H",
             "text": "2019 Extraordinary Ability in Sciences, O-1 Visa | "
                     "U.S. Citizen & Immigration Service (USCIS)",
             "extracted_fields": {
                 "award_name": "Extraordinary Ability in Sciences, O-1 Visa",
                 "granting_body": "U.S. Citizen & Immigration Service (USCIS)",
                 "date": "2019"}}
    keeper = {"taxonomy_code": "H", "text": "",
              "extracted_fields": {"award_name": "Teaching Award",
                                   "granting_body": "Example University",
                                   "date": "2018"}}
    run_pii_pass({"H": [entry, keeper]}, routed_codes=RENDER_ROUTED_CODES,
                 section_names=TAXONOMY_TO_SECTION)
    assert _render_honors([entry, keeper]) == [
        ["Teaching Award", "Example University", "2018"]]


def test_org_not_fabricated_from_the_award_name_itself():
    """#887: with no granting body in the text, an award named after a
    college or university must not become its own organization."""
    gen = WCMTemplateGenerator(verbose=False)
    # the award name minus its trailing award word (no digit, no role word:
    # only the equals-the-award-name guard can catch this one)
    assert gen._extract_organization_from_award(
        "College of Education Outstanding Thesis Award") == ""
    # a digit in the built organization (a year is the date column's)
    assert gen._extract_organization_from_award(
        "2016 College of Education 2015 Outstanding Thesis Award") == ""
    assert gen._extract_organization_from_award(
        "College of Education 2015 Outstanding Thesis Award") == ""
    # a role word ends the built organization
    assert gen._extract_organization_from_award(
        "University of Western Ontario Representative, "
        "Ontario Undergraduate Student Alliance") == ""


def test_org_not_fabricated_from_the_award_first_comma_clause_1412():
    """#1412 (RCBKFG FLYBMX 288/289): the short-proper-noun fallback took
    the clause an award name opens with as its organization. A first clause
    with no institutional keyword is the award's own words, not a grantor."""
    gen = WCMTemplateGenerator(verbose=False)
    # the last clause is too long for the fallback, so it reached the first
    assert gen._extract_organization_from_award(
        "First Place, Imaginary Regional Student Poster Competition Day") == ""
    assert gen._extract_organization_from_award(
        "Cum Laude, Made-up School of Example Studies and Arts") == ""
    # a first clause in no award words is still the grantor
    assert gen._extract_organization_from_award(
        "Imaginary Civic Network, Forum for Example Causes") == \
        "Imaginary Civic Network"
    # a first clause that names an institution is still the grantor
    assert gen._extract_organization_from_award(
        "Made-up University, Teaching Award") == "Made-up University"
    assert gen._extract_organization_from_award(
        "Made-up University Award, Forum for Example Causes") == \
        "Made-up University Award"
    # a later keyword-free clause is still the proper-noun fallback, award
    # word or not: only the clause the name opens with is the award's own
    assert gen._extract_organization_from_award(
        "Some Prize, Weill Cornell") == "Weill Cornell"
    assert gen._extract_organization_from_award(
        "Some Prize, Imaginary Award Committee") == "Imaginary Award Committee"


def test_org_not_fabricated_from_a_clause_of_the_award_name_1711():
    """#1711 (NECRGC XUTTKJ 181/192, SPIHVQ 760): the short-proper-noun
    fallback took a comma clause of the award's own name -- its last clause
    when that ends in an award noun, a middle clause of a list -- as the
    granting body."""
    gen = WCMTemplateGenerator(verbose=False)
    # an eponymous award's tail ends in an award noun: never a grantor
    for text in ("Jane Q. Doe, MD Mentor Award",
                 "Resident Supervisor Award: Jane Q. Doe, MD Mentor Award",
                 "Jane Q. Doe, Imaginary Prize",
                 "Jane Q. Doe, Example Fellowship",
                 "Jane Q. Doe, Example Runner-up",
                 "Jane Q. Doe, Example Medal"):
        assert gen._extract_organization_from_award(text) == "", text
    # a keyword-free middle clause is a list inside the award's name
    assert gen._extract_organization_from_award(
        "Imaginary Workshop on Grant Writing, Preparation, and Submissions "
        "in Example Research") == ""
    assert gen._extract_organization_from_award(
        "Example Cup, Leadership, Teamwork, or Spirit Shown") == ""
    # ...and so is the first clause of a three-clause list
    assert gen._extract_organization_from_award(
        "Advising, Mentoring, & Teaching Recognition Award") == ""
    assert gen._extract_organization_from_award(
        "Best Doctors, Top Physicians, Selected by Their Peers") == ""
    # a clause of nothing but degree post-nominals is a person's name tail
    for text in ("Example Award, Jane Q. Doe, MD",
                 "Example Award, Jane Q. Doe, M.D.",
                 "Example Award, Jane Q. Doe, PhD",
                 "Example Award, Jane Q. Doe, MD PhD",
                 "Example Award, Jane Q. Doe, DO"):
        assert gen._extract_organization_from_award(text) == "", text


def test_org_extraction_keeps_a_real_grantor_after_1711_guards():
    """The #1711 guards leave a real grantor alone: a last clause that ends
    in no award noun, a degree token that only opens an organization's name,
    a grantor after a trailing year, and an award-noun clause that names an
    institution."""
    gen = WCMTemplateGenerator(verbose=False)
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell") == "Weill Cornell"
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell, 2010") == "Weill Cornell"
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell, NY") == "Weill Cornell"
    # a trailing lowercase description is passed over, like a year
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell, voted by the students") == \
        "Weill Cornell"
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell, 2010, voted by the students") == \
        "Weill Cornell"
    # a description whose first word only begins with "or"/"and" is still one
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell, organized by the students") == \
        "Weill Cornell"
    # only a WHOLE last word that is an award noun counts, not a part of one
    assert gen._extract_organization_from_award(
        "Teaching Award, Imaginary Medal-Works") == "Imaginary Medal-Works"
    # an empty clause is not a degree clause
    assert gen._extract_organization_from_award(
        "Teaching Award, Weill Cornell,, ") == "Weill Cornell"
    assert gen._extract_organization_from_award(
        "Best Paper, American Heart Association") == \
        "American Heart Association"
    assert gen._extract_organization_from_award(
        "Example Award, MD Anderson") == "MD Anderson"
    assert gen._extract_organization_from_award(
        "Example Award, MD Anderson Cancer Center") == \
        "MD Anderson Cancer Center"
    assert gen._extract_organization_from_award(
        "Some Prize, Imaginary Award Committee") == "Imaginary Award Committee"


def test_a_comma_in_the_award_name_keeps_the_name_whole_1711():
    """#1711 end to end: the Name cell keeps the whole award name and the
    Organization cell stays empty, not split across the two."""
    entry = {"taxonomy_code": "H", "text": "",
             "extracted_fields": {
                 "award_name": "Jane Q. Doe, MD Mentor Award",
                 "granting_body": None, "date": "2019"}}
    assert _render_honors([entry]) == [[
        "Jane Q. Doe, MD Mentor Award", "", "2019"]]


def test_award_name_keeps_an_org_it_ends_on_after_a_preposition_1412():
    """#1412 (RCBKFG KUUKNJ 139): stripping a granting body that is the
    object of the name's last phrase left the award cell ending in "of". The name now renders whole."""
    entry = {"taxonomy_code": "H", "text": "",
             "extracted_fields": {
                 "award_name": "Distinguished Alumnus of Made-up Medical College, Narnia",
                 "granting_body": "Made-up Medical College", "date": "1982"}}
    assert _render_honors([entry]) == [[
        "Distinguished Alumnus of Made-up Medical College, Narnia",
        "Made-up Medical College", "1982"]]


def test_org_extraction_still_finds_a_real_grantor_after_887_guards():
    gen = WCMTemplateGenerator(verbose=False)
    # a grantor that is a PREFIX of the award name (not the whole name)
    assert gen._extract_organization_from_award(
        "Society of Fictional Medicine Mentoring Award") == \
        "Society of Fictional Medicine"
    assert gen._extract_organization_from_award(
        "Best Talk Award, Made-up University") == "Made-up University"


def test_single_award_with_no_granting_body_renders_empty_org_887():
    entry = {"taxonomy_code": "H",
             "text": "College of Education 2015 Outstanding Thesis Award",
             "extracted_fields": {
                 "award_name": "College of Education 2015 Outstanding "
                               "Thesis Award",
                 "granting_body": None, "date": "2016"}}
    rows = _render_honors([entry])
    assert rows == [["College of Education 2015 Outstanding Thesis Award",
                     "", "2016"]]


# --------------------------------------------------------------------------
# #1245 / EBYSBC class E15: dates under keys outside the H schema, and one
# stage-4 award record split into several rows. Every value is invented.
# --------------------------------------------------------------------------

def _h_entry(text, **fields):
    return {"taxonomy_code": "H", "text": text, "extracted_fields": fields}


def test_a_start_end_range_renders_when_the_schema_date_is_empty():
    """KDAZOM-04 / VVRTUC-03 / OTBUCZ-01 shape: stage 4 put the range under
    start_date/end_date and left `date` empty; the text's "2011-13" is not a
    year the text fallback can read, so the cell used to be empty."""
    rows = _render_honors([
        _h_entry("Imaginary Teaching Prize, Fictional Board 2011-13",
                 award_name="Imaginary Teaching Prize",
                 granting_body="Fictional Board", date=None,
                 start_date="2011", end_date="2013"),
        _h_entry("Pretend Honor Roll, Made-up School 2017-",
                 award_name="Pretend Honor Roll", granting_body="Made-up School",
                 start_date="2017-09", end_date="present"),
    ])
    assert rows == [
        ["Pretend Honor Roll", "Made-up School", "2017-Present"],
        ["Imaginary Teaching Prize", "Fictional Board", "2011-2013"],
    ]


def test_an_end_date_with_no_start_closes_the_range_date_opens():
    """XWNZWW-09 shape: `date` plus an `end_date` and no `start_date` is one range. Next to a
    start_date of its own, `date` is the award's date and wins."""
    rows = _render_honors([
        _h_entry("Named in Pretend Directory 2001-2004",
                 award_name="Named in Pretend Directory", granting_body="Imaginary Press",
                 date="2001", end_date="2004"),
        _h_entry("2016 Made-up Fellow, Fictional Institute (2016-2017)",
                 award_name="Made-up Fellow", granting_body="Fictional Institute",
                 date="2016", start_date="2016", end_date="2017"),
    ])
    assert rows == [
        ["Made-up Fellow", "Fictional Institute", "2016"],
        ["Named in Pretend Directory", "Imaginary Press", "2001-2004"],
    ]


def test_a_dates_list_renders_every_date_it_holds():
    """HZGJFM-02 (a list of spans) and BZZNRL-06 (a '; ' string): the text
    fallback kept only the last year, or none for a "1993-94" span."""
    rows = _render_honors([
        _h_entry("Pretend Top Clinician 1993-94, 2008-2009",
                 award_name="Pretend Top Clinician", date=None,
                 dates=[{"start_date": "1993", "end_date": "1994"},
                        {"start_date": "2008", "end_date": "2009"}]),
        _h_entry("Imaginary Teaching Award, 2012, 2013",
                 award_name="Imaginary Teaching Award", date=None,
                 dates="2012; 2013"),
        _h_entry("Fictional Prize",
                 award_name="Fictional Prize", additional_dates="2008; 2007"),
    ])
    assert [row[2] for row in rows] == [
        "1993-1994, 2008-2009", "2012, 2013", "2008, 2007"]


def test_each_line_of_one_award_won_in_several_years_keeps_its_own_year():
    """ZDCXIV-03: every line naming stage 4's award rendered stage 4's one
    date. ZDCXIV-02: the organization's own line above them rendered as an
    award of its own."""
    rows = _render_honors([_h_entry(
        "Fictional College of Pretend Studies\n"
        "Made-up Faculty Award — 06/11/2019\n"
        "Made-up Faculty Award — 05/14/2018",
        award_name="Made-up Faculty Award",
        granting_body="Fictional College of Pretend Studies",
        date="2019-06-11", additional_dates="2018-05-14")])
    assert rows == [
        ["Made-up Faculty Award", "Fictional College of Pretend Studies", "2019"],
        ["Made-up Faculty Award", "Fictional College of Pretend Studies", "2018"],
    ]


def test_a_line_keeps_its_own_trailing_year_or_date_column():
    """The same rule for the other two ways a line dates itself: a trailing
    year (`_split_award_year`) and a '|' date column (the parsed cell)."""
    for text in ("Made-up Faculty Award, 2019\nMade-up Faculty Award, 2018",
                 "Made-up Faculty Award | 2019\nMade-up Faculty Award | 2018"):
        rows = _render_honors([_h_entry(
            text, award_name="Made-up Faculty Award", date="2019",
            additional_dates="2018")])
        assert [row[2] for row in rows] == ["2019", "2018"], text


def test_the_granting_body_line_is_matched_whatever_its_case_or_full_stop():
    rows = _render_honors([_h_entry(
        "FICTIONAL COLLEGE OF PRETEND STUDIES.\n"
        "Made-up Faculty Award — 06/11/2019\n"
        "Made-up Faculty Award — 05/14/2018",
        award_name="Made-up Faculty Award",
        granting_body="Fictional College of Pretend Studies",
        date="2019-06-11", additional_dates="2018-05-14")])
    assert [row[0] for row in rows] == ["Made-up Faculty Award"] * 2


def test_one_award_written_over_several_lines_is_one_row():
    """ZDCXIV-02: "Award", then "Organization — date", or a poster title and
    its authors above the dated line, rendered one row per line."""
    rows = _render_honors([
        _h_entry("Imaginary Pharmacy Scholarship\n"
                 "Fictional University, College of Pretend — 07/09/2010",
                 award_name="Imaginary Pharmacy Scholarship",
                 granting_body="Fictional University, College of Pretend",
                 date="2010-07-09"),
        _h_entry("Pretend Poster Prize Runner-up\n"
                 "A Made-up Study of Widget Costs\n"
                 "Doe J, Roe R, Poe Q\n"
                 "Imaginary Congress of Widgets, Nowhere — 03/2012",
                 award_name="Pretend Poster Prize Runner-up",
                 granting_body="Imaginary Congress of Widgets", date="2012"),
        _h_entry("Imaginary Counseling Contest\nFirst Prize — 02/2009",
                 award_name="First Prize, Imaginary Counseling Contest",
                 granting_body="Pretend Society", date="2009-02"),
    ])
    assert rows == [
        ["Pretend Poster Prize Runner-up", "Imaginary Congress of Widgets", "2012"],
        ["Imaginary Pharmacy Scholarship",
         "Fictional University, College of Pretend", "2010"],
        ["First Prize, Imaginary Counseling Contest", "Pretend Society", "2009"],
    ]


def test_a_list_whose_only_date_is_on_the_extracted_award_keeps_every_award():
    """The guard on the one-row reading: stage 4 extracted the last award of
    a fused list, and that award's own line is the only dated one. A curly
    quote in the source must not hide that line (stage 4 wrote a straight
    one), which would collapse the list into its last award. The year is
    written in free text as well as in a '|' date column: the column alone
    is caught earlier, by the date-cell guard, and would not reach this one."""
    for dated_line in ("Doe’s Neighbourhood Service Prize | 2015",
                       "Doe’s Neighbourhood Service Prize, 2015"):
        rows = _render_honors([_h_entry(
            "Pretend Mentoring Certificate\n"
            "Imaginary Service Medal\n" + dated_line,
            award_name="Doe's Neighbourhood Service Prize", date="2015")])
        assert [row[0] for row in rows] == [
            "Pretend Mentoring Certificate", "Imaginary Service Medal",
            "Doe's Neighbourhood Service Prize"], dated_line
        assert rows[2][2] == "2015"


def test_an_undated_list_is_still_read_as_a_list():
    rows = _render_honors([_h_entry(
        "Pretend Teaching Prize\nImaginary Service Medal",
        award_name="Pretend Teaching Prize")])
    assert [row[0] for row in rows] == [
        "Pretend Teaching Prize", "Imaginary Service Medal"]


def test_a_list_naming_years_the_record_does_not_have_stays_a_list():
    """Stage 4's award is on an undated line here, so only the dates tie
    the other lines to it -- and they are not the record's own."""
    rows = _render_honors([_h_entry(
        "Pretend Teaching Prize\nImaginary Service Medal 2014\n"
        "Fictional Research Cup 2011",
        award_name="Pretend Teaching Prize", date="2016")])
    assert sorted(row[0] for row in rows) == [
        "Fictional Research Cup", "Imaginary Service Medal",
        "Pretend Teaching Prize"]


def test_a_list_sharing_one_year_on_its_own_line_keeps_every_award():
    """The years-on-their-own-lines shape (#229): one bare year, before or
    after the awards, can be every award's. Every date the text names is
    the record's and stage 4's award line is undated, so without the
    date-cell guard this read as one award and lost the other line."""
    for text in ("2015\nPretend Medal\nImaginary Prize",
                 "Pretend Medal\nImaginary Prize\n2015"):
        rows = _render_honors([_h_entry(
            text, award_name="Imaginary Prize", date="2015")])
        assert [row[0] for row in rows] == ["Pretend Medal", "Imaginary Prize"]
        assert [row[2] for row in rows] == ["2015", "2015"]


def test_a_list_with_a_row_dated_in_its_own_column_keeps_every_award():
    """A tab- or '|'-joined row whose last cell is its date is an award with
    its own date column (the module's delimiter contract), not a
    continuation line of stage 4's award."""
    for text in ("Pretend Medal\t2015\nImaginary Prize",
                 "Pretend Medal | 2015\nImaginary Prize"):
        rows = _render_honors([_h_entry(
            text, award_name="Imaginary Prize", date="2015")])
        assert [row[0] for row in rows] == ["Pretend Medal", "Imaginary Prize"]


def test_a_pipe_cell_that_is_not_a_date_does_not_keep_one_award_split():
    """A '|' shape the parser does not recognise (a leading date) hands its
    second cell over as a "year". That cell is a note, not a date cell, so
    the award, its note and its citation line stay one row."""
    rows = _render_honors([_h_entry(
        "2012 | Pretend supervisor of the winning entry\n"
        "Pretend Poster Prize, Imaginary Congress\n"
        "Doe J, Roe R. A Made-up Study. Imaginary Congress 2012",
        award_name="Pretend Poster Prize",
        granting_body="Imaginary Congress", date="2012")])
    assert rows == [["Pretend Poster Prize", "Imaginary Congress", "2012"]]


def test_an_entry_with_no_extracted_award_is_never_read_as_one_award():
    """With no stage-4 award there is no line to tie the others to: the
    one-award reading would put the whole raw text, newline and all, into
    one name cell."""
    rows = _render_honors([_h_entry(
        "Pretend Medal\nImaginary Prize 2015", award_name=None, date="2015")])
    assert [row[0] for row in rows] == ["Pretend Medal", "Imaginary Prize"]
    assert rows[1][2] == "2015"


def _award_over_program_line(second_line, **fields):
    record = {"award_name": "Made-up Award for Pretend Outreach",
              "granting_body": "Fictional Agency", "date": "2005", **fields}
    return _render_honors([_h_entry(
        "Fictional Agency Made-up Award for Pretend Outreach (2005)\n" + second_line,
        **record)])


def test_the_dated_program_line_of_an_award_joins_its_name():
    """#1245 (X6 KJJVVO 286): the award's own line is dated, so the entry
    read as a list, and the program the award went to rendered as a second
    award row with an empty organization."""
    assert _award_over_program_line("Imaginary Outreach Program — 02/07/2005") == [
        ["Made-up Award for Pretend Outreach - Imaginary Outreach Program",
         "Fictional Agency", "2005"]]


def test_a_second_line_naming_an_award_noun_stays_its_own_award():
    # The list reading keeps a line's own text, date included, as before.
    rows = _award_over_program_line("Imaginary Service Medal — 02/07/2005")
    assert [row[0] for row in rows] == [
        "Made-up Award for Pretend Outreach", "Imaginary Service Medal — 02/07/2005"]


def test_an_undated_second_line_stays_its_own_row():
    rows = _award_over_program_line("Imaginary Outreach Program")
    assert [row[0] for row in rows] == [
        "Made-up Award for Pretend Outreach", "Imaginary Outreach Program"]


def test_a_second_line_dated_another_year_stays_its_own_row():
    rows = _award_over_program_line("Imaginary Outreach Program — 02/07/2009")
    assert [row[0] for row in rows] == [
        "Made-up Award for Pretend Outreach", "Imaginary Outreach Program — 02/07/2009"]


def test_only_two_lines_one_naming_the_award_with_no_date_cell_join():
    for second_line in ("Imaginary Outreach Program — 02/07/2005\n"
                        "Pretend Clinic Initiative — 03/09/2005",
                        "Imaginary Outreach Program\t2005"):
        rows = _award_over_program_line(second_line)
        assert len(rows) == len(second_line.split("\n")) + 1, second_line
    # Both lines name stage 4's award: one award won twice, two rows.
    rows = _award_over_program_line("Made-up Award for Pretend Outreach — 02/07/2005")
    assert [row[0] for row in rows] == ["Made-up Award for Pretend Outreach"] * 2


def test_a_dated_granting_body_line_is_not_a_program():
    rows = _award_over_program_line("Fictional Agency — 02/07/2005")
    assert all(not row[0].endswith("-") for row in rows)
    assert rows[0] == ["Made-up Award for Pretend Outreach", "Fictional Agency", "2005"]


def test_two_distinct_same_year_honors_stay_separate():
    """PR #1847 review: a second line with no award noun and the record's
    own year is a separate honor unless it ends in a program noun."""
    for second_line in ("Chief Resident — 06/2005",
                        "Dean's List — 2005",
                        "Imaginary Outreach Program Coordinator — 02/07/2005"):
        rows = _award_over_program_line(second_line)
        assert len(rows) == 2, second_line
        assert rows[0][0] == "Made-up Award for Pretend Outreach"
        assert rows[1][0].startswith(second_line.split(" — ")[0]), second_line


def test_a_second_line_that_is_only_a_date_is_not_a_program():
    rows = _award_over_program_line("— 02/07/2005")
    assert rows[0][0] == "Made-up Award for Pretend Outreach"
    assert all(" - " not in row[0] for row in rows)


def test_a_dated_granting_body_ending_in_a_program_noun_is_not_a_program():
    rows = _award_over_program_line("Imaginary Outreach Program — 02/07/2005",
                                    granting_body="Imaginary Outreach Program")
    assert [row[0] for row in rows][0] == "Made-up Award for Pretend Outreach"
    assert len(rows) == 2


@pytest.mark.parametrize("text, award_name", [
    ("Phi Beta Kappa 2005\nUniversity Honors Program — 2005", "Phi Beta Kappa"),
    ("Young Investigator Award, 2005\nNIH Loan Repayment Program, 2005",
     "Young Investigator Award"),
    ("Teaching Award 2005\nSelected for Leadership Development Program, 2005",
     "Teaching Award"),
    ("Best Poster 2005\nResearch Project 2005", "Best Poster"),
])
def test_two_same_year_honors_dated_alike_stay_two_rows(text, award_name):
    """PR #1847 review round 3: a second line ending in a program noun and
    dated with the award's own year is a second honor, as on dev."""
    rows = _render_honors([_h_entry(text, award_name=award_name,
                                    granting_body="", date="2005")])
    assert len(rows) == 2, rows
    assert rows[0][0] == award_name
    assert " - " not in rows[0][0] + rows[1][0]


def test_an_award_line_dated_to_the_day_does_not_take_a_program_line():
    rows = _render_honors([_h_entry(
        "Fictional Agency Made-up Award for Pretend Outreach — 01/03/2005\n"
        "Imaginary Outreach Program — 02/07/2005",
        award_name="Made-up Award for Pretend Outreach",
        granting_body="Fictional Agency", date="2005")])
    assert len(rows) == 2
    assert all(" - " not in row[0] for row in rows)


def test_a_dated_program_line_naming_an_award_noun_stays_its_own_award():
    # The award-noun half: the line ends in a program noun and is dated to
    # the day, so only its "Scholarship" keeps it a second honor.
    rows = _award_over_program_line("Imaginary Scholarship Program — 02/07/2005")
    assert len(rows) == 2
    assert rows[0][0] == "Made-up Award for Pretend Outreach"


def test_a_program_line_led_by_a_year_is_not_a_program():
    # `_PROGRAM_LINE_RE` takes no digit in the program's name.
    for second_line in ("2005 Imaginary Outreach Program 02/07/2005",
                        "2005 Imaginary Outreach Program",
                        "2005 Imaginary Outreach Program — 02/07/2005",
                        "Imaginary 2005 Outreach Program — 02/07/2005"):
        rows = _award_over_program_line(second_line)
        assert len(rows) == 2, second_line
        assert all(" - " not in row[0] for row in rows), second_line


@pytest.mark.parametrize("text, award_name", [
    ("Young Investigator Award, 2005\nNIH Loan Repayment Program, 07/2005",
     "Young Investigator Award"),
    ("Teaching Award 2005\nLeadership Development Program 09/2005", "Teaching Award"),
    ("Teaching Award, 2005\nSelected, Leadership Development Program (06/2005)",
     "Teaching Award"),
])
def test_a_program_line_dated_to_the_month_stays_its_own_honor(text, award_name):
    """PR #1847 review round 4: a month-dated program line under a
    year-dated award is a second honor, as on dev; only the day-dated
    "<Program> — dd/mm/yyyy" shape joins."""
    rows = _render_honors([_h_entry(text, award_name=award_name,
                                    granting_body="", date="2005")])
    assert len(rows) == 2, rows
    assert rows[0][0] == award_name
    assert rows[1][0] == text.split("\n")[1]


@pytest.mark.parametrize("second_line", [
    "Imaginary Outreach Program, 02/07/2005",
    "Imaginary Outreach Program (02/07/2005)",
    "Imaginary Outreach Program 02/07/2005",
    "Selected, Imaginary Outreach Program — 02/07/2005",
    "Imaginary Outreach Program — 07/2005",
    "• Imaginary Outreach Program — 02/07/2005",
])
def test_only_the_dashed_day_dated_program_shape_joins(second_line):
    rows = _award_over_program_line(second_line)
    assert len(rows) == 2, rows
    assert rows[0][0] == "Made-up Award for Pretend Outreach"
    assert all(" - " not in row[0] for row in rows)


def test_a_join_that_would_split_on_a_comma_is_not_made():
    # With no granting body, "Award, University - Program" would re-split
    # into award "Award" and organization "University - Program".
    rows = _render_honors([_h_entry(
        "Made-up Award, Fictional University (2005)\n"
        "Imaginary Outreach Program — 02/07/2005",
        award_name="Made-up Award, Fictional University", granting_body="", date="2005")])
    assert [row[0] for row in rows] == [
        "Made-up Award, Fictional University", "Imaginary Outreach Program — 02/07/2005"]


def test_a_comma_in_the_award_joins_when_the_granting_body_is_known():
    # A known granting body means the name is never re-split, so the
    # comma in stage 4's award name does not block the join.
    rows = _render_honors([_h_entry(
        "Made-up Award, Fictional University (2005)\n"
        "Imaginary Outreach Program — 02/07/2005",
        award_name="Made-up Award, Fictional University",
        granting_body="Fictional Agency", date="2005")])
    assert rows == [["Made-up Award, Fictional University - Imaginary Outreach Program",
                     "Fictional Agency", "2005"]]


def test_an_award_line_ending_in_a_date_cell_does_not_take_a_program_line():
    # The tab-joined award line is a date cell (`_has_date_cell`): its
    # year is that row's own, so the entry is a list.
    rows = _render_honors([_h_entry(
        "Fictional Agency Made-up Award for Pretend Outreach\t2005\n"
        "Imaginary Outreach Program — 02/07/2005",
        award_name="Made-up Award for Pretend Outreach",
        granting_body="Fictional Agency", date="2005")])
    assert [row[0] for row in rows] == [
        "Made-up Award for Pretend Outreach", "Imaginary Outreach Program — 02/07/2005"]


def test_a_fanned_out_child_is_one_row_with_its_own_dates():
    """ZCTARO-04: fan-out gives the first record the built line "Award | Org |
    YYYY-MM | YYYY-MM", which the '|' parser split into four rows, and the
    parent kept the first award's year from the raw text."""
    from unified_pipeline.stage4.schemas import STAGE4_RECORDS_KEY
    first = {"award_name": "Pretend Investigator",
             "granting_body": "Imaginary Foundation, Nowhere",
             "date": None, "start_date": "1998-03", "end_date": "2002-08"}
    last = {"award_name": "Senior Made-up Scholar",
            "granting_body": "Fictional Coalition",
            "date": None, "start_date": "2003-09", "end_date": "2007-05"}
    entry = {"taxonomy_code": "H", "element_idx_start": 7,
             "text": "1998.03 - 2002.08 Pretend Investigator, Imaginary "
                     "Foundation, Nowhere 2003.09 - 2007.05 Senior "
                     "Made-up Scholar, Fictional Coalition",
             "extracted_fields": {**last, STAGE4_RECORDS_KEY: [first, last]}}
    grouped = WCMTemplateGenerator(verbose=False)._group_entries_by_code([entry])
    assert _render_honors(grouped["H"]) == [
        ["Senior Made-up Scholar", "Fictional Coalition", "2003-2007"],
        ["Pretend Investigator", "Imaginary Foundation, Nowhere", "1998-2002"],
    ]



def test_the_last_stage4_record_does_not_reparse_its_siblings_lines():
    """#1556 a (OIEPQD UNKNBU 277): the last of stage 4's records keeps the
    parent's whole text, one award per line. Re-parsed, it printed every
    sibling again: three records rendered as five rows."""
    from unified_pipeline.stage4.schemas import STAGE4_RECORDS_KEY
    records = [{"award_name": name, "granting_body": "Imaginary Society",
                "date": year}
               for name, year in (("Pretend Medal", "2015"),
                                  ("Made-up Prize", "2014"),
                                  ("Fictional Fellowship", "2013"))]
    entry = {"taxonomy_code": "H", "element_idx_start": 7,
             "text": "Pretend Medal, Imaginary Society, 2015\n"
                     "Made-up Prize, Imaginary Society, 2014\n"
                     "Fictional Fellowship, Imaginary Society, 2013",
             "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}
    grouped = WCMTemplateGenerator(verbose=False)._group_entries_by_code([entry])
    assert _render_honors(grouped["H"]) == [
        ["Pretend Medal", "Imaginary Society", "2015"],
        ["Made-up Prize", "Imaginary Society", "2014"],
        ["Fictional Fellowship", "Imaginary Society", "2013"],
    ]


def test_the_last_stage4_record_falls_back_to_its_own_line():
    """#1556 a: a last record with no award name falls back to its own line
    for the name cell, not to the parent's whole text, which would print
    the sibling award inside that cell."""
    from unified_pipeline.stage4.schemas import STAGE4_RECORDS_KEY
    records = [{"award_name": "Pretend Medal",
                "granting_body": "Imaginary Society", "date": "2015"},
               {"award_name": None, "granting_body": "Fictional Board",
                "date": "2013"}]
    entry = {"taxonomy_code": "H", "element_idx_start": 7,
             "text": "Pretend Medal, Imaginary Society, 2015\n"
                     "Fictional Board, 2013",
             "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}
    grouped = WCMTemplateGenerator(verbose=False)._group_entries_by_code([entry])
    rows = _render_honors(grouped["H"])
    assert len(rows) == 2
    assert [row for row in rows if "Pretend Medal" in " ".join(row)] == [
        ["Pretend Medal", "Imaginary Society", "2015"]]

def test_a_date_only_fragment_keeps_an_empty_name_cell():
    """A stage-2 date continuation coded H: the range now comes from stage
    4's fields, and the fragment is still peeled off the name cell rather
    than repeated there."""
    rows = _render_honors([_h_entry("1971- 1972.", award_name=None,
                                    start_date="1971", end_date="1972")])
    assert rows == [["", "", "1971-1972"]]
