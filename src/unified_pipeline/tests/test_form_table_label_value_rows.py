"""#811: a 2-column `Label: | value` sub-header row is content, not a header.

`extract_unified_elements`'s per-row walk classifies a table row as a header
whenever its first cell scores as header-like (`looks_like_section_header`,
e.g. a short colon-terminated label like "Name:"). That branch used to read
only the first cell and silently drop every other cell in the row -- so a
form-style label|value row (personal-data blocks: Name/Office Address/Work
Email) lost its value column entirely, which emptied `cv_owner` downstream
(web172, web207 in the 2026-09-11 corpus batch).

A row is now content (not a header) only when it has more than one cell AND
at least one trailing cell carries non-blank text. A single-cell header row,
or a multi-cell row whose trailing cells are all blank, is unaffected.

Deterministic, no LLM, no DB, no network, no PII (synthetic values only).
Run with:

    python3 -m pytest src/unified_pipeline/tests/test_form_table_label_value_rows.py -p no:cacheprovider
"""
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[1]  # src/unified_pipeline
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from core.docx_structure_extractor import extract_unified_elements  # noqa: E402


def _build_form_table_docx(path: str) -> None:
    """One table: a merged single-cell header row ('PUBLICATIONS') above
    three 2-column sub-header-shaped rows -- two form rows (label + non-blank
    value) and one row whose value cell is blank."""
    doc = Document()
    table = doc.add_table(rows=4, cols=2)

    header_cell = table.cell(0, 0).merge(table.cell(0, 1))
    header_cell.text = "PUBLICATIONS"

    table.cell(1, 0).text = "Name:"
    table.cell(1, 1).text = "Synthetic Person"

    table.cell(2, 0).text = "Office Address:"
    table.cell(2, 1).text = "123 Synthetic Ave"

    table.cell(3, 0).text = "Certification:"
    table.cell(3, 1).text = ""  # blank value cell -- negative case

    doc.save(path)


def test_form_label_value_rows_survive_as_table_content(tmp_path):
    """Positive case: 'Name:' | value and 'Office Address:' | value are each
    header-like in cell 0 but carry a non-blank value cell, so both cells
    must survive together inside a table_content element."""
    docx_path = str(tmp_path / "form_table.docx")
    _build_form_table_docx(docx_path)

    elements = extract_unified_elements(docx_path)["elements"]
    table_contents = [e for e in elements if e.get("type") == "table_content"]
    content_text = "\n".join(e["text"] for e in table_contents)

    assert "Name:" in content_text
    assert "Synthetic Person" in content_text
    assert "Office Address:" in content_text
    assert "123 Synthetic Ave" in content_text


def test_single_cell_header_row_is_still_table_header(tmp_path):
    """Positive case: the merged single-cell 'PUBLICATIONS' row (nothing in
    row[1:] to lose) keeps today's table_header behaviour unchanged."""
    docx_path = str(tmp_path / "form_table.docx")
    _build_form_table_docx(docx_path)

    elements = extract_unified_elements(docx_path)["elements"]
    table_headers = [e for e in elements if e.get("type") == "table_header"]

    assert any(e["text"] == "PUBLICATIONS" for e in table_headers)


def test_row_header_with_blank_value_cell_stays_table_header(tmp_path):
    """Negative case: a 2-column row whose second cell is BLANK ('Certification:'
    | '') keeps today's table_header behaviour -- only a non-blank trailing
    cell should flip a label row into content."""
    docx_path = str(tmp_path / "form_table.docx")
    _build_form_table_docx(docx_path)

    elements = extract_unified_elements(docx_path)["elements"]
    table_headers = [e for e in elements if e.get("type") == "table_header"]
    table_contents = [e for e in elements if e.get("type") == "table_content"]
    content_text = "\n".join(e["text"] for e in table_contents)

    assert any(e["text"] == "Certification:" for e in table_headers)
    assert "Certification:" not in content_text
