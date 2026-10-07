"""#208: single-column layout tables must explode into per-paragraph elements;
multi-column (real data) tables must NOT. Deterministic, no LLM."""
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[1]  # src/unified_pipeline
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))
from core.docx_structure_extractor import extract_unified_elements  # noqa: E402
from docx import Document  # noqa: E402


def _build_docx(path):
    doc = Document()
    doc.add_paragraph("TEACHING ACTIVITIES")
    # 1-cell LAYOUT box fusing 3 paragraphs (the #208 shape)
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    cell.paragraphs[0].text = "Uniformed Services University"
    cell.add_paragraph("Workshop on assessment methods")
    cell.add_paragraph("Korea University College of Medicine")
    # real 2-column DATA table
    d = doc.add_table(rows=2, cols=2)
    d.rows[0].cells[0].text = "2020"; d.rows[0].cells[1].text = "PhD"
    d.rows[1].cells[0].text = "2015"; d.rows[1].cells[1].text = "MS"
    doc.save(path)


def _texts(path):
    els = extract_unified_elements(path)["elements"]
    return [e["text"] for e in els if e.get("text", "").strip()]


def test_layout_table_explodes_but_data_table_stays(tmp_path):
    p = str(tmp_path / "cv.docx")
    _build_docx(p)
    texts = _texts(p)

    # each layout paragraph survives as its OWN line (no fused blob)
    assert "Uniformed Services University" in texts
    assert "Workshop on assessment methods" in texts
    assert "Korea University College of Medicine" in texts
    # nothing fused: no single element holds two of the layout paragraphs
    assert not any("Uniformed" in t and "Korea" in t for t in texts)

    # multi-column data table is NOT exploded -- its cells stay joined together
    assert any("2020" in t and "PhD" in t for t in texts)
    assert any("2015" in t and "MS" in t for t in texts)



def _build_gridspan_docx(path):
    """#1229: the same layout box, but on a 2-column grid with its one cell
    gridSpan-merged across both columns (python-docx: len(row.cells) == 2)."""
    doc = Document()
    box = doc.add_table(rows=1, cols=2)
    cell = box.rows[0].cells[0].merge(box.rows[0].cells[1])
    cell.paragraphs[0].text = "TEACHING ACTIVITIES"
    cell.add_paragraph("Lecture on cardiac physiology")
    cell.add_paragraph("Seminar on clinical reasoning")
    # a merged title row over a real 2-cell data row is still a data table
    data = doc.add_table(rows=2, cols=2)
    data.rows[0].cells[0].merge(data.rows[0].cells[1]).text = "Degrees"
    data.rows[1].cells[0].text = "2019"
    data.rows[1].cells[1].text = "MD"
    doc.save(path)


def test_gridspan_merged_layout_box_explodes_like_unmerged(tmp_path):
    p = str(tmp_path / "cv.docx")
    _build_gridspan_docx(p)
    texts = _texts(p)

    # every box paragraph is its own element, once (no " | " copy of the cell)
    for line in ("TEACHING ACTIVITIES", "Lecture on cardiac physiology",
                 "Seminar on clinical reasoning"):
        assert texts.count(line) == 1, (line, texts)
    assert not any("Lecture" in t and "Seminar" in t for t in texts)

    # the table with a 2-cell row keeps the joined data path
    assert any("2019" in t and "MD" in t for t in texts)


def test_zero_row_table_is_not_a_layout_box(tmp_path):
    # all() over no rows is True; an empty table keeps the table path
    p = str(tmp_path / "cv.docx")
    doc = Document()
    doc.add_table(rows=0, cols=2)
    doc.save(p)
    els = extract_unified_elements(p)["elements"]
    assert [e["type"] for e in els] == ["table"]

if __name__ == "__main__":   # ponytail: runnable without pytest
    import os
    import tempfile
    p = os.path.join(tempfile.mkdtemp(), "cv.docx")
    _build_docx(p)
    texts = _texts(p)
    assert "Workshop on assessment methods" in texts
    assert not any("Uniformed" in t and "Korea" in t for t in texts)
    assert any("2020" in t and "PhD" in t for t in texts)
    print("OK:", len(texts), "lines ->", texts)
