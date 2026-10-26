"""Supplementary prose as tracked-deleted sub-points (#1205), stage6/supplementary.py.

What is offered (prose the document provably lacks and that holds no rendered
field value), where it goes (under the entry's own line, in the entry's own
section), and the markup: a deletion whose paragraph mark or row is deleted
too, so accepting all changes leaves the document as if the pass never ran,
and rejecting them shows the prose under its line. All fixture text is
synthetic.

    python3 -m pytest src/unified_pipeline/tests/test_stage6_supplementary.py -p no:cacheprovider
"""

import copy
import json
import sys
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Twips  # noqa: E402

from unified_pipeline.stage6 import supplementary as sp  # noqa: E402
from unified_pipeline.stage6.fan_out import (  # noqa: E402
    FANNED_OUT_FROM,
    LAST_STAGE4_RECORD,
)
from unified_pipeline.stage6.formatting import (  # noqa: E402
    is_entry_end,
    mark_entry_end,
)

_DATE = "2026-10-08T00:00:00Z"

PROSE = ("Responsible for coordinating the departmental curriculum committee "
         "and supervising graduate trainees")
OTHER_PROSE = "Developed laboratory teaching materials for undergraduate physiology courses"


def _haystack(*lines):
    return sp._Haystack("\x00".join(sp._squash(line) for line in lines),
                        [sp._tokens(line) for line in lines])


def _entry(text, code="D1", **fields):
    return {"taxonomy_code": code, "text": text, "extracted_fields": fields}


def _appointment(text=f"Associate Professor\tExample University\t{PROSE}"):
    return _entry(text, title="Associate Professor", institution="Example University",
                  start_date="2010", end_date="2015")


def _deleted_marks(element):
    return [d for d in element.iter(qn("w:del")) if d.get(qn("w:author")) == sp.SUBPOINT_AUTHOR]


# --- the switch ---------------------------------------------------------------

@pytest.mark.parametrize("value, on", [("1", True), (" 1 ", True), ("", True), (None, True),
                                       ("true", True), ("yes", True), (1, True), (True, True),
                                       ("0", False), (" 0 ", False), ("false", False),
                                       ("No", False), ("OFF", False), (0, False), (False, False)])
def test_the_flag_is_on_unless_turned_off(value, on):
    # On by default since the owner's 2026-10-09 review (#1205); "0" is the
    # ops switch, and an unquoted YAML false/no/0 reads as off too.
    assert sp.SUBPOINT_FLAG_DEFAULT == "1"
    assert sp.subpoints_enabled(value) is on


def test_the_codes_are_teaching_grants_appointments_service_honors_mentees_degrees_leadership_committees_talks_training_memberships_clinical_leadership_editorial_boards_and_publications():
    assert sorted(sp.SUBPOINT_CODES) == ["B1", "C", "D1", "D2", "D3", "H", "I", "K1", "K2", "K3",
                                         "K4", "K5", "L3", "M2A", "M2B", "M2C", "M2D", "N3A",
                                         "N3B", "O", "P", "Q1", "Q2", "Q4B", "Q4C", "R", "S1", "S2"]
    assert set(sp.PUBLICATION_SUBPOINT_CODES) == {"S1", "S2"}
    # S3-S9 share the section and are not offered.
    assert not sp.SUBPOINT_CODES & {"S0", "S3", "S4", "S5", "S6", "S7", "S8", "S9"}
    # The owner decision's exclusions stay out.
    assert not sp.SUBPOINT_CODES & {"A", "T", "M1", "E", "G", "J", "N1", "N4", "S0", "L1",
                                    "L2", "F1", "Q4D", "B2"}


# --- what is prose --------------------------------------------------------------

def test_prose_fragments_split_on_tabs_newlines_and_pipes():
    text = f"Associate Professor | Example University\n2010-2015\t{PROSE}"
    assert sp.prose_fragments(text) == [PROSE]


@pytest.mark.parametrize("fragment", [
    "Associate Professor",                                 # under SUBPOINT_MIN_TOKENS
    "2010 - 2015 Associate Professor of Medicine, Example University Hospital",  # a record
    "May 2010 - present Visiting Professor, Example Institute of Science",
    'Email: HYPERLINK "mailto:someone@example.org" someone@example.org',
    "Contact the example programme office at someone@example.org for details",
    "Recording available at https://example.org/webinar/example-session-recording",
    "Recording available at www.example.org/webinar/example-session-recording",
])
def test_short_lines_dated_record_lines_and_contact_lines_are_not_prose(fragment):
    assert sp.prose_fragments(fragment) == []


def test_a_paragraph_broken_at_printed_lines_is_offered_whole():
    lines = ["I taught the anatomy laboratory for the", "graduate physiology students of the",
             "department and wrote the course manual."]
    assert sp.prose_fragments("\t".join(lines)) == [" ".join(lines)]


def test_a_fragment_mostly_holding_field_values_is_the_record_not_prose():
    entry = _entry("Example Foundation, Principal Investigator 2005-2007\t" + PROSE, code="M2B",
                   agency="Example Foundation", pi_role="Principal Investigator")
    assert sp.unrendered_prose(entry, "M2B", _haystack("Unrelated rendered line")) == [PROSE]


def test_prose_stage_4_put_in_a_field_no_renderer_reads_is_still_offered():
    # `duties` is no D1 rendered field (`_RENDERED_FIELDS`), so holding the
    # prose there does not make it the record: it reaches the page nowhere.
    entry = _appointment()
    entry["extracted_fields"]["duties"] = PROSE
    assert sp.unrendered_prose(entry, "D1", _haystack("Associate Professor | Example University")) \
        == [PROSE]


def test_a_rendered_fragment_is_not_offered():
    assert sp.unrendered_prose(_appointment(), "D1", _haystack(PROSE)) == []


def test_an_absent_fragment_is_offered():
    assert sp.unrendered_prose(_appointment(), "D1",
                               _haystack("Associate Professor | Example University")) == [PROSE]


def test_a_code_rendered_from_its_text_judges_every_scalar_field():
    # K2 has no `_RENDERED_FIELDS` row: all scalar fields count as values.
    entry = _entry("Clinical Preceptor, Example Teaching Hospital, Medical Students\t" + PROSE,
                   code="K2", role="Clinical Preceptor", institution="Example Teaching Hospital",
                   learners="Medical Students")
    assert sp.unrendered_prose(entry, "K2", _haystack("Unrelated rendered line")) == [PROSE]


def test_a_label_table_vouches_for_a_line_whose_facts_sit_in_its_rows():
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row, (label, value) in zip(table.rows, [("Agency:", "Example Health Institute"),
                                                ("Role:", "Principal Investigator"),
                                                ("Title:", "Example Pathways Research")],
                                   strict=True):
        row.cells[0].text, row.cells[1].text = label, value
    line = "Example Health Institute Research Award, Principal Investigator 1990-1994"
    rows = ["Agency: | Example Health Institute", "Role: | Principal Investigator"]
    haystack = sp.rendered_haystack(rows, doc)
    assert sp._record_rendered(line, haystack.text, haystack.line_tokens) is True
    plain = sp._Haystack(haystack.text, [sp._tokens(row) for row in rows])
    assert sp._record_rendered(line, plain.text, plain.line_tokens) is False


# --- where it goes -----------------------------------------------------------------

def _candidate(text, position=0, paragraph=None):
    return sp.AnchorCandidate(paragraph, text, sp._tokens(text), position)


def test_best_anchor_takes_the_line_holding_most_of_the_entry():
    tokens = sp._tokens("Associate Professor Example University")
    lines = [_candidate("Associate Professor of Surgery"),
             _candidate("Associate Professor | Example University | 2010-2015"),
             _candidate("Example University Hospital")]
    assert sp.best_anchor(tokens, lines) is lines[1]


def test_best_anchor_keeps_the_first_of_two_equal_lines():
    tokens = sp._tokens("Associate Professor Example University")
    lines = [_candidate("Associate Professor Example University"),
             _candidate("Associate Professor Example University")]
    assert sp.best_anchor(tokens, lines) is lines[0]


def test_best_anchor_needs_half_the_tokens():
    tokens = sp._tokens("Associate Professor Example University Hospital")
    assert sp.best_anchor(tokens, [_candidate("Associate Professor")]) is None


def test_best_anchor_needs_two_shared_tokens():
    # One token is all of the entry's, but one word is no anchor (an owner's
    # surname matched the Personal Data name line).
    assert sp.best_anchor({"examplename"}, [_candidate("Name: Examplename")]) is None
    assert sp.SUBPOINT_ANCHOR_MIN_TOKENS == 2


def test_best_anchor_needs_two_shared_tokens_even_at_half_the_entry():
    tokens = sp._tokens("Associate Examplename")
    assert sp.best_anchor(tokens, [_candidate("Name: Examplename")]) is None
    assert sp.best_anchor(tokens, [_candidate("Associate Examplename")]) is not None


def test_plan_places_an_entry_under_its_own_line_in_its_own_section():
    lines = [_candidate("Associate Professor | Example University | 2010-2015", position=5),
             _candidate("Associate Professor Example University, elsewhere", position=50)]
    planned = sp.plan_subpoints({"D1": [_appointment()]}, _haystack(lines[0].text), lines,
                                {"D1": sp.SectionSpan(1, 10)})
    assert [(p.anchor, p.paragraphs) for p in planned] == [(lines[0], (PROSE,))]


def test_plan_records_the_code_each_entry_is_planned_under_and_its_index():
    # The group an entry is planned under is the code it rendered as: a K2
    # rerouted into N3B keeps its stored code, but rendered as a mentee.
    lines = _mentee_tables("Ana Li")
    trainee = dict(_postdoc(), taxonomy_code="K2", element_idx_start=1477)
    (planned,) = sp.plan_subpoints({"N3B": [trainee]}, _haystack(), lines,
                                   {"N3B": sp.SectionSpan(0, None)})
    assert (planned.code, planned.element_idx_start) == ("N3B", 1477)


def test_a_placement_reads_back_from_its_sidecar_record():
    placement = sp.SubpointPlacement(1477, "N3B", (PROSE, OTHER_PROSE))
    assert sp.SubpointPlacement.from_record(json.loads(json.dumps(placement._asdict()))) == placement
    # A record missing a key, or holding a stray shape, reads as empty, not as text.
    assert sp.SubpointPlacement.from_record({"paragraphs": PROSE}) == sp.SubpointPlacement(None, "", ())
    assert sp.SubpointPlacement.from_record({"paragraphs": [7]}).paragraphs == ("7",)


def test_plan_finds_no_anchor_outside_the_section():
    line = _candidate("Associate Professor | Example University | 2010-2015", position=50)
    assert sp.plan_subpoints({"D1": [_appointment()]}, _haystack(line.text), [line],
                             {"D1": sp.SectionSpan(1, 10)}) == []


def test_plan_skips_a_code_with_no_section_and_a_code_out_of_scope():
    line = _candidate("Associate Professor | Example University | 2010-2015", position=5)
    entry = _appointment()
    assert sp.plan_subpoints({"D1": [entry]}, _haystack(line.text), [line], {}) == []
    assert sp.plan_subpoints({"Q4D": [{**entry, "taxonomy_code": "Q4D"}]}, _haystack(line.text),
                             [line], {"Q4D": sp.SectionSpan(0, None)}) == []


def test_plan_offers_a_fragment_once():
    line = _candidate("Associate Professor | Example University | 2010-2015", position=5)
    planned = sp.plan_subpoints({"D1": [_appointment(), _appointment()]}, _haystack(line.text),
                                [line], {"D1": sp.SectionSpan(0, None)})
    assert len(planned) == 1


def test_plan_skips_an_entry_whose_line_is_a_row_of_an_excluded_table():
    doc = Document()
    excluded, other = doc.add_table(rows=1, cols=3), doc.add_table(rows=1, cols=3)
    for table in (excluded, other):
        for cell, text in zip(table.rows[0].cells, ["Associate Professor", "Example University",
                                                    "2010-2015"], strict=True):
            cell.text = text
    (in_excluded, elsewhere) = [_candidate(c.text, position=i, paragraph=c.paragraph)
                                for i, c in enumerate(sp.anchor_candidates(doc))]
    args = ({"D1": [_appointment(), _appointment()]}, _haystack(in_excluded.text),
            [in_excluded, elsewhere], {"D1": sp.SectionSpan(0, None)})
    assert len(sp.plan_subpoints(*args)) == 1
    assert sp.plan_subpoints(*args, excluded_tables=[excluded._tbl]) == []
    (planned,) = sp.plan_subpoints(*args, excluded_tables=[other._tbl])
    assert planned.anchor is in_excluded


def _mentee_tables(*names):
    """Anchor candidates of one two-row mentee table per name, each table at
    the same level."""
    doc = Document()
    for name in names:
        table = doc.add_table(rows=2, cols=2)
        for row, (label, value) in zip(table.rows, [("Name:", name),
                                                    ("Site/Position:", "Postdoctoral Associate")],
                                       strict=True):
            row.cells[0].text, row.cells[1].text = label, value
    return list(sp.anchor_candidates(doc))


def _postdoc(name="Ana Li"):
    return _entry(f"{name}, Postdoctoral Associate\t{PROSE}", code="N3B", mentee_name=name,
                  mentee_level="Postdoctoral Associate")


def test_plan_places_a_mentee_only_in_the_table_naming_it():
    # The first table's level row holds every token the entry renders: only
    # its Name: row says it is another mentee's.
    lines = _mentee_tables("Casey Sample", "Ana Li")
    spans = {"N3B": sp.SectionSpan(0, None)}
    (planned,) = sp.plan_subpoints({"N3B": [_postdoc()]}, _haystack(), lines, spans)
    assert planned.anchor is lines[3]
    # Stage 4's spacing is not the cell's: the row is read with its spaces collapsed.
    (planned,) = sp.plan_subpoints({"N3B": [_postdoc(" Ana \n Li ")]}, _haystack(), lines, spans)
    assert planned.anchor is lines[3]
    # `name` is read before `mentee_name`, as the section reads it.
    by_name = _postdoc("Casey Sample")
    by_name["extracted_fields"]["name"] = "Ana Li"
    (planned,) = sp.plan_subpoints({"N3B": [by_name]}, _haystack(), lines, spans)
    assert planned.anchor is lines[3]
    for nobody in (_postdoc("Jordan Example"), {**_postdoc(), "extracted_fields": None}):
        assert sp.plan_subpoints({"N3B": [nobody]}, _haystack(), lines, spans) == []


def test_plan_places_a_mentee_naming_no_one_in_no_table():
    # Not even one whose Name: row is blank.
    lines = _mentee_tables("")
    nameless = _postdoc()
    del nameless["extracted_fields"]["mentee_name"]
    assert sp.plan_subpoints({"N3B": [nameless]}, _haystack(), lines,
                             {"N3B": sp.SectionSpan(0, None)}) == []


def _record_of_line(name, text, index, count):
    """Record `index` of `count` that fan-out made of one source line: the
    last keeps the whole line (`LAST_STAGE4_RECORD`), the others their own."""
    entry = dict(_postdoc(name), text=text, element_idx_start=7)
    marker = {"key": "stage4_records", "index": index, "count": count}
    if index == count - 1:
        entry[LAST_STAGE4_RECORD] = {**marker, "text": None}
    else:
        entry[FANNED_OUT_FROM] = marker
    return entry


def test_plan_reads_the_last_record_of_a_paired_line_from_its_own_segment():
    # The first record's prose is withheld with it (#892): the last record,
    # which keeps the whole line, does not offer it under its own table.
    first_line, last_line = f"Ana Li, Postdoctoral Associate, {PROSE}", "Casey Sample, Postdoc"
    first = _record_of_line("Ana Li", first_line, 0, 2)
    first["_pii_dropped_fields"] = ["site_position"]
    first["_pii_withheld"] = True  # the pass sets both
    lines = _mentee_tables("Ana Li", "Casey Sample")
    spans = {"N3B": sp.SectionSpan(0, None)}
    last = _record_of_line("Casey Sample", f"{first_line}\t{last_line}", 1, 2)
    assert sp.plan_subpoints({"N3B": [first, last]}, _haystack(), lines, spans) == []
    # Prose in its own segment, the line's last, is its own.
    last = _record_of_line("Casey Sample", f"{first_line}\t{last_line}, {OTHER_PROSE}", 1, 2)
    (planned,) = sp.plan_subpoints({"N3B": [first, last]}, _haystack(), lines, spans)
    assert planned.anchor is lines[2] and planned.paragraphs == (f"{last_line}, {OTHER_PROSE}",)


def _degree_rows(*rows):
    """A line naming the field and institution, then a degrees table, one
    row of cells per item; and their anchor candidates."""
    doc = Document()
    doc.add_paragraph("Example Marine Studies at Norvale University")
    table = doc.add_table(rows=len(rows), cols=3)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _master(text=f"M.S., Example Marine Studies, Norvale University, 2001\t{PROSE}"):
    return _entry(text, code="B1", degree="M.S.", discipline="Example Marine Studies",
                  institution="Norvale University", year="2001")


def test_plan_places_a_degree_only_under_a_row_written_for_it():
    # The line and the Ph.D. row hold every distinctive word the M.S. row
    # does: words alone take the first, and only the writer's record says
    # which row is whose.
    table, lines = _degree_rows(["Ph.D., Example Marine Studies", "Norvale University", "2005"],
                                ["M.S., Example Marine Studies", "Norvale University", "2001"])
    entry, spans = _master(), {"B1": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "B1"), lines) is lines[0]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    # The education writer records the row's first-cell paragraph.
    first_cell = table.rows[1].cells[0].paragraphs[0]._p
    for recorded in ([table.rows[1]._tr], [first_cell]):
        (planned,) = sp.plan_subpoints({"B1": [entry]}, _haystack(), lines, spans,
                                       rows_written=written(recorded))
        assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    # No row recorded for it, a paragraph outside any table, or no ledger: none.
    stray = Document().add_paragraph("M.S., Example Marine Studies")._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({"B1": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_degrees_one_own_row_whatever_its_words():
    # "PhD", a four-letter name and a year hold no distinctive word: the
    # word test finds no anchor, and the writer's record alone says the row
    # is the entry's (#1205 review: the prose reached the document in
    # neither arm).
    table, lines = _degree_rows(["PhD", "Vaux", "2010"])
    entry = _entry(f"PhD, Vaux, 2010\t{PROSE}", code="B1", degree="PhD", institution="Vaux",
                   year="2010")
    spans = {"B1": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "B1"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"B1": [entry]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)
    # Two rows recorded for one entry: its words choose between them.
    table, lines = _degree_rows(["M.S.", "Norvale University", "2001"],
                                ["M.S., Example Marine Studies", "Norvale University", "2001"])
    (planned,) = sp.plan_subpoints({"B1": [_master()]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [row._tr for row in table.rows])
    assert planned.anchor is lines[2]
    # Any other code still needs the words, with one line in scope or more.
    line = _candidate("Unrelated rendered line", position=0)
    assert sp.plan_subpoints({"D1": [_appointment()]}, _haystack(), [line],
                             {"D1": sp.SectionSpan(0, None)}) == []


@pytest.mark.parametrize("fields, planned", [
    ({"degree": "M.S."}, 1),
    # Its row is a field or an institution alone (EBYSBC VGHNZD 1861: a
    # thesis citation coded B1): its prose is of no degree there.
    ({"degree": None}, 0),
    ({"degree": "  "}, 0),
    ({}, 0),
])
def test_plan_gives_an_entry_naming_no_degree_no_sub_point(fields, planned):
    table, lines = _degree_rows(["Example Marine Studies", "Norvale University", "2001"])
    entry = _master()
    entry["extracted_fields"] = {**{k: v for k, v in entry["extracted_fields"].items()
                                    if k != "degree"}, **fields}
    assert len(sp.plan_subpoints({"B1": [entry]}, _haystack(), lines,
                                 {"B1": sp.SectionSpan(0, None)},
                                 rows_written=lambda e: [table.rows[0]._tr])) == planned
    # Fields that are not a mapping name no degree either.
    entry["extracted_fields"] = ["M.S."]
    assert not sp._names_degree(entry)


@pytest.mark.parametrize("field", sp._DEGREE_CELL_FIELDS)
def test_a_degrees_cell_fields_count_as_its_rendered_values(field):
    # The degree's cell prints them, though `_RENDERED_FIELDS` cannot list
    # them (no B1 schema field is any of them). Another code's are its own.
    entry = _entry("M.S.", code="B1", degree="M.S.", **{field: "Estuarine Sediment Geochemistry"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "B1")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="H"), "H")


def _degree_record(text, index, count):
    """Record `index` of `count` that fan-out made of one source line."""
    entry = dict(_master(text), element_idx_start=7)
    marker = {"key": "stage4_records", "index": index, "count": count}
    if index == count - 1:
        entry[LAST_STAGE4_RECORD] = {**marker, "text": None}
    else:
        entry[FANNED_OUT_FROM] = marker
    return entry


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Ph.D. line", f"M.S. line, {PROSE}"], 1, f"M.S. line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Ph.D. line", PROSE, "M.S. line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Ph.D. line, {PROSE}", "M.S. line"], 0, f"Ph.D. line, {PROSE}\tM.S. line"),
])
def test_a_degree_reads_its_prose_from_its_own_line(segments, index, read):
    entry = _degree_record("\t".join(segments), index, 2)
    assert sp._own_record_line(entry) == read
    assert sp._prose_text("B1", entry, {}) == read
    assert sp._prose_text("D1", entry, {}) is None


_ROLE = "Director, Example Outreach Programme"


def _role_rows(*rows):
    """A line naming the role and institution, then a leadership table, one
    row of cells per item; and their anchor candidates."""
    doc = Document()
    doc.add_paragraph(f"{_ROLE} at Norvale University")
    table = doc.add_table(rows=len(rows), cols=3)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _leadership_role(text=f"2008-2012\t{_ROLE}, Norvale University\t{PROSE}", **fields):
    return _entry(text, code="O", **{"leadership_role": _ROLE, "institution": "Norvale University",
                                     "start_date": "2008", "end_date": "2012", **fields})


def test_plan_places_a_leadership_role_only_under_a_row_written_for_it():
    # One title held twice at one institution: the rows differ only in their
    # dates, which are digits, and words alone take the first line holding them.
    table, lines = _role_rows([_ROLE, "Norvale University", "2016-2020"],
                              [_ROLE, "Norvale University", "2008-2012"])
    entry, spans = _leadership_role(), {"O": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "O"), lines) is lines[0]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({"O": [entry]}, _haystack(), lines, spans,
                                   rows_written=written([table.rows[1]._tr]))
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == ("O", id(entry))
    # No row recorded for it, a paragraph outside any table, or no ledger: none.
    stray = Document().add_paragraph(_ROLE)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({"O": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_roles_one_own_row_whatever_its_words():
    # "Dean" and an acronym hold no distinctive word: the word test finds no
    # anchor, and the writer's record alone says the row is the entry's.
    table, lines = _role_rows(["Dean", "NVU", "2010"])
    entry = _entry(f"2010\tDean, NVU\t{PROSE}", code="O", leadership_role="Dean",
                   institution="NVU", start_date="2010")
    spans = {"O": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "O"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"O": [entry]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)
    # Several rows recorded for one entry (a flattened table the writer split):
    # its words choose between them.
    table, lines = _role_rows(["Chair, Example Admissions Committee", "Norvale University", "2001"],
                              [_ROLE, "Norvale University", "2008-2012"])
    (planned,) = sp.plan_subpoints({"O": [_leadership_role()]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [row._tr for row in table.rows])
    assert planned.anchor is lines[2]


@pytest.mark.parametrize("field", sp._LEADERSHIP_ROW_FIELDS)
def test_a_leadership_rows_fallback_fields_count_as_its_rendered_values(field):
    # The row prints them when the schema's role or institution is empty,
    # though `_RENDERED_FIELDS` cannot list them. Another code's are its own.
    entry = _entry("Dean", code="O", **{field: "Estuarine Sediment Geochemistry"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "O")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="H"), "H")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="B1"), "B1")


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Chair line", f"Director line, {PROSE}"], 1, f"Director line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Chair line", PROSE, "Director line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Chair line, {PROSE}", "Director line"], 0, f"Chair line, {PROSE}\tDirector line"),
])
def test_a_leadership_role_reads_its_prose_from_its_own_line(segments, index, read):
    entry = _degree_record("\t".join(segments), index, 2)
    entry["taxonomy_code"] = "O"
    assert sp._prose_text("O", entry, {}) == read


_CHAIR = "Chair, Example Admissions Committee"


@pytest.mark.parametrize("text, records", [
    # A bare date column is a neighbouring part's date, not a record.
    (f"2008-2012\t{_ROLE}, Norvale University\t{PROSE}", False),
    (f"2005-2010 {_ROLE}, Norvale University\t{PROSE}", False),
    (f"2001-2004 {_CHAIR}, Norvale University\t{OTHER_PROSE}\n"
     f"2005-2010 {_ROLE}, Norvale University\t{PROSE}", True),
    (f"2001 MS, Example Marine Studies\t{OTHER_PROSE}\t2005 PhD, Example Marine Studies", True),
], ids=["date_column", "one_record", "two_roles", "two_degrees"])
def test_a_line_holds_dated_records_when_two_parts_open_with_their_dates(text, records):
    assert sp._holds_dated_records(text) is records


@pytest.mark.parametrize("code, first, second, fields, row", [
    # Two roles on one line, each opening with its dates and followed by its
    # prose; stage 4 returned the second (#1205 review, probe P12).
    ("O", f"2001-2004 {_CHAIR}, Norvale University", f"2005-2010 {_ROLE}, Norvale University",
     {"leadership_role": _ROLE, "institution": "Norvale University", "start_date": "2005",
      "end_date": "2010"},
     [_ROLE, "Norvale University", "2005-2010"]),
    # Two degrees, the same way.
    ("B1", "2001 MS, Example Marine Studies, Norvale University",
     "2005 PhD, Example Marine Studies, Norvale University",
     {"degree": "PhD", "discipline": "Example Marine Studies", "institution": "Norvale University",
      "year": "2005"},
     ["PhD, Example Marine Studies", "Norvale University", "2005"]),
], ids=["O", "B1"])
def test_plan_offers_no_prose_from_a_line_of_two_dated_records(code, first, second, fields, row):
    # Which record a prose line is about cannot be told: the first record's
    # would go under the second's row, with the second's own.
    table, lines = _role_rows(row)
    spans = {code: sp.SectionSpan(0, None)}

    def plan(text):
        return sp.plan_subpoints({code: [_entry(text, code=code, **fields)]}, _haystack(), lines,
                                 spans, rows_written=lambda e: [table.rows[0]._tr])

    (planned,) = plan(f"{second}\t{PROSE}")
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)
    two_records = f"{first}\t{OTHER_PROSE}\t{second}\t{PROSE}"
    assert sp._own_record_line(_entry(two_records, code=code, **fields)) == ""
    assert plan(two_records) == []


_SPOUSE = "Zelphine Qorvath"


@pytest.mark.parametrize("protect", [
    # The PII pass dropped a field named for protected data ("spouse_name"):
    # it marks the entry withheld and lists no `_pii_dropped_fields`, and the
    # prose names the spouse with no label (#1205 review, probe P7).
    lambda e: e.update(_pii_withheld=True, text=e["text"] + f", with my spouse {_SPOUSE}"),
    # Prose the full Personal Data / Appendix scope withholds, which a routed
    # code's scope leaves in the text.
    lambda e: e.update(text=e["text"] + f"\tMarried ({_SPOUSE}), co-founded the outreach clinics"),
    # A phone number in the middle of a line, which no scope withholds (probe P9).
    lambda e: e.update(text=e["text"] + ". Home phone: 212-555-0147"),
], ids=["withheld", "full_scope", "home_phone"])
def test_plan_offers_no_prose_of_an_entry_holding_protected_data(protect):
    table, lines = _role_rows([_ROLE, "Norvale University", "2008-2012"])
    spans = {"O": sp.SectionSpan(0, None)}
    entry = _leadership_role()
    (planned,) = sp.plan_subpoints({"O": [entry]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.paragraphs == (PROSE,)

    protect(entry)
    assert sp.plan_subpoints({"O": [entry]}, _haystack(), lines, spans,
                             rows_written=lambda e: [table.rows[0]._tr]) == []


_COMMITTEE = "Example Curriculum Steering Committee"
_COMMITTEE_ROW = [f"{_COMMITTEE}, Norvale University", "Member", "2008-2012"]


def _committee_rows(*rows):
    """A line naming the committee and institution, then a committee table,
    one row of cells per item; and their anchor candidates."""
    doc = Document()
    doc.add_paragraph(f"{_COMMITTEE}, Norvale University")
    table = doc.add_table(rows=len(rows), cols=3)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _committee(text=f"{_COMMITTEE}, Norvale University, Member, 2008-2012\t{PROSE}", **fields):
    return _entry(text, code="P", **{"committee_name": _COMMITTEE, "institution": "Norvale University",
                                     "role": "Member", "start_date": "2008", "end_date": "2012",
                                     **fields})


def test_plan_places_a_committee_only_under_its_one_own_row():
    # One committee in two terms: the rows differ only in their dates, which
    # are digits, and words alone take the first row holding them.
    table, lines = _committee_rows([*_COMMITTEE_ROW[:2], "2016-2020"], _COMMITTEE_ROW)
    entry, spans = _committee(), {"P": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "P"), lines) is lines[1]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({"P": [entry]}, _haystack(), lines, spans,
                                   rows_written=written([table.rows[1]._tr]))
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == ("P", id(entry))
    # Both rows recorded for it (a list of records, a re-parsed block): which
    # one a line is about cannot be shown. No row recorded for it, a paragraph
    # outside any table, or no ledger: none either.
    stray = Document().add_paragraph(_COMMITTEE)._p
    for rows_written in (written([row._tr for row in table.rows]), written([]), written([stray]),
                         None):
        assert sp.plan_subpoints({"P": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_committees_one_own_row_whatever_its_words():
    # An acronym and a short role hold too few distinctive words for the word
    # test, and the writer's record alone says the row is the entry's.
    table, lines = _committee_rows(["IRB, NVU", "Chair", "2010"])
    entry = _entry(f"IRB, NVU, Chair, 2010\t{PROSE}", code="P", committee_name="IRB",
                   institution="NVU", role="Chair", start_date="2010")
    spans = {"P": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "P"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"P": [entry]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)


@pytest.mark.parametrize("field", sp._ADMINISTRATIVE_ROW_FIELDS)
def test_a_committee_rows_fallback_fields_count_as_its_rendered_values(field):
    # The row prints them when the schema's committee_name is empty, though
    # `_RENDERED_FIELDS` cannot list them. Another code's are its own.
    entry = _entry("Member", code="P", **{field: "Estuarine Sediment Geochemistry"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "P")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="O"), "O")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="Q2"), "Q2")


def test_the_committee_fallback_fields_are_the_writers_off_schema_names():
    # The writer's precedence, schema field first, and the two names the
    # schema does not define.
    assert sp.ADMINISTRATIVE_ACTIVITY_FIELDS == ("committee_name", "committee", "activity")
    assert sp._ADMINISTRATIVE_ROW_FIELDS == ("committee", "activity")


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its
    # own. Neither names a committee body (`_names_another_committee`).
    (["Chair line", f"Office line, {PROSE}"], 1, f"Office line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Chair line", PROSE, "Office line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Chair line, {PROSE}", "Office line"], 0, f"Chair line, {PROSE}\tOffice line"),
])
def test_a_committee_reads_its_prose_from_its_own_line(segments, index, read):
    entry = _degree_record("\t".join(segments), index, 2)
    entry["taxonomy_code"] = "P"
    assert sp._prose_text("P", entry, {}) == read


_ADMISSIONS = "Example Admissions Committee, Norvale University"


@pytest.mark.parametrize("text, own", [
    # One committee and its years, then its remit: its own line.
    (f"{_COMMITTEE}, Norvale University, Member, 2008-2012\t{PROSE}", True),
    # The role and years in a parenthesis, the writer's repair shape.
    (f"{_COMMITTEE} (Member 2008-2012)\t{PROSE}", True),
    # A run of committees each followed by its years, stage 2's tabs between
    # them: one line, so the writer prints one row, of the last.
    (f"{_ADMISSIONS}, 2001-2004\t{OTHER_PROSE}\t{_COMMITTEE}, Norvale University, 2008-2012"
     f"\t{PROSE}", False),
    # Two committees in a pipe table, the years in a parenthesis and a column.
    (f"{_ADMISSIONS} (Chair 2001-2004) | {OTHER_PROSE}\n{_COMMITTEE} | 2008-2012 | {PROSE}", False),
    # Each opening with its years (`_holds_dated_records`).
    (f"2001-2004 {_ADMISSIONS}\t{OTHER_PROSE}\t2008-2012 {_COMMITTEE}\t{PROSE}", False),
], ids=["one_record", "parenthetical", "trailing_years", "pipe_rows", "leading_years"])
def test_a_committee_line_holding_several_dated_records_offers_none(text, own):
    """Which committee a prose line is about cannot be told: the others'
    records would be offered as the extracted one's prose, under its row."""
    entry = _committee(text)
    assert sp._own_committee_line(entry) == (text if own else "")
    table, lines = _committee_rows(_COMMITTEE_ROW)
    planned = sp.plan_subpoints({"P": [entry]}, _haystack(), lines, {"P": sp.SectionSpan(0, None)},
                                rows_written=lambda e: [table.rows[0]._tr])
    assert [p.paragraphs for p in planned] == ([(PROSE,)] if own else [])


def test_only_a_committee_reads_its_parts_for_several_records():
    # The writer's own test, per line, sees one record on a tab-joined run,
    # and so do the degree and leadership rules, which read date-led parts only.
    # The other record names no committee: its years alone say it is one.
    text = (f"Example Admissions Office, Norvale University, 2001-2004\t{OTHER_PROSE}"
            f"\t{_COMMITTEE}, Norvale University, 2008-2012")
    assert not sp._looks_like_multiple_records([text])
    own_name = sp._committee_name_tokens(_committee(text)["extracted_fields"])
    assert not any(sp._names_another_committee(part, own_name) for part in text.split("\t"))
    assert sp._own_committee_line(_committee(text)) == ""
    for code in ("O", "B1"):
        assert sp._prose_text(code, dict(_committee(text), taxonomy_code=code), {}) == text


_KESTREL = "Example Admissions Appeals Committee, Kestrel Medical Society"


@pytest.mark.parametrize("text", [
    # An undated committee before the extracted one, which carries the years.
    # It shares a word of the extracted name ("Example"), not all of them.
    f"Member, {_KESTREL}\t{_COMMITTEE}, Norvale University, Member, 2008-2012\t{PROSE}",
    # Two undated ones, a panel among them.
    f"Chair, Example Formulary Review Panel, Osprey Regional Hospital\tMember, {_KESTREL}"
    f"\t{_COMMITTEE}, Norvale University, Member, 2008-2012\t{PROSE}",
    # One dated one, the extracted committee's years in its fields only.
    f"Member, {_KESTREL} (2001-2004)\tMember, {_COMMITTEE}, Norvale University\t{PROSE}",
], ids=["undated_sibling", "two_undated_siblings", "dated_sibling_own_years_in_fields"])
def test_a_committee_line_naming_another_committee_offers_none(text):
    """At most one part is dated, so the writer's #660 test reads one record
    and one row prints, of the committee stage 4 extracted: the other
    committee's record would be offered as its prose, under its row."""
    entry = _committee(text)
    assert not sp._looks_like_multiple_records(text.split("\t"))
    assert sp._own_committee_line(entry) == ""
    table, lines = _committee_rows(_COMMITTEE_ROW)
    assert sp.plan_subpoints({"P": [entry]}, _haystack(), lines, {"P": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [table.rows[0]._tr]) == []


@pytest.mark.parametrize("prose", [
    # It names its own committee again, every word of the name.
    f"As secretary to the {_COMMITTEE}, {PROSE.lower()}",
    # "the Committee" refers back to the one the line named: all its words are the name's.
    f"{PROSE} for the Committee",
], ids=["names_its_own", "refers_back"])
def test_a_committee_line_naming_only_its_own_committee_offers_its_prose(prose):
    text = f"{_COMMITTEE}, Norvale University, Member, 2008-2012\t{prose}"
    entry = _committee(text)
    assert sp._own_committee_line(entry) == text
    table, lines = _committee_rows(_COMMITTEE_ROW)
    (planned,) = sp.plan_subpoints({"P": [entry]}, _haystack(), lines,
                                   {"P": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.paragraphs == (prose,)


@pytest.mark.parametrize("field", ["committee", "activity"])
def test_a_committees_own_name_is_read_from_the_fields_its_row_prints(field):
    # Stage 4's off-schema names, which the row prints when the schema's is empty.
    text = f"Member, {_COMMITTEE}, Norvale University, 2008-2012\t{PROSE}"
    assert sp._own_committee_line(_entry(text, code="P", **{field: _COMMITTEE})) == text
    other = f"Member, {_KESTREL}\t{text}"
    assert sp._own_committee_line(_entry(other, code="P", **{field: _COMMITTEE})) == ""


@pytest.mark.parametrize("body", ["Committee", "Subcommittee", "Council", "Board", "Panel",
                                  "Commission", "Senate", "Task Force", "Taskforce",
                                  "Working Group", "Workgroup"])
def test_a_part_naming_a_committee_body_not_the_entrys_own_is_another_record(body):
    part = f"Member, Example Admissions {body}, Kestrel Medical Society"
    assert sp._names_another_committee(part, {"curriculum", "steering"})
    assert not sp._names_another_committee(part, {"admissions"})
    # In lower case it is no name, nor is a word that only begins with one.
    assert not sp._names_another_committee(part.lower(), {"curriculum", "steering"})
    assert not sp._names_another_committee(part.replace(body, f"{body}ist"), {"curriculum"})


@pytest.mark.parametrize("part, other", [
    # The name wrapped onto the next part, or typed with a typo: most of its words.
    ("Committee, Norvale University, 2008-2012", False),
    ("Elected member of Example Curriculum Steerng Committee, Norvale University", False),
    # Another committee shares the word for its kind and one more ("Example"), no more.
    ("Member, Example Admissions Appeals Committee, Kestrel Medical Society", True),
    # A committee on a subject is named with it.
    ("Member, Committee on the Admissions Appeals, Kestrel Medical Society", True),
    # The part holds the whole name: its record, and a sub-committee one of its details.
    (f"Member, {_COMMITTEE} (Example Admissions Subcommittee)", False),
], ids=["wrapped", "typo", "another", "committee_on", "sub_committee"])
def test_another_committee_shares_no_more_than_half_its_words_with_the_entrys(part, other):
    own_name = sp._committee_name_tokens({"committee_name": _COMMITTEE})
    assert sp._names_another_committee(part, own_name) is other


def test_a_committee_on_a_subject_is_named_with_every_word_of_it():
    own_name = sp._committee_name_tokens({"committee_name": "Kestrel Admissions Committee"})
    assert sp._names_another_committee("Committee on Admissions Appeals Review", own_name)


def test_a_committee_line_naming_a_committee_its_fields_do_not_name_offers_none():
    # Stage 4 named no committee, or only in a list of bare names: whose the
    # named one is cannot be shown. A line naming none is its own.
    named = f"Member, {_KESTREL}, 2001-2004\t{PROSE}"
    plain = f"Member, Example Admissions Office, 2001-2004\t{PROSE}"
    for fields in ({"role": "Member", "start_date": "2001"}, {"committee_name": [_KESTREL]}):
        assert sp._own_committee_line(_entry(named, code="P", **fields)) == ""
        assert sp._own_committee_line(_entry(plain, code="P", **fields)) == plain
    assert sp._own_committee_line({"taxonomy_code": "P", "text": plain}) == plain


@pytest.mark.parametrize("field", ["committee_name", "committee", "activity"])
def test_a_committee_list_of_several_records_offers_none_though_one_names_no_committee(field):
    """The writer prints no row for a record naming no committee (#256): a
    list of a nameless record and a named one, under any of its committee
    fields, prints one row, of the named one, and the other's text would be
    offered as its prose. No part names a committee body here, so the list
    alone says the entry holds two."""
    programme = "Example Clerkship Steering Programme"
    text = f"Member, 2001-2004, {OTHER_PROSE}\t{programme}, Norvale University, Member"
    named = {"committee_name": programme, "institution": "Norvale University", "role": "Member"}
    nameless = {"role": "Member", "start_date": "2001", "end_date": "2004"}
    entry = _entry(text, code="P", **{field: [nameless, named]})
    assert sp._own_committee_line(entry) == ""
    table, lines = _committee_rows([f"{programme}, Norvale University", "Member", ""])
    assert sp.plan_subpoints({"P": [entry]}, _haystack(), lines, {"P": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [table.rows[0]._tr]) == []
    # A list of one record is one record.
    assert sp._own_committee_line(_entry(text, code="P", **{field: [named]})) == text


def test_plan_offers_no_prose_of_a_committee_the_pii_pass_withheld():
    table, lines = _committee_rows(_COMMITTEE_ROW)
    spans = {"P": sp.SectionSpan(0, None)}
    entry = _committee()
    (planned,) = sp.plan_subpoints({"P": [entry]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.paragraphs == (PROSE,)

    entry["_pii_withheld"] = True
    assert sp.plan_subpoints({"P": [entry]}, _haystack(), lines, spans,
                             rows_written=lambda e: [table.rows[0]._tr]) == []


_TALK = "Example Advances in Coastal Wetland Hydrology"
_TALK_ROW = [_TALK, "Example Grand Rounds, Norvale University", "2016"]


def _talk_rows(*rows):
    """A line naming the talk and its venue, then a talk table, one row of
    cells per item; and their anchor candidates."""
    doc = Document()
    doc.add_paragraph(f"{_TALK}, Norvale University")
    table = doc.add_table(rows=len(rows), cols=3)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _talk(text=f"{_TALK}, Example Grand Rounds, Norvale University, 2016\t{PROSE}", **fields):
    return _entry(text, code="R", **{"title": _TALK, "event_name": "Example Grand Rounds",
                                     "location": "Norvale University", "date": "2016", **fields})


def _plan_talk(entry, *rows):
    table, lines = _talk_rows(*(rows or (_TALK_ROW,)))
    return sp.plan_subpoints({"R": [entry]}, _haystack(), lines, {"R": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [table.rows[0]._tr] if e is entry else [])


def test_plan_places_a_talk_only_under_its_own_row():
    # One talk given twice: the rows differ only in their years, which are
    # digits, and words alone take the first row holding them.
    table, lines = _talk_rows([*_TALK_ROW[:2], "2019"], _TALK_ROW)
    entry, spans = _talk(), {"R": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "R"), lines) is lines[1]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({"R": [entry]}, _haystack(), lines, spans,
                                   rows_written=written([table.rows[1]._tr]))
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == ("R", id(entry))
    # No row recorded for it, a paragraph outside any table, or no ledger: none.
    stray = Document().add_paragraph(_TALK)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({"R": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_talks_one_own_row_whatever_its_words():
    # An acronym title at an acronym's venue holds no distinctive word.
    entry = _entry(f"IBD, NVU, 2016\t{PROSE}", code="R", title="IBD", location="NVU", date="2016")
    table, lines = _talk_rows(["IBD", "NVU", "2016"])
    assert sp.best_anchor(sp._value_tokens(entry, "R"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"R": [entry]}, _haystack(), lines, {"R": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)


@pytest.mark.parametrize("field", sp._PRESENTATION_ROW_FIELDS)
def test_a_talk_rows_fallback_fields_count_as_its_rendered_values(field):
    # The row prints them when the schema's are empty, though
    # `_RENDERED_FIELDS` cannot list them. Another code's are its own.
    entry = _entry("Invited talk", code="R", **{field: "Estuarine Sediment Geochemistry"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "R")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="H"), "H")


def test_the_talk_row_fields_are_the_writers_cells_and_their_off_schema_names():
    # The writer's precedence per cell, and the names the schema does not define.
    from unified_pipeline.stage6 import formatting
    from unified_pipeline.stage6.sections import presentations

    assert sp.PRESENTATION_TITLE_FIELDS == ("title", "presentation_title")
    assert sp.PRESENTATION_VENUE_FIELDS == ("institution", "location", "venue")
    assert sp.PRESENTATION_EVENT_FIELDS == ("event_name", "conference_name")
    assert sp.PRESENTATION_DATE_FIELDS == ("year", "date", "start_date")
    assert sp._PRESENTATION_ROW_FIELDS == ("presentation_title", "institution", "venue",
                                           "conference_name", "year", "start_date", "end_date")
    assert presentations._TALK_DATE_KEYS is sp.PRESENTATION_DATE_FIELDS
    assert presentations._TALK_EVENT_KEYS is sp.PRESENTATION_EVENT_FIELDS
    # The Dates cell's fields are the ones its builder reads.
    assert sp.PRESENTATION_DATE_FIELDS is formatting.PRESENTATION_DATE_FIELDS
    assert sp.PRESENTATION_END_DATE_FIELD is formatting.PRESENTATION_END_DATE_FIELD


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Wetland line", f"Estuary line, {PROSE}"], 1, f"Estuary line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Wetland line", PROSE, "Estuary line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Wetland line, {PROSE}", "Estuary line"], 0, f"Wetland line, {PROSE}\tEstuary line"),
])
def test_a_talk_reads_its_prose_from_its_own_line(segments, index, read):
    entry = _degree_record("\t".join(segments), index, 2)
    entry["taxonomy_code"] = "R"
    assert sp._prose_text("R", entry, {}) == read


_OTHER_TALK = "Estuary Sediment Transport in Tidal Marshes, Kestrel Institute"


@pytest.mark.parametrize("text, own", [
    # One talk and its year, then its audience: its own line.
    (f"{_TALK}, Norvale University, 2016\t{PROSE}", True),
    # Its year opens the line, and its full date closes it: one talk.
    (f"2016 {_TALK}, Norvale University, March 2016\t{PROSE}", True),
    # Two talks, each ending in its year, stage 2's tabs between them.
    (f"{_OTHER_TALK}, 2014\t{OTHER_PROSE}\t{_TALK}, Norvale University, 2016\t{PROSE}", False),
    # Each opening with its year (`_holds_dated_records`'s shape).
    (f"2014 {_OTHER_TALK}\t{OTHER_PROSE}\t2016 {_TALK}\t{PROSE}", False),
    # Each opening with its year, in one paragraph: one part.
    (f"2014 {_OTHER_TALK}: {OTHER_PROSE} 2016 {_TALK}, Norvale University. {PROSE}", False),
    # The same with a month before each year.
    (f"May 2014 {_OTHER_TALK}. {OTHER_PROSE} March 2016 {_TALK}. {PROSE}", False),
    # Each ending in its year in brackets.
    (f"{_OTHER_TALK}, Osprey Bay (2014)\t{_TALK}, Norvale University (2016).\t{PROSE}", False),
], ids=["one_talk", "year_then_date", "trailing_years", "leading_years", "one_paragraph",
        "month_years", "bracketed_years"])
def test_a_talk_line_holding_several_dated_talks_offers_none(text, own):
    """Which talk a prose line is about cannot be told: the other's text
    would be offered as the extracted one's prose, under its row."""
    entry = _talk(text)
    assert sp._own_talk_line(entry, "R") == (text if own else "")
    assert [p.paragraphs for p in _plan_talk(entry)] == ([(PROSE,)] if own else [])


def test_a_talk_with_no_date_kept_and_one_dated_part_is_one_talk():
    # The Dates cell prints nothing, and one part opens with a date: one talk.
    text = f"2016 {_TALK}, Norvale University\t{PROSE}"
    entry = _entry(text, code="R", title=_TALK, location="Norvale University")
    assert sp._talk_dates(entry["extracted_fields"], "R") == 0
    assert sp._own_talk_line(entry, "R") == text


@pytest.mark.parametrize("dates, own", [
    # The Dates cell prints both years: the two dated parts are this talk's.
    ({"date": "2014; 2016"}, True),
    ({"date": "2016", "additional_dates": ["2014"]}, True),
    ({"year": None, "date": None, "additional_dates": ["2014", "2016"]}, True),
    # It prints one: the other dated part is a talk the row does not print.
    ({"date": "2016"}, False),
    # Three dated parts, two dates printed: one more talk than the row prints.
    ({"date": "2014; 2016"}, "three"),
], ids=["date_list", "further_span", "spans_only", "one_date", "one_too_many"])
def test_one_talk_given_on_several_dates_is_one_record(dates, own):
    parts = [f"2014 {_TALK}, Kestrel Institute", f"2016 {_TALK}, Norvale University", PROSE]
    if own == "three":
        parts.insert(2, f"2018 {_OTHER_TALK}")
    text = "\t".join(parts)
    entry = _talk(text, **dates)
    entry["extracted_fields"] = {k: v for k, v in entry["extracted_fields"].items() if v is not None}
    assert sp._own_talk_line(entry, "R") == (text if own is True else "")


@pytest.mark.parametrize("fields, dates", [
    ({}, 0),
    ({"date": "2016"}, 1),
    ({"year": "2016", "date": "2010"}, 1),
    ({"date": "2014; 2016"}, 2),
    # A list repeating one year prints it once.
    ({"date": "2016-03-01; 2016-05-02"}, 1),
    ({"start_date": "2016", "additional_dates": ["2014", "2018"]}, 3),
    # A further span repeating the talk's own date is not printed again.
    ({"date": "2016", "additional_dates": ["2016"]}, 1),
    # The string "None" is no date.
    ({"year": "None", "additional_dates": ["2016"]}, 1),
    ({"year": "None"}, 0),
])
def test_a_talks_dates_are_counted_as_its_dates_cell_prints_them(fields, dates):
    from unified_pipeline.stage6.sections import presentations

    assert sp._talk_dates(fields, "R") == dates
    cell = presentations._talk_date_cell(fields)
    assert (len(cell.split(", ")) if cell else 0) == dates


@pytest.mark.parametrize("dates", [
    {"date": "2016", "additional_dates": ["2016"]},
    {"year": "None", "date": None, "additional_dates": ["2016"]},
], ids=["span_repeats_its_date", "none_string"])
def test_a_date_its_dates_cell_does_not_print_marks_no_talk_of_its_own(dates):
    # The cell prints one date, so the other dated talk is not this one.
    text = "\t".join([f"{_OTHER_TALK}, 2014", f"{_TALK}, Norvale University, 2016", PROSE])
    entry = _talk(text, **dates)
    entry["extracted_fields"] = {k: v for k, v in entry["extracted_fields"].items() if v is not None}
    assert sp._own_talk_line(entry, "R") == ""
    assert _plan_talk(entry) == []


@pytest.mark.parametrize("line, records", [
    ("Example lecture, Norvale University, 2016", 1),
    ("2016 Example lecture, Norvale University", 1),
    # A year inside the meeting's name: one record.
    ("Example Summit 2016, Norvale University", 1),
    ("2014 Example lecture\t2016 Example seminar", 2),
    ("Example lecture, 2014\tExample seminar, 2016", 2),
    ("2014 Example lecture, Kestrel Institute 2016 Example seminar, Norvale University", 2),
    ("May 2014 Example lecture. March 2016 Example seminar. April 2018 Example panel", 3),
    # A numeric month before each year (EBYSBC DPEHSZ 227's shape).
    ("11/2014 Example lecture, Kestrel Institute 1/2016 Example seminar", 2),
    # A range opens one record, whatever its two years.
    ("2010 - 2012 Example lecture series", 1),
    # A year in brackets ending a part, before a full stop or not; a range in
    # brackets, anywhere.
    ("Example lecture, Norvale University (2016)", 1),
    ("Example lecture, Kestrel Institute (2014)\tExample seminar, Norvale University (2016).", 2),
    ("Example lecture series (2010-2012)\tExample seminar, 2016", 2),
    # A bracketed year inside a part does not end it.
    ("Example lecture (2014), Kestrel Institute\tExample seminar, 2016", 1),
])
def test_a_talk_lines_dates_mark_its_records(line, records):
    assert sp._talk_records(line) == records


@pytest.mark.parametrize("line, another", [
    # Its own title, quoted, punctuated or cased differently.
    (f"“{_TALK.lower()}.” Example Grand Rounds", False),
    (f'"{_TALK}", Example Grand Rounds', False),
    # Its meeting, quoted.
    (f"{_TALK}, “Example Grand Rounds”", False),
    # A word in quotes is no title.
    (f"{_TALK}, a “Student Choice” lecture", False),
    # Another title, in either kind of quotes, of three words or more.
    (f"“{_TALK}”. “Estuary Sediment Transport in Tidal Marshes”", True),
    (f'{_TALK}. "Estuary Sediment Transport in Tidal Marshes", Kestrel Institute', True),
    (f"{_TALK}. “Tidal Marsh Sediments”", True),
    # Half its words the talk's own: another; more than half: its own.
    (f"{_TALK}. “Coastal Wetland Restoration Funding”", True),
    ("“Coastal Wetland Hydrology Funding Review”", False),
], ids=["own_lower", "own_straight", "own_meeting", "short_quote", "another_curly",
        "another_straight", "another_three_words", "half_shared", "most_shared"])
def test_a_quoted_title_not_the_talks_own_is_another_talk(line, another):
    assert sp._quotes_another_talk(line, _talk()["extracted_fields"]) is another


def test_a_quoted_title_of_a_talk_with_no_title_or_meeting_cannot_be_shown_its_own():
    line = f"“{_TALK}”, Norvale University"
    assert sp._quotes_another_talk(line, {"location": "Norvale University"})
    assert not sp._quotes_another_talk("Norvale University seminar", {})


@pytest.mark.parametrize("fields, listed", [
    ({"additional_presentations": [{"title": "Estuary Sediment Transport", "date": "2014"}]}, True),
    ({"stage4_records": [{"location": "Kestrel Institute"}, {"event_name": "Example Seminar"}]}, True),
    # A record naming its venue alone is a talk too.
    ({"additional_presentations": [{"location": "Kestrel Institute", "date": "2014"}]}, True),
    # Date spans are no talks, nor is a list of strings.
    ({"additional_dates": [{"start_date": "2014", "end_date": "2015"}]}, False),
    ({"date": ["2014", "2016"]}, False),
    ({"title": _TALK}, False),
])
def test_a_list_of_talks_under_a_field_is_several_records(fields, listed):
    assert sp._lists_talks(fields) is listed
    text = f"{_TALK}, Norvale University, 2016\t{PROSE}"
    assert sp._own_talk_line(_talk(text, **fields), "R") == ("" if listed else text)


def test_only_a_talk_reads_its_line_for_several_talks():
    # The degree and leadership rules read date-led parts only, and the
    # committee rule names no talk: a line of two talks in one paragraph is
    # one part to each of them.
    text = f"2014 {_OTHER_TALK}: {OTHER_PROSE} 2016 {_TALK}, Norvale University. {PROSE}"
    assert sp._own_talk_line(_talk(text), "R") == ""
    for code in ("O", "B1", "P"):
        assert sp._prose_text(code, dict(_talk(text), taxonomy_code=code), {}) == text


def test_plan_offers_no_prose_of_a_talk_the_pii_pass_withheld():
    entry = _talk()
    (planned,) = _plan_talk(entry)
    assert planned.paragraphs == (PROSE,)

    entry["_pii_withheld"] = True
    assert _plan_talk(entry) == []


_RESIDENCY = "Example Pharmacy Residency"
_RESIDENCY_ROW = [_RESIDENCY, "Norvale General Hospital", "07/09-06/10"]


def _training_rows(*rows):
    """A line naming the hospital, then a training table, one row of cells
    per item; and their anchor candidates."""
    doc = Document()
    doc.add_paragraph("Norvale General Hospital")
    table = doc.add_table(rows=len(rows), cols=3)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _training(text=f"{_RESIDENCY} July 2009-June 2010\tNorvale General Hospital\t{PROSE}",
              **fields):
    return _entry(text, code="C", **{"training_type": _RESIDENCY,
                                     "institution": "Norvale General Hospital",
                                     "start_date": "2009-07", "end_date": "2010-06", **fields})


def _plan_training(entry, *rows):
    table, lines = _training_rows(*(rows or (_RESIDENCY_ROW,)))
    return sp.plan_subpoints({"C": [entry]}, _haystack(), lines, {"C": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [table.rows[0]._tr] if e is entry else [])


def test_plan_places_a_training_position_only_under_its_own_row():
    # A second-year residency at the same hospital renders first and holds
    # every distinctive word of the first year's row: only its dates, digits,
    # tell them apart, and words alone take the first row holding them.
    table, lines = _training_rows([_RESIDENCY, "Norvale General Hospital", "07/10-06/11"],
                                  _RESIDENCY_ROW)
    entry, spans = _training(), {"C": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "C"), lines) is lines[1]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({"C": [entry]}, _haystack(), lines, spans,
                                   rows_written=written([table.rows[1]._tr]))
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == ("C", id(entry))
    # No row recorded for it, a paragraph outside any table, or no ledger: none.
    stray = Document().add_paragraph(_RESIDENCY)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({"C": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_training_positions_one_own_row_whatever_its_words():
    # A type and an institution of acronyms hold no distinctive word.
    entry = _entry(f"PGY-1, NGH, 2009\t{PROSE}", code="C", training_type="PGY-1",
                   institution="NGH", start_date="2009")
    table, lines = _training_rows(["PGY-1", "NGH", "2009"])
    assert sp.best_anchor(sp._value_tokens(entry, "C"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"C": [entry]}, _haystack(), lines, {"C": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)


@pytest.mark.parametrize("field", ["title", "field_of_study", "location"])
def test_a_training_rows_fallback_fields_count_as_its_rendered_values(field):
    # The row prints them when the schema's are empty, though
    # `_RENDERED_FIELDS` cannot list them. Another code's are its own.
    entry = _entry("Fellowship", code="C", **{field: "Estuarine Sediment Geochemistry"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "C")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="H"), "H")


def test_the_training_row_fields_are_the_writers_cells_and_its_other_fields():
    # The writer's precedence per cell, read from here, and the fields it
    # prints whatever the rest hold: C's rendered fields outside the two
    # cells, and `resolution`'s location fallback.
    from unified_pipeline.stage6.resolution import institution
    from unified_pipeline.stage6.sections import postdoc_training

    assert sp.POSTDOC_TYPE_FIELDS == ("training_type", "title")
    assert sp.POSTDOC_FIELD_OF_STUDY_FIELDS == ("field_of_study", "specialty")
    assert sp.POSTDOC_DATE_FIELDS == ("start_date", "end_date")
    assert sp._POSTDOC_PLAIN_FIELDS == ("end_date", "institution", "role", "start_date", "location")
    assert postdoc_training.POSTDOC_TYPE_FIELDS is sp.POSTDOC_TYPE_FIELDS
    assert postdoc_training.POSTDOC_FIELD_OF_STUDY_FIELDS is sp.POSTDOC_FIELD_OF_STUDY_FIELDS
    assert postdoc_training.POSTDOC_DATE_FIELDS is sp.POSTDOC_DATE_FIELDS
    assert postdoc_training.first_field_value is sp.first_field_value
    assert sp.FALLBACK_LOCATION_FIELD is institution.FALLBACK_LOCATION_FIELD


@pytest.mark.parametrize("cell", [sp.POSTDOC_TYPE_FIELDS, sp.POSTDOC_FIELD_OF_STUDY_FIELDS])
def test_a_training_cell_counts_only_the_alias_it_prints(cell):
    # The cell prints the first alias holding a value; the other one's words
    # are not the row's, so a fragment restating them is not held back.
    printed, shadowed = cell
    entry = _training(**{printed: "Estuarine Geochemistry", shadowed: "Saltmarsh Restoration Track"})
    assert sp.first_field_value(entry["extracted_fields"], cell) == "Estuarine Geochemistry"
    tokens = sp._value_tokens(entry, "C")
    assert {"estuarine", "geochemistry"} <= tokens and not {"saltmarsh", "restoration"} & tokens
    entry["extracted_fields"][printed] = ""
    assert sp.first_field_value(entry["extracted_fields"], cell) == "Saltmarsh Restoration Track"
    assert {"saltmarsh", "restoration"} <= sp._value_tokens(entry, "C")
    assert sp.first_field_value({}, cell) is None


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Residency line", f"Fellowship line, {PROSE}"], 1, f"Fellowship line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Residency line", PROSE, "Fellowship line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Residency line, {PROSE}", "Fellowship line"], 0, f"Residency line, {PROSE}\tFellowship line"),
])
def test_a_training_position_reads_its_prose_from_its_own_line(segments, index, read):
    entry = dict(_degree_record("\t".join(segments), index, 2), taxonomy_code="C")
    assert sp._prose_text("C", entry, {}) == read


def test_a_training_line_holding_several_dated_records_offers_none():
    # A residency and a fellowship on one line, each opening with its dates
    # and followed by its prose; stage 4 returned the fellowship. Which one a
    # prose line is about cannot be told.
    fellowship = "2010-2012 Example Critical Care Fellowship, Norvale General Hospital"
    dates = {"start_date": "2010-07", "end_date": "2012-06"}
    own = _training(f"{fellowship}\t{PROSE}", **dates)
    assert [p.paragraphs for p in _plan_training(own)] == [(PROSE,)]
    two = _training(f"2007-2010 {_RESIDENCY}, Norvale General Hospital\t{OTHER_PROSE}\t"
                    f"{fellowship}\t{PROSE}", **dates)
    assert sp._own_record_line(two) == ""
    assert _plan_training(two) == []


@pytest.mark.parametrize("fields, listed", [
    # Stage 4 listed positions under a key the schema does not define, and
    # fan-out kept the list whole: the row prints the entry's own fields only.
    ({"positions": [{"training_type": "Example Surgical Internship",
                     "institution": "Kestrel Institute"}]}, True),
    ({"positions": [{"institution": "Kestrel Institute"}]}, True),
    ({"positions": [{"field_of_study": "Example Marine Studies"}]}, True),
    # A list of date spans, or of bare strings, names no position.
    ({"additional_dates": [{"start_date": "2004", "end_date": "2006"}]}, False),
    ({"positions": ["Example Surgical Internship, Kestrel Institute"]}, False),
], ids=["type", "institution", "field", "date_spans", "strings"])
def test_a_list_of_positions_under_a_field_offers_none(fields, listed):
    text = f"{_RESIDENCY} July 2009-June 2010\tNorvale General Hospital\t{PROSE}"
    entry = _training(text, **fields)
    assert sp._own_training_line(entry) == ("" if listed else text)
    assert [p.paragraphs for p in _plan_training(entry)] == ([] if listed else [(PROSE,)])
    # Only a training position reads its fields for listed positions.
    assert sp._prose_text("O", dict(entry, taxonomy_code="O"), {}) == text


@pytest.mark.parametrize("fields", [None, ["Example Surgical Internship"], "Residency"])
def test_a_training_position_whose_fields_are_not_a_mapping_reads_its_own_line(fields):
    # Its row prints no dates, so a line holding a year holds no year it
    # prints (`_dates_another_position`).
    text = f"{_RESIDENCY}\tNorvale General Hospital\t{PROSE}"
    entry = {"taxonomy_code": "C", "text": text, "extracted_fields": fields}
    assert sp._own_training_line(entry) == text
    assert sp._own_training_line(dict(entry, text=f"{text}\t2009-2010")) == ""


_FELLOWSHIP = "Example Critical Care Fellowship, Kestrel Institute"


@pytest.mark.parametrize("line, start, end, offered", [
    # Another position, its dates ending its part, as a range or in months,
    # or inside it; or the one before, opening it.
    (f"{_RESIDENCY}, Norvale General Hospital, 2009-2010\t{_FELLOWSHIP}, 2010-2013", "2009-07", "2010-06", False),
    (f"{_RESIDENCY}, Norvale General Hospital, 2009-2010\t{_FELLOWSHIP}, July 2010 - June 2013",
     "2009-07", "2010-06", False),
    (f"{_RESIDENCY}, Norvale General Hospital, 2009-2010\t{_FELLOWSHIP} (2010-2013)", "2009-07", "2010-06", False),
    (f"2006-2009 {_FELLOWSHIP}\t{_RESIDENCY}, Norvale General Hospital, 2009-2010", "2009-07", "2010-06", False),
    # A year inside the span is the position's own: a chief year, its own
    # dates written twice, a running position's later year.
    (f"{_RESIDENCY}, Norvale General Hospital, 2008-2012\tChief resident -2011-12- {PROSE}",
     "2008-07", "2012-06", True),
    (f"{_RESIDENCY} July 2009 to July 2010\tNorvale General Hospital, 2009-July 2010- {PROSE}",
     "July 2009", "July 2010", True),
    (f"{_RESIDENCY}, Norvale General Hospital, 2009-present\tChief resident 2013, {PROSE}", "2009", "present", True),
    # No end: the start's year alone; no start: the end's.
    (f"{_RESIDENCY}, Norvale General Hospital, 2009\t{_FELLOWSHIP}, 2010", "2009", "", False),
    (f"{_RESIDENCY}, Norvale General Hospital, 2010\t{PROSE}", "", "2010-06", True),
    (f"{_RESIDENCY}, Norvale General Hospital, 2010\t{_FELLOWSHIP}, 2007", "", "2010-06", False),
    # No dates: any year is another position's. Digits of a number are none.
    (f"{_RESIDENCY}, Norvale General Hospital, 2010\t{PROSE}", "", "", False),
    (f"{_RESIDENCY}, Norvale General Hospital\tLicence: Q207731, 5550192837\t{PROSE}", "", "", True),
], ids=["ending", "months", "inside", "opening", "chief_year", "written_twice", "running",
        "no_end", "no_start", "no_start_other", "no_dates", "number"])
def test_a_training_line_holding_a_year_outside_its_rows_span_offers_none(line, start, end, offered):
    # The #221 recovery reads a sub-point line as rendered, so another
    # position offered as this one's prose would be in the accepted document
    # nowhere (#1205 review).
    entry = _training(line, start_date=start, end_date=end)
    assert sp._own_training_line(entry) == (line if offered else "")
    assert bool(_plan_training(entry)) is offered


def test_plan_offers_no_prose_of_a_training_position_the_pii_pass_withheld():
    entry = _training()
    (planned,) = _plan_training(entry)
    assert planned.paragraphs == (PROSE,)

    entry["_pii_withheld"] = True
    assert _plan_training(entry) == []


def test_families_joining_after_r_are_offered_after_every_family_before_them():
    # Q4B and Q4C sort before R by name; they joined after L3. S1 and S2,
    # which joined after Q4B/Q4C, would otherwise come before C.
    assert sorted(["S2", "Q4C", "L3", "I", "R", "S1", "Q4B", "C", "D1", "B1", "Q2", "K1", "M2A"],
                  key=sp._offer_order) \
        == ["B1", "D1", "K1", "M2A", "Q2", "R", "C", "I", "L3", "Q4B", "Q4C", "S1", "S2"]
    assert sp._LATER_FAMILIES == (sp.POSTDOCTORAL_CODES, sp.MEMBERSHIP_CODES,
                                  sp.CLINICAL_LEADERSHIP_CODES, sp.EDITORIAL_BOARD_SUBPOINT_CODES,
                                  sp.PUBLICATION_SUBPOINT_CODES)


def test_a_line_a_training_position_and_an_appointment_both_hold_stays_the_appointments():
    # Both a residency and the appointment after it name one duty, each on
    # its own line (YUYVIG LTTYWI 76 and 80's shape). The appointment, a
    # family shipped before C, keeps it; the residency keeps its other line.
    duty = "Provided anaesthesia services for the hospital's clinical caseload"
    table, lines = _training_rows(_RESIDENCY_ROW)
    residency = _training(f"{_RESIDENCY} July 2009-June 2010\tNorvale General Hospital\t"
                          f"{duty}\t{PROSE}")
    appointment = _appointment(f"Associate Professor\tExample University\t{duty}")
    appointment_line = _candidate("Associate Professor | Example University | 2010-2015", position=5)
    spans = {"C": sp.SectionSpan(0, 3), "D1": sp.SectionSpan(5, None)}
    planned = sp.plan_subpoints({"C": [residency], "D1": [appointment]}, _haystack(),
                                [*lines, appointment_line], spans,
                                rows_written=lambda e: [table.rows[0]._tr] if e is residency else [])
    assert [(p.code, p.paragraphs) for p in planned] == [("D1", (duty,)), ("C", (PROSE,))]
    # With no appointment holding it, the residency offers both.
    (alone,) = sp.plan_subpoints({"C": [residency]}, _haystack(), lines, spans,
                                 rows_written=lambda e: [table.rows[0]._tr])
    assert alone.paragraphs == (duty, PROSE)


_SOCIETY = "Example Society of Estuarine Ecology"
_SOCIETY_ROW = [f"Member, {_SOCIETY}", "2009-2014"]
_OFFICE = "Member, Estuarine Sediment Working Group, 2003-2008"


def _membership_rows(*rows):
    """A line naming the society, then a memberships table, one row of
    cells per item; and their anchor candidates."""
    doc = Document()
    doc.add_paragraph(_SOCIETY)
    table = doc.add_table(rows=len(rows), cols=2)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _membership(text=f"Member, {_SOCIETY}, 2009-2014\t{PROSE}", **fields):
    return _entry(text, code="I", **{"organization": _SOCIETY, "membership_type": "Member",
                                     "start_date": "2009", "end_date": "2014", **fields})


def _plan_membership(entry, *rows, record_lines=None):
    table, lines = _membership_rows(*(rows or (_SOCIETY_ROW,)))
    return sp.plan_subpoints({"I": [entry]}, _haystack(), lines, {"I": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [table.rows[0]._tr] if e is entry else [],
                             record_lines=record_lines)


def test_plan_places_a_membership_only_under_its_own_row():
    # A later membership of the same society and type renders first and
    # holds every distinctive word of the earlier one's row: only its years,
    # digits, tell them apart, and words alone take the first row.
    table, lines = _membership_rows([f"Member, {_SOCIETY}", "2015-Present"], _SOCIETY_ROW)
    entry, spans = _membership(), {"I": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "I"), lines) is lines[1]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({"I": [entry]}, _haystack(), lines, spans,
                                   rows_written=written([table.rows[1]._tr]))
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == ("I", id(entry))
    # No row recorded for it, a paragraph outside any table, or no ledger: none.
    stray = Document().add_paragraph(_SOCIETY)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({"I": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_memberships_one_own_row_whatever_its_words():
    # An acronym's society and a short type hold one distinctive word.
    entry = _entry(f"Member, ESEE, 2009\t{PROSE}", code="I", organization="ESEE",
                   membership_type="Member", start_date="2009")
    table, lines = _membership_rows(["Member, ESEE", "2009"])
    assert sp.best_anchor(sp._value_tokens(entry, "I"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"I": [entry]}, _haystack(), lines, {"I": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)


def test_a_membership_the_writer_wrote_several_rows_for_offers_none():
    # The writer parsed several memberships out of the entry's line, a row
    # each (#476): which of them a prose line is about cannot be shown.
    table, lines = _membership_rows(_SOCIETY_ROW, ["Fellow, Kestrel Academy of Tidal Studies", "2012"])
    entry, spans = _membership(), {"I": sp.SectionSpan(0, None)}
    assert sp.plan_subpoints({"I": [entry]}, _haystack(), lines, spans,
                             rows_written=lambda e: [row._tr for row in table.rows]) == []
    (planned,) = sp.plan_subpoints({"I": [entry]}, _haystack(), lines, spans,
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1]


@pytest.mark.parametrize("fields, cells", [
    ({}, ("", "")),
    # The schema names the organization: the schema's two only.
    ({"organization": " Example Guild ", "membership_type": "Member", "institution": "Kestrel Union",
      "role": "Chair"}, ("Member", "Example Guild")),
    ({"organization": "Example Guild", "role": "Chair"}, ("", "Example Guild")),
    # It does not: the first of each precedence holding text.
    ({"organization": "  ", "institution": "", "society": "Example Guild", "role": "Chair"},
     ("Chair", "Example Guild")),
    ({"membership_type": "Fellow", "fellowship_designation": "Patron", "organization_name": "Guild"},
     ("Fellow", "Guild")),
], ids=["none", "schema", "schema_no_type", "off_schema", "schema_type_first"])
def test_membership_cell_values_follow_the_writers_precedence(fields, cells):
    assert sp.membership_cell_values(fields) == cells


def test_the_membership_row_fields_are_the_writers():
    # The writer's precedence per cell part, read from here, and the fields
    # it prints whatever the rest hold.
    from unified_pipeline.stage6.sections import memberships

    assert sp.MEMBERSHIP_ORGANIZATION_FIELDS == ("organization", "institution", "organization_name",
                                                 "society")
    assert sp.MEMBERSHIP_TYPE_FIELDS == ("membership_type", "fellowship_designation", "role")
    assert sp._MEMBERSHIP_PLAIN_FIELDS == ("end_date", "start_date")
    assert memberships.MEMBERSHIP_ORGANIZATION_FIELDS is sp.MEMBERSHIP_ORGANIZATION_FIELDS
    assert memberships.MEMBERSHIP_TYPE_FIELDS is sp.MEMBERSHIP_TYPE_FIELDS
    assert memberships.membership_cell_value is sp.membership_cell_value


@pytest.mark.parametrize("field", ["institution", "organization_name", "society"])
def test_an_off_schema_organization_counts_only_when_the_schema_names_none(field):
    # The row prints it when the schema's organization is empty, though
    # `_RENDERED_FIELDS` cannot list it; beside the schema's it prints nothing.
    entry = _entry("Member", code="I", **{field: "Estuarine Sediment Geochemistry Guild"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "I")
    entry["extracted_fields"]["organization"] = "Kestrel Ornithology Union"
    tokens = sp._value_tokens(entry, "I")
    assert {"kestrel", "ornithology"} <= tokens and not {"estuarine"} & tokens


@pytest.mark.parametrize("field", ["fellowship_designation", "role"])
def test_an_off_schema_type_counts_only_when_the_row_prints_it(field):
    entry = _entry("Member", code="I", institution="Kestrel Ornithology Union",
                   **{field: "Distinguished Founding Patron"})
    assert {"distinguished", "founding", "patron"} <= sp._value_tokens(entry, "I")
    # The schema's type comes first ...
    entry["extracted_fields"]["membership_type"] = "Elected Fellow"
    tokens = sp._value_tokens(entry, "I")
    assert {"elected", "fellow"} <= tokens and not {"distinguished"} & tokens
    # ... and beside the schema's organization only the schema's type prints.
    del entry["extracted_fields"]["membership_type"]
    entry["extracted_fields"]["organization"] = "Kestrel Ornithology Union"
    assert not {"distinguished"} & sp._value_tokens(entry, "I")


@pytest.mark.parametrize("line, own, another", [
    # One sentence listing the membership's society among others (EBYSBC NJIKGI 65's shape).
    (f"Member of the {_SOCIETY}, the Kestrel Statistical Association, and the Norvale "
     "Academy of Tidal Studies.", _SOCIETY, True),
    # A chapter's line fused with another college's fellowship (YUYVIG FCAAUV 30's shape).
    ("Fellow Kestrel College of Physicians, Theta Chapter, Example Honor Guild 1986",
     "Theta Chapter, Example Honor Guild", True),
    # Interleaved columns naming other colleges (EBYSBC TXTATQ 22's shape).
    ("Elected Fellow of the Kestrel College of Clinical Pharmacy, appointments across Norvale "
     "University, Osprey College of Critical Care", "Norvale Academies of Practice", True),
    # Another college of the same place: half its words are the kind and the
    # place, shared with the membership's own (EBYSBC TXTATQ 22's two colleges).
    ("Fellow Kestrel College of Physicians, appointed to its ethics council",
     "Kestrel College of Clinical Pharmacy", True),
    # #1205 review of I: three words, two of them the kind and the place it
    # shares with the membership's own. The kind's word is every such
    # society's, so the name's own words share half.
    ("Norvale College of Physicians, elected to its ethics and professionalism council",
     "Norvale College of Cardiology", True),
    ("Member, Norvale Medical Association", "Norvale Heart Association", True),
    ("Norvale Society of Hematology", "Norvale Society of Clinical Oncology", True),
    # A federation, a union, and a kind in the plural are societies too.
    ("Member, Norvale Federation for Tidal Research and Clinical Investigation", _SOCIETY, True),
    ("Federation of Kestrel Societies for Experimental Biology", _SOCIETY, True),
    ("Kestrel Union Against Marsh Erosion", _SOCIETY, True),
    # Its own society, worded another way: two of its three own words.
    ("Member, Example Society of Estuarine Ecologists, its northern chapter", _SOCIETY, False),
    # Its own society by the kind's word alone.
    ("Chaired the Society's annual meeting and edited its proceedings", _SOCIETY, False),
    # Its own society, the type run into its name.
    (f"Fellow {_SOCIETY}; served on its council", _SOCIETY, False),
    # A committee, working group, task force or chapter of it is its detail.
    (f"{_OFFICE}; Chair, Northern Chapter; Task Force on Marsh Restoration", _SOCIETY, False),
    # A word only lowercase, an institute or a university is no society.
    ("Membership goes beyond induction into an honor society", "Example Honor Guild", False),
    ("Member of the Cancer Centre of Norvale University, funded by the Kestrel Cancer Institute",
     "Cancer Centre of Norvale University", False),
    ("Collegiate and Societal outreach to estuarine communities", _SOCIETY, False),
    # With no organization to hold, a named society is another's.
    ("Member of the Kestrel Statistical Association", "", True),
    (PROSE, "", False),
], ids=["sentence_list", "chapter_and_college", "interleaved", "same_kind_and_place",
        "three_words_kind_and_place", "association_of_its_place", "longer_own_name", "federation",
        "plural", "union", "own_worded_otherwise", "own_by_its_kind", "own_with_type",
        "own_details", "lowercase", "institute", "word_begins", "no_own_name", "none_named"])
def test_a_membership_line_naming_another_society(line, own, another):
    assert sp._names_another_society(line, sp._tokens(own)) is another


def test_a_membership_line_naming_another_society_offers_none():
    listed = (f"Member of the {_SOCIETY}, the Kestrel Statistical Association, and the Norvale "
              "Academy of Tidal Studies.")
    entry = _membership(f"{listed}\t{PROSE}")
    assert sp._own_membership_line(entry) == ""
    assert _plan_membership(entry) == []
    # The same prose on a line naming only its own society is offered.
    assert [p.paragraphs for p in _plan_membership(_membership())] == [(PROSE,)]


@pytest.mark.parametrize("line, own, own_type, another", [
    # An open term's end before the name, and the row's own type after a tab,
    # are words the name pattern takes in, not the society's.
    ("2016-Present Society for the Restoration of Tidal Marshes",
     "Society for the Restoration of Tidal Marshes", "Member", False),
    ("Norvale Association of Limnologists\tAssociate Member, 2004-2016",
     "Norvale Association of Limnologists", "Associate Member", False),
    # Set aside, they hide no other society.
    ("2015-Present Norvale College of Physicians", "Norvale College of Cardiology", "Member", True),
    ("Fellow Norvale College of Physicians", "Norvale College of Cardiology", "Fellow", True),
], ids=["open_term", "own_type_after_a_tab", "open_term_other", "own_type_other"])
def test_a_named_society_is_read_without_an_open_term_or_the_rows_type(line, own, own_type, another):
    assert sp._names_another_society(line, sp._tokens(own), sp._tokens(own_type)) is another


def test_a_membership_line_holding_its_own_type_after_a_tab_offers_its_prose():
    own = "Norvale Association of Limnologists"
    entry = _membership(f"{own}\tAssociate Member, 2004-2016\t{PROSE}", organization=own,
                        membership_type="Associate Member", start_date="2004", end_date="2016")
    assert sp._own_membership_line(entry) == entry["text"]
    assert [p.paragraphs for p in _plan_membership(entry, [f"Associate Member, {own}", "2004-2016"])] \
        == [(PROSE,)]


@pytest.mark.parametrize("kind", ["Society", "Societies", "Association", "Associations", "College",
                                  "Colleges", "Academy", "Academies", "Federation", "Federations",
                                  "Union", "Unions"])
def test_every_kind_of_society_names_another_membership(kind):
    assert sp._names_another_society(f"Member, Kestrel {kind} of Marsh Science", sp._tokens(_SOCIETY))


def test_a_membership_line_naming_another_college_of_its_place_offers_none():
    """#1205 review of I: the line goes on into another college of the same
    kind and place. Offered, Reject All would print that membership under
    this one's row."""
    own = "Norvale College of Cardiology"
    other = "Norvale College of Physicians, elected to its ethics and professionalism council"
    entry = _membership(f"Member, {own}, 2009-2014\t{other}", organization=own)
    assert sp._own_membership_line(entry) == ""
    assert _plan_membership(entry, [f"Member, {own}", "2009-2014"]) == []


@pytest.mark.parametrize("fields, listed", [
    # Stage 4 listed memberships under a key the schema does not define, and
    # fan-out kept the list whole: the row prints the entry's own fields only.
    ({"memberships": [{"organization": "Kestrel Statistical Association", "start_date": "2001"}]}, True),
    ({"memberships": [{"society": "Kestrel Statistical Association"}]}, True),
    # A list of date spans, or of bare strings, names no membership.
    ({"additional_dates": [{"start_date": "2004", "end_date": "2006"}]}, False),
    ({"memberships": ["Kestrel Statistical Association"]}, False),
], ids=["organization", "alias", "date_spans", "strings"])
def test_a_list_of_memberships_under_a_field_offers_none(fields, listed):
    entry = _membership(**fields)
    assert sp._own_membership_line(entry) == ("" if listed else entry["text"])
    assert [p.paragraphs for p in _plan_membership(entry)] == ([] if listed else [(PROSE,)])


@pytest.mark.parametrize("fields", [None, ["Example Society of Estuarine Ecology"], "Member"])
def test_a_membership_whose_fields_are_not_a_mapping_reads_its_own_line(fields):
    # It names no organization, so a society its line names is another's.
    text = f"Member, ESEE\t{PROSE}"
    entry = {"taxonomy_code": "I", "text": text, "extracted_fields": fields}
    assert sp._own_membership_line(entry) == text
    assert sp._own_membership_line(dict(entry, text=f"Member, {_SOCIETY}\t{PROSE}")) == ""


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Member line", f"Fellow line, {PROSE}"], 1, f"Fellow line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Member line", PROSE, "Fellow line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Member line, {PROSE}", "Fellow line"], 0, f"Member line, {PROSE}\tFellow line"),
])
def test_a_membership_reads_its_prose_from_its_own_line(segments, index, read):
    entry = dict(_degree_record("\t".join(segments), index, 2), taxonomy_code="I")
    assert sp._prose_text("I", entry, {}) == read


@pytest.mark.parametrize("fragment, records, held", [
    (_OFFICE, [_OFFICE.lower().replace(" ", "  ")], True),
    # Part of a record line, or a fragment holding one.
    ("Member, Estuarine Sediment Working Group", [_OFFICE], True),
    (f"{_OFFICE}. {PROSE}", [_OFFICE], True),
    (PROSE, [_OFFICE], False),
    (PROSE, [], False),
    (PROSE, [""], False),
], ids=["spacing_and_case", "part", "holds", "other", "no_records", "empty_record"])
def test_a_fragment_held_by_a_record_line(fragment, records, held):
    assert sp._held_by_records(fragment, records) is held


def test_a_membership_part_the_recovery_reads_as_a_record_is_left_to_it():
    """A dated office of the society the #221 recovery reads as one of the
    entry's records: dev re-emits it to the Appendix, and as a sub-point the
    recovery would read it as rendered, so it would be in the accepted
    document nowhere. The prose beside it is still offered."""
    entry = _membership(f"Member, {_SOCIETY}, 2009-2014\t{_OFFICE}\t{PROSE}")
    assert [p.paragraphs for p in _plan_membership(entry)] == [(_OFFICE, PROSE)]
    calls = []

    def record_lines(code, e):
        calls.append((code, e))
        # Read as rendered and as no record of its own: a membership's
        # record lines are the recovery's whatever its verdict.
        return [sp.RecordLine(_OFFICE, False, False)] if e is entry else []

    (planned,) = _plan_membership(entry, record_lines=record_lines)
    assert planned.paragraphs == (PROSE,)
    assert calls == [("I", entry)]
    # A line holding only records offers nothing.
    only = _membership(f"Member, {_SOCIETY}, 2009-2014\t{_OFFICE}")
    assert _plan_membership(only, record_lines=lambda code, e: [sp.RecordLine(_OFFICE, True, True)]) == []


def test_only_a_membership_leaves_a_line_of_its_own_to_the_recovery():
    # PROSE is no record of its own: for an appointment it is duty prose,
    # the sub-point's (owner decision 2026-10-08), though the recovery
    # re-emits it with sub-points off.
    appointment = _appointment()
    line = _candidate("Associate Professor | Example University | 2010-2015")
    (planned,) = sp.plan_subpoints({"D1": [appointment]}, _haystack(), [line],
                                   {"D1": sp.SectionSpan(0, None)},
                                   record_lines=lambda code, e: [sp.RecordLine(PROSE, True, False)])
    assert planned.paragraphs == (PROSE,)


#: A sub-record of the appointment with its own years, and prose beside it.
_SUB_RECORD = "Director, Example Admissions Committee, 2012-2014"


@pytest.mark.parametrize("code", ["D1", "K1", "M2A"])
def test_every_family_leaves_a_dated_record_line_to_the_recovery(code):
    """#1205: the #221 recovery re-emits a dated record line of the entry
    with sub-points off. Offered, the sub-point held it, so the recovery read
    it as rendered and Accept All lost it."""
    entry = _entry(f"Associate Professor, Example University, 2010-2015\t{_SUB_RECORD}\t{PROSE}",
                   code=code, title="Associate Professor", institution="Example University")
    line = _candidate("Associate Professor | Example University | 2010-2015")
    spans = {code: sp.SectionSpan(0, None)}
    calls = []

    def record_lines(c, e):
        calls.append(c)
        return [sp.RecordLine("Associate Professor, Example University, 2010-2015", False, False),
                sp.RecordLine(_SUB_RECORD, True, True)]

    (offered,) = sp.plan_subpoints({code: [entry]}, _haystack(), [line], spans)
    assert offered.paragraphs == (_SUB_RECORD, PROSE)
    (planned,) = sp.plan_subpoints({code: [entry]}, _haystack(), [line], spans,
                                   record_lines=record_lines)
    assert planned.paragraphs == (PROSE,) and calls == [code]


def test_record_lines_are_asked_only_for_an_entry_with_prose_to_offer():
    # The generator's verdict on each record line costs a recovery check:
    # an entry with no fragment to offer does not need it.
    calls = []
    entry = _entry("Associate Professor\tExample University", title="Associate Professor",
                   institution="Example University")
    planned = sp.plan_subpoints({"D1": [entry]}, _haystack(), [_candidate("Associate Professor")],
                                {"D1": sp.SectionSpan(0, None)},
                                record_lines=lambda code, e: calls.append(code) or [])
    assert planned == [] and calls == []


#: The entry's own years (`_own_years`) are 2010 and 2015.
_OWN = _entry("unused", title="Associate Professor", institution="Example University",
              start_date="2010", end_date="2015", additional_roles=["Director, 2012-2014"])


@pytest.mark.parametrize("line, tab_part, sub_record", [
    (_SUB_RECORD, True, True),                                    # a tab part the recovery split out
    ("Fellow of the Example Society of Estuarine Ecology, 1985", True, True),  # a lone year
    (_SUB_RECORD, False, False),                                  # a row's text, not one of its tab parts
    ("2012-2014 Director, Example Admissions Committee", False, True),   # date-led
    ("July 2008 - present\tCourtesy Professor, Example College", False, True),  # month-led
    (f"2010-2015\t{PROSE}", False, False),                        # the entry's own years: its duties
    ("Associate Professor, Example University, 2010-2015", True, False),
    (f"2010-2015\tAssociate Professor\t{PROSE} since 2012", False, True),  # a year of its own inside
    (f"Main Annex, Example Campus.\t{PROSE}, 2012", False, False),  # a row by its shape alone
    (f"Main Annex, Example Campus.\t{PROSE} (2012-2014)", False, False),
    ("Member | Example Committee on Curriculum | Example University", False, False),
], ids=["tab_part", "lone_year", "not_a_tab_part", "date_led", "month_led", "own_years",
        "own_years_tab_part", "year_inside", "row_ends_on_a_year", "row_ends_on_a_range", "undated_row"])
def test_a_sub_record_is_dated_by_the_recoverys_reading_with_a_year_of_its_own(line, tab_part, sub_record):
    """#1205 review: a record of its own is a line the recovery dates (a tab
    part it splits out, or a line a date opens) with a year the entry's
    fields do not hold; a list field's years are another record's."""
    assert sp.is_sub_record(line, {line} if tab_part else set(), _OWN) is sub_record


_DUTY_ROW = f"Main Annex, Example Campus.\t{PROSE}"
_OWN_LINE = f"2010-2015\tAssociate Professor\t{PROSE}"


@pytest.mark.parametrize("fragment, code, records, left", [
    (_SUB_RECORD, "D1", [sp.RecordLine(_SUB_RECORD, True, True)], True),
    (f"{_SUB_RECORD}, as its first chair", "K1", [sp.RecordLine(_SUB_RECORD, True, True)], True),
    ("Director, Example Admissions Committee", "M2A",
     [sp.RecordLine("2012-2014\tDirector, Example Admissions Committee", True, True)], True),
    ("Director, Example Admissions Committee", "D1",
     [sp.RecordLine("2012-2014\tDirector, Example Admissions Committee", False, True)], False),
    (PROSE, "D1", [sp.RecordLine(_OWN_LINE, True, False)], False),
    (PROSE, "D1", [sp.RecordLine(_DUTY_ROW, True, False)], False),
    (PROSE, "I", [sp.RecordLine(_DUTY_ROW, False, False)], True),
    (_SUB_RECORD, "D1", [], False),
    (PROSE, "D1", [sp.RecordLine(_SUB_RECORD, True, True)], False),
], ids=["sub_record", "holds_a_sub_record", "part_of_a_lost_sub_record", "part_of_a_rendered_sub_record",
        "part_of_its_own_line", "duties_row", "membership_row", "no_records", "other_record"])
def test_a_fragment_left_to_the_recovery(fragment, code, records, left):
    """#1205 review: a fragment merely part of a sub-record the recovery
    reads as rendered is offered; held back, it was in neither document
    (FINSIS 511)."""
    assert sp._left_to_recovery(fragment, code, records) is left


#: Another entry's row holding the same prose among cells of its own.
_OTHER_ENTRYS_ROW = f"Example Advisory Panel | {PROSE} | Example Water Authority"


@pytest.mark.parametrize("line, prose, held", [
    (_DUTY_ROW, [PROSE], True),             # the duties' own row
    (PROSE, [OTHER_PROSE, PROSE], True),
    (PROSE, [PROSE, OTHER_PROSE], True),
    (_OTHER_ENTRYS_ROW, [PROSE], True),     # another entry's row, read document-wide
    (_DUTY_ROW, [OTHER_PROSE], False),
    (_DUTY_ROW, [], False),
], ids=["duties_row", "the_prose_last", "the_prose_first", "another_entrys_row", "other_prose", "no_sub_point"])
def test_the_recovery_reads_a_sub_point_for_any_entrys_line_it_is_part_of(line, prose, held):
    """#1205: read only for the sub-point's own entry, another entry's row
    holding the duties was re-emitted with them, and Reject All printed them
    twice. A sub-record is the caller's to keep out (`is_sub_record`;
    `test_the_recovery_reads_no_sub_point_as_a_record_of_its_own`)."""
    assert sp.held_by_subpoint(line, prose) is held


def test_plan_offers_no_prose_of_a_membership_the_pii_pass_withheld():
    entry = _membership()
    (planned,) = _plan_membership(entry)
    assert planned.paragraphs == (PROSE,)

    entry["_pii_withheld"] = True
    assert _plan_membership(entry) == []


def test_a_line_a_membership_and_a_service_row_both_hold_stays_the_service_rows():
    # A society member who also sat on an advisory board names the board's
    # remit on both lines. The board, a family shipped before I, keeps it;
    # the membership keeps its other line.
    remit = "Advised the federal agency on coastal wetland restoration funding priorities"
    table, lines = _membership_rows(_SOCIETY_ROW)
    membership = _membership(f"Member, {_SOCIETY}, 2009-2014\t{remit}\t{PROSE}")
    service = _entry(f"Member, Coastal Advisory Board, 2010\t{remit}", code="Q2",
                     organization="Coastal Advisory Board", role="Member", start_date="2010")
    service_line = _candidate("Member | Coastal Advisory Board | 2010", position=5)
    spans = {"I": sp.SectionSpan(0, 3), "Q2": sp.SectionSpan(5, None)}
    planned = sp.plan_subpoints({"I": [membership], "Q2": [service]}, _haystack(),
                                [*lines, service_line], spans,
                                rows_written=lambda e: [table.rows[0]._tr] if e is membership else [])
    assert [(p.code, p.paragraphs) for p in planned] == [("Q2", (remit,)), ("I", (PROSE,))]
    (alone,) = sp.plan_subpoints({"I": [membership]}, _haystack(), lines, spans,
                                 rows_written=lambda e: [table.rows[0]._tr])
    assert alone.paragraphs == (remit, PROSE)


_CLINICAL_ROLE = "Director, Example Wound Care Program"
_CLINICAL_LINE = f"{_CLINICAL_ROLE}, Norvale General Hospital, 2010-2015"
_LATER_CLINICAL_LINE = f"{_CLINICAL_ROLE}, Norvale General Hospital, 2016-Present"


def _clinical_lines(*texts):
    """The clinical-practice section's lines as its writer leaves them, one
    body paragraph per role; and their anchor candidates."""
    doc = Document()
    paragraphs = [doc.add_paragraph(text) for text in texts]
    return paragraphs, list(sp.anchor_candidates(doc))


def _clinical_role(text=f"{_CLINICAL_LINE}\t{PROSE}", **fields):
    return _entry(text, code="L3", **{"leadership_role": _CLINICAL_ROLE,
                                      "institution": "Norvale General Hospital",
                                      "start_date": "2010", "end_date": "2015", **fields})


def _plan_clinical_role(entry, *texts, record_lines=None):
    paragraphs, lines = _clinical_lines(*(texts or (_CLINICAL_LINE,)))
    return sp.plan_subpoints({"L3": [entry]}, _haystack(), lines, {"L3": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [paragraphs[0]._p] if e is entry else [],
                             record_lines=record_lines)


def test_plan_places_a_clinical_leadership_role_only_under_its_own_line():
    # The same role at the same hospital from 2016 renders first and holds
    # every distinctive word of the earlier one's line: only the years,
    # digits, tell them apart, and words alone take the first line.
    paragraphs, lines = _clinical_lines(_LATER_CLINICAL_LINE, _CLINICAL_LINE)
    entry, spans = _clinical_role(), {"L3": sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, "L3"), lines) is lines[0]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({"L3": [entry]}, _haystack(), lines, spans,
                                   rows_written=written([paragraphs[1]._p]))
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == ("L3", id(entry))
    # Nothing recorded for it, a paragraph no candidate is, or no ledger: none.
    stray = Document().add_paragraph(_CLINICAL_LINE)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({"L3": [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_places_a_clinical_leadership_role_under_its_own_table_row():
    # The table form, a "Title | Institution/Location | Dates" table a
    # template ships under the heading: the row the writer recorded takes
    # the prose, not the later role's row of the same words.
    table, lines = _training_rows(["2016-Present", _CLINICAL_ROLE, "Norvale General Hospital"],
                                  ["2010-2015", _CLINICAL_ROLE, "Norvale General Hospital"])
    entry = _clinical_role()
    assert sp.best_anchor(sp._value_tokens(entry, "L3"), lines) is lines[1]
    (planned,) = sp.plan_subpoints({"L3": [entry]}, _haystack(), lines, {"L3": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[1]._tr] if e is entry else [])
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)


def test_plan_takes_a_clinical_leadership_roles_one_own_line_whatever_its_words():
    # A role and a hospital of acronyms hold no distinctive word.
    entry = _entry(f"MD, NGH, 2010\t{PROSE}", code="L3", leadership_role="MD",
                   institution="NGH", start_date="2010")
    paragraphs, lines = _clinical_lines("Example clinical practice line", "MD, NGH, 2010")
    assert sp.best_anchor(sp._value_tokens(entry, "L3"), lines) is None
    (planned,) = sp.plan_subpoints({"L3": [entry]}, _haystack(), lines, {"L3": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [paragraphs[1]._p])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)


@pytest.mark.parametrize("field", ["role", "title", "organization"])
def test_a_clinical_leadership_lines_fallback_fields_count_as_its_rendered_values(field):
    # The line prints them when the schema's role or institution is empty,
    # though `_RENDERED_FIELDS` cannot list them. Another code's are its own.
    entry = _entry("Example role", code="L3", **{field: "Estuarine Sediment Geochemistry"})
    assert {"estuarine", "sediment", "geochemistry"} <= sp._value_tokens(entry, "L3")
    assert not {"estuarine"} & sp._value_tokens(dict(entry, taxonomy_code="H"), "H")


def test_the_clinical_leadership_line_fields_are_the_writers():
    # The writer's precedence per name, read from here, and the fields it
    # prints whatever the rest hold: L3's rendered fields outside the two.
    from unified_pipeline.stage6.sections import clinical_practice

    assert sp.CLINICAL_LEADERSHIP_ROLE_FIELDS == ("role", "leadership_role", "title")
    assert sp.CLINICAL_LEADERSHIP_INSTITUTION_FIELDS == ("institution", "organization")
    assert sp._CLINICAL_LEADERSHIP_PLAIN_FIELDS == ("end_date", "start_date", "unit_program")
    assert clinical_practice.clinical_leadership_cell_values is sp.clinical_leadership_cell_values


@pytest.mark.parametrize("names", [sp.CLINICAL_LEADERSHIP_ROLE_FIELDS,
                                   sp.CLINICAL_LEADERSHIP_INSTITUTION_FIELDS], ids=["role", "institution"])
def test_a_clinical_leadership_name_counts_only_the_alias_it_prints(names):
    # The line prints the first alias holding a value; the other one's words
    # are not the line's, so a fragment restating them is not held back.
    printed, shadowed = names[:2]
    entry = _clinical_role(**{printed: "Estuarine Geochemistry", shadowed: "Saltmarsh Restoration Track"})
    tokens = sp._value_tokens(entry, "L3")
    assert {"estuarine", "geochemistry"} <= tokens and not {"saltmarsh", "restoration"} & tokens
    entry["extracted_fields"][printed] = ""
    assert {"saltmarsh", "restoration"} <= sp._value_tokens(entry, "L3")
    assert sp.clinical_leadership_cell_values({}) == (None, None)


@pytest.mark.parametrize("field", ["description", "program"])
def test_a_description_the_table_form_adds_is_offered_as_the_roles_prose(field):
    # The line form, the one the WCM template takes, prints neither: a
    # fragment restating one is the role's prose. The table form prints it,
    # and the document then holds it (`_record_rendered`).
    entry = _clinical_role(**{field: PROSE})
    assert not sp._tokens(PROSE) & sp._value_tokens(entry, "L3")
    assert [p.paragraphs for p in _plan_clinical_role(entry)] == [(PROSE,)]
    assert sp.unrendered_prose(entry, "L3", _haystack(f"2010-2015 {_CLINICAL_ROLE} {PROSE}")) == []


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Chief line", f"Director line, {PROSE}"], 1, f"Director line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Chief line", PROSE, "Director line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Chief line, {PROSE}", "Director line"], 0, f"Chief line, {PROSE}\tDirector line"),
])
def test_a_clinical_leadership_role_reads_its_prose_from_its_own_line(segments, index, read):
    entry = dict(_degree_record("\t".join(segments), index, 2), taxonomy_code="L3")
    assert sp._prose_text("L3", entry, {}) == read


def test_a_clinical_leadership_line_holding_several_dated_records_offers_none():
    # Two roles on one line, each opening with its dates and followed by its
    # prose; stage 4 returned the first. Which one a prose line is about
    # cannot be told.
    own = _clinical_role(f"2010-2015 {_CLINICAL_ROLE}, Norvale General Hospital\t{PROSE}")
    assert [p.paragraphs for p in _plan_clinical_role(own)] == [(PROSE,)]
    two = _clinical_role(f"2010-2015 {_CLINICAL_ROLE}, Norvale General Hospital\t{PROSE}\t"
                         f"2016-2020 Chief, Example Limb Preservation Service\t{OTHER_PROSE}")
    assert sp._own_record_line(two) == ""
    assert _plan_clinical_role(two) == []


def test_a_clinical_leadership_role_leaves_a_dated_record_line_to_the_recovery():
    # #2010's rule, every family's: a record of its own beside the role's,
    # which the #221 recovery re-emits, stays the recovery's. One it reads
    # as rendered, and a line that is no record of its own (the role's
    # description, owner decision 2026-10-08), stay the sub-point.
    sub_record = "Chair, Limb Salvage Quality Committee, 2012-2014"
    entry = _clinical_role(f"{_CLINICAL_LINE}\t{sub_record}\t{PROSE}")
    (offered,) = _plan_clinical_role(entry)
    assert offered.paragraphs == (sub_record, PROSE)
    (planned,) = _plan_clinical_role(entry, record_lines=lambda code, e: [
        sp.RecordLine(_CLINICAL_LINE, False, False), sp.RecordLine(sub_record, True, True)])
    assert planned.paragraphs == (PROSE,)
    (kept,) = _plan_clinical_role(entry, record_lines=lambda code, e: [
        sp.RecordLine(sub_record, False, True), sp.RecordLine(PROSE, True, False)])
    assert kept.paragraphs == (sub_record, PROSE)


def test_plan_offers_no_prose_of_a_clinical_leadership_role_the_pii_pass_withheld():
    entry = _clinical_role()
    (planned,) = _plan_clinical_role(entry)
    assert planned.paragraphs == (PROSE,)

    entry["_pii_withheld"] = True
    assert _plan_clinical_role(entry) == []


def test_a_line_a_clinical_leadership_role_and_an_institutional_role_both_hold_stays_the_institutional_roles():
    # A CV can list one directorship both under institutional leadership (O)
    # and under clinical leadership (L3), each with the same line of what it
    # oversaw. The institutional role, a family shipped before L3, keeps it;
    # the clinical role keeps its other line, and alone it offers both.
    doc = Document()
    clinical_line = doc.add_paragraph(_CLINICAL_LINE)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, [_ROLE, "Norvale University", "2008-2012"], strict=True):
        cell.text = text
    lines = list(sp.anchor_candidates(doc))
    clinical = _clinical_role(f"{_CLINICAL_LINE}\t{PROSE}\t{OTHER_PROSE}")
    institutional = _leadership_role()
    owned = {id(clinical): [clinical_line._p], id(institutional): [table.rows[0]._tr]}
    spans = {"L3": sp.SectionSpan(0, 1), "O": sp.SectionSpan(1, None)}
    planned = sp.plan_subpoints({"L3": [clinical], "O": [institutional]}, _haystack(), lines, spans,
                                rows_written=lambda e: owned.get(id(e), []))
    assert [(p.code, p.paragraphs) for p in planned] == [("O", (PROSE,)), ("L3", (OTHER_PROSE,))]
    (alone,) = sp.plan_subpoints({"L3": [clinical]}, _haystack(), lines, spans,
                                 rows_written=lambda e: owned.get(id(e), []))
    assert alone.paragraphs == (PROSE, OTHER_PROSE)


_SENTENCE_TAIL = "supervised across the regional wound clinics and satellite programmes."


def test_a_clinical_leadership_fragment_opening_in_lower_case_is_not_offered():
    # The tail of a sentence whose start is in a part holding the line's
    # names (a PDF's two columns read across, YUYVIG RLADNC 524): alone it
    # reads as nothing. Capitalised, the same words are offered, and another
    # family's tail is offered as before.
    entry = _clinical_role(f"{_CLINICAL_LINE}\t{PROSE}.\t{_SENTENCE_TAIL}")
    assert [p.paragraphs for p in _plan_clinical_role(entry)] == [(f"{PROSE}.",)]
    capitalised = _SENTENCE_TAIL.capitalize()
    entry = _clinical_role(f"{_CLINICAL_LINE}\t{PROSE}.\t{capitalised}")
    assert [p.paragraphs for p in _plan_clinical_role(entry)] == [(f"{PROSE}.", capitalised)]
    table, lines = _role_rows([_ROLE, "Norvale University", "2008-2012"])
    institutional = _leadership_role(f"2008-2012\t{_ROLE}, Norvale University\t{_SENTENCE_TAIL}")
    (planned,) = sp.plan_subpoints({"O": [institutional]}, _haystack(), lines, {"O": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.paragraphs == (_SENTENCE_TAIL,)


def _clinical_subpoint(*paragraphs, code="L3"):
    """A clinical leadership role's planned sub-point under its line."""
    _paragraphs, (line,) = _clinical_lines(_CLINICAL_LINE)
    return sp.SubPoint(line, paragraphs, 0, code, 1482)


@pytest.mark.parametrize("part, held", [
    (OTHER_PROSE, True),                                   # a paragraph of the sub-point
    (_CLINICAL_LINE, True),                                # the role's line
    ("Norvale General Hospital, 2010 - 2015", True),       # every word of it the line's
    ("Budget: $4M", False),                                # too short to be offered
    ("Staff:", False),                                     # a label
    ("12 nurses", False),                                  # a count the line does not print
    ("Norvale General Hospital, 2009-2016", False),        # years the line does not print
    ("2013: opened the limb salvage clinic", False),       # a date-led line
    ("Department of Surgery, Norvale General Hospital", False),  # a unit the line does not print
])
def test_a_clinical_leadership_sub_point_stands_for_its_overflow_only_holding_every_part(part, held):
    # The content overflow writes the role's whole text in place of its
    # line and skips an entry with a sub-point: a part neither holds would
    # be in neither document (#1205 review).
    subpoint = _clinical_subpoint(PROSE, OTHER_PROSE)
    text = f"{_CLINICAL_LINE}\t{PROSE}\t{part}"
    assert sp.yields_to_overflow(subpoint, text) is not held
    # A paragraph the planner refolded from parts the source broke apart, and
    # a source table's row, its cells joined with pipes.
    assert not sp.yields_to_overflow(_clinical_subpoint(f"{PROSE} {OTHER_PROSE}"),
                                     f"{_CLINICAL_LINE}\t{PROSE}\n{OTHER_PROSE}")
    assert not sp.yields_to_overflow(subpoint, f"{_CLINICAL_LINE} | {PROSE} | {OTHER_PROSE}")
    # No overflow for the entry, or another family's sub-point: it stands.
    assert not sp.yields_to_overflow(subpoint, None)
    assert not sp.yields_to_overflow(subpoint._replace(code="D1"), text)


_JOURNAL = "Journal of Estuarine Sediment Studies"
_SEAT = "Associate Editor"
_SEAT_ROW = [f"{_SEAT}, {_JOURNAL}", "2012-2016"]


def _board_rows(*rows):
    """A line naming the journal, then an Editorial Board Membership table
    ("Board / Organization Name | Dates"), one row of cells per item; and
    their anchor candidates."""
    doc = Document()
    doc.add_paragraph(f"Reviewer for the {_JOURNAL}")
    table = doc.add_table(rows=len(rows), cols=2)
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    return table, list(sp.anchor_candidates(doc))


def _board_seat(text=f"{_SEAT}, {_JOURNAL}, 2012-2016\t{PROSE}", code="Q4B", **fields):
    return _entry(text, code=code, **{"journal_name": _JOURNAL, "role": _SEAT,
                                      "start_date": "2012", "end_date": "2016", **fields})


def _plan_board_seat(entry, *rows, record_lines=None):
    table, lines = _board_rows(*(rows or (_SEAT_ROW,)))
    code = entry["taxonomy_code"]
    return sp.plan_subpoints({code: [entry]}, _haystack(), lines, {code: sp.SectionSpan(0, None)},
                             rows_written=lambda e: [table.rows[0]._tr] if e is entry else [],
                             record_lines=record_lines)


@pytest.mark.parametrize("code", ["Q4B", "Q4C"])
def test_plan_places_an_editorial_board_seat_only_under_its_own_row(code):
    # A later seat on the same journal's board renders first and holds every
    # distinctive word of the earlier one's row: only the years, digits,
    # tell them apart, and words alone take the first row.
    table, lines = _board_rows([f"{_SEAT}, {_JOURNAL}", "2017-Present"], _SEAT_ROW)
    entry, spans = _board_seat(code=code), {code: sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, code), lines) is lines[1]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({code: [entry]}, _haystack(), lines, spans,
                                   rows_written=written([table.rows[1]._tr]))
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)
    assert (planned.code, planned.entry_id) == (code, id(entry))
    # No row recorded for it, a paragraph outside any table, or no ledger: none.
    stray = Document().add_paragraph(_JOURNAL)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({code: [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_an_editorial_board_seats_one_own_row_whatever_its_words():
    # An acronym journal and a short role hold one distinctive word.
    entry = _entry(f"Editor, JESS, 2012\t{PROSE}", code="Q4B", journal_name="JESS", role="Editor",
                   start_date="2012")
    table, lines = _board_rows(["Editor, JESS", "2012"])
    assert sp.best_anchor(sp._value_tokens(entry, "Q4B"), lines[1:]) is None
    (planned,) = sp.plan_subpoints({"Q4B": [entry]}, _haystack(), lines, {"Q4B": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [table.rows[0]._tr])
    assert planned.anchor is lines[1] and planned.paragraphs == (PROSE,)


@pytest.mark.parametrize("fields, text, role", [
    # Stage 4's role, whatever the line says.
    ({"role": "Section Editor"}, f"{_SEAT}, {_JOURNAL}", "Section Editor"),
    # A role the line writes as a cell of its own, between pipes.
    ({}, f"2012 | {_SEAT} | {_JOURNAL}", _SEAT),
    ({}, f"2012 | Editorial Board Member | {_JOURNAL}", "Editorial Board Member"),
    ({}, f"2012 | editorial board | {_JOURNAL}", "editorial board"),
    # One of the writer's line roles, as written, in the writer's order.
    ({}, f"Guest Editor, special issue of the {_JOURNAL}", "Guest Editor"),
    ({}, f"Guest Editor and {_SEAT}, {_JOURNAL}", _SEAT),
    # Any other wording names none: the row prints no role.
    ({}, f"Member, Editorial Advisory Board, {_JOURNAL}", ""),
    ({}, f"Editorial board member, {_JOURNAL}", ""),
    ({}, f"associate editor, {_JOURNAL}", ""),
    ({"role": ""}, _JOURNAL, ""),
    # A role stage 4 returned as a list reads as the cell prints it; a null one names none.
    ({"role": [_SEAT, "Guest Editor"]}, _JOURNAL, f"{_SEAT}; Guest Editor"),
    ({"role": None}, f"Guest Editor, special issue of the {_JOURNAL}", "Guest Editor"),
], ids=["stage4", "cell", "cell_member", "cell_lower", "line_role", "first_line_role",
        "advisory", "lower_case_board", "lower_case_role", "none", "list", "null"])
def test_an_editorial_board_rows_role_is_stage_4s_or_the_one_its_line_names(fields, text, role):
    assert sp.editorial_board_role(fields, text) == role


def test_an_editorial_board_seat_with_several_recorded_rows_is_read_by_words():
    # The writer records one row per seat. Were two recorded for one entry,
    # its words would choose between them, as for a talk or a training
    # position, rather than none being taken.
    table, lines = _board_rows(["Guest Editor, Kestrel Tidal Review", "2014"], _SEAT_ROW)
    entry = _board_seat()
    (planned,) = sp.plan_subpoints({"Q4B": [entry]}, _haystack(), lines, {"Q4B": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [row._tr for row in table.rows])
    assert planned.anchor is lines[2] and planned.paragraphs == (PROSE,)


def test_the_editorial_board_row_fields_are_the_writers():
    # The writer reads the role and the organization's precedence from here.
    from unified_pipeline.stage6.sections import service

    assert service.EDITORIAL_BOARD_CODES is sp.EDITORIAL_BOARD_SUBPOINT_CODES
    assert set(sp.EDITORIAL_BOARD_SUBPOINT_CODES) == {"Q4B", "Q4C"}
    assert service.SECTION_EDITOR_CODE in sp.EDITORIAL_BOARD_SUBPOINT_CODES
    assert sp.OTHER_SERVICE_ORGANIZATION_FIELDS == ("organization", "committee_name", "agency")
    assert service.OTHER_SERVICE_ORGANIZATION_FIELDS is sp.OTHER_SERVICE_ORGANIZATION_FIELDS
    assert service.editorial_board_role is sp.editorial_board_role
    # The words a generic seat restates hold the heading of the table the
    # writer fills, the first it looks up.
    for code in sp.EDITORIAL_BOARD_SUBPOINT_CODES:
        heading = service.OTHER_SERVICE_SECTION_ROUTING[code][1][0]
        assert sp._tokens(heading) <= sp._EDITORIAL_BOARD_HEADING_WORDS, heading


@pytest.mark.parametrize("code", ["Q4B", "Q4C"])
def test_an_editorial_board_rows_role_and_organization_count_as_its_rendered_values(code):
    # The row prints them, though `_RENDERED_FIELDS` cannot list Q4C's role
    # or either code's organization.
    entry = _entry("Example line", code=code, role="Statistical Consultant Editor",
                   organization="Norvale Marine Press", committee_name="Estuarine Committee")
    tokens = sp._value_tokens(entry, code)
    assert {"statistical", "consultant", "norvale", "marine", "press"} <= tokens
    # Only the first organization of the writer's precedence prints.
    assert not {"estuarine", "committee"} & tokens
    del entry["extracted_fields"]["organization"]
    assert {"estuarine", "committee"} <= sp._value_tokens(entry, code)
    # The role the writer reads from the line counts when stage 4 names none.
    entry = _entry(f"2012 | {_SEAT} | {_JOURNAL}", code=code, journal_name=_JOURNAL)
    assert {"associate", "editor"} <= sp._value_tokens(entry, code)
    # Another code's line names no editorial role.
    assert not {"associate"} & sp._value_tokens(dict(entry, taxonomy_code="Q4D"), "Q4D")


def test_a_line_restating_an_editorial_board_row_with_the_role_it_prints_is_not_offered():
    """The row prints the role stage 4 names, or the one the writer reads
    from the line: a line restating the row with that role is the row
    reformatted, and Reject All would print it twice."""
    line = f"Member, Editorial Advisory Board, {_JOURNAL}, 2012-2016"
    shown = _entry(line, code="Q4C", journal_name=_JOURNAL, role="Member, Editorial Advisory Board",
                   start_date="2012", end_date="2016")
    assert sp.unrendered_prose(shown, "Q4C", _haystack()) == []
    read = _entry(f"{_SEAT}, {_JOURNAL}, 2012-2016", code="Q4C", journal_name=_JOURNAL,
                  start_date="2012", end_date="2016")
    assert sp.editorial_board_role(read["extracted_fields"], read["text"]) == _SEAT
    assert sp.unrendered_prose(read, "Q4C", _haystack()) == []


@pytest.mark.parametrize("line, journal", [
    (f"Honorary Member, Editorial Advisory Board, {_JOURNAL}, 2012-2016", _JOURNAL),
    # One word beyond the heading's is the kind of board or the seat's standing.
    ("Member, Editorial Advisory Board, Marsh Hydrology, 2012-2016", "Marsh Hydrology"),
    ("Editorial Board inaugural member, JESS, 2012-2016", "JESS"),
], ids=["honorary", "advisory", "inaugural"])
def test_a_line_restating_an_editorial_board_row_with_a_role_it_does_not_print_is_offered(line, journal):
    """A Q4C row prints no role when stage 4 names none and the writer reads
    none from the line. The role is the kind of board and the seat's
    standing, which no cell or heading holds: the line is offered as
    written."""
    entry = _entry(line, code="Q4C", journal_name=journal, start_date="2012", end_date="2016")
    assert sp.editorial_board_role(entry["extracted_fields"], line) == ""
    assert [p.paragraphs for p in _plan_board_seat(entry, [journal, "2012-2016"])] == [(line,)]


@pytest.mark.parametrize("code", ["Q4B", "Q4C"])
@pytest.mark.parametrize("line, journal, start, end", [
    # An acronym journal holds no distinctive word, and a two-word one two
    # of the line's five: neither line mostly holds the row's values.
    ("Editorial Board Member, JESS, 2015-present", "JESS", "2015", "present"),
    ("Member, Editorial Board, Marsh Hydrology, 2010-2014", "Marsh Hydrology", "2010", "2014"),
    ("Editorial Board Membership: Member, Marsh Hydrology, 2010-2014", "Marsh Hydrology", "2010", "2014"),
], ids=["acronym", "two_words", "every_heading_word"])
def test_a_line_restating_an_editorial_board_row_with_only_its_headings_words_is_not_offered(
        line, journal, start, end, code):
    """A generic seat names only the words of the Editorial Board
    Membership heading the row sits under, and the row prints no role: the
    line says what the row and its heading say, and Reject All would print
    the journal and the years twice with nothing added."""
    entry = _entry(line, code=code, journal_name=journal, start_date=start, end_date=end)
    assert sp.prose_fragments(line) == [line]
    assert sp.editorial_board_role(entry["extracted_fields"], line) == ""
    assert not sp._holds_field_values(line, sp._value_tokens(entry, code))
    assert sp.unrendered_prose(entry, code, _haystack()) == []
    assert _plan_board_seat(entry, [journal, f"{start}-{end}"]) == []
    # The seat's prose on a segment of its own is still offered, and under
    # another code's heading the line is not the row's.
    entry["text"] = f"{line}\t{PROSE}"
    assert sp.unrendered_prose(entry, code, _haystack()) == [PROSE]
    assert sp.unrendered_prose(entry, "Q4A", _haystack()) == [line, PROSE]


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    (["Board line", f"Editor line, {PROSE}"], 1, f"Editor line, {PROSE}"),
    # The last of two records on a line of three segments: none of it.
    (["Board line", PROSE, "Editor line"], 1, ""),
    # An earlier record: the line fan-out gave it, whatever it holds.
    ([f"Board line, {PROSE}", "Editor line"], 0, f"Board line, {PROSE}\tEditor line"),
])
def test_an_editorial_board_seat_reads_its_prose_from_its_own_line(segments, index, read):
    for code in ("Q4B", "Q4C"):
        entry = dict(_degree_record("\t".join(segments), index, 2), taxonomy_code=code)
        assert sp._prose_text(code, entry, {}) == read


def test_an_editorial_board_line_holding_several_dated_records_offers_none():
    # Two seats on one line, each opening with its dates and followed by its
    # prose; stage 4 returned the first. Which one a prose line is about
    # cannot be told.
    own = _board_seat(f"2012-2016 {_SEAT}, {_JOURNAL}\t{PROSE}")
    assert [p.paragraphs for p in _plan_board_seat(own)] == [(PROSE,)]
    two = _board_seat(f"2012-2016 {_SEAT}, {_JOURNAL}\t{PROSE}\t"
                      f"2017-2020 Guest Editor, Kestrel Tidal Review\t{OTHER_PROSE}")
    assert sp._own_record_line(two) == ""
    assert _plan_board_seat(two) == []


def test_an_editorial_board_seat_leaves_a_dated_record_line_to_the_recovery():
    # #2010's rule, every family's: a record of its own beside the seat's,
    # which the #221 recovery re-emits, stays the recovery's. One it reads
    # as rendered, and a line that is no record of its own (the seat's
    # duties, owner decision 2026-10-08), stay the sub-point.
    sub_record = "Guest Editor, Special Issue on Marsh Hydrology, 2014"
    entry = _board_seat(f"{_SEAT}, {_JOURNAL}, 2012-2016\t{sub_record}\t{PROSE}")
    (offered,) = _plan_board_seat(entry)
    assert offered.paragraphs == (sub_record, PROSE)
    (planned,) = _plan_board_seat(entry, record_lines=lambda code, e: [
        sp.RecordLine(f"{_SEAT}, {_JOURNAL}, 2012-2016", False, False), sp.RecordLine(sub_record, True, True)])
    assert planned.paragraphs == (PROSE,)
    (kept,) = _plan_board_seat(entry, record_lines=lambda code, e: [
        sp.RecordLine(sub_record, False, True), sp.RecordLine(PROSE, True, False)])
    assert kept.paragraphs == (sub_record, PROSE)


def test_plan_offers_no_prose_of_an_editorial_board_seat_the_pii_pass_withheld():
    entry = _board_seat()
    (planned,) = _plan_board_seat(entry)
    assert planned.paragraphs == (PROSE,)

    entry["_pii_withheld"] = True
    assert _plan_board_seat(entry) == []


def test_a_line_a_talk_and_an_editorial_board_seat_both_hold_stays_the_talks():
    # A guest editor who also lectured on the special issue's theme names it
    # on both lines. The talk, a family shipped before Q4B and named after
    # it, keeps it; the seat keeps its other line, and alone it offers both.
    theme = "Introduced the special issue's theme of sediment budgets in restored tidal marshes"
    seat_table, seat_lines = _board_rows(_SEAT_ROW)
    talk_table, talk_lines = _talk_rows(_TALK_ROW)
    seat = _board_seat(f"{_SEAT}, {_JOURNAL}, 2012-2016\t{theme}\t{PROSE}")
    talk = _talk(f"{_TALK}, Example Grand Rounds, Norvale University, 2016\t{theme}")
    lines = [*seat_lines, *(line._replace(position=line.position + 5) for line in talk_lines)]
    owned = {id(seat): [seat_table.rows[0]._tr], id(talk): [talk_table.rows[0]._tr]}
    spans = {"Q4B": sp.SectionSpan(0, 3), "R": sp.SectionSpan(5, None)}
    planned = sp.plan_subpoints({"Q4B": [seat], "R": [talk]}, _haystack(), lines, spans,
                                rows_written=lambda e: owned.get(id(e), []))
    assert [(p.code, p.paragraphs) for p in planned] == [("R", (theme,)), ("Q4B", (PROSE,))]
    (alone,) = sp.plan_subpoints({"Q4B": [seat]}, _haystack(), lines, spans,
                                 rows_written=lambda e: owned.get(id(e), []))
    assert alone.paragraphs == (theme, PROSE)


# --- S1/S2 publications -----------------------------------------------------------

_PAPER_AUTHORS = "Quill AM, Rook BT, Marblegate Q"
_PAPER_TITLE = "Sediment budgets of restored tidal marshes"
_PAPER_JOURNAL = "Journal of Estuarine Sediment Studies"
_PAPER_CITATION = f"{_PAPER_AUTHORS}. {_PAPER_TITLE}. {_PAPER_JOURNAL}. 2014;12(3):45-52."
_LATER_PAPER_CITATION = (f"{_PAPER_AUTHORS}. {_PAPER_TITLE} after a decade. {_PAPER_JOURNAL}. "
                         "2018;16(1):1-9.")
_PAPER_PROSE = "Designed the sampling study, analysed the sediment cores and wrote the manuscript"


def _paper(text=f"{_PAPER_CITATION}\t{_PAPER_PROSE}", code="S1", **fields):
    return _entry(text, code=code, **{"authors": _PAPER_AUTHORS, "title": _PAPER_TITLE,
                                      "journal": _PAPER_JOURNAL, "year": "2014", "volume": "12",
                                      "issue": "3", "pages": "45-52", **fields})


def _citations(*texts):
    """The bibliography's citation paragraphs as its writer leaves them, one
    body paragraph per paper; and their anchor candidates."""
    doc = Document()
    paragraphs = [doc.add_paragraph(text) for text in texts]
    return paragraphs, list(sp.anchor_candidates(doc))


def _plan_paper(entry, *texts, record_lines=None):
    texts = texts or (_PAPER_CITATION,)
    paragraphs, lines = _citations(*texts)
    code = entry["taxonomy_code"]
    return sp.plan_subpoints({code: [entry]}, _haystack(*texts), lines, {code: sp.SectionSpan(0, None)},
                             rows_written=lambda e: [paragraphs[0]._p] if e is entry else [],
                             record_lines=record_lines)


@pytest.mark.parametrize("code", ["S1", "S2"])
def test_plan_places_a_publications_prose_only_under_its_own_citation(code):
    # The same group's later paper in the same journal renders first and
    # holds every distinctive word of the earlier one's citation: words
    # alone take the first paragraph.
    paragraphs, lines = _citations(_LATER_PAPER_CITATION, _PAPER_CITATION)
    entry, spans = _paper(code=code), {code: sp.SectionSpan(0, None)}
    assert sp.best_anchor(sp._value_tokens(entry, code), lines) is lines[0]

    def written(recorded):
        return lambda e: recorded if e is entry else []

    (planned,) = sp.plan_subpoints({code: [entry]}, _haystack(_LATER_PAPER_CITATION, _PAPER_CITATION),
                                   lines, spans, rows_written=written([paragraphs[1]._p]))
    assert planned.anchor is lines[1] and planned.paragraphs == (_PAPER_PROSE,)
    assert (planned.code, planned.entry_id) == (code, id(entry))
    # Nothing recorded for it (a split, declined or superseded entry, a
    # dedup loser), a paragraph no candidate is, or no ledger: none.
    stray = Document().add_paragraph(_PAPER_CITATION)._p
    for rows_written in (written([]), written([stray]), None):
        assert sp.plan_subpoints({code: [entry]}, _haystack(), lines, spans,
                                 rows_written=rows_written) == []


def test_plan_takes_a_publications_one_own_citation_whatever_its_words():
    # A citation of acronyms and initials holds no distinctive word.
    entry = _entry(f"Quill AM. MRI of JESS. EJSS. 2014;1:2.\t{_PAPER_PROSE}", code="S1",
                   authors="Quill AM", title="MRI of JESS", journal="EJSS", year="2014")
    paragraphs, lines = _citations(_LATER_PAPER_CITATION, "Quill AM. MRI of JESS. EJSS. 2014;1:2.")
    assert sp.best_anchor(sp._value_tokens(entry, "S1"), lines) is None
    (planned,) = sp.plan_subpoints({"S1": [entry]}, _haystack(), lines, {"S1": sp.SectionSpan(0, None)},
                                   rows_written=lambda e: [paragraphs[1]._p])
    assert planned.anchor is lines[1] and planned.paragraphs == (_PAPER_PROSE,)


def _tracked(paragraph, tag, text):
    """A tracked insertion or deletion of `text` appended to `paragraph`."""
    change = sp.Revision(0, _DATE).element(tag)
    run = change.makeelement(qn("w:r"), {})
    piece = run.makeelement(qn("w:delText" if tag == "w:del" else "w:t"), {})
    piece.text = text
    run.append(piece)
    change.append(run)
    paragraph._p.append(change)
    return paragraph


def test_a_citation_written_as_a_tracked_change_takes_no_sub_point():
    # The enrichment writer's fallback shape: the CV's line deleted, then
    # the citation as plain runs. Reject All gives the line back, prose and
    # all, so a sub-point would print the prose twice.
    entry = _paper()
    paragraphs, lines = _citations(_PAPER_CITATION)
    _tracked(paragraphs[0], "w:del", entry["text"])
    assert sp._holds_tracked_change(paragraphs[0]._p)
    assert sp.plan_subpoints({"S1": [entry]}, _haystack(_PAPER_CITATION), lines, {"S1": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [paragraphs[0]._p]) == []
    # An insertion beside plain runs: Reject All would drop part of the
    # citation and keep the prose.
    paragraphs, lines = _citations(_PAPER_CITATION)
    _tracked(paragraphs[0], "w:ins", " Erratum in Coastal Studies Review 2015.")
    assert sp._holds_tracked_change(paragraphs[0]._p)
    assert sp.plan_subpoints({"S1": [entry]}, _haystack(_PAPER_CITATION), lines, {"S1": sp.SectionSpan(0, None)},
                             rows_written=lambda e: [paragraphs[0]._p]) == []
    # An enriched citation, every run inside w:del or w:ins: python-docx
    # reads no text, so it is no candidate at all.
    doc = Document()
    enriched = _tracked(_tracked(doc.add_paragraph(), "w:del", entry["text"]), "w:ins", _PAPER_CITATION)
    assert sp._holds_tracked_change(enriched._p) and list(sp.anchor_candidates(doc)) == []
    # A plain citation holds none, and takes the prose.
    assert not sp._holds_tracked_change(_citations(_PAPER_CITATION)[0][0]._p)
    assert [p.paragraphs for p in _plan_paper(_paper())] == [(_PAPER_PROSE,)]


_FULL_AUTHOR_NAMES = "Annette Marisol Quill, Bertram Tobias Rook, and Quenby Marblegate (2014)."


@pytest.mark.parametrize("authors, fragment, offered", [
    # The author list in full names beside the citation's initials (EBYSBC HFAJCC's shape).
    (_PAPER_AUTHORS, _FULL_AUTHOR_NAMES, False),
    (_PAPER_AUTHORS, "Annette M. Quill, Bertram T. Rook, Quenby Marblegate, Cecilia Fenn", False),
    # Two of its three authors, neither of them the first, is the list too.
    (_PAPER_AUTHORS, "Bertram Tobias Rook, Quenby Marblegate (2014).", False),
    # A surname's particles, written in lower case, are a name's.
    ("van der Berg K, Quill AM", "Katrien Marieke van der Berg and Annette Marisol Quill", False),
    # Naming two authors in a sentence, or going on into the title, is prose.
    (_PAPER_AUTHORS, "Designed the cores study with Rook and Marblegate, and wrote the manuscript", True),
    (_PAPER_AUTHORS, "Annette Quill, Bertram Rook and Quenby Marblegate cored the marshes together", True),
    # Naming one of several authors is not their list.
    (_PAPER_AUTHORS, "Bertram Tobias Rook Lecture Award, Example Estuarine Society", True),
    # Two authors' surnames heading a prize and a meeting are no list's names.
    (_PAPER_AUTHORS, "Quill Early Career Prize, Rook Foundation Annual Meeting", True),
    # A sole author named: the list is that author.
    ("Quill A", "Annette Marisol Quill, Example Estuarine Laboratory Director", False),
    # No author field to compare with.
    (None, _FULL_AUTHOR_NAMES, True),
])
def test_a_publications_author_list_in_another_form_is_not_prose(authors, fragment, offered):
    entry = _paper(f"{_PAPER_CITATION}\t{fragment}", authors=authors)
    assert sp.unrendered_prose(entry, "S1", _haystack(_PAPER_CITATION)) == ([fragment] if offered else [])
    # Only a publication's: another family's line naming the same people is its own.
    other = dict(entry, taxonomy_code="H")
    assert sp.unrendered_prose(other, "H", _haystack(_PAPER_CITATION)) == [fragment]


@pytest.mark.parametrize("heading, offered", [
    # A heading of the CV's own bibliography stage 2 joined to the entry (EBYSBC HTNNHG 153).
    ("3. Short Reports, Methods Notes, Letters", False),
    ("B. Invited Reviews, Editorials and Commentaries", False),
    ("IV. Peer-Reviewed Original Research Articles:", False),
    ("(c) Estuarine Research Reports & Coastal Monographs", False),
    # A word in lower case, a phrase too long, a phrase naming no kind of publication: prose.
    ("Cited in Estuarine Reviews and Coastal Reports", True),
    ("Featured Estuarine Sediment Transport Reviews", True),
    ("Estuarine Science Highlights, Coastal Reports", True),
])
def test_a_bibliography_heading_fused_onto_a_publication_is_not_prose(heading, offered):
    entry = _paper(f"{heading}\t{_PAPER_CITATION}")
    assert sp.unrendered_prose(entry, "S2", _haystack(_PAPER_CITATION)) == ([heading] if offered else [])
    # Only a publication's.
    assert sp.unrendered_prose(dict(entry, taxonomy_code="H"), "H", _haystack(_PAPER_CITATION)) == [heading]


_NEXT_PAPER = "22. Fenn CD, Quill AM. Saltmarsh accretion under rising tides. Coastal Studies Review."


def test_a_publication_line_holding_several_numbered_citations_offers_none():
    # The CV's next paper fused into the entry (EBYSBC CYOFWJ 344): its
    # citation would print as this one's prose, under this one's citation.
    fused = _paper(f"21. {_PAPER_CITATION}\t{_NEXT_PAPER}\tCoastal Stud Rev 9:101-12, 2005.")
    assert sp._own_citation_line(fused) == ""
    assert _plan_paper(fused) == []
    # One numbered citation, then its prose: offered.
    (planned,) = _plan_paper(_paper(f"21. {_PAPER_CITATION}\t{_PAPER_PROSE}"))
    assert planned.paragraphs == (_PAPER_PROSE,)
    # A numbered list of contributions names no author: not citations.
    steps = ("1) Designed the sampling study across three restored marshes",
             "2) Analysed the sediment cores and wrote the manuscript")
    (planned,) = _plan_paper(_paper("\t".join((_PAPER_CITATION, *steps))))
    assert planned.paragraphs == steps
    # Nor does one whose verb an acronym follows, then a word in lower case.
    steps = ("1) Designed MRI protocols for the sediment cores of three restored marshes",
             "2) Analysed CT scans of every core and wrote the manuscript")
    assert not any(map(sp._NUMBERED_CITATION_RE.match, steps))
    (planned,) = _plan_paper(_paper("\t".join((_PAPER_CITATION, *steps))))
    assert planned.paragraphs == steps
    # An author's initials followed by "and" or "et al." still open a citation.
    assert sp._NUMBERED_CITATION_RE.match("22. Fenn CD and Quill AM. Saltmarsh accretion")
    assert sp._NUMBERED_CITATION_RE.match("22. Fenn CD et al. Saltmarsh accretion")


_OTHER_PAPER = "Fenn CD, Quill AM. Saltmarsh accretion under rising tides. Coastal Studies Review"


@pytest.mark.parametrize("separator", ["\t", "\n", " | "], ids=["tab", "newline", "pipe"])
@pytest.mark.parametrize("locator", [". 2011;9:101-12.", ". 2014;9:101-12.", " 9:101-12, 2011"],
                         ids=["other_year", "same_year", "trailing_year"])
def test_a_publication_line_holding_another_papers_citation_offers_none(separator, locator):
    # The CV's next paper joined to the entry's line with no list number
    # (#1205 review): offered, its citation would print as this paper's
    # prose, and the doctor would read that paper as rendered.
    other = f"{_OTHER_PAPER}{locator}"
    entry = _paper(f"{_PAPER_CITATION}{separator}{other}")
    assert sp._own_citation_line(entry) == entry["text"]
    assert sp._cites_another_paper(" ".join(other.split()), entry)
    assert sp.unrendered_prose(entry, "S1", _haystack(_PAPER_CITATION), entry["text"]) == []
    assert _plan_paper(entry) == []
    # Nor the prose beside it: which paper it is about cannot be told.
    assert _plan_paper(_paper(f"{_PAPER_CITATION}{separator}{other}{separator}{_PAPER_PROSE}")) == []
    # Only a publication's: another family's line naming people is its own.
    other_family = dict(entry, taxonomy_code="H")
    assert sp.unrendered_prose(other_family, "H", _haystack(_PAPER_CITATION)) == [" ".join(other.split())]


def _wrapped_own_citation():
    """The entry's own citation broken at its printed lines, its authors and
    year on the first (EBYSBC XWNZWW 606's shape): written unlike the
    rendered citation, the first line is offered, and reads as a citation
    whose title is not the entry's."""
    return _paper("\n".join((
        "• Quill AM, Rook BT, Marblegate Q, et al. for the Example Marsh Restoration Trial 2014",
        "(EMRT) Study investigators (Q Marblegate listed as investigator) “Sediment budgets of",
        f"restored tidal marshes” {_PAPER_JOURNAL} 2014;12:45-52.")))


def _untitled_paper():
    """A paper stage 4 found no title for: whose citation a fragment is
    cannot be told."""
    return _paper(f"{_PAPER_CITATION}\t{_OTHER_PAPER}. 2011;9:101-12.", title=None)


@pytest.mark.parametrize("make", [_wrapped_own_citation, _untitled_paper], ids=["wrapped_own", "untitled"])
def test_a_line_whose_citation_cannot_be_told_from_another_papers_offers_none(make):
    assert _plan_paper(make()) == []


@pytest.mark.parametrize("fragment", [
    # A prize, a meeting or a society written in capitals gives no initials,
    # or fewer than the people it names.
    "Runner-up, Example Coastal Prize, Example Estuarine Society Annual Meeting, 2015",
    "Featured Article, Example Coastal Society Newsletter, 2016",
    "Rook BT Medal, Example Estuarine Society Annual Meeting, 2015",
    # Authors named in a sentence: no author list opens it.
    "Analysed the sediment cores and drafted the methods with Fenn CD, Rook BT, Quill AM, 2015",
    # Authors with no year: the next printed line of a citation.
    "Fenn CD, Quill AM. Saltmarsh accretion under rising tides",
])
def test_a_publication_fragment_that_is_no_other_papers_citation_is_offered(fragment):
    entry = _paper(f"{_PAPER_CITATION}\t{fragment}")
    assert not sp._cites_another_paper(fragment, entry)
    assert [p.paragraphs for p in _plan_paper(entry)] == [(fragment,)]


def test_a_publication_line_restating_its_own_citation_in_another_form_is_offered_as_written():
    # The CV's own line, list number, author markers and journal abbreviation
    # and all, then what the owner did: it opens like a citation and holds
    # the entry's title, so it is the entry's own (#1205 comment 6282590626,
    # option 1), not another paper's.
    line = ("39. Quill AM*, Rook BT, Marblegate Q. Sediment Budgets of Restored Tidal Marshes. J Estuar "
            "Sediment Stud. 2014 Mar;12(3):45-52. *Designed the sampling study, analysed the sediment "
            "cores, drafted every figure and wrote the manuscript")
    entry = _paper(line)
    assert sp._opens_with_author_list(line) and not sp._cites_another_paper(line, entry)
    assert [p.paragraphs for p in _plan_paper(entry)] == [(line,)]


@pytest.mark.parametrize("fragment, offered", [
    # The legend of the bibliography's author markers (YUYVIG SLYBST 159's shape).
    ("(† joint first authors; ‡ joint senior authors)", False),
    ("(† joint first authors; ‡ joint senior authors;)", False),
    ("* joint first authors, # corresponding author", False),
    # One marker and its note can be the paper's own.
    ("*Joint first author with the second author", True),
    ("*Joint first author with the second author; designed the sampling study", True),
    ("*Designed the sampling study; *analysed the cores", True),
    # A note that opens with a marker and goes on: not a legend.
    ("*Designed the sampling study; † analysed the cores; wrote the manuscript", True),
])
def test_an_author_marker_legend_fused_onto_a_publication_is_not_prose(fragment, offered):
    entry = _paper(f"{_PAPER_CITATION}\t{fragment}")
    assert sp.unrendered_prose(entry, "S1", _haystack(_PAPER_CITATION)) == ([fragment] if offered else [])
    # Only a publication's.
    assert sp.unrendered_prose(dict(entry, taxonomy_code="H"), "H", _haystack(_PAPER_CITATION)) == [fragment]


def test_a_publication_line_restating_its_citation_with_a_note_is_offered_as_written():
    # The CV's own citation format, with where it was presented: the note is
    # the paper's prose, and the line is offered as written (#1205 comment
    # 6282590626, option 1).
    line = ("Quill, A. M., Rook, B. T. and Marblegate, Q.: Sediment Budgets of Restored Tidal Marshes "
            "(Presented at the Example Coastal Ecology Congress, Norvale, May 2013; winner of the "
            "Congress Student Paper Prize and featured in the Congress newsletter)")
    (planned,) = _plan_paper(_paper(line))
    assert planned.paragraphs == (line,)


@pytest.mark.parametrize("segments, index, read", [
    # The last of two records, its line's two segments pairing with them: its own.
    ([_LATER_PAPER_CITATION, f"{_PAPER_CITATION} {_PAPER_PROSE}"], 1, f"{_PAPER_CITATION} {_PAPER_PROSE}"),
    # The last of two records on a line of three segments: none of it.
    ([_LATER_PAPER_CITATION, _PAPER_PROSE, _PAPER_CITATION], 1, ""),
])
def test_a_publication_reads_its_prose_from_its_own_line(segments, index, read):
    entry = dict(_degree_record("\t".join(segments), index, 2), taxonomy_code="S1")
    assert sp._prose_text("S1", entry, {}) == read


def test_a_publication_leaves_a_dated_record_line_to_the_recovery():
    # #2010's rule, every family's: a record of its own beside the paper's,
    # which the #221 recovery re-emits, stays the recovery's. One it reads
    # as rendered, and a line that is no record of its own, stay the
    # sub-point.
    sub_record = "Featured Article, Example Coastal Society Newsletter, 2015-2016"
    entry = _paper(f"{_PAPER_CITATION}\t{sub_record}\t{_PAPER_PROSE}")
    (offered,) = _plan_paper(entry)
    assert offered.paragraphs == (sub_record, _PAPER_PROSE)
    (planned,) = _plan_paper(entry, record_lines=lambda code, e: [
        sp.RecordLine(_PAPER_CITATION, False, False), sp.RecordLine(sub_record, True, True)])
    assert planned.paragraphs == (_PAPER_PROSE,)
    (kept,) = _plan_paper(entry, record_lines=lambda code, e: [
        sp.RecordLine(sub_record, False, True), sp.RecordLine(_PAPER_PROSE, True, False)])
    assert kept.paragraphs == (sub_record, _PAPER_PROSE)


def test_plan_offers_no_prose_of_a_publication_the_pii_pass_withheld():
    entry = _paper()
    (planned,) = _plan_paper(entry)
    assert planned.paragraphs == (_PAPER_PROSE,)

    entry["_pii_withheld"] = True
    assert _plan_paper(entry) == []


def test_a_line_a_training_position_and_a_publication_both_hold_stays_the_training_positions():
    # A fellow's research project, named under the fellowship and under the
    # paper it led to. The training position, a family shipped before S1,
    # keeps it though S1 would sort first; the paper keeps its other line,
    # and alone it offers both.
    project = "Research project on saltmarsh accretion rates, supervised by the fellowship programme director"
    table, training_lines = _training_rows(_RESIDENCY_ROW)
    paragraphs, paper_lines = _citations(_PAPER_CITATION)
    training = _training(f"{_RESIDENCY} July 2009-June 2010\tNorvale General Hospital\t{project}")
    paper = _paper(f"{_PAPER_CITATION}\t{project}\t{_PAPER_PROSE}")
    lines = [*training_lines, *(line._replace(position=line.position + 5) for line in paper_lines)]
    owned = {id(training): [table.rows[0]._tr], id(paper): [paragraphs[0]._p]}
    spans = {"C": sp.SectionSpan(0, 3), "S1": sp.SectionSpan(5, None)}
    planned = sp.plan_subpoints({"S1": [paper], "C": [training]}, _haystack(_PAPER_CITATION), lines, spans,
                                rows_written=lambda e: owned.get(id(e), []))
    assert [(p.code, p.paragraphs) for p in planned] == [("C", (project,)), ("S1", (_PAPER_PROSE,))]
    (alone,) = sp.plan_subpoints({"S1": [paper]}, _haystack(_PAPER_CITATION), lines, spans,
                                 rows_written=lambda e: owned.get(id(e), []))
    assert alone.paragraphs == (project, _PAPER_PROSE)


def test_only_a_mentee_code_is_placed_over_both_mentee_headings():
    # The section files a past mentee whose range is open under Current Mentees.
    lines = _mentee_tables("Ana Li")
    planned = sp.plan_subpoints({"N3B": [_postdoc()]}, _haystack(), lines,
                                {"N3A": sp.SectionSpan(0, 5), "N3B": sp.SectionSpan(5, None)})
    assert [p.anchor for p in planned] == [lines[1]]
    line = _candidate("Associate Professor | Example University | 2010-2015", position=0)
    assert sp.plan_subpoints({"D1": [_appointment()]}, _haystack(line.text), [line],
                             {"N3A": sp.SectionSpan(0, 5), "D1": sp.SectionSpan(5, None)}) == []


def test_section_span_bounds():
    span = sp.SectionSpan(3, 7)
    assert [span.holds(_candidate("x", position=i)) for i in (2, 3, 6, 7)] == [False, True, True, False]
    assert sp.SectionSpan(3, None).holds(_candidate("x", position=10_000))


def test_anchor_candidates_read_a_row_as_one_line_and_stop_at_the_appendix():
    doc = Document()
    doc.add_paragraph("Academic Appointments")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Associate Professor", "Example University",
                                                "2010-2015"], strict=True):
        cell.text = text
    appendix = doc.add_paragraph("T. APPENDIX")
    doc.add_paragraph("Associate Professor Example University 2010-2015")
    found = list(sp.anchor_candidates(doc, appendix._p))
    assert [c.text for c in found] == ["Academic Appointments",
                                       "Associate Professor Example University 2010-2015"]
    assert found[1].paragraph._p is table.rows[0].cells[0].paragraphs[0]._p
    assert [c.position for c in found] == [0, 1]


# --- the markup ------------------------------------------------------------------------

def _numbered_paragraph(doc, text, level="0"):
    para = doc.add_paragraph(text)
    ppr = para._p.get_or_add_pPr()
    numpr = ppr._add_numPr()
    numpr.get_or_add_ilvl().val = int(level)
    numpr.get_or_add_numId().val = 1
    return para


def _accept_all(body):
    """Word's Accept All Changes on this module's markup: a deleted row goes,
    deleted runs go, and a paragraph whose mark is deleted merges into the
    next one (which keeps its own properties)."""
    body = copy.deepcopy(body)
    for tr in [tr for tr in body.iter(qn("w:tr")) if tr.find(f"{qn('w:trPr')}/{qn('w:del')}") is not None]:
        tr.getparent().remove(tr)
    for d in [d for d in body.iter(qn("w:del")) if d.getparent().tag not in (qn("w:rPr"), qn("w:trPr"))]:
        d.getparent().remove(d)
    for p in [p for p in body.iter(qn("w:p")) if p.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None]:
        nxt = p.getnext()
        assert nxt is not None and nxt.tag == qn("w:p"), "a deleted mark with nothing to merge into"
        assert not [c for c in p if c.tag != qn("w:pPr")], "deleted paragraph kept content"
        p.getparent().remove(p)
    return body


def _texts(body):
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in body.iter(qn("w:p"))]


def _reject_all_texts(body):
    return ["".join(t.text or "" for t in p.iter(qn("w:t"), qn("w:delText")))
            for p in body.iter(qn("w:p"))]


def _subpoint(paragraph, *texts):
    return sp.SubPoint(_candidate(paragraph.text, paragraph=paragraph), texts, 0, "D1", 0)


def _plain_entry_paragraph(doc, text, left_twips):
    para = doc.add_paragraph(text)
    para.paragraph_format.left_indent = Twips(left_twips)
    return para


@pytest.mark.parametrize(("anchor_left", "expected_left"), [(1080, "1800"), (1800, "2520")])
def test_a_sub_point_under_an_unbulleted_entry_goes_one_indent_step_deeper(anchor_left,
                                                                           expected_left):
    # #1943: a teaching entry has no w:numPr, only a left indent on the
    # bullet-text grid, so the sub-point deepens that indent instead of ilvl.
    doc = Document()
    anchor = _plain_entry_paragraph(doc, "Course Director, Example Course", anchor_left)
    doc.add_paragraph("Next line")

    sp.write_subpoint(_subpoint(anchor, PROSE), sp.Revision(40, _DATE))

    new = anchor._p.getnext()
    assert new.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None
    assert new.find(f"{qn('w:pPr')}/{qn('w:ind')}").get(qn("w:left")) == expected_left
    assert anchor._p.pPr.find(qn("w:ind")).get(qn("w:left")) == str(anchor_left)


def test_a_sub_point_under_an_off_grid_indent_keeps_the_anchor_indent():
    # Only the unbulleted-entry grid deepens; any other indent is copied as is.
    doc = Document()
    anchor = _plain_entry_paragraph(doc, "Course Director, Example Course", 720)
    doc.add_paragraph("Next line")

    sp.write_subpoint(_subpoint(anchor, PROSE), sp.Revision(40, _DATE))

    assert anchor._p.getnext().find(f"{qn('w:pPr')}/{qn('w:ind')}").get(qn("w:left")) == "720"


def _gap_flags(anchor):
    """Whether the anchor and each paragraph after it carry the entry gap."""
    flags = []
    for p in [anchor._p, *anchor._p.itersiblings(qn("w:p"))]:
        flags.append(is_entry_end(p.pPr))
    return flags


def test_sub_points_under_an_entrys_last_line_pass_the_entry_gap_to_the_last_one():
    # An entry's last line carries the gap to the next entry
    # (`mark_entry_end`). Sub-points written under it stay tight to each
    # other, and the last one carries the gap, so it still lands before the
    # next entry; the anchor keeps it, for the document with changes accepted.
    doc = Document()
    anchor = _plain_entry_paragraph(doc, "Course Director, Example Course", 1080)
    mark_entry_end(anchor)
    doc.add_paragraph("Next entry")

    revision = sp.Revision(0, _DATE)
    sp.write_subpoint(_subpoint(anchor, PROSE, OTHER_PROSE), revision)
    assert _gap_flags(anchor) == [True, False, True, False]

    sp.write_subpoint(_subpoint(anchor, "A third sub-point of prose"), revision)
    assert _gap_flags(anchor) == [True, False, False, True, False]
    after = anchor._p.getnext().pPr
    assert after.find(qn("w:spacing")) is None or after.find(qn("w:spacing")).get(qn("w:after")) is None


def test_sub_points_under_a_line_without_the_entry_gap_copy_its_properties_unchanged():
    doc = Document()
    anchor = _plain_entry_paragraph(doc, "Course Director, Example Course", 1080)
    anchor.paragraph_format.space_after = Twips(240)
    doc.add_paragraph("Next line")

    revision = sp.Revision(0, _DATE)
    sp.write_subpoint(_subpoint(anchor, PROSE, OTHER_PROSE), revision)
    sp.write_subpoint(_subpoint(anchor, "A third sub-point of prose"), revision)

    subpoints = list(anchor._p.itersiblings(qn("w:p")))[:3]
    for p in subpoints:
        assert not is_entry_end(p.pPr)
        assert p.pPr.find(qn("w:spacing")).get(qn("w:after")) == "240"


@pytest.mark.parametrize("anchor_kind", ["bullet", "citation"])
def test_a_body_sub_point_is_a_deleted_paragraph_one_level_deeper(anchor_kind):
    doc = Document()
    anchor = _numbered_paragraph(doc, "Course Director, Example Course" if anchor_kind == "bullet"
                                 else "1. Author A. Example article. J Example. 2010;1:1-2.")
    doc.add_paragraph("Next line")
    before = _texts(doc.element.body)
    revision = sp.Revision(40, _DATE)

    sp.write_subpoint(_subpoint(anchor, PROSE, OTHER_PROSE), revision)

    new = anchor._p.getnext()
    assert new.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:ilvl')}").get(qn("w:val")) == "1"
    assert len(_deleted_marks(new)) == 2          # the mark and the run
    assert list(new.iter(qn("w:t"))) == []
    assert _reject_all_texts(doc.element.body)[1:3] == [PROSE, OTHER_PROSE]
    assert _texts(_accept_all(doc.element.body)) == before
    assert revision.next_id == 44
    ids = [d.get(qn("w:id")) for d in _deleted_marks(doc.element.body)]
    assert ids == ["40", "41", "42", "43"]
    assert {d.get(qn("w:date")) for d in _deleted_marks(doc.element.body)} == {_DATE}


def test_a_second_sub_point_under_one_line_follows_the_first():
    doc = Document()
    anchor = doc.add_paragraph("Course Director, Example Course")
    doc.add_paragraph("Next line")
    revision = sp.Revision(0, _DATE)
    sp.write_subpoint(_subpoint(anchor, PROSE), revision)
    sp.write_subpoint(_subpoint(anchor, OTHER_PROSE), revision)
    assert _reject_all_texts(doc.element.body)[:4] == [
        "Course Director, Example Course", PROSE, OTHER_PROSE, "Next line"]


def _followed_by(doc, neighbour):
    """A bullet anchor with a table or the section properties straight after it."""
    anchor = _numbered_paragraph(doc, "Course Director, Example Course")
    if neighbour == "table":
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "Table cell text"
        doc.add_paragraph("After the table")
    return anchor


@pytest.mark.parametrize("neighbour", ["table", "sectPr"])
def test_a_sub_point_before_a_table_or_the_section_end_keeps_its_last_mark(neighbour):
    doc = Document()
    anchor = _followed_by(doc, neighbour)
    before = _texts(doc.element.body)
    assert anchor._p.getnext().tag == qn("w:tbl" if neighbour == "table" else "w:sectPr")

    sp.write_subpoint(_subpoint(anchor, PROSE, OTHER_PROSE), sp.Revision(0, _DATE))

    first, last = anchor._p.getnext(), anchor._p.getnext().getnext()
    assert first.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None
    assert last.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is None
    assert last.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None
    assert _reject_all_texts(doc.element.body)[1:3] == [PROSE, OTHER_PROSE]
    accepted = _accept_all(doc.element.body)
    # One empty plain paragraph is left, not an empty bullet, and the table text is its own.
    assert _texts(accepted) == before[:1] + [""] + before[1:]
    assert accepted.findall(qn("w:p"))[1].find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None


def test_a_second_sub_point_before_a_table_deletes_the_first_ones_kept_mark():
    doc = Document()
    anchor = _followed_by(doc, "table")
    before = _texts(doc.element.body)
    revision = sp.Revision(0, _DATE)
    sp.write_subpoint(_subpoint(anchor, PROSE), revision)
    sp.write_subpoint(_subpoint(anchor, OTHER_PROSE), revision)
    assert _reject_all_texts(doc.element.body)[:3] == [
        "Course Director, Example Course", PROSE, OTHER_PROSE]
    first = anchor._p.getnext()
    assert first.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None
    assert first.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:ilvl')}").get(qn("w:val")) == "1"
    assert _texts(_accept_all(doc.element.body)) == before[:1] + [""] + before[1:]


def _appointments_table(doc):
    table = doc.add_table(rows=3, cols=3)
    for row, title in zip(table.rows, ["Title", "Associate Professor", "Assistant Professor"],
                          strict=True):
        row.cells[0].text, row.cells[1].text, row.cells[2].text = title, "Example University", "2010"
    doc.add_paragraph("After the table")
    return table


def test_a_table_sub_point_is_a_deleted_row_spanning_the_table_under_its_row():
    doc = Document()
    table = _appointments_table(doc)
    before = _texts(doc.element.body)

    sp.write_subpoint(_subpoint(table.rows[1].cells[0].paragraphs[0], PROSE), sp.Revision(0, _DATE))

    rows = table._tbl.findall(qn("w:tr"))
    assert len(rows) == 4
    new = rows[2]
    assert new.find(f"{qn('w:trPr')}/{qn('w:del')}").get(qn("w:author")) == sp.SUBPOINT_AUTHOR
    cells = new.findall(qn("w:tc"))
    assert len(cells) == 1
    assert cells[0].find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}").get(qn("w:val")) == "3"
    assert list(new.iter(qn("w:t"))) == []
    assert "".join(t.text for t in new.iter(qn("w:delText"))) == PROSE
    assert _texts(_accept_all(doc.element.body)) == before


def test_a_label_table_takes_its_sub_point_at_its_end():
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row, label in zip(table.rows, ["Title:", "Agency:", "Notes:"], strict=True):
        row.cells[0].text, row.cells[1].text = label, "Example value"
    sp.write_subpoint(_subpoint(table.rows[0].cells[0].paragraphs[0], PROSE), sp.Revision(0, _DATE))
    rows = table._tbl.findall(qn("w:tr"))
    assert rows[-1].find(f"{qn('w:trPr')}/{qn('w:del')}") is not None
    assert [r.find(f"{qn('w:trPr')}/{qn('w:del')}") is None for r in rows[:3]] == [True] * 3


def test_a_table_whose_rows_are_not_all_labels_is_not_a_label_table():
    doc = Document()
    table = _appointments_table(doc)
    assert sp._is_label_table(table._tbl) is False
