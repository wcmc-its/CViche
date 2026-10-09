"""Tests for doctor/docx_diff.py and its CLI, scripts/docx_review_diff.py (#1587).

Every document is built here from synthetic text: a delivered file, then a
corrected copy with one kind of reviewer edit, made both as a pending tracked
revision and as a plain (accepted) edit where the two must agree.

    python3 -m pytest src/unified_pipeline/tests/test_doctor_docx_diff.py -p no:cacheprovider
"""
import copy
import importlib.util
import io
import json
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_ROOT = Path(__file__).resolve().parents[3]
_SRC = _ROOT / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.doctor import docx_diff as dd  # noqa: E402
from unified_pipeline.stage6.formatting import add_cviche_box  # noqa: E402

UID = "TSTREV"
CITE_A = "Quorvane T, Plesk M. Heliotropic drift in vexillary marmosets. J Synth Imag. 2019;12:34-56."
CITE_B = "Quorvane T, Abernoth R. Saltwick lattices under brennic loading. Synth Rep. 2021;3:7-19."
CITE_C = "Quorvane T. Tremulant gradients across fennish moorland. Moor Synth. 2016;8:101-110."
GRANT = "Brennic Foundation Award: Saltwick lattice modelling, 2018-2022, Principal Investigator"
LOST = "Grimsby Varnell Lectureship: Vexillary marmoset cognition symposium, Lisbonne 2017"


def _ins_run(paragraph, text):
    """Append `text` to `paragraph` as a pending tracked insertion."""
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "901")
    ins.set(qn("w:author"), "Synthetic Reviewer")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    ins.append(run)
    paragraph._p.append(ins)


def _track_delete(paragraph):
    """Turn every run of `paragraph` into a pending tracked deletion."""
    for run in list(paragraph._p.iterchildren(qn("w:r"))):
        for t in run.iterchildren(qn("w:t")):
            t.tag = qn("w:delText")
        wrapper = OxmlElement("w:del")
        wrapper.set(qn("w:id"), "902")
        wrapper.set(qn("w:author"), "Synthetic Reviewer")
        run.addprevious(wrapper)
        wrapper.append(run)


def _wrap_runs(paragraph, tag):
    """Move every run of `paragraph` inside one `tag` element (w:moveFrom / w:moveTo)."""
    wrapper = OxmlElement(tag)
    wrapper.set(qn("w:id"), "903")
    wrapper.set(qn("w:author"), "Synthetic Reviewer")
    runs = list(paragraph._p.iterchildren(qn("w:r")))
    runs[0].addprevious(wrapper)
    for run in runs:
        wrapper.append(run)


def _delivered(*, cite_a_as_insertion=False, with_box=False):
    """RESEARCH (a grant table), then BIBLIOGRAPHY with three citations."""
    doc = Document()
    doc.add_paragraph("RESEARCH")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = GRANT.split(": ", 1)
    doc.add_paragraph("BIBLIOGRAPHY")
    if cite_a_as_insertion:
        _ins_run(doc.add_paragraph(), CITE_A)  # how stage 6 writes an enriched citation
    else:
        doc.add_paragraph(CITE_A)
    doc.add_paragraph(CITE_B)
    doc.add_paragraph(CITE_C)
    if with_box:
        add_cviche_box(doc, "CViche note: delete this box before sending")
    return doc


def _clone(doc):
    """An independent copy, the way a reviewer gets one: saved and reopened."""
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return Document(buffer)


def _paragraph(doc, startswith):
    return next(p for p in doc.paragraphs if dd._accepted_text(p._p).startswith(startswith))


def _diff(before, after, stage_json=None):
    entries = dd.entry_tokens(stage_json) if stage_json else []
    return dd.diff_blocks(dd.read_blocks(before), dd.read_blocks(after), entries)


def _kinds(changes):
    return [c.change_type for c in changes]


def test_identical_copy_has_no_changes():
    assert _diff(_delivered(), _clone(_delivered())) == []


def test_stage6_tracked_insertion_reads_as_present_and_rejecting_it_is_a_deletion():
    """#1167: python-docx's paragraph.text is '' for a citation held in w:ins."""
    delivered = _delivered(cite_a_as_insertion=True)
    assert _paragraph(delivered, CITE_A[:20]).text == ""
    assert _diff(delivered, _clone(delivered)) == []

    rejected = _clone(delivered)
    para = _paragraph(rejected, CITE_A[:20])
    para._p.getparent().remove(para._p)
    changes = _diff(delivered, rejected)
    assert _kinds(changes) == [dd.CHANGE_DELETED]
    assert changes[0].before_chars == len(CITE_A)


def test_pending_and_accepted_deletion_are_the_same_change():
    delivered = _delivered()
    pending = _clone(delivered)
    _track_delete(_paragraph(pending, CITE_B[:20]))
    accepted = _clone(delivered)
    para = _paragraph(accepted, CITE_B[:20])
    para._p.getparent().remove(para._p)

    assert dd.count_revisions(pending).deletions == 1
    assert _diff(delivered, pending) == _diff(delivered, accepted)
    change, = _diff(delivered, pending)
    assert (change.change_type, change.section_before) == (dd.CHANGE_DELETED, "BIBLIOGRAPHY")


def test_pending_and_plain_insertion_are_the_same_change():
    delivered = _delivered()
    pending, plain = _clone(delivered), _clone(delivered)
    _ins_run(pending.add_paragraph(), LOST)
    plain.add_paragraph(LOST)

    assert dd.count_revisions(pending).insertions == 1
    change, = _diff(delivered, pending)
    assert _diff(delivered, plain) == [change]
    assert (change.change_type, change.before_chars, change.after_chars) == (
        dd.CHANGE_ADDED, 0, len(LOST))


def test_value_edit_is_one_change_with_both_lengths():
    corrected = _clone(_delivered())
    para = _paragraph(corrected, CITE_C[:20])
    edited = CITE_C.replace("2016", "2015")
    para.text = edited
    change, = _diff(_delivered(), corrected)
    assert change.change_type == dd.CHANGE_VALUE_EDITED
    assert (change.before_chars, change.after_chars) == (len(CITE_C), len(edited))
    assert (change.block_before, change.block_after) == (5, 5)


def test_unrelated_replacement_is_a_deletion_and_an_addition():
    corrected = _clone(_delivered())
    _paragraph(corrected, CITE_C[:20]).text = LOST
    assert sorted(_kinds(_diff(_delivered(), corrected))) == [dd.CHANGE_ADDED, dd.CHANGE_DELETED]


def test_cut_and_paste_into_another_section_is_one_move():
    corrected = _clone(_delivered())
    para = _paragraph(corrected, CITE_B[:20])
    body = para._p.getparent()
    body.remove(para._p)
    _paragraph(corrected, "RESEARCH")._p.addnext(para._p)
    change, = _diff(_delivered(), corrected)
    assert (change.change_type, change.section_before, change.section_after) == (
        dd.CHANGE_MOVED, "BIBLIOGRAPHY", "RESEARCH")


def test_tracked_move_reads_moveto_and_drops_movefrom():
    corrected = _clone(_delivered())
    source = _paragraph(corrected, CITE_B[:20])
    moved = copy.deepcopy(source._p)
    _wrap_runs(source, "w:moveFrom")
    _paragraph(corrected, "RESEARCH")._p.addnext(moved)
    from docx.text.paragraph import Paragraph
    _wrap_runs(Paragraph(moved, source._parent), "w:moveTo")

    assert dd.count_revisions(corrected).moves == 1
    assert _kinds(_diff(_delivered(), corrected)) == [dd.CHANGE_MOVED]


def test_removing_the_cviche_box_is_not_a_correction():
    assert _diff(_delivered(with_box=True), _delivered()) == []


def test_only_template_headings_are_reported_as_sections():
    doc = Document()
    doc.add_paragraph(CITE_A)                   # before any heading
    doc.add_paragraph("RESEARCH")
    doc.add_paragraph("QUORVANE SYNTHETIC LAB")  # all caps, but CV text
    doc.add_paragraph(CITE_B)
    assert [b.section for b in dd.read_blocks(doc)] == [None, "RESEARCH", "RESEARCH", "RESEARCH"]


def test_letter_prefixed_appendix_heading_is_a_section():
    doc = _delivered()
    doc.add_paragraph("T. APPENDIX")
    doc.add_paragraph(LOST)
    assert dd.read_blocks(doc)[-1].section == "APPENDIX"


_STAGE4 = {"entries": [
    {"element_idx_start": 12, "text": GRANT, "taxonomy_code": "I1"},
    {"element_idx_start": 30.0, "text": CITE_A, "taxonomy_code": "S1"},
    {"element_idx_start": 31, "text": CITE_B, "taxonomy_code": "S1"},
    {"element_idx_start": 474.1, "text": LOST, "taxonomy_code": "J"},
    {"element_idx_start": None, "text": CITE_C, "taxonomy_code": "S1"},
]}


def test_changes_map_to_entries_from_the_side_that_holds_the_text():
    corrected = _clone(_delivered())
    para = _paragraph(corrected, CITE_A[:20])
    para._p.getparent().remove(para._p)
    corrected.add_paragraph(LOST)            # a lost record restored by the reviewer
    _paragraph(corrected, CITE_C[:20]).text = CITE_C.replace("2016", "2015")
    changes = {c.change_type: c for c in _diff(_delivered(), corrected, _STAGE4)}
    assert changes[dd.CHANGE_DELETED].element_idx == 30
    assert changes[dd.CHANGE_ADDED].element_idx == 474   # truncated like doctor_vs_autopsy
    assert changes[dd.CHANGE_VALUE_EDITED].element_idx is None  # its entry has no index


def test_an_edit_maps_by_the_delivered_text_not_the_reviewers():
    """The delivered block is what stage 6 rendered from an entry; the
    reviewer's rewrite can read closer to a different one."""
    erratum = CITE_A + " Erratum lodged: Quorvane heliotropic correction."
    stage = {"entries": [{"element_idx_start": 30, "text": CITE_A},
                         {"element_idx_start": 40, "text": erratum}]}
    corrected = _clone(_delivered())
    _paragraph(corrected, CITE_A[:20]).text = erratum
    change, = _diff(_delivered(), corrected, stage)
    assert (change.change_type, change.element_idx) == (dd.CHANGE_VALUE_EDITED, 30)


def test_a_w_t_inside_w_del_is_still_deleted():
    """Word writes a deleted run's text as w:delText; some writers leave w:t."""
    doc = Document()
    para = doc.add_paragraph(CITE_A)
    wrapper = OxmlElement("w:del")
    run = next(para._p.iterchildren(qn("w:r")))
    run.addprevious(wrapper)
    wrapper.append(run)
    assert dd.read_blocks(doc) == []


def test_a_block_two_entries_match_equally_maps_to_none():
    entries = dd.entry_tokens({"entries": [{"element_idx_start": 1, "text": CITE_A},
                                           {"element_idx_start": 2, "text": CITE_A}]})
    assert dd.map_entry(CITE_A, entries) is None
    assert dd.map_entry(CITE_A, entries[:1]) == 1


def _load_script(name):
    spec = importlib.util.spec_from_file_location(name, _ROOT / "scripts" / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _cli_run(tmp_path):
    delivered, corrected = _delivered(), _clone(_delivered())
    para = _paragraph(corrected, CITE_A[:20])
    para._p.getparent().remove(para._p)
    corrected.add_paragraph(LOST)
    delivered.save(tmp_path / f"{UID}_wcm.docx")
    corrected.save(tmp_path / "corrected.docx")
    (tmp_path / "stage4.json").write_text(json.dumps(_STAGE4), encoding="utf-8")
    labels = tmp_path / "labels"
    rc = _load_script("docx_review_diff").main([
        str(tmp_path / f"{UID}_wcm.docx"), str(tmp_path / "corrected.docx"), "--uid", UID,
        "--entries", str(tmp_path / "stage4.json"), "--report", str(tmp_path / "report.json"),
        "--label", str(labels)])
    return rc, labels


def test_cli_label_loads_in_the_autopsy_label_store(tmp_path):
    rc, labels = _cli_run(tmp_path)
    assert rc == 0
    store = _load_script("doctor_vs_autopsy").load_labels(labels)
    assert store.uids == {UID}
    assert sorted((v.id, v.batch, v.severity, tuple(v.idxs)) for v in store.verified) == [
        (f"{UID}-R01", dd.REVIEW_LABEL_BATCH, "medium", (30,)),
        (f"{UID}-R02", dd.REVIEW_LABEL_BATCH, "high", (474,))]
    label = json.loads((labels / f"{UID}.json").read_text())
    assert {f["class"] for f in label["findings"]} == {"review_deleted", "review_added"}
    assert {f["batch_class"] for f in label["findings"]} == {"deleted", "added"}
    sva = _load_script("score_vs_autopsy")
    run = sva.Run(uid=UID, score=0.0, raw=0.0, band="", caps=[], penalties={}, lint_counts={})
    sva.apply_label(run, label, labels / f"{UID}.json")  # raises on an unknown severity
    assert (run.high, run.medium) == (1, 1)


def test_cli_outputs_carry_no_document_text(tmp_path):
    _cli_run(tmp_path)
    written = (tmp_path / "report.json").read_text() + (tmp_path / "labels" / f"{UID}.json").read_text()
    for word in ("Quorvane", "Heliotropic", "Varnell", "Lectureship", "Brennic"):
        assert word not in written
    report = json.loads((tmp_path / "report.json").read_text())
    assert report["by_type"] == {"deleted": 1, "added": 1, "moved": 0, "value_edited": 0}


def test_cli_unreadable_docx_exits_2(tmp_path, capsys):
    (tmp_path / "bad.docx").write_text("not a docx", encoding="utf-8")
    _delivered().save(tmp_path / "good.docx")
    cli = _load_script("docx_review_diff")
    assert cli.main([str(tmp_path / "good.docx"), str(tmp_path / "bad.docx"), "--uid", UID]) == 2
    assert "docx_review_diff:" in capsys.readouterr().err


@pytest.mark.parametrize("kind", dd.CHANGE_TYPES)
def test_every_change_type_has_a_label_severity(kind):
    assert dd.REVIEW_SEVERITY[kind] in {"high", "medium", "low"}
