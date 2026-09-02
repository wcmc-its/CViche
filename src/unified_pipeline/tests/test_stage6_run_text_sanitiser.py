"""#552: a control character in LLM-written text kills the render, and the
documented citation fallback does not rescue it.

lxml's own ``Element.text`` setter raises ``ValueError: All strings must be
XML compatible...`` on the XML-1.0-invalid C0 subset ``[\\x00-\\x08\\x0b\\x0c
\\x0e-\\x1f]`` -- the same range python-docx's ``Run.text`` setter rejects.
DEL (``\\x7f``) and the C1 range (``\\x80-\\x9f``) are valid XML and are
deliberately left alone; see ``test_del_and_c1_controls_are_valid_xml_and_
preserved`` below, which proves that in-test rather than asserting it in
prose. Three places in stage 6 build revision XML by hand and assign
straight into that setter, bypassing python-docx entirely:

- ``WCMTemplateGenerator._add_track_change_insertion``
  (stage_6_word_template.py, ``t.text = ...``)
- ``WCMTemplateGenerator._add_track_change_deletion``
  (stage_6_word_template.py, ``delText.text = ...``)
- ``BibliographySection._add_citation_with_bold_author_as_insertion``'s
  nested ``create_run_element`` (bibliography.py, ``t.text = ...``)

A fourth site, ``BibliographySection._add_citation_with_bold_author``, calls
python-docx's ``add_run``/``Run.text``, which reaches the same lxml setter
(#711 T2.3). All four now route through the shared ``BibliographySection.
_sanitize_run_text`` (reached via the mixin as ``self._sanitize_run_text``),
which strips that range and explicitly preserves ``\\t``, ``\\n``, ``\\r`` --
both valid XML and required by ``_clean_inline_tabs``'s label/value contract
(test_cell_separators.py:34-37).

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_run_text_sanitiser.py -p no:cacheprovider
"""

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6.sections.bibliography import (  # noqa: E402
    BibliographySection,
    _CONTROL_CHAR_PATTERN,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402

# One representative codepoint from each stripped sub-range, plus the three
# preserved whitespace controls that must survive untouched.
CONTROL_CHARS = ['\x00', '\x07', '\x0b', '\x0c', '\x0e', '\x1f']
PRESERVED_WHITESPACE = ['\t', '\n', '\r']
# DEL and the C1 range: valid XML 1.0, deliberately left alone (T2.1).
DEL_AND_C1_CHARS = ['\x7f', '\x85', '\x9f']


def _generator(emit_track_changes=True):
    gen = WCMTemplateGenerator(verbose=False, emit_track_changes=emit_track_changes)
    gen.doc = Document()
    return gen


def _blank_paragraph(gen):
    return gen.doc.add_paragraph()


# --- the shared helper itself ---


@pytest.mark.parametrize("char", CONTROL_CHARS)
def test_sanitize_run_text_strips_every_stripped_subrange_member(char):
    assert BibliographySection._sanitize_run_text(f"before{char}after") == "beforeafter"


@pytest.mark.parametrize("char", PRESERVED_WHITESPACE)
def test_sanitize_run_text_preserves_tab_newline_cr(char):
    assert BibliographySection._sanitize_run_text(f"before{char}after") == f"before{char}after"


def test_sanitize_run_text_leaves_ordinary_text_untouched():
    text = "Doe J, Smith A. A study of things. Journal. 2024;10(2):100-110."
    assert BibliographySection._sanitize_run_text(text) == text


@pytest.mark.parametrize("char", DEL_AND_C1_CHARS)
def test_del_and_c1_controls_are_valid_xml_and_preserved(char):
    # `_sanitize_run_text` leaves the character untouched...
    text = f"a{char}b"
    assert BibliographySection._sanitize_run_text(text) == text

    # ...and this is *why*: lxml's own `.text` setter (what a raw `w:t`
    # assignment reaches) accepts it without raising, unlike the stripped
    # C0 subset above.
    t = OxmlElement('w:t')
    t.text = text  # must not raise
    assert t.text == text


def test_control_char_pattern_never_matches_preserved_whitespace():
    # Regression guard on the regex itself, independent of the function that
    # wraps it: \x09/\x0a/\x0d must never be added to the stripped range.
    for char in PRESERVED_WHITESPACE:
        assert not _CONTROL_CHAR_PATTERN.search(char)
    for char in CONTROL_CHARS:
        assert _CONTROL_CHAR_PATTERN.search(char)


# --- site 1: stage_6_word_template.py:_add_track_change_insertion ---


def test_track_change_insertion_survives_control_character():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_insertion(para, "Enrolled 12\x0b patients")  # must not raise

    ins_elem = para._p.find(qn("w:ins"))
    assert ins_elem is not None
    t_elem = ins_elem.find(qn("w:r")).find(qn("w:t"))
    assert t_elem.text == "Enrolled 12 patients"


def test_track_change_insertion_preserves_tab_in_text():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_insertion(para, "Label\tValue")

    ins_elem = para._p.find(qn("w:ins"))
    t_elem = ins_elem.find(qn("w:r")).find(qn("w:t"))
    assert t_elem.text == "Label\tValue"


# --- site 2: stage_6_word_template.py:_add_track_change_deletion ---


def test_track_change_deletion_survives_control_character():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_deletion(para, "Enrolled 12\x0c patients")  # must not raise

    del_elem = para._p.find(qn("w:del"))
    assert del_elem is not None
    delText_elem = del_elem.find(qn("w:r")).find(qn("w:delText"))
    assert delText_elem.text == "Enrolled 12 patients"


def test_track_change_deletion_preserves_newline_in_text():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_deletion(para, "Line one\nLine two")

    del_elem = para._p.find(qn("w:del"))
    delText_elem = del_elem.find(qn("w:r")).find(qn("w:delText"))
    assert delText_elem.text == "Line one\nLine two"


# --- site 3: bibliography.py:_add_citation_with_bold_author_as_insertion ---


def test_citation_insertion_survives_control_character_without_falling_back():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)
    citation = "1. Doe J, Smith\x0b A. A study. Journal. 2024;10(2):100-110."

    gen._add_citation_with_bold_author_as_insertion(
        para, citation, "Smith A", "", author="PubMed Enrichment"
    )  # must not raise

    # The tracked-insertion path itself must have succeeded -- not the
    # except-block fallback to the plain (non-tracked) writer, which is the
    # documented "safety net" #552 found does not actually work for this
    # input class (it re-raises the identical error on the same string).
    assert gen.stats["track_changes_added"] == 1
    ins_elem = para._p.find(qn("w:ins"))
    assert ins_elem is not None
    rendered = "".join(
        (t.text or "") for r in ins_elem.findall(qn("w:r")) for t in r.findall(qn("w:t"))
    )
    assert "\x0b" not in rendered
    assert rendered == citation.replace("\x0b", "")


def test_citation_insertion_preserves_tab_in_bolded_author_split():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)
    citation = "1. Doe J, Smith A.\tA study with a tab. Journal. 2024;10(2):100-110."

    gen._add_citation_with_bold_author_as_insertion(
        para, citation, "Smith A", "", author="PubMed Enrichment"
    )

    ins_elem = para._p.find(qn("w:ins"))
    rendered = "".join(
        (t.text or "") for r in ins_elem.findall(qn("w:r")) for t in r.findall(qn("w:t"))
    )
    assert "\t" in rendered


# --- the xml:space="preserve" guard reads the SANITIZED string ---
#
# Round-2 finding. All three sites assign the sanitized text but originally
# tested the RAW text in their `startswith(' ') or endswith(' ')` guard. A
# control character in front of a leading space (`"\x0b Smith"`) is stripped,
# so the rendered run does begin with a space while the raw string does not --
# the guard skipped xml:space="preserve" and Word collapsed that space,
# gluing the run to its neighbour. Whitespace-only damage, but it is silent,
# which is exactly what the sanitiser was added to avoid.

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def test_insertion_preserves_space_exposed_by_stripping_a_control_char():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_insertion(para, "\x0b Smith")

    t_elem = para._p.find(qn("w:ins")).find(qn("w:r")).find(qn("w:t"))
    assert t_elem.text == " Smith"
    assert t_elem.get(XML_SPACE) == "preserve"


def test_deletion_preserves_space_exposed_by_stripping_a_control_char():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_deletion(para, "Smith \x0b")

    delText_elem = para._p.find(qn("w:del")).find(qn("w:r")).find(qn("w:delText"))
    assert delText_elem.text == "Smith "
    assert delText_elem.get(XML_SPACE) == "preserve"


def test_citation_insertion_preserves_space_exposed_by_stripping_a_control_char():
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)
    # No target name matches, so the whole citation is one run -- the raw
    # string starts and ends with a control character, the sanitized one with
    # a space.
    citation = "\x0b Doe J. A study. Journal. 2024;10(2):100-110. \x0b"

    gen._add_citation_with_bold_author_as_insertion(
        para, citation, "Nobody Q", "", author="PubMed Enrichment"
    )

    ins_elem = para._p.find(qn("w:ins"))
    t_elems = [t for r in ins_elem.findall(qn("w:r")) for t in r.findall(qn("w:t"))]
    assert len(t_elems) == 1
    assert t_elems[0].text == " Doe J. A study. Journal. 2024;10(2):100-110. "
    assert t_elems[0].get(XML_SPACE) == "preserve"


def test_ordinary_leading_space_still_preserved_at_all_three_sites():
    # Control arm: the guard's pre-existing behaviour for text that needs no
    # sanitising at all must be unchanged by the reordering.
    gen = _generator(emit_track_changes=True)

    para_ins = _blank_paragraph(gen)
    gen._add_track_change_insertion(para_ins, " Smith ")
    t_elem = para_ins._p.find(qn("w:ins")).find(qn("w:r")).find(qn("w:t"))
    assert t_elem.get(XML_SPACE) == "preserve"

    para_del = _blank_paragraph(gen)
    gen._add_track_change_deletion(para_del, " Smith ")
    delText_elem = para_del._p.find(qn("w:del")).find(qn("w:r")).find(qn("w:delText"))
    assert delText_elem.get(XML_SPACE) == "preserve"

    para_cit = _blank_paragraph(gen)
    gen._add_citation_with_bold_author_as_insertion(
        para_cit, " Doe J. A study. Journal. 2024;1:1-2. ", "Nobody Q", ""
    )
    t_elems = [
        t for r in para_cit._p.find(qn("w:ins")).findall(qn("w:r")) for t in r.findall(qn("w:t"))
    ]
    assert t_elems[0].get(XML_SPACE) == "preserve"


def test_text_without_boundary_spaces_still_gets_no_xml_space_attribute():
    # The other control arm: reordering must not start setting the attribute
    # on runs that never needed it.
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_track_change_insertion(para, "Enrolled 12\x0b patients")

    t_elem = para._p.find(qn("w:ins")).find(qn("w:r")).find(qn("w:t"))
    assert t_elem.text == "Enrolled 12 patients"
    assert t_elem.get(XML_SPACE) is None


# --- T2.2: the production entry point, not just the writers directly ---
#
# `_fill_bibliography` is the actual LLM-derived-data entry point; the tests
# above call the writers directly. This proves the sanitiser is reached on
# the real wire, against the real WCM template, for both the enriched
# (tracked-insertion + deletion) and non-enriched (plain, T2.3) paths.


def _real_template_generator():
    """A real generator instance against the actual WCM template -- the
    production entry point for `_fill_bibliography`, not a synthetic
    Document() double (matches test_stage6_bibliography_round2.py's own
    `_generator`)."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    return gen


def test_fill_bibliography_sanitises_control_character_in_enriched_entry():
    gen = _real_template_generator()
    # `formatted_citation` uses \x1f, not \x0b: \x0b is one of the handful of
    # codepoints `str.splitlines()` treats as a line break (so is \x0c), and
    # `split_fused_citation_entries` (normalization/records.py:63) calls
    # `.splitlines()` on this exact field upstream of the writers this test
    # targets -- a \x0b here would fuse-split the entry into two before the
    # sanitiser is ever reached, which is a different (untouched) code path.
    # `text` (the original, track-change-deletion side) is never
    # `.splitlines()`-ed, so it keeps \x0b as specified.
    formatted_citation = "Doe J\x1f, Smith A. A study. Journal. 2024;10(2):100-110."
    original_text = "Doe J. A stu\x0bdy (original). Journal. 2024;10(2):100-110."
    entry = {
        "extracted_fields": {
            "formatted_citation": formatted_citation,
            "formatting_source": "stage_5d_llm",
            "target_name": "Smith A",
            "year": 2024,
        },
        "enrichment_status": "enriched",
        "text": original_text,
        "enrichment_source": "pubmed",
    }

    gen._fill_bibliography({"S1": [entry]}, cv_owner={}, document_uid="")  # must not raise

    header_idx = gen._find_paragraph_with_text("Peer-reviewed Research Articles:")
    assert header_idx is not None
    para = gen.doc.paragraphs[header_idx + 2]  # header, blank separator, citation

    ins_elem = para._p.find(qn("w:ins"))
    assert ins_elem is not None
    ins_text = "".join(
        (t.text or "") for r in ins_elem.findall(qn("w:r")) for t in r.findall(qn("w:t"))
    )
    assert ins_text == f"1. {formatted_citation}".replace("\x1f", "")

    del_elem = para._p.find(qn("w:del"))
    assert del_elem is not None
    del_text = "".join(
        (d.text or "") for r in del_elem.findall(qn("w:r")) for d in r.findall(qn("w:delText"))
    )
    assert del_text == original_text.replace("\x0b", "")

    assert "\x1f" not in gen.doc.element.xml
    assert "\x0b" not in gen.doc.element.xml


def test_fill_bibliography_sanitises_control_character_in_plain_entry():
    # T2.3's integration proof: the non-enriched path goes through
    # `_add_citation_with_bold_author`, not the tracked-insertion writer.
    # \x1f, not \x0b -- see the comment in the enriched-entry test above on
    # why \x0b in `formatted_citation` hits an unrelated upstream fuse-split.
    gen = _real_template_generator()
    formatted_citation = (
        "Doe J, Smith A. A stu\x1fdy without enrichment. Journal. 2024;10(2):100-110."
    )
    entry = {
        "extracted_fields": {
            "formatted_citation": formatted_citation,
            "formatting_source": "stage_5d_llm",
            "target_name": "Smith A",
            "year": 2024,
        },
    }
    assert "enrichment_status" not in entry

    gen._fill_bibliography({"S1": [entry]}, cv_owner={}, document_uid="")  # must not raise

    header_idx = gen._find_paragraph_with_text("Peer-reviewed Research Articles:")
    assert header_idx is not None
    para = gen.doc.paragraphs[header_idx + 2]

    assert para._p.find(qn("w:ins")) is None
    assert para._p.find(qn("w:del")) is None
    assert "".join(r.text for r in para.runs) == f"1. {formatted_citation}".replace("\x1f", "")
    bold_run = next(r for r in para.runs if r.bold)
    assert bold_run.text == "Smith A"

    assert "\x1f" not in gen.doc.element.xml


# --- T2.3: the plain writer now sanitises too ---


@pytest.mark.parametrize(
    "citation, target_name",
    [
        pytest.param(
            "Doe J\x0b, Smith A. A study. Journal. 2024;10(2):100-110.",
            "Smith A",
            id="before",
        ),
        pytest.param(
            "Doe J, Smith A\x0b. A study. Journal. 2024;10(2):100-110.",
            "Smith A",
            id="adjacent_to_name",
        ),
        pytest.param(
            "Doe J, Smith A. A stu\x0bdy. Journal. 2024;10(2):100-110.",
            "Smith A",
            id="after",
        ),
    ],
)
def test_plain_writer_sanitises_control_characters_and_keeps_bold(citation, target_name):
    gen = _generator()
    para = _blank_paragraph(gen)

    gen._add_citation_with_bold_author(para, citation, target_name, "")  # must not raise

    assert "".join(r.text for r in para.runs) == citation.replace("\x0b", "")
    bold_run = next(r for r in para.runs if r.bold)
    assert bold_run.text == target_name


def test_insertion_writer_with_track_changes_disabled_sanitises_via_plain_path():
    # emit_track_changes=False routes `_add_citation_with_bold_author_as_
    # insertion` straight into `_add_citation_with_bold_author` (bibliography
    # .py:346-348) -- the same route the fallback `except` block takes.
    gen = _generator(emit_track_changes=False)
    para = _blank_paragraph(gen)
    citation = "Doe J\x0b, Smith A. A study. Journal. 2024;10(2):100-110."

    gen._add_citation_with_bold_author_as_insertion(
        para, citation, "Smith A", "", author="PubMed Enrichment"
    )  # must not raise

    assert para._p.find(qn("w:ins")) is None
    assert "".join(r.text for r in para.runs) == citation.replace("\x0b", "")


# --- T2.4: every run segment of the tracked writer, independently ---


@pytest.mark.parametrize(
    "citation, target_name",
    [
        pytest.param(
            "Doe J\x0b, Smith A. A study. Journal. 2024;10(2):100-110.",
            "Smith A",
            id="before",
        ),
        pytest.param(
            "Doe J, Smi\x0bth A. A study. Journal. 2024;10(2):100-110.",
            # The citation's own copy of the name carries the control
            # character; target_name must too, or `in` no longer matches it
            # (`_citation_author_split`) and nothing gets bolded at all.
            "Smi\x0bth A",
            id="name",
        ),
        pytest.param(
            "Doe J, Smith A. A stu\x0bdy. Journal. 2024;10(2):100-110.",
            "Smith A",
            id="after",
        ),
    ],
)
def test_citation_insertion_sanitises_each_run_segment(citation, target_name):
    gen = _generator(emit_track_changes=True)
    para = _blank_paragraph(gen)

    gen._add_citation_with_bold_author_as_insertion(
        para, citation, target_name, "", author="PubMed Enrichment"
    )  # must not raise

    ins_elem = para._p.find(qn("w:ins"))
    assert ins_elem is not None
    runs = ins_elem.findall(qn("w:r"))

    t_texts = []
    bold_runs = 0
    bold_text = None
    for r in runs:
        rPr = r.find(qn("w:rPr"))
        is_bold = rPr is not None and rPr.find(qn("w:b")) is not None
        t = r.find(qn("w:t"))
        assert t is not None
        text = t.text or ""
        assert "\x0b" not in text
        t_texts.append(text)
        if is_bold:
            bold_runs += 1
            bold_text = text

    assert "".join(t_texts) == citation.replace("\x0b", "")
    assert bold_runs == 1
    assert bold_text == target_name.replace("\x0b", "")


if __name__ == "__main__":
    import pytest as _pytest

    raise SystemExit(_pytest.main([__file__, "-v"]))
