"""Section K (teaching) newline-blind fix (#476, PR1 of the accuracy wave),
and the delimiter contract review thread 3927100368 asked to have pinned.

Section K reconstructs an entry's bullet text at two call sites that read two
DIFFERENT variables, so the section's delimiter decision is made per call site
rather than as one migration:

The RAW-text side -- `original_lines = entry_lines(original_text)`, inside
`_teaching_entry_lines`' `if formatted_text and original_text:` branch --
answers "did Stage 5c fuse multiple distinct source entries into one formatted
blob?" ("Stage 5c may have over-combined"). Switching it to `entry_fragments`
was tried and reverted after reading what it actually does to the farm: 456 of
722 K entries are genuinely ONE teaching record whose raw text happens to be
tab-joined ("Title\\tDate\\tDescription"), which Stage 5c already reformats
correctly into one well-structured bullet with its own markdown sub-bullets
("**2020-2024** - **Course Director**...\\n  - Selected initiatives...").
`entry_fragments` makes every one of those look like a multi-item entry, which
takes the OTHER branch and replaces one good, LLM-formatted bullet with several
raw, unformatted fragment bullets (title, date and description as three
separate bullets with no attempt at prose). That is a regression, not a fix,
and `test_regression_guard_raw_text_...` below pins it by showing what the
discarded change would have done. The same reading is why the last fallback
(raw text, no formatted_text, no course_title) also keeps the newline-only
split: tab-splitting it turns 19 farm entries into 41 bullets, 16 of them a
bare date, when `_clean_inline_tabs` already renders the untouched row as
"Title: Date — Description".

The Stage-5c-prose side -- `_item_parts(...)` inside the `elif formatted_text:`
branch, reached only when `original_text` is EMPTY -- has nothing to compare
against, so the "did 5c over-combine" question doesn't apply; it just decides
how many `'. '`-joined pieces of Stage 5c's OWN prose to combine. It splits on
newlines and tabs, and deliberately NOT on `|`: a pipe joins the cells of one
row, and `_clean_inline_tabs` already renders it as " — " inside a single
bullet (review thread 3927100368, item 4). 0 farm K entries reach this branch
with either text present, so the positive control is synthetic (§6.5 hole,
disclosed in the PR body).

This branch turns out to be unreachable through `_insert_teaching_entry` as
written, for a reason that has nothing to do with #476:
`_is_structural_label(entry)` (parsing/text.py, unrelated, pre-existing)
returns True -- and the function returns immediately -- whenever
`entry.get('text', '')` is falsy, which is exactly the condition the
`elif formatted_text:` branch needs to be reached at all. The positive and
negative controls below patch `_is_structural_label` out for the tests that
need to exercise the branch itself, which is the honest way to pin what the
changed line does without either fabricating a reachable-looking fixture or
silently declining to test it the way the other four sections' positive
controls are tested end to end.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_teaching_fragments.py -p no:cacheprovider
"""

import logging
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.core.render_check import entry_fragments, entry_lines  # noqa: E402
from unified_pipeline.stage6.sections import teaching  # noqa: E402
from unified_pipeline.stage6.sections.teaching import (  # noqa: E402
    TEACHING_SECTION_HEADERS,
    _item_parts,
    _teaching_entry_lines,
)
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    RENDER_ROUTED_CODES,
    WCMTemplateGenerator,
)


MULTILINE_CASES = [
    "First point\nSecond point",
    "**2020** - Course Director\n  - Detail one\n  - Detail two",
]


def _generator_with_trailing_paragraph():
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen.doc.add_paragraph("HEADING")
    gen.doc.add_paragraph("SENTINEL-END")
    return gen


# --- (a) no line the old newline split produced is lost ---------------------

def test_item_parts_keeps_every_entry_lines_part():
    for case in MULTILINE_CASES:
        old = entry_lines(case)
        assert len(old) > 1
        assert _item_parts(case) == old, f"diverged on {case!r}"


# --- (b) the delimiter contract: tab and newline split, '|' does not --------

def test_item_parts_splits_tabs_but_not_pipes():
    """Review thread 3927100368 item 4. `|` joins the cells of ONE item
    ("Role | Institution") and `_clean_inline_tabs` renders it as " — " in a
    single bullet, so making it an item separator would split one teaching
    record into two. A future swap of `_item_parts` for `entry_fragments`
    would do exactly that; this is the test that stops it.
    """
    mixed = "Role A | Institution A\tRole B | Institution B"
    assert _item_parts(mixed) == ["Role A | Institution A", "Role B | Institution B"]
    # entry_fragments is the function NOT to use here, and this is why.
    assert len(entry_fragments(mixed)) == 4


# --- (c) positive control: the tab-joined blob gains both parts -------------

def test_positive_control_pipe_blind_formatted_text_gains_both_parts(monkeypatch):
    """`elif formatted_text:` branch: original_text is empty, formatted_text
    is a single-line, tab-joined blob. entry_lines returns ONE part, so
    combined_text was that whole unsplit blob before #476.

    `_is_structural_label` is patched out (see module docstring): it always
    returns True when `entry.get('text', '')` is falsy, which is exactly
    the condition this branch needs -- unrelated, pre-existing gate, not
    something #476 touches.
    """
    monkeypatch.setattr(teaching, "_is_structural_label", lambda entry: False)
    entry = {
        "text": "",
        "extracted_fields": {
            "formatted_text": "Course Director, Internal Medicine Clerkship\tPreceptor, Ambulatory Clinic",
        },
    }
    gen = _generator_with_trailing_paragraph()
    gen._insert_teaching_entry(1, entry, is_first_visible=False)
    texts = [p.text for p in gen.doc.paragraphs]
    matches = [t for t in texts if "Course Director" in t]
    assert len(matches) == 1, texts
    assert matches[0] == "Course Director, Internal Medicine Clerkship. Preceptor, Ambulatory Clinic"


def test_pipe_joined_formatted_text_stays_one_item(monkeypatch):
    """Same branch, the other half of the contract: a pipe is not an item
    separator, so the blob stays one bullet and `_clean_inline_tabs` renders
    the pipe as " — " rather than the two ". "-joined items entry_fragments
    would have produced.
    """
    monkeypatch.setattr(teaching, "_is_structural_label", lambda entry: False)
    entry = {
        "text": "",
        "extracted_fields": {
            "formatted_text": "Course Director, Internal Medicine Clerkship | Preceptor, Ambulatory Clinic",
        },
    }
    gen = _generator_with_trailing_paragraph()
    gen._insert_teaching_entry(1, entry, is_first_visible=False)
    matches = [p.text for p in gen.doc.paragraphs if "Course Director" in p.text]
    assert len(matches) == 1
    assert matches[0] == "Course Director, Internal Medicine Clerkship — Preceptor, Ambulatory Clinic"


# --- (d) negative control: single-part text is unchanged --------------------

def test_negative_control_single_part_unchanged(monkeypatch):
    monkeypatch.setattr(teaching, "_is_structural_label", lambda entry: False)
    entry = {
        "text": "",
        "extracted_fields": {"formatted_text": "Course Director, Internal Medicine Clerkship"},
    }
    gen = _generator_with_trailing_paragraph()
    gen._insert_teaching_entry(1, entry, is_first_visible=False)
    texts = [p.text for p in gen.doc.paragraphs]
    assert "Course Director, Internal Medicine Clerkship" in texts


# --- raw-text regression guard: pins the DECLINED migration -----------------

def test_regression_guard_raw_text_entry_fragments_would_wrongly_fragment_the_farm_shape():
    """Not a behaviour pin -- a documentation-as-code guard for why the raw
    text paths keep `entry_lines`. The real farm shape: tab-joined single
    record, Stage 5c already reformats it well. entry_lines correctly sees ONE
    line (uses the good formatted_text); entry_fragments would see 3+ fragments
    (would wrongly take the "multi-item, use raw lines" branch and discard
    the LLM formatting). If this test ever fails, the farm shape changed and
    the discarded migration needs re-evaluating, not silently redone.
    """
    original_text = (
        "Leadership in education, Duke Psychiatry Residency Program"
        "\tJuly 2020 – June 2024"
        "\tSelected Initiatives: Co-designed the curriculum."
    )
    assert len(entry_lines(original_text)) == 1
    assert len(entry_fragments(original_text)) > 1


# ===========================================================================
# Review thread 3931638847: the twenty-six cases the reviewer listed, plus the
# paragraph-order regression thread 3927100368 item 3 asked for.
#
# Two levels on purpose. `_teaching_entry_lines` is pure, so every text rule --
# the three-way source fallback, the delimiter contract, the field joins -- is
# pinned on the function itself, with no document in the way. `_fill_teaching`
# is then exercised over a real python-docx document, because the things that
# can silently break there are POSITIONS: `section_idx + 1`, repeated insertion
# at one index, and the per-code blank line. A mock document would assert the
# calls we already know we make; a real one asserts the page.
# ===========================================================================


def _generator(*paragraph_texts):
    """A generator whose document is exactly these paragraphs, in this order.

    `_find_paragraph_with_text` is a case-insensitive substring scan over
    `self.doc.paragraphs`, so a "heading" is simply a paragraph written here.
    That is what lets one test give a code its own heading, the next give it
    only the shared EDUCATIONAL CONTRIBUTIONS fallback, and the next give it
    nothing at all -- the three routing outcomes, without three fixtures.
    """
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    for text in paragraph_texts:
        gen.doc.add_paragraph(text)
    return gen


def _template_generator():
    """A generator over the shipped WCM template.

    Used only by the routing test, which has to prove the five REAL heading
    strings are the ones found -- a synthetic document would prove
    `TEACHING_SECTION_HEADERS` matches itself.
    """
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    return gen


def _paragraphs(gen):
    """Every paragraph's text, blank spacer lines included."""
    return [p.text for p in gen.doc.paragraphs]


def _visible(gen):
    """Paragraph texts with the blank spacer lines dropped."""
    return [text for text in _paragraphs(gen) if text.strip()]


def _index_of(texts, needle):
    """Index of the first paragraph containing `needle`, case-insensitively."""
    lowered = needle.lower()
    for i, text in enumerate(texts):
        if lowered in text.lower():
            return i
    raise AssertionError(f"{needle!r} not found in {texts!r}")


def _entry(code, text, hierarchy=None, **fields):
    """One stage-5c-shaped teaching entry. Keyword args become extracted_fields.

    `hierarchy` is a top-level entry key, not an extracted field: it carries the
    source CV's own section labels above this entry, and `_is_structural_label`
    reads it to decide whether an all-caps run is a header or is legitimate
    all-caps content (a name, an initialism, an org written in caps).
    """
    entry = {"taxonomy_code": code, "text": text, "extracted_fields": dict(fields)}
    if hierarchy is not None:
        entry["hierarchy"] = list(hierarchy)
    return entry


# --- 1. end-to-end routing, against the real template ----------------------

def test_fill_teaching_routes_all_k_codes_to_correct_sections():
    """Case 1. Each K code lands under its OWN heading, not a sibling's.

    The five headings sit two paragraphs apart in the shipped template
    (85/87/89/91/93, under EDUCATIONAL CONTRIBUTIONS at 82), so an off-by-one
    in `section_idx + 1` -- or a heading lookup that matches an already
    inserted bullet instead of the heading -- moves a whole code's entries
    under the neighbouring subsection without raising anything.
    """
    codes = ["K1", "K2", "K3", "K4", "K5"]
    gen = _template_generator()
    gen._fill_teaching({code: [_entry(code, f"MARKER-{code}",
                                      formatted_text=f"MARKER-{code}")]
                        for code in codes})

    texts = _paragraphs(gen)
    headings = ["Didactic teaching", "Clinical teaching", "Administrative teaching",
                "Continuing education", "Other education/outreach"]
    following = headings[1:] + [None]
    for code, heading, next_heading in zip(codes, headings, following):
        marker = _index_of(texts, f"MARKER-{code}")
        assert _index_of(texts, heading) < marker, f"{code} landed above its heading"
        if next_heading is not None:
            assert marker < _index_of(texts, next_heading), \
                f"{code} landed under {next_heading!r}"


def test_each_k_code_matches_its_alternate_header():
    """Case 20. Every code carries a LIST of candidate headings because the
    template wording has drifted; only the first is exercised by the shipped
    template, so the alternates are otherwise untested and a typo in one would
    surface as a whole code silently taking the EDUCATIONAL CONTRIBUTIONS
    fallback on a revised template.
    """
    gen = _generator("Didactic", "bedside teaching", "leadership role",
                     "professional education", "community education or patient",
                     "SENTINEL-END")
    gen._fill_teaching({code: [_entry(code, f"MARKER-{code}",
                                      formatted_text=f"MARKER-{code}")]
                        for code in ["K1", "K2", "K3", "K4", "K5"]})

    assert _visible(gen) == [
        "Didactic", "MARKER-K1",
        "bedside teaching", "MARKER-K2",
        "leadership role", "MARKER-K3",
        "professional education", "MARKER-K4",
        "community education or patient", "MARKER-K5",
        "SENTINEL-END",
    ]


# --- 2. the shared fallback heading, and no heading at all -----------------

def test_fill_teaching_falls_back_to_educational_contributions():
    """Case 2. A code whose own headings all miss uses the parent heading
    rather than dropping its entries."""
    gen = _generator("EDUCATIONAL CONTRIBUTIONS", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "Grand rounds", formatted_text="Grand rounds")]})
    assert _visible(gen) == ["EDUCATIONAL CONTRIBUTIONS", "Grand rounds", "SENTINEL-END"]


def test_fill_teaching_skips_code_when_no_section_header_exists(caplog):
    """Case 2. No heading and no fallback heading: the entries cannot be
    placed. They are dropped -- but reported, not swallowed (§5.4), which is
    the only signal a template revision has removed a heading."""
    gen = _generator("SOMETHING ELSE ENTIRELY", "SENTINEL-END")
    with caplog.at_level(logging.WARNING, logger=teaching.__name__):
        gen._fill_teaching({"K1": [_entry("K1", "Grand rounds",
                                          formatted_text="Grand rounds")]})
    assert _visible(gen) == ["SOMETHING ELSE ENTIRELY", "SENTINEL-END"]
    assert "no template heading for teaching code K1" in caplog.text


def test_two_codes_sharing_the_fallback_heading_stack_in_documented_order():
    """Review thread 3927100368 item 1: the generic EDUCATIONAL CONTRIBUTIONS
    fallback with MULTIPLE codes taking it at once.

    K1-K5 resolve their heading independently, so on a template revision that
    drops a subsection heading two codes land in one place. The module
    docstring states what that produces; this is that statement as an
    assertion, so a template revision cannot silently merge and interleave two
    teaching categories:

    - each code's block is CONTIGUOUS (never interleaved),
    - blocks stack in REVERSE code order, because every block is inserted
      directly under the shared heading and pushes the previous one down,
    - each block keeps its own blank spacer, `is_first_visible` being computed
      per code.
    """
    gen = _generator("EDUCATIONAL CONTRIBUTIONS", "SENTINEL-END")
    gen._fill_teaching({
        "K1": [_entry("K1", "K1 first", formatted_text="K1 first"),
               _entry("K1", "K1 second", formatted_text="K1 second")],
        "K2": [_entry("K2", "K2 only", formatted_text="K2 only")],
    })
    assert _paragraphs(gen) == [
        "EDUCATIONAL CONTRIBUTIONS",
        "",          # K2's spacer
        "K2 only",
        "",          # K1's spacer
        "K1 first",
        "K1 second",
        "SENTINEL-END",
    ]


# --- 3. the exact paragraph order the section contracts to ------------------

def test_paragraph_order_for_one_two_and_three_entries():
    """Review thread 3927100368 item 3, and case 19. `_insert_bulleted_entry`
    inserts BEFORE the index it is given, and every entry in a code is given
    the SAME index, so the list is walked in reverse to come out forwards.
    Nothing about that raises when it breaks: a helper change that inserted
    after the index instead would silently reverse the section, and a change
    to the shared index would scatter entries outside it.
    """
    dated = [("Newest", "2022"), ("Middle", "2020"), ("Oldest", "2018")]
    for count in (1, 2, 3):
        gen = _generator("Didactic teaching", "SENTINEL-END")
        gen._fill_teaching({"K1": [
            _entry("K1", label, formatted_text=label, end_date=year)
            for label, year in dated[:count]
        ]})
        expected = ["Didactic teaching", ""] + [label for label, _ in dated[:count]]
        assert _paragraphs(gen) == expected + ["SENTINEL-END"], f"{count} entries"


def test_paragraph_order_for_a_multiline_entry():
    """Review thread 3927100368 item 3, the multiline half. One entry becomes
    several adjacent bullets in SOURCE order, and the blank spacer goes above
    the bullet that ends up visually first -- not above the one inserted
    first, which is the last line."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry(
        "K1", "Line one\nLine two\nLine three",
        formatted_text="Stage 5c fused all three into one sentence.")]})
    assert _paragraphs(gen) == ["Didactic teaching", "",
                                "Line one", "Line two", "Line three",
                                "SENTINEL-END"]


def test_entries_render_reverse_chronologically():
    """Case 17. Input order is deliberately neither sorted nor reversed, so a
    sort that silently became a no-op would still fail here."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [
        _entry("K1", "Oldest", formatted_text="Oldest", end_date="2018"),
        _entry("K1", "Newest", formatted_text="Newest", end_date="2022"),
        _entry("K1", "Middle", formatted_text="Middle", end_date="2020"),
    ]})
    assert _visible(gen) == ["Didactic teaching", "Newest", "Middle", "Oldest",
                             "SENTINEL-END"]


def test_only_the_first_visible_entry_gets_a_blank_line_above_it():
    """Case 18. One spacer per code, above the entry that ends up visually
    first -- `is_first_visible` is the LAST index of the reversed list, so an
    off-by-one there puts the blank line in the middle of the section or
    gives every entry one."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [
        _entry("K1", "Newer", formatted_text="Newer", end_date="2022"),
        _entry("K1", "Older", formatted_text="Older", end_date="2018"),
    ]})
    assert _paragraphs(gen) == ["Didactic teaching", "", "Newer", "Older",
                                "SENTINEL-END"]


def test_one_entry_under_two_k_codes_renders_under_both():
    """Case 23. The same entry object routed under two codes is rendered once
    per code -- `_fill_teaching` reads `entries_by_code` per code and holds no
    cross-code seen-set, so nothing dedupes the second copy. Pinned because
    the opposite (a shared seen-set) would look like a reasonable cleanup and
    would drop content from whichever code is processed second.
    """
    shared = _entry("K1", "Shared activity", formatted_text="Shared activity")
    gen = _generator("Didactic teaching", "Clinical teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [shared], "K2": [shared]})
    assert _visible(gen) == ["Didactic teaching", "Shared activity",
                             "Clinical teaching", "Shared activity",
                             "SENTINEL-END"]


def test_zero_entries_renders_nothing():
    """Case 21. Both shapes of "no entries": no codes at all, and a code
    present with an empty list."""
    gen = _generator("Didactic teaching", "EDUCATIONAL CONTRIBUTIONS", "SENTINEL-END")
    before = _paragraphs(gen)
    gen._fill_teaching({})
    gen._fill_teaching({"K1": [], "K5": []})
    assert _paragraphs(gen) == before


def test_unknown_k_code_renders_nothing_and_is_left_for_the_appendix():
    """Case 22. A K code outside K1-K5 is not counted and not rendered.

    That is not silent loss: `RENDER_ROUTED_CODES` is what marks a code as
    "already placed", and its K membership is exactly K1-K5, so anything else
    under K stays unrouted and the appendix picks it up. The two set
    assertions are the actual guard -- adding a sixth K code to
    `RENDER_ROUTED_CODES` without a `TEACHING_SECTION_HEADERS` row would make
    it vanish from both the section and the appendix.
    """
    gen = _generator("Didactic teaching", "EDUCATIONAL CONTRIBUTIONS", "SENTINEL-END")
    before = _paragraphs(gen)
    gen._fill_teaching({"K9": [_entry("K9", "Unmapped", formatted_text="Unmapped")]})
    assert _paragraphs(gen) == before

    assert set(TEACHING_SECTION_HEADERS) <= RENDER_ROUTED_CODES
    assert {code for code in RENDER_ROUTED_CODES
            if code.startswith("K")} == set(TEACHING_SECTION_HEADERS)


# --- the two content gates on `_insert_teaching_entry` ----------------------

def test_all_caps_structural_label_entry_is_not_rendered():
    """Case 6. A source CV's own section header extracted as an entry is
    furniture: the WCM template supplies the structure.

    The entry carries the hierarchy label it echoes, because that is what a
    real one looks like and because case alone is not enough: #665 item 3
    (merged as 7e82fb8) narrowed `_is_structural_label` to require either a
    hierarchy echo or a stage-4 extraction that found nothing, after the
    unconditional all-caps rule was found to drop legitimate all-caps content.
    A fixture with no hierarchy and a populated `formatted_text` satisfies
    neither signal, so it would render -- correctly, under the current rule.
    """
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "DIRECT TEACHING AND PRECEPTING",
                                      hierarchy=["Direct Teaching and Precepting"],
                                      formatted_text="Direct teaching and precepting")]})
    assert _visible(gen) == ["Didactic teaching", "SENTINEL-END"]


def test_orphan_fragment_is_not_rendered_but_a_dated_sibling_is():
    """Case 7. A short entry with no date, audience, location, formatted text
    or title is a dangling continuation line. The second half is the control
    that keeps the gate honest: the SAME text with a date renders, so this
    pins the gate rather than merely the length."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "Curriculum development")]})
    assert _visible(gen) == ["Didactic teaching", "SENTINEL-END"]

    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "Curriculum development", date="2020")]})
    assert _visible(gen) == ["Didactic teaching", "Curriculum development",
                             "SENTINEL-END"]


def test_completely_unrenderable_entry_is_reported_and_dropped(caplog):
    """Case 16. Every raw line was a bare column label and no field survived,
    so there is nothing to bullet. It is dropped -- with a warning naming the
    code and the text, because a renderer that silently discards an entry is
    how content loss goes unnoticed for a release (§5.4)."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    with caplog.at_level(logging.WARNING, logger=teaching.__name__):
        gen._fill_teaching({"K1": [_entry("K1", "Title\nInstitution", date="2020")]})
    assert _visible(gen) == ["Didactic teaching", "SENTINEL-END"]
    assert "teaching entry produced no renderable line (K1)" in caplog.text


# --- every `_teaching_entry_lines` branch, on the pure function -------------

def test_every_text_reconstruction_branch():
    """Case 3, as a single table. Each row is one of the six outcomes
    `_teaching_entry_lines` can reach; the rows below this one then pin each
    outcome's own rules in detail. Kept together so a new branch added without
    a decision about the others is visible in one place.
    """
    assert _teaching_entry_lines(
        {"formatted_text": "Fused."}, "Line one\nLine two") == ["Line one", "Line two"]
    assert _teaching_entry_lines(
        {"formatted_text": "**Bold** text"}, "Line one") == ["Bold text"]
    assert _teaching_entry_lines(
        {"formatted_text": "First\tSecond"}, "") == ["First. Second"]
    assert _teaching_entry_lines(
        {"course_title": "Intro to Medicine"}, "") == ["Intro to Medicine"]
    assert _teaching_entry_lines(
        {}, "Attending rounds") == ["Attending rounds"]
    assert _teaching_entry_lines(
        {"institution": "Weill Cornell"}, "Title\nDates") == ["Weill Cornell"]
    assert _teaching_entry_lines({}, "Title\nDates") == []


def test_multiline_original_text_uses_original_lines_not_formatted_text():
    """Cases 4 and 26, and review thread 3927100368 item 5.

    `if len(original_lines) > 1:` is a data-integrity rule, not an
    optimisation: Stage 5c routinely fuses a multi-item teaching block into
    one sentence, and this branch says the source's own line breaks are the
    more faithful record. The assertion that matters is the NEGATIVE one --
    the over-combined sentence must not reach the page at all. 37 of the 722
    K entries in the local corpus take this branch, rendering 278 bullets
    rather than 37.
    """
    fused = "Stage 5c fused all three into one sentence."
    lines = _teaching_entry_lines({"formatted_text": fused},
                                  "Line one\nLine two\nLine three")
    assert lines == ["Line one", "Line two", "Line three"]
    assert fused not in lines

    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "Line one\nLine two\nLine three",
                                      formatted_text=fused)]})
    rendered = _visible(gen)
    assert rendered == ["Didactic teaching", "Line one", "Line two", "Line three",
                        "SENTINEL-END"]
    assert not any(fused in text for text in rendered)


def test_single_line_original_text_uses_the_formatted_text():
    """Case 5. The other side of the same `if`: one raw line means Stage 5c
    had nothing to over-combine, so its formatting is kept -- with markdown
    stripped, since a Word run has no markdown, and ISO dates normalized."""
    assert _teaching_entry_lines(
        {"formatted_text": "**2020-2024** - Course Director"},
        "2020-2024 Course Director") == ["2020-2024 - Course Director"]

    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "2020-2024 Course Director",
                                      formatted_text="**2020-2024** - Course Director")]})
    assert _visible(gen) == ["Didactic teaching", "2020-2024 - Course Director",
                             "SENTINEL-END"]


def test_mixed_newline_tab_pipe_teaching_input_renders_one_bullet_per_line():
    """Case 26 at the render, the delimiter contract end to end. Newlines
    separate items; a tab and a pipe INSIDE a line stay inside their bullet,
    where `_clean_inline_tabs` punctuates them ("Label: Value", "A — B").
    Splitting on either would turn this two-item entry into four bullets."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry(
        "K1", "Lecturer\tCardiology 101\nPreceptor | Ambulatory Clinic",
        formatted_text="Lecturer for Cardiology 101 and preceptor.")]})
    assert _visible(gen) == ["Didactic teaching",
                             "Lecturer: Cardiology 101",
                             "Preceptor — Ambulatory Clinic",
                             "SENTINEL-END"]


# --- the structured-field fallback -----------------------------------------

def test_course_code_and_course_title_fallback():
    """Case 8. No Stage 5c text at all, so the bullet is built from fields."""
    assert _teaching_entry_lines({"course_code": "MED 101",
                                  "course_title": "Intro to Medicine",
                                  "institution": "Weill Cornell",
                                  "role": "Lecturer"}, "") == \
        ["MED 101: Intro to Medicine, Weill Cornell (Lecturer)"]


def test_course_title_without_a_course_code():
    """Case 10. No code means no "CODE: " prefix -- not an empty one."""
    assert _teaching_entry_lines({"course_title": "Intro to Medicine",
                                  "institution": "Weill Cornell"}, "") == \
        ["Intro to Medicine, Weill Cornell"]


def test_list_valued_course_code_and_title_are_semicolon_joined():
    """Case 9, review thread 3927100368 item 6. Extraction emits both fields
    as a str and as a list of str for the same field, which is why
    `_TeachingFields` types them `str | list[str]`. Without the join the
    f-string would render a Python list literal into the CV. All four
    combinations, because the two normalizations are separate statements and
    losing either one is a distinct defect.
    """
    both = _teaching_entry_lines({"course_code": ["MED 101", "MED 102"],
                                  "course_title": ["Intro to Medicine",
                                                   "Advanced Medicine"]}, "")
    assert both == ["MED 101; MED 102: Intro to Medicine; Advanced Medicine"]

    code_only = _teaching_entry_lines({"course_code": ["MED 101", "MED 102"],
                                       "course_title": "Intro to Medicine"}, "")
    assert code_only == ["MED 101; MED 102: Intro to Medicine"]

    title_only = _teaching_entry_lines({"course_code": "MED 101",
                                        "course_title": ["Intro to Medicine",
                                                         "Advanced Medicine"]}, "")
    assert title_only == ["MED 101: Intro to Medicine; Advanced Medicine"]

    assert _teaching_entry_lines({"course_title": ["Solo course"]}, "") == ["Solo course"]


def test_institution_and_role_already_in_the_title_are_not_repeated():
    """Case 11. Both appends are guarded by a substring test, so a title that
    already names the institution or the role does not get it twice."""
    assert _teaching_entry_lines({"course_title": "Intro to Medicine, Weill Cornell (Lecturer)",
                                  "institution": "Weill Cornell",
                                  "role": "Lecturer"}, "") == \
        ["Intro to Medicine, Weill Cornell (Lecturer)"]


# --- the raw-text fallback --------------------------------------------------

def test_raw_original_text_fallback():
    """Case 12. No formatted text and no course title: the raw lines are the
    last real source of content, split on newlines only (the module docstring
    gives the farm counts behind that choice)."""
    assert _teaching_entry_lines({"institution": "Weill Cornell"},
                                 "Attending rounds\nJournal club") == \
        ["Attending rounds", "Journal club"]

    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "Attending rounds\nJournal club",
                                      institution="Weill Cornell")]})
    assert _visible(gen) == ["Didactic teaching", "Attending rounds",
                             "Journal club", "SENTINEL-END"]


def test_structural_column_labels_are_filtered_from_raw_lines():
    """Case 13. A raw line that is nothing but a source table's column label
    is furniture. Case-insensitively, and only when the label is the WHOLE
    line -- "Title of course" is content."""
    assert _teaching_entry_lines(
        {"institution": "Weill Cornell"},
        "Title\nAttending rounds\nINSTITUTION\n dates \nRole") == ["Attending rounds"]
    assert _teaching_entry_lines(
        {}, "Title of course: Attending rounds") == ["Title of course: Attending rounds"]


def test_structural_column_labels_constant_is_the_sections_vocabulary():
    """Review thread 3927100368 item 9. The labels are a module-level frozenset
    with one definition rather than a list literal rebuilt per line; pinned so
    the vocabulary cannot be edited at one of two call sites."""
    assert teaching._STRUCTURAL_COLUMN_LABELS == frozenset(
        {"title", "institution", "dates", "role"})


def test_long_semicolon_line_splits_into_one_bullet_per_item():
    """Case 14. A long raw line carrying ';' is a run-together list of
    activities, so it becomes one bullet each."""
    long_line = ("Lecture on cardiology for second-year medical students; "
                 "Seminar on renal physiology for third-year medical students")
    assert len(long_line) > teaching._SEMICOLON_SPLIT_MIN_CHARS
    assert _teaching_entry_lines({"institution": "Weill Cornell"}, long_line) == [
        "Lecture on cardiology for second-year medical students",
        "Seminar on renal physiology for third-year medical students",
    ]


def test_short_semicolon_line_stays_one_bullet():
    """Case 14, the control. Below the threshold a ';' is punctuation inside a
    single activity, and splitting would shatter it."""
    short_line = "Lecture on cardiology; seminar on renal physiology"
    assert len(short_line) <= teaching._SEMICOLON_SPLIT_MIN_CHARS
    assert _teaching_entry_lines({"institution": "Weill Cornell"}, short_line) == \
        [short_line]


def test_institution_and_role_fallback_when_no_lines_remain():
    """Case 15. Every raw line was a bare column label, so the filter above
    left nothing -- but institution and role are still sitting in the fields,
    and an entry with them is worth more than a dropped entry (#574)."""
    assert _teaching_entry_lines({"institution": "Weill Cornell", "role": "Lecturer"},
                                 "Title\nInstitution\nDates\nRole") == \
        ["Weill Cornell, Lecturer"]
    assert _teaching_entry_lines({"role": "Lecturer"},
                                 "Title\nInstitution") == ["Lecturer"]
    assert _teaching_entry_lines({}, "Title\nInstitution") == []


# --- whitespace on both text sources ---------------------------------------

def test_formatted_text_that_strips_to_nothing_keeps_the_raw_line():
    """Case 24, and the defect writing it exposed.

    A `formatted_text` that reduces to nothing once ISO dates are normalized
    and markdown is stripped ("   ", "**  **") used to be handed straight to
    `_insert_bulleted_entry`. That did two wrong things at once: it put an
    empty Word list paragraph on the page, and it discarded the raw line the
    entry still carried, because this branch is only reached when there IS
    raw text. The raw line is now used instead, and when there is no usable
    raw line either the entry reconstructs to nothing and the caller's
    "produced no renderable line" warning fires.
    """
    assert _teaching_entry_lines({"formatted_text": "   "}, "Attending rounds") == \
        ["Attending rounds"]
    assert _teaching_entry_lines({"formatted_text": "**  **"}, "Attending rounds") == \
        ["Attending rounds"]
    assert _teaching_entry_lines({"formatted_text": "   "}, "  \n  ") == []
    assert _teaching_entry_lines({"formatted_text": "   "}, "") == []


def test_blank_formatted_text_renders_the_raw_line_not_an_empty_bullet():
    """Case 24 at the render: the same defect seen as paragraphs. The old
    behaviour produced ["Didactic teaching", "", "", "SENTINEL-END"] -- a
    spacer, then a bullet with no words in it, and the entry's real text
    nowhere on the page."""
    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "Attending rounds, Weill Cornell",
                                      formatted_text="   ",
                                      institution="Weill Cornell")]})
    assert _paragraphs(gen) == ["Didactic teaching", "",
                                "Attending rounds, Weill Cornell", "SENTINEL-END"]


def test_whitespace_only_original_text_uses_the_formatted_text():
    """Case 25. `entry_lines` drops blank lines, so whitespace-only raw text
    is zero lines, not one -- it must not be mistaken for a multi-item entry
    and must not suppress the formatted text.

    At the render it never gets that far: `_is_structural_label` (parsing, not
    this section, and unrelated to #476) returns True for any entry whose text
    strips to empty, so the entry is dropped before reconstruction. Both
    halves are asserted because they disagree, and a reader of the helper test
    alone would predict the wrong page. 0 of the 722 K entries in the local
    corpus have whitespace-only text.
    """
    assert _teaching_entry_lines({"formatted_text": "Grand rounds"}, "  \n  ") == \
        ["Grand rounds"]

    gen = _generator("Didactic teaching", "SENTINEL-END")
    gen._fill_teaching({"K1": [_entry("K1", "  \n  ", formatted_text="Grand rounds")]})
    assert _visible(gen) == ["Didactic teaching", "SENTINEL-END"]


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-q"]))
