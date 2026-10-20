"""The R block must name the meeting that hosted an invited talk.

`event_name` was read nowhere in stage_6_word_template.py, so stage 4 extracted
it and stage 6 dropped it: 2,194 of 2,956 values across 79 of the 100 distinct
corpus CVs reached no part of the rendered document. All of them are R entries,
which never pass through stage 5d, so the loss is entirely in the renderer.

The WCM faculty template fixes the block at three columns -- Title |
Institution/Location | Dates (yyyy) -- so the meeting gets no column of its own
and rides with the venue: "ASMBS 2022 Presidential Grand Rounds, Dallas, TX".

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_presentation_event_name.py -p no:cacheprovider
"""

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


def _venue(**fields):
    """Render one R entry through the real filler and return its middle cell."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None          # skip the LLM geographic classifier
    fields.setdefault("title", "Hepatic Vagotomy in Obese Patients")
    fields.setdefault("year", "2022")
    gen._fill_presentations(
        [{"taxonomy_code": "R", "text": fields["title"], "extracted_fields": fields}])
    for table in gen.doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            if cells and cells[0] == fields["title"]:
                return cells[1]
    raise AssertionError("the entry did not render into any R table")


def test_event_name_precedes_the_location():
    # web069, real values: extracted at stage 4 and absent from the render.
    assert _venue(institution="Dallas, TX",
                  event_name="ASMBS 2022 Presidential Grand Rounds") == \
        "ASMBS 2022 Presidential Grand Rounds, Dallas, TX"


def test_event_name_stands_alone_when_no_venue_was_extracted():
    assert _venue(institution="",
                  event_name="Alfredo Lopez-S Lectureship in Nutrition") == \
        "Alfredo Lopez-S Lectureship in Nutrition"


def test_event_name_already_in_the_venue_is_not_repeated():
    assert _venue(institution="SLS 2019, New Orleans, LA",
                  event_name="SLS 2019") == "SLS 2019, New Orleans, LA"


def test_event_name_already_in_the_title_is_not_repeated():
    assert _venue(title="SAGES 2022 keynote: Reimagining the Future of Surgery",
                  institution="Denver, CO", event_name="SAGES 2022") == "Denver, CO"


def test_an_s8_conference_name_is_read_as_the_event():
    # #1570: an S8 abstract a hierarchy mismatch moved to R keeps its meeting
    # (YUYVIG SQMWHM 1047 rendered "Orlando, FL" alone). Invented values.
    assert _venue(location="Larchmere, ZZ",
                  conference_name="Annual Meeting of the Quillfern Society") == \
        "Annual Meeting of the Quillfern Society, Larchmere, ZZ"


def test_an_event_name_outranks_a_conference_name():
    assert _venue(institution="Larchmere, ZZ", event_name="Quillfern Lecture",
                  conference_name="Quillfern Annual Meeting") == \
        "Quillfern Lecture, Larchmere, ZZ"


def test_venue_is_untouched_when_no_event_name_was_extracted():
    assert _venue(institution="Baton Rouge, LA") == "Baton Rouge, LA"


def test_untitled_entry_renders_its_whole_text_not_the_first_150_characters():
    """#983: an R entry with no extracted title renders as its own text;
    `text[:150]` cut it mid-word."""
    text = ("Grand rounds on the Zorblax method for ferrous metallurgy and its applications to "
            "ceremonial bunting standards, with a review of historical pennant practice worldwide")
    assert len(text) > 150
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None
    gen._fill_presentations([{"taxonomy_code": "R", "text": text,
                              "extracted_fields": {"year": "2022"}}])
    firsts = [row.cells[0].text for table in gen.doc.tables for row in table.rows]
    assert text in firsts


def test_role_leads_the_venue():
    """#475: the speaker's role was never in the R schema, then never rendered."""
    assert _venue(institution="Northgate Institute, Springfield",
                  role="Visiting Professor") == \
        "Visiting Professor, Northgate Institute, Springfield"


def test_role_leads_event_name_and_venue():
    assert _venue(institution="Springfield", event_name="Zorblax Lecture",
                  role="Visiting Professor") == \
        "Visiting Professor, Zorblax Lecture, Springfield"


def test_role_stands_alone_when_no_venue_was_extracted():
    assert _venue(institution="", role="Visiting Professor") == "Visiting Professor"


def test_role_already_in_the_venue_is_not_repeated():
    assert _venue(institution="Visiting Professor, Northgate Institute",
                  role="Visiting Professor") == "Visiting Professor, Northgate Institute"


def test_role_already_in_the_title_is_not_prepended():
    assert _venue(title="Visiting Professor Lecture: Zorblax Methods",
                  institution="Northgate Institute", role="Visiting Professor") == \
        "Northgate Institute"


def test_role_duplicate_check_ignores_case():
    assert _venue(institution="Visiting Professor, Northgate Institute",
                  role="visiting professor") == "Visiting Professor, Northgate Institute"


def test_string_none_role_is_not_rendered():
    assert _venue(institution="Springfield", role="None") == "Springfield"


def _cells(**fields):
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None
    fields.setdefault("year", "2022")
    gen._fill_presentations(
        [{"taxonomy_code": "R", "text": "2022\t* raw source line, Springfield",
          "extracted_fields": fields}])
    for table in gen.doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            if "2022" in cells:
                return cells
    raise AssertionError("row not rendered")


def test_empty_title_keeps_the_rest_of_the_raw_text_and_shows_the_role():
    """The LLM emptied title but filled role: what the source line says beyond
    role, venue and year may be the talk title (#475), so it still renders --
    without the leading year the Dates cell shows (EBYSBC E23)."""
    cells = _cells(title="", role="Visiting Professor",
                   location="Northgate Institute, Springfield")
    assert cells[0] == "* raw source line, Springfield"
    assert cells[1] == "Visiting Professor, Northgate Institute, Springfield"


def _untitled_row(text, **fields):
    """Render one title-less R entry and return its cells."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None
    gen._fill_presentations(
        [{"taxonomy_code": "R", "text": text, "extracted_fields": fields}])
    for table in gen.doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            if cells[-1] and cells[-1] == fields["date"][:4]:
                return cells
    raise AssertionError("row not rendered")


def test_untitled_talk_with_event_name_is_titled_by_the_event_not_the_source_line():
    """EBYSBC MRJDWE-08: a numbered 'N. YYYY. <series>. <venue>' line filled the
    Title cell, repeating the venue and year of the next two cells."""
    cells = _untitled_row("4. 2018. Zorblax Seminar series. Northgate Diabetes Center",
                          date="2018", event_name="Zorblax Seminar series",
                          location="Northgate Diabetes Center")
    assert cells == ["Zorblax Seminar series", "Northgate Diabetes Center", "2018"]


def test_untitled_talk_is_titled_by_role_and_event_name():
    cells = _untitled_row("2021 Fern Graduate Program, Annual Retreat Presenter, CO",
                          date="2021", role="Presenter",
                          event_name="Fern Graduate Program, Annual Retreat", location="CO")
    assert cells == ["Presenter, Fern Graduate Program, Annual Retreat", "CO", "2021"]


def test_untitled_talk_with_role_only_drops_venue_and_trailing_date():
    cells = _untitled_row("Visiting Professor, Northgate School\t  March  1976",
                          date="1976-03", role="Visiting Professor",
                          location="Northgate School")
    assert cells == ["Visiting Professor", "Northgate School", "1976"]


def test_untitled_talk_drops_the_leading_year_and_the_venue_from_the_source_line():
    """EBYSBC GJXIWD-04: '<year> Invited Speaker - <program>, <venue>' with no
    role field -- the year and venue already fill their own cells."""
    cells = _untitled_row("2019 Invited Speaker - Zorblax Program, Northgate University, CO",
                          date="2019", event_name="Zorblax Program",
                          location="Northgate University, CO")
    assert cells == ["Invited Speaker - Zorblax Program", "Northgate University, CO", "2019"]


def test_untitled_talk_with_only_a_venue_leaves_the_title_blank():
    """EBYSBC RNKYST-06: the line is year and venue only; both have cells."""
    cells = _untitled_row("2003: Northgate University", date="2003",
                          location="Northgate University")
    assert cells == ["", "Northgate University", "2003"]


def test_leading_year_is_kept_when_the_dates_cell_does_not_show_it():
    """No date field: the year in the line is the only place the year renders."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None
    gen._fill_presentations([{"taxonomy_code": "R",
                              "text": "2010 Northgate School: Widget Studies",
                              "extracted_fields": {}}])
    firsts = [row.cells[0].text for table in gen.doc.tables for row in table.rows]
    assert "2010 Northgate School: Widget Studies" in firsts


def test_venue_is_removed_only_as_a_whole_word():
    """A two-letter state venue must not cut into a word that starts with it."""
    cells = _untitled_row("2020 Colloquium on Copper Widgets, CO", date="2020", location="CO")
    assert cells == ["Colloquium on Copper Widgets", "CO", "2020"]


def test_role_is_stripped_before_it_is_rendered():
    assert _venue(institution="Northgate Institute", role=" Visiting Professor ") == \
        "Visiting Professor, Northgate Institute"


def _rendered_date(**fields):
    """Render one R entry through the real filler and return its Dates cell."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None          # skip the LLM geographic classifier
    fields.setdefault("title", "Example Talk on Widget Calibration")
    gen._fill_presentations(
        [{"taxonomy_code": "R", "text": fields["title"], "extracted_fields": fields}])
    for table in gen.doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            if cells and cells[0] == fields["title"]:
                return cells[-1]
    raise AssertionError("the entry did not render into any R table")


def test_a_multi_day_dict_date_renders_as_its_year_not_its_repr():
    """#1233: stage 4 returns `date` as {start_date, end_date} for a multi-day
    event; the whole repr landed in the Dates cell."""
    assert _rendered_date(date={"start_date": "2021-05-17", "end_date": "2021-05-19"}) == "2021"


def test_a_dict_date_spanning_a_new_year_renders_both_years():
    assert _rendered_date(date={"start_date": "2021-12-30", "end_date": "2022-01-02"}) == "2021-2022"


def test_a_string_date_is_unchanged_by_the_dict_handling():
    assert _rendered_date(date="2021-05-17") == "2021"


def test_a_dict_dated_row_sorts_into_date_order_not_to_the_bottom():
    """#1233: the dict date keyed (0, 0, 0), so the row it fixed still landed
    below every dated row. Newest first, wherever the dict-dated row falls."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None          # skip the LLM geographic classifier
    dates = ["2020-03-01", {"start_date": "2021-05-17", "end_date": "2021-05-19"}, "2019-06-01"]
    gen._fill_presentations([
        {"taxonomy_code": "R", "text": f"Example Talk {i}",
         "extracted_fields": {"title": f"Example Talk {i}", "date": date}}
        for i, date in enumerate(dates)])
    rows = [[c.text for c in row.cells] for table in gen.doc.tables for row in table.rows]
    assert [r[-1] for r in rows if r[0].startswith("Example Talk")] == ["2021", "2020", "2019"]


def test_untitled_talk_drops_a_closing_date_with_its_weekday_and_no_stranded_stop():
    cells = _untitled_row("Zorblax Network Conference, Springfield, ZZ. Saturday July 24, 2010.",
                          date="2010", location="Springfield, ZZ")
    assert cells == ["Zorblax Network Conference", "Springfield, ZZ", "2010"]


def test_untitled_talk_leaves_no_stranded_full_stop_after_the_venue():
    cells = _untitled_row("2020 Invited Speaker, Fern Center, Online.", date="2020",
                          location="Fern Center, Online")
    assert cells == ["Invited Speaker", "Fern Center, Online", "2020"]


def test_closing_date_is_kept_when_the_dates_cell_shows_another_year():
    cells = _untitled_row("Zorblax Lecture, Fern Hall, March 1999", date="2001",
                          location="Fern Hall")
    assert cells[0] == "Zorblax Lecture, March 1999"


def test_venue_inside_a_phrase_is_kept():
    """Cutting the venue out of 'hosted by <venue>' would leave the phrase
    dangling; only a venue set off as its own segment is removed."""
    cells = _untitled_row("2020 Fern Symposium hosted by Northgate University (online)",
                          date="2020", location="Northgate University")
    assert cells[0] == "Fern Symposium hosted by Northgate University (online)"


def test_list_number_is_dropped_from_a_kept_source_line():
    cells = _untitled_row("(7) Zorblax workshop talk, Northgate Hall", date="1990",
                          location="Northgate Hall")
    assert cells[0] == "Zorblax workshop talk"


def test_joining_words_alone_do_not_keep_the_source_line():
    """'at' is all that is left once role and venue are out: the role alone
    is the title, not a line that repeats the venue."""
    cells = _untitled_row("2015 Invited Speaker at Northgate University", date="2015",
                          role="Invited Speaker", location="Northgate University")
    assert cells == ["Invited Speaker", "Northgate University", "2015"]


def test_an_untitled_last_split_record_is_titled_from_its_own_line():
    """#1445 (EOAHMI BRUSUZ 265): the last of stage 4's records carries the
    parent's whole line, both talks, and an untitled one was titled with it,
    so the first talk printed twice."""
    from unified_pipeline.stage4.schemas import FIELD_SCHEMAS, STAGE4_RECORDS_KEY
    from unified_pipeline.stage6.fan_out import fan_out_multi_record_entries
    records = [{"title": None, "location": "Ashby, Varnoria", "event_name": "Zqhubble Institute"},
               {"title": None, "location": "Kestrel, Japan", "event_name": "2nd Zqriken Symposium"}]
    entry = {"taxonomy_code": "R", "element_idx_start": 5,
             "text": "Zqhubble Institute (Ashby, Varnoria) 2nd Zqriken Symposium (Kestrel, Japan)",
             "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records}}
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None
    gen._fill_presentations(fan_out_multi_record_entries(
        [entry], FIELD_SCHEMAS, records_key=STAGE4_RECORDS_KEY))
    rows = sorted([c.text for c in row.cells] for table in gen.doc.tables for row in table.rows
                  if any("Zq" in c.text for c in row.cells))
    assert rows == [["2nd Zqriken Symposium", "Kestrel, Japan", ""],
                    ["Zqhubble Institute", "Ashby, Varnoria", ""]]


def test_an_open_ended_talk_renders_its_range_not_its_start_year():
    """#1346 (OTBUCZ 116/119/125): a course given "2006 to present" has
    start_date 2006 and end_date present; the Dates cell showed only 2006."""
    assert _rendered_date(start_date="2006", end_date="present") == "2006-Present"


def test_a_closed_multi_year_range_renders_both_years():
    assert _rendered_date(start_date="1981-01-09", end_date="1983-05-20") == "1981-1983"


def test_a_multi_day_talk_within_one_year_still_renders_one_year():
    assert _rendered_date(start_date="2014-09-27", end_date="2014-10-02") == "2014"


def test_an_end_date_extends_a_year_given_under_date():
    assert _rendered_date(date="2015", end_date="2017") == "2015-2017"


def test_a_string_none_end_date_is_ignored():
    assert _rendered_date(start_date="2006", end_date="None") == "2006"


def test_a_dict_date_does_not_take_a_second_end_date():
    assert _rendered_date(date={"start_date": "2021-12-30", "end_date": "2022-01-02"},
                          end_date="2022-01-02") == "2021-2022"


def test_a_talk_dated_only_under_additional_dates_renders_them():
    """#1245 (X6 VPMMFM 531): `date` null and three dates under the
    off-schema `additional_dates`; the Dates cell was empty."""
    assert _rendered_date(date=None, additional_dates="1999-09-30; 1999-10-14; 2000-01-05") \
        == "1999, 2000"


def test_a_further_date_follows_the_talks_own_date_unless_it_repeats_it():
    assert _rendered_date(date="2018", additional_dates="2018-05; 2020") == "2018, 2020"


def test_a_talk_without_further_dates_keeps_its_own_cell():
    assert _rendered_date(date="2018") == "2018"


def test_a_list_of_dates_under_date_renders_each_year_not_the_iso_string():
    """#1245 (OIEPQD GOWQKA 531/533): stage 4 wrote "1999-09-30; 1999-10-14;
    1999-10-26" into `date`, and the yyyy column showed that string."""
    assert _rendered_date(date="1999-09-30; 1999-10-14; 1999-10-26") == "1999"
    assert _rendered_date(date="1997-10-15; 1998-10-22") == "1997, 1998"



# #1436: an event-status note on the source line has no R field, so stage 4
# drops it and a postponed talk rendered as delivered (NDMRSO TVZDVF 451/494).
_STATUS_TITLE = "Calibrating widgets at scale"


def _status_title(text, **fields):
    """The Title cell of one titled R entry whose source line is `text`."""
    fields.setdefault("title", _STATUS_TITLE)
    fields.setdefault("date", "2020-04-07")
    return _untitled_row(text, **fields)[0]


def test_a_trailing_postponed_parenthetical_stays_on_the_title():
    text = ("10. Doe J. Calibrating widgets at scale. Department of Widgets, "
            "Example University, April 7th, 2020 (postponed due to COVID-19).")
    assert _status_title(text) == f"{_STATUS_TITLE} (postponed due to COVID-19)"


def test_an_unbracketed_postponed_clause_stays_on_the_title():
    text = "Calibrating widgets at scale. Widget Society, Springfield. Postponed due to a hurricane."
    assert _status_title(text) == f"{_STATUS_TITLE} (Postponed due to a hurricane)"


def test_a_bracketed_cancellation_stays_whole():
    text = "Calibrating widgets at scale [Poster session canceled]. Widget Society Meeting."
    assert _status_title(text) == f"{_STATUS_TITLE} (Poster session canceled)"


def test_a_virtual_due_to_note_stays_on_the_title():
    text = "Calibrating widgets at scale. Widget Society, 2020 (virtual due to COVID-19)."
    assert _status_title(text) == f"{_STATUS_TITLE} (virtual due to COVID-19)"


def test_a_source_line_without_a_status_word_leaves_the_title_unchanged():
    text = "Calibrating widgets at scale. Widget Society, Springfield, April 7th, 2020."
    assert _status_title(text) == _STATUS_TITLE


def test_a_virtual_talk_with_no_status_clause_leaves_the_title_unchanged():
    text = "Calibrating widgets at scale. Virtual seminar, Widget Society, 2020."
    assert _status_title(text) == _STATUS_TITLE


def test_a_status_word_already_in_the_title_is_not_repeated():
    title = "Postponed surgery after the pandemic"
    assert _status_title(f"{title}. Widget Society, 2020.", title=title) == title


def test_a_status_word_already_in_the_venue_is_not_added_to_the_title():
    text = "Calibrating widgets at scale. Widget Society (cancelled), 2020."
    assert _status_title(text, location="Widget Society (cancelled)") == _STATUS_TITLE


# #1346: title-less talks whose source line still filled the Title cell with
# the date, event and venue the other two cells show.

def test_untitled_talk_drops_a_leading_numeric_date_and_its_talk_label():
    """PZTNWT 118: '<mm/dd/yyyy> <event><TAB><venue>. Invited talk.' was the
    whole Title cell."""
    cells = _untitled_row("12/15/2025 2025 Zqgpcr Workshop\tKestrel Island, Varnoria. Invited talk.",
                          date="2025-12-15", event_name="2025 Zqgpcr Workshop",
                          location="Kestrel Island, Varnoria")
    assert cells == ["2025 Zqgpcr Workshop, Invited talk", "Kestrel Island, Varnoria", "2025"]


def test_untitled_talk_with_only_a_venue_and_a_talk_label_is_titled_by_the_label():
    """PZTNWT 121: the venue came tab-separated where stage 4 wrote it with a
    comma, so the whole line stayed in the Title cell."""
    cells = _untitled_row("03/17/2025 Zqplanck Institute\tAshby, Varnoria. Invited talk.",
                          date="2025-03-17", location="Zqplanck Institute, Ashby, Varnoria")
    assert cells == ["Invited talk", "Zqplanck Institute, Ashby, Varnoria", "2025"]


def test_untitled_talk_drops_a_two_digit_year_date_and_the_pieces_of_its_event():
    """VPMMFM 524: '<mm/dd/yy>, <venue>,: <event part>, panelist, <event part>'."""
    cells = _untitled_row("10/27/04, Ashby, ZZ,:  Zqfib Roundtable, panelist, ZZ chapter ACX",
                          date="2004-10-27", role="Panelist", location="Ashby, ZZ",
                          event_name="Zqfib Roundtable, ZZ chapter ACX")
    assert cells == ["Panelist, Zqfib Roundtable, ZZ chapter ACX", "Ashby, ZZ", "2004"]


def test_a_numeric_date_is_kept_when_the_dates_cell_shows_another_year():
    cells = _untitled_row("10/27/04 Zqwidget Calibration at Scale, Northgate Hall",
                          date="2006", location="Northgate Hall")
    assert cells[0] == "10/27/04 Zqwidget Calibration at Scale"


def test_untitled_talk_keeps_a_line_with_words_beyond_the_label():
    """Only label words ('Invited talk') stop the line from counting as a
    possible title; the series name is not shown elsewhere."""
    cells = _untitled_row("08/19/2020 Zqcaltech, Kestrel, ZZ.\tSummer Zqpostdoc Series. Virtual. Invited talk.",
                          date="2020-08-19", location="Zqcaltech, Kestrel, ZZ")
    assert cells[0] == "Summer Zqpostdoc Series. Virtual. Invited talk."


def test_untitled_talk_counts_virtual_as_a_label_word():
    """PZTNWT 158 (#1346): with the series in event_name, 'Virtual. Invited
    talk.' is all that is left, so the stripped line no longer fills the cell."""
    cells = _untitled_row("08/19/2020 Zqcaltech, Kestrel, ZZ.\tSummer Zqpostdoc Series. Virtual. Invited talk.",
                          date="2020-08-19", event_name="Summer Zqpostdoc Series",
                          location="Zqcaltech, Kestrel, ZZ")
    assert cells == ["Summer Zqpostdoc Series, Virtual Invited talk", "Zqcaltech, Kestrel, ZZ", "2020"]


def test_untitled_talk_with_no_role_or_event_keeps_a_line_left_with_only_virtual():
    """1BFKKK 106 (#1346): with no role or event, a bare 'Virtual' would read
    as the talk's title and quiet the missing-title lint, so the line stays."""
    cells = _untitled_row("Zqemory University, 2022 (Virtual)", date="2022",
                          location="Zqemory University")
    assert cells == ["2022 (Virtual)", "Zqemory University", "2022"]


def test_untitled_talk_with_a_role_counts_virtual_as_a_label_word():
    """A role alone is enough beside 'Virtual': the row is titled by the
    role and the label, not the stripped line (#1346)."""
    cells = _untitled_row("2019 Zqkestrel University. Visiting Professor. Virtual.",
                          date="2019", role="Visiting Professor", location="Zqkestrel University")
    assert cells[0] == "Visiting Professor, Virtual"


def test_untitled_talk_takes_out_an_event_before_the_venue_it_contains():
    """PZTNWT 132 (#1346): the venue 'Zqcold Harbor' went first and broke the
    event 'Zqcold Harbor Meeting on Zqproteins', so the line stayed."""
    cells = _untitled_row("04/27/2024 Zqcold Harbor Meeting on Zqproteins\tZqcold Harbor. Invited talk.",
                          date="2024-04-27", event_name="Zqcold Harbor Meeting on Zqproteins",
                          location="Zqcold Harbor")
    assert cells == ["Zqcold Harbor Meeting on Zqproteins, Invited talk", "Zqcold Harbor", "2024"]


def test_untitled_talk_matches_an_event_whose_punctuation_the_source_spaces_differently():
    """PZTNWT 151 (#1346): stage 4 wrote 'SFB960/RCB Colloquium' for the
    source's 'SFB960 /RCB colloquium', so the event did not match and the
    stripped line filled the Title cell."""
    cells = _untitled_row("08/12/2021 Zqregen University, Varnoria. SFB960 /RCB colloquium.\tVirtual. Invited talk.",
                          date="2021-08-12", event_name="SFB960/RCB Colloquium, Invited Talk",
                          location="Zqregen University, Varnoria (Virtual)")
    assert cells == ["SFB960/RCB Colloquium, Invited Talk",
                     "Zqregen University, Varnoria (Virtual)", "2021"]


def test_untitled_talk_still_needs_whitespace_between_two_words_of_its_event():
    """The spacing tolerance is for punctuation only: event 'Zq Hall' does not
    match the run-together 'ZqHall', so the line stays (#1346)."""
    cells = _untitled_row("2019 ZqHall. Invited talk.", date="2019", event_name="Zq Hall")
    assert cells[0] == "ZqHall. Invited talk."


def test_a_titled_talk_renders_its_title_whatever_its_source_line_holds():
    """Blast radius (#1346): a row with a stage-4 title is untouched."""
    cells = _untitled_row("12/15/2025 Calibrating Zqwidgets\tBig Island. Invited talk.",
                          date="2025-12-15", title="Calibrating Zqwidgets",
                          location="Big Island")
    assert cells == ["Calibrating Zqwidgets", "Big Island", "2025"]


def test_a_two_digit_year_numeric_date_opening_a_title_is_dropped():
    """VPMMFM 524 (#1346): '<mm/dd/yy>' names the year the Dates cell shows."""
    cells = _untitled_row("10/27/04 Zqwidget Calibration at Scale, Northgate Hall",
                          date="2004-10-27", location="Northgate Hall")
    assert cells[0] == "Zqwidget Calibration at Scale"


def test_untitled_talk_drops_the_bracketed_piece_of_its_event():
    """An event written 'A (B)' appears in the line as 'A, B': both pieces
    are shown, so the title is the event (#1346)."""
    cells = _untitled_row("05/01/2019 Zqfib Roundtable, ZZ chapter. Northgate Hall",
                          date="2019-05-01", event_name="Zqfib Roundtable (ZZ chapter)",
                          location="Northgate Hall")
    assert cells == ["Zqfib Roundtable (ZZ chapter)", "Northgate Hall", "2019"]


@pytest.mark.parametrize("label", ["Invited lecture", "Invited presentation"])
def test_untitled_talk_counts_lecture_and_presentation_as_label_words(label):
    """'lecture' and 'presentation' are talk labels like 'talk' (#1346)."""
    cells = _untitled_row(f"2019 Zq Event, Zq Hall. {label}.", date="2019",
                          role="Speaker", event_name="Zq Event", location="Zq Hall")
    assert cells[0] == f"Speaker, Zq Event, {label}"


def test_every_talk_row_is_recorded_against_its_own_entry():
    """#534, #1205: the row-receipt ledger holds each talk's one row, in
    whichever scope table it renders, and no other entry's: the sub-point
    pass anchors a talk's prose only under it."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen.cv_owner_location = None          # skip the LLM geographic classifier
    national = {"taxonomy_code": "R", "text": "Zqwidget Calibration",
                "extracted_fields": {"title": "Zqwidget Calibration", "year": "2020"}}
    abroad = {"taxonomy_code": "R", "text": "Zqgear Alignment",
              "hierarchy": ["INVITATIONS TO SPEAK/PRESENT", "International"],
              "extracted_fields": {"title": "Zqgear Alignment", "year": "2019"}}
    gen._fill_presentations([national, abroad])

    rows = {}
    for entry, title in ((national, "Zqwidget Calibration"), (abroad, "Zqgear Alignment")):
        (row,) = gen._row_receipts.elements_for(entry)
        assert row.tag == qn("w:tr")
        assert "".join(t.text or "" for t in row.iter(qn("w:t"))).startswith(title)
        rows[title] = row
    assert rows["Zqwidget Calibration"].getparent() is not rows["Zqgear Alignment"].getparent()


@pytest.mark.parametrize("fields, venue", [
    ({"institution": "Zq Institute", "location": "Zq City", "venue": "Zq Hall"}, "Zq Institute"),
    ({"location": "Zq City", "venue": "Zq Hall"}, "Zq City"),
    ({"venue": "Zq Hall"}, "Zq Hall"),
])
def test_the_venue_cell_takes_institution_then_location_then_venue(fields, venue):
    """The writer's one precedence (`supplementary.PRESENTATION_VENUE_FIELDS`),
    which the #1205 sub-point pass counts as the row's."""
    assert _venue(**fields) == venue

