"""Section H (honors) newline-blind fix (#476, PR1 of the accuracy wave).

`honors.py:122` was one of the ten `entry_lines`-based call sites. Its own
per-part loop already disambiguates a tab ("does the last tab part look like
a year?" -- Award\\tYear splits, Award\\tDate\\tDescription rejoins into one
award) and a pipe ("Award | Year") *within a single already-newline-split
line*; the bug entry_lines caused was narrower than the other four sections':
a fully blind entry (no newline AT ALL) never reached that per-part logic as
more than one opaque blob, so a genuinely fused multi-award entry joined only
by '|' silently kept just its first award and dropped the rest (the OLD
`elif '|' in line:` branch reads only `parts[0]`/`parts[1]`, ignoring
anything past the first pipe).

`_entry_parts` fixes exactly that: when `entry_lines` already returns more
than one part (a genuine newline-separated entry), it is returned unchanged
-- migrating those to `entry_fragments` wholesale was tried and reverted
after reading the actual corpus diff it produced (see the function's own
docstring): 2054_Opresko_Cv's "1994 | American Chemical Society Award,\\n..."
already has 2 lines and collided with the pre-existing (separate) "DATE |
AWARD" mislabeling bug in the loop below, adding a spurious duplicate row;
2068_Yount_Cv's single blind entry mixes a tab AND a '|', and splitting the
'|' first defeats the tab-rejoin disambiguation, again producing a spurious
row ("Research.com" on its own). Both are why the fix is scoped to a
single-line, tab-free entry only -- disclosed here, not silently dropped.
There is no farm H entry of that exact shape (the farm's 9 tab-free blind '|'
entries are all "YEAR | Award, Org", already rendered correctly today via the
extracted fields' own single-award path -- this fix changes nothing for
them, hence CHANGED 0 in the render gate); the positive control below is a
synthetic fixture pinning the behaviour for when the farm does grow one.

The same emptiness of the farm hid a regression in the first cut of this
fix, which is why the last three tests exist: the pipeline's own table-cell
join writes ONE award as "Award | Organization | Year", and splitting every
'|' re-emitted the organization as a second, empty award row that origin/dev
never produced. `_entry_parts` now filters out the parts stage 4 already
extracted as column cells and requires the bare years to pair one-to-one
with the awards, so a column tuple is returned whole. The third test is the
counterweight: the two-award, two-year entry the fix exists for must still
split.

Review round 1 (#733) added a second half to this file: regression cover for
the parsing rules that review changed -- the two- and three-column '|' forms,
trailing year ranges, strong-before-generic organization keywords, the
case-sensitive proper-noun check, the narrower-table fallback, and the parser
and record the rest of it now goes through.

Round 1's second thread asked for the wider behaviour matrix, which is the
third part of this file. Writing it found three rendering losses that had
never failed a test, all of them silent:

  * the shared `_is_table_header_entry` calls a cell header-like when it
    contains one of the section's header words, and for honors "award" is one
    of them -- so a fused multi-award list ("Award A | 2024 | Award B | 2023 |
    Award C | 2022") counted as a header row and was dropped whole, rendering
    NOTHING. Three of the review's own pipe and tab shapes did that;
  * a column-header line fused into a multi-line entry was dropped by the
    parser and then rendered anyway, because the single-award record fell back
    to the entry's raw text with that line still in it;
  * a whitespace-only field or text rendered a visibly blank row, because the
    emptiness guard cannot see that " " is truthy.

Each is pinned by a test below whose docstring names it.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_honors_fragments.py -p no:cacheprovider
"""

import contextlib
import io
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from unified_pipeline.core.render_check import entry_lines  # noqa: E402
from unified_pipeline.stage6.parsing import _is_table_header_entry  # noqa: E402
from unified_pipeline.stage6.sections.honors import (  # noqa: E402
    _ENTRY_HEADER_KEYWORDS,
    _FALLBACK_SCHEMA_STAT,
    _MAX_RAW_AWARD_CHARS,
    HonorRecord,
    _entry_parts,
    _extract_organization_from_award,
    _honor_columns,
    _is_date_column,
    _is_honors_header_entry,
    _merge_first_cell_continuation,
    _split_award_year,
    parse_honor_entry,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


# Genuinely multi-line, farm-derived shapes: entry_lines already returns >1
# part for each of these, so _entry_parts must return them byte-identical.
MULTILINE_CASES = [
    # 2054_Opresko_Cv H entry shape: a pipe INSIDE the first of two newline
    # lines. Splitting further on '|' here is what produced the spurious row
    # -- _entry_parts must not touch it.
    "1994 | American Chemical Society Award,\nLehigh Valley Chapter of ACS",
    "1998-1999 | Mentored Investigator Award,\nThe Four Diamonds Fund of the Milton S. Hershey",
    "Award One\nAward Two",
    "2020 AICT Outstanding Article Award\n2018 Prior Award",
]

# Single-line (blind) shapes that must ALSO pass through untouched: a tab
# anywhere in the line (038WKA's 3-tab-field award, 2082_Dr_Scot's
# mid-sentence wrap artifact, 2068_Yount_Cv's tab+pipe combination).
BLIND_TAB_CASES = [
    "Award Name\tOctober 2024\tSelected from a national pool of applicants.",
    "1997 – National Institute for Staff and Organizational Development (NISOD) Award for\tTeaching Excellence. Valencia Community College, Orlando, Florida.",
    "2022\tWorld's Best Social Scientists, Ranking | Research.com",
]


_TEMPLATE_HEADERS = ["Name of award", "Organization", "Date awarded (yyyy)"]


def _honors_document(cols=3, heading="H. HONORS AND AWARDS",
                     table_before_heading=False, stale_rows=0,
                     with_table=True, verbose=False):
    """A minimal document with an H heading and its table: (generator, table).

    Deliberately hand-built rather than the bundled template, because the
    shapes below need a two-column variant, a missing heading, a missing table
    and a decoy table ahead of the heading -- none of which the real template
    can be made to have. `test_production_template_fills_its_own_honors_table`
    covers the real one.
    """
    gen = WCMTemplateGenerator(verbose=verbose)
    gen.doc = Document()
    if table_before_heading:
        gen.doc.add_paragraph("G. AN EARLIER SECTION")
        gen.doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "decoy"
    if heading is not None:
        gen.doc.add_paragraph(heading)
    if not with_table:
        return gen, None
    table = gen.doc.add_table(rows=1 + stale_rows, cols=cols)
    for i, header in enumerate(_TEMPLATE_HEADERS[:cols]):
        table.rows[0].cells[i].text = header
    for row in range(1, 1 + stale_rows):
        table.rows[row].cells[0].text = f"stale row {row}"
    return gen, table


def _fill_honors_into(entries, cols=3, **kwargs):
    """(rows, generator) after filling a `cols`-column H table."""
    gen, table = _honors_document(cols=cols, **kwargs)
    gen._fill_honors(entries)
    return [[c.text for c in r.cells] for r in table.rows[1:]], gen


def _render_honors(entries, cols=3):
    rows, _ = _fill_honors_into(entries, cols)
    return rows


def _raw(text):
    """An entry stage 4 extracted nothing from -- the fallback parse path."""
    return {"text": text, "extracted_fields": {}}


# --- (a) no line the old newline split produced is lost ---------------------

def test_multiline_entries_are_returned_unchanged():
    for case in MULTILINE_CASES:
        old = entry_lines(case)
        assert len(old) > 1, f"fixture {case!r} is not actually multi-line"
        assert _entry_parts(case) == old, f"diverged on already-multiline {case!r}"


def test_blind_lines_with_a_tab_are_also_returned_whole():
    for case in BLIND_TAB_CASES:
        old = entry_lines(case)
        assert old == [case]
        assert _entry_parts(case) == old, f"tab-bearing blind line split: {case!r}"


# --- (b) positive control: fails on dev today --------------------------------

def test_positive_control_pipe_blind_multi_award_entry_gains_both_rows():
    """Today's `elif '|' in line:` branch reads only the first two pipe
    segments of a single-line entry and silently drops the rest -- this is a
    single line (no '\\n'), so `entry_lines` returns exactly one opaque part
    and the multi-award branch (which needs >1 award_lines) never triggers;
    the entry falls to the single-award branch, which renders the whole raw
    text truncated at 150 chars as one garbled row. Fails on dev.
    """
    entry = {
        "text": "Best Teaching Award | 2020 | Excellence in Mentorship Award | 2018",
        "extracted_fields": {},
    }
    rows = _render_honors([entry])
    assert len(rows) == 2, f"expected 2 awards, got {rows}"
    assert rows[0][0] == "Best Teaching Award"
    assert rows[0][2] == "2020"
    assert rows[1][0] == "Excellence in Mentorship Award"
    assert rows[1][2] == "2018"


# --- (c) negative control: a single-part entry is unchanged ------------------

def test_negative_control_single_award_unchanged():
    entry = {
        "text": "2020 AECT Distinguished Service Award, Purdue University",
        "extracted_fields": {
            "award_name": "AECT Distinguished Service Award",
            "granting_body": "Purdue University",
            "date": "2020",
        },
    }
    rows = _render_honors([entry])
    assert len(rows) == 1
    assert rows[0][0] == "AECT Distinguished Service Award"
    assert rows[0][1] == "Purdue University"
    assert rows[0][2] == "2020"


def test_negative_control_farm_pipe_blind_year_first_shape_unchanged():
    """2054_Opresko_Cv's actual shape: a single-line 'YEAR | Award, Org'
    entry. One bare year and one award-looking part survive `_entry_parts`'
    column filter, which is one award -- below the two a multi-award list
    needs -- so the entry is returned whole and renders from the clean
    extracted fields exactly as it does on dev."""
    entry = {
        "text": "1991 | Freshman Chemistry Achievement Award, CRC Press",
        "extracted_fields": {
            "award_name": "Freshman Chemistry Achievement Award",
            "granting_body": "CRC Press",
            "date": "1991",
        },
    }
    rows = _render_honors([entry])
    assert len(rows) == 1
    assert rows[0][0] == "Freshman Chemistry Achievement Award"
    assert rows[0][1] == "CRC Press"
    assert rows[0][2] == "1991"


# --- regression: the pipeline's own three-column cell join ------------------

# "Award | Organization | Year" is how the pipeline's own table-cell join
# writes ONE award. Splitting it re-emitted the organization as a second,
# empty award row. Both variants below render exactly one row on origin/dev.
COLUMN_JOIN_TEXT = "Best Teaching Award | Purdue University | 2020"


def test_column_join_with_fields_is_not_split_into_two_awards():
    entry = {
        "text": COLUMN_JOIN_TEXT,
        "extracted_fields": {
            "award_name": "Best Teaching Award",
            "granting_body": "Purdue University",
            "date": "2020",
        },
    }
    assert _entry_parts(COLUMN_JOIN_TEXT, ("Purdue University", "2020")) == [
        COLUMN_JOIN_TEXT]
    rows = _render_honors([entry])
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]], rows


def test_column_join_without_fields_is_not_split_into_two_awards():
    """The year-balance half of the guard, with nothing extracted to filter
    with: two award-looking parts and a single bare year is a column tuple,
    not two awards that each carry a year.

    The row it renders is asserted in full since the three-column fix
    (#733 review item 1): before that fix this text reached the table as one
    row whose award cell was the whole raw string, pipes included, with the
    organization column left empty.
    """
    assert _entry_parts(COLUMN_JOIN_TEXT) == [COLUMN_JOIN_TEXT]
    rows = _render_honors([_raw(COLUMN_JOIN_TEXT)])
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]], rows


def test_pipe_blind_multi_award_split_survives_the_guard():
    """The guard must not swallow the case the fix exists for: two awards
    that each carry their own year still split."""
    text = "Best Teaching Award | 2020 | Excellence in Mentorship Award | 2018"
    assert _entry_parts(text, ("", "")) == [
        "Best Teaching Award", "2020",
        "Excellence in Mentorship Award", "2018"]


# =========================================================================
# Review round 1 (#733): regression cover for the parsing rules the review
# changed. The reviewer's wider behaviour matrix -- sorting, verbose mode,
# duplicate awards, production-template integration and the rest -- is a
# separate pass; what follows is only the cover the source items themselves
# asked for.
#
# A note that explains the parser-level tests below: the shared entry filter
# `_is_table_header_entry` (stage6/parsing/text.py) classifies a SHORT
# pipe-joined entry as a source table's header row when half its cells carry
# a header word, and "award" is one of them. "Best Teaching Award | 2020" is
# therefore dropped by the renderer before the parser ever sees it. That is a
# pre-existing filter outside this section, so the two-column shapes it eats
# are pinned on `parse_honor_entry`/`_honor_columns` directly and the
# rendering assertions use shapes the filter passes.
# =========================================================================


# --- item 1: two- and three-column '|' forms --------------------------------

def test_three_column_pipe_form_fills_all_three_cells():
    """`Award | Organization | Year` put the ORGANIZATION in the date column
    and lost the date: the old branch read cell 2 as the year unconditionally.
    """
    rows = _render_honors([_raw("Best Teaching Award | Purdue University | 2020")])
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]]


def test_three_column_pipe_form_accepts_a_month_and_year_date_cell():
    """A date cell is not only a bare year. Reading it as one split this text
    into three awards, because none of its three cells was a bare year."""
    rows = _render_honors(
        [_raw("Best Teaching Award | Purdue University | August 2020")])
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]]


def test_two_column_pipe_form_reads_the_second_cell_as_the_date():
    """`Award | Year`, the two-column form, still means award and date."""
    assert _honor_columns("Best Teaching Award | 2020") == HonorRecord(
        "Best Teaching Award", "", "2020")
    assert parse_honor_entry(_raw("Best Teaching Award | 2020")) == [
        HonorRecord("Best Teaching Award", "", "2020")]


def test_two_column_pipe_form_does_not_call_an_organization_a_date():
    """The other two-column form. `parts[1]` was previously assumed to be a
    year whatever it said, which put the granting body in the date column."""
    assert _honor_columns("Best Teaching Award | Purdue University") == \
        HonorRecord("Best Teaching Award", "Purdue University", "")


def test_pipe_shapes_that_name_no_columns_are_left_alone():
    """Anything outside the two documented forms keeps the historical
    first-cell/second-cell reading rather than being guessed at."""
    assert _honor_columns("2020 | Best Teaching Award") is None      # leading date
    assert _honor_columns("Award A | 2024 | Award B") is None        # odd count
    assert _honor_columns("A | B | C | 2020") is None                # four cells
    assert _honor_columns("Award | 2019 | 2020") is None             # middle date


def test_a_pipe_inside_an_award_name_is_read_as_a_column_boundary():
    """The limitation the module docstring states: '|' is structural, so this
    is one award, an organization and a date -- not one award whose name
    contains a pipe, and not two awards."""
    rows = _render_honors(
        [_raw("Excellence in Research | Teaching Award | 2024")])
    assert rows == [["Excellence in Research", "Teaching Award", "2024"]]


def test_extracted_fields_beat_the_pipe_columns():
    """The column reading is a fallback: whatever stage 4 extracted wins."""
    rows = _render_honors([{
        "text": "Best Teaching Award | Purdue University | 2020",
        "extracted_fields": {"award_name": "Best Teaching Award",
                             "granting_body": "Purdue Research Foundation",
                             "date": "2019"},
    }])
    assert rows == [["Best Teaching Award", "Purdue Research Foundation",
                     "2019"]]


# --- item 3: trailing year ranges -------------------------------------------

def test_trailing_year_range_parses_with_either_dash():
    """The contract named ranges but only the LEADING half accepted one, so
    'Award 2015-2017' kept a dangling '2015 -' in the award name."""
    assert _split_award_year("Distinguished Service Award 2015-2017") == (
        "Distinguished Service Award", "2015-2017")
    assert _split_award_year("Distinguished Service Award 2015\u20132017") == (
        "Distinguished Service Award", "2015\u20132017")
    assert _split_award_year("Distinguished Service Award (2015-2017)") == (
        "Distinguished Service Award", "2015-2017")
    assert _split_award_year("Distinguished Service Award 2015 - present") == (
        "Distinguished Service Award", "2015 - present")


def test_leading_and_trailing_year_forms_agree():
    """Both ends accept the same year grammar, which is what 'consistent'
    means here: same award name out, same year out."""
    for leading, trailing in [
            ("2015-2017 Long Fellowship", "Long Fellowship 2015-2017"),
            ("2015\u20132017 Long Fellowship", "Long Fellowship 2015\u20132017"),
            ("2020 Long Fellowship", "Long Fellowship 2020")]:
        assert _split_award_year(leading) == _split_award_year(trailing), (
            leading, trailing)


def test_trailing_year_range_reaches_the_date_cell():
    rows = _render_honors([_raw("Distinguished Service Award 2015-2017")])
    assert rows == [["Distinguished Service Award", "", "2015-2017"]]


# --- item 4: strong organization keywords before the generic one ------------

def test_generic_org_keyword_inside_an_award_name_does_not_win():
    """'Center' is generic enough to sit inside an award's own name, and
    anchoring on the LAST keyword of one flat list let it beat the body that
    granted the award purely by coming later."""
    assert _extract_organization_from_award(
        "Purdue University Clinical Simulation Center Award") == \
        "Purdue University"
    assert _extract_organization_from_award(
        "Purdue University Teaching Center Award") == "Purdue University"
    assert _extract_organization_from_award(
        "Best Teaching Award, Purdue University, Cancer Center") == \
        "Purdue University"


def test_generic_org_keyword_still_wins_when_nothing_stronger_appears():
    assert _extract_organization_from_award(
        "Innovation Award, Simulation Center") == "Simulation Center"


def test_generic_org_keyword_still_extends_a_name_it_continues():
    """The ambiguity cuts both ways: here 'Center' is the tail of the
    hospital's own name, not part of the award, so the span keeps running to
    it. Award vocabulary in between is what tells the two apart."""
    assert _extract_organization_from_award(
        "Faculty Teaching Award, New York Presbyterian Hospital "
        "Weill Cornell Center") == \
        "New York Presbyterian Hospital Weill Cornell Center"


def test_no_institutional_keyword_yields_no_organization():
    assert _extract_organization_from_award("Nothing institutional here") == ""
    assert _extract_organization_from_award("") == ""


# --- item 7: the proper-noun check is case-sensitive ------------------------

def test_named_body_needs_a_real_proper_noun():
    """[A-Z] is the signal that what follows 'of'/'for' is a proper name.
    Compiling the whole pattern with re.IGNORECASE made it match a lowercase
    letter, so the branch accepted text with no proper name in it at all."""
    assert _extract_organization_from_award(
        "Society of Fictional Medicine Mentoring Award") == \
        "Society of Fictional Medicine"
    assert _extract_organization_from_award(
        "society of fictional medicine Mentoring Award") == ""


def test_named_body_keyword_itself_stays_case_insensitive():
    """Only the proper-name half became case-sensitive: a CV that writes the
    institutional keyword in lower case still matches, as long as the name
    after it is capitalized."""
    assert _extract_organization_from_award(
        "society of Fictional Medicine Mentoring Award") == \
        "society of Fictional Medicine"


# --- item 2: a template variant with fewer columns --------------------------

_ORG_ENTRY = {
    "text": "Best Teaching Award, Purdue University, 2020",
    "extracted_fields": {"award_name": "Best Teaching Award",
                         "granting_body": "Purdue University",
                         "date": "2020"},
}


def test_two_column_table_keeps_the_granting_body():
    """A two-column variant mapped cell 2 to the date and dropped the
    granting body outright, producing a valid-looking DOCX missing it."""
    rows, gen = _fill_honors_into([_ORG_ENTRY], cols=2)
    assert rows == [["Best Teaching Award, Purdue University", "2020"]]
    assert gen.stats[_FALLBACK_SCHEMA_STAT] == 1


def test_one_column_table_keeps_the_granting_body_and_the_date():
    rows, gen = _fill_honors_into([_ORG_ENTRY], cols=1)
    assert rows == [["Best Teaching Award, Purdue University (2020)"]]
    assert gen.stats[_FALLBACK_SCHEMA_STAT] == 1


def test_the_full_template_records_no_schema_fallback():
    rows, gen = _fill_honors_into([_ORG_ENTRY])
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]]
    assert gen.stats.get(_FALLBACK_SCHEMA_STAT, 0) == 0


# --- items 5 and 6: a pure parser over a typed record -----------------------

def test_parse_honor_entry_is_pure_and_typed():
    """The parser needs no document, no generator and no template: it is
    reachable on a plain dict and answers in HonorRecords."""
    records = parse_honor_entry({
        "text": "2020 AECT Distinguished Service Award, Purdue University",
        "extracted_fields": {},
    })
    assert records == [HonorRecord("AECT Distinguished Service Award",
                                   "Purdue University", "2020")]
    assert all(isinstance(r, HonorRecord) for r in records)


def test_malformed_fields_never_reach_a_table_cell():
    """`extracted_fields` is not a contract anything enforces: a key can be
    absent, None or a number. Coercing at the parser boundary is what keeps
    None out of a cell -- and out of _strip_org_tail, which raised on it."""
    records = parse_honor_entry({
        "text": "Best Poster Award, Imaginary Society of Things, 2021.",
        "extracted_fields": {"award_name": None, "granting_body": None,
                             "date": None},
    })
    assert records == [HonorRecord("Best Poster Award",
                                   "Imaginary Society of Things", "2021")]
    # a non-string field is coerced rather than written into the cell as-is
    assert parse_honor_entry({
        "text": "Best Poster Award",
        "extracted_fields": {"award_name": "Best Poster Award", "date": 2021},
    }) == [HonorRecord("Best Poster Award", "", "2021")]


def test_an_entry_with_no_text_renders_nothing_rather_than_a_blank_row():
    """A None text used to raise TypeError slicing it; an empty row is not
    the honest answer either."""
    assert parse_honor_entry({"text": None, "extracted_fields": None}) == [
        HonorRecord("", "", "")]
    assert _render_honors([{"text": None, "extracted_fields": None}]) == []


# =========================================================================
# Review round 1 (#733), thread 2: the behaviour matrix the review listed.
# Empty entries, `_entry_parts`, `_fill_honors`, rendering, the production
# template, a missing heading, header entries, the extracted-fields fallback
# matrix, date formats, sorting, multiple entries, the odd pipe structures
# named verbatim, tabs, the newline-and-pipe interaction, annotations, table
# accuracy, table lookup, the long-raw-text fallback, None and malformed
# fields, whitespace, duplicates, missing text and verbose mode.
#
# Three of these pin fixes this round made rather than behaviour that was
# already right; their docstrings say which. The rest pin what the section
# already did, which is what a matrix is for.
# =========================================================================


# --- empty and missing input ------------------------------------------------

def test_no_entries_leaves_the_table_alone():
    """`_fill_honors([])` returns before `_clear_table_data`, so an empty
    honors list cannot blank rows a template shipped with, and does not count
    a table as populated."""
    rows, gen = _fill_honors_into([], stale_rows=1)
    assert rows == [["stale row 1", "", ""]]
    assert gen.stats['tables_populated'] == 0
    assert gen.stats['entries_inserted'] == 0


def test_missing_section_heading_renders_nothing():
    """Neither "HONORS" nor "AWARDS" appears: the section is absent from this
    template variant, so nothing is written and no table is claimed."""
    rows, gen = _fill_honors_into([_ORG_ENTRY], heading="Z. SOME OTHER SECTION")
    assert rows == []
    assert gen.stats['tables_populated'] == 0


def test_heading_with_no_table_after_it_renders_nothing():
    """The heading is found but `_find_table_after_paragraph` returns None --
    a variant that writes honors as prose. The renderer returns rather than
    reaching into a table it does not have."""
    gen, table = _honors_document(with_table=False)
    gen._fill_honors([_ORG_ENTRY])
    assert table is None
    assert gen.stats['tables_populated'] == 0
    assert gen.stats['entries_inserted'] == 0


def test_entries_with_nothing_in_them_render_no_rows():
    """The emptiness guard in `_fill_honors`. The whitespace-only cases are
    the ones it could not see before this round: " " is truthy, so a
    whitespace-only text or field reached a cell and rendered a visibly blank
    row instead of nothing at all (#733 review)."""
    rows, gen = _fill_honors_into([
        {},
        {"text": ""},
        {"text": "   "},
        {"text": None, "extracted_fields": None},
        {"text": "  ", "extracted_fields": {"award_name": "  "}},
    ])
    assert rows == []
    assert gen.stats['entries_inserted'] == 0


def test_padded_fields_reach_the_cells_stripped():
    """The other half of that fix: a padded field renders stripped, and a
    whitespace-only one counts as absent rather than as content."""
    assert parse_honor_entry({
        "text": "2018 Fictional Merit Award",
        "extracted_fields": {"award_name": " Padded Award ",
                             "granting_body": "   ", "date": " 2018 "},
    }) == [HonorRecord("Padded Award", "", "2018")]


# --- `_entry_parts` unit tests ----------------------------------------------

def test_entry_parts_on_empty_and_delimiterless_text():
    """The degenerate inputs. `entry_lines` coerces None and drops blank
    parts, and `_entry_parts` inherits both."""
    assert _entry_parts("") == []
    assert _entry_parts("   ") == []
    assert _entry_parts(None) == []
    assert _entry_parts("Single award, no delimiter") == [
        "Single award, no delimiter"]


def test_entry_parts_strips_parts_and_drops_empty_ones():
    """One rule covers four of the review's odd shapes: a doubled, leading or
    trailing pipe contributes no empty part, and every part comes back
    stripped, so all four yield the same four parts."""
    for text in ["Award A | 2024 || Award B | 2023",
                 " Award A   |   2024   |   Award B   |   2023 ",
                 "Award A | 2024 | Award B | 2023 |",
                 "| Award A | 2024 | Award B | 2023"]:
        assert _entry_parts(text) == [
            "Award A", "2024", "Award B", "2023"], text


def test_entry_parts_declines_the_shapes_it_cannot_pair():
    """Fewer than two award-looking parts, or dates that do not pair one to
    one with them: the pipes are column separators, not record separators, so
    the line comes back whole."""
    for text in ["Award A | 2024 | Award B",
                 "Best Teaching Award | 2020",
                 "Best Teaching Award | Purdue University | 2020"]:
        assert _entry_parts(text) == [text], text


def test_entry_parts_uses_the_extracted_columns_to_refuse_a_split():
    """A part stage 4 already extracted as a column cell is a column, not an
    award -- without that filter the pipeline's own "Award | Organization |
    Year" cell join split into two awards."""
    text = "Best Teaching Award | Purdue University | 2020"
    assert _entry_parts(text, ("Purdue University", "2020")) == [text]


# The review's own list, with the rows each shape must render. The first five
# are the fused multi-award lists #476 exists to recover; the sixth is the
# shape the parser deliberately declines to guess at.
ODD_PIPE_SHAPES = [
    ("three awards",
     "Award A | 2024 | Award B | 2023 | Award C | 2022",
     [["Award A", "", "2024"], ["Award B", "", "2023"],
      ["Award C", "", "2022"]]),
    ("empty fragments",
     "Award A | 2024 || Award B | 2023",
     [["Award A", "", "2024"], ["Award B", "", "2023"]]),
    ("extra whitespace",
     " Award A   |   2024   |   Award B   |   2023 ",
     [["Award A", "", "2024"], ["Award B", "", "2023"]]),
    ("trailing pipe",
     "Award A | 2024 | Award B | 2023 |",
     [["Award A", "", "2024"], ["Award B", "", "2023"]]),
    ("leading pipe",
     "| Award A | 2024 | Award B | 2023",
     [["Award A", "", "2024"], ["Award B", "", "2023"]]),
    ("odd number of parts",
     "Award A | 2024 | Award B",
     [["Award A | 2024 | Award B", "", ""]]),
]


# --- header rows must not render --------------------------------------------

def test_a_real_column_header_row_never_renders():
    """The source CV's own header row, extracted as if it were data. It names
    a date column but carries no date, which is what still identifies it."""
    for text in ["Name of award\tOrganization\tDate awarded",
                 "Name of award | Organization | Date awarded",
                 "Honor | Granting body | Year",
                 "Award | Organization | Date awarded (yyyy)"]:
        assert _is_honors_header_entry(text), text
        assert _render_honors([_raw(text)]) == [], text


def test_a_fused_multi_award_list_is_not_a_header_row():
    """The first of this round's three rendering losses. The shared
    `_is_table_header_entry` counts a cell as header-like when it contains one
    of the section's header words, and "award" is one of them -- so every
    award cell of a fused list counted and the entry was dropped before the
    parser saw it, rendering nothing at all. Three of the review's own shapes
    did that, and the assertion below names which."""
    shared_check_says_header = [
        name for name, text, _ in ODD_PIPE_SHAPES
        if _is_table_header_entry(text, _ENTRY_HEADER_KEYWORDS)]
    assert shared_check_says_header == [
        "three awards", "extra whitespace", "odd number of parts"], \
        shared_check_says_header
    for name, text, _ in ODD_PIPE_SHAPES:
        assert not _is_honors_header_entry(text), name


def test_a_header_row_fused_into_a_multi_line_entry_does_not_reach_a_cell():
    """The second. `_parse_honor_lines` drops a header LINE, but the
    single-award record then fell back to the entry's RAW text -- that line
    still in it -- so the source table's header rendered as the award name
    (#733 review)."""
    assert _render_honors([
        _raw("Name of award\tDate awarded\nBest Award\t2020")]) == [
        ["Best Award", "", "2020"]]


# --- the odd pipe structures, named verbatim (ODD_PIPE_SHAPES, above) -------

def test_odd_pipe_structures_render_the_rows_they_name():
    """Three of these -- the three the assertion above names -- rendered no
    rows at all before this round, whatever `_entry_parts` made of them."""
    for name, text, expected in ODD_PIPE_SHAPES:
        assert _render_honors([_raw(text)]) == expected, name


def test_an_unpairable_pipe_shape_renders_as_its_own_text():
    """The documented choice for "Award A | 2024 | Award B": two award-looking
    parts and one date pair no year to an award, and nothing in the text says
    which award 2024 belongs to. The line renders whole rather than guessed
    at -- the property that matters is that no part of it is dropped."""
    rows = _render_honors([_raw("Award A | 2024 | Award B")])
    assert rows == [["Award A | 2024 | Award B", "", ""]]
    assert "Award B" in rows[0][0] and "2024" in rows[0][0]


# --- tabs, and the newline-and-pipe interaction -----------------------------

def test_tab_plus_multiple_pipes_renders_one_award():
    """Asserted through to the rendered row, not just `_entry_parts`. The tab
    cells join because the last one is not a bare year, and the joined line is
    then read as the three columns it now has. The leading "2020" stays inside
    the award's name: peeling it would mean choosing between it and the 2019
    the date column states, which nothing in the text settles."""
    assert _entry_parts("2020\tAward A | Award B | 2019") == [
        "2020\tAward A | Award B | 2019"]
    assert _render_honors([_raw("2020\tAward A | Award B | 2019")]) == [
        ["2020 Award A", "Award B", "2019"]]


def test_each_newline_part_is_rendered_as_expected():
    """`_entry_parts` returning the right structure does not prove
    `_fill_honors` interprets it: this entry's parts were always correct and
    it still rendered nothing, because the shared header check ate the whole
    entry first."""
    text = "Award A | 2024\nAward B | 2023"
    assert _entry_parts(text) == ["Award A | 2024", "Award B | 2023"]
    assert _render_honors([_raw(text)]) == [["Award A", "", "2024"],
                                            ["Award B", "", "2023"]]


def test_newline_separated_awards_keep_their_own_leading_years():
    """The plainest multi-line shape: one award per line, each with its year
    in front of it."""
    assert _render_honors([
        _raw("2020 Award A\n2018 Award B\n2016 Award C")]) == [
        ["Award A", "", "2020"], ["Award B", "", "2018"],
        ["Award C", "", "2016"]]


# --- the extracted-fields fallback matrix -----------------------------------

# One raw text carrying all three fields, against every combination of what
# stage 4 may or may not have extracted from it.
_MATRIX_TEXT = "2018 Fictional Merit Award, Imaginary University"

FALLBACK_MATRIX = [
    ({},
     HonorRecord("Fictional Merit Award", "Imaginary University", "2018")),
    ({"award_name": "Fictional Merit Award"},
     HonorRecord("Fictional Merit Award", "", "2018")),
    ({"granting_body": "Imaginary University"},
     HonorRecord("Fictional Merit Award", "Imaginary University", "2018")),
    ({"date": "2018"},
     HonorRecord("2018 Fictional Merit Award", "Imaginary University", "2018")),
    ({"award_name": "Fictional Merit Award", "date": "2018"},
     HonorRecord("Fictional Merit Award", "", "2018")),
    ({"award_name": "Fictional Merit Award",
      "granting_body": "Imaginary University"},
     HonorRecord("Fictional Merit Award", "Imaginary University", "2018")),
    ({"granting_body": "Imaginary University", "date": "2018"},
     HonorRecord("2018 Fictional Merit Award", "Imaginary University", "2018")),
    ({"award_name": "Fictional Merit Award",
      "granting_body": "Imaginary University", "date": "2018"},
     HonorRecord("Fictional Merit Award", "Imaginary University", "2018")),
]


def test_extracted_fields_fallback_matrix():
    """Each extracted field replaces the corresponding piece of the parse, and
    only that piece; what stage 4 left out is still parsed out of the text.
    Note row 2: the organization is extracted from the award NAME, so an
    entry that supplies the name and nothing else loses the granting body the
    raw text names -- pinned here because it is the one asymmetry in the
    matrix."""
    for fields, expected in FALLBACK_MATRIX:
        assert parse_honor_entry({"text": _MATRIX_TEXT,
                                  "extracted_fields": fields}) == [expected], \
            fields


def test_an_extracted_date_leaves_the_name_s_own_leading_year_in_place():
    """A residual, pinned rather than fixed. `_split_award_year` runs only
    when stage 4 extracted no date, so a text that leads with a year plus an
    extracted date renders the year twice -- once in the award cell, once in
    the date cell. No farm honors entry has that combination (0 of 361), and
    peeling the year would mean deciding what to do when it disagrees with the
    extracted date, so it is left where the source wrote it."""
    assert parse_honor_entry({
        "text": _MATRIX_TEXT, "extracted_fields": {"date": "2018"}}) == [
        HonorRecord("2018 Fictional Merit Award", "Imaginary University",
                    "2018")]


def test_alternate_field_key_names_are_read():
    """`organization` stands in for `granting_body` and `year` for `date`."""
    assert parse_honor_entry({
        "text": "ignored",
        "extracted_fields": {"award_name": "A Prize",
                             "organization": "Imaginary University",
                             "year": "2001"},
    }) == [HonorRecord("A Prize", "Imaginary University", "2001")]


def test_a_falsy_extracted_fields_is_treated_as_absent():
    """`entry.get('extracted_fields') or {}` covers every falsy malformation
    the upstream record can carry -- the key missing, None, an empty dict, an
    empty list, an empty string -- and the entry falls back to its raw text
    rather than raising on one of them."""
    for fields in [None, {}, [], ""]:
        assert parse_honor_entry({"text": "2020 Some Award",
                                  "extracted_fields": fields}) == [
            HonorRecord("Some Award", "", "2020")], fields


# --- date formats -----------------------------------------------------------

DATE_CASES = [
    ("2020", "2020"),
    ("March 2019", "2019"),
    ("2019-03-01", "2019"),
    ("03/2019", "2019"),
    ("2015-2017", "2015-2017"),
    ("2015–2017", "2015–2017"),
    ("present", "Present"),
    ("n.d.", "n.d."),
]


def test_every_date_shape_reduces_to_the_templates_year_column():
    """H's column is "Date awarded (yyyy)", so `format_date_for_section`
    reduces a month or a full date to its year. A range keeps both ends in
    whichever dash the source used, and a string it cannot parse is passed
    through rather than blanked -- an unparseable date is still information."""
    for given, expected in DATE_CASES:
        assert _render_honors([{
            "text": "X Award",
            "extracted_fields": {"award_name": "X Award", "date": given},
        }]) == [["X Award", "", expected]], given


# --- sorting, multiple entries, duplicates ----------------------------------

def test_entries_render_most_recent_first():
    """`sort_entries_reverse_chronological` orders the ENTRIES, before any of
    them is parsed into rows."""
    entry = lambda name, date: {                          # noqa: E731
        "text": name, "extracted_fields": {"award_name": name, "date": date}}
    assert _render_honors([entry("Older Award", "2015"),
                           entry("Newest Award", "2024"),
                           entry("Middle Award", "2019")]) == [
        ["Newest Award", "", "2024"], ["Middle Award", "", "2019"],
        ["Older Award", "", "2015"]]


def test_undated_entries_keep_the_order_they_arrived_in():
    """Nothing to sort by, so the sort is stable and the CV's own order
    survives -- an undated honors list is not silently reshuffled."""
    assert _render_honors([_raw("Alpha Award"), _raw("Beta Award"),
                           _raw("Gamma Award")]) == [
        ["Alpha Award", "", ""], ["Beta Award", "", ""],
        ["Gamma Award", "", ""]]


def test_multiple_entries_of_different_shapes_all_render():
    """A fused list, a column tuple and a plain extracted entry in one
    document: four rows from three entries, and `entries_inserted` counts
    rows, not entries."""
    rows, gen = _fill_honors_into([
        _raw("Award A | 2024 | Award B | 2023"),
        _raw("Solo Award | Imaginary University | 2022"),
        {"text": "Plain Award",
         "extracted_fields": {"award_name": "Plain Award", "date": "2021"}},
    ])
    assert rows == [["Plain Award", "", "2021"],
                    ["Award A", "", "2024"],
                    ["Award B", "", "2023"],
                    ["Solo Award", "Imaginary University", "2022"]]
    assert gen.stats['entries_inserted'] == 4


def test_duplicate_awards_both_render():
    """Deduplication is not this section's job: two identical entries render
    two identical rows, so a duplicate that reached stage 6 stays visible
    rather than being quietly halved."""
    rows, gen = _fill_honors_into([dict(_ORG_ENTRY), dict(_ORG_ENTRY)])
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]] * 2
    assert gen.stats['entries_inserted'] == 2


# --- annotations ------------------------------------------------------------

def test_annotations_inside_an_award_name_are_kept():
    """A parenthetical is part of the award's name, not a delimiter -- 44 of
    the farm's 361 honors entries carry one, and a trailing year is still
    peeled off from behind it."""
    assert _render_honors([_raw(
        "Best Teaching Award (nominated), Purdue University, 2020")]) == [
        ["Best Teaching Award (nominated)", "Purdue University", "2020"]]
    assert _render_honors([_raw("Dean's List (Fall semester) 2019")]) == [
        ["Dean's List (Fall semester)", "", "2019"]]


def test_an_annotation_after_the_year_keeps_the_year_in_the_name():
    """A residual, pinned rather than fixed. `_TRAILING_YEAR_RE` allows a
    closing bracket or a full stop after the year, not a whole trailing
    annotation, so the date cell stays empty and the year stays in the award
    cell. No farm honors entry has an annotation after a trailing year (0 of
    361); the 44 that carry one carry it inside the name, above."""
    assert _render_honors([_raw(
        "Best Award, Purdue University, 2020 [declined]")]) == [
        ["Best Award, Purdue University, 2020 [declined]",
         "Purdue University", ""]]


# --- table accuracy, table lookup, the long-raw-text cap --------------------

def test_the_header_row_survives_and_stale_rows_do_not():
    """`_clear_table_data(keep_header=True)`: the template's own header row
    stays, and rows a previous run left behind are removed rather than
    appended to."""
    rows, gen = _fill_honors_into([_ORG_ENTRY], stale_rows=2)
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]]
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['entries_inserted'] == 1


def test_every_rendered_cell_is_arial_11():
    """`_add_honors_row` fonts every run of every cell it writes, including
    the two it may have left empty."""
    gen, table = _honors_document()
    gen._fill_honors([_ORG_ENTRY])
    runs = [run for cell in table.rows[1].cells
            for para in cell.paragraphs for run in para.runs]
    assert len(runs) == 3
    for run in runs:
        assert run.font.name == "Arial"
        assert run.font.size == Pt(11)
        assert run.bold is False
        assert run.italic is False


def test_the_table_filled_is_the_one_after_the_heading():
    """`_find_table_after_paragraph`: a table earlier in the document is not
    the honors table, however much it looks like one."""
    rows, gen = _fill_honors_into([_ORG_ENTRY], table_before_heading=True)
    assert rows == [["Best Teaching Award", "Purdue University", "2020"]]
    decoy = gen.doc.tables[0]
    assert len(decoy.rows) == 1
    assert [c.text for c in decoy.rows[0].cells] == ["decoy", ""]


def test_a_long_raw_entry_is_capped_not_spilled():
    """An entry stage 4 extracted nothing from renders as its own text, so
    the cap is the only thing stopping a runaway blob filling the cell."""
    text = "Recognition of Sustained Contribution " * 8
    rows = _render_honors([_raw(text)])
    assert len(rows) == 1
    assert rows[0][0] == text.strip()[:_MAX_RAW_AWARD_CHARS]
    assert len(rows[0][0]) == _MAX_RAW_AWARD_CHARS


# --- verbose mode -----------------------------------------------------------

def test_verbose_mode_reports_the_entry_count_and_each_skipped_header():
    """The two `print()`s in `_fill_honors`. Neither is read by
    `orchestrator.py`'s progress regexes nor by `run_corpus_batch.sh`'s
    summary greps, so neither is a parsed contract -- but both are pinned here
    so that rewording one is a test failure rather than a silent change."""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        gen, _ = _honors_document(verbose=True)
        gen._fill_honors([
            _ORG_ENTRY,
            _raw("Name of award | Organization | Date awarded")])
    out = buf.getvalue()
    assert "Filling Honors (2 entries)..." in out
    assert ("Skipping header entry: 'Name of award | Organization | "
            "Date awarded") in out


def test_quiet_mode_prints_nothing():
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        gen, _ = _honors_document(verbose=False)
        gen._fill_honors([_ORG_ENTRY])
    assert buf.getvalue() == ""


# --- production-template integration ----------------------------------------

def test_production_template_fills_its_own_honors_table():
    """The bundled WCM template rather than a hand-built document: its honors
    heading is found by substring ("HONORS, AWARDS"), the table after it is
    the template's own three-column one, and both a plain entry and a fused
    multi-award list land in it -- so no schema fallback is recorded."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen._fill_honors([
        _ORG_ENTRY,
        _raw("Award A | 2024 | Award B | 2023 | Award C | 2022")])

    table = gen._find_table_after_paragraph(
        gen._find_paragraph_with_text("HONORS"))
    rows = [[c.text for c in row.cells] for row in table.rows]
    assert rows[0] == _TEMPLATE_HEADERS
    assert rows[1:] == [["Best Teaching Award", "Purdue University", "2020"],
                        ["Award A", "", "2024"],
                        ["Award B", "", "2023"],
                        ["Award C", "", "2022"]]
    assert gen.stats.get(_FALLBACK_SCHEMA_STAT, 0) == 0
    assert gen.stats['tables_populated'] == 1
    assert gen.stats['entries_inserted'] == 4


# =========================================================================
# #828: a full mm/dd/yyyy (or "30 October 2017" / "October 30, 2017") date
# cell in a 3-column award row is a date, not a second award. Synthetic
# names throughout -- none of these strings are corpus content.
# =========================================================================

_FULL_DATE_ENTRY = {
    "text": "Fictional Excellence Award | Imaginary Testing Society | 10/30/2017",
    "extracted_fields": {
        "award_name": "Fictional Excellence Award",
        "granting_body": "Imaginary Testing Society",
        "date": "2017-10-30",
    },
}


def test_full_date_cell_is_not_read_as_a_second_award():
    """The issue's own shape: a 3-column row whose date cell is a full
    mm/dd/yyyy date, not the bare year the template's column header asks
    for. Before the fix neither `_is_date_column` nor the claimed-value
    filter recognised '10/30/2017' as the date column, so two award-looking
    parts (the award and the date) survived and the row split into three."""
    rows = _render_honors([_FULL_DATE_ENTRY])
    assert rows == [["Fictional Excellence Award", "Imaginary Testing Society",
                     "2017"]]


def test_alternate_full_date_cell_shapes_each_yield_one_record():
    """`_DATE_COLUMN_RE` was widened to accept two more full-date shapes
    (#828); each renders the same single row as the mm/dd/yyyy form above."""
    for cell in ("30 October 2017", "October 30, 2017"):
        assert _is_date_column(cell), cell
        entry = {
            "text": f"Fictional Excellence Award | Imaginary Testing Society | {cell}",
            "extracted_fields": {
                "award_name": "Fictional Excellence Award",
                "granting_body": "Imaginary Testing Society",
                "date": "2017-10-30",
            },
        }
        rows = _render_honors([entry])
        assert rows == [["Fictional Excellence Award",
                         "Imaginary Testing Society", "2017"]], cell


def test_month_year_and_bare_year_date_cells_remain_one_record():
    """The two shapes `_DATE_COLUMN_RE` already accepted on dev still
    collapse a 3-column row to one part after the widening -- the fix adds
    shapes, it does not narrow the existing ones."""
    for cell in ("10/2017", "2017"):
        text = ("Fictional Excellence Award | Imaginary Testing Society | "
                f"{cell}")
        assert _entry_parts(
            text, ("Imaginary Testing Society", "2017")) == [text], cell


def test_two_genuinely_separate_full_date_awards_still_split():
    """Negative control (MUST NOT in the ticket): the widened date-column
    check must not swallow a real two-award list -- each award still carries
    its own full-date cell and `_MIN_AWARDS_FOR_SPLIT` is untouched."""
    text = ("Fictional Award One | 10/30/2017 | "
            "Fictional Award Two | 5/1/2018")
    assert _entry_parts(text, ("", "")) == [
        "Fictional Award One", "10/30/2017",
        "Fictional Award Two", "5/1/2018"]


def test_claimed_date_that_differs_from_the_cell_is_still_a_date_column():
    """Choice stated for the report: whether a cell counts as the date
    column is decided by its SHAPE (`_is_date_column`), not by whether it
    matches the extracted `date` -- '10/30/2017' is a date column even when
    stage 4 extracted a different date (2016-01-01) for this entry, so the
    row still renders as one record rather than being second-guessed against
    a mismatched extracted value."""
    entry = {
        "text": "Fictional Excellence Award | Imaginary Testing Society | 10/30/2017",
        "extracted_fields": {
            "award_name": "Fictional Excellence Award",
            "granting_body": "Imaginary Testing Society",
            "date": "2016-01-01",
        },
    }
    rows = _render_honors([entry])
    assert rows == [["Fictional Excellence Award", "Imaginary Testing Society",
                     "2016"]]


# --- two-digit-year date cells (#828 review round 2, F-R2-2) ----------------

def test_two_digit_year_date_cell_is_recognised():
    """A `mm/dd/yy` cell (the incident's fourth H entry, redacted here) is
    also a date column, same as the four-digit-year shapes above."""
    for cell in ("10/30/17", "1/5/17", "01/05/17"):
        assert _is_date_column(cell), cell


def test_two_digit_year_date_cell_requires_a_plausible_month():
    """A bare `dd/dd` (two components, no year at all) is not a date --
    `12/34` never had a third slash-separated component to be a year -- and
    neither is a three-component string whose first number is not a
    plausible month (`13/45/17`): the two-digit-year alternative requires
    its first component to be 1-12 before it matches at all, since a bare
    two-digit year is otherwise ambiguous with almost any other slash-joined
    pair of small numbers."""
    for cell in ("12/34", "13/45/17"):
        assert not _is_date_column(cell), cell


def test_two_digit_year_full_date_cell_is_not_read_as_a_second_award():
    """The two-digit-year sibling of `test_full_date_cell_is_not_read_as_a_
    second_award`: a 3-column row whose date cell is `mm/dd/yy` still
    collapses to one record. The rendered date column is stage 4's own
    extracted (ISO, full-year) `date` field, same as the mm/dd/yyyy case
    above -- what changes here is only that `_is_date_column('10/30/17')`
    now recognises the raw cell too, so the row is not split into three."""
    entry = {
        "text": "Fictional Legacy Award | Imaginary Testing Society | 10/30/17",
        "extracted_fields": {
            "award_name": "Fictional Legacy Award",
            "granting_body": "Imaginary Testing Society",
            "date": "2017-10-30",
        },
    }
    rows = _render_honors([entry])
    assert rows == [["Fictional Legacy Award", "Imaginary Testing Society",
                     "2017"]]


# --- item 3: the first cell's second paragraph is a continuation ------------
#
# Round 2's implementation targeted the wrong newline position (verifier
# finding F-R2-1): the docx extractor joins a CELL's own paragraphs with '\n'
# *inside* the cell, then joins the ROW's cells with ' | ' (see
# `core/docx_structure_extractor.py`), so a parenthetical second paragraph
# under the award name arrives as line 0 = the plain award title (no '|' of
# its own) and line 1 = "(description) | Organization | Date" -- not the
# reverse. Every fixture below uses that real shape.

def test_second_paragraph_in_the_award_cell_joins_the_award_name():
    """A source table cell's second paragraph -- the parenthetical under the
    award name -- arrives as its own newline-separated line ahead of the
    row's organization and date cells, and reads like a second award. When
    the remainder line already has an organization and a date cell, it is
    joined into the award name instead of splitting (#828 review round 3)."""
    text = ("Fictional Team Award\n"
            "(Co-recipient team award) | Imaginary Testing Society | 10/30/2017")
    rows = _render_honors([_raw(text)])
    assert rows == [["Fictional Team Award (Co-recipient team award)",
                     "Imaginary Testing Society", "2017"]]


def test_continuation_merge_does_not_apply_to_two_real_awards():
    """The guard's negative path: two genuinely separate multi-line awards,
    neither naming its own organization and date, must still split --
    unaffected by the merge added for #828."""
    text = "Fictional Award One\nFictional Award Two"
    assert _entry_parts(text) == ["Fictional Award One", "Fictional Award Two"]


def test_continuation_merge_requires_both_organization_and_date():
    """A remainder line with only an organization (no date) is not the
    3-column shape the merge is scoped to, so the two lines are still read
    as a genuinely separate fused-list pair."""
    text = ("Fictional Award\n"
            "(continuation) | Imaginary Testing Society")
    assert _entry_parts(text) == [
        "Fictional Award", "(continuation) | Imaginary Testing Society"]


# --- one negative test per `_merge_first_cell_continuation` guard ------------
# Each names the guard it targets. Guard order and structure deliberately
# mirror the pre-R3 code with only the two lines' roles swapped -- see the
# function's own docstring for which of these are independently killable by
# their own test versus provably redundant with the `columns is None`
# fallback (disclosed, same pattern as the pre-existing equivalent-guard
# note this file already carries for `_honor_columns`'s own '|' check).

def test_continuation_merge_requires_an_organization():
    """A remainder line with only a date (no organization) is not the
    3-column shape the merge is scoped to -- mirrors the date-only guard
    above but for the organization half of `_honor_columns`."""
    lines = ["Fictional Award", "(continuation) | 10/30/2017"]
    assert _merge_first_cell_continuation(lines) == lines


def test_continuation_merge_rejects_a_bare_year_remainder_line():
    """A remainder line that is nothing but a year is a loose year cell for
    a genuinely separate award, not a continuation row. (This guard is
    provably redundant with the `columns is None` fallback below it -- a
    bare year has no '|' of its own and can never satisfy `_honor_columns`
    either way -- kept for the same reason the pre-existing `_honor_columns`
    '|' check is kept: harmless, and it fails closed the same way if the
    fallback's shape ever changes.)"""
    lines = ["Fictional Award", "2018"]
    assert _merge_first_cell_continuation(lines) == lines


def test_continuation_merge_rejects_a_header_row_remainder_line():
    """A remainder line that is itself the source table's fused header row
    (its own org and date cells, but two header keywords) is not a
    continuation -- unlike the bare-year guard above, this one is NOT
    redundant with `columns is None`: "Name of Award | Organization |
    10/30/2017" parses into a full org+date record on its own and would
    otherwise merge as if it were real data."""
    lines = ["Fictional Award",
             "Name of Award | Organization | 10/30/2017"]
    assert _merge_first_cell_continuation(lines) == lines


def test_continuation_merge_rejects_a_first_line_with_its_own_pipe():
    """A first line already carrying its own '|' is a genuinely separate,
    already-complete 3-column award row, not a plain continuation-target
    title -- folding the next line into it would swallow a whole second
    award."""
    lines = ["Fictional Award One | Imaginary Testing Society | 10/30/2017",
             "Fictional Award Two | Imaginary Testing Society Two | 5/1/2018"]
    assert _merge_first_cell_continuation(lines) == lines


def test_continuation_merge_rejects_a_first_line_with_a_tab():
    """A first line carrying a tab is column-structured on its own (a
    tab-joined award/date/note triple), not the plain continuation-target
    title the merge is scoped to."""
    lines = ["Fictional Award\tSome Note",
             "(continuation) | Imaginary Testing Society | 10/30/2017"]
    assert _merge_first_cell_continuation(lines) == lines


def test_continuation_merge_rejects_a_remainder_line_with_a_tab():
    """A remainder line carrying a tab is column-structured on its own, not
    a clean '|'-joined row the continuation merge is scoped to."""
    lines = ["Fictional Award",
             "(continuation)\t| Imaginary Testing Society | 10/30/2017"]
    assert _merge_first_cell_continuation(lines) == lines


def test_continuation_merge_does_not_apply_past_exactly_two_lines():
    """Three or more lines is not the narrow two-line shape the merge is
    scoped to, whatever the first two look like -- returned unchanged."""
    lines = ["Fictional Award",
             "(continuation) | Imaginary Testing Society | 10/30/2017",
             "A third line"]
    assert _merge_first_cell_continuation(lines) == lines


# --- proof: the ticket's four synthetic rows, one record each (#828 R3) -----

def test_four_synthetic_incident_shaped_rows_yield_four_records():
    """Four synthetic entries built in the shape of the incident's own four
    H (taxonomy-coded) entries found during round-2 verification (no corpus
    content below) each collapse to exactly one record, for a total of four
    -- the fix's actual close of the incident's row-4 symptom, not merely
    the widened regex's 4-to-7 partial credit."""
    entries = [
        # fields[98]-shaped: parenthetical continuation + mm/dd/yyyy date.
        _raw("Fictional Team Award\n"
             "(Co-recipient team award) | Imaginary Testing Society | 10/30/2017"),
        # fields[99]/[100]-shaped: plain 3-column row, full mm/dd/yyyy date.
        _raw("Fictional Excellence Award | Imaginary Testing Society | 10/30/2017"),
        {
            "text": "Fictional Merit Award | Imaginary Research Council | 5/1/2018",
            "extracted_fields": {
                "award_name": "Fictional Merit Award",
                "granting_body": "Imaginary Research Council",
                "date": "2018-05-01",
            },
        },
        # fields[184]-shaped: plain 3-column row, two-digit-year date.
        _raw("Fictional Legacy Award | Imaginary Testing Society | 10/30/17"),
    ]
    records = [r for entry in entries for r in parse_honor_entry(entry)]
    assert len(records) == 4, records


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-q"]))
