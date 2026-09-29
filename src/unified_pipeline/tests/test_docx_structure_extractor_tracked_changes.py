"""Guard: the docx readers must not drop w:ins/w:smartTag/w:sdt-wrapped runs
mid-word (#557).

python-docx's Paragraph.text / _Cell.text only concatenate w:r elements that
are DIRECT CHILDREN of w:p. When Word splits a word across runs -- routine
for tracked-change insertions -- a run nested inside w:ins is silently
skipped, dropping 1-2 characters mid-word ("Down Syndrome" reads back as
"Down yndrome"). get_paragraph_text/get_cell_text walk the paragraph's XML
directly and must not reproduce that loss.

    python3 -m pytest src/unified_pipeline/tests/test_docx_structure_extractor_tracked_changes.py -p no:cacheprovider

Self-contained: no DB, no template. Requires only python-docx.
"""

import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402

from unified_pipeline.core.docx_structure_extractor import (  # noqa: E402
    extract_table_metadata,
    extract_unified_elements,
    get_cell_text,
    get_paragraph_text,
    get_table_first_cell_text,
)


def _splice_tracked_insertion(second_run, inserted_text):
    """Insert a w:ins-wrapped run containing `inserted_text` immediately before
    `second_run`'s underlying <w:r> -- reproducing the shape Word produces when
    a tracked-change edit splits a word across runs."""
    ins = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:t>{inserted_text}</w:t></w:r></w:ins>'
    )
    second_run._r.addprevious(ins)


def _down_syndrome_paragraph(add_run_to):
    """Build 'Down ' + <w:ins>S</w:ins> + 'yndrome' on whatever python-docx
    object exposes .add_run() (a Paragraph or a cell's first paragraph)."""
    add_run_to.add_run("Down ")
    tail_run = add_run_to.add_run("yndrome")
    _splice_tracked_insertion(tail_run, "S")


def test_get_paragraph_text_reads_through_tracked_insertion():
    doc = Document()
    para = doc.add_paragraph()
    _down_syndrome_paragraph(para)

    assert get_paragraph_text(para) == "Down Syndrome"


def test_get_cell_text_reads_through_tracked_insertion():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _down_syndrome_paragraph(cell.paragraphs[0])

    assert get_cell_text(cell) == "Down Syndrome"


def test_extract_table_metadata_does_not_mangle_tracked_insertion(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _down_syndrome_paragraph(cell.paragraphs[0])

    metadata = extract_table_metadata(table, idx="table_0")
    cell_text = metadata["data"][0][0]["text"]
    assert cell_text == "Down Syndrome"
    assert "yndrome" not in cell_text.replace("Syndrome", "")


def test_get_table_first_cell_text_does_not_mangle_tracked_insertion():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _down_syndrome_paragraph(cell.paragraphs[0])

    assert get_table_first_cell_text(table) == "Down Syndrome"


def test_get_paragraph_text_tab_char_parameterized():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("a")
    para.add_run().add_tab()
    para.add_run("b")

    assert get_paragraph_text(para) == "a b"
    assert get_paragraph_text(para, tab_char='\t') == "a\tb"


def _splice_wrapped_run(anchor_run, wrapped_xml):
    """Insert an arbitrary wrapper element containing its own <w:r><w:t>
    immediately before `anchor_run`'s underlying <w:r> -- same mechanism as
    _splice_tracked_insertion, generalized for smartTag/sdt/hyperlink."""
    anchor_run._r.addprevious(parse_xml(wrapped_xml))


def test_get_paragraph_text_reads_through_smarttag():
    # Word's legacy "smart tag" auto-recognition (addresses, names, dates)
    # wraps a run without it being a direct <w:p> child -- same drop risk as
    # w:ins, per get_paragraph_text's own docstring.
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("before ")
    after = para.add_run("after")
    _splice_wrapped_run(
        after,
        f'<w:smartTag {nsdecls("w")} w:uri="urn:schemas-microsoft-com:office:smarttags" '
        f'w:element="place"><w:r><w:t>WRAPPED </w:t></w:r></w:smartTag>',
    )

    assert get_paragraph_text(para) == "before WRAPPED after"


def test_get_paragraph_text_reads_through_sdt():
    # Structured document tags (content controls) nest their run inside
    # <w:sdt><w:sdtContent>, not as a direct <w:p> child.
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("before ")
    after = para.add_run("after")
    _splice_wrapped_run(
        after,
        f'<w:sdt {nsdecls("w")}><w:sdtContent>'
        f'<w:r><w:t>WRAPPED </w:t></w:r></w:sdtContent></w:sdt>',
    )

    assert get_paragraph_text(para) == "before WRAPPED after"


def test_get_paragraph_text_reads_through_hyperlink():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("before ")
    after = para.add_run("after")
    _splice_wrapped_run(
        after,
        f'<w:hyperlink {nsdecls("w")}><w:r><w:t>WRAPPED </w:t></w:r></w:hyperlink>',
    )

    assert get_paragraph_text(para) == "before WRAPPED after"


def test_get_paragraph_text_converts_br_and_cr_to_newline():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("line1").add_break(WD_BREAK.LINE)  # <w:br/>
    cr_run = para.add_run("")
    cr_run._r.append(parse_xml(f'<w:cr {nsdecls("w")}/>'))
    para.add_run("line2")

    assert get_paragraph_text(para) == "line1\n\nline2"


def test_extract_unified_elements_reads_through_tracked_insertion(tmp_path):
    # Integration regression: the fix must reach extract_unified_elements()
    # itself -- the actual stage 1b/2 production entry point -- not just the
    # lower-level helpers exercised by the tests above.
    doc = Document()
    para = doc.add_paragraph()
    _down_syndrome_paragraph(para)
    docx_path = tmp_path / "tracked_change.docx"
    doc.save(str(docx_path))

    result = extract_unified_elements(str(docx_path))

    texts = [el["text"] for el in result["elements"] if el.get("type") == "paragraph"]
    assert "Down Syndrome" in texts
