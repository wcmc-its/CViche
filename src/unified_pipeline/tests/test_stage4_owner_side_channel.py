"""#456: `extract_cv_owner_name`'s side-channel fallback tier.

The tier consults `extract_owner_side_channel` (sdt/header/footer text the
main body walk never sees) only when the body-derived pass left BOTH
`last_name` and `full_name` empty and a real `docx_path` was given -- and
only before `fallback_from_uid`, never overriding a body-derived name.

    python3 -m pytest src/unified_pipeline/tests/test_stage4_owner_side_channel.py -p no:cacheprovider

Self-contained: `call_llm` is always stubbed at the module attribute -- no
Bedrock/OpenAI. `extract_owner_side_channel` is stubbed too, except in the
corrupt-file test and the #1655 tests, which read a synthetic .docx built
in-test (`letterhead_docx`). Synthetic names only.
"""

import io
import json
import logging
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import unified_pipeline.stage4.extraction as extraction  # noqa: E402
import unified_pipeline.stage4.owner_name as owner_name  # noqa: E402
import unified_pipeline.stage_4_field_extractor as stage_4_field_extractor  # noqa: E402


def _llm_result(content, total_tokens=10, cost=0.001):
    return {"content": content, "total_tokens": total_tokens, "cost": cost}


def _name_reply(**fields):
    base = {
        "first_name": "", "middle_name": "", "last_name": "",
        "suffix": "", "full_name": "", "full_name_with_credentials": "",
    }
    base.update(fields)
    return json.dumps(base)


def _touch(tmp_path, name="cv.docx"):
    """A file that exists -- content is irrelevant, `extract_owner_side_channel`
    is always stubbed in these tests, so nothing ever actually parses it."""
    p = tmp_path / name
    p.write_bytes(b"not a real docx; only Path.is_file() is checked")
    return str(p)


# ---------------------------------------------------------------------------
# Body tier succeeds -> side channel never consulted
# ---------------------------------------------------------------------------

def test_body_tier_success_never_consults_side_channel(monkeypatch, tmp_path):
    calls = {"call_llm": 0, "side_channel": 0}

    def fake_call_llm(**kwargs):
        calls["call_llm"] += 1
        return _llm_result(_name_reply(first_name="Jane", last_name="Public", full_name="Jane Public"))

    def fake_side_channel(path):
        calls["side_channel"] += 1
        return {"sdt_lines": [], "header_lines": [], "footer_lines": []}

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(owner_name, "extract_owner_side_channel", fake_side_channel)

    result = owner_name.extract_cv_owner_name(
        "2024_Test_CV", [{"text": "Jane Public, MD, Professor of Surgery"}],
        docx_path=_touch(tmp_path),
    )

    assert result["last_name"] == "Public"
    assert calls["call_llm"] == 1
    assert calls["side_channel"] == 0, "a body-derived name must never trigger the side channel"


# ---------------------------------------------------------------------------
# Body tier empty + side channel has lines -> prompt runs over them, name used
# ---------------------------------------------------------------------------

def test_side_channel_used_when_body_tier_finds_nothing(monkeypatch, tmp_path):
    prompts = []

    def fake_call_llm(**kwargs):
        prompts.append(kwargs["messages"][-1]["content"])
        if len(prompts) == 1:
            return _llm_result(_name_reply())  # body tier: nothing found
        return _llm_result(_name_reply(first_name="Sam", last_name="Rivera", full_name="Sam Rivera"))

    def fake_side_channel(path):
        return {"sdt_lines": ["Sam Rivera, PhD"], "header_lines": [], "footer_lines": []}

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(owner_name, "extract_owner_side_channel", fake_side_channel)

    result = owner_name.extract_cv_owner_name(
        "2024_Test_CV", [{"text": "a narrative paragraph with no name in it at all"}],
        docx_path=_touch(tmp_path),
    )

    assert result["last_name"] == "Rivera"
    assert len(prompts) == 2
    assert "Sam Rivera, PhD" in prompts[1]


# ---------------------------------------------------------------------------
# Body tier empty + side channel empty -> fallback_from_uid, unchanged
# ---------------------------------------------------------------------------

def test_side_channel_empty_falls_through_to_uid_fallback_unchanged(monkeypatch, tmp_path):
    def fake_call_llm(**kwargs):
        return _llm_result(_name_reply())

    def fake_side_channel(path):
        return {"sdt_lines": [], "header_lines": [], "footer_lines": []}

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(owner_name, "extract_owner_side_channel", fake_side_channel)

    with_docx = owner_name.extract_cv_owner_name(
        "2024_Rivera_CV", [{"text": "narrative body text"}], docx_path=_touch(tmp_path),
    )
    without_docx = owner_name.extract_cv_owner_name(
        "2024_Rivera_CV", [{"text": "narrative body text"}],
    )

    assert with_docx == without_docx == {
        "first_name": "", "middle_name": "", "last_name": "Rivera",
        "suffix": "", "full_name": "", "full_name_with_credentials": "",
    }


# ---------------------------------------------------------------------------
# docx_path=None -> identical to pre-#456 behaviour
# ---------------------------------------------------------------------------

def test_docx_path_none_matches_pre_456_behavior(monkeypatch):
    def fake_call_llm(**kwargs):
        return _llm_result(_name_reply())

    side_channel_calls = []
    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: side_channel_calls.append(p) or {"sdt_lines": [], "header_lines": [], "footer_lines": []},
    )

    result = owner_name.extract_cv_owner_name("2024_Doe_CV", [{"text": "narrative text"}])

    assert result["last_name"] == "Doe"
    assert side_channel_calls == []


# ---------------------------------------------------------------------------
# docx_path to a missing file -> no raise, falls through
# ---------------------------------------------------------------------------

def test_missing_docx_path_falls_through_without_raising(monkeypatch, tmp_path):
    def fake_call_llm(**kwargs):
        return _llm_result(_name_reply())

    side_channel_calls = []
    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: side_channel_calls.append(p) or {"sdt_lines": ["should never be read"], "header_lines": [], "footer_lines": []},
    )

    missing_path = str(tmp_path / "does_not_exist.docx")

    result = owner_name.extract_cv_owner_name(
        "2024_Doe_CV", [{"text": "narrative text"}], docx_path=missing_path,
    )

    assert result["last_name"] == "Doe"
    assert side_channel_calls == [], "extract_owner_side_channel must not be called for a nonexistent path"


# ---------------------------------------------------------------------------
# The `:104` no-entries path also consults the side channel
# ---------------------------------------------------------------------------

def test_no_entries_path_consults_side_channel(monkeypatch, tmp_path):
    """mapped_entries=[] hits the no-first_entries branch directly -- no
    body-tier LLM call at all -- and the side channel must still run before
    the uid fallback."""
    prompts = []

    def fake_call_llm(**kwargs):
        prompts.append(kwargs["messages"][-1]["content"])
        return _llm_result(_name_reply(first_name="Alex", last_name="Nguyen", full_name="Alex Nguyen"))

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: {"sdt_lines": ["Alex Nguyen, MD"], "header_lines": [], "footer_lines": []},
    )

    result = owner_name.extract_cv_owner_name("web999", [], docx_path=_touch(tmp_path))

    assert result["last_name"] == "Nguyen"
    assert len(prompts) == 1, "only the side-channel prompt should run -- no body-tier call with zero entries"
    assert "Alex Nguyen, MD" in prompts[0]


def test_no_entries_path_with_no_side_channel_content_uses_uid_fallback(monkeypatch, tmp_path):
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: {"sdt_lines": [], "header_lines": [], "footer_lines": []},
    )
    called_llm = []
    monkeypatch.setattr(owner_name, "call_llm", lambda **kwargs: called_llm.append(1))

    result = owner_name.extract_cv_owner_name("web999", [], docx_path=_touch(tmp_path))

    assert result["last_name"] == ""  # 'web999' is opaque -- #457's contract
    assert called_llm == [], "no side-channel lines means no LLM call at all"


# ---------------------------------------------------------------------------
# Line cap and char cap enforced
# ---------------------------------------------------------------------------

def test_side_channel_line_and_char_caps_are_enforced(monkeypatch, tmp_path):
    long_line = "X" * 500
    # Fixed-width, zero-padded markers so no marker is a substring of another
    # (unlike "line 1" inside "line 10") -- keeps the membership checks below exact.
    many_lines = [f"UNIQLINE{i:03d}TAG" for i in range(50)]

    prompts = []

    def fake_call_llm(**kwargs):
        prompts.append(kwargs["messages"][-1]["content"])
        return _llm_result(_name_reply())

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: {"sdt_lines": [long_line] + many_lines, "header_lines": [], "footer_lines": []},
    )

    owner_name.extract_cv_owner_name("web998", [], docx_path=_touch(tmp_path))

    assert len(prompts) == 1
    prompt_text = prompts[0]

    # The long line is truncated to OWNER_SIDE_CHANNEL_MAX_CHARS, not passed whole.
    assert "X" * owner_name.OWNER_SIDE_CHANNEL_MAX_CHARS in prompt_text
    assert "X" * (owner_name.OWNER_SIDE_CHANNEL_MAX_CHARS + 1) not in prompt_text

    # Only OWNER_SIDE_CHANNEL_MAX_LINES lines total make it into the prompt --
    # the long_line occupies one slot, leaving MAX_LINES - 1 of many_lines.
    included = sum(1 for marker in many_lines if marker in prompt_text)
    assert included == owner_name.OWNER_SIDE_CHANNEL_MAX_LINES - 1


# ---------------------------------------------------------------------------
# #456-R2 F4: a corrupt-but-existing docx_path must not raise -- it must log
# a warning and fall through to fallback_from_uid, matching the docstring's
# own claim (which was previously false: a real non-zip file propagated
# PackageNotFoundError straight out of the tier and failed the web driver).
# ---------------------------------------------------------------------------

def test_corrupt_docx_path_falls_through_without_raising(monkeypatch, tmp_path, caplog):
    def fake_call_llm(**kwargs):
        return _llm_result(_name_reply())

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    # extract_owner_side_channel is NOT stubbed here -- this test exercises
    # the real function against a genuinely corrupt file (an ASCII file with
    # a .docx extension), so Document(docx_path) really raises
    # PackageNotFoundError and the tier's own try/except must catch it.

    corrupt_path = tmp_path / "corrupt.docx"
    corrupt_path.write_text("not a docx at all, just ascii text")

    with caplog.at_level("WARNING", logger="unified_pipeline.stage4.owner_name"):
        result = owner_name.extract_cv_owner_name(
            "2024_Rivera_CV", [{"text": "narrative body text"}], docx_path=str(corrupt_path),
        )

    assert result["last_name"] == "Rivera"  # fallback_from_uid, unchanged
    warnings = [r.message for r in caplog.records if "owner side channel unreadable" in r.message]
    assert len(warnings) == 1
    assert warnings[0].startswith("2024_Rivera_CV: owner side channel unreadable:")


# ---------------------------------------------------------------------------
# #456-R2 F3, mutant m5: header/sdt/footer lines must ALL reach the prompt,
# in that order (header first since #1655) -- not just whichever channel
# happened to be tested alone above. `combined = channel['sdt_lines']` (dropping header/footer) survived
# the full suite before this test existed.
# ---------------------------------------------------------------------------

def test_side_channel_combines_header_sdt_and_footer_lines_in_order(monkeypatch, tmp_path):
    prompts = []

    def fake_call_llm(**kwargs):
        prompts.append(kwargs["messages"][-1]["content"])
        return _llm_result(_name_reply())

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: {
            "sdt_lines": ["SDTMARK Owner"],
            "header_lines": ["HDRMARK Owner"],
            "footer_lines": ["FTRMARK Owner"],
        },
    )

    owner_name.extract_cv_owner_name("web997", [], docx_path=_touch(tmp_path))

    assert len(prompts) == 1
    prompt_text = prompts[0]
    assert "SDTMARK Owner" in prompt_text
    assert "HDRMARK Owner" in prompt_text
    assert "FTRMARK Owner" in prompt_text
    # header -> sdt -> footer is the contract order (#1655,
    # `_owner_side_channel_content_lines`).
    assert (
        prompt_text.index("HDRMARK Owner")
        < prompt_text.index("SDTMARK Owner")
        < prompt_text.index("FTRMARK Owner")
    )


# ---------------------------------------------------------------------------
# #456-R2 F3, mutants m7a/m7b: `docx_path` must reach `extract_cv_owner_name`
# across BOTH hops of the wire -- process_cv -> extract_fields_from_mapped_
# entries (m7a) and extract_fields_from_mapped_entries -> extract_cv_owner_
# name (m7b). Either kwarg silently dropped left the full 3257-test suite
# green before these two tests existed.
# ---------------------------------------------------------------------------

def test_extraction_passes_docx_path_to_extract_cv_owner_name(monkeypatch):
    """m7b: extract_fields_from_mapped_entries -> extract_cv_owner_name."""
    captured = {}

    def fake_owner_name(document_uid, mapped_entries, docx_path=None, usage=None):
        captured["docx_path"] = docx_path
        return {
            "first_name": "", "middle_name": "", "last_name": "",
            "suffix": "", "full_name": "", "full_name_with_credentials": "",
        }

    monkeypatch.setattr(extraction, "extract_cv_owner_name", fake_owner_name)
    monkeypatch.setattr(extraction, "extract_fields_batch", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("no entries -- extract_fields_batch must not be called")
    ))

    extraction.extract_fields_from_mapped_entries(
        [], document_uid="web996", docx_path="/synthetic/path/web996.docx",
    )

    assert captured["docx_path"] == "/synthetic/path/web996.docx"


def test_process_cv_passes_its_docx_path_to_extract_fields_from_mapped_entries(monkeypatch, tmp_path):
    """m7a: process_cv -> extract_fields_from_mapped_entries.

    The stage-3b JSON lookup is stubbed (a synthetic, empty entries payload)
    so this test needs no real corpus file: `Path.exists` is patched to
    report True only for the one expected stage-3b path (delegating to the
    real implementation for every other path, tmp_path's own files
    included), and this module's bare `open` name is shadowed (module-level
    only -- never touches the `open` builtin other modules see) to hand back
    that synthetic payload for that one path.
    """
    cv_path = str(tmp_path / "synthetic_cv.docx")
    Path(cv_path).write_bytes(b"not a real docx; process_cv never opens this file itself")
    document_uid = Path(cv_path).stem

    expected_stage3b = (
        Path(stage_4_field_extractor.__file__).parent
        / "outputs" / "stage_3b_classified_entries" / f"{document_uid}_classified.json"
    )
    expected_output_dir = Path(stage_4_field_extractor.__file__).parent / "outputs" / "stage_4_field_extraction"
    expected_output_path = expected_output_dir / f"{document_uid}_fields.json"

    real_exists = Path.exists
    real_mkdir = Path.mkdir

    def fake_exists(self):
        if self == expected_stage3b:
            return True
        return real_exists(self)

    def fake_mkdir(self, *args, **kwargs):
        # process_cv unconditionally mkdir's its real output dir -- never
        # actually touch the repo's real outputs/ tree from a test.
        if self == expected_output_dir:
            return None
        return real_mkdir(self, *args, **kwargs)

    def fake_open(path, mode="r", *args, **kwargs):
        p = Path(path)
        if p == expected_stage3b:
            return io.StringIO(json.dumps({"entries": []}))
        if p == expected_output_path:
            return io.StringIO()  # process_cv's own output write, discarded
        raise AssertionError(f"unexpected open() inside process_cv test: {path}")

    captured = {}

    def fake_extract(mapped_entries, **kwargs):
        captured["docx_path"] = kwargs.get("docx_path")
        return {
            "entries": [], "cv_owner": {}, "cv_owner_location": {},
            "total_cost": 0.0, "total_tokens": 0,
            "stats": {"extracted": 0, "skipped": 0},
        }

    monkeypatch.setattr(Path, "exists", fake_exists)
    monkeypatch.setattr(Path, "mkdir", fake_mkdir)
    monkeypatch.setattr(stage_4_field_extractor, "open", fake_open, raising=False)
    monkeypatch.setattr(stage_4_field_extractor, "extract_fields_from_mapped_entries", fake_extract)

    stage_4_field_extractor.process_cv(cv_path)

    assert captured["docx_path"] == cv_path


def test_process_cv_stamps_over_the_unfiltered_3b_list_before_dropping_fragments(monkeypatch, tmp_path):
    """#985 B1: a sub-heading 3b flagged is_fragment is filtered out before
    extraction, so the stamp must run on the FULL list, where that fragment
    still ends the previous heading's run."""
    document_uid = "synthetic_cv_985"
    stage3b = Path(stage_4_field_extractor.__file__).parent / "outputs" / "stage_3b_classified_entries" / f"{document_uid}_classified.json"
    hier = ["Service"]
    entries = [
        {"text": "Alpha University:", "taxonomy_code": "T", "hierarchy": hier},
        {"text": "Member, Committee X", "taxonomy_code": "P", "hierarchy": hier},
        {"text": "Beta University:", "taxonomy_code": "O", "hierarchy": hier, "is_fragment": True},
        {"text": "Chair, Committee Y", "taxonomy_code": "P", "hierarchy": hier},
    ]
    real_exists = Path.exists
    monkeypatch.setattr(Path, "exists", lambda self: True if self == stage3b else real_exists(self))
    monkeypatch.setattr(Path, "mkdir", lambda self, *a, **k: None)
    monkeypatch.setattr(stage_4_field_extractor, "open", lambda path, mode="r", *a, **k: (
        io.StringIO(json.dumps({"entries": entries})) if Path(path) == stage3b else io.StringIO()), raising=False)
    captured = {}

    def fake_extract(mapped_entries, **kwargs):
        captured["stamps"] = [(e["text"], e.get("context_heading")) for e in mapped_entries]
        return {"entries": [], "cv_owner": {}, "cv_owner_location": {}, "total_cost": 0.0, "total_tokens": 0, "stats": {"extracted": 0, "skipped": 0}}

    monkeypatch.setattr(stage_4_field_extractor, "extract_fields_from_mapped_entries", fake_extract)
    stage_4_field_extractor.process_cv(str(tmp_path / f"{document_uid}.docx"))
    assert captured["stamps"] == [("Alpha University:", None), ("Member, Committee X", "Alpha University"), ("Chair, Committee Y", None)]


# ---------------------------------------------------------------------------
# #1655: a letterhead that lives only in the first-page header. Real .docx
# built here with python-docx -- a body-level w:sdt of 25 paragraphs and a
# first-page header carrying a synthetic owner's name and contact block.
# ---------------------------------------------------------------------------

LETTERHEAD = [
    "Quinn Synthetic, MD, PhD",
    "Department of Imaginary Medicine",
    "100 Example Avenue, Room 5, Testville, NY 10000",
    "Phone: (212) 555-0100  Fax: (212) 555-0101",
    "Email: quinn.synthetic@example.org",
]
SDT_LINES = [f"Licence course LINE{i:02d}" for i in range(25)]


def letterhead_docx(tmp_path, header=LETTERHEAD, *, default_header=(), footer=(), sdt=SDT_LINES):
    """A docx whose body opens with "Biography" and holds a body-level w:sdt
    of `sdt` paragraphs, with `header` in the first-page header
    (w:titlePg), `default_header` in the default header and `footer` in the
    default footer. Imported by test_stage4_extraction.py's #1655 wire test."""
    doc = Document()
    doc.add_paragraph("Biography")
    anchor = doc.add_paragraph("A synthetic narrative paragraph about research interests.")
    if sdt:
        ps = "".join(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in sdt)
        anchor._p.addprevious(parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtContent>{ps}</w:sdtContent></w:sdt>'))
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    for container, texts in ((section.first_page_header, header), (section.header, default_header),
                             (section.footer, footer)):
        if texts:
            container.paragraphs[0].text = texts[0]
            for text in texts[1:]:
                container.add_paragraph(text)
    path = tmp_path / "letterhead.docx"
    doc.save(str(path))
    return str(path)


def test_header_name_reaches_the_prompt_past_a_25_paragraph_sdt(monkeypatch, tmp_path):
    """#1655 fix 1 (RSFOYB): the 25 sdt lines used to fill the whole
    OWNER_SIDE_CHANNEL_MAX_LINES budget, so the header never reached the
    prompt and `cv_owner` came back empty. The real extract_owner_side_channel
    runs here; only call_llm is stubbed."""
    prompts = []

    def fake_call_llm(**kwargs):
        prompt = kwargs["messages"][-1]["content"]
        prompts.append(prompt)
        if "Quinn Synthetic" in prompt:
            return _llm_result(_name_reply(first_name="Quinn", last_name="Synthetic", full_name="Quinn Synthetic"))
        return _llm_result(_name_reply())

    monkeypatch.setattr(owner_name, "call_llm", fake_call_llm)

    result = owner_name.extract_cv_owner_name(
        "web990", [{"text": "Biography"}], docx_path=letterhead_docx(tmp_path))

    assert result["last_name"] == "Synthetic"
    assert len(prompts) == 2  # body tier, then side channel
    assert "Quinn Synthetic, MD, PhD" in prompts[1]


def test_side_channel_names_header_as_first_channel_when_both_exist(monkeypatch, tmp_path):
    """The tier's log line names the channel that leads the prompt."""
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: {"sdt_lines": ["S"], "header_lines": ["H"], "footer_lines": ["F"]})

    lines, channel = owner_name._owner_side_channel_content_lines("web990", _touch(tmp_path))

    assert (lines, channel) == (["H", "S", "F"], "header")


def test_side_channel_names_sdt_as_first_channel_without_a_header(monkeypatch, tmp_path):
    monkeypatch.setattr(
        owner_name, "extract_owner_side_channel",
        lambda p: {"sdt_lines": ["S"], "header_lines": [], "footer_lines": ["F"]})

    assert owner_name._owner_side_channel_content_lines("web990", _touch(tmp_path)) == (["S", "F"], "sdt")


def test_header_footer_contact_entry_holds_header_and_footer_lines_only(tmp_path):
    """#1655 fix 2: the synthetic A entry carries the letterhead and the
    footer, in that order, and none of the sdt content."""
    path = letterhead_docx(tmp_path, footer=["Updated by the synthetic owner"])

    entry = owner_name.header_footer_contact_entry("web990", path)

    assert entry == {
        "text": "\n".join([*LETTERHEAD, "Updated by the synthetic owner"]),
        "taxonomy_code": owner_name.PERSONAL_DATA_CODE,
        "element_idx_start": owner_name.HEADER_FOOTER_ELEMENT_IDX,
        "element_idx_end": owner_name.HEADER_FOOTER_ELEMENT_IDX,
        owner_name.OWNER_CONTACT_SOURCE_KEY: owner_name.OWNER_CONTACT_SOURCE_HEADER_FOOTER,
    }
    assert "LINE00" not in entry["text"]


def test_header_footer_contact_entry_sends_a_repeated_letterhead_once(tmp_path):
    """A first-page and a default header repeating one letterhead are two
    parts with the same lines; each line is sent once."""
    path = letterhead_docx(tmp_path, default_header=LETTERHEAD)

    entry = owner_name.header_footer_contact_entry("web990", path)

    assert entry is not None
    assert entry["text"].split("\n") == LETTERHEAD


def test_header_footer_contact_entry_none_for_a_header_without_contact(tmp_path):
    """A name and a running title carry no email or phone: no entry, so no
    extra LLM call and nothing for the Appendix."""
    path = letterhead_docx(tmp_path, header=["Quinn Synthetic, MD", "Curriculum Vitae"])

    assert owner_name.header_footer_contact_entry("web990", path) is None


def test_header_footer_contact_entry_none_for_sdt_contact_alone(tmp_path):
    """A content control's email is body content, not a letterhead."""
    path = letterhead_docx(tmp_path, header=(), sdt=["Email: quinn.synthetic@example.org"])

    assert owner_name.header_footer_contact_entry("web990", path) is None


def test_header_footer_contact_entry_none_without_a_docx(monkeypatch, tmp_path):
    """No path, or one that is not a file: nothing is opened at all."""
    opened = []
    monkeypatch.setattr(owner_name, "extract_owner_side_channel", opened.append)

    assert owner_name.header_footer_contact_entry("web990", None) is None
    assert owner_name.header_footer_contact_entry("web990", str(tmp_path / "missing.docx")) is None
    assert opened == []


def test_header_footer_contact_entry_none_for_a_corrupt_docx(tmp_path, caplog):
    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_text("not a docx at all")

    with caplog.at_level(logging.WARNING, logger="unified_pipeline.stage4.owner_name"):
        assert owner_name.header_footer_contact_entry("web990", str(corrupt)) is None

    assert any(r.message.startswith("web990: header/footer unreadable for owner contact:")
               for r in caplog.records)


def test_header_footer_contact_entry_caps_the_lines_sent(tmp_path):
    many = [f"Email: owner{i:02d}@example.org" for i in range(30)]
    path = letterhead_docx(tmp_path, header=many, sdt=())

    entry = owner_name.header_footer_contact_entry("web990", path)

    assert entry is not None
    assert entry["text"].split("\n") == many[:owner_name.OWNER_SIDE_CHANNEL_MAX_LINES]


def test_contact_signal_reads_an_email_or_a_phone_and_not_a_year_range():
    assert owner_name.OWNER_CONTACT_SIGNAL_RE.search("quinn@example.org")
    assert owner_name.OWNER_CONTACT_SIGNAL_RE.search("Tel 212-555-0100")
    assert owner_name.OWNER_CONTACT_SIGNAL_RE.search("(212) 555-0100")
    assert not owner_name.OWNER_CONTACT_SIGNAL_RE.search("Curriculum Vitae 2001-2005")
    assert not owner_name.OWNER_CONTACT_SIGNAL_RE.search("100 Example Avenue, NY 10000")
