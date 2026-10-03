"""#456: `extract_owner_side_channel` reads text the main body walk drops --
a body-level `w:sdt` (Word content-control) wrapping whole paragraphs, and
`word/header*.xml` / `word/footer*.xml` letterhead -- as a SEPARATE structure
stage 4's owner-name pass may consult, never merged into the element stream.

The landmine this file's last test guards against: the issue's own comment
warns that the "obvious fix" (descending into `w:sdt` and emitting its
paragraphs inline into `extract_unified_elements`/`extract_docx_structure`)
silently corrupts every CV, because `doc.paragraphs` -- which
`stage_2_entry_extraction.py:1118/1153/1226/1245` indexes as a fallback --
only enumerates direct-body `CT_P` children. `sdt_lines` must never change
what those two functions emit.

    python3 -m pytest src/unified_pipeline/tests/test_docx_owner_side_channel.py -p no:cacheprovider

Self-contained: no DB, no network, no LLM, no PII, synthetic names only.
"""

import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.enum.section import WD_HEADER_FOOTER, WD_SECTION  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402

from unified_pipeline.core.docx_structure_extractor import (  # noqa: E402
    extract_owner_side_channel,
    extract_unified_elements,
)


def _splice_body_level_sdt(anchor_para, texts: list[str]) -> None:
    """Insert a body-level <w:sdt><w:sdtContent> wrapping one <w:p> per text
    in `texts`, as a SIBLING immediately before `anchor_para`'s underlying
    <w:p> -- the actual #456 shape (a content control wrapping whole
    paragraphs), not the inline run-level sdt already covered by
    test_docx_structure_extractor_tracked_changes.py's
    test_get_paragraph_text_reads_through_sdt. Takes an existing paragraph as
    the anchor (rather than creating one) so callers control the document's
    exact paragraph count -- needed by the stream-identity test below."""
    ps = "".join(f'<w:p><w:r><w:t>{t}</w:t></w:r></w:p>' for t in texts)
    sdt_xml = f'<w:sdt {nsdecls("w")}><w:sdtContent>{ps}</w:sdtContent></w:sdt>'
    anchor_para._p.addprevious(parse_xml(sdt_xml))


def _splice_cell_level_sdt(doc: Document, text: str):
    """A one-cell table whose single paragraph is wrapped in a body-level
    (not inline-run) `w:sdt`, nested inside the cell -- the issue's own
    "check w:sdt at table and row level too" ask. Returns the table so a
    caller can also assert the main element stream is unaffected."""
    table = doc.add_table(rows=1, cols=1)
    cell_p = table.cell(0, 0).paragraphs[0]._p
    sdt_xml = f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    cell_p.addprevious(parse_xml(sdt_xml))
    return table


def _save(doc: Document, tmp_path: Path, name: str = "f.docx") -> str:
    docx_path = tmp_path / name
    doc.save(str(docx_path))
    return str(docx_path)


# ---------------------------------------------------------------------------
# sdt_lines
# ---------------------------------------------------------------------------

def test_body_level_sdt_paragraph_is_collected(tmp_path):
    doc = Document()
    doc.add_paragraph("before")
    after = doc.add_paragraph("after")
    _splice_body_level_sdt(after, ["Synthetic Owner Name"])

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["sdt_lines"] == ["Synthetic Owner Name"]
    assert channel["header_lines"] == []
    assert channel["footer_lines"] == []


def test_sdt_line_reads_a_non_breaking_hyphen_as_a_hyphen(tmp_path):
    # A hyphenated surname written with Word's <w:noBreakHyphen/>: no '-' in
    # any w:t, so a w:t-only walk would hand stage 4 the fused name.
    doc = Document()
    after = doc.add_paragraph("after")
    sdt_xml = (
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p>'
        '<w:r><w:t>Synthetic Testname</w:t></w:r>'
        '<w:r><w:noBreakHyphen/></w:r>'
        '<w:r><w:t>Example</w:t></w:r>'
        '</w:p></w:sdtContent></w:sdt>'
    )
    after._p.addprevious(parse_xml(sdt_xml))

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["sdt_lines"] == ["Synthetic Testname-Example"]


def test_sdt_nested_inside_a_table_cell_is_also_collected(tmp_path):
    doc = Document()
    doc.add_paragraph("cover text")
    _splice_cell_level_sdt(doc, "Cell Owner Name")

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert "Cell Owner Name" in channel["sdt_lines"]


def test_sdt_lines_are_deduped_and_in_document_order(tmp_path):
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    _splice_body_level_sdt(anchor, ["First Line", "Second Line"])

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["sdt_lines"] == ["First Line", "Second Line"]


def test_forty_paragraph_sdt_is_stable_across_repeated_runs(tmp_path):
    """#456 verifier finding 1: an `id()`-keyed dedup set over lxml element
    proxies drops lines nondeterministically (a 40-paragraph body-level sdt
    was observed dropping to 35). `body.iter(w:p)` already yields each node
    exactly once by construction, so no dedup set is needed; this pins that
    all 40 lines survive, identically, across repeated independent calls."""
    doc = Document()
    doc.add_paragraph("before")
    after = doc.add_paragraph("after")
    texts = [f"L{i:03d}" for i in range(40)]
    _splice_body_level_sdt(after, texts)
    path = _save(doc, tmp_path)

    for _ in range(3):
        channel = extract_owner_side_channel(path)
        assert channel["sdt_lines"] == texts


# ---------------------------------------------------------------------------
# header_lines / footer_lines
# ---------------------------------------------------------------------------

def test_header_paragraph_is_collected(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    doc.sections[0].header.add_paragraph("Synthetic Letterhead Name")

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["Synthetic Letterhead Name"]
    assert channel["sdt_lines"] == []
    assert channel["footer_lines"] == []


def test_footer_paragraph_is_collected(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    doc.sections[0].footer.add_paragraph("Synthetic Contact Block")

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["footer_lines"] == ["Synthetic Contact Block"]


def test_sdt_line_carrying_a_dob_is_scrubbed_before_stage4_ever_sees_it(tmp_path):
    """#847 round 2: this side channel feeds `stage4/owner_name.py`'s LLM
    fallback tier directly -- an sdt-wrapped Personal Data block reached
    that prompt with the raw value before this fix."""
    doc = Document()
    anchor = doc.add_paragraph("after")
    _splice_body_level_sdt(anchor, ["Date of Birth: 01/02/1970"])

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["sdt_lines"] == ["Date of Birth: [withheld]"]


def test_sdt_value_in_the_paragraph_after_a_bare_dob_label_is_scrubbed(tmp_path):
    """#847 residual: a content control whose "Date of Birth:" paragraph is
    blank and whose value is the NEXT paragraph -- each line was scrubbed
    alone, so the value reached stage 4's owner-name prompt."""
    doc = Document()
    anchor = doc.add_paragraph("after")
    _splice_body_level_sdt(anchor, ["Jane Roe", "Date of Birth:", "01/02/1970", "Citizenship: US"])

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["sdt_lines"] == ["Jane Roe", "Date of Birth:", "[withheld]", "Citizenship: US"]


def test_header_value_on_the_line_after_a_bare_dob_label_is_scrubbed(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    header = doc.sections[0].header
    header.paragraphs[0].text = "Date of Birth:"
    header.add_paragraph("January 2, 1970")

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["Date of Birth:", "[withheld]"]


def test_sdt_next_line_scrub_keeps_a_line_that_does_not_open_with_a_whole_date(tmp_path):
    """The same cross-boundary rule as the body stream: after a bare
    "Date of Birth:", a line with its own text before the date, or a bare
    year, is another field and keeps its value."""
    doc = Document()
    anchor = doc.add_paragraph("after")
    _splice_body_level_sdt(anchor, [
        "Date of Birth:", "Appointed 07/01/2005",
        "Date of Birth:", "2001",
        "Date of Birth:", "Date of Appointment: 07/01/2005",
    ])

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["sdt_lines"] == [
        "Date of Birth:", "Appointed 07/01/2005",
        "Date of Birth:", "2001",
        "Date of Birth:", "Date of Appointment: 07/01/2005",
    ]


def test_header_line_carrying_an_ssn_is_scrubbed(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    doc.sections[0].header.add_paragraph("SSN: 123-45-6789")

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["SSN: [withheld]"]
    assert not any(c.isdigit() for line in channel["header_lines"] for c in line)


def test_first_page_only_header_is_collected_when_default_header_is_empty(tmp_path):
    """#456 round-2 verifier finding 5: 6 corpus files carry the owner name
    ONLY in the section's first-page header, with an empty (unwritten,
    still-linked) default header. `section.header` alone misses these;
    `first_page_header` must be read too."""
    doc = Document()
    doc.add_paragraph("body")
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.is_linked_to_previous = False
    doc.sections[0].first_page_header.add_paragraph("Synthetic First Page Owner")
    # default header deliberately left untouched (still linked, no definition)

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["Synthetic First Page Owner"]


def test_three_sections_distinct_headers_plus_one_first_page_header(tmp_path):
    """Determinism proof (#456-R2 F1): three sections with distinct 5-line
    default headers, one of which also has a distinct 5-line first-page
    header -> 20 lines total, identical across 3 independent calls."""
    doc = Document()
    doc.add_paragraph("s0 body")
    doc.sections[0].header.is_linked_to_previous = False
    for i in range(5):
        doc.sections[0].header.add_paragraph(f"S0H{i}")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("s1 body")
    doc.sections[1].header.is_linked_to_previous = False
    for i in range(5):
        doc.sections[1].header.add_paragraph(f"S1H{i}")
    doc.sections[1].different_first_page_header_footer = True
    doc.sections[1].first_page_header.is_linked_to_previous = False
    for i in range(5):
        doc.sections[1].first_page_header.add_paragraph(f"S1FPH{i}")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("s2 body")
    doc.sections[2].header.is_linked_to_previous = False
    for i in range(5):
        doc.sections[2].header.add_paragraph(f"S2H{i}")

    path = _save(doc, tmp_path)
    expected = (
        [f"S0H{i}" for i in range(5)]
        + [f"S1H{i}" for i in range(5)]
        + [f"S1FPH{i}" for i in range(5)]
        + [f"S2H{i}" for i in range(5)]
    )
    for _ in range(3):
        channel = extract_owner_side_channel(path)
        assert channel["header_lines"] == expected
        assert len(channel["header_lines"]) == 20


def test_second_section_header_linked_to_previous_is_not_duplicated(tmp_path):
    """A section whose header inherits the prior section's definition
    (`is_linked_to_previous`) must contribute its lines exactly once, not
    once per section that shares the part."""
    doc = Document()
    doc.add_paragraph("s0 body")
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].header.add_paragraph("Shared Header Owner")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("s1 body")
    # section 1's header is left linked to previous (default) -- it must not
    # add a second copy of "Shared Header Owner"

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["Shared Header Owner"]


def test_two_non_linked_sections_sharing_the_same_header_part_are_deduped_by_partname(tmp_path):
    """#456-R2 verifier Finding 1: two sections that are each explicitly
    NOT linked to the previous section (`is_linked_to_previous` False on
    both -- each has its own `w:headerReference`) can still have their
    `w:headerReference` r:ids resolve to the SAME header part, e.g. when a
    document is hand-assembled or re-saved by a tool that reuses a
    relationship id across sections rather than cloning the part. The
    dedup in `_header_footer_paragraph_lines` keys on the OPC part name
    (`str(container.part.partname)`), not on `is_linked_to_previous`, so
    this case -- distinct from the "linked" test above -- must also
    collapse to one copy of the header's lines."""
    doc = Document()
    doc.add_paragraph("s0 body")
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].header.add_paragraph("Shared Part Owner")
    shared_rId = doc.sections[0]._sectPr.get_headerReference(WD_HEADER_FOOTER.PRIMARY).rId

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("s1 body")
    # Section 1 gets its OWN headerReference element (so is_linked_to_previous
    # is False, unlike the "linked" test above) but it is pointed at section
    # 0's r:id rather than a freshly-added part -- the same-partname,
    # not-linked shape the ticket's Finding 1 asks for.
    doc.sections[1]._sectPr.add_headerReference(WD_HEADER_FOOTER.PRIMARY, shared_rId)

    assert doc.sections[0].header.is_linked_to_previous is False
    assert doc.sections[1].header.is_linked_to_previous is False
    assert str(doc.sections[0].header.part.partname) == str(doc.sections[1].header.part.partname)

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["Shared Part Owner"]


def test_header_with_only_a_page_field_is_skipped(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    p = doc.sections[0].header.add_paragraph()
    p._p.append(parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
    ))

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == []


def test_header_with_a_page_field_plus_real_text_keeps_the_text(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    doc.sections[0].header.add_paragraph("Synthetic Faculty Member")
    p = doc.sections[0].header.add_paragraph()
    p._p.append(parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
    ))

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel["header_lines"] == ["Synthetic Faculty Member"]


# ---------------------------------------------------------------------------
# no sdt / no header / no footer at all
# ---------------------------------------------------------------------------

def test_a_plain_docx_with_none_of_the_three_returns_all_empty(tmp_path):
    doc = Document()
    doc.add_paragraph("plain body text, nothing structured")

    channel = extract_owner_side_channel(_save(doc, tmp_path))

    assert channel == {"sdt_lines": [], "header_lines": [], "footer_lines": []}


# ---------------------------------------------------------------------------
# THE landmine test: the element stream must be byte-identical with vs
# without a body-level sdt. This is the test that stops someone "fixing"
# #456 by inlining sdt paragraphs into the main body walk.
# ---------------------------------------------------------------------------

def test_element_stream_is_unchanged_by_a_body_level_sdt(tmp_path):
    doc_without = Document()
    doc_without.add_paragraph("before")
    doc_without.add_paragraph("after")
    path_without = _save(doc_without, tmp_path, "without_sdt.docx")

    doc_with = Document()
    doc_with.add_paragraph("before")
    after_with = doc_with.add_paragraph("after")
    _splice_body_level_sdt(after_with, ["Synthetic Owner Name"])
    path_with = _save(doc_with, tmp_path, "with_sdt.docx")

    result_without = extract_unified_elements(path_without)
    result_with = extract_unified_elements(path_with)

    def _stream(result):
        return [
            (el.get("type"), el.get("text"), el.get("unified_idx"), el.get("para_idx"))
            for el in result["elements"]
        ]

    assert _stream(result_with) == _stream(result_without)
    assert result_with["meta"]["num_elements"] == result_without["meta"]["num_elements"]
    assert result_with["meta"]["num_paragraphs"] == result_without["meta"]["num_paragraphs"]
    assert "Synthetic Owner Name" not in [el.get("text") for el in result_with["elements"]]

    # para_idx must still line up with doc.paragraphs (the #456 issue
    # comment's own verified probe, reproduced as a permanent regression
    # guard): the sdt paragraph is not one of doc.paragraphs' direct-body
    # CT_P children, so it must not have shifted any subsequent para_idx.
    reopened = Document(path_with)
    para_idxs = [
        el["para_idx"] for el in result_with["elements"]
        if el.get("type") in ("paragraph", "empty")
    ]
    assert para_idxs == list(range(len(reopened.paragraphs)))
