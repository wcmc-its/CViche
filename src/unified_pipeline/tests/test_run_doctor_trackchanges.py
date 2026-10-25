"""Guard: the doctor's docx readers must see tracked-change INSERTIONS (#249).

Stage 6 inserts LLM-enriched content (research summaries, reformatted citations)
as tracked changes (``<w:ins>``). python-docx's ``.text`` skips those runs, so
the old readers under-reported what rendered and false-flagged content as
'unrendered' (classified_unrendered, unrendered_records, dead_sections). The
readers now read every ``w:t`` under an element — including inside ``<w:ins>``,
excluding deleted ``<w:delText>`` — i.e. the accepted-changes view a reader sees.

    python3 -m pytest src/unified_pipeline/tests/test_run_doctor_trackchanges.py -p no:cacheprovider

Self-contained: no DB, no template. Requires only python-docx.
"""

import json
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from unified_pipeline.run_doctor import _docx_text, read_docx_blocks, run_doctor


def test_docx_text_includes_insertions_excludes_deletions():
    p = parse_xml(
        f'<w:p {nsdecls("w")}>'
        '<w:r><w:t>plain </w:t></w:r>'
        '<w:ins><w:r><w:t>inserted</w:t></w:r></w:ins>'
        '<w:del><w:r><w:delText> deleted</w:delText></w:r></w:del>'
        '</w:p>'
    )
    # inserted text is present (was invisible before); deleted text is not
    assert _docx_text(p) == "plain inserted"


def test_docx_text_default_view_drops_tabs_and_breaks():
    """The lints' view is text only; whitespace elements are opt-in (#461)."""
    p = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:t>a</w:t><w:tab/><w:br/><w:t>b</w:t></w:r></w:p>'
    )
    assert _docx_text(p) == "ab"


def test_docx_text_with_whitespace_renders_tab_break_and_cr():
    p = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:t>a</w:t><w:tab/><w:t>b</w:t><w:br/>'
        '<w:t>c</w:t><w:cr/><w:t>d</w:t></w:r></w:p>'
    )
    assert _docx_text(p, with_whitespace=True) == "a\tb\nc\nd"


def test_docx_text_with_whitespace_skips_tab_stops_deletions_and_page_breaks():
    p = parse_xml(
        f'<w:p {nsdecls("w")}><w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
        '<w:r><w:t>a</w:t><w:br w:type="page"/></w:r>'
        '<w:del><w:r><w:tab/><w:br/><w:delText>x</w:delText></w:r></w:del>'
        '<w:ins><w:r><w:tab/><w:t>b</w:t></w:r></w:ins></w:p>'
    )
    assert _docx_text(p, with_whitespace=True) == "a\tb"


def test_docx_text_empty_paragraph():
    p = parse_xml(f'<w:p {nsdecls("w")}></w:p>')
    assert _docx_text(p) == ""


def test_docx_text_deleted_view_reads_only_the_tracked_deletions():
    """#1223: a tracked deletion is still in the file a reader opens (Word's
    default markup view shows it struck through), so the protected-data lint
    reads it through `deleted=True`: the other half of the default view."""
    p = parse_xml(
        f'<w:p {nsdecls("w")}>'
        '<w:r><w:t>plain </w:t></w:r>'
        '<w:ins><w:r><w:t>inserted</w:t></w:r></w:ins>'
        '<w:del><w:r><w:delText>deleted one</w:delText></w:r></w:del>'
        '<w:del><w:r><w:delText>, two</w:delText></w:r></w:del>'
        '</w:p>'
    )
    assert _docx_text(p, deleted=True) == "deleted one, two"
    with pytest.raises(ValueError):
        _docx_text(p, with_whitespace=True, deleted=True)


def test_read_docx_blocks_deleted_view_is_block_for_block(tmp_path):
    """Block i of the deleted view is block i of the default view -- the same
    paragraph or table -- holding only that block's deleted text."""
    doc = Document()
    doc.add_paragraph("kept")
    doc.add_paragraph()._p.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="1" w:author="a" w:date="2026-07-25T00:00:00Z">'
        '<w:r><w:delText>struck</w:delText></w:r></w:del>'))
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Label:"
    table.rows[0].cells[1].paragraphs[0]._p.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="2" w:author="a" w:date="2026-07-25T00:00:00Z">'
        '<w:r><w:delText>old value</w:delText></w:r></w:del>'))
    nested = table.rows[0].cells[0].add_table(rows=1, cols=1)
    nested.rows[0].cells[0].paragraphs[0]._p.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="3" w:author="a" w:date="2026-07-25T00:00:00Z">'
        '<w:r><w:delText>nested old</w:delText></w:r></w:del>'))
    path = tmp_path / "x.docx"
    doc.save(path)
    accepted = read_docx_blocks(str(path))
    deleted = read_docx_blocks(str(path), deleted=True)
    assert [kind for kind, _ in deleted] == [kind for kind, _ in accepted]
    assert [text for _, text in deleted] == ["", "struck", "nested old\nold value"]
    # (python-docx adds an empty paragraph after a nested table: the "\n")
    assert [text for _, text in accepted] == ["kept", "", "Label:\n"]


def _tracked_changes_run(root, uid, cv_owner):
    """Minimal run: a populated stage-4 artifact plus a stage-6 docx whose whole
    body is a tracked INSERTION."""
    stage4 = root / "stage_4_field_extraction"
    stage4.mkdir(parents=True)
    (stage4 / f"{uid}_fields.json").write_text(json.dumps(
        {"document_uid": uid, "cv_owner": cv_owner, "entries": []}))

    doc = Document()
    doc.add_paragraph()._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-07-25T00:00:00Z">'
        '<w:r><w:t>Miriam Shapiro, M.D. — Curriculum Vitae</w:t></w:r></w:ins>'))
    out = root / "stage_6_wcm_documents"
    out.mkdir(parents=True)
    doc.save(out / f"{uid}_wcm.docx")
    return out / f"{uid}_wcm.docx"


def test_hard_fail_gates_do_not_fire_on_tracked_changes(tmp_path):
    """The #437 hard-fail lints decide from the stage-4 JSON, never from the
    rendered docx, so a document delivered entirely as tracked insertions cannot
    make them report an undeliverable run -- the false-positive shape this file
    exists to pin."""
    root = tmp_path / "outputs"
    docx = _tracked_changes_run(root, "TRACKED", {
        "first_name": "Miriam", "last_name": "Shapiro",
        "full_name": "Miriam Shapiro"})
    # not vacuous: the docx really does carry its text inside <w:ins>
    assert any("Miriam Shapiro" in text for _, text in read_docx_blocks(str(docx)))

    payload = run_doctor(root, "TRACKED")
    gate_lints = ("owner_contact_missing", "pipeline_errors_present")
    assert not [f for f in payload["findings"] if f["lint"] in gate_lints], \
        "tracked changes must not manufacture a hard-fail finding"
    assert payload["counts"]["ERROR"] == 0

    # ...and the lint is not merely inert on this fixture: the same
    # tracked-changes document with an emptied cv_owner still reports ERROR.
    empty = tmp_path / "empty-owner"
    _tracked_changes_run(empty, "TRACKED", {"first_name": "", "last_name": "",
                                            "full_name": ""})
    broken = run_doctor(empty, "TRACKED")
    assert broken["worst_severity"] == "ERROR"
    assert [f["lint"] for f in broken["findings"] if f["severity"] == "ERROR"] \
        == ["owner_contact_missing"]


if __name__ == "__main__":
    test_docx_text_includes_insertions_excludes_deletions()
    test_docx_text_default_view_drops_tabs_and_breaks()
    test_docx_text_with_whitespace_renders_tab_break_and_cr()
    test_docx_text_with_whitespace_skips_tab_stops_deletions_and_page_breaks()
    test_docx_text_empty_paragraph()
    print("OK")


# --- sub-points (#1205): stage 6's tracked-deleted source prose ---------------

_SUBPOINT_PROSE = "Coordinates the example curriculum committee and supervises graduate trainees"


def _doc_with_a_sub_point_row_and_another_deletion():
    """A table row with a sub-point row under it (stage6/supplementary.py's
    writer), plus a body paragraph with a formatter's tracked deletion."""
    from unified_pipeline.stage6.supplementary import Revision, write_row_subpoint

    doc = Document()
    doc.add_paragraph().add_run("Lecturer")._r.addnext(parse_xml(
        f'<w:del {nsdecls("w")} w:id="90" w:author="Formatter" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:delText>old wording</w:delText></w:r></w:del>'))
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Lecturer"
    table.rows[0].cells[1].text = "Example College"
    line = write_row_subpoint(table.rows[0]._tr, (_SUBPOINT_PROSE,), Revision(1, "2026-01-01T00:00:00Z"))
    return doc, line


def test_docx_subpoint_lines_reads_only_the_sub_points():
    from unified_pipeline.doctor.shared import docx_subpoint_lines

    doc, line = _doc_with_a_sub_point_row_and_another_deletion()
    assert docx_subpoint_lines(doc) == [_SUBPOINT_PROSE]
    # The row's line, as a reader who rejects the deletion sees the entry.
    assert line == f"Lecturer Example College {_SUBPOINT_PROSE}"


def test_read_docx_subpoints_feeds_offered_as_subpoint(tmp_path):
    from unified_pipeline.doctor.shared import offered_as_subpoint
    from unified_pipeline.run_doctor import read_docx_subpoints

    doc, _line = _doc_with_a_sub_point_row_and_another_deletion()
    path = tmp_path / "out.docx"
    doc.save(str(path))
    subpoints = read_docx_subpoints(str(path))
    # Verbatim -- also inside a record line most of whose words are elsewhere --
    # and reworded on the same line's words.
    assert offered_as_subpoint(f"2019-2021\tLecturer\t{_SUBPOINT_PROSE}", subpoints)
    assert offered_as_subpoint("Lecturer, Northfield Regional Hospital, Department of "
                               f"Anaesthesiology Research\t{_SUBPOINT_PROSE}", subpoints)
    assert offered_as_subpoint("supervises graduate trainees, coordinates curriculum committee",
                               subpoints)
    # The formatter's deletion is not a sub-point; nor is unrelated prose.
    assert not offered_as_subpoint("old wording", subpoints)
    assert not offered_as_subpoint("Chairs the departmental finance review panel", subpoints)
    assert not offered_as_subpoint(_SUBPOINT_PROSE, None)


def _sidecar_placements(*placements):
    from unified_pipeline.stage6.supplementary import SUBPOINT_PLACEMENTS_KEY

    return {"warnings": [], SUBPOINT_PLACEMENTS_KEY: [p._asdict() for p in placements]}


_OTHER_SUBPOINT = "Organises the regional examiners' workshop and moderates assessment panels"


def test_a_sub_point_vouches_only_for_its_own_entry():
    """#1205: the sidecar's placements say whose each sub-point is. With an
    entry's index, only its own count; without one, or without placements
    (an older sidecar, or none), every sub-point in the document does."""
    from unified_pipeline.doctor.shared import (
        docx_subpoint_lines,
        offered_as_subpoint,
        subpoint_placements,
        subpoint_text,
    )
    from unified_pipeline.stage6.supplementary import SubpointPlacement

    doc, _line = _doc_with_a_sub_point_row_and_another_deletion()
    lines = [*docx_subpoint_lines(doc), _OTHER_SUBPOINT]
    placements = subpoint_placements(_sidecar_placements(
        SubpointPlacement(29, "D1", (_SUBPOINT_PROSE,)), SubpointPlacement(30, "K1", (_OTHER_SUBPOINT,))))
    placed = subpoint_text(lines, placements)
    assert offered_as_subpoint(_SUBPOINT_PROSE, placed, 29)
    assert not offered_as_subpoint(_SUBPOINT_PROSE, placed, 30)
    assert not offered_as_subpoint(_SUBPOINT_PROSE, placed, 31)
    assert offered_as_subpoint(_SUBPOINT_PROSE, placed)
    for older in (None, {"warnings": []}):
        assert subpoint_placements(older) is None
        assert offered_as_subpoint(_SUBPOINT_PROSE, subpoint_text(lines, subpoint_placements(older)), 30)
    # A sidecar recording no sub-point is a record too: none is any entry's.
    assert subpoint_placements(_sidecar_placements()) == ()
    assert not offered_as_subpoint(_SUBPOINT_PROSE, subpoint_text(lines, ()), 29)
    # The document says what is offered: a placement whose text a corrected
    # copy no longer holds as a sub-point vouches for nothing.
    assert not offered_as_subpoint(_SUBPOINT_PROSE, subpoint_text([_OTHER_SUBPOINT], placements), 29)


def test_run_doctor_reads_the_placements_from_the_render_sidecar(tmp_path):
    """The wire: run_doctor hands the sidecar's placements to the sub-point
    view, so a record line only another entry's sub-point matches is
    reported (EBYSBC VGHNZD's shape, invented), and with a sidecar from
    before the placements it stays quiet, as it did."""
    from unified_pipeline.stage6.supplementary import (
        Revision,
        SubpointPlacement,
        write_row_subpoint,
    )

    shown = "2014-2016 Quimby Farrow, graduate student, marsupial cardiology rotation"
    lost = "2017-2019 Quimby Farrow, doctoral thesis on wallaby circulation outcomes"
    others = "Mentored outcomes research on wallaby circulation for a doctoral thesis"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = shown
    write_row_subpoint(table.rows[0]._tr, (others,), Revision(1, "2026-01-01T00:00:00Z"))
    stage6 = tmp_path / "stage_6_wcm_documents"
    stage6.mkdir()
    doc.save(str(stage6 / "VSHAPE_wcm.docx"))
    stage4 = tmp_path / "stage_4_field_extraction"
    stage4.mkdir()
    (stage4 / "VSHAPE_fields.json").write_text(json.dumps({"entries": [
        {"element_idx_start": 463, "taxonomy_code": "N3A", "text": "Casey Sample, postdoc"},
        {"element_idx_start": 477, "taxonomy_code": "N3A", "text": f"{shown}\n{lost}"}]}))
    sidecar = stage6 / "VSHAPE_render_warnings.json"

    def unrendered(report):
        sidecar.write_text(json.dumps(report))
        return [f["message"] for f in run_doctor(tmp_path, "VSHAPE")["findings"]
                if f["lint"] == "unrendered_records"]

    assert unrendered(_sidecar_placements(SubpointPlacement(463, "N3A", (others,)))) == [
        "entry 477 (N3A): 1 of 2 records absent from output"]
    assert unrendered({"warnings": []}) == []


def test_subpoint_owners_turns_the_per_entry_view_round():
    """Each sub-point line of the document with the entries whose own it is
    (`subpoint_owners`, which `source_line_coverage` reads): a line two
    placements wrote is both entries'; one no placement wrote is no entry's;
    with no record of whose each one is, every line's owners are unknown
    (None), so it counts for every entry."""
    from unified_pipeline.core.text_norm import squash
    from unified_pipeline.doctor.shared import (
        subpoint_owners,
        subpoint_placements,
        subpoint_text,
    )
    from unified_pipeline.stage6.supplementary import SubpointPlacement

    unplaced = "Chairs the departmental finance review panel"
    lines = [_SUBPOINT_PROSE, _OTHER_SUBPOINT, unplaced]
    placements = subpoint_placements(_sidecar_placements(
        SubpointPlacement(29, "D1", (_SUBPOINT_PROSE,)), SubpointPlacement(30, "K1", (_OTHER_SUBPOINT,)),
        SubpointPlacement(31, "K1", (_OTHER_SUBPOINT,))))
    assert subpoint_owners(subpoint_text(lines, placements)) == {
        squash(_SUBPOINT_PROSE): frozenset({29}), squash(_OTHER_SUBPOINT): frozenset({30, 31}),
        squash(unplaced): frozenset()}
    assert subpoint_owners(subpoint_text([], placements)) == {}
    assert subpoint_owners(subpoint_text(lines)) == dict.fromkeys(map(squash, lines))
    assert subpoint_owners(None) == {}


def test_run_doctor_hands_source_line_coverage_the_sub_points_and_their_owners(tmp_path):
    """The wire (#1205, BMAMWE 1368's shape, invented): run_doctor gives
    source_line_coverage the docx's sub-points with the sidecar's
    placements, so a note only another entry's sub-point holds is
    reported, and with a sidecar from before the placements it is not."""
    from unified_pipeline.stage6.supplementary import (
        Revision,
        SubpointPlacement,
        write_row_subpoint,
    )

    prose = "Featured on the cover of the spring marsupial issue, with an editorial on wallaby welfare"
    note = "Cover feature of the spring marsupial issue, with an editorial on wallaby welfare"
    source = Document()
    source.add_paragraph(note)
    source.save(str(tmp_path / "WSHAPE.docx"))
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Farrow Q. Wallaby circulation. J Marsupial Med. 2019;3:1-5."
    write_row_subpoint(table.rows[0]._tr, (prose,), Revision(1, "2026-01-01T00:00:00Z"))
    stage6 = tmp_path / "stage_6_wcm_documents"
    stage6.mkdir()
    doc.save(str(stage6 / "WSHAPE_wcm.docx"))
    stage3b = tmp_path / "stage_3b_classified_entries"
    stage3b.mkdir()
    (stage3b / "WSHAPE_classified.json").write_text(json.dumps({"entries": [
        {"element_idx_start": 20, "taxonomy_code": "S1", "text": f"Farrow Q. Wallaby circulation.\t{prose}"},
        {"element_idx_start": 30, "taxonomy_code": "S7", "text": note}]}))
    sidecar = stage6 / "WSHAPE_render_warnings.json"

    def coverage(report):
        sidecar.write_text(json.dumps(report))
        return [f["message"].split(":")[0] for f in run_doctor(tmp_path, "WSHAPE")["findings"]
                if f["lint"] == "source_line_coverage"]

    assert coverage(_sidecar_placements(SubpointPlacement(20, "S1", (prose,)))) == ["entry 30"]
    assert coverage(_sidecar_placements(SubpointPlacement(30, "S7", (prose,)))) == []
    assert coverage({"warnings": []}) == []
