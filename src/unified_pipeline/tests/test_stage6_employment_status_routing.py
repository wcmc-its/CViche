"""Employment Status entries route to the row their OWN label names (#571).

`_fill_employment_status` used to write every entry into the first paragraph
containing 'employer' -- on the WCM template always "Name of Current
Employer(s):" -- so a Position/Title entry and a Dates entry both overwrote the
employer row, last write wins, and the surviving value was filed under a label
that is not its own. It also counted every one of those writes as inserted.

These tests drive the real fill path (`_fill_passthrough_sections`) against a
real python-docx Document, once with the issue's synthetic multi-row block and
once with the committed October-2022 WCM template itself.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_employment_status_routing.py -p no:cacheprovider
"""

import logging
import sys
from pathlib import Path

import docx
import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6.sections.passthrough import PassthroughSection  # noqa: E402
from unified_pipeline.stage_6_word_template import TEMPLATE_PATH  # noqa: E402


class _StubGenerator(PassthroughSection):
    """Just enough of WCMTemplateGenerator to drive the passthrough fill."""

    def __init__(self, doc):
        self.doc = doc
        self.verbose = False
        self.stats = {"entries_inserted": 0, "tables_populated": 0}

    def _find_paragraph_with_text(self, search_text):
        for i, para in enumerate(self.doc.paragraphs):
            if search_text.lower() in para.text.lower():
                return i
        return None


def _entry(text: str) -> dict:
    return {"text": text, "hierarchy": ["E. EMPLOYMENT STATUS"], "extracted_fields": {}}


def _block_document():
    """The issue's reproduction block: header plus three labelled rows."""
    doc = docx.Document()
    for line in (
        "E. EMPLOYMENT STATUS",
        "Name of Current Employer(s):",
        "Position/Title:",
        "Dates of Employment:",
    ):
        doc.add_paragraph(line)
    return doc


def _texts(doc) -> list[str]:
    return [p.text for p in doc.paragraphs]


def test_each_entry_lands_in_its_own_row():
    doc = _block_document()
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([
        _entry("Name of Current Employer(s): Weill Cornell Medicine"),
        _entry("Position/Title: Professor of Medicine"),
        _entry("Dates of Employment: 2011-Present"),
    ])

    assert "Name of Current Employer(s):\tWeill Cornell Medicine" in _texts(doc)
    assert "Position/Title:\tProfessor of Medicine" in _texts(doc)
    assert "Dates of Employment:\t2011-Present" in _texts(doc)
    assert generator.stats["entries_inserted"] == 3


def test_dates_entry_does_not_overwrite_the_employer_row():
    """The exact defect: a lone Dates entry used to land in the employer row."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([_entry("Dates of Employment: 2011-Present")])

    texts = _texts(doc)
    assert "Name of Current Employer(s):" in texts  # untouched
    assert "Dates of Employment:\t2011-Present" in texts


def test_loosely_worded_source_label_still_reaches_the_employer_row():
    """Corpus wording: 'Name of Employer(s)' differs from the template's label."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections(
        [_entry("Name of Employer(s):  Weill Cornell Medical College")])

    assert "Name of Current Employer(s):\tWeill Cornell Medical College" in _texts(doc)
    assert generator.stats["entries_inserted"] == 1


def test_unknown_label_is_not_written_not_counted_and_is_logged(caplog):
    doc = _block_document()
    generator = _StubGenerator(doc)

    with caplog.at_level(logging.WARNING):
        generator._fill_passthrough_sections([_entry("Favorite Color: blue")])

    assert _texts(doc) == [
        "E. EMPLOYMENT STATUS",
        "Name of Current Employer(s):",
        "Position/Title:",
        "Dates of Employment:",
    ]
    assert generator.stats["entries_inserted"] == 0
    assert any("Favorite Color" in record.getMessage() for record in caplog.records)


def test_row_missing_from_template_is_not_counted_and_is_logged(caplog):
    """A matching label class with no row in the window must not fake a write."""
    doc = docx.Document()
    doc.add_paragraph("E. EMPLOYMENT STATUS")
    doc.add_paragraph("Name of Current Employer(s):")
    for _ in range(15):  # push the look-alike row beyond the scan window
        doc.add_paragraph("filler")
    doc.add_paragraph("Date of preparation:")
    generator = _StubGenerator(doc)

    with caplog.at_level(logging.WARNING):
        generator._fill_passthrough_sections(
            [_entry("Dates of Employment: 2011-Present")])

    assert generator.stats["entries_inserted"] == 0
    assert not any("2011-Present" in p.text for p in doc.paragraphs)
    assert any("Dates of Employment" in record.getMessage() for record in caplog.records)


def test_row_at_scan_window_boundary_is_found():
    """A row at the last paragraph the scan window covers is still routed."""
    doc = docx.Document()
    doc.add_paragraph("E. EMPLOYMENT STATUS")
    doc.add_paragraph("Name of Current Employer(s):")
    for _ in range(10):
        doc.add_paragraph("filler")
    doc.add_paragraph("Dates of Employment:")  # anchor offset 11, last in-window
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([_entry("Dates of Employment: 2011-Present")])

    assert "Dates of Employment:\t2011-Present" in _texts(doc)
    assert generator.stats["entries_inserted"] == 1


def test_row_just_outside_scan_window_is_not_found():
    """One paragraph further than the boundary case, the row is unreachable."""
    doc = docx.Document()
    doc.add_paragraph("E. EMPLOYMENT STATUS")
    doc.add_paragraph("Name of Current Employer(s):")
    for _ in range(11):
        doc.add_paragraph("filler")
    doc.add_paragraph("Dates of Employment:")  # anchor offset 12, first out-of-window
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([_entry("Dates of Employment: 2011-Present")])

    assert not any("2011-Present" in text for text in _texts(doc))
    assert generator.stats["entries_inserted"] == 0


@pytest.mark.parametrize(
    ("entry_text", "expected", "must_not_change"),
    [
        (
            "Name of Employer(s): Weill Cornell",
            "Name of Current Employer(s):\tWeill Cornell",
            ["Position/Title:", "Dates of Employment:"],
        ),
        (
            "Position/Title: Professor",
            "Position/Title:\tProfessor",
            ["Name of Current Employer(s):", "Dates of Employment:"],
        ),
        (
            "Dates of Employment: 2011-Present",
            "Dates of Employment:\t2011-Present",
            ["Name of Current Employer(s):", "Position/Title:"],
        ),
    ],
)
def test_entry_is_not_written_to_another_employment_row(entry_text, expected, must_not_change):
    """Each row class writes ONLY its own row -- the other two stay untouched."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([_entry(entry_text)])

    texts = _texts(doc)
    assert expected in texts
    for original in must_not_change:
        assert original in texts


@pytest.mark.parametrize(
    "label",
    [
        "Name of Employer(s):",
        "NAME OF EMPLOYER(S):",
        "Name   of   Employer(s):",
        "  Name of Employer(s):  ",
    ],
)
def test_employer_label_normalization_variants(label):
    """_squash() promises whitespace- and case-insensitive matching; prove it."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([_entry(f"{label} Weill Cornell")])

    assert "Name of Current Employer(s):\tWeill Cornell" in _texts(doc)


@pytest.mark.parametrize(
    "label",
    ["Date of Birth", "Date Submitted", "Project Title", "Application Status"],
)
def test_unrelated_labels_are_not_routed_as_employment_rows(label, caplog):
    """A generic word ("date", "title", "status") alone must not match a row --
    only that word paired with "employ" (or "position") does (#571 review)."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    with caplog.at_level(logging.WARNING):
        generator._fill_passthrough_sections([_entry(f"{label}: some value")])

    assert not any("some value" in text for text in _texts(doc))
    assert generator.stats["entries_inserted"] == 0
    assert any(label in record.getMessage() for record in caplog.records)


def test_real_wcm_template_routes_employer_and_status_rows():
    """On the committed October-2022 template, the two real rows fill independently."""
    doc = docx.Document(str(TEMPLATE_PATH))
    generator = _StubGenerator(doc)

    generator._fill_passthrough_sections([
        _entry("Name of Employer(s): Weill Cornell Medical College"),
        _entry("Employment Status: Full-time salaried by Weill Cornell"),
    ])

    texts = _texts(doc)
    assert "Name of Current Employer(s):\tWeill Cornell Medical College" in texts
    assert ("Current Employment Status (Please choose one, list here, delete the "
            "others):\tFull-time salaried by Weill Cornell") in texts
    assert generator.stats["entries_inserted"] == 2


# --- #807: E-coded entries under a foreign heading ---------------------------


def _e_coded(text: str, heading: str = "Other Example Activities") -> dict:
    """An E-coded entry filed under a heading that is not Employment Status --
    the over-segmented copy stage 3b's dedup can leave behind."""
    return {"text": text, "hierarchy": [heading], "taxonomy_code": "E", "extracted_fields": {}}


def test_e_coded_entry_under_foreign_heading_is_written_and_consumed():
    doc = _block_document()
    generator = _StubGenerator(doc)
    entry = _e_coded("Position/Title: Example Professor")

    consumed = generator._fill_passthrough_sections([entry])

    assert consumed == [entry]
    assert "Position/Title:\tExample Professor" in _texts(doc)
    assert generator.stats["entries_inserted"] == 1


def test_e_coded_foreign_entry_with_unknown_label_is_not_written_or_consumed():
    """The known-label guard (#571) still gates a code-selected entry."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    consumed = generator._fill_passthrough_sections([_e_coded("Favorite Color: blue")])

    assert consumed == []
    assert not any("blue" in text for text in _texts(doc))
    assert generator.stats["entries_inserted"] == 0


@pytest.mark.parametrize("text", ["Position Example Professor", "Position/Title: ab"])
def test_e_coded_foreign_entry_without_a_label_value_shape_is_not_written(text):
    doc = _block_document()
    generator = _StubGenerator(doc)

    assert generator._fill_passthrough_sections([_e_coded(text)]) == []
    assert generator.stats["entries_inserted"] == 0


def test_only_e_code_is_selected_under_a_foreign_heading():
    """A T-coded entry with an employment-shaped label under a foreign heading
    is not E's: selection is by the E code, not by the label alone."""
    doc = _block_document()
    generator = _StubGenerator(doc)
    entry = _e_coded("Position/Title: Example Professor")
    entry["taxonomy_code"] = "T"

    assert generator._fill_passthrough_sections([entry]) == []
    assert generator.stats["entries_inserted"] == 0


@pytest.mark.parametrize(
    "heading",
    ["G. INSTITUTIONAL/HOSPITAL AFFILIATION", "Hospital Affiliations", "J. PERCENT EFFORT"],
)
def test_e_coded_entry_under_another_passthrough_heading_is_not_claimed(heading):
    """A heading G or J owns is not E's to claim: the entry would otherwise be
    written by two writers."""
    doc = _block_document()
    generator = _StubGenerator(doc)

    # E's writer alone: the stub carries no G/J template to drive theirs.
    assert generator._fill_employment_status(
        [_e_coded("Position/Title: Example Professor", heading)]) == []
    assert generator.stats["entries_inserted"] == 0


def test_foreign_copy_of_a_written_row_is_a_consumed_duplicate_not_a_second_write():
    doc = _block_document()
    generator = _StubGenerator(doc)
    heading_entry = _entry("Position/Title: Example Professor")
    copy = _e_coded("Position/Title: Example Professor")

    # The foreign copy comes FIRST in the input: the heading entry still wins.
    consumed = generator._fill_passthrough_sections([copy, heading_entry])

    assert consumed == [heading_entry, copy]
    assert _texts(doc).count("Position/Title:\tExample Professor") == 1
    assert generator.stats["entries_inserted"] == 1


def test_conflicting_foreign_copy_neither_overwrites_nor_is_consumed(caplog):
    doc = _block_document()
    generator = _StubGenerator(doc)
    heading_entry = _entry("Position/Title: Example Professor")
    conflict = _e_coded("Position/Title: Example Lecturer")

    with caplog.at_level(logging.WARNING):
        consumed = generator._fill_passthrough_sections([conflict, heading_entry])

    assert consumed == [heading_entry]
    assert "Position/Title:\tExample Professor" in _texts(doc)
    assert not any("Example Lecturer" in text for text in _texts(doc))
    assert any("already written" in r.getMessage() for r in caplog.records)


def test_foreign_copies_of_one_row_write_once_and_the_conflict_stays_unconsumed():
    """With no heading-matched entry at all, the first foreign copy claims the
    row; a second with the same value is a duplicate, one with another value is
    left for the Appendix."""
    doc = _block_document()
    generator = _StubGenerator(doc)
    first = _e_coded("Position/Title: Example Professor")
    same = _e_coded("Position/Title: Example Professor", "Another Example Heading")
    other = _e_coded("Position/Title: Example Lecturer")

    consumed = generator._fill_passthrough_sections([first, other, same])

    assert consumed == [first, same]
    assert generator.stats["entries_inserted"] == 1
