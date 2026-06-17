"""Regression guard for Stage 6 output-rendering options (issue #153).

The WCM template generator used to emit Word track changes (``w:ins``/``w:del``)
AND classification comments (``commentReference`` + a ``comments.xml`` part)
UNCONDITIONALLY. Issue #153 wires the two long-dormant Run columns
(``show_track_changes`` ON by default, ``show_pipeline_comments`` OFF) into the
generator via ``emit_track_changes`` / ``emit_comments`` flags.

These tests drive the low-level emit methods directly on a fresh document,
save the ``.docx``, then unzip and inspect ``word/document.xml`` and the
comments part. They assert the tracked-change / comment markup is present or
omitted per flag combination. Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_render_options.py -p no:cacheprovider

Self-contained: no DB, no FastAPI app, no backend conftest. Requires only
``python-docx`` and the bundled WCM template (auto-located by the generator).
"""

import sys
import zipfile
from pathlib import Path

import pytest

# Ensure the repo's ``src`` directory is importable regardless of cwd/rootdir.
_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402

from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


def _emit_sample_content(gen: WCMTemplateGenerator):
    """Drive the gated emit paths on a small in-memory document.

    Exercises every method issue #153 gates:
      - a tracked-change pair (deletion of the original + insertion of the new),
      - a standalone insertion,
      - a Word comment.
    The generator's ``self.doc`` is the live document we later save/inspect.
    """
    gen.doc = Document()
    para = gen.doc.add_paragraph()
    # Pair: original "Old citation" deleted, "New citation" inserted.
    gen._add_track_change_pair(para, "Old citation", "New citation", author="Tester")
    # Standalone insertion (e.g. an enriched bibliography line).
    gen._add_track_change_insertion(para, "Inserted text", author="Tester")
    # A classification comment on the paragraph.
    gen._add_word_comment(para, "Classified as S1 (confidence: high)", author="Classification")
    # Mirror the real generate() flow: comments are flushed to comments.xml here.
    gen._finalize_comments()


def _render(tmp_path: Path, *, emit_track_changes: bool, emit_comments: bool) -> dict:
    """Generate a docx for a flag combination and return its inspected markup.

    Returns a dict with the document.xml body and whether a comments part exists.
    """
    gen = WCMTemplateGenerator(
        verbose=False,
        emit_track_changes=emit_track_changes,
        emit_comments=emit_comments,
    )
    _emit_sample_content(gen)

    out = tmp_path / f"tc{int(emit_track_changes)}_cm{int(emit_comments)}.docx"
    gen.doc.save(str(out))

    with zipfile.ZipFile(out) as zf:
        names = zf.namelist()
        document_xml = zf.read("word/document.xml").decode("utf-8")
        has_comments_part = any("comments" in n.lower() for n in names)

    return {"document_xml": document_xml, "has_comments_part": has_comments_part}


def test_defaults_match_db_columns():
    """Constructor defaults mirror the Run column defaults: TC on, comments off."""
    gen = WCMTemplateGenerator(verbose=False)
    assert gen.emit_track_changes is True
    assert gen.emit_comments is False


def test_track_changes_on_emits_ins_and_del(tmp_path):
    result = _render(tmp_path, emit_track_changes=True, emit_comments=False)
    xml = result["document_xml"]
    assert "w:ins" in xml, "expected tracked insertion (w:ins) when track changes ON"
    assert "w:del" in xml, "expected tracked deletion (w:del) when track changes ON"
    # The accepted/new text must still be present.
    assert "New citation" in xml
    assert "Inserted text" in xml


def test_track_changes_off_omits_ins_and_del_but_keeps_final_text(tmp_path):
    result = _render(tmp_path, emit_track_changes=False, emit_comments=False)
    xml = result["document_xml"]
    assert "w:ins" not in xml, "track changes OFF must not emit w:ins"
    assert "w:del" not in xml, "track changes OFF must not emit w:del"
    # Insertions degrade to plain runs (final/accepted text survives)...
    assert "New citation" in xml
    assert "Inserted text" in xml
    # ...and the superseded/deleted original is gone.
    assert "Old citation" not in xml


def test_comments_off_omits_comment_markup(tmp_path):
    result = _render(tmp_path, emit_track_changes=True, emit_comments=False)
    xml = result["document_xml"]
    assert "commentReference" not in xml, "comments OFF must not emit commentReference"
    assert "commentRangeStart" not in xml
    assert result["has_comments_part"] is False, "comments OFF must not create a comments.xml part"


def test_comments_on_emits_comment_markup_and_part(tmp_path):
    result = _render(tmp_path, emit_track_changes=True, emit_comments=True)
    xml = result["document_xml"]
    assert "commentReference" in xml, "comments ON must emit commentReference"
    assert result["has_comments_part"] is True, "comments ON must create a comments.xml part"


@pytest.mark.parametrize("emit_tc", [True, False])
@pytest.mark.parametrize("emit_cm", [True, False])
def test_all_flag_combinations_produce_valid_docx(tmp_path, emit_tc, emit_cm):
    """Every combination must produce a readable, openable .docx."""
    result = _render(tmp_path, emit_track_changes=emit_tc, emit_comments=emit_cm)
    xml = result["document_xml"]

    # Track-change markup gated correctly.
    assert ("w:ins" in xml) is emit_tc
    assert ("w:del" in xml) is emit_tc
    # Comment markup gated correctly.
    assert ("commentReference" in xml) is emit_cm
    assert result["has_comments_part"] is emit_cm
    # The final accepted text survives regardless of either flag.
    assert "New citation" in xml
    assert "Inserted text" in xml


def test_saved_docx_reopens_cleanly_with_flags_off(tmp_path):
    """A plain-rendered document must round-trip through python-docx without error."""
    gen = WCMTemplateGenerator(verbose=False, emit_track_changes=False, emit_comments=False)
    _emit_sample_content(gen)
    out = tmp_path / "plain.docx"
    gen.doc.save(str(out))

    reopened = Document(str(out))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "New citation" in text
    assert "Inserted text" in text
    assert "Old citation" not in text
