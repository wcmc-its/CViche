"""Review-round-1 contract tests for `unified_pipeline/doctor/lints/extraction.py`
(PR #723, review threads 3916403496 and 3916473829).

T1 item 1 ("`lint_dedup_drops` returns strings, not findings"): the return
is one `_finding(...)` call with `suspect[:6]` as its evidence argument, so it
has always returned a one-element `list[dict]` (blamed to 5eb47a4, 2026-07-07,
before the #493 module split); `test_lint_dedup_drops_returns_structured_findings`
below pins the shape so a future edit cannot change it silently.

T1 item 2 / T2 item 1 (`_RENDERED_BUT_NOT_IN_RENDER_ROUTED_CODES` is a second,
hand-maintained routing source of truth): accepted as a real but pre-existing
and already-guarded design tradeoff (extraction.py's own comment discloses
it). Neither existing test in test_taxonomy_code_render_coverage.py iterates
the set's members programmatically against their render hooks --
`test_every_render_exception_code_has_a_documented_render_hook` below does.

Run:
    python3 -m pytest src/unified_pipeline/tests/test_doctor_extraction_lint_contracts.py -q -p no:cacheprovider
"""
import itertools
import json
from datetime import datetime
from collections import Counter
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.core.template_boilerplate import (  # noqa: E402
    is_near_template_instruction,
    is_template_instruction,
)
from unified_pipeline.doctor.lints.extraction import (  # noqa: E402
    CLASSIFIED_UNRENDERED_WARN_ENTRIES,
    DEDUP_SAFE_CONTAINMENT,
    INVENTED_RECORD_LICENSURE_CODE,
    INVENTED_RECORD_MIN_VALUES,
    UNDER_EXTRACTION_MAX_PCT,
    UNDER_EXTRACTION_MIN_CHARS,
    UNDER_EXTRACTION_MIN_RECORDS,
    _alphanumeric_tokens,
    _entry_rendered,
    _entry_status,
    _FUNDING_SECTIONS,
    _is_invented_record,
    _MIN_EXACT_LEN,
    _nonempty_field_values,
    _RENDERED_BUT_NOT_IN_RENDER_ROUTED_CODES,
    _rendered_row_value_sets,
    _shared_entry_pieces,
    lint_bucket_status,
    lint_classified_unrendered,
    lint_dedup_drops,
    lint_invented_records,
    lint_under_extraction,
    lint_wrong_start_date,
)
from unified_pipeline.doctor.lints import extraction as extraction_lints  # noqa: E402
from unified_pipeline.doctor.lints.extraction import (  # noqa: E402
    DEGREE_YEAR_LEAD,
    FIELD_EVIDENCE_MAX_VALUES,
    FIELD_EVIDENCE_VALUE_CHARS,
    IMPLAUSIBLE_YEAR_CEILING,
    IMPLAUSIBLE_YEAR_FLOOR,
    MULTI_RECORD_MIN_NEW,
    MULTI_RECORD_NEW_SHARE,
    MULTI_RECORD_PROSE_MIN_LOWERCASE,
    lint_implausible_year,
    lint_multi_record_coverage,
    lint_offschema_fields,
    lint_year_not_in_source,
)
from unified_pipeline.doctor.lints.extraction import (  # noqa: E402
    GRANT_ORPHAN_MIN_WORDS,
    _short_end_year,
    lint_grant_boundary,
    lint_grant_bucket,
)
from unified_pipeline.doctor.shared import (  # noqa: E402
    _LINE_SENTINEL, _haystacks, _piece_in_template, _template_haystack, docx_body_blocks)
from unified_pipeline.segmentation_regression import _norm  # noqa: E402
from unified_pipeline.stage4.schemas import STAGE4_RECORDS_KEY  # noqa: E402


# ==========================================================================
# D1 (T1.1) -- lint_dedup_drops returns structured findings, not strings.

def test_lint_dedup_drops_returns_structured_findings():
    poorly_covered = {"dedup_decisions": [
        {"code": "N1", "metric": "jaccard",
         "dropped_text": "alpha beta gamma delta", "kept_text": "alpha"},
    ]}
    result = lint_dedup_drops(poorly_covered)
    assert len(result) == 1
    finding = result[0]
    assert isinstance(finding, dict)
    # Shape per doctor/shared.py's _finding().
    assert set(finding.keys()) == {"lint", "severity", "message", "evidence",
                                   "status", "reason"}
    assert finding["lint"] == "dedup_drops"
    assert finding["severity"] == "WARN"
    assert isinstance(finding["message"], str)
    assert isinstance(finding["evidence"], list)
    assert all(isinstance(item, str) for item in finding["evidence"])
    assert len(finding["evidence"]) <= 6

    # No suspects at all -> [].
    assert lint_dedup_drops({"dedup_decisions": []}) == []

    # A drop fully covered by the kept entry is not a suspect -> [].
    well_covered = {"dedup_decisions": [
        {"code": "N1", "metric": "jaccard",
         "dropped_text": "alpha beta", "kept_text": "alpha beta"},
    ]}
    assert lint_dedup_drops(well_covered) == []

    # 7 suspects: evidence caps at 6, the message still counts all 7.
    seven = {"dedup_decisions": [
        {"code": f"N{i}", "metric": "jaccard",
         "dropped_text": f"alpha{i} beta{i} gamma{i} delta{i}",
         "kept_text": f"alpha{i}"}
        for i in range(7)
    ]}
    result7 = lint_dedup_drops(seven)
    assert len(result7) == 1
    assert len(result7[0]["evidence"]) == 6
    assert "7 dedup drop(s)" in result7[0]["message"]


# ==========================================================================
# D2 (T1.2 / T2.1) -- the render-exception set is a second source of truth
# by design (extraction.py's own comment discloses it); pin every member
# against the render hook it depends on, the way the two existing tests in
# test_taxonomy_code_render_coverage.py pin the set as a whole but do not
# iterate its members individually against generate()'s source.

_EXCEPTION_HOOKS = {
    "E": "self._fill_passthrough_sections(",
    "G": "self._fill_passthrough_sections(",
    "J": "self._fill_passthrough_sections(",
}


def test_every_render_exception_code_has_a_documented_render_hook():
    assert set(_EXCEPTION_HOOKS) == _RENDERED_BUT_NOT_IN_RENDER_ROUTED_CODES, (
        "a code was added to or removed from "
        "_RENDERED_BUT_NOT_IN_RENDER_ROUTED_CODES without updating this "
        "test's hook map -- add its render hook here")
    src = (_SRC / "unified_pipeline" / "stage_6_word_template.py").read_text()
    for code, hook in _EXCEPTION_HOOKS.items():
        assert hook in src, (
            f"{code}'s render hook {hook!r} is gone from generate() -- if "
            f"{code} no longer renders, remove it from "
            "_RENDERED_BUT_NOT_IN_RENDER_ROUTED_CODES so "
            "lint_taxonomy_code_coverage catches the gap")


# ==========================================================================
# D3 (T2.2) / D6 (T2.6, bucket_status half) -- lint_bucket_status verdicts.

_FUNDING_TITLES = dict(_FUNDING_SECTIONS)


def _blocks_under(code, text):
    """One rendered funding subsection: its real header title from
    _FUNDING_SECTIONS, followed by one content paragraph."""
    return [("p", _FUNDING_TITLES[code].title()), ("p", text)]


def test_lint_bucket_status_skips_when_unlocatable_under_any_bucket():
    # Status implies a rebucket (M2A -> M2B), but the entry text is too
    # short to verify under any bucket (all three verdicts None) -> no
    # finding, output is empty so every bucket's haystack is empty too.
    stage4 = {"entries": [
        {"taxonomy_code": "M2A", "element_idx_start": 1,
         "text": "Pending review.",
         "extracted_fields": {"status": "Completed"}},
    ]}
    assert lint_bucket_status(stage4, []) == []


_LOCATABLE_GRANT_TEXT = ("Alpha Grant Foundation Study of Longevity "
                          "Research Program")


def test_lint_bucket_status_warns_when_rendered_under_the_wrong_bucket():
    # Same entry, now locatable: status implies M2B but it actually
    # rendered under M2A -> one WARN naming M2A.
    stage4 = {"entries": [
        {"taxonomy_code": "M2A", "element_idx_start": 5,
         "text": _LOCATABLE_GRANT_TEXT,
         "extracted_fields": {"status": "Completed"}},
    ]}
    findings = lint_bucket_status(stage4, _blocks_under("M2A", _LOCATABLE_GRANT_TEXT))
    assert len(findings) == 1
    message = findings[0]["message"]
    assert "M2B" in message and "M2A" in message
    assert "current research funding" in message


def test_lint_bucket_status_unrecognised_extracted_status_does_not_mask_text():
    # Stage 4's stray value ("Awarded") must not hide the text's "Not funded".
    stage4 = {"entries": [
        {"taxonomy_code": "M2A", "element_idx_start": 7,
         "text": _LOCATABLE_GRANT_TEXT + "\nStatus: Not funded",
         "extracted_fields": {"status": "Awarded"}},
    ]}
    findings = lint_bucket_status(
        stage4, _blocks_under("M2A", _LOCATABLE_GRANT_TEXT))
    assert len(findings) == 1
    assert "M2C" in findings[0]["message"]


def test_lint_bucket_status_silent_when_rebucketed_correctly():
    # Rendered under the bucket the status implies (M2B) -> stage 6 already
    # rebucketed it correctly, no finding.
    stage4 = {"entries": [
        {"taxonomy_code": "M2A", "element_idx_start": 5,
         "text": _LOCATABLE_GRANT_TEXT,
         "extracted_fields": {"status": "Completed"}},
    ]}
    findings = lint_bucket_status(stage4, _blocks_under("M2B", _LOCATABLE_GRANT_TEXT))
    assert findings == []


def test_lint_bucket_status_warns_when_target_unlocatable_but_another_bucket_hits():
    # D6 mixed verdict case: the target bucket (M2B) is unlocatable (None,
    # not False), but the entry IS locatable under a different bucket
    # (M2A) -- must still warn, and must name the bucket it actually found
    # it under, not the unlocatable target.
    text = "Award 12345 total 67890 sum"  # only 2 long-word tokens: unverifiable via the token path
    stage4 = {"entries": [
        {"taxonomy_code": "M2C", "element_idx_start": 9, "text": text,
         "extracted_fields": {"status": "Completed"}},
    ]}
    findings = lint_bucket_status(stage4, _blocks_under("M2A", text))
    assert len(findings) == 1
    assert "M2A" in findings[0]["message"]


# ==========================================================================
# D4 (T2.3) -- _entry_status: what the status label regex actually captures.

@pytest.mark.parametrize("entry,expected", [
    ({"text": "Status: Funded"}, "Funded"),
    ({"text": "status - Not funded"}, "Not funded"),
    # The capture runs to the next '|' or newline, not to the first hyphen.
    ({"text": "Status: Funded - 2020"}, "Funded - 2020"),
    ({"text": "Status: Funded | PI: Smith"}, "Funded"),
    ({"text": "Some header line\nStatus: Pending review\nmore text"},
     "Pending review"),
    # A recognised extracted_fields.status wins over any text label (#720).
    ({"text": "Status: Funded", "extracted_fields": {"status": "Completed"}},
     "Completed"),
    ({"text": "Status: Completed",
      "extracted_fields": {"status": "Under review"}}, "Under review"),
    ({"text": "x", "extracted_fields": {"status": "Not funded"}}, "Not funded"),
    # An unrecognised extracted status falls back to the text label (#720).
    ({"text": "Status: Not funded", "extracted_fields": {"status": "Awarded"}},
     "Not funded"),
    ({"text": "Status: Pending review",
      "extracted_fields": {"status": "zzz"}}, "Pending review"),
    # Unrecognised and no text label: nothing to report.
    ({"text": "no label", "extracted_fields": {"status": "zzz"}}, None),
    ({"text": "Status: Funded", "extracted_fields": {"status": ""}}, "Funded"),
    ({"text": "No label here at all"}, None),
])
def test_entry_status_label_extraction(entry, expected):
    assert _entry_status(entry) == expected


# ==========================================================================
# D5 (T2.5) -- adversarial _entry_rendered cases.

def test_entry_rendered_whitespace_and_punctuation_differences_still_match():
    entry_text = ("A Longitudinal Study of Diabetes Outcomes in Rural "
                  "Populations")
    rendered_blocks = [
        ("p", "A   Longitudinal Study of Diabetes Outcomes, in Rural "
              "Populations!!!"),
    ]
    h = _haystacks(rendered_blocks)
    assert _entry_rendered(entry_text, h.text, h.tokens) is True


def test_entry_rendered_low_overlap_long_entry_is_definitively_false():
    entry_text = ("Alpha Bravo Charlie Delta Echo Foxtrot Golf Hotel India "
                  "Juliet")
    unrelated_blocks = [
        ("p", "Kilo Lima Mike November Oscar Papa Quebec Romeo Sierra "
              "Tango"),
    ]
    h = _haystacks(unrelated_blocks)
    assert _entry_rendered(entry_text, h.text, h.tokens) is False


def test_entry_rendered_single_long_word_is_unverifiable():
    entry_text = "Pneumonoultramicroscopicsilicovolcanoconiosis"
    h = _haystacks([("p", "unrelated text here")])
    assert _entry_rendered(entry_text, h.text, h.tokens) is None


_BOILERPLATE_HAYSTACK_BLOCKS = [
    ("p", "Jane Smith, MD is affiliated with the Department of "
          "Medicine at a different institution entirely, working on "
          "unrelated topics."),
]


def test_entry_rendered_shared_institutional_boilerplate_is_not_rendered():
    """#744: 'Department of Medicine' occurs in two entries of the document,
    so its verbatim hit in an unrelated paragraph is not evidence."""
    siblings = [{"text": "Department of Medicine"},
                {"text": "Department of Medicine\nDivision of Cardiology"}]
    shared = _shared_entry_pieces(siblings)
    h = _haystacks(_BOILERPLATE_HAYSTACK_BLOCKS)
    assert _entry_rendered("Department of Medicine", h.text, h.tokens,
                           shared) is False


def test_entry_rendered_unshared_short_piece_still_counts_as_rendered():
    """Control for the above: with no sibling repeating the piece (and the
    piece absent from the template) the verbatim hit still counts."""
    shared = _shared_entry_pieces([{"text": "Department of Medicine"},
                                   {"text": "Division of Cardiology"}])
    h = _haystacks(_BOILERPLATE_HAYSTACK_BLOCKS)
    assert _entry_rendered("Department of Medicine", h.text, h.tokens,
                           shared) is True


def test_entry_rendered_template_scaffolding_piece_is_not_rendered():
    """#744: 'Full Name of Board' is a cell of the pristine WCM template, so
    finding it in the output proves nothing about this entry."""
    h = _haystacks([("p", "Certified by the Full Name of Board of Surgery")])
    assert _entry_rendered("Full Name of Board", h.text, h.tokens) is False


def test_shared_entry_pieces_counts_distinct_entries_not_occurrences():
    """One entry repeating a fragment twice is one entry: not shared."""
    entries = [{"text": "Department of Medicine\nDepartment of Medicine"},
               {"text": "Unrelated Other Entry Text Here"}]
    assert _shared_entry_pieces(entries) == frozenset()


def test_shared_entry_pieces_identical_text_duplicates_are_not_boilerplate():
    """Two entries with identical text are one record listed twice (stage 4
    dedups them), so their common pieces are content, not boilerplate."""
    entries = [{"text": "Ad hoc Reviewer, Journal of Neurosurgery"},
               {"text": "Ad hoc Reviewer, Journal of Neurosurgery"}]
    assert _shared_entry_pieces(entries) == frozenset()


def test_entry_rendered_boilerplate_label_with_short_value_is_unverifiable():
    """A template label plus a value too short to be a piece: the value is
    content _entry_rendered cannot see, so the verdict is None, not False."""
    h = _haystacks([("p", "Office telephone: on file")])
    assert _entry_rendered("Office telephone: | 212 555 0100",
                           h.text, h.tokens) is None


def test_entry_rendered_boilerplate_piece_absent_from_output_stays_unverifiable():
    """The discount only revokes a hit that made the entry look rendered; a
    shared piece the output never contained is no more definitive than before."""
    shared = _shared_entry_pieces([{"text": "Epic Implementation team"},
                                   {"text": "Epic Implementation team\nLead"}])
    h = _haystacks([("p", "Unrelated paragraph about something else")])
    assert _entry_rendered("Epic Implementation team", h.text, h.tokens,
                           shared) is None


def test_entry_rendered_distinctive_piece_absent_beside_boilerplate_hit_is_unverifiable():
    """A boilerplate hit does not make an entry lost when another piece is
    distinctive: that piece is simply absent, which (with too few tokens) is
    the pre-#744 'cannot verify' answer."""
    text = "Department of Medicine\nab cd ef gh ij kl mn op"
    shared = _shared_entry_pieces([{"text": "Department of Medicine"},
                                   {"text": text}])
    h = _haystacks(_BOILERPLATE_HAYSTACK_BLOCKS)
    assert _entry_rendered(text, h.text, h.tokens, shared) is None


def test_entry_rendered_template_piece_matches_inside_a_longer_template_line():
    """The template check is containment: an entry piece is a 40-char prefix
    of a longer template paragraph, not a whole line. The output glues the
    words so only the verbatim piece, never a token, could vouch for it."""
    text = "Percent Effort and Institutional Responsibilities"
    h = _haystacks([("p", "PercentEffortandInstitutionalResponsibilities")])
    assert _entry_rendered(text, h.text, h.tokens) is False


def test_template_piece_must_sit_inside_one_template_line():
    """A piece spanning the end of one template line and the start of the
    next is not scaffolding: 5 farm pieces match only across a boundary."""
    lines = _template_haystack().split(_LINE_SENTINEL)
    spans = [a[-8:] + b[:8] for a, b in zip(lines, lines[1:])
             if len(a) >= 8 and len(b) >= 8]
    spans = [p for p in spans
             if p not in _template_haystack().replace(_LINE_SENTINEL, "|")]
    assert spans
    assert not any(_piece_in_template(p) for p in spans)


def test_lint_bucket_status_does_not_trust_a_shared_boilerplate_hit():
    """#744 at bucket_status: the entry's only piece also opens a sibling
    entry, so its verbatim hit under the target bucket is not evidence."""
    text = "Department of Medicine"
    stage4 = {"entries": [
        {"taxonomy_code": "M2A", "element_idx_start": 5, "text": text,
         "extracted_fields": {"status": "Completed"}},
        {"taxonomy_code": "M2A", "element_idx_start": 6,
         "text": text + "\nDivision of Cardiology"},
    ]}
    findings = lint_bucket_status(stage4, _blocks_under("M2B", text))
    assert len(findings) == 1
    assert "no funding heading at all" in findings[0]["message"]


# ==========================================================================
# D6 (T2.6) -- None/False/True semantics at lint_classified_unrendered.

_UNVERIFIABLE_TEXT = "Pending review."
_UNRENDERED_TEXT = ("Distinguished Career Achievement Award for Excellence "
                    "in Translational Science")
_RENDERED_TEXT = ("Groundbreaking Multicenter Trial of Novel Cardiac "
                  "Regeneration Therapy")
_OUTPUT_BLOCKS = [("p", "SOME HEADER"), ("p", _RENDERED_TEXT)]


def _entries(*texts, code="X"):
    return {"entries": [
        {"text": t, "taxonomy_code": code, "element_type": "paragraph"}
        for t in texts
    ]}


def test_lint_classified_unrendered_silent_when_all_entries_unverifiable():
    findings = lint_classified_unrendered(
        _entries(_UNVERIFIABLE_TEXT), _OUTPUT_BLOCKS)
    assert findings == []


def test_lint_classified_unrendered_warns_on_false_plus_none():
    findings = lint_classified_unrendered(
        _entries(_UNVERIFIABLE_TEXT, _UNRENDERED_TEXT), _OUTPUT_BLOCKS)
    assert len(findings) == 1
    assert "2 classified" in findings[0]["message"]


def test_lint_classified_unrendered_skips_m1_the_research_summary_replaces():
    """YTPMZK: stage 6 renders the research summary in place of M1's own
    text, so M1 never renders verbatim when the summary does -- and goes to
    the Appendix, where it IS seen, when it doesn't. The same unrendered text
    under any other code is still flagged."""
    assert lint_classified_unrendered(_entries(_UNRENDERED_TEXT, code="M1"), _OUTPUT_BLOCKS) == []
    assert len(lint_classified_unrendered(_entries(_UNRENDERED_TEXT, code="M2A"), _OUTPUT_BLOCKS)) == 1


def test_lint_classified_unrendered_does_not_trust_a_shared_boilerplate_hit():
    """#744 at the lint: two 3b entries share the piece 'Department of
    Medicine'; the output holds only that piece, so neither code's entry has
    real evidence and both codes are flagged (pre-fix: [])."""
    stage3b = {"entries": [
        {"text": "Department of Medicine", "taxonomy_code": "C",
         "element_type": "paragraph"},
        {"text": "Department of Medicine\nDivision of Cardiology",
         "taxonomy_code": "D", "element_type": "paragraph"},
    ]}
    findings = lint_classified_unrendered(
        stage3b, [("p", "Department of Medicine")])
    assert sorted(f["message"].split(":")[0] for f in findings[1:]) == [
        "taxonomy code C", "taxonomy code D"]


def test_lint_classified_unrendered_silent_when_any_entry_verified_rendered():
    findings = lint_classified_unrendered(
        _entries(_RENDERED_TEXT, _UNRENDERED_TEXT), _OUTPUT_BLOCKS)
    assert findings == []


# ==========================================================================
# D7 (T2.4) -- exact thresholds.

def test_dedup_safe_containment_exact_boundary():
    assert DEDUP_SAFE_CONTAINMENT == 0.9
    dropped = " ".join(f"tok{i}" for i in range(10))
    kept_90pct = " ".join(f"tok{i}" for i in range(9))   # 9/10 = 0.9
    kept_80pct = " ".join(f"tok{i}" for i in range(8))   # 8/10 = 0.8
    report_90 = {"dedup_decisions": [
        {"code": "A", "metric": "j", "dropped_text": dropped,
         "kept_text": kept_90pct}]}
    report_80 = {"dedup_decisions": [
        {"code": "A", "metric": "j", "dropped_text": dropped,
         "kept_text": kept_80pct}]}
    assert lint_dedup_drops(report_90) == []          # exactly at the safe line
    assert len(lint_dedup_drops(report_80)) == 1       # just below -> suspect


def _distinctive_missing_entry(i):
    return {"text": (f"Distinguished Career Achievement Award Number {i} "
                     "for Excellence in Translational Science Research"),
            "taxonomy_code": "X", "element_type": "paragraph"}


@pytest.mark.parametrize("lost,expected_severity", [
    (CLASSIFIED_UNRENDERED_WARN_ENTRIES - 1, "INFO"),
    (CLASSIFIED_UNRENDERED_WARN_ENTRIES, "WARN"),
    (CLASSIFIED_UNRENDERED_WARN_ENTRIES + 1, "WARN"),
])
def test_classified_unrendered_warn_entries_boundary(lost, expected_severity):
    assert CLASSIFIED_UNRENDERED_WARN_ENTRIES == 2
    stage3b = {"entries": [_distinctive_missing_entry(i) for i in range(lost)]}
    findings = lint_classified_unrendered(stage3b, _OUTPUT_BLOCKS)
    assert len(findings) == 1
    assert findings[0]["severity"] == expected_severity


def test_classified_unrendered_one_run_level_warn_over_per_code_info():
    """#719: several codes lost -> ONE run-level finding carries the magnitude
    severity (entries lost across all codes, #438) and each per-code finding
    is INFO evidence. Pre-fix every per-code finding was stamped WARN, so the
    code that lost one entry read WARN because code A lost three."""
    lost_a = [_distinctive_missing_entry(i) for i in range(3)]
    lost_b = [{**_distinctive_missing_entry(9), "taxonomy_code": "B"}]
    for e in lost_a:
        e["taxonomy_code"] = "A"
    findings = lint_classified_unrendered(
        {"entries": lost_a + lost_b}, _OUTPUT_BLOCKS)
    run_level, *per_code = findings
    assert run_level["lint"] == "classified_unrendered"
    assert run_level["severity"] == "WARN"
    assert run_level["message"] == ("4 classified entries across 2 taxonomy "
                                    "codes appear nowhere in the output document")
    assert run_level["evidence"] == ["taxonomy code A", "taxonomy code B"]
    assert [f["message"].split(":")[0] for f in per_code] == [
        "taxonomy code A", "taxonomy code B"]
    assert [f["severity"] for f in per_code] == ["INFO", "INFO"]


def test_classified_unrendered_run_level_finding_tracks_the_magnitude(monkeypatch):
    """#719: the run-level finding's severity is the #438 magnitude, not a
    fixed WARN -- below the entries-lost threshold it is INFO too. The
    threshold is raised because two lost codes already meet the real one."""
    import unified_pipeline.doctor.lints.extraction as extraction
    monkeypatch.setattr(extraction, "CLASSIFIED_UNRENDERED_WARN_ENTRIES", 3)
    stage3b = {"entries": [{**_distinctive_missing_entry(i), "taxonomy_code": c}
                           for i, c in enumerate("AB")]}
    findings = lint_classified_unrendered(stage3b, _OUTPUT_BLOCKS)
    assert [f["severity"] for f in findings] == ["INFO", "INFO", "INFO"]


_UNDER_EXTRACTION_RECORD_LINE = ("{year} Distinguished award for excellence "
                                 "in service to community broadly and "
                                 "repeatedly.")


def _padded_under_extraction_text(n_records, target_len):
    lines = [_UNDER_EXTRACTION_RECORD_LINE.format(year=2000 + i)
             for i in range(n_records)]
    text = "\n".join(lines)
    assert len(text) <= target_len, "record lines alone exceed target_len"
    pad = target_len - len(text)
    if pad > 0:
        text = text + "\n" + ("z" * (pad - 1))
    assert len(text) == target_len
    return text


def _under_extraction_entry(pct, total_len, n_records):
    return {"element_type": "paragraph", "element_idx_start": 1,
            "extraction_coverage": {"extraction_coverage_percent": pct},
            "text": _padded_under_extraction_text(n_records, total_len)}


def test_under_extraction_max_pct_boundary():
    assert UNDER_EXTRACTION_MAX_PCT == 40.0
    below = _under_extraction_entry(UNDER_EXTRACTION_MAX_PCT - 0.01,
                                    UNDER_EXTRACTION_MIN_CHARS + 50, 3)
    at = _under_extraction_entry(UNDER_EXTRACTION_MAX_PCT,
                                 UNDER_EXTRACTION_MIN_CHARS + 50, 3)
    assert len(lint_under_extraction({"entries": [below]})) == 1
    assert lint_under_extraction({"entries": [at]}) == []  # >= threshold is excluded


def test_under_extraction_min_chars_boundary():
    assert UNDER_EXTRACTION_MIN_CHARS == 800
    at = _under_extraction_entry(30.0, UNDER_EXTRACTION_MIN_CHARS, 3)
    above = _under_extraction_entry(30.0, UNDER_EXTRACTION_MIN_CHARS + 1, 3)
    assert lint_under_extraction({"entries": [at]}) == []  # <= threshold is excluded
    assert len(lint_under_extraction({"entries": [above]})) == 1


def test_under_extraction_min_records_boundary():
    assert UNDER_EXTRACTION_MIN_RECORDS == 2
    below = _under_extraction_entry(30.0, UNDER_EXTRACTION_MIN_CHARS + 50,
                                    UNDER_EXTRACTION_MIN_RECORDS - 1)
    at = _under_extraction_entry(30.0, UNDER_EXTRACTION_MIN_CHARS + 50,
                                 UNDER_EXTRACTION_MIN_RECORDS)
    assert lint_under_extraction({"entries": [below]}) == []  # < threshold is excluded
    assert len(lint_under_extraction({"entries": [at]})) == 1



def _stage4_records_entry(code, field):
    """An entry whose two stage-4 records carry its whole text; its own
    field (the last record, as stage 4 leaves it) carries one line."""
    text = _padded_under_extraction_text(4, UNDER_EXTRACTION_MIN_CHARS + 50)
    lines = text.split("\n")
    records = [{field: "\n".join(lines[:2])}, {field: "\n".join(lines[2:])}]
    return {"element_type": "paragraph", "element_idx_start": 1, "taxonomy_code": code,
            "text": text, "extracted_fields": {**records[-1], STAGE4_RECORDS_KEY: records},
            "extraction_coverage": {"extraction_coverage_percent": 100.0}}


def test_under_extraction_ignores_stage4_records_no_renderer_writes():
    """#1299: T has no fan-out renderer, so its stage-4 records never
    render and must not count toward coverage."""
    entry = _stage4_records_entry("T", "title")
    entry["extracted_fields"] = {"title": "zz", STAGE4_RECORDS_KEY:
                                 entry["extracted_fields"][STAGE4_RECORDS_KEY]}
    assert len(lint_under_extraction({"entries": [entry]})) == 1


def test_under_extraction_credits_stage4_records_that_fan_out():
    assert lint_under_extraction({"entries": [_stage4_records_entry("H", "award_name")]}) == []


_OVERLAP_WORDS = [''.join(c) for c in itertools.islice(
    itertools.product('abcdefghijklmnopqrstuvwxyz', repeat=5), 100)]
_OVERLAP_ENTRY_TEXT = " ".join(_OVERLAP_WORDS)


def _overlap_haystack(matched_n):
    # Reversed order so the entry's own leading squashed piece (its first
    # ~8 words) never appears as a literal substring here -- this isolates
    # the token-overlap path from verbatim containment.
    matched = list(reversed(_OVERLAP_WORDS[:matched_n]))
    return _haystacks([("p", "SOME HEADER"),
                       ("p", " ".join(matched) + " filler filler filler")])


def test_entry_rendered_token_overlap_exact_boundary():
    from unified_pipeline.doctor.shared import RENDER_TOKEN_OVERLAP
    assert RENDER_TOKEN_OVERLAP == 0.7
    at_boundary = _overlap_haystack(70)     # 70/100 = 0.70
    just_below = _overlap_haystack(69)      # 69/100 = 0.69
    assert _entry_rendered(_OVERLAP_ENTRY_TEXT, at_boundary.text,
                           at_boundary.tokens) is True
    assert _entry_rendered(_OVERLAP_ENTRY_TEXT, just_below.text,
                           just_below.tokens) is False


# ==========================================================================
# D8 (T2.7) -- dedup tokenization: repetition, Unicode, digits, punctuation.

def test_alphanumeric_tokens_repeated_tokens_keep_multiplicity():
    # #718: the token view is a multiset, so a dropped entry that repeats a
    # token is NOT fully covered by a kept entry that says it once.
    dropped = {"code": "A", "metric": "j",
              "dropped_text": "grant grant grant", "kept_text": "grant"}
    result = lint_dedup_drops({"dedup_decisions": [dropped]})
    assert len(result) == 1
    assert "33% covered" in result[0]["evidence"][0]
    assert _alphanumeric_tokens("grant grant grant") == Counter({"grant": 3})
    assert _alphanumeric_tokens("grant") == Counter({"grant": 1})


def test_lint_dedup_drops_repeated_tokens_fully_present_still_safe():
    # The true-duplicate case is unchanged: same multiset -> 100%.
    same = {"code": "A", "metric": "j",
            "dropped_text": "grant grant funded", "kept_text": "funded grant grant x"}
    assert lint_dedup_drops({"dedup_decisions": [same]}) == []


# #666: a fully-covered drop whose extracted name differs from the kept entry's
# and is on the page as no cell of its own.

def _named_decision(dropped_name, kept_name):
    return {"code": "Q4D", "metric": "containment=1.00",
            "dropped_text": dropped_name, "kept_text": kept_name,
            "dropped_fields": {"journal_name": dropped_name},
            "kept_fields": {"journal_name": kept_name}}


def test_dedup_drops_info_when_dropped_name_absent_from_page():
    report = {"dedup_decisions": [_named_decision("Example Optics", "European Example Optics")]}
    blocks = [("p", "Reviewer"), ("table", "European Example Optics\nJournal Reviewer")]
    result = lint_dedup_drops(report, blocks)
    assert len(result) == 1
    assert result[0]["severity"] == "INFO" and result[0]["lint"] == "dedup_drops"
    assert "Q4D" in result[0]["evidence"][0]


def test_dedup_drops_quiet_when_dropped_name_is_its_own_cell():
    report = {"dedup_decisions": [_named_decision("Example Optics", "European Example Optics")]}
    blocks = [("table", "European Example Optics\nExample Optics")]
    assert lint_dedup_drops(report, blocks) == []


def test_dedup_drops_quiet_when_names_match_after_normalisation():
    report = {"dedup_decisions": [_named_decision("Example Optics", "example  optics.")]}
    assert lint_dedup_drops(report, [("p", "unrelated")]) == []


def test_dedup_drops_info_needs_rendered_blocks_and_fields():
    report = {"dedup_decisions": [_named_decision("Example Optics", "European Example Optics")]}
    assert lint_dedup_drops(report) == []
    legacy = {"dedup_decisions": [{"code": "Q4D", "metric": "m",
                                   "dropped_text": "Example Optics",
                                   "kept_text": "European Example Optics"}]}
    assert lint_dedup_drops(legacy, [("p", "unrelated")]) == []


def test_dedup_drops_info_uses_institution_for_appointments():
    decision = {"code": "D2", "metric": "jaccard=0.82",
                "dropped_text": "Example Hospital Example Title 2001",
                "kept_text": "Sample Hospital Example Hospital Example Title 2001",
                "dropped_fields": {"institution": "Example Hospital"},
                "kept_fields": {"institution": "Sample Hospital"}}
    result = lint_dedup_drops({"dedup_decisions": [decision]},
                              [("table", "Sample Hospital\nExample Title")])
    assert [f["severity"] for f in result] == ["INFO"]


def test_dedup_drops_quiet_when_only_the_dropped_entry_fills_a_name():
    decision = _named_decision("Example Optics", "European Example Optics")
    decision["kept_fields"] = {}
    assert lint_dedup_drops({"dedup_decisions": [decision]}, [("p", "x")]) == []


def test_dedup_drops_info_evidence_is_capped():
    report = {"dedup_decisions": [
        _named_decision(f"Example Optics {n}", f"European Example Optics {n} Letters")
        for n in range(extraction_lints.DEDUP_EVIDENCE_LIMIT + 3)]}
    result = lint_dedup_drops(report, [("p", "x")])
    assert len(result[0]["evidence"]) == extraction_lints.DEDUP_EVIDENCE_LIMIT
    assert str(extraction_lints.DEDUP_EVIDENCE_LIMIT + 3) in result[0]["message"]


def test_dedup_drops_reports_warn_and_info_together():
    poorly = {"code": "N1", "metric": "jaccard",
              "dropped_text": "alpha beta gamma delta", "kept_text": "alpha"}
    report = {"dedup_decisions": [poorly, _named_decision("Example Optics", "European Example Optics")]}
    result = lint_dedup_drops(report, [("p", "x")])
    assert [f["severity"] for f in result] == ["WARN", "INFO"]


# #666 (EBYSBC E4): a drop that names another occasion -- a month, a day, a
# part or a numeral the kept entry lacks -- is another record even when its
# words are all in the kept text. Synthetic dates and titles.

def _drop(dropped, kept, code="K4", metric="jaccard=0.95", **extra):
    return {"code": code, "metric": metric, "dropped_text": dropped,
            "kept_text": kept, **extra}


@pytest.mark.parametrize("dropped, kept", [
    # the same lecture in another month: every dropped word is in the kept text
    ("3/2031 Example grand rounds talk", "9/2031 Example grand rounds talk, 3 hours"),
    # another day of the same month, a two-digit year
    ("5/9/31 Example teaching session", "5/16/31 Example teaching session 9"),
])
def test_dedup_drops_warns_on_a_date_the_kept_entry_lacks(dropped, kept):
    result = lint_dedup_drops({"dedup_decisions": [_drop(dropped, kept)]})
    assert [f["severity"] for f in result] == ["WARN"]
    assert "a date the kept entry does not carry" in result[0]["evidence"][0]


def test_dedup_drops_quiet_when_the_kept_entry_carries_the_date():
    decision = _drop("3/2031 Example grand rounds talk",
                     "3/2031 Example grand rounds talk, Example Hospital")
    assert lint_dedup_drops({"dedup_decisions": [decision]}) == []


@pytest.mark.parametrize("dropped, kept, reason", [
    ("Example Series Part I, noon talk", "Example Series Part II, noon talk", "part 1"),
    ("Example Program II", "Example Program III", "numeral II"),
    ("Example Course I\tExample Society", "Example Course II\tExample Society", "numeral I"),
    ("Example trial, Phase II", "Example trial, Phase III", "phase 2"),
])
def test_occasion_apart_reads_parts_and_numerals(dropped, kept, reason):
    assert extraction_lints._occasion_apart(dropped, kept) == f"{reason} the kept entry lacks"


@pytest.mark.parametrize("dropped, kept", [
    ("I taught an example course", "Taught an example course"),     # a pronoun
    ("Doe V, Roe X. Example paper", "Doe V. Example paper"),       # initials
    ("Example Series Part I", "Example Series Parts I and II"),     # a listing covers it
])
def test_occasion_apart_ignores_pronouns_initials_and_listed_parts(dropped, kept):
    assert extraction_lints._occasion_apart(dropped, kept) is None


def test_dedup_drops_occasion_test_also_checks_record_rule_drops():
    """A record-rule drop skips the word-coverage test (its texts share few
    words by construction) but not the occasion test."""
    decision = _drop("Example Series Part I", "Example Series Part II overview",
                     metric="record=example_rule")
    result = lint_dedup_drops({"dedup_decisions": [decision]})
    assert [f["severity"] for f in result] == ["WARN"]
    assert "record=example_rule, part 1" in result[0]["evidence"][0]


def test_dedup_drops_info_on_a_rank_word_even_when_the_name_renders():
    """SEKQUI-01's shape: the dropped rank renders elsewhere (another row), so
    the cell test alone stays quiet; the rank word tells the records apart."""
    decision = _drop("Assistant Professor, Example Dept", "Clinical Assistant Professor, "
                     "Example Dept, 2001", code="D1", metric="containment=1.00",
                     dropped_fields={"title": "Assistant Professor"},
                     kept_fields={"title": "Clinical Assistant Professor"})
    blocks = [("table", "Assistant Professor\nClinical Assistant Professor")]
    result = lint_dedup_drops({"dedup_decisions": [decision]}, blocks)
    assert [f["severity"] for f in result] == ["INFO"]
    assert lint_dedup_drops({"dedup_decisions": [decision]}) == result


def test_dedup_drops_evidence_names_the_dropped_entry_from_stage_5d():
    decision = _drop("3/2031 Example grand rounds talk", "9/2031 Example grand rounds talk, 3")
    stage_5d = {"entries": [
        {"element_idx_start": 12, "taxonomy_code": "K4",
         "text": "3/2031  Example grand rounds talk", "extracted_fields": {}},
        {"element_idx_start": 14, "taxonomy_code": "K4",
         "text": "9/2031 Example grand rounds talk, 3", "extracted_fields": {}}]}
    result = lint_dedup_drops({"dedup_decisions": [decision]}, None, stage_5d)
    assert result[0]["evidence"][0].startswith("entry 12: K4 (")


def test_dedup_drops_evidence_names_the_entry_by_its_formatted_citation():
    """A citation's dropped text is its stage-5d formatted_citation, not its
    raw text."""
    decision = _drop("3/2031 Example grand rounds talk", "9/2031 Example grand rounds talk, 3")
    stage_5d = {"entries": [
        {"element_idx_start": 12, "taxonomy_code": "K4", "text": "raw source line 12",
         "extracted_fields": {"formatted_citation": "3/2031 Example grand rounds talk"}}]}
    result = lint_dedup_drops({"dedup_decisions": [decision]}, None, stage_5d)
    assert result[0]["evidence"][0].startswith("entry 12: K4 (")


def test_dedup_drops_evidence_names_no_entry_when_two_carry_the_text():
    decision = _drop("3/2031 Example grand rounds talk", "9/2031 Example grand rounds talk, 3")
    twin = {"taxonomy_code": "K4", "text": "3/2031 Example grand rounds talk",
            "extracted_fields": {}}
    stage_5d = {"entries": [dict(twin, element_idx_start=12),
                            dict(twin, element_idx_start=20)]}
    result = lint_dedup_drops({"dedup_decisions": [decision]}, None, stage_5d)
    assert result[0]["evidence"][0].startswith("K4 (")


def test_dedup_drops_record_rule_drop_skips_the_coverage_warn_not_the_name_info():
    """A stage-6 record-rule drop matched on fields: its low text coverage
    is not evidence, but a dropped name absent from the page still is."""
    header = {"code": "D1", "metric": "record=undated_appointment",
              "dropped_text": "Title: Associate Professor",
              "kept_text": "2031-present Associate Prof., Dept. of Widgets"}
    assert lint_dedup_drops({"dedup_decisions": [header]}, [("p", "x")]) == []
    named = dict(header, dropped_fields={"title": "Associate Professor"},
                 kept_fields={"title": "Associate Prof."})
    result = lint_dedup_drops({"dedup_decisions": [named]}, [("p", "Associate Prof.")])
    assert [f["severity"] for f in result] == ["INFO"]
    text_drop = dict(header, metric="containment=0.80")
    assert [f["severity"] for f in lint_dedup_drops({"dedup_decisions": [text_drop]})] == ["WARN"]


def test_alphanumeric_tokens_unicode_diaeresis_folds_to_one_token():
    # #541: _norm folds combining marks and _DEDUP_TOKEN_RE is Unicode-aware,
    # so "müller" is one token (it used to split into "m" + "ller").
    assert _alphanumeric_tokens("Müller") == Counter({"muller": 1})


def test_alphanumeric_tokens_digits_are_tokens_punctuation_is_not():
    assert _alphanumeric_tokens("Grant 2020") == Counter({"grant": 1, "2020": 1})
    assert _alphanumeric_tokens("grant, funded!") == Counter({"grant": 1, "funded": 1})


# ==========================================================================
# D9 -- lint_invented_records (#829): a record built from the WCM template's
# own labels rather than real content, reaching the delivered document.
#
# "Full Name of Board" / "Certificate #" / "Dates of Certification" are the
# WCM board-certification template's own column labels, already committed in
# `template_boilerplate_phrases.json` -- template scaffolding text, not a
# verbatim line from any real CV. Same for the licensure paragraph below,
# which is the tracked template's own instruction wording with one word
# swapped ("ABC" for the real template's own hospital name) so the near-match
# path is exercised without reproducing any real CV's text.

_BOARD_LABEL = "Full Name of Board"
_CERT_NUMBER_LABEL = "Certificate #"
_CERT_DATE_LABEL = "Dates of Certification"

_HEADER_ROW = [_BOARD_LABEL, f"{_CERT_NUMBER_LABEL} \n(indicate if board eligible)",
              f"{_CERT_DATE_LABEL} \n(yyyy–yyyy)"]
_INVENTED_CERT_ROW = [_BOARD_LABEL, _CERT_NUMBER_LABEL, ""]
_REAL_CERT_ROW = ["American Board of Internal Medicine", "123456", "2015"]

_NEAR_MATCH_LICENSURE_TEXT = (
    "(Every doctor appointed to the ABC Hospital staff, except interns and "
    "aliens in the US via non-immigrant visas, must have a New York State "
    "license or a temporary certificate in lieu of the license.)")


def _invented_cert_entry(element_idx_start=1):
    return {"taxonomy_code": "F2", "element_type": "table_row",
            "element_idx_start": element_idx_start,
            "text": f"{_BOARD_LABEL} | {_CERT_NUMBER_LABEL} | {_CERT_DATE_LABEL}",
            "extracted_fields": {
                "certifying_board": _BOARD_LABEL,
                "certificate_number": _CERT_NUMBER_LABEL,
                "year_certified": None, "recertification_date": None}}


def _real_cert_entry(element_idx_start=1):
    return {"taxonomy_code": "F2", "element_type": "table_row",
            "element_idx_start": element_idx_start,
            "text": "American Board of Internal Medicine | 123456 | 2015",
            "extracted_fields": {
                "certifying_board": "American Board of Internal Medicine",
                "certificate_number": "123456",
                "year_certified": "2015", "recertification_date": None}}


def _invented_licensure_entry(element_idx_start=1, element_type="break",
                              text=_NEAR_MATCH_LICENSURE_TEXT):
    return {"taxonomy_code": "F1", "element_type": element_type,
            "element_idx_start": element_idx_start, "text": text,
            "extracted_fields": {"state_country": "New York State",
                                 "license_number": None, "issue_date": None,
                                 "expiration_date": None}}


def test_invented_record_min_values_is_2_and_licensure_code_is_f1():
    assert INVENTED_RECORD_MIN_VALUES == 2
    assert INVENTED_RECORD_LICENSURE_CODE == "F1"
    assert _MIN_EXACT_LEN == 25


# -- _nonempty_field_values ------------------------------------------------

def test_nonempty_field_values_flattens_lists_and_drops_blanks_and_dicts():
    fields = {
        "single": "Alpha",
        "blank": "",
        "whitespace_only": "   ",
        "absent": None,
        "listed": ["Beta", "", None, "Gamma"],
        "nested": {"unexpected": "should not be flattened"},
    }
    assert _nonempty_field_values(fields) == ["Alpha", "Beta", "Gamma"]
    assert _nonempty_field_values({}) == []


# -- _is_invented_record ----------------------------------------------------

def test_is_invented_record_true_on_the_board_certification_header_row():
    assert _is_invented_record(_invented_cert_entry()["extracted_fields"])


def test_is_invented_record_false_on_real_board_certification_values():
    assert not _is_invented_record(_real_cert_entry()["extracted_fields"])


def test_is_invented_record_false_below_the_min_values_floor():
    # One matched label alone, however distinctive, is not enough evidence
    # (the G/Institutional-Affiliation shape: a single unfilled prompt label
    # legitimately echoes back with a blank companion cell).
    assert not _is_invented_record({"affiliation_type": _BOARD_LABEL})
    # Isolate the count guard from the length guard: "Board / Organization
    # Name" alone is exactly 25 chars (== _MIN_EXACT_LEN), long enough on
    # its own to pass the length floor, so this fails ONLY on value count.
    long_single_label = "Board / Organization Name"
    assert len(long_single_label) == _MIN_EXACT_LEN
    assert not _is_invented_record({"granting_body": long_single_label})


def test_is_invented_record_false_below_the_combined_length_floor():
    # "Total"/"100%" are themselves registered template phrases (the blank
    # %-effort table's own worked example), and a real, fully-filled J table
    # legitimately ends in a "Total | 100%" row -- 2 matched values, but only
    # 9 combined chars, well under _MIN_EXACT_LEN (25).
    assert not _is_invented_record({"activity": "Total", "percent_effort": "100%"})


def test_is_invented_record_true_at_the_exact_combined_length_floor():
    # "Organization" (12 chars) and "Bibliography" (12 chars) are each their
    # own registered template label -- raw lengths sum to 24, and the "|"
    # the code joins them with brings the combined length to exactly 25
    # (== _MIN_EXACT_LEN), the boundary the `<` comparison must accept.
    # A `<` -> `<=` mutant on the length check, or a join separator swapped
    # from "|" to "" (which drops this to 24), each flip this to False.
    a, b = "Organization", "Bibliography"
    assert len(a) + len(b) == 24
    assert len(f"{a}|{b}") == _MIN_EXACT_LEN
    assert _is_invented_record({"field_one": a, "field_two": b})


def test_is_invented_record_false_one_char_below_the_combined_length_floor():
    # Same two-label shape, one char short of the floor: "Organization"
    # (12) and "Institution" (11) join to exactly 24 chars, isolating the
    # `<` boundary from the other side.
    a, b = "Organization", "Institution"
    assert len(f"{a}|{b}") == _MIN_EXACT_LEN - 1
    assert not _is_invented_record({"field_one": a, "field_two": b})


def test_is_invented_record_false_on_a_mixed_real_and_label_record():
    # Isolates the `all(...)` guard itself (not just the two floors above):
    # one real value plus one template label clears both the value-count
    # floor (2) and the combined-length floor (>= 25, here 49) but must
    # stay False -- `any(...)` in place of `all(...)` would wrongly call
    # this fabricated, and nothing else in this file rebuilds a
    # certifying_board/certificate_number pair with one side real (#829).
    mixed = {"certifying_board": "American Board of Internal Medicine",
             "certificate_number": _CERT_NUMBER_LABEL}
    assert len(mixed) >= INVENTED_RECORD_MIN_VALUES
    assert len("|".join(mixed.values())) >= _MIN_EXACT_LEN
    assert not _is_invented_record(mixed)


def test_is_invented_record_false_with_no_values_at_all():
    assert not _is_invented_record({"a": None, "b": ""})


# -- _rendered_row_value_sets ------------------------------------------------

def test_rendered_row_value_sets_distinguishes_header_from_data_row():
    rows = _rendered_row_value_sets([[_HEADER_ROW, _INVENTED_CERT_ROW]])
    invented_key = frozenset({_BOARD_LABEL.lower(), _CERT_NUMBER_LABEL.lower()})
    header_key = frozenset(_norm(v) for v in _HEADER_ROW)
    assert invented_key in rows
    assert header_key in rows
    # The header row's 3rd cell carries a parenthetical the data row's
    # matching cell does not -- they must not collide into one set.
    assert invented_key != header_key


def test_rendered_row_value_sets_ignores_empty_cells_and_empty_tables():
    assert _rendered_row_value_sets([]) == set()
    assert _rendered_row_value_sets([[["", "", ""]]]) == set()
    assert _rendered_row_value_sets([[["Alpha", "", "Beta"]]]) == {
        frozenset({"alpha", "beta"})}


# -- lint_invented_records: part (a), the rendered-header-record shape -----

def test_lint_invented_records_warns_on_a_rendered_header_record():
    stage4 = {"entries": [_invented_cert_entry(element_idx_start=7)]}
    table_rows = [[_HEADER_ROW, _INVENTED_CERT_ROW]]
    findings = lint_invented_records(stage4, table_rows)
    assert len(findings) == 1
    finding = findings[0]
    assert set(finding.keys()) == {"lint", "severity", "message", "evidence",
                                   "status", "reason"}
    assert finding["lint"] == "invented_records"
    assert finding["severity"] == "WARN"
    assert "F2" in finding["message"] and "7" in finding["message"]
    assert "#829" in finding["message"]
    # The evidence is the populated field values, not just the finding's
    # key set -- an `[]` in place of the list comprehension would still
    # pass every assertion above (#829's own contract test only checks
    # keys), silently dropping what a reviewer sees. The two None-valued
    # fields (`year_certified`, `recertification_date`) contribute nothing.
    assert finding["evidence"] == [
        "certifying_board: Full Name of Board",
        "certificate_number: Certificate #",
    ]


def test_lint_invented_records_silent_when_the_header_record_never_rendered():
    # Stage 4 still fabricates the entry, but stage 6 correctly suppressed
    # it (#959's own fix, for this one section) -- only the real header row
    # made it into the document, no second data row.
    stage4 = {"entries": [_invented_cert_entry()]}
    table_rows = [[_HEADER_ROW]]
    assert lint_invented_records(stage4, table_rows) == []


def test_lint_invented_records_silent_on_real_certification_content():
    stage4 = {"entries": [_real_cert_entry()]}
    table_rows = [[_HEADER_ROW, _REAL_CERT_ROW]]
    assert lint_invented_records(stage4, table_rows) == []


def test_lint_invented_records_skips_taxonomy_code_t():
    entry = _invented_cert_entry()
    entry["taxonomy_code"] = "T"
    assert lint_invented_records({"entries": [entry]}, [[_HEADER_ROW, _INVENTED_CERT_ROW]]) == []


def test_lint_invented_records_handles_extracted_fields_none_without_raising():
    # Stage 4 can write extracted_fields: null for an entry field extraction
    # skipped outright (not just individual field values of None, the shape
    # test_lint_invented_records_silent_on_single_label_affiliation_prompt
    # covers below). `lint_invented_records`'s own read is
    # `e.get("extracted_fields") or {}`, not `e.get("extracted_fields", {})`
    # -- the latter returns None (the key IS present) rather than {} when the
    # value itself is None, and passing that None into `_is_invented_record`
    # would raise AttributeError on `.values()`, which `_run_lint` in
    # run_doctor.py turns into a doctor ERROR finding instead of the correct
    # silent skip.
    entry = {"taxonomy_code": "F2", "element_type": "table_row",
            "element_idx_start": 5, "text": "unextracted row",
            "extracted_fields": None}
    assert lint_invented_records({"entries": [entry]}, []) == []


def test_lint_invented_records_silent_on_single_label_affiliation_prompt():
    # The G/Institutional-Affiliation shape: one unfilled prompt label
    # rendered with a blank companion cell is the template's own intentional
    # "show the prompt, blank if unfilled" convention, not a fabricated
    # multi-field record.
    entry = {"taxonomy_code": "G", "element_type": "break",
            "element_idx_start": 3, "text": f"{_BOARD_LABEL}:",
            "extracted_fields": {"organization": None,
                                 "affiliation_type": _BOARD_LABEL}}
    table_rows = [[[_BOARD_LABEL, ""]]]
    assert lint_invented_records({"entries": [entry]}, table_rows) == []


def test_lint_invented_records_silent_on_a_real_percent_effort_total_row():
    entry = {"taxonomy_code": "J", "element_type": "table_row",
            "element_idx_start": 4, "text": "Total | 100%",
            "extracted_fields": {"activity": "Total", "percent_effort": "100%"}}
    table_rows = [[["Activity", "% Effort"], ["Total", "100%"]]]
    assert lint_invented_records({"entries": [entry]}, table_rows) == []


# -- lint_invented_records: part (b), the F1 invented-licence shape --------

def test_lint_invented_records_warns_on_an_f1_entry_matching_a_known_instruction_exactly():
    known_instruction = ("Licensure: Every physician appointed to the "
                         "Hospital staff, except interns, and aliens in "
                         "the US via non-immigrant visas, must have a New "
                         "York State license or a temporary certificate in "
                         "lieu of the license.")
    entry = _invented_licensure_entry(text=known_instruction)
    findings = lint_invented_records({"entries": [entry]}, [])
    assert len(findings) == 1
    assert findings[0]["lint"] == "invented_records"
    assert findings[0]["severity"] == "WARN"
    assert "F1" in findings[0]["message"] and "#829" in findings[0]["message"]
    # The evidence is the source text itself (truncated), not just the
    # message -- an `[]` in place of `[text[:120]]` would still pass every
    # assertion above. `known_instruction` is 207 chars, so this also pins
    # the truncation, not just that evidence is non-empty.
    assert findings[0]["evidence"] == [known_instruction[:120]]


def test_lint_invented_records_warns_on_an_f1_entry_with_foreign_template_instruction_text():
    # #530: another institution's instruction line (invented text) is False
    # under both WCM-only helpers, so only the foreign detector claims it.
    text = "D. Sample Licensure (list state, license number and dates of issue)"
    assert not is_template_instruction(text)
    assert not is_near_template_instruction(text)
    findings = lint_invented_records({"entries": [_invented_licensure_entry(text=text)]}, [])
    assert len(findings) == 1
    assert findings[0]["lint"] == "invented_records"
    assert "F1" in findings[0]["message"]
    assert findings[0]["evidence"] == [text]


def test_lint_invented_records_warns_on_an_f1_entry_matching_via_the_pipe_split_path_only():
    # A table-row-shaped F1 text ("... | ") reaches `is_template_instruction`
    # via its pipe-split rule (b), not the whole-string exact match in rule
    # (a) -- the trailing "|" survives normalization, so the joined string
    # never equals the known instruction verbatim. `is_near_template_instruction`
    # refuses any text containing "|" outright, so this exercises the exact
    # branch (`is_template_instruction(...)`) with the near-match branch
    # provably False, isolating the `or` in `lint_invented_records` from the
    # near-match test above (#829).
    known_instruction = ("Licensure: Every physician appointed to the "
                         "Hospital staff, except interns, and aliens in "
                         "the US via non-immigrant visas, must have a New "
                         "York State license or a temporary certificate in "
                         "lieu of the license.")
    text = f"{known_instruction} | "
    assert is_template_instruction(text)
    assert not is_near_template_instruction(text)
    entry = _invented_licensure_entry(text=text)
    findings = lint_invented_records({"entries": [entry]}, [])
    assert len(findings) == 1
    assert findings[0]["lint"] == "invented_records"
    assert "F1" in findings[0]["message"]


def test_lint_invented_records_warns_on_an_f1_entry_near_matching_a_known_instruction():
    entry = _invented_licensure_entry()  # the "ABC Hospital" near-variant
    findings = lint_invented_records({"entries": [entry]}, [])
    assert len(findings) == 1
    assert findings[0]["lint"] == "invented_records"


def test_lint_invented_records_fires_on_f1_regardless_of_element_type():
    # A5IZ6Q's own invented licence was extracted as element_type "break" --
    # pin that this lint does not filter it out (#829).
    for element_type in ("break", "header", "paragraph", "table_row"):
        entry = _invented_licensure_entry(element_type=element_type)
        assert len(lint_invented_records({"entries": [entry]}, [])) == 1, element_type


def test_lint_invented_records_ignores_the_same_text_under_a_different_code():
    entry = _invented_licensure_entry()
    entry["taxonomy_code"] = "F2"
    assert lint_invented_records({"entries": [entry]}, []) == []


def test_lint_invented_records_silent_on_real_f1_content():
    entry = {"taxonomy_code": "F1", "element_type": "table_row",
            "element_idx_start": 2,
            "text": "New York | 123456 | 06/01/2010",
            "extracted_fields": {"state_country": "New York",
                                 "license_number": "123456",
                                 "issue_date": "06/01/2010",
                                 "expiration_date": None}}
    assert lint_invented_records({"entries": [entry]}, []) == []


def test_lint_invented_records_reports_both_shapes_in_one_run():
    stage4 = {"entries": [_invented_cert_entry(element_idx_start=1),
                          _invented_licensure_entry(element_idx_start=2)]}
    table_rows = [[_HEADER_ROW, _INVENTED_CERT_ROW]]
    findings = lint_invented_records(stage4, table_rows)
    assert len(findings) == 2
    assert any("F2" in f["message"] for f in findings)
    assert any("F1" in f["message"] for f in findings)


# ==========================================================================
# #729 -- lint_wrong_start_date.

def _dated_entry(text, start, end=None, code="D1", idx=7):
    return {"taxonomy_code": code, "element_idx_start": idx, "text": text,
            "extracted_fields": {"start_date": start, "end_date": end}}


def test_wrong_start_date_fires_on_the_fsmb_shape_as_warn():
    findings = lint_wrong_start_date(
        {"entries": [_dated_entry("Example Board | 2025-2026", "2026")]})
    assert len(findings) == 1
    assert findings[0]["lint"] == "wrong_start_date"
    assert findings[0]["severity"] == "WARN"
    assert "2025-2026" in findings[0]["message"]
    assert "2026-Present" in findings[0]["message"]
    assert "7" in findings[0]["message"]


def test_wrong_start_date_is_info_when_start_is_not_the_range_end():
    findings = lint_wrong_start_date(
        {"entries": [_dated_entry("Example Board | 2018-2020", "Sep 2018")]})
    assert [f["severity"] for f in findings] == ["INFO"]


@pytest.mark.parametrize("entry", [
    _dated_entry("Example Board | 2025-2026", "2026", end="2026"),  # end set
    _dated_entry("Example Board | 2025-Present", "2026"),           # no closed range
    _dated_entry("Example Board | 2025-2026 present", "2026"),       # ongoing marker
    _dated_entry("Board | 2020-2021 | 2025-2026", "2026"),          # two ranges
    _dated_entry("Course 5130-1020", "2026"),                        # implausible range
    _dated_entry("Example Board | 2025-2026", "2026", code="S1"),    # no date schema
])
def test_wrong_start_date_silent_when_a_condition_fails(entry):
    assert lint_wrong_start_date({"entries": [entry]}) == []


def test_wrong_start_date_leaves_the_extracted_value_alone():
    entry = _dated_entry("Example Board | 2025-2026", "2026")
    lint_wrong_start_date({"entries": [entry]})
    assert entry["extracted_fields"] == {"start_date": "2026", "end_date": None}


def test_dedup_and_render_tokens_are_unicode_aware():
    """#541: accented Latin stays whole, Cyrillic/Greek produce tokens,
    ASCII is unchanged. Invented names."""
    from unified_pipeline.doctor.shared import _long_word_tokens
    assert _alphanumeric_tokens("Zoë Brändström") == Counter(
        {"zoe": 1, "brandstrom": 1})
    assert _alphanumeric_tokens("report_2019") == Counter(
        {"report": 1, "2019": 1})
    assert _alphanumeric_tokens("Иван Петров") == Counter(
        {"иван": 1, "петров": 1})
    assert _long_word_tokens("Zoë Brändström") == {"brandstrom"}
    assert _long_word_tokens("Ελένη Παπαδοπούλου") == {"ελενη", "παπαδοπουλου"}
    assert _long_word_tokens("Alpha beta gamma12") == {"alpha", "gamma"}
    assert _long_word_tokens("hello_world") == {"hello", "world"}
    # #722: CJK never forms a token (excluded, not measured).
    assert _long_word_tokens("東京大学医学部教授 한국어로된논문제목") == set()
    assert _long_word_tokens("東京大学医学部 Blorvane") == {"blorvane"}


# ==========================================================================
# lint_offschema_fields: a value under a key no renderer reads. Every value
# below is invented.

def _fields_entry(code, fields, text="Example entry", idx=11):
    return {"taxonomy_code": code, "element_idx_start": idx, "text": text,
            "extracted_fields": fields}


def _offschema(*entries):
    return lint_offschema_fields({"entries": list(entries)})


# One non-date schema value per code, so the entry renders from its fields
# rather than falling back to its raw text (the lint skips that case).
_ANCHOR = {"R": {"title": "Talk one"}, "I": {"organization": "Society A"},
           "D1": {"title": "Lecturer"}, "P": {"institution": "Example College"},
           "N3B": {"mentee_name": "A. Mentee"}, "B1": {"degree": "BA"},
           "K1": {"course_title": "Course"}, "S1": {"title": "Paper"},
           "S8": {"title": "Paper"}, "R1": {"title": "Talk"}, "M2C": {"title": "Grant"},
           "H": {"award_name": "Prize"}, "A": {"name": "Owner"},
           "T": {"description": "x"}, "G": {"organization": "Example Org"}}


def _anchored(code, fields, **kwargs):
    return _fields_entry(code, {**_ANCHOR[code], **fields}, **kwargs)


def test_offschema_record_dict_sharing_a_schema_key_is_warn():
    findings = _offschema(_anchored("R", {
        "title": "Talk one",
        "additional_entry": {"title": "Talk two", "location": "Springfield"}}))
    assert len(findings) == 1
    assert findings[0]["lint"] == "offschema_fields"
    assert findings[0]["severity"] == "WARN"
    assert "`additional_entry`" in findings[0]["message"]
    assert "1 R entry" in findings[0]["message"]
    assert "1 whole record under it" in findings[0]["message"]
    assert findings[0]["evidence"] == [
        'entry 11: {"title": "Talk two", "location": "Springfield"}']


def test_offschema_dict_sharing_no_schema_key_is_info():
    findings = _offschema(_anchored("R", {"extra": {"colour": "blue"}}))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_numbered_schema_field_is_warn():
    findings = _offschema(_anchored(
        "I", {"organization": "Society A", "organization_2": "Society B"}))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "`organization_2`" in findings[0]["message"]


def test_offschema_numbered_key_off_the_schema_is_info():
    findings = _offschema(_anchored("I", {"widget_2": "Society B"}))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_single_record_list_is_warn():
    findings = _offschema(_anchored("D1", {"appointments": [
        {"title": "Lecturer", "institution": "Example College"}]}))
    assert [f["severity"] for f in findings] == ["WARN"]


def test_offschema_list_of_strings_is_info():
    findings = _offschema(_anchored("D1", {"keywords": ["one", "two"]}))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_scalar_is_info_and_names_the_fact():
    findings = _offschema(_anchored("B1", {"degree": "BA", "honors": "with distinction"}))
    assert [f["severity"] for f in findings] == ["INFO"]
    assert "each holds one fact" in findings[0]["message"]
    assert "missing from the output" in findings[0]["message"]
    assert findings[0]["evidence"] == ["entry 11: with distinction"]


def test_offschema_a_record_list_fan_out_splits_is_not_reported():
    """Stage 6 renders each record of this list as its own child entry."""
    entry = _anchored("P", {"committees": [
        {"committee_name": "Alpha Board", "role": "Chair"},
        {"committee_name": "Beta Panel", "role": "Member"}]},
        text="Alpha Board Chair\tBeta Panel Member")
    assert _offschema(entry) == []


def test_offschema_a_record_list_fan_out_declines_is_warn():
    """Same list, but the text holds a token no rendered field carries, so
    fan-out keeps the entry whole and the list is never read."""
    entry = _anchored("P", {"committees": [
        {"committee_name": "Alpha Board", "role": "Chair"},
        {"committee_name": "Beta Panel", "role": "Member"}]},
        text="Alpha Board Chair\tBeta Panel Member Gamma")
    assert [f["severity"] for f in _offschema(entry)] == ["WARN"]


def test_offschema_stage4_records_fan_out_splits_is_not_reported():
    """Stage 4 keeps every record the LLM returned under `stage4_records`, and
    stage 6 hands that key to fan-out, which renders each as its own row."""
    entry = _anchored("I", {"organization": "Society B", STAGE4_RECORDS_KEY: [
        {"organization": "Society A", "membership_type": "Member"},
        {"organization": "Society B", "membership_type": "Fellow"}]},
        text="Societies: Society A Member\tSociety B Fellow")
    # "Societies" is held by no field, so the generic list rule would decline;
    # only the stage-4 records path splits this entry.
    assert _offschema(entry) == []


def test_offschema_stage4_records_fan_out_declines_is_warn():
    """A record with no value its section writes makes fan-out decline the
    whole list, so the records are never read."""
    entry = _anchored("I", {"organization": "Society B", STAGE4_RECORDS_KEY: [
        {"notes": "Founding member"},
        {"organization": "Society B", "membership_type": "Fellow"}]},
        text="Founding member\tSociety B Fellow")
    findings = _offschema(entry)
    assert [f["severity"] for f in findings] == ["WARN"]
    assert f"`{STAGE4_RECORDS_KEY}`" in findings[0]["message"]


@pytest.mark.parametrize("entry", [
    _anchored("I", {"organization": "Society A"}),             # built-in schema key
    _anchored("R", {"scope": "National"}),                     # config-only, extract false
    _anchored("N3B", {"thesis_title": "Example thesis"}),     # built-in only
    _anchored("R", {"pmid": "123"}),                           # stage-4 identifiers
    _anchored("R", {"pmcid": "PMC1"}),
    _anchored("R", {"doi": "10.1/x"}),
    _anchored("R1", {"extra": "x"}),                           # no schema at all
    _anchored("M2C", {"percent_effort": "5%"}),                # stage-4 effort
    _anchored("H", {"date_range": "2001-2003"}),               # date-named
    _anchored("H", {"start_date_1": "2001"}),                  # numbered date
    _anchored("B1", {"dates_attended": {"start_date": "1990"}}),
    _anchored("I", {"extra": ""}),                             # blank values
    _anchored("I", {"extra": None}),
    _anchored("I", {"extra": []}),
    _anchored("I", {"extra": {}}),
    _anchored("A", {"school": "Example School"}),              # personal data
    _anchored("T", {"extra": "x"}),                            # text-rendered
    _anchored("G", {"extra": "x"}),
    _anchored("S1", {"other_id": "x"}),                        # publication (5d)
    _anchored("S1", {"notes": "x"}),
    _anchored("K1", {"description": "x"}),                     # teaching (5c)
    _fields_entry("I", {"extra": "x"}),                            # renders its text:
    _fields_entry("I", {"start_date": "2001", "extra": "x"}),      # a one-fact value
    _fields_entry("I", {"organization": "", "end_date": "2001", "extra": "x"}),  # skipped
    _anchored("S8", {"journal_or_source": "x"}),
    _fields_entry("I", "not an object"),                           # malformed fields
    _fields_entry("I", None),
])
def test_offschema_silent_when_a_renderer_or_a_rule_accounts_for_the_key(entry):
    assert _offschema(entry) == []


@pytest.mark.parametrize("fields", [
    {"start_date": "2001", "appointments": [
        {"title": "Lecturer", "institution": "Example College"}]},  # list of objects
    {"start_date": "2001", "organization_2": "Society B"},           # numbered field
    {"end_date": "2001", "extra": {"organization": "Society B"}},    # schema-key object
])
def test_offschema_record_on_an_entry_rendering_its_text_is_warn(fields):
    """The raw text an entry with only dates falls back to does not carry a
    whole record, so the one-fact skip never covers a record-shaped value."""
    findings = _offschema(_fields_entry("I", fields))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "1 whole record under it" in findings[0]["message"]


def test_offschema_entry_rendering_its_text_reports_its_record_but_not_its_fact():
    findings = _offschema(_fields_entry("D1", {
        "start_date": "2001", "note": "visiting",
        "appointments": [{"title": "Lecturer"}, {"title": "Reader"}]}))
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "appointments")]


def test_field_lint_constants_are_pinned():
    assert FIELD_EVIDENCE_VALUE_CHARS == 100
    assert FIELD_EVIDENCE_MAX_VALUES == 3
    assert DEGREE_YEAR_LEAD == 10


def test_offschema_a_key_merely_containing_date_is_not_a_date_key():
    findings = _offschema(_anchored("I", {"candidate": "x", "update": "y"}))
    assert [f["message"].split("`")[1] for f in findings] == ["candidate", "update"]


def test_offschema_list_mixing_objects_and_strings_is_info():
    findings = _offschema(_anchored("D1", {"appointments": [
        {"title": "Lecturer"}, "Reader"]}))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_config_comment_entries_and_null_fields_are_tolerated(
        monkeypatch, tmp_path):
    config = tmp_path / "schemas.json"
    config.write_text(json.dumps({"schemas": {
        "__NOTE_example": "a comment string, not a schema",
        "I": {"fields": {"region": {"extract": False}}},
        "R": {"fields": None}}}))
    monkeypatch.setattr(extraction_lints, "FIELD_SCHEMA_CONFIG_PATH", config)
    assert _offschema(_anchored("I", {"region": "North"})) == []
    # `scope` is declared only by the real config file, not the built-ins.
    assert len(_offschema(_anchored("R", {"scope": "National"}))) == 1


def test_offschema_a_rendered_field_outside_both_schemas_is_not_reported(monkeypatch):
    """`fan_out._RENDERED_FIELDS` is the third source of 'something reads
    this key'; no code needs it today, so plant one."""
    monkeypatch.setattr(extraction_lints, "_RENDERED_FIELDS",
                        {"I": frozenset({"chapter"})})
    assert _offschema(_anchored("I", {"chapter": "Local"})) == []
    assert len(_offschema(_anchored("I", {"region": "Local"}))) == 1


def test_offschema_one_finding_per_code_and_key_sorted_with_capped_evidence():
    entries = [_anchored("N3B", {"outcome": f"result {i}"}, idx=i)
               for i in range(FIELD_EVIDENCE_MAX_VALUES + 2)]
    entries.append(_anchored("B1", {"honors": "cum laude"}, idx=90))
    entries.append(_anchored("B1", {"honors": "magna"}, idx=91))
    findings = _offschema(*entries)
    assert [f["message"].split(":")[0] for f in findings] == [
        "2 B1 entries", f"{FIELD_EVIDENCE_MAX_VALUES + 2} N3B entries"]
    assert len(findings[1]["evidence"]) == FIELD_EVIDENCE_MAX_VALUES
    assert findings[1]["evidence"][0] == "entry 0: result 0"


def test_offschema_warn_when_any_value_of_the_key_is_a_record():
    findings = _offschema(
        _anchored("R", {"additional_entry": "free text"}, idx=1),
        _anchored("R", {"additional_entry": {"title": "Talk"}}, idx=2),
        _anchored("R", {"additional_entry": {"title": "Talk 2"}}, idx=3))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "2 whole records under it" in findings[0]["message"]


def test_offschema_evidence_value_is_truncated():
    long_value = "x" * (FIELD_EVIDENCE_VALUE_CHARS + 50)
    findings = _offschema(_anchored("I", {"extra": long_value}))
    assert findings[0]["evidence"] == [
        "entry 11: " + "x" * FIELD_EVIDENCE_VALUE_CHARS]


def test_offschema_reads_the_config_file_and_fails_closed_without_it(monkeypatch, tmp_path):
    monkeypatch.setattr(extraction_lints, "FIELD_SCHEMA_CONFIG_PATH",
                        tmp_path / "absent.json")
    with pytest.raises(FileNotFoundError):
        _offschema(_anchored("I", {"organization": "Society A"}))


def test_offschema_leaves_the_entry_alone():
    entry = _anchored("P", {"committees": [
        {"committee_name": "Alpha Board", "role": "Chair"},
        {"committee_name": "Beta Panel", "role": "Member"}]},
        text="Alpha Board Chair\tBeta Panel Member")
    before = json.dumps(entry, sort_keys=True)
    _offschema(entry)
    assert json.dumps(entry, sort_keys=True) == before


def test_offschema_record_list_counts_each_record():
    """#1245: a list of three appointments is three records lost, not one."""
    findings = _offschema(
        _anchored("D1", {"appointments": [
            {"title": "Lecturer"}, {"title": "Reader"}, {"title": "Fellow"}]}, idx=1),
        _anchored("D1", {"appointments": [{"title": "Tutor"}]}, idx=2))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["message"].startswith("2 D1 entries: `appointments`")
    assert "4 whole records under it" in findings[0]["message"]


# --------------------------------------------------------------------------
# lint_offschema_fields graded against the rendered document (#1245). Blocks
# are the `read_docx_blocks` shape: ("p", text) or ("table", its lines joined
# by newlines). Every name and value is invented.

def _graded(blocks, *entries):
    return lint_offschema_fields({"entries": list(entries)}, blocks)


def _row(*cells):
    """One table row as `_table_lines` writes it: each cell, then the row
    joined across its cells."""
    return [*cells, " | ".join(cells)]


def _table(*rows):
    return ("table", "\n".join(line for row in rows for line in row))


def _offschema_degree(**extra):
    return _fields_entry("B1", {"degree": "BA", "institution": "Example College",
                                **extra}, text="BA, Example College, Townsville, Exland")


def test_offschema_value_its_own_row_shows_is_not_reported():
    """Stage 5b writes a B1 `location` into the institution cell, re-punctuated."""
    entry = _offschema_degree(location="Townsville, Exland")
    blocks = [_table(_row("BA", "Example College, Townsville City, Exland", "2001"))]
    assert _graded(blocks, entry) == []
    # Without the document the value is reported, as before #1245.
    assert [f["severity"] for f in _offschema(entry)] == ["INFO"]


def test_offschema_value_the_document_lacks_and_the_text_states_is_warn():
    entry = _fields_entry("B1", {"degree": "BA", "institution": "Example College",
                                 "honors": "with distinction"},
                          text="BA, Example College, with distinction")
    findings = _graded([_table(_row("BA", "Example College", "2001"))], entry)
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "1 of them nowhere in the document" in findings[0]["message"]
    assert findings[0]["evidence"] == ["entry 11: with distinction"]


def test_offschema_a_number_no_line_shows_is_reported():
    """A value with no string in it is matched as written, not waved
    through as shown."""
    entry = _fields_entry("M2C", {"title": "Example Grant", "share": 0.25},
                          text="Example Grant, share 0.25")
    findings = _graded([_table(_row("Title:", "Example Grant"))], entry)
    assert [f["severity"] for f in findings] == ["WARN"]


def test_offschema_a_whole_record_is_reported_even_where_its_row_shows_it():
    """A record under an off-schema key is not graded: a second talk fused
    onto the first one's row is still a record stage 6 never wrote as one
    (#1187), so it stays WARN."""
    entry = _anchored("R", {"event_name": "Example Meeting", "additional_entry": {
        "title": "Talk two", "event_name": "Example Meeting"}})
    blocks = [_table(_row("Talk one; Talk two", "Example Meeting", "2010"))]
    findings = _graded(blocks, entry)
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "1 whole record under it" in findings[0]["message"]


def test_offschema_value_the_text_never_states_stays_info():
    """The model's own remark on an entry is not content the CV lost."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil",
                                  "note": "no further detail was given"},
                          text="Ann Pupil, summer student")
    findings = _graded([_table(_row("Name:", "Ann Pupil"))], entry)
    assert [f["severity"] for f in findings] == ["INFO"]
    assert "1 of them nowhere in the document" in findings[0]["message"]


def test_offschema_value_the_text_states_in_other_words_is_warn():
    """Stage 4 rewrites a date inside a value ("05/2015" as "May 2015"):
    most of its words in the entry's text still state it."""
    entry = _fields_entry("F2", {"certifying_board": "Example Board",
                                 "notes": "Board recertified May 2015"},
                          text="Example Board; Board recertified 05/2015")
    findings = _graded([_table(_row("Example Board", "2010"))], entry)
    assert [f["severity"] for f in findings] == ["WARN"]


def test_offschema_list_value_is_a_loss_only_when_the_text_states_every_item():
    """One item the entry's text never states (the model's own) keeps the
    whole value INFO."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil",
                                  "awards": ["Gold Ribbon", "likely a school prize"]},
                          text="Ann Pupil; Gold Ribbon")
    findings = _graded([_table(_row("Name:", "Ann Pupil"))], entry)
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_value_shown_only_with_another_record_is_info():
    lost = _fields_entry("N3B", {"mentee_name": "Ann Pupil", "advisor": "Bea Guide"},
                         text="Ann Pupil, advised by Bea Guide", idx=1)
    other = _fields_entry("N3B", {"mentee_name": "Cy Pupil",
                                  "site_position": "Lab of Bea Guide"}, idx=2)
    blocks = [_table(_row("Name:", "Ann Pupil")),
              _table(_row("Name:", "Cy Pupil"), _row("Site/Position:", "Lab of Bea Guide"))]
    findings = _graded(blocks, lost, other)
    assert [f["severity"] for f in findings] == ["INFO"]
    assert "the document shows it, but not with its record" in findings[0]["message"]


def test_offschema_a_form_table_is_one_record():
    """A "Label: | value" table is one record's form: a value in one of its
    rows is shown with the name in another. A table of records is not."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil",
                                  "awards": "Gold Ribbon, Science Fair"},
                          text="Ann Pupil; Gold Ribbon, Science Fair")
    form = _table(_row("Name:", "Ann Pupil"),
                  _row("Project/Accomplishments:", "Awards: Gold Ribbon, Science Fair"))
    assert _graded([form], entry) == []
    listing = _table(_row("Ann Pupil", "2001"),
                     _row("Bo Pupil", "Gold Ribbon, Science Fair"))
    assert [f["severity"] for f in _graded([listing], entry)] == ["INFO"]


def test_offschema_a_form_table_matches_inside_one_row_only():
    """Read across rows, "Delta Group" ending one and "Award detail:"
    opening the next would carry an anchor neither row holds."""
    entry = _fields_entry("I", {"organization": "Delta Group Award",
                                "standing": "charter member"},
                          text="Delta Group Award, charter member")
    form = _table(_row("Body:", "Delta Group"), _row("Award detail:", "charter member"))
    findings = _graded([form], entry)
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_a_form_table_shows_a_value_only_within_one_row():
    """A form table is one unit, not one bag of words: a value whose words
    are spread over its rows is not shown in it."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil", "awards": "Gold Ribbon"},
                          text="Ann Pupil; Gold Ribbon")
    form = _table(_row("Name:", "Ann Pupil"), _row("Project:", "Gold study"),
                  _row("Award source:", "Ribbon Fund"))
    findings = _graded([form], entry)
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "1 of them nowhere in the document" in findings[0]["message"]


def test_offschema_an_anchor_inside_a_sibling_value_finds_no_row():
    """A value inside another record's longer value cannot say which row is
    this record's: the sibling's row shows the institution this one lost."""
    lost = _fields_entry("N3A", {"site_position": "Doctoral Program",
                                 "institution": "Lakeside University"}, idx=1)
    sibling = _fields_entry("N3A", {
        "site_position": "MD, Doctoral Program - Lakeside University"}, idx=2)
    blocks = [_table(_row("Site/Position:", "MD, Doctoral Program - Lakeside University")),
              _table(_row("Site/Position:", "Doctoral Program"))]
    findings = _graded(blocks, lost, sibling)
    assert [(f["severity"], f["evidence"]) for f in findings] == [
        ("INFO", ["entry 1: Lakeside University"])]


def test_offschema_nested_anchor_values_count_once():
    """A level inside the program name is the same evidence, not a second."""
    lost = _fields_entry("N3A", {"mentee_level": "MSc", "site_position": "MSc Program",
                                 "institution": "Lakeside University"}, idx=1)
    sibling = _fields_entry("N3A", {
        "site_position": "BSc, MSc Program - Lakeside University"}, idx=2)
    blocks = [_table(_row("Site/Position:", "BSc, MSc Program - Lakeside University"))]
    findings = _graded(blocks, lost, sibling)
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_values_shared_with_other_records_find_their_row_together():
    """Degree and institution each sit in another record too, but together
    only on the degree's own row."""
    degree = _fields_entry("B1", {"degree": "PhD", "institution": "Northfield University",
                                  "location": "Northfield, Exland"}, idx=1)
    award = _fields_entry("H", {"award_name": "Scholarship (PhD)",
                                "granting_body": "Northfield University, Exland"}, idx=2)
    blocks = [_table(_row("PhD", "Northfield University, Northfield, Exland", "1991"),
                     _row("MA", "Southfield University, Southfield, Exland", "1986")),
              _table(_row("Scholarship (PhD)", "Northfield University, Exland", "1986"))]
    assert _graded(blocks, degree, award) == []


def test_offschema_a_row_carrying_one_shared_value_is_not_the_records():
    """Without a distinctive value, a row is the record's only when it
    carries a majority of the record's values: the medal's row names the
    same university, and the honor on it is the medal's."""
    degree = _fields_entry("B1", {"degree": "PhD", "institution": "Northfield University",
                                  "honors": "with distinction"},
                           text="PhD, Northfield University, with distinction", idx=1)
    medal = _fields_entry("H", {"award_name": "Northfield University Medal",
                                "granting_body": "Northfield University"}, idx=2)
    blocks = [_table(_row("PhD", "Northfield University", "1991")),
              _table(_row("Northfield University Medal, with distinction",
                          "Northfield University", "1990"))]
    findings = _graded(blocks, degree, medal)
    assert [(f["severity"], f["evidence"]) for f in findings] == [
        ("INFO", ["entry 1: with distinction"])]
    assert "the document shows it, but not with its record" in findings[0]["message"]


def test_offschema_template_wording_never_marks_a_records_row():
    """A value that is the template's own wording ("Mentor") is on the
    template's lines whatever record they belong to, so it cannot say which
    row is this record's, even when no other record holds it."""
    entry = _fields_entry("N2", {"role": "Mentor", "sponsor": "Example Foundation"},
                          text="Mentor, Example Foundation")
    other = _table(_row("Program:", "Example Foundation Scholars"),
                   _row("Mentor:", "Dee Guide"))
    findings = _graded([other], entry)
    assert [f["severity"] for f in findings] == ["INFO"]
    assert "the document shows it, but not with its record" in findings[0]["message"]


def test_offschema_a_date_under_a_non_date_key_never_marks_a_records_row():
    """`expected_completion` is not date-named, but "May 2024" is a date:
    another record's row that shows it is not this record's row."""
    lost = _fields_entry("N3A", {"mentee_name": "Ann Pupil", "expected_completion": "May 2024",
                                 "advisor": "Bea Guide"},
                         text="Ann Pupil, advised by Bea Guide", idx=1)
    other = _fields_entry("N3A", {"mentee_name": "Cy Pupil", "start_date": "05/2024",
                                  "site_position": "Lab of Bea Guide"}, idx=2)
    blocks = [_table(_row("Ann Pupil", "Doctoral student")),
              _table(_row("Cy Pupil", "May 2024", "Lab of Bea Guide"))]
    findings = _graded(blocks, lost, other)
    assert [(f["severity"], f["evidence"]) for f in findings] == [
        ("INFO", ["entry 1: Bea Guide"])]


def _honor(**dates):
    return _fields_entry("H", {"award_name": "Silver Lamp Award",
                               "granting_body": "Example Society", **dates},
                         text="Silver Lamp Award, Example Society, 2004-2007")


def test_offschema_single_date_code_range_its_row_lacks_is_warn():
    """H declares `date` alone; a range under start/end renders nowhere."""
    findings = _graded([_table(_row("Silver Lamp Award", "Example Society"))],
                       _honor(start_date="2004", end_date="2007"))
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "end_date"), ("WARN", "start_date")]
    assert "1 of them not on its record's line" in findings[0]["message"]


def test_offschema_single_date_code_range_its_row_shows_is_not_reported():
    shown = _table(_row("Silver Lamp Award", "Example Society", "2004-2007"))
    assert _graded([shown], _honor(start_date="2004", end_date="2007")) == []
    start_only = _table(_row("Silver Lamp Award", "Example Society", "2004"))
    findings = _graded([start_only], _honor(start_date="2004", end_date="2007"))
    assert [f["message"].split("`")[1] for f in findings] == ["end_date"]


def test_offschema_date_with_no_four_digit_year_is_not_reported():
    row = _row("Silver Lamp Award", "Example Society")
    assert _graded([_table(row)], _honor(end_date="present")) == []


def test_offschema_date_whose_record_has_no_row_is_not_reported():
    """Its year on some other line would be a coincidence, not its record."""
    row = _row("Bronze Bowl", "Other Society", "2004-2007")
    assert _graded([_table(row)], _honor(start_date="2004", end_date="2007")) == []


def test_offschema_date_is_judged_on_its_own_row_not_on_another_records():
    """The record's own row lacks the range; another record's row carrying
    the same years does not show it."""
    blocks = [_table(_row("Silver Lamp Award", "Example Society"),
                     _row("Bronze Bowl", "Other Society", "2004-2007"))]
    findings = _graded(blocks, _honor(start_date="2004", end_date="2007"))
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "end_date"), ("WARN", "start_date")]


def test_offschema_date_key_renaming_a_schema_date_is_not_reported():
    """D1 declares start and end dates; its `date` is a rename a renderer
    may read, so the lint stays out of it even with the document."""
    entry = _fields_entry("D1", {"title": "Lecturer", "institution": "Example College",
                                 "date": "2001"})
    assert _graded([_table(_row("Lecturer", "Example College"))], entry) == []


def test_offschema_a_list_of_periods_on_any_code_is_judged():
    """A second term under `additional_dates` is not a rename of start/end."""
    entry = _fields_entry("O", {
        "leadership_role": "Chair", "institution": "Example College",
        "start_date": "2019", "end_date": "2020",
        "additional_dates": [{"start_date": "2022", "end_date": "2023"}]})
    findings = _graded([_table(_row("Chair", "Example College", "2019-2020"))], entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "additional_dates")]
    assert "each holds one fact" in findings[0]["message"]


def test_offschema_a_further_span_pair_is_judged_as_a_date_on_its_row():
    """EBYSBC E22 (HZGJFM-03): `additional_period_start` is the year of a
    second term. With its row missing the year it is a lost date (WARN, on
    the record's line), not a fact shown elsewhere (INFO); once the row
    carries the span it is not reported."""
    entry = _fields_entry("P", {
        "committee_name": "Example Panel", "role": "Member",
        "start_date": "1990", "end_date": "1993",
        "additional_period_start": "2001", "additional_period_end": "present"})
    lost = [_table(_row("Example Panel", "Member", "1990-1993")),
            ("paragraph", "Elsewhere: 2001")]
    findings = _graded(lost, entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "additional_period_start")]
    assert "its span did not reach the record's date cell" in findings[0]["message"]
    assert "1 of them not on its record's line" in findings[0]["message"]
    shown = [_table(_row("Example Panel", "Member", "1990-1993, 2001-Present"))]
    assert _graded(shown, entry) == []


def test_offschema_a_further_span_is_not_reported_without_a_document():
    """EBYSBC E22: like any date, a further span is judged only against the
    rendered document. Without one, `additional_period_start` (a key the
    date-name pattern does not match) and `additional_periods` are left out
    rather than reported as unread facts."""
    entry = _anchored("P", {
        "start_date": "1990", "end_date": "1993",
        "additional_period_start": "2001", "additional_period_end": "2004",
        "additional_periods": [{"start_date": "2008", "end_date": "2009"}]})
    assert _offschema(entry) == []


def test_offschema_a_list_of_further_periods_is_dates_not_whole_records():
    """`additional_periods` holds `{start_date, end_date}` objects: spans of
    the record, graded on its row, never counted as records of their own."""
    entry = _fields_entry("O", {
        "leadership_role": "Chair", "institution": "Example College",
        "start_date": "2019", "end_date": "2020",
        "additional_periods": [{"start_date": "2022", "end_date": "2023"}]})
    findings = _graded([_table(_row("Chair", "Example College", "2019-2020"))], entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "additional_periods")]
    assert "whole record" not in findings[0]["message"]
    shown = _table(_row("Chair", "Example College", "2019-2020, 2022-2023"))
    assert _graded([shown], entry) == []


def test_offschema_a_span_key_on_a_code_no_date_cell_extends_is_unread():
    """EBYSBC E22: only the O/P/Q1/Q2/D/K date cells read the span keys. On
    H, which renders no further span, `additional_dates` is a key no
    renderer reads, and the message says so."""
    entry = _fields_entry("H", {
        "award_name": "Example Prize", "granting_body": "Example Society",
        "date": "2001", "additional_dates": "2004"})
    findings = _graded([_table(_row("Example Prize", "Example Society", "2001"))], entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "additional_dates")]
    assert "no renderer reads it" in findings[0]["message"]
    assert "date cell" not in findings[0]["message"]


def test_offschema_row_found_when_its_name_is_split_across_cells():
    """The honors renderer moves an award name's tail into its own cell, and
    the document writes a typographic apostrophe stage 4 wrote plainly."""
    entry = _fields_entry("H", {"award_name": "Reader's Choice Prize, First Edition",
                                "date": "1995", "end_date": "1998"})
    row = _row("Reader’s Choice Prize", "First Edition", "1995")
    findings = _graded([_table(row)], entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "end_date")]


def _owner(**extra):
    return _fields_entry("A", {"name": "Pat Owner", **extra}, idx=3)


_PERSONAL_TABLE = _table(_row("Office address:", "1 Example Way"),
                         _row("Work email:", "pat@example.org"))


def test_offschema_personal_data_value_shown_anywhere_is_not_reported():
    assert _graded([_PERSONAL_TABLE], _owner(institutional_email="pat@example.org")) == []


def test_offschema_personal_data_list_value_shown_anywhere_is_not_reported():
    """Every line is a Personal Data value's own, so a list of numbers the
    document shows is shown, though no row names the owner."""
    table = _table(_row("Office telephone:", "555-0100"))
    assert _graded([table], _owner(office_phones=["555-0100"])) == []


def test_offschema_personal_data_value_is_never_quoted():
    findings = _graded([_PERSONAL_TABLE], _owner(office_phone="555-0100", fax="555-0199"))
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("INFO", "fax"), ("WARN", "office_phone")]
    assert all(f["evidence"] == ["entry 3: (Personal Data value, not quoted)"]
               for f in findings)
    assert "555" not in json.dumps(findings)


def test_offschema_personal_data_entry_with_no_schema_value_is_still_judged():
    """Personal Data is a fixed table, not a raw-text fallback: an A entry
    whose every value is off its schema has lost them all."""
    entry = _fields_entry("A", {"office_phone": "555-0100"}, idx=3)
    findings = _graded([_PERSONAL_TABLE], entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "office_phone")]


@pytest.mark.parametrize("withheld", [
    {"place_of_birth": "Exampletown"}, {"spouse_name": "Lee Owner"},
    {"marital_status": "married"}, {"home_address": "2 Example Lane"},
])
def test_offschema_personal_data_the_policy_withholds_is_not_reported(withheld):
    assert _graded([_PERSONAL_TABLE], _owner(**withheld)) == []


def test_offschema_owner_contact_on_another_record_is_shown_by_personal_data():
    """A page header's email fused into a record is the owner's own,
    rendered in Personal Data -- not a fact the record lost."""
    role = _fields_entry("O", {"leadership_role": "Chair", "institution": "Example College",
                               "work_email": "pat@example.org"},
                         text="Chair, Example College\tEmail: pat@example.org", idx=5)
    blocks = [_PERSONAL_TABLE, _table(_row("Chair", "Example College", "2018"))]
    assert _graded(blocks, _owner(email="pat@example.org"), role) == []
    # The same address held by no Personal Data entry is only shown elsewhere.
    assert [f["severity"] for f in _graded(blocks, _owner(), role)] == ["INFO"]


def test_offschema_date_list_is_shown_only_when_its_row_shows_every_year():
    """An H `dates` of two years whose row shows one of them has lost the
    other: one year on the row is not the value shown."""
    entry = _fields_entry("H", {"award_name": "Silver Lamp Award",
                                "granting_body": "Example Society",
                                "dates": ["2019", "2021"]},
                          text="Silver Lamp Award, Example Society, 2019, 2021")
    one_year = _table(_row("Silver Lamp Award", "Example Society", "2021"))
    findings = _graded([one_year], entry)
    assert [(f["severity"], f["message"].split("`")[1]) for f in findings] == [
        ("WARN", "dates")]
    both = _table(_row("Silver Lamp Award", "Example Society", "2019, 2021"))
    assert _graded([both], entry) == []


def test_offschema_value_the_text_states_only_in_part_stays_info():
    """Sharing some words with the entry's text (the mentee's name) is not
    stating the value: below `RENDER_TOKEN_OVERLAP` it stays INFO."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil",
                                  "note": "Ann Pupil later left the program"},
                          text="Ann Pupil, summer student")
    findings = _graded([_table(_row("Name:", "Ann Pupil"))], entry)
    assert [f["severity"] for f in findings] == ["INFO"]
    assert "1 of them nowhere in the document" in findings[0]["message"]


@pytest.mark.parametrize("block", [
    _table(_row("Ann Pupil", "2001"), _row("Awards:", "Gold Ribbon")),
    ("p", "Name: | Ann Pupil\nAwards: | Gold Ribbon"),
], ids=["table-with-an-unlabelled-row", "paragraph"])
def test_offschema_only_a_wholly_labelled_table_is_a_form(block):
    """One labelled row does not make a table one record's form, and a
    paragraph is never one: the value's row is not the name's."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil", "awards": "Gold Ribbon"},
                          text="Ann Pupil; Gold Ribbon")
    findings = _graded([block], entry)
    assert [f["severity"] for f in findings] == ["INFO"]
    assert "the document shows it, but not with its record" in findings[0]["message"]


def test_offschema_date_key_on_a_code_declaring_no_date_is_not_reported():
    """Personal Data declares no date at all: a date key there is not the
    single-date shape `_is_offschema_date` admits."""
    assert _graded([_PERSONAL_TABLE], _owner(start_date="2004")) == []


def test_offschema_a_date_value_never_marks_a_records_row():
    """A date-named key's value, even one with words in it, is not an anchor:
    the row it sits on need not be the record's."""
    entry = _fields_entry("H", {"award_name": "Silver Lamp Award",
                                "date": "Founders Day 2004",
                                "citation": "for long service"},
                          text="Silver Lamp Award, Founders Day 2004, for long service")
    blocks = [_table(_row("Silver Lamp Award", "Example Society")),
              _table(_row("Founders Day 2004", "for long service"))]
    findings = _graded(blocks, entry)
    assert [f["severity"] for f in findings] == ["INFO"]


def _degree_beside_an_award(degree, institution="Lakeside College", **extra):
    degree_entry = _fields_entry("B1", {"degree": degree, "institution": institution,
                                        "honors": "with honors", **extra},
                                 text=f"{degree}, Lakeside College, with honors", idx=1)
    award = _fields_entry("H", {"award_name": "Lakeside Prize",
                                "granting_body": "Lakeside College"}, idx=2)
    return degree_entry, award


@pytest.mark.parametrize("degree, other_row", [
    ("MS Ed", ("Lakeside College", "Rams Edge", "with honors")),
    ("MS Ed", ("Lakeside College", "Ed Hall", "with honors")),
    ("MPhil", ("MPhil prize", "with honors")),
], ids=["short-anchor-inside-a-word", "short-anchor-one-word-of-two",
        "short-anchor-alone"])
def test_offschema_a_short_anchor_marks_a_row_only_as_all_its_words(degree, other_row):
    """A short anchor ("MS Ed") is carried only as whole words, all of them,
    and is never distinctive alone: a row with "Rams Edge", "Ed Hall" or
    another "MPhil" is not the degree's row."""
    blocks = [_table(_row(degree, "Lakeside College", "2001")), _table(_row(*other_row))]
    findings = _graded(blocks, *_degree_beside_an_award(degree))
    assert [(f["severity"], f["evidence"]) for f in findings] == [
        ("INFO", ["entry 1: with honors"])]


def test_offschema_an_anchor_only_its_object_holds_is_not_distinctive():
    """An institution written as an object is not among the record's own
    string values, so another record holding it makes it shared."""
    blocks = [_table(_row("BA", "Lakeside College", "2001")),
              _table(_row("Lakeside Prize", "Lakeside College, with honors"))]
    findings = _graded(blocks, *_degree_beside_an_award(
        "BA", institution={"name": "Lakeside College"}))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_an_anchor_under_two_keys_still_counts_once():
    """The same value under two keys is one anchor, not none."""
    degree = _fields_entry("B1", {"degree": "PhD", "discipline": "PhD",
                                  "institution": "Northfield University",
                                  "honors": "with distinction"}, idx=1)
    medal = _fields_entry("H", {"award_name": "Northfield Medal",
                                "granting_body": "Northfield University"}, idx=2)
    blocks = [_table(_row("PhD", "Northfield University, with distinction", "1991")),
              _table(_row("Northfield Medal", "Northfield University"))]
    assert _graded(blocks, degree, medal) == []


def test_offschema_a_row_carrying_most_shared_anchors_is_the_records():
    """Two of three shared anchors is a majority: the row is the record's."""
    degree = _fields_entry("B1", {"degree": "PhD", "institution": "Northfield University",
                                  "advisor": "Bea Guide", "honors": "with distinction"},
                           idx=1)
    medal = _fields_entry("H", {"award_name": "Northfield Medal",
                                "granting_body": "Northfield University",
                                "description": "nominated by Bea Guide"}, idx=2)
    blocks = [_table(_row("PhD", "Northfield University, with distinction", "1991")),
              _table(_row("Northfield Medal", "Northfield University"))]
    assert _graded(blocks, degree, medal) == []


def test_offschema_list_value_is_shown_only_when_its_row_shows_every_item():
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil",
                                  "awards": ["Gold Ribbon", "Blue Medal"]},
                          text="Ann Pupil; Gold Ribbon; Blue Medal")
    findings = _graded([_table(_row("Ann Pupil", "Gold Ribbon"))], entry)
    assert [f["severity"] for f in findings] == ["WARN"]


def test_offschema_a_string_with_no_words_is_never_stated_by_the_text():
    """A value with no letters or digits in one of its strings is never
    WARN: the entry's text cannot be said to state it."""
    entry = _fields_entry("N3B", {"mentee_name": "Ann Pupil",
                                  "awards": ["Gold Ribbon", "\u2014"]},
                          text="Ann Pupil; Gold Ribbon")
    findings = _graded([_table(_row("Ann Pupil", "2001"))], entry)
    assert [f["severity"] for f in findings] == ["INFO"]


def test_offschema_a_rendered_field_is_an_anchor(monkeypatch):
    """A key only `fan_out._RENDERED_FIELDS` names still writes the
    record's row, so it finds that row."""
    monkeypatch.setattr(extraction_lints, "_RENDERED_FIELDS",
                        {"I": frozenset({"chapter"})})
    entry = _fields_entry("I", {"chapter": "Delta Chapter Lodge",
                                "standing": "charter member"},
                          text="Delta Chapter Lodge, charter member")
    assert _graded([_table(_row("Delta Chapter Lodge", "charter member"))], entry) == []


def test_offschema_constants_are_pinned():
    assert extraction_lints.PERSONAL_DATA_ROW_WORDS == frozenset(
        {"address", "phone", "telephone"})
    assert extraction_lints._SINGLE_DATE_KEY == "date"


# ==========================================================================
# lint_implausible_year: a two-digit year given the wrong century.

def _implausible(*entries):
    return lint_implausible_year({"entries": list(entries)})


def test_implausible_year_fires_on_a_two_digit_year_given_the_wrong_century():
    findings = _implausible(_fields_entry(
        "R", {"date": "1902-11"}, text="Invited talk, Example City 11/02", idx=40))
    assert len(findings) == 1
    assert findings[0]["lint"] == "implausible_year"
    assert findings[0]["severity"] == "WARN"
    assert "entry 40 (R): date=1902" in findings[0]["message"]
    assert f"before {IMPLAUSIBLE_YEAR_FLOOR}" in findings[0]["message"]
    assert findings[0]["evidence"] == ["Invited talk, Example City 11/02"]


def test_implausible_year_floor_is_1930_and_exclusive():
    assert IMPLAUSIBLE_YEAR_FLOOR == 1930
    fires = _implausible(_fields_entry("D1", {"start_date": "1929"}, text="7/29"))
    silent = _implausible(_fields_entry("D1", {"start_date": "1930"}, text="7/30"))
    assert len(fires) == 1 and silent == []


def test_implausible_year_lists_every_bad_field_of_one_entry_in_one_finding():
    findings = _implausible(_fields_entry(
        "P", {"start_date": "1902", "end_date": "1904", "year": "2012"},
        text="Committee 9/02-4/04, renewed 2012"))
    assert len(findings) == 1
    assert "start_date=1902, end_date=1904 --" in findings[0]["message"]
    assert "year=" not in findings[0]["message"]


@pytest.mark.parametrize("entry", [
    _fields_entry("H", {"date": "1925"}, text="Society prize, 1925"),       # in the text
    _fields_entry("H", {"date": "1925"}, text="Prize 192529"),              # fused range
    _fields_entry("H", {"title": "1902"}, text="11/02"),                    # not a date key
    _fields_entry("H", {"candidate": "1902"}, text="11/02"),
    _fields_entry("H", {"date": "19021"}, text=""),                         # not four digits
    _fields_entry("H", {"date": "21902"}, text=""),
    _fields_entry("A", {"date_of_birth": "1902"}, text="11/02"),           # personal data
    _fields_entry("H", {"date": "2002-11"}, text="11/02"),                 # right century
    _fields_entry("H", {"date": True}, text=""),                           # not a year
    _fields_entry("H", "not an object", text=""),
])
def test_implausible_year_silent(entry):
    assert _implausible(entry) == []


@pytest.mark.parametrize("value", [
    {"start_date": "1905"},   # nested object
    ["1905"],                 # list
    1905,                     # integer
])
def test_implausible_year_reads_nested_and_integer_values(value):
    assert len(_implausible(_fields_entry("B1", {"dates_attended": value}, text="'05"))) == 1


def _degree(year, text=None, code="B1"):
    return _fields_entry(code, {"degree": "MD", "year": year},
                         text=f"MD, Example University {year}" if text is None else text)


def test_implausible_year_floor_rises_to_ten_years_before_the_earliest_degree():
    floor = 1990 - DEGREE_YEAR_LEAD
    old = _fields_entry("H", {"date": str(floor - 1)}, text="Example prize")
    edge = _fields_entry("H", {"date": str(floor)}, text="Example prize")
    findings = _implausible(_degree("1995"), _degree("1990"), old, edge)
    assert len(findings) == 1
    assert f"date={floor - 1}" in findings[0]["message"]
    assert f"earliest degree year, 1990" in findings[0]["message"]


@pytest.mark.parametrize("degree", [
    _degree("1990", text="MD, Example University"),       # degree year not written
    _degree("1990", code="B2"),                           # not an academic degree
    _degree("1925"),                                      # degree itself pre-floor
])
def test_implausible_year_degree_floor_needs_a_written_plausible_b1_year(degree):
    later = _fields_entry("H", {"date": "1975"}, text="Prize '75")
    assert _implausible(degree, later) == []


def test_implausible_year_text_match_is_bounded_on_the_left():
    """1902 inside a longer number is not the year written."""
    findings = _implausible(_fields_entry("R", {"date": "1902"}, text="No. 21902, 11/02"))
    assert len(findings) == 1


def test_implausible_year_evidence_is_the_truncated_entry_text():
    text = "11/02 " + "y" * (FIELD_EVIDENCE_VALUE_CHARS + 20)
    findings = _implausible(_fields_entry("R", {"date": "1902"}, text=text))
    assert findings[0]["evidence"] == [text[:FIELD_EVIDENCE_VALUE_CHARS]]


def test_implausible_year_degree_floor_applies_only_above_the_fixed_floor():
    """A degree exactly DEGREE_YEAR_LEAD years above the fixed floor sets no
    floor of its own: the fixed floor and its reason stand."""
    findings = _implausible(
        _degree(str(IMPLAUSIBLE_YEAR_FLOOR + DEGREE_YEAR_LEAD)),
        _fields_entry("H", {"date": "1929"}, text="'29"))
    assert len(findings) == 1
    assert "no two-digit year resolves below it" in findings[0]["message"]


def test_implausible_year_a_pre_floor_degree_year_does_not_mask_a_later_one():
    """A degree year that is itself a wrong century is not the earliest
    degree: the next written B1 year sets the floor."""
    findings = _implausible(_degree("1925"), _degree("1990"),
                            _fields_entry("H", {"date": "1975"}, text="Example prize"))
    assert len(findings) == 1
    assert "earliest degree year, 1990" in findings[0]["message"]


def test_implausible_year_degree_floor_never_drops_below_the_fixed_floor():
    findings = _implausible(_degree("1935"),
                            _fields_entry("H", {"date": "1929"}, text="'29"))
    assert len(findings) == 1
    assert f"before {IMPLAUSIBLE_YEAR_FLOOR}" in findings[0]["message"]


def test_implausible_year_a_two_digit_year_in_its_pivot_century_is_stated():
    """s7ab QZWBKQ's shape: a CV writing every date as m/yy. "9/68" read as
    1968 is the right century, and a B1 "5/72" sets the degree floor (a
    Ph.D. year written in four digits used to set it alone, decades late)."""
    findings = _implausible(
        _fields_entry("B1", {"dates_attended": {"start_date": "1968-09", "end_date": "1972-05"}},
                      text="9/68-5/72 Example College; B.A."),
        _fields_entry("B1", {"dates_attended": {"start_date": "1995-08", "end_date": "1999-05"}},
                      text="8/1995-5/1999 Example Institute; Ph.D."),
        _fields_entry("C", {"start_date": "1978", "end_date": "1979"},
                      text="7/78-6/79 Internship, Example Hospital"))
    assert findings == []


def test_implausible_year_degree_floor_reads_a_two_digit_degree_year():
    """The B.A. "9/68" sets the floor (1958), not the four-digit Ph.D. year
    alone (1990): an undated 1975 is above the floor, so it is no wrong
    century -- year_not_in_source reports it instead."""
    entries = [
        _fields_entry("B1", {"dates_attended": {"start_date": "1968-09", "end_date": "1972-05"}},
                      text="9/68-5/72 Example College; B.A."),
        _fields_entry("B1", {"year": "2000"}, text="Ph.D., Example Institute, 2000"),
        _fields_entry("H", {"date": "1975"}, text="Example prize", idx=60)]
    assert _implausible(*entries) == []
    assert len(lint_year_not_in_source({"entries": entries})) == 1


def test_implausible_year_a_range_shorthand_end_is_written():
    """"1975-79" writes 1979, below a 1990 degree's floor: not a wrong
    century."""
    findings = _implausible(_degree("1990"), _fields_entry(
        "P", {"start_date": "1975", "end_date": "1979"}, text="Example committee 1975-79"))
    assert findings == []


def test_implausible_year_judges_the_date_a_formatter_rendered():
    """EBYSBC MIFYLG 234's shape: stage 4 wrote 1900 for a written-out date;
    5c's text, which stage 6 renders instead of the field, has the text's
    year. Judged only when the formatter's text drops the stage-4 year."""
    entry = _fields_entry("K4", {"date": "1900-03-03"},
                          text="12. Example lecture to fellows, March 3, 2014.", idx=34)
    entry["element_idx_end"] = 35

    def stage5d(formatted):
        return {"entries": [{"element_idx_start": 34, "element_idx_end": 35,
                             "extracted_fields": {"formatted_text": formatted}}]}
    assert len(lint_implausible_year({"entries": [entry]})) == 1
    assert lint_implausible_year({"entries": [entry]},
                                 stage5d("**2014-03-03** - Example lecture")) == []
    assert len(lint_implausible_year({"entries": [entry]},
                                     stage5d("**1900-03-03** - Example lecture"))) == 1
    # A span two 5d entries share says nothing about which one this is.
    shared = stage5d("**2014-03-03** - Example lecture")
    shared["entries"] *= 2
    assert len(lint_implausible_year({"entries": [entry]}, shared)) == 1


def test_implausible_year_ceiling_is_info_for_a_source_typo_warn_otherwise():
    """EBYSBC VNUAHA 175's shape: the source itself mistypes the end year.
    Stage 4 copied it faithfully; a reviewer should still see it."""
    assert IMPLAUSIBLE_YEAR_CEILING == 2100
    typo = _implausible(_fields_entry("P", {"start_date": "2013", "end_date": "7013"},
                                      text="Example committee, 2013-7013", idx=175))
    assert len(typo) == 1 and typo[0]["severity"] == "INFO"
    assert "entry 175 (P): end_date=7013 -- after 2100" in typo[0]["message"]
    made_up = _implausible(_fields_entry("P", {"end_date": "2201"},
                                         text="Example committee, 2010-2012"))
    assert len(made_up) == 1 and made_up[0]["severity"] == "WARN"
    assert _implausible(_fields_entry("P", {"end_date": "2100"}, text="until 2100")) == []


def test_implausible_year_leaves_the_value_alone():
    entry = _fields_entry("R", {"date": "1902-11"}, text="11/02")
    _implausible(entry)
    assert entry["extracted_fields"] == {"date": "1902-11"}


# ==========================================================================
# lint_multi_record_coverage: an entry holding several records that stage 4
# returned as one (#1243). Every name, place and value below is invented.

_TWO_ROLES = ("Lecturer, Northfield University School of Medicine, 2001-2005 "
              "Visiting Instructor of Pathology, Lakeside Hospital Institute, 2006-2008")
_FIRST_ROLE = {"title": "Lecturer", "institution": "Northfield University School of Medicine",
               "start_date": "2001", "end_date": "2005"}
_FIRST_ROLE_ROW = "Lecturer | Northfield University School of Medicine | 2001-2005"


def _multi_record(text, fields, code="D1", lines=(_FIRST_ROLE_ROW,)):
    entry = {"taxonomy_code": code, "element_idx_start": 7, "text": text,
             "extracted_fields": fields}
    return lint_multi_record_coverage({"entries": [entry]}, [("p", line) for line in lines])


def test_multi_record_second_dated_clause_absent_is_warn():
    findings = _multi_record(_TWO_ROLES, _FIRST_ROLE)
    assert findings == [{
        "lint": "multi_record_coverage", "severity": "WARN",
        "message": "entry 7 (D1): 2 record-shaped clauses, one stage-4 record; "
                   "1 other clause(s) on no line of the output (#1243)",
        "evidence": ["Visiting Instructor of Pathology, Lakeside Hospital Institute, 2006-2008"],
        "status": "ran", "reason": ""}]


def test_multi_record_second_clause_on_a_rendered_line_is_info():
    findings = _multi_record(_TWO_ROLES, _FIRST_ROLE, lines=(
        _FIRST_ROLE_ROW,
        "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2006-2008"))
    assert [f["severity"] for f in findings] == ["INFO"]
    assert findings[0]["message"].endswith("the other clause(s) are in the output (#1243)")
    assert findings[0]["evidence"] == [
        "Visiting Instructor of Pathology, Lakeside Hospital Institute, 2006-2008"]


@pytest.mark.parametrize("lines", [
    # the clause's words on one line, its years on another
    (_FIRST_ROLE_ROW, "Visiting Instructor of Pathology | Lakeside Hospital Institute",
     "Term 2006-2008"),
    # the same, as two rows of one table block
    (_FIRST_ROLE_ROW,
     "Visiting Instructor of Pathology\tLakeside Hospital Institute\nTerm 2006-2008"),
    # a year inside a longer number is not that year
    (_FIRST_ROLE_ROW,
     "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2008 | Room 120064"),
    # a year with a digit before it, or after it, is not that year
    (_FIRST_ROLE_ROW,
     "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2008 | Room 12006"),
    (_FIRST_ROLE_ROW,
     "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2008 | Room 20061"),
])
def test_multi_record_one_line_must_carry_the_clause_words_and_years(lines):
    """The render check reads line by line, a table block's rows included:
    words on one line and years on another are not one rendered record."""
    findings = _multi_record(_TWO_ROLES, _FIRST_ROLE, lines=lines)
    assert [f["severity"] for f in findings] == ["WARN"]


def test_multi_record_through_joins_two_years_into_one_date():
    text = ("Lecturer, Northfield University School of Medicine, 2001 through 2005\t"
            "Visiting Instructor of Pathology, Lakeside Hospital Institute, 2006-2008")
    findings = _multi_record(text, {
        "title": "Visiting Instructor of Pathology", "institution": "Lakeside Hospital Institute",
        "start_date": "2006", "end_date": "2008"}, lines=(
        "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2006-2008",))
    assert findings[0]["message"].startswith("entry 7 (D1): 2 record-shaped clauses")
    assert findings[0]["evidence"] == [
        "Lecturer, Northfield University School of Medicine, 2001 through 2005"]


@pytest.mark.parametrize("prefix", ["", "June ", "July ", "The "])
def test_multi_record_a_part_opening_with_a_date_gives_each_date_the_words_after_it(prefix):
    """A month or a filler word before the first year still opens the part
    with a date."""
    text = f"{prefix}2013 Example Workshop, Lakeport; June 2014 Sample Seminar, Riverton"
    findings = _multi_record(text, {
        "title": "Example Workshop", "location": "Lakeport", "date": "2013"},
        code="R", lines=())
    assert findings[0]["evidence"] == ["2014 Sample Seminar, Riverton"]


@pytest.mark.parametrize("closing", [")", ".)", " )", ", 2)"])
def test_multi_record_a_year_that_closes_an_aside_leaves_it(closing):
    """'(Poster Prize 2012.)': the words after the closing parenthesis are
    outside the aside, so they are a clause of their own."""
    text = (f"2011 Example Workshop, Lakeport (Poster Prize 2012{closing} "
            "Sample Seminar on Wound Care, Riverton")
    findings = _multi_record(text, {
        "title": "Example Workshop", "location": "Lakeport", "date": "2011"},
        code="R", lines=())
    assert [f["severity"] for f in findings] == ["WARN"]


@pytest.mark.parametrize("separator", ["\n", "\t", " | "])
def test_multi_record_each_strongly_separated_part_reads_its_own_layout(separator):
    """A line, cell or pipe-joined cell ends a part: a part that opens with a
    date gives it the words after it even when the part before it put its
    date last."""
    text = f"Example Workshop, Lakeport, 2005{separator}2006 Sample Seminar, Riverton"
    findings = _multi_record(text, {
        "title": "Example Workshop", "location": "Lakeport", "date": "2005"},
        code="R", lines=())
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["evidence"] == ["2006 Sample Seminar, Riverton"]


def test_multi_record_a_three_letter_word_before_the_first_date_is_a_payload():
    """'Art Lab, <years>; <committee>, <years>' puts each date after its
    record: a three-letter name before the first date is words, not a date
    prefix."""
    text = "Art Lab, 2001-2003; Sample Seminar Committee, Riverton, 2004-2006"
    findings = _multi_record(text, {
        "committee_name": "Art Lab", "start_date": "2001", "end_date": "2003"},
        code="P", lines=())
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["evidence"] == ["Sample Seminar Committee, Riverton, 2004-2006"]


def test_multi_record_an_open_parenthesis_after_a_late_date_is_no_aside():
    """Only a part that opens with a date reads a parenthesis left open after
    a later date as that clause's aside; a part that puts its dates last
    keeps the clause before the date."""
    text = ("Lecturer (Northfield University School of Medicine, 2001-2005), "
            "Visiting Instructor of Pathology (Lakeside Hospital Institute, 2006-2008; adjunct)")
    findings = _multi_record(text, _FIRST_ROLE, lines=())
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "2 record-shaped clauses" in findings[0]["message"]


def test_multi_record_the_record_s_own_clause_is_never_reported():
    """The clause the record stands for is skipped even when it carries
    words the record lacks."""
    findings = _multi_record(_TWO_ROLES, {
        "title": "Lecturer", "institution": "Northfield", "start_date": "2001",
        "end_date": "2005"}, lines=())
    assert "1 other clause(s)" in findings[0]["message"]
    assert findings[0]["evidence"] == [
        "Visiting Instructor of Pathology, Lakeside Hospital Institute, 2006-2008"]


@pytest.mark.parametrize("talk", [
    "2014 Lecture on wound healing outcomes",
    # a date word is not a lowercase word
    "2014 Seminar on wound healing outcomes, ongoing",
])
def test_multi_record_a_short_lowercase_clause_is_not_prose(talk):
    """Three lowercase 4+-letter words, and short words like 'on', are below
    MULTI_RECORD_PROSE_MIN_LOWERCASE: a talk title in sentence case is a
    record."""
    text = f"2013 Example Workshop, Lakeport\t{talk}"
    findings = _multi_record(text, {
        "title": "Example Workshop", "location": "Lakeport", "date": "2013"},
        code="R", lines=())
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["evidence"] == [talk]


def test_multi_record_the_first_clause_is_never_an_aside():
    """An aside qualifies the clause before it; the first clause has none."""
    text = "(2010 Example Workshop, Lakeport) 2011 Sample Seminar, Riverton"
    findings = _multi_record(text, {
        "title": "Sample Seminar", "location": "Riverton", "date": "2011"},
        code="R", lines=())
    assert findings[0]["evidence"] == ["2010 Example Workshop, Lakeport"]


def test_multi_record_a_sentence_case_title_in_a_capitalised_clause_is_not_prose():
    """Four lowercase words are not a sentence when capitalised ones
    outnumber them."""
    text = ("2013 Example Workshop, Lakeport\t"
            "2014 Northfield Surgical Society Annual Meeting, Riverton: wound repair after burns")
    findings = _multi_record(text, {
        "title": "Example Workshop", "location": "Lakeport", "date": "2013"},
        code="R", lines=())
    assert [f["severity"] for f in findings] == ["WARN"]


def test_multi_record_clause_verdict_comes_before_the_fused_values():
    """An entry with two dated clauses and two mentees gets the clause
    finding, which names what is missing."""
    text = ("2015-2017 Ada Quill, Resident, Lakeside Hospital\t"
            "2018-2020 Bo Renner, Fellow, Northfield University")
    findings = _multi_record(text, {
        "mentee_name": "Ada Quill and Bo Renner", "site_position": "Lakeside Hospital",
        "start_date": "2015", "end_date": "2017"}, code="N3B", lines=())
    assert len(findings) == 1
    assert findings[0]["message"].startswith("entry 7 (N3B): 2 record-shaped clauses")


def test_multi_record_same_words_beside_other_years_do_not_vouch():
    """The same title held at another time is a different record."""
    findings = _multi_record(_TWO_ROLES, _FIRST_ROLE, lines=(
        _FIRST_ROLE_ROW,
        "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2011-2014"))
    assert [f["severity"] for f in findings] == ["WARN"]


def test_multi_record_the_record_may_stand_for_any_clause():
    """Stage 4 kept the second role; the first is the one left out."""
    findings = _multi_record(_TWO_ROLES, {
        "title": "Visiting Instructor of Pathology", "institution": "Lakeside Hospital Institute",
        "start_date": "2006", "end_date": "2008"}, lines=(
        "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2006-2008",))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["evidence"] == ["Lecturer, Northfield University School of Medicine, 2001-2005"]


def test_multi_record_render_check_looks_for_the_words_the_record_lacks():
    """A line carrying the clause's shared words (its institution) and a
    word or two of its own does not render the left-out role."""
    text = ("Lecturer, Northfield University School of Medicine, 2001-2005\t"
            "Visiting Instructor of Pathology, Northfield University School of Medicine, 2006-2008")
    findings = _multi_record(text, _FIRST_ROLE, lines=(
        _FIRST_ROLE_ROW,
        "Pathology Seminar | Northfield University School of Medicine | 2006-2008"))
    assert [f["severity"] for f in findings] == ["WARN"]


def test_multi_record_dates_with_no_words_between_them_stay_one_clause():
    """'2006; 2008' are two years of one role: both must be on the line."""
    text = ("Lecturer, Northfield University School of Medicine, 2001-2005\t"
            "Visiting Instructor of Pathology, Lakeside Hospital Institute, 2006; 2008")
    findings = _multi_record(text, _FIRST_ROLE, lines=(
        _FIRST_ROLE_ROW, "Visiting Instructor of Pathology | Lakeside Hospital Institute | 2006"))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert "2 record-shaped clauses" in findings[0]["message"]


def test_multi_record_and_joins_two_years_into_one_date():
    """'2003 and 2005' is one date of one workshop, not a clause of its own."""
    text = "2003 and 2005 Example Workshop, Lakeport; 2006 Sample Seminar, Riverton"
    findings = _multi_record(text, {
        "title": "Example Workshop", "location": "Lakeport", "date": "2003; 2005"},
        code="R", lines=())
    assert findings[0]["message"].startswith("entry 7 (R): 2 record-shaped clauses")
    assert findings[0]["evidence"] == ["2006 Sample Seminar, Riverton"]


def test_multi_record_two_digit_end_year_stays_in_its_date():
    """'1991-93' is one date: its '93' does not open the next clause."""
    text = ("Junior Delegate, Planning Council, 1991-93; "
            "Senior Delegate, Curriculum Board, 1993-95")
    findings = _multi_record(text, {
        "role": "Junior Delegate", "committee_name": "Planning Council",
        "start_date": "1991", "end_date": "1993"}, code="P", lines=())
    assert findings[0]["evidence"] == ["Senior Delegate, Curriculum Board, 1993-95"]


def test_multi_record_reads_track_changed_insertions():
    """The clause the record left out, rendered as a tracked insertion, is in
    the output: the lint reads the doctor's own docx view, which includes
    `w:ins` text."""
    doc = Document()
    doc.add_paragraph(_FIRST_ROLE_ROW)
    doc.add_paragraph()._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-10-02T00:00:00Z">'
        '<w:r><w:t>Visiting Instructor of Pathology, Lakeside Hospital Institute, '
        '2006-2008</w:t></w:r></w:ins>'))
    entry = {"taxonomy_code": "D1", "element_idx_start": 7, "text": _TWO_ROLES,
             "extracted_fields": _FIRST_ROLE}
    findings = lint_multi_record_coverage({"entries": [entry]}, docx_body_blocks(doc))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_multi_record_year_led_clauses_take_the_words_after_their_year():
    text = ("2013 Society of Example Medicine, Session chair, Riverton "
            "2013 Academy of Sample Surgery, Grand rounds on wound healing, Lakeport")
    findings = _multi_record(text, {
        "event_name": "Society of Example Medicine", "role": "Session chair",
        "location": "Riverton", "date": "2013"}, code="R", lines=())
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["evidence"] == [
        "2013 Academy of Sample Surgery, Grand rounds on wound healing, Lakeport"]


def test_multi_record_evidence_is_capped_and_truncated():
    talks = " ".join((
        "2013 Alder Grove Workshop, Lakeport", "2013 Birch Hollow Seminar, Riverton",
        "2013 Cedar Ridge Colloquium, Northfield", "2013 Douglas Fjord Lecture, Eastvale",
        "2013 Elmwood Basin Symposium, Westbrook"))
    findings = _multi_record(talks, {
        "title": "Alder Grove Workshop", "location": "Lakeport", "date": "2013"},
        code="R", lines=())
    assert "4 other clause(s)" in findings[0]["message"]
    assert len(findings[0]["evidence"]) == FIELD_EVIDENCE_MAX_VALUES
    long_clause = "2013 " + "Wordy " * 40 + "Symposium"
    findings = _multi_record(f"2015 Short Talk, Lakeport {long_clause}", {
        "title": "Short Talk", "location": "Lakeport", "date": "2015"}, code="R", lines=())
    assert len(findings[0]["evidence"][0]) == FIELD_EVIDENCE_VALUE_CHARS


@pytest.mark.parametrize("separator", ["\t", "; "])
def test_multi_record_undated_title_and_institution_parts(separator):
    """A hospital post with the concurrent faculty rank in the next cell, one
    date range for both (the TAUBPU-01 shape)."""
    text = (f"2009-2013 Attending Surgeon, Lakeside Hospital, Riverton{separator}"
            "Assistant Professor, Northfield University College of Medicine")
    findings = _multi_record(text, {
        "title": "Attending Surgeon", "institution": "Lakeside Hospital, Riverton",
        "start_date": "2009", "end_date": "2013"}, code="D2",
        lines=("Attending Surgeon | Lakeside Hospital, Riverton | 2009-2013",))
    assert [f["severity"] for f in findings] == ["WARN"]
    assert findings[0]["evidence"] == [
        "Assistant Professor, Northfield University College of Medicine"]


@pytest.mark.parametrize("second_part", [
    "Assistant Professor of Surgery",  # a title, no institution
    "Northfield University College of Medicine",  # an institution, no title
    # a PI named in a credit line is not the owner's role
    "Principal Investigator, Northfield University Research Institute",
])
def test_multi_record_an_undated_part_needs_a_title_and_an_institution(second_part):
    text = f"Attending Surgeon, Lakeside Hospital, Riverton\t{second_part}"
    assert _multi_record(text, {
        "title": "Attending Surgeon", "institution": "Lakeside Hospital, Riverton"},
        code="D2", lines=()) == []


def test_multi_record_new_years_count_toward_the_minimum():
    """One new word and one new year: a nationwide term after a regional one."""
    text = ("Society of Example Medicine\tTreasurer (Western Region), 2002-2004\t"
            "Treasurer (Nationwide), 2004-2008")
    findings = _multi_record(text, {
        "role": "Treasurer (Western Region)", "organization": "Society of Example Medicine",
        "start_date": "2002", "end_date": "2004"}, code="Q1", lines=(
        "Treasurer (Western Region) | Society of Example Medicine | 2002-2004",))
    assert [f["severity"] for f in findings] == ["WARN"]


def test_multi_record_a_clause_with_one_new_word_is_checked_by_all_its_words():
    """A line carrying only the clause's one new word, beside its years, does
    not render it: with fewer than MULTI_RECORD_MIN_NEW new words the render
    check looks for all of them."""
    text = ("Society of Example Medicine\tTreasurer (Western Region), 2002-2004\t"
            "Treasurer (Nationwide), 2004-2008")
    findings = _multi_record(text, {
        "role": "Treasurer (Western Region)", "organization": "Society of Example Medicine",
        "start_date": "2002", "end_date": "2004"}, code="Q1", lines=(
        "Treasurer (Western Region) | Society of Example Medicine | 2002-2004",
        "Nationwide Board | 2004-2008"))
    assert [f["severity"] for f in findings] == ["WARN"]


@pytest.mark.parametrize("text, fields", [
    # only new years: the same role held again is not reported
    ("Member, Library Committee, 2014-2015\tMember, Library Committee, 2009-2010",
     {"role": "Member", "committee_name": "Library Committee",
      "start_date": "2014", "end_date": "2015"}),
    # one new word and one new year out of six: below MULTI_RECORD_NEW_SHARE
    ("Imaging Search Committee, Department of Radiology Science, 2012-2013\t"
     "Genetics Search Committee, Department of Radiology Science, 2011",
     {"committee_name": "Imaging Search Committee",
      "institution": "Department of Radiology Science", "start_date": "2012",
      "end_date": "2013"}),
    # the record holds both clauses' words
    ("Member, Society of Example Medicine, 2001-2003; Chair, Society of Example Medicine, "
     "2004-2006", {"role": "Member; Chair", "organization": "Society of Example Medicine",
                   "start_date": "2001", "end_date": "2006"}),
    # several years of one record: dates with no words between them
    ("Visiting Professor, Northfield University, 2013 and 2015, 2017",
     {"title": "Visiting Professor", "institution": "Northfield University",
      "start_date": "2013"}),
    # the same role again: its only new words are date words
    ("Member, Society of Example Medicine, 2001-2005\t"
     "Member, Society of Example Medicine, January 2006 - present",
     {"role": "Member", "committee_name": "Society of Example Medicine",
      "start_date": "2001", "end_date": "2005"}),
    # one new word and no new year: below MULTI_RECORD_MIN_NEW
    ("Member, Society of Example Medicine, 2005\tFellow, 2005",
     {"role": "Member", "committee_name": "Society of Example Medicine", "start_date": "2005"}),
    # a connective is not a distinctive word
    ("Mentor, Northfield University, 2019\tMentor through Lakeside, 2019",
     {"role": "Mentor", "committee_name": "Northfield University", "start_date": "2019"}),
    # an open end ("ongoing", "present") is a date word, not a new word
    ("Treasurer, Example Society, 2001-2005\tTreasurer, 2006 - ongoing",
     {"role": "Treasurer", "organization": "Example Society", "start_date": "2001",
      "end_date": "2005"}),
    ("Treasurer, Example Society, 2001-2005\t2006 - present Treasurer",
     {"role": "Treasurer", "organization": "Example Society", "start_date": "2001",
      "end_date": "2005"}),
    # two new words out of four words and two years: the years count toward
    # MULTI_RECORD_NEW_SHARE's denominator
    ("Member, Library Committee, 2014-2015\tSenior Chair, Library Committee, 2014-2015",
     {"role": "Member", "committee_name": "Library Committee", "start_date": "2014",
      "end_date": "2015"}),
    # a year inside a longer number is no date
    ("Lecturer, Northfield University School of Medicine, 2001-2005, Suite 42019",
     {"title": "Lecturer", "institution": "Northfield University School of Medicine",
      "start_date": "2001", "end_date": "2005"}),
])
def test_multi_record_silent_when_the_other_clause_is_no_new_record(text, fields):
    assert _multi_record(text, fields, code="P", lines=()) == []


@pytest.mark.parametrize("text, fields, code", [
    # an aside that never closes inside its clause
    ("Residency in Family Medicine, Northfield University 2004-2007 "
     "(on the combined track, ended 2008)",
     {"training_type": "Residency", "specialty": "Family Medicine",
      "institution": "Northfield University", "start_date": "2004", "end_date": "2007"}, "C"),
    # an aside an earlier clause opened, after a year-led clause
    ("2010-2013 Bo Renner, Resident, Lakeside Hospital (won the Example Award 2012, "
     "Annual Meeting, Lakeport).",
     {"mentee_name": "Bo Renner", "site_position": "Lakeside Hospital",
      "mentee_level": "Resident", "start_date": "2010", "end_date": "2013"}, "N3B"),
    # a sentence about the record
    ("2014 Example Leadership Award, Lakeside Clinic\tgiven for steady dedication and "
     "service to the clinic and its patients over many years 2001-2014",
     {"award_name": "Example Leadership Award", "granting_body": "Lakeside Clinic",
      "date": "2014"}, "H"),
    # how the owner took part, not another record
    ("Invited Speaker, 2019\tExample Symposium on Wound Healing, Lakeport, 2019",
     {"title": "Wound Healing", "event_name": "Example Symposium on Wound Healing",
      "location": "Lakeport", "date": "2019"}, "R"),
    # an aside that closes with nothing but a year after it
    ("Member, Example Society, 2001-2005 (Section on Rural Practice) 2003",
     {"role": "Member", "organization": "Example Society", "start_date": "2001",
      "end_date": "2005"}, "I"),
    # a bracketed note on the record
    ("2011 Wound Care Update, Example Institute, Lakeport\t"
     "June 2011. [6.5 continuing education credits]",
     {"title": "Wound Care Update", "event_name": "Example Institute",
      "location": "Lakeport", "date": "2011"}, "R"),
    # a part's opening label names a field, not a value
    ("Fellow, Society of Example Medicine, 2010-2014\t"
     "Committees: Society of Example Medicine, 2012",
     {"membership_type": "Fellow", "organization": "Society of Example Medicine",
      "start_date": "2010", "end_date": "2014"}, "I"),
])
def test_multi_record_silent_on_a_clause_that_is_the_records_own_detail(text, fields, code):
    assert _multi_record(text, fields, code=code, lines=()) == []


@pytest.mark.parametrize("code", ["S1", "K1", "T", "A", "J", "N4"])
def test_multi_record_skips_the_codes_offschema_skips(code):
    assert _multi_record(_TWO_ROLES, _FIRST_ROLE, code=code) == []


def test_multi_record_skips_an_entry_stage4_returned_several_records_for():
    """Stage 4 returned two records, so this is not the one-record shape, even
    where its second record holds fewer of the text's words than the text."""
    fields = {**_FIRST_ROLE, STAGE4_RECORDS_KEY: [
        {"title": "Visiting Instructor"}, dict(_FIRST_ROLE)]}
    assert _multi_record(_TWO_ROLES, fields) == []


def test_multi_record_is_silent_on_a_record_list_fan_out_splits():
    """Fan-out splits a list only when its values hold every word of the
    text, so every clause is covered and the lint needs no fan-out skip."""
    entry_fields = {"committee_name": "Alpha Board", "committees": [
        {"committee_name": "Alpha Board", "role": "Chair"},
        {"committee_name": "Beta Panel", "role": "Member"}]}
    text = "Alpha Board Chair 2001-2003\tBeta Panel Member 2004-2006"
    assert _multi_record(text, entry_fields, code="P", lines=()) == []


@pytest.mark.parametrize("names", [
    "Ada Quill; Bo Renner; Cy Tamsin",
    "Ada Quill, Bo Renner and Cy Tamsin",
])
@pytest.mark.parametrize("lines, severity, tail", [
    # every mentee is on a line: fused into the record's row
    (("Ada Quill, Bo Renner, Cy Tamsin | Resident | 2021",),
     "INFO", "each is on a line of the output"),
    # the third mentee is on no line
    (("Ada Quill, Bo Renner | Resident | 2021",), "WARN", "1 on no line of the output"),
])
def test_multi_record_several_mentees_in_one_record(names, lines, severity, tail):
    findings = _multi_record(f"Cohort of 2021: {names}", {
        "mentee_name": names, "start_date": "2021"}, code="N3B", lines=lines)
    assert [f["severity"] for f in findings] == [severity]
    assert findings[0]["message"] == (
        f"entry 7 (N3B): stage 4 returned one record naming 3 mentees; {tail} (#1243)")


def test_multi_record_a_mentee_name_with_no_long_word_counts_as_on_the_output():
    """A name with no 5+-letter word cannot be looked for on a line."""
    findings = _multi_record("Cohort of 2021: Al Fox and Ed Roe", {
        "mentee_name": "Al Fox and Ed Roe", "start_date": "2021"}, code="N3B",
        lines=("Residents | 2021",))
    assert [f["severity"] for f in findings] == ["INFO"]


def test_multi_record_a_mentee_name_that_is_not_a_string_is_skipped():
    assert _multi_record("2021 Ada Quill and Bo Renner", {
        "mentee_name": ["Ada Quill", "Bo Renner"], "start_date": "2021"},
        code="N3B", lines=()) == []


@pytest.mark.parametrize("names", ["Ada Quill", "Quill, Ada", "Ada Quill, PhD"])
def test_multi_record_one_mentee_is_silent(names):
    assert _multi_record(f"2021 {names}", {"mentee_name": names, "start_date": "2021"},
                         code="N3B", lines=()) == []


@pytest.mark.parametrize("lines, severity, tail", [
    # both years on the record's row
    (("Ph.D., Master of Arts | Northfield University | 1999; 1995",),
     "INFO", "each is on a line of the output"),
    # 1995 only on a line that is not the record's
    (("Ph.D., Master of Arts | Northfield University | 1999", "Example Lecture Series, 1995"),
     "WARN", "1 on no line of the output"),
])
def test_multi_record_two_degree_years_in_one_record(lines, severity, tail):
    text = "Doctoral Program: Northfield University\tDegree: Ph.D., 1999; Master of Arts, 1995"
    findings = _multi_record(text, {
        "degree": "Ph.D., Master of Arts", "institution": "Northfield University",
        "year": "1999; 1995"}, code="B1", lines=lines)
    assert [f["severity"] for f in findings] == [severity]
    assert findings[0]["message"] == (
        f"entry 7 (B1): stage 4 returned one record for 2 degree years; {tail} (#1243)")


@pytest.mark.parametrize("code, fields", [
    # one degree year
    ("B1", {"degree": "Ph.D.", "institution": "Northfield University", "year": "1999"}),
    # two years in `year` outside a degree record
    ("H", {"award_name": "Example Teaching Award", "granting_body": "Lakeside Clinic",
           "year": "2001; 2003"}),
])
def test_multi_record_degree_years_need_two_years_in_a_degree_record(code, fields):
    text = "Example Teaching Award, Lakeside Clinic, 2001, 2003" if code == "H" else (
        "Doctor of Philosophy, Northfield University, 1999")
    assert _multi_record(text, fields, code=code, lines=()) == []


_TWO_LICENCES = "Northland, Pharmacist 12-3-45678\nSouthland, Pharmacist 098765"


@pytest.mark.parametrize("fields, lines, severity, lost", [
    ({"state_country": "Northland", "license_number": "12-3-45678"},
     ("Northland | 12-3-45678",), "WARN", 1),
    ({"state_country": "Northland", "license_number": "12-3-45678"},
     ("Northland | 12-3-45678", "Southland | 098765"), "INFO", 0),
    ({"state_country": "Northland; Southland", "license_number": "12-3-45678",
      "notes": "also 098765"}, (), "INFO", 0),
    # a number split over two lines is on neither
    ({"state_country": "Northland", "license_number": "12-3-45678"},
     ("Northland | 12-3-45678", "Suite 09", "8765 Example Road"), "WARN", 1),
    # the same, as rows of one table block
    ({"state_country": "Northland", "license_number": "12-3-45678"},
     ("Northland | 12-3-45678\nSuite 09\n8765 Example Road",), "WARN", 1),
])
def test_multi_record_two_licence_numbers_in_one_record(fields, lines, severity, lost):
    findings = _multi_record(_TWO_LICENCES, fields, code="F1", lines=lines)
    assert [f["severity"] for f in findings] == [severity]
    assert f"for 2 numbers; {lost} in neither the record nor the output" in findings[0]["message"]


@pytest.mark.parametrize("lines, severity, lost", [
    (("Example Valve | 1234567",), "WARN", 1),
    (("Example Valve | 1234567", "Example Valve, second filing | 7654321"), "INFO", 0),
])
def test_multi_record_two_patent_numbers_in_one_record(lines, severity, lost):
    findings = _multi_record("Example Valve 1234567; Example Valve 7654321", {
        "title": "Example Valve", "patent_number": "1234567"}, code="M2D", lines=lines)
    assert [f["severity"] for f in findings] == [severity]
    assert f"for 2 numbers; {lost} in neither the record nor the output" in findings[0]["message"]


@pytest.mark.parametrize("text, fields, code", [
    # one number
    ("Northland, Pharmacist 12-3-45678",
     {"state_country": "Northland", "license_number": "12-3-45678"}, "F1"),
    # one number, filed nowhere: a field loss, not a second record
    ("Northland, Pharmacist 12-3-45678", {"state_country": "Northland"}, "F1"),
    # one number beside years: a year is not an identifier
    ("Northland, Pharmacist 12-3-45678, 2004-2010",
     {"state_country": "Northland", "license_number": "12-3-45678",
      "issue_date": "2004", "expiration_date": "2010"}, "F1"),
    # the number field holds both
    (_TWO_LICENCES, {"state_country": "Northland; Southland",
                     "license_number": "12-3-45678; 098765"}, "F1"),
    # two long numbers on a code with no number field (a grant and its total)
    ("Example Outcomes Study, R01 CA123456, Example Agency, total costs 250000",
     {"grant_number": "R01 CA123456", "title": "Example Outcomes Study",
      "agency": "Example Agency", "total_costs": "250000"}, "M2B"),
])
def test_multi_record_licence_numbers_silent_when_not_several_records(text, fields, code):
    assert _multi_record(text, fields, code=code, lines=()) == []


def test_multi_record_constants_are_pinned():
    assert (MULTI_RECORD_MIN_NEW, MULTI_RECORD_NEW_SHARE,
            MULTI_RECORD_PROSE_MIN_LOWERCASE) == (2, 0.4, 4)

# ==========================================================================
# lint_year_not_in_source: a plausible year the entry's text never states.

def _not_in_source(*entries, stage5d=None):
    return lint_year_not_in_source({"entries": list(entries)}, stage5d)


@pytest.mark.parametrize("code, fields, text, flagged", [
    # EBYSBC HFAJCC 392's shape: an mm/dd/yy start read as another year.
    ("M2A", {"start_date": "1997-04-01", "end_date": "2024-03-31"},
     "Example Agency (PI) 04/01/19 \u2013 03/31/24 Example aims", "start_date=1997"),
    # EBYSBC VVRTUC 43's: the year of the entry before it in the batch.
    ("D1", {"start_date": "1981-03", "end_date": "current"},
     "Example Professor (current) March 2012 -", "start_date=1981"),
    # EBYSBC YYVHNN 480's: "3-24-11-26" is Mar 2024 to Nov 2026; no 2011.
    # The start year is the target. No year is read out of a dash-joined
    # run of numbers, so the right end year, written "26", is listed too.
    ("M2B", {"start_date": "2011-03-24", "end_date": "2026"},
     "Co-I Example contract 10% 3-24-11-26 $1.000", "start_date=2011, end_date=2026"),
])
def test_year_not_in_source_fires_on_a_year_the_text_never_states(code, fields, text, flagged):
    findings = _not_in_source(_fields_entry(code, fields, text=text, idx=43))
    assert len(findings) == 1
    assert findings[0]["lint"] == "year_not_in_source"
    assert findings[0]["severity"] == "WARN"
    assert findings[0]["message"].startswith(f"entry 43 ({code}): {flagged} -- ")
    assert findings[0]["evidence"] == [text[:FIELD_EVIDENCE_VALUE_CHARS]]


@pytest.mark.parametrize("fields, text", [
    ({"start_date": "2005"}, "Example course, 2005"),                 # four digits
    ({"start_date": "1997-04"}, "Example review 04/081997"),          # fused on the left
    ({"end_date": "2024"}, "Example 20232024Total"),                  # fused table cells
    ({"start_date": "1987"}, "Example board l987"),                   # scanned l for 1
    ({"end_date": "2034-05-31"}, "Example grant 6/1/29\u20135/31/34"),  # any century
    ({"start_date": "2006-03-14"}, "Example lecture 3/14/06"),
    ({"start_date": "2008"}, "Example award 0415/08"),
    ({"start_date": "2012-06"}, "Example talk 6//12"),
    ({"start_date": "2004"}, "Example talk, June '04"),
    ({"start_date": "1997", "end_date": "1997"}, "Example award FY97 $1000"),
    ({"date": "2009-05-21"}, "Example interview 5.21.09"),
    ({"start_date": "1964", "end_date": "1966"}, "Example mentee 64-66"),
    ({"start_date": "2011", "end_date": "2012"}, "Example committee 2011-2"),
    ({"start_date": "1985", "end_date": "1989"}, "Example committee 1985-89"),
    ({"start_date": "1993", "end_date": "1994"}, "Example committee 1993-994"),
    ({"start_date": "1998", "end_date": "2002"}, "Example committee 1998-02"),
    ({"start_date": "1986", "end_date": "1988"}, "1986,88 Example reviewer"),
    ({"start_date": "2006"}, "EXM 1234 Example seminar (Fall Sem., 06)"),
    ({"title": "1973"}, "Example title"),                             # not a date key
])
def test_year_not_in_source_silent_when_the_text_states_the_year(fields, text):
    assert _not_in_source(_fields_entry("C", fields, text=text)) == []


def test_year_not_in_source_skips_personal_data_and_publications():
    """Code A is personal data. A citation's year renders through 5d and
    PubMed; on the farm each publication hit was a split citation or a
    corrected typo, never a wrong year."""
    for code in ("A", "S1", "S8"):
        assert _not_in_source(_fields_entry(code, {"year": "2007"}, text="Example")) == []


def test_year_not_in_source_leaves_the_floor_and_ceiling_to_implausible_year():
    """One defect, one lint: a year below the owner's floor or above the
    ceiling is implausible_year's, so it is not reported twice."""
    below = _fields_entry("R", {"date": "1902-11"}, text="Example talk 11/02")
    above = _fields_entry("R", {"date": "2201"}, text="Example talk")
    assert _not_in_source(below, above) == []
    assert len(_implausible(below, above)) == 2


def test_year_not_in_source_skips_a_year_a_formatter_replaced():
    entry = _fields_entry("K1", {"start_date": "2016"}, text="Example course 2017-2019", idx=8)
    entry["element_idx_end"] = 8
    stage5d = {"entries": [{"element_idx_start": 8, "element_idx_end": 8,
                            "extracted_fields": {"formatted_text": "2017-2019 Example course"}}]}
    assert len(_not_in_source(entry)) == 1
    assert _not_in_source(entry, stage5d=stage5d) == []


def test_year_not_in_source_leaves_the_value_alone():
    entry = _fields_entry("D1", {"start_date": "1981-03"}, text="March 2012 -")
    before = json.dumps(entry, sort_keys=True)
    _not_in_source(entry)
    assert json.dumps(entry, sort_keys=True) == before


# ==========================================================================
# grant_boundary (#1226, EBYSBC E5) -- a grant list cut one line or row off.
# Invented grants; no corpus text.

_GRANT_HEADING = ["Research Support", "Past"]


def _grant(idx, text, heading=_GRANT_HEADING, code="M2B", **fields):
    return {"element_idx_start": idx, "taxonomy_code": code, "hierarchy": list(heading),
            "text": text, "extracted_fields": fields}


def _boundary(*entries):
    return lint_grant_boundary({"entries": list(entries)})


def _flagged(findings):
    return [int(f["message"].split()[1]) for f in findings]


_WHOLE_GRANT = dict(agency="Example Fund", grant_number="R01 XX000001",
                    title="Study of example things", start_date="2001", end_date="2004")


def test_grant_boundary_is_silent_on_whole_records():
    assert _boundary(
        _grant(10, "Agency: Example Fund\tGrant Title: Study one\tPI: A. Person", **_WHOLE_GRANT),
        _grant(12, "Agency: Other Fund\tGrant Title: Study two\tPI: A. Person", **_WHOLE_GRANT),
        _grant(14, "Agency: Third Fund\tGrant Title: Study three\tPI: A. Person", **_WHOLE_GRANT),
    ) == []


def test_grant_boundary_flags_a_record_split_between_two_neighbours():
    findings = _boundary(
        _grant(20, "Agency: Example Fund\tI.D.# Fellowship", agency="Example Fund",
               grant_number="Fellowship"),
        _grant(22, "Title: Study of example things\tP.I.: A. Person\t2001-2004",
               title="Study of example things", pi_name="A. Person", start_date="2001"))
    assert [(f["lint"], f["severity"]) for f in findings] == [("grant_boundary", "WARN")]
    assert findings[0]["message"].startswith("entry 20 (M2B):")
    assert "entry 22" in findings[0]["message"]


def test_grant_boundary_reads_the_split_in_either_order():
    findings = _boundary(
        _grant(30, "Study of example things", title="Study of example things"),
        _grant(31, "Example Fund (R01 XX000001)", agency="Example Fund",
               grant_number="R01 XX000001"))
    assert _flagged(findings) == [30]


def test_grant_boundary_reads_a_number_alone_as_the_source_half():
    findings = _boundary(_grant(32, "R01 XX000001", grant_number="R01 XX000001"),
                         _grant(33, "Study of example things", title="Study of example things"))
    assert _flagged(findings) == [32]


def test_grant_boundary_reads_an_empty_title_as_no_title():
    findings = _boundary(_grant(34, "Example Fund", agency="Example Fund", title=""),
                         _grant(35, "Study of example things", title="Study of example things"))
    assert _flagged(findings) == [34]


def test_grant_boundary_does_not_pair_a_number_with_a_record_naming_its_sponsor():
    """A titled entry that names its own sponsor is a whole record, not the
    title half of the number-only entry before it."""
    assert _boundary(_grant(36, "R01 XX000001", grant_number="R01 XX000001"),
                     _grant(37, "Example Fund\tStudy of example things", agency="Example Fund",
                            title="Study of example things")) == []


def test_grant_boundary_does_not_pair_a_sponsor_line_that_carries_its_title():
    assert _boundary(_grant(38, "Example Fund\tStudy one", agency="Example Fund",
                            title="Study one"),
                     _grant(39, "Study two", title="Study two")) == []


def test_grant_boundary_pairs_a_title_half_that_carries_its_number():
    """The title half is judged by its missing sponsor alone: a title with the
    grant number but no agency still completes the sponsor line before it."""
    findings = _boundary(_grant(40, "Example Fund", agency="Example Fund"),
                         _grant(41, "Study of example things (R01 XX000001)",
                                title="Study of example things", grant_number="R01 XX000001"))
    assert _flagged(findings) == [40]


def test_grant_boundary_reads_an_end_date_alone_as_a_period():
    """A sponsor line with only an end date carries a period, so it is a
    whole record's head, not a source half."""
    assert _boundary(_grant(42, "Example Fund, ending 2004", agency="Example Fund",
                            end_date="2004"),
                     _grant(43, "Study of example things", title="Study of example things")) == []


def test_grant_boundary_does_not_pair_a_sponsor_line_that_carries_its_period():
    """A one-line record with a sponsor and dates but no title (a fellowship)
    next to a titled one with no sponsor is two whole records."""
    assert _boundary(
        _grant(40, "Example Training Program (1978-84)", title="Example Training Program",
               start_date="1978", end_date="1984"),
        _grant(41, "Graduate Fellow, Example Foundation (1976-78)", agency="Example Foundation",
               start_date="1976", end_date="1978")) == []


@pytest.mark.parametrize("opening", ["PI: A. Person", "P.I.: A. Person", "PI Name: A. Person",
                                     "Personnel: A. Person", "Percent Effort: 5%",
                                     "% Effort: 5%", "Effort: 5%"])
def test_grant_boundary_flags_an_entry_opening_with_a_personnel_line(opening):
    findings = _boundary(_grant(50, f"{opening}\t2001-2004\tExample Fund\tStudy",
                                **_WHOLE_GRANT))
    assert _flagged(findings) == [50]
    assert "closes a record" in findings[0]["message"]


def test_grant_boundary_reads_a_personnel_line_only_at_the_start():
    assert _boundary(_grant(55, "Example Fund\tStudy\tPI: A. Person\t2001-2004",
                            **_WHOLE_GRANT)) == []


def _labelled(idx, opening="Agency", **fields):
    fields = fields or _WHOLE_GRANT
    order = {"Agency": "Agency: Example Fund\tGrant Title: Study\tAmount: $10",
             "Grant Title": "Grant Title: Study\tAmount: $10\tAgency: Next Fund",
             "Industry grant": "Industry grant: Example Co\tGrant Title: Study"}
    return _grant(idx, order[opening], **fields)


def test_grant_boundary_flags_an_entry_opening_with_a_label_its_siblings_carry_inside():
    findings = _boundary(_labelled(60), _labelled(62), _labelled(64, "Grant Title"),
                         _labelled(66, "Grant Title"))
    assert _flagged(findings) == [64, 66]
    assert "'agency:'" in findings[0]["message"] and "'grant title:'" in findings[0]["message"]


def test_grant_boundary_spares_a_record_opening_with_a_label_of_its_own():
    """A different kind of record (an industry grant) opens with a label the
    list's records never carry inside: its own record head, not a drift."""
    assert _boundary(_labelled(70), _labelled(72), _labelled(74, "Industry grant")) == []


def test_grant_boundary_needs_the_first_label_to_be_the_lists_shape():
    """One entry opening with the first label is not a record shape."""
    assert _boundary(_labelled(80), _labelled(82, "Grant Title"), _labelled(84, "Grant Title"),
                     _labelled(86, "Grant Title")) == []


def test_grant_boundary_reads_a_label_shape_only_from_a_labelled_first_entry():
    """Unlabelled records give no opening label to drift from: as measured
    on the farm, only a labelled list is judged by its labels."""
    plain = "Example Fund\tStudy of things\tRole: Lead"
    assert _boundary(_grant(94, plain, **_WHOLE_GRANT), _grant(95, plain, **_WHOLE_GRANT),
                     _grant(96, plain, **_WHOLE_GRANT),
                     _grant(97, "Role: Lead\tExample Fund\tStudy", **_WHOLE_GRANT)) == []


def test_grant_boundary_reads_a_label_shape_held_by_a_third_of_the_list():
    """Two of six entries open with the first label: a drifted list keeps
    under half its entries on it, so a third is the bar."""
    findings = _boundary(_labelled(200), _labelled(202),
                         *(_labelled(idx, "Grant Title") for idx in (204, 206, 208, 210)))
    assert _flagged(findings) == [204, 206, 208, 210]


def test_grant_boundary_needs_a_third_of_the_list_on_the_first_label():
    """Two of seven is under a third: no record shape."""
    assert _boundary(_labelled(220), _labelled(222),
                     *(_labelled(idx, "Grant Title") for idx in range(224, 234, 2))) == []


def test_grant_boundary_needs_three_entries_for_a_label_shape():
    assert _boundary(_labelled(90), _labelled(92, "Grant Title")) == []


def test_grant_boundary_flags_a_stray_description_after_the_first_entry():
    words = " ".join(["word"] * GRANT_ORPHAN_MIN_WORDS)
    findings = _boundary(_grant(100, "Example Fund\tStudy\t2001-2004", **_WHOLE_GRANT),
                         _grant(102, words, pi_role="PI"),
                         _grant(104, "Other Fund\tStudy\t2005-2008", **_WHOLE_GRANT))
    assert _flagged(findings) == [102]
    assert "tail" in findings[0]["message"]


def test_grant_boundary_reads_twelve_words_as_a_description():
    """The word-count arm alone, with no detail field: twelve words are a
    stray description, eleven a sub-heading."""
    def run(count):
        return _boundary(_grant(240, "Example Fund\tStudy", **_WHOLE_GRANT),
                         _grant(242, " ".join(["word"] * count)),
                         _grant(244, "Other Fund\tStudy", **_WHOLE_GRANT))
    assert _flagged(run(12)) == [242]
    assert run(11) == []


def test_grant_boundary_needs_two_detail_fields_for_a_stray_tail():
    assert _boundary(_grant(250, "Example Fund\tStudy", **_WHOLE_GRANT),
                     _grant(252, "2002", start_date="2002"),
                     _grant(254, "Other Fund\tStudy", **_WHOLE_GRANT)) == []


def test_grant_boundary_flags_a_stray_tail_of_detail_fields():
    findings = _boundary(_grant(110, "Example Fund\tStudy", **_WHOLE_GRANT),
                         _grant(112, "A. Person\t2001-2004", pi_name="A. Person",
                                start_date="2001"),
                         _grant(114, "Other Fund\tStudy", **_WHOLE_GRANT))
    assert _flagged(findings) == [112]


def test_grant_boundary_spares_a_short_subheading_and_the_first_entry():
    """A year or sub-heading line inside a list, or a list that opens with
    one, is not another record's tail."""
    words = " ".join(["word"] * GRANT_ORPHAN_MIN_WORDS)
    assert _boundary(_grant(120, words, pi_role="PI"),
                     _grant(122, "Example Fund\tStudy", **_WHOLE_GRANT),
                     _grant(124, "2002"),
                     _grant(126, "Funded Training Grants", status="Funded"),
                     _grant(128, "Other Fund\tStudy", **_WHOLE_GRANT)) == []


def test_grant_boundary_flags_a_lone_title_closing_a_list():
    findings = _boundary(_grant(130, "Example Fund\tStudy", **_WHOLE_GRANT),
                         _grant(132, "Other Fund\tStudy", **_WHOLE_GRANT),
                         _grant(134, "A study title alone", title="A study title alone"))
    assert _flagged(findings) == [134]
    assert "only a title" in findings[0]["message"]


def test_grant_boundary_spares_a_closing_title_when_no_sibling_holds_more():
    """A list whose every entry extracted only a title is a list of titles,
    not one record that lost its title."""
    assert _boundary(*(_grant(idx, f"Study {idx}", title=f"Study {idx}")
                       for idx in (135, 136, 137))) == []


def test_grant_boundary_needs_three_entries_for_a_closing_title():
    assert _boundary(_grant(138, "Example Fund\tStudy", **_WHOLE_GRANT),
                     _grant(139, "A study title alone", title="A study title alone")) == []


def test_grant_boundary_spares_a_lone_title_inside_a_list():
    assert _boundary(_grant(140, "Example Fund\tStudy", **_WHOLE_GRANT),
                     _grant(142, "Proposals at an example school",
                            title="Proposals at an example school"),
                     _grant(144, "Other Fund\tStudy", **_WHOLE_GRANT)) == []


def test_grant_boundary_judges_each_list_on_its_own():
    """A heading change or a non-grant entry closes a list: halves on either
    side of it are not paired."""
    source = dict(agency="Example Fund", grant_number="R01 XX000001")
    title = dict(title="Study of example things")
    assert _boundary(_grant(150, "Example Fund R01", **source),
                     _grant(152, "Study of example things", ["Research Support", "Current"],
                            code="M2A", **title)) == []
    assert _boundary(_grant(160, "Example Fund R01", **source),
                     {"element_idx_start": 161, "taxonomy_code": "S1", "text": "A paper",
                      "extracted_fields": {}},
                     _grant(162, "Study of example things", **title)) == []


def test_grant_boundary_reports_an_entry_once_with_every_shape():
    findings = _boundary(_labelled(170), _labelled(172),
                         _grant(174, "PI Name: A. Person\tAmount: $10", pi_name="A. Person",
                                total_funding="$10"))
    assert _flagged(findings) == [174]
    assert "closes a record" in findings[0]["message"] and "tail" in findings[0]["message"]


def test_grant_boundary_ignores_non_grant_codes():
    assert _boundary(_grant(180, "PI: A. Person\tStudy", code="D1")) == []


# ==========================================================================
# grant_bucket (#1343, EBYSBC E7) -- a grant filed in a subsection its own
# record contradicts.

_BUCKET_TEXT = "Example Research Foundation Study of Example Longevity Measures Program"


def _bucket(heading, rendered_under, year=2026, code="M2B", **fields):
    entry = _grant(200, _BUCKET_TEXT, heading, code=code, **fields)
    return lint_grant_bucket({"entries": [entry]}, _blocks_under(rendered_under, _BUCKET_TEXT),
                             current_year=year)


@pytest.mark.parametrize("heading", [["Grants Applied"], ["Grant Applications"],
                                     ["Research", "Proposals Submitted"], ["Not Funded"],
                                     ["Submitted But Not Funded Research Grants"],
                                     ["Non-Funded Proposals"], ["Non Funded Proposals"],
                                     ["Grant Applications", "Federal"]])
@pytest.mark.parametrize("rendered_under", ["M2A", "M2B"])
def test_grant_bucket_flags_an_application_rendered_as_an_award(heading, rendered_under):
    findings = _bucket(heading, rendered_under)
    assert [(f["lint"], f["severity"]) for f in findings] == [("grant_bucket", "WARN")]
    assert findings[0]["message"].startswith("entry 200 (M2B):")
    assert heading[-1] in findings[0]["message"]


def test_grant_bucket_spares_an_application_rendered_under_pending():
    assert _bucket(["Grants Applied"], "M2C") == []


@pytest.mark.parametrize("heading", [["Current and Pending Support"], ["Funded Grants"],
                                     ["Research Support"]])
def test_grant_bucket_spares_a_heading_that_also_files_awards_or_is_silent(heading):
    assert _bucket(heading, "M2B") == []


def test_grant_bucket_spares_an_application_also_rendered_under_pending():
    entry = _grant(205, _BUCKET_TEXT, ["Grants Applied"])
    blocks = _blocks_under("M2B", _BUCKET_TEXT) + _blocks_under("M2C", _BUCKET_TEXT)
    assert lint_grant_bucket({"entries": [entry]}, blocks, current_year=2026) == []


def test_grant_bucket_spares_a_grant_it_cannot_locate():
    entry = _grant(210, _BUCKET_TEXT, ["Grants Applied"])
    assert lint_grant_bucket({"entries": [entry]}, [], current_year=2026) == []


@pytest.mark.parametrize("start, end, ended", [
    ("9/1/07", "8/30/1", 2011),     # a one-digit year past the start year's own
    ("7/1/08", "6/30/9", 2009),
    ("2015-01", "12/31/25", 2025),  # stage 6's own reading, still Current
    ("2010", "June 2019", 2019),
])
def test_grant_bucket_flags_a_current_grant_whose_end_date_reads_past(start, end, ended):
    findings = _bucket(["Research Support", "Ongoing"], "M2A", code="M2A",
                       start_date=start, end_date=end)
    assert len(findings) == 1
    assert f"reads as {ended}, before 2026" in findings[0]["message"]


@pytest.mark.parametrize("start, end", [
    ("", "8/30/1"),             # no start year to read a truncated year against
    ("9/1/07", "Present"),
    ("9/1/07", ""),
    ("2020", "6/30/28"),        # ends in the future
    ("9/1/24", "8/30/6"),       # truncated, but 2026 is not before this year
])
def test_grant_bucket_spares_a_current_grant_not_known_to_have_ended(start, end):
    assert _bucket(["Research Support", "Ongoing"], "M2A", code="M2A",
                   start_date=start, end_date=end) == []


def test_grant_bucket_spares_an_ended_grant_rendered_under_past():
    assert _bucket(["Research Support", "Ongoing"], "M2B", code="M2A",
                   start_date="9/1/07", end_date="8/30/1") == []


@pytest.mark.parametrize("rendered_under", [("M2C",), ("M2A", "M2B")])
def test_grant_bucket_judges_an_ended_grant_only_when_current_alone(rendered_under):
    """An ended grant under Pending is bucket_status's question, and one
    also rendered under Past is filed there too."""
    entry = _grant(220, _BUCKET_TEXT, ["Ongoing"], code="M2A",
                   start_date="9/1/07", end_date="8/30/1")
    blocks = [block for code in rendered_under for block in _blocks_under(code, _BUCKET_TEXT)]
    assert lint_grant_bucket({"entries": [entry]}, blocks, current_year=2026) == []


def test_grant_bucket_judges_the_end_date_against_the_year_it_is_given():
    fields = dict(start_date="2015-01", end_date="12/31/25")
    assert len(_bucket(["Ongoing"], "M2A", year=2026, code="M2A", **fields)) == 1
    assert _bucket(["Ongoing"], "M2A", year=2025, code="M2A", **fields) == []


def test_grant_bucket_defaults_to_this_year():
    entry = _grant(230, _BUCKET_TEXT, ["Ongoing"], code="M2A", start_date="1990",
                   end_date="2001")
    findings = lint_grant_bucket({"entries": [entry]}, _blocks_under("M2A", _BUCKET_TEXT))
    assert len(findings) == 1 and "reads as 2001" in findings[0]["message"]


def test_grant_bucket_default_year_is_this_year_not_next():
    """A Current grant ending this year is not yet ended; one that ended last
    year is. Both are read against the real current year."""
    this_year = datetime.now().year
    current = _grant(231, _BUCKET_TEXT, ["Ongoing"], code="M2A", start_date="2010",
                     end_date=str(this_year))
    assert lint_grant_bucket({"entries": [current]}, _blocks_under("M2A", _BUCKET_TEXT)) == []
    ended = _grant(232, _BUCKET_TEXT, ["Ongoing"], code="M2A", start_date="2010",
                   end_date=str(this_year - 1))
    findings = lint_grant_bucket({"entries": [ended]}, _blocks_under("M2A", _BUCKET_TEXT))
    assert len(findings) == 1 and f"reads as {this_year - 1}" in findings[0]["message"]


def test_grant_bucket_ignores_non_grant_codes():
    assert _bucket(["Grants Applied"], "M2B", code="D1") == []


@pytest.mark.parametrize("end, start, expected", [
    ("8/30/1", "9/1/07", 2011), ("5/30/13", "5/1/08", 2013), ("6/30/0", "2010", 2010),
    ("2/28/05", "1999", 2005), ("8/30", "2001", 2030), ("", "2001", None),
    ("8/30/1", "", None), ("8-30-1", "2007", 2011), ("8.30.1", "2007", 2011),
    ("6/30/13", "1995", 2013),  # two digits read modulo 100
])
def test_short_end_year(end, start, expected):
    assert _short_end_year(end, start, 2026) == expected


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
