"""Tests for #565: `generate()`'s per-section exception boundary.

Before this fix, `generate()` ran its 21 `_fill_*` calls with no exception
boundary; any one of them raising aborted the entire render -- no docx at
all, even when every other section would have rendered cleanly (#442 lost
two full documents this way). `_render_section` now isolates each section:
a raise costs that section only, is logged with its traceback, and is
recorded onto the render-warnings sidecar at ERROR severity so the failure
is loud rather than a silently incomplete document.

Self-contained: no DB, no network, no PII. Uses the bundled WCM template and
python-docx, synthetic entries only. The LLM-driven appendix-reconsider pass
is neutralized (same idiom as test_m1_appendix_fallback.py) so tests stay
deterministic and credential-free. Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_section_boundary.py -p no:cacheprovider
"""

import json
import logging
import sys
from pathlib import Path

import pytest
from docx import Document

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.doctor.lints.render import lint_stage6_warnings  # noqa: E402
from unified_pipeline.stage6.sections.appendix import RecoveredLine  # noqa: E402
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    RENDER_ROUTED_CODES,
    WCMTemplateGenerator,
)

# The 21-name flat dispatch this replaced (`stage_6_word_template.py:812-833`
# on the pre-fix tree), in the exact order `generate()` called them.
# test_dispatch_wire_preserves_call_order pins this list: dropping a section
# from the table, or reordering one, fails here even though every other
# test in this file would still pass.
PINNED_DISPATCH_ORDER = [
    "_fill_personal_data",
    "_fill_researcher_profiles",
    "_fill_education",
    "_fill_other_education",
    "_fill_postdoc_training",
    "_fill_positions",
    "_fill_licensure",
    "_fill_board_certification",
    "_fill_honors",
    "_fill_memberships",
    "_fill_teaching",
    "_fill_research_summary",
    "_fill_research_support",
    "_fill_patents",
    "_fill_mentoring",
    "_fill_clinical_practice",
    "_fill_leadership",
    "_fill_administrative_activities",
    "_fill_service",
    "_fill_presentations",
    "_fill_bibliography",
]

_APPENDIX_HEADER = "T. APPENDIX"


def _new_generator() -> WCMTemplateGenerator:
    gen = WCMTemplateGenerator(verbose=False)
    # Neutralize the LLM-driven appendix-reconsider pass: keeps every test
    # here deterministic and credential-free (test_m1_appendix_fallback.py
    # idiom).
    gen._reconsider_appendix_entries = lambda: None
    return gen


def _render(gen: WCMTemplateGenerator, tmp_path: Path, entries: list,
            document_uid: str = "TESTAA") -> Path:
    """Drive the real generate() path and return the output docx path."""
    data = {"document_uid": document_uid, "entries": entries}
    ip = tmp_path / "in.json"
    op = tmp_path / "out.docx"
    ip.write_text(json.dumps(data))
    gen.generate(str(ip), str(op), research_summary_path=None)
    return op


def _sidecar(tmp_path: Path, document_uid: str = "TESTAA") -> dict:
    return json.loads((tmp_path / f"{document_uid}_render_warnings.json").read_text())


def _full_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for tb in doc.tables for row in tb.rows for c in row.cells]
    return "\n".join(parts)


def _appendix_paragraphs(doc: Document) -> list:
    """Just the T. APPENDIX paragraphs -- a specific slot, not 'somewhere in
    the document' (mirrors test_stage6_passthrough_appendix_exclusion.py)."""
    paragraphs = [p.text for p in doc.paragraphs]
    for i, text in enumerate(paragraphs):
        if text.strip() == _APPENDIX_HEADER:
            return paragraphs[i:]
    return []


def _personal_data_entry() -> dict:
    return {"text": "Name: Jane Q. Public, MD", "taxonomy_code": "A",
            "extracted_fields": {}, "element_idx_start": 0}


# --------------------------------------------------------------- positive control

def test_nonfatal_section_failure_isolated_to_that_section(tmp_path, caplog):
    """The core fix: `_fill_honors` raising costs only Honors. The document
    is still written, a later section's content still reaches its slot, and
    the failure is both logged (with traceback) and recorded at ERROR."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("boom")

    gen._fill_honors = _boom

    entries = [
        _personal_data_entry(),
        {"text": "The DISTINCTIVE_LATER_TOKEN unmapped content",
         "taxonomy_code": "T", "extracted_fields": {}, "element_idx_start": 5},
    ]
    with caplog.at_level(logging.ERROR):
        op = _render(gen, tmp_path, entries)

    assert op.exists(), "one section raising must not cost the whole document"

    # A later section (the appendix, fed by the T-coded entry) still
    # rendered into its specific slot.
    doc = Document(str(op))
    assert "DISTINCTIVE_LATER_TOKEN" in "\n".join(_appendix_paragraphs(doc)), \
        "a section after the failed one did not reach its slot"

    sidecar = _sidecar(tmp_path)
    failures = [w for w in sidecar["warnings"] if w["check"] == "section_render_failed"]
    assert len(failures) == 1
    record = failures[0]
    assert record["section"] == "honors"
    assert record["severity"] == "ERROR"
    assert "boom" in record["message"]
    assert record["evidence"], "evidence must not be empty"
    assert all(len(line) <= 200 for line in record["evidence"])

    # Logged with its traceback (§5.4 -- never just str(exc)), at ERROR.
    error_records = [r for r in caplog.records if r.levelno == logging.ERROR]
    assert any("honors" in r.getMessage() for r in error_records)
    assert any(r.exc_info is not None for r in error_records), \
        "the failure must be logged with logger.exception (traceback attached)"


def test_log_validation_warnings_banner_and_message_at_warning(tmp_path, caplog):
    """`_log_validation_warnings` (#839's pure-move helper, round 3 NOTE):
    a section failure logs the VALIDATION WARNINGS banner and the failure
    message itself at WARNING, not just `_render_section`'s ERROR log."""
    gen = _new_generator()
    gen._fill_honors = lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("boom"))
    with caplog.at_level(logging.WARNING):
        _render(gen, tmp_path, [_personal_data_entry()])
    warn_msgs = [r.getMessage() for r in caplog.records if r.levelno == logging.WARNING]
    assert any("VALIDATION WARNINGS" in m for m in warn_msgs)
    assert any("boom" in m for m in warn_msgs)


# --------------------------------------------------------------- fatal contract

def test_fatal_personal_data_failure_still_raises_no_docx(tmp_path):
    """`_fill_personal_data` is the one deliberate exception to the boundary
    (#565's judgement call): a document with no owner on it is worse than a
    failed run, so its raise must still propagate and no docx is written."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("fatal-boom")

    gen._fill_personal_data = _boom

    with pytest.raises(RuntimeError, match="fatal-boom"):
        _render(gen, tmp_path, [_personal_data_entry()])

    assert not (tmp_path / "out.docx").exists()


# --------------------------------------------------------------- return-value fallbacks

def test_research_summary_failure_falls_back_to_false_m1_reaches_appendix(tmp_path):
    """`_fill_research_summary` raising must not crash the render, and its
    return-value fallback (False) must route M1 entries to the appendix --
    the same #317/C0ZGFW safety net as when there is no summary at all
    (test_m1_appendix_fallback.py), now also triggered by a section raise."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("research-summary-boom")

    gen._fill_research_summary = _boom

    entries = [
        _personal_data_entry(),
        {"text": "The M1_FALLBACK_TOKEN research program ran 2010 to 2014",
         "taxonomy_code": "M1", "extracted_fields": {"narrative": "ran 2010 to 2014"},
         "element_idx_start": 5},
    ]
    op = _render(gen, tmp_path, entries)
    assert op.exists()

    doc = Document(str(op))
    assert "M1_FALLBACK_TOKEN" in _full_text(doc), \
        "M1 entry was dropped when research_summary_rendered did not fall back to False"

    sidecar = _sidecar(tmp_path)
    failures = [w for w in sidecar["warnings"] if w["check"] == "section_render_failed"]
    assert len(failures) == 1
    assert failures[0]["section"] == "research_summary"


def test_passthrough_failure_entries_not_silently_dropped(tmp_path):
    """`_fill_passthrough_sections` raising must not crash the render, and
    its return-value fallback (empty set of consumed ids) must not silently
    drop passthrough-eligible entries -- they still reach the appendix as
    unconsumed content, not nowhere."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("passthrough-boom")

    gen._fill_passthrough_sections = _boom

    entries = [
        _personal_data_entry(),
        {"text": "PASSTHROUGH_FALLBACK_TOKEN Employment Status: Active",
         "taxonomy_code": "E", "extracted_fields": {}, "element_idx_start": 5},
    ]
    op = _render(gen, tmp_path, entries)
    assert op.exists(), "passthrough failing must not cost the whole document"

    doc = Document(str(op))
    assert "PASSTHROUGH_FALLBACK_TOKEN" in "\n".join(_appendix_paragraphs(doc)), \
        "an unconsumed passthrough-eligible entry must still reach the appendix"

    sidecar = _sidecar(tmp_path)
    failures = [w for w in sidecar["warnings"] if w["check"] == "section_render_failed"]
    assert len(failures) == 1
    assert failures[0]["section"] == "passthrough_sections"


# --------------------------------------------------------------- dispatch wire

def test_dispatch_wire_preserves_call_order(tmp_path):
    """Spy every _fill_* in the pinned 21-name list and assert the CALL
    ORDER matches dev exactly. This is the test that kills a section
    silently falling out of the table (m6), or the table being reordered."""
    gen = _new_generator()
    call_order = []

    def _recorder(name, return_value=None):
        def _f(*_a, **_k):
            call_order.append(name)
            return return_value
        return _f

    for name in PINNED_DISPATCH_ORDER:
        return_value = False if name == "_fill_research_summary" else None
        setattr(gen, name, _recorder(name, return_value))

    _render(gen, tmp_path, [])

    assert call_order == PINNED_DISPATCH_ORDER
    assert gen._section_failures == [], "no section here raised; nothing should be recorded"


# --------------------------------------------------------------- two failures

def test_two_section_failures_both_recorded_rest_still_renders(tmp_path):
    """Two independent sections raising must produce two independent
    records naming both labels, with the document still written."""
    gen = _new_generator()

    def _boom(label):
        def _f(*_a, **_k):
            raise RuntimeError(f"boom-{label}")
        return _f

    gen._fill_honors = _boom("honors")
    gen._fill_leadership = _boom("leadership")

    op = _render(gen, tmp_path, [_personal_data_entry()])
    assert op.exists()

    sidecar = _sidecar(tmp_path)
    failures = [w for w in sidecar["warnings"] if w["check"] == "section_render_failed"]
    assert len(failures) == 2
    assert {f["section"] for f in failures} == {"honors", "leadership"}
    assert all(f["severity"] == "ERROR" for f in failures)


# --------------------------------------------------------------- _write_render_warnings_sidecar unit tests

def test_write_render_warnings_sidecar_writes_expected_json(tmp_path):
    """Positive path: the sidecar is written with the exact merged shape
    generate() relies on."""
    gen = _new_generator()
    output_path = str(tmp_path / "out.docx")
    gen._write_render_warnings_sidecar(
        output_path, "TESTAA", [{"check": "x"}], [{"code": "A"}])
    sidecar = json.loads((tmp_path / "TESTAA_render_warnings.json").read_text())
    assert sidecar == {
        "document_uid": "TESTAA",
        "warnings": [{"check": "x"}],
        "dedup_decisions": [{"code": "A"}],
    }


def test_write_render_warnings_sidecar_is_fail_soft(tmp_path, caplog):
    """Negative path: a sidecar write failure (here, a nonexistent parent
    directory) must never raise -- the render already succeeded and must
    not be failed by a sidecar problem. It is logged at ERROR with the
    traceback attached (§5.4): a missing sidecar leaves the doctor unable
    to tell a clean run from a pre-sidecar build, so it must be loud."""
    gen = _new_generator()
    bad_output_path = str(tmp_path / "missing_dir" / "out.docx")
    with caplog.at_level(logging.ERROR):
        gen._write_render_warnings_sidecar(bad_output_path, "TESTAA", [], [])
    error_records = [r for r in caplog.records if r.levelno == logging.ERROR]
    assert any("could not write render-warnings sidecar" in r.getMessage()
               for r in error_records)
    assert any(r.exc_info is not None for r in error_records), \
        "the sidecar failure must be logged with logger.exception (traceback attached)"


# --------------------------------------------------------------- lint severity path

def test_lint_stage6_warnings_severity_path():
    report = {"warnings": [
        {"message": "x", "severity": "ERROR"},
        {"message": "y"},
    ]}
    findings = lint_stage6_warnings(report)
    assert [f["severity"] for f in findings] == ["ERROR", "WARN"]


def test_lint_stage6_warnings_unknown_severity_falls_back_to_warn():
    report = {"warnings": [{"message": "x", "severity": "CATASTROPHIC"}]}
    findings = lint_stage6_warnings(report)
    assert findings[0]["severity"] == "WARN"


def test_appendix_failure_with_recovered_codes_still_writes_the_document(tmp_path, caplog):
    """#531 on top of #565: when `_fill_appendix` raises, `_render_section`
    returns None, and the diversion-warning builder is later handed the list of
    written appendix entries together with the codes `_recover_unrendered_records`
    reported. Without the `or []` on that return the builder iterates None and
    the run dies AFTER the docx was saved -- an isolated appendix failure turned
    back into a failed run. Pins: document written, exactly one
    `section_render_failed` for the appendix, no appendix_diversion record
    derived from a None `written`."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("appendix boom")

    gen._fill_appendix = _boom
    gen._recover_unrendered_records = lambda *_a, **_k: [RecoveredLine("T", "Recovered sample line")]

    entries = [
        _personal_data_entry(),
        {"text": "Unmapped content that would go to the appendix",
         "taxonomy_code": "T", "extracted_fields": {}, "element_idx_start": 5},
    ]
    with caplog.at_level(logging.ERROR):
        op = _render(gen, tmp_path, entries)

    assert op.exists(), "an appendix raise must not cost the document"
    sidecar = _sidecar(tmp_path)
    failures = [w for w in sidecar["warnings"] if w["check"] == "section_render_failed"]
    assert [f["section"] for f in failures] == ["appendix"]
    diversions = [(w["code"], w["reason"]) for w in sidecar["warnings"]
                  if w["check"] == "appendix_diversion"]
    assert diversions == [("T", "recovered_unrendered")], \
        "only the recovered code is reported; nothing is derived from the failed appendix"


# --------------------------------------------------------------- #842: failed section -> Appendix

def test_failed_section_entries_fall_to_appendix_with_diversion_warning(tmp_path):
    """#842: on dev, a failed section's entries are absent from BOTH the
    section and the Appendix -- their codes are in RENDER_ROUTED_CODES, so
    the unmapped sweep never sees them, and no renderer ran to place them.
    Once `_render_section` discards the failed section's codes from
    mapped_codes, they fall to the Appendix as numbered lines, and
    appendix.py's existing `renderer_declined` reason (#838) covers the
    warning for free -- no new reason vocabulary."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("boom")

    gen._fill_honors = _boom

    entries = [
        _personal_data_entry(),
        {"text": "HONORS_TOKEN_842 Some Award", "taxonomy_code": "H",
         "extracted_fields": {}, "element_idx_start": 1},
    ]
    op = _render(gen, tmp_path, entries)
    assert op.exists()

    doc = Document(str(op))
    assert "HONORS_TOKEN_842" in "\n".join(_appendix_paragraphs(doc)), \
        "a failed section's entry must fall to the Appendix, not vanish"

    sidecar = _sidecar(tmp_path)
    section_failures = [w for w in sidecar["warnings"] if w["check"] == "section_render_failed"]
    assert [f["section"] for f in section_failures] == ["honors"]

    diversions = [w for w in sidecar["warnings"]
                  if w["check"] == "appendix_diversion" and w["code"] == "H"]
    assert len(diversions) == 1
    assert diversions[0]["reason"] == "renderer_declined"
    assert diversions[0]["count"] == 1
    assert diversions[0]["message"] == (
        "H: 1 entry diverted to the Appendix — not placed by the section "
        "routed for H (see any section_render_failed record for that section)"
    ), "the M1-specific 'no research summary rendered' text must not leak onto other codes (#842 r2)"


def test_failed_multi_code_section_all_codes_fall_to_appendix(tmp_path):
    """A section that owns several codes (service: Q1..Q4D) must divert
    ALL of them when it raises, not just one -- kills a mutant that only
    unions in a single code from the failed section's set."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("boom")

    gen._fill_service = _boom

    entries = [
        _personal_data_entry(),
        {"text": "Q1_TOKEN_842 leadership role", "taxonomy_code": "Q1",
         "extracted_fields": {}, "element_idx_start": 1},
        {"text": "Q4B_TOKEN_842 grant review", "taxonomy_code": "Q4B",
         "extracted_fields": {}, "element_idx_start": 2},
    ]
    op = _render(gen, tmp_path, entries)
    assert op.exists()

    appendix_text = "\n".join(_appendix_paragraphs(Document(str(op))))
    assert "Q1_TOKEN_842" in appendix_text
    assert "Q4B_TOKEN_842" in appendix_text

    sidecar = _sidecar(tmp_path)
    diversion_warnings = [w for w in sidecar["warnings"]
                           if w["check"] == "appendix_diversion" and w["reason"] == "renderer_declined"]
    diversions = {w["code"] for w in diversion_warnings}
    assert diversions == {"Q1", "Q4B"}
    messages = {w["code"]: w["message"] for w in diversion_warnings}
    assert messages["Q1"] == (
        "Q1: 1 entry diverted to the Appendix — not placed by the section "
        "routed for Q1 (see any section_render_failed record for that section)"
    )
    assert messages["Q4B"] == (
        "Q4B: 1 entry diverted to the Appendix — not placed by the section "
        "routed for Q4B (see any section_render_failed record for that section)"
    )


def test_failed_mentoring_section_diverts_n4_to_appendix(tmp_path):
    """#587: N4 is owned by the mentoring section's dispatch codes, so a
    raising `_fill_mentoring` drops N4 from `mapped_codes` like N1-N3B --
    kills a mutant that adds N4 to RENDER_ROUTED_CODES but not to the
    section's own code set (the entry would then vanish: unrendered AND
    excluded from the Appendix)."""
    gen = _new_generator()

    def _boom(*_a, **_k):
        raise RuntimeError("boom")

    gen._fill_mentoring = _boom
    entries = [
        _personal_data_entry(),
        {"text": "N4_TOKEN_587 outcome narrative", "taxonomy_code": "N4",
         "extracted_fields": {}, "element_idx_start": 1},
    ]
    op = _render(gen, tmp_path, entries)

    assert "N4_TOKEN_587" in "\n".join(_appendix_paragraphs(Document(str(op))))
    diversions = [w for w in _sidecar(tmp_path)["warnings"]
                  if w["check"] == "appendix_diversion" and w["code"] == "N4"]
    assert [(w["reason"], w["count"]) for w in diversions] == [("renderer_declined", 1)]


def test_no_failure_routed_entry_stays_out_of_appendix(tmp_path):
    """Mutant-killer: when no section fails, `_failed_section_codes` must
    stay empty and a routed H entry renders in its own slot, NOT the
    Appendix -- kills a mutant that discards a section's codes
    unconditionally instead of only on failure."""
    gen = _new_generator()

    entries = [
        _personal_data_entry(),
        {"text": "HONORS_TOKEN_842_CLEAN Some Award", "taxonomy_code": "H",
         "extracted_fields": {}, "element_idx_start": 1},
    ]
    op = _render(gen, tmp_path, entries)
    assert gen._failed_section_codes == set()

    appendix_text = "\n".join(_appendix_paragraphs(Document(str(op))))
    assert "HONORS_TOKEN_842_CLEAN" not in appendix_text, \
        "a section that did not fail must not divert its codes to the Appendix"


def test_failed_section_codes_reset_between_renders(tmp_path):
    """Per-call reset (#842): a failure in one `generate()` call must not
    leak into the next call's `mapped_codes` discard -- same generator
    instance, two successive renders."""
    gen = _new_generator()
    real_fill_honors = gen._fill_honors

    def _boom(*_a, **_k):
        raise RuntimeError("boom")

    gen._fill_honors = _boom
    first_entries = [
        _personal_data_entry(),
        {"text": "HONORS_TOKEN_842_FIRST Some Award", "taxonomy_code": "H",
         "extracted_fields": {}, "element_idx_start": 1},
    ]
    _render(gen, tmp_path, first_entries, document_uid="TESTAA")
    assert gen._failed_section_codes == {"H"}

    gen._fill_honors = real_fill_honors
    second_entries = [
        _personal_data_entry(),
        {"text": "HONORS_TOKEN_842_SECOND Some Award", "taxonomy_code": "H",
         "extracted_fields": {}, "element_idx_start": 1},
    ]
    _render(gen, tmp_path, second_entries, document_uid="TESTBB")
    assert gen._failed_section_codes == set(), \
        "a prior render's failed codes must not leak into this one"

    sidecar = _sidecar(tmp_path, document_uid="TESTBB")
    h_diversions = [w for w in sidecar["warnings"]
                    if w["check"] == "appendix_diversion" and w["code"] == "H"]
    assert h_diversions == [], "the second render's H entry must not be diverted"


def test_dispatch_codes_cover_render_routed_codes_exactly(tmp_path):
    """Completeness self-check (#842): the union of every dispatch entry's
    `codes` covers exactly `RENDER_ROUTED_CODES - {'A'}` (A is Personal
    Data, outside the `_render_section` boundary by deliberate judgement
    call), and the sets are pairwise disjoint. Catches a forgotten or
    duplicated code before it costs a real failure. Drives the real
    `generate()` with a spy on `_render_section` (the harness already spies
    `_fill_*` methods the same way) rather than hoisting the dispatch table
    to module level just for this test."""
    gen = _new_generator()
    recorded: list[tuple[str, frozenset]] = []
    real_render_section = gen._render_section

    def _spy(label, fn, codes=frozenset(), **kwargs):
        recorded.append((label, codes))
        return real_render_section(label, fn, codes, **kwargs)

    gen._render_section = _spy
    _render(gen, tmp_path, [_personal_data_entry()])

    code_sets = [codes for _label, codes in recorded if codes]
    union: frozenset[str] = frozenset().union(*code_sets)
    assert union == RENDER_ROUTED_CODES - {'A'}
    assert sum(len(c) for c in code_sets) == len(union), \
        "dispatch code sets must be pairwise disjoint"


def test_each_dispatched_section_logs_a_progress_bar_line(tmp_path, caplog, progress_patterns):
    """Stage 6 printed nothing orchestrator.PROGRESS_PATTERNS could read, so
    the web bar sat at its placeholder through the render. Every dispatched
    section -- failed or not, via _render_section's finally -- logs
    ``[done/total] sections rendered``, read the way orchestrator.py reads it."""
    gen = _new_generator()
    with caplog.at_level(logging.INFO, logger="unified_pipeline.stage_6_word_template"):
        _render(gen, tmp_path, [_personal_data_entry()])

    seen = []
    for r in caplog.records:
        match = next((m for p in progress_patterns if (m := p.search(r.getMessage()))), None)
        if match:
            seen.append((int(match.group(1)), int(match.group(2))))
    total = seen[0][1]
    assert total > 1
    assert seen == [(n, total) for n in range(1, total + 1)]
