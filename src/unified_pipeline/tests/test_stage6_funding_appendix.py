"""Regression guards for the stage-6 funding/appendix fixes (#209, #210, #213).

Grounded in run 89HQVQ: a source CV packed 8 grants into one table cell; only
the first was field-extracted, the other 7 were dumped in the Appendix as
"• [M2A] ..." bullets, and every grant (including "Under review" and "Not
Funded" ones) was bucketed as Current Research Funding.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_funding_appendix.py -p no:cacheprovider

Self-contained: no DB, no LLM calls. The appendix test loads the bundled WCM
template like test_stage6_strip_instruction_box.py does.
"""

import json
import sys
from pathlib import Path

# Ensure the repo's ``src`` directory is importable regardless of cwd/rootdir.
_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from unified_pipeline.core.template_boilerplate import (
    is_source_boilerplate,  # noqa: E402
)
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    WCMTemplateGenerator,
    grant_status_rebucket_target,
    parse_reclassified_segments,
    segment_already_rendered,
)


def _num_level(para):
    """The paragraph's w:ilvl, or None when it carries no list markup.

    Mirrors test_stage6_k_list_bullets.py's helper of the same name -- the
    appendix segments converted to real Word list paragraphs by #483.
    """
    pPr = para._p.pPr
    numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
    if numPr is None:
        return None
    return numPr.find(qn("w:ilvl")).get(qn("w:val"))


# ---------------------------------------------------------------- #210

def test_status_awarded_stays_put():
    target, note = grant_status_rebucket_target("Awarded 2026")
    assert target is None and note is None


def test_status_under_review_goes_pending():
    target, note = grant_status_rebucket_target("Submitted 2026, Under review")
    assert target == "M2C"
    assert "Pending" in note


def test_status_not_funded_goes_pending_with_review_note():
    # Ordering matters: this status also contains "Submitted".
    target, note = grant_status_rebucket_target("Submitted 2025-2026, Not Funded")
    assert target == "M2C"
    assert "confirm" in note.lower()


def test_status_completed_goes_past():
    target, _ = grant_status_rebucket_target("Completed 2023")
    assert target == "M2B"


def test_status_empty_or_none_is_noop():
    assert grant_status_rebucket_target("") == (None, None)
    assert grant_status_rebucket_target(None) == (None, None)


def test_status_awarded_pending_contract_not_moved():
    # "pending" appears, but the grant is awarded — must not move.
    target, _ = grant_status_rebucket_target("Awarded, pending contract")
    assert target is None


# ---------------------------------------------------------------- #575

def test_status_in_review_goes_pending():
    for status in ("In review", "In Review"):
        target, note = grant_status_rebucket_target(status)
        assert target == "M2C", status
        assert "Pending" in note


def test_status_awaiting_or_under_consideration_goes_pending():
    for status in (
        "Awaiting sponsor decision",
        "Awaiting decision",
        "Under consideration",
    ):
        target, note = grant_status_rebucket_target(status)
        assert target == "M2C", status
        assert "Pending" in note


def test_status_awaiting_non_decision_not_moved():
    # "Awaiting" a post-award step (no 'award' substring) is not a pending
    # application — only an awaited *decision* rebuckets.
    for status in ("Awaiting contract execution", "Awaiting IRB approval"):
        target, _ = grant_status_rebucket_target(status)
        assert target is None, status


def test_status_awaiting_award_setup_not_moved():
    # "Awaiting" appears, but the grant is awarded — the 'award' guard holds.
    target, _ = grant_status_rebucket_target("Awaiting award setup")
    assert target is None


# ---------------------------------------------------------------- #209

_FSMB_FIELDS = {
    "title": "Improving Access to Healthcare in Rural and Underserved Area",
    "agency": "Federation of State Medical Boards (FSMB) Foundation",
    "status": "Awarded 2026",
    "pi_name": "Shapiro, M.",
}

_FSMB_SEGMENT = (
    "Federation of State Medical Boards (FSMB) Foundation Grant | Shapiro, M. "
    "(PI), Jung, E. (Co-PI) | Improving Access to Healthcare in Rural and "
    "Underserved Areas | Role: Co-PI | Amount: $75,000 | Status: Awarded 2026."
)

_TEMPLETON_SEGMENT = (
    "John Templeton Foundation Online Funding Inquiry (OFI) | Jung, E. (PI) | "
    "Developing Intellectual Humility in Clinical Judgment | Role: PI | "
    "Amount: $1,400,000 | Status: Submitted 2026, Under review."
)


def test_rendered_segment_detected():
    assert segment_already_rendered(_FSMB_SEGMENT, _FSMB_FIELDS)


def test_unrendered_sibling_not_flagged():
    assert not segment_already_rendered(_TEMPLETON_SEGMENT, _FSMB_FIELDS)


def test_generic_fields_never_match():
    # A shared status string must not mark sibling records as rendered.
    fields = {"status": "Submitted 2026, Under review"}
    assert not segment_already_rendered(_TEMPLETON_SEGMENT, fields)


def test_empty_fields_no_match():
    assert not segment_already_rendered(_FSMB_SEGMENT, {})
    assert not segment_already_rendered(_FSMB_SEGMENT, None)


# ---------------------------------------------------------------- #213

def test_source_boilerplate_positives():
    for text in (
        "CURRICULUM VITAE",
        "Curriculum Vitae",
        "Last Updated - JUN 2026",
        "Updated: 06/2026",
        "Revised January 2024",
        "Page 3 of 12",
    ):
        assert is_source_boilerplate(text), text


def test_source_boilerplate_negatives():
    for text in (
        "Updated the curriculum for the MS4 elective",
        "Revised institutional guidelines for resident supervision",
        "Prepared expert testimony for federal court",
        "CURRICULUM VITAE\nEulho Jung",  # multi-line: real preamble block
        "Curriculum Vitae Committee, School of Medicine",
        "",
        None,
    ):
        assert not is_source_boilerplate(text), text


def test_appendix_renders_without_codes_and_drops_noise():
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    before = len(gen.doc.paragraphs)

    gen._add_remaining_to_appendix([
        ("Some grant text | Role: PI | Status: Under review", "M2A", 10.0),
        ("", "T", 0.0),                # empty — must be skipped
        ("CURRICULUM VITAE", "T", 0.0),  # source furniture — must be skipped
    ])

    # Only the paragraphs this call added -- the blank template already
    # carries 32 pre-existing ilvl=0 list paragraphs (its own section
    # headings), so filtering the whole document by numPr level would also
    # match those.
    added = gen.doc.paragraphs[before:]
    bullets = [p for p in added if _num_level(p) == "0"]
    assert [p.text for p in bullets] == [
        "Some grant text | Role: PI | Status: Under review"
    ]
    assert not any(p.text.startswith("•") for p in added)
    assert not any("[M2A]" in p.text for p in added)


def test_reconsider_routes_unrendered_siblings_home(monkeypatch):
    """The 89HQVQ failure, end to end (LLM stubbed): of an 8-grant mega-entry
    only FSMB was extracted/rendered. The reconsider pass must route the
    unrendered siblings to their sections — including same-code ones — and
    must not duplicate the already-rendered FSMB anywhere."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)

    entry = {
        "text": "eight grants in one blob",
        "taxonomy_code": "M2A",
        "extracted_fields": _FSMB_FIELDS,
    }
    gen._appendix_pending = [(entry, 12.0)]

    segments = [
        (_FSMB_SEGMENT, "M2A"),       # already rendered — must vanish
        (_TEMPLETON_SEGMENT, "M2C"),  # cross-code — routed
        ("SDRME Review Paper Grant | Jung, E. (Co-PI) | Metacognitive "
         "Reflection Training | Status: Under review.", "M2A"),  # same-code — routed
    ]
    monkeypatch.setattr(gen, "_reclassify_entry_segments", lambda text, code: segments)

    routed, appended = [], []
    monkeypatch.setattr(
        gen, "_insert_reconsidered_segment",
        # Returns True: a failed insert now falls back to the appendix (#221).
        lambda text, code: routed.append((text, code)) or True,
    )
    monkeypatch.setattr(
        gen, "_add_remaining_to_appendix", lambda remaining: appended.extend(remaining)
    )

    gen._reconsider_appendix_entries()

    assert (_TEMPLETON_SEGMENT, "M2C") in routed
    assert any(code == "M2A" and "SDRME" in text for text, code in routed)
    assert not any("FSMB" in text for text, _ in routed)
    assert not any("FSMB" in item[0] for item in appended)
    assert appended == []


# ---------------------------------------------------------------- #264

_PROSE_REPLY = (
    "**All segments are retained under M2B: the grant details, project goals "
    "and associated publications directly relate to the funded project.\n"
    "M2B: Synthetic Grant Alpha | Example Agency | 2010-2012"
)


def test_parse_reclassified_segments_well_formed():
    reply = ("M2B: Synthetic Grant Alpha | Example Agency | 2010-2012\n"
             "K2: Participate in clinical teaching conferences\n"
             "KEEP: Attending Physician, Example Hospital, 2004-Present\n"
             "q4a: Editorial board member, Example Journal")
    assert parse_reclassified_segments(reply, "M2B") == [
        ("Synthetic Grant Alpha | Example Agency | 2010-2012", "M2B"),
        ("Participate in clinical teaching conferences", "K2"),
        ("Attending Physician, Example Hospital, 2004-Present", "M2B"),
        ("Editorial board member, Example Journal", "Q4A"),
    ]


_LIVE_SHAPE_REPLY = (
    "Here is the analysis of the CV content:\n\n"
    "M2B: R01 XX000000 Synthetic Study of Signaling, Example Institute, 2010-2015, PI\n"
    "M2B: Project goals: characterize pathway X\n"
    "M2B: Associated publications: Doe J et al 2014\n"
    "K2: Participate in clinical teaching conferences\n\n"
    "**Rationale:**\n"
    "- All segments remain classified as **M2B** because:\n"
    "  - Both grants have clearly defined end dates\n"
    "- No segments meet the threshold of **clearly** belonging elsewhere\n\n"
    "> **Note:** If your institution treats these as active, reclassify them."
)


def test_parse_reclassified_segments_drops_commentary_keeps_code_lines():
    """The real reply shape (preamble + code lines + rationale + note): only
    the code lines are segments. Refusing the whole reply over the preamble
    sent 13 of 18 live replies back to the appendix (#264)."""
    assert parse_reclassified_segments(_LIVE_SHAPE_REPLY, "M2B") == [
        ("R01 XX000000 Synthetic Study of Signaling, Example Institute, 2010-2015, PI", "M2B"),
        ("Project goals: characterize pathway X", "M2B"),
        ("Associated publications: Doe J et al 2014", "M2B"),
        ("Participate in clinical teaching conferences", "K2"),
    ]


def test_parse_reclassified_segments_prose_line_is_dropped_not_a_segment():
    """Commentary with a colon used to parse as a segment whose 'code' was the
    sentence, then rendered as an appendix bullet (#264)."""
    assert parse_reclassified_segments(_PROSE_REPLY, "M2B") == [
        ("Synthetic Grant Alpha | Example Agency | 2010-2012", "M2B")]


@pytest.mark.parametrize("reply", [
    # commentary only, no valid code line: nothing usable
    "All segments are retained under M2B: the grant details, project goals relate.",
    "Here is the analysis of the CV content:\n\n**Rationale:** all stay put.",
    "Note: segments are retained under M2B and the grant details relate.",
    "M2Z: Synthetic Grant Alpha, Example Agency, 2010-2012",
])
def test_parse_reclassified_segments_no_code_line_is_none(reply):
    assert parse_reclassified_segments(reply, "M2B") is None


@pytest.mark.parametrize("prefix", ["ALL", "Note", "M2Z", "**All segments", "> **Note"])
def test_parse_reclassified_segments_non_taxonomy_prefix_never_a_segment(prefix):
    reply = (f"{prefix}: segments are retained under M2B and details relate.\n"
             "M2B: Synthetic Grant Alpha, Example Agency, 2010-2012")
    assert parse_reclassified_segments(reply, "M2B") == [
        ("Synthetic Grant Alpha, Example Agency, 2010-2012", "M2B")]


def test_parse_reclassified_segments_keep_short_and_empty_edges():
    reply = ("  K1: Synthetic course lecture series, Example University  \n"
             "KEEP: Synthetic Grant Alpha, Example Agency, 2010-2012\n"
             "K2: too short\n")
    # KEEP resolves to the original code; a segment of <=10 chars is dropped;
    # surrounding whitespace on a line is stripped.
    assert parse_reclassified_segments(reply, "M2B") == [
        ("Synthetic course lecture series, Example University", "K1"),
        ("Synthetic Grant Alpha, Example Agency, 2010-2012", "M2B")]
    # KEEP with an unknown original code has no home: code None.
    assert parse_reclassified_segments(
        "KEEP: Synthetic Grant Alpha, Example Agency", "?") == [
        ("Synthetic Grant Alpha, Example Agency", None)]
    # Nothing usable -> None, not [].
    assert parse_reclassified_segments("K2: too short", "M2B") is None
    assert parse_reclassified_segments("no colons here", "M2B") is None


def test_reclassify_prose_reply_never_reaches_document(monkeypatch):
    """End to end with call_llm stubbed to return commentary with no code
    line: the original entry text is what lands in the appendix, never the
    model's prose."""
    import unified_pipeline.stage_6_word_template as st6

    monkeypatch.setattr(st6, "call_llm", lambda **kw: {"content": (
        "**All segments are retained under M2B: the grant details, project "
        "goals and associated publications directly relate to the project.")})
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    original = "Synthetic Grant Alpha, Example Agency, 2010-2012, funded project"
    gen._appendix_pending = [({"text": original, "taxonomy_code": "M2B",
                               "extracted_fields": {}}, 10.0)]

    gen._reconsider_appendix_entries()

    texts = [p.text for p in gen.doc.paragraphs]
    assert not any("All segments are retained" in t or "grant details" in t
                   for t in texts)
    assert any(original in t for t in texts)


def test_reclassify_live_shape_reply_renders_segments_without_commentary(monkeypatch):
    """End to end with the real reply shape: the code lines are used and none
    of the preamble / rationale / note text reaches the document."""
    import unified_pipeline.stage_6_word_template as st6

    monkeypatch.setattr(st6, "call_llm", lambda **kw: {"content": _LIVE_SHAPE_REPLY})
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    original = "Synthetic Study of Signaling, Example Institute, 2010-2015, PI; project goals"
    gen._appendix_pending = [({"text": original, "taxonomy_code": "M2B",
                               "extracted_fields": {}}, 10.0)]

    gen._reconsider_appendix_entries()

    joined = "\n".join(p.text for p in gen.doc.paragraphs)
    for leak in ("Here is the analysis", "Rationale", "If your institution",
                 "clearly", "threshold"):
        assert leak not in joined
    # Text only the PARSED segments carry: the fallback that keeps the
    # original entry text can never produce it, so this pins the call site.
    assert "R01 XX000000" in joined
    assert "Associated publications: Doe J et al 2014" in joined
    assert gen.stats["appendix_segments_reconsidered"] >= 1


def test_reclassify_keep_line_routes_under_original_code(monkeypatch):
    """End to end: a 'KEEP: ...' segment of an M2B entry routes as M2B (the
    original code), not as '?' / the appendix (#209, #264)."""
    import unified_pipeline.stage_6_word_template as st6

    monkeypatch.setattr(st6, "call_llm", lambda **kw: {"content": (
        "KEEP: Synthetic Grant Beta, Example Agency, 2011-2013, completed")})
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    gen._appendix_pending = [({"text": "blob of text", "taxonomy_code": "M2B",
                               "extracted_fields": {}}, 10.0)]
    routed = []
    monkeypatch.setattr(
        gen, "_insert_reconsidered_segment",
        lambda text, code: routed.append((text, code)) or True)
    appended = []
    monkeypatch.setattr(
        gen, "_add_remaining_to_appendix", lambda r: appended.extend(r) or [])

    gen._reconsider_appendix_entries()

    assert routed == [
        ("Synthetic Grant Beta, Example Agency, 2011-2013, completed", "M2B")]
    assert appended == []


@pytest.mark.parametrize("line", [
    "- K2: Participate in clinical teaching conferences",
    "* K2: Participate in clinical teaching conferences",
    "\u2022 K2: Participate in clinical teaching conferences",
    "> K2: Participate in clinical teaching conferences",
    "# K2: Participate in clinical teaching conferences",
    "1. K2: Participate in clinical teaching conferences",
    "**K2:** Participate in clinical teaching conferences",
    "- **K2:** Participate in clinical teaching conferences",
])
def test_parse_reclassified_segments_strips_markdown_prefix(line):
    assert parse_reclassified_segments(line, "M2B") == [
        ("Participate in clinical teaching conferences", "K2")]


def test_parse_reclassified_segments_markdown_prefix_commentary_still_dropped():
    reply = ("- **Note:** everything stays put and details relate.\n"
             "1. Rationale: grants are complete.\n"
             "- KEEP: Synthetic Grant Alpha, Example Agency, 2010-2012")
    assert parse_reclassified_segments(reply, "M2B") == [
        ("Synthetic Grant Alpha, Example Agency, 2010-2012", "M2B")]


def test_all_noise_batch_creates_no_appendix():
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)
    before = len(gen.doc.paragraphs)

    gen._add_remaining_to_appendix([("", "T", 0.0), ("Page 2 of 9", "T", 0.0)])

    assert len(gen.doc.paragraphs) == before
    assert not any("T. APPENDIX" in p.text for p in gen.doc.paragraphs)


def test_a_trial_enrollment_status_is_not_a_completed_award():
    # #291: a trial closed to accrual is still running; only its end date may
    # move it to Past Funding. A plain "Completed"/"Closed" still does.
    for status in ("Closed to accrual", "Enrollment completed", "Accrual completed",
                   "Recruitment completed", "Closed to new patients"):
        assert grant_status_rebucket_target(status) == (None, None), status
    assert grant_status_rebucket_target("Closed")[0] == "M2B"
    assert grant_status_rebucket_target("Completed 2021")[0] == "M2B"


# -------------------------------------------------------------- A5IZ6Q (#420)
#
# Wire tests, real generate() against the bundled WCM template: a stage-2
# structurally-recovered table row (#420's `recover_unclaimed_table_rows`)
# that duplicates content already RENDERED must not also land in the
# Appendix. Synthetic reproduction of the A5IZ6Q incident shape -- one fused
# M2B grant entry plus single-field `recovered_row` siblings, each one line
# of the same fused text the model's delimiter already captured whole.
#
# Round 3 (simplify, LEAD directive): `recovered_row_already_rendered`
# (`stage6/dedup.py`) drops a recovered row only when every non-trivial
# value cell of its own raw text is already printed somewhere in the
# document's rendered body -- normalized (casefold, whitespace-collapsed,
# word-boundary matched), and never scoped to a specific parent entry.
# Earlier rounds' parent-block scoping is gone: it was itself a source of
# content loss (two grants sharing a value could resolve to the same
# block -- see the shared-agency test below), and the simpler rule can only
# ever KEEP more content, never drop content that isn't visible elsewhere.

_RECOVERY_GRANT_TEXT = (
    "Award Source: | Fictional Research Foundation\n"
    "Project title: | Synthetic Tools for Data Curation\n"
    "Annual direct costs: | $15,000.00\n"
    "Duration of support: | 00/2021-00/2022\n"
    "Name of Principal Investigator: | A. Researcher\n"
    "Your percent (%) effort: | 1%"
)

_RECOVERY_OWNER_ENTRY = {"text": "Name: A. Researcher", "taxonomy_code": "A",
                         "extracted_fields": {}, "element_idx_start": 0}

_RECOVERY_GRANT_ENTRY = {
    "text": _RECOVERY_GRANT_TEXT,
    "taxonomy_code": "M2B",
    "element_idx_start": 300,
    "element_idx_end": 302,
    "extracted_fields": {
        "title": "Synthetic Tools for Data Curation",
        "agency": "Fictional Research Foundation",
        "pi_name": "A. Researcher",
        "pi_role": "PI",
        "start_date": "2021",
        "end_date": "2022",
        "total_funding": "$15,000.00",
        "percent_effort": "1%",
    },
}


def _recovered_row(suffix: str, text: str) -> dict:
    # `parent_idx` is carried here only because real stage-2 output always
    # sets it (#420) -- `recovered_row_already_rendered` reads no such
    # field; the value is never resolved back to a parent entry any more.
    return {"text": text, "taxonomy_code": "T", "recovered_row": True,
            "parent_idx": 300, "element_idx_start": f"300.{suffix}",
            "extracted_fields": {}, "hierarchy": ["Past Funding"]}


def _all_text(doc: Document) -> str:
    lines = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


def test_recovered_rows_dropped_when_their_values_already_rendered(tmp_path):
    """Five of the grant's six recovered rows are verbatim duplicates of a
    VALUE the grant's own table renders unchanged and must not repeat in the
    Appendix. The sixth (Duration) carries a raw "00/2021-00/2022" value
    that stage 6 reformats to "2021-2022" on the way to the render slot --
    this function does no date parsing (LEAD directive: normalized
    casefold/whitespace-collapse/word-boundary matching only), so that one
    value is never confirmed and the row stays. Disclosed trade-off, not a
    bug: duplication, never loss (see recovered_row_already_rendered's
    docstring)."""
    entries = [_RECOVERY_OWNER_ENTRY, _RECOVERY_GRANT_ENTRY,
               _recovered_row("0", "Award Source: | Fictional Research Foundation"),
               _recovered_row("1", "Project title: | Synthetic Tools for Data Curation"),
               _recovered_row("2", "Annual direct costs: | $15,000.00"),
               _recovered_row("3", "Duration of support: | 00/2021-00/2022"),
               _recovered_row("4", "Name of Principal Investigator: | A. Researcher"),
               _recovered_row("5", "Your percent (%) effort: | 1%")]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T420A", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))

    # The grant rendered in the body...
    assert "Fictional Research Foundation" in full_text
    # ...and every field the grant table actually renders unchanged appears
    # exactly once: the five matching recovered rows did not repeat it.
    assert full_text.count("Fictional Research Foundation") == 1
    assert full_text.count("Synthetic Tools for Data Curation") == 1
    assert full_text.count("$15,000.00") == 1
    assert full_text.count("A. Researcher") == 1
    assert full_text.count("Your percent (%) effort") == 1
    # The Duration row's raw "00/2021-00/2022" never renders verbatim (the
    # grant table shows the reformatted "2021-2022"), so it stays -- one
    # duplicate line, not zero, and the Appendix exists because of it.
    assert "T. APPENDIX" in full_text
    assert "Duration of support: — 00/2021-00/2022" in full_text
    # Confirms the trailing boundary too: the reformatted value never shows
    # the raw one's leading "00/" anywhere.
    assert full_text.count("00/2021-00/2022") == 1


def test_recovered_row_with_unrendered_value_still_reaches_appendix(tmp_path):
    """The negative case: a recovered row whose value is not printed ANYWHERE
    in the rendered document (a row the model's delimiter genuinely
    skipped -- #420's backstop exists for exactly this) must still
    surface."""
    missed_row = _recovered_row(
        "5", "Non-financial support: | Conference travel support")
    entries = [_RECOVERY_OWNER_ENTRY, _RECOVERY_GRANT_ENTRY, missed_row]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T420B", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))
    assert "T. APPENDIX" in full_text
    assert "Conference travel support" in full_text


def test_recovered_row_whose_field_never_reached_a_render_slot_still_reaches_appendix(tmp_path):
    """The gap a blind review of the original A5IZ6Q fix caught, still live
    under the round-3 rule: a recovered row's raw text can echo the fused
    parent's own text while the field it carries never reaches a render
    slot at all. Stage 6's grant table is fixed-slot (CLAUDE.md "Stage 6
    drops unnamed fields") and has no grant-number row at all unless
    extracted_fields carries one. Here the parent's raw text has a
    grant-number line but extracted_fields does not, so no renderer ever
    prints it -- the value is not confirmed anywhere and the row must still
    reach the Appendix."""
    grant_text = (
        "Award Source: | Fictional Research Foundation\n"
        "Project title: | Synthetic Tools for Data Curation\n"
        "Grant number: | R01-ZZ98765\n"
        "Duration of support: | 2021-2022"
    )
    grant_entry = {
        "text": grant_text,
        "taxonomy_code": "M2B",
        "element_idx_start": 400,
        "element_idx_end": 400,
        "extracted_fields": {
            "title": "Synthetic Tools for Data Curation",
            "agency": "Fictional Research Foundation",
            "start_date": "2021",
            "end_date": "2022",
            # No grant_number key -- the renderer never receives it, even
            # though the raw text above carries the exact same line.
        },
    }
    row = {"text": "Grant number: | R01-ZZ98765", "taxonomy_code": "T",
           "recovered_row": True, "parent_idx": 400,
           "element_idx_start": "400.0", "extracted_fields": {},
           "hierarchy": ["Past Funding"]}
    entries = [_RECOVERY_OWNER_ENTRY, grant_entry, row]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T420D", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))
    assert "T. APPENDIX" in full_text
    assert "R01-ZZ98765" in full_text
    # The grant itself still rendered, and did not gain a grant-number cell.
    assert "Fictional Research Foundation" in full_text


# ---------------------------------------------- shared-agency cross-vouching
#
# The exact repro that sank round 2's parent-block scoping (`_parent_rendered_block`
# picked the FIRST rendered block carrying ANY one of a parent's identifying
# values -- a shared agency resolved two different grants to the SAME
# block, so a grant whose own table rendered nothing for a field could still
# lose its recovered row to a same-agency sibling's render). Round 3 removes
# scoping entirely; these tests prove the removal does not reopen the loss.

def test_shared_agency_grant_with_unrendered_recovered_rows_stays_in_appendix(tmp_path):
    """Two M2B grants share an agency ('National Institutes of Health').
    Grant A's own extracted_fields carry dates and percent_effort and render
    normally. Grant B's RAW text also carries a duration line and a
    percent-effort line -- so recover_unclaimed_table_rows splits them out
    as B's own recovered rows -- but B's extracted_fields carry only title
    and agency, so B's OWN grant table renders neither value (label only,
    blank cell). B's duration row never verbatim-matches (stage 6 reformats
    the date), and B's effort row ('5%') must not cross-match A's rendered
    '25%' merely because they share a digit. Both of B's rows must survive
    to the Appendix."""
    grant_a = {
        "text": ("Award Source: | National Institutes of Health\n"
                 "Project title: | Alpha Sequencing Initiative\n"
                 "Duration of support: | 00/2019-00/2020\n"
                 "Your percent (%) effort: | 25%"),
        "taxonomy_code": "M2B",
        "element_idx_start": 600,
        "element_idx_end": 600,
        "extracted_fields": {
            "title": "Alpha Sequencing Initiative",
            "agency": "National Institutes of Health",
            "start_date": "2019",
            "end_date": "2020",
            "percent_effort": "25%",
        },
    }
    grant_b = {
        "text": ("Award Source: | National Institutes of Health\n"
                 "Project title: | Beta Imaging Cohort\n"
                 "Duration of support: | 00/2019-00/2020\n"
                 "Your percent (%) effort: | 5%"),
        "taxonomy_code": "M2B",
        "element_idx_start": 700,
        "element_idx_end": 700,
        "extracted_fields": {
            "title": "Beta Imaging Cohort",
            "agency": "National Institutes of Health",
            # No start_date/end_date/percent_effort: B's own table renders
            # neither value, even though the raw text above (and B's own
            # recovered-row siblings below) carries them verbatim.
        },
    }
    row_b_duration = {"text": "Duration of support: | 00/2019-00/2020",
                      "taxonomy_code": "T", "recovered_row": True,
                      "parent_idx": 700, "element_idx_start": "700.0",
                      "extracted_fields": {}, "hierarchy": ["Past Funding"]}
    row_b_effort = {"text": "Your percent (%) effort: | 5%",
                    "taxonomy_code": "T", "recovered_row": True,
                    "parent_idx": 700, "element_idx_start": "700.1",
                    "extracted_fields": {}, "hierarchy": ["Past Funding"]}
    entries = [_RECOVERY_OWNER_ENTRY, grant_a, grant_b, row_b_duration, row_b_effort]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T946A", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))
    # Both grants rendered in the body, under the same agency.
    assert "Alpha Sequencing Initiative" in full_text
    assert "Beta Imaging Cohort" in full_text
    assert full_text.count("National Institutes of Health") == 2
    # B's own recovered rows were NOT dropped -- they reached the Appendix
    # (_clean_inline_tabs rejoins the raw " | " cell separator as " — " for
    # display), proving neither was found rendered anywhere.
    assert "T. APPENDIX" in full_text
    assert "Duration of support: — 00/2019-00/2020" in full_text
    # And the word-boundary half of the same repro: B's "5%" appears only
    # once, as its own Appendix bullet -- never merged with A's unrelated
    # "25%" cell (a plain, non-boundary substring check would have read A's
    # rendered "25%" as confirming B's "5%" and dropped this row).
    assert full_text.count("Your percent (%) effort: — 5%") == 1


def test_recovered_row_dropped_when_confirmed_only_by_a_different_records_render(tmp_path):
    """Round 3's accepted trade-off, at the wire level. Removing per-parent
    scoping means a drop no longer asks WHICH record printed a value, only
    whether it is already in the rendered document. Here grant_entry's own
    raw text never mentions percent effort at all -- recover_unclaimed_table_rows
    never produced the row below from this grant -- but grant_entry's
    extracted_fields legitimately carry percent_effort '5%', which DOES
    render in its own table. The recovered row's value is therefore
    confirmed, wherever it came from, and the row is dropped: the value is
    genuinely visible in the document, which is the only guarantee this
    drop makes."""
    grant_entry = {
        "text": ("Award Source: | Fictional Research Foundation\n"
                 "Project title: | Synthetic Tools for Data Curation\n"
                 "Duration of support: | 2021-2022"),
        "taxonomy_code": "M2B",
        "element_idx_start": 500,
        "element_idx_end": 500,
        "extracted_fields": {
            "title": "Synthetic Tools for Data Curation",
            "agency": "Fictional Research Foundation",
            "start_date": "2021",
            "end_date": "2022",
            "percent_effort": "5%",  # renders even though the raw text above
                                     # never mentions percent effort at all
        },
    }
    row = {"text": "Your percent (%) effort: | 5%",
           "taxonomy_code": "T", "recovered_row": True,
           "parent_idx": 500, "element_idx_start": "500.0",
           "extracted_fields": {}, "hierarchy": ["Past Funding"]}
    entries = [_RECOVERY_OWNER_ENTRY, grant_entry, row]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T946B", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))
    # The value rendered once, from the grant table's own row...
    assert full_text.count("Your percent (%) effort: | 5%") == 1
    # ...and the recovered row was dropped: no Appendix duplicate, no
    # Appendix section at all (this is the only unmapped entry in the run).
    assert "T. APPENDIX" not in full_text
    assert "Your percent (%) effort: — 5%" not in full_text


# ---------------------------------------------- instruction-box vouching (round 4)
#
# `_drop_recovered_row_duplicates` runs before `_remove_instruction_box`
# (generate() strips the box only after every content-search fill has run --
# see that call site's comment), so its "already rendered" haystack would,
# without the round-4 fix, briefly still include the box's own prompt text --
# which literally reads "...enter 'Not Applicable' or 'N/A'"... "'Local'
# refers to the home institution"... "please record 04/2022". A recovered
# row whose real value happens to equal one of those fragments would be
# vouched for by the box, dropped here, and then the box itself deleted
# before save: the value would appear nowhere in the output. These use the
# REAL bundled template (not a synthetic fixture) because the bug is in what
# that specific template's box text contains.

def test_recovered_row_not_falsely_confirmed_by_the_instruction_box_date_example(tmp_path):
    """'04/2022' is the box's own worked example of the mm/yyyy date format.
    A recovered row whose real value is literally '04/2022' must not be
    dropped merely because the box happens to contain that string."""
    missed_row = _recovered_row("5", "Duration of support: | 04/2022")
    entries = [_RECOVERY_OWNER_ENTRY, _RECOVERY_GRANT_ENTRY, missed_row]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T959A", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))
    # The box itself is gone (strip_template_instructions defaults on)...
    assert "delete this instruction box" not in full_text.lower()
    # ...and the recovered row's real value still made it into the output.
    assert "T. APPENDIX" in full_text
    assert "Duration of support: — 04/2022" in full_text


def test_recovered_row_not_falsely_confirmed_by_the_instruction_box_local_example(tmp_path):
    """'Local' is the box's own definition text ("'Local' refers to the home
    institution"). Same failure mode as the date example above, with a
    non-numeric value."""
    missed_row = _recovered_row("5", "Geographic scope: | Local")
    entries = [_RECOVERY_OWNER_ENTRY, _RECOVERY_GRANT_ENTRY, missed_row]

    gen = WCMTemplateGenerator(verbose=False, recover_unrendered_records=False)
    gen._reconsider_appendix_entries = lambda: []
    data = {"document_uid": "T959B", "entries": entries}
    input_path = tmp_path / "in.json"
    output_path = tmp_path / "out.docx"
    input_path.write_text(json.dumps(data))
    gen.generate(str(input_path), str(output_path), research_summary_path=None)

    full_text = _all_text(Document(str(output_path)))
    assert "delete this instruction box" not in full_text.lower()
    assert "T. APPENDIX" in full_text
    assert "Geographic scope: — Local" in full_text


def test_rendered_output_lines_excludes_instruction_box_only_when_asked():
    """Unit-level pin on the flag itself: `_recover_unrendered_records`'s
    call (the pre-existing, out-of-scope one) must keep seeing the box's
    text by default, while a caller that opts in does not."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document(gen.template_path)

    default_lines = gen._rendered_output_lines()
    excluded_lines = gen._rendered_output_lines(exclude_instruction_box=True)

    box_phrase = "when preparing the wcm cv template"
    assert any(box_phrase in ln.lower() for ln in default_lines)
    assert not any(box_phrase in ln.lower() for ln in excluded_lines)
    # Real content (the personal-data table) is unaffected either way.
    assert any("work email" in ln.lower() for ln in excluded_lines)


def _append_tracked_insertion(paragraph, text: str) -> None:
    """A w:ins run, the shape `_add_track_change_insertion` writes."""
    ins = paragraph._p.makeelement(qn("w:ins"), {qn("w:id"): "1", qn("w:author"): "Test"})
    run = ins.makeelement(qn("w:r"), {})
    t = run.makeelement(qn("w:t"), {})
    t.text = text
    run.append(t)
    ins.append(run)
    paragraph._p.append(ins)


def test_rendered_output_lines_reads_tracked_insertions_only_when_asked():
    """python-docx's `paragraph.text` skips w:ins runs, which is how the
    research summary and enriched citations are written (E27). A caller that
    opts in sees them in body paragraphs and table cells; the default is
    unchanged for every other caller."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    body_para = gen.doc.add_paragraph("Plain lead ")
    _append_tracked_insertion(body_para, "Inserted body words")
    cell_para = gen.doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    _append_tracked_insertion(cell_para, "Inserted cell words")

    default_lines = gen._rendered_output_lines()
    tracked_lines = gen._rendered_output_lines(include_tracked_insertions=True)

    assert "Plain lead " in default_lines
    assert not any("Inserted" in ln for ln in default_lines)
    assert "Plain lead Inserted body words" in tracked_lines
    assert "Inserted cell words" in tracked_lines
