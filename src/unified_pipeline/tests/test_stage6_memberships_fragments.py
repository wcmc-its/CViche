"""Section I (memberships) newline-blind fix (#476, PR1 of the accuracy wave).

`memberships.py:115` was one of the ten `entry_lines`-based call sites,
gating `_parse_multi_membership_entry` on `len(lines) > 2`. A fully blind
entry (no literal newline at all) always reads as exactly one line under
`entry_lines`, so a "Type1 | Org1 | Date1 | Type2 | Org2 | Date2" entry fused
without a single newline never reached the multi-membership parser and fell
to the single-membership path, which uses `extracted_fields` -- describing
only the first membership.

`_entry_parts` fixes exactly the blind case (entry_lines already returns
more than one part unchanged) by adding '|' as a boundary when the whole
entry is one line. Migrating this call site to `entry_fragments` outright
(tab included) was tried and reverted after reading the actual render-gate
diff: 038WKA and 11 near-duplicate "Jonathan Nahmias" uids (12 total) carry
a single
membership with three tab-separated fields ("Member of the American
Psychiatric Association\\tFebruary 2018 - Present\\tI have attended and
presented..."); `_parse_multi_membership_entry`'s classifier needs a
membership-type part to be <=3 words and a date part to match a strict
numeric pattern, both of which that entry's fields fail, so splitting its
tabs turned one membership into three garbled "organizations" with no
dates. Tab stays unsplit here, matching what `entry_lines` already did.

The call site also gates on the PARSED membership count (`len(memberships)
> 1`), not just the raw part count: a single membership whose fields split
into exactly three parts ("Fellow | American Academy of Pediatrics |
1/1997-present", the farm's actual blind-'|' shape on 1FRABQ and 11 other
uids) crosses the >2 part-count threshold but must still resolve to one
membership.

That stricter gate applies ONLY to entries this fix newly splits. Applied to
every entry it also silently changed a path #476 never touched: an
already-multi-line entry that parses to exactly one membership rendered the
parsed row before #476 and would have fallen through to the extracted-fields
row instead. No farm entry has that shape, so no gate could see it; the
`len(lines) > 2` half of the call site's condition preserves it, and
`test_multiline_single_membership_still_renders_the_parsed_row` pins it
against the measured dev output.

The review round on this PR added the rest of the file. Grouped by what each
block pins:

- the `stats['entries_inserted']` double count (`_add_table_row` and both of
  its callers each incremented it, so one rendered membership reported two);
- `_add_table_row` silently dropping values past the table's last column;
- the two date paths formatting dates differently -- both now end at one
  `format_date_range(..., 'I')` call;
- `_entry_parts` splitting on a '|' that is punctuation inside one field
  rather than a membership boundary;
- membership-type suppression by bare substring rather than whole words;
- the case-sensitive "Date" test on the fallback table's second header;
- `original_text[:150]` as the organization fallback.

Block (s) is the second round: the 66-CV render gate read an organization
cell that the new fallback had shortened to "February" on an entry whose
whole text is "February 2018 - Present". The recovery had lifted the trailing
date by matching the year alone and left the month behind as the
"organization" -- the raw-text truncation the fallback replaced, reintroduced
by the fallback itself. The block pins the general rule rather than that one
cell: a recovered organization may give up a tail only when the date column
now carries every word of it.

Block (t) is the third round, and the same class of loss one level up. The
render gate read 2082_Dr_Scot's memberships table missing "Faculty Associate"
on two of three rows. The fallback's structured-alias branch resolved the
organization and the membership type TOGETHER: an `institution` field named
the organization, the branch returned on it, and the `role` field sitting
beside it was never read -- so the role reached no cell of the document, and
the third row, which takes a different path, kept its role and left the table
internally inconsistent. Block (s)'s no-loss check would have caught it, but
round 2 had exempted the alias branch from it on purpose. The block pins both
halves of the repair: the two values are resolved independently, and the
exemption is gone.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_memberships_fragments.py -p no:cacheprovider
"""

import ast
import re
import sys
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import pytest  # noqa: E402
from docx import Document  # noqa: E402

from unified_pipeline.core.render_check import entry_fragments, entry_lines  # noqa: E402
from unified_pipeline.stage6.sections import memberships as memberships_module  # noqa: E402
from unified_pipeline.stage6.sections.memberships import (  # noqa: E402
    MembershipsRowShapeError,
    _entry_parts,
    _organization_fallback,
    _split_trailing_dates,
    _type_already_named,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


# Genuinely multi-line, farm-derived shape (2071_Zuschlag_Cv's actual I
# entry): entry_lines already returns >1 part, each carrying its own unsplit
# tabs -- _entry_parts must not touch it further.
MULTILINE_CASES = [
    "2009-2012\t\tStudent Osteopathic Surgical Association\t\t\n\n2009-2012\t\tFlorida Osteopathic Medical Association",
    "Member\nElected Member | Org1\nOrg2 | date1\ndate2",
]

# Single-line (blind) shapes with a tab that must ALSO pass through
# untouched (038WKA's real farm shape).
BLIND_TAB_CASE = (
    "Member of the American Psychiatric Association\tFebruary 2018 – Present"
    "\tI have attended and presented at the American Psychiatric Association "
    "Annual Meeting (see section on presentations) and remain up to date "
    "with their newsletter."
)

MEMBERSHIPS_HEADING = "I. PROFESSIONAL ORGANIZATIONS AND SOCIETY MEMBERSHIPS"
DEFAULT_HEADER = ("Organization", "Date (yyyy-yyyy)")


def _generator_with_table(header=DEFAULT_HEADER, table_before_heading=False):
    """A real WCMTemplateGenerator over a hand-built section I table.

    `table_before_heading` puts the table ahead of the heading paragraph, so
    `_find_table_after_paragraph` misses and `_fill_memberships` takes the
    "Organization" fallback search instead -- the only way to reach the
    second-header guard.
    """
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    if table_before_heading:
        table = gen.doc.add_table(rows=1, cols=len(header))
        gen.doc.add_paragraph(MEMBERSHIPS_HEADING)
    else:
        gen.doc.add_paragraph(MEMBERSHIPS_HEADING)
        table = gen.doc.add_table(rows=1, cols=len(header))
    for i, label in enumerate(header):
        table.rows[0].cells[i].text = label
    return gen, table


def _fill(entries, header=DEFAULT_HEADER, table_before_heading=False):
    """Render `entries` into section I; returns (generator, rendered rows)."""
    gen, table = _generator_with_table(header, table_before_heading)
    gen._fill_memberships(entries)
    return gen, [[c.text for c in r.cells] for r in table.rows[1:]]


def _render_memberships(entries):
    return _fill(entries)[1]


# --- (0) #862: zero entries still clears the placeholder row ----------------

def test_no_entries_clears_the_placeholder_row():
    """#862: `_fill_memberships([])` now clears the table's blank
    placeholder row instead of returning before `_clear_table_data` runs --
    the WCM template ships that row, and it used to survive into the
    delivered document on every CV with zero I entries. Still does not
    count the table as populated, since nothing was written."""
    gen, table = _generator_with_table()
    table.add_row()
    table.rows[1].cells[0].text = "stale placeholder"
    gen._fill_memberships([])
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    assert rows == []
    assert gen.stats['tables_populated'] == 0


def test_no_entries_with_a_foreign_fallback_table_is_left_alone():
    """The zero-entry path reaches the same fallback guard the entries path
    already uses (`test_fallback_table_without_a_date_header_is_still_rejected`):
    a table without a date-shaped second header column is rejected, so a
    foreign table sitting where memberships would fall back to is not
    cleared just because there happen to be no entries."""
    gen, table = _generator_with_table(
        header=("Organization", "Certificate #"), table_before_heading=True)
    table.add_row()
    table.rows[1].cells[0].text = "foreign stale row"
    gen._fill_memberships([])
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    assert rows == [["foreign stale row", ""]]
    assert gen.stats['tables_populated'] == 0


def test_no_entries_with_a_foreign_table_after_the_heading_is_left_alone():
    """The primary lookup is a forward paragraph scan with no table
    identity: a template variant could put someone else's table right
    after the memberships heading. With zero entries there is no
    data-driven signal to catch that, so `_is_memberships_table` refuses
    to clear a table whose second header column is not date-shaped."""
    gen, table = _generator_with_table(header=("Organization", "Certificate #"))
    table.add_row()
    table.rows[1].cells[0].text = "foreign stale row"
    gen._fill_memberships([])
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    assert rows == [["foreign stale row", ""]]
    assert gen.stats['tables_populated'] == 0


# --- (a) no line the old newline split produced is lost ---------------------

def test_multiline_entries_are_returned_unchanged():
    for case in MULTILINE_CASES:
        old = entry_lines(case)
        assert len(old) > 1, f"fixture {case!r} is not actually multi-line"
        assert _entry_parts(case) == old, f"diverged on already-multiline {case!r}"


def test_blind_line_with_a_tab_is_also_returned_whole():
    old = entry_lines(BLIND_TAB_CASE)
    assert old == [BLIND_TAB_CASE]
    assert _entry_parts(BLIND_TAB_CASE) == old


# --- (b) positive control: fails on dev today --------------------------------

def test_positive_control_pipe_blind_multi_membership_entry_gains_both_rows():
    """Today's `entry_lines` returns one opaque part for this text (no
    '\\n'), so `len(lines) > 2` is never true and `_parse_multi_membership_
    entry` never runs; the entry falls to the single-membership branch,
    which uses `extracted_fields` describing only the first membership and
    silently drops the second. Fails on dev.

    The dates read "2015-Present"/"2018-Present" rather than the raw
    "2015-present"/"2018-present" the parser hands back because the
    multi-membership path now formats its dates through the same
    `format_date_range(..., 'I')` call the single path always used."""
    entry = {
        "text": "Member | State Medical Society | 2015-present | Fellow | National Surgical Association | 2018-present",
        "extracted_fields": {
            "organization": "State Medical Society",
            "start_date": "2015",
            "end_date": "present",
        },
    }
    rows = _render_memberships([entry])
    assert len(rows) == 2, f"expected 2 memberships, got {rows}"
    assert rows[0][0] == "Member, State Medical Society"
    assert rows[0][1] == "2015-Present"
    assert rows[1][0] == "Fellow, National Surgical Association"
    assert rows[1][1] == "2018-Present"


# --- (c) negative controls: unaffected shapes render exactly as today -------

def test_negative_control_single_membership_unchanged():
    entry = {
        "text": "Fellow, American College of Surgeons",
        "extracted_fields": {
            "organization": "American College of Surgeons",
            "membership_type": "Fellow",
            "start_date": "2015",
            "end_date": "Present",
        },
    }
    rows = _render_memberships([entry])
    assert len(rows) == 1
    assert rows[0][0] == "Fellow, American College of Surgeons"


def test_negative_control_farm_pipe_blind_three_field_single_membership_unchanged():
    """1FRABQ's actual shape: one membership whose fields are pipe-joined
    into exactly three parts. `_entry_parts` DOES split it, but the
    len(memberships) > 1 gate keeps it on the single-membership path
    (today's rendering: the organization field alone, membership_type field
    absent, formatted date range) -- unchanged from dev."""
    entry = {
        "text": "Fellow | American Academy of Pediatrics | 1/1997-present",
        "extracted_fields": {
            "organization": "American Academy of Pediatrics",
            "start_date": "1997-01-01",
            "end_date": "present",
        },
    }
    rows = _render_memberships([entry])
    assert len(rows) == 1
    assert rows[0][0] == "American Academy of Pediatrics"


def test_multiline_single_membership_still_renders_the_parsed_row():
    """An already-multi-line entry that `_parse_multi_membership_entry`
    resolves to exactly ONE membership is a path #476 must not touch: it is
    not newline-blind, so the parsed-count gate added for the blind case must
    not reach it. Measured on origin/dev this renders the PARSED row
    ('Fellow, American Academy of Pediatrics'), not the extracted-fields row
    ('American Academy of Pediatrics') -- the extracted fields here
    deliberately differ from the parse so the two paths are distinguishable
    in the assertion. The date reads '1997-Present' on both paths now that
    they share one formatter."""
    entry = {
        "text": "Fellow\nAmerican Academy of Pediatrics\n1997-present",
        "extracted_fields": {
            "organization": "American Academy of Pediatrics",
            "start_date": "1997-01-01",
            "end_date": "present",
        },
    }
    rows = _render_memberships([entry])
    assert rows == [["Fellow, American Academy of Pediatrics", "1997-Present"]], rows


# --- (d) review item 1: entries_inserted is counted once per rendered row ----

def test_one_membership_reports_exactly_one_inserted_entry():
    """`_add_table_row` incremented stats['entries_inserted'] and so did both
    of its callers in `_fill_memberships`, so a single rendered membership
    reported two inserted entries -- a live defect in the run's own stats,
    not a style point. Measured over the 66-CV farm the section reported 424
    inserted entries for 212 rendered rows; it now reports 212."""
    entry = {
        "text": "Fellow, American College of Surgeons",
        "extracted_fields": {
            "organization": "American College of Surgeons",
            "membership_type": "Fellow",
            "start_date": "2015",
            "end_date": "Present",
        },
    }
    gen, rows = _fill([entry])
    assert len(rows) == 1
    assert gen.stats["entries_inserted"] == 1


def test_multi_membership_entry_reports_one_inserted_entry_per_rendered_row():
    """The other call site: two rendered rows out of one source entry must
    report two, not four."""
    entry = {
        "text": "Member | State Medical Society | 2015-present | Fellow | National Surgical Association | 2018-present",
        "extracted_fields": {"organization": "State Medical Society"},
    }
    gen, rows = _fill([entry])
    assert len(rows) == 2
    assert gen.stats["entries_inserted"] == 2


# --- (e) review item 2: surplus values are refused, not dropped -------------

def test_row_writer_rejects_more_values_than_the_table_has_columns():
    """`_add_table_row` wrote values only while `i < len(row.cells)`, so a
    caller handing it more values than the table has columns lost the
    surplus with no error and no log line. That is a caller/template
    programming error, so it now raises rather than emitting a document that
    quietly lost a column."""
    gen, table = _generator_with_table()
    with pytest.raises(MembershipsRowShapeError):
        gen._add_table_row(table, ["organization", "dates", "surplus"])
    assert len(table.rows) == 1, "a rejected row must not be left in the table"


def test_row_writer_accepts_one_value_per_column():
    gen, table = _generator_with_table()
    gen._add_table_row(table, ["American College of Surgeons", "2015-Present"])
    assert [c.text for c in table.rows[1].cells] == [
        "American College of Surgeons", "2015-Present"]


def test_every_add_table_row_call_site_passes_exactly_two_values():
    """The proof that no CV content can trigger the error above: `data` is a
    two-element list literal at both call sites, so its length is fixed at
    import time and nothing extraction produces can change it. The section's
    table has two columns on both paths that reach the writer -- the WCM
    template's own table, and a fallback table the header guard rejects
    unless it has at least two columns."""
    tree = ast.parse(Path(memberships_module.__file__).read_text(encoding="utf-8"))
    lengths = []
    for node in ast.walk(tree):
        if (isinstance(node, ast.Call)
                and isinstance(node.func, ast.Attribute)
                and node.func.attr == "_add_table_row"):
            data_arg = node.args[1]
            assert isinstance(data_arg, ast.List), \
                f"line {node.lineno}: data is not a list literal, length is not static"
            lengths.append(len(data_arg.elts))
    assert lengths == [2, 2], lengths


# --- (f) review item 3: one date formatting path ----------------------------

def test_both_membership_paths_format_the_same_dates_identically():
    """The single path normalized dates through `format_date_range` while the
    multi path wrote whatever `_parse_multi_membership_entry` returned, so
    the same membership rendered "1/2002-present" or "2002-Present" purely
    according to which branch its entry took."""
    multi = {
        "text": ("Fellow | American Society of Anesthesiologists | 1/2002-present | "
                 "Member | National Board of Physicians | 3/2018-present"),
        "extracted_fields": {"organization": "American Society of Anesthesiologists"},
    }
    single = {
        "text": "Fellow, American Society of Anesthesiologists",
        "extracted_fields": {
            "organization": "American Society of Anesthesiologists",
            "membership_type": "Fellow",
            "start_date": "1/2002",
            "end_date": "present",
        },
    }
    multi_rows = _render_memberships([multi])
    single_rows = _render_memberships([single])
    assert len(multi_rows) == 2
    assert multi_rows[0] == single_rows[0], (multi_rows, single_rows)
    assert multi_rows[0][1] == "2002-Present"
    assert multi_rows[1][1] == "2018-Present"


# --- (g) review item 4: '|' is only a boundary when the parts parse ---------

def test_entry_parts_keeps_a_pipe_inside_one_field_unsplit():
    """A '|' surviving extraction inside one organization or source field is
    punctuation, not a membership boundary. The parts here supply no date and
    no membership type, so the split is refused and the entry keeps the
    single opaque part it had before #476 -- rather than becoming two
    organizations."""
    case = "Society of Critical Care Medicine | Anesthesia and Analgesia section"
    assert _entry_parts(case) == [case]


def test_entry_parts_still_splits_when_the_parts_supply_a_full_record():
    case = "Fellow | American Academy of Pediatrics | 1/1997-present"
    assert _entry_parts(case) == [
        "Fellow", "American Academy of Pediatrics", "1/1997-present"]


def test_three_organizations_joined_by_pipes_are_not_split_into_three_rows():
    """The rendering consequence of the gate above. Three organization-shaped
    parts and nothing else is not a membership record list, so the entry
    stays one row instead of the three it would have produced."""
    entry = {
        "text": ("American Thoracic Society | European Respiratory Society | "
                 "Society of Critical Care Medicine"),
        "extracted_fields": {"organization": "American Thoracic Society"},
    }
    rows = _render_memberships([entry])
    assert len(rows) == 1, rows
    assert rows[0][0] == "American Thoracic Society"


# --- (h) review item 5: membership type matching is token-aware -------------

def test_membership_type_is_not_suppressed_by_a_bare_substring_match():
    """"associate" is a substring of "Associated", so the substring test
    dropped the membership type from an organization that merely contains
    the same letters inside a longer word."""
    entry = {
        "text": "Associate, Associated Medical Schools of New York",
        "extracted_fields": {
            "organization": "Associated Medical Schools of New York",
            "membership_type": "Associate",
            "start_date": "2019",
            "end_date": "Present",
        },
    }
    rows = _render_memberships([entry])
    assert rows[0][0] == "Associate, Associated Medical Schools of New York"


def test_membership_type_already_named_in_the_organization_is_still_suppressed():
    """The behaviour the module docstring promises: "Fellow, American College
    of Surgeons" must not become "Fellow, Fellow of the American College of
    Surgeons"."""
    assert _type_already_named("Fellow", "Fellow of the American College of Surgeons")
    assert _type_already_named("Elected Member", "Elected Member, Institute of Medicine")
    assert not _type_already_named("Associate", "Associated Medical Schools of New York")
    assert not _type_already_named("Member", "Remembrance Society")


# --- (i) review item 6: header matching is case/whitespace normalized -------

def test_fallback_table_date_header_is_matched_case_and_whitespace_folded():
    """The fallback required a literal "Date" in the second header cell, so a
    template or source variation spelling it "DATES" had its valid
    memberships table rejected and rendered nothing at all."""
    entry = {
        "text": "Fellow, American College of Surgeons",
        "extracted_fields": {
            "organization": "American College of Surgeons",
            "membership_type": "Fellow",
            "start_date": "2015",
            "end_date": "Present",
        },
    }
    for header in (("Organization", "DATES"),
                   ("Organization", "  date   awarded "),
                   ("Organization", "Date (yyyy-yyyy)")):
        gen, rows = _fill([entry], header=header, table_before_heading=True)
        assert gen.stats["tables_populated"] == 1, header
        assert rows == [["Fellow, American College of Surgeons", "2015-Present"]], header


def test_fallback_table_without_a_date_header_is_still_rejected():
    """Normalizing must not widen the guard into accepting another section's
    table: "Organization" alone matches several, which is why the second
    header is checked at all. "Update" contains the letters d-a-t-e and is
    still not a date column."""
    entry = {"text": "Fellow, American College of Surgeons", "extracted_fields": {}}
    for header in (("Organization", "Certificate #"), ("Organization", "Update")):
        gen, rows = _fill([entry], header=header, table_before_heading=True)
        assert gen.stats["tables_populated"] == 0, header
        assert rows == [], header


# --- (j) review item 7: a dedicated organization fallback -------------------

def test_missing_organization_is_recovered_without_the_raw_text_blob():
    """`original_text[:150]` put the membership type, the dates and any
    trailing source text into the organization cell, cut mid-word. The
    dedicated extractor takes the first fragment that is neither a date nor a
    bare membership type, lifts the trailing dates into the date column, and
    lets the normal path prefix the type."""
    entry = {
        "text": "Elected Member\tAmerican Association for Thoracic Surgery\t2005 - Present",
        "extracted_fields": {},
    }
    gen, rows = _fill([entry])
    assert rows == [
        ["Elected Member, American Association for Thoracic Surgery", "2005-Present"]], rows
    assert "\t" not in rows[0][0]


def test_organization_fallback_is_counted_and_not_silent():
    entry = {
        "text": "American Association for Thoracic Surgery\t2005 - Present",
        "extracted_fields": {},
    }
    gen, rows = _fill([entry])
    assert gen.stats["membership_organization_fallbacks"] == 1
    assert rows == [["American Association for Thoracic Surgery", "2005-Present"]], rows


def test_organization_fallback_prefers_a_structured_alias_field():
    """Two farm uids carry the organization under `institution` with
    `organization` absent; reading it beats parsing the raw text, which here
    would have kept the tab and the date in the organization cell."""
    recovered = _organization_fallback(
        "2016-Present\tAmerican College of Physicians",
        {"institution": "American College of Physicians"},
    )
    assert recovered.organization == "American College of Physicians"
    assert recovered.membership_type == ""


def test_organization_fallback_never_swallows_the_date_only_first_fragment():
    """A flattened two-column source row arrives as "2022\\t\\t\\tSociety...":
    taking fragment zero blindly would render the year as the organization
    and drop the society, which is what a naive first-fragment rule does."""
    recovered = _organization_fallback(
        "2022\t\t\tAmerican Association for Thoracic Surgery", {})
    assert recovered.organization == "American Association for Thoracic Surgery"
    assert recovered.dates == "2022"


def test_organization_fallback_truncates_on_a_word_boundary():
    long_name = "American " + ("Society " * 30) + "of Medicine"
    recovered = _organization_fallback(long_name, {})
    assert len(recovered.organization) <= 150
    assert not recovered.organization.endswith(" ")
    assert long_name.startswith(recovered.organization)
    assert recovered.organization.split()[-1] == "Society"


# --- (k) review thread 2, item 1: three or more memberships -----------------

def test_three_pipe_separated_memberships_render_exactly_three_rows():
    """The reviewer's own example, verbatim. Two recovered memberships do not
    prove the parser handles an arbitrary number of groups: a regression that
    consumed only the first two would still pass the two-membership control."""
    entry = {
        "text": ("Member | Org A | 2010-present |\n"
                 "Fellow | Org B | 2015-present |\n"
                 "Board Member | Org C | 2020-present"),
        "extracted_fields": {"organization": "Org A"},
    }
    rows = _render_memberships([entry])
    assert len(rows) == 3, rows
    assert [r[0] for r in rows] == ["Member, Org A", "Fellow, Org B", "Board Member, Org C"]
    assert [r[1] for r in rows] == ["2010-Present", "2015-Present", "2020-Present"]


def test_three_memberships_fused_into_one_blind_line_render_three_rows():
    """The same three groups with every newline gone -- the shape this fix
    exists for, and the one where the count could silently cap at two.

    Long organization names here; the short-acronym blind form is the next
    test (#758)."""
    entry = {
        "text": ("Member | State Medical Society | 2010-present | "
                 "Fellow | National Surgical Association | 2015-present | "
                 "Board Member | American Board of Pediatrics | 2020-present"),
        "extracted_fields": {"organization": "State Medical Society"},
    }
    rows = _render_memberships([entry])
    assert len(rows) == 3, rows
    assert [r[0] for r in rows] == [
        "Member, State Medical Society",
        "Fellow, National Surgical Association",
        "Board Member, American Board of Pediatrics",
    ]
    assert [r[1] for r in rows] == ["2010-Present", "2015-Present", "2020-Present"]


def test_three_memberships_with_short_acronym_orgs_on_one_blind_line():
    """#758: the blind form of the reviewer's three-group example with
    organizations of five characters or fewer, all of which used to be dropped."""
    entry = {
        "text": ("Member | AMA | 2010-present | "
                 "Fellow | ASCO | 2015-present | "
                 "Board Member | NIH | 2020-present"),
        "extracted_fields": {"organization": "AMA"},
    }
    rows = _render_memberships([entry])
    assert [r[0] for r in rows] == ["Member, AMA", "Fellow, ASCO", "Board Member, NIH"]
    assert [r[1] for r in rows] == ["2010-Present", "2015-Present", "2020-Present"]


# --- (l) review thread 2, item 2: malformed pipe structure ------------------

def test_malformed_pipe_structure_neither_raises_nor_fabricates_a_membership():
    """An incomplete trailing group -- a type and an organization with no
    date -- is a realistic extraction artefact, and this PR changed which
    malformed inputs reach `_parse_multi_membership_entry`.

    Exactly two memberships render: both are genuinely present in the text,
    and the second's date cell is left EMPTY rather than borrowed from the
    first. Inventing that date, or emitting a third row for the leftover
    fragments, is what fabrication would look like here."""
    entry = {
        "text": "Member | State Medical Society | 2015-present | Fellow | National Association",
        "extracted_fields": {"organization": "State Medical Society"},
    }
    rows = _render_memberships([entry])
    assert len(rows) == 2, rows
    assert rows[0] == ["Member, State Medical Society", "2015-Present"]
    assert rows[1] == ["Fellow, National Association", ""]


# --- (m) review thread 2, item 3: the same case with no extracted fields ----

def test_three_field_single_membership_with_no_extracted_fields():
    """The three-field control above can pass without exercising the changed
    parsing at all, because the renderer reads `extracted_fields` on the
    single-membership path. With no structured fields the entry can only be
    rendered from the text this PR re-segments, and it must still resolve to
    exactly ONE membership."""
    entry = {
        "text": "Fellow | American Academy of Pediatrics | 1/1997-present",
        "extracted_fields": {},
    }
    gen, rows = _fill([entry])
    assert len(rows) == 1, rows
    assert rows[0] == ["Fellow, American Academy of Pediatrics", "1997-Present"]
    assert gen.stats["entries_inserted"] == 1


# --- (n) review thread 2, item 4: _entry_parts directly ---------------------

def test_entry_parts_splits_the_blind_pipe_case():
    """A focused assertion on the new boundary detection, so a failure here
    localizes to segmentation rather than to downstream parsing or Word
    rendering; the rendering regressions above stay as integration cover."""
    case = ("Member | State Medical Society | 2015-present | "
            "Fellow | National Surgical Association | 2018-present")
    assert _entry_parts(case) == [
        "Member", "State Medical Society", "2015-present",
        "Fellow", "National Surgical Association", "2018-present",
    ]


# --- (o) review thread 2, item 5: empty pipe fragments ----------------------

def test_empty_pipe_fragments_make_no_empty_records_and_do_not_shift_pairing():
    """Repeated separators and empty columns are ordinary extraction output.
    Since '|' is now a parsing boundary, an empty fragment must not become an
    empty membership or push an organization onto the wrong date."""
    case = ("Member || State Medical Society | 2015-present ||| "
            "Fellow |  | National Surgical Association | 2018-present |")
    assert "" not in _entry_parts(case)
    rows = _render_memberships([{"text": case, "extracted_fields": {}}])
    assert rows == [
        ["Member, State Medical Society", "2015-Present"],
        ["Fellow, National Surgical Association", "2018-Present"],
    ], rows


# --- (p) review thread 2, item 6: rendered ordering -------------------------

def test_memberships_render_in_reverse_chronological_order():
    """`_fill_memberships` sorts before writing, which none of the
    `_entry_parts` assertions can see: they stop at segmentation."""
    older = {
        "text": "Member, American Thoracic Society",
        "extracted_fields": {
            "organization": "American Thoracic Society",
            "membership_type": "Member",
            "start_date": "2005", "end_date": "2010",
        },
    }
    newer = {
        "text": "Fellow, American College of Surgeons",
        "extracted_fields": {
            "organization": "American College of Surgeons",
            "membership_type": "Fellow",
            "start_date": "2015", "end_date": "2020",
        },
    }
    rows = _render_memberships([older, newer])
    assert [r[0] for r in rows] == [
        "Fellow, American College of Surgeons",
        "Member, American Thoracic Society",
    ], rows
    assert [r[1] for r in rows] == ["2015-2020", "2005-2010"]


# --- (q) review thread 2, item 7: tabs stay unsplit alongside a pipe --------

def test_tabs_stay_unsplit_in_an_entry_that_also_carries_a_pipe():
    """The module contract is that tab is NOT a boundary here -- a tab-split
    membership loses its organization to the classifier (see this file's
    docstring). Pinning a mixed tab-and-pipe entry stops a later refactor
    from collapsing the conditional into an unconditional `entry_fragments()`
    call, which would split both."""
    case = "Member\tsince 2010 | American Thoracic Society | 2010-present"
    parts = _entry_parts(case)
    assert parts == ["Member\tsince 2010", "American Thoracic Society", "2010-present"]
    assert "\t" in parts[0]
    assert parts != entry_fragments(case), \
        "an unconditional entry_fragments() call would also split the tab"


# --- (r) review thread 2, item 8: empty and separator-only entries ----------

@pytest.mark.parametrize("text", ["", "   ", "\n", "\n\n  \n", "\t", "\t\t", "|", " | | "])
def test_blank_and_separator_only_entries_render_no_rows(text):
    """Defined behaviour for a source record with no content: no row at all.
    A blank row in the WCM table is a visible defect, and the entry has
    nothing to lose by being dropped -- the drop is counted so the run can
    still say it happened."""
    gen, rows = _fill([{"text": text, "extracted_fields": {}}])
    assert rows == [], (text, rows)
    assert gen.stats["entries_inserted"] == 0
    assert gen.stats["membership_entries_blank"] == 1


def test_a_blank_entry_does_not_suppress_a_real_one():
    real = {
        "text": "Fellow, American College of Surgeons",
        "extracted_fields": {
            "organization": "American College of Surgeons",
            "membership_type": "Fellow",
            "start_date": "2015", "end_date": "Present",
        },
    }
    gen, rows = _fill([{"text": "  ", "extracted_fields": {}}, real])
    assert rows == [["Fellow, American College of Surgeons", "2015-Present"]], rows
    assert gen.stats["entries_inserted"] == 1


# --- (s) the render gate's regression: a recovered cell that lost characters -

# The organization cell the raw-text fallback rendered before the dedicated
# extractor existed. Every shape below is checked against it: the recovery is
# allowed to move a date out of this cell, never to shorten it into nothing.
_RAW_TEXT_FALLBACK_MAX = 150


def _words(text):
    return re.findall(r"[0-9a-z]+", text.lower())


def _is_shortened_prefix(cell, text):
    """True when `cell` is `text` cut short -- the defect shape, before the
    date carve-out is applied."""
    return cell != text and text.startswith(cell)


def _dropped_words_are_all_in_the_date_cell(cell, text, date_cell):
    """The one shortening the ticket allows: the tail the organization cell
    gave up is a date, and the date column now carries it."""
    tail = text[len(cell):]
    date_words = set(_words(date_cell))
    return all(word in date_words for word in _words(tail))


_MONTH_WORDS = {
    "january", "february", "march", "april", "may", "june", "july", "august",
    "september", "october", "november", "december",
    "jan", "feb", "mar", "apr", "jun", "jul", "aug", "sep", "sept", "oct",
    "nov", "dec",
}


def _is_a_date_word(word):
    return word.isdigit() or word in _MONTH_WORDS or word in {"present", "pres"}


def _cell_still_names_something(cell):
    """A shortened cell has to be left holding an actual name. "February",
    all that survived the render gate's regression, is a piece of the date it
    was cut away from."""
    return any(not _is_a_date_word(word) for word in _words(cell))


# (raw text, extracted_fields) shapes taken from the farm's section I, with
# organization names generalized. Each one reaches the fallback because no
# `organization` field is extracted.
NO_ORGANIZATION_FIELD_SHAPES = [
    # The render gate's regression: the whole value is a date, and the date
    # column already carries it from structured fields.
    ("February 2018 – Present",
     {"organization": None, "start_date": "February 2018", "end_date": "Present"}),
    ("July 2023 – Present",
     {"organization": None, "start_date": "2023-07", "end_date": "Present"}),
    # A flattened two-column row whose date half is written in a shape the
    # date test does not read, so the date half looks like an organization.
    ("2003-             \tAmerican Association for Thoracic Surgery, Associate Member", {}),
    ("2018-Pres\t\tMember - American College of Physicians", {}),
    ("2008-12\t\tMember, American Thoracic Society", {}),
    ("2004-08, 17, 2023\tMember, American Public Health Association", {}),
    # A tab inside the organization name itself.
    ("Psi Chi\t, National Honors Society in Psychology\t\t\t1995", {}),
    # A date inside a bracketed aside.
    ("American Academy of Forensic Sciences (elected Fellow, 1980)", {}),
    # The shape the recovery is FOR: the date is its own trailing field.
    ("American College of Physicians, Member\t\t\t2024-Present", {}),
    ("2022\t\t\tAmerican Association for Thoracic Surgery", {}),
]


@pytest.mark.parametrize("text,fields", NO_ORGANIZATION_FIELD_SHAPES)
def test_no_recovered_organization_cell_is_a_shortened_prefix(text, fields):
    """The regression the 66-CV render gate caught on this PR.

    The round-1 organization recovery lifted a trailing date off a fragment
    by matching the year alone, so an entry whose whole text is a date left
    the month behind: "February 2018 - Present" rendered an organization cell
    reading "February", strictly shorter than the cell the section rendered
    before, and exactly the raw-text-truncation failure the recovery was
    written to stop.

    The rule this pins is the general one, not the single cell: a recovered
    organization may give up a tail only when the date column now carries
    every word of it."""
    _, rows = _fill([{"text": text, "extracted_fields": fields}])
    assert len(rows) == 1, rows
    cell, date_cell = rows[0]
    if _is_shortened_prefix(cell, text):
        assert _dropped_words_are_all_in_the_date_cell(cell, text, date_cell), (
            f"organization cell lost {text[len(cell):]!r}, which the date cell "
            f"{date_cell!r} does not carry")
        assert _cell_still_names_something(cell), (
            f"organization cell was shortened to {cell!r}, which is nothing "
            f"but a piece of the date")


def test_the_render_gate_regression_cell_verbatim():
    """The exact pair of cells the gate read, pinned as values rather than as
    an invariant: an entry whose whole text is a month-name date, with the
    date column already filled from structured fields."""
    entry = {
        "text": "February 2018 – Present",
        "extracted_fields": {
            "organization": None, "start_date": "February 2018", "end_date": "Present",
        },
    }
    _, rows = _fill([entry])
    assert rows == [["February 2018 – Present", "2018-Present"]], rows


def test_a_written_out_month_belongs_to_the_date_not_the_organization():
    for text in ("February 2018 – Present", "July 2023 – Present",
                 "Sept. 2004 - Present", "May 2001"):
        remainder, dates = _split_trailing_dates(text)
        assert remainder == "", (text, remainder)
        assert dates, text


def test_a_month_word_is_only_a_date_when_a_year_follows_it():
    """The month is read as part of the date only immediately in front of a
    year, so an organization whose name merely starts with those letters keeps
    its name."""
    assert _split_trailing_dates("Marching Band Alumni 2018") == (
        "Marching Band Alumni", "2018")
    assert _split_trailing_dates("Society of May") == ("Society of May", "")


def test_recovery_that_would_drop_a_word_renders_the_raw_text_instead():
    """A fragment split that mislays part of the entry is refused outright:
    the cell goes back to the raw text, which is what the section rendered
    before the recovery existed. Better a blob than a missing organization."""
    text = "2003-             	American Association for Thoracic Surgery, Associate Member"
    recovered = _organization_fallback(text, {})
    assert recovered.organization == text.strip()
    assert recovered.dates == ""


def test_a_date_move_that_loses_nothing_is_still_made():
    """The guard must not cost the recovery its point: where every word of the
    entry still lands somewhere, the date leaves the organization cell."""
    recovered = _organization_fallback(
        "American College of Physicians, Member			2024-Present", {})
    assert recovered.organization == "American College of Physicians, Member"
    assert recovered.dates == "2024-Present"


def test_a_lift_that_would_unbalance_a_bracket_is_refused():
    """Lifting a date out of a bracketed aside left the organization holding
    an opening bracket it never closes."""
    text = "American Academy of Forensic Sciences (elected Fellow, 1980)"
    assert _split_trailing_dates(text) == (text, "")
    assert _organization_fallback(text, {}).organization == text


# --- (t) round 3: a structured alias must not silence the field beside it ---

# 2082_Dr_Scot's two section I entries verbatim in shape: stage 4 extracted
# them under another taxonomy code's field schema (course_code/institution/
# role, stage4/schemas.py), so the organization arrives under `institution`
# and the membership role under `role` -- section I's own `organization` and
# `membership_type` keys are both absent.
ALIAS_AND_ROLE_ENTRIES = [
    {
        "text": "2016-present\tFaculty Associate. Center for Southeast Asian Studies",
        "extracted_fields": {
            "course_code": None, "course_title": None,
            "institution": "Center for Southeast Asian Studies",
            "role": "Faculty Associate", "enrollment": None,
            "start_date": "2016", "end_date": "present",
        },
    },
    {
        "text": "2012-2020\tFaculty Associate-Center for Nonprofit and NGO Studies",
        "extracted_fields": {
            "course_code": None, "course_title": None,
            "institution": "Center for Nonprofit and NGO Studies",
            "role": "Faculty Associate", "enrollment": None,
            "start_date": "2012", "end_date": "2020",
        },
    },
]


def test_a_role_beside_an_institution_alias_still_reaches_the_organization_cell():
    """The alias branch resolved organization AND membership type together:
    `institution` named the organization, the branch returned on it, and the
    `role` sitting beside it was never read -- so "Faculty Associate" was
    rendered nowhere in the document. The date correctly moving to the date
    column is the improvement; losing the role with it was not."""
    rows = _render_memberships(ALIAS_AND_ROLE_ENTRIES)
    assert rows == [
        ["Faculty Associate, Center for Southeast Asian Studies", "2016-Present"],
        ["Faculty Associate, Center for Nonprofit and NGO Studies", "2012-2020"],
    ], rows


def test_the_membership_type_alias_is_read_independently_of_the_organization_one():
    """Each value prefers its own structured alias; neither short-circuits the
    other. `fellowship_designation` is section I's own second designation
    field and is read the same way."""
    for type_key in ("role", "fellowship_designation", "membership_type"):
        recovered = _organization_fallback(
            "2016-present\tFaculty Associate. Center for Southeast Asian Studies",
            {"institution": "Center for Southeast Asian Studies",
             type_key: "Faculty Associate"},
        )
        assert recovered.membership_type == "Faculty Associate", type_key
        assert recovered.organization == "Center for Southeast Asian Studies", type_key


def test_a_work_title_is_not_read_as_a_membership_type():
    """`title` names a work -- a paper, a talk, a patent -- on every schema
    that emits it, and no schema emits `position` at all, so neither is a
    membership designation. Reading `title` here would prefix a publication
    name onto an organization cell."""
    recovered = _organization_fallback(
        "American College of Physicians",
        {"institution": "American College of Physicians",
         "title": "A Randomized Trial of Something", "position": "Chair"},
    )
    assert recovered.membership_type == ""
    assert recovered.organization == "American College of Physicians"


def test_an_alias_naming_less_than_the_entry_is_refused_like_any_other_recovery():
    """The no-loss check covers the alias branch too. An alias names ONE
    field, so it can be right about the organization and still leave the rest
    of the entry unaccounted for; when it does, the row goes back to the raw
    text rather than dropping the remainder."""
    text = "Elected Fellow\tAmerican College of Physicians, Cardiology Council"
    recovered = _organization_fallback(
        text, {"institution": "American College of Physicians"})
    assert recovered.organization == text.strip()
    assert recovered.membership_type == ""


def test_the_entrys_own_date_fields_count_as_covered_by_the_date_column():
    """The caller renders `start_date`/`end_date` into the date cell, so the
    no-loss check has to count them -- otherwise an alias-supplied row whose
    dates live in the fields is refused for "losing" words the row prints."""
    recovered = _organization_fallback(
        "2016 - present American College of Physicians",
        {"institution": "American College of Physicians",
         "start_date": "2016", "end_date": "present"},
    )
    assert recovered.organization == "American College of Physicians"


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))


@pytest.mark.parametrize("trailer", ["NY", "USA", "(AHA)"])
def test_short_line_after_an_organization_does_not_become_a_membership(trailer):
    """#758 follow-up: a short line trailing an organization (a state, a
    country, the organization's own acronym) is not a second organization.
    Counting it shifted the next membership's type and date onto it."""
    entry = {
        "text": (f"Member\nAmerican Heart Association\n{trailer}\n2010-present\n"
                 "Fellow\nAmerican College of Cardiology\n2012-present"),
        "extracted_fields": {},
    }
    rows = _render_memberships([entry])
    assert rows == [["Member, American Heart Association", "2010-Present"],
                    ["Fellow, American College of Cardiology", "2012-Present"]]


def test_short_organization_right_after_a_type_is_kept_on_the_multi_line_path():
    entry = {"text": "Member\nAMA\n2010-present\nFellow\nACP\n2012-present",
             "extracted_fields": {}}
    rows = _render_memberships([entry])
    assert rows == [["Member, AMA", "2010-Present"], ["Fellow, ACP", "2012-Present"]]
