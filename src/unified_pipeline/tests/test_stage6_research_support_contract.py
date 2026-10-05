"""The M2 research-support contract: bucket rules, field rules, rendered rows.

Companion to `test_stage6_grant_number_only_sparsity.py`, which pins the
sparsity guard alone. This file pins the rest of section M2 -- the twenty gaps
named in the review of PR #734 (thread 3932451691) plus the explicit rendering
contract asked for in thread 3932312407 item 8.

Two levels, deliberately:

* the module-level classifiers (`filter_role_effort_headers`,
  `match_effort_for_title`, `apply_effort_to_grants`, `rebucket_grants_by_status`,
  `reclassify_past_m2a_grants`, `resolve_pi_name`, `normalize_percent_effort`)
  take plain data and return plain data, so a bucket or effort decision is
  pinned without building a DOCX;
* every claim about what a reader of the finished CV actually sees goes through
  `_create_grant_table` / `_fill_research_support` and is read back off the
  generated document -- never off a stdout line. A log message is not the
  behaviour; the document is.

The generator fixture is the same minimal `__new__` shell the sparsity file
uses: a blank `docx.Document` plus the three WCM funding headers as plain
paragraphs, which is all `_fill_research_support` looks for. No template file,
no DB, no LLM.

Every person named below is invented -- "Ada Testowner" for the CV owner,
"Jane Smith" for an extracted PI. This module's own subject is that a grant's
values are CV content, so its fixtures do not get to be an exception: a real
researcher's name checked into a test is the same disclosure whether or not
their grants came with it.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_research_support_contract.py -p no:cacheprovider
"""

import ast
import logging
import sys
from pathlib import Path

import docx
import pytest
from docx.table import Table

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage6.sections import research_support  # noqa: E402
from unified_pipeline.stage6.sections.research_support import (  # noqa: E402
    ROLE_EFFORT_DIGIT_RE,
    ROLE_EFFORT_FIELD_ROW_RE,
    ROLE_EFFORT_HEADER_RE,
    ROLE_EFFORT_HEADER_WINDOW,
    UnsupportedRebucketTargetError,
    apply_effort_to_grants,
    claim_goal_rows,
    fill_major_goals_from_text,
    filter_role_effort_headers,
    grant_end_year,
    is_role_effort_header,
    match_effort_for_title,
    normalize_percent_effort,
    parse_major_goals,
    promote_open_ended_m2b_grants,
    rebucket_grants_by_status,
    reclassify_past_m2a_grants,
    resolve_pi_name,
    _is_cv_owner,
    _pi_name_from_label,
)
from unified_pipeline.stage6.normalization import (  # noqa: E402
    grant_heading_rebucket_target,
    grant_status_is_empty_section_label,
    grant_status_rebucket_target,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402

CURRENT = 'Current Research Funding'
COMPLETED = 'Past (Completed) Funding'
PENDING = 'Pending Funding'
HEADERS = (CURRENT, COMPLETED, PENDING)

# Every date rule below is judged against this year, passed in rather than read
# off the clock, so the boundary cases stay true after 2026 (thread 3932312407
# item 7).
TEST_YEAR = 2026


def _generator(emit_comments=False):
    """A generator with just enough state to build one table or fill one section."""
    gen = WCMTemplateGenerator.__new__(WCMTemplateGenerator)
    gen._bullet_paras = set()
    gen.doc = docx.Document()
    gen.verbose = False
    gen.stats = {'tables_populated': 0, 'entries_inserted': 0, 'comments_added': 0}
    gen.emit_comments = emit_comments
    gen._comments = []
    gen._comment_id = 1
    gen._overflow_entries = []
    gen._declined_grant_entries = []  # #839: _create_grant_table's decline path appends here
    return gen


def _sectioned_generator(emit_comments=False):
    """A generator whose document already carries the three WCM funding headers."""
    gen = _generator(emit_comments=emit_comments)
    for header in HEADERS:
        gen.doc.add_paragraph().add_run(header).bold = True
    return gen


def _entry(code='M2A', text='grant', **fields):
    return {'text': text, 'taxonomy_code': code, 'extracted_fields': dict(fields)}


def _cells(table):
    return {row.cells[0].text: row.cells[1].text for row in table.rows}


def _rows(table):
    return [(row.cells[0].text, row.cells[1].text) for row in table.rows]


def _header_positions(gen):
    """Body index of each WCM funding header paragraph."""
    body = list(gen.doc.element.body)
    positions = {}
    for para in gen.doc.paragraphs:
        text = para.text.strip()
        if text in HEADERS and text not in positions:
            positions[text] = body.index(para._element)
    return positions


def _elements_under(gen, header):
    """Body elements between one funding header and the next one."""
    body = list(gen.doc.element.body)
    positions = _header_positions(gen)
    start = positions[header]
    later = [pos for pos in positions.values() if pos > start]
    end = min(later) if later else len(body)
    return body[start + 1:end]


def _tables_under(gen, header):
    return [Table(el, gen.doc) for el in _elements_under(gen, header)
            if el.tag.endswith('}tbl')]


def _tags_under(gen, header):
    return [el.tag.split('}')[-1] for el in _elements_under(gen, header)]


def _titles_under(gen, header):
    return [_cells(table).get('Project title:') for table in _tables_under(gen, header)]


# --- item 1: status-based rebucketing decides the bucket ------------------------

def test_under_review_status_moves_an_m2a_grant_to_pending():
    """An M2A grant whose status reads "Under review" renders under Pending.

    `grant_status_rebucket_target` was called on every grant and its answer was
    never asserted anywhere -- only the stdout line was. Pins the resulting
    bucket instead (review thread 3932451691 item 1).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Under Review Project', agency='NIH',
                        status='Under review', start_date='01/2026')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == []
    assert _titles_under(gen, COMPLETED) == []
    assert _titles_under(gen, PENDING) == ['Under Review Project']


def test_under_review_status_moves_an_m2b_grant_to_pending():
    """The rebucketer reads M2B as a source too, not only M2A.

    Same gap as above from the other side: a completed-funding record whose own
    status says it is still under review belongs under Pending (review thread
    3932451691 item 1).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2B': [_entry('M2B', title='Resubmitted Project', agency='NIH',
                        status='Under review', start_date='01/2020')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, COMPLETED) == []
    assert _titles_under(gen, PENDING) == ['Resubmitted Project']


# --- item 2: a past end date moves the grant in the DOCUMENT --------------------

def test_past_m2a_grant_is_rendered_under_completed_funding():
    """The past-end-date rule is pinned on the document, not the stdout line.

    The existing verbose test asserts "Reclassified to M2B..." reaches stdout,
    which a refactor could keep printing while rendering the grant in the wrong
    place. This asserts where the table lands (review thread 3932451691 item 2).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Ended Project Study', agency='NIH',
                        start_date='01/2018', end_date='06/2020')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == []
    assert _titles_under(gen, COMPLETED) == ['Ended Project Study']


# --- item 3: open-ended end dates are never reclassified ------------------------

@pytest.mark.parametrize('end_date', ['Present', 'current', 'ONGOING', '', 'TBD'])
def test_current_end_date_does_not_move_m2a_grant(end_date):
    """"Present"/"current"/"ongoing"/"" keep a grant in Current Research Funding.

    Without a test, a refactor of the date parser could silently file every
    active grant under Past (Completed) Funding (review thread 3932451691
    item 3).

    What this pins is the outcome, not one branch: an open-ended end date is
    held out of the reclassifier twice over -- by the `OPEN_ENDED_END_DATES`
    membership test and, since none of those four literals contains a
    four-digit year, by the `\\d{4}` search behind it. No input can tell the
    two apart while the literal list stays digit-free, so removing either one
    leaves this test green. "TBD" is the second guard on its own: not in the
    literal list, still no year to parse.
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Ongoing Project Study', agency='NIH',
                        start_date='01/2019', end_date=end_date)]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Ongoing Project Study']
    assert _titles_under(gen, COMPLETED) == []


# --- item 4: a future end date stays current ------------------------------------

def test_future_m2a_end_date_remains_current():
    """The other side of the boundary: an end date after `current_year` stays M2A.

    Judged against the injected year, so the test does not start failing when
    the wall clock passes 2030 (review thread 3932451691 item 4).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Future Project Study', agency='NIH',
                        start_date='01/2026', end_date='2030-12')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Future Project Study']
    assert _titles_under(gen, COMPLETED) == []


def test_end_date_in_the_current_year_remains_current():
    """The comparison is `end_year < current_year`, so this year is still current."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='This Year Project', agency='NIH',
                        start_date='01/2026', end_date=f'06/{TEST_YEAR}')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['This Year Project']


# --- item 5: grant_number in Award Source, never twice --------------------------

def test_grant_number_is_appended_to_award_source_when_absent():
    """A grant number nowhere else in the record is appended to Award Source.

    The WCM block has no grant-number row, so the identifier rides in Award
    Source; production-visible and previously untested (review thread
    3932451691 item 5).
    """
    fields = {'agency': 'NIH', 'title': 'Cancer Immunology Project',
              'grant_number': 'R01 CA123456'}
    table = _generator()._create_grant_table(fields, 'M2A')
    assert _cells(table)['Award Source:'] == 'NIH (R01 CA123456)'


def test_a_pending_grant_renders_its_requested_amount_and_submission_date():
    """#1299: M2C's schema names the amount `total_funding_requested` and the
    date `submission_date`. No row read either, so 13 of 14 pending-grant
    amounts in one CV rendered nowhere. A request is not an award, so the
    amount does not go under "Total award:"."""
    fields = {'agency': 'NIH', 'title': 'Pending Kestrel Project',
              'total_funding_requested': '250000', 'submission_date': '2025-06'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2C'))
    assert cells[research_support.TOTAL_REQUESTED_LABEL] == '$250,000'
    assert cells[research_support.SUBMISSION_DATE_LABEL]
    assert research_support.TOTAL_AWARD_LABEL not in cells


def test_a_grant_without_the_pending_fields_gets_neither_row():
    cells = _cells(_generator()._create_grant_table(
        {'agency': 'NIH', 'title': 'Pending Kestrel Project'}, 'M2C'))
    assert research_support.TOTAL_REQUESTED_LABEL not in cells
    assert research_support.SUBMISSION_DATE_LABEL not in cells


@pytest.mark.parametrize('code', ['M2A', 'M2B', 'M2C'])
def test_the_owners_share_of_the_award_gets_its_own_row(code):
    """#817 (EBYSBC E14): stage 4 keeps the owner's part of an award under the
    off-schema `share_total`, which no row read. It renders under its own
    label, never as the annual direct costs."""
    fields = {'agency': 'NIH', 'title': 'Shared Kestrel Project',
              'total_funding': '900000', 'share_total': '$120,500'}
    cells = _cells(_generator()._create_grant_table(fields, code))
    assert cells[research_support.SHARE_LABEL] == '$120,500'
    assert cells[research_support.TOTAL_AWARD_LABEL] == '$900,000'
    assert research_support.ANNUAL_COSTS_LABEL not in cells


@pytest.mark.parametrize('share', [None, '', '   ', False])
def test_a_grant_without_a_share_gets_no_share_row(share):
    cells = _cells(_generator()._create_grant_table(
        {'agency': 'NIH', 'title': 'Shared Kestrel Project', 'share_total': share}, 'M2C'))
    assert research_support.SHARE_LABEL not in cells


def test_grant_number_alone_becomes_the_award_source_when_there_is_no_agency():
    """With no agency at all the identifier is the whole Award Source value."""
    fields = {'grant_number': 'R01 CA123456', 'title': 'Cancer Immunology Project'}
    table = _generator()._create_grant_table(fields, 'M2A')
    assert _cells(table)['Award Source:'] == 'R01 CA123456'


def test_grant_number_already_in_the_agency_is_not_duplicated():
    """`_create_grant_table`'s Award Source guard: a number already inside the
    agency text is not re-appended."""
    fields = {'agency': 'NIH R01 CA123456', 'title': 'Cancer Immunology Project',
              'grant_number': 'R01 CA123456'}
    table = _generator()._create_grant_table(fields, 'M2A')
    assert _cells(table)['Award Source:'] == 'NIH R01 CA123456'


def test_grant_number_already_in_the_title_is_not_duplicated():
    """The same guard reads the title too, and is case-folded, not case-sensitive."""
    fields = {'agency': 'NIH', 'title': 'Cancer Immunology Project (r01 ca123456)',
              'grant_number': 'R01 CA123456'}
    table = _generator()._create_grant_table(fields, 'M2A')
    assert _cells(table)['Award Source:'] == 'NIH'


# --- #839: a decline appends to _declined_grant_entries -------------------------

def test_declined_sparse_entry_is_appended_to_declined_grant_entries():
    """No title, no substantive info, no grant-number info: `_create_grant_table`
    returns None and records the entry so `generate()` can still route it to
    the Appendix instead of dropping it outright."""
    gen = _generator()
    fields = {}
    entry = {'text': 'DECLINED_SPARSE_TOKEN', 'taxonomy_code': 'M2A', 'extracted_fields': fields}
    table = gen._create_grant_table(fields, 'M2A', entry)
    assert table is None
    assert gen._declined_grant_entries == [entry]


def test_non_declined_entry_is_not_appended_to_declined_grant_entries():
    """A well-formed grant renders a table and leaves the decline list empty."""
    gen = _generator()
    fields = {'agency': 'NIH', 'title': 'Cancer Immunology Project'}
    entry = {'text': 'Cancer Immunology Project', 'taxonomy_code': 'M2A', 'extracted_fields': fields}
    table = gen._create_grant_table(fields, 'M2A', entry)
    assert table is not None
    assert gen._declined_grant_entries == []


def test_declined_entry_with_no_entry_arg_is_not_appended():
    """The `isinstance(entry, dict)` guard's negative path (verify_r2 NOTE 2):
    `entry=None` still declines (returns None) but must not append -- there
    is nothing dict-shaped to append. Kills the mutant that drops the guard
    and appends unconditionally."""
    gen = _generator()
    table = gen._create_grant_table({}, 'M2A', None)
    assert table is None
    assert gen._declined_grant_entries == []


# --- item 6: PI auto-fill from the CV owner -------------------------------------

def test_principal_investigator_role_auto_fills_cv_owner():
    """A Principal Investigator role with no extracted PI name names the CV owner.

    Driven through `_fill_research_support` so the owner resolution wire
    (`_get_cv_owner_name`) is covered as well (review thread 3932451691 item 6).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Owner Led Project', agency='NIH',
                        pi_role='Principal Investigator', start_date='01/2019',
                        end_date='Present')]},
        cv_owner={'first_name': 'Ada', 'last_name': 'Testowner'},
        current_year=TEST_YEAR)

    cells = _cells(_tables_under(gen, CURRENT)[0])
    assert cells['Name of Principal Investigator:'] == 'Ada Testowner'
    assert cells['Your role:'] == 'Principal Investigator'


def test_non_pi_role_does_not_auto_fill_cv_owner():
    """A Co-I gets nothing; the long-form auto-fill needs "principal" and "investigator"."""
    assert resolve_pi_name({}, '', 'Co-Investigator', 'Ada Testowner') == ''
    assert resolve_pi_name({}, '', 'Principal Investigator', 'Ada Testowner') == 'Ada Testowner'


def test_bare_pi_role_auto_fills_cv_owner_through_the_grant_table():
    """Stage 4 writes the owner's role as a bare "PI" since #1410 (JIJRSN-01).

    Driven through `_fill_research_support`, the same wire as the long form, so
    the owner reaches the rendered "Name of Principal Investigator:" cell.
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Owner Led Trial', agency='NIH',
                        pi_role='PI', start_date='01/2019', end_date='Present')]},
        cv_owner={'first_name': 'Ada', 'last_name': 'Testowner'},
        current_year=TEST_YEAR)

    cells = _cells(_tables_under(gen, CURRENT)[0])
    assert cells['Name of Principal Investigator:'] == 'Ada Testowner'
    assert cells['Your role:'] == 'PI'


@pytest.mark.parametrize('role', [
    'PI', 'pi', ' PI ', 'P.I.', 'P.I', 'P. I.', 'p.i.', '(PI)', '( PI )', '(PI) ', 'PI:', 'PI :',
    'Sole PI', 'sole P.I.',
])
def test_bare_pi_role_forms_auto_fill_cv_owner(role):
    assert resolve_pi_name({}, '', role, 'Ada Testowner') == 'Ada Testowner'


@pytest.mark.parametrize('role', [
    'Co-PI', 'co-PI', 'Co-P.I.', 'Co PI', 'Co-I', 'Sub-PI', 'Subcontract PI', 'MPI', 'M-PI',
    'Multi-PI', 'Multiple PI', 'Contact MPI', 'Site PI', 'PI-DDN', 'Pilot', 'Pi Beta', 'P', 'I',
])
def test_pi_roles_that_are_not_the_owner_as_pi_do_not_auto_fill(role):
    """A whole-role match: a qualifier that makes someone else the PI never fires."""
    assert resolve_pi_name({}, '', role, 'Ada Testowner') == ''


def test_long_form_rule_is_unchanged_for_co_principal_investigator():
    """The "principal" + "investigator" test still fires on its long forms.

    That includes "Co-Principal Investigator" and "Site Principal
    Investigator", as before this change; "Co-PI" does not (above). The bare
    form is not widened to match the long form's reach.
    """
    for role in ('Co-Principal Investigator', 'Site Principal Investigator'):
        assert resolve_pi_name({}, '', role, 'Ada Testowner') == 'Ada Testowner'
    assert resolve_pi_name({}, '', 'Principle Investigator', 'Ada Testowner') == ''


def test_a_named_pi_beats_the_owner_on_a_bare_pi_role():
    """The bare-PI auto-fill runs last: an extracted or labelled PI wins."""
    assert resolve_pi_name({'pi_name': 'Jane Smith'}, '', 'PI', 'Ada Testowner') == 'Jane Smith'
    assert resolve_pi_name({}, 'Ellison Foundation (PI: Holloway)', 'PI',
                           'Ada Testowner') == 'Holloway'
    assert resolve_pi_name({}, 'PI', 'PI', '') == ''


def test_extracted_pi_name_beats_the_cv_owner():
    """An extracted PI name wins over the owner auto-fill."""
    assert resolve_pi_name({'pi_name': 'Jane Smith'}, '', 'Principal Investigator',
                           'Ada Testowner') == 'Jane Smith'


def test_null_role_with_a_cv_owner_does_not_crash():
    """A JSON-null `role` reached the owner auto-fill as None and raised there.

    `_create_grant_table` reads the role as
    `fields.get('pi_role') or fields.get('role', '')`, which yields None -- not
    '' -- on a record that carries `role` present and JSON-null while `pi_role`
    is falsy. `resolve_pi_name` then called `None.lower()`, aborting the whole
    stage-6 run for that CV. Both ends are pinned: the classifier direct, and
    the caller wire that manufactures the None. The PI cell stays empty, which
    is what a falsy role rendered before, so nothing that renders today moves.
    """
    assert resolve_pi_name({'role': None}, '', None, 'Ada Testowner') == ''

    table = _generator()._create_grant_table(
        {'title': 'Null Role Project', 'agency': 'NIH', 'role': None},
        'M2A', owner_name='Ada Testowner')
    assert _cells(table)['Name of Principal Investigator:'] == ''


# --- item 7: PI extraction from a pipe-delimited source row ---------------------

def test_pi_name_is_extracted_from_four_part_pipe_text():
    """The trailing cell of "Agency | Amount | Dates | PI Name" is read as the PI.

    Several heuristics stacked in one branch and none of them were covered
    (review thread 3932451691 item 7).
    """
    raw = 'NIH | $100,000 | 2019-2021 | Jane Smith'
    assert resolve_pi_name({}, raw, '', '') == 'Jane Smith'


def test_pi_name_accepts_a_middle_initial():
    """"First M. Last" is the second form the whitespace pattern admits."""
    raw = 'NIH | $100,000 | 2019-2021 | Jane M. Smith'
    assert resolve_pi_name({}, raw, '', '') == 'Jane M. Smith'


@pytest.mark.parametrize('raw, reason', [
    ('NIH | $100,000 | Jane Smith', 'fewer than four pipe-separated parts'),
    ('NIH | NIH | NIH | NIH', 'every part identical (merged source cells)'),
    ('NIH | $1 | 2019 | Grant Smith', 'trailing cell carries a project keyword'),
    ('NIH | $1 | 2019 | Bartholomewmaximilianconstantine Vandersteenhuysenrichards',
     'trailing cell is 50 characters or longer'),
    ('NIH | $1 | 2019 | 2019-2021', 'trailing cell is digits and separators only'),
    ('NIH | $1 | 2019 | Smith, Jane', '"Last, First" is not a form the pattern accepts'),
])
def test_pi_name_is_not_extracted_from_a_rejected_trailing_cell(raw, reason):
    """Each documented reject of the pipe parser leaves the PI name empty.

    The guards overlap by construction -- a digits-only cell also fails the name
    pattern -- so each case pins the outcome the parser is documented to
    produce, not which single guard produced it. "Smith, Jane" is the comment
    fix from thread 3932312407 item 4: the comma form is explicitly out of
    scope, and this is the regression test that says so (review thread
    3932451691 item 7).
    """
    assert resolve_pi_name({}, raw, '', '') == '', reason


# --- #982: PI resolution -- no co_investigators fallback, a "PI:" label first ---

def test_co_investigators_never_fills_the_pi_cell():
    """web207: `co_investigators` held the CV owner and rendered as the PI (#982)."""
    fields = {'co_investigators': 'Tanaka, CoI'}
    assert resolve_pi_name(fields, '', 'Co-Investigator', 'Ada Testowner') == ''


def test_pi_label_in_the_source_text_fills_the_pi_before_the_owner_auto_fill():
    """A named PI wins over the owner even when the role says the owner is a PI."""
    raw = 'PI: Holloway | $329,925 | 5% Tanaka, CoI'
    assert resolve_pi_name({}, raw, 'Principal Investigator', 'Ada Testowner') == 'Holloway'


@pytest.mark.parametrize('raw, expected', [
    ('P30CA047904 Cancer Center Support Grant (PI: Keller) | NCI', 'Keller'),
    ('PI:Moss', 'Moss'),
    ('PI: R Voss', 'R Voss'),
    ('PI: De Luca', 'De Luca'),
    ('(PI: Dr. Min Wu, Department of Radiation Oncology, UMB)', 'Dr. Min Wu'),
    ('Sponsor: Tanaka\nPI: Bradley\nRole: Co-I', 'Bradley'),
    # "(PI Name)" and "[PI Name" without a colon
    ('Quadrangle Seminar Fund (PI Paul Reyes), 2003', 'Paul Reyes'),
    ('(PI Lan Mai Thu Pham, CONVERGE trainee graduate)', 'Lan Mai Thu Pham'),
    ('$3,000,000 / [PI Dr. Ann B. Cole, Co-Is Dr. Omar Haddad]', 'Dr. Ann B. Cole'),
    ('Ellison Foundation (PI Lee and Park)', 'Lee and Park'),
    # "Principal Investigator(s): Name(s)"
    ('Role: Postdoctoral Fellow\tPrincipal Investigator: Sam Ortiz\tTitle: X', 'Sam Ortiz'),
    ('Principal Investigators: Ann Cole & Beth Hale\tFunding agency: A',
     'Ann Cole & Beth Hale'),
    # "Name (PI)", "Names (MPI)", "Name (Principal Investigator)"
    ('National Cancer Institute (R01)\tNguyen (PI)\t07/01/25', 'Nguyen'),
    ('NIA (R21)\tLee & Park (MPI)\t12/01/24', 'Lee & Park'),
    ('Collaborating with Dr. Ana Cruz (Principal Investigator) and others', 'Dr. Ana Cruz'),
    ('1R21EB002742-01 Srinivasan (PI)  9/1/03 - 8/31/05', 'Srinivasan'),
    ('R21EB002742-01 Srinivasan (PI)  9/1/03 - 8/31/05', 'Srinivasan'),
    # "Name, Principal Investigator" / "Name, PI"
    ('$144,301. M. Silva, Principal Investigator, A Byrne, Co-Investigator.', 'M. Silva'),
    ('Rosa Diaz, PI.   Co-investigators: Ivan T. Roth', 'Rosa Diaz'),
    # "[PIs A, B, ...]"
    ('$1,350,000 [PIs Dr. Ann Cole, Dr. Ben Hale, et al.]',
     'Dr. Ann Cole, Dr. Ben Hale'),
])
def test_pi_label_name_shapes(raw, expected):
    assert resolve_pi_name({}, raw, '', '') == expected


@pytest.mark.parametrize('raw, reason', [
    ('Co-PI: Eric Stone', 'a Co-PI is not the PI'),
    ('Subcontract-PI: $103,886 (direct costs)', 'an amount, not a name'),
    ('06/01/2021-05/31/2026 PI: 75%\tNIAID', 'an effort figure, not a name'),
    ('PI:\tHart Role on Grant:', 'a tab ends the field: the next cell is not the name'),
    ('PI: smith', 'a lower-case word is not a name'),
    ('(PI Lan Mai Thu Pham Vo Ha)', 'a name past the token limit is not truncated to a wrong one'),
    ('(PI: Jones R01 renewal)', 'an award-number token after a surname is not part of a name'),
    ('Rehabilitation Institute of Chicago, Principal Investigator.',
     '"Name, Principal Investigator" needs two tokens: Chicago is the owner\'s place'),
    ('1P20CA086278-01A1Frost (PI)', 'an award number glued to a name is not a name'),
    ('Manual Therapy grant (co-principal investigator) Smith', 'no PI label at all'),
])
def test_pi_label_that_is_not_a_pi_name_is_left_alone(raw, reason):
    assert resolve_pi_name({}, raw, '', '') == '', reason


def test_a_pi_label_naming_the_cv_owner_by_surname_renders_the_owner_full_name():
    """"Srinivasan (PI)" on Srinivasan's own CV is the owner, so the cell keeps the
    owner's full name, as the owner auto-fill rendered it before the label was read."""
    assert resolve_pi_name({}, 'Duke (PI: Srinivasan) 2017', 'PI', 'Priya Srinivasan') \
        == 'Priya Srinivasan'
    assert resolve_pi_name({}, 'Duke (PI: Someone Else) 2017', 'PI', 'Priya Srinivasan') \
        == 'Someone Else'


@pytest.mark.parametrize('raw, owner', [
    ('PI: Nora Smith', 'Nora M. Quinn'),
    ('PI: Mark Anderson', 'Andrew Mark Jones'),
])
def test_a_pi_sharing_only_a_given_name_with_the_cv_owner_is_not_the_owner(raw, owner):
    """A shared first or middle name is not the owner: the named PI is kept."""
    assert resolve_pi_name({}, raw, '', owner) == raw[len('PI: '):]


@pytest.mark.parametrize('raw', ['PI: Quinn', 'PI: Quinn, Nora', 'PI: Nora Quinn'])
def test_a_pi_with_the_cv_owners_surname_is_the_owner(raw):
    assert resolve_pi_name({}, raw, '', 'Nora M. Quinn') == 'Nora M. Quinn'


@pytest.mark.parametrize('name, owner, expected', [
    ('Quinn, Nora', 'Nora M. Quinn', False),
    ('Quinn, Holloway', 'Nora M. Quinn', False),
    ('Nora Quinn', 'Nora M. Quinn', True),
    ('Dr. Nora Quinn, Dr. Ann Lee', 'Nora M. Quinn', False),
    ('Quinn', '', False),
    ('Al Li', 'Bo Ng', False),
    ('Ana Rivera', 'Ana Rivera-Ortiz', True),
    ('Rivera AO', 'Ana Rivera-Ortiz', True),
    ('Ana Ortiz', 'Ana Rivera-Ortiz', True),
    ('Ana Smith', 'Ana Rivera-Ortiz', False),
    ('Lee & Park', 'Amy Park', False),
])
def test_is_cv_owner_matches_one_surname_and_never_a_list(name, owner, expected):
    assert _is_cv_owner(name, owner) is expected


@pytest.mark.parametrize('raw', ['Harvard (PIs Quinn, Holloway) 2019', '[PIs Nora Quinn, Bo Li]'])
def test_a_pi_list_naming_the_cv_owner_keeps_every_pi(raw):
    """A "PIs" list that includes the owner is not collapsed to the owner: that
    would drop the other PI, the loss #982 is about."""
    assert resolve_pi_name({}, raw, '', 'Nora M. Quinn') == raw.split('PIs ')[1].rstrip(']').split(')')[0]


def test_an_and_joined_pair_naming_the_cv_owner_keeps_both_pis():
    assert resolve_pi_name({}, 'Ellison Foundation (PI Lee and Park)', '', 'Amy Park') == 'Lee and Park'


@pytest.mark.parametrize('raw, expected', [
    ("PI: Mary O'Brien", "Mary O'Brien"),
    ('PI: Mary O\u2019Brien', 'Mary O\u2019Brien'),
    ('PI: Ana Rivera-Ortiz', 'Ana Rivera-Ortiz'),
])
def test_pi_name_tokens_keep_apostrophes_and_hyphens(raw, expected):
    assert resolve_pi_name({}, raw, '', '') == expected


def test_a_hyphenated_pi_label_naming_the_cv_owner_renders_the_owner():
    assert resolve_pi_name({}, 'PI: Ana Rivera', '', 'Ana Rivera-Ortiz') == 'Ana Rivera-Ortiz'


def test_a_name_glued_to_an_award_number_by_a_hyphen_is_not_a_name():
    assert resolve_pi_name({}, 'R01-Frost (PI)', '', '') == ''


def test_pi_name_token_limit_is_exactly_three_extra_tokens():
    """Four extra tokens is over the limit and is skipped, not truncated."""
    assert resolve_pi_name({}, '(PI Ann Bo Cy Di)', '', '') == 'Ann Bo Cy Di'
    assert resolve_pi_name({}, '(PI Ann Bo Cy Di Ed)', '', '') == ''


def test_a_joined_pair_naming_the_cv_owner_is_not_collapsed_to_the_owner():
    assert resolve_pi_name({}, 'NIA\tLee & Park (MPI)', '', 'Amy Lee') == 'Lee & Park'


def test_an_explicit_pi_label_beats_a_name_suffix_shape():
    """Shapes are tried in order: "PI: <name>" outranks "<name>, Principal Investigator"."""
    raw = 'Ada Lovelace, Principal Investigator; renewal (PI: Grace Hopper)'
    assert resolve_pi_name({}, raw, '', '') == 'Grace Hopper'


_OWN_TITLE = 'Stress Study \u2013 Mouse Models'
_OWN_TITLE_TEXT = '1997 Example Foundation, "Stress Study" \u2013 Mouse Models, PI, Total $7,000.'


def test_a_pi_label_capturing_the_grants_own_title_words_is_not_a_pi():
    """#1403 (EQADVR 396): "<title>, PI" matched the "Name, PI" shape with the
    title's last two words. Those words are the title, not a PI. With no
    role the cell stays empty; with the bare "PI" role the owner fills it
    (#1446's auto-fill), never the title words."""
    fields = {'title': _OWN_TITLE, 'pi_role': 'PI'}
    assert _pi_name_from_label(_OWN_TITLE_TEXT) == 'Mouse Models'
    assert resolve_pi_name(fields, _OWN_TITLE_TEXT, '', 'Ada Testowner') == ''
    assert resolve_pi_name(fields, _OWN_TITLE_TEXT, 'PI', 'Ada Testowner') == 'Ada Testowner'


@pytest.mark.parametrize('fields', [
    {'title': _OWN_TITLE.lower()},
    {'title': None, 'study_title': _OWN_TITLE},
])
def test_the_own_title_guard_is_case_folded_and_reads_study_title(fields):
    assert resolve_pi_name(fields, _OWN_TITLE_TEXT, '', '') == ''


@pytest.mark.parametrize('title', ['Stress Study', 'Mouse Modelsx Study', 'AMouse Models Study',
                                   None, ['a', 'list']])
def test_a_pi_label_outside_the_title_still_names_the_pi(title):
    """Only whole title words veto the capture: a title without them, one
    that holds them inside a longer word, no title, or a non-string title
    (stage 4 stores raw LLM JSON) keeps it."""
    assert resolve_pi_name({'title': title}, _OWN_TITLE_TEXT, '', '') == 'Mouse Models'


@pytest.mark.parametrize('title', ['Support Grant (PI: Lee)', 'Support Grant [PI Lee]',
                                   'Support Grant, Lee (PI)'])
def test_a_pi_named_by_a_label_inside_the_title_is_still_the_pi(title):
    """A stage-4 title can keep the source's "(PI: <name>)" parenthetical. The
    name follows a PI label there, so it is the PI, not title words: the guard
    must not veto it (farm M2B entry 61.1)."""
    raw = '08/01/2020-07/31/2025 | P30CA000000/ ' + title + ' | NCI | $1,000 (Total) | 5%'
    assert resolve_pi_name({'title': title, 'pi_role': '5%'}, raw, '5%', 'Ada Testowner') == 'Lee'


def test_a_title_word_capture_falls_through_to_the_owner_auto_fill():
    """With the capture refused, a "Principal Investigator" role still fills
    the owner, as it does for any grant whose text names no PI."""
    assert resolve_pi_name({'title': _OWN_TITLE}, _OWN_TITLE_TEXT, 'Principal Investigator',
                           'Ada Testowner') == 'Ada Testowner'


def test_a_grant_whose_title_tail_matches_the_pi_shape_renders_the_owner_not_title_words():
    """The rendered wire: the PI cell never names the title words. The role is
    a bare "PI", so the owner auto-fill (#1446) names the CV owner."""
    table = _generator()._create_grant_table(
        {'title': _OWN_TITLE, 'agency': 'Example Foundation', 'pi_role': 'PI',
         'start_date': '1997'},
        'M2B', entry={'text': _OWN_TITLE_TEXT}, owner_name='Ada Testowner')
    cells = _cells(table)
    assert cells['Name of Principal Investigator:'] == 'Ada Testowner'
    assert cells['Your role:'] == 'PI'


def test_a_null_source_text_resolves_to_no_pi_rather_than_raising():
    """`entry['text']` can be a JSON null; the label parse must not hand it to `re`."""
    assert resolve_pi_name({}, None, '', 'Ada Testowner') == ''


def test_extracted_pi_name_beats_the_pi_label():
    assert resolve_pi_name({'pi_name': 'Jane Smith'}, 'PI: Someone Else', '', '') == 'Jane Smith'


def test_pi_label_beats_the_pipe_row_trailing_cell():
    raw = 'NIH | $100,000 | 2019-2021 | Jane Smith (PI: Ada Lovelace)'
    assert resolve_pi_name({}, raw, '', '') == 'Ada Lovelace'


def test_a_grant_whose_text_names_another_pi_does_not_render_the_owner():
    """The rendered wire: co-investigator owner, `PI: <other>` in the text."""
    gen = _generator()
    table = gen._create_grant_table(
        {'title': 'Structure-specific nuclease study', 'agency': 'NIH',
         'co_investigators': 'Tanaka, CoI', 'annual_funding': '329925'},
        'M2A', entry={'text': 'PI: Holloway | $329,925 | 5% Tanaka, CoI'},
        owner_name='Mia Tanaka')
    cells = _cells(table)
    assert cells['Name of Principal Investigator:'] == 'Holloway'
    assert cells['Annual direct costs:'] == '$329,925'


# --- item 8: percent-effort extraction, parsed and applied ----------------------

def test_role_effort_header_is_parsed_into_the_effort_lookup():
    """The reviewer's own worked example, asserted on the parsed lookup.

    "Individual's role in project including percent effort" is a source-table
    header, not a grant: its body lines become the effort lookup and the row
    itself is dropped from the entries to render (review thread 3932451691
    item 8).
    """
    lookup = {}
    entries = [{'text': "Individual's role in project including percent effort\n"
                        'Project Alpha 0.01\n'
                        'Project Beta .08FTE\n'
                        'Project Gamma 25'}]
    kept, messages = filter_role_effort_headers(entries, lookup)

    assert kept == []
    assert lookup == {'project alpha': '1%', 'project beta': '8%', 'project gamma': '25%'}
    assert messages == ['  Filtered role/effort header entry, extracted 3 effort values']


# --- #1227: only the template's header row is filtered, never a valued grant row --

# Each of these has the header phrase inside its first 60 characters and also
# carries a value a reader of the CV would look for. The first four are
# label|value rows (a 2-column grant form, with and without the colon, and a
# pipe-only row whose value has no digit); then a free-text role line; then
# values written with no colon or pipe at all (a number after the label, in
# brackets, after a tab, or around the phrase); the last is the phrase on a
# later line of an entry whose first line is a different label|value row.
VALUED_ROLE_EFFORT_ROWS = [
    'Percent Effort: | 35\nTotal Direct Costs: | $12,345',
    'Percent Effort: | 45',
    'Percent Effort | 45',
    'Percent Effort | Part time',
    'Role and percent effort:   Example Co-I (unfunded)',
    'Percent effort 35%',
    'Percent effort 5%',
    'Percent effort (35%)',
    'Percent Effort\t35',
    'Role in project PI, 35% effort',
    'Example Agency Grant; role in project Co-I 55%',
    'Grant A12345 role in project PI',
    'Years Inclusive: | 01/01/2031 - 12/31/2031\nPercent Effort: | 65\n'
    'Total Direct Costs: | $67,890',
]


@pytest.mark.parametrize('text', VALUED_ROLE_EFFORT_ROWS)
def test_a_grant_row_carrying_a_value_is_not_a_role_effort_header(text):
    """A label|value row whose label is "Percent Effort" is a grant, not the header.

    Every one of these was dropped with no Appendix line and no log: the filter
    looked at the label wording in the first 60 characters and nothing else
    (#1227).
    """
    lookup = {}
    entries = [{'text': text}]
    kept, messages = filter_role_effort_headers(entries, lookup)

    assert kept == entries
    assert lookup == {}
    assert messages == []


@pytest.mark.parametrize('text', [
    "Individual's role in project including percent effort",
    "Individual's role in project including percent effort\n",
    "Individual's role in project including percent effort\n"
    'Project Alpha: percent effort | 0.01',
    'Source | Title | Individual\'s role in project including percent effort',
    "Your role in the project (including percent effort):",
    'Percent Effort:',
    'Percent Effort: |',
    'Percent Effort: |\n',
    'Percent Effort: |\n\n   \n',
])
def test_a_bare_header_phrase_or_blank_label_is_still_filtered(text):
    """The template header sentence, and a label whose value cell is blank, go.

    The blank label carries nothing to lose when nothing is below it; dropping
    it keeps an empty "Percent Effort:" line out of the Appendix, as before.
    """
    kept, _messages = filter_role_effort_headers([{'text': text}], {})

    assert kept == []


# A first line that is a blank label ("Percent Effort:", "Percent Effort: |"),
# with the fragment's own value rows under it. Stage 2 splits a label|value grant
# into fragments that begin at the effort row, so the rows below the label are the
# grant's cost and dates. Dropping the fragment as a header lost them (#1227).
BLANK_LABEL_WITH_VALUED_ROWS = [
    'Percent Effort: |\nTotal Direct Costs: | $12,345',
    'Percent Effort:\nTotal Direct Costs: | $12,345',
    'Percent Effort: |\nYears Inclusive: | 01/01/2031 - 12/31/2031\n'
    'Total Direct Costs: | $12,345',
    "Your role in the project (including percent effort):\nTotal Direct Costs: | $12,345",
    # Rows that END in a bare number have the "<project> <effort>" shape, but the
    # number is not a percent effort (a year, a cost with no `$` or separators),
    # so they are not pairs the filter reads.
    'Percent Effort: |\nYears Inclusive: | 2031 - 2033',
    'Percent Effort: |\nTotal Direct Costs: | 12345',
    # A value cell that wrapped onto its own, indented line: a number with no
    # project name in front of it is not a pair once the line is stripped.
    'Percent Effort: |\n    35',
]


@pytest.mark.parametrize('text', BLANK_LABEL_WITH_VALUED_ROWS)
def test_a_blank_label_with_valued_rows_below_it_is_not_a_role_effort_header(text):
    """The rows under the label are not `<project> <effort>` pairs, so it is a grant.

    The filter only reads pairs, with an effort that normalizes, out of a header;
    anything else under the first line is content it would drop unread.
    """
    lookup = {}
    entries = [{'text': text}]
    kept, messages = filter_role_effort_headers(entries, lookup)

    assert kept == entries
    assert lookup == {}
    assert messages == []


def test_one_row_that_is_not_a_pair_keeps_the_whole_entry():
    """Every row under the first line has to be a pair, not just one of them."""
    lookup = {}
    entries = [{'text': "Individual's role in project including percent effort\n"
                        'Project Alpha 0.01\n'
                        'Total Direct Costs: | $12,345'}]
    kept, messages = filter_role_effort_headers(entries, lookup)

    assert kept == entries
    assert lookup == {}
    assert messages == []


@pytest.mark.parametrize('text', [
    'Funding Agency: | Example Agency\nPercent Effort: |',
    # Pair-shaped rows below, and the first digit past character 60 of the entry,
    # so neither the digit rule nor the pair rule can be what keeps it: only
    # reading the phrase off the FIRST line does.
    'Funding Agency: | Example Agency\nPercent Effort Example Long Project Name Here 0.5',
])
def test_a_phrase_on_a_later_line_is_not_a_role_effort_header(text):
    """The 60-character window ends with the first line, not 60 characters in."""
    entries = [{'text': text}]
    kept, _messages = filter_role_effort_headers(entries, {})

    assert kept == entries


def test_a_blank_label_with_pairs_below_it_is_a_header_and_the_pairs_are_read():
    """A header whose first line ends in a colon, over its project/effort pairs.

    The pairs are harvested and the entry goes, whitespace-only rows between
    them included.
    """
    lookup = {}
    entries = [{'text': 'Your role in the project (including percent effort):\n'
                        '\n'
                        '   \n'
                        'Project Alpha 0.01\n'
                        '  Project Beta 25  '}]
    kept, messages = filter_role_effort_headers(entries, lookup)

    assert kept == []
    assert lookup == {'project alpha': '1%', 'project beta': '25%'}
    assert messages == ['  Filtered role/effort header entry, extracted 2 effort values']


# The phrase is inside the first 60 characters, and the first value is past them.
# The first has a digit and a colon, the second only a digit, the third only a
# colon, so each of the two value checks is the only thing that keeps one of them.
VALUE_PAST_THE_WINDOW_ROWS = [
    'Role in project and percent effort (calendar months per year): 1.2',
    'Percent effort on the collaborative multi-site award per year 35%',
    'Role in project and percent effort for the collaborative multi-site award: Example Co-I',
]


@pytest.mark.parametrize('text', VALUE_PAST_THE_WINDOW_ROWS)
def test_a_value_past_the_search_window_is_not_a_role_effort_header(text):
    """The digit and colon/pipe checks read the whole first line, not 60 characters.

    The 60 characters bound the phrase search only; a long label that puts its
    value after them is still a grant row, and used to be dropped silently.
    """
    window = text[:ROLE_EFFORT_HEADER_WINDOW]
    assert ROLE_EFFORT_HEADER_RE.search(window)
    assert not ROLE_EFFORT_DIGIT_RE.search(window)
    assert not ROLE_EFFORT_FIELD_ROW_RE.search(window)

    entries = [{'text': text}]
    kept, _messages = filter_role_effort_headers(entries, {})

    assert kept == entries


@pytest.mark.parametrize('padding, is_header', [
    # 'role in project' is 15 characters. The numbers are literals on purpose: a
    # test built from ROLE_EFFORT_HEADER_WINDOW would move with the constant and
    # pin nothing.
    (45, True),
    (46, False),
])
def test_the_header_phrase_has_to_end_inside_the_search_window(padding, is_header):
    """The window is 60 characters: a phrase ending at 60 counts, one ending at 61 does not."""
    text = 'Example Agency'.ljust(padding) + 'role in project'
    assert len(text) == (60 if is_header else 61)

    assert is_role_effort_header(text) is is_header


def test_a_phrase_past_the_search_window_does_not_make_a_header():
    """A first line that only reaches the phrase after 60 characters is a grant."""
    text = 'Example funding program for pediatric research, with the work ' \
           'plan and percent effort noted'
    assert text.lower().index('percent effort') > 60
    kept, _messages = filter_role_effort_headers([{'text': text}], {})

    assert kept == [{'text': text}]


def test_a_header_row_is_filtered_and_the_valued_row_beside_it_is_kept():
    """Both in one bucket: the header's pairs are harvested, the grant row stays.

    The real template header row carries stage-4 fields too (a `title` and a
    `percent_effort` read off its first project line), so "no extracted fields"
    cannot be the test; the header is told apart by its text.
    """
    header = {'text': "Individual's role in project including percent effort\n"
                      'Project Alpha 0.01',
              'extracted_fields': {'title': 'Project Alpha', 'percent_effort': '1%'}}
    grant_row = {'text': 'Percent Effort: | 35\nTotal Direct Costs: | $12,345',
                 'extracted_fields': {'percent_effort': '35%', 'total_funding': '$12,345'}}
    lookup = {}
    kept, _messages = filter_role_effort_headers([header, grant_row], lookup)

    assert kept == [grant_row]
    assert lookup == {'project alpha': '1%'}


def test_a_valued_effort_row_reaches_the_appendix_through_the_section():
    """The wire: `_fill_research_support` hands the row on, `_create_grant_table`
    declines it (no title, agency or grant number), and it lands on the #839
    decline list that `generate()` sends to the Appendix instead of vanishing."""
    gen = _sectioned_generator()
    fragment = _entry('M2A', text='Percent Effort: | 35\nTotal Direct Costs: | $12,345',
                      percent_effort='35%', total_funding='$12,345')
    gen._fill_research_support({'M2A': [fragment]}, current_year=TEST_YEAR)

    assert [e['text'] for e in gen._declined_grant_entries] == [fragment['text']]


def test_a_blank_effort_label_with_value_rows_reaches_the_appendix_through_the_section():
    """The wire for the blank-label fragment: cost and dates are not dropped.

    Same route as the valued row above: no title, agency or grant number, so
    `_create_grant_table` declines it and the whole fragment, cost row included,
    is on the list `generate()` sends to the Appendix."""
    gen = _sectioned_generator()
    fragment = _entry('M2A', text='Percent Effort: |\nTotal Direct Costs: | $12,345',
                      total_funding='$12,345')
    gen._fill_research_support({'M2A': [fragment]}, current_year=TEST_YEAR)

    assert [e['text'] for e in gen._declined_grant_entries] == [fragment['text']]


@pytest.mark.parametrize('text, fields', [
    ('Percent Effort: |\nYears Inclusive: | 2031 - 2033',
     {'start_date': '2031', 'end_date': '2033'}),
    ('Percent Effort: |\nTotal Direct Costs: | 12345', {'total_funding': '12345'}),
    ('Role in project and percent effort (calendar months per year): 1.2',
     {'percent_effort': '1.2'}),
    ('Percent effort on the collaborative multi-site award per year 35%',
     {'percent_effort': '35%'}),
])
def test_a_value_the_header_shape_check_used_to_miss_reaches_the_appendix(text, fields):
    """The wire for the year-range / bare-cost fragments and the long-label rows.

    Each was read as a header and removed with nothing harvested and nothing on
    the list `generate()` sends to the Appendix."""
    gen = _sectioned_generator()
    fragment = _entry('M2A', text=text, **fields)
    gen._fill_research_support({'M2A': [fragment]}, current_year=TEST_YEAR)

    assert [e['text'] for e in gen._declined_grant_entries] == [text]


def test_a_titled_grant_that_starts_with_an_effort_row_is_rendered():
    """The wire, rendered: a whole grant whose first row is `Percent Effort`.

    The title makes it a real grant, so it must come out as a table under the
    Current header rather than being filtered as a header."""
    gen = _sectioned_generator()
    entry = _entry('M2A', text='Percent Effort: | 35\nTitle of Grant: | Example Titled Study',
                   title='Example Titled Study', percent_effort='35%')
    gen._fill_research_support({'M2A': [entry]}, current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Example Titled Study']


@pytest.mark.parametrize('raw, expected', [
    ('0.01', '1%'),
    ('.08', '8%'),
    ('25', '25%'),
    ('0.015', '1.5%'),   # truncated to "1%" before the fix
    ('1.5', '1.5%'),     # truncated to "1%" before the fix
    ('1', '100%'),       # at the fractional ceiling: still read as a fraction
    ('100', '100%'),
])
def test_percent_effort_keeps_its_stated_precision(raw, expected):
    """Fractional percentages survive the conversion instead of being truncated.

    The two `int()` calls this replaced truncated on both branches:
    `int(0.015 * 100)` rendered a 1.5% effort as "1%", and `int(1.5)` rendered
    a 1.5% effort as "1%" as well (review thread 3932312407 item 3, thread
    3932451691 item 8).
    """
    assert normalize_percent_effort(raw) == expected


@pytest.mark.parametrize('raw', ['0', '0.0', '-1', '101', '150', 'abc', ''])
def test_out_of_range_percent_effort_is_dropped(raw):
    """A figure outside (0%, 100%] is not a percent effort and is not rendered.

    Before the fix "0" rendered as "0%" and "150" as "150%" (review thread
    3932312407 item 3).
    """
    assert normalize_percent_effort(raw) is None


@pytest.mark.parametrize('raw', ['nan', 'NaN', '-NaN', 'sNaN'])
def test_a_nan_effort_figure_is_discarded_rather_than_raised(raw):
    """A NaN got past the constructor guard and raised off the next comparison.

    `Decimal('abc')` raises InvalidOperation from the constructor and is caught;
    `Decimal('NaN')` constructs, and then `value <= FRACTIONAL_EFFORT_CEILING`
    signals the same InvalidOperation out of the function. The docstring
    promised None for anything unparseable, so this pins the promise for the
    input that broke it. `PROJECT_EFFORT_LINE_RE` cannot produce these strings,
    so no rendered byte depends on it.
    """
    assert normalize_percent_effort(raw) is None


@pytest.mark.parametrize('raw', ['Infinity', '-Infinity'])
def test_an_infinite_effort_figure_needs_no_extra_guard(raw):
    """Infinity is why the NaN guard is `is_nan()` and not `is_finite()`.

    Unlike NaN it compares without signalling, so the range check already
    discards it in both signs; pinned so a later widening of that guard is a
    deliberate choice rather than an accident.
    """
    assert normalize_percent_effort(raw) is None


@pytest.mark.parametrize('raw', ['150', '-1', 'abc'])
def test_a_discarded_effort_figure_is_never_written_to_the_log(raw, caplog):
    """The discard is logged structurally: what happened, never which figure.

    The figure is a cell of the CV's own role/effort table, so it is CV
    content. The out-of-range branch used to log the value itself -- "Discarded
    an out-of-range percent effort figure: 150" -- in a module whose docstring
    states the no-PII logging rule and whose other logger calls carry counts
    and key names only.

    Values chosen so the assertion cannot pass by accident: none of them is a
    substring of the range the message names.
    """
    with caplog.at_level(logging.DEBUG, logger=research_support.__name__):
        assert normalize_percent_effort(raw) is None

    assert caplog.records, 'the discard is still expected to be logged'
    for record in caplog.records:
        assert raw not in record.getMessage()


def test_extracted_effort_reaches_the_rendered_percent_effort_cell():
    """The header's effort figure lands in the matching grant's rendered cell.

    The end-to-end half of item 8: parsing the lookup is not the behaviour, the
    populated "Your percent (%) effort:" row is (review thread 3932451691
    item 8).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [
            {'text': "Individual's role in project including percent effort\n"
                     'Project Alpha 0.01\n'
                     'Project Beta .08FTE'},
            _entry('M2A', title='Project Alpha', agency='NIH', start_date='01/2019',
                   end_date='Present'),
            _entry('M2A', title='Project Beta', agency='NSF', start_date='01/2019',
                   end_date='Present'),
        ]},
        current_year=TEST_YEAR)

    rendered = {_cells(t)['Project title:']: _cells(t)['Your percent (%) effort:']
                for t in _tables_under(gen, CURRENT)}
    assert rendered == {'Project Alpha': '1%', 'Project Beta': '8%'}


def test_role_effort_header_row_is_not_rendered_as_a_grant():
    """The header row itself leaves no table behind."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [{'text': "Individual's role in project including percent effort\n"
                          'Project Alpha 0.01'}]},
        current_year=TEST_YEAR)

    assert _tables_under(gen, CURRENT) == []


# --- item 9: an already-extracted effort wins -----------------------------------

def test_extracted_percent_effort_takes_precedence():
    """A percent effort already on the record is never overwritten by the lookup.

    `if not title or fields.get('percent_effort'): continue` -- the lookup only
    fills a gap (review thread 3932451691 item 9).
    """
    entries = [_entry('M2A', title='Project Alpha', percent_effort='12%')]
    apply_effort_to_grants(entries, {'project alpha': '8%'})
    assert entries[0]['extracted_fields']['percent_effort'] == '12%'


def test_extracted_percent_effort_takes_precedence_in_the_rendered_cell():
    """The same precedence, read off the document rather than the record."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [
            {'text': "Individual's role in project including percent effort\n"
                     'Project Alpha .08FTE'},
            _entry('M2A', title='Project Alpha', agency='NIH', percent_effort='12%',
                   start_date='01/2019', end_date='Present'),
        ]},
        current_year=TEST_YEAR)

    assert _cells(_tables_under(gen, CURRENT)[0])['Your percent (%) effort:'] == '12%'


# --- item 10: project-name matching, both directions, no false positives --------

def test_effort_matches_the_exactly_named_project_in_both_directions():
    """Neither of a pair of prefix-related project names takes the other's effort.

    The old rule was substring-only and took the first hit in insertion order,
    so "Project Alpha Extended" was handed "Project Alpha"'s 1% -- a wrong
    percent effort on a rendered CV, not merely a testing gap (review thread
    3932451691 item 10 / thread 3932312407 item 6).
    """
    lookup = {'project alpha': '1%', 'project alpha extended': '8%'}
    assert match_effort_for_title('project alpha', lookup) == '1%'
    assert match_effort_for_title('project alpha extended', lookup) == '8%'


def test_prefix_related_projects_render_their_own_percent_effort():
    """The same pair, end to end, read off the two generated tables."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [
            {'text': "Individual's role in project including percent effort\n"
                     'Project Alpha 0.01\n'
                     'Project Alpha Extended .08FTE'},
            _entry('M2A', title='Project Alpha', agency='NIH', start_date='01/2019',
                   end_date='Present'),
            _entry('M2A', title='Project Alpha Extended', agency='NIH',
                   start_date='01/2019', end_date='Present'),
        ]},
        current_year=TEST_YEAR)

    rendered = {_cells(t)['Project title:']: _cells(t)['Your percent (%) effort:']
                for t in _tables_under(gen, CURRENT)}
    assert rendered == {'Project Alpha': '1%', 'Project Alpha Extended': '8%'}


def test_a_single_substring_match_is_still_used():
    """Substring matching survives as the explicit fallback, not as the first rule.

    A source table really does write "Project Alpha" in the effort header and
    "Project Alpha: aims 1-3" in the grant list.
    """
    assert match_effort_for_title('project alpha: aims 1-3',
                                  {'project alpha': '1%'}) == '1%'


def test_ambiguous_substring_match_assigns_no_effort():
    """Two project names substring-matching one title assign nothing at all.

    Guessing between them is how the wrong percentage reached a CV in the first
    place (review thread 3932451691 item 10).
    """
    lookup = {'alpha': '1%', 'alpha study': '8%'}
    assert match_effort_for_title('alpha study of xyz', lookup) is None


# --- item 11: a numeric role is a percent effort --------------------------------

@pytest.mark.parametrize('role', ['10', '10%', '0.5', '0.5%'])
def test_numeric_pi_role_is_treated_as_percent_effort(role):
    """A `pi_role` that is only a number is the misextracted percent effort.

    The role cell empties and the effort cell takes the raw value (review
    thread 3932451691 item 11).
    """
    fields = {'title': 'Numeric Role Project', 'pi_role': role, 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Your percent (%) effort:'] == role
    assert cells['Your role:'] == ''


@pytest.mark.parametrize('role', ['10 percent', 'approximately 10%', 'Co-Investigator'])
def test_descriptive_role_stays_a_role(role):
    """Anything that is not purely numeric stays in the role cell, untouched."""
    fields = {'title': 'Descriptive Role Project', 'pi_role': role, 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Your role:'] == role
    assert cells['Your percent (%) effort:'] == ''


def test_numeric_role_conversion_does_not_run_when_an_effort_exists():
    """With a percent effort already extracted the role is left alone."""
    fields = {'title': 'Numeric Role Project', 'pi_role': '10',
              'percent_effort': '12%', 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Your role:'] == '10'
    assert cells['Your percent (%) effort:'] == '12%'


# --- item 12: title fallbacks ---------------------------------------------------

@pytest.mark.parametrize('field', ['title', 'trial_title', 'study_title', 'text'])
def test_grant_title_fallbacks(field):
    """Each of the four title keys alone produces the "Project title:" value.

    Clinical-trial records arrive under `trial_title` / `study_title` / `text`
    rather than `title`, and none of those paths was covered (review thread
    3932451691 item 12).
    """
    fields = {field: 'Longitudinal Cohort Study', 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Project title:'] == 'Longitudinal Cohort Study'


def test_title_precedence_is_title_then_trial_then_study_then_text():
    """The four keys are an `or` chain, in that order."""
    fields = {'title': 'A Title Long Enough', 'trial_title': 'A Trial Title',
              'study_title': 'A Study Title', 'text': 'Some Raw Text',
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Project title:'] == 'A Title Long Enough'

    del fields['title']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Project title:'] \
        == 'A Trial Title'
    del fields['trial_title']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Project title:'] \
        == 'A Study Title'
    del fields['study_title']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Project title:'] \
        == 'Some Raw Text'


# --- item 13: agency fallbacks --------------------------------------------------

@pytest.mark.parametrize('field', ['agency', 'funding_source', 'sponsor'])
def test_award_source_agency_fallbacks(field):
    """Each of the three agency keys alone produces the Award Source value.

    Industry-sponsored records use `sponsor` and some extractions use
    `funding_source`; neither path was covered (review thread 3932451691
    item 13).
    """
    fields = {'title': 'Sponsored Cohort Study', field: 'Acme Pharmaceuticals',
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Award Source:'] == 'Acme Pharmaceuticals'


def test_agency_precedence_is_agency_then_funding_source_then_sponsor():
    """The three keys are an `or` chain, in that order."""
    fields = {'title': 'Sponsored Cohort Study', 'agency': 'NIH',
              'funding_source': 'Foundation', 'sponsor': 'Acme',
              'start_date': '01/2019'}
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Award Source:'] == 'NIH'
    del fields['agency']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Award Source:'] \
        == 'Foundation'
    del fields['funding_source']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Award Source:'] == 'Acme'


# --- item 14: funding fallback --------------------------------------------------

@pytest.mark.parametrize('field', ['annual_direct_costs', 'annual_funding'])
def test_annual_direct_costs_fallbacks(field):
    """Either yearly key alone fills the "Annual direct costs:" row.

    `annual_funding` is the key the stage-4 M2 schema actually emits, and it
    was never read, so 16 corpus records rendered an empty Annual cell (#982).
    """
    fields = {'title': 'Funded Cohort Study', field: '100000', 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Annual direct costs:'] == '$100,000'
    assert 'Total award:' not in cells


def test_total_funding_alone_is_labelled_total_award_not_annual():
    """A total with no yearly figure is a total: it never sits under "Annual".

    web36's "Institutional Award: $2,638,299" used to render as Annual direct
    costs (#982). The row keeps its slot, so the block is still eight rows.
    """
    fields = {'title': 'Funded Cohort Study', 'total_funding': '2638299',
              'start_date': '01/2019'}
    table = _generator()._create_grant_table(fields, 'M2A')
    cells = _cells(table)
    assert cells['Total award:'] == '$2,638,299'
    assert 'Annual direct costs:' not in cells
    assert len(table.rows) == 8


def test_annual_and_total_both_render_as_two_rows_and_neither_is_dropped():
    """web207 R01CA284633: annual $329,925 under Annual, the $1,649,625 total under
    its own label directly after it -- the old render showed the total (mislabelled)
    and dropped the annual figure, so dropping the total now would just swap the loss."""
    fields = {'title': 'Funded Cohort Study', 'annual_funding': '329925',
              'total_funding': '1649625', 'start_date': '01/2019'}
    table = _generator()._create_grant_table(fields, 'M2A')
    rows = _rows(table)
    costs_at = [label for label, _ in rows].index('Annual direct costs:')
    assert rows[costs_at] == ('Annual direct costs:', '$329,925')
    assert rows[costs_at + 1] == ('Total award:', '$1,649,625')
    assert len(rows) == 9


def test_a_total_equal_to_the_annual_amount_is_one_row():
    """The same figure twice is one fact; `total_funding` falling back to the
    yearly amount must not manufacture a Total award row."""
    fields = {'title': 'Funded Cohort Study', 'annual_funding': '100000',
              'total_funding': '100000', 'start_date': '01/2019'}
    table = _generator()._create_grant_table(fields, 'M2A')
    assert len(table.rows) == 8
    assert 'Total award:' not in _cells(table)


def test_annual_direct_costs_beats_annual_funding_and_blank_falls_through_to_it():
    """The older key is read first; a blank one does not shadow `annual_funding`."""
    both = {'title': 'Funded Cohort Study', 'annual_direct_costs': '100000',
            'annual_funding': '200000', 'start_date': '01/2019'}
    assert _cells(_generator()._create_grant_table(both, 'M2A'))['Annual direct costs:'] \
        == '$100,000'
    blank = dict(both, annual_direct_costs='  ')
    assert _cells(_generator()._create_grant_table(blank, 'M2A'))['Annual direct costs:'] \
        == '$200,000'


def test_numeric_annual_funding_without_a_total_does_not_crash():
    """The title-duplicate check compares `total_funding`, which falls back to the
    yearly amount, so a JSON number there reached `.strip()` once the yearly key
    was read at all (#982)."""
    fields = {'title': 'Funded Cohort Study', 'annual_funding': 329925,
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Annual direct costs:'] == '$329,925'


def test_no_amount_at_all_keeps_the_annual_label_with_an_empty_cell():
    fields = {'title': 'Funded Cohort Study', 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Annual direct costs:'] == ''


def test_annual_direct_costs_wins_when_both_funding_keys_are_present():
    """`annual_direct_costs` is what the row is for; `total_funding` is the fallback."""
    fields = {'title': 'Funded Cohort Study', 'annual_direct_costs': '100000',
              'total_funding': '500000', 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Annual direct costs:'] == '$100,000'


# --- item 15: currency formatting -----------------------------------------------

@pytest.mark.parametrize('value, expected', [
    ('100000', '$100,000'),
    ('$100,000', '$100,000'),
    ('100000.50', '$100,000.50'),
    ('100,000', '$100,000'),
    ('not a number', 'not a number'),
])
def test_annual_direct_costs_are_formatted(value, expected):
    """The rendered costs cell follows `_format_currency`'s real contract.

    Whole numbers lose their decimals, an already-$-prefixed value is passed
    through verbatim, and an unparseable value is rendered as it arrived rather
    than dropped (`stage6/formatting/values.py:101-134`). Expected strings are
    read off that implementation, not invented (review thread 3932451691
    item 15).
    """
    fields = {'title': 'Funded Cohort Study', 'annual_direct_costs': value,
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Annual direct costs:'] == expected


# --- item 16: cross-field duplication cleanup -----------------------------------

def test_title_duplicated_into_agency_clears_the_award_source():
    """When stage 4 puts the title in `agency` too, Award Source renders empty.

    Untested branch; the same content in two rows reads as two facts (review
    thread 3932451691 item 16).
    """
    fields = {'title': 'Repeated Content Study', 'agency': 'repeated content study',
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Award Source:'] == ''
    assert cells['Project title:'] == 'Repeated Content Study'


def test_title_duplicated_into_funding_clears_the_costs_cell():
    """The funding half of the same cleanup zeroes the costs row."""
    fields = {'title': 'Repeated Content Study', 'total_funding': 'Repeated Content Study',
              'annual_direct_costs': 'Repeated Content Study', 'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Annual direct costs:'] == ''
    assert cells['Project title:'] == 'Repeated Content Study'


def test_rendering_does_not_mutate_the_callers_extracted_fields():
    """Rendering leaves the caller's record exactly as it arrived.

    The duplicate-funding cleanup used to write `fields['total_funding'] = ''`
    and `fields['annual_direct_costs'] = ''` straight into the pipeline record,
    making rendering non-idempotent and changing what later consumers of the
    entry saw (thread 3932312407 item 1, review thread 3932451691 item 16).
    """
    fields = {'title': 'Repeated Content Study', 'total_funding': 'Repeated Content Study',
              'annual_direct_costs': 'Repeated Content Study', 'start_date': '01/2019'}
    submitted = dict(fields)

    _generator()._create_grant_table(fields, 'M2A', {'text': 'x', 'taxonomy_code': 'M2A',
                                                    'extracted_fields': fields})

    assert fields == submitted


def test_section_fill_does_not_mutate_the_callers_entries():
    """The same immutability one level up, across the whole section fill.

    Classification writes `percent_effort` and `reclassification_note`, both on
    this section's own copies (`copy_entries_for_render`).
    """
    entries_by_code = {
        'M2A': [
            {'text': "Individual's role in project including percent effort\n"
                     'Project Alpha 0.01'},
            _entry('M2A', title='Project Alpha', agency='NIH', start_date='01/2018',
                   end_date='06/2020'),
        ],
    }
    submitted = [dict(e, extracted_fields=dict(e.get('extracted_fields') or {}))
                 if e.get('extracted_fields') is not None else dict(e)
                 for e in entries_by_code['M2A']]

    gen = _sectioned_generator()
    gen._fill_research_support(entries_by_code, current_year=TEST_YEAR)

    assert entries_by_code['M2A'] == submitted


# --- item 17: the optional major-goals row --------------------------------------

@pytest.mark.parametrize('goals, present', [
    ('1234567890', False),   # exactly 10 characters: absent
    ('12345678901', True),   # 11 characters: present
])
def test_major_goals_row_boundary(goals, present):
    """`len(goals.strip()) > 10` is the row's boundary, pinned on both sides.

    A strict `>` -- an off-by-one here either drops a real goals row or renders
    a stub one (review thread 3932451691 item 17).
    """
    fields = {'title': 'Goals Boundary Study', 'major_goals': goals,
              'start_date': '01/2019'}
    labels = [label for label, _ in _rows(_generator()._create_grant_table(fields, 'M2A'))]
    assert ('Major project goals:' in labels) is present


@pytest.mark.parametrize('field', ['major_goals', 'description', 'narrative'])
def test_major_goals_fallback_order(field):
    """Each of the three goals keys alone fills the optional row."""
    fields = {'title': 'Goals Fallback Study', field: 'Aim 1: characterise the cohort',
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Major project goals:'] == 'Aim 1: characterise the cohort'


def test_major_goals_precedence_is_goals_then_description_then_narrative():
    """The three keys are an `or` chain, in that order."""
    fields = {'title': 'Goals Precedence Study', 'major_goals': 'Stated major goals',
              'description': 'A description here', 'narrative': 'A narrative here',
              'start_date': '01/2019'}
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Major project goals:'] \
        == 'Stated major goals'
    del fields['major_goals']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Major project goals:'] \
        == 'A description here'
    del fields['description']
    assert _cells(_generator()._create_grant_table(fields, 'M2A'))['Major project goals:'] \
        == 'A narrative here'


def test_a_goal_identical_to_the_title_is_not_rendered_twice():
    """A goal that is (normalized) the same text already rendered as the title
    is not a second fact -- it is the title read twice, once by the title
    fallback chain and once by the goals label (#829 corpus scan, BYFQBG#82:
    stage 4 put the goal text straight into `title`, and the same entry's own
    "Major Goals:" label parsed to the identical string, so the block rendered
    it twice). Comparing case-insensitively and whitespace-stripped, matching
    the title/agency and title/funding duplicate checks above.
    """
    fields = {'title': 'Map the pollinator corridors of the Example Valley',
              'major_goals': '  MAP THE POLLINATOR CORRIDORS OF THE EXAMPLE VALLEY  ',
              'start_date': '01/2019'}
    labels = [label for label, _ in _rows(_generator()._create_grant_table(fields, 'M2A'))]
    assert 'Major project goals:' not in labels


@pytest.mark.parametrize('goals, expected', [
    # Differs from the title outright.
    ('Map the pollinator corridors and survey nesting sites',
     'Map the pollinator corridors and survey nesting sites'),
    # Extends the title: the title is a literal prefix, plus more text. A
    # containment check (`title in goals`) would wrongly read this as a
    # repeat of the title rather than the equality check the guard is meant
    # to be (round-4 review: a mutant swapping `==` for `in` survived with no
    # test covering this shape).
    ('Map the pollinator corridors of the Example Valley and survey nesting sites',
     'Map the pollinator corridors of the Example Valley and survey nesting sites'),
    # A strict substring of the title is not a repeat either (`goals in title`).
    ('Map the pollinator corridors', 'Map the pollinator corridors'),
])
def test_a_goal_that_differs_from_the_title_still_renders(goals, expected):
    """Only an exact (normalized) match is suppressed -- a goal that extends or
    differs from the title is still new information and still renders.
    """
    fields = {'title': 'Map the pollinator corridors of the Example Valley',
              'major_goals': goals,
              'start_date': '01/2019'}
    cells = _cells(_generator()._create_grant_table(fields, 'M2A'))
    assert cells['Major project goals:'] == expected


def test_null_goal_and_title_fields_do_not_crash_the_goals_row():
    """stage 4 can emit a key with a null value; `fields.get(k, '')` then
    returns None, so the guard must not call .strip() on it.
    """
    no_goal = {'title': 'Map the pollinator corridors of the Example Valley',
               'narrative': None, 'start_date': '01/2019'}
    labels = [label for label, _ in _rows(_generator()._create_grant_table(no_goal, 'M2A'))]
    assert 'Major project goals:' not in labels

    no_title = {'text': None, 'agency': 'Example Fund', 'start_date': '01/2019',
                'major_goals': 'Survey nesting sites across the valley'}
    cells = _cells(_generator()._create_grant_table(no_title, 'M2A'))
    assert cells['Major project goals:'] == 'Survey nesting sites across the valley'


def test_a_goal_repeating_a_whitespace_padded_title_is_still_suppressed():
    """The guard strips both sides before comparing. `title` can carry
    surrounding whitespace (nothing upstream of `_create_grant_table` trims
    it when the raw field has no `|` for `_deduplicate_repeated_content` to
    act on) -- dropping `.strip()` on the title side alone left every test
    green (round-4 review).
    """
    fields = {'title': '  Map the pollinator corridors of the Example Valley  ',
              'major_goals': 'MAP THE POLLINATOR CORRIDORS OF THE EXAMPLE VALLEY',
              'start_date': '01/2019'}
    labels = [label for label, _ in _rows(_generator()._create_grant_table(fields, 'M2A'))]
    assert 'Major project goals:' not in labels


def test_the_goal_repeat_guard_compares_the_rendered_title_not_the_raw_field():
    """The guard's `title` is the computed value -- deduplicated, with the
    trial_title/study_title/text fallback chain already applied -- that is
    actually rendered as 'Project title:', not the raw `fields.get('title')`.
    A grant whose title comes only from a fallback field has no `'title'` key
    at all, so comparing against the raw field would never suppress a repeat
    here (round-4 review: this mutant also left every test green).
    """
    fields = {'trial_title': 'Map the pollinator corridors of the Example Valley',
              'major_goals': 'MAP THE POLLINATOR CORRIDORS OF THE EXAMPLE VALLEY',
              'start_date': '01/2019'}
    labels = [label for label, _ in _rows(_generator()._create_grant_table(fields, 'M2A'))]
    assert 'Major project goals:' not in labels


# --- #958: major goals from the source text -------------------------------------

_GOAL = 'Map the pollinator corridors of the Example Valley'


@pytest.mark.parametrize('text, expected', [
    # The plain label, tab-separated, as the last line of a table-form grant.
    (f'Award Source: | Example Fund\nThe major goals of this project are:\t{_GOAL}', _GOAL),
    # The WCM template's own label, in a goals row of its own.
    (f'(Optional - The major goals of this project are): | {_GOAL}', _GOAL),
    (f'(Optional - The major goals of this project are:) | {_GOAL}', _GOAL),
    (f'The major goals of this project are: | {_GOAL}', _GOAL),
    (f'The major goals of this project are | {_GOAL}', _GOAL),
    (f'The major goals of this project are\t{_GOAL}', _GOAL),
    (f'THE MAJOR GOALS OF THIS PROJECT ARE: {_GOAL}', _GOAL),
    # No separator: the phrase opens the faculty member's sentence, kept whole.
    ('Example Study\tThe major goals of this project are to map the corridors.',
     'The major goals of this project are to map the corridors.'),
    # ...and a paragraph-form grant's next field, the role, is not the goal.
    ('The major goals of this project are to map the corridors.\tRole: PI',
     'The major goals of this project are to map the corridors.'),
    (f'The major goals of this project are: {_GOAL}\tYour role: Co-PI', _GOAL),
    # Any other tab is a wrapped source line and stays, verbatim.
    ('The major goals of this project are:  Survey the valley; oversaw\tfield work  ',
     'Survey the valley; oversaw\tfield work'),
    # The goal ends with its line.
    (f'The major goals of this project are: {_GOAL}\nAnnual direct costs: | $5,000', _GOAL),
    # Measured wording variants (#829).
    # A bare label -- no "of (this|the) project/program" noun at all -- still
    # needs its separator to read as a label.
    (f'Major Goals: {_GOAL}', _GOAL),
    (f'Major Goals of Project: {_GOAL}', _GOAL),
    (f'Major Goals of the Project: {_GOAL}', _GOAL),
    # Singular "goal ... is", and "program" in place of "project" -- both keep
    # the sentence-form fallback the plural "goals ... are" case already has.
    ('Example Study\tThe major goal of this project is to map the pollinator corridors.',
     'The major goal of this project is to map the pollinator corridors.'),
    ('Example Study\tThe major goals of this program are to map the pollinator corridors.',
     'The major goals of this program are to map the pollinator corridors.'),
    # Singular "goal ... is" with an explicit separator -- the goal alone, not
    # the whole label sentence. "are" and "is" are both live alternatives in
    # `proj`; the plural cases above only exercise "are" before a separator,
    # and the "is" sentence-form case above never reaches `proj` for "is" at
    # all (it flows straight into `rest` whether or not `proj` matches it).
    (f'The major goal of this project is: {_GOAL}', _GOAL),
    (f'The major goals of this program is: {_GOAL}', _GOAL),
    # The "gals" typo (A5IZ6Q) with an explicit separator -- the goal only,
    # not the misspelled label.
    (f'Example Grant\tThe major gals of this project: {_GOAL}', _GOAL),
    # A stray, unanchored "major goal(s)" mention earlier in the text must not
    # shadow a real, anchored label that follows it (regression: `.search()`
    # stopped at the first mention and returned None here, dropping the real
    # label further down) -- on its own line, and on the *same* line, where a
    # naive fix (skip the unanchored match, `finditer` for the next one) still
    # fails: the unanchored match's own greedy `rest` group has already
    # swallowed the real label as part of the span being skipped.
    (f'Our major goals include improving efficiencies.\n'
     f'The major goals of this project are: {_GOAL}', _GOAL),
    (f'Major goals and aims. The major goals of this project are: {_GOAL}', _GOAL),
    # Two *anchored* labels in one text -- the first wins, matching dev's own
    # `.search()` semantics (which also stops at the first match). Nothing
    # above exercises two anchored labels together, so a selection bug that
    # picks the last one instead of the first (round-4 review) left every
    # test green.
    (f'The major goals of this project are: {_GOAL}\n'
     f'Major Goals: A later, different goal entirely', _GOAL),
])
def test_major_goals_are_parsed_verbatim_from_the_source_text(text, expected):
    assert parse_major_goals(text) == expected


@pytest.mark.parametrize('text', [
    'The major goals of this project are: |',
    '(Optional - The major goals of this project are):',
    'The major goals of this project are\nAward Source: | Example Fund',
    'Award Source: | Example Fund\nProject title: | Example Study',
    '',
    None,
    # A bare label with an empty value is still no goal.
    'Major Goals:',
    'Major Goals of Project:',
    # "Major goal(s)" with neither the "of (this|the) project/program" anchor
    # nor an explicit separator is grant content, not a label -- otherwise it
    # would be read as an unbounded whole-sentence claim.
    'Our major goals include improving efficiencies across the department.',
    'The committee highlighted major goals for the coming year during the review.',
    # A5IZ6Q's measured typo is the plural "gals"; the singular "gal" is not a
    # measured variant, so the anchor deliberately does not accept it.
    f'The major gal of this project is: {_GOAL}',
])
def test_an_empty_or_absent_goals_label_is_no_goal(text):
    """An empty label renders nothing -- and never borrows the next line."""
    assert parse_major_goals(text) is None


def test_a_grant_gains_the_goal_stated_in_its_own_text():
    grant = _entry('M2B', text=f'Award Source: | Example Fund\n'
                               f'The major goals of this project are:\t{_GOAL}',
                   title='Example Corridor Study', agency='Example Fund')

    fill_major_goals_from_text([grant])

    assert grant['extracted_fields']['major_goals'] == _GOAL


def test_a_stage4_goal_is_not_replaced_by_the_text():
    grant = _entry('M2B', text=f'The major goals of this project are: {_GOAL}',
                   major_goals='Goal as stage 4 extracted it')

    fill_major_goals_from_text([grant])

    assert grant['extracted_fields']['major_goals'] == 'Goal as stage 4 extracted it'


def test_a_goal_parsed_from_the_grants_own_text_that_repeats_its_title_does_not_double_render():
    """End-to-end shape of BYFQBG#82 (#829 blocking item 1): stage 4 set `title`
    to the goal text itself, `major_goals` was left empty, and the same entry's
    raw text carries a "Major Goals:" label after the role -- so
    `fill_major_goals_from_text` parses that label into `major_goals` with the
    identical string. Rendering must not show the goal a second time under
    "Major project goals:" once it already appears as "Project title:".
    """
    grant = _entry('M2B', text=f'Role: PI\nMajor Goals: {_GOAL}',
                   title=_GOAL, agency='Example Fund', start_date='01/2019')

    fill_major_goals_from_text([grant])
    assert grant['extracted_fields']['major_goals'] == _GOAL  # parsed, as #958 promises

    table = _generator()._create_grant_table(grant['extracted_fields'], 'M2B')
    cells = _cells(table)
    assert cells['Project title:'] == _GOAL
    assert 'Major project goals:' not in [label for label, _ in _rows(table)]


def _grant(start, end, text='grant', **fields):
    entry = _entry('M2B', text=text, title='Example Corridor Study',
                   agency='Example Fund', start_date='01/2019', **fields)
    entry.update(element_idx_start=start, element_idx_end=end)
    return entry


def _goal_row(parent_idx, goal=_GOAL):
    return {'text': f'The major goals of this project are: | {goal}'.rstrip(),
            'taxonomy_code': 'T', 'element_type': 'table_row', 'recovered_row': True,
            'parent_idx': parent_idx, 'extracted_fields': {}}


@pytest.mark.parametrize('parent_idx', [236, 237, 238])
def test_a_goals_row_inside_the_grant_range_is_claimed_by_that_grant(parent_idx):
    """Both range ends included: the real rows hang off the grant's LAST element."""
    grant, row = _grant(236, 238), _goal_row(parent_idx)
    submitted = dict(row)

    claimed = claim_goal_rows([grant], [row])

    assert claimed == [(row, grant)]
    assert grant['extracted_fields']['major_goals'] == _GOAL
    assert row == submitted


@pytest.mark.parametrize('parent_idx', [235, 239])
def test_a_goals_row_outside_every_grant_range_is_left_alone(parent_idx):
    grant = _grant(236, 238)

    assert claim_goal_rows([grant], [_goal_row(parent_idx)]) == []
    assert 'major_goals' not in grant['extracted_fields']


def test_a_goals_row_inside_two_grant_ranges_is_ambiguous_and_left_alone():
    first, second = _grant(236, 238), _grant(238, 240)

    assert claim_goal_rows([first, second], [_goal_row(238)]) == []
    assert 'major_goals' not in first['extracted_fields']


def test_a_goals_row_does_not_replace_a_different_goal_the_grant_already_has():
    grant = _grant(236, 238, major_goals='A different stated goal')

    assert claim_goal_rows([grant], [_goal_row(238)]) == []
    assert grant['extracted_fields']['major_goals'] == 'A different stated goal'


def test_a_goals_row_repeating_the_grant_own_goal_is_claimed():
    """A5IZ6Q's shape: the grant's text already carries the row's goal."""
    grant = _grant(236, 238, text=f'The major goals of this project are:\t{_GOAL}')
    row = _goal_row(238)
    fill_major_goals_from_text([grant])

    assert claim_goal_rows([grant], [row]) == [(row, grant)]


def test_an_empty_goals_row_is_not_claimed():
    grant = _grant(236, 238)

    assert claim_goal_rows([grant], [_goal_row(238, goal='')]) == []
    assert 'major_goals' not in grant['extracted_fields']


def test_a_goals_line_with_no_parent_table_is_not_claimed():
    row = _goal_row(238)
    del row['parent_idx']

    assert claim_goal_rows([_grant(236, 238)], [row]) == []


@pytest.mark.parametrize('start, end', [(None, 240), (236, None), ('236', '238')])
def test_a_grant_without_an_integer_element_range_claims_nothing(start, end):
    assert claim_goal_rows([_grant(start, end)], [_goal_row(238)]) == []


def test_section_fill_renders_the_claimed_goal_and_returns_the_row():
    grant, row = _grant(236, 238), _goal_row(238)
    t_entries = [row]

    gen = _sectioned_generator()
    claimed = gen._fill_research_support({'M2B': [grant], 'T': t_entries},
                                         current_year=TEST_YEAR)

    assert claimed == [row] and claimed[0] is row
    assert _cells(_tables_under(gen, COMPLETED)[0])['Major project goals:'] == _GOAL
    assert 'major_goals' not in grant['extracted_fields']


@pytest.mark.parametrize('code, header', [
    ('M2A', CURRENT), ('M2B', COMPLETED), ('M2C', PENDING)])
def test_section_fill_renders_the_goal_from_the_grant_own_text(code, header):
    grant = _grant(236, 238, text=f'The major goals of this project are:\t{_GOAL}')
    grant['taxonomy_code'] = code

    gen = _sectioned_generator()
    claimed = gen._fill_research_support({code: [grant]}, current_year=TEST_YEAR)

    assert claimed == []
    assert _cells(_tables_under(gen, header)[0])['Major project goals:'] == _GOAL


def test_a_row_claimed_by_a_declined_grant_is_not_returned():
    """A grant too sparse to render took no goal with it: the row stays Appendix."""
    sparse = _entry('M2B', text='x')
    sparse.update(element_idx_start=236, element_idx_end=238)

    gen = _sectioned_generator()
    claimed = gen._fill_research_support({'M2B': [sparse], 'T': [_goal_row(238)]},
                                         current_year=TEST_YEAR)

    assert claimed == []
    assert gen._declined_grant_entries


def test_a_row_claimed_by_a_grant_under_a_missing_header_is_not_returned():
    gen = _generator()  # no funding headers: nothing renders
    claimed = gen._fill_research_support({'M2B': [_grant(236, 238)], 'T': [_goal_row(238)]},
                                         current_year=TEST_YEAR)

    assert claimed == []


# --- item 18: grant duration ----------------------------------------------------

def test_grant_duration_uses_start_and_end_dates():
    """M2 dates render mm/yy-mm/yy, the format the WCM template asks for."""
    gen = _generator()
    assert gen._format_grant_duration({'start_date': '07/2019', 'end_date': '06/2021'},
                                      'M2A') == '07/19-06/21'


def test_grant_duration_falls_back_to_date():
    """Clinical-trial records carry `date` instead of `start_date`."""
    gen = _generator()
    assert gen._format_grant_duration({'date': '07/2019', 'end_date': '06/2021'},
                                      'M2A') == '07/19-06/21'
    assert gen._format_grant_duration({'date': '07/2019'}, 'M2A') == '07/19-Present'


def test_grant_duration_with_no_end_date_reads_as_present():
    """An open-ended grant renders "-Present", and an empty record renders nothing."""
    gen = _generator()
    assert gen._format_grant_duration({'start_date': '07/2019'}, 'M2A') == '07/19-Present'
    assert gen._format_grant_duration({}, 'M2A') == ''


@pytest.mark.parametrize('text, expected', [
    ('2011 Fictional Pilot Study Award', '2011'),
    ('2011- Fictional Pilot Study Award', '2011-Present'),
])
def test_grant_table_start_only_duration_reads_the_entry_source_text(text, expected):
    """Class 13 (2026-10-02): a start-only grant rendered "<year>-Present"
    whatever its source said; the Duration row now reads the entry's text."""
    fields = {'agency': 'Fictional Fund', 'title': 'Fictional Pilot Study',
              'start_date': '2011'}
    table = _generator()._create_grant_table(fields, 'M2B', entry=_entry('M2B', text=text,
                                                                         **fields))
    assert _cells(table)['Duration of support:'] == expected


@pytest.mark.parametrize('code', ['M2A', 'M2B', 'M2C'])
def test_grant_duration_is_identical_across_the_three_m2_codes(code):
    """M2A/M2B/M2C share one date format, so the bucket cannot change the string.

    `DATE_FORMATS` maps all three to 'mm/yy' (`stage6/formatting/dates.py`), so
    this pins what the code really does rather than assuming the buckets differ
    (review thread 3932451691 item 18).
    """
    fields = {'start_date': '07/2019', 'end_date': '06/2021'}
    assert _generator()._format_grant_duration(fields, code) == '07/19-06/21'


# --- item 19: the reclassification note reaches the document --------------------

def test_reclassification_note_is_attached_as_a_document_comment():
    """The past-date note is attached to the table, not merely printed.

    `_add_entry_comments` is what makes the note observable to a reviewer
    opening the CV; the existing test only read stdout (review thread
    3932451691 item 19).
    """
    gen = _sectioned_generator(emit_comments=True)
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Ended Project Study', agency='NIH',
                        start_date='01/2018', end_date='06/2020')]},
        current_year=TEST_YEAR)

    notes = [c['text'] for c in gen._comments if c['author'] == 'Reclassification']
    assert notes == ['Reclassified from Current (M2A) to Completed (M2B): '
                     f'end date 06/2020 is before {TEST_YEAR}']
    assert gen.stats['comments_added'] == 1

    # The comment must hang off the grant's own table, not off some other
    # paragraph: the reference run lives inside the reclassified table's XML.
    table = _tables_under(gen, COMPLETED)[0]
    assert 'commentReference' in table._element.xml


def test_status_rebucket_note_is_attached_as_a_document_comment():
    """The status rule's note reaches the document by the same path."""
    gen = _sectioned_generator(emit_comments=True)
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Under Review Project', agency='NIH',
                        status='Under review', start_date='01/2026')]},
        current_year=TEST_YEAR)

    notes = [c['text'] for c in gen._comments if c['author'] == 'Reclassification']
    assert notes == ["Reclassified to Pending (M2C): status is 'Under review'"]


def test_no_reclassification_leaves_no_reclassification_comment():
    """A grant nothing moved carries no note -- the comment is not unconditional."""
    gen = _sectioned_generator(emit_comments=True)
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Ongoing Project Study', agency='NIH',
                        start_date='01/2019', end_date='Present')]},
        current_year=TEST_YEAR)

    assert [c for c in gen._comments if c['author'] == 'Reclassification'] == []


# --- item 20: structural cases --------------------------------------------------

def test_missing_section_header_renders_nothing_and_does_not_raise():
    """A document without the WCM funding headers is skipped, not crashed into."""
    gen = _generator()  # no header paragraphs at all
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Homeless Project Study', agency='NIH',
                        start_date='01/2019')]},
        current_year=TEST_YEAR)

    assert gen.doc.tables == []
    assert gen.stats['tables_populated'] == 0


def test_existing_template_table_is_removed_even_when_the_category_is_empty():
    """The blank template table under a header goes even with no grants to add.

    Leaving it behind ships an empty two-column stub in the finished CV.
    """
    gen = _generator()
    gen.doc.add_paragraph().add_run(CURRENT).bold = True
    placeholder = gen.doc.add_table(rows=1, cols=2)
    placeholder.cell(0, 0).text = 'Award Source:'
    gen.doc.add_paragraph().add_run(COMPLETED).bold = True
    gen.doc.add_paragraph().add_run(PENDING).bold = True

    gen._fill_research_support({'M2A': [], 'M2B': [], 'M2C': []}, current_year=TEST_YEAR)

    assert gen.doc.tables == []


def test_existing_template_table_is_replaced_by_the_grant_tables():
    """With grants to render the placeholder still goes and only real tables remain."""
    gen = _generator()
    gen.doc.add_paragraph().add_run(CURRENT).bold = True
    placeholder = gen.doc.add_table(rows=1, cols=2)
    placeholder.cell(0, 0).text = 'Award Source:'
    gen.doc.add_paragraph().add_run(COMPLETED).bold = True
    gen.doc.add_paragraph().add_run(PENDING).bold = True

    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Replacement Project Study', agency='NIH',
                        start_date='01/2019', end_date='Present')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Replacement Project Study']
    assert len(gen.doc.tables) == 1


def test_empty_category_renders_no_table():
    """An empty bucket contributes nothing under its header."""
    gen = _sectioned_generator()
    gen._fill_research_support({'M2A': [], 'M2B': [], 'M2C': []}, current_year=TEST_YEAR)

    assert gen.doc.tables == []
    assert gen.stats['entries_inserted'] == 0


def test_grants_are_sorted_reverse_chronologically():
    """Within a bucket the most recent grant renders first."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [
            _entry('M2A', title='Older Project Study', agency='NIH',
                   start_date='01/2015', end_date='12/2028'),
            _entry('M2A', title='Newer Project Study', agency='NIH',
                   start_date='01/2021', end_date='12/2030'),
        ]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Newer Project Study', 'Older Project Study']


def test_spacing_paragraph_sits_between_grants_but_not_after_the_last():
    """Three grants produce table, spacer, table, spacer, table -- and stop."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title=f'Spaced Project {n} Study', agency='NIH',
                        start_date='01/2019', end_date=f'12/20{30 - n}')
                 for n in (1, 2, 3)]},
        current_year=TEST_YEAR)

    assert _tags_under(gen, CURRENT) == ['tbl', 'p', 'tbl', 'p', 'tbl']


def test_a_single_grant_gets_no_spacing_paragraph():
    """The "except the last one" rule means one grant gets no spacer at all."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Lonely Project Study', agency='NIH',
                        start_date='01/2019', end_date='Present')]},
        current_year=TEST_YEAR)

    assert _tags_under(gen, CURRENT) == ['tbl']


def test_each_grant_table_lands_after_its_own_section_header():
    """Every bucket's tables sit between its own header and the next one.

    Body order is the only thing that decides which WCM heading a grant appears
    under (review thread 3932451691 item 20).
    """
    gen = _sectioned_generator()
    gen._fill_research_support(
        {
            'M2A': [_entry('M2A', title='Current Project Study', agency='NIH',
                           start_date='01/2019', end_date='Present')],
            'M2B': [_entry('M2B', title='Completed Project Study', agency='NIH',
                           start_date='01/2010', end_date='12/2012')],
            'M2C': [_entry('M2C', title='Pending Project Study', agency='NIH',
                           start_date='01/2027')],
        },
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Current Project Study']
    assert _titles_under(gen, COMPLETED) == ['Completed Project Study']
    assert _titles_under(gen, PENDING) == ['Pending Project Study']


@pytest.mark.parametrize('status, expected_code', [
    ('Not funded', 'M2C'),      # kept under Pending rather than dropped
    ('Under review', 'M2C'),
    ('Completed', 'M2B'),
    ('Awarded', None),          # no move: an award is not a rebucketing signal
    ('Review completed', None),  # a process step, not an ended award (#982)
    ('Site visit completed', None),
    ('Project completed', 'M2B'),
    ('', None),
    (None, None),
])
def test_grant_status_rebucket_target_branches(status, expected_code):
    """Each branch of the status rule, including the no-move default.

    The rule the whole rebucketer is built on and it had no direct test
    (review thread 3932451691 item 20).
    """
    target, note = grant_status_rebucket_target(status)
    assert target == expected_code
    assert (note is not None) is (expected_code is not None)


def test_not_funded_grants_are_kept_under_pending_with_a_review_note():
    """"Not funded" is filed under Pending with a note, never silently dropped."""
    gen = _sectioned_generator(emit_comments=True)
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Declined Project Study', agency='NIH',
                        status='Not funded', start_date='01/2026')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, PENDING) == ['Declined Project Study']
    notes = [c['text'] for c in gen._comments if c['author'] == 'Reclassification']
    assert notes and 'confirm whether to keep this entry on the CV' in notes[0]


def test_rebucketing_leaves_a_grant_alone_when_the_status_names_its_own_bucket():
    """A completed status on an M2B grant is not a move."""
    entry = _entry('M2B', title='Completed Project Study', status='Completed')
    current, completed, pending, messages = rebucket_grants_by_status([], [entry], [])

    assert (current, completed, pending) == ([], [entry], [])
    assert messages == []


def test_unsupported_rebucket_target_raises_a_named_error(monkeypatch):
    """An unrenderable bucket code fails with a domain error, not a bare KeyError.

    `bucket_lists` holds only M2B and M2C. Before the guard, a status rule
    returning anything else died on `bucket_lists[target]` naming neither the
    code nor the grant (thread 3932312407 item 2). Preventive: the real rule
    cannot return M2D today, so the rule is stubbed to prove the guard.

    The message identifies the grant by its position, never by its title: it
    reaches a traceback and from there the pod logs, and a grant title is CV
    content. That is this module's stated logging rule, and an exception is
    the one place it would be easy to break by accident.
    """
    monkeypatch.setattr(research_support, 'grant_status_rebucket_target',
                        lambda status: ('M2D', 'stub note') if status else (None, None))
    entry = _entry('M2A', title='Misrouted Project Study', status='Under review')

    with pytest.raises(UnsupportedRebucketTargetError) as excinfo:
        rebucket_grants_by_status([entry], [], [])

    message = str(excinfo.value)
    assert 'M2D' in message
    assert 'M2A' in message and '#0' in message
    assert 'Misrouted Project Study' not in message


def test_reclassify_past_m2a_grants_leaves_its_inputs_alone():
    """The pure classifier returns new lists rather than editing the caller's."""
    entry = _entry('M2A', title='Ended Project Study', end_date='06/2020')
    m2a = [entry]
    m2b = []

    current, completed, messages = reclassify_past_m2a_grants(m2a, m2b, TEST_YEAR)

    assert m2a == [entry] and m2b == []
    assert current == [] and completed == [entry]
    assert messages == ["  Reclassified to M2B: 'Ended Project Study...' (ended 2020)"]


# --- #981: the heading is the status when stage 4 extracted none ---------------

def _under(heading, code='M2A', **fields):
    """A grant record filed under a hierarchy heading, with no status field."""
    entry = _entry(code, **fields)
    entry['hierarchy'] = list(heading)
    return entry


@pytest.mark.parametrize('heading, expected_code', [
    (['GRANT SUPPORT', 'Pending applications'], 'M2C'),
    (['GRANTS', 'GRANT APPLICATIONS IN REVIEW'], 'M2C'),
    (['Non-funded applications'], 'M2C'),
    (['GRANTS', 'GRANT APPLICATIONS AWAITING FINAL ADMINISTRATIVE APPROVAL'], 'M2C'),
    (['NOT FUNDED'], 'M2C'),
    (['Submitted, Not Funded'], 'M2C'),
    (['Grants', 'Completed Research Support'], 'M2B'),
    (['Current Grant Support'], None),       # "Grant Support" alone is no status
    (['Research Support', 'Active'], None),
    (['Current and Pending Support'], None),  # names two buckets: silent
    (['Past and Present Funding'], None),
    (['Pending and Completed Grants'], None),
    (['Impending Renewals'], None),           # whole words only
    (['Reunfunded Items'], None),             # no boundary before the word
    (['Unfundedness Report'], None),          # no boundary after the word
    (['Pendingx Applications'], None),        # no boundary after a pending word
    ([], None),
])
def test_grant_heading_rebucket_target(heading, expected_code):
    target, note = grant_heading_rebucket_target(heading)
    assert target == expected_code
    assert (note is not None) is (expected_code is not None)


def test_heading_note_says_the_text_came_from_the_heading():
    _, note = grant_heading_rebucket_target(['Pending applications'])
    assert note == "Reclassified to Pending (M2C): section heading is 'Pending applications'"


def test_a_grant_without_a_status_moves_on_its_heading():
    """3b coded it completed; the CV filed it under "Pending applications"."""
    entry = _under(['Pending applications'], code='M2B', title='Filed Pending')

    current, completed, pending, _ = rebucket_grants_by_status([], [entry], [])

    assert (current, completed, pending) == ([], [], [entry])
    assert entry['reclassification_note'].startswith('Reclassified to Pending (M2C)')


def test_a_status_field_beats_the_heading():
    entry = _under(['Pending applications'], code='M2B', title='Awarded',
                   status='Completed')

    current, completed, pending, _ = rebucket_grants_by_status([], [entry], [])

    assert (current, completed, pending) == ([], [entry], [])


def test_not_funded_heading_keeps_the_grant_under_pending_with_a_review_note():
    gen = _sectioned_generator(emit_comments=True)
    gen._fill_research_support(
        {'M2B': [_under(['NOT FUNDED'], code='M2B', title='Declined Heading Study',
                        agency='NIH', start_date='01/2020')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, PENDING) == ['Declined Heading Study']
    notes = [c['text'] for c in gen._comments if c['author'] == 'Reclassification']
    assert notes and 'confirm whether to keep this entry on the CV' in notes[0]


def test_a_current_grant_support_heading_moves_nothing():
    entry = _under(['Current Grant Support'], title='Running Study')

    current, completed, pending, messages = rebucket_grants_by_status([entry], [], [])

    assert (current, completed, pending, messages) == ([entry], [], [], [])


# --- #981: a completed-coded grant that is still running is current ------------

@pytest.mark.parametrize('end_date', ['present', 'Ongoing', '12/31/2026', '06/2028'])
def test_an_m2b_grant_still_running_is_promoted_to_current(end_date):
    entry = _under(['Grants'], code='M2B', title='Running Study', end_date=end_date)

    current, completed, messages = promote_open_ended_m2b_grants([], [entry], TEST_YEAR)

    assert (current, completed) == ([entry], [])
    assert entry['reclassification_note'].startswith('Reclassified from Completed (M2B)')
    assert messages == [f"  Reclassified to M2A: 'Running Study...' (ends {end_date})"]


@pytest.mark.parametrize('fields', [
    {'end_date': '12/31/2025'},               # ended last year
    {'end_date': ''},                         # nothing says it is running
    {},
    {'end_date': '2028', 'status': 'Completed'},  # an explicit status beats the date
])
def test_an_m2b_grant_that_is_not_running_stays_completed(fields):
    entry = _under(['Grants'], code='M2B', title='Past Study', **fields)

    current, completed, messages = promote_open_ended_m2b_grants([], [entry], TEST_YEAR)

    assert (current, completed, messages) == ([], [entry], [])


@pytest.mark.parametrize('heading', [
    ['GRANT SUPPORT', 'PAST GRANT SUPPORT'],
    ['GRANTS', 'Summary of Major Previous Grants'],
    ['Grants: Prior'],
])
def test_an_m2b_grant_under_a_past_heading_is_not_promoted_on_an_open_end_date(heading):
    """"ongoing" in a grant filed under "Past Grant Support" does not outvote the heading."""
    entry = _under(heading, code='M2B', title='Past Study', end_date='ongoing')

    current, completed, _ = promote_open_ended_m2b_grants([], [entry], TEST_YEAR)

    assert (current, completed) == ([], [entry])


def test_an_m2b_grant_under_a_past_and_present_heading_is_promoted_on_its_date():
    entry = _under(['Past and Present Funding'], code='M2B', title='Running Study',
                   end_date='2028')

    current, completed, _ = promote_open_ended_m2b_grants([], [entry], TEST_YEAR)

    assert (current, completed) == ([entry], [])


def test_an_m2b_grant_whose_heading_says_completed_stays_completed():
    entry = _under(['Completed Grants'], code='M2B', title='Past Study', end_date='2028')

    current, completed, _ = promote_open_ended_m2b_grants([], [entry], TEST_YEAR)

    assert (current, completed) == ([], [entry])


def test_a_grant_already_moved_by_a_rule_is_not_promoted_back():
    """A status rebucket or past-date demotion leaves a note; that move stands."""
    entry = _under(['Grants'], code='M2B', title='Moved', end_date='2028')
    entry['reclassification_note'] = 'Reclassified from Current (M2A) to Completed (M2B)'

    current, completed, _ = promote_open_ended_m2b_grants([], [entry], TEST_YEAR)

    assert (current, completed) == ([], [entry])


def test_promoting_leaves_its_inputs_alone():
    entry = _under(['Grants'], code='M2B', title='Running Study', end_date='present')
    m2a, m2b = [], [entry]

    promote_open_ended_m2b_grants(m2a, m2b, TEST_YEAR)

    assert m2a == [] and m2b == [entry]


def test_fill_research_support_files_a_running_completed_coded_grant_under_current():
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2B': [_under(['Grants'], code='M2B', title='Running Study', agency='NIH',
                        start_date='03/2024', end_date='12/2028')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, CURRENT) == ['Running Study']
    assert _titles_under(gen, COMPLETED) == []


# --- thread 3932312407 item 8: the rendering contract ---------------------------

_CONTRACT_FIELDS = {
    'agency': 'National Institutes of Health',
    'title': 'Metabolic Reprogramming in Heart Failure',
    'annual_direct_costs': '100000',
    'non_financial_support': 'Core facility access',
    'start_date': '07/2019',
    'end_date': '06/2021',
    'pi_name': 'Jane Smith',
    'pi_role': 'Principal Investigator',
    'percent_effort': '20%',
}

_CONTRACT_ROWS = [
    ('Award Source:', 'National Institutes of Health'),
    ('Project title:', 'Metabolic Reprogramming in Heart Failure'),
    ('Annual direct costs:', '$100,000'),
    ('Non-financial support:', 'Core facility access'),
    ('Duration of support:', '07/19-06/21'),
    ('Name of Principal Investigator:', 'Jane Smith'),
    ('Your role:', 'Principal Investigator'),
    ('Your percent (%) effort:', '20%'),
]


def test_grant_table_is_exactly_the_eight_wcm_rows_in_order():
    """The WCM grant block: eight rows, these labels, this order, these values.

    The row list is written inline in `_create_grant_table`, so a field change
    could add, drop or reorder a row of the WCM structure with nothing failing
    (thread 3932312407 item 8).
    """
    table = _generator()._create_grant_table(dict(_CONTRACT_FIELDS), 'M2A')

    assert len(table.rows) == 8
    assert _rows(table) == _CONTRACT_ROWS


def test_major_goals_is_the_ninth_row_and_only_when_present():
    """The optional ninth row appends after the eight; nothing else shifts."""
    fields = dict(_CONTRACT_FIELDS, major_goals='Aim 1: characterise the cohort')
    table = _generator()._create_grant_table(fields, 'M2A')

    assert len(table.rows) == 9
    assert _rows(table) == _CONTRACT_ROWS + [
        ('Major project goals:', 'Aim 1: characterise the cohort')]


def test_absent_values_render_as_empty_cells_not_missing_rows():
    """A sparse-but-admissible grant still renders all eight rows, some empty."""
    fields = {'title': 'Minimal Project Study', 'start_date': '01/2019'}
    table = _generator()._create_grant_table(fields, 'M2A')

    assert [label for label, _ in _rows(table)] == [label for label, _ in _CONTRACT_ROWS]
    assert _cells(table)['Non-financial support:'] == ''


def _fields_get_keys():
    """Every literal key this module reads off a `fields` mapping."""
    tree = ast.parse(Path(research_support.__file__).read_text(encoding='utf-8'))
    keys = set()
    for node in ast.walk(tree):
        if (isinstance(node, ast.Call)
                and isinstance(node.func, ast.Attribute)
                and node.func.attr == 'get'
                and isinstance(node.func.value, ast.Name)
                and node.func.value.id == 'fields'
                and node.args
                and isinstance(node.args[0], ast.Constant)
                and isinstance(node.args[0].value, str)):
            keys.add(node.args[0].value)
    return keys


def test_every_field_the_module_reads_is_declared_on_the_grant_record_type():
    """A `fields.get('x')` that `GrantFields` does not name makes the diagnostic lie.

    `CONSUMED_GRANT_FIELDS` is derived from `GrantFields` and drives the
    debug line that reports which stage-4 keys reached no row. A slot added to
    `_create_grant_table` without its key added to the record type would leave
    that key reported as dropped forever -- the diagnostic reporting a loss
    that is not happening, which is worse than no diagnostic. Nothing enforced
    the pairing before; this does, off the module's own source.
    """
    keys = _fields_get_keys()

    assert keys, 'the AST walk found no fields.get(...) calls at all'
    assert keys <= research_support.CONSUMED_GRANT_FIELDS, (
        'read but not declared on GrantFields: '
        f'{sorted(keys - research_support.CONSUMED_GRANT_FIELDS)}')


def test_the_grant_record_type_declares_nothing_the_module_never_reads():
    """The other direction: a declared key no reader wants is dead contract.

    All 30 keys, `status` included, are reached through a `fields.get(...)` in
    this module -- `status` from the bucket rules rather than from a rendered
    row. A key left on the record type after its reader is deleted would go on
    suppressing that key's line in the unconsumed-fields diagnostic, silently.
    """
    declared_but_unread = research_support.CONSUMED_GRANT_FIELDS - _fields_get_keys()

    assert declared_but_unread == set(), sorted(declared_but_unread)
    assert len(research_support.CONSUMED_GRANT_FIELDS) == 30


# --- #291: clinical trials render in the grant block ------------------------------

def test_a_trial_record_renders_its_nct_id_sponsor_and_role():
    """Stage 4 maps a trial onto the grant fields (NCT id -> grant_number,
    sponsor -> agency), so all three must reach the table."""
    fields = {'title': 'Phase II trial of an invented compound', 'grant_number': 'NCT00000001',
              'agency': 'Invented Pharma', 'pi_role': 'Site PI', 'start_date': '2019'}
    text = ' '.join(' '.join(r) for r in _rows(_generator()._create_grant_table(fields, 'M2A')))
    assert 'NCT00000001' in text and 'Invented Pharma' in text and 'Site PI' in text


def test_a_stored_pre_291_trial_record_keeps_its_nct_id():
    """Stage-4 output from before #291 carries the id as `nct_number`."""
    fields = {'trial_title': 'Phase III trial of an invented device', 'nct_number': 'NCT00000002',
              'sponsor': 'Invented Devices', 'start_date': '2018'}
    text = ' '.join(' '.join(r) for r in _rows(_generator()._create_grant_table(fields, 'M2A')))
    assert 'NCT00000002' in text and 'Invented Devices' in text


# --- #982: stage 4 now keeps status and notes, and the block renders them --------

def test_status_and_notes_render_as_their_own_rows_after_the_effort_row():
    """web39's withdrawn grants lost the status word: no row named `status`."""
    fields = {'title': 'Withdrawn Cohort Study', 'agency': 'NIH', 'start_date': '01/2019',
              'percent_effort': '10%', 'status': 'withdrawn', 'notes': 'Sponsor closed the call'}
    rows = _rows(_generator()._create_grant_table(fields, 'M2B'))
    labels = [label for label, _ in rows]
    assert rows[labels.index('Status:')] == ('Status:', 'withdrawn')
    assert rows[labels.index('Notes:')] == ('Notes:', 'Sponsor closed the call')
    assert labels.index('Your percent (%) effort:') < labels.index('Status:') < labels.index('Notes:')


def test_a_grant_with_no_status_or_notes_keeps_the_eight_row_template_block():
    fields = {'title': 'Funded Cohort Study', 'start_date': '01/2019', 'status': None, 'notes': '  '}
    table = _generator()._create_grant_table(fields, 'M2A')
    assert len(table.rows) == 8
    assert 'Status:' not in _cells(table) and 'Notes:' not in _cells(table)


def test_a_note_that_only_repeats_the_title_is_dropped():
    fields = {'title': 'Funded Cohort Study', 'notes': 'funded cohort study', 'start_date': '01/2019'}
    assert 'Notes:' not in _cells(_generator()._create_grant_table(fields, 'M2A'))


def test_withdrawn_status_reaches_the_document_through_the_section_fill():
    """The wire: `_fill_research_support` (bucket rules included) to the table cell."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2B': [_entry('M2B', title='Withdrawn Cohort Study', agency='NIH',
                        status='withdrawn', notes='Update: withdrawn', start_date='01/2020',
                        end_date='06/2021')]},
        current_year=TEST_YEAR)
    (table,) = _tables_under(gen, COMPLETED)
    assert _cells(table)['Status:'] == 'withdrawn'
    assert _cells(table)['Notes:'] == 'Update: withdrawn'


# --- #982: an unrecognised status falls back to the heading --------------------

@pytest.mark.parametrize('status', ['withdrawn', 'Funded', 'NCE', 'Awarded 2021'])
def test_an_unrecognised_status_falls_back_to_the_heading(status):
    """The vocabulary knows none of these; they must not silence the #981 heading rule."""
    entry = _under(['Pending applications'], code='M2B', title='Filed Pending', status=status)

    assert research_support.explicit_status_target(entry)[0] == 'M2C'
    current, completed, pending, _ = rebucket_grants_by_status([], [entry], [])
    assert (current, completed, pending) == ([], [], [entry])


def test_an_unrecognised_status_under_a_silent_heading_moves_nothing():
    entry = _under(['Current Grant Support'], title='Running Study', status='withdrawn')

    assert research_support.explicit_status_target(entry) == (None, None)


def test_a_review_completed_status_does_not_file_a_pending_grant_as_completed():
    """Judgement call: "Review completed" names no bucket, so the Pending heading decides."""
    entry = _under(['Pending applications'], title='Under Review Study',
                   status='Review completed')

    target, note = research_support.explicit_status_target(entry)

    assert target == 'M2C'
    assert 'section heading' in note


def test_a_recognised_status_beats_a_disagreeing_heading():
    """Judgement call: the status is the grant's own word, so it wins a disagreement."""
    pending_status = _under(['Completed Research Support'], code='M2B',
                            title='Resubmitted', status='Under review')
    completed_status = _under(['Pending applications'], code='M2C',
                              title='Ended', status='Project completed')

    assert research_support.explicit_status_target(pending_status)[0] == 'M2C'
    assert research_support.explicit_status_target(completed_status)[0] == 'M2B'


# --- #982: co_investigators renders in its own row -----------------------------

def test_co_investigators_render_in_their_own_row_and_leave_the_pi_cell_alone():
    fields = {'title': 'Shared Cohort Study', 'pi_name': 'Ada Lovelace-Test',
              'co_investigators': 'Tanaka, CoI; Reyes, CoTwo', 'start_date': '01/2019'}
    rows = _rows(_generator()._create_grant_table(fields, 'M2A'))
    labels = [label for label, _ in rows]

    assert rows[labels.index('Co-Investigators:')] == (
        'Co-Investigators:', 'Tanaka, CoI; Reyes, CoTwo')
    assert rows[labels.index('Name of Principal Investigator:')][1] == 'Ada Lovelace-Test'
    assert labels.index('Your percent (%) effort:') < labels.index('Co-Investigators:')


def test_co_investigators_identical_to_the_pi_are_not_repeated():
    fields = {'title': 'Solo Study', 'pi_name': 'Tanaka, CoI',
              'co_investigators': 'tanaka, coi', 'start_date': '01/2019'}

    assert 'Co-Investigators:' not in _cells(_generator()._create_grant_table(fields, 'M2A'))


def test_no_co_investigators_keeps_the_eight_row_block():
    fields = {'title': 'Solo Study', 'co_investigators': None, 'start_date': '01/2019'}

    assert len(_generator()._create_grant_table(fields, 'M2A').rows) == 8


# --- EBYSBC E7: an empty-section label is no status ----------------------------

@pytest.mark.parametrize('status', [
    'PENDING \u2013 none', 'Pending - none.', 'Pending: None', 'CURRENT \u2014 none',
    'none', 'None.', 'N/A', 'n/a',
])
def test_an_empty_section_label_is_no_status(status):
    """A CV's "PENDING - none" label, left inside the next grant's lines, names no bucket."""
    assert grant_status_is_empty_section_label(status)
    assert grant_status_rebucket_target(status) == (None, None)


@pytest.mark.parametrize('status', [
    'Pending', 'Submitted/Under review', 'Pending (none yet)', 'Funded', 'No-cost extension',
    'Not funded', 'Nonetheless pending', '',
])
def test_a_real_status_is_not_an_empty_section_label(status):
    assert not grant_status_is_empty_section_label(status)


def test_a_funded_grant_carrying_an_empty_pending_label_stays_completed():
    """The E7 shape: 3b filed an ended, funded grant under an Active heading as
    completed; its status is the CV's "PENDING - none" label, which must not
    move it to Pending."""
    entry = _under(['Research Support', 'Active'], code='M2B', title='Funded Cohort Study',
                   status='PENDING \u2013 none', total_funding='900,000',
                   start_date='2014', end_date='2019')

    current, completed, pending, messages = rebucket_grants_by_status(
        [], [entry], [])

    assert (current, completed, pending, messages) == ([], [entry], [], [])


def test_an_empty_section_label_renders_no_status_row():
    label_only = {'title': 'Funded Cohort Study', 'status': 'PENDING \u2013 none',
                  'start_date': '01/2015'}
    real_status = {'title': 'Funded Cohort Study', 'status': 'Pending',
                   'start_date': '01/2015'}

    assert 'Status:' not in _cells(_generator()._create_grant_table(label_only, 'M2B'))
    assert _cells(_generator()._create_grant_table(real_status, 'M2C'))['Status:'] == 'Pending'


def test_an_empty_pending_label_reaches_neither_bucket_nor_row_through_the_section_fill():
    """The wire: `_fill_research_support` keeps the grant under Past with no Status row."""
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2B': [_under(['Research Support', 'Active'], code='M2B',
                        title='Funded Cohort Study', agency='NIH',
                        status='PENDING \u2013 none', total_funding='900,000',
                        start_date='2014', end_date='2019')]},
        current_year=TEST_YEAR)

    (table,) = _tables_under(gen, COMPLETED)
    assert _cells(table)['Project title:'] == 'Funded Cohort Study'
    assert 'Status:' not in _cells(table)
    assert _tables_under(gen, PENDING) == []


# --- EBYSBC E7: an explicit pending status still beats an amount and old dates --

@pytest.mark.parametrize('status', ['Under review', 'Submitted', 'Pending'])
@pytest.mark.parametrize('heading', [('Research Support',), ('Research Support', 'Active')])
@pytest.mark.parametrize('code', ['M2A', 'M2B'])
def test_an_explicit_pending_status_moves_an_amounted_ended_grant_to_pending(
        status, heading, code):
    """#210: an explicit status beats date inference. Stage 6 has no award signal
    independent of the 3b code: an M2A/M2B schema's `total_funding` is filled for
    an application too, so an amount and an ended range under a generic heading
    must not keep a grant whose own status says it awaits a decision out of
    Pending (that would render a requested budget as Total award)."""
    entry = _under(list(heading), code=code, title='Old Application Study',
                   status=status, total_funding='$250,000',
                   start_date='2010', end_date='2013')
    buckets = {'M2A': [], 'M2B': []}
    buckets[code].append(entry)

    current, completed, pending, _ = rebucket_grants_by_status(
        buckets['M2A'], buckets['M2B'], [])

    assert (current, completed, pending) == ([], [], [entry])


# --- EBYSBC E7: a two-digit-year end date is read --------------------------------

@pytest.mark.parametrize('end_date, expected', [
    ('4/15/14', 2014),
    ('03/31/11', 2011),
    ('9-30-28', 2028),
    ('6/30/31', 2031),
    ('8/31/36', 2036),
    ('8/31/37', 1937),
    ('6/30/95', 1995),
    ('13/45/14', None),
    ('2016-10-21', 2016),
    ('June 2015', 2015),
    ('2015-16', 2015),
    ('7/31/1', None),
    ('12/13', None),
    ('present', None),
    ('', None),
    (None, None),
])
def test_grant_end_year(end_date, expected):
    """A two-digit year reads against TEST_YEAR (2026) plus the ten-year horizon,
    not the shared past-leaning pivot of 30."""
    assert grant_end_year(end_date, TEST_YEAR) == expected


@pytest.mark.parametrize('end_date', ['6/30/31', '8/31/32'])
def test_a_current_grant_ending_past_the_shared_pivot_stays_current(end_date):
    """A five-year award starting mid-2026 ends 6/30/31: 2031, not 1931."""
    running = _entry('M2A', title='New Award Study', start_date='7/1/26', end_date=end_date)

    current, completed, messages = reclassify_past_m2a_grants([running], [], TEST_YEAR)

    assert (current, completed, messages) == ([running], [], [])


def test_a_current_grant_with_a_two_digit_year_end_in_the_past_moves_to_completed():
    ended = _entry('M2A', title='Ended Study', start_date='4/1/09', end_date='4/15/14')
    running = _entry('M2A', title='Running Study', start_date='10/1/24', end_date='9/30/28')

    current, completed, messages = reclassify_past_m2a_grants([ended, running], [], TEST_YEAR)

    assert (current, completed) == ([running], [ended])
    assert ended['reclassification_note'].endswith(f'4/15/14 is before {TEST_YEAR}')


def test_a_completed_grant_with_a_two_digit_year_end_still_running_is_promoted():
    running = _entry('M2B', title='Running Study', start_date='10/1/24', end_date='9/30/28')

    current, completed, _ = promote_open_ended_m2b_grants([], [running], TEST_YEAR)

    assert (current, completed) == ([running], [])


# --- #1343 (RCBKFG CAOACN 589): a one-digit end year reads off the start ----

@pytest.mark.parametrize('start_date, end_date, expected', [
    ('9/1/07', '8/30/1', 2011),
    ('2007-09-01', '8/30/1', 2011),
    ('9/1/07', '8/30/9', 2009),
    ('9/1/17', '8/30/1', 2021),
    ('9/1/07', '13/30/1', None),
    ('', '8/30/1', None),
    (None, '8/30/1', None),
    ('9/1/7', '8/30/1', None),
])
def test_grant_end_year_reads_a_one_digit_year_against_the_start(start_date, end_date, expected):
    assert grant_end_year(end_date, TEST_YEAR, start_date) == expected


def test_a_current_grant_with_a_one_digit_year_end_in_the_past_moves_to_completed():
    ended = _entry('M2A', title='Ended Study', start_date='9/1/07', end_date='8/30/1')

    current, completed, _ = reclassify_past_m2a_grants([ended], [], TEST_YEAR)

    assert (current, completed) == ([], [ended])


def test_a_completed_grant_with_a_one_digit_year_end_still_running_is_promoted():
    running = _entry('M2B', title='Running Study', start_date='7/1/25', end_date='6/30/9')

    current, completed, _ = promote_open_ended_m2b_grants([], [running], TEST_YEAR)

    assert (current, completed) == ([running], [])


def test_a_two_digit_year_end_reaches_past_funding_through_the_section_fill():
    gen = _sectioned_generator()
    gen._fill_research_support(
        {'M2A': [_entry('M2A', title='Ended Study', agency='NIH',
                        start_date='4/1/09', end_date='4/15/14')]},
        current_year=TEST_YEAR)

    assert _titles_under(gen, COMPLETED) == ['Ended Study']
    assert _tables_under(gen, CURRENT) == []
