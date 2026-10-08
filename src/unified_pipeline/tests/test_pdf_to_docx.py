"""Tests for core/pdf_to_docx.py and scripts/pdf_to_docx.py (#806).

No PDF writer is a dependency, so fixtures are built from hand-written PDF
bytes (`_make_pdf`): base-14 Helvetica / Helvetica-Bold / Courier, optional
1x1 image.
"""

import hashlib
import struct
import subprocess
import sys
from pathlib import Path

import pdfplumber
import pytest
from docx import Document

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.core.docx_structure_extractor import extract_unified_elements
from unified_pipeline.core.pdf_to_docx import (
    GUTTER_EDGE_FRAC,
    LIST_MARKER_RE,
    _blank_gap_em,
    _column_gaps,
    _is_bold,
    _Line,
    _Para,
    _Run,
    _runs,
    _write_docx,
    convert_pdf_to_docx,
)
from unified_pipeline.run_doctor import iter_header_candidates

_REPO = Path(__file__).resolve().parents[3]
_CLI = _REPO / "scripts" / "pdf_to_docx.py"

_LONG = ("Completed the doctoral programme and residency with a thesis on the "
         "long term outcomes of")
_LONG_TAIL = "surgical patients in the region."
# Same width as _LONG, but the last word is not a connector ("of").
_LONG_STOP = _LONG[:-2] + "at"


_PDF_PAD = bytes.fromhex("28BF4E5E4E758A4164004E56FFFA01082E2E00B6D0683E802F0CA9FE6453697A")


def _rc4(key: bytes, data: bytes) -> bytes:
    box, j = list(range(256)), 0
    for i in range(256):
        j = (j + box[i] + key[i % len(key)]) % 256
        box[i], box[j] = box[j], box[i]
    out, i, j = bytearray(), 0, 0
    for byte in data:
        i = (i + 1) % 256
        j = (j + box[i]) % 256
        box[i], box[j] = box[j], box[i]
        out.append(byte ^ box[(box[i] + box[j]) % 256])
    return bytes(out)


def _encrypt_dict(user_password: str, obj_num: int) -> tuple[bytes, bytes]:
    """PDF standard security handler, revision 2 (RC4-40): an /Encrypt object
    body and the trailer entries that make pdfminer demand `user_password`."""
    pad = lambda pw: (pw.encode() + _PDF_PAD)[:32]  # noqa: E731
    owner = _rc4(hashlib.md5(pad("owner-pw")).digest()[:5], pad(user_password))
    perms, doc_id = -4, b"0123456789abcdef"
    key = hashlib.md5(pad(user_password) + owner + struct.pack("<i", perms) + doc_id).digest()[:5]
    user = _rc4(key, _PDF_PAD)
    body = (f"<< /Filter /Standard /V 1 /R 2 /P {perms} /O <{owner.hex()}> "
            f"/U <{user.hex()}> >>").encode()
    return body, f"/Encrypt {obj_num} 0 R /ID [<{doc_id.hex()}> <{doc_id.hex()}>]".encode()


#: `bold` in a `_make_pdf` op picks the font: False Helvetica, True
#: Helvetica-Bold, COURIER Courier.
COURIER = "courier"
_FONT_NUMBER = {False: 1, True: 2, COURIER: 3}


def _make_pdf(pages, image_pages=(), user_password=None) -> bytes:
    """pages: list of lists of (bold, size, x, y, text). image_pages: 0-based
    indexes of pages that also carry a 1x1 image."""
    objs: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        4: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        5: (b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 "
            b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\n"
            b"stream\n\x80\nendstream"),
        6: b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
    }
    kids = []
    for i, ops in enumerate(pages):
        page_id, content_id = 7 + 2 * i, 8 + 2 * i
        kids.append(f"{page_id} 0 R")
        body = "".join(
            f"BT /F{_FONT_NUMBER[bold]} {size} Tf {x} {y} Td ({text}) Tj ET\n"
            for bold, size, x, y, text in ops)
        if i in image_pages:
            body += "q 200 0 0 200 100 300 cm /Im0 Do Q\n"
        stream = body.encode("latin-1")
        objs[content_id] = (b"<< /Length %d >>\nstream\n" % len(stream)
                            + stream + b"\nendstream")
        objs[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R "
            f"/F2 4 0 R /F3 6 0 R >> /XObject << /Im0 5 0 R >> >> >>").encode()
    trailer_extra = b""
    if user_password is not None:
        # xref subsections are contiguous, so the object takes the next number.
        objs[len(objs) + 2], trailer_extra = _encrypt_dict(user_password, len(objs) + 2)
    objs[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for num in sorted(objs):
        out += b"%010d 00000 n \n" % offsets[num]
    out += b"trailer\n<< /Size %d /Root 1 0 R %s >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objs) + 1, trailer_extra, xref)
    return bytes(out)


def _convert(tmp_path, pages, image_pages=(), name="cv"):
    pdf, docx = tmp_path / f"{name}.pdf", tmp_path / f"{name}.docx"
    pdf.write_bytes(_make_pdf(pages, image_pages))
    report = convert_pdf_to_docx(pdf, docx)
    return report, Document(str(docx))


def _cv_page():
    """Header, a two-line wrapped paragraph, a gap, a column-aligned row."""
    return [
        (True, 12, 72, 700, "EDUCATION"),
        (False, 10, 72, 680, _LONG),
        (False, 10, 72, 668, _LONG_TAIL),
        (False, 10, 72, 630, "2019"),
        (False, 10, 260, 630, "Award"),
    ]


def test_bold_header_is_a_bold_run_with_size(tmp_path):
    _, doc = _convert(tmp_path, [_cv_page()])
    first = doc.paragraphs[0]
    assert first.text == "EDUCATION"
    assert all(r.bold for r in first.runs)
    assert first.runs[0].font.size.pt == 12
    assert not doc.paragraphs[1].runs[0].bold


def test_wrapped_lines_merge_into_one_paragraph(tmp_path):
    _, doc = _convert(tmp_path, [_cv_page()])
    assert doc.paragraphs[1].text == f"{_LONG} {_LONG_TAIL}"


def test_vertical_gap_emits_blank_paragraph(tmp_path):
    report, doc = _convert(tmp_path, [_cv_page()])
    assert [p.text for p in doc.paragraphs][2] == ""
    assert report.blank_paragraphs == 1
    assert report.paragraphs == 4


def test_column_gap_becomes_tab(tmp_path):
    _, doc = _convert(tmp_path, [_cv_page()])
    assert doc.paragraphs[3].text == "2019\tAward"


def test_normal_word_gap_stays_a_space(tmp_path):
    _, doc = _convert(tmp_path, [[(False, 10, 72, 700, "Assistant Professor")]])
    assert doc.paragraphs[0].text == "Assistant Professor"


def test_short_lines_do_not_merge(tmp_path):
    page = [(False, 10, 72, 700, _LONG), (False, 10, 72, 688, "Short entry"),
            (False, 10, 72, 676, "Another short entry")]
    _, doc = _convert(tmp_path, [page])
    assert [p.text for p in doc.paragraphs] == [f"{_LONG} Short entry", "Another short entry"]


def _furniture_pages():
    pages = [
        [(False, 9, 72, 765, "Confidential CV"), (False, 10, 72, 600, f"Body {n}"),
         (False, 9, 280, 25, f"Page {n} of 3")]
        for n in (1, 2, 3)]
    pages[0].append((False, 10, 72, 700, _LONG))  # sets the right margin
    return pages


def test_running_header_and_page_number_removed(tmp_path):
    """The running header keeps its page-1 copy; page numbers all go."""
    report, doc = _convert(tmp_path, _furniture_pages())
    assert [p.text for p in doc.paragraphs if p.text] == ["Confidential CV", _LONG, "Body 1", "Body 2", "Body 3"]
    assert report.pages == 3


def test_top_band_repeat_keeps_its_first_copy_and_page_numbers_all_go(tmp_path):
    """A name at the top of 3 pages is kept once, on page 1; a page number
    in the bottom band is dropped on every page."""
    pages = [[(False, 14, 72, 750, "Jane Q. Public"), (False, 10, 72, 600, f"Body {n}"),
              (False, 9, 280, 25, f"Page {n} of 3")] for n in (1, 2, 3)]
    _, doc = _convert(tmp_path, pages)
    texts = [t for t in _texts(doc) if t]
    assert texts[0] == "Jane Q. Public"
    assert " ".join(texts).count("Jane Q. Public") == 1
    assert "Page" not in " ".join(texts)


def test_repeated_body_line_is_kept(tmp_path):
    pages = [[(False, 10, 72, 400, "Present")] for _ in range(3)]
    _, doc = _convert(tmp_path, pages)
    assert " ".join(p.text for p in doc.paragraphs).split() == ["Present"] * 3


def _footer_page(number, two_column):
    """A body, and a footer printed as "November 2020", a column gap, and
    the page number."""
    body = _two_column_page(right_offset=5) if two_column else [(False, 10, 72, 600, f"Body {number}")]
    return body + [(False, 9, 72, 25, "November 2020"), (False, 9, 500, 25, str(number))]


def test_running_footer_cut_by_a_gutter_is_still_furniture(tmp_path):
    """RLADNC (#1584): the footer prints whole on pages 1 and 3; the gutter
    of two-column page 2 cuts it into "November 2020" and "2", each on
    one page only. Both parts still match the whole footer's copies."""
    pages = [_footer_page(1, False), _footer_page(2, True), _footer_page(3, False)]
    _, doc = _convert(tmp_path, pages)
    texts = [t for t in _texts(doc) if t]
    assert "November" not in " ".join(texts) and "2" not in texts
    assert _order(doc, "Side", "Main") == ["Side"] * 8 + ["Main"] * 8


def test_top_band_line_cut_by_a_gutter_keeps_both_parts_on_page_one(tmp_path):
    """A running header cut at page 1's gutter keeps both of its parts
    there, once; its whole copies on later pages go."""
    header = [(False, 9, 72, 765, "Running Name"), (False, 9, 400, 765, "Curriculum Vitae")]
    pages = [header + _two_column_page(right_offset=5)] + [
        header + [(False, 10, 72, 600, f"Body {n}")] for n in (1, 2)]
    _, doc = _convert(tmp_path, pages)
    text = " ".join(_texts(doc))
    assert (text.count("Running Name"), text.count("Curriculum Vitae")) == (1, 1)


def test_single_page_is_never_treated_as_furniture(tmp_path):
    _, doc = _convert(tmp_path, [[(False, 9, 72, 765, "Only header")]])
    assert doc.paragraphs[0].text == "Only header"


def test_image_only_page_reported(tmp_path):
    pages = [[(False, 10, 72, 700, "This first page carries plenty of readable text.")], []]
    report, _ = _convert(tmp_path, pages, image_pages=(1,))
    assert report.image_only_pages == [2]
    assert report.chars > 0


def test_textless_page_without_image_is_not_image_only(tmp_path):
    pages = [[(False, 10, 72, 700, "This first page carries plenty of readable text.")], []]
    report, _ = _convert(tmp_path, pages)
    assert report.image_only_pages == []


def test_control_characters_do_not_break_the_docx(tmp_path):
    # pdfminer usually maps such glyphs to "(cid:N)"; a font ToUnicode map can
    # still yield a raw control char, which python-docx rejects.
    out = tmp_path / "c.docx"
    _write_docx([_Para(runs=[_Run("ab\x01cd\x0c", False, 10.0)])], out)
    assert Document(str(out)).paragraphs[0].text == "abcd"


def test_unreadable_pdf_raises(tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf at all")
    with pytest.raises(Exception) as excinfo:
        convert_pdf_to_docx(bad, tmp_path / "bad.docx")
    assert "encrypted" not in str(excinfo.value)
    assert not (tmp_path / "bad.docx").exists()


def test_wire_reader_and_doctor_see_converted_file(tmp_path):
    """extract_unified_elements + iter_header_candidates on the converted docx:
    the bold ALL-CAPS header is a candidate and elements keep source order."""
    page = [
        (True, 12, 72, 720, "EDUCATION"),
        (False, 10, 72, 700, "Doctor of Medicine, Example University"),
        (False, 10, 72, 640, "2019"),
        (False, 10, 260, 640, "Award"),
        (True, 12, 72, 580, "PUBLICATIONS"),
        (False, 10, 72, 560, "First paper title"),
    ]
    pdf, docx = tmp_path / "w.pdf", tmp_path / "w.docx"
    pdf.write_bytes(_make_pdf([page]))
    convert_pdf_to_docx(pdf, docx)
    assert iter_header_candidates(str(docx)) == ["EDUCATION", "PUBLICATIONS"]
    texts = [e.get("text") for e in extract_unified_elements(str(docx))["elements"]]
    order = [t for t in texts if t]
    assert order.index("EDUCATION") < order.index("Doctor of Medicine, Example University") \
        < order.index("PUBLICATIONS") < order.index("First paper title")
    assert "" in texts  # blank paragraphs survive into the element stream


def _texts(doc):
    return [p.text for p in doc.paragraphs]


def _width(tmp_path, text, size=10):
    """Rendered width of `text` in base-14 Helvetica, measured by pdfplumber."""
    probe = tmp_path / "probe.pdf"
    probe.write_bytes(_make_pdf([[(False, size, 0, 700, text)]]))
    with pdfplumber.open(probe) as pdf:
        return max(w["x1"] for w in pdf.pages[0].extract_words())


def test_superscript_stays_on_its_line(tmp_path):
    page = [(False, 10, 72, 700, "Figure"), (False, 7, 104, 703, "1")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["Figure 1"]


def test_mixed_sizes_on_one_baseline_are_one_line_of_two_runs(tmp_path):
    page = [(False, 14, 72, 684, "Big"), (False, 10, 100, 684, "small")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["Big small"]
    assert [(r.text, r.font.size.pt) for r in doc.paragraphs[0].runs] == [("Big ", 14), ("small", 10)]


def test_runs_split_on_weight_and_merge_within_a_weight(tmp_path):
    page = [(True, 10, 72, 700, "Name:"), (False, 10, 110, 700, "value"),
            (False, 10, 140, 700, "more")]
    _, doc = _convert(tmp_path, [page])
    assert [(r.text, r.bold) for r in doc.paragraphs[0].runs] == [("Name: ", True), ("value more", False)]


def test_words_of_mixed_size_are_ordered_left_to_right(tmp_path):
    """Same baseline, tops 10pt apart, right-hand word drawn first: only a
    bottom-keyed line grouping keeps them on one line, and only a sort by x
    puts them in reading order."""
    page = [(False, 20, 110, 684, "Big"), (False, 10, 72, 684, "small")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["small Big"]


def test_words_are_ordered_left_to_right(tmp_path):
    page = [(False, 10, 260, 630, "Award"), (False, 10, 72, 630, "2019")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["2019\tAward"]


def test_bold_line_does_not_absorb_a_plain_continuation(tmp_path):
    page = [(True, 10, 72, 700, _LONG), (False, 10, 72, 688, "tail")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [_LONG, "tail"]


def test_mixed_weight_line_does_not_merge_with_an_all_bold_line(tmp_path):
    page = [(True, 10, 72, 700, "Lead:"), (False, 10, 104, 700, _LONG), (True, 10, 72, 688, "Tail")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [f"Lead: {_LONG}", "Tail"]


def test_tabbed_line_is_never_merged_with_its_neighbours(tmp_path):
    before = [(False, 10, 72, 700, _LONG), (False, 10, 72, 688, "2019"),
              (False, 10, 260, 688, "Award")]
    after = [(False, 10, 72, 700, "Label"), (False, 10, 200, 700, _LONG),
             (False, 10, 72, 688, "tail")]
    _, doc_before = _convert(tmp_path, [before], name="b")
    _, doc_after = _convert(tmp_path, [after], name="a")
    assert _texts(doc_before) == [_LONG, "2019\tAward"]
    assert _texts(doc_after) == [f"Label\t{_LONG}", "tail"]


def test_size_change_starts_a_new_paragraph(tmp_path):
    page = [(False, 10, 72, 700, _LONG), (False, 14, 72, 684, "Bigger tail")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [_LONG, "Bigger tail"]


def test_line_short_of_the_margin_is_not_a_wrapped_line(tmp_path):
    """The margin is measured from the leftmost line; P ends 36pt short of it,
    which is outside the slack (33pt) only if `left` is the real left edge."""
    right = _width(tmp_path, _LONG) + 72
    short = "short text"
    x = right - 36 - (_width(tmp_path, short) - 0)
    pages = [[(False, 10, x, 700, short), (False, 10, 72, 688, "next")],
             [(False, 10, 72, 700, _LONG)]]
    _, doc = _convert(tmp_path, pages)
    assert _texts(doc)[:2] == [short, "next"]


@pytest.mark.parametrize("fontname, bold", [
    ("ABCDEF+Arial-BoldMT", True), ("Foo-Black", True), ("FooHEAVY", True),
    ("ABCDEF+CMBX12", True), ("ABCDEF+CMSSBX10", True), ("CMBX10", True),
    ("NimbusRomNo9L-Medi", True), ("Foo-Semibold", True), ("Foo-DemiBold", True),
    ("Foo-Demi", True),
    ("Acmbx12", False), ("CMBXfoo", False), ("Arial-bxtra", False),
    ("HelveticaNeue-Medium", False), ("Arial-ItalicMT", False),
    ("TimesNewRomanPSMT", False), ("ABCDEF+CMSS10", False), ("CMR12", False),
    ("ABCDEF+CIDFont+F1", False)])
def test_bold_detection_by_font_name(fontname, bold):
    assert _is_bold(fontname) is bold


def test_furniture_needs_half_the_pages(tmp_path):
    pages = [[(False, 9, 72, 765, "Rare header" if n < 2 else "Other"),
              (False, 10, 72, 600, f"Body {n}")] for n in range(5)]
    _, doc = _convert(tmp_path, pages)
    assert _texts(doc).count("Rare header") == 2


def test_two_page_furniture_keeps_only_its_first_copy(tmp_path):
    pages = [[(False, 9, 72, 765, "Both pages"), (False, 10, 72, 600, f"Body {n}")]
             for n in range(2)]
    _, doc = _convert(tmp_path, pages)
    assert " ".join(_texts(doc)).count("Both pages") == 1


def test_furniture_matches_across_small_position_jitter(tmp_path):
    pages = [[(False, 9, 72, 765 - 0.4 * n, "Jittery header"), (False, 10, 72, 600, f"Body {n}")]
             for n in range(3)]
    _, doc = _convert(tmp_path, pages)
    assert " ".join(_texts(doc)).count("Jittery header") == 1


def test_same_text_at_different_positions_is_not_furniture(tmp_path):
    pages = [[(False, 9, 72, 765 - 25 * n, "Moving title"), (False, 10, 72, 600, f"Body {n}")]
             for n in range(2)]
    _, doc = _convert(tmp_path, pages)
    assert " ".join(_texts(doc)).count("Moving title") == 2


def test_page_with_text_and_an_image_is_not_image_only(tmp_path):
    pages = [[(False, 10, 72, 700, "A page with real text and a logo image on it.")]]
    report, _ = _convert(tmp_path, pages, image_pages=(0,))
    assert report.image_only_pages == []


@pytest.mark.parametrize("blank_above", [True, False])
def test_page_break_after_a_blank_separated_entry_gets_a_blank(tmp_path, blank_above):
    """Entries separated by blank lines that straddle a page break keep a
    blank between them; tight entries do not get one."""
    second_y = 670 if blank_above else 688
    pages = [[(False, 10, 72, 700, "Grant one"), (False, 10, 72, second_y, "Grant two")],
             [(False, 10, 72, 700, "Grant three"), _FILL]]
    _, doc = _convert(tmp_path, pages)
    assert _texts(doc)[:-2] == (["Grant one", "", "Grant two", "", "Grant three"] if blank_above
                                else ["Grant one", "Grant two", "Grant three"])


def test_tight_entry_after_a_blank_separated_one_gets_no_blank_on_the_same_page(tmp_path):
    page = [(False, 10, 72, 700, "Grant one"), (False, 10, 72, 670, "Grant two"),
            (False, 10, 72, 658, "Grant three"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:4] == ["Grant one", "", "Grant two", "Grant three"]


def test_page_break_is_not_a_vertical_gap(tmp_path):
    pages = [[(False, 10, 72, 750, "Top of page one")], [(False, 10, 72, 60, "Bottom of page two")]]
    report, _ = _convert(tmp_path, pages)
    assert report.blank_paragraphs == 0


@pytest.mark.parametrize("big_first", [True, False])
def test_gap_after_a_mixed_size_line_uses_its_largest_size_and_lowest_edge(tmp_path, big_first):
    """Gap 13.6pt: under the 14pt line's blank threshold (14) but over 10 (the
    smaller word's) and, measured from the smaller word's bottom, over 14.
    Either word order: the size is the max, not the first run's."""
    big, small = (False, 14, 72, 700, "Big"), (False, 10, 100, 700, "small")
    if not big_first:  # the small word now sits to the left
        big, small = (False, 14, 100, 700, "Big"), (False, 10, 72, 700, "small")
    page = [big, small, (False, 10, 72, 675.57, "next")]
    report, _ = _convert(tmp_path, [page])
    assert (report.paragraphs, report.blank_paragraphs) == (2, 0)


def test_gap_before_a_mixed_size_line_is_measured_from_its_highest_top(tmp_path):
    """Gap 8pt to the 14pt word's top; 11pt if measured to the 10pt word's."""
    page = [(False, 10, 72, 700, "plain"), (False, 14, 72, 678.83, "Big"),
            (False, 10, 100, 678.83, "small")]
    report, _ = _convert(tmp_path, [page])
    assert (report.paragraphs, report.blank_paragraphs) == (2, 0)


def _run_cli(*args):
    return subprocess.run([sys.executable, str(_CLI), *map(str, args)],
                          capture_output=True, text=True,
                          env={"PYTHONPATH": str(_SRC), "PATH": ""})


@pytest.mark.parametrize("prefix, text, breaks", [
    ("", "12. Second entry", True), ("", "3\\) Second entry", True),
    ("", "[4] Second entry", True), ("", "\\267 Second entry", True),
    ("- ", "- Second entry", True), ("* ", "* Second entry", True),
    ("\\261 ", "\\261 Second entry", True),
    ("", "12.Smith and colleagues", True), ("", "3\\)Second entry", True),
    ("", "1.5 mg is the dose", False), ("", "12.smith lower case", False), ("", "-5 units of it", False),
    ("", "2019 was the year", False), ("", "1234. long number", False),
    ("", "- present", False), ("", "* present", False),
    ("", "\\261 present", False), ("", "\\320 present", False), ("- ", "* Second entry", False)])
def test_list_marker_line_starts_a_new_entry(tmp_path, prefix, text, breaks):
    """A dash/asterisk marker only splits inside a list that began with the
    same marker; a lone one is a wrapped range ("- present")."""
    page = [(False, 10, 72, 700, prefix + _LONG + " cohorts."), (False, 10, 72, 688, text)]
    _, doc = _convert(tmp_path, [page])
    assert (len(doc.paragraphs) == 2) is breaks


@pytest.mark.parametrize("ending", [
    "529-", "529\\261", "529\\320", "smith,", "trials \\(", "Note:", "Smith &",
    "the trial and", "a study of", "in the", "held in"])
def test_marker_after_a_mid_phrase_line_is_a_wrap(tmp_path, ending):
    """The wrapped page range "...529-" / "45. PMID 1234" is one citation."""
    page = [(False, 10, 72, 700, f"{_LONG} (3):{ending}"), (False, 10, 72, 688, "45. PMID 1234")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 1


_FILL = (False, 10, 72, 400, _LONG)  # a wide line elsewhere: the short lines are not "full"


def test_open_year_ranges_stay_separate_entries(tmp_path):
    lines = ["Member, Committee A, 2004-", "Member, Committee B, 2009-", "Chair, Committee C, 2003-2006"]
    page = [(False, 10, 72, 700 - 12 * i, t) for i, t in enumerate(lines)] + [_FILL]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:3] == lines


@pytest.mark.parametrize("dash", ["-", "\\261", "\\320"])
@pytest.mark.parametrize("following, merges", [
    ("present", True), ("Present", True), ("current", True), ("now", True),
    ("9 months", True), ("June 2010", True), ("May 2010", True), ("Sept. 2011", True),
    ("Dec 2011", True),
    ("Marine Corps veteran", False), ("Chair, Board", False), ("Januaryish", False),
    ("Nowhere Fund", False), ("Presenter, X", False)])
def test_open_year_range_wraps_only_into_a_continuation(tmp_path, dash, following, merges):
    page = [(False, 10, 72, 700, f"Member, Committee A, 2004{dash}"),
            (False, 10, 72, 688, following), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert (len(_texts(doc)[0].split()) > 4 and following in _texts(doc)[0]) is merges


@pytest.mark.parametrize("prev_text, following, merges", [
    ("Pages A12004-", "Marine Corps veteran", True),       # not a year (no word boundary)
    ("Room 1234-", "Marine Corps veteran", True),          # not a 19xx/20xx year
    ("Chair, 2003-2006, pp. 45-", "Marine Corps veteran", True),  # year-dash is not at the end
    ("Member, 2004 -", "Marine Corps veteran", False),     # space before the dash
    ("Member, 2004-", "Chair, 2009 - present", False),     # "present" is not at the start
    ("Member, 2004-", "PRESENT", True), ("Member, 2004-", "JUNE 2010", True),
    ("Member, 2004-", "March 2010", True), ("Member, 2004-", "Mar 2010", True),
    ("Member, 2004-", "2009", True)])
def test_open_range_edge_cases(tmp_path, prev_text, following, merges):
    page = [(False, 10, 72, 700, prev_text), (False, 10, 72, 688, following), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert (following in _texts(doc)[0]) is merges


def test_outdent_after_an_open_year_range_starts_an_entry(tmp_path):
    second = _LONG[:-6] + " 2004-"  # about the width of _LONG, ends in an open range
    page = [(False, 10, 72, 700, _LONG), (False, 10, 90, 688, second),
            (False, 10, 72, 676, "Next entry begins here")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 2


def test_page_range_before_a_marker_line_still_merges_but_an_open_range_does_not(tmp_path):
    page_range = [(False, 10, 72, 700, f"{_LONG} (3):529-"), (False, 10, 72, 688, "45. PMID 1234")]
    open_range = [(False, 10, 72, 700, f"{_LONG} Member, 2004-"), (False, 10, 72, 688, "\\267 Member, 2009-")]
    _, doc_a = _convert(tmp_path, [page_range], name="a")
    _, doc_b = _convert(tmp_path, [open_range], name="b")
    assert len(doc_a.paragraphs) == 1
    assert len(doc_b.paragraphs) == 2


@pytest.mark.parametrize("rule", ["-----", "\\261\\261\\261\\261", "- - - -"])
def test_rule_line_is_not_a_line_ending_in_a_dash(tmp_path, rule):
    page = [(False, 10, 72, 700, rule), (False, 10, 72, 688, "1904 in print"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[1] == "1904 in print"


def test_full_width_rule_line_does_not_exempt_a_marker_line(tmp_path):
    page = [(False, 10, 72, 700, "-" * 130), (False, 10, 72, 688, "12. Next entry")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 2


@pytest.mark.parametrize("ending, merges", [
    ("529-", True), ("529\\261", True), ("529\\320", True), ("529.", False), ("529 -x", False)])
def test_short_line_ending_in_a_dash_still_wraps_on(tmp_path, ending, merges):
    """A page range broken at its dash stops short of the margin (the widest
    line on the page sets it) but is still one paragraph with what follows."""
    page = [(False, 10, 72, 700, f"Pages {ending}"), (False, 10, 72, 688, "1904 in print"),
            (False, 10, 72, 400, _LONG)]
    _, doc = _convert(tmp_path, [page])
    assert (_texts(doc)[0].endswith("1904 in print")) is merges


def test_marker_after_a_line_that_merely_contains_a_connector_word_splits(tmp_path):
    page = [(False, 10, 72, 700, f"{_LONG} begin"), (False, 10, 72, 688, "45. PMID 1234")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 2


def test_real_dash_list_splits_and_a_lone_dash_wrap_merges(tmp_path):
    lst = [(False, 10, 72, 700, f"- {_LONG} cohorts."), (False, 10, 72, 688, f"- {_LONG} cohorts."),
           (False, 10, 72, 676, f"- {_LONG} cohorts.")]
    wrap = [(False, 10, 72, 700, _LONG + " cohorts."), (False, 10, 72, 688, "- present")]
    _, doc_list = _convert(tmp_path, [lst], name="l")
    _, doc_wrap = _convert(tmp_path, [wrap], name="w")
    assert len(doc_list.paragraphs) == 3
    assert _texts(doc_wrap) == [f"{_LONG} cohorts. - present"]


@pytest.mark.parametrize("char", ["\uf0b7", "\u2500", "\u2022"])
def test_bullet_glyphs_are_list_markers(char):
    assert LIST_MARKER_RE.match(f"{char} item")


def _bulleted_items(font, glyph, x, drop):
    """A heading, then two items at x=90 whose bullet `glyph` sits at `x`,
    `drop` points below the item's baseline."""
    page = [(True, 12, 72, 700, "SERVICE")]
    for i, item in enumerate(("First item", "Second item")):
        y = 650 - 14 * i
        page += [(font, 10, x, y - drop, glyph), (False, 10, 90, y, item)]
    return page


@pytest.mark.parametrize("font, glyph, drop, first, second", [
    # Word's level-2 bullet on its item's line (XACIVX, DYLJXC) ...
    (COURIER, "o", 0.5, "\u25e6 First item", "\u25e6 Second item"),
    # ... and 4.5pt lower, its own line cluster (SXPHOG, #1584).
    (COURIER, "o", 4.5, "\u25e6 First item", "\u25e6 Second item"),
    # A bullet glyph off its item's baseline is that item's marker too.
    (False, "\267", 4.5, "\u2022 First item", "\u2022 Second item"),
    # A Helvetica "o" is a word, not Word's bullet.
    (False, "o", 0.5, "o First item", "o Second item"),
])
def test_bullet_glyph_marks_the_item_it_sits_left_of(tmp_path, font, glyph, drop, first, second):
    _, doc = _convert(tmp_path, [_bulleted_items(font, glyph, 72, drop)])
    assert [t for t in _texts(doc) if t] == ["SERVICE", first, second]


@pytest.mark.parametrize("x, drop", [
    (300, 4.5),  # right of its neighbour line's first word: not its marker
    (72, 8.0),   # overlaps the line above by under half its own height
])
def test_bullet_glyph_that_marks_no_line_stays_as_printed(tmp_path, x, drop):
    _, doc = _convert(tmp_path, [_bulleted_items(COURIER, "o", x, drop)])
    words = " ".join(_texts(doc)).split()
    assert words.count("o") == 2 and "\u25e6" not in words
    # Not folded into the item's line: at x=300 a fold lands the glyph after
    # the item ("First item\to"), which the word count above cannot see.
    texts = _texts(doc)
    assert "First item" in texts and not any("item\to" in t for t in texts)


def test_first_line_indent_paragraph_is_not_split(tmp_path):
    page = [(False, 10, 90, 700, _LONG), (False, 10, 72, 688, _LONG),
            (False, 10, 72, 676, "tail of the paragraph")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 1


@pytest.mark.parametrize("last_word, splits", [("of", False), ("in", False), ("at", True), ("9.", True)])
def test_outdent_after_a_mid_phrase_line_does_not_start_an_entry(tmp_path, last_word, splits):
    second = _LONG[:-2] + last_word  # same width as _LONG: still a full line
    page = [(False, 10, 72, 700, _LONG), (False, 10, 90, 688, second),
            (False, 10, 72, 676, "Next entry begins here")]
    _, doc = _convert(tmp_path, [page])
    assert (len(doc.paragraphs) == 2) is splits


def test_indentation_is_judged_from_the_second_line_only(tmp_path):
    """72, 72, 90, 72: line 2 fixes the paragraph as not hanging; a later
    deeper line must not turn it into a hanging-indent entry."""
    page = [(False, 10, x, y, _LONG_STOP) for x, y in ((72, 700), (72, 688), (90, 676), (72, 664))]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 1


def test_hanging_indent_needs_a_return_to_the_first_line_x0(tmp_path):
    """Returning 6pt right of the first line's x0 is outside the 3pt
    tolerance, so no new entry; returning to 74pt is inside it."""
    def entry(x):
        return [(False, 10, 72, 700, _LONG), (False, 10, 90, 688, _LONG_STOP),
                (False, 10, x, 676, "Next entry begins here")]
    _, far = _convert(tmp_path, [entry(78)], name="far")
    _, near = _convert(tmp_path, [entry(74)], name="near")
    assert len(far.paragraphs) == 1
    assert len(near.paragraphs) == 2


def test_outdented_line_after_a_hanging_indent_is_a_new_entry(tmp_path):
    page = [(False, 10, 72, 700, _LONG), (False, 10, 90, 688, _LONG_STOP),
            (False, 10, 72, 676, "Next entry begins here")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [f"{_LONG} {_LONG_STOP}", "Next entry begins here"]


def _spaced_line(tmp_path, gaps, size=10):
    """A line of 'aa' words whose measured gaps are exactly `gaps`."""
    w = _width(tmp_path, "aa", size)
    xs = [72.0]
    for gap in gaps:
        xs.append(xs[-1] + w + gap)
    return [(False, size, x, 700, "aa") for x in xs]


def _words(gaps, size=4.0, width=11.0):
    xs = [0.0]
    for gap in gaps:
        xs.append(xs[-1] + width + gap)
    return [{"x0": x, "x1": x + width, "size": size} for x in xs]


def test_full_width_justified_line_with_uniform_gaps_has_no_tabs(tmp_path):
    page = _spaced_line(tmp_path, [89, 89, 89, 89]) + [(False, 10, 72, 400, _LONG)]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[0] == "aa aa aa aa aa"


def test_full_width_line_still_tabs_an_outlier_gap(tmp_path):
    page = _spaced_line(tmp_path, [10, 10, 10, 350]) + [(False, 10, 72, 400, _LONG)]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[0] == "aa aa aa aa\taa"


def test_row_short_of_the_margin_tabs_every_wide_gap(tmp_path):
    """4 columns, gaps all near the median: not justified (not full width)."""
    page = _spaced_line(tmp_path, [30, 24, 43, 43]) + [(False, 10, 72, 400, _LONG)]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[0] == "aa\taa\taa\taa\taa"


def test_outlier_test_is_off_unless_asked_for():
    assert _column_gaps(_words([20, 20, 20, 20])) == [True] * 4
    assert _column_gaps(_words([20, 20, 20, 20]), outlier_test=True) == [False] * 4


def test_gap_at_exactly_twice_the_median_is_an_outlier():
    assert _column_gaps(_words([10, 10, 10, 20]), outlier_test=True) == [False, False, False, True]
    assert _column_gaps(_words([10, 10, 10, 19.9]), outlier_test=True) == [False] * 4


def test_gap_over_twice_the_median_is_a_column_gap_even_at_2_25x():
    assert _column_gaps(_words([20, 20, 20, 45]), outlier_test=True)[-1] is True


def test_outlier_is_measured_against_the_median_not_the_smallest_gap():
    assert _column_gaps(_words([16, 50, 50, 50]), outlier_test=True) == [False] * 4


def test_few_gaps_skip_the_median_test():
    assert _column_gaps(_words([40, 40]), outlier_test=True) == [True, True]
    assert _column_gaps(_words([40, 40, 40]), outlier_test=True) == [False] * 3


def test_line_of_three_columns_keeps_both_tabs(tmp_path):
    _, doc = _convert(tmp_path, [_spaced_line(tmp_path, [40, 40])])
    assert _texts(doc) == ["aa\taa\taa"]


def test_tab_threshold_uses_the_size_of_the_word_before_the_gap(tmp_path):
    big = _width(tmp_path, "Big", 20)
    after_big = [(False, 20, 72, 700, "Big"), (False, 10, 72 + big + 20, 700, "small")]
    small = _width(tmp_path, "small", 10)
    before_big = [(False, 10, 72, 700, "small"), (False, 20, 72 + small + 20, 700, "Big")]
    _, doc_a = _convert(tmp_path, [after_big], name="a")
    _, doc_b = _convert(tmp_path, [before_big], name="b")
    assert _texts(doc_a) == ["Big small"]      # 20pt gap < 1.5 x 20pt
    assert _texts(doc_b) == ["small\tBig"]     # 20pt gap > 1.5 x 10pt


def test_password_protected_pdf_raises_a_clear_error(tmp_path):
    pdf = tmp_path / "locked.pdf"
    pdf.write_bytes(_make_pdf([_cv_page()], user_password="secret"))
    with pytest.raises(ValueError, match="encrypted/password-protected"):
        convert_pdf_to_docx(pdf, tmp_path / "locked.docx")


def test_cli_names_the_password_protected_cause(tmp_path):
    (tmp_path / "in").mkdir()
    (tmp_path / "in" / "locked.pdf").write_bytes(_make_pdf([_cv_page()], user_password="secret"))
    res = _run_cli(tmp_path / "in", tmp_path / "out")
    assert res.returncode == 1
    assert "conversion failed: locked.pdf (ValueError: encrypted/password-protected PDF)" in res.stderr
    assert res.stdout.strip() == "locked\t\t\t\t\t\terror"


def test_cli_one_tsv_line_per_file_and_continues_past_failure(tmp_path):
    src, out = tmp_path / "in", tmp_path / "out"
    src.mkdir()
    (src / "a.pdf").write_bytes(_make_pdf([_cv_page()]))
    (src / "b.pdf").write_bytes(b"garbage")
    (src / "c.pdf").write_bytes(_make_pdf([[(False, 10, 72, 700, "x" * 30)], []], image_pages=(1,)))
    (src / "notes.txt").write_text("ignored")
    res = _run_cli(src, out)
    rows = [line.split("\t") for line in res.stdout.splitlines()]
    assert res.returncode == 1
    assert [r[0] for r in rows] == ["a", "b", "c"]
    assert rows[0] == ["a", "1", "4", "1", "", str(len(_LONG) + len(_LONG_TAIL) + 1
                                                   + len("EDUCATION") + len("2019\tAward")), "ok"]
    assert rows[1] == ["b", "", "", "", "", "", "error"]
    assert rows[2][4] == "2" and rows[2][6] == "ok"
    assert all(len(r) == 7 for r in rows)
    assert "b.pdf" in res.stderr and "a.pdf" not in res.stderr
    assert (out / "a.docx").exists() and (out / "c.docx").exists()
    assert not (out / "b.docx").exists()


def test_cli_all_ok_exits_zero_for_single_file(tmp_path):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(_make_pdf([_cv_page()]))
    res = _run_cli(pdf, tmp_path / "out")
    assert res.returncode == 0 and res.stdout.count("\n") == 1
    assert res.stderr == ""


# --- Round 2 (#806): the live A/B's defects --------------------------------


def _bold_width(tmp_path, text, size=10):
    probe = tmp_path / "probe_bold.pdf"
    probe.write_bytes(_make_pdf([[(True, size, 0, 700, text)]]))
    with pdfplumber.open(probe) as pdf:
        return max(w["x1"] for w in pdf.pages[0].extract_words())


_MAIN = "Main text of the main column that runs on {}"


def _two_column_page(right_offset, n_side=8, n_main=8, top=650):
    """A sidebar (x=72) and a main column (x=260), 14pt pitch; the main
    column's baselines sit `right_offset` points lower."""
    side = [(False, 10, 72, top - 14 * i, f"Side {i}") for i in range(n_side)]
    main = [(False, 10, 260, top - right_offset - 14 * i, _MAIN.format(i)) for i in range(n_main)]
    return side + main


def _order(doc, *keys):
    return [w for w in " ".join(_texts(doc)).split() if w in keys]


def _gutter_middle(tmp_path):
    """Where `_widest_gutter` puts the middle of `_two_column_page`'s gutter:
    the strip runs from GUTTER_EDGE_FRAC in from the text's left edge to
    the main column's x0."""
    right = 260 + _width(tmp_path, _MAIN.format(0))
    return (int(72 + GUTTER_EDGE_FRAC * (right - 72)) + 260) / 2


def test_two_column_page_reads_the_left_column_first(tmp_path):
    """Each column's equal-width lines are full against their OWN column's
    right edge, so each column is one paragraph; a column start never
    continues the sidebar's last line."""
    _, doc = _convert(tmp_path, [_two_column_page(right_offset=5)])
    assert _texts(doc) == [" ".join(f"Side {i}" for i in range(8)),
                           " ".join(_MAIN.format(i) for i in range(8))]


def test_row_aligned_columns_stay_one_tabbed_line_per_row(tmp_path):
    """A date column beside its entries is a table, not two columns."""
    _, doc = _convert(tmp_path, [_two_column_page(right_offset=0, n_side=10, n_main=10)])
    assert _texts(doc)[0] == f"Side 0\t{_MAIN.format(0)}"


def test_rows_offset_under_a_line_height_still_split_into_columns(tmp_path):
    """3pt apart the two columns' lines group into one visual line each, but
    they are not row-aligned (over ROW_ALIGN_TOLERANCE_PT) and the gap at
    the gutter is a gutter, not a word gap: two columns."""
    _, doc = _convert(tmp_path, [_two_column_page(right_offset=3)])
    assert _order(doc, "Side", "Main") == ["Side"] * 8 + ["Main"] * 8


@pytest.mark.parametrize("in_gutter", [False, True])
def test_full_width_line_splits_the_columns_into_regions(tmp_path, in_gutter):
    """Between full-width lines each region reads left column, then right.
    The heading's long word starts before the gutter, or inside it short of
    its middle: neither gap is a column gap by the gutter."""
    # Column gaps before the gutter band and in the main column, none by
    # the gutter: still full width.
    label_end = 72 + _width(tmp_path, "2019")
    heading = [(False, 10, 72, 520, "2019"),
               (False, 10, 2 * _gutter_middle(tmp_path) - 260 + 5 if in_gutter else label_end + 16, 520,
                "Fullwidthheadingthatcrossesthegutter"),
               (False, 10, 450, 520, "2020")]
    below = [(False, 10, 72, 500 - 14 * i, f"Low {i}") for i in range(6)] + \
            [(False, 10, 260, 495 - 14 * i, f"Deep {i}") for i in range(6)]
    _, doc = _convert(tmp_path, [_two_column_page(right_offset=5) + heading + below])
    words = [w[:4] for w in " ".join(_texts(doc)).split()]
    assert [w for w in words if w in ("Side", "Main", "Full", "Low", "Deep")] == (
        ["Side"] * 8 + ["Main"] * 8 + ["Full"] + ["Low"] * 6 + ["Deep"] * 6)


@pytest.mark.parametrize("clustered", [False, True])
def test_long_sidebar_line_ending_near_the_main_column_is_a_sidebar_line(tmp_path, clustered):
    """A sidebar at x=50 with one long line ending 8pt short of the main
    column at x=240 (on its own row, or on a main line's row): not full
    width, so the main column still reads as one paragraph after the
    whole sidebar."""
    long_side = "Sidebar line that runs on"
    side = [(False, 10, 50, 650 - 14 * i, f"Side {i}") for i in range(8)]
    side[2] = (False, 10, 232 - _width(tmp_path, long_side), side[2][3] - (3 if clustered else 0), long_side)
    main = [(False, 10, 240, 645 - 14 * i, _MAIN.format(i)) for i in range(8)]
    _, doc = _convert(tmp_path, [side + main])
    assert _order(doc, "Side", "Sidebar", "Main")[-8:] == ["Main"] * 8
    assert " ".join(_MAIN.format(i) for i in range(8)) in _texts(doc)


def test_gutter_tolerates_exactly_a_tenth_of_the_lines_crossing(tmp_path):
    """20 body lines, 2 full-width headings across the gutter: 2 <= 10%."""
    wide = "Heading " * 8
    page = (_two_column_page(right_offset=5, n_side=5, n_main=5)
            + [(False, 10, 72, 570, "H1" + wide)]
            + _two_column_page(right_offset=5, n_side=4, n_main=4, top=540)
            + [(False, 10, 72, 470, "H2" + wide)])
    _, doc = _convert(tmp_path, [page])
    words = [w[:2] if w.startswith("H") else w for w in " ".join(_texts(doc)).split()]
    assert [w for w in words if w in ("Side", "Main", "H1", "H2")] == (
        ["Side"] * 5 + ["Main"] * 5 + ["H1"] + ["Side"] * 4 + ["Main"] * 4 + ["H2"])


def test_the_widest_strip_is_the_gutter(tmp_path):
    """A 14pt strip inside the sidebar comes first, left to right; the wide
    one past it is the gutter."""
    first = "Sidebarentry"
    x = 72 + _width(tmp_path, first) + 14
    page = _two_column_page(right_offset=5)
    page[:8] = [ln for i in range(8) for ln in ((False, 10, 72, 650 - 14 * i, first),
                                                 (False, 10, x, 650 - 14 * i, f"x{i}"))]
    _, doc = _convert(tmp_path, [page])
    assert _order(doc, first, "Main") == [first] * 8 + ["Main"] * 8


def test_column_line_poking_into_the_gutter_does_not_cross_it(tmp_path):
    """One sidebar line runs 40pt further right than the others (into the
    gutter strip, which tolerates one crossing line of 16) but stops well
    short of the main column: still a sidebar line, so the page keeps one
    region and reads every sidebar line before any main line."""
    page = _two_column_page(right_offset=5)
    page[3] = (False, 10, 72, page[3][3], "Side 3 runs on and on")
    _, doc = _convert(tmp_path, [page])
    assert _order(doc, "Side", "Main") == ["Side"] * 8 + ["Main"] * 8


def test_edge_band_lines_do_not_count_against_the_gutter(tmp_path):
    """Three full-width lines in the top band (a name block) would be 3 of
    19 lines crossing; the gutter is judged on the body band alone."""
    header = [(False, 10, 72, 760 - 12 * i, "Header line " * 20) for i in range(3)]
    _, doc = _convert(tmp_path, [header + _two_column_page(right_offset=5)])
    assert _order(doc, "Side", "Main") == ["Side"] * 8 + ["Main"] * 8


def test_indented_main_column_line_is_not_full_width(tmp_path):
    """A main-column line indented 20pt starts off the column's edge but
    spans nothing across the gutter: still a main-column line."""
    page = _two_column_page(right_offset=5)
    page[11] = (False, 10, 280, page[11][3], "Indented main line")
    _, doc = _convert(tmp_path, [page])
    assert _order(doc, "Side", "Main", "Indented") == ["Side"] * 8 + ["Main"] * 3 + ["Indented"] + ["Main"] * 4


@pytest.mark.parametrize("past_edge", [2.9, 3.6, -6.0])
def test_sidebar_line_and_main_heading_on_one_baseline_split(tmp_path, past_edge):
    """A small sidebar URL running past the gutter's middle shares a baseline
    with a large main-column heading starting just inside or just outside
    the edge tolerance (or 6pt into the gutter), 20pt after the URL: split at
    that column gap (over 1.5em, under 3em), URL in the sidebar."""
    url = "https://example.org/a/long/path"
    x = 260 + past_edge
    page = _two_column_page(right_offset=5) + [
        (False, 9.4, x - 20 - _width(tmp_path, url, 9.4), 665, url), (False, 16.9, x, 665, "Heading")]
    _, doc = _convert(tmp_path, [page])
    assert not any(url in t and "Heading" in t for t in _texts(doc))
    assert _order(doc, "Side", url, "Main", "Heading") == [url] + ["Side"] * 8 + ["Heading"] + ["Main"] * 8


def test_last_gap_by_the_gutter_is_the_cut(tmp_path):
    """A sidebar line with a word in the gutter strip past its middle, then
    a main-column word on the same baseline: both gaps qualify; the cut is
    the last, before the main word, so the gutter word stays in the sidebar."""
    page = _two_column_page(right_offset=5)
    y = page[3][3]
    page[3:4] = [(False, 10, 72, y, "Side 3"), (False, 10, 200, y, "pokes"), (False, 10, 260, y, "Extra")]
    _, doc = _convert(tmp_path, [page])
    assert any("pokes" in t and "Side" in t and "Extra" not in t for t in _texts(doc))
    assert _order(doc, "Side", "Extra", "Main")[:8] == ["Side"] * 8


def test_edge_band_lines_on_one_side_or_split_by_a_gap_join_their_columns(tmp_path):
    """Bottom-band lines are not spanning names: one alone in the sidebar,
    and a sidebar/main pair split by a column gap, go to their columns."""
    page = _two_column_page(right_offset=5) + [
        (False, 10, 72, 75, "Footside"), (False, 10, 72, 62, "Pairside"), (False, 10, 260, 62, "Pairmain")]
    _, doc = _convert(tmp_path, [page])
    order = _order(doc, "Side", "Footside", "Pairside", "Main", "Pairmain")
    assert order == ["Side"] * 8 + ["Footside", "Pairside"] + ["Main"] * 8 + ["Pairmain"]


def test_name_in_the_top_band_spanning_the_gutter_stays_whole_and_first(tmp_path):
    """A 37.6pt name line in the top band, one of its words on the right
    column's edge: never split by a gutter found without it."""
    q = 72 + _width(tmp_path, "Janet", 37.6) + 10
    assert 260 - (q + _width(tmp_path, "Q.", 37.6)) < 1.5 * 37.6  # no column gap
    name = [(False, 37.6, 72, 740, "Janet"), (False, 37.6, q, 740, "Q."), (False, 37.6, 260, 740, "Public")]
    _, doc = _convert(tmp_path, [name + _two_column_page(right_offset=5)])
    assert _texts(doc)[0] == "Janet Q. Public"
    assert _order(doc, "Side", "Main") == ["Side"] * 8 + ["Main"] * 8


def test_narrow_gap_is_not_a_gutter(tmp_path):
    """Columns 11pt apart (under GUTTER_MIN_WIDTH_PT) are read row by row."""
    side_w = _width(tmp_path, "Side 0")
    page = [(False, 10, 72, 650 - 14 * i, f"Side {i}") for i in range(8)] + \
           [(False, 10, 72 + side_w + 11, 645 - 14 * i, f"Main text {i}") for i in range(8)]
    _, doc = _convert(tmp_path, [page])
    words = " ".join(_texts(doc)).split()
    assert words.index("Main") < words.index("Side", 1)


def _read_row_by_row(doc):
    order = _order(doc, "Side", "Main")
    return order.index("Main") < len(order) - 1 - order[::-1].index("Side")


def test_short_label_column_is_not_a_sidebar(tmp_path):
    """5 left lines beside 8 (under COLUMN_BALANCE_MIN_FRAC): a label column,
    read row by row."""
    _, doc = _convert(tmp_path, [_two_column_page(right_offset=5, n_side=5)])
    assert _read_row_by_row(doc)


def test_full_width_lines_count_in_neither_column(tmp_path):
    """6 sidebar lines beside 9 is under COLUMN_BALANCE_MIN_FRAC; a heading
    that starts in the sidebar and runs across must not lift that to 7 of 9."""
    page = _two_column_page(right_offset=5, n_side=6, n_main=9) + [
        (False, 10, 72, 480, "Headingstraddlingwordacrossthegutterandon")]
    _, doc = _convert(tmp_path, [page])
    assert _read_row_by_row(doc)


def test_column_needs_five_lines_even_when_balanced(tmp_path):
    """4 beside 5 is balanced (0.8) but under COLUMN_MIN_LINES. The heading
    is one word across the gutter, so it adds a line to neither column."""
    page = _two_column_page(right_offset=5, n_side=4, n_main=5) + [
        (False, 10, 72, 400, "Fullwidthheading" * 4)]
    _, doc = _convert(tmp_path, [page])
    assert _read_row_by_row(doc)


def test_row_alignment_is_judged_on_the_column_with_fewer_lines(tmp_path):
    """14 left lines, 10 right lines of which 6 sit on left rows: the right
    column is 60% aligned (a date column), the left only 43%."""
    left = [(False, 10, 72, 650 - 14 * i, f"Side {i}") for i in range(14)]
    right = [(False, 10, 260, 650 - 14 * i - (0 if i < 6 else 5), _MAIN.format(i)) for i in range(10)]
    _, doc = _convert(tmp_path, [left + right])
    assert _texts(doc)[0] == f"Side 0\t{_MAIN.format(0)}"


@pytest.mark.parametrize("column", ["left", "right"])
def test_column_of_tabbed_rows_is_part_of_a_table(tmp_path, column):
    """A 'column' whose lines carry column gaps is a table cut down its
    middle, read row by row."""
    if column == "left":
        extra = [(False, 10, 72 + _width(tmp_path, "Side 0") + 20, 650 - 14 * i, "x") for i in range(8)]
    else:
        extra = [(False, 10, 260 + _width(tmp_path, _MAIN.format(0)) + 20, 645 - 14 * i, "x")
                 for i in range(8)]
    _, doc = _convert(tmp_path, [_two_column_page(right_offset=5) + extra])
    assert _read_row_by_row(doc)


def test_right_column_first_on_the_page_may_continue_the_last_page(tmp_path):
    """A page whose first region is right-column lines only: its first line
    is not a column start, so it can continue page 1's last paragraph."""
    page2 = ([(False, 10, 260, 670 - 14 * i, f"Top {i}") for i in range(2)]
             + [(False, 10, 72, 630, "Fullwidthheading" * 4)]
             + _two_column_page(right_offset=5, top=600))
    _, doc = _convert(tmp_path, [[(False, 10, 72, 100, _LONG)], page2])
    assert _texts(doc)[0].startswith(f"{_LONG} Top 0")


def test_wrap_inside_a_column_is_judged_against_the_column_edge(tmp_path):
    """'Side' stops 50pt short of the sidebar's right edge; 'Continuation'
    would not fit there (it would fit before the page's right margin)."""
    side = ["Side alpha beta", "Side", "Continuation", "Side 3", "Side 4", "Side 5", "Side 6", "Side 7"]
    page = _two_column_page(right_offset=5)
    page[:8] = [(False, 10, 72, 650 - 14 * i, t) for i, t in enumerate(side)]
    _, doc = _convert(tmp_path, [page])
    assert "Side Continuation" in " ".join(_texts(doc))


def test_label_wrap_inside_a_column_uses_the_column_width(tmp_path):
    """A right-column line with a leading date is full-width in its column
    (so re-laid out for justification); its wrap hangs under the text."""
    page = _two_column_page(right_offset=5)
    text = _MAIN.format(0)
    page[8] = (False, 10, 260, 645, "2019")
    page += [(False, 10, 302, 645, text), (False, 10, 302, 631, "tail words")]
    page[9:16] = [(False, 10, 260, 617 - 14 * i, _MAIN.format(i)) for i in range(7)]
    _, doc = _convert(tmp_path, [page])
    assert f"2019\t{text} tail words" in _texts(doc)


@pytest.mark.parametrize("first_main_is_row", [False, True])
def test_column_start_below_the_sidebar_adds_no_blank(tmp_path, first_main_is_row):
    """The main column starts 30pt below the sidebar's last line: a column
    start, not a vertical gap (also when the first main line is a row)."""
    page = _two_column_page(right_offset=5, top=650)
    page[8:] = [(False, 10, 260, 510 - 14 * i, _MAIN.format(i)) for i in range(8)]
    if first_main_is_row:
        page[8:9] = [(False, 10, 260, 510, "2019"), (False, 10, 320, 510, "cellb"),
                     (False, 10, 380, 510, "cellc"), (False, 10, 380, 503, "wrap")]
    report, _ = _convert(tmp_path, [page])
    assert report.blank_paragraphs == 0


def test_row_at_the_foot_of_the_sidebar_does_not_take_the_main_column(tmp_path):
    page = _two_column_page(right_offset=5)
    page[7] = (False, 10, 72, page[7][3], "a")
    page += [(False, 10, 100, 552, "b"), (False, 10, 128, 552, "c")]
    _, doc = _convert(tmp_path, [page])
    assert not any("Main" in t and "\t" in t for t in _texts(doc))


def _spaceless_line(tmp_path, y, words, gap, size=10, bold=False):
    """Each word drawn on its own, `gap` points apart, with no space glyph."""
    out, x = [], 72.0
    for word in words:
        out.append((bold, size, x, y, word))
        x += (_bold_width if bold else _width)(tmp_path, word, size) + gap
    return out


_SPACELESS_WORDS = ["alpha", "beta", "gamma", "delta", "epsilon"]


def test_page_without_space_glyphs_splits_words_on_a_relative_gap(tmp_path):
    page = [w for i in range(9) for w in _spaceless_line(tmp_path, 700 - 30 * i, _SPACELESS_WORDS, 2.8)]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[0] == " ".join(_SPACELESS_WORDS)


def test_page_with_space_glyphs_keeps_the_fixed_tolerance(tmp_path):
    """Letter-spaced capitals 2pt apart (over 0.15 x 12pt) stay one word on
    a page that has real spaces."""
    prose = [(False, 10, 72, 700 - 30 * i, "a line of ordinary words with spaces") for i in range(8)]
    caps, x = [], 72.0
    for ch in "SECTION":
        caps.append((False, 12, x, 400, ch))
        x += _width(tmp_path, ch, 12) + 2.0
    _, doc = _convert(tmp_path, [prose + caps])
    assert "SECTION" in _texts(doc)


def test_too_few_characters_is_never_spaceless(tmp_path):
    page = _spaceless_line(tmp_path, 700, _SPACELESS_WORDS, 2.8)
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["".join(_SPACELESS_WORDS)]


@pytest.mark.parametrize("label", ["2019-2020", "Jan 2019"])
def test_wrap_under_the_text_after_a_leading_date_merges(tmp_path, label):
    page = [(False, 10, 72, 700, label), (False, 10, 150, 700, _LONG),
            (False, 10, 150, 688, "tail words")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [f"{label}\t{_LONG} tail words"]


@pytest.mark.parametrize("x, merges", [(150, True), (153, True), (154, False), (72, False)])
def test_wrap_after_a_label_must_align_with_the_text_column(tmp_path, x, merges):
    page = [(False, 10, 72, 700, "2019"), (False, 10, 150, 700, _LONG),
            (False, 10, x, 688, "tail words")]
    _, doc = _convert(tmp_path, [page])
    assert (len(doc.paragraphs) == 1) is merges


@pytest.mark.parametrize("label, merges", [
    ("12.", True), ("3\\)", True), ("Sept. 2011 - Dec 2012", True), ("05/2019 - present", True),
    ("2004 -", True), ("Note", False), ("Boston", False), ("1234", False)])
def test_only_a_number_or_date_label_is_a_leading_label(tmp_path, label, merges):
    page = [(False, 10, 72, 700, label), (False, 10, 200, 700, _LONG),
            (False, 10, 200, 688, "tail words")]
    _, doc = _convert(tmp_path, [page])
    assert (len(doc.paragraphs) == 1) is merges


def test_label_past_a_fifth_of_the_width_is_not_a_leading_label(tmp_path):
    """The label ends past 20% of the text width: a label/value line, whose
    next line is its own entry even when aligned under the value."""
    label = "Department of Internal Medicine"
    x = 72 + _width(tmp_path, label) + 20
    page = [(False, 10, 72, 700, label), (False, 10, x, 700, _LONG),
            (False, 10, x, 688, "tail words")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 2


def test_ragged_right_wrap_merges_when_the_next_word_would_not_fit(tmp_path):
    """Prev stops 40pt short of the margin (outside the 8% slack) but the
    next line's first word is wider than that gap: a soft wrap."""
    right = 72 + _width(tmp_path, _LONG)
    word = "Internationalization"
    short = "short line of text"
    x = right - 40 - _width(tmp_path, short)
    assert _width(tmp_path, word) > 40
    page = [(False, 10, 72, 400, _LONG), (False, 10, x, 700, short), (False, 10, x, 688, word)]
    _, doc = _convert(tmp_path, [page])
    assert f"{short} {word}" in _texts(doc)


_SURNAME = "Abcdefghijklmnopqrstuvwxy"  # 25 letters


@pytest.mark.parametrize("ending, merges", [
    (".", False), (";", False), ("\\)", False), (",", True), (" and", True), (":", True)])
def test_entry_ending_a_sentence_does_not_absorb_a_long_surname(tmp_path, ending, merges):
    """An entry ending ~40pt short of the margin, then an entry starting
    with a 25-letter surname that would not have fitted there: a finished
    entry (. ; or )) is not wrapped; a comma, a connector or a colon is."""
    right = 72 + _width(tmp_path, _LONG)
    entry = "Smith J, Jones K. A short title" + ending
    x = right - 40 - _width(tmp_path, entry)
    assert _width(tmp_path, _SURNAME) > 40
    page = [(False, 10, 72, 400, _LONG), (False, 10, x, 700, entry),
            (False, 10, x, 688, _SURNAME + " L, Other M.")]
    _, doc = _convert(tmp_path, [page])
    assert (entry.replace("\\)", ")") in _texts(doc)) is not merges


@pytest.mark.parametrize("sep", [" ", ": "])
@pytest.mark.parametrize("ragged", [False, True])
def test_date_led_one_line_entries_stay_separate(tmp_path, ragged, sep):
    """Three degree/appointment lines, each a date range and then words:
    full width (merge by width) or ending 30pt short before a date token
    that would not fit there (merge by ragged wrap) -- three entries."""
    entries = [f"{2010 + i}-{2014 + i}{sep}Residency in Internal Medicine, Boston, MA" for i in range(3)]
    right = 72 + _width(tmp_path, _LONG)
    x = right - 30 - _width(tmp_path, entries[0]) if ragged else 72
    page = [(False, 10, x, 700 - 12 * i, t) for i, t in enumerate(entries)]
    page.append((False, 10, 72, 400, _LONG) if ragged else (False, 10, 72, 400, "x"))
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:3] == entries


def test_open_ended_date_ranges_lead_entries_too(tmp_path):
    entries = [f"{2010 + i}-present Professor of Medicine, Boston, MA" for i in range(2)]
    page = [(False, 10, 72, 700 - 12 * i, t) for i, t in enumerate(entries)]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == entries


def test_date_led_wrap_after_a_connector_still_merges(tmp_path):
    first = "2010-2014 " + _LONG[:-3] + " and"
    page = [(False, 10, 72, 700, first), (False, 10, 72, 688, "2015-2018 Fellowship in Cardiology")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [f"{first} 2015-2018 Fellowship in Cardiology"]


@pytest.mark.parametrize("tail", ["Nov 2019. (Invited talk)", "2019, Boston, MA."])
def test_single_date_wrap_in_a_date_led_paragraph_merges(tmp_path, tail):
    """A wrapped citation tail starting with one date is not a new entry."""
    first = "2010-2014 " + _LONG[:-12] + " Journal"
    page = [(False, 10, 72, 700, first), (False, 10, 72, 688, tail.replace("(", "\\(").replace(")", "\\)"))]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [f"{first} {tail}"]


def test_wrapped_date_range_tail_still_merges(tmp_path):
    """A date-led grant whose wrap is only "2019-2021." (no words after the
    date) is the same entry."""
    first = "2018-2021 " + _LONG[:-6] + " Funded"
    page = [(False, 10, 72, 700, first), (False, 10, 72, 688, "2019-2021.")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == [f"{first} 2019-2021."]


def test_date_led_entry_after_a_plain_paragraph_still_wraps(tmp_path):
    """Only a paragraph that itself began with a date vetoes the merge."""
    page = [(False, 10, 72, 700, _LONG_STOP), (False, 10, 72, 688, "2015-2018 Fellowship in Cardiology")]
    _, doc = _convert(tmp_path, [page])
    assert len(doc.paragraphs) == 1


def test_ragged_right_line_whose_next_word_fits_is_not_a_wrap(tmp_path):
    right = 72 + _width(tmp_path, _LONG)
    short = "short line of text"
    x = right - 40 - _width(tmp_path, short)
    page = [(False, 10, 72, 400, _LONG), (False, 10, x, 700, short), (False, 10, x, 688, "at home")]
    _, doc = _convert(tmp_path, [page])
    assert short in _texts(doc)


def test_right_margin_ignores_a_few_lines_that_stick_out(tmp_path):
    """51 lines end at the margin, one sticks out 30pt: a line 10pt short of
    the real margin is full; measured from the stray line it would not be,
    and "tail" would fit after it."""
    base = 72 + _width(tmp_path, _LONG)
    lines = [(False, 10, 72 + (30 if i == 25 else 0), 760 - 12 * i, _LONG) for i in range(51)]
    short = _LONG[:-12]
    page2 = [(False, 10, base - 10 - _width(tmp_path, short), 700, short), (False, 10, 72, 688, "tail")]
    _, doc = _convert(tmp_path, [lines, page2])
    assert _texts(doc)[-1].endswith(f"{short} tail")


def test_right_margin_counts_each_printed_line_of_a_table_row(tmp_path):
    """50 printed lines, one sticking out 30pt: the margin ignores it. Two
    of them are one reassembled table row; counted as one line, the stray
    line would set the margin (#1583)."""
    base = 72 + _width(tmp_path, _LONG)
    lines = [(False, 10, 72 + (30 if i == 25 else 0), 760 - 12 * i, _LONG) for i in range(46)]
    row = (_table_row(200, [(72, "2019"), (150, "Title of"), (400, "Boston")])
           + _table_row(188, [(150, "the award")]))
    short = _LONG[:-12]
    page2 = [(False, 10, base - 10 - _width(tmp_path, short), 700, short), (False, 10, 72, 688, "tail")]
    _, doc = _convert(tmp_path, [lines + row, page2])
    assert "2019\tTitle of the award\tBoston" in _texts(doc)
    assert _texts(doc)[-1].endswith(f"{short} tail")


def _table_row(y, cells):
    return [(False, 10, x, y, text) for x, text in cells if text]


def test_wrapped_table_cells_reassemble_into_one_row(tmp_path):
    page = (_table_row(700, [(72, "2019"), (150, "Title of"), (400, "Boston")])
            + _table_row(688, [(150, "the award")])
            + _table_row(676, [(400, "MA USA")])
            + _table_row(664, [(72, "2020"), (150, "Next title"), (400, "Paris")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tTitle of the award\tBoston MA USA", "2020\tNext title\tParis"]


def test_row_with_an_empty_first_cell_is_its_own_row(tmp_path):
    """A vertically merged first cell prints once (its own text wrapped:
    "Inst"): the next row starts in cell 2. It is a row of its own, its
    first cell left empty, not absorbed into the row above."""
    page = (_table_row(700, [(72, "Alpha"), (150, "Professor"), (400, "2001-2005")])
            + _table_row(688, [(72, "Inst")])
            + _table_row(676, [(150, "Lecturer"), (400, "1998-2001")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["Alpha Inst\tProfessor\t2001-2005", "\tLecturer\t1998-2001"]


def test_cells_wrapping_mid_sentence_stay_in_their_row(tmp_path):
    """#1583: every cell of a row wraps, so each printed line after the first
    has an empty first cell. Where the cells it fills continue the row's
    cells mid-sentence, it is the same row, not the next one."""
    page = (_table_row(700, [(72, "2007"), (150, "Sample"), (260, "Leader in the"),
                             (380, "Method:")])
            + _table_row(688, [(150, "tracking"), (260, "design,"), (380, "supports the")])
            + _table_row(676, [(150, "Program"), (260, "Review"),
                               (380, "review of results.")])
            + _table_row(664, [(72, "2009"), (150, "Second item"), (260, "Built"),
                               (380, "A new method.")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == [
        "2007\tSample tracking Program\tLeader in the design, Review"
        "\tMethod: supports the review of results.",
        "2009\tSecond item\tBuilt\tA new method."]


def test_cell_empty_in_the_row_does_not_count_as_continued(tmp_path):
    """The second row leaves its third cell empty. The line below fills that
    cell lowercase, but an empty cell has no text to continue, so the line
    continues none of the row's cells and is the next row."""
    page = (_table_row(700, [(72, "2019"), (150, "Prize."), (260, "Lead."), (400, "Boston.")])
            + _table_row(688, [(150, "Medal."), (400, "Paris.")])
            + _table_row(676, [(150, "Gold"), (260, "co-lead")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:3] == ["2019\tPrize.\tLead.\tBoston.", "\tMedal.\t\tParis.", "\tGold\tco-lead"]


@pytest.mark.parametrize("range_above, year_below, same_row", [
    ("2011-", "2014", True), ("2011-", "2014-", False), ("since 2011-", "2014-", False)])
def test_open_year_range_cell_continues_only_into_its_end(tmp_path, range_above, year_below,
                                                         same_row):
    """"2011-" over "2014" is one range wrapped; over "2014-" it is the next
    row's own open range, under the same first cell. The range is found at
    the end of a cell of several words too ("since 2011-")."""
    page = (_table_row(700, [(72, "Example Org"), (200, range_above), (300, "Faculty")])
            + _table_row(688, [(200, year_below), (300, "Coach")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    expected = ([f"Example Org\t{range_above} {year_below}\tFaculty Coach"] if same_row
                else [f"Example Org\t{range_above}\tFaculty", f"\t{year_below}\tCoach"])
    assert _texts(doc)[:len(expected)] == expected


def test_cell_row_continuing_under_half_its_cells_is_the_next_row(tmp_path):
    page = (_table_row(700, [(72, "2007"), (150, "Sample"), (260, "Leader"), (380, "Uses the")])
            + _table_row(688, [(150, "Program"), (260, "Lead"), (380, "results.")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2007\tSample\tLeader\tUses the", "\tProgram\tLead\tresults."]


def test_lowercase_line_continues_a_cell_with_no_connector_end(tmp_path):
    """"Senior research" ends on no connector and no sentence end; "fellow"
    below it opens lowercase, so that cell wraps. It is one of the line's two
    filled cells, half, so the line is the row's own wrap."""
    page = (_table_row(700, [(72, "2019"), (150, "Senior research"), (400, "Lead")])
            + _table_row(688, [(150, "fellow"), (400, "Boston")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[0] == "2019\tSenior research fellow\tLead Boston"


def test_cell_ending_a_sentence_is_not_wrapped_by_a_lowercase_line(tmp_path):
    page = (_table_row(700, [(72, "2019"), (150, "Award."), (400, "Boston.")])
            + _table_row(688, [(150, "runner up"), (400, "and Paris")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tAward.\tBoston.", "\trunner up\tand Paris"]


def test_cell_ending_in_a_colon_is_wrapped_by_the_line_below(tmp_path):
    page = (_table_row(700, [(72, "2019"), (150, "Role:"), (400, "Site:")])
            + _table_row(688, [(150, "co-lead"), (400, "Boston")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[0] == "2019\tRole: co-lead\tSite: Boston"


def test_indented_note_across_cells_is_not_spread_into_them(tmp_path):
    """A note under a row, indented but running across several cells with
    no column gap, is its own paragraph."""
    note = "Note: this appointment was held jointly with the partner institution"
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
            + _table_row(688, [(90, note)]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tTitle\tBoston", note]


def test_cell_continuation_must_end_before_the_next_cell(tmp_path):
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
            + _table_row(688, [(360, "Internationalization")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tTitle\tBoston", "Internationalization"]


def test_line_with_a_gap_inside_one_cell_ends_the_row(tmp_path):
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
            + _table_row(688, [(400, "MA"), (440, "USA")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tTitle\tBoston", "MA\tUSA"]


def test_gapless_line_from_a_later_anchor_across_cells_is_not_a_row(tmp_path):
    line = "A remark that starts under the title and runs on across the city column"
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
            + _table_row(688, [(150, line)]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tTitle\tBoston", line]


def test_line_with_its_own_first_cell_is_not_a_merged_cell_row(tmp_path):
    """One column gap, from the first anchor to the third: not a row of the
    table above (that needs an empty first cell), so what follows it is not
    a cell continuation."""
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
            + _table_row(688, [(72, "2020"), (400, "Paris")])
            + _table_row(676, [(400, "France")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:3] == ["2019\tTitle\tBoston", "2020\tParis", "France"]


def test_empty_middle_cell_keeps_its_tab(tmp_path):
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (280, "Dept"), (380, "Boston")])
            + _table_row(688, [(150, "Lecturer"), (380, "Paris")]) + [_FILL])
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["2019\tTitle\tDept\tBoston", "\tLecturer\t\tParis"]


def test_row_needs_two_column_gaps_and_close_spacing(tmp_path):
    one_gap = _table_row(700, [(72, "2019"), (400, "Boston")]) + _table_row(688, [(150, "MA")])
    spaced = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
              + _table_row(680, [(400, "MA")]))
    _, doc_a = _convert(tmp_path, [one_gap], name="a")
    _, doc_b = _convert(tmp_path, [spaced], name="b")
    assert _texts(doc_a) == ["2019\tBoston", "MA"]
    assert [t for t in _texts(doc_b) if t] == ["2019\tTitle\tBoston", "MA"]


def test_row_continuation_word_goes_to_the_cell_its_x0_falls_in(tmp_path):
    """A word 3pt left of a cell anchor still belongs to that cell; 4pt left
    belongs to the cell before."""
    page = (_table_row(700, [(72, "2019"), (150, "Title"), (400, "Boston")])
            + _table_row(688, [(397, "near")]) + _table_row(676, [(300, "far")]))
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["2019\tTitle far\tBoston near"]


def _gapped_groups(groups, blank_gap):
    """`groups` three-line groups of short lines, 2pt apart inside a group,
    `blank_gap` points between groups."""
    page, y = [], 760.0
    for g in range(groups):
        for i in range(3):
            page.append((False, 10, 72, y, f"Group {g} line {i}"))
            y -= 12
        y -= blank_gap - 2
    return page


def test_blank_threshold_follows_the_documents_blank_line_gap(tmp_path):
    """A Word blank line measured 0.9em: under the 1.0em fallback, over the
    midpoint of this document's line gap (0.2em) and blank gap (0.9em)."""
    report, doc = _convert(tmp_path, [_gapped_groups(12, 9)])
    assert report.blank_paragraphs == 11


def test_blank_threshold_falls_back_without_a_clear_second_mode(tmp_path):
    report, _ = _convert(tmp_path, [_gapped_groups(10, 9)])  # 9 gaps < BLANK_MODE_MIN_COUNT
    assert report.blank_paragraphs == 0


def _gap_pages(gaps, block_at=()):
    """One page of 10pt lines separated by `gaps` (in ems); lines whose
    index is in `block_at` are column starts."""
    lines, top = [], 100.0
    for index, gap in enumerate([0.0] + gaps):
        top += gap * 10
        lines.append(_Line(top=top, bottom=top + 10, x0=72, x1=100, size=10.0,
                           runs=[_Run("x", False, 10.0)], tabs=[], block_start=index in block_at))
        top += 10
    return [lines]


@pytest.mark.parametrize("gaps, block_at, expected", [
    ([0.2] * 24 + [0.9] * 11, (), 0.55),                        # midway between the modes
    ([-0.5] * 30 + [0.2] * 20 + [0.9] * 11, (), 0.55),          # overlapping lines ignored
    ([0.35] * 30 + [0.2] * 20 + [0.9] * 11, range(1, 31), 0.55),  # column starts ignored
    ([0.6] * 30 + [1.3] * 12, (), 1.0),                         # double-spaced: fallback
    ([0.2] * 20 + [2.4] * 12, (), 1.0),                         # never above BLANK_GAP_EM
    ([0.0] * 20 + [0.5] * 12, (), 0.45),                        # never below PARAGRAPH_GAP_EM
    ([0.2] * 30 + [0.6] * 10 + [1.2] * 20, (), 0.7),            # the most frequent wide gap
    ([0.1] * 5 + [0.2] * 30 + [0.9] * 11, (), 0.55),            # the most frequent line gap
])
def test_blank_gap_threshold(gaps, block_at, expected):
    assert _blank_gap_em(_gap_pages(gaps, block_at)) == pytest.approx(expected)


@pytest.mark.parametrize("kern", [0.0, -0.5])
def test_raised_superscript_folds_into_its_line(tmp_path, kern):
    """Also when kerning tucks it 0.5pt under the end of its word."""
    x = 72 + _width(tmp_path, "21") + kern
    page = [(False, 10, 72, 700, "21"), (False, 7.9, x, 704.5, "st"),
            (False, 10, x + 10, 700, "century")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["21st century"]


def test_small_text_away_from_the_line_is_not_a_superscript(tmp_path):
    """Small text raised on the same row but not touching a word (a sidebar
    set smaller) stays its own line."""
    page = [(False, 10, 72, 700, "Main line"), (False, 7.9, 300, 704.5, "aside"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert sorted(_texts(doc)[:2]) == ["Main line", "aside"]


def test_superscript_before_its_word_folds_in(tmp_path):
    x = 72 + _width(tmp_path, "a", 7.9)
    page = [(False, 7.9, 72, 704.5, "a"), (False, 10, x, 700, "Smith")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["aSmith"]


def test_same_size_raised_word_is_not_a_superscript(tmp_path):
    """Touching and overlapping, but as big as its neighbour: its own line."""
    x = 72 + _width(tmp_path, "Line")
    page = [(False, 10, 72, 700, "Line"), (False, 10, x, 704.5, "next"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert sorted(_texts(doc)[:2]) == ["Line", "next"]


def test_small_cluster_folds_only_if_every_word_touches(tmp_path):
    x = 72 + _width(tmp_path, "21")
    page = [(False, 10, 72, 700, "21"), (False, 7.9, x, 704.5, "st"),
            (False, 7.9, 300, 704.5, "aside"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert "21" in _texts(doc) and "st\taside" in _texts(doc)


def test_small_text_touching_but_barely_overlapping_is_not_folded(tmp_path):
    """Raised 9pt: it touches the word but overlaps its line by under half
    of its own height."""
    x = 72 + _width(tmp_path, "Word")
    page = [(False, 10, 72, 700, "Word"), (False, 7.9, x, 709, "up"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert sorted(_texts(doc)[:2]) == ["Word", "up"]


def test_lowered_subscript_folds_into_the_line_above(tmp_path):
    """9pt lowered 4.9pt under 12pt: bottoms 4.3pt apart (its own cluster,
    sorted after its host), overlapping the host by over half its height."""
    x = 72 + _width(tmp_path, "CO", 12)
    page = [(False, 12, 72, 700, "CO"), (False, 9, x, 695.1, "2"), (False, 12, x + 10, 700, "level")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["CO2 level"]


def test_same_size_line_close_above_is_not_a_superscript(tmp_path):
    page = [(False, 10, 72, 700, "first"), (False, 10, 72, 690, "second"), _FILL]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc)[:2] == ["first", "second"]


@pytest.mark.parametrize("gap", [0.0, 0.5])
def test_font_change_inside_a_word_adds_no_space(tmp_path, gap):
    x = 72 + _bold_width(tmp_path, "Word") + gap
    page = [(True, 10, 72, 700, "Word"), (False, 10, x, 700, ", then more")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["Word, then more"]


def test_font_change_across_a_real_space_keeps_it(tmp_path):
    x = 72 + _bold_width(tmp_path, "Word") + 1.1
    page = [(True, 10, 72, 700, "Word"), (False, 10, x, 700, "next")]
    _, doc = _convert(tmp_path, [page])
    assert _texts(doc) == ["Word next"]


def test_combining_diacritic_is_composed():
    words = [{"text": "José", "fontname": "Helvetica", "size": 10.0}]
    assert [r.text for r in _runs(words, [])] == ["José"]


def test_indent_is_measured_from_the_pages_text_left(tmp_path):
    page = [(False, 10, 72, 700, "Flush"), (False, 10, 108, 660, "Indented")]
    _, doc = _convert(tmp_path, [page])
    assert doc.paragraphs[0].paragraph_format.left_indent is None
    assert doc.paragraphs[-1].paragraph_format.left_indent.pt == 36


def test_furniture_window_is_by_position_not_by_text_alone(tmp_path):
    """Repeated at the top of pages 1-2 (page 1's copy kept), the same text
    25pt lower on page 3 (still in the edge band) is not in the window:
    kept too."""
    pages = [[(False, 9, 72, y, "Running Name Header"), (False, 10, 72, 600, f"Body {n}")]
             for n, y in enumerate((765, 765, 740))]
    _, doc = _convert(tmp_path, pages)
    assert " ".join(_texts(doc)).count("Running Name Header") == 2


def test_body_line_inside_a_window_that_overhangs_the_edge_band_is_kept(tmp_path):
    """The header's window runs 12pt down from its top, past the band's
    edge; the same text just below the band is a body line."""
    pages = [[(False, 9, 72, 701, "Running Name Header"), (False, 10, 72, 600, f"Body {n}")]
             for n in range(2)]  # top 83.9: in the band (95.0), window to 95.9
    pages.append([(False, 9, 72, 689.5, "Running Name Header"),  # top 95.4: body
                  (False, 10, 72, 600, "Body 2")])
    _, doc = _convert(tmp_path, pages)
    assert " ".join(_texts(doc)).count("Running Name Header") == 2  # page 1's and page 3's


@pytest.mark.parametrize("drift, removed", [(0, True), (11, True), (13, False)])
def test_running_header_drift_within_the_window_is_furniture(tmp_path, drift, removed):
    """Three pages need two repeats; `drift` apart, only 12pt or less pair up."""
    tops = (0, drift, 2 * drift)
    pages = [[(False, 9, 72, 760 - t, "Running Name Header"), (False, 10, 72, 600, f"Body {n}")]
             for n, t in enumerate(tops)]
    _, doc = _convert(tmp_path, pages)
    assert " ".join(_texts(doc)).count("Running Name Header") == (1 if removed else 3)
