"""Regression test for issue #315: stage-1 segmentation crashed on any
table-bearing CV.

``extract_docx_structure`` stamps table elements with a *string* ``idx``
(``"table_N"``, which ``stage_2`` requires) while paragraphs carry an *int*
``idx``. ``chunk_section`` range-compared the int section boundaries against
every element's ``idx``, so the first table raised
``TypeError: '<=' not supported between instances of 'int' and 'str'``.

Pure fixture: no LLM, no DB, no PII. The docx is synthesized in-memory. The
test builds a section large enough to force chunking, with a table inside it,
and asserts (a) no crash, (b) chunking engaged, and (c) the table content is
retained (guards against the tempting-but-wrong "drop non-int-idx elements"
one-liner). Run with:

    python3 -m pytest src/unified_pipeline/tests/test_word_chunked_table_idx.py -p no:cacheprovider
"""

import sys
from pathlib import Path

from docx import Document

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.segmentation.word_chunked import (  # noqa: E402
    chunk_section,
    table_to_text,
    MAX_CHARS_PER_SECTION,
)
from unified_pipeline.core.docx_structure_extractor import (  # noqa: E402
    extract_docx_structure,
)

TABLE_TOKEN = "DISTINCTIVE_TABLE_TOKEN"


def _build_docx(path: Path) -> None:
    """A header, enough padded paragraphs to exceed MAX_CHARS_PER_SECTION, a
    table inside the section, and a trailing paragraph."""
    doc = Document()
    doc.add_paragraph("INVITATIONS TO SPEAK")

    # Deterministic padding: push the section well past the 12000-char chunk
    # threshold so chunk_section actually splits.
    filler = "Presentation on point-of-care ultrasound at a national venue. " * 5  # ~300 chars
    n_paras = (MAX_CHARS_PER_SECTION // len(filler)) + 5
    for _ in range(n_paras):
        doc.add_paragraph(filler)

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = f"{TABLE_TOKEN} Keynote"
    table.cell(0, 1).text = "2019"
    table.cell(1, 0).text = "Grand Rounds"
    table.cell(1, 1).text = "2020"

    doc.add_paragraph("EDUCATION")
    doc.save(str(path))


def test_chunk_section_handles_table_string_idx(tmp_path):
    path = tmp_path / "table_bearing.docx"
    _build_docx(path)

    structure = extract_docx_structure(str(path))
    elements = structure["elements"]

    # Precondition: at least one table element with a *string* idx exists — this
    # is exactly what crashed the pre-fix comprehension.
    str_idx_tables = [
        e for e in elements
        if e.get("type") == "table" and isinstance(e.get("idx"), str)
    ]
    assert str_idx_tables, "fixture must contain a table element with a string idx"

    # Pre-fix this raised TypeError (int <= str); post-fix it must not.
    chunks = chunk_section(elements, 0, len(elements))

    # Chunking engaged (section exceeded MAX_CHARS_PER_SECTION).
    assert len(chunks) >= 2, f"expected the large section to split; got {len(chunks)} chunk(s)"

    flat = [e for chunk in chunks for e in chunk]

    # Table content retained — NOT dropped by a naive isinstance-int filter.
    surviving_tables = [e for e in flat if e.get("type") == "table"]
    assert surviving_tables, "table element was dropped from the chunked section"

    joined = "\n".join(table_to_text(e) for e in surviving_tables)
    assert TABLE_TOKEN in joined, "distinctive table content lost after chunking"
