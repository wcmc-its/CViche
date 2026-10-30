"""#983: a multi-record entry fans out into one entry per record.

`stage6/fan_out.py` is pure, so every case below is a dict in, dicts out. The
wiring (that `_group_entries_by_code` calls it before the PII pass and dedup)
is pinned in `test_stage6_render_drops.py`, next to the other tests of that
method, and the dedup guard for its children in `test_stage6_dedup_safety.py`.

Every value is invented. The shapes are the ones the 148-CV render set
carries: web199's H `awards`, web218's P `committees` / `entries`, web205's O
`additional_roles`, web241's N3B `mentees`, web181's K2 `mentees(name, year)`
and web228's K4 `sessions`.

    python3 -m pytest src/unified_pipeline/tests/test_stage6_fan_out.py -p no:cacheprovider
"""
import copy
import itertools
import json
import re
import sys
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.stage4.schemas import FIELD_SCHEMAS  # noqa: E402
from unified_pipeline.stage4.schemas import STAGE4_RECORDS_KEY as _RECORDS  # noqa: E402
from unified_pipeline.stage6 import fan_out  # noqa: E402
from unified_pipeline.stage6.fan_out import (  # noqa: E402
    FANNED_OUT_FROM,
    LAST_STAGE4_RECORD,
    fan_out_multi_record_entries,
)

# The drift guard the stage-6 package cannot import stage4 for: the codes this
# file builds entries for still have the fields the fixtures rely on.
_SCHEMA_FIELDS_USED = {
    'H': {'award_name', 'granting_body', 'date'},
    'P': {'committee_name', 'role', 'institution', 'start_date', 'end_date'},
    'O': {'leadership_role', 'institution', 'start_date', 'end_date'},
    'N3B': {'mentee_name', 'mentee_level', 'start_date', 'end_date'},
    'K2': {'teaching_role', 'start_date', 'end_date'},
    'K4': {'activity_title', 'date'},
}


def _fan(*entries):
    return fan_out_multi_record_entries(list(entries), FIELD_SCHEMAS)


def _award(name, body='Hollis College', date='1986'):
    return {'award_name': name, 'granting_body': body, 'date': date}


def _honors(text, awards, **scalars):
    return {'taxonomy_code': 'H', 'text': text, 'element_idx_start': 7,
            'extracted_fields': {'awards': awards, **scalars}}


_THREE_HONORS = _honors(
    '1986    Kappa Delta Honor Society, Hollis College\tBeta Sigma Honor Society, '
    'Hollis College\tPhi Rho History Honor Society, Hollis College',
    [_award('Kappa Delta Honor Society'), _award('Beta Sigma Honor Society'),
     _award('Phi Rho History Honor Society')])


def test_schema_fields_the_fixtures_use_still_exist():
    for code, fields in _SCHEMA_FIELDS_USED.items():
        assert fields <= set(FIELD_SCHEMAS[code]['fields']), code


class TestFanOut:
    def test_one_child_per_record_in_order(self):
        children = _fan(_THREE_HONORS)
        assert [c['extracted_fields']['award_name'] for c in children] == [
            'Kappa Delta Honor Society', 'Beta Sigma Honor Society',
            'Phi Rho History Honor Society']
        assert 'awards' not in children[0]['extracted_fields']

    def test_text_is_the_matching_tab_segment_when_the_counts_line_up(self):
        children = _fan(_THREE_HONORS)
        assert [c['text'] for c in children] == [
            '1986    Kappa Delta Honor Society, Hollis College',
            'Beta Sigma Honor Society, Hollis College',
            'Phi Rho History Honor Society, Hollis College']

    def test_text_is_built_from_the_record_when_a_record_wrapped(self):
        # Four tab segments for two records: the second award wrapped across
        # two lines, so pairing by position would give it the wrong tail.
        parent = _honors(
            '1988    Gamma Omega Honorary, Ashby University\tGraduate School\t'
            'Award Winner, Trenton Society\tStudent Workshop',
            [_award('Gamma Omega Honorary', 'Ashby University Graduate School', '1988'),
             _award('Award Winner, Trenton Society Student Workshop', 'Trenton Society', '1988')])
        children = _fan(parent)
        assert [c['text'] for c in children] == [
            'Gamma Omega Honorary | Ashby University Graduate School | 1988',
            'Award Winner, Trenton Society Student Workshop | Trenton Society | 1988']

    def test_provenance_names_the_list_and_position(self):
        children = _fan(_THREE_HONORS)
        assert [c[FANNED_OUT_FROM] for c in children] == [
            {'key': 'awards', 'index': 0, 'count': 3},
            {'key': 'awards', 'index': 1, 'count': 3},
            {'key': 'awards', 'index': 2, 'count': 3}]

    def test_item_wins_over_parent_scalar_and_scalars_are_inherited(self):
        parent = _honors('Alpha Prize, Hollis College\tBeta Prize, Hollis College',
                         [{'award_name': 'Alpha Prize', 'date': '1990'},
                          {'award_name': 'Beta Prize'}],
                         date='1985', granting_body='Hollis College')
        first, second = _fan(parent)
        assert first['extracted_fields'] == {
            'date': '1990', 'granting_body': 'Hollis College', 'award_name': 'Alpha Prize'}
        assert second['extracted_fields']['date'] == '1985'

    def test_other_entry_keys_are_copied_and_the_parent_is_not_mutated(self):
        original = copy.deepcopy(_THREE_HONORS)
        children = _fan(_THREE_HONORS)
        assert _THREE_HONORS == original
        assert all(c['element_idx_start'] == 7 and c['taxonomy_code'] == 'H'
                   for c in children)
        children[0]['extracted_fields']['award_name'] = 'changed'
        assert children[1]['extracted_fields']['award_name'] == 'Beta Sigma Honor Society'

    def test_entries_around_a_fanned_one_keep_their_place_and_identity(self):
        before = {'taxonomy_code': 'H', 'text': 'Solo Prize', 'extracted_fields': {'award_name': 'Solo Prize'}}
        after = {'taxonomy_code': 'T', 'text': 'anything', 'extracted_fields': {}}
        out = _fan(before, _THREE_HONORS, after)
        assert len(out) == 5
        assert out[0] is before and out[-1] is after

    def test_parent_remark_and_formatter_keys(self):
        parent = _honors('Alpha Prize, Hollis College\tBeta Prize, Hollis College',
                         [_award('Alpha Prize'), _award('Beta Prize')],
                         notes='Entry contains two tab-separated honors')
        children = _fan(parent)
        assert all('notes' not in c['extracted_fields'] for c in children)


class TestWhatIsNotFannedOut:
    def _unchanged(self, entry):
        assert _fan(entry) == [entry]

    def test_a_list_of_strings(self):
        self._unchanged(_honors('Alpha Prize\tBeta Prize', ['Alpha Prize', 'Beta Prize']))

    def test_the_schemas_own_field(self):
        # `award_name` IS a schema field: a list there is stage 4's value, not a
        # record list under a foreign key.
        self._unchanged({'taxonomy_code': 'H', 'text': 'Alpha Prize Beta Prize',
                         'extracted_fields': {'award_name': [_award('Alpha Prize'), _award('Beta Prize')]}})

    def test_a_list_holding_an_empty_record(self):
        # The text is covered by the second record alone, so only the guard on
        # an empty record keeps the entry whole.
        self._unchanged(_honors('Beta Prize, Hollis College', [{}, _award('Beta Prize')]))

    def test_records_that_share_no_key_with_the_schema(self):
        # web181's K2 `mentees: [{name, year}]`: nothing the K2 renderer reads.
        self._unchanged({'taxonomy_code': 'K2', 'text': 'Mentored Ana Cruz 2019\tLee Park 2020',
                         'extracted_fields': {'mentees': [{'name': 'Ana Cruz', 'year': '2019'},
                                                          {'name': 'Lee Park', 'year': '2020'}]}})

    def test_records_carrying_a_key_no_renderer_reads(self):
        # web228's K4 `sessions`: `title` is the session's own name. The text is
        # written so the fields cover it: only the key guard keeps it whole.
        sessions = [{'date': '2020-04-22', 'title': 'Definition and Clinical Evaluation'},
                    {'date': '2020-04-29', 'title': 'Complexity and Impact'}]
        self._unchanged({'taxonomy_code': 'K4',
                         'text': 'Practice Improvement Project 2020\tPractice Improvement Project 2020',
                         'extracted_fields': {'activity_title': 'Practice Improvement Project',
                                              'sessions': sessions}})

    def test_an_entry_a_stage5_formatter_already_rendered_whole(self):
        parent = _honors('Alpha Prize, Hollis College\tBeta Prize, Hollis College',
                         [_award('Alpha Prize'), _award('Beta Prize')],
                         formatted_text='- Alpha Prize\n- Beta Prize')
        self._unchanged(parent)

    def test_an_entry_with_two_record_lists(self):
        # Covered by the first list alone: only the one-list rule keeps it whole.
        records = [{'committee_name': 'Alpha Board'}, {'committee_name': 'Beta Board'}]
        entry = {'taxonomy_code': 'P', 'text': 'Alpha Board\tBeta Board',
                 'extracted_fields': {'committees': records, 'entries': list(records)}}
        self._unchanged(entry)

    def test_a_code_with_no_schema_and_a_non_mapping_field_bag(self):
        self._unchanged({'taxonomy_code': 'ZZ', 'text': 't', 'extracted_fields': {'x': [{'a': 1}, {'a': 2}]}})
        self._unchanged({'taxonomy_code': 'H', 'text': 't', 'extracted_fields': None})
        self._unchanged({'taxonomy_code': 'H', 'text': 't'})


_COMMITTEES = [
    {'committee_name': 'Alpha Curriculum Oversight Council', 'role': 'Member',
     'institution': 'Ashby University'},
    {'committee_name': 'Beta Admissions Selection Council', 'role': 'Chair',
     'institution': 'Ashby University'}]


class TestTextCoverage:
    """Every token of the entry's text must be held by a field the renderer
    writes: a child renders from those fields alone, so any other token is
    content the output would no longer have."""

    _RECORDS = ('Alpha Curriculum Oversight Council member{a}\t'
                'Beta Admissions Selection Council chair{b}')

    def _committee_entry(self, text, committees=_COMMITTEES, **scalars):
        return {'taxonomy_code': 'P', 'text': text,
                'extracted_fields': {'committees': committees, **scalars}}

    def test_text_the_fields_hold_is_fanned_out(self):
        assert len(_fan(self._committee_entry(self._RECORDS.format(a='', b='')))) == 2

    def test_prose_around_the_records_keeps_the_entry_whole(self):
        prose = ('I serve on the Alpha Curriculum Oversight Council as a member and I chair the '
                 'Beta Admissions Selection Council, which reviews proposals every semester')
        entry = self._committee_entry(prose)
        assert _fan(entry) == [entry]

    @pytest.mark.parametrize('qualifier', ['(weekly)', '(1-3 committees/yr)', 'UXB'])
    def test_one_word_no_field_holds_keeps_the_entry_whole(self, qualifier):
        # No tolerance: web218's "(1-3 committees/yr)" and its institution
        # abbreviation were lost.
        entry = self._committee_entry(self._RECORDS.format(a=' ' + qualifier, b=''))
        assert _fan(entry) == [entry]

    def test_a_field_the_renderer_never_writes_holds_nothing(self):
        # `description` is a P schema field the P renderer never writes, so
        # words only it holds would be lost (web240's case before #985, with
        # `institution`).
        entry = self._committee_entry(self._RECORDS.format(a=' curriculum review', b=''),
                                      description='curriculum review')
        assert 'description' in FIELD_SCHEMAS['P']['fields']
        assert 'description' not in fan_out._RENDERED_FIELDS['P']
        assert _fan(entry) == [entry]

    def test_an_institution_p_renders_holds_its_words(self):
        # #985: P appends `institution` to the name cell, so its words are held.
        text = self._RECORDS.format(a=' Ashby University', b=' Ashby University')
        entry = self._committee_entry(text, committees=[
            dict(record, institution='Ashby University') for record in _COMMITTEES])
        assert len(_fan(entry)) == 2

    def test_a_key_outside_the_schema_holds_nothing(self):
        entry = self._committee_entry(self._RECORDS.format(a=' weekly', b=''), cadence='weekly')
        assert _fan(entry) == [entry]

    def test_a_word_the_text_repeats_must_be_held_as_often(self):
        # 'Council' once in the fields, twice in the text.
        text = 'Alpha Board member\tBeta Board chair\tCouncil Council'
        entry = self._committee_entry(text, [{'committee_name': 'Alpha Board', 'role': 'member'},
                                             {'committee_name': 'Beta Board', 'role': 'chair'}],
                                      role='Council')
        assert _fan(entry) == [entry]

    def test_a_scalar_counts_once_per_child(self):
        # The scalar `granting_body` is written on each of the two children,
        # so a text that names it twice is covered.
        entry = _honors('Alpha Prize, Hollis College\tBeta Prize, Hollis College',
                        [{'award_name': 'Alpha Prize'}, {'award_name': 'Beta Prize'}],
                        granting_body='Hollis College')
        assert len(_fan(entry)) == 2

    def test_connectives_are_not_counted(self):
        text = 'Chair of the Board\tMember of Alpha in Beta for Gamma on a Delta at an Epsilon to Zeta and Eta'
        entry = self._committee_entry(text, [{'committee_name': 'Board', 'role': 'Chair'},
                                             {'committee_name': 'Alpha Beta Gamma Delta Epsilon Zeta Eta',
                                              'role': 'Member'}])
        assert len(_fan(entry)) == 2

    def test_a_year_no_field_holds_keeps_the_entry_whole(self):
        entry = self._committee_entry(self._RECORDS.format(a=' 1999', b=''))
        assert _fan(entry) == [entry]

    def test_a_year_a_field_holds_is_fine(self):
        entry = self._committee_entry(self._RECORDS.format(a=' 1999', b=''), start_date='1999')
        assert len(_fan(entry)) == 2

    def test_a_month_the_date_column_does_not_show_is_not_content(self):
        # EBYSBC E6: P's date column is `yyyy`, so "07/2009" is written as
        # "2009" whether the entry is fanned out or not; a date is compared by
        # its year (it used to keep the entry whole).
        entry = self._committee_entry(self._RECORDS.format(a=' 07/2009', b=''), start_date='07/2009')
        assert len(_fan(entry)) == 2

    def test_a_day_the_date_column_does_not_show_still_keeps_the_entry_whole(self):
        # Only month/year is reduced to its year: a full m/d/y date is left as
        # written, so its day is a token no yyyy column holds.
        entry = self._committee_entry(self._RECORDS.format(a=' 07/15/2009', b=''), start_date='2009')
        assert _fan(entry) == [entry]

    def test_a_date_range_the_column_shows_in_full_is_fine(self):
        entry = self._committee_entry(self._RECORDS.format(a=' 2009 to 2011', b=''),
                                      start_date='2009-07', end_date='2011-03')
        assert len(_fan(entry)) == 2

    def test_a_date_looking_value_in_a_non_date_field_is_held_as_written(self):
        # Only date fields go through the date formatter: `role` "July 2009"
        # would lose "July" if it did.
        entry = self._committee_entry('Alpha Board July 2009\tBeta Board chair', [
            {'committee_name': 'Alpha Board', 'role': 'July 2009'},
            {'committee_name': 'Beta Board', 'role': 'chair'}])
        assert len(_fan(entry)) == 2

    @pytest.mark.parametrize('role', [['chair', 'member'], {'title': 'chair member'}])
    def test_a_list_or_mapping_value_is_held_by_its_leaves(self, role):
        entry = self._committee_entry('Alpha Board chair member\tBeta Board', [
            {'committee_name': 'Alpha Board', 'role': role}, {'committee_name': 'Beta Board'}])
        assert len(_fan(entry)) == 2

    def test_a_missing_value_holds_no_word(self):
        # `None` must not be read as the word "None".
        entry = self._committee_entry('Alpha Board none\tBeta Board', [
            {'committee_name': 'Alpha Board', 'role': None}, {'committee_name': 'Beta Board'}])
        assert _fan(entry) == [entry]

    def test_an_entry_with_no_text_has_nothing_to_cover(self):
        entry = {'taxonomy_code': 'P', 'extracted_fields': {'committees': _COMMITTEES}}
        assert len(_fan(entry)) == 2

    def test_a_code_the_renderer_map_does_not_name_is_never_fanned_out(self, monkeypatch):
        entry = self._committee_entry(self._RECORDS.format(a='', b=''))
        monkeypatch.setattr(fan_out, '_RENDERED_FIELDS', {})
        assert _fan(entry) == [entry]


class TestChildrenShareNothingWithTheirParent:
    def test_nested_values_are_copies(self):
        parent = _honors('Alpha Prize, Hollis College\tBeta Prize, Hollis College',
                         [dict(_award('Alpha Prize'), description=['first']), _award('Beta Prize')],
                         venue=['hall'])
        parent['hierarchy'] = ['HONORS']
        first, second = _fan(parent)
        first['hierarchy'].append('changed')
        first['extracted_fields']['description'].append('changed')
        first['extracted_fields']['venue'].append('changed')
        assert second['hierarchy'] == ['HONORS'] and parent['hierarchy'] == ['HONORS']
        assert parent['extracted_fields']['awards'][0]['description'] == ['first']
        assert second['extracted_fields']['venue'] == ['hall']
        assert parent['extracted_fields']['venue'] == ['hall']


def test_built_text_leaves_out_empty_values():
    parent = _honors('Alpha Prize\tBeta\tPrize',
                     [{'award_name': 'Alpha Prize', 'granting_body': '', 'date': '1988'},
                      {'award_name': 'Beta Prize', 'granting_body': None, 'date': '1989'}])
    assert [c['text'] for c in _fan(parent)] == ['Alpha Prize | 1988', 'Beta Prize | 1989']


_LATER_POSTS = [
    {'leadership_role': 'Leader, Genomics Program', 'start_date': '2015', 'end_date': '2022'},
    {'leadership_role': 'Deputy Director', 'start_date': '2022', 'end_date': 'present'}]


class TestParentOwnRecord:
    """web205's O: the Co-Leader post is in the parent's scalars, the later
    posts are `additional_roles`. web36's P stage-5d shape repeats the first
    post in both."""

    def _o(self, text, **scalars):
        return {'taxonomy_code': 'O', 'text': text,
                'extracted_fields': {'additional_roles': _LATER_POSTS, **scalars}}

    def test_one_more_segment_than_records_means_the_parent_is_a_record(self):
        parent = self._o('2012- Co-Leader, Genomics Program\t2015-2022 Leader, Genomics Program\t'
                         '2022- Deputy Director',
                         leadership_role='Co-Leader, Genomics Program', start_date='2012')
        out = _fan(parent)
        assert [c['extracted_fields']['leadership_role'] for c in out] == [
            'Co-Leader, Genomics Program', 'Leader, Genomics Program', 'Deputy Director']
        assert out[0]['text'] == '2012- Co-Leader, Genomics Program'
        assert out[0]['extracted_fields']['start_date'] == '2012'
        assert [c[FANNED_OUT_FROM]['count'] for c in out] == [3, 3, 3]

    def test_as_many_segments_as_records_means_the_list_holds_them_all(self):
        parent = self._o('2015-2022 Leader, Genomics Program\t2022- Deputy Director',
                         leadership_role='Leader, Genomics Program', start_date='2015')
        out = _fan(parent)
        assert len(out) == 2

    def test_no_identity_value_means_no_extra_record(self):
        # `mentee_level` is shared context, not a record of its own, even when
        # the text has a header segment ("NIH T32 (Post-graduate):").
        parent = {'taxonomy_code': 'N3B',
                  'text': 'Post-graduate:\tAna Cruz (2013)\tLee Park (2014)',
                  'extracted_fields': {'mentee_level': 'Post-graduate', 'funding_source': 'Fellowship',
                                       'mentees': [{'mentee_name': 'Ana Cruz', 'start_date': '2013'},
                                                   {'mentee_name': 'Lee Park', 'start_date': '2014'}]}}
        out = _fan(parent)
        assert [c['extracted_fields']['mentee_name'] for c in out] == ['Ana Cruz', 'Lee Park']
        assert all(c['extracted_fields']['mentee_level'] == 'Post-graduate' for c in out)

    def test_a_parent_an_item_repeats_is_not_emitted_twice(self):
        parent = {'taxonomy_code': 'P',
                  'text': 'Interim Director, North Wing\tInterim Director, South Wing',
                  'extracted_fields': {'role': 'Interim Director', 'additional_roles': [
                      {'role': 'Interim Director', 'committee_name': 'North Wing'},
                      {'role': 'Interim Director', 'committee_name': 'South Wing'}, ]}}
        assert len(_fan(parent)) == 2

    def test_counts_that_say_neither_fall_back_to_the_repeat_test(self):
        # Four segments for two records: a wrapped line, so not decisive.
        distinct = self._o('2012- Co-Leader, Genomics\tProgram\t2015-2022 Leader, Genomics\tProgram',
                           leadership_role='Co-Leader, Genomics Program', start_date='2012')
        assert len(_fan(distinct)) == 3
        repeated = self._o('2015-2022 Leader, Genomics\tProgram\t2022- Deputy\tDirector',
                           leadership_role='Deputy Director', start_date='2022')
        assert len(_fan(repeated)) == 2


    def test_a_date_the_items_do_not_repeat_is_not_an_identity(self):
        # Four segments for two records (a wrapped line, counts undecided) and
        # a parent whose only scalars are dates that differ from both items':
        # the dates are context, so no extra record.
        parent = {'taxonomy_code': 'P',
                  'text': 'Alpha Board of\tGovernors\tBeta Board of\tTrustees',
                  'extracted_fields': {'start_date': '2010', 'committees': [
                      {'committee_name': 'Alpha Board of Governors', 'start_date': '2016'},
                      {'committee_name': 'Beta Board of Trustees', 'start_date': '2017'}]}}
        assert len(_fan(parent)) == 2

    @pytest.mark.parametrize('blank', [None, '', '   ', [], {}])
    def test_a_blank_identity_field_is_not_an_identity(self, blank):
        # One header segment more than records, and a parent whose only
        # identity-shaped scalar is blank: not a record of its own.
        parent = {'taxonomy_code': 'P', 'text': 'Council:\tAlpha Board\tBeta Board',
                  'extracted_fields': {'committee_name': blank, 'role': 'Council', 'committees': [
                      {'committee_name': 'Alpha Board'}, {'committee_name': 'Beta Board'}]}}
        assert len(_fan(parent)) == 2


@pytest.mark.parametrize('key,expected', [('start_date', True), ('date', True), ('year', True),
                                          ('end_date', True), ('dates_attended', True),
                                          ('dates_attended_start_date', True), ('role', False), ('mentee_name', False)])
def test_date_keys_never_count_as_a_parents_identity(key, expected):
    from unified_pipeline.stage6.fan_out import _is_date_key
    assert _is_date_key(key) is expected


# --- the renderer map, checked against the renderers themselves ---------------

_NEUTRAL_TEXT = 'Xyneutral probe line'


def _rendered_markers(tmp_path, code, extra_fields=frozenset()):
    """The schema fields of `code` (and `extra_fields`) whose marker reaches the
    .docx when one entry of that code is rendered with a unique marker in every
    field, and whether the entry's own (neutral) text reaches it. The text must
    not carry the markers: a section that writes the text would then "render"
    every field."""
    markers = {name: 'Zq' + name.replace('_', 'x') + 'Z'
               for name in {*FIELD_SCHEMAS[code]['fields'], *extra_fields} if name != 'narrative'}
    body = _rendered_body(tmp_path, code, markers)
    return {name for name, marker in markers.items() if marker in body}, _NEUTRAL_TEXT in body


def _rendered_body(tmp_path, code, fields):
    """Every `w:t` of the .docx one entry of `code` with `fields` renders to."""
    from docx import Document

    from unified_pipeline.stage_6_word_template import WCMTemplateGenerator

    entry = {'taxonomy_code': code, 'element_idx_start': 0, 'text': _NEUTRAL_TEXT,
             'extracted_fields': dict(fields)}
    source, target = tmp_path / 'in.json', tmp_path / 'out.docx'
    source.write_text(json.dumps({'document_uid': 'TESTAA', 'entries': [entry]}))
    generator = WCMTemplateGenerator(verbose=False)
    generator.generate(str(source), str(target), research_summary_path=None)
    document = Document(str(target))
    return ' '.join(t.text or '' for t in document.element.body.iter(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))


# Fields the map keeps although a probe entry of the code does not render
# them. A Past (Completed) Funding table leaves out the effort row (#1944),
# but an M2B record re-bucketed into Current or Pending still renders it, and
# the supplementary pass reads the map to tell a grant's own record line from
# its prose: dropping `percent_effort` here offered six record and goals lines
# on three corpus CVs as new sub-points.
_KEPT_THOUGH_UNRENDERED = {'M2B': {'percent_effort'}}


@pytest.mark.parametrize('code', sorted(fan_out._RENDERED_FIELDS))
def test_rendered_fields_match_what_each_section_writes(tmp_path, code):
    """`_RENDERED_FIELDS[code]` is exactly the set of fields the code's section
    writes: a field the map claims but the renderer skips would let the
    coverage test pass on content the output drops (web240's `institution`),
    and one it omits would needlessly keep entries whole."""
    expected = set(fan_out._RENDERED_FIELDS[code]) - _KEPT_THOUGH_UNRENDERED.get(code, set())
    assert _rendered_markers(tmp_path, code) == (expected, False)


def _offschema_date_fields(code):
    """The date fields stage 4 may emit for `code` that its `FIELD_SCHEMAS`
    entry does not name: the active (config-file) schema's and the field
    descriptions' (#1556 d: Q4D's start/end)."""
    from unified_pipeline.stage4.schemas import FIELD_DESCRIPTIONS, get_active_schemas

    active = get_active_schemas().get(code, {}).get('fields', [])
    named = set(active) | set(FIELD_DESCRIPTIONS.get(code, {}))
    return {key for key in named - set(FIELD_SCHEMAS[code]['fields']) if fan_out._is_date_key(key)}


@pytest.mark.parametrize('code', sorted(
    code for code in fan_out._RENDERED_FIELDS if _offschema_date_fields(code)))
def test_offschema_date_fields_a_section_writes_are_the_ones_fan_out_counts(tmp_path, code):
    """A date field outside the schema that the section writes must be a
    written date to `_is_written_date`, or `_repeated_rows` reads two spans of
    one name as one row (#1556 d); one it does not write must not be."""
    extra = _offschema_date_fields(code)
    written = _rendered_markers(tmp_path, code, extra)[0] & extra
    markers = {key: 'Zq' + key for key in extra}
    assert written == {key for key in extra
                       if fan_out._is_written_date(code, key, markers[key], markers)}


# #1901: years no template text holds, one per field.
_H_DATES = {'date': '1987', 'start_date': '1993', 'end_date': '1999'}


@pytest.mark.parametrize('date_key', ['date', 'year'])
@pytest.mark.parametrize('present', [
    combo for n in range(1, 4) for combo in itertools.combinations(sorted(_H_DATES), n)],
    ids='+'.join)
def test_h_bounds_fan_out_counts_are_the_ones_honors_writes(tmp_path, present, date_key):
    """H's schema names no start/end, but `honors._record_date` writes them
    unless the `date` (or `year`) wins, so `_is_written_date` must say which
    bound the date cell shows, or `_repeated_rows` reads two spans of one
    award as one row (#1901, NECRGC EYTHNB 164) or two rows printed alike as
    two."""
    fields = {'award_name': 'Zqaward of Lanterns',
              **{date_key if key == 'date' else key: _H_DATES[key] for key in present}}
    body = _rendered_body(tmp_path, 'H', fields)
    bounds = [key for key in ('start_date', 'end_date') if key in present]
    assert [key for key in bounds if _H_DATES[key] in body] == [
        key for key in bounds if fan_out._is_written_date('H', key, fields[key], fields)]


@pytest.mark.parametrize('code', ['B2', 'Q4D'])
def test_only_b1_writes_a_nested_dates_attended(code):
    """B2 and Q4D write a flat start/end only; a `dates_attended` on one of
    their records is not in the row, so it must not tell two rows apart."""
    nested = {'start_date': '1990', 'end_date': '1992'}
    assert not fan_out._is_written_date(code, 'dates_attended', nested, {})
    assert fan_out._is_written_date('B1', 'dates_attended', nested, {})


@pytest.mark.parametrize('code', sorted(fan_out._TEXT_RENDERED_CODES))
def test_text_rendered_codes_write_the_text_and_no_field(tmp_path, code):
    """A section that writes the entry's text is never fanned out: its children
    would render only their built text and drop the parent's scalars (a G
    entry's organization). If one starts writing fields, move it into the map."""
    assert code not in fan_out._RENDERED_FIELDS
    assert _rendered_markers(tmp_path, code) == (set(), True)


def test_stopwords_are_exactly_the_connectives():
    """A content word added here would let a fan-out drop it silently."""
    assert fan_out._STOPWORDS == {'a', 'an', 'and', 'at', 'for', 'in', 'of', 'on', 'the', 'to'}


def test_a_text_rendered_entry_with_a_record_list_is_left_whole():
    """The verifier's G repro: the parent's organization is in no child's text."""
    entry = {'taxonomy_code': 'G', 'element_idx_start': 0,
             'text': 'Kestrel Harbor Hospital\tAttending Physician 2010\tConsultant 2012',
             'extracted_fields': {'organization': 'Kestrel Harbor Hospital', 'affiliations': [
                 {'affiliation_type': 'Attending Physician', 'start_date': '2010'},
                 {'affiliation_type': 'Consultant', 'start_date': '2012'}]}}
    assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS) == [entry]


def test_the_renderer_map_names_only_schema_fields_and_never_narrative():
    for code, names in fan_out._RENDERED_FIELDS.items():
        assert names <= set(FIELD_SCHEMAS[code]['fields']), code
        assert 'narrative' not in names, code


def test_every_code_with_a_schema_is_in_the_renderer_map_except_text_codes_and_personal_data():
    assert set(FIELD_SCHEMAS) - set(fan_out._RENDERED_FIELDS) == fan_out._TEXT_RENDERED_CODES | {'A'}


# --- #1187: B1 degrees carrying dates_attended, and the warning for a refusal ---

def _degree(degree, institution, year, **dates):
    return {'degree': degree, 'institution': institution, 'year': year, **dates}


def _two_degrees(dates_attended):
    first = _degree('BSc', 'Harrowfield University', '2005', **dates_attended[0])
    second = _degree('MSc', 'Birch Hollow University', '2008', **dates_attended[1])
    return {'taxonomy_code': 'B1', 'element_idx_start': 3,
            'text': 'BSc Harrowfield University 2001 2005 2005\tMSc Birch Hollow University 2006 2008 2008',
            'extracted_fields': {'degrees': [first, second]}}


_NESTED_DATES = [{'dates_attended': {'start_date': '2001', 'end_date': '2005'}},
                 {'dates_attended': {'start_date': '2006', 'end_date': '2008'}}]


class TestB1DatesAttended:
    def test_two_degrees_with_dates_attended_fan_out(self):
        children = _fan(_two_degrees(_NESTED_DATES))
        assert [c['extracted_fields']['degree'] for c in children] == ['BSc', 'MSc']
        assert children[1]['extracted_fields']['dates_attended'] == _NESTED_DATES[1]['dates_attended']

    def test_a_degree_key_nobody_reads_still_keeps_the_list_whole(self):
        entry = _two_degrees(_NESTED_DATES)
        entry['extracted_fields']['degrees'][0]['gpa'] = '3.9'
        assert _fan(entry) == [entry]

    def test_a_string_dates_attended_holds_its_tokens_when_no_start_end_exists(self):
        # The B1 renderer writes a string `dates_attended` into the Dates cell
        # when no start/end builds a range, so the 2001/2006 start years the
        # entry's text names are held and the degrees fan out.
        string_dates = [{'dates_attended': '2001-2005'}, {'dates_attended': '2006-2008'}]
        children = _fan(_two_degrees(string_dates))
        assert [c['extracted_fields']['degree'] for c in children] == ['BSc', 'MSc']

    def test_a_string_dates_attended_holds_nothing_beside_a_start_date(self):
        # With a flat start/end the renderer builds the range from those and
        # never writes the string, so its tokens are not held.
        entry = _two_degrees([{'dates_attended': '2001-2005', 'start_date': '2005'},
                              {'dates_attended': '2006-2008', 'start_date': '2008'}])
        children = [dict(item) for item in entry['extracted_fields']['degrees']]
        assert not fan_out._fields_carry_text(entry, children)
        # B1 skips that test, so the degrees still fan out; the years it
        # leaves unheld are what the warning reports.
        assert [c['extracted_fields']['degree'] for c in _fan(entry)] == ['BSc', 'MSc']

    def test_discipline_is_a_rendered_field_so_a_degree_with_one_fans_out(self):
        entry = _two_degrees(_NESTED_DATES)
        entry['text'] += ' Zzfield'
        entry['extracted_fields']['degrees'][0]['discipline'] = 'Zzfield'
        assert [c['extracted_fields']['degree'] for c in _fan(entry)] == ['BSc', 'MSc']

    def test_a_section_label_and_a_place_no_field_holds_do_not_keep_the_degrees_whole(self):
        # B1 writes no text, so declining on a word no field holds lost the
        # degrees too (#1187: four corpus entries).
        entry = _two_degrees(_NESTED_DATES)
        entry['text'] = 'Qualifications:\t' + entry['text'] + ', Zzcountry'
        children = _fan(entry)
        assert [c['extracted_fields']['degree'] for c in children] == ['BSc', 'MSc']

    def test_another_code_still_keeps_an_entry_whole_for_a_word_no_field_holds(self):
        entry = {'taxonomy_code': 'Q2', 'element_idx_start': 3,
                 'text': 'Zzlabel: Alpha Council 2001\tBeta Council 2004',
                 'extracted_fields': {'committees': [
                     {'committee_name': 'Alpha Council', 'start_date': '2001'},
                     {'committee_name': 'Beta Council', 'start_date': '2004'}]}}
        assert _fan(entry) == [entry]

    def test_both_degrees_render(self, tmp_path):
        self._assert_degrees_render(tmp_path, _two_degrees(_NESTED_DATES))

    def test_both_degrees_render_beside_a_section_label_and_a_place(self, tmp_path):
        # The text shape the coverage skip is for (#1187 "Check" line).
        entry = _two_degrees(_NESTED_DATES)
        entry['text'] = 'Qualifications:\t' + entry['text'] + ', Zzcountry'
        self._assert_degrees_render(tmp_path, entry)

    @staticmethod
    def _assert_degrees_render(tmp_path, entry):
        import json

        from docx import Document

        from unified_pipeline.stage_6_word_template import WCMTemplateGenerator
        source, target = tmp_path / 'in.json', tmp_path / 'out.docx'
        source.write_text(json.dumps({'document_uid': 'TESTAA', 'entries': [entry]}))
        WCMTemplateGenerator(verbose=False).generate(
            str(source), str(target), research_summary_path=None)
        body = ' '.join(t.text or '' for t in Document(str(target)).element.body.iter(
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        for name in ('Harrowfield University', 'Birch Hollow University', 'BSc', 'MSc'):
            assert name in body


class TestRejectedListWarning:
    def _warnings(self, *entries):
        warnings = []
        fan_out_multi_record_entries(list(entries), FIELD_SCHEMAS, warnings)
        return warnings

    def test_an_unknown_item_key_warns_with_key_names_only(self):
        entry = _two_degrees(_NESTED_DATES)
        entry['extracted_fields']['degrees'][0]['gpa'] = '3.9'
        (warning,) = self._warnings(entry)
        assert warning['check'] == fan_out.REJECTED_LIST_CHECK
        assert (warning['code'], warning['severity']) == ('B1', 'WARN')
        assert 'gpa' in warning['evidence']
        assert '3.9' not in str(warning)

    def test_a_degree_stage_4_left_out_of_a_fanned_out_list_warns(self):
        # The text names three degrees and the list holds two. B1 skips the
        # coverage test, so the list fans out, and the third degree's year is
        # reported instead of dropped without a signal (#1187).
        entry = _two_degrees(_NESTED_DATES)
        entry['text'] += '\tPhD Quillmere University 2012'
        assert len(_fan(entry)) == 2
        (warning,) = self._warnings(entry)
        assert (warning['check'], warning['code'], warning['severity']) == (
            fan_out.REJECTED_LIST_CHECK, 'B1', 'WARN')
        assert warning['evidence'] == ['B1.degrees: 1 entry']
        assert '2012' not in str(warning) and 'Quillmere' not in str(warning)

    def test_a_stage4_records_list_with_attendance_years_does_not_warn(self):
        # Years held only by attendance ranges in the text are not a missing
        # record when stage 4 split the reply itself (#1187).
        records = [{'degree': 'BS', 'institution': 'Harrowfield University', 'year': '1999'},
                   {'degree': 'MD', 'institution': 'Birch Hollow University', 'year': '2003'}]
        entry = _stage4_entry(records, code='B1', text=(
            'Harrowfield University 1995-1999 BS\tBirch Hollow University 1999-2003 MD'))
        warnings = []
        children = fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings, _RECORDS)
        assert len(children) == 2
        assert warnings == []

    def test_words_no_field_holds_without_a_year_do_not_warn(self):
        entry = _two_degrees(_NESTED_DATES)
        entry['text'] = 'Qualifications:\t' + entry['text'] + ', Zzcountry'
        assert len(_fan(entry)) == 2
        assert self._warnings(entry) == []

    def test_a_year_no_record_holds_does_not_warn_for_a_code_that_tests_coverage(self):
        # Q2 still keeps the entry whole on any unheld token, and then the
        # existing declined-list warning is the only one raised.
        entry = {'taxonomy_code': 'Q2', 'element_idx_start': 3,
                 'text': 'Alpha Council 2001\tBeta Council 2004\tGlade Panel 2009',
                 'extracted_fields': {'committees': [
                     {'committee_name': 'Alpha Council', 'start_date': '2001'},
                     {'committee_name': 'Beta Council', 'start_date': '2004'}]}}
        (warning,) = self._warnings(entry)
        assert 'was split' not in warning['message']

    def test_a_declined_record_list_warns_without_stray_keys(self):
        entry = _two_degrees([{}, {}])
        entry['extracted_fields']['formatted_text'] = 'Rendered whole by stage 5c'
        (warning,) = self._warnings(entry)
        assert warning['evidence'] == ['B1.degrees: 1 entry']

    def test_a_code_that_keeps_the_entry_text_does_not_warn(self):
        entry = _two_degrees([{}, {}])
        entry['text'] = entry['text'] + ' Zzunheld'
        entry['taxonomy_code'] = 'K4'
        assert self._warnings(entry) == []

    def test_one_warning_per_code_and_key_counts_the_entries(self):
        entry = _two_degrees(_NESTED_DATES)
        entry['extracted_fields']['degrees'][0]['gpa'] = '3.9'
        (warning,) = self._warnings(entry, copy.deepcopy(entry))
        assert warning['evidence'][0] == 'B1.degrees: 2 entries'

    def test_a_fanned_out_or_single_record_entry_does_not_warn(self):
        single = {'taxonomy_code': 'B1', 'text': 'BSc Harrowfield University 2005',
                  'extracted_fields': _degree('BSc', 'Harrowfield University', '2005')}
        assert self._warnings(_two_degrees(_NESTED_DATES), single) == []

    def test_no_warnings_list_is_backward_compatible(self):
        assert len(_fan(_two_degrees(_NESTED_DATES))) == 2

    def test_the_generator_puts_the_warning_in_the_sidecar_list(self):
        from unified_pipeline.stage_6_word_template import WCMTemplateGenerator
        entry = _two_degrees(_NESTED_DATES)
        entry['extracted_fields']['degrees'][0]['gpa'] = '3.9'
        generator = WCMTemplateGenerator(verbose=False)
        generator._group_entries_by_code([entry])
        assert [w['check'] for w in generator._section_failures] == [fan_out.REJECTED_LIST_CHECK]


# --- the list stage 4 keeps when one reply held several items for an entry ---
# Invented values. The entry's scalars are the LAST record, as stage 4 leaves
# them; `_RECORDS` is the name stage 6 hands in (`STAGE4_RECORDS_KEY`).


def _committee(name, role='Member', start='2001', end='2003'):
    return {'committee_name': name, 'role': role, 'institution': 'Ashby University',
            'start_date': start, 'end_date': end}


def _stage4_entry(records, text=None, code='P', **entry_keys):
    return {'taxonomy_code': code, 'element_idx_start': 4,
            'text': text or 'Committees: Glade Board, Fern Council and Moss Panel, Ashby University',
            'extracted_fields': {**records[-1], _RECORDS: records}, **entry_keys}


_THREE_COMMITTEES = [_committee('Glade Board', 'Chair', '1999', '2001'),
                     _committee('Fern Council'), _committee('Moss Panel', start='2004', end='2006')]


def _fan4(*entries):
    return fan_out_multi_record_entries(list(entries), FIELD_SCHEMAS, records_key=_RECORDS)


class TestStage4Records:
    def test_one_child_per_record_and_the_last_is_the_parent(self):
        parent = _stage4_entry(copy.deepcopy(_THREE_COMMITTEES))
        children = _fan4(parent)
        assert [c['extracted_fields']['committee_name'] for c in children] == [
            'Glade Board', 'Fern Council', 'Moss Panel']
        last = children[-1]
        assert last['text'] == parent['text']
        assert last['extracted_fields'] == _THREE_COMMITTEES[-1]
        assert FANNED_OUT_FROM not in last
        # The default text's "Committees:" label is held by no field, so the
        # parent's whole line stays the last record's fallback (#1445).
        assert last[LAST_STAGE4_RECORD] == {'key': _RECORDS, 'index': 2, 'count': 3,
                                            fan_out.OWN_TEXT_KEY: None}
        assert [c[FANNED_OUT_FROM] for c in children[:-1]] == [
            {'key': _RECORDS, 'index': 0, 'count': 3}, {'key': _RECORDS, 'index': 1, 'count': 3}]
        assert all(LAST_STAGE4_RECORD not in c for c in children[:-1])

    def test_a_list_of_bare_s8_meetings_stays_whole(self):
        # #1570, NDMRSO QFQLNF 181's shape: records holding only a meeting, a
        # place and a status. S8 renders the meeting and place, but a record
        # with no title or year goes to the Appendix, so splitting printed the
        # first meeting there twice. Invented values.
        records = [{'conference_name': 'Quillfern Meeting', 'location': 'Larchmere, ZZ',
                    'status': 'pending'},
                   {'conference_name': 'Sable Symposium', 'location': 'Rookby, ZZ',
                    'status': 'pending'}]
        parent = _stage4_entry(records, code='S8', text=(
            'Quillfern Meeting (Larchmere, ZZ)-pending Sable Symposium (Rookby, ZZ)-pending'))
        assert _fan4(copy.deepcopy(parent)) == [parent]
        titled = [dict(record, title=f'Talk {i}') for i, record in enumerate(records)]
        assert len(_fan4(_stage4_entry(titled, code='S8'))) == 2

    def test_an_earlier_record_inherits_no_scalar_from_the_last(self):
        records = [{'committee_name': 'Glade Board'}, _committee('Moss Panel')]
        first, last = _fan4(_stage4_entry(records))
        assert first['extracted_fields'] == {'committee_name': 'Glade Board'}
        assert last['extracted_fields']['role'] == 'Member'

    def test_a_stage5_annotation_stays_on_the_last_record_only(self):
        # Every entry key stage 5 writes, named here rather than read off the
        # module, so dropping one from `_STAGE5_ENTRY_KEYS` fails this test.
        annotations = {'enriched_fields': ['city'], 'enrichment_data': {'pmid': '1'},
                       'enrichment_rejected': True, 'enrichment_source': 'llm',
                       'enrichment_status': 'ok',
                       'institution_enrichment': {'cleaned_name': 'Ashby University'},
                       'crossref_match': {'source': 'cv_doi', 'doi': '10.1000/x'},
                       'citation_note': 'This citation has no article title.'}
        parent = _stage4_entry(copy.deepcopy(_THREE_COMMITTEES), **annotations)
        children = _fan4(parent)
        for child in children[:-1]:
            assert not set(child) & set(annotations)
        assert {key: children[-1][key] for key in annotations} == annotations

    def test_a_scalar_stage5_added_to_the_parent_stays_on_the_last_record(self):
        parent = _stage4_entry(copy.deepcopy(_THREE_COMMITTEES))
        parent['extracted_fields']['city'] = 'Ashby'
        children = _fan4(parent)
        assert children[-1]['extracted_fields']['city'] == 'Ashby'
        assert 'city' not in children[0]['extracted_fields']

    def test_tab_segments_are_the_earlier_records_texts_when_the_counts_line_up(self):
        text = 'Chair, Glade Board 1999-2001\tMember, Fern Council 2001-2003\tMember, Moss Panel 2004-2006'
        children = _fan4(_stage4_entry(copy.deepcopy(_THREE_COMMITTEES), text=text))
        assert [c['text'] for c in children] == [
            'Chair, Glade Board 1999-2001', 'Member, Fern Council 2001-2003', text]

    def test_a_record_key_outside_the_schema_does_not_decline_the_list(self):
        records = copy.deepcopy(_THREE_COMMITTEES)
        records[0]['site_note'] = 'east wing'
        assert len(_fan4(_stage4_entry(records))) == 3

    def test_text_the_records_do_not_hold_does_not_decline_the_list(self):
        # The parent alone never held it either: the last child keeps the text.
        text = 'Zzunheld Qqunheld: Glade Board, Fern Council, Moss Panel'
        assert len(_fan4(_stage4_entry(copy.deepcopy(_THREE_COMMITTEES), text=text))) == 3

    @pytest.mark.parametrize('why', ['text_rendered_code', 'formatted', 'empty_record'])
    def test_a_declined_list_leaves_the_entry_as_it_was(self, why):
        records = copy.deepcopy(_THREE_COMMITTEES)
        entry = _stage4_entry(records, code='K2' if why == 'text_rendered_code' else 'P')
        if why == 'formatted':
            entry['extracted_fields']['formatted_text'] = 'Glade Board; Fern Council; Moss Panel'
        if why == 'empty_record':
            records[1] = {'description': 'nothing a P row writes'}
        assert _fan4(entry) == [entry]

    @pytest.mark.parametrize('value', [['Glade Board', 'Fern Council'], [_committee('Glade Board')]],
                             ids=['strings', 'one_record'])
    def test_a_value_that_is_not_two_or_more_records_is_left_alone(self, value):
        entry = _stage4_entry(copy.deepcopy(_THREE_COMMITTEES))
        entry['extracted_fields'][_RECORDS] = value
        assert _fan4(entry) == [entry]

    def test_a_declined_list_is_not_split_by_the_generic_rules(self):
        # The second record writes nothing in a P row, so this list declines;
        # the generic rules, handed it, would split it into an empty row.
        records = [{'committee_name': 'Glade Board', 'role': 'Chair'}, {'description': 'Zz'}]
        entry = _stage4_entry(records, text='Glade Board Chair')
        assert len(fan_out_multi_record_entries([entry], FIELD_SCHEMAS)) == 2
        assert _fan4(entry) == [entry]

    def test_a_record_list_in_the_scalars_is_left_to_the_generic_rules(self):
        # The last record carried its own off-schema list: that list splits as
        # it did before the stage-4 list existed, and the stage-4 list goes.
        last = {'role': 'Member', 'committees': [_committee('Fern Council'), _committee('Moss Panel')]}
        entry = {'taxonomy_code': 'P', 'text': 'Member, Fern Council\tMember, Moss Panel',
                 'extracted_fields': {**last, _RECORDS: [_committee('Glade Board'), last]}}
        children = _fan4(entry)
        assert [c['extracted_fields']['committee_name'] for c in children] == [
            'Fern Council', 'Moss Panel']
        assert all(_RECORDS not in c['extracted_fields'] for c in children)
        assert [c[FANNED_OUT_FROM]['key'] for c in children] == ['committees', 'committees']

    def test_the_parent_is_not_mutated(self):
        parent = _stage4_entry(copy.deepcopy(_THREE_COMMITTEES), institution_enrichment={'city': 'Ashby'})
        original = copy.deepcopy(parent)
        children = _fan4(parent)
        children[-1]['extracted_fields']['role'] = 'changed'
        children[-1]['institution_enrichment']['city'] = 'changed'
        children[0]['extracted_fields']['role'] = 'changed'
        assert parent == original


# --- EBYSBC E6 (#1187): record lists the coverage test used to decline ------
# Invented values (names, titles and dates alike), in the shapes the EBYSBC
# batch carried: a D1 `appointments` list behind a section label with month
# names and 'YY years, a Q2 `committees` list with YYYY-YY year ends, a Q2
# `additional_entries` list dated M/YY, a D1 list whose end is "current", a
# numbered Q1 line.


def _appointment(title, start, end, institution='Ashby University'):
    return {'title': title, 'institution': institution, 'start_date': start, 'end_date': end}


class TestCoverageComparesWhatTheColumnShows:
    @pytest.mark.parametrize('text,tokens', [
        ("March '88-May '89", {'1988': 1, '1989': 1}),
        ('Sept. 2011 - Aug 2013', {'2011': 1, '2013': 1}),
        ('3/02-11/03', {'2002': 1, '2003': 1}),
        ('11/1988', {'1988': 1}),
        ('1991-1993; 1993-96; 1998-01', {'1991': 1, '1993': 2, '1996': 1, '1998': 1, '2001': 1}),
        ('2016-08', {'2016': 1, '08': 1}),
        ('6/3/2012', {'6': 1, '3': 1, '2012': 1}),
        ('13/04', {'13': 1, '04': 1}),
        ('2016 - current, now, ongoing, Present', {'2016': 1, 'present': 4}),
        ('May Hollis Fund', {'may': 1, 'hollis': 1, 'fund': 1}),
    ], ids=['month-names-and-apostrophe-years', 'abbreviated-months', 'month-slash-yy',
            'month-slash-yyyy', 'yyyy-yy-ends', 'iso-month-is-not-a-range', 'full-date-left-whole',
            'not-a-month', 'open-ends', 'month-name-without-a-number-is-a-word'])
    def test_tokens(self, text, tokens):
        assert fan_out._tokens(text) == tokens

    def test_month_names_and_apostrophe_years_against_an_mm_yy_column(self):
        # D1 writes `mm/yy` ("03/88"): both sides read 1988.
        entry = {'taxonomy_code': 'D1',
                 'text': "March '88-May '89 Lecturer, Ashby University\t"
                         "June '89-present Reader, Ashby University",
                 'extracted_fields': {'appointments': [
                     _appointment('Lecturer', '1988-03', '1989-05'),
                     _appointment('Reader', '1989-06', 'present')]}}
        assert [c['extracted_fields']['title'] for c in _fan(entry)] == ['Lecturer', 'Reader']

    def test_an_open_end_written_current_is_held_by_present(self):
        entry = {'taxonomy_code': 'D1',
                 'text': '2016 - current Reader, Ashby University\tVisiting Fellow, Birch Hollow University',
                 'extracted_fields': {'appointments': [
                     _appointment('Reader', '2016', 'current'),
                     _appointment('Visiting Fellow', '2016', 'current', 'Birch Hollow University')]}}
        assert len(_fan(entry)) == 2

    def test_two_digit_year_ends(self):
        committees = [{'committee_name': name, 'organization': 'Glade Society',
                       'start_date': start, 'end_date': end}
                      for name, start, end in [('Alpha Council', '1991', '1993'),
                                               ('Beta Council', '1993', '1996'),
                                               ('Gamma Council', '2004', '2008')]]
        entry = {'taxonomy_code': 'Q2',
                 'text': 'Glade Society: Alpha Council 1991-1993, Beta Council 1993-96, '
                         'Gamma Council 2004-08',
                 'extracted_fields': {'start_date': '1991', 'end_date': '1993',
                                      'committees': committees}}
        assert [c['extracted_fields']['committee_name'] for c in _fan(entry)] == [
            'Alpha Council', 'Beta Council', 'Gamma Council']

    def test_month_slash_two_digit_year(self):
        entry = {'taxonomy_code': 'Q2',
                 'text': 'Glade Society Alpha Working Group, 3/02-11/03; Beta Working Group, 6/04-2/05',
                 'extracted_fields': {
                     'committee_name': 'Alpha Working Group', 'organization': 'Glade Society',
                     'start_date': '2002-03', 'end_date': '2003-11',
                     'additional_entries': [{'committee_name': 'Beta Working Group',
                                             'organization': 'Glade Society',
                                             'start_date': '2004-06', 'end_date': '2005-02'}]}}
        assert [c['extracted_fields']['committee_name'] for c in _fan(entry)] == [
            'Alpha Working Group', 'Beta Working Group']

    def test_a_list_number_is_not_content(self):
        entry = {'taxonomy_code': 'Q1',
                 'text': '2. Glade Society: Secretary (2003-2006), Chair-elect (2008-2010).',
                 'extracted_fields': {'organization': 'Glade Society', 'additional_roles': [
                     {'role': 'Secretary', 'start_date': '2003', 'end_date': '2006'},
                     {'role': 'Chair-elect', 'start_date': '2008', 'end_date': '2010'}]}}
        assert [c['extracted_fields']['role'] for c in _fan(entry)] == ['Secretary', 'Chair-elect']

    def test_a_number_inside_the_text_is_still_content(self):
        entry = {'taxonomy_code': 'Q1',
                 'text': 'Glade Society Secretary 2003-2006 2. Chair-elect 2008-2010',
                 'extracted_fields': {'organization': 'Glade Society', 'additional_roles': [
                     {'role': 'Secretary', 'start_date': '2003', 'end_date': '2006'},
                     {'role': 'Chair-elect', 'start_date': '2008', 'end_date': '2010'}]}}
        assert _fan(entry) == [entry]


_TWO_APPOINTMENTS = [_appointment('Lecturer', '1988-03', '1989-05'),
                     _appointment('Senior Lecturer', '1989-06', '1994-02')]


class TestSectionLabel:
    _TEXT = ("March '88-May '89 Lecturer, Ashby University\t"
             "June '89-Feb '94 Senior Lecturer, Ashby University")

    def _d1(self, label, items=_TWO_APPOINTMENTS):
        return {'taxonomy_code': 'D1', 'text': label + self._TEXT,
                'extracted_fields': {'appointments': copy.deepcopy(items)}}

    @pytest.mark.parametrize('label', ['Positions      ', 'Positions:\t', 'Teaching Posts: '])
    def test_a_label_before_the_first_date_is_dropped(self, label):
        assert len(_fan(self._d1(label))) == 2

    def test_a_label_not_followed_by_a_date_is_content(self):
        entry = self._d1('Positions      ')
        entry['text'] = entry['text'].replace(
            "March '88-May '89 Lecturer, Ashby University",
            "Lecturer, Ashby University, March '88-May '89", 1)
        assert _fan(entry) == [entry]

    def test_a_label_that_may_be_the_role_a_record_left_out_is_content(self):
        items = [dict(item) for item in _TWO_APPOINTMENTS]
        del items[1]['title']
        entry = self._d1('Tutor      ', items)
        entry['text'] = entry['text'].replace(" Senior Lecturer,", '', 1)
        assert _fan(entry) == [entry]
        del entry['extracted_fields']['appointments'][1]
        entry['text'] = entry['text'].split('\t')[0]
        assert len(_fan(entry)) == 1

    def test_a_label_further_in_is_content(self):
        entry = self._d1('')
        entry['text'] = entry['text'].replace('\t', '\tPositions      ')
        assert _fan(entry) == [entry]


class TestSharedContextLines:
    """EBYSBC E6 residual (#1187, KYOPUV 32 / WTKSYX 32): undated lines of
    context every record shares, around a dated list, are not a record's own
    words. Invented values."""
    _UNIT = 'Unit of Kestrel Studies'
    _PLACE = 'Ashby University, Fernvale'
    _DATED = ("March '88-May '89 Lecturer\t"
              "June '89-Feb '94 Senior Lecturer\t"
              "March '94-present Reader")

    def _d1(self, text, institution=_PLACE, titles=('Lecturer', 'Senior Lecturer', 'Reader')):
        dates = [('1988-03', '1989-05'), ('1989-06', '1994-02'), ('1994-03', 'present')]
        return {'taxonomy_code': 'D1', 'text': text,
                'extracted_fields': {'appointments': [
                    _appointment(title, start, end, institution)
                    for title, (start, end) in zip(titles, dates)]}}

    def test_a_trailing_unit_and_institution_line_fan_out(self):
        entry = self._d1(f'{self._DATED}\t{self._UNIT}\t{self._PLACE}')
        assert [c['extracted_fields']['title'] for c in _fan(entry)] == [
            'Lecturer', 'Senior Lecturer', 'Reader']

    def test_a_leading_institution_and_unit_line_fan_out(self):
        assert len(_fan(self._d1(f'{self._PLACE}\t{self._UNIT}\t{self._DATED}'))) == 3

    def test_without_an_institution_line_every_child_holds_the_unit_is_content(self):
        entry = self._d1(f'{self._DATED}\t{self._UNIT}\t{self._PLACE}',
                         institution='Ashby University')
        assert _fan(entry) == [entry]

    def test_a_context_line_longer_than_a_unit_name_is_content(self):
        unit = 'Unit of Kestrel Studies and the Joint Program in Heron Ecology'
        entry = self._d1(f'{self._DATED}\t{unit}\t{self._PLACE}')
        assert _fan(entry) == [entry]

    def test_an_undated_line_between_records_is_content(self):
        first, rest = self._DATED.split('\t', 1)
        entry = self._d1(f'{first}\t{self._UNIT}\t{rest}\t{self._PLACE}')
        assert _fan(entry) == [entry]

    def test_a_records_own_word_is_still_compared(self):
        entry = self._d1(f'{self._DATED}\t{self._UNIT}\t{self._PLACE}'.replace(
            'Senior Lecturer', 'Senior Visiting Lecturer'))
        assert _fan(entry) == [entry]

    def test_one_records_periods_are_not_a_list_of_records(self):
        # VNUAHA 142's shape (EBYSBC E22): two terms of one post.
        entry = self._d1(f'{self._DATED}\t{self._UNIT}\t{self._PLACE}',
                         titles=('Reader', 'Reader', 'Reader'))
        entry['text'] = entry['text'].replace('Senior Lecturer', 'Reader').replace(
            'Lecturer', 'Reader')
        assert _fan(entry) == [entry]

    @pytest.mark.parametrize('segments,children,expected', [
        (['1990 Alpha', '1991 Beta', 'Unit', 'Ashby'], 2, {2, 3}),
        (['1990 Alpha', '1991 Beta', '1992 Gamma', 'Unit', 'Ashby'], 2, set()),
        (['1990 Alpha', 'Unit', 'Ashby'], 1, set()),
        (['1990 Alpha', '1991 Beta', '', 'Ashby'], 2, {2, 3}),
        (['1990 Alpha', '1991 Beta', 'Unit'], 2, set()),
    ], ids=['one-dated-segment-per-child', 'a-dated-segment-no-child-holds',
            'a-lone-record', 'an-empty-line-is-no-anchor-but-goes', 'no-anchor-line'])
    def test_which_segments_are_context(self, segments, children, expected):
        child_fields = [{'title': f'T{i}', 'institution': 'Ashby'} for i in range(children)]
        rendered = fan_out._RENDERED_FIELDS['D1']
        assert fan_out._shared_context_segments(segments, child_fields, rendered, 'D1') == expected

    def test_an_open_end_dates_a_segment(self):
        child_fields = [{'title': f'T{i}', 'institution': 'Ashby'} for i in range(2)]
        rendered = fan_out._RENDERED_FIELDS['D1']
        assert fan_out._shared_context_segments(
            ['1990 Alpha', 'present Beta', 'Unit', 'Ashby'], child_fields, rendered, 'D1') == {2, 3}

    def test_an_institution_one_child_holds_is_no_anchor(self):
        child_fields = [{'title': 'T0', 'institution': 'Ashby'},
                        {'title': 'T1', 'institution': 'Birch'}]
        rendered = fan_out._RENDERED_FIELDS['D1']
        assert fan_out._shared_context_segments(
            ['1990 Alpha', '1991 Beta', 'Unit', 'Ashby'], child_fields, rendered, 'D1') == set()

    def test_an_empty_edge_line_alone_is_no_anchor(self):
        child_fields = [{'title': f'T{i}', 'institution': 'Ashby'} for i in range(2)]
        rendered = fan_out._RENDERED_FIELDS['D1']
        assert fan_out._shared_context_segments(
            ['1990 Alpha', '1991 Beta', '', 'Unit'], child_fields, rendered, 'D1') == set()


class TestOneItemList:
    def test_beside_the_parents_own_role_it_is_a_second_record(self):
        # MRJDWE 101's shape; the one segment holds both records.
        entry = {'taxonomy_code': 'Q1',
                 'text': '4. Glade Society board member - 2004-2007; Secretary 2007-2009.',
                 'extracted_fields': {'role': 'Board member', 'organization': 'Glade Society',
                                      'start_date': '2004', 'end_date': '2007',
                                      'additional_roles': [{'role': 'Secretary', 'start_date': '2007',
                                                            'end_date': '2009'}]}}
        first, second = _fan(entry)
        assert first['extracted_fields'] == {'role': 'Board member', 'organization': 'Glade Society',
                                             'start_date': '2004', 'end_date': '2007'}
        assert second['extracted_fields'] == {'role': 'Secretary', 'organization': 'Glade Society',
                                              'start_date': '2007', 'end_date': '2009'}
        assert [c[FANNED_OUT_FROM]['count'] for c in (first, second)] == [2, 2]

    def test_under_a_parent_holding_only_its_dates_it_is_the_record(self):
        # The pilot's BFSUMA shape (#1187): the item is the whole record.
        entry = {'taxonomy_code': 'D1',
                 'text': 'Ashby University, School of Nursing\tTutor in Nursing, 1983-1984',
                 'extracted_fields': {'appointments': [_appointment(
                     'Tutor in Nursing', '1983', '1984', 'Ashby University, School of Nursing')],
                     'start_date': '1983', 'end_date': '1984'}}
        (child,) = _fan(entry)
        assert child['extracted_fields']['title'] == 'Tutor in Nursing'
        assert child['text'] == entry['text']
        assert child[FANNED_OUT_FROM] == {'key': 'appointments', 'index': 0, 'count': 1}

    def test_a_list_of_dates_is_the_parents_second_period(self):
        entry = {'taxonomy_code': 'O',
                 'text': 'Acting Dean, Ashby University\tMarch 2011-May 2012\tOctober 2013-April 2014',
                 'extracted_fields': {'leadership_role': 'Acting Dean', 'institution': 'Ashby University',
                                      'start_date': '2011-03', 'end_date': '2012-05',
                                      'additional_dates': [{'start_date': '2013-10',
                                                            'end_date': '2014-04'}]}}
        children = _fan(entry)
        assert [(c['extracted_fields']['leadership_role'], c['extracted_fields']['start_date'])
                for c in children] == [('Acting Dean', '2011-03'), ('Acting Dean', '2013-10')]

    def test_an_award_list_of_one_renders_from_its_fields(self):
        (child,) = _fan(_honors('Alpha Prize, Hollis College', [_award('Alpha Prize')]))
        assert child['extracted_fields'] == _award('Alpha Prize')

    def test_an_item_that_repeats_the_parent_is_not_emitted_twice(self):
        entry = {'taxonomy_code': 'Q1', 'text': 'Glade Society Secretary 2007-2009',
                 'extracted_fields': {'role': 'Secretary', 'organization': 'Glade Society',
                                      'additional_roles': [{'role': 'Secretary', 'start_date': '2007',
                                                            'end_date': '2009'}]}}
        assert len(_fan(entry)) == 1

    def test_the_stage4_records_list_still_needs_two(self):
        entry = _stage4_entry([_committee('Glade Board')])
        assert _fan4(entry) == [entry]

    def test_a_one_item_list_in_the_last_record_leaves_the_stage4_list_to_split(self):
        # Only a list of two or more in the scalars hands the entry to the
        # generic rules, which would drop the stage-4 records.
        last = dict(_committee('Moss Panel'), additional_roles=[{'role': 'Chair'}])
        children = _fan4(_stage4_entry([_committee('Glade Board'), last]))
        assert [c['extracted_fields']['committee_name'] for c in children] == ['Glade Board', 'Moss Panel']

    def test_a_one_item_list_with_a_stray_key_warns(self):
        entry = {'taxonomy_code': 'D3', 'text': 'Consultant, Glade Agency 1990-1992',
                 'extracted_fields': {'consultantships': [{
                     'title': 'Consultant', 'organization': 'Glade Agency', 'start_date': '1990',
                     'end_date': '1992', 'description': 'advised on staffing'}]}}
        warnings = []
        assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings) == [entry]
        assert warnings[0]['evidence'] == ['D3.consultantships: 1 entry', 'description']


_LICENCES = [{'state_country': 'Arden', 'issue_date': '2014'},
             {'state_country': 'Brinmoor', 'issue_date': '2014'}]


class TestNestedRecordList:
    def _f1(self, text='Corvale 2012\tArden 2014\tBrinmoor 2014', **extra):
        return {'taxonomy_code': 'F1', 'text': text,
                'extracted_fields': {'state_country': 'Corvale', 'issue_date': '2012',
                                     'additional_info': {'remote_licences': copy.deepcopy(_LICENCES),
                                                         **extra}}}

    def test_a_list_one_level_down_fans_out(self):
        children = _fan(self._f1())
        assert [c['extracted_fields']['state_country'] for c in children] == [
            'Corvale', 'Arden', 'Brinmoor']
        assert [c[FANNED_OUT_FROM]['key'] for c in children] == ['additional_info.remote_licences'] * 3
        assert all('additional_info' not in c['extracted_fields'] for c in children)

    def test_the_rest_of_the_object_stays_with_every_child(self):
        children = _fan(self._f1(license_type='Medical'))
        assert all(c['extracted_fields']['additional_info'] == {'license_type': 'Medical'}
                   for c in children)

    def test_a_declined_nested_list_warns_with_its_path(self):
        warnings = []
        entry = self._f1(text='Corvale 2012 Remote Permits Arden 2014 Brinmoor 2014')
        assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings) == [entry]
        assert warnings[0]['evidence'] == ['F1.additional_info.remote_licences: 1 entry']

    def test_a_nested_list_with_a_stray_key_warns_naming_it(self):
        entry = self._f1()
        entry['extracted_fields']['additional_info']['remote_licences'][0]['portal'] = 'web'
        warnings = []
        assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings) == [entry]
        assert warnings[0]['evidence'] == ['F1.additional_info.remote_licences: 1 entry', 'portal']

    def test_two_levels_down_is_not_read(self):
        entry = self._f1()
        entry['extracted_fields']['additional_info'] = {
            'more': {'remote_licences': copy.deepcopy(_LICENCES)}}
        assert _fan(entry) == [entry]


class TestDeclinedListWarnsForEveryFieldRenderedCode:
    @pytest.mark.parametrize('code,list_item', [
        ('D1', {'title': 'Instructor'}), ('D2', {'title': 'Attending'}),
        ('D3', {'title': 'Consultant'}), ('F1', {'state_country': 'Arden'}),
        ('I', {'organization': 'Glade Society'}), ('Q1', {'role': 'Chair'}),
        ('Q2', {'committee_name': 'Alpha Council'}),
        # Silent until #1187: each writes the parent's fields and none of a list.
        ('P', {'committee_name': 'Glade Board'}), ('Q3', {'panel_name': 'Glade Panel'}),
        ('R', {'event_name': 'Glade Meeting'})])
    def test_a_declined_list_warns(self, code, list_item):
        entry = {'taxonomy_code': code, 'text': 'Zzunheld prose no field carries',
                 'extracted_fields': {'records': [dict(list_item), dict(list_item)]}}
        warnings = []
        assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings) == [entry]
        assert [(w['check'], w['code']) for w in warnings] == [(fan_out.REJECTED_LIST_CHECK, code)]

    @pytest.mark.parametrize('code,list_key', [('H', 'dates'), ('H', 'additional_dates'),
                                               ('O', 'additional_dates')])
    def test_a_declined_date_span_list_the_section_renders_does_not_warn(self, code, list_key):
        # H and O write these lists into their date cell, so a declined one
        # loses nothing and a warning would cost quality-score points.
        entry = {'taxonomy_code': code, 'text': 'Zzunheld prose no field carries',
                 'extracted_fields': {list_key: [{'start_date': '2001', 'end_date': '2002'},
                                                 {'start_date': '2005', 'end_date': '2006'}]}}
        warnings = []
        assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings) == [entry]
        assert warnings == []


def _declined(*entries, records_key=None):
    """`(warnings, declined)` of one fan-out pass over `entries`."""
    warnings, declined = [], []
    fan_out_multi_record_entries(list(entries), FIELD_SCHEMAS, warnings,
                                 records_key=records_key, declined=declined)
    return warnings, declined


_FANOUT_LIST_SHAPE = re.compile(r"holds (?:several )?records that were not split")


class TestDeclinedListCheckedAfterRender:
    """#1187: a declined list of these codes is in the document when the
    parent holds no name of its own (its row then prints the entry's text),
    so it warns only once the render shows a record missing."""

    @pytest.mark.parametrize('code,list_item', [
        ('B2', {'program_name': 'Glade Program'}), ('C', {'training_type': 'Fellowship'}),
        ('F2', {'certifying_board': 'Glade Board'}), ('L3', {'leadership_role': 'Director'}),
        ('M2A', {'title': 'Glade Study'}), ('M2B', {'title': 'Glade Study'}),
        ('M2C', {'title': 'Glade Study'}), ('M2D', {'title': 'Glade Device'}),
        ('N2', {'grant_title': 'Glade Training Grant'}), ('N3', {'mentee_name': 'Quillon Ashgrove'}),
        ('N3A', {'mentee_name': 'Quillon Ashgrove'}), ('N3B', {'mentee_name': 'Quillon Ashgrove'}),
        ('Q4', {'journal_name': 'Glade Review'}), ('Q4A', {'journal_name': 'Glade Review'}),
        ('Q4B', {'journal_name': 'Glade Review'}), ('Q4C', {'journal_name': 'Glade Review'}),
        ('Q4D', {'journal_name': 'Glade Review'}), ('H', {'award_name': 'Glade Prize'})])
    def test_a_declined_list_is_handed_back_not_warned(self, code, list_item):
        entry = {'taxonomy_code': code, 'text': 'Zzunheld prose no field carries',
                 'extracted_fields': {'records': [dict(list_item), dict(list_item)]}}
        warnings, declined = _declined(entry)
        assert warnings == []
        assert [(d.code, d.key, d.records, d.stage4) for d in declined] == [
            (code, 'records', (list_item, list_item), False)]

    def test_a_directly_warned_code_is_not_handed_back(self):
        entry = {'taxonomy_code': 'P', 'text': 'Zzunheld prose no field carries',
                 'extracted_fields': {'records': [{'committee_name': 'Glade Board'}] * 2}}
        warnings, declined = _declined(entry)
        assert [w['code'] for w in warnings] == ['P'] and declined == []

    def test_stray_item_keys_are_named_and_values_never(self):
        entry = {'taxonomy_code': 'N3B', 'text': 'Quillon Ashgrove; Tamsin Oakhurst',
                 'extracted_fields': {'mentee_name': 'Pellam Fairweather', 'mentees': [
                     {'mentee_name': 'Quillon Ashgrove', 'site_position': 'Glade Lab'},
                     {'mentee_name': 'Tamsin Oakhurst', 'site_position': 'Glade Lab'}]}}
        (found,) = _declined(entry)[1]
        (warning,) = fan_out.unrendered_list_warnings([found], ['Pellam Fairweather'])
        assert (warning['check'], warning['code'], warning['severity']) == (
            fan_out.REJECTED_LIST_CHECK, 'N3B', 'WARN')
        assert warning['evidence'] == ['N3B.mentees: 1 entry', 'site_position']
        assert _FANOUT_LIST_SHAPE.search(warning['message'])
        assert 'Quillon' not in str(warning) and 'Glade' not in str(warning)

    def test_records_the_document_holds_do_not_warn(self):
        found = fan_out.DeclinedList('Q4D', 'journals', ('start_date',), (
            {'journal_name': 'Quillmere Journal', 'start_date': '2009'},
            {'journal_name': 'Ashgrove Review', 'start_date': '2010'}))
        document = ['Reviewer: Quillmere Journal (2009-2012); Ashgrove Review (2010-2014)']
        assert fan_out.unrendered_list_warnings([found], document) == []

    def test_a_record_whose_words_are_spread_over_lines_is_not_in_the_document(self):
        # A journal's name words also in citations (VYNARH 138): no one line
        # prints the journal, so it is missing.
        found = fan_out.DeclinedList('Q4D', 'journals', (), ({'journal_name': 'Quillmere Ashgrove Glade Review'},))
        lines = ['Doe J. Quillmere study. 2001', 'Roe R. Ashgrove trial. 2002',
                 'Poe P. Glade cohort. 2003', 'Fern Review. 2004']
        assert len(fan_out.unrendered_list_warnings([found], lines)) == 1

    def test_one_record_missing_warns_once_per_code_and_key(self):
        found = fan_out.DeclinedList('Q4D', 'journals', (), (
            {'journal_name': 'Quillmere Journal'}, {'journal_name': 'Ashgrove Review'}))
        warnings = fan_out.unrendered_list_warnings([found, found], ['Quillmere Journal'])
        assert [w['evidence'] for w in warnings] == [['Q4D.journals: 2 entries']]

    @pytest.mark.parametrize('document,missing', [
        ('Quillmere', False),            # 1 of 2 words: half is enough
        ('Glade', True),                 # 0 of 2
    ])
    def test_a_record_counts_as_rendered_from_half_its_words(self, document, missing):
        found = fan_out.DeclinedList('Q4D', 'journals', (), ({'journal_name': 'Quillmere Journal'},))
        assert bool(fan_out.unrendered_list_warnings([found], [document])) is missing

    def test_a_record_with_no_word_has_nothing_to_lose(self):
        found = fan_out.DeclinedList('Q4D', 'journals', (), ({'journal_name': '—'},))
        assert fan_out.unrendered_list_warnings([found], []) == []

    def test_dates_are_compared_by_year(self):
        found = fan_out.DeclinedList('N3B', 'mentees', (), (
            {'mentee_name': 'Quillon', 'start_date': '2013-09', 'end_date': 'current'},))
        assert fan_out.unrendered_list_warnings([found], ['Quillon 2013-Present']) == []

    def test_a_k1_or_s_list_beside_a_named_parent_is_handed_back(self):
        # web204 S3 1802 / 1814: a book's chapters, kept whole beside its own
        # title, which is all the row prints.
        chapters = [{'authors': 'Gorman, B. S.', 'chapter': 'Glade Tides'},
                    {'authors': 'Pellam, R.', 'chapter': 'Fern Lattice'}]
        entry = {'taxonomy_code': 'S3', 'text': 'Glade Handbook. Chapters: Glade Tides; Fern Lattice',
                 'extracted_fields': {'title': 'Glade Handbook', 'chapters_authored': chapters}}
        warnings, declined = _declined(entry)
        assert warnings == []
        assert [(d.code, d.key, d.stray, d.stage4) for d in declined] == [
            ('S3', 'chapters_authored', ('chapter',), False)]

    def test_a_k1_list_is_handed_back_beside_its_declined_stage4_list(self):
        # Stage 4's records list is handed back once, as itself.
        records = [{'course_title': 'Glade Anatomy', 'role': 'Lecturer'}, {'institution': None}]
        courses = [{'course': 'Fern Physiology'}, {'course': 'Moss Histology'}]
        entry = _stage4_entry(records, code='K1', text='Glade Anatomy; Fern Physiology; Moss Histology',
                              extracted_fields={**records[-1], _RECORDS: records, 'courses': courses})
        _, declined = _declined(entry, records_key=_RECORDS)
        assert sorted((d.key, d.stage4) for d in declined) == [('courses', False), (_RECORDS, True)]

    def test_words_every_record_shares_do_not_vouch_for_one(self):
        # web204 S3 1814: every chapter names the owner, so a publication line
        # naming the owner and one title word does not print the chapter.
        found = fan_out.DeclinedList('S3', 'chapters_authored', ('chapter',), (
            {'authors': 'Allison Quillmere', 'chapter': 'Glade Tides'},
            {'authors': 'Allison Quillmere', 'chapter': 'Fern Lattice Theory'}))
        lines = ['Allison Quillmere. Glade Tides.', 'Allison Quillmere. Fern cohort. 2001']
        (warning,) = fan_out.unrendered_list_warnings([found], lines)
        assert warning['evidence'][0] == 'S3.chapters_authored: 1 entry'

    def test_initials_do_not_vouch_for_a_record(self):
        found = fan_out.DeclinedList('S3', 'chapters_authored', ('chapter',), (
            {'authors': 'Gorman, B. S.', 'chapter': 'Glade Tides'},
            {'authors': 'Pellam, R.', 'chapter': 'Fern Lattice'}))
        lines = ['Pellam R. Fern Lattice.', 'Gorman B S. Moss cohort.']
        assert len(fan_out.unrendered_list_warnings([found], lines)) == 1

    @pytest.mark.parametrize('rows,missing', [
        (['Title: Glade', 'PI: Quillon', 'Amount: Ashgrove'], False),
        (['Title: Glade', 'Fern', 'Moss', 'PI: Quillon', 'Ashby', 'Pine', 'Amount: Ashgrove'], True),
    ])
    def test_a_record_written_one_field_per_row_spans_its_field_count(self, rows, missing):
        # web26 M2A: a grant block writes one field per row; a whole table's
        # text would vouch for words strewn over unrelated rows.
        found = fan_out.DeclinedList('M2A', 'sub_awards', (), (
            {'sub_title': 'Glade', 'pi_name': 'Quillon', 'funding': 'Ashgrove', 'notes': None},))
        assert bool(fan_out.unrendered_list_warnings([found], [], [rows])) is missing

    def test_no_declined_argument_hands_nothing_back(self):
        entry = {'taxonomy_code': 'N3B', 'text': 'Zzunheld prose no field carries',
                 'extracted_fields': {'records': [{'mentee_name': 'Quillon'}] * 2}}
        warnings = []
        assert fan_out_multi_record_entries([entry], FIELD_SCHEMAS, warnings) == [entry]
        assert warnings == []


def _h_spans():
    return [{'start_date': '2001', 'end_date': '2002'}, {'start_date': '2005', 'end_date': '2006'}]


class TestHonorsDateListIsNotADeclinedList:
    """#1187: H's date cell writes a dates-only list, but only the first of
    `H_DATE_LIST_KEYS` that holds one, and only beside no date of its own."""

    def _h(self, **fields):
        return {'taxonomy_code': 'H', 'text': 'Zzunheld prose no field carries',
                'extracted_fields': {'award_name': 'Glade Prize', **fields}}

    @pytest.mark.parametrize('key', ['dates', 'additional_dates'])
    def test_the_list_the_date_cell_writes_is_not_handed_back(self, key):
        assert _declined(self._h(**{key: _h_spans()})) == ([], [])

    @pytest.mark.parametrize('own', ['date', 'year', 'start_date', 'end_date'])
    def test_a_date_list_beside_the_records_own_date_is(self, own):
        _, declined = _declined(self._h(dates=_h_spans(), **{own: '1999'}))
        assert [d.key for d in declined] == ['dates']

    def test_only_the_first_date_list_is_written(self):
        _, declined = _declined(self._h(dates=_h_spans(),
                                        additional_dates=_h_spans()))
        assert [d.key for d in declined] == ['additional_dates']

    def test_a_date_list_under_another_key_is(self):
        _, declined = _declined(self._h(periods=_h_spans()))
        assert [d.key for d in declined] == ['periods']

    def test_a_list_carrying_more_than_dates_is(self):
        spans = [dict(span, award_name='Glade Medal') for span in _h_spans()]
        _, declined = _declined(self._h(dates=spans))
        assert [d.key for d in declined] == ['dates']

    def test_the_section_reads_the_same_keys(self):
        from unified_pipeline.stage6.sections import honors
        assert honors._DATE_LIST_KEYS is fan_out.H_DATE_LIST_KEYS


class TestDeclinedStage4RecordsList:
    """#1187: a K1 or S* entry whose stage-4 records `_stage4_children`
    declined renders the last record only, unless stage 5 rendered all."""

    def _k1(self, **fields):
        # The last record renders nothing K1 writes, so the split is declined.
        records = [{'course_title': 'Glade Anatomy', 'role': 'Lecturer'}, {'institution': None}]
        return _stage4_entry(records, code='K1', text='Glade Anatomy, Lecturer; Fern Physiology',
                             extracted_fields={**records[-1], _RECORDS: records, **fields})

    def test_a_declined_k1_list_is_handed_back(self):
        warnings, (found,) = _declined(self._k1(), records_key=_RECORDS)
        assert warnings == []
        assert (found.code, found.key, found.stage4, len(found.records)) == ('K1', _RECORDS, True, 2)

    def test_its_warning_has_the_list_shape(self):
        _, declined = _declined(self._k1(), records_key=_RECORDS)
        (warning,) = fan_out.unrendered_list_warnings(declined, ['Fern Physiology'])
        assert warning['evidence'] == [f'K1.{_RECORDS}: 1 entry']
        assert _FANOUT_LIST_SHAPE.search(warning['message'])

    def test_a_k1_list_5c_rendered_whole_is_not(self):
        assert _declined(self._k1(formatted_text='Glade Anatomy\nFern Physiology'),
                         records_key=_RECORDS) == ([], [])

    def _s8(self, citation):
        records = [{'title': 'Glade Poster', 'conference_name': 'Fern Meeting'},
                   {'title': 'Moss Talk', 'conference_name': 'Ashby Meeting'}]
        return _stage4_entry(records, code='S8', text='Glade Poster, Fern Meeting; Moss Talk, Ashby Meeting',
                             extracted_fields={**records[-1], _RECORDS: records,
                                               'formatted_citation': citation})

    def test_a_citation_writing_every_title_renders_the_list_whole(self):
        assert _declined(self._s8('Glade Poster. Fern Meeting. Moss Talk. Ashby Meeting.'),
                         records_key=_RECORDS) == ([], [])

    def test_a_citation_writing_no_title_does_not(self):
        _, declined = _declined(self._s8('Fern Meeting; Ashby Meeting.'), records_key=_RECORDS)
        assert [(d.code, d.stage4) for d in declined] == [('S8', True)]

    def test_a_split_list_is_not_handed_back(self):
        records = [{'course_title': 'Glade Anatomy'}, {'course_title': 'Fern Physiology'}]
        entry = _stage4_entry(records, code='K1', text='Glade Anatomy\tFern Physiology')
        assert _declined(entry, records_key=_RECORDS) == ([], [])

    def test_a_code_outside_the_set_is_not_handed_back(self):
        records = [{'committee_name': 'Glade Board'}, {'role': None}]
        entry = _stage4_entry(records, code='O', text='Glade Board; Fern Council')
        assert _declined(entry, records_key=_RECORDS)[1] == []


class TestGeneratorChecksDeclinedListsAfterRender:
    def _render(self, tmp_path, entry):
        from unified_pipeline.stage_6_word_template import WCMTemplateGenerator
        source, target = tmp_path / 'in.json', tmp_path / 'out.docx'
        source.write_text(json.dumps({'document_uid': 'TESTAA', 'entries': [entry]}))
        generator = WCMTemplateGenerator(verbose=False)
        generator.generate(str(source), str(target), research_summary_path=None)
        sidecar = json.loads(next(tmp_path.glob('*render_warnings.json')).read_text())
        return [w for w in sidecar['warnings'] if w.get('check') == fan_out.REJECTED_LIST_CHECK]

    def test_a_list_lost_beside_a_named_parent_warns(self, tmp_path):
        entry = {'taxonomy_code': 'N3B', 'element_idx_start': 5,
                 'text': 'Pellam Fairweather, PhD 2010-2012; Quillon Ashgrove 2013; Tamsin Oakhurst 2014',
                 'extracted_fields': {
                     'mentee_name': 'Pellam Fairweather', 'mentee_level': 'PhD',
                     'start_date': '2010', 'end_date': '2012',
                     'mentees': [{'mentee_name': 'Quillon Ashgrove', 'start_date': '2013', 'site_position': 'Lab'},
                                 {'mentee_name': 'Tamsin Oakhurst', 'start_date': '2014', 'site_position': 'Lab'}]}}
        assert [(w['code'], w['evidence'][0]) for w in self._render(tmp_path, entry)] == [
            ('N3B', 'N3B.mentees: 1 entry')]

    def test_a_list_the_text_renders_does_not(self, tmp_path):
        entry = {'taxonomy_code': 'Q4D', 'element_idx_start': 5,
                 'text': 'Reviewer: Quillmere Journal (2009-2012); Ashgrove Review (2010-2014)',
                 'extracted_fields': {'journals': [
                     {'journal_name': 'Quillmere Journal', 'start_date': '2009', 'end_date': '2012'},
                     {'journal_name': 'Ashgrove Review', 'start_date': '2010', 'end_date': '2014'}]}}
        assert self._render(tmp_path, entry) == []


class TestParentKeepsItsOwnPart:
    """EBYSBC E34 (CTWLTR 55): the parent's `organization` joined every
    consultancy, two of which the items render again."""

    def _q2(self, organization):
        return {'taxonomy_code': 'Q2',
                'text': 'Adviser: Alpha Fund, 2004; Beta Trust 2006; Gamma Forum 2006-present',
                'extracted_fields': {'role': 'Adviser', 'organization': organization,
                                     'start_date': '2004', 'end_date': '2004',
                                     'additional_roles': [
                                         {'role': 'Adviser', 'organization': 'Beta Trust',
                                          'start_date': '2006', 'end_date': '2006'},
                                         {'role': 'Adviser', 'organization': 'Gamma Forum',
                                          'start_date': '2006', 'end_date': 'present'}]}}

    def test_parts_an_item_names_leave_the_parent(self):
        children = _fan(self._q2('Alpha Fund; Beta Trust; Gamma Forum'))
        assert [c['extracted_fields']['organization'] for c in children] == [
            'Alpha Fund', 'Beta Trust', 'Gamma Forum']

    def test_a_value_every_part_of_which_an_item_holds_is_kept(self):
        entry = self._q2('Beta Trust; Gamma Forum')
        entry['text'] += ' Alpha Fund'
        entry['extracted_fields']['committee_name'] = 'Alpha Fund'
        assert _fan(entry)[0]['extracted_fields']['organization'] == 'Beta Trust; Gamma Forum'

    def test_an_item_inherits_the_whole_value_it_lacks(self):
        entry = self._q2('Alpha Fund; Beta Trust; Gamma Forum')
        del entry['extracted_fields']['additional_roles'][1]['organization']
        children = _fan(entry)
        assert children[0]['extracted_fields']['organization'] == 'Alpha Fund; Gamma Forum'
        assert children[2]['extracted_fields']['organization'] == 'Alpha Fund; Beta Trust; Gamma Forum'


def _scope_generator(monkeypatch, abroad):
    """A generator whose scope classifier answers International when the
    activity it is asked about names `abroad`, and Regional otherwise, so a
    National heading that wins is visible."""
    import unified_pipeline.stage_6_word_template as s6

    def classify(*args, **kwargs):
        activity = next(line for line in kwargs['messages'][1]['content'].splitlines()
                        if line.startswith('**Activity'))
        return {'content': json.dumps({'scope': 'International' if abroad in activity else 'Regional'})}

    monkeypatch.setattr(s6, 'call_llm', classify)
    generator = s6.WCMTemplateGenerator(verbose=False)
    generator.cv_owner_location = {'primary_location': {'institution': 'Ashby University',
                                                        'city': 'Ashby', 'state': 'AB'}}
    return generator


def _meeting(place, year):
    return {'committee_name': f'Poster Talk, Autumn Meeting, {place}', 'role': 'Moderator',
            'organization': '', 'start_date': year, 'end_date': year}


class TestInheritedScope:
    """EBYSBC E34 (BZZNRL 137): a fanned child keeps its heading's scope
    unless the classifier, asked about that record alone, puts it abroad.
    #1579 (YUYVIG): any other entry under a scope heading takes it, and the
    classifier is not asked."""

    @pytest.mark.parametrize('hierarchy,scope', [
        (['SERVICE', 'National'], 'National'), (['NATIONAL SERVICE ROLES'], 'National'),
        (['International', 'Talks'], 'International'), (['Regional'], 'Regional'),
        (['International Activities', 'National'], 'National'),
        (['Talks: National, International'], None), (['Editorial Work'], None), ([], None),
        # #1579: the heading forms the YUYVIG CVs file their talks and committees under.
        (['R. INVITED TALKS', 'Regional*'], 'Regional'), (['R. INVITED TALKS', 'National*'], 'National'),
        (['R. INVITED TALKS', 'International*'], 'International'), (['INVITED LECTURER', 'STATE'], 'Regional'),
        (['Scholarship', 'Presentations (State & Local)'], 'Regional'), (['LOCAL'], 'Regional'),
        (['Statewide Review Panels'], 'Regional'), (['Provincial Committees'], 'Regional'),
        (['Federal Advisory Boards'], 'National'), (['Talks', 'International/National'], None),
        (['Committees for State and Federal Agencies'], None)])
    def test_the_nearest_heading_naming_one_scope(self, hierarchy, scope):
        assert fan_out.inherited_scope({'hierarchy': hierarchy}) == scope

    def test_a_fanned_out_record_inherits_too(self):
        assert fan_out.inherited_scope({FANNED_OUT_FROM: {}, 'hierarchy': ['National']}) == 'National'
        assert fan_out.inherited_scope({LAST_STAGE4_RECORD: {}, 'hierarchy': ['National']}) == 'National'

    @pytest.mark.parametrize('marker,heading,own,scope', [
        (FANNED_OUT_FROM, 'National', 'Regional', 'National'),
        (LAST_STAGE4_RECORD, 'National', 'Regional', 'National'),
        (FANNED_OUT_FROM, 'National', 'International', 'International'),
        (FANNED_OUT_FROM, 'Regional', 'National', 'Regional'),
        (FANNED_OUT_FROM, 'International', 'National', 'International'),
        (FANNED_OUT_FROM, 'Editorial Work', 'Regional', 'Regional'),
        (None, 'Editorial Work', 'International', 'International')])
    def test_a_record_keeps_its_heading_unless_the_classifier_puts_it_abroad(
            self, marker, heading, own, scope):
        entry = {'hierarchy': ['SERVICE', heading], **({marker: {}} if marker else {})}
        assert fan_out.record_scope(entry, lambda _: own) == scope

    @pytest.mark.parametrize('heading,scope', [
        ('Regional*', 'Regional'), ('National', 'National'), ('International*', 'International'),
        ('STATE', 'Regional'), ('Presentations (State & Local)', 'Regional')])
    def test_an_entry_under_a_scope_heading_takes_it_without_the_classifier(self, heading, scope):
        # #1579 IZABPD h0041, SQMWHM h1322, SXPHOG h1271: the classifier's
        # distance answer re-bucketed talks the CV had already sorted.
        asked = []
        entry = {'hierarchy': ['R. INVITED TALKS', heading]}
        assert fan_out.record_scope(entry, lambda e: asked.append(e) or 'International') == scope
        assert asked == []

    def test_an_ambiguous_heading_falls_back_to_the_classifier(self):
        asked = []
        entry = {'hierarchy': ['R. INVITED TALKS', 'International/National']}
        assert fan_out.record_scope(entry, lambda e: asked.append(e) or 'Regional') == 'Regional'
        assert asked == [entry]

    def test_the_last_stage4_record_is_its_own_fields_not_the_whole_line(self):
        entry = _stage4_entry(copy.deepcopy(_THREE_COMMITTEES), text='Glade Board\tFern Council and Moss Panel')
        children = _fan4(entry)
        assert fan_out.record_text(children[-1]) == 'Moss Panel | Member | Ashby University | 2004 | 2006'
        assert fan_out.record_text(children[0]) == children[0]['text']
        assert fan_out.record_text({'text': None}) == ''

    def test_every_stage4_record_keeps_the_heading_unless_it_is_abroad(self, monkeypatch):
        # BZZNRL 137's shape: the whole line names the meeting abroad first,
        # and the last record, the parent itself, is a meeting at home.
        meetings = [_meeting('Varnor, Kesh', '2010'), _meeting('Corvale, Westmark', '2012'),
                    _meeting('Lindell, Westmark', '2013')]
        text = 'Moderator:\t' + '\t'.join(f"{m['committee_name']}, {m['start_date']}" for m in meetings)
        children = _fan4(_stage4_entry(meetings, text=text, code='Q2', hierarchy=['SERVICE', 'National']))
        generator = _scope_generator(monkeypatch, abroad='Varnor')
        assert [generator._classify_geographic_scope(c) for c in children] == [
            'International', 'National', 'National']
        # Classified on the whole line, as the parent was, the last record goes
        # abroad; under its National heading it now takes that heading (#1579).
        whole = {key: value for key, value in children[-1].items() if key != LAST_STAGE4_RECORD}
        assert generator._classify_geographic_scope({**whole, 'hierarchy': ['SERVICE']}) == 'International'
        assert generator._classify_geographic_scope(whole) == 'National'

    def test_a_listed_child_keeps_the_heading_unless_it_is_abroad(self, monkeypatch):
        entry = {'taxonomy_code': 'Q2', 'hierarchy': ['SERVICE', 'National'],
                 'text': 'Panelist: Spring Forum, Ashby\tSpring Forum, Varnor',
                 'extracted_fields': {'committee_name': 'Spring Forum, Ashby', 'role': 'Panelist',
                                      'additional_entries': [{'committee_name': 'Spring Forum, Varnor',
                                                              'role': 'Panelist'}]}}
        generator = _scope_generator(monkeypatch, abroad='Varnor')
        assert [generator._classify_geographic_scope(c) for c in _fan(entry)] == [
            'National', 'International']
        # An entry that was not fanned out takes its heading, even where the
        # classifier would put it abroad (#1579), and its own answer without one.
        assert generator._classify_geographic_scope({**entry, 'text': 'Spring Forum, Varnor'}) == 'National'
        assert generator._classify_geographic_scope(
            {**entry, 'hierarchy': ['SERVICE'], 'text': 'Spring Forum, Ashby'}) == 'Regional'


@pytest.mark.parametrize('context', ['', '\tUnit of Kestrel Studies\tAshby University'],
                         ids=['label-led', 'trailing-unit-and-institution'])
def test_every_appointment_of_a_label_led_list_renders(tmp_path, context):
    import json

    from docx import Document

    from unified_pipeline.stage_6_word_template import WCMTemplateGenerator
    entry = {'taxonomy_code': 'D1', 'element_idx_start': 3,
             'text': "Positions      March '88-May '89 Zqlecturer, Ashby University\t"
                     "June '89-present Zqreader, Ashby University" + context,
             'extracted_fields': {'appointments': [
                 _appointment('Zqlecturer', '1988-03', '1989-05'),
                 _appointment('Zqreader', '1989-06', 'present')]}}
    source, target = tmp_path / 'in.json', tmp_path / 'out.docx'
    source.write_text(json.dumps({'document_uid': 'TESTAA', 'entries': [entry]}))
    WCMTemplateGenerator(verbose=False).generate(str(source), str(target), research_summary_path=None)
    body = ' '.join(t.text or '' for t in Document(str(target)).element.body.iter(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    assert 'Zqlecturer' in body and 'Zqreader' in body


# --- #1445: residue of the stage-4 records split (EOAHMI recheck) ----------
# Invented values, in the shapes the EOAHMI runs carried: a Q1 role held over
# three terms (DUTAVD 82), a board list whose undated last line took the first
# line's term (WYMVGU 865), one talk given on several dates in one year
# (QTATUP 980).


def _q1_record(role, start, end, organization=None):
    return {'role': role, 'organization': organization, 'start_date': start, 'end_date': end}


class TestStage4RecordOwnText:
    def test_the_last_record_falls_back_to_its_own_line_when_the_fields_hold_the_text(self):
        records = [_q1_record('Zqhistorian', '1992', '1993'),
                   _q1_record('Zqhistorian', '1994', '1996'),
                   _q1_record('Zqhistorian', '1997', '1998')]
        children = _fan4(_stage4_entry(records, code='Q1',
                                       text='Zqhistorian, 1992-1993, 1994-1996, 1997-1998'))
        last = children[-1]
        assert last['text'] == 'Zqhistorian, 1992-1993, 1994-1996, 1997-1998'
        assert last[LAST_STAGE4_RECORD][fan_out.OWN_TEXT_KEY] == 'Zqhistorian | 1997 | 1998'
        assert [fan_out.fallback_text(c) for c in children] == [
            'Zqhistorian | 1992 | 1993', 'Zqhistorian | 1994 | 1996', 'Zqhistorian | 1997 | 1998']

    def test_the_own_line_is_the_last_tab_segment_when_the_segments_line_up(self):
        text = 'Chair, Glade Board 1999-2001\tMember, Fern Council 2001-2003\tMember, Moss Panel 2004-2006'
        records = [{**r, 'institution': None} for r in copy.deepcopy(_THREE_COMMITTEES)]
        children = _fan4(_stage4_entry(records, text=text))
        assert fan_out.fallback_text(children[-1]) == 'Member, Moss Panel 2004-2006'

    def test_text_no_field_holds_keeps_the_whole_line_as_the_fallback(self):
        # BRUSUZ 196: the venue names are in no record, so the whole line is
        # the only place they are kept.
        records = [{'location': 'Ashby'}, {'location': 'Varnor'}]
        text = 'Zqfoundation Center (Ashby) Zqcapital Partners (Varnor)'
        children = _fan4(_stage4_entry(records, code='R', text=text))
        assert children[-1][LAST_STAGE4_RECORD][fan_out.OWN_TEXT_KEY] is None
        assert fan_out.fallback_text(children[-1]) == text

    def test_an_entry_that_was_not_split_falls_back_to_its_text(self):
        assert fan_out.fallback_text({'text': 'Glade Board'}) == 'Glade Board'
        assert fan_out.fallback_text({'text': None}) == ''

    def test_is_split_record(self):
        children = _fan4(_stage4_entry(copy.deepcopy(_THREE_COMMITTEES)))
        assert [fan_out.is_split_record(c) for c in children] == [True, True, True]
        assert [fan_out.is_split_record(c) for c in _fan(_THREE_HONORS)] == [True] * 3
        assert not fan_out.is_split_record({'text': 'Glade Board'})


class TestBuiltTextIsMarked:
    """#1556 e: a child whose text is the line built from its fields says so
    (`has_built_text`), so a renderer does not re-parse the built line's
    separators as a source pipe column. A child that keeps a line of the CV
    -- its tab segment, or the parent's whole text -- does not."""

    def test_stage4_children_with_built_lines_are_marked_and_the_last_is_not(self):
        children = _fan4(_stage4_entry(copy.deepcopy(_THREE_COMMITTEES)))
        assert [fan_out.has_built_text(c) for c in children] == [True, True, False]

    def test_stage4_children_with_their_own_tab_segments_are_not_marked(self):
        text = 'Chair, Glade Board 1999-2001\tMember, Fern Council 2001-2003\tMember, Moss Panel 2004-2006'
        children = _fan4(_stage4_entry(copy.deepcopy(_THREE_COMMITTEES), text=text))
        assert not any(fan_out.has_built_text(c) for c in children)

    def test_generic_children_follow_the_same_rule(self):
        assert not any(fan_out.has_built_text(c) for c in _fan(_THREE_HONORS))
        wrapped = _honors('Alpha Prize, Hollis College\tGraduate School\tBeta Prize',
                          [_award('Alpha Prize Graduate School'), _award('Beta Prize')])
        assert [fan_out.has_built_text(c) for c in _fan(wrapped)] == [True, True]

    def test_a_lone_record_keeps_the_whole_text_unmarked(self):
        # Two tab segments, one record: the child's text is the CV's own text.
        entry = {'taxonomy_code': 'D1',
                 'text': 'Varnor College, School of Botany\tLecturer in Botany, 1977-1979',
                 'extracted_fields': {'appointments': [_appointment(
                     'Lecturer in Botany', '1977', '1979', 'Varnor College, School of Botany')],
                     'start_date': '1977', 'end_date': '1979'}}
        (child,) = _fan(entry)
        assert child['text'] == entry['text']
        assert not fan_out.has_built_text(child)

    def test_an_entry_that_was_not_split_is_not_marked(self):
        assert not fan_out.has_built_text({'text': 'Glade Board | Member'})


class TestStage4RecordDates:
    _BOARD_TEXT = ('1993-2004 Member, Zqboard, Ashby Clinic\t1996-97 Vice-President\t'
                   '1997-99 President\tPast-President')

    def _board(self, last_dates=('1993', '2004')):
        return [_q1_record('Member', '1993', '2004', 'Ashby Clinic'),
                _q1_record('Vice-President', '1996', '1997', 'Ashby Clinic'),
                _q1_record('President', '1997', '1999', 'Ashby Clinic'),
                _q1_record('Past-President', *last_dates, 'Ashby Clinic')]

    def test_an_undated_line_loses_dates_it_took_from_a_line_other_than_the_one_above(self):
        children = _fan4(_stage4_entry(self._board(), code='Q1', text=self._BOARD_TEXT))
        last = children[-1]['extracted_fields']
        assert (last['role'], last['start_date'], last['end_date']) == ('Past-President', None, None)
        assert [c['extracted_fields']['start_date'] for c in children[:-1]] == ['1993', '1996', '1997']

    def test_an_earlier_undated_record_loses_them_too(self):
        records = self._board()
        records[2], records[3] = records[3], records[2]
        text = '1993-2004 Member, Zqboard, Ashby Clinic\t1996-97 Vice-President\tPast-President\t1997-99 President'
        children = _fan4(_stage4_entry(records, code='Q1', text=text))
        assert children[2]['extracted_fields']['start_date'] is None
        assert children[3]['extracted_fields']['start_date'] == '1997'

    @pytest.mark.parametrize('second_date', ['2015', '2015-05-10'],
                             ids=['year', 'a-day-the-year-column-does-not-show'])
    def test_an_undated_line_keeps_the_dates_of_the_line_above_it(self, second_date):
        # A year printed once over two talks.
        records = [{'title': 'Zqtalk one', 'date': '2015'}, {'title': 'Zqtalk two', 'date': second_date}]
        children = _fan4(_stage4_entry(records, code='R', text='2015 Zqtalk one\tZqtalk two'))
        assert [c['extracted_fields']['date'] for c in children] == ['2015', second_date]

    @pytest.mark.parametrize('text', [
        'Past-President\t1993-2004 Member, Zqboard\t1996-97 Vice-President\t1997-99 President',
        '1993-2004 Member, Zqboard\tVice-President, President, Past-President',
    ], ids=['undated-line-above-every-dated-one', 'lines-do-not-line-up'])
    def test_dates_are_left_as_stage4_wrote_them(self, text):
        records = self._board()
        if text.startswith('Past'):
            records = [records[3], *records[:3]]
        children = _fan4(_stage4_entry(records, code='Q1', text=text))
        assert all(c['extracted_fields']['start_date'] for c in children)


    def test_a_record_named_beside_its_dates_keeps_them_off_step(self):
        # OIEPQD TVQZIP 116 (#1342): four lines, four records, out of step
        # with each other. The co-leader's 2020-2021 is on the line that names
        # it, so it stays; the undated member after it gets none.
        text = ('2018.01 - Zqchair\tDepartment of Zqpathology\tAdjunct Professor, Zqsurgery\t'
                'Co-leader of Zqtherapeutics program 2020-2021 Member of Zqcancer Center')
        records = [{'title': 'Zqchair', 'institution': 'Department of Zqpathology',
                    'start_date': '2018-01', 'end_date': 'present'},
                   {'title': 'Adjunct Professor', 'institution': 'Zqsurgery'},
                   {'title': 'Co-leader of Zqtherapeutics program', 'institution': None,
                    'start_date': '2020', 'end_date': '2021'},
                   {'title': 'Member', 'institution': 'Zqcancer Center'}]
        children = _fan4(_stage4_entry(records, code='D1', text=text))
        dates = {c['extracted_fields']['title']: (c['extracted_fields'].get('start_date'),
                                                  c['extracted_fields'].get('end_date'))
                 for c in children}
        assert dates['Co-leader of Zqtherapeutics program'] == ('2020', '2021')
        assert dates['Member'] == (None, None)

    @pytest.mark.parametrize('records, text', [
        # "President" is written inside its sibling's "Vice President" line.
        ([_q1_record('Vice President', '2010', '2012'), _q1_record('Treasurer', '2014', '2015'),
          _q1_record('President', '2010', '2012')],
         'Vice President 2010-2012\tTreasurer 2014-2015\tPresident'),
        # "Member" sits at the end of a sibling's line, not at its start.
        ([_q1_record('Executive Committee Member', '2010', '2012'), _q1_record('Treasurer', '2014', '2015'),
          _q1_record('Member', '2010', '2012')],
         'Executive Committee Member 2010-2012\tTreasurer 2014-2015\tMember'),
        # "Glade" opens its sibling's "Glade Society Chair" line, a value of that sibling.
        ([_q1_record('Chair', '2010', '2012', 'Glade Society'), _q1_record('Treasurer', '2014', '2015'),
          _q1_record('Member', '2010', '2012', 'Glade')],
         'Glade Society Chair 2010-2012\tTreasurer 2014-2015\tGlade Member'),
        # The sibling stored under another spelling: its line still is not this record's.
        ([_q1_record('VP', '2010', '2012'), _q1_record('Treasurer', '2014', '2015'),
          _q1_record('President', '2010', '2012')],
         'Vice President 2010-2012\tTreasurer 2014-2015\tPresident'),
        # "President" does not open "President-elect", a hyphenated word.
        ([_q1_record('Pres.-elect', '2010', '2012'), _q1_record('Treasurer', '2014', '2015'),
          _q1_record('President', '2010', '2012')],
         'President-elect 2010-2012\tTreasurer 2014-2015\tPresident'),
        # A date-first sibling line: the year that opens it is not a name,
        # though that sibling's reply dropped it.
        ([_q1_record('Vice President', None, None), _q1_record('Treasurer', '2014', None),
          _q1_record('President', '2010', None)],
         '2010 Vice President\t2014 Treasurer\tPresident'),
        # The sibling's reply left out the words its line opens with, so the
        # later record's organization opens that line; its dates equal the
        # sibling's, so they are a copy.
        ([_q1_record('Council Chair', '2010', '2012'), _q1_record('Treasurer', '2014', '2015'),
          _q1_record('Member', '2010', '2012', 'American Heart Association')],
         'American Heart Association Council Chair 2010-2012\tTreasurer 2014-2015\tMember'),
    ], ids=['inside-a-sibling-name', 'end-of-a-sibling-line', 'start-of-a-sibling-value',
            'sibling-spelled-otherwise', 'hyphenated-word', 'date-first-sibling-line',
            'sibling-dates-equal'])
    def test_a_name_on_a_siblings_line_does_not_keep_its_copied_date(self, records, text):
        # #1342 round 3: each last record's 2010-2012 is copied from the
        # first line, which names a sibling; dev clears it, and so must this.
        assert fan_out._unsourced_date_records(records, text, 'Q1') == {2}


class TestStage4RepeatedRows:
    def _talk(self, date, title='Zqsafety talk'):
        return {'title': title, 'location': 'Ashby', 'date': date,
                'event_name': 'Zqteleconference series', 'role': None}

    def test_records_a_year_only_column_prints_identically_render_once(self):
        records = [self._talk(d) for d in ('2003-07-30', '2003-09-08', '2003-12-03')]
        children = _fan4(_stage4_entry(records, code='R', text='Zqsafety talk, Ashby, 2003'))
        assert len(children) == 1
        assert children[0]['extracted_fields']['date'] == '2003-12-03'
        assert LAST_STAGE4_RECORD in children[0]

    def test_only_the_repeated_records_go(self):
        records = [self._talk('2003-07-30'), self._talk('2004-01-05'),
                   self._talk('2003-09-01', title='Zqother talk'), self._talk('2003-12-03')]
        children = _fan4(_stage4_entry(records, code='R', text='Zqsafety talk and Zqother talk'))
        assert [(c['extracted_fields']['title'], c['extracted_fields']['date']) for c in children] == [
            ('Zqsafety talk', '2004-01-05'), ('Zqother talk', '2003-09-01'),
            ('Zqsafety talk', '2003-12-03')]

    def test_a_value_no_cell_prints_does_not_tell_rows_apart(self):
        records = [{**self._talk('2011-03'), 'target_name': 'Zqfirst'},
                   {**self._talk('2011-09'), 'target_name': 'Zqsecond'}]
        assert len(_fan4(_stage4_entry(records, code='R', text='Zqsafety talk, 2011'))) == 1

    def test_a_month_the_date_column_shows_keeps_both_records(self):
        records = [{'training_type': 'Resident', 'institution': 'Ashby Hospital',
                    'start_date': start, 'end_date': end}
                   for start, end in (('2001-07', '2002-06'), ('2002-07', '2002-12'))]
        records[1]['start_date'] = '2001-09'
        records[1]['end_date'] = '2002-06'
        assert len(_fan4(_stage4_entry(records, code='C', text='Resident, Ashby Hospital'))) == 2

    # #1556 (d), OIEPQD XAYKXA entry 134: Q4D's and B2's schema names no
    # start/end, but their sections write the range, so two spans are two rows.
    # #1901, NECRGC EYTHNB entry 164: H's section writes the range too when
    # no `date` is set.
    @pytest.mark.parametrize('code,name_key', [('Q4D', 'journal_name'), ('B2', 'program_name'),
                                               ('H', 'award_name')])
    def test_two_spans_the_date_column_shows_keep_both_records(self, code, name_key):
        records = [{name_key: 'Zqjournal of Lanterns', 'start_date': '2009', 'end_date': 'present'},
                   {name_key: 'Zqjournal of Lanterns', 'start_date': '2016', 'end_date': '2019'}]
        children = _fan4(_stage4_entry(
            records, code=code, text='Zqjournal of Lanterns reviewer since 2009, board 2016-2019'))
        assert [(c['extracted_fields']['start_date'], c['extracted_fields']['end_date'])
                for c in children] == [('2009', 'present'), ('2016', '2019')]

    @pytest.mark.parametrize('code,name_key', [('Q4D', 'journal_name'), ('B2', 'program_name'),
                                               ('H', 'award_name')])
    def test_one_span_listed_twice_still_renders_once(self, code, name_key):
        records = [{name_key: 'Zqjournal of Lanterns', 'start_date': '2016', 'end_date': '2019'}] * 2
        children = _fan4(_stage4_entry(
            copy.deepcopy(records), code=code, text='Zqjournal of Lanterns 2016-2019'))
        assert len(children) == 1

    def test_two_h_spans_behind_one_shown_date_still_render_once(self):
        """H's cell shows the `date` over a range beside a start of its own
        (`honors._record_date`), so two such rows print alike."""
        records = [{'award_name': 'Zqaward of Lanterns', 'date': '2016',
                    'start_date': start, 'end_date': '2019'} for start in ('2016', '2017')]
        children = _fan4(_stage4_entry(records, code='H', text='Zqaward of Lanterns 2016'))
        assert len(children) == 1

    def test_two_b1_attendance_spans_keep_both_records(self):
        records = [{'degree': 'Zqcertificate', 'institution': 'Ashby College', 'year': None,
                    'dates_attended': {'start_date': start, 'end_date': end}}
                   for start, end in (('1990', '1992'), ('1995', '1997'))]
        children = _fan4(_stage4_entry(
            records, code='B1', text='Zqcertificate, Ashby College, 1990-1992 and 1995-1997'))
        assert len(children) == 2


# --- #1243: stage 5d's one citation on an entry stage 4 split into records ---
# NDMRSO VYRDHN 607/613: two tab-joined articles, split by stage 4; 5d wrote
# one citation (of the first) and the split was declined, so the second
# article never rendered. Invented authors, titles and journals.


def _article(title, journal, authors='Zqowner A, Brack T'):
    return {'authors': authors, 'year': '2006', 'title': title, 'journal': journal,
            'volume': '5', 'issue': '1', 'pages': '10-19'}


_TWO_ARTICLES = [_article('Lanterns in the orchard model.', 'Zqjournal Alpha'),
                 _article('Copper weathering of the north arch.', 'Zqjournal Beta')]
_TWO_ARTICLES_TEXT = ('Zqowner A, Brack T. 2006. Lanterns in the orchard model. Zqjournal Alpha 5(1):10-19.'
                      '\tZqowner A, Brack T. 2006. Copper weathering of the north arch. '
                      'Zqjournal Beta 5(1):10-19.')


def _cited_entry(citation):
    entry = _stage4_entry(copy.deepcopy(_TWO_ARTICLES), code='S1', text=_TWO_ARTICLES_TEXT)
    entry['extracted_fields'].update({'formatted_citation': citation,
                                      'formatting_source': 'stage_5d_llm'})
    return entry


_FIRST_CITATION = ('Zqowner A, Brack T. "Lanterns in the Orchard Model." '
                   'Zqjournal Alpha. 2006;5(1):10-19.')


class TestStage5dCitationOnASplitEntry:
    def test_the_split_fans_out_and_the_citation_stays_with_its_record(self):
        first, last = _fan4(_cited_entry(_FIRST_CITATION))
        assert first['extracted_fields']['formatted_citation'] == _FIRST_CITATION
        assert first['extracted_fields']['formatting_source'] == 'stage_5d_llm'
        assert last['extracted_fields']['title'] == 'Copper weathering of the north arch.'
        assert 'formatted_citation' not in last['extracted_fields']
        assert 'formatting_source' not in last['extracted_fields']

    def test_a_citation_of_the_last_record_stays_on_the_last_child(self):
        citation = 'Zqowner A, Brack T. Copper weathering of the north arch. Zqjournal Beta. 2006.'
        first, last = _fan4(_cited_entry(citation))
        assert 'formatted_citation' not in first['extracted_fields']
        assert last['extracted_fields']['formatted_citation'] == citation

    @pytest.mark.parametrize('citation', [
        'Zqowner A. An unrelated essay. Zqjournal Gamma. 2006.',
        'Lanterns in the orchard model; Copper weathering of the north arch. Zqmeeting; Ashby.',
    ], ids=['no_title', 'both_titles'])
    def test_a_citation_naming_no_one_record_declines_the_split(self, citation):
        # EOAHMI GHCIXA S8: a citation of every record renders them all, with
        # a venue and a place no S8 field holds; splitting it lost both.
        entry = _cited_entry(citation)
        assert _fan4(entry) == [entry]

    def test_both_articles_reach_the_document(self, tmp_path):
        from docx import Document

        from unified_pipeline.stage_6_word_template import WCMTemplateGenerator

        source, target = tmp_path / 'in.json', tmp_path / 'out.docx'
        source.write_text(json.dumps({'document_uid': 'TESTAA',
                                      'entries': [_cited_entry(_FIRST_CITATION)]}))
        WCMTemplateGenerator(verbose=False).generate(
            str(source), str(target), research_summary_path=None)
        body = ' '.join(t.text or '' for t in Document(str(target)).element.body.iter(
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        assert 'Lanterns in the Orchard Model' in body
        assert 'Copper weathering of the north arch' in body
        assert body.count('Lanterns in the') == 1


# --- #1243 (HNQBOI 19): one title stage 4 split into a record per year ---
# Beside a record of another title; stage 5d cited the split title once, with
# the year of one of its records. Invented titles.
_ARCH = 'Copper Weathering of the North Arch'
_ORCHARD = 'Lanterns in the Orchard Model'


def _yearly_entry(citation, years=('2009', '2008'), second_title=_ARCH, title_key='title', code='S8'):
    # A DOI's digits are no date: only the date fields' years name a record.
    records = [{'authors': None, 'year': years[0], title_key: _ARCH, 'doi': '10.1000/zq.2015.1'},
               {'authors': None, 'year': years[1], title_key: second_title, 'doi': '10.1000/zq.2016.2'},
               {'authors': None, 'year': '2007', title_key: _ORCHARD, 'doi': None}]
    text = '\t'.join(f'{record[title_key]}, {record["year"]}.' for record in records)
    entry = _stage4_entry(records, code=code, text=text)
    entry['extracted_fields'].update({'formatted_citation': citation, 'formatting_source': 'stage_5d_llm'})
    return entry


class TestStage5dCitationOfOneTitleSplitByYear:
    @pytest.mark.parametrize('year, cited', [('2009', 0), ('2008', 1)])
    @pytest.mark.parametrize('title_key, code', [('title', 'S8'), ('chapter_title', 'S4')])
    def test_the_citation_goes_to_the_record_of_its_year(self, year, cited, title_key, code):
        citation = f'{_ARCH}; {year}.'
        children = _fan4(_yearly_entry(citation, title_key=title_key, code=code))
        assert [child['extracted_fields']['year'] for child in children] == ['2009', '2008', '2007']
        assert [child['extracted_fields'].get('formatted_citation') for child in children] == [
            citation if i == cited else None for i in range(3)]
        # A citation of the one other title goes to its record, year or no year.
        orchard = _fan4(_yearly_entry(f'{_ORCHARD}. Ashby Meeting.'))
        assert [child['extracted_fields'].get('formatted_citation') for child in orchard] == [
            None, None, f'{_ORCHARD}. Ashby Meeting.']

    @pytest.mark.parametrize('years, second_title, citation', [
        (('2009', '2009'), _ARCH, f'{_ARCH}; 2009.'),
        (('2009', '2008'), 'Copper Weathering of the South Arch',
         f'{_ARCH}; Copper Weathering of the South Arch; 2009.'),
        (('2009', '2008'), _ARCH, f'{_ARCH}; 2009, 2008.'),
        (('2009', None), _ARCH, f'{_ARCH}; 2009.'),
    ], ids=['same_year', 'different_titles', 'every_year', 'undated_twin'])
    def test_a_citation_its_years_do_not_tie_to_one_record_declines_the_split(
            self, years, second_title, citation):
        # BCTOGR 2166: one article split twice in one year; a citation of two
        # titles is of the whole list; EYTHNB 89: a citation writing every
        # record's year renders them all; an undated twin fits any citation.
        entry = _yearly_entry(citation, years, second_title)
        assert _fan4(entry) == [entry]
        # The same entry with its two years apart, cited with one, fans out.
        assert len(_fan4(_yearly_entry(f'{_ARCH}; 2009.'))) == 3

    @pytest.mark.parametrize('citation', [f'{_ARCH}. Ashby Journal. doi:10.1000/zq.2009.1',
                                          f'{_ARCH}. https://ashby.example.org/2008/arch',
                                          f'{_ARCH}. www.ashby.example.org/2009/arch'])
    def test_the_digits_of_a_doi_or_a_link_name_no_record(self, citation):
        # #1243 review: a citation with no year names neither record, whatever
        # its DOI or link writes.
        entry = _yearly_entry(citation)
        assert _fan4(entry) == [entry]

    def test_a_year_with_a_letter_names_its_record(self):
        # "2009a", an author-year list's second paper of the year, is 2009.
        citation = f'{_ARCH}; 2009a.'
        assert [child['extracted_fields'].get('formatted_citation') for child in _fan4(_yearly_entry(citation))] == [
            citation, None, None]
