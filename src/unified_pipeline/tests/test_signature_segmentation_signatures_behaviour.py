"""Behaviour tests for the pure-Python signature half of
segmentation/signature_based_segmentation.py (CViche issue #704 round 2, packet 2).

Covers: extract_all_paragraphs, FormatSignature.to_hash, rgb_to_hex, pt_to_inches,
extract_border_info, extract_paragraph_signature, group_by_signature,
compute_prominence_score, rescue_locked_headers, ensure_personal_data_first
(including its nested is_document_title), and remove_duplicate_children.

None of these functions call call_llm (classify_signature_groups_with_llm,
normalize_hierarchy_with_llm, parse_normalized_hierarchy,
validate_headers_vs_entries, build_hierarchy_from_classifications and
segment_cv_with_signatures are the LLM half and are covered by a sibling test
file) -- so no call_llm stub is needed here; importing the module does print a
harmless "auth_config.yaml not found" warning from unified_pipeline.llm_client's
own import-time config load, which is expected and not asserted on.

Untestable without network/real corpus: none in this scope -- every function
here is pure Python/XML manipulation over paragraph and hierarchy dicts.
"""

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from unified_pipeline.segmentation import signature_based_segmentation as sbs  # noqa: E402


# ---------------------------------------------------------------------------
# Small shared XML helpers (paragraph-level borders and shading are not
# exposed by python-docx's high-level API, only via raw pPr children).
# ---------------------------------------------------------------------------

def _add_paragraph_border(paragraph, side: str, *, sz: str | None = "8", color: str = "000000") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    if sz is not None:
        el.set(qn("w:sz"), sz)
    el.set(qn("w:color"), color)
    el.set(qn("w:space"), "1")
    pBdr.append(el)


def _add_shading(paragraph, fill: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _default_format_signature(**overrides) -> sbs.FormatSignature:
    base = dict(
        font_family="Arial", font_size_pt=12.0, bold=False, italic=False, underline="none",
        all_caps=False, small_caps=False, color="#000000", alignment="left",
        indent_left_in=0.0, indent_right_in=0.0, line_spacing=1.0, space_before_pt=0.0,
        space_after_pt=0.0, has_border_top=False, has_border_bottom=False, has_border_left=False,
        has_border_right=False, border_top_width_pt=0.0, border_bottom_width_pt=0.0,
        background_color=None, style_name="Normal", has_mixed_formatting=False,
    )
    base.update(overrides)
    return sbs.FormatSignature(**base)


def _prominence_paragraphs(count: int, position: float, canonical, *, matches_locked: bool = False,
                            case_type: str = "MixedCase") -> list[dict]:
    return [
        {
            "position_in_doc": position,
            "text_metadata": {"matches_locked_header": matches_locked, "case_type": case_type},
            "format_signature": canonical,
        }
        for _ in range(count)
    ]


# ---------------------------------------------------------------------------
# extract_all_paragraphs
# ---------------------------------------------------------------------------

def test_extract_all_paragraphs_includes_table_cell_text_in_document_order():
    doc = Document()
    doc.add_paragraph("Body para")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "TL"
    table.cell(0, 1).text = "TR"
    table.cell(1, 0).text = "BL"
    table.cell(1, 1).text = "BR"

    paras = sbs.extract_all_paragraphs(doc)
    texts = [p.text for p in paras]

    assert texts == ["Body para", "TL", "TR", "BL", "BR"]


def test_extract_all_paragraphs_dedupes_consecutive_identical_body_text():
    doc = Document()
    doc.add_paragraph("Body para 1")
    doc.add_paragraph("Repeated Line")
    doc.add_paragraph("Repeated Line")
    doc.add_paragraph("After")

    texts = [p.text for p in sbs.extract_all_paragraphs(doc)]

    assert texts == ["Body para 1", "Repeated Line", "After"]


def test_extract_all_paragraphs_keeps_consecutive_empty_paragraphs():
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("Nonempty")

    texts = [p.text for p in sbs.extract_all_paragraphs(doc)]

    assert texts == ["", "", "Nonempty"]


def test_extract_all_paragraphs_dedupes_horizontally_merged_table_cell():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "TL"
    table.cell(0, 1).text = "TR"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "MERGED"
    table.cell(1, 0).text = "BL"
    table.cell(1, 1).text = "BR"

    texts = [p.text for p in sbs.extract_all_paragraphs(doc)]

    assert texts.count("MERGED") == 1
    assert texts == ["MERGED", "BL", "BR"]


def test_extract_all_paragraphs_dedupes_horizontally_merged_empty_table_cell():
    # An EMPTY merged cell is the case that isolates the seen_cells id() guard:
    # the consecutive-identical-text dedup in the paragraph branch below never
    # kicks in for empty text (it always appends when `not current_text`), so
    # only the seen_cells guard can prevent the merged cell's second grid
    # position from being emitted a second time.
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = ""
    table.cell(0, 1).text = ""
    table.cell(0, 0).merge(table.cell(0, 1)).text = ""
    table.cell(1, 0).text = "BL"
    table.cell(1, 1).text = "BR"

    texts = [p.text for p in sbs.extract_all_paragraphs(doc)]

    assert texts == ["", "BL", "BR"]


def test_extract_all_paragraphs_table_unmerged_empty_cells_both_kept():
    # Isolates the table-branch `if current_text != last_text or not current_text`
    # fallback from the seen_cells merged-cell guard: with NO merge involved,
    # only the "or not current_text" half can save the second empty cell --
    # two unmerged empty cells in a row would otherwise look like a duplicate
    # of the previous empty cell and get dropped.
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = ""
    table.cell(0, 1).text = ""

    texts = [p.text for p in sbs.extract_all_paragraphs(doc)]

    assert texts == ["", ""]


# ---------------------------------------------------------------------------
# FormatSignature.to_hash
# ---------------------------------------------------------------------------

def test_to_hash_identical_signatures_hash_equal():
    a = _default_format_signature(font_size_pt=14.0, bold=True)
    b = _default_format_signature(font_size_pt=14.0, bold=True)

    assert a.to_hash() == b.to_hash()


def test_to_hash_differing_field_hashes_differently():
    a = _default_format_signature(bold=True)
    b = _default_format_signature(bold=False)

    assert a.to_hash() != b.to_hash()


def test_to_hash_rounds_floats_to_two_decimals_before_hashing():
    a = _default_format_signature(indent_left_in=0.101)
    b = _default_format_signature(indent_left_in=0.104)
    c = _default_format_signature(indent_left_in=0.2)

    assert a.to_hash() == b.to_hash()
    assert a.to_hash() != c.to_hash()

    # Pin the precision at exactly 2 decimals, not merely "at least 1": 0.11 and
    # 0.14 both round to 0.1 at 1-decimal precision (which the assertions above
    # alone do not rule out) but stay distinct (0.11 vs 0.14) at 2 decimals.
    d = _default_format_signature(indent_left_in=0.11)
    e = _default_format_signature(indent_left_in=0.14)
    assert d.to_hash() != e.to_hash()


# ---------------------------------------------------------------------------
# rgb_to_hex
# ---------------------------------------------------------------------------

def test_rgb_to_hex_none_returns_black():
    assert sbs.rgb_to_hex(None) == "#000000"


def test_rgb_to_hex_converts_rgb_color():
    assert sbs.rgb_to_hex(RGBColor(0x1A, 0x2B, 0x3C)) == "#1a2b3c"


# ---------------------------------------------------------------------------
# pt_to_inches
# ---------------------------------------------------------------------------

def test_pt_to_inches_none_returns_zero():
    assert sbs.pt_to_inches(None) == 0.0


def test_pt_to_inches_converts_points():
    assert sbs.pt_to_inches(Pt(72)) == 1.0


def test_pt_to_inches_falls_back_to_zero_without_inches_attribute():
    assert sbs.pt_to_inches(5) == 0.0


# ---------------------------------------------------------------------------
# extract_border_info
# ---------------------------------------------------------------------------

def test_extract_border_info_no_pPr_returns_all_none():
    doc = Document()
    p = doc.add_paragraph("No formatting at all")

    borders = sbs.extract_border_info(p)

    assert borders == {"top": None, "bottom": None, "left": None, "right": None}


def test_extract_border_info_top_and_bottom_borders_with_explicit_width():
    doc = Document()
    p = doc.add_paragraph("Bordered header")
    _add_paragraph_border(p, "top", sz="8", color="000000")
    _add_paragraph_border(p, "bottom", sz="8", color="000000")

    borders = sbs.extract_border_info(p)

    assert borders["top"] == {"type": "single", "pt": 1.0, "color": "000000", "space_twips": "1"}
    assert borders["bottom"]["pt"] == 1.0
    assert borders["left"] is None
    assert borders["right"] is None


def test_extract_border_info_defaults_width_to_half_point_without_sz():
    doc = Document()
    p = doc.add_paragraph("Border without explicit size")
    _add_paragraph_border(p, "left", sz=None, color="FF0000")

    borders = sbs.extract_border_info(p)

    assert borders["left"]["pt"] == 0.5
    assert borders["left"]["color"] == "FF0000"


def test_extract_border_info_defaults_type_to_single_without_val():
    # Every other border fixture in this file sets w:val="single" explicitly
    # via _add_paragraph_border, which never exercises the .get(..., 'single')
    # fallback -- build a border element with only w:sz, no w:val attribute.
    doc = Document()
    p = doc.add_paragraph("Border with size but no explicit val")
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    pPr.append(pBdr)
    top = OxmlElement("w:top")
    top.set(qn("w:sz"), "8")
    pBdr.append(top)

    borders = sbs.extract_border_info(p)

    assert borders["top"]["type"] == "single"


# ---------------------------------------------------------------------------
# extract_paragraph_signature
# ---------------------------------------------------------------------------

def test_extract_paragraph_signature_bold_all_caps_header_like():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("EDUCATION")
    run.bold = True
    run.font.size = Pt(16)

    sig = sbs.extract_paragraph_signature(p, 0, 10)

    assert sig["format_signature"].bold is True
    assert sig["format_signature"].font_size_pt == 16.0
    assert sig["text_metadata"]["case_type"] == "ALL_CAPS"
    assert sig["text_metadata"]["all_caps"] is True


def test_extract_paragraph_signature_case_type_variants():
    doc = Document()
    cases = {
        "Title Case Words Here": "TitleCaseWords",
        "lowercase only text": "lowercase",
        "Mixed case Here now": "MixedCase",
    }
    for text, expected in cases.items():
        p = doc.add_paragraph(text)
        sig = sbs.extract_paragraph_signature(p, 0, 1)
        assert sig["text_metadata"]["case_type"] == expected, text


def test_extract_paragraph_signature_format_based_all_caps_overrides_mixed_case_text():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("Mixed Case Text")
    run.font.all_caps = True

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["format_signature"].all_caps is True
    assert sig["text_metadata"]["case_type"] == "ALL_CAPS"


def test_extract_paragraph_signature_mixed_formatting_true_when_runs_differ():
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Bold Part ")
    r1.bold = True
    r2 = p.add_run("Not Bold Part")
    r2.bold = False

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["format_signature"].has_mixed_formatting is True


def test_extract_paragraph_signature_mixed_formatting_false_when_runs_uniform():
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Bold One ")
    r1.bold = True
    r2 = p.add_run("Bold Two")
    r2.bold = True

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["format_signature"].has_mixed_formatting is False


def test_extract_paragraph_signature_mixed_formatting_false_at_exactly_three_chars():
    # The mixed-formatting check only fires once cumulative non-empty run text
    # exceeds 3 chars (`text_char_count > 3`). Two runs totalling EXACTLY 3
    # chars ("ab" + "c") must never cross that threshold, so a differing-bold
    # second run must NOT flip has_mixed_formatting to True here.
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("ab")
    r1.bold = True
    r2 = p.add_run("c")
    r2.bold = False

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["format_signature"].has_mixed_formatting is False


def test_extract_paragraph_signature_alignment_variants():
    doc = Document()
    p_default = doc.add_paragraph("No explicit alignment")
    p_center = doc.add_paragraph("Centered")
    p_center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_justify = doc.add_paragraph("Justified")
    p_justify.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    assert sbs.extract_paragraph_signature(p_default, 0, 3)["format_signature"].alignment == "left"
    assert sbs.extract_paragraph_signature(p_center, 1, 3)["format_signature"].alignment == "center"
    assert sbs.extract_paragraph_signature(p_justify, 2, 3)["format_signature"].alignment == "justify"


def test_extract_paragraph_signature_indent_spacing_and_line_spacing():
    doc = Document()
    p = doc.add_paragraph("Indented and spaced")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)

    fs = sbs.extract_paragraph_signature(p, 0, 1)["format_signature"]

    assert fs.indent_left_in == 0.5
    assert fs.indent_right_in == 0.25
    assert fs.line_spacing == 1.5
    assert fs.space_before_pt == 12.0
    assert fs.space_after_pt == 6.0


def test_extract_paragraph_signature_borders_flow_into_format_signature():
    doc = Document()
    p = doc.add_paragraph("Bordered header")
    _add_paragraph_border(p, "top", sz="8")
    _add_paragraph_border(p, "bottom", sz="8")

    fs = sbs.extract_paragraph_signature(p, 0, 1)["format_signature"]

    assert fs.has_border_top is True
    assert fs.has_border_bottom is True
    assert fs.has_border_left is False
    assert fs.border_top_width_pt == 1.0


def test_extract_paragraph_signature_background_shading():
    doc = Document()
    p_shaded = doc.add_paragraph("Shaded")
    _add_shading(p_shaded, "FFFF00")
    p_auto = doc.add_paragraph("Auto fill means no real shading")
    _add_shading(p_auto, "auto")

    assert sbs.extract_paragraph_signature(p_shaded, 0, 2)["format_signature"].background_color == "#FFFF00"
    assert sbs.extract_paragraph_signature(p_auto, 1, 2)["format_signature"].background_color is None


def test_extract_paragraph_signature_style_based_inheritance():
    doc = Document()
    style = doc.styles.add_style("MyHeaderStyle", WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    style.font.size = Pt(18)
    style.font.name = "Georgia"
    p = doc.add_paragraph(style="MyHeaderStyle")
    p.add_run("Styled Header")

    fs = sbs.extract_paragraph_signature(p, 0, 1)["format_signature"]

    assert fs.font_family == "Georgia"
    assert fs.font_size_pt == 18.0
    assert fs.bold is True
    assert fs.style_name == "MyHeaderStyle"


def test_extract_paragraph_signature_direct_run_level_font_name_italic_and_color():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("Direct formatting")
    run.font.name = "Consolas"
    run.italic = True
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)

    fs = sbs.extract_paragraph_signature(p, 0, 1)["format_signature"]

    assert fs.font_family == "Consolas"
    assert fs.italic is True
    assert fs.color == "#123456"


def test_extract_paragraph_signature_style_based_italic_and_small_caps_inheritance():
    doc = Document()
    style = doc.styles.add_style("ItalicSmallCapsStyle", WD_STYLE_TYPE.PARAGRAPH)
    style.font.italic = True
    style.font.small_caps = True
    p = doc.add_paragraph(style="ItalicSmallCapsStyle")
    p.add_run("Styled Italic SmallCaps")

    fs = sbs.extract_paragraph_signature(p, 0, 1)["format_signature"]

    assert fs.italic is True
    assert fs.small_caps is True


def test_extract_paragraph_signature_style_based_all_caps_on_non_upper_text():
    doc = Document()
    style = doc.styles.add_style("AllCapsStyle", WD_STYLE_TYPE.PARAGRAPH)
    style.font.all_caps = True
    p = doc.add_paragraph(style="AllCapsStyle")
    p.add_run("Mixed Case Via Style")

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["format_signature"].all_caps is True
    assert sig["text_metadata"]["case_type"] == "ALL_CAPS"


def test_extract_paragraph_signature_whitespace_only_run_falls_back_to_first_run():
    # No run has non-whitespace text, so the "first non-empty run" scan finds
    # nothing and the function must fall back to para.runs[0] -- if that
    # fallback is skipped, bold/font_size_pt never pick up this run's
    # formatting and silently stay at their pre-run defaults (False / 12.0).
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("   ")
    run.bold = True
    run.font.size = Pt(20)

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["text"] == ""
    assert sig["format_signature"].bold is True
    assert sig["format_signature"].font_size_pt == 20.0


def test_extract_paragraph_signature_mixed_formatting_skips_empty_runs():
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Bold Part ")
    r1.bold = True
    p.add_run("")  # empty run must be skipped, not counted toward dominant formatting
    r3 = p.add_run("Not Bold Part")
    r3.bold = False

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["format_signature"].has_mixed_formatting is True


def test_extract_paragraph_signature_underline_and_small_caps():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("Underlined SmallCaps")
    run.underline = True
    run.font.small_caps = True

    fs = sbs.extract_paragraph_signature(p, 0, 1)["format_signature"]

    assert fs.underline == "True"
    assert fs.small_caps is True


def test_extract_paragraph_signature_matches_locked_header_strips_parenthetical_and_punctuation():
    doc = Document()
    p_match = doc.add_paragraph("Honors and Awards (Recent Years):")
    p_no_match = doc.add_paragraph("Random Body Sentence Here")

    assert sbs.extract_paragraph_signature(p_match, 0, 2)["text_metadata"]["matches_locked_header"] is True
    assert sbs.extract_paragraph_signature(p_no_match, 1, 2)["text_metadata"]["matches_locked_header"] is False


def test_extract_paragraph_signature_position_in_doc_and_zero_total_paras():
    doc = Document()
    p = doc.add_paragraph("Some text")

    sig_mid = sbs.extract_paragraph_signature(p, 5, 10)
    sig_zero_total = sbs.extract_paragraph_signature(p, 0, 0)

    assert sig_mid["position_in_doc"] == 0.5
    assert sig_zero_total["position_in_doc"] == 0.0


def test_extract_paragraph_signature_tokens_and_single_line_with_soft_break():
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Line one has five")
    r1.add_break()
    p.add_run("Line two")

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["text_metadata"]["single_line"] is False
    assert sig["text_metadata"]["tokens_count"] == len(sig["text_metadata"]["tokens"])
    assert "\n" in sig["text"]


def test_extract_paragraph_signature_signature_hash_matches_to_hash():
    doc = Document()
    p = doc.add_paragraph("Consistency check")

    sig = sbs.extract_paragraph_signature(p, 0, 1)

    assert sig["signature_hash"] == sig["format_signature"].to_hash()


# ---------------------------------------------------------------------------
# group_by_signature
# ---------------------------------------------------------------------------

def test_group_by_signature_groups_matching_hashes_together():
    paras = [
        {"signature_hash": "h1", "text": "a"},
        {"signature_hash": "h1", "text": "b"},
        {"signature_hash": "h2", "text": "c"},
    ]

    groups = sbs.group_by_signature(paras)

    assert set(groups.keys()) == {"h1", "h2"}
    assert [p["text"] for p in groups["h1"]] == ["a", "b"]
    assert [p["text"] for p in groups["h2"]] == ["c"]


# ---------------------------------------------------------------------------
# compute_prominence_score
# ---------------------------------------------------------------------------

def test_compute_prominence_score_header_group_scores_higher_than_body_group():
    header_sig = _default_format_signature(
        font_size_pt=16.0, bold=True, all_caps=True, has_border_top=True, space_before_pt=14.0,
    )
    body_sig = _default_format_signature(
        font_size_pt=10.0, bold=False, all_caps=False, indent_left_in=0.5, has_mixed_formatting=True,
    )
    header_group = _prominence_paragraphs(5, 0.05, header_sig, case_type="ALL_CAPS")
    body_group = _prominence_paragraphs(200, 0.5, body_sig, case_type="MixedCase")

    header_score = sbs.compute_prominence_score(header_group)
    body_score = sbs.compute_prominence_score(body_group)

    assert header_score > body_score


def test_compute_prominence_score_early_position_bonus_is_exactly_weighted():
    canonical = _default_format_signature()
    early_group = _prominence_paragraphs(5, 0.05, canonical)
    late_group = _prominence_paragraphs(5, 0.5, canonical)

    delta = sbs.compute_prominence_score(early_group) - sbs.compute_prominence_score(late_group)

    assert delta == pytest.approx(0.05)


def test_compute_prominence_score_font_size_term_clamps_at_16pt():
    # 12pt -> 0.0, 16pt -> 1.0 (0.15 weight); anything above 16pt clamps, so
    # a 20pt header scores exactly the same as a 16pt one.
    at_16 = _prominence_paragraphs(5, 0.5, _default_format_signature(font_size_pt=16.0))
    at_20 = _prominence_paragraphs(5, 0.5, _default_format_signature(font_size_pt=20.0))
    at_12 = _prominence_paragraphs(5, 0.5, _default_format_signature(font_size_pt=12.0))

    assert sbs.compute_prominence_score(at_20) == pytest.approx(sbs.compute_prominence_score(at_16))
    assert sbs.compute_prominence_score(at_16) - sbs.compute_prominence_score(at_12) == pytest.approx(0.15)


def test_compute_prominence_score_space_before_term_clamps_at_12pt():
    at_12 = _prominence_paragraphs(5, 0.5, _default_format_signature(space_before_pt=12.0))
    at_24 = _prominence_paragraphs(5, 0.5, _default_format_signature(space_before_pt=24.0))
    at_0 = _prominence_paragraphs(5, 0.5, _default_format_signature(space_before_pt=0.0))

    assert sbs.compute_prominence_score(at_24) == pytest.approx(sbs.compute_prominence_score(at_12))
    assert sbs.compute_prominence_score(at_12) - sbs.compute_prominence_score(at_0) == pytest.approx(0.10)


def test_compute_prominence_score_indent_penalty_clamps_at_half_inch():
    at_half = _prominence_paragraphs(5, 0.5, _default_format_signature(indent_left_in=0.5))
    at_one = _prominence_paragraphs(5, 0.5, _default_format_signature(indent_left_in=1.0))
    at_zero = _prominence_paragraphs(5, 0.5, _default_format_signature(indent_left_in=0.0))

    assert sbs.compute_prominence_score(at_one) == pytest.approx(sbs.compute_prominence_score(at_half))
    assert sbs.compute_prominence_score(at_zero) - sbs.compute_prominence_score(at_half) == pytest.approx(0.20)


def test_compute_prominence_score_title_case_casing_rank_is_exactly_weighted():
    canonical = _default_format_signature()
    title_case_group = _prominence_paragraphs(5, 0.1, canonical, case_type="TitleCaseWords")
    unranked_group = _prominence_paragraphs(5, 0.1, canonical, case_type="lowercase")

    delta = sbs.compute_prominence_score(title_case_group) - sbs.compute_prominence_score(unranked_group)

    assert delta == pytest.approx(0.7 * 0.10)


def test_compute_prominence_score_low_count_bonus_is_exactly_weighted():
    canonical = _default_format_signature()
    low_count_group = _prominence_paragraphs(5, 0.1, canonical)
    high_count_group = _prominence_paragraphs(50, 0.1, canonical)

    delta = sbs.compute_prominence_score(low_count_group) - sbs.compute_prominence_score(high_count_group)

    assert delta == pytest.approx(0.10)


def test_compute_prominence_score_centered_alignment_is_exactly_weighted():
    left_canonical = _default_format_signature(alignment="left")
    center_canonical = _default_format_signature(alignment="center")
    left_group = _prominence_paragraphs(5, 0.1, left_canonical)
    center_group = _prominence_paragraphs(5, 0.1, center_canonical)

    delta = sbs.compute_prominence_score(center_group) - sbs.compute_prominence_score(left_group)

    assert delta == pytest.approx(0.05)


def test_compute_prominence_score_indent_penalty_is_exactly_weighted():
    no_indent_canonical = _default_format_signature(indent_left_in=0.0)
    full_indent_canonical = _default_format_signature(indent_left_in=0.5)
    no_indent_group = _prominence_paragraphs(5, 0.1, no_indent_canonical)
    full_indent_group = _prominence_paragraphs(5, 0.1, full_indent_canonical)

    delta = sbs.compute_prominence_score(no_indent_group) - sbs.compute_prominence_score(full_indent_group)

    assert delta == pytest.approx(0.20)


def test_compute_prominence_score_mixed_formatting_penalty_is_exactly_weighted():
    # Isolates the CRITICAL mixed-formatting penalty (the function's largest
    # single term) from every other term by holding everything else fixed.
    uniform_canonical = _default_format_signature(has_mixed_formatting=False)
    mixed_canonical = _default_format_signature(has_mixed_formatting=True)
    uniform_group = _prominence_paragraphs(5, 0.5, uniform_canonical)
    mixed_group = _prominence_paragraphs(5, 0.5, mixed_canonical)

    delta = sbs.compute_prominence_score(uniform_group) - sbs.compute_prominence_score(mixed_group)

    assert delta == pytest.approx(0.8)


def test_compute_prominence_score_locked_header_ratio_is_exactly_weighted():
    # Isolates the locked_header_ratio term: every _prominence_paragraphs call
    # elsewhere in this file leaves matches_locked=False, so this is the only
    # test that ever sets it True and proves the 0.15 weight is applied.
    canonical = _default_format_signature()
    all_locked_group = _prominence_paragraphs(5, 0.5, canonical, matches_locked=True)
    none_locked_group = _prominence_paragraphs(5, 0.5, canonical, matches_locked=False)

    delta = sbs.compute_prominence_score(all_locked_group) - sbs.compute_prominence_score(none_locked_group)

    assert delta == pytest.approx(0.15)


def test_compute_prominence_score_italic_is_exactly_weighted():
    non_italic_canonical = _default_format_signature(italic=False)
    italic_canonical = _default_format_signature(italic=True)
    non_italic_group = _prominence_paragraphs(5, 0.5, non_italic_canonical)
    italic_group = _prominence_paragraphs(5, 0.5, italic_canonical)

    delta = sbs.compute_prominence_score(italic_group) - sbs.compute_prominence_score(non_italic_group)

    assert delta == pytest.approx(0.10)


def test_compute_prominence_score_font_size_is_exactly_weighted():
    # Isolates the normalized_font_size term (weight 0.15): 12pt normalizes to
    # 0.0, 16pt+ normalizes (and clamps) to 1.0, so the full 0.15 weight must
    # show up in the delta with every other field held at its default.
    small_canonical = _default_format_signature(font_size_pt=12.0)
    large_canonical = _default_format_signature(font_size_pt=16.0)
    small_group = _prominence_paragraphs(5, 0.5, small_canonical)
    large_group = _prominence_paragraphs(5, 0.5, large_canonical)

    delta = sbs.compute_prominence_score(large_group) - sbs.compute_prominence_score(small_group)

    assert delta == pytest.approx(0.15)


def test_compute_prominence_score_border_is_exactly_weighted_and_bottom_alone_counts():
    # Isolates the has_border term (weight 0.25, the function's "strong
    # signal"): only has_border_bottom is set here (top stays False), so this
    # also pins that has_border is `has_border_top OR has_border_bottom`, not
    # top alone -- a mutant that drops the bottom check from the `or` would
    # leave this delta at 0.0 instead of 0.25.
    bordered_canonical = _default_format_signature(has_border_bottom=True)
    plain_canonical = _default_format_signature()
    bordered_group = _prominence_paragraphs(5, 0.5, bordered_canonical)
    plain_group = _prominence_paragraphs(5, 0.5, plain_canonical)

    delta = sbs.compute_prominence_score(bordered_group) - sbs.compute_prominence_score(plain_group)

    assert delta == pytest.approx(0.25)


def test_compute_prominence_score_bold_is_exactly_weighted():
    bold_canonical = _default_format_signature(bold=True)
    plain_canonical = _default_format_signature(bold=False)
    bold_group = _prominence_paragraphs(5, 0.5, bold_canonical)
    plain_group = _prominence_paragraphs(5, 0.5, plain_canonical)

    delta = sbs.compute_prominence_score(bold_group) - sbs.compute_prominence_score(plain_group)

    assert delta == pytest.approx(0.15)


def test_compute_prominence_score_space_before_is_exactly_weighted():
    spaced_canonical = _default_format_signature(space_before_pt=12.0)
    unspaced_canonical = _default_format_signature(space_before_pt=0.0)
    spaced_group = _prominence_paragraphs(5, 0.5, spaced_canonical)
    unspaced_group = _prominence_paragraphs(5, 0.5, unspaced_canonical)

    delta = sbs.compute_prominence_score(spaced_group) - sbs.compute_prominence_score(unspaced_group)

    assert delta == pytest.approx(0.10)


def test_compute_prominence_score_mixed_case_casing_rank_is_exactly_weighted():
    # The existing title-case-vs-lowercase test pins the 0.7 rank; this pins
    # the separate MixedCase rank (0.3), the branch a MixedCase body group
    # relies on to differ from lowercase at all.
    canonical = _default_format_signature()
    mixed_case_group = _prominence_paragraphs(5, 0.5, canonical, case_type="MixedCase")
    lowercase_group = _prominence_paragraphs(5, 0.5, canonical, case_type="lowercase")

    delta = sbs.compute_prominence_score(mixed_case_group) - sbs.compute_prominence_score(lowercase_group)

    assert delta == pytest.approx(0.10 * 0.3)


def test_compute_prominence_score_early_position_boundary_at_exactly_0_2_gets_no_bonus():
    # The documented boundary is avg_position < 0.2 (strict). A group sitting
    # exactly AT 0.2 must NOT get the bonus; a group just under 0.2 must. A
    # `<=` mutant would grant the bonus to both, collapsing this delta to 0.0.
    canonical = _default_format_signature()
    at_boundary_group = _prominence_paragraphs(5, 0.2, canonical)
    just_under_group = _prominence_paragraphs(5, 0.19, canonical)

    delta = sbs.compute_prominence_score(just_under_group) - sbs.compute_prominence_score(at_boundary_group)

    assert delta == pytest.approx(0.05)


def test_compute_prominence_score_low_count_boundary_at_exactly_3_gets_bonus():
    # The documented range is 3 <= count <= 30. count=2 must NOT get the
    # bonus; count=3 (the lower boundary) must. A mutant widening the lower
    # bound to 1 <= count <= 30 would grant the bonus to count=2 too,
    # collapsing this delta to 0.0.
    canonical = _default_format_signature()
    at_lower_bound_group = _prominence_paragraphs(3, 0.5, canonical)
    below_lower_bound_group = _prominence_paragraphs(2, 0.5, canonical)

    delta = sbs.compute_prominence_score(at_lower_bound_group) - sbs.compute_prominence_score(below_lower_bound_group)

    assert delta == pytest.approx(0.10)


# ---------------------------------------------------------------------------
# rescue_locked_headers
# ---------------------------------------------------------------------------

def _rescue_group(text: str, *, level: str, matches_locked: bool):
    para = {"text": text, "text_metadata": {"matches_locked_header": matches_locked}}
    sig_hash = "h-" + text
    signature_groups = {sig_hash: [para]}
    classifications = {sig_hash: {"level": level, "confidence": 0.5, "reasoning": "seed"}}
    return signature_groups, classifications, para


def test_rescue_locked_headers_promotes_matching_not_header_paragraph():
    groups, classifications, para = _rescue_group("Teaching", level="NOT_HEADER", matches_locked=True)

    sbs.rescue_locked_headers(groups, classifications)

    assert para["classification"] == "H1"
    assert para["is_rescued_locked_header"] is True
    # build_hierarchy_from_classifications reads this straight onto the node.
    assert para["classification_confidence"] == 0.90


def test_rescue_locked_headers_leaves_non_matching_paragraph_untouched():
    groups, classifications, para = _rescue_group("Not a header at all", level="NOT_HEADER", matches_locked=False)

    sbs.rescue_locked_headers(groups, classifications)

    assert "classification" not in para


def test_rescue_locked_headers_promotes_all_of_more_than_ten_matching_paragraphs():
    # Exercises the "> 10 rescued" truncation-message branch; the real
    # assertion is still on returned data, not the print -- every one of the
    # 11 distinct-signature paragraphs must be promoted, not just the first 10.
    signature_groups = {}
    classifications = {}
    paras = []
    for i in range(11):
        para = {"text": f"Teaching {i}", "text_metadata": {"matches_locked_header": True}}
        sig_hash = f"h-{i}"
        signature_groups[sig_hash] = [para]
        classifications[sig_hash] = {"level": "NOT_HEADER", "confidence": 0.5, "reasoning": "seed"}
        paras.append(para)

    sbs.rescue_locked_headers(signature_groups, classifications)

    assert all(p["classification"] == "H1" for p in paras)


def test_rescue_locked_headers_skips_groups_already_classified_as_header():
    groups, classifications, para = _rescue_group("Teaching", level="H1", matches_locked=True)

    sbs.rescue_locked_headers(groups, classifications)

    assert "classification" not in para


def test_rescue_locked_headers_excludes_exact_document_title_text():
    groups, classifications, para = _rescue_group("CV", level="NOT_HEADER", matches_locked=True)

    sbs.rescue_locked_headers(groups, classifications)

    assert "classification" not in para


def test_rescue_locked_headers_excludes_overlong_text_even_if_matched():
    long_text = "Teaching " * 20  # > 100 chars
    groups, classifications, para = _rescue_group(long_text, level="NOT_HEADER", matches_locked=True)

    sbs.rescue_locked_headers(groups, classifications)

    assert "classification" not in para


def test_rescue_locked_headers_length_boundary_exactly_100_chars_still_rescued():
    # The cap is `len(text) > 100` (strict): exactly 100 chars must still be
    # rescued. A `>=` mutant would exclude this boundary text too, so this is
    # the only test that distinguishes ">" from ">=" -- the existing overlong
    # fixture (180 chars) is caught by either version of the check.
    text_100 = "A" * 100
    assert len(text_100) == 100
    groups, classifications, para = _rescue_group(text_100, level="NOT_HEADER", matches_locked=True)

    sbs.rescue_locked_headers(groups, classifications)

    assert para["classification"] == "H1"
    assert para["is_rescued_locked_header"] is True


def test_rescue_locked_headers_bug_document_title_with_trailing_punctuation_not_excluded():
    doc = Document()
    p = doc.add_paragraph("Curriculum Vitae:")
    sig = sbs.extract_paragraph_signature(p, 0, 1)
    assert sig["text_metadata"]["matches_locked_header"] is True  # precondition for the bug

    sig_hash = sig["signature_hash"]
    signature_groups = {sig_hash: [sig]}
    classifications = {sig_hash: {"level": "NOT_HEADER", "confidence": 0.5, "reasoning": "seed"}}

    sbs.rescue_locked_headers(signature_groups, classifications)

    assert "classification" not in sig  # correct behaviour: document titles must never be rescued


def _assert_title_not_rescued(title_text: str) -> None:
    doc = Document()
    p = doc.add_paragraph(title_text)
    sig = sbs.extract_paragraph_signature(p, 0, 1)
    assert sig["text_metadata"]["matches_locked_header"] is True  # precondition for the bug

    sig_hash = sig["signature_hash"]
    signature_groups = {sig_hash: [sig]}
    classifications = {sig_hash: {"level": "NOT_HEADER", "confidence": 0.5, "reasoning": "seed"}}

    sbs.rescue_locked_headers(signature_groups, classifications)

    assert "classification" not in sig  # document titles must never be rescued


def test_rescue_locked_headers_excludes_title_with_trailing_parenthetical():
    # #871: the signature step strips a trailing parenthetical before matching
    # LOCKED_CV_HEADERS, but rescue_locked_headers' EXCLUDED_HEADERS check did
    # not, so "Curriculum Vitae (CV)" escaped the exclusion (web152).
    _assert_title_not_rescued("Curriculum Vitae (CV)")


def test_rescue_locked_headers_excludes_title_with_semicolon():
    _assert_title_not_rescued("Resume;")


def test_rescue_locked_headers_excludes_title_with_period():
    _assert_title_not_rescued("Curriculum Vitae .")


def test_rescue_locked_headers_excludes_title_with_space_before_colon():
    # #871: rescue_locked_headers had no trailing .strip() after rstrip(':;.'),
    # so a title with whitespace before the colon left a trailing space that
    # never matched EXCLUDED_HEADERS.
    _assert_title_not_rescued("Curriculum Vitae :")


def test_rescue_locked_headers_still_rescues_real_header_with_parenthetical():
    # Negative case: a real section header with a parenthetical that is NOT
    # in EXCLUDED_HEADERS must still be rescued -- the parenthetical strip
    # must not swallow legitimate headers.
    doc = Document()
    p = doc.add_paragraph("Publications (peer reviewed)")
    sig = sbs.extract_paragraph_signature(p, 0, 1)
    sig["text_metadata"]["matches_locked_header"] = True  # force the rescue path

    sig_hash = sig["signature_hash"]
    signature_groups = {sig_hash: [sig]}
    classifications = {sig_hash: {"level": "NOT_HEADER", "confidence": 0.5, "reasoning": "seed"}}

    sbs.rescue_locked_headers(signature_groups, classifications)

    assert sig["classification"] == "H1"
    assert sig["is_rescued_locked_header"] is True


def test_normalise_header_text_strips_parenthetical_colon_semicolon_and_period():
    # Direct unit coverage for the shared normaliser (#857's verifier: a
    # mutant stripping only ':' survived rstrip(':;.') coverage -- pin ';'
    # and '.' explicitly alongside ':' and the parenthetical strip.
    assert sbs._normalise_header_text("Curriculum Vitae (CV)") == "curriculum vitae"
    assert sbs._normalise_header_text("Curriculum Vitae:") == "curriculum vitae"
    assert sbs._normalise_header_text("Resume;") == "resume"
    assert sbs._normalise_header_text("Curriculum Vitae .") == "curriculum vitae"
    # C-871 round 2 (verify-r1 finding 1, mutant m07): trailing punctuation
    # BEFORE the parenthetical requires the strip BEFORE the parenthetical
    # strip too, else the leftover space from the removed "(selected)"
    # segment survives into the lowered/rstripped result. Pins strip order.
    assert sbs._normalise_header_text("Publications: (selected)") == "publications"


def test_rescue_locked_headers_tolerates_none_or_blank_text():
    # A flagged paragraph with text=None or all-whitespace text must not
    # crash on `.strip()`/`.lower()` and must never be promoted to a header.
    groups_none, classifications_none, para_none = _rescue_group("placeholder", level="NOT_HEADER", matches_locked=True)
    para_none["text"] = None
    groups_blank, classifications_blank, para_blank = _rescue_group("  ", level="NOT_HEADER", matches_locked=True)

    sbs.rescue_locked_headers(groups_none, classifications_none)
    sbs.rescue_locked_headers(groups_blank, classifications_blank)

    assert "classification" not in para_none
    assert "classification" not in para_blank


# ---------------------------------------------------------------------------
# ensure_personal_data_first (and its nested is_document_title)
# ---------------------------------------------------------------------------

def _header(text: str, *, paragraph_index=None, has_index: bool = True) -> dict:
    node = {"text": text, "level": "H1", "children": []}
    if has_index:
        node["paragraph_index"] = paragraph_index
    return node


def test_ensure_personal_data_first_removes_document_title_and_keeps_next_as_first():
    hierarchy = [_header("Curriculum Vitae", paragraph_index=0), _header("Education", paragraph_index=1)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Education"]


def test_ensure_personal_data_first_inserts_synthetic_header_when_first_section_is_late():
    hierarchy = [_header("Education", paragraph_index=10)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["PERSONAL DATA", "Education"]
    assert result[0]["level"] == "H1"


def test_ensure_personal_data_first_no_insertion_when_first_index_is_early():
    hierarchy = [_header("Name And Title", paragraph_index=0)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Name And Title"]


def test_ensure_personal_data_first_chunked_mode_recognizes_personal_data_variant():
    hierarchy = [_header("Contact Information", has_index=False)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Contact Information"]


def test_ensure_personal_data_first_chunked_mode_strips_a_trailing_colon_before_matching():
    # CV headers routinely carry a trailing colon; the chunked-mode variant
    # match must normalise it away or "Contact Information:" would get a
    # synthetic PERSONAL DATA inserted in front of it.
    hierarchy = [_header("Contact Information:", has_index=False)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Contact Information:"]


def test_ensure_personal_data_first_chunked_mode_inserts_when_no_personal_data_variant():
    hierarchy = [_header("Education", has_index=False)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["PERSONAL DATA", "Education"]


def test_ensure_personal_data_first_document_title_exact_match_is_removed():
    hierarchy = [_header("CV", paragraph_index=0), _header("Publications", paragraph_index=1)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Publications"]


def test_ensure_personal_data_first_document_title_substring_match_is_removed():
    hierarchy = [_header("My Professional Portfolio", paragraph_index=0), _header("Grants", paragraph_index=1)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Grants"]


def test_ensure_personal_data_first_inserts_synthetic_header_when_every_node_was_a_title():
    # When every top-level node is filtered out as a document title, `filtered`
    # is empty, so Step 2 is skipped entirely and Step 3 takes its
    # filtered-is-empty branch (no first_real_idx to report at all).
    hierarchy = [_header("Curriculum Vitae", paragraph_index=0)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["PERSONAL DATA"]


def test_ensure_personal_data_first_document_title_trailing_colon_is_stripped_before_matching():
    # is_document_title normalizes with .rstrip(':') before comparing against
    # EXACT_TITLE_MATCHES -- "CV:" must still be recognized as a document
    # title, not merely the bare "cv". Without the rstrip, "CV:" survives
    # filtering, lands at paragraph_index 0 (< 3), and the early-return in
    # Step 2 short-circuits before PERSONAL DATA insertion is ever considered.
    hierarchy = [_header("CV:", paragraph_index=0), _header("Education", paragraph_index=1)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Education"]


def test_ensure_personal_data_first_document_title_trailing_period_is_stripped_before_matching():
    # is_document_title normalizes with .rstrip('.') as well as .rstrip(':')
    # before comparing against EXACT_TITLE_MATCHES -- "Resume." must still be
    # recognized as a document title, not merely the bare "resume". Without
    # the period strip, "Resume." survives filtering, lands at paragraph_index
    # 0 (< 3), and the early-return in Step 2 short-circuits before PERSONAL
    # DATA insertion is ever considered.
    hierarchy = [_header("Resume.", paragraph_index=0), _header("Education", paragraph_index=1)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Education"]


def test_ensure_personal_data_first_boundary_index_three_is_not_early():
    # The documented "early header" boundary is paragraph_index < 3 (i.e. 0,
    # 1, or 2 count as early). Index 3 is the first value that must NOT count
    # as early and therefore must still get a synthetic PERSONAL DATA header
    # inserted ahead of it. The existing 0-vs-10 tests do not pin this edge --
    # both are far from the boundary.
    hierarchy = [_header("Education", paragraph_index=3)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["PERSONAL DATA", "Education"]


def test_ensure_personal_data_first_boundary_index_two_is_early():
    # The other side of the same boundary: index 2 is still "early" and must
    # NOT get a synthetic header inserted.
    hierarchy = [_header("Education", paragraph_index=2)]

    result = sbs.ensure_personal_data_first(hierarchy)

    assert [h["text"] for h in result] == ["Education"]


# ---------------------------------------------------------------------------
# remove_duplicate_children
# ---------------------------------------------------------------------------

def test_remove_duplicate_children_removes_child_matching_parent_paragraph_index():
    headers = [
        {
            "text": "Education", "paragraph_index": 5,
            "children": [{"text": "Education", "paragraph_index": 5, "children": []}],
        },
    ]

    result = sbs.remove_duplicate_children(headers)

    assert result[0]["children"] == []


def test_remove_duplicate_children_only_compares_against_immediate_parent():
    # The grandchild shares its paragraph_index with the TOP-level ancestor, not
    # with its immediate parent (the middle child) -- so it must survive.
    headers = [
        {
            "text": "Education", "paragraph_index": 5,
            "children": [
                {
                    "text": "Subheading", "paragraph_index": 6,
                    "children": [{"text": "Education", "paragraph_index": 5, "children": []}],
                },
            ],
        },
    ]

    result = sbs.remove_duplicate_children(headers)

    grandchildren = result[0]["children"][0]["children"]
    assert [g["text"] for g in grandchildren] == ["Education"]


def test_remove_duplicate_children_same_text_under_different_parents_both_kept():
    headers = [
        {"text": "Parent A", "paragraph_index": 1, "children": [{"text": "Same Label", "paragraph_index": 2, "children": []}]},
        {"text": "Parent B", "paragraph_index": 3, "children": [{"text": "Same Label", "paragraph_index": 4, "children": []}]},
    ]

    result = sbs.remove_duplicate_children(headers)

    assert [c["text"] for c in result[0]["children"]] == ["Same Label"]
    assert [c["text"] for c in result[1]["children"]] == ["Same Label"]


def test_remove_duplicate_children_keeps_top_level_headers_lacking_paragraph_index():
    # dedupe_recursive's top-level call has parent_index=None; headers from
    # chunked extraction (the shape ensure_personal_data_first's has_index=False
    # tests exercise) carry no 'paragraph_index' key at all, so header.get(...)
    # is also None. The "parent_index is not None and" guard is what stops
    # None == None from reading as a false duplicate-of-parent match at the
    # top level -- without it every such header would be dropped.
    headers = [
        {"text": "Education", "children": []},
        {"text": "Grants", "children": []},
    ]

    result = sbs.remove_duplicate_children(headers)

    assert [h["text"] for h in result] == ["Education", "Grants"]
