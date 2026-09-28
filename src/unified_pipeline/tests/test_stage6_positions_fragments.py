"""Section D (positions) tab/pipe child rows (#476, PR1 of the accuracy wave).

`positions.py` is not an `entry_lines` call site -- it never called the
ten-times-duplicated newline splitter the issue is about. What it has instead
is a narrower loss of the same shape: a D1/D2/D3 entry whose text carries a
tab-joined, bullet-prefixed child appointment after its header row renders
only the header. The confirmed instance on the corpus farm is the -DAZFA D2
entry below (`stage_5_enrichment/-DAZFA_..._enriched.json`): stage 4 extracted
fields for the header only, and the two child promotions --
"Attending Physician | 07/2002 - 06/2003" and "Attending Physician & Assistant
Director | 06/2003 - 12/2006" -- rendered nowhere. `_tab_joined_child_
fragments` (positions.py) recovers them from `entry_fragments`, gated on a
literal tab in the raw text.

MNZ7IA and ZZLKMA render the same Borman source pre-split by stage 4 into
three separate D2 entries (no tab in any one entry's text) and are the model
for what -DAZFA's rendering should now look like; `test_negative_control_
already_split_entry_is_unchanged` pins that this fix does not touch them.

The corpus fixture is one production shape, so it is the regression case and
not the specification: groups (b) and (c) state the cardinality contract on
synthetic fixtures instead -- parent + N valid children renders N + 1 rows,
in source order, for N = 1, 2 and 3, and a child that is malformed, that is
ordinary tab-separated prose, or that repeats its parent adds no row at all.

Groups (d) and (e) cover the rest of the section's repair work rather than the
fragment scanner: that one record yields exactly one rendered row and one
increment of `entries_inserted`, and that the two passes which rewrite stage-4
output -- institution propagation and appointment merging -- refuse to act
without evidence rather than falling back on document adjacency.

Group (f) is the other side of (e). A record can reach the renderer with no
title, no employer and no dates -- stage 3b routes a stray line to a D code
and stage 4 finds no field in it, or the employer that would have been carried
onto it is refused at a source-structure boundary -- and rendering it put
three blank cells into a delivered CV. The rule the group pins is two-sided: a
record with nothing to put in any column renders no row at all, and a row
carrying even one populated cell is kept.

The farm's own instances of it are the three the `_normalized_positions`
docstring names. Two, NGFNYQ and SO2IVQ, are a fieldless D1 record leading its
code list, so no employer was ever carried onto it and it was already three
blank cells before this branch. The third, 6NGAYQ, is this group's fixture
shape and the boundary half of it: a fieldless record ~200 elements past the
one row it could have inherited from, under a different source heading. It is
the only record on the farm where the propagation boundary stop fires.

Both counts need a real `generate()` of every farm document to see. 6NGAYQ's
record is stored under code `I` and reaches the D1 list only because stage 6
reroutes it mid-render, so a harness that groups records by the code they were
stored with misses it and reports two blank rows and no boundary stop at all.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_stage6_positions_fragments.py -p no:cacheprovider
"""

import copy
import logging
import sys
from pathlib import Path

import pytest

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from unified_pipeline.core.render_check import entry_fragments, entry_lines  # noqa: E402
from unified_pipeline.stage6.sections.positions import (  # noqa: E402
    _child_position_records,
    _employers_match,
    _institution_from_raw_text,
    _position_row_cells,
    _tab_joined_child_fragments,
)
from unified_pipeline.stage_6_word_template import WCMTemplateGenerator  # noqa: E402


# The exact -DAZFA D2 entry text (stage_5_enrichment JSON), en-dash and bullet
# glyph verbatim.
BORMAN_D2_TEXT = (
    "Attending Physician, Department of Emergency Medicine | Lincoln Hospital, Bronx, NY | 07/2002 – 12/2006"
    "\t• Attending Physician | 07/2002 – 06/2003"
    "\t• Attending Physician & Assistant Director | 06/2003 – 12/2006"
)

# Farm-derived shapes with no tab-joined child (newline-only or pipe-only, or
# a stage-4-already-split header/child that carries a bullet but no tab).
NO_CHILD_CASES = [
    "Program\tOchsner Clinic Foundation\t2009 - 2018",
    "Role | Organization | 2013 - present",
    "Attending Physician | Lincoln Hospital, Bronx, NY | 07/2023 – Present",
    # MNZ7IA/ZZLKMA's own already-split child entry: a bullet immediately
    # followed by a pipe-date, but NO tab -- the gate this function's own
    # docstring names as the reason it must not fire here.
    "• Attending Physician | 07/2002 – 06/2003",
]

# The same four shapes as whole records: text, the fields stage 4 extracted
# for them, and the single row each must still render (test item 6 -- the
# helper returning no children does not by itself prove the record renders).
NO_CHILD_RENDER_CASES = [
    (NO_CHILD_CASES[0],
     {"title": "Program Director", "institution": "Quexley Clinic Foundation",
      "start_date": "2009", "end_date": "2018"},
     ["Program Director", "Quexley Clinic Foundation", "2009-2018"]),
    (NO_CHILD_CASES[1],
     {"title": "Consulting Physician", "institution": "Quexley Health Network",
      "start_date": "2013", "end_date": "present"},
     ["Consulting Physician", "Quexley Health Network", "2013-Present"]),
    (NO_CHILD_CASES[2],
     {"title": "Attending Physician", "institution": "Quexley General Hospital, Crab Hollow, ZQ",
      "start_date": "2023-07", "end_date": "Present"},
     ["Attending Physician", "Quexley General Hospital, Crab Hollow, ZQ", "07/23-Present"]),
    (NO_CHILD_CASES[3],
     {"title": "Attending Physician", "institution": None,
      "start_date": "2002-07", "end_date": "2003-06"},
     ["Attending Physician", "", "07/02-06/03"]),
]

# Synthetic employer and template anchors for the cardinality fixtures. No
# corpus name: these state the accepted input shape, and a production UID
# should not be the only specification of it (test item 8).
SYNTHETIC_EMPLOYER = "Quexley General Hospital, Crab Hollow, ZQ"
TABLE_ANCHORS = {
    "D1": "Academic Appointments",
    "D2": "Hospital Appointments",
    "D3": "Other Professional Positions",
}
GENERIC_TABLE_ANCHOR = "PROFESSIONAL POSITIONS"


def _positions_generator(anchor="Hospital Appointments"):
    """A generator whose document holds just one section anchor and an empty
    three-column table, ready for `_fill_positions`."""
    gen = WCMTemplateGenerator(verbose=False)
    gen.doc = Document()
    gen.doc.add_paragraph("D. POSITIONS")
    gen.doc.add_paragraph(anchor)
    table = gen.doc.add_table(rows=1, cols=3)
    for i, header in enumerate(["Title", "Institution/Location", "Dates"]):
        table.rows[0].cells[i].text = header
    return gen, table


def _rendered_rows(table):
    return [[c.text for c in r.cells] for r in table.rows[1:]]


def _render_positions(entries, code="D2"):
    gen, table = _positions_generator(TABLE_ANCHORS[code])
    gen._fill_positions({code: entries})
    return _rendered_rows(table)


def _position_entry(idx, title, institution, hierarchy, start="", end=""):
    """A minimal D2 record: document position, title, employer and the source
    heading path propagation reads."""
    return {
        "element_idx_start": idx,
        "taxonomy_code": "D2",
        "hierarchy": hierarchy,
        "extracted_fields": {
            "title": title,
            "institution": institution,
            "start_date": start,
            "end_date": end,
        },
    }


def _institution_of(entry):
    return (entry.get("extracted_fields") or {}).get("institution") or ""


def _borman_entry():
    return {
        "text": BORMAN_D2_TEXT,
        "taxonomy_code": "D2",
        "extracted_fields": {
            "title": "Attending Physician, Department of Emergency Medicine",
            "institution": "Lincoln Hospital, Bronx, NY",
            "start_date": "2002-07",
            "end_date": "2006-12",
        },
    }


def _tab_joined_entry(children, title="Attending Physician",
                      institution=SYNTHETIC_EMPLOYER, start="2002-07",
                      end="2006-12", header_dates="07/2002 – 12/2006",
                      code="D2", idx=1):
    """A synthetic parent record whose text carries `children` -- each a
    (title, start, end) triple -- as bullet-prefixed, tab-joined fragments in
    the -DAZFA shape."""
    text = f"{title} | {institution} | {header_dates}" + "".join(
        f"\t• {child_title} | {child_start} – {child_end}"
        for child_title, child_start, child_end in children)
    return {
        "text": text,
        "taxonomy_code": code,
        "element_idx_start": idx,
        "extracted_fields": {
            "title": title,
            "institution": institution,
            "start_date": start,
            "end_date": end,
        },
    }


def _synthetic_row(title, dates, institution=SYNTHETIC_EMPLOYER):
    return [title, institution, dates]


# --- (a) no content the old split saw is lost or fabricated -----------------

def test_tab_joined_child_fragments_recovers_exactly_the_two_children():
    children = _tab_joined_child_fragments(BORMAN_D2_TEXT)
    assert children == [
        ("Attending Physician", ("07/2002", "06/2003")),
        ("Attending Physician & Assistant Director", ("06/2003", "12/2006")),
    ]
    # Every entry_lines "line" -- here there is exactly one, the whole blob,
    # since the text carries no literal newline -- is still fully accounted
    # for: nothing in the recovered children is absent from entry_fragments.
    assert entry_lines(BORMAN_D2_TEXT) == [BORMAN_D2_TEXT]
    frags = entry_fragments(BORMAN_D2_TEXT)
    for title, (start, end) in children:
        assert any(title in f for f in frags)
        assert any(start in f for f in frags)
        assert any(end in f for f in frags)


@pytest.mark.parametrize("count", [1, 2, 3])
def test_tab_joined_child_fragments_recovers_every_child_at_any_cardinality(count):
    """Pins the scanner's cardinality contract, which the two-child corpus
    fixture cannot state on its own: N bullet-prefixed children in the text
    come back as N children, whatever N is.

    A `>1 child` guard, or a loop that consumed one fragment too many after a
    match, would still satisfy the -DAZFA regression while dropping data.
    """
    children = [("Attending Physician", "07/2002", "06/2003"),
                ("Senior Attending Physician", "06/2003", "12/2004"),
                ("Interim Director", "01/2005", "12/2006")][:count]
    entry = _tab_joined_entry(children)

    assert _tab_joined_child_fragments(entry["text"]) == [
        (title, (start, end)) for title, start, end in children]


def test_no_tab_no_children_even_with_a_leading_bullet():
    """The gate is a literal tab, not just bullet-then-date-range -- this is
    what keeps MNZ7IA/ZZLKMA's already-split entries (a bullet-prefixed
    title immediately followed by its own pipe-date, no tab) untouched."""
    for case in NO_CHILD_CASES:
        assert _tab_joined_child_fragments(case) == [], f"unexpected children for {case!r}"


@pytest.mark.parametrize("tail", [
    "\tSupervised residents in the emergency department",   # ordinary prose
    "\t• Directed the residency programme year-round",      # bullet, no dates
    "\t• Interim Director | Summer 2005",                   # bullet, unparseable dates
    "\t• Interim Director 01/2005 – 12/2006",               # no pipe: one fragment
    "\t• | 01/2005 – 12/2006",                              # dates, no title
    "\t01/2005 – 12/2006",                                  # dates, no bullet
])
def test_a_tab_fragment_that_is_not_a_child_appointment_yields_nothing(tail):
    """Pins the boundary the tab gate draws (test items 4 and 5).

    The production gate is tab presence, so everything after a tab reaches
    the scanner. A fragment only becomes a child when it is bullet-prefixed,
    carries a title, and is immediately followed by a fragment that is
    wholly a date range: arbitrary tab-separated descriptions, and children
    missing a title or a parseable date, are dropped whole rather than
    recovered half-populated.
    """
    text = f"Attending Physician | {SYNTHETIC_EMPLOYER} | 07/2002 – 12/2006{tail}"
    assert _tab_joined_child_fragments(text) == []


def test_a_malformed_child_does_not_take_its_valid_siblings_with_it():
    """A child with an unparseable date range is skipped where it stands and
    the scan continues: the well-formed child after it is still recovered.

    Pinned because the scanner advances by one fragment on a non-match and by
    two on a match; an unconditional two-step would swallow the next child.
    """
    text = (f"Attending Physician | {SYNTHETIC_EMPLOYER} | 07/2002 – 12/2006"
            "\t• Interim Director | Summer 2005"
            "\t• Senior Attending Physician | 06/2003 – 12/2004")

    assert _tab_joined_child_fragments(text) == [
        ("Senior Attending Physician", ("06/2003", "12/2004")),
    ]


# --- (b) positive control: fails on dev today --------------------------------

def test_positive_control_borman_entry_gains_both_child_rows():
    """The -DAZFA D2 entry renders as three complete rows, asserted cell for
    cell: a row count plus a title check would pass on a regression that put
    the wrong institution or dates in a row, and two legitimate appointments
    may share a title, so per-row uniqueness is not the contract either.
    """
    assert _render_positions([_borman_entry()]) == [
        ["Attending Physician, Department of Emergency Medicine",
         "Lincoln Hospital, Bronx, NY", "07/02-12/06"],
        ["Attending Physician", "Lincoln Hospital, Bronx, NY", "07/02-06/03"],
        ["Attending Physician & Assistant Director",
         "Lincoln Hospital, Bronx, NY", "06/03-12/06"],
    ]


# --- (c) the cardinality contract, on synthetic fixtures ---------------------

def test_parent_with_one_child_renders_two_rows():
    """Parent + 1 valid child = 2 rows. The corpus fixture has two children,
    so an off-by-one that dropped the first child would still pass it."""
    entry = _tab_joined_entry([("Senior Attending Physician", "06/2003", "12/2006")])

    assert _render_positions([entry]) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
        _synthetic_row("Senior Attending Physician", "06/03-12/06"),
    ]


def test_parent_with_three_children_renders_four_rows():
    """Parent + 3 valid children = 4 rows, each carrying its own title and
    dates and the parent's employer. A `>1 child` guard would pass the
    two-child corpus fixture and drop the third child here."""
    entry = _tab_joined_entry([
        ("Attending Physician", "07/2002", "06/2003"),
        ("Senior Attending Physician", "06/2003", "12/2004"),
        ("Interim Director", "01/2005", "12/2006"),
    ])

    assert _render_positions([entry]) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
        _synthetic_row("Attending Physician", "07/02-06/03"),
        _synthetic_row("Senior Attending Physician", "06/03-12/04"),
        _synthetic_row("Interim Director", "01/05-12/06"),
    ]


def test_children_render_in_source_order_not_date_order():
    """Children follow their parent in the order the source text lists them.

    The section sorts records reverse-chronologically, and children are
    recovered after that sort, so a CV that lists its promotions newest-first
    keeps that order instead of having it reversed under the parent.
    """
    entry = _tab_joined_entry([
        ("Interim Director", "01/2005", "12/2006"),
        ("Senior Attending Physician", "06/2003", "12/2004"),
        ("Attending Physician", "07/2002", "06/2003"),
    ])

    assert [row[0] for row in _render_positions([entry])] == [
        "Attending Physician",           # the parent header row
        "Interim Director",
        "Senior Attending Physician",
        "Attending Physician",
    ]


@pytest.mark.parametrize("tail", [
    "\tSupervised residents in the emergency department",
    "\t• Directed the residency programme year-round",
    "\t• Interim Director | Summer 2005",
    "\t• | 01/2005 – 12/2006",
])
def test_no_valid_child_renders_the_parent_row_alone(tail):
    """The render-level half of the boundary: a tab fragment that is not a
    child appointment must not become a phantom position row.

    Asserted through `_fill_positions` rather than the scanner, because the
    row count is decided in normalization and a phantom row would be visible
    only here.
    """
    entry = _tab_joined_entry([])
    entry["text"] += tail

    assert _render_positions([entry]) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
    ]


def test_a_child_that_repeats_its_parent_is_not_rendered_twice():
    """A child whose title and formatted dates are the parent's own is
    dropped: stage 4 sometimes promotes a bullet-prefixed fragment to its own
    record, and that record arrives here as its own parent. This is the one
    documented exception to parent + N children = N + 1 rows."""
    entry = _tab_joined_entry([("Attending Physician", "07/2002", "12/2006")])

    assert _render_positions([entry]) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
    ]


def test_parent_without_an_institution_still_renders_its_child():
    """A child copies the parent's employer fields verbatim, so a parent with
    no employer yields a child with none -- an empty Institution cell on both
    rows, never a fabricated one, and never a child dropped for lack of it."""
    entry = _tab_joined_entry(
        [("Senior Attending Physician", "06/2003", "12/2006")], institution="")

    assert _render_positions([entry]) == [
        ["Attending Physician", "", "07/02-12/06"],
        ["Senior Attending Physician", "", "06/03-12/06"],
    ]


def test_multiple_parents_keep_their_own_children():
    """Two tab-joined parents in one table: each child follows its own parent
    and carries that parent's employer. A recovery pass that collected
    children globally, or attached them to the wrong record, would show up
    here as a row under the wrong employer."""
    recent = _tab_joined_entry(
        [("Senior Attending Physician", "06/2003", "12/2006")], idx=1)
    earlier = _tab_joined_entry(
        [("Staff Physician", "01/1997", "12/1998")],
        title="Resident Physician", institution="Norvale University Hospital, Crab Hollow, ZQ",
        start="1995-07", end="1998-12", header_dates="07/1995 – 12/1998", idx=2)

    assert _render_positions([recent, earlier]) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
        _synthetic_row("Senior Attending Physician", "06/03-12/06"),
        ["Resident Physician", "Norvale University Hospital, Crab Hollow, ZQ", "07/95-12/98"],
        ["Staff Physician", "Norvale University Hospital, Crab Hollow, ZQ", "01/97-12/98"],
    ]


@pytest.mark.parametrize("code", ["D1", "D2", "D3"])
def test_a_recovered_child_renders_in_its_own_code_table(code):
    """The same record shape through each of the three template tables.

    `_fill_positions` looks each table up by its own anchor paragraph and
    normalizes each code list separately, so child recovery has to be proved
    on all three routes, not only on the D2 one the corpus fixture uses.
    """
    entry = _tab_joined_entry(
        [("Senior Attending Physician", "06/2003", "12/2006")], code=code)

    assert _render_positions([entry], code=code) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
        _synthetic_row("Senior Attending Physician", "06/03-12/06"),
    ]


def test_a_recovered_child_renders_into_the_generic_positions_table():
    """The fourth route: a template with none of the three subsection anchors
    falls back to one combined PROFESSIONAL POSITIONS table, which normalizes
    D1 + D2 + D3 together. Recovery has to survive that path too."""
    gen, table = _positions_generator(GENERIC_TABLE_ANCHOR)
    entry = _tab_joined_entry([("Senior Attending Physician", "06/2003", "12/2006")])

    gen._fill_positions({"D2": [entry]})

    assert _rendered_rows(table) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
        _synthetic_row("Senior Attending Physician", "06/03-12/06"),
    ]


def test_a_d3_child_inherits_the_organization_field():
    """D3 records carry their employer as `organization`, not `institution`.
    A child copies both, so the D3 Institution cell is the parent's whatever
    field name it arrived under."""
    organization = "Quexley Medical Society, Crab Hollow, ZQ"
    entry = _tab_joined_entry(
        [("Board Chair", "01/2004", "12/2006")], title="Board Member",
        institution=organization, code="D3")
    entry["extracted_fields"] = {
        "title": "Board Member",
        "organization": organization,
        "start_date": "2002-07",
        "end_date": "2006-12",
    }

    assert _render_positions([entry], code="D3") == [
        ["Board Member", organization, "07/02-12/06"],
        ["Board Chair", organization, "01/04-12/06"],
    ]


def test_a_child_keeps_the_parent_location_field():
    """Pins the parent's `location` field onto the recovered child's row.

    `_get_institution_location` falls back to `extracted_fields['location']`
    when stage-5b enrichment names no city, and the child record copied
    institution, organization and department but not that field. The child row
    then named the employer without its city and state while the parent row
    directly above it kept them -- two rows of one appointment disagreeing
    about where it was.
    """
    entry = _tab_joined_entry([("Senior Attending Physician", "06/2003", "12/2004")],
                              institution="Norvale University Hospital")
    entry["extracted_fields"]["location"] = "Crab Hollow, ZQ"

    assert _render_positions([entry]) == [
        ["Attending Physician", "Norvale University Hospital, Crab Hollow, ZQ", "07/02-12/06"],
        ["Senior Attending Physician", "Norvale University Hospital, Crab Hollow, ZQ", "06/03-12/04"],
    ]


# The employer-field shapes a stage-4 position record arrives in. The
# Institution cell reads four fields and any subset of them can be populated,
# so the parent/child equality is asserted across the subsets rather than
# against whichever one a single fixture happens to use.
EMPLOYER_FIELD_SHAPES = [
    {"institution": "Norvale University Hospital"},
    {"institution": "Norvale University Hospital", "location": "Crab Hollow, ZQ"},
    {"organization": "Quexley Medical Society", "location": "Crab Hollow, ZQ"},
    {"institution": "Norvale University Hospital",
     "department": "Department of Emergency Medicine",
     "location": "Crab Hollow, ZQ"},
    {"location": "Crab Hollow, ZQ"},
]


@pytest.mark.parametrize("employer_fields", EMPLOYER_FIELD_SHAPES)
def test_a_child_institution_cell_is_the_parents_whichever_field_carries_it(employer_fields):
    """`_child_position_records` claims a child's Institution cell comes out
    identical to its parent's. That holds only while the record copies every
    field the cell reads, so the claim is asserted against the renderer itself
    -- the two cells, run lists and all -- across the employer-field subsets a
    stage-4 record arrives in, rather than against a list of field names that
    can silently fall behind the cell.
    """
    parent = _tab_joined_entry([("Senior Attending Physician", "06/2003", "12/2004")])
    parent["extracted_fields"] = {"title": "Attending Physician",
                                  "start_date": "2002-07", "end_date": "2006-12",
                                  **employer_fields}

    children = _child_position_records(parent)
    parent_cell = _position_row_cells(parent)[1]

    assert len(children) == 1
    assert any(text.strip() for text, _, _ in parent_cell)
    assert _position_row_cells(children[0])[1] == parent_cell


def test_recovery_does_not_mutate_the_supplied_entries():
    """Stage-4 records are pipeline data later stages read, so recovering a
    child must not write into the record it came from (test item 9).

    Scoped to the recovery path, which is what these entries exercise: the
    two repair passes deliberately write into the records they repair --
    institution propagation fills a blank employer in place, and the merge
    copies dates onto the row it keeps, both asserted in group (e). This
    fixture gives every record its own employer, title and dates so neither
    pass has anything to do and any difference is the recovery path's.
    """
    entries = [
        _tab_joined_entry([("Senior Attending Physician", "06/2003", "12/2006")], idx=1),
        _position_entry(2, "Chief Resident", "Norvale University Hospital, Crab Hollow, ZQ",
                        ["Hospital Appointments"], "1999-07", "2001-06"),
    ]
    before = copy.deepcopy(entries)

    rows = _render_positions(entries, code="D2")

    assert len(rows) == 3
    assert entries == before


# --- (c) negative control: a single-part entry is unchanged ------------------

def test_negative_control_single_part_entry_unchanged():
    entry = {
        "text": "Attending Physician | Weill Cornell Medicine, New York, NY | 07/2020 – Present",
        "taxonomy_code": "D2",
        "extracted_fields": {
            "title": "Attending Physician",
            "institution": "Weill Cornell Medicine, New York, NY",
            "start_date": "2020-07",
            "end_date": "Present",
        },
    }
    rows = _render_positions([entry])
    assert len(rows) == 1
    assert rows[0][0] == "Attending Physician"
    assert rows[0][2] == "07/20-Present"


def test_negative_control_already_split_entry_is_unchanged():
    """MNZ7IA/ZZLKMA's shape: the child was already promoted to its own D2
    entry by stage 4 (institution None, own start/end dates, bullet still in
    the raw text but no tab). Must render as exactly one row, not two."""
    entry = {
        "text": "• Attending Physician | 07/2002 – 06/2003",
        "taxonomy_code": "D2",
        "extracted_fields": {
            "title": "Attending Physician",
            "institution": None,
            "start_date": "2002-07",
            "end_date": "2003-06",
        },
    }
    rows = _render_positions([entry])
    assert len(rows) == 1
    assert rows[0][0] == "Attending Physician"
    assert rows[0][2] == "07/02-06/03"


@pytest.mark.parametrize("text,fields,expected", NO_CHILD_RENDER_CASES)
def test_an_entry_with_no_tab_joined_child_renders_exactly_one_row(text, fields, expected):
    """The render-level complement to `test_no_tab_no_children_even_with_a_
    leading_bullet`: the scanner returning no children proves the helper is
    quiet, not that the record still renders. Each of these shapes must come
    out as the single row it was before the recovery path existed.
    """
    entry = {"text": text, "taxonomy_code": "D2", "extracted_fields": dict(fields)}

    assert _render_positions([entry]) == [expected]


def test_an_already_split_child_beside_a_tab_joined_parent_is_not_duplicated():
    """Both shapes in one table (test item 7): a stage-4 record that is already
    its own row, and a parent whose child still has to be recovered from the
    text. The recovery path must add the parent's child once and leave the
    already-split record alone -- not duplicate it, not give it the other
    employer, not drop it.
    """
    parent = _tab_joined_entry([("Senior Attending Physician", "06/2003", "12/2006")], idx=1)
    already_split = _position_entry(
        2, "Chief Resident", "Norvale University Hospital, Crab Hollow, ZQ",
        ["Hospital Appointments"], "1999-07", "2001-06")

    assert _render_positions([parent, already_split]) == [
        _synthetic_row("Attending Physician", "07/02-12/06"),
        _synthetic_row("Senior Attending Physician", "06/03-12/06"),
        ["Chief Resident", "Norvale University Hospital, Crab Hollow, ZQ", "07/99-06/01"],
    ]


# --- (d) the rendering layer owns the row count ------------------------------

def test_entries_inserted_counts_one_per_rendered_row():
    """Pins the row-count invariant: `entries_inserted` is incremented exactly
    once per physical row, child rows included.

    A parent with two recovered children must leave the counter at 3. One
    (children not counted) or five (children counted at two levels) both mean
    the statistic no longer describes the document.
    """
    gen, table = _positions_generator()
    gen._fill_positions({"D2": [_borman_entry()]})
    assert len(table.rows) - 1 == 3
    assert gen.stats["entries_inserted"] == 3


@pytest.mark.parametrize("count", [1, 2, 3])
def test_entries_inserted_counts_child_rows_at_every_cardinality(count):
    """The same invariant across N: a parent with N valid children leaves the
    counter at N + 1, and the table holds exactly that many rows. Pinning one
    cardinality cannot tell a per-row increment from a per-record one."""
    children = [("Attending Physician", "07/2002", "06/2003"),
                ("Senior Attending Physician", "06/2003", "12/2004"),
                ("Interim Director", "01/2005", "12/2006")][:count]
    gen, table = _positions_generator()

    gen._fill_positions({"D2": [_tab_joined_entry(children)]})

    assert len(table.rows) - 1 == count + 1
    assert gen.stats["entries_inserted"] == count + 1


# --- (e) the repair passes fail closed ---------------------------------------

def test_institution_propagation_stops_at_unrelated_entries():
    """Pins that a carried-forward employer stops at a source-structure
    boundary instead of running to the end of the code list.

    Document order alone used to keep an employer in scope for every later
    record, so a record under a different source heading inherited it. The
    record under the other heading must stay empty, and so must the record
    after it -- the carry is dropped at the boundary, not resumed across it.
    """
    parent = _position_entry(1, "Chief of Service", "Lincoln Hospital",
                             ["Hospital Appointments"], "2001", "2004")
    sibling = _position_entry(2, "Attending Physician", "", ["Hospital Appointments"])
    unrelated = _position_entry(3, "Course Director", "", ["Teaching"])
    after = _position_entry(4, "Preceptor", "", ["Hospital Appointments"])

    WCMTemplateGenerator._propagate_institution_to_subentries(
        [parent, sibling, unrelated, after])

    assert _institution_of(sibling) == "Lincoln Hospital"
    assert _institution_of(unrelated) == ""
    assert _institution_of(after) == ""


@pytest.mark.parametrize("malformed_enrichment", [["a"], "abc", 7])
def test_propagation_survives_a_non_mapping_enrichment(malformed_enrichment, caplog):
    """#743: stage 5b is an LLM output, and `institution_enrichment` can come
    back as a list, a string, or another non-mapping instead of a dict. The
    parent-to-subentry carry used to call `dict(last_enrichment)` on whatever
    the parent held, which raises on a non-mapping (`ValueError` for a
    single-character list, `TypeError` for an int) and fails the
    Positions section, which `_render_section` then sends to the Appendix. It must instead treat the malformed value like no enrichment and
    keep propagating the institution name."""
    parent = _position_entry(1, "Chief of Service", "Lincoln Hospital",
                             ["Hospital Appointments"], "2001", "2004")
    parent["institution_enrichment"] = malformed_enrichment
    sub_entry = _position_entry(2, "Attending Physician", "",
                                ["Hospital Appointments"])

    with caplog.at_level(logging.WARNING):
        WCMTemplateGenerator._propagate_institution_to_subentries(
            [parent, sub_entry])

    assert _institution_of(sub_entry) == "Lincoln Hospital"
    assert "institution_enrichment" not in sub_entry
    assert "not a mapping" in caplog.text


def test_raw_text_employer_candidate_must_name_an_employer():
    """Pins the last-resort employer scan against City/State-shaped text.

    It used to take any fragment containing a "City, ST" pattern wherever it
    sat, so a bare location, a "Surname, Forename", or a whole date-prefixed
    record line became the Institution cell of a record that had no employer.
    A candidate must name something in front of a location tail, and must
    carry no year.
    """
    employer = "Attending Physician\tLincoln Hospital, Bronx, NY\t07/2002"
    assert _institution_from_raw_text(employer) == "Lincoln Hospital, Bronx, NY"

    for text in (
        "Attending Physician\tBronx, NY",                       # bare location
        "Reviewer\tMarblegate, Quenby",                         # a person
        "Aug 2019-Dec 2023, Associate Director, Institute of "
        "Speculative Metrics, Norvale University, Crab Hollow, ZQ\n"
        "Jan 2024-Present, Vice Chair",                         # record lines
    ):
        assert _institution_from_raw_text(text) == "", f"accepted {text!r}"


def test_merge_needs_employer_evidence_not_adjacency():
    """Pins that the merge refuses a pair whose employers disagree.

    A title-only row next to a bare dates row used to take that row's dates
    and delete it whatever the two said about their employers, which invents
    an appointment the CV never claimed. Two employers with no word in common
    must leave both rows standing and the titled row date-less.
    """
    dated = _position_entry(1, "", "Quexley College", ["Hospital Appointments"],
                            "2019-09", "2021-06")
    titled = _position_entry(2, "Program Director", "Norvale University Medical College",
                             ["Hospital Appointments"])

    merged = WCMTemplateGenerator._merge_grouped_appointments([dated, titled])

    assert len(merged) == 2
    # the titled row keeps its own employer, and gains no dates from the other
    assert _institution_of(titled) == "Norvale University Medical College"
    assert not (titled["extracted_fields"].get("start_date")
                or titled["extracted_fields"].get("end_date"))


def test_merge_refuses_a_pair_that_names_no_employer_at_all():
    """Pins the other half of the same rule: a missing employer is not a match.

    Two records that neither name an employer nor inherited one from a parent
    row used to merge on adjacency alone, because the header test read two
    unknown employers as equal.
    """
    dated = _position_entry(1, "", "", ["Hospital Appointments"], "2002-07", "2006-12")
    titled = _position_entry(2, "Attending Physician", "", ["Hospital Appointments"])

    merged = WCMTemplateGenerator._merge_grouped_appointments([dated, titled])

    assert len(merged) == 2
    assert not (titled["extracted_fields"].get("start_date")
                or titled["extracted_fields"].get("end_date"))


# Employer strings a source CV produces that carry no word at all: an empty
# field, whitespace, and the dash or ampersand a table cell is left as when the
# CV had nothing to put there. Field extraction passes these through as an
# institution, so they arrive at the merge looking like a named employer.
NAMELESS_EMPLOYERS = ["", "   ", "-", "–", "—", "&", ", ,"]


@pytest.mark.parametrize("nameless", NAMELESS_EMPLOYERS)
def test_an_employer_with_no_word_in_it_matches_nothing(nameless):
    """Pins `_employers_match`'s first guard: an employer that names nobody is
    no evidence, in either argument position and against itself.

    The comparison is on word sets, so a value with no word in it produces an
    empty set, and every empty set is a subset of every other. Without the
    guard the subset test reads "unknown" as "equal to anything" -- the exact
    reading review item 2 asked to be replaced with a fail-closed one.
    """
    assert not _employers_match(nameless, "Lincoln Hospital")
    assert not _employers_match("Lincoln Hospital", nameless)
    assert not _employers_match(nameless, nameless)


def test_two_named_employers_still_match_by_word_subset():
    """The guard's positive control: refusing the nameless ones must not cost
    the sub-unit match the comparison exists for."""
    assert _employers_match("Lincoln Hospital",
                            "Lincoln Hospital, Department of Emergency Medicine")


def test_a_placeholder_employer_does_not_license_a_merge():
    """The same guard at the merge decision rather than in the helper.

    A source table that leaves its employer cell as a dash hands field
    extraction an institution with no word in it. That is a missing employer
    wearing a name, and it is no evidence that the row beside it is the same
    appointment: both rows must survive carrying what they came with. Read as
    a match, it copies the dates off one row onto the other and deletes it --
    an appointment the CV never claimed, which is the loss review item 2 is
    about.
    """
    dated = _position_entry(1, "", "Lincoln Hospital", ["Hospital Appointments"],
                            "2002-07", "2006-12")
    titled = _position_entry(2, "Attending Physician", "—", ["Hospital Appointments"])

    assert _render_positions([dated, titled], code="D2") == [
        ["", "Lincoln Hospital", "07/02-12/06"],
        ["Attending Physician", "—", ""],
    ]


def _rule_1_pair_resting_on_inheritance():
    """A title-only row that inherited employer X, then a bare-dates row
    naming employer Y: Rule 1 (adjacent pair) merges them on the inheritance
    alone, because the CV cannot say whether Y is a unit inside X."""
    return [
        _position_entry(1, "Chief of Service", "Quexley General Hospital",
                        ["Hospital Appointments"], "1998-01", "2001-12"),
        _position_entry(2, "Attending Physician", "", ["Hospital Appointments"]),
        _position_entry(3, "", "Norvale University Medical College",
                        ["Hospital Appointments"], "2002-07", "2006-12"),
    ]


def _rule_2_pair_resting_on_inheritance():
    """The same thin evidence through Rule 2 (header + children): the header
    named no employer and inherited one, and the role beneath it names a
    different employer."""
    return [
        _position_entry(1, "Chief of Service", "Quexley General Hospital",
                        ["Hospital Appointments"], "1998-01", "2001-12"),
        _position_entry(2, "", "", ["Hospital Appointments"], "2002-07", "2006-12"),
        _position_entry(3, "Attending Physician", "Norvale University Medical College",
                        ["Hospital Appointments"]),
    ]


def _rule_1_pair_with_matching_employers():
    return [
        _position_entry(1, "Attending Physician",
                        "Lincoln Hospital, Department of Emergency Medicine",
                        ["Hospital Appointments"]),
        _position_entry(2, "", "Lincoln Hospital", ["Hospital Appointments"],
                        "2002-07", "2006-12"),
    ]


def _rule_2_pair_with_matching_employers():
    return [
        _position_entry(1, "", "Lincoln Hospital", ["Hospital Appointments"],
                        "2002-07", "2006-12"),
        _position_entry(2, "Attending Physician", "", ["Hospital Appointments"]),
    ]


@pytest.mark.parametrize("build_entries", [_rule_1_pair_resting_on_inheritance,
                                           _rule_2_pair_resting_on_inheritance])
def test_a_merge_resting_on_inheritance_alone_is_counted(build_entries, caplog):
    """Pins the one merge rule that still fires without a matching employer,
    and the tally that makes it visible, at both merge sites.

    A record inherits its employer from the nearest preceding row that named
    one, which is not always the row it is later compared against, so a row
    carrying employer X can merge with one naming employer Y. The merge
    happens: refusing it on the name mismatch was tried and put back the
    title-less dated row that the #156 fixture -- a ward, "Medical/Surgical
    Unit", inside an inherited "New York Presbyterian Hospital" -- exists to
    keep out, and the stage-4 fields cannot tell that ward from a different
    employer. So the run reports how many of its merges rest on this.
    """
    entries = build_entries()
    WCMTemplateGenerator._propagate_institution_to_subentries(entries)

    with caplog.at_level("DEBUG", logger="unified_pipeline.stage6.sections.positions"):
        merged = WCMTemplateGenerator._merge_grouped_appointments(entries, verbose=True)

    assert len(merged) == len(entries) - 1
    assert "1 merged row(s) matched no employer name" in caplog.text


def test_the_unmatched_employer_tally_counts_every_such_merge(caplog):
    """The tally is a count, not a flag.

    A header that inherited one employer over two roles that both name another
    rests on the thin rule twice, and the run has to say two -- every other
    fixture in this group produces exactly one, which a tally hard-wired to
    that number, or a boolean dressed up as one, would satisfy just as well.
    """
    entries = [
        _position_entry(1, "Chief of Service", "Quexley General Hospital",
                        ["Hospital Appointments"], "1998-01", "2001-12"),
        _position_entry(2, "", "", ["Hospital Appointments"], "2002-07", "2006-12"),
        _position_entry(3, "Attending Physician", "Norvale University Medical College",
                        ["Hospital Appointments"]),
        _position_entry(4, "Associate Attending Physician",
                        "Norvale University Medical College", ["Hospital Appointments"]),
    ]
    WCMTemplateGenerator._propagate_institution_to_subentries(entries)

    with caplog.at_level("DEBUG", logger="unified_pipeline.stage6.sections.positions"):
        merged = WCMTemplateGenerator._merge_grouped_appointments(entries, verbose=True)

    assert len(merged) == 3
    assert "2 merged row(s) matched no employer name" in caplog.text


@pytest.mark.parametrize("build_entries", [_rule_1_pair_with_matching_employers,
                                           _rule_2_pair_with_matching_employers])
def test_a_merge_whose_employers_match_is_not_counted(build_entries, caplog):
    """The tally's other half, at both merge sites: a row merging with an
    employer it matches -- a sub-unit of it, or one it inherited from that
    very row -- is not reported, so the count measures the thin-evidence
    merges only and not merging in general."""
    entries = build_entries()
    WCMTemplateGenerator._propagate_institution_to_subentries(entries)

    with caplog.at_level("DEBUG", logger="unified_pipeline.stage6.sections.positions"):
        merged = WCMTemplateGenerator._merge_grouped_appointments(entries, verbose=True)

    assert len(merged) == len(entries) - 1
    assert "matched no employer name" not in caplog.text


def test_merge_still_joins_a_sub_unit_of_the_same_employer():
    """The other half of the same rule: a record naming a sub-unit of the
    header's employer is still one appointment with it, so requiring evidence
    does not cost the merge the case it was written for.
    """
    dated = _position_entry(1, "", "Lincoln Hospital", ["Hospital Appointments"],
                            "2002-07", "2006-12")
    titled = _position_entry(
        2, "Attending Physician",
        "Lincoln Hospital, Department of Emergency Medicine",
        ["Hospital Appointments"])

    merged = WCMTemplateGenerator._merge_grouped_appointments([dated, titled])

    assert len(merged) == 1
    assert merged[0]["extracted_fields"]["title"] == "Attending Physician"
    assert merged[0]["extracted_fields"]["start_date"] == "2002-07"
    assert merged[0]["extracted_fields"]["end_date"] == "2006-12"


def test_merge_still_joins_a_sub_position_that_inherited_its_employer():
    """The inheritance case the employer match must not cost: a title-only
    sub-position that inherited the header's own employer still merges with
    it, so a CV that indents its roles under one employer heading is
    unaffected by the stricter rule.
    """
    header = _position_entry(1, "", "Lincoln Hospital", ["Hospital Appointments"],
                             "2002-07", "2006-12")
    sub_position = _position_entry(2, "Attending Physician", "", ["Hospital Appointments"])
    entries = [header, sub_position]
    WCMTemplateGenerator._propagate_institution_to_subentries(entries)

    merged = WCMTemplateGenerator._merge_grouped_appointments(entries)

    assert len(merged) == 1
    assert merged[0]["extracted_fields"]["title"] == "Attending Physician"
    assert merged[0]["extracted_fields"]["start_date"] == "2002-07"
    assert merged[0]["extracted_fields"]["end_date"] == "2006-12"


# --- (f) no all-blank row, and no populated row dropped ----------------------

def _fieldless_record(idx, text, hierarchy, code="D1"):
    """A record stage 3b routed to a D code and stage 4 found no field in:
    no title, no employer, no dates. The shape that reaches the renderer with
    nothing to put in any of the three columns."""
    return {
        "element_idx_start": idx,
        "taxonomy_code": code,
        "hierarchy": hierarchy,
        "text": text,
        "extracted_fields": {},
    }


APPOINTMENTS_HEADING = ["ACADEMIC APPOINTMENTS AND OTHER WORK EXPERIENCE"]
LEADERSHIP_HEADING = ["ADMINISTRATIVE AND ACADEMIC LEADERSHIP"]


def _appointment_and_distant_prose():
    """The 6NGAYQ D1 shape, with synthetic names and text: one real
    appointment under the appointments heading, and a fieldless stray record
    ~200 elements later under the leadership heading that reaches the D1 list
    anyway.

    6NGAYQ's element-217 record is the only one on the whole farm where the
    propagation boundary stop fires, so the two passes this fixture runs
    through are the two that disagreed about a real corpus record. What is
    reproduced from it: the two element indices, the leadership heading, and a
    record with no title, employer or dates arriving in the D1 list from
    somewhere other than the appointments heading. What is not: the employer,
    the titles and the sentence are invented, the appointments heading is
    shortened, and the farm record arrives as D1 by stage 6's own mismatch
    correction (stage 3b stored it as `I`) rather than out of stage 3b."""
    appointment = _position_entry(
        19, "Faculty member of the residency program",
        "Northgate Hospitals Psychiatry Residency Program",
        APPOINTMENTS_HEADING, "January 2025", "Present")
    appointment["taxonomy_code"] = "D1"
    prose = _fieldless_record(217, "I remain subscribed to their newsletter.",
                              LEADERSHIP_HEADING)
    return appointment, prose


def test_a_stray_sentence_under_another_heading_renders_no_row():
    """Pins the 6NGAYQ regression, both halves at once: a fieldless record on
    the far side of a source-structure boundary must render no row at all.

    Two defects meet on this record. The employer used to propagate to it in
    document order across ~200 elements and a different source heading, so it
    rendered a row claiming an employer the sentence never named. Stopping the
    carry then leaves the record with nothing in any column, and rendering it
    puts three blank cells in the delivered CV -- a row a reader cannot read
    as anything. Neither is correct output: the record must not become a row,
    and it must not become a row with a borrowed employer either.
    """
    appointment, prose = _appointment_and_distant_prose()

    rows = _render_positions([appointment, prose], code="D1")

    assert rows == [
        ["Faculty member of the residency program",
         "Northgate Hospitals Psychiatry Residency Program", "01/25-Present"],
    ]


def test_the_stray_sentence_inherits_no_employer():
    """The first half of the same regression, read off the record rather than
    the rendered table: the boundary check must still refuse the carry, so a
    later blank-row guard can never be what is hiding a false employer."""
    appointment, prose = _appointment_and_distant_prose()

    WCMTemplateGenerator._propagate_institution_to_subentries([appointment, prose])

    assert _institution_of(prose) == ""


@pytest.mark.parametrize("fields,expected", [
    ({"title": "Attending Physician"},
     ["Attending Physician", "", ""]),
    ({"institution": "Northgate Hospital"},
     ["", "Northgate Hospital", ""]),
    ({"start_date": "2002-07", "end_date": "2006-12"},
     ["", "", "07/02-12/06"]),
])
def test_a_row_with_a_single_populated_cell_is_still_rendered(fields, expected):
    """The guard's other side: only an all-blank row is dropped.

    A legitimately sparse record -- a title the CV gave no employer or dates
    for, an employer with neither, a bare date span -- still carries content a
    reader needs, so it keeps its row. Dropping on "incomplete" rather than on
    "empty" would delete real appointments.
    """
    entry = {"element_idx_start": 1, "taxonomy_code": "D2",
             "text": "", "extracted_fields": dict(fields)}

    assert _render_positions([entry]) == [expected]


def test_a_blank_record_is_reported_rather_than_dropped_silently():
    """A dropped row is a content defect upstream, so the run has to say it
    happened: a verbose line and a stats counter, not a silent skip."""
    gen, table = _positions_generator(TABLE_ANCHORS["D1"])
    gen.verbose = True
    appointment, prose = _appointment_and_distant_prose()

    gen._fill_positions({"D1": [appointment, prose]})

    assert len(_rendered_rows(table)) == 1
    assert gen.stats["blank_position_rows_skipped"] == 1


def test_entries_inserted_still_counts_one_per_rendered_row():
    """Review item 7's invariant across the new drop: the counter follows the
    physical rows, so a record that renders no row increments nothing."""
    gen, table = _positions_generator(TABLE_ANCHORS["D1"])
    appointment, prose = _appointment_and_distant_prose()

    gen._fill_positions({"D1": [appointment, prose]})

    assert gen.stats["entries_inserted"] == len(_rendered_rows(table)) == 1


def test_a_blank_parent_still_yields_its_recovered_children():
    """The guard is applied per record, not per entry: a parent stage 4 found
    no field for is dropped, and the child appointments recovered from its own
    text still render. Dropping the entry wholesale would lose them.
    """
    parent = _tab_joined_entry([("Attending Physician", "07/2002", "06/2003"),
                                ("Assistant Director", "06/2003", "12/2006")])
    parent["extracted_fields"] = {}

    assert _render_positions([parent]) == [
        _synthetic_row("Attending Physician", "07/02-06/03", institution=""),
        _synthetic_row("Assistant Director", "06/03-12/06", institution=""),
    ]


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))


def test_a_city_named_employer_still_gets_its_location():
    """#897: stage 5b hands back `cleaned_name` (location stripped) plus the
    city/state; the row builder then refused to append the location because
    the city WORD was inside the employer's own name. Every "New York
    University" / "New York Presbyterian" appointment rendered without a
    location. The location is appended (as the enrichment tracked change)
    unless the employer string already ends in it. The source states no
    location on purpose: since #899 a source-stated location is written
    plain, not as a tracked change, so the city is left to stage 5b here."""
    entry = _position_entry(1, "Attending Physician",
                            "Crab Hollow University Hospital",
                            ("PROFESSIONAL POSITIONS",), "2018-10", "current")
    entry["institution_enrichment"] = {
        "cleaned_name": "Crab Hollow University Hospital",
        "city": "Crab Hollow", "state": "New York", "country_code": "US",
    }

    gen, table = _positions_generator(TABLE_ANCHORS["D2"])
    gen._fill_positions({"D2": [entry]})
    cell = table.rows[1].cells[1]

    # the name is plain text; the location is the enrichment tracked change,
    # which `cell.text` cannot see (it reads direct `w:r` children only)
    assert cell.text == "Crab Hollow University Hospital"
    inserted = "".join(t.text or "" for p in cell.paragraphs
                       for ins in p._p.findall(qn("w:ins")) for t in ins.iter(qn("w:t")))
    assert inserted == ", Crab Hollow, NY"
