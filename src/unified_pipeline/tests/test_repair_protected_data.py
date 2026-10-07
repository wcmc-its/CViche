"""Tests for `repair/protected_data.py` (#1389): removing the protected personal
data `protected_data_in_output` still finds in a rendered WCM document.

Every value below is synthetic -- no real name, date, or number from any
corpus CV. Each test renders through the real stage-6 generator and then
writes the leak into the saved document, the way a withhold miss reaches it.

Run with:

    python3 -m pytest src/unified_pipeline/tests/test_repair_protected_data.py -p no:cacheprovider
"""

import ast
import copy
import json
import sys
import zipfile
from pathlib import Path

_SRC = Path(__file__).resolve().parents[2]
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from unified_pipeline.doctor.lints.protected_data import (
    lint_protected_data_in_output,  # noqa: E402
)
from unified_pipeline.doctor.shared import docx_body_blocks  # noqa: E402
from unified_pipeline.repair import protected_data as repair  # noqa: E402
from unified_pipeline.stage6.pii_pass import PII_REDACTED_NOTICE  # noqa: E402
from unified_pipeline.stage_6_word_template import (  # noqa: E402
    WCMTemplateGenerator,
    run_stage6,
)

_UID = "TESTRP"
_ENTRIES = [
    {"text": "Jane Q. Public, MD", "taxonomy_code": "A", "element_idx_start": 0,
     "extracted_fields": {"name": "Jane Q. Public, MD"}},
    {"text": "Hobbies include sailing", "taxonomy_code": "T", "element_idx_start": 1,
     "extracted_fields": {}},
]
#: Synthetic values the leaks below carry; none may survive the repair.
_VALUES = ("Pat Example", "Robin Example", "1970", "123-45-6789", "Elm Street")


@pytest.fixture(autouse=True)
def _no_llm_appendix_pass(monkeypatch):
    monkeypatch.setattr(WCMTemplateGenerator, "_reconsider_appendix_entries", lambda self: None)


def _input(tmp_path: Path) -> Path:
    path = tmp_path / "in.json"
    path.write_text(json.dumps({"document_uid": _UID, "entries": _ENTRIES}))
    return path


def _render(tmp_path: Path) -> Path:
    out = tmp_path / f"{_UID}_wcm.docx"
    WCMTemplateGenerator(verbose=False).generate(str(_input(tmp_path)), str(out),
                                                 research_summary_path=None)
    return out


def _paragraph_after(anchor: Paragraph, xml_runs: str) -> Paragraph:
    """A copy of `anchor` (its numbering and style) holding only `xml_runs`."""
    new = copy.deepcopy(anchor._p)
    for child in list(new):
        if child.tag != qn("w:pPr"):
            new.remove(child)
    for element in parse_xml(f"<w:x {nsdecls('w')}>{xml_runs}</w:x>"):
        new.append(element)
    anchor._p.addnext(new)
    return Paragraph(new, anchor._parent)


def _leak_everywhere(docx_path: Path) -> None:
    """One leak per path the repair takes: a cut inside a paragraph, a body
    paragraph that is nothing but the leak, a value split across runs, a
    tracked insertion, a tracked deletion, and a Personal Data value cell."""
    doc = Document(str(docx_path))
    sailing = next(p for p in doc.paragraphs if "sailing" in p.text)
    sailing.runs[-1].text += "; Husband: Pat Example"
    born = _paragraph_after(sailing, '<w:r><w:t xml:space="preserve">Born January </w:t></w:r>'
                                     '<w:r><w:t>2, 1970</w:t></w:r>')
    inserted = _paragraph_after(born, '<w:r><w:t xml:space="preserve">Travel </w:t></w:r>'
                                      '<w:ins w:id="900" w:author="x"><w:r>'
                                      '<w:t>Spouse: Robin Example</w:t></w:r></w:ins>')
    _paragraph_after(inserted, '<w:r><w:t xml:space="preserve">Note </w:t></w:r>'
                               '<w:del w:id="901" w:author="x"><w:r>'
                               '<w:delText>Spouse: Pat Example</w:delText></w:r></w:del>')
    bibliography = next(p for p in doc.paragraphs if p.text.strip().upper().endswith("BIBLIOGRAPHY"))
    _paragraph_after(bibliography, "<w:r><w:t>Social Security Number: 123-45-6789</w:t></w:r>")
    home_row = next(row for table in doc.tables for row in table.rows
                    if row.cells[0].text.strip().lower().startswith("home address"))
    home_row.cells[1].paragraphs[0].add_run("12 Elm Street, Springfield")
    doc.save(str(docx_path))


def _findings(docx_path: Path) -> list[dict]:
    doc = Document(str(docx_path))
    return lint_protected_data_in_output(docx_body_blocks(doc), docx_body_blocks(doc, deleted=True))


def _every_xml_part(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as z:
        return "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))


def test_every_leak_the_lint_finds_is_removed_and_the_lint_then_finds_none(tmp_path):
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    found = len(_findings(docx_path))
    assert found == 6  # one per leak

    result = repair.repair_protected_data(docx_path)

    assert (result.found, result.remaining) == (found, 0)
    assert _findings(docx_path) == []
    xml = _every_xml_part(docx_path)  # body, deletions and comments alike
    assert [v for v in _VALUES if v in xml] == []
    assert sorted(r.where for r in result.removals) == [
        "Appendix", "Appendix", "Appendix", "Appendix", "Personal Data table", "body paragraph"]


def test_a_cut_leaves_the_rest_of_its_paragraph_and_no_dangling_separator(tmp_path):
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    repair.repair_protected_data(docx_path)
    texts = [p.text for p in Document(str(docx_path)).paragraphs]
    assert "1. Hobbies include sailing" in texts or "Hobbies include sailing" in texts
    assert "Travel" in texts and "Note " in texts  # a cut at the end leaves no trailing space


def test_a_body_paragraph_that_was_only_the_leak_carries_the_withheld_notice(tmp_path):
    """The withhold's own presentation (`stage6/pii_pass.py`): a paragraph
    left empty says data was withheld, rather than vanishing silently."""
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    repair.repair_protected_data(docx_path)
    texts = [p.text for p in Document(str(docx_path)).paragraphs]
    assert texts.count(PII_REDACTED_NOTICE) == 2  # the date of birth and the SSN paragraphs


def test_each_cut_paragraph_gets_one_comment_naming_categories_never_values(tmp_path):
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    result = repair.repair_protected_data(docx_path)
    comments = [c for c in Document(str(docx_path)).comments if c.text.startswith("CViche removed")]
    assert len(comments) == len(result.removals) == 6
    assert all(c.author == "CViche" for c in comments)
    assert any("(spouse)" in c.text for c in comments)
    assert not [v for v in _VALUES for c in comments if v in c.text]


def test_the_report_names_where_and_categories_but_no_value(tmp_path):
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    result = repair.repair_and_report(docx_path)
    report = json.loads(repair.repairs_report_path(docx_path).read_text())
    assert report == result.to_json()
    assert repair.repairs_report_path(docx_path).name == f"{_UID}_repairs.json"
    assert (report["found"], report["remaining"], len(report["removed"])) == (6, 0, 6)
    assert not [v for v in _VALUES if v in json.dumps(report)]


def test_a_document_with_no_finding_is_not_rewritten(tmp_path):
    docx_path = _render(tmp_path)
    before = docx_path.read_bytes()
    result = repair.repair_and_report(docx_path)
    assert result is not None and result.to_json() == {
        "repair": "protected_data", "found": 0, "removed": [], "remaining": 0}
    assert docx_path.read_bytes() == before


def test_a_finding_it_cannot_locate_is_reported_left_and_the_document_untouched(tmp_path, monkeypatch):
    """The repair never claims a fix the lint does not see."""
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    before = docx_path.read_bytes()
    monkeypatch.setattr(repair, "_plan", lambda doc: {})
    result = repair.repair_protected_data(docx_path)
    assert (result.found, result.removals, result.remaining) == (6, [], 6)
    assert docx_path.read_bytes() == before


def test_a_failing_repair_leaves_the_render_and_writes_no_report(tmp_path, monkeypatch):
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    before = docx_path.read_bytes()

    def boom(_path):
        raise RuntimeError("synthetic")
    monkeypatch.setattr(repair, "repair_protected_data", boom)

    assert repair.repair_and_report(docx_path) is None
    assert docx_path.read_bytes() == before
    assert not repair.repairs_report_path(docx_path).exists()


def test_a_report_that_cannot_be_written_is_logged_and_the_repair_stands(tmp_path, caplog):
    """One sidecar must never fail a run: an OSError on the report write
    (here a directory squatting on its name) is logged, not raised."""
    docx_path = _render(tmp_path)
    _leak_everywhere(docx_path)
    repair.repairs_report_path(docx_path).mkdir()

    with caplog.at_level("WARNING", logger=repair.__name__):
        result = repair.repair_and_report(docx_path)

    assert result is not None and (result.found, len(result.removals), result.remaining) == (6, 6, 0)
    assert _findings(docx_path) == []
    assert [r.exc_info[0] for r in caplog.records if "report not written" in r.getMessage()] == [
        IsADirectoryError]


def test_run_stage6_survives_a_report_it_cannot_write(tmp_path):
    out_path = tmp_path / f"{_UID}_wcm.docx"
    repair.repairs_report_path(out_path).mkdir()
    out = run_stage6(str(_input(tmp_path)), str(out_path), verbose=False,
                     discover_original_doc=False, repair_protected_data=True)
    assert Path(out).is_file()


@pytest.mark.parametrize("flag, ran", [(False, False), (True, True)])
def test_run_stage6_repairs_only_when_asked(tmp_path, flag, ran):
    out = run_stage6(str(_input(tmp_path)), str(tmp_path / f"{_UID}_wcm.docx"), verbose=False,
                     discover_original_doc=False, repair_protected_data=flag)
    assert repair.repairs_report_path(out).exists() is ran


@pytest.mark.parametrize("value, on", [
    ("1", True), (" 1 ", True), ("0", False), ("", False), (None, False), ("true", False)])
def test_the_flag_is_on_only_for_1(value, on):
    assert repair.repair_flag_on(value) is on


def test_the_loose_dea_shape_is_cut_only_where_its_block_names_dea():
    """The lint's #1217 gate, applied per block: in a table, "DEA number:" sits
    in another cell than the value."""
    licensure = "licensure, board certification"
    assert repair._leaks("AB12CD345", licensure, dea_label_in_block=True)
    assert repair._leaks("AB12CD345", licensure, dea_label_in_block=False) == []
    assert repair._leaks("AB1234567", licensure, dea_label_in_block=False)


_FORBIDDEN_IMPORTS = ("unified_pipeline.stage_6_word_template", "unified_pipeline.run_doctor",
                      "unified_pipeline.quality_score", "app", "web_interface")


def test_the_repair_package_imports_only_in_its_written_direction():
    """§1.1: `repair/__init__.py` says this package may import the doctor and
    stage6, never stage_6_word_template (which imports it lazily), run_doctor,
    quality_score, or the web backend."""
    package = _SRC / "unified_pipeline" / "repair"
    targets = set()
    for module in package.glob("*.py"):
        for node in ast.walk(ast.parse(module.read_text())):
            if isinstance(node, ast.ImportFrom) and node.level == 0 and node.module:
                targets.add(node.module)
            elif isinstance(node, ast.ImportFrom) and node.level:
                targets.add(f"unified_pipeline.repair.{node.module or ''}")
            elif isinstance(node, ast.Import):
                targets.update(alias.name for alias in node.names)
    assert not [t for t in targets for bad in _FORBIDDEN_IMPORTS if t == bad or t.startswith(bad + ".")]


#: Invented title text carrying the lint's independent shapes (#1392 review):
#: a citation subtitle, a cohort "born in <Place>, <years>", a talk title.
_TITLE_TEXTS = (
    "Doe J, Roe R. Anxiety in Children: Case Illustration and Systematic Review. "
    "J Invented Pediatr. 2019;12(3):45-67.",
    "Roe R, Doe J. Outcomes among veterans born in Springfield, 1989-1993. "
    "Invented Med J. 2021;4:1-9.",
    "Foster Children - Placement Stability and Wellbeing",
)


@pytest.mark.parametrize("text", _TITLE_TEXTS)
def test_an_independent_shape_in_appendix_title_text_is_not_cut(text):
    assert repair._leaks(text, "appendix", dea_label_in_block=False) == []


@pytest.mark.parametrize("text", [
    "Married: Pat Example, MD", "Grandchildren – Ann, Bea, Cy", "Citizenship: USA; three children",
    "Born in Springfield, 1970.", "212-555-0199 (home)",
    "Caring for my three children, including 5 years at home 2001-2006"])
def test_a_personal_line_in_the_appendix_is_still_cut(text):
    assert repair._leaks(text, "appendix", dea_label_in_block=False)


@pytest.mark.parametrize("text", [
    # NDMRSO (#1223): a bare "Birth" label before a whole date, as BNYLDF idx 24
    # and HUOGDE idx 8 rendered it; synthetic values.
    "Birth: March 4, 1970 in Exampleville, Examplestan",
    "BIRTH March 4, 1970: Exampleville",
])
def test_a_bare_birth_line_in_the_appendix_is_cut_with_its_place(text):
    leaks = repair._merged(repair._leaks(text, "appendix", dea_label_in_block=False))
    assert [(leak.start, leak.end) for leak in leaks] == [(0, len(text))]


def test_title_text_in_the_appendix_stays_a_finding_and_is_left_whole(tmp_path):
    docx_path = _render(tmp_path)
    doc = Document(str(docx_path))
    anchor = next(p for p in doc.paragraphs if "sailing" in p.text)
    for text in _TITLE_TEXTS:
        anchor = _paragraph_after(anchor, f"<w:r><w:t>{text}</w:t></w:r>")
    doc.save(str(docx_path))
    found = len(_findings(docx_path))
    assert found == len(_TITLE_TEXTS)

    result = repair.repair_protected_data(docx_path)

    assert (result.found, result.removals, result.remaining) == (found, [], found)
    texts = [p.text for p in Document(str(docx_path)).paragraphs]
    assert [t for t in _TITLE_TEXTS if t not in texts] == []


def test_a_cut_at_the_end_of_a_paragraph_leaves_no_trailing_space(tmp_path):
    docx_path = _render(tmp_path)
    doc = Document(str(docx_path))
    sailing = next(p for p in doc.paragraphs if "sailing" in p.text)
    _paragraph_after(sailing, '<w:r><w:t xml:space="preserve">Telephone: 212-555-0100 (work) </w:t></w:r>'
                              "<w:r><w:t>212-555-0199 (home)</w:t></w:r>")
    doc.save(str(docx_path))

    repair.repair_protected_data(docx_path)

    phone = next(p.text for p in Document(str(docx_path)).paragraphs if p.text.startswith("Telephone"))
    assert phone == "Telephone: 212-555-0100 (work)"


def test_the_title_gate_is_the_appendix_s_only():
    """The Personal Data block holds no titles: a long citation-shaped line
    there is still cut."""
    text = "Children: Ann and Bea. The family moved to Springfield in 1975 and later to Rivertown."
    assert repair._independent_leaks(text, "appendix") == []
    assert repair._independent_leaks(text, "personal data")
