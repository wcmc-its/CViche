"""PDF -> docx converter for CV ingest (#806, PR 1: converter only, no wiring).

Every pipeline reader (stages 1a/1b/2, run_doctor) opens the source through
python-docx and indexes off the paragraph stream, so a PDF has to become a
docx that looks like a native one to those readers: one paragraph per logical
line group, an EMPTY paragraph where the page has a vertical gap (stage 1a
keeps blanks as "" and later stages index positions off them), bold carried
onto runs (`run_doctor.iter_header_candidates` reads `all(r.bold ...)`), and
`\\t` where the page has a column gap.

Uses only pdfplumber + python-docx. Unreadable or encrypted PDFs raise; this
module never swallows an error.

Dependencies: third-party only (pdfplumber, python-docx). Imports nothing
from the rest of unified_pipeline.
"""

import math
import re
import statistics
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import pdfplumber
from docx import Document
from docx.shared import Pt
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfplumber.page import Page
from pdfplumber.utils import cluster_objects
from pdfplumber.utils.exceptions import PdfminerException

# --- Word extraction -----------------------------------------------------

# ponytail: a page is "spaceless" by its space-glyph count alone. Ceiling: a
# spaceless page that ALSO letter-spaces a heading wider than the ratio
# splits that heading into letters. Upgrade path: cluster each line's
# inter-character gaps and split at the gap between the two clusters.
#: A page whose space glyphs are fewer than this fraction of its characters
#: separates words by position only (a ScanSoft export measured 0 spaces in
#: 1,047 characters, word gaps 2.7-3.0pt, which glue at pdfplumber's fixed
#: 3pt tolerance).
SPACELESS_MAX_SPACE_FRAC = 0.01
#: ...judged only on a page with at least this many characters.
SPACELESS_MIN_CHARS = 200
#: On a spaceless page a gap wider than this fraction of the font size
#: splits two words. Relative, not fixed: a fixed 1.5pt split letter-spaced
#: capitals on other pages, so it applies only to spaceless pages.
SPACELESS_X_TOLERANCE_RATIO = 0.15

# --- Line grouping -------------------------------------------------------

#: Words whose `bottom` differs by <= this many points share a visual line.
#: `bottom` (not `top`) because words of different sizes on one baseline
#: share a bottom but not a top. Also absorbs a superscript raised about a
#: third of an em (measured: 3pt raise on 10pt text moves `bottom` 3.7pt).
LINE_BOTTOM_TOLERANCE_PT = 4.0
# ponytail: a superscript is any small cluster that touches and overlaps a
# neighbour line. Ceiling: small text set tight against a larger line (a
# footnote-sized word butting a heading) folds into it. Upgrade path: also
# require the cluster's baseline to sit above (superscript) or below
# (subscript) the host's by a measured offset.
#: A line cluster whose every word is at most this fraction of its
#: neighbour line's size, and which overlaps that line vertically, is a
#: raised superscript (measured: 7.9pt ordinals 4.0pt above a 10pt
#: baseline, one line too far for LINE_BOTTOM_TOLERANCE_PT).
SUPERSCRIPT_MAX_SIZE_RATIO = 0.85
#: ...by at least this fraction of the small cluster's own height...
SUPERSCRIPT_MIN_OVERLAP_FRAC = 0.5
#: ...and each of whose words starts or ends within this many ems (of the
#: host's size) of a host word: small text elsewhere on the row (a sidebar
#: set in a smaller size) is another line, not a superscript.
SUPERSCRIPT_ATTACH_EM = 0.25
# ponytail: a sub-0.1em gap is read as one word split by a font change.
# Ceiling: two words set in different fonts with a space narrower than that
# glue together. Upgrade path: look for a space glyph between them in
# `page.chars` instead of measuring the gap.
#: Two adjacent words closer than this many ems are one word that
#: pdfplumber split on a font or size change ("Word" bold + "," regular):
#: joined with no space. A real word space measures ~0.25em.
FONT_CHANGE_JOIN_EM = 0.1

# --- Two-column pages ----------------------------------------------------

# ponytail: two columns at most, found as ONE vertical gutter over the
# whole body band. Ceiling: three columns, or a page that is two-column for
# only part of its height (its full-width part crosses the gutter), reads
# row by row as before. Upgrade path: find gutters per vertical region
# between full-width lines, recursively.
#: A gutter is a vertical strip at least this wide...
GUTTER_MIN_WIDTH_PT = 12.0
#: ...crossed by at most this fraction of the page's body lines (full-width
#: headings may cross it; a row-by-row table crosses it on most rows).
GUTTER_MAX_CROSSING_FRAC = 0.1
#: The gutter is searched for only this far in from each text edge, so a
#: page margin never counts as one.
GUTTER_EDGE_FRAC = 0.15
#: Each column needs at least this many body lines.
COLUMN_MIN_LINES = 5
#: ...and the column with fewer lines at least this fraction of the other's:
#: a sidebar is a text flow as long as the page, while a label or date
#: column prints one line per entry beside the entry's wrapped lines.
#: Measured over the 276 PDFs' candidate pages: a sidebar 0.87-0.97; label
#: and date columns (all mis-split without this) 0.14-0.55.
COLUMN_BALANCE_MIN_FRAC = 0.7
#: Two line bottoms within this many points are on one row.
ROW_ALIGN_TOLERANCE_PT = 1.5
#: If at least this fraction of the lines of the column with fewer lines
#: sit on a row of the other column, the "columns" are a date/label column
#: beside its entries, not two text flows: the page is not split.
#: Measured: a sidebar layout 0.25-0.37; right-hand date columns and Word
#: tables 0.6-1.0 (45 of 92 candidate pages at exactly 1.0).
ROW_ALIGNED_MIN_FRAC = 0.6
#: A column in which at least this fraction of lines carry a column gap is
#: part of a multi-column table cut down its middle: the page is not split.
#: Measured: sidebar layouts 0.0; a Word table split between its third and
#: fourth columns 0.35.
COLUMN_MAX_TAB_FRAC = 0.2

# --- Spacing thresholds, in ems of the previous line's font size ---------
# ponytail: PARAGRAPH_GAP_EM is guessed from Word/LaTeX exports of the
# sample CVs (single spacing = ~0.2em between line boxes). Ceiling: a CV
# with double-spaced body text reads as one paragraph per line. Upgrade
# path: derive it per document from the modal line pitch, as
# `_blank_gap_em` does for the blank threshold -- whose own ceiling is a CV
# whose most common wide gap is paragraph spacing, not a blank line: it
# gets a blank paragraph between paragraphs (one Word twin: +60).

#: Gap above which a new paragraph starts (below it, lines may merge).
PARAGRAPH_GAP_EM = 0.45
#: Fallback gap above which an empty paragraph is emitted before the new
#: one, and the ceiling of the per-document threshold.
BLANK_GAP_EM = 1.0
#: The per-document blank threshold histograms line gaps in bins this wide.
GAP_BIN_EM = 0.05
#: A gap mode above PARAGRAPH_GAP_EM counts as the document's blank-line
#: gap only with at least this many gaps in its bin.
BLANK_MODE_MIN_COUNT = 10
#: Horizontal gap between two words above which they MAY be joined with "\t".
TAB_GAP_EM = 1.5
#: On a FULL-WIDTH line (the only kind a justified paragraph produces) a
#: gap must also be at least this many times the line's median word gap to
#: be a tab: a justified line has uniformly wide gaps (no tabs), a
#: label/value line has one outlier gap (a tab). Every other line tabs any
#: gap over TAB_GAP_EM, so 4+ column rows keep their tabs.
TAB_OUTLIER_RATIO = 2.0
#: A line with at most this many word gaps skips the median test (with one
#: or two gaps the median is the gap itself, so nothing could be an outlier).
TAB_FEW_GAPS_MAX = 2
#: A line that starts with a list marker begins a new entry, never a
#: wrapped continuation: "12." / "3)" (1-3 digits, then whitespace or an
#: uppercase letter: "12.Smith", while "1.5 mg" is not a marker), "[12]",
#: a bullet glyph (U+F0B7 is Word's Symbol-font bullet), or a
#: dash/asterisk/hyphen followed by whitespace. Exceptions below.
LIST_MARKER_RE = re.compile(
    r"^(?:\d{1,3}[.)](?:\s|(?=[A-Z]))|\[\d{1,3}\]|[•▪◦‣●○■□\uf0b7\u2500]|[–—*-]\s)")
#: The dash/asterisk markers, captured. A lone "- present" is a wrapped
#: range, so these only veto a merge when the paragraph itself began with
#: the same marker (a real dash list is consistent).
DASH_MARKER_RE = re.compile(r"^([–—*-])\s")
#: A previous line ending like this is mid-phrase (a wrapped page range
#: "529-" / "45.", a date range, "and"): the next line continues it even if
#: it looks like a marker.
CONNECTOR_END_RE = re.compile(r"(?:[-–—,(:&]|\b(?:and|of|the|in))$", re.IGNORECASE)
#: A line ending like this finished a sentence or an entry: the next
#: line's first word not fitting after it is no evidence of a wrap
#: (a short entry above a long surname), unless it also ends like a
#: connector (CONNECTOR_END_RE, which includes ":").
SENTENCE_END_RE = re.compile(r"[.;:)]$")
_MONTH = (r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?"
          r"|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)")
_DATE = rf"(?:{_MONTH}\.?,?\s+|\d{{1,2}}/(?:\d{{1,2}}/)?)?(?:19|20)\d{{2}}"
#: A tabbed line's label is a LEADING label (so the next line may hang
#: under the text after its tab) only if it is a list number ("12", "3.",
#: "4)"), a year or date ("2019", "Jan 2019", "05/2019"), or a date range
#: ("2019-2020", "Jan 2019 - present", or open: "2019 -").
LEADING_LABEL_RE = re.compile(
    rf"^(?:\d{{1,3}}[.)]?|{_DATE}(?:\s*[-–—]\s*(?:{_DATE}|present|current|now)?)?)[.:,]?$",
    re.IGNORECASE)
#: A line starting with a date (or a date range: it starts with a date).
DATE_START_RE = re.compile(rf"^{_DATE}", re.IGNORECASE)
#: A line starting with a date RANGE and then words is a date-led entry
#: ("2015-2019 Residency, ..."): not a wrapped "2019-2021." tail (no words)
#: and not a wrapped citation tail that starts with one date ("Nov 2019.
#: (...)", "2019, City" -- 4 of 4 single-date vetoes on the twins split one
#: source paragraph).
DATE_LED_ENTRY_RE = re.compile(
    rf"^{_DATE}\s*[-–—]\s*(?:{_DATE}|present|current|now)[.:,]?\s+\S", re.IGNORECASE)
#: A line ending in a hyphen or dash is broken mid-word/mid-range ("1895-" /
#: "1904."), so it wraps on even though it stops short of the right margin.
WRAP_DASH_RE = re.compile(r"[-–—]$")
#: ...except an OPEN YEAR RANGE ("Member, Committee A, 2004-"), which is a
#: complete entry: its dash only counts as a wrap (and a connector) when
#: the next line continues the range: a digit, a month name, or
#: present/current/now.
OPEN_RANGE_RE = re.compile(r"\b(?:19|20)\d{2}\s*[-–—]$")
RANGE_CONTINUATION_RE = re.compile(
    r"^(?:\d|(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?"
    r"|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\b"
    r"|(?:present|current|now)\b)", re.IGNORECASE)
#: A line of only dashes/underscores/spaces is a horizontal rule, not a
#: line that ends in a dash.
RULE_LINE_RE = re.compile(r"^[-–—_\s]+$")
#: Hanging-indent test: a paragraph whose first line starts LEFT of its
#: later lines by more than this many points is a hanging-indent entry;
#: the next line returning to within this of the first line's x0 starts a
#: new entry. A first-line indent (first line right of later lines) is a
#: plain paragraph and never splits. Also the alignment tolerance for a
#: wrapped line under the text after a leading label, and for table cells.
OUTDENT_TOLERANCE_PT = 3.0
# ponytail: "the next line's first word would have fit" assumes the text
# is set ragged-right or justified to one right margin per column.
# Ceiling: a paragraph with its own right indent reads every line as
# short, so its lines stay separate. Upgrade path: per-block right edges.
#: A line wraps into the next when the next line's first word, after a word
#: space of this many ems, would not have fit before the right margin.
WRAP_SPACE_EM = 0.25
#: The right margin ignores this fraction of the document's lines that end
#: furthest right (measured: 2 of ~1,600 lines of a Word export end 8pt past
#: its margin, which made every wrap 2-8pt short look like it would fit).
RIGHT_MARGIN_OUTLIER_FRAC = 0.02
#: A line ends "full" (so the next one is a wrapped continuation, whatever
#: its first word) when its right edge is within this fraction of the text
#: width of the right margin. Also the only kind of line justification
#: stretches.
FULL_LINE_SLACK_FRAC = 0.08
#: A tabbed line may wrap only when every tab follows a LEADING label (a
#: list number or date) ending within this fraction of the text width.
LEADING_LABEL_MAX_FRAC = 0.2
#: Two lines with font sizes further apart than this never merge.
MERGE_SIZE_TOLERANCE_PT = 0.6

# --- Output --------------------------------------------------------------

#: A paragraph starting less than this many points right of its page's text
#: left gets no left indent.
INDENT_MIN_PT = 1.0

# --- Page furniture ------------------------------------------------------

#: A line is running furniture if it repeats at the same position on at
#: least this fraction of the pages (and on at least FURNITURE_MIN_PAGES).
FURNITURE_PAGE_FRACTION = 0.5
FURNITURE_MIN_PAGES = 2
#: Only lines within this fraction of the page height from the top or
#: bottom edge can be furniture, so a repeated body line is never dropped.
FURNITURE_BAND_FRAC = 0.12
#: Repeats of one (digit-masked) edge-band text count as the same position
#: when their tops all lie within this many points: a running name block
#: measured tops 47.3/50.5/57.0 on three pages; a title that moves 25pt
#: between pages is content, not furniture.
FURNITURE_DRIFT_PT = 12.0

# --- Reporting -----------------------------------------------------------

#: A page with fewer characters than this AND at least one image is
#: reported as image-only (scanned page; #536).
IMAGE_ONLY_MAX_CHARS = 20

# ponytail: bold is guessed from the font NAME. Ceiling: generic subset
# names ("CIDFont+F1") carry no weight and read as regular. Upgrade path:
# read the font descriptor (pdfminer `PDFFont.descriptor` StemV / FontWeight
# / the ForceBold flag) instead of the name.
#: Bold face names (matched case-insensitively): Bold/SemiBold/DemiBold,
#: Black, Heavy, Demi, URW "-Medi" (NimbusRomNo9L-Medi), and LaTeX bold
#: extended CMBX12 / CMSSBX10 (`bx` before the size, after the subset prefix).
#: "Medium" (HelveticaNeue-Medium) is deliberately NOT bold: `-medi(?!um)`.
BOLD_FONT_RE = re.compile(r"bold|black|heavy|demi|-medi(?!um)|(?:^|\+)cm[a-z]*bx\d")

#: XML 1.0 forbids these; python-docx raises ValueError on them.
_XML_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_DIGITS_RE = re.compile(r"\d+")


@dataclass(frozen=True)
class ConversionReport:
    pages: int
    paragraphs: int
    blank_paragraphs: int
    image_only_pages: list[int]
    chars: int


@dataclass
class _Run:
    text: str
    bold: bool
    size: float


@dataclass
class _Line:
    top: float
    bottom: float
    x0: float
    x1: float
    size: float
    runs: list[_Run]
    tabs: list[int]  # i: the gap after words[i] is a column gap
    words: list[dict] = field(default_factory=list)  # x-sorted, for a re-layout
    #: (left, right) text edges of the column this line sits in, on a
    #: two-column page; None means the document's margins.
    bounds: tuple[float, float] | None = None
    #: First line of a column after the page's first: its gap to the line
    #: before is not a vertical gap, and it always starts a paragraph.
    block_start: bool = False
    #: A reassembled table row: never re-laid-out, never merged.
    is_row: bool = False

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)

    @property
    def has_tab(self) -> bool:
        return bool(self.tabs)

    @property
    def all_bold(self) -> bool:
        return all(r.bold for r in self.runs if r.text.strip())


@dataclass
class _Para:
    """One output paragraph; `runs` is empty for a blank paragraph."""
    runs: list[_Run] = field(default_factory=list)
    first_x0: float = 0.0            # x0 of the paragraph's first line
    later_x0: float | None = None    # x0 of its second line, once there is one
    marker: str = ""                 # dash/asterisk marker its first line began with
    indent: float = 0.0              # first_x0 minus its page's text left
    date_led: bool = False           # its first line started with a date or date range


def _is_bold(fontname: str) -> bool:
    return BOLD_FONT_RE.search(fontname.lower()) is not None


def _column_gaps(words: list[dict], outlier_test: bool = False) -> list[bool]:
    """For each pair of adjacent words: is the gap between them a column gap
    (a tab), as opposed to a word space? With `outlier_test` (a full-width
    line, which may be justified) a stretched-but-uniform gap is not one."""
    gaps = [b["x0"] - a["x1"] for a, b in zip(words, words[1:])]
    median = statistics.median(gaps) if gaps else 0.0
    exempt = not outlier_test or len(gaps) <= TAB_FEW_GAPS_MAX
    return [gap > TAB_GAP_EM * a["size"] and (exempt or gap >= TAB_OUTLIER_RATIO * median)
            for gap, a in zip(gaps, words)]


def _word_separator(left: dict, right: dict, column_gap: bool) -> str:
    """"\\t" across a column gap, "" across a font change inside one word,
    else a space."""
    if column_gap:
        return "\t"
    if right["x0"] - left["x1"] < FONT_CHANGE_JOIN_EM * left["size"]:
        return ""
    return " "


def _runs(words: list[dict], seps: list[str]) -> list[_Run]:
    """Runs split on (bold, size) change; seps[i] follows words[i] and stays
    with the run on its left."""
    runs: list[_Run] = []
    for index, w in enumerate(words):
        bold, size = _is_bold(w["fontname"]), round(w["size"], 1)
        if runs:
            runs[-1].text += seps[index - 1]
        text = unicodedata.normalize("NFC", w["text"])
        if runs and (runs[-1].bold, runs[-1].size) == (bold, size):
            runs[-1].text += text
        else:
            runs.append(_Run(text, bold, size))
    return runs


def _make_line(words: list[dict], seps: list[str], tabs: list[int]) -> _Line:
    runs = _runs(words, seps)
    return _Line(
        top=min(w["top"] for w in words),
        bottom=max(w["bottom"] for w in words),
        x0=min(w["x0"] for w in words),
        x1=max(w["x1"] for w in words),
        size=max(r.size for r in runs),
        runs=runs,
        tabs=tabs,
        words=words,
    )


def _build_line(words: list[dict], outlier_test: bool = False) -> _Line:
    """One visual line from same-line words: words joined by a space, a tab
    across a wide column gap, or nothing across a font change."""
    words = sorted(words, key=lambda w: w["x0"])
    column_gap = _column_gaps(words, outlier_test)
    seps = [_word_separator(a, b, gap) for a, b, gap in zip(words, words[1:], column_gap)]
    return _make_line(words, seps, [i for i, gap in enumerate(column_gap) if gap])


# --- Words and lines of one page -----------------------------------------

def _is_spaceless(page: Page) -> bool:
    chars = page.chars
    spaces = sum(1 for c in chars if c["text"] == " ")
    return len(chars) >= SPACELESS_MIN_CHARS and spaces < SPACELESS_MAX_SPACE_FRAC * len(chars)


def _page_words(page: Page) -> list[dict]:
    if _is_spaceless(page):
        return page.extract_words(extra_attrs=["fontname", "size"],
                                  x_tolerance_ratio=SPACELESS_X_TOLERANCE_RATIO)
    return page.extract_words(extra_attrs=["fontname", "size"])


def _touches(word: dict, host: list[dict]) -> bool:
    """`word` starts or ends within SUPERSCRIPT_ATTACH_EM of a host word's
    end or start (either way: kerning can tuck it slightly under)."""
    return any(abs(word["x0"] - h["x1"]) <= SUPERSCRIPT_ATTACH_EM * h["size"]
               or abs(h["x0"] - word["x1"]) <= SUPERSCRIPT_ATTACH_EM * h["size"] for h in host)


def _superscript_host(clusters: list[list[dict]], index: int) -> int | None:
    """Index of the neighbour cluster that cluster `index` is a raised
    superscript of, or None."""
    small = clusters[index]
    top, bottom = min(w["top"] for w in small), max(w["bottom"] for w in small)
    best, best_overlap = None, SUPERSCRIPT_MIN_OVERLAP_FRAC * (bottom - top)
    for host in (index - 1, index + 1):
        if not 0 <= host < len(clusters):
            continue
        words = clusters[host]
        host_size = max(w["size"] for w in words)
        if any(w["size"] > SUPERSCRIPT_MAX_SIZE_RATIO * host_size for w in small) \
                or not all(_touches(w, words) for w in small):
            continue
        overlap = min(bottom, max(w["bottom"] for w in words)) - max(top, min(w["top"] for w in words))
        if overlap >= best_overlap:
            best, best_overlap = host, overlap
    return best


def _fold_superscripts(clusters: list[list[dict]]) -> list[list[dict]]:
    """Fold each small-font cluster that overlaps a neighbour line into it."""
    clusters = [list(c) for c in clusters]
    index = 0
    while index < len(clusters):
        host = _superscript_host(clusters, index)
        if host is None:
            index += 1
            continue
        clusters[host].extend(clusters.pop(index))  # the next cluster is now at `index`
    return clusters


def _page_lines(page: Page) -> list[_Line]:
    words = _page_words(page)
    clusters = cluster_objects(words, lambda w: w["bottom"], LINE_BOTTOM_TOLERANCE_PT)
    lines = [_build_line(c) for c in _fold_superscripts(clusters)]
    gutter = _find_gutter(lines, float(page.height))
    if gutter is not None:
        lines = _column_order(lines, gutter, float(page.height))
    return _join_table_rows(lines)


# --- Two-column pages ----------------------------------------------------

def _in_edge_band(line: _Line, page_height: float) -> bool:
    band = FURNITURE_BAND_FRAC * page_height
    return line.top < band or line.bottom > page_height - band


def _widest_gutter(lines: list[_Line]) -> tuple[float, float] | None:
    """Widest vertical strip that at most GUTTER_MAX_CROSSING_FRAC of the
    lines cross, away from the text edges, or None."""
    left = min(ln.x0 for ln in lines)
    right = max(ln.x1 for ln in lines)
    crossing: Counter = Counter()
    for ln in lines:
        crossing.update({x for w in ln.words for x in range(int(w["x0"]), int(w["x1"]) + 1)})
    allowed = GUTTER_MAX_CROSSING_FRAC * len(lines)
    lo, hi = int(left + GUTTER_EDGE_FRAC * (right - left)), int(right - GUTTER_EDGE_FRAC * (right - left))
    best: tuple[float, float] | None = None
    start = None
    for x in range(lo, hi + 2):
        if x <= hi and crossing[x] <= allowed:
            start = x if start is None else start
            continue
        if start is not None and x - start >= GUTTER_MIN_WIDTH_PT and (
                best is None or x - start > best[1] - best[0]):
            best = (float(start), float(x))
        start = None
    return best


def _gutter_cut(line: _Line, gutter: tuple[float, float]) -> int | None:
    """Index of the first right-column word when the line has a column gap
    (over TAB_GAP_EM) by the gutter: from a word ending before the right
    column to one starting past the gutter's middle; the last such gap. A
    sidebar line and a main-column line on one baseline split there, even
    when a sidebar word runs past the middle."""
    middle = (gutter[0] + gutter[1]) / 2
    cuts = [k + 1 for k, (a, b) in enumerate(zip(line.words, line.words[1:]))
            if b["x0"] - a["x1"] > TAB_GAP_EM * a["size"]
            and a["x1"] <= gutter[1] and b["x0"] >= middle]
    return cuts[-1] if cuts else None


def _split_at(line: _Line, gutter: tuple[float, float]) -> tuple[_Line | None, _Line | None]:
    """The line's left-column and right-column words, as two lines: at a
    column gap by the gutter if it has one, else where the right column
    starts (the gutter's end), so a left-column word that pokes into the
    gutter (even past its middle) stays in the left column."""
    cut = _gutter_cut(line, gutter)
    if cut is None:
        edge = gutter[1] - OUTDENT_TOLERANCE_PT
        cut = sum(1 for w in line.words if w["x0"] < edge)
    left, right = line.words[:cut], line.words[cut:]
    return (_build_line(left) if left else None), (_build_line(right) if right else None)


def _row_aligned(left: list[_Line], right: list[_Line]) -> bool:
    """Do the lines of the column with fewer lines sit on rows of the other
    column (a date column beside its entries), rather than run
    independently of it (a sidebar beside the main text)?"""
    small, big = (left, right) if len(left) <= len(right) else (right, left)
    aligned = sum(1 for ln in small
                  if any(abs(ln.bottom - other.bottom) <= ROW_ALIGN_TOLERANCE_PT for other in big))
    return aligned >= ROW_ALIGNED_MIN_FRAC * len(small)


def _has_cell_rows(column: list[_Line]) -> bool:
    """A column whose lines carry column gaps is part of a table."""
    return sum(1 for ln in column if ln.has_tab) >= COLUMN_MAX_TAB_FRAC * len(column)


def _find_gutter(lines: list[_Line], height: float) -> tuple[float, float] | None:
    """A vertical gutter splitting the page's body band into two
    independent text columns, or None."""
    body = [ln for ln in lines if not _in_edge_band(ln, height)]
    if len(body) < COLUMN_MIN_LINES:  # (a visual line may hold both columns' lines)
        return None
    gutter = _widest_gutter(body)
    if gutter is None:
        return None
    # A full-width line belongs to neither column: its halves would skew
    # both the balance and the row-alignment tests.
    halves = [_split_at(ln, gutter) for ln in body if not _crosses(ln, gutter)]
    left = [a for a, _ in halves if a is not None]
    right = [b for _, b in halves if b is not None]
    fewer, more = sorted((len(left), len(right)))
    if fewer < COLUMN_MIN_LINES or fewer < COLUMN_BALANCE_MIN_FRAC * more:
        return None
    if _row_aligned(left, right) or _has_cell_rows(left) or _has_cell_rows(right):
        return None
    return gutter


def _crosses(line: _Line, gutter: tuple[float, float]) -> bool:
    """A full-width line: a word spans the gutter's middle and the text runs
    on into the right column, with no column gap by the gutter and not as a
    right-column line (no word starts on the right column's edge). A
    sidebar line that pokes into the gutter, alone or beside a right-column
    line, does not cross it."""
    middle = (gutter[0] + gutter[1]) / 2
    return (_gutter_cut(line, gutter) is None
            and any(w["x0"] < middle < w["x1"] for w in line.words)
            and line.x1 > gutter[1] + OUTDENT_TOLERANCE_PT
            and not any(abs(w["x0"] - gutter[1]) <= OUTDENT_TOLERANCE_PT for w in line.words))


def _flush_region(out: list[_Line], left: list[_Line], right: list[_Line]) -> None:
    """Append a region's left column, then its right column, to `out`; the
    right column's first line starts a block unless it is the page's first."""
    if right:
        right[0].block_start = bool(left or out)
    out.extend(left + right)
    left.clear()
    right.clear()


def _set_bounds(column: list[_Line]) -> None:
    """Give each line of a column the column's text edges."""
    if column:
        edges = (min(ln.x0 for ln in column), max(ln.x1 for ln in column))
        for ln in column:
            ln.bounds = edges


def _column_order(lines: list[_Line], gutter: tuple[float, float], height: float) -> list[_Line]:
    """Reading order for a two-column page: between full-width lines (which
    cross the gutter, or sit in the edge band and span it with no column
    gap: a name block the gutter was found without), the left column's
    lines, then the right column's. Each column line carries its column's
    text edges as `bounds`."""
    out: list[_Line] = []
    left: list[_Line] = []
    right: list[_Line] = []
    columns: tuple[list[_Line], list[_Line]] = ([], [])
    for ln in lines:
        parts = _split_at(ln, gutter)
        spans_edge_band = (_in_edge_band(ln, height) and None not in parts
                           and _gutter_cut(ln, gutter) is None)
        if spans_edge_band or _crosses(ln, gutter):
            _flush_region(out, left, right)
            out.append(ln)
            continue
        for part, region, column in zip(parts, (left, right), columns):
            if part is not None:
                region.append(part)
                column.append(part)
    _flush_region(out, left, right)
    for column in columns:
        _set_bounds(column)
    return out


# --- Table rows ----------------------------------------------------------

# ponytail: every printed line of a cell joins with a space. Ceiling: a
# cell with hard line breaks (institution / department / city on their own
# lines) becomes one run-on cell. Upgrade path: keep a break where the
# cell's previous line stopped short of the cell's right edge (the next
# column's anchor), as `_next_word_would_not_fit` does for paragraphs.

def _cell_anchors(line: _Line) -> list[float] | None:
    """x0 of each cell of a line with at least two column gaps, else None."""
    if len(line.tabs) < 2:
        return None
    return [line.words[0]["x0"]] + [line.words[i + 1]["x0"] for i in line.tabs]


def _cell_of(word: dict, anchors: list[float]) -> int:
    """Index of the cell whose anchor range contains the word's x0."""
    return max(i for i, a in enumerate(anchors) if i == 0 or a - OUTDENT_TOLERANCE_PT <= word["x0"])


def _near(prev: _Line, line: _Line) -> bool:
    return not line.block_start and line.top - prev.bottom <= PARAGRAPH_GAP_EM * prev.size


def _in_one_cell(line: _Line, anchors: list[float]) -> bool:
    """`line` is wrapped text of one cell: no column gap of its own, and
    every word inside that one cell's x-range."""
    if line.has_tab:
        return False
    cell = _cell_of(line.words[0], anchors)  # its words run left to right from here
    return cell + 1 == len(anchors) or line.x1 <= anchors[cell + 1] - OUTDENT_TOLERANCE_PT


def _is_cell_row(line: _Line, anchors: list[float]) -> bool:
    """`line` is the table's next row with an empty first cell (a
    vertically merged cell): its own column gaps, and every cell it starts
    sits on one of the row's later anchors."""
    if not line.tabs:
        return False
    starts = [line.words[0]["x0"]] + [line.words[i + 1]["x0"] for i in line.tabs]
    return all(any(abs(x - a) <= OUTDENT_TOLERANCE_PT for a in anchors[1:]) for x in starts)


def _row_line(lines: list[_Line], anchors: list[float]) -> _Line:
    """One tab-separated line from a row's printed lines: each word goes to
    the cell whose anchor range contains its x0. An empty cell stays empty
    (a leading empty cell is a leading tab); nothing is guessed into it."""
    cells: list[list[dict]] = [[] for _ in anchors]
    for ln in lines:
        for w in ln.words:
            cells[_cell_of(w, anchors)].append(w)
    words: list[dict] = []
    seps: list[str] = []
    tabs: list[int] = []
    empty = lead = 0
    for cell in cells:
        if not cell:
            empty += 1
            continue
        if words:
            seps.append("\t" * (empty + 1))
            tabs.append(len(words) - 1)
        else:
            lead = empty
        empty = 0
        seps.extend([" "] * (len(cell) - 1))
        words.extend(cell)
    row = _make_line(words, seps, tabs)
    row.runs[0].text = "\t" * lead + row.runs[0].text
    row.is_row = True
    row.block_start = lines[0].block_start
    return row


def _emit_row(out: list[_Line], row: list[_Line], anchors: list[float]) -> None:
    """A row printed on one line from its first cell is kept as printed;
    anything else is reassembled."""
    if len(row) == 1 and row[0].x0 <= anchors[0] + OUTDENT_TOLERANCE_PT:
        out.append(row[0])
    else:
        out.append(_row_line(row, anchors))


def _join_table_rows(lines: list[_Line]) -> list[_Line]:
    """Reassemble table rows whose cells wrapped onto several printed lines
    into one `cell\\tcell` line per row. After a line with at least two
    column gaps, a close line that fits inside one cell continues its row;
    a close line whose own cells sit on the row's anchors is the next row
    (first cell empty); anything else ends the table."""
    out: list[_Line] = []
    row: list[_Line] = []
    anchors: list[float] | None = None
    for line in lines:
        if anchors and _near(row[-1], line):
            if _in_one_cell(line, anchors):
                row.append(line)
                continue
            if _is_cell_row(line, anchors):
                _emit_row(out, row, anchors)
                row = [line]
                continue
        if anchors:
            _emit_row(out, row, anchors)
        anchors = _cell_anchors(line)
        row = [line]
        if anchors is None:
            out.append(line)
    if anchors:
        _emit_row(out, row, anchors)
    return out


# --- Page furniture ------------------------------------------------------

def _furniture_text(line: _Line) -> str:
    # Digits are masked so "Page 3 of 9" repeats as one line across pages.
    return _DIGITS_RE.sub("#", line.text.strip())


def _furniture_windows(pages: list[list[_Line]], heights: list[float]) -> dict[str, list[tuple[float, float]]]:
    """For each edge-band text: the top ranges (FURNITURE_DRIFT_PT wide) in
    which it repeats on enough pages to be running furniture."""
    seen: defaultdict[str, list[tuple[float, int]]] = defaultdict(list)
    for number, (lines, height) in enumerate(zip(pages, heights)):
        for ln in lines:
            if _in_edge_band(ln, height):
                seen[_furniture_text(ln)].append((ln.top, number))
    needed = max(FURNITURE_MIN_PAGES, math.ceil(FURNITURE_PAGE_FRACTION * len(pages)))
    windows: defaultdict[str, list[tuple[float, float]]] = defaultdict(list)
    for text, hits in seen.items():
        for top, _ in hits:
            pages_in = {n for t, n in hits if top <= t <= top + FURNITURE_DRIFT_PT}
            if len(pages_in) >= needed:
                windows[text].append((top, top + FURNITURE_DRIFT_PT))
    return windows


def _is_furniture(line: _Line, height: float, windows: dict[str, list[tuple[float, float]]]) -> bool:
    return _in_edge_band(line, height) and any(
        lo <= line.top <= hi for lo, hi in windows.get(_furniture_text(line), ()))


def _drop_furniture(pages: list[list[_Line]], heights: list[float]) -> list[list[_Line]]:
    """Pages without their running headers/footers. Only edge-band lines
    can match, so a body line is never dropped. A running line in the TOP
    band keeps its first copy: the owner's name and CV title usually sit
    there on page 1, and later stages read the name from body text. Bottom
    band repeats (page numbers, footers) are dropped on every page."""
    windows = _furniture_windows(pages, heights)
    first_kept: set[str] = set()
    out = []
    for lines, height in zip(pages, heights):
        page = []
        for ln in lines:
            if _is_furniture(ln, height, windows):
                text = _furniture_text(ln)
                if ln.top >= FURNITURE_BAND_FRAC * height or text in first_kept:
                    continue
                first_kept.add(text)
            page.append(ln)
        out.append(page)
    return out


# --- Paragraph assembly --------------------------------------------------

def _edges(line: _Line, left: float, right: float) -> tuple[float, float]:
    return line.bounds if line.bounds is not None else (left, right)


def _is_full_line(line: _Line, left: float, right: float) -> bool:
    left, right = _edges(line, left, right)
    return line.x1 >= right - FULL_LINE_SLACK_FRAC * (right - left)


def _next_word_would_not_fit(prev: _Line, line: _Line, right: float) -> bool:
    """`line` is a soft wrap of `prev`: its first word, after a space, would
    not have fitted before the right margin."""
    first = line.words[0]
    return prev.x1 + WRAP_SPACE_EM * prev.size + (first["x1"] - first["x0"]) > right


def _ragged_wrap(prev: _Line, line: _Line, right: float) -> bool:
    """`line`'s first word would not have fitted on `prev`, and `prev` did
    not end a sentence or entry (or ends like a connector)."""
    text = prev.text.rstrip()
    ended = SENTENCE_END_RE.search(text) is not None and CONNECTOR_END_RE.search(text) is None
    return not ended and _next_word_would_not_fit(prev, line, right)


def _hangs_from_label(prev: _Line, line: _Line, left: float, right: float) -> bool:
    """`prev`'s tab follows a leading label (LEADING_LABEL_RE: a list number,
    a year, a date or a date range, ending in the first
    LEADING_LABEL_MAX_FRAC of the width) and `line` starts under the text
    after it: a hanging indent. (`prev` has one tab: a line with two is a
    table row, whose wraps `_join_table_rows` already took.)"""
    tab = prev.tabs[-1]
    label = " ".join(w["text"] for w in prev.words[:tab + 1])
    if LEADING_LABEL_RE.match(label) is None:
        return False
    if prev.words[tab]["x1"] > left + LEADING_LABEL_MAX_FRAC * (right - left):
        return False
    return abs(line.x0 - prev.words[tab + 1]["x0"]) <= OUTDENT_TOLERANCE_PT


def _dash_wraps(prev: _Line, line: _Line) -> bool:
    """`prev` ends in a dash that means "continued on the next line": not a
    rule line, and not an open year range unless `line` continues it."""
    text = prev.text
    if WRAP_DASH_RE.search(text) is None or RULE_LINE_RE.match(text):
        return False
    return not OPEN_RANGE_RE.search(text) or RANGE_CONTINUATION_RE.match(line.text) is not None


def _ends_mid_phrase(prev: _Line, line: _Line) -> bool:
    """`prev` stops mid-phrase, so `line` continues it: a connector ending,
    where a dash counts only under `_dash_wraps`."""
    if WRAP_DASH_RE.search(prev.text):
        return _dash_wraps(prev, line)
    return CONNECTOR_END_RE.search(prev.text) is not None


def _marker_vetoes(prev: _Line, line: _Line, para: _Para) -> bool:
    """`line` starts a new list entry: it opens with a marker, `prev` does
    not end mid-phrase, and a dash/asterisk marker matches the marker the
    paragraph itself began with."""
    if not LIST_MARKER_RE.match(line.text) or _ends_mid_phrase(prev, line):
        return False
    dash = DASH_MARKER_RE.match(line.text)
    return dash is None or dash.group(1) == para.marker


def _date_label_vetoes(prev: _Line, line: _Line, para: _Para) -> bool:
    """`line` starts the next date-led entry: a date range and then words,
    in a paragraph whose own first line started with a date, after a
    `prev` that does not end mid-phrase. (One date-led entry per line is a
    degree / training / appointment list, whatever the line lengths.)"""
    return (para.date_led and DATE_LED_ENTRY_RE.match(line.text) is not None
            and not _ends_mid_phrase(prev, line))


def _outdent_splits(prev: _Line, line: _Line, para: _Para) -> bool:
    """`line` returns to the first-line x0 of a hanging-indent paragraph and
    `prev` does not end mid-phrase."""
    return (not _ends_mid_phrase(prev, line)
            and para.later_x0 is not None
            and para.later_x0 - para.first_x0 > OUTDENT_TOLERANCE_PT
            and line.x0 <= para.first_x0 + OUTDENT_TOLERANCE_PT)


def _continues(prev: _Line, line: _Line, left: float, right: float,
               para: _Para) -> bool:
    """True when `line` looks like a wrapped continuation of `prev` (the
    last line of `para`): same weight and size, no tabs (bar a leading
    label that `line` hangs under), not a new list entry or hanging-indent
    entry, and `prev` ran to the right margin or `line`'s first word would
    not have fitted on it.
    ponytail: a hanging indent needs two lines to be seen, so a one-line
    entry followed by another entry merges unless a marker separates them."""
    if _marker_vetoes(prev, line, para) or _date_label_vetoes(prev, line, para) \
            or _outdent_splits(prev, line, para):
        return False
    left, right = _edges(prev, left, right)
    if line.has_tab or (prev.has_tab and not _hangs_from_label(prev, line, left, right)):
        return False
    if prev.all_bold != line.all_bold or abs(prev.size - line.size) > MERGE_SIZE_TOLERANCE_PT:
        return False
    return (_is_full_line(prev, left, right) or _ragged_wrap(prev, line, right)
            or _dash_wraps(prev, line))


def _margins(pages: list[list[_Line]]) -> tuple[float, float]:
    """Left edge over every kept line; right edge ignoring the few lines
    that stick out past the text block's right margin."""
    lines = [ln for page in pages for ln in page]
    rights = sorted(ln.x1 for ln in lines)
    return min(ln.x0 for ln in lines), rights[-1 - int(RIGHT_MARGIN_OUTLIER_FRAC * len(rights))]


def _untab_justified(pages: list[list[_Line]]) -> list[list[_Line]]:
    """Re-lay-out full-width tabbed lines with the outlier test: only a
    line that reaches the right margin can be a justified, stretched one."""
    if not any(pages):
        return pages
    left, right = _margins(pages)
    out = []
    for lines in pages:
        page = []
        for ln in lines:
            if ln.has_tab and not ln.is_row and _is_full_line(ln, left, right):
                relaid = _build_line(ln.words, outlier_test=True)
                relaid.bounds, relaid.block_start = ln.bounds, ln.block_start
                ln = relaid
            page.append(ln)
        out.append(page)
    return out


def _blank_gap_em(pages: list[list[_Line]]) -> float:
    """The document's blank-paragraph gap threshold, in ems: midway between
    its line-gap mode and its next gap mode above PARAGRAPH_GAP_EM (a Word
    blank line measured ~0.9em, just under BLANK_GAP_EM). BLANK_GAP_EM when
    there is no clear second mode; never above it."""
    bins: Counter = Counter()
    for lines in pages:
        for a, b in zip(lines, lines[1:]):
            gap = b.top - a.bottom
            if gap >= 0 and not b.block_start:
                bins[round(gap / a.size / GAP_BIN_EM)] += 1
    if not bins:
        return BLANK_GAP_EM
    line_mode = max(bins, key=lambda k: (bins[k], -k))
    wide = {k: n for k, n in bins.items() if k * GAP_BIN_EM > PARAGRAPH_GAP_EM}
    if line_mode * GAP_BIN_EM > PARAGRAPH_GAP_EM or not wide:
        return BLANK_GAP_EM
    blank_mode = max(wide, key=lambda k: (wide[k], -k))
    if wide[blank_mode] < BLANK_MODE_MIN_COUNT:
        return BLANK_GAP_EM
    midpoint = (line_mode + blank_mode) / 2 * GAP_BIN_EM
    return min(BLANK_GAP_EM, max(PARAGRAPH_GAP_EM, midpoint))


def _blank_before_last(paras: list[_Para]) -> bool:
    """The last paragraph was blank-separated from the one before it: a
    page break that starts a new paragraph after it keeps that rhythm (the
    gap across a page break cannot be measured)."""
    return len(paras) >= 2 and not paras[-2].runs


def _assemble(pages: list[list[_Line]], blank_gap_em: float = BLANK_GAP_EM) -> list[_Para]:
    """Lines -> paragraph stream, with an empty paragraph per large gap."""
    if not any(pages):
        return []
    left, right = _margins(pages)
    paras: list[_Para] = []
    prev: _Line | None = None
    for lines in pages:
        page_left = min((ln.x0 for ln in lines), default=left)
        for index, line in enumerate(lines):
            same_page = index > 0 and not line.block_start
            gap = line.top - prev.bottom if (prev and same_page) else 0.0
            # A column start never continues the column before it: in a CV the
            # left column is a sidebar, not the start of a flowing text.
            merge = prev is not None and not line.block_start \
                and gap <= PARAGRAPH_GAP_EM * prev.size \
                and _continues(prev, line, left, right, paras[-1])
            if merge:
                paras[-1].runs[-1].text += " "
                paras[-1].runs.extend(line.runs)
                if paras[-1].later_x0 is None:
                    paras[-1].later_x0 = line.x0
            else:
                if prev is not None and (gap > blank_gap_em * prev.size
                                         or (index == 0 and _blank_before_last(paras))):
                    paras.append(_Para())
                dash = DASH_MARKER_RE.match(line.text)
                paras.append(_Para(runs=list(line.runs), first_x0=line.x0,
                                   marker=dash.group(1) if dash else "",
                                   indent=line.x0 - page_left,
                                   date_led=DATE_START_RE.match(line.text) is not None))
            prev = line
    return paras


def _write_docx(paras: list[_Para], docx_path: Path) -> None:
    doc = Document()
    for para in paras:
        p = doc.add_paragraph()
        if para.indent >= INDENT_MIN_PT:
            p.paragraph_format.left_indent = Pt(round(para.indent, 1))
        for run in para.runs:
            r = p.add_run(_XML_ILLEGAL_RE.sub("", run.text))
            r.bold = run.bold
            r.font.size = Pt(run.size)
    doc.save(str(docx_path))


def _is_image_only(page: Page) -> bool:
    return len(page.chars) < IMAGE_ONLY_MAX_CHARS and len(page.images) > 0


def _read_pages(pdf_path: str | Path) -> tuple[list[list[_Line]], list[float], list[int]]:
    """Each page's lines in reading order, page heights, image-only pages."""
    pages: list[list[_Line]] = []
    heights: list[float] = []
    image_only: list[int] = []
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            for number, page in enumerate(pdf.pages, start=1):
                pages.append(_page_lines(page))
                heights.append(float(page.height))
                if _is_image_only(page):
                    image_only.append(number)
    except PdfminerException as exc:
        # pdfplumber wraps everything in a message-less PdfminerException;
        # name the one cause a caller can act on, re-raise the rest as is.
        if exc.args and isinstance(exc.args[0], PDFPasswordIncorrect):
            raise ValueError("encrypted/password-protected PDF") from exc
        raise
    return pages, heights, image_only


def convert_pdf_to_docx(pdf_path: str | Path, docx_path: str | Path) -> ConversionReport:
    """Convert `pdf_path` to a docx at `docx_path`; raises on an unreadable
    or encrypted PDF."""
    pages, heights, image_only = _read_pages(pdf_path)
    kept = _drop_furniture(pages, heights)
    # ponytail: no docx tables in v1. pdfplumber invents tables from
    # tab-aligned columns (10 on one CV), which sends a PDF-sourced CV down
    # a different reader path than its docx twin. Ceiling: a real ruled
    # table becomes tab-separated paragraphs, one per row. Upgrade path: emit
    # a table only for pdfplumber tables with ruling lines (strategy "lines")
    # and >= 2 columns, and check element order with extract_unified_elements.
    paras = _assemble(_untab_justified(kept), _blank_gap_em(kept))
    _write_docx(paras, Path(docx_path))
    blanks = sum(1 for p in paras if not p.runs)
    chars = sum(len(r.text) for p in paras for r in p.runs)
    return ConversionReport(
        pages=len(pages),
        paragraphs=len(paras),
        blank_paragraphs=blanks,
        image_only_pages=image_only,
        chars=chars,
    )
