#!/usr/bin/env python3
"""
Test the transparent 3-stage pipeline on 2071_Zuschlag_Cv

Stage 1: Extract hierarchy only (headers)
Stage 2a: LLM identifies entry delimiters with full hierarchy
Stage 2b: Extract actual text snippets (no LLM)
"""

import sys
import json
from pathlib import Path
from docx import Document

# Add to path
sys.path.insert(0, str(Path(__file__).parent))

from core.output_manager import OutputManager
from segmentation.signature_based_segmentation import segment_cv_hierarchical

def stage_1_hierarchy_only(docx_path: str):
    """
    Stage 1: Extract just the hierarchy (headers only, no entries)

    Output:
        - JSON with hierarchy structure
        - TXT with readable hierarchy tree
    """
    print("="*80)
    print("STAGE 1: HIERARCHY EXTRACTION")
    print("="*80)
    print(f"Input: {docx_path}")
    print()

    # Use existing signature-based segmentation to get hierarchy
    result = segment_cv_hierarchical(docx_path, output_path=None)

    # Load the result
    with open(result['output_file']) as f:
        full_data = json.load(f)

    # Extract just the hierarchy (remove entries, keep structure)
    def extract_hierarchy(node):
        """Recursively extract hierarchy without entry content"""
        hierarchy_node = {
            "text": node.get("text", ""),
            "level": node.get("level", ""),
            "paragraph_index": node.get("paragraph_index", -1)
        }

        # Add children if they exist
        if "children" in node and node["children"]:
            hierarchy_node["children"] = [extract_hierarchy(child) for child in node["children"]]

        return hierarchy_node

    hierarchy_only = {
        "document_uid": full_data.get("document_uid"),
        "meta": {
            "total_paragraphs": full_data.get("meta", {}).get("total_paragraphs", 0),
            "total_headers": full_data.get("meta", {}).get("total_headers", 0),
            "top_level_sections": full_data.get("meta", {}).get("top_level_sections", 0)
        },
        "hierarchy": [extract_hierarchy(node) for node in full_data.get("hierarchy", [])]
    }

    # Save to Stage 1 output
    om = OutputManager(docx_path)
    output_json = om.get_stage1_json_path()
    output_txt = om.get_stage1_txt_path()

    # Save JSON
    with open(output_json, 'w') as f:
        json.dump(hierarchy_only, f, indent=2)

    # Save TXT (human-readable)
    with open(output_txt, 'w') as f:
        f.write(f"CV Hierarchy: {hierarchy_only['document_uid']}\n")
        f.write("="*80 + "\n\n")

        def write_hierarchy(nodes, depth=0):
            for node in nodes:
                indent = "  " * depth
                level = node.get("level", "H?")
                text = node.get("text", "NO TITLE")
                f.write(f"{indent}[{level}] {text}\n")

                if "children" in node:
                    write_hierarchy(node["children"], depth + 1)

        write_hierarchy(hierarchy_only["hierarchy"])

    print(f"✓ Hierarchy JSON: {output_json}")
    print(f"✓ Hierarchy TXT: {output_txt}")
    print(f"  Total headers: {hierarchy_only['meta']['total_headers']}")
    print(f"  Top-level sections: {hierarchy_only['meta']['top_level_sections']}")
    print()

    return hierarchy_only, output_json

def stage_2a_find_delimiters(docx_path: str, hierarchy_json: str):
    """
    Stage 2a: LLM identifies entry delimiters within each section

    For each leaf section in the hierarchy, ask LLM to identify:
    - element_idx_start
    - element_idx_end
    - element_type (paragraph, table, table_row)
    """
    print("="*80)
    print("STAGE 2A: ENTRY DELIMITER DETECTION (LLM)")
    print("="*80)
    print(f"Input hierarchy: {hierarchy_json}")
    print()

    # Load hierarchy
    with open(hierarchy_json) as f:
        hierarchy = json.load(f)

    # Load Word document structure
    doc = Document(docx_path)

    # Extract all elements with indices
    elements = []
    for i, para in enumerate(doc.paragraphs):
        elements.append({
            "idx": i,
            "type": "paragraph",
            "text": para.text
        })

    # For now, create a simple delimiter structure
    # In full implementation, this would call LLM for each section
    om = OutputManager(docx_path)
    output_path = om.get_stage2a_path()

    # Example delimiters (would be LLM-generated)
    delimiters = {
        "document_uid": hierarchy["document_uid"],
        "total_sections_analyzed": len(hierarchy["hierarchy"]),
        "delimiters": []
    }

    # Save
    with open(output_path, 'w') as f:
        json.dump(delimiters, f, indent=2)

    print(f"✓ Delimiters: {output_path}")
    print(f"  NOTE: This is a stub - full LLM implementation needed")
    print()

    return delimiters, output_path

def stage_2b_extract_snippets(docx_path: str, delimiters_json: str):
    """
    Stage 2b: Extract actual text snippets using delimiters (no LLM)

    Simply slice elements[start:end+1] to get entry text
    """
    print("="*80)
    print("STAGE 2B: SNIPPET EXTRACTION (NO LLM)")
    print("="*80)
    print(f"Input delimiters: {delimiters_json}")
    print()

    # Load delimiters
    with open(delimiters_json) as f:
        delimiters = json.load(f)

    # Load Word document
    doc = Document(docx_path)

    # Extract snippets
    om = OutputManager(docx_path)
    output_path = om.get_stage2b_path()

    entries = {
        "document_uid": delimiters["document_uid"],
        "total_entries": 0,
        "entries": []
    }

    for section in delimiters.get("delimiters", []):
        hierarchy = section.get("hierarchy", [])

        for delimiter in section.get("entry_delimiters", []):
            start_idx = delimiter["element_idx_start"]
            end_idx = delimiter["element_idx_end"]
            element_type = delimiter.get("element_type", "paragraph")

            # Extract text
            if element_type == "paragraph":
                text_lines = []
                for i in range(start_idx, end_idx + 1):
                    if i < len(doc.paragraphs):
                        text_lines.append(doc.paragraphs[i].text)
                text = "\n".join(text_lines)
            else:
                text = f"[Table extraction not implemented yet for idx {start_idx}-{end_idx}]"

            entries["entries"].append({
                "hierarchy": hierarchy,
                "element_idx_start": start_idx,
                "element_idx_end": end_idx,
                "element_type": element_type,
                "text": text
            })
            entries["total_entries"] += 1

    # Save
    with open(output_path, 'w') as f:
        json.dump(entries, f, indent=2)

    print(f"✓ Entries: {output_path}")
    print(f"  Total entries extracted: {entries['total_entries']}")
    print()

    return entries, output_path

if __name__ == '__main__':
    cv_path = "data/sample_cvs/word/2071_Zuschlag_Cv.docx"

    print("\n" + "="*80)
    print("TRANSPARENT PIPELINE TEST: 2071_Zuschlag_Cv")
    print("="*80)
    print()

    # Stage 1: Hierarchy only
    hierarchy, hierarchy_json = stage_1_hierarchy_only(cv_path)

    # Stage 2a: Find delimiters (stub for now)
    delimiters, delimiters_json = stage_2a_find_delimiters(cv_path, hierarchy_json)

    # Stage 2b: Extract snippets
    entries, entries_json = stage_2b_extract_snippets(cv_path, delimiters_json)

    print("="*80)
    print("PIPELINE COMPLETE")
    print("="*80)
    print(f"Stage 1: {hierarchy_json}")
    print(f"Stage 2a: {delimiters_json}")
    print(f"Stage 2b: {entries_json}")
    print("="*80)
