"""Keep every WCM template slot, and write "N/A" in the ones nothing filled (#1942).

Decided 2026-10-19 (Paul), from a WCM reviewer's corrected copy of run XAXVYB:
every subsection and table of the WCM template stays in the delivered CV, and
one with no entries says so. An empty paragraph subsection ("Pending Funding",
"Other education/outreach activities") gets one "N/A" paragraph under its
heading. A template table nothing wrote into stays in place with "N/A" in its
first value cell (a record table: `is_record_table`). This reverses #845 (2026-09-14), which removed unfilled
placeholder tables because "the delivered document is a finished record, not
a form": the reviewer reads a missing table as one CViche deleted, and asks for
it back with N/A in it.

The first value cell, not every one: a faculty member filling in the template
by hand writes N/A once, in the table's first answer ("Award Source: | N/A",
the shape of the XAXVYB source itself), and leaves the rest blank. That reads
as "this whole record does not apply" where a column of N/A would read as six
separate answers. For a column-header table (Title | Institution | Dates) the
first value cell is the first cell of the blank row under the header; a
section that cleared its table down to the header row and had nothing to add
(Honors, Invitations to Speak) gets the template's blank row back for it.

Emptiness is bookkeeping, not text search. `snapshot_template` records the
text of every paragraph and table of the template as loaded, before any
section renders. After every section has rendered, a paragraph slot is empty
when nothing between its heading and the next heading is new or changed, and
a template table is unfilled when none of its cells holds text the template's
cell did not. The slot headings are matched only on paragraphs the template
itself shipped, so a content line that happens to read "Books:" is never one.

No `stage_6_word_template` import (a back-edge, `stage6/__init__.py`).
"""
from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.text.paragraph import Paragraph

from ..core.template_boilerplate import (
    normalize_template_text,
    template_section_headers,
)
from .formatting.docx import _insert_after, _set_font

#: What an empty template slot says (#1942).
NOT_APPLICABLE = "N/A"

#: The heading `_write_appendix_header` opens the Appendix with. Nothing at or
#: after it is a template slot.
APPENDIX_HEADING = "T. APPENDIX"

#: The template's subsections that hold paragraphs, not a table, as the start
#: of each heading's text (`key_files/wcm_cv_template_faculty_october_2022_final.docx`).
#: Each gets "N/A" when nothing rendered under it. "Past (Completed) Funding"
#: is one: the template gives it no table, and its grants render as new ones.
PARAGRAPH_SLOT_PREFIXES = (
    "Didactic teaching",
    "Clinical teaching",
    "Administrative teaching",
    "Continuing education and professional education",
    "Other education/outreach activities",
    "Clinical Practice",
    "Clinical Innovations",
    "Clinical Leadership",
    "Research Activities:",
    "Past (Completed) Funding",
    "Pending Funding",
    "Patents & Inventions",
    "Leadership and mentoring in programs",
    "Peer-reviewed Research Articles:",
    "Reviews and Editorials:",
    "Books:",
    "Chapters:",
    "Non-peer-reviewed Research Publications:",
    "Case Reports",
    "In review",
    "Abstracts",
    "Other (media, podcasts, etc.):",
)

#: Headings over a template table. The table carries their N/A, so they are no
#: paragraph slot, but each ends the scope of the slot above it.
_TABLE_SLOT_PREFIXES = (
    "Current Research Funding",
    "Institutional Training Grants and Mentored Trainee Grants",
    "Mentees",
    "Current Mentees:",
    "Past Mentees:",
)

#: The first label of the template's record tables laid out label | value: a
#: grant ("Award Source: ...", under Current Research Funding and under
#: Institutional Training Grants) and a mentee ("Name"). The other label |
#: value tables -- Personal Data, DEA/NPI, Hospital Affiliation -- and the
#: Percent Effort grid are forms, not lists of records: a blank answer there
#: is not a slot with no entries, and gets no N/A.
_GRANT_TABLE_FIRST_LABEL = "Award Source"
_MENTEE_TABLE_FIRST_LABEL = "Name"

_SECTION_HEADERS = template_section_headers()


@dataclass(frozen=True)
class TemplateSnapshot:
    """The template's own body text, per element, before any section renders.

    Keyed by the lxml element itself: holding it keeps lxml's proxy alive, so
    the same body node keeps returning this same object for the lifetime of
    the snapshot.
    """
    paragraphs: Mapping[BaseOxmlElement, str]
    tables: Mapping[BaseOxmlElement, tuple[tuple[str, ...], ...]]
    #: A copy of each table's first all-blank row, for a table a section
    #: cleared down to its header and then had nothing to add.
    blank_rows: Mapping[BaseOxmlElement, BaseOxmlElement]


@dataclass(frozen=True)
class SlotFillCounts:
    """How many slots `fill_empty_slots` wrote N/A into."""
    paragraphs: int
    tables: int


def _text(element: BaseOxmlElement) -> str:
    """Every `w:t` under *element*, joined: tracked insertions in, deletions out."""
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


def _table_grid(tbl: BaseOxmlElement) -> tuple[tuple[str, ...], ...]:
    """*tbl*'s cell texts, row by row, one entry per `w:tc`."""
    return tuple(tuple(_text(tc) for tc in tr.iterchildren(qn("w:tc")))
                 for tr in tbl.iterchildren(qn("w:tr")))


def snapshot_template(body: BaseOxmlElement) -> TemplateSnapshot:
    """Record *body*'s paragraphs and tables as the template ships them."""
    paragraphs = {el: _text(el) for el in body.iterchildren(qn("w:p"))}
    tables = {el: _table_grid(el) for el in body.iterchildren(qn("w:tbl"))}
    blank_rows = {}
    for tbl in tables:
        blank = next((tr for tr in tbl.iterchildren(qn("w:tr")) if not _text(tr).strip()), None)
        if blank is not None:
            blank_rows[tbl] = deepcopy(blank)
    return TemplateSnapshot(paragraphs, tables, blank_rows)


def _heading_text(text: str) -> str:
    return " ".join(text.split())


def _is_paragraph_slot(element: BaseOxmlElement, snapshot: TemplateSnapshot) -> bool:
    """True for a template-shipped paragraph that heads a paragraph slot."""
    if element not in snapshot.paragraphs:
        return False
    return _heading_text(snapshot.paragraphs[element]).startswith(PARAGRAPH_SLOT_PREFIXES)


def _ends_scope(element: BaseOxmlElement, snapshot: TemplateSnapshot) -> bool:
    """True where a paragraph slot's scope stops: the next heading or table the
    template shipped, or the Appendix."""
    if element.tag == qn("w:tbl"):
        return element in snapshot.tables
    if element.tag != qn("w:p"):
        return False
    if _heading_text(_text(element)) == APPENDIX_HEADING:
        return True
    if element not in snapshot.paragraphs:
        return False
    heading = _heading_text(snapshot.paragraphs[element])
    return (heading.startswith(PARAGRAPH_SLOT_PREFIXES + _TABLE_SLOT_PREFIXES)
            or normalize_template_text(heading) in _SECTION_HEADERS)


#: What a paragraph with no text can still show: a picture or an embedded object.
_PARAGRAPH_CONTENT_TAGS = (qn("w:drawing"), qn("w:pict"), qn("w:object"))


def _is_blank_paragraph(paragraph: BaseOxmlElement) -> bool:
    """No text and nothing else to show: a spacer a section writer added
    (WG6JFG's Abstracts slot had one more than the template), not content."""
    return not _text(paragraph).strip() and next(
        paragraph.iter(*_PARAGRAPH_CONTENT_TAGS), None) is None


def _holds_nothing_written(element: BaseOxmlElement, snapshot: TemplateSnapshot) -> bool:
    """True for a body element that holds nothing a section wrote: one the
    template shipped and nothing has changed, or a blank paragraph. Anything
    that is neither a paragraph nor a table (a bookmark, `w:sectPr`) holds no
    content either."""
    if element.tag == qn("w:p"):
        if _is_blank_paragraph(element):
            return True
        return element in snapshot.paragraphs and _text(element) == snapshot.paragraphs[element]
    if element.tag == qn("w:tbl"):
        return element in snapshot.tables and _table_grid(element) == snapshot.tables[element]
    return True


def paragraph_slot_is_empty(heading: BaseOxmlElement, snapshot: TemplateSnapshot) -> bool:
    """True when nothing new or changed sits between *heading* and the next
    heading: the template's own instruction and blank lines are all there is."""
    for element in heading.itersiblings():
        if _ends_scope(element, snapshot):
            return True
        if not _holds_nothing_written(element, snapshot):
            return False
    return True


def is_record_table(template_grid: tuple[tuple[str, ...], ...]) -> bool:
    """True when the template ships this table as a blank record: column
    headers over blank rows (Title | Institution | Dates), or a grant or
    mentee block (`_GRANT_TABLE_FIRST_LABEL`, `_MENTEE_TABLE_FIRST_LABEL`)."""
    if not template_grid or not template_grid[0]:
        return False
    header, rows = template_grid[0], template_grid[1:]
    if rows and all(cell.strip() for cell in header) and not any(
            cell.strip() for row in rows for cell in row):
        return True
    first = header[0].strip()
    return first == _MENTEE_TABLE_FIRST_LABEL or first.startswith(_GRANT_TABLE_FIRST_LABEL)


def table_is_unfilled(tbl: BaseOxmlElement, template_grid: tuple[tuple[str, ...], ...]) -> bool:
    """True when no cell of *tbl* holds text its template cell did not: every
    non-empty cell is still the template's label, and no row was added."""
    grid = _table_grid(tbl)
    if len(grid) > len(template_grid):
        return False
    for row, template_row in zip(grid, template_grid, strict=False):
        for column, text in enumerate(row):
            if text and (column >= len(template_row) or text != template_row[column]):
                return False
    return True


def _first_empty_cell(tbl: BaseOxmlElement) -> BaseOxmlElement | None:
    """The first `w:tc` with no text, reading row by row, or None."""
    for tc in tbl.iter(qn("w:tc")):
        if not _text(tc).strip():
            return tc
    return None


def _write_not_applicable(paragraph: BaseOxmlElement) -> None:
    run = Paragraph(paragraph, None).add_run(NOT_APPLICABLE)
    _set_font(run)


def _fill_unfilled_table(tbl: BaseOxmlElement, blank_row: BaseOxmlElement | None) -> bool:
    """Write N/A into *tbl*'s first empty cell. A table a section cleared down
    to its header (`_clear_table_data`) gets the template's blank row back
    first. False when there is nowhere to write."""
    tc = _first_empty_cell(tbl)
    if tc is None and blank_row is not None:
        tbl.append(deepcopy(blank_row))
        tc = _first_empty_cell(tbl)
    if tc is None:
        return False
    paragraph = tc.find(qn("w:p"))
    if paragraph is None:
        paragraph = OxmlElement("w:p")
        tc.append(paragraph)
    _write_not_applicable(paragraph)
    return True


def _insert_not_applicable_after(heading: BaseOxmlElement) -> None:
    paragraph = OxmlElement("w:p")
    _write_not_applicable(paragraph)
    _insert_after(heading, paragraph)


def fill_empty_slots(body: BaseOxmlElement, snapshot: TemplateSnapshot) -> SlotFillCounts:
    """Write N/A into every empty paragraph slot and unfilled template table of
    *body* above the Appendix. Runs once, after every section has rendered."""
    paragraphs = tables = 0
    for element in list(body.iterchildren()):
        if element.tag == qn("w:p") and _heading_text(_text(element)) == APPENDIX_HEADING:
            break
        if element in snapshot.tables:
            template_grid = snapshot.tables[element]
            if (is_record_table(template_grid) and table_is_unfilled(element, template_grid)
                    and _fill_unfilled_table(element, snapshot.blank_rows.get(element))):
                tables += 1
        elif _is_paragraph_slot(element, snapshot) and paragraph_slot_is_empty(element, snapshot):
            _insert_not_applicable_after(element)
            paragraphs += 1
    return SlotFillCounts(paragraphs, tables)
