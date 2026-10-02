"""Section F1: licensure, plus the DEA and NPI numbers (#398).

The licensure table is four columns -- state, number, date of issue, date of
last registration -- but the section's real work is that three different kinds
of identifier arrive under one taxonomy code. A state medical licence belongs in
the table; a DEA registration and an NPI belong in a separate two-row table
further down the template, which is why `_fill_dea_npi` lives here rather than
anywhere else: it is the second half of one section's output.

#821 (settled 2026-09-14): a DEA number is withheld with notice, like every
other #820/#821 protected-data category, but enforced HERE rather than by
`stage6/normalization/pii.py`'s policy table -- `_resolve_licensure` never
lets an entry it classifies as DEA reach `identifiers.dea`, whether the
classification came from a label or from the number's shape alone. The NPI
is a public identifier and still renders.

The section runs as four steps, and only the last one touches the document:

    raw stage-4 entries
        -> _normalize_licensure_entry   one LicensureEntry per raw dict
        -> _classify_licensure_entry    licence vs NPI vs DEA
        -> _resolve_licensure           LicensureResult (rows + identifiers)
        -> _fill_licensure              writes the rows into the Word table

The first three take no `docx` object and no `self.stats`, so classification
and field resolution are testable without a document, and a change to Word
table handling cannot reach them (#624 review). `_fill_licensure` reads no
`extracted_fields` at all -- every stage-4 field alias is resolved in exactly
one place, `_normalize_licensure_entry`.

Sorting the three kinds out is done by LABEL first and shape second, because
shape alone cannot do it (#573):

- the `license_type` field, when stage 4 extracted one;
- else the word "NPI" or "DEA" in the entry's raw text, at a word start and
  not running on into letters (a fused "NPI15180546000" still counts) -- the
  old substring test fired on any entry mentioning "Dean"; the agency's full
  name ("Drug Enforcement Administration") counts as the DEA label, in the
  raw text or in the state field it was extracted into (#1217);
- else, and only when the entry names no state, the number's shape: an NPI
  is 10 or 11 digits, a DEA-LIKE number two letters plus seven alphanumerics
  (wider than the real two-letters-plus-seven-DIGITS format on purpose --
  this tiebreak withholds, so it fails safe by over-matching; the tight
  `_DEA_NUMBER_RE` is what may DELETE a line. See both regexes' comment). An
  all-numeric 10-11 digit state licence number is indistinguishable from an
  NPI by shape, which is why a stated jurisdiction disables the tiebreak.

Either match consumes the entry -- it is pulled out of the licence list, not
copied -- so a DEA number never also appears as a row in the state table. A DEA
number that the reader fused into ANOTHER licence entry's number cell (stacked
paragraphs read as one row) is cut out of that cell instead, and withheld with
the same notice (#1217). The unstructured fallback is guarded by the same two
label tests for the same reason. The NPI and DEA slots are single-valued; a
second, different candidate is ignored with a warning rather than silently
overwriting the first (#573).

`_fill_dea_npi` finds its table by scanning every table in the document for a
cell containing "DEA number", not by position. The template revision history has
moved it, and matching on content survives that. It then matches each row by
its own label, so a template that lists NPI first still fills correctly.

The row writer copes with a template whose licensure table has been narrowed:
four columns get the full record, two get state and number, and anything
narrower is left alone rather than raising.

`_fill_licensure` is pinned on the class surface by
`tests/test_stage6_import_surface.py`, and `tests/test_stage6_unrendered_recovery.py`
calls it on an instance. The mixin keeps it resolving through the MRO, which is
what that guard checks.
"""
import logging
import re
from collections.abc import Mapping, MutableMapping, Sequence
from dataclasses import dataclass, replace
from typing import Any

try:
    from docx.table import Table
except ImportError as exc:
    raise ImportError(
        "python-docx is required for stage 6. Install with: pip install python-docx lxml"
    ) from exc

from ..formatting import _clear_table_data, _set_font, format_date_for_section
from ..normalization import CAT_DEA, WithheldItem
from ..sorting import sort_entries_reverse_chronological

logger = logging.getLogger(__name__)

# NPI / DEA detection (#573). A label is the acronym at a word start and
# not running on into letters: the old substring test
# ('DEA' in text.upper()) consumed any licence row whose text mentioned
# "Dean". A trailing DIGIT is allowed because extraction sometimes fuses
# the label into the number itself ("NPI15180546000" on corpus CV HU4DXA),
# and a plain \b would reject exactly those. Shape tests are a tiebreak
# only -- see _classify_licensure_entry.
_NPI_LABEL_RE = re.compile(r'\bNPI(?![A-Za-z])', re.IGNORECASE)
# The DEA label is the acronym OR the agency's full name (#1217): a row whose
# first column spells out "Drug Enforcement Administration" (or the common
# misnomer "Agency") carries no "DEA" token at all, so the acronym test alone
# let its number render as an ordinary licence row.
_DEA_LABEL_RE = re.compile(
    r'\bDEA(?![A-Za-z])|\bDrug\s+Enforcement\s+(?:Administration|Agency)\b',
    re.IGNORECASE)
_NPI_SHAPE_RE = re.compile(r'^\d{10,11}$')

# Two DEA shapes, deliberately different widths, because the two callers
# fail in OPPOSITE directions (#821 R3 F-A).
#
# `_DEA_NUMBER_RE` is the real format -- two letters then SEVEN DIGITS (the
# first letter is the registrant type, the second the registrant's surname
# initial, the last digit a checksum). It is the only one of the two that
# may drive a DELETION: `_strip_dea_lines` removes a whole line of an
# entry's raw text on it, so a predicate wider than the thing it is named
# for silently destroys content -- the loose shape below matches every
# nine-letter English word ("Wisconsin", "Certified", "Emergency"), which
# took a fused entry's real state-licence lines out of the #221 recovery
# pass (#821 R3 F-A, found by the round-2 verifier).
#
# `_DEA_LIKE_SHAPE_RE` is that same shape widened to any alphanumeric tail,
# and it stays wide on purpose: its ONLY use is the classifier's tiebreak
# (`_classify_licensure_entry`, reached only for an entry that carries no
# NPI/DEA label at all, names no state, and DOES carry a `license_number`),
# where a match WITHHOLDS the entry instead of rendering it. Failing safe
# there means over-matching -- a DEA number mis-transcribed with a letter
# for a digit ("AB123456O") must still be withheld, and the cost of a false
# positive is one licence row withheld with a notice, not a leak. Corpus
# check before splitting the two (both farms, 196 F1 entries): zero
# `license_number` values match the loose shape but not the tight one, so
# the split changes no rendered output today.
_DEA_NUMBER_SHAPE = r'[A-Za-z]{2}\d{7}'
_DEA_NUMBER_RE = re.compile(rf'^{_DEA_NUMBER_SHAPE}$')
_DEA_LIKE_SHAPE_RE = re.compile(r'^[A-Za-z]{2}[A-Za-z0-9]{7}$')

# A DEA-number token inside a fused `license_number` cell (#1217), cut out by
# `_strip_dea_number_tokens`. The tight shape, as for `_strip_dea_lines`: this
# predicate deletes, so it must not match a nine-letter word. A leading '#' (a
# "#BP1234567" the source wrote that way) goes with the token.
_FUSED_DEA_TOKEN_RE = re.compile(
    rf'(?<![A-Za-z0-9])#?{_DEA_NUMBER_SHAPE}(?![A-Za-z0-9])')
_NUMBER_SEPARATOR_RE = re.compile(r'[\s,;/]+')
_SEPARATOR_RUN_RE = re.compile(r'(?:\s*[,;/]\s*){2,}')

# Classification results for one F1 entry.
KIND_LICENSE = 'license'
KIND_NPI = 'npi'
KIND_DEA = 'dea'

# How much of an unstructured entry's raw text becomes the state cell when
# there is no jurisdiction and no number to render instead.
_FALLBACK_TEXT_CHARS = 100


@dataclass(frozen=True)
class LicensureEntry:
    """One F1 entry with the stage-4 field aliases already resolved.

    Built only by `_normalize_licensure_entry`. Nothing downstream reads
    `extracted_fields`, so there is one place -- and one place only -- where
    'state_country' vs 'state' vs 'jurisdiction' is decided.
    """

    state: str = ''
    number: str = ''
    issue_date: str = ''
    expiration_date: str = ''
    license_type: str = ''
    original_text: str = ''


@dataclass(frozen=True)
class LicenseRecord:
    """One row of the licensure table, dates already formatted for F1.

    `last_registration_date` is named for the WCM template column it fills;
    stage 4's extraction schema has no separate field for it, so the closest
    available value -- `expiration_date` -- is what gets put there.
    """

    state: str = ''
    number: str = ''
    issue_date: str = ''
    last_registration_date: str = ''


@dataclass(frozen=True)
class IdentifierSet:
    """The two single-valued identifier slots of the DEA/NPI table.

    `None` means no candidate was seen at all; it is written to the document
    as an empty string, which blanks any stale template value.

    `dea` is ALWAYS `None` (#821): every entry `_resolve_licensure`
    classifies as DEA is withheld, never claimed into this slot -- see
    `LicensureResult.dea_withheld`.
    """

    dea: str | None = None
    npi: str | None = None


@dataclass(frozen=True)
class LicensureResult:
    """Everything section F1 renders, decided without touching a document.

    `dea_withheld` is #821's decision surfacing out of a pure function: True
    when at least one entry classified as DEA. `_fill_licensure` is the one
    that touches `self._pii_result` and the document, so it is the one that
    records the `WithheldItem` and skips writing a value -- this field is
    what tells it to.
    """

    licenses: tuple[LicenseRecord, ...]
    identifiers: IdentifierSet
    dea_withheld: bool = False


def _classify_licensure_entry(state: str, license_number: str,
                              license_type: str, original_text: str) -> str:
    """Classify one F1 entry as KIND_NPI, KIND_DEA or KIND_LICENSE.

    Routes on the label first: the `license_type` field when stage 4
    extracted one, else an explicit NPI/DEA word in the raw text. The
    number-shape tests run only for unlabelled entries that name no state,
    because a 10-11 digit state licence number is indistinguishable from an
    NPI by shape alone -- shape-first classification consumed real state
    licences (#573).
    """
    label = str(license_type or '').lower()
    text = str(original_text or '')
    # Bounded the same way as the raw-text checks below (#658), same reason:
    # a bare `'npi'`/`'dea' in label` substring test matches the acronym
    # embedded in an unrelated word (e.g. a `license_type` of "Idea for
    # renewal"), the same class of false positive the raw-text
    # `_NPI_LABEL_RE`/`_DEA_LABEL_RE` fix (#573) exists to prevent for
    # "Dean".
    if _NPI_LABEL_RE.search(label):
        return KIND_NPI
    if _DEA_LABEL_RE.search(label):
        return KIND_DEA
    if _NPI_LABEL_RE.search(text):
        return KIND_NPI
    if _DEA_LABEL_RE.search(text) or _DEA_LABEL_RE.search(str(state or '')):
        return KIND_DEA
    if not license_number:
        return KIND_LICENSE
    if not state:
        if _NPI_SHAPE_RE.match(license_number):
            return KIND_NPI
        if _DEA_LIKE_SHAPE_RE.match(license_number):
            return KIND_DEA
    return KIND_LICENSE


def _claim_identifier_slot(slot_name: str, current: str | None,
                           candidate: str, original_text: str) -> str:
    """Fill the single-valued NPI or DEA slot.

    When a second, different candidate arrives, keep the first (entries are
    already sorted most-recent-first) and warn instead of silently
    overwriting (#573)."""
    if current and current != candidate:
        logger.warning(
            "F1: second %s candidate %r (entry text %r) ignored; keeping %r",
            slot_name, candidate, original_text[:40], current)
        return current
    return candidate


def _normalize_licensure_entry(entry: Mapping[str, Any]) -> LicensureEntry:
    """Resolve one raw stage-4 F1 dict into a typed record.

    The `or` chains are load-bearing and are deliberately not
    `.get(key, default)`: stage 4 writes an alias as an empty string about as
    often as it omits the key, and only an `or` chain falls through on both.
    Swapping in a `.get` default would keep the empty string and stop the
    next alias ever being consulted.

    `extracted_fields` is sometimes explicitly `None` rather than absent,
    which a bare `.get('extracted_fields', {})` does not cover.
    """
    fields = entry.get('extracted_fields', {}) or {}
    return LicensureEntry(
        state=(fields.get('state_country') or fields.get('state')
               or fields.get('jurisdiction') or ''),
        number=str(fields.get('license_number')
                   or fields.get('number') or '').strip(),
        issue_date=fields.get('issue_date') or fields.get('date') or '',
        expiration_date=fields.get('expiration_date') or '',
        license_type=fields.get('license_type') or '',
        original_text=str(entry.get('text') or ''),
    )


def _license_record(entry: LicensureEntry) -> LicenseRecord | None:
    """The table row for an entry already classified as a licence.

    `None` means the entry contributes no row at all -- it named no
    jurisdiction, no number and carried no raw text.

    The raw-text fallback runs no NPI/DEA label re-check:
    `_classify_licensure_entry` already tests `original_text` for both labels
    before returning KIND_LICENSE, so reaching this branch already proves
    neither label is present -- re-testing here would only duplicate the
    classifier's own answer (#573 review).
    """
    if entry.state or entry.number:
        return LicenseRecord(
            state=entry.state,
            number=entry.number,
            issue_date=(format_date_for_section(entry.issue_date, 'F1')
                        if entry.issue_date else ''),
            last_registration_date=(
                format_date_for_section(entry.expiration_date, 'F1')
                if entry.expiration_date else ''),
        )
    if entry.original_text:
        return LicenseRecord(state=entry.original_text[:_FALLBACK_TEXT_CHARS])
    return None


def _strip_dea_lines(text: str) -> str:
    """Remove only the DEA label/value line(s) of a (possibly fused) F1
    entry's raw text, keeping every other line intact (#821 R2 F4).

    `_resolve_licensure` used to blank a DEA-classified entry's `text`
    entirely; that also erased any real state-licence record lines fused
    into the SAME entry (one entry, classified once, by label or shape
    anywhere in its text -- see `_classify_licensure_entry`), which took
    them out of reach of `generate()`'s #221 post-render recovery pass too,
    since that pass reads `entry['text']` directly. A line is stripped when
    it names DEA (`_DEA_LABEL_RE`) or carries a token that is a DEA NUMBER
    (`_DEA_NUMBER_RE`: two letters then seven digits, matched per
    alphanumeric token so it cannot also consume a same-length state
    licence number sitting on the SAME line as unrelated content).

    The token test is the tight shape, never the classifier's deliberately
    loose `_DEA_LIKE_SHAPE_RE` (#821 R3 F-A): that one matches any nine
    alphanumerics opening with two letters, i.e. every nine-letter English
    word, so a fused entry's "Wisconsin ... licence" or "Certified ..."
    line was dropped here as if it were a DEA registration -- and with it
    every licence number on that line, silently, exactly the content loss
    this function was written to stop. See the two regexes' own comment for
    why the classifier keeps the wide one."""
    kept = []
    for line in str(text or '').split('\n'):
        if _DEA_LABEL_RE.search(line):
            continue
        if any(_DEA_NUMBER_RE.match(tok)
               for tok in re.findall(r'[A-Za-z0-9]+', line)):
            continue
        kept.append(line)
    return '\n'.join(kept)


def _strip_dea_number_tokens(number: str) -> str:
    """Cut DEA-number tokens out of a fused, multi-token licence number cell.

    The reader sometimes leaves several stacked credentials' numbers in ONE
    licence entry's number cell ("<licence>, <licence>, <npi>, <dea>"); the
    DEA registration that rode along is not the entry's own, so
    `_classify_licensure_entry` (one verdict per entry) cannot see it (#1217).
    Only the tight DEA shape is cut (`_FUSED_DEA_TOKEN_RE`), and only from a
    cell of two or more tokens: a lone token beside a stated jurisdiction is
    still a state licence by the classifier's own rule, since a licence
    number can share the shape. Returns `number` unchanged when no token
    matched.
    """
    if len(_NUMBER_SEPARATOR_RE.split(number.strip())) < 2:
        return number
    if not _FUSED_DEA_TOKEN_RE.search(number):
        return number
    cut = _SEPARATOR_RUN_RE.sub(', ', _FUSED_DEA_TOKEN_RE.sub('', number))
    return cut.strip(' ,;/')


def _resolve_licensure(
        entries: Sequence[MutableMapping[str, Any]]) -> LicensureResult:
    """Decide the whole of section F1 without touching a document.

    Sorts, normalizes, classifies, routes each entry to the licence table or
    to one of the two identifier slots, and formats the dates. Takes no
    `docx` object and no stats counter, so a change to Word table handling
    cannot reach classification and vice versa (#624 review).

    A DEA-classified entry never reaches `identifiers.dea` (#821: withheld,
    never rendered) -- `dea_withheld` records that it happened instead, by
    classification kind rather than by re-detecting protected data with a
    second vocabulary: `_classify_licensure_entry` already tells DEA apart
    from NPI and from an ordinary licence, by label OR by number shape, and
    that is the only place in the pipeline that can (`stage6/normalization/
    pii.py`'s policy table intentionally does not reach an F1 entry's text --
    see that row's own comment). `_claim_identifier_slot`'s single-valued,
    keep-first-and-warn dedup does not apply to DEA any more: nothing is
    ever claimed, so there is nothing to conflict.

    A DEA-classified `raw` entry's `text` has its DEA line(s) stripped
    (`_strip_dea_lines`) in place, the one exception to "no side effects"
    this function has: `generate()`'s #221 post-render recovery
    (`_recover_unrendered_records`) re-scans every entry's raw text line by
    line afterwards and re-inserts any line it cannot verify rendered,
    reading `entry['text']` directly rather than this section's own output
    -- exactly the class of leak `stage6/pii_pass.py` closes for every OTHER
    category by mutating `entry['text']` in place before anything downstream
    can read it (see that module's docstring). A single-line DEA entry never
    reaches that recovery pass at all (it requires >= 2 fused record lines
    to even be considered), but a DEA number fused into a multi-line F1
    entry alongside real licence lines does, so the DEA content is stripped
    here too, at the only point that knows an entry was classified as
    DEA -- narrowly, so the fused licence lines survive for that pass to
    find, rather than being discarded along with the DEA line (#821 R2 F4:
    a whole-text blank silently dropped a fused entry's real licence lines
    from #221 recovery).
    """
    npi_number: str | None = None
    dea_withheld = False
    licenses: list[LicenseRecord] = []

    for raw in sort_entries_reverse_chronological(entries):
        entry = _normalize_licensure_entry(raw)

        # NPI/DEA vs state licence: label first, shape as tiebreak (#573)
        kind = _classify_licensure_entry(entry.state, entry.number,
                                         entry.license_type,
                                         entry.original_text)
        if kind == KIND_NPI:
            npi_number = _claim_identifier_slot(
                'NPI', npi_number, entry.number, entry.original_text)
            continue
        if kind == KIND_DEA:
            dea_withheld = True
            raw['text'] = _strip_dea_lines(entry.original_text)
            continue

        number = _strip_dea_number_tokens(entry.number)
        if number != entry.number:
            # A DEA number fused into this licence's number cell (#1217):
            # the entry is a licence, the token is not. Withhold it with
            # the same notice, and out of the raw text the #221 recovery
            # pass re-reads, exactly as for a DEA-classified entry.
            dea_withheld = True
            raw['text'] = _strip_dea_lines(entry.original_text)
            entry = replace(entry, number=number)

        record = _license_record(entry)
        if record is not None:
            licenses.append(record)

    return LicensureResult(
        licenses=tuple(licenses),
        identifiers=IdentifierSet(dea=None, npi=npi_number),
        dea_withheld=dea_withheld,
    )


#: The section name the #821 Word comment names a withheld DEA number
#: under -- the same vocabulary `stage6/pii_pass.py`'s `_section_label`
#: would produce for code F1 (`TAXONOMY_TO_SECTION['F1'] = 'licensure'`,
#: title-cased), named here directly rather than imported: this package is
#: imported BY `stage_6_word_template`, not the reverse (see `pii_pass.py`'s
#: own module docstring on the same back-edge constraint).
_LICENSURE_SECTION_LABEL = "Licensure"

# The WCM template's own licensure header row (`key_files/
# wcm_cv_template_faculty_october_2022_final.docx`), normalized. On the
# zero-entry path `_fill_licensure` has no data-driven signal that
# `_find_table_after_paragraph` landed on ITS table rather than some other
# one a template variant placed right after the same heading (#862's
# positive-shape guard: every forward table scan needs one) -- the header
# row is the only thing left to check.
_LICENSURE_HEADER_CELLS = (
    'state', 'number', 'date of issue (mm/dd/yyyy)',
    'date of last registration (mm/dd/yyyy) – (mm/dd/yyyy)',
)


def _is_licensure_table(table: Table) -> bool:
    """True when `table`'s own header row is the licensure table's."""
    if not table.rows:
        return False
    cells = tuple(' '.join(c.text.split()).lower() for c in table.rows[0].cells)
    return cells == _LICENSURE_HEADER_CELLS[:len(cells)]


class LicensureSection:
    """Section F1 writers, mixed into `WCMTemplateGenerator`."""

    def _fill_licensure(self, entries: Sequence[Mapping[str, Any]]) -> None:
        """Fill F1. LICENSURE section.

        WCM template has table with: State | Number | Date of issue | Date of last registration
        Also fills DEA and NPI numbers in a separate table (Table 9).

        Rendering only: `_resolve_licensure` has already resolved the stage-4
        field aliases, routed NPI/DEA and formatted the dates, so nothing
        below reads `extracted_fields`.

        No entries still locates the licensure table and clears it (#862,
        same class as #708's board-certification fix): the WCM template
        ships a blank placeholder data row in this table, and this used to
        return before `_clear_table_data` ever ran, so that row survived
        into the delivered document on every CV with zero F1 entries. The
        clear now always runs once the table is found and its header row
        is confirmed as the licensure table's own (`_is_licensure_table`).
        There is nothing to resolve or fill from an empty entry list, so
        the function still returns right after -- the DEA/NPI table (a
        separate template table `_fill_dea_npi` writes) is untouched on
        this path, exactly as before.
        """

        if self.verbose and entries:
            logger.info("Filling Licensure (%s entries)...", len(entries))

        # Find Licensure section
        section_idx = self._find_header_paragraph("Licensure")
        if section_idx is None:
            section_idx = self._find_header_paragraph("LICENSURE")
        if section_idx is None:
            return

        table = self._find_table_after_paragraph(section_idx)
        if not table:
            return

        if not entries and not _is_licensure_table(table):
            # A forward paragraph scan, not a table-identity lookup -- on a
            # template variant this could land on a table that only
            # happens to sit after the same heading. With no entries there
            # is no data-driven signal to catch that, so refuse rather
            # than clear someone else's table (#862).
            return

        _clear_table_data(table, keep_header=True)
        if not entries:
            return
        self.stats['tables_populated'] += 1

        result = _resolve_licensure(entries)

        # Fill regular licenses table
        for record in result.licenses:
            if self._write_license_row(table, record):
                self.stats['entries_inserted'] += 1

        # #821: a DEA number is withheld with notice, never written --
        # `result.identifiers.dea` is already None (`_resolve_licensure`
        # never claims one), so this only records the fact on the SAME
        # list the document-wide notice paragraph and Word comment read
        # (`self._pii_result`, populated earlier in `generate()` by
        # `run_pii_pass` for every other category -- see that row's
        # comment in `stage6/normalization/pii.py` for why this one is
        # recorded here instead).
        if result.dea_withheld:
            self._pii_result.withheld.append(
                WithheldItem(CAT_DEA, _LICENSURE_SECTION_LABEL, None))

        # Fill DEA/NPI table (Table 9 in template)
        self._fill_dea_npi(result.identifiers.dea, result.identifiers.npi)

    def _write_license_row(self, table, record: LicenseRecord) -> bool:
        """Append one licence row. False means nothing was written.

        Copes with a template whose licensure table has been narrowed: four
        columns get the full record, two get state and number, and anything
        narrower is left alone rather than raising -- and is not counted as
        inserted, because nothing reached the document.

        `table` is the caller's `docx` table and is deliberately left
        unannotated: no module under `stage6/` imports `docx`, and annotating
        it would be the first one to.
        """
        num_cols = len(table.columns)
        if num_cols < 2:
            return False

        row = table.add_row()
        if num_cols >= 4:
            row.cells[0].text = record.state or ''
            row.cells[1].text = record.number or ''
            row.cells[2].text = record.issue_date or ''
            row.cells[3].text = record.last_registration_date or ''
        else:
            row.cells[0].text = record.state or ''
            row.cells[1].text = record.number or ''

        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    _set_font(run)
        return True

    def _fill_dea_npi(self, dea_number: str | None,
                      npi_number: str | None) -> None:
        """Fill DEA and NPI numbers in their dedicated table.

        The WCM template has a 2-row table:
        Row 0: DEA number: (optional) | [value]
        Row 1: NPI number: (optional) | [value]

        Always finds the table and sets both cells (blanking whichever value
        is absent) rather than returning early when both are empty, so a
        second call on a reused document/generator can't leave a stale
        value from an earlier call in place.
        """
        # Find the DEA/NPI table by looking for a table containing "DEA number"
        dea_npi_table = None
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if 'dea number' in cell.text.lower():
                        dea_npi_table = table
                        break
                if dea_npi_table:
                    break
            if dea_npi_table:
                break

        if not dea_npi_table:
            return

        # Fill in the values
        for row in dea_npi_table.rows:
            if len(row.cells) >= 2:
                label = row.cells[0].text.lower()
                # Bounded the same way as `_classify_licensure_entry` (#658)
                # -- a template cell relabelled from its current "DEA
                # Number"/"NPI Number" text should not risk matching on an
                # embedded substring.
                if _DEA_LABEL_RE.search(label):
                    row.cells[1].text = dea_number or ''
                    for para in row.cells[1].paragraphs:
                        for run in para.runs:
                            _set_font(run)
                elif _NPI_LABEL_RE.search(label):
                    row.cells[1].text = npi_number or ''
                    for para in row.cells[1].paragraphs:
                        for run in para.runs:
                            _set_font(run)
