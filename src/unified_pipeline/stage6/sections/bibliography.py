"""Section S: bibliography -- publications S1 through S9 (#398).

Nine subsections, one loop. Each taxonomy code maps to a template header string
that must match the official WCM template exactly (they are matched by text, not
by index), numbering restarts at 1 under every header, and entries are inserted
before the paragraph that follows the header so the list builds downward.

Two things happen to the entry list before it is numbered:

- `sort_entries_reverse_chronological` puts newest first.
- `split_fused_citation_entries` (#208) undoes upstream fusion. Several
  citations that arrived collapsed into one entry would otherwise render as a
  single number with `<w:br/>` continuation lines -- visually a list, but with
  one citation numbered and the rest not.

The two `_add_citation_with_bold_author*` helpers exist because the CV owner's
name has to be bold inside the citation, and Word offers no way to say that
without splitting the text into runs. The author-matching and citation-splitting
rule lives once, in `_citation_author_split` (#572); the writers are thin tails
over two different substrates:

- `_add_citation_with_bold_author` writes plain runs via python-docx, and is
  what a non-enriched citation uses.
- `_add_citation_with_bold_author_as_insertion` builds the same runs as raw
  `w:r` elements inside a `w:ins`, because an enriched citation is rendered as a
  tracked insertion paired with a deletion of the original text, and python-docx
  cannot add a run *into* a revision element. It falls back to the plain path if
  the XML build raises. Raw control characters, the one input class known to
  make the XML build raise, are stripped inside the tracked-insertion writer
  itself (#552) -- and the plain writer now sanitizes too (#711), so all
  four run-text writes (the two here plus stage_6_word_template.py's
  `_add_track_change_insertion` / `_add_track_change_deletion`) reject the
  same input class the same way instead of raising.

Author bolding targets `target_name` from `_format_citation`, falling back to
the CV owner's last name -- taken from `cv_owner`, or recovered from the
document uid when the pipeline never resolved an owner -- and last to the
plausible spelling `normalization/owner_alias.py` infers from the citations
none of those names matched (#1393).
"""
import inspect
import logging
import re
import unicodedata
from collections import Counter
from collections.abc import Callable, Mapping
from dataclasses import dataclass, replace
from datetime import datetime
from difflib import SequenceMatcher
from typing import Any

try:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
except ImportError as exc:
    raise ImportError(
        "python-docx is required for stage 6. Install with: pip install python-docx lxml"
    ) from exc

from unified_pipeline.core.text_norm import is_placeholder_title

from ..formatting import _format_citation, _set_font
from ..normalization import split_fused_citation_entries
from ..normalization.owner_alias import (
    OwnerAlias,
    author_segment,
    fold_name,
    infer_owner_alias,
    name_parts,
)
from ..normalization.publication import resolve_publication
from ..parsing import _extract_last_name_from_uid, _strip_appended_initials
from ..sorting import sort_entries_reverse_chronological

logger = logging.getLogger(__name__)

# Tracked-insertion citation runs are built as raw w:r XML -- python-docx has
# no API to add a run *into* a w:ins element, so this path cannot call
# _set_font() like every other run in this file (#625). A prior fix (#625
# round 2) named these two constants so a diverging value would at least be
# visible at both call sites; that only holds if someone remembers to update
# both by hand. #662 item 4 closes the actual gap: read straight out of
# _set_font's own default parameters, so a policy change in
# stage6/formatting/docx.py reaches this path automatically instead of
# silently drifting out of step. w:sz is expressed in half-points.
_SET_FONT_DEFAULTS = inspect.signature(_set_font).parameters
TRACKED_INSERTION_FONT_NAME = _SET_FONT_DEFAULTS['name'].default
TRACKED_INSERTION_FONT_SIZE_PT = _SET_FONT_DEFAULTS['size'].default
TRACKED_INSERTION_FONT_SIZE_HALF_POINTS = str(TRACKED_INSERTION_FONT_SIZE_PT * 2)

# #552: raw `w:t` / `w:delText` assignment is lxml's own `.text` setter, not
# python-docx's `Run.text` -- it raises ValueError on the same control-code
# range python-docx itself rejects ("All strings must be XML compatible: ...
# no NULL bytes or control characters"), for any of the three run-text
# writes in stage 6 that build revision XML by hand (the other two are
# stage_6_word_template.py's `_add_track_change_insertion` /
# `_add_track_change_deletion`; this module's own is `create_run_element`
# below). Source .docx text cannot carry these codepoints (lxml rejects them
# at parse time), so the exposure is LLM-written fields -- a stage-4.5/5c/5d
# field reaching a citation or a tracked-change run. `\t`, `\n` and `\r` are
# valid XML characters and must NOT be stripped: `_clean_inline_tabs`'s
# label/value contract (test_cell_separators.py:34-37) depends on tab
# characters surviving into rendered text.
_CONTROL_CHAR_PATTERN = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _enrichment_field_text(value: Any) -> str:
    """Coerce one raw enrichment field to plain text.

    A missing key already becomes '' via `dict.get`'s own default before
    this is ever called; this covers the other two cases: an explicit
    ``None`` (some stage-5 writers set the key rather than omit it), and a
    field that came back as something other than a string. Never raises --
    stringifying an unexpected shape keeps the citation rendering instead of
    aborting it."""
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    return str(value)


@dataclass
class _CitationEnrichment:
    """One bibliography entry's enrichment fields, the way this renderer
    actually reads them (review thread 3850155382): `enrichment_status`,
    `enrichment_source`, and the original `text` a track-changes deletion
    needs, read together at one boundary instead of three separate
    `pub.get(...)` calls scattered through the loop body below. Mirrors
    `_CommitteeRecord` in `administrative_activities.py` from this same PR
    round -- a dataclass local to this renderer, not a stage6-wide
    convention.

    `from_raw` never raises: an unfamiliar `pub` shape, or a field that came
    back as something other than a string, degrades to the empty-string
    default for that field rather than aborting the citation."""
    enrichment_status: str = ''
    enrichment_source: str = ''
    text: str = ''
    in_press_note: str = ''
    in_press_superseded: bool = False

    @classmethod
    def from_raw(cls, raw) -> _CitationEnrichment:
        """Build a record from one raw stage-5 publication entry."""
        if not isinstance(raw, Mapping):
            return cls()
        return cls(
            enrichment_status=_enrichment_field_text(raw.get('enrichment_status')),
            enrichment_source=_enrichment_field_text(raw.get('enrichment_source')),
            text=_enrichment_field_text(raw.get('text')),
            in_press_note=_enrichment_field_text(raw.get('in_press_note')),
            in_press_superseded=raw.get('in_press_superseded') is True,
        )


# One author-token matcher for both bolding paths (#662 item 6). A name is a
# run of tokens joined by any run of whitespace or hyphen; apostrophes and
# hyphens each accept their typographic variants; and the match is a whole
# author token, not a substring: it may not start or end inside a longer name
# ("Wu" in "Wuertz", "Diaz" in "Alvarez-Diaz", "Neil" in "O'Neil").
_NAME_APOSTROPHES = "'\u2019\u02bc"
_NAME_HYPHENS = r"\-\u2010\u2011\u2013"
_NAME_SEPARATOR = rf"[\s{_NAME_HYPHENS}]+"
_NAME_CHAR = r"[\w\u0300-\u036f]"
_NAME_JOIN = rf"[{_NAME_APOSTROPHES}{_NAME_HYPHENS}]"
_NAME_START_GUARD = rf"(?<!{_NAME_CHAR})(?<!{_NAME_CHAR}{_NAME_JOIN})"
# A possessive ("Marrow's") still ends the name.
_NAME_END_GUARD = rf"(?!{_NAME_CHAR}|(?![{_NAME_APOSTROPHES}]s(?!\w)){_NAME_JOIN}\w)"
_MAX_INITIALS = 3
_SURNAME_PARTICLES = (
    "van", "von", "der", "den", "de", "la", "le", "da", "di", "del", "della",
    "du", "dos", "das", "do", "bin", "ibn", "al", "el", "ter", "ten", "zu", "st",
)
# A run of particles that opens an author: at the start of the citation or
# right after a list separator, so a first name such as "Al" mid-author never
# gets pulled into the bold run.
_PARTICLES_BEFORE = re.compile(
    r"(?:^|(?<=[,;&])\s*|(?<=\band)\s+)"
    rf"((?:(?:{'|'.join(_SURNAME_PARTICLES)})\.?\s+)+)$",
    re.IGNORECASE,
)
_WHITESPACE_RUN = re.compile(r"\s*")
# #1393: a name word right after / right before a matched surname, in the same
# author (whitespace only between), for widening a spaced compound.
_NAME_WORD = rf"[^\W\d_]{_NAME_CHAR}*(?:[{_NAME_APOSTROPHES}]{_NAME_CHAR}+)?"
_OWNER_PART_AFTER = re.compile(rf"[ \t]+({_NAME_WORD})(?!{_NAME_CHAR}|{_NAME_JOIN})")
_OWNER_PART_BEFORE = re.compile(rf"(?<!{_NAME_CHAR})(?<!{_NAME_JOIN})({_NAME_WORD})[ \t]+$")
# The first initial after a surname ("Ruiz MJ", "Ruiz, M. J."), and a given
# name opening the same author before it ("Mateo Ruiz", "M. E. Ruiz").
_INITIAL_AFTER = re.compile(rf"\s*,?\s*([A-Z])\.?(?:\s?[A-Z]\.?){{0,2}}(?!{_NAME_CHAR})")
_GIVEN_NAME_BEFORE = re.compile(
    rf"(?:^\s*(?:\d+[.)]\s*)?|[,;&]\s*|\band\s+)({_NAME_WORD})\.?(?:\s+[A-Z]\.?)*\s+$")


def _name_pattern(name: str) -> str | None:
    """Regex for `name` as one whole author token, or None if it is blank.

    NFC and NFD spellings are both accepted, so a decomposed accent in the
    citation still matches a composed one in the owner's name."""
    forms = []
    for form in dict.fromkeys(
            (unicodedata.normalize('NFC', name), unicodedata.normalize('NFD', name))):
        tokens = [t for t in re.split(_NAME_SEPARATOR, form.strip()) if t]
        if not tokens:
            return None
        forms.append(_NAME_SEPARATOR.join(
            ''.join(f"[{_NAME_APOSTROPHES}]" if ch in _NAME_APOSTROPHES else re.escape(ch)
                    for ch in token)
            for token in tokens))
    return f"{_NAME_START_GUARD}(?:{'|'.join(forms)}){_NAME_END_GUARD}"


def _find_name_span(citation: str, name: str) -> tuple[int, int] | None:
    """Span of the first whole-token, case-insensitive occurrence of `name`."""
    pattern = _name_pattern(name)
    match = re.search(pattern, citation, re.IGNORECASE) if pattern else None
    return match.span() if match else None


def _trailing_initials_end(citation: str, end: int) -> int:
    """End of the author token after its surname: `end` extended over up to
    `_MAX_INITIALS` capital initials ("Wende ME"), else `end` unchanged. The
    initials must be upper case and stand alone, so "Wende and" and "Wende
    Michael" stop at the surname."""
    start = _WHITESPACE_RUN.match(citation, end).end()
    stop = start
    while stop < len(citation) and citation[stop].isalpha() and citation[stop].isupper():
        stop += 1
    if not 0 < stop - start <= _MAX_INITIALS:
        return end
    if stop < len(citation) and re.match(_NAME_CHAR, citation[stop]):
        return end
    return stop


def _widen_over_owner_parts(citation: str, start: int, end: int,
                           owner_parts: frozenset[str]) -> tuple[int, int]:
    """`start, end` widened over the adjacent space-separated words of the
    same author that are owner name parts, so a spaced compound bolds whole
    ("Garza Ruiz M", not "Garza" or "Ruiz M") wherever it is matched (#1393)."""
    if not owner_parts:
        return start, end
    while (after := _OWNER_PART_AFTER.match(citation, end)) \
            and fold_name(after.group(1)) in owner_parts:
        end = after.end()
    while (before := _OWNER_PART_BEFORE.search(citation, 0, start)) \
            and fold_name(before.group(1)) in owner_parts:
        start = before.start(1)
    return start, end


def _find_surname_span(citation: str, surname: str,
                       owner_parts: frozenset[str] = frozenset()) -> tuple[int, int] | None:
    """Span of a surname plus trailing initials, widened over adjacent owner
    name parts and leftward over surname particles ("de la Cruz M")."""
    span = _find_name_span(citation, surname)
    if span is None:
        return None
    return _surname_span_around(citation, *span, owner_parts)


def _surname_span_around(citation: str, start: int, end: int,
                         owner_parts: frozenset[str]) -> tuple[int, int]:
    """A matched surname's `start:end` grown to the whole author token."""
    start, end = _widen_over_owner_parts(citation, start, end, owner_parts)
    particles = _PARTICLES_BEFORE.search(citation, 0, start)
    if particles:
        start = particles.start(1)
    return start, _trailing_initials_end(citation, end)


def _author_initial(citation: str, start: int, end: int) -> str:
    """The initial of the author whose surname spans `start:end`: the first
    trailing initial ("Ruiz M", "Ruiz, M."), else the first letter of a given
    name before it in the same author ("Mateo Ruiz"), else ''."""
    after = _INITIAL_AFTER.match(citation, end)
    if after:
        return after.group(1)
    before = _GIVEN_NAME_BEFORE.search(citation, 0, start)
    return before.group(1)[0].upper() if before else ''


def _find_alias_span(citation: str, alias: OwnerAlias) -> tuple[int, int] | None:
    """Span of the first alias match, in the author list, on an author whose
    initial is the owner's: an alias never bolds a co-author who shares the
    surname, or half of it, but not the initials (#1393)."""
    segment = author_segment(citation)
    segment_end = citation.find(segment) + len(segment)
    for surname in alias.surnames:
        pattern = _name_pattern(surname)
        for match in re.finditer(pattern, citation[:segment_end], re.IGNORECASE) if pattern else ():
            start, end = _widen_over_owner_parts(citation, *match.span(), alias.parts)
            if _author_initial(citation, start, end) in alias.initials:
                return _surname_span_around(citation, start, end, alias.parts)
    return None


def _target_surname(target_name: str) -> str:
    """`target_name` without its trailing capital initials ("Pell-Rowan FM"
    -> "Pell-Rowan", "Tarn-Ellery, K." -> "Tarn-Ellery"), or unchanged
    when it does not end in some. Stage 4 records the initials it saw in one
    citation form; another form of the same author may carry more or fewer.
    A "Surname, Given" form keeps what precedes the comma ("Wende, Michael
    E." -> "Wende", #1393); a bare trailing comma is not that form. Spaced
    initials all go ("Wende M. E." -> "Wende")."""
    surname, comma, given = target_name.partition(",")
    if comma and surname.strip() and given.strip():
        return surname.strip()
    parts = re.split(r"[\s,]+", target_name.strip())
    cut = len(parts)
    while cut > 1 and _is_initials_word(parts[cut - 1]):
        cut -= 1
    return " ".join(parts[:cut]) if cut < len(parts) else target_name


def _is_initials_word(word: str) -> bool:
    """One word of up to `_MAX_INITIALS` capitals, dotted or not ("ME",
    "M.", "M.E.") -- spaced initials ("I. I.") are several such words."""
    letters = word.replace(".", "")
    return 0 < len(letters) <= _MAX_INITIALS and letters.isalpha() and letters.isupper()


# The render-warnings sidecar check for citations whose owner bold differs
# from what stage 4's names alone give (`_warn_owner_bold_from_bibliography`).
OWNER_BOLD_FROM_BIBLIOGRAPHY_CHECK = "owner_bold_from_bibliography"

# A stage-4 target_name longer than this, or holding a colon, is not an author
# name (the corpus has a book title there) and is skipped for the owner surname.
_MAX_TARGET_NAME_TOKENS = 4
# A target surname this close (difflib ratio) to a known owner name part, with
# the same first letter and an owner initial, is the owner misspelt
# ("Wesssely S", "Ramanaujam N" in the corpus score 0.93-0.95; a transposition
# "Oprekso" 0.86), not a co-author (#1394).
MIN_TARGET_SIMILARITY = 0.8


def _looks_like_author_name(target_name: str | None) -> bool:
    """False for a blank target_name or one shaped like a title, not a name."""
    return bool(target_name) and ':' not in target_name \
        and len(target_name.split()) <= _MAX_TARGET_NAME_TOKENS


def _citation_author_split(citation: str, target_name: str | None,
                           cv_owner_last_name: str,
                           owner_alias: OwnerAlias | None = None) -> tuple[str, str, str]:
    """Split a citation around the author name that should render bold.

    The one home for the author-matching rule shared by the plain and the
    tracked-insertion citation writers (#572): a change to the rule must reach
    both writers, or enriched and non-enriched citations in the same
    bibliography bold different text.

    Both lookups use one matcher (#662 item 6): whole author token, case
    insensitive, Unicode aware. It rejects a match that starts or ends inside a
    longer name (including across a hyphen or apostrophe), treats hyphen and
    space as interchangeable inside a name and typographic apostrophes as
    equal, and accepts composed or decomposed accents.

    Prefers `target_name` when it appears in the citation as written; then
    tries its surname alone, so "Quill J" still finds "Quill JD"; otherwise
    falls back to `cv_owner_last_name`, and last to `owner_alias`, the
    spellings `_infer_bibliography_owner_alias` read from the whole
    bibliography (#1393) -- matched only in the author list and only on an
    author whose initial is the owner's. Every surname match is widened over
    adjacent owner name parts, so a spaced compound bolds whole. A
    `target_name` shaped like a title
    (a colon, or more than `_MAX_TARGET_NAME_TOKENS` words) is skipped. A surname match takes trailing capital
    initials, e.g. "Wende ME", "Wende M", and any leading particles ("de la
    Cruz M"). The comma form ("Wende, M") bolds only the surname, pinned by
    test_citation_author_split_additional_dimensions.

    Returns `(before, name_to_bold, after)`, where `name_to_bold` is the
    citation's own text. When nothing matches, `name_to_bold` is '' and the
    whole citation is in `before`.
    """
    alias = owner_alias or OwnerAlias()
    if not _looks_like_author_name(target_name) or not _target_is_owner(target_name, alias):
        target_name = None
    span = _find_name_span(citation, target_name) if target_name else None
    if span is not None:
        span = _widen_over_owner_parts(citation, *span, alias.parts)
    if span is None and target_name:
        span = _find_surname_span(citation, _target_surname(target_name), alias.parts)
    if span is None and cv_owner_last_name:
        span = _find_surname_span(citation, cv_owner_last_name, alias.parts)
    if span is None:
        span = _find_alias_span(citation, alias)
    if span is None:
        return citation, '', ''
    start, end = span
    return citation[:start], citation[start:end], citation[end:]


def _lacks_title_and_year(entry: dict) -> bool:
    """True when the entry has neither a year nor a title, a 5d bracketed
    placeholder counting as none (#446: EBYSBC HFAJCC-05, XWNZWW-02, and the
    author-only split heads of NDXXAD-02). Numbered as a citation it reads as
    a paper that does not exist."""
    pub = resolve_publication(entry)
    return not pub.year and (not pub.title.strip()
                             or is_placeholder_title(pub.title, entry.get('text') or ''))


def _resolve_uid_owner_surname(uid: str, publications: list[dict]) -> str:
    """The CV owner's surname as the citations spell it, from the document uid.

    A uid such as "2003_Quennevillejs_Cv" can carry the owner's initials glued to
    the surname, but "Smith" ends in the same lower-case tail, and no rule on
    the uid alone can tell the two apart (#665 item 1). So the uid's last name
    is used as written when any citation names that author, and is stripped of
    possible appended initials only when it is not found -- the citations are
    the evidence the uid cannot give.
    """
    raw = _extract_last_name_from_uid(uid)
    citations = (_format_citation(pub, 0)[0] for pub in publications)
    if not raw or any(_find_surname_span(c, raw) for c in citations):
        return raw
    return _strip_appended_initials(raw)


def _name_initial(name: str) -> str:
    """The given-name initial in a stage-4 name: after a comma ("Wende,
    Michael E."), the trailing initials ("Eil R"), or the first of several
    words ("Michael E. Wende"); '' for a lone word."""
    surname, comma, given = name.partition(",")
    if comma:
        return next((ch.upper() for ch in given if ch.isalpha()), '')
    words = name.split()
    if len(words) < 2:
        return ''
    if _target_surname(name) != name:
        return words[-1][0]
    return words[0][0].upper()


@dataclass(frozen=True)
class _OwnerAnchor:
    """What stage 4 says about the owner before the bibliography is read:
    the name parts of `cv_owner.last_name` and of the bibliography-majority
    target surname, and the owner's first initials (cv_owner first and full
    name, the majority target)."""
    parts: frozenset[str]
    initials: frozenset[str]

    def names_owner(self, target_name: str) -> bool:
        """A target sharing a surname part, or the first initial, with the
        owner -- the owner's other surname ("Married V" on a "Maiden" CV)
        agrees by initial."""
        return bool(name_parts(_target_surname(target_name)) & self.parts) \
            or _name_initial(target_name) in self.initials


def _majority_target(targets: list[str]) -> str:
    """The target_name whose surname more than half of `targets` carry, or ''."""
    folded = Counter(fold_name(_target_surname(t)) for t in targets)
    if not folded:
        return ''
    top, count = folded.most_common(1)[0]
    if count * 2 <= len(targets):
        return ''
    return next(t for t in targets if fold_name(_target_surname(t)) == top)


def _owner_anchor(targets: list[str], owner: Mapping) -> _OwnerAnchor:
    """The `_OwnerAnchor` of one CV, from its cv_owner and its targets."""
    majority = _majority_target(targets)
    parts = name_parts(str(owner.get('last_name') or '')) | name_parts(_target_surname(majority))
    initials = {str(owner.get('first_name') or '')[:1].upper(),
                _name_initial(str(owner.get('full_name') or '')), _name_initial(majority)}
    return _OwnerAnchor(frozenset(parts), frozenset(initials - {''}))


def _target_is_owner(target_name: str, alias: OwnerAlias) -> bool:
    """False for a `target_name` that is not the owner, once the bibliography
    has confirmed the owner (`alias.surnames`): its surname shares no part
    with a name known to be the owner's, and is not a close spelling of one
    (same first letter, `MIN_TARGET_SIMILARITY`) carried with an owner
    initial. Stage 4 set that citation's target to a co-author, and bolding
    it would mark the co-author as the owner (web207, web172; #1394). A
    misspelt owner ("Wesssely S") still counts. With nothing confirmed,
    every target is taken, as before."""
    if not alias.surnames:
        return True
    target_parts = name_parts(_target_surname(target_name))
    if target_parts & alias.known_parts:
        return True
    return _name_initial(target_name) in alias.initials and any(
        t[0] == k[0] and SequenceMatcher(None, t, k).ratio() >= MIN_TARGET_SIMILARITY
        for t in target_parts for k in alias.known_parts)


def _infer_bibliography_owner_alias(publications: list[dict], cv_owner: Mapping | None,
                                    cv_owner_last_name: str) -> OwnerAlias:
    """The owner's surname spellings, read from every citation in the
    bibliography (#1393). Stage 4's names are one input, not the answer: the
    last name, each `target_name` surname, and the owner's initials from the
    first name, full name and `target_name`."""
    owner = cv_owner if isinstance(cv_owner, Mapping) else {}
    citations: list[str] = []
    targets: list[str] = []
    for pub in publications:
        citation, target_name, _ = _format_citation(pub, 0)
        citations.append(citation)
        if _looks_like_author_name(target_name):
            targets.append(target_name)
    anchor = _owner_anchor(targets, owner)
    # Only the targets that name the owner: stage 4 sometimes sets one
    # citation's target to a co-author (web207), whose surname and initial
    # must not become the owner's. All of them when the anchor is empty.
    owner_targets = [t for t in targets if anchor.names_owner(t)] or targets
    surnames = [cv_owner_last_name] + [_target_surname(t) for t in owner_targets]
    initials = [*anchor.initials, *(_name_initial(t) for t in owner_targets)]
    alias = infer_owner_alias(citations, surnames, initials)
    alias = replace(alias, known_parts=alias.parts | anchor.parts | name_parts(cv_owner_last_name))
    if any(fold_name(s) != fold_name(cv_owner_last_name) for s in alias.surnames):
        logger.info("Owner surname read from the bibliography as %r (stage 4 last name %r)",
                    alias.surnames, cv_owner_last_name)
    return alias


class BibliographySection:
    """Section S writers, mixed into `WCMTemplateGenerator`."""

    @staticmethod
    def _sanitize_run_text(text: str) -> str:
        """Strip the control characters lxml's `.text` setter rejects from
        one run's text, before it reaches a raw `w:t`/`w:delText` element.

        The one sanitiser for stage 6's four run-text writes (#552, #711) --
        `WCMTemplateGenerator._add_track_change_insertion` and
        `_add_track_change_deletion` in stage_6_word_template.py reach this
        through the mixin (`self._sanitize_run_text`); this module's own
        `create_run_element`, below, and `_add_citation_with_bold_author`
        call it directly. `\\t`, `\\n` and `\\r` are valid XML and are left
        untouched.
        """
        return _CONTROL_CHAR_PATTERN.sub('', text)

    def _fill_bibliography(self, entries_by_code: dict[str, list[dict]], cv_owner: dict,
                           document_uid: str = '') -> list[dict]:
        """Fill bibliography section with formatted citations.

        Uses the official WCM template section headers and restarts numbering
        within each subsection. Returns the entries it declined as no citation
        (`_lacks_title_and_year`, #446), for `generate()` to send to the Appendix.
        """
        # Map taxonomy codes to WCM template section header text
        # These must match the exact text in the official WCM template
        section_headers = {
            'S1': 'Peer-reviewed Research Articles:',
            'S2': 'Reviews and Editorials:',
            'S3': 'Books:',
            'S4': 'Chapters:',
            'S5': 'Non-peer-reviewed Research Publications:',
            'S6': 'Case Reports',
            'S7': 'In review',
            'S8': 'Abstracts',
            'S9': 'Other (media, podcasts, etc.):',
        }

        pub_codes = ['S1', 'S2', 'S3', 'S4', 'S5', 'S6', 'S7', 'S8', 'S9']

        # Get CV owner last name for fallback bolding
        cv_owner_last_name = ''
        if cv_owner and cv_owner.get('last_name'):
            cv_owner_last_name = cv_owner['last_name']
        elif document_uid:
            # Fallback: extract from document_uid (e.g., "2015_Wende" -> "Wende")
            cv_owner_last_name = _resolve_uid_owner_surname(
                document_uid,
                [pub for code in pub_codes for pub in entries_by_code.get(code, [])])

        all_pubs = [pub for code in pub_codes for pub in entries_by_code.get(code, [])]
        owner_alias = _infer_bibliography_owner_alias(all_pubs, cv_owner, cv_owner_last_name)
        self._warn_owner_bold_from_bibliography(all_pubs, cv_owner_last_name, owner_alias)

        # Count total publications
        total_pubs = sum(len(entries_by_code.get(code, [])) for code in pub_codes)

        logger.info("Filling Bibliography (%d publications)...", total_pubs)

        declined: list[dict] = []
        # Process each publication type
        for code in pub_codes:
            pubs = entries_by_code.get(code, [])
            if not pubs:
                continue

            header_text = section_headers.get(code, '')
            if not header_text:
                continue

            # Find the section header in the template
            section_idx = self._find_template_label(header_text)
            if section_idx is None:
                logger.warning("Could not find section header %r", header_text)
                continue

            logger.info("%s: %d entries -> '%s...'", code, len(pubs), header_text[:30])

            # Sort reverse chronologically (most recent first)
            pubs_sorted = sort_entries_reverse_chronological(pubs)
            # Un-fuse any entry that collapsed several citations into one
            # (#208), so each is numbered instead of rendering as unnumbered
            # <w:br/> continuation lines under one number.
            pubs_sorted = split_fused_citation_entries(pubs_sorted)

            # Resolve the insertion anchor once. Every paragraph for this
            # section is inserted immediately before this one anchor object,
            # rather than by re-indexing self.doc.paragraphs (which re-walks
            # the document body on every access) after each insertion -- that
            # re-indexing is what made the old insert_idx pattern fragile.
            # A bibliography header is allowed to be the template's last
            # paragraph, in which case there is no following paragraph to
            # anchor on; new content is then appended to the end of the
            # document instead (#625).
            paragraphs = self.doc.paragraphs
            anchor = paragraphs[section_idx + 1] if section_idx + 1 < len(paragraphs) else None

            def _insert_before_anchor(text: str = "") -> Paragraph:
                if anchor is not None:
                    return anchor.insert_paragraph_before(text)
                return self.doc.add_paragraph(text)

            # Add blank line before first citation in each section
            if pubs_sorted:
                _insert_before_anchor("")

            citation_num = 0
            for pub in pubs_sorted:
                if _lacks_title_and_year(pub):
                    # #446: a split head, a section label or a 5d placeholder,
                    # not a citation; the Appendix keeps it, with a warning.
                    declined.append(pub)
                    continue
                enrichment = _CitationEnrichment.from_raw(pub)
                if enrichment.in_press_superseded:
                    # Takes no citation number: accepting the deletion leaves none.
                    self._insert_superseded_citation(_insert_before_anchor, enrichment)
                    continue
                citation_num += 1
                citation_text, target_name, enriched_fields = _format_citation(pub, citation_num)
                original_text = enrichment.text
                enrichment_status = enrichment.enrichment_status

                # Insert citation paragraph before the anchor
                para = _insert_before_anchor("")

                # Check if this entry was enriched - if so, use track changes
                if enrichment_status == 'enriched' and original_text:
                    # Show original as deleted, enriched citation as inserted
                    # First add the deletion (original text)
                    self._add_track_change_deletion(para, original_text, author="PubMed Enrichment")
                    # Then add the insertion (enriched citation) with bold author
                    self._add_citation_with_bold_author_as_insertion(
                        para, citation_text, target_name, cv_owner_last_name,
                        author="PubMed Enrichment", owner_alias=owner_alias
                    )
                else:
                    # No enrichment - add citation normally with target name bolded
                    self._add_citation_with_bold_author(para, citation_text, target_name,
                                                        cv_owner_last_name, owner_alias)

                # Add comment explaining enrichment (no inline text)
                if enriched_fields:
                    enrichment_source = enrichment.enrichment_source
                    if enrichment_source:
                        comment = f"Data enriched from {enrichment_source.upper()}. Fields updated: {', '.join(enriched_fields)}"
                        self._add_word_comment(para, comment, author="PubMed Enrichment")

                # Stage 5 found an "in press" paper published. This comment
                # is for the faculty member deciding on the tracked change,
                # not a pipeline comment, so it ignores `emit_comments`.
                if enrichment.in_press_note:
                    self._add_word_comment(para, enrichment.in_press_note,
                                           author="PubMed Enrichment", always=True)

                # Add comments from upstream pipeline processes
                self._add_entry_comments(para, pub)

                self.stats['entries_inserted'] += 1
        return declined

    def _insert_superseded_citation(self, insert_paragraph: Callable[[], Paragraph],
                                    enrichment: _CitationEnrichment) -> None:
        """An "in press" entry stage 5 found already listed as published:
        the CV's text as one tracked deletion carrying the note, so the faculty
        member accepts or rejects dropping it. With track changes off the
        accepted state is shown, which is no paragraph at all."""
        if not self.emit_track_changes or not enrichment.text:
            return
        para = insert_paragraph()
        self._add_track_change_deletion(para, enrichment.text, author="PubMed Enrichment")
        self._add_word_comment(para, enrichment.in_press_note,
                               author="PubMed Enrichment", always=True)

    def _warn_owner_bold_from_bibliography(self, publications: list[dict], cv_owner_last_name: str,
                                           owner_alias: OwnerAlias) -> None:
        """Record, as an INFO render warning, how many citations bold other
        text than stage 4's names alone would: the owner's spelling read
        from the bibliography, a compound widened, or a co-author target
        refused (#1393, #1394). Counts only, like the stat-style records
        (geographic-classification and reclassify failures). Some self-check
        records quote a paragraph, but this one would quote the owner's name,
        and the sidecar reaches the Teams card; the bolds are in the document."""
        changed = 0
        for pub in publications:
            citation, target_name, _ = _format_citation(pub, 0)
            changed += _citation_author_split(citation, target_name, cv_owner_last_name)[1] \
                != _citation_author_split(citation, target_name, cv_owner_last_name, owner_alias)[1]
        if not changed:
            return
        self._section_failures.append({
            "check": OWNER_BOLD_FROM_BIBLIOGRAPHY_CHECK,
            "code": None,
            "section": "bibliography",
            "message": (f"{changed} citation(s) bold the CV owner as the bibliography "
                        "spells the name, or refuse a co-author target, where stage 4's "
                        "names alone would bold otherwise"),
            "evidence": [f"{OWNER_BOLD_FROM_BIBLIOGRAPHY_CHECK}={changed}",
                         f"owner_spellings={len(owner_alias.surnames)}"],
            "severity": "INFO",
        })

    def _add_citation_with_bold_author(self, para: Paragraph, citation: str, target_name: str | None,
                                       cv_owner_last_name: str = '',
                                       owner_alias: OwnerAlias | None = None) -> None:
        """
        Add citation text to paragraph, bolding the target author name.

        If target_name is not found, falls back to searching for
        cv_owner_last_name, then owner_alias.
        """
        para.clear()

        # python-docx's `add_run`/`Run.text` reaches the same lxml `.text`
        # setter the raw-XML writers use and raises on the same control-code
        # range (#552). This writer is also the tracked writer's
        # `emit_track_changes=False` path and its exception fallback, so an
        # unsanitised control character here would still abort the citation.
        # Sanitise before the split, on both writers, so a control character in
        # the citation or the target name cannot make the plain and tracked
        # paths bold different text (#552 round 2).
        citation = self._sanitize_run_text(citation)
        if target_name:
            target_name = self._sanitize_run_text(target_name)

        before, name_to_bold, after = _citation_author_split(
            citation, target_name, cv_owner_last_name, owner_alias)

        if name_to_bold:
            # Add before (normal)
            if before:
                run1 = para.add_run(before)
                _set_font(run1)

            # Add target name (bold)
            run2 = para.add_run(name_to_bold)
            _set_font(run2, bold=True)
            self.stats['target_names_bolded'] += 1

            # Add after (normal)
            if after:
                run3 = para.add_run(after)
                _set_font(run3)
        else:
            # No target name to bold
            run = para.add_run(citation)
            _set_font(run)

    def _add_citation_with_bold_author_as_insertion(self, para: Paragraph, citation: str,
                                                     target_name: str | None, cv_owner_last_name: str = '',
                                                     author: str = "PubMed Enrichment",
                                                     owner_alias: OwnerAlias | None = None) -> None:
        """
        Add citation as a track change insertion, bolding the target author name.

        This creates proper Word track change structure with w:ins element,
        and includes bold formatting for the target author within the insertion.
        """
        # Sanitise before the split, on both writers, so a control character in
        # the citation or the target name cannot make the plain and tracked
        # paths bold different text (#552 round 2).
        citation = self._sanitize_run_text(citation)
        if target_name:
            target_name = self._sanitize_run_text(target_name)

        # Issue #153: when track changes are disabled, render the citation as a
        # plain (non-tracked) paragraph with the target author bolded.
        if not self.emit_track_changes:
            self._add_citation_with_bold_author(para, citation, target_name,
                                                cv_owner_last_name, owner_alias)
            return
        try:
            revision_id = str(self._revision_id)
            self._revision_id += 1

            # Create the insertion element
            ins = OxmlElement('w:ins')
            ins.set(qn('w:id'), revision_id)
            ins.set(qn('w:author'), author)
            ins.set(qn('w:date'), datetime.now().strftime('%Y-%m-%dT%H:%M:%SZ'))

            before, name_to_bold, after = _citation_author_split(
                citation, target_name, cv_owner_last_name, owner_alias)

            def create_run_element(text: str, bold: bool = False) -> Any:
                """Create a w:r element with text and optional bold."""
                run_elem = OxmlElement('w:r')
                rPr = OxmlElement('w:rPr')
                rFonts = OxmlElement('w:rFonts')
                rFonts.set(qn('w:ascii'), TRACKED_INSERTION_FONT_NAME)
                rFonts.set(qn('w:hAnsi'), TRACKED_INSERTION_FONT_NAME)
                rPr.append(rFonts)
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), TRACKED_INSERTION_FONT_SIZE_HALF_POINTS)
                rPr.append(sz)
                if bold:
                    b = OxmlElement('w:b')
                    rPr.append(b)
                run_elem.append(rPr)
                t = OxmlElement('w:t')
                # The xml:space guard reads the sanitized string, not the raw
                # one: stripping a control character can expose a leading or
                # trailing space that the raw text did not start or end with,
                # and without xml:space="preserve" Word collapses it.
                clean_text = self._sanitize_run_text(text)
                t.text = clean_text
                if clean_text.startswith(' ') or clean_text.endswith(' '):
                    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                run_elem.append(t)
                return run_elem

            if name_to_bold:
                # Add before (normal)
                if before:
                    ins.append(create_run_element(before, bold=False))

                # Add target name (bold)
                ins.append(create_run_element(name_to_bold, bold=True))

                # Add after (normal)
                if after:
                    ins.append(create_run_element(after, bold=False))
            else:
                # No target name to bold - single run
                ins.append(create_run_element(citation, bold=False))

            # Append insertion to paragraph. Stats are updated only now, once
            # the tracked-insertion path is known to have actually succeeded.
            # Incrementing target_names_bolded earlier (previously right
            # after the bold run was built, before the XML build finished)
            # meant an exception raised later here still left the increment
            # in place, and the except block's fallback to
            # _add_citation_with_bold_author counted the same citation a
            # second time (#625).
            para._p.append(ins)
            self.stats['track_changes_added'] += 1
            if name_to_bold:
                self.stats['target_names_bolded'] += 1

        except Exception as e:
            logger.warning("Could not add citation as insertion: %s", e)
            # Fall back to normal citation
            self._add_citation_with_bold_author(para, citation, target_name,
                                                cv_owner_last_name, owner_alias)
