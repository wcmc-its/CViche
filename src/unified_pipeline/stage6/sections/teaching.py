"""Section K: teaching activities, K1 through K5 (#398).

Five subsections, and the writer's only structural job is finding each one:

    K1  Didactic teaching
    K2  Clinical teaching
    K3  Administrative teaching
    K4  Continuing education and professional education
    K5  Other education / outreach activities

Each K-code carries a LIST of candidate header strings rather than one, because
the template wording has drifted across revisions ("Clinical teaching" vs
"bedside teaching", "outreach activities" vs "community education or patient").
A code whose headers all miss falls back to the parent "EDUCATIONAL
CONTRIBUTIONS" heading rather than dropping its entries.

Entries render as a FLAT bulleted list under the template's own header. The
source CV's sub-headings are deliberately not carried over: the WCM subsections
already are the taxonomy, and re-emitting the original hierarchy produced two
competing levels of grouping.

The one exception is INSIDE an entry (#423): a table row whose source cell is a
title followed by bullets renders the bullets one list level in. The reader
flattens the cell, so `SourceCellLevels` re-reads the level from the source
docx; see its docstring for how a row is found and when it is refused.

Insertion runs backwards. `_insert_bulleted_entry` inserts BEFORE the index it
is given, so the list is walked in reverse and each entry lands above the one
inserted before it, leaving reverse-chronological order on the page. That is
also why `is_first_visible` is computed as the LAST index of the reversed list
-- the entry that ends up visually first is the one that gets the blank line
above it.

WHAT HAPPENS WHEN TWO CODES SHARE ONE HEADING
Every K-code resolves its heading independently, so when several codes all miss
their own wording they all fall back to the SAME "EDUCATIONAL CONTRIBUTIONS"
paragraph. That is intended -- losing the entries would be worse -- and the
resulting page order is fixed, not incidental, so it is written down here
rather than left to be rediscovered:

- Each code's entries stay CONTIGUOUS. Two codes are never interleaved: a
  code's whole block is inserted before the next code is looked up.
- Within a block, order is reverse-chronological, newest first, exactly as
  under a dedicated heading.
- Blocks stack in REVERSE code order. Codes are processed K1..K5 in
  `TEACHING_SECTION_HEADERS` order and every block is inserted directly under
  the shared heading, pushing the previous block down, so the page reads K5,
  K4, K3, K2, K1 top to bottom.
- Each block still gets its own blank line above it, because
  `is_first_visible` is computed per code.

Against the shipped template this is dormant: all five K headings exist
(`key_files/wcm_cv_template_faculty_october_2022_final.docx` carries them at
paragraphs 85/87/89/91/93, under EDUCATIONAL CONTRIBUTIONS at 82), so no code
takes the fallback. It describes what a template revision would get.

DELIMITER CONTRACT
`|` is NEVER an item separator in this section. It joins fields inside one item
("Role | Institution"), so

    "Role A | Institution A\\tRole B | Institution B"

reconstructs as TWO items, not four. Newlines always separate items. Tabs
separate items only in Stage 5c's own `formatted_text` (`_item_parts`).

The reason a separator can be left in place is that `_insert_bulleted_entry`
already renders both of them readably: `_clean_inline_tabs`
(`normalization/text.py:258`) rejoins a residual "|" with " — " and a residual
tab as "Label: Value — Value". So in RAW source text a tab marks the CELLS of
one table row, and leaving it alone yields one correctly punctuated bullet;
splitting on it would replace that bullet with fragments. The farm says the
same thing in counts: 456 of 722 K entries are a single record shaped
"Title\\tDate\\tDescription", and tab-splitting the raw fallback would turn 19
entries into 41 bullets, 16 of them a bare date. Raw text therefore keeps the
newline-only split (`entry_lines`) on both of the paths that read it.

WHERE THE TEXT COMES FROM
`_teaching_entry_lines` is the whole reconstruction, as a pure function: fields
in, the bullet strings out, no Word object touched. It picks between three
sources of text, in order:

- `formatted_text` from stage 5c, when the raw entry is a single item. ISO dates
  the LLM left behind are normalized first, and markdown is stripped, since Word
  runs have no markdown. A `formatted_text` that strips down to nothing is not
  bulleted blank; the raw line is used instead, and the entry is reported when
  there is no raw line either.
- the RAW lines, when the raw entry has several of them. Stage 5c routinely
  merges a multi-item teaching block into one sentence, so the original line
  breaks are the more faithful record; only one of the bullets carries the
  entry's comments, so an entry is not annotated N times. The exception is ONE
  table row whose cells wrap over paragraphs (#987): the extractor joins a
  cell's own paragraphs with a newline, so that row's raw text has several
  lines although it is one course, and those lines are not items. Such a row
  is rejoined into ONE bullet of its own raw text (every cell's paragraphs
  joined with a space; stage 5c's text is not used, it drops words); see
  `wrapped_row_text`.
- the extracted fields (course code, title, institution, role), when stage 5c
  produced nothing at all. Both code and title arrive as lists often enough that
  each is joined before use.

`_insert_teaching_entry` keeps only what is not text: the two content gates, the
warning for an entry that reconstructs to nothing, and the insertion order.

Structural labels and orphan fragments are dropped at the top -- a bare "Course
Title" or a dangling continuation line is source-table furniture, not a
teaching activity.
"""
import logging
import re
import zipfile
from typing import TypedDict

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import _Row

from ..formatting import (
    DATE_SPAN_SEPARATOR,
    envelope_date_spans,
    extra_date_spans,
    normalize_iso_dates_in_text,
)
from ..parsing.dates import _parse_date_components
from ..normalization import _strip_markdown_for_word
from ..parsing import _is_orphan_fragment, _is_structural_label
from ..sorting import sort_entries_reverse_chronological
from unified_pipeline.core.docx_structure_extractor import get_paragraph_text
from unified_pipeline.core.render_check import CELL_SEPARATOR, entry_lines, wrapped_row_text

logger = logging.getLogger(__name__)


class _TeachingFields(TypedDict, total=False):
    """`extracted_fields` as section K actually reads it (review thread
    3927100368), in place of an unparameterized `Dict` that told a reader
    nothing about which keys exist or what they hold.

    `total=False` because stage 4's JSON has no schema enforcing any of these
    keys: this documents the contract the renderer relies on, it is not a
    validation layer -- see `normalization/fields.py` for why this codebase
    coerces at read time instead of rejecting.

    `course_code` and `course_title` are `str | list[str]` because extraction
    emits both shapes for the same field; the `'; '.join(...)` normalization in
    `_teaching_entry_lines` exists for the list case, and this union is the
    fact that makes it necessary rather than defensive noise.
    """
    formatted_text: str
    start_date: str
    course_code: str | list[str]
    course_title: str | list[str]
    institution: str
    role: str


class _TeachingEntry(TypedDict, total=False):
    """One stage-4/5c entry as the K writers read it. Same `total=False`
    reasoning as `_TeachingFields`."""
    text: str
    extracted_fields: _TeachingFields
    taxonomy_code: str
    element_type: str
    element_idx_start: int
    element_idx_end: int


# K-code to candidate header strings, tried in order (see module docstring and
# docs/CODING_STANDARDS.md §8.2). K1-K5: didactic, clinical, administrative,
# continuing education, and outreach teaching, respectively.
TEACHING_SECTION_HEADERS: dict[str, list[str]] = {
    'K1': ['Didactic teaching', 'Didactic'],
    'K2': ['Clinical teaching', 'bedside teaching'],
    'K3': ['Administrative teaching', 'leadership role'],
    'K4': ['Continuing education', 'professional education'],
    'K5': ['outreach activities', 'Other education/outreach', 'community education or patient'],
}

# The source CV's own table column labels. A raw line that is nothing but one of
# these is furniture, not a teaching activity (#574). One frozenset rather than
# a list literal rebuilt per line: it is this section's vocabulary for "column
# header", and it now has one definition (§8.2).
_STRUCTURAL_COLUMN_LABELS = frozenset({'title', 'institution', 'dates', 'role'})

# A raw line at least this long that also carries a ';' is a run-together list
# of activities rather than one activity; below it, a ';' is punctuation inside
# a single item. Named because the bare 100 said nothing about which of the two
# it was guarding against (§8.2).
_SEMICOLON_SPLIT_MIN_CHARS = 100

# Stage 5c's own date for a record starting in `year`: the year, or a range
# opening with it ("1982-1986", "1982 - present"). Not inside a longer number.
_OWN_DATE_TOKEN = (r'(?<!\d){year}(?:\s*[-\u2013\u2014]\s*(?:\d{{2,4}}|present|current))?'
                   r'(?!\d)')


def _item_parts(text: str | None) -> list[str]:
    """Stripped, non-empty ITEMS of `text`: newlines and tabs separate items,
    `|` does not (#476, review thread 3927100368).

    The module docstring's delimiter contract in one function, so the rule is
    testable without a Word document and cannot drift between call sites. Note
    what this is NOT: `entry_fragments` splits on `|` as well, and swapping it
    in here would silently turn "Role | Institution" into two items instead of
    the one bullet `_clean_inline_tabs` renders it as ("Role — Institution").

    Its one caller is `_teaching_entry_lines`' Stage-5c-prose branch (the
    second `if formatted_text:`), which runs only when `original_text` is
    empty -- Stage 5c's own prose carries no source-table furniture, so a tab
    in it is a fused item rather than a cell boundary. Raw source text is
    deliberately split by `entry_lines` instead; the module docstring says
    why, with the farm counts.
    """
    return [part.strip()
            for line in str(text or '').split('\n')
            for part in line.split('\t')
            if part.strip()]



def _with_extra_spans(text: str, fields: _TeachingFields, taxonomy_code: str) -> str:
    """Stage 5c's `text` with the record's other spans (`extra_date_spans`)
    written after its own date: "1982-1986, 1989-1995, 2004-2012 - <role>".
    Stage 5c formats from the fields and drops the spans stage 4 filed under
    `additional_periods`/`additional_dates` (EBYSBC E22: EQADVR 33/34). The
    own date is the first year-or-range token naming the record's start
    year; a text with none is left as it is, since there is no date to
    extend. When that own range is only the envelope of the spans
    (`envelope_date_spans`: UYFRTL 33/34/48/49), the spans replace it."""
    envelope = envelope_date_spans(fields, taxonomy_code)
    extras = envelope or extra_date_spans(fields, taxonomy_code)
    if not extras:  # also when the start has no year: both are [] then
        return text
    start_year = _parse_date_components(str(fields.get('start_date') or '').strip())[0]
    own_date = re.search(_OWN_DATE_TOKEN.format(year=start_year), text, re.IGNORECASE)
    if own_date is None:
        return text
    if envelope:
        return text[:own_date.start()] + DATE_SPAN_SEPARATOR.join(envelope) + text[own_date.end():]
    cut = own_date.end()
    return DATE_SPAN_SEPARATOR.join([text[:cut], *extras]) + text[cut:]


def _teaching_entry_lines(fields: _TeachingFields, original_text: str,
                          row_text: str | None = None,
                          taxonomy_code: str = '') -> list[str]:
    """The bullet strings one teaching entry renders as, in page order.

    Pure: no document, no logging, no insertion -- extracted out of
    `_insert_teaching_entry` (review thread 3927100368) so the delimiter rule
    and the three-way source fallback are testable on their own, and so the
    renderer is left with one insertion loop instead of seven.

    An empty list means the entry reconstructs to nothing at all; the caller
    logs that (§5.4 -- it is recorded, never swallowed). It is never a list
    holding one empty string: a `formatted_text` of nothing but whitespace or
    markup used to be bulleted blank, which both put an empty list paragraph
    on the page AND discarded whatever raw line the entry still carried, so
    such a `formatted_text` now falls through to the raw line and, when there
    is none, to the empty list the caller reports.

    `row_text` (`wrapped_row_text`, #987) is the raw text of a row whose raw
    lines are one course's wrapped cells rather than several items: it is the
    one bullet in their place. Stage 5c's text is NOT used for it, because
    5c drops words the raw cells carry. With no `formatted_text` the field
    fallback below already applies, as it does for every entry.
    """
    formatted_text = fields.get('formatted_text', '') or ''
    if formatted_text:
        # Normalize any raw ISO dates the LLM left in formatted text
        formatted_text = normalize_iso_dates_in_text(formatted_text)

    if formatted_text and original_text:
        # Multi-item entry: use original lines (Stage 5c may have over-combined).
        # Newline-only on purpose -- see the module's delimiter contract.
        original_lines = entry_lines(original_text)
        if len(original_lines) > 1:
            return [row_text] if row_text else original_lines
        stripped = _strip_markdown_for_word(formatted_text, preserve_newlines=True)
        if stripped.strip():
            return [_with_extra_spans(stripped, fields, taxonomy_code)]
        # Stage 5c sent a formatted_text with no words in it ("   ", "**  **").
        # Bulleting it renders an empty list paragraph and throws away the raw
        # line this entry still has, so take the raw line instead; when the raw
        # text was blank too this is [] and the caller warns (#476).
        return original_lines

    if formatted_text:
        parts = _item_parts(_strip_markdown_for_word(formatted_text, preserve_newlines=True))
        if len(parts) > 1:
            return ['. '.join(parts)]
        # Same rule with no raw text to fall back to: `parts[:1]` is [] when the
        # formatted text reduced to nothing, so the caller reports the entry
        # rather than the renderer emitting a blank bullet for it.
        return parts[:1]

    # No Stage 5c formatting - fall back to building text from fields
    course_code = fields.get('course_code', '')
    course_title = fields.get('course_title', '')
    institution = fields.get('institution', '')
    role = fields.get('role', '')

    if isinstance(course_title, list):
        course_title = '; '.join(course_title)
    if isinstance(course_code, list):
        course_code = '; '.join(course_code)

    if course_title:
        text = f"{course_code}: {course_title}" if course_code else course_title
        if institution and institution not in text:
            text += f", {institution}"
        if role and role not in text:
            text += f" ({role})"
        return [text]

    lines: list[str] = []
    for line in original_text.split('\n'):
        line = line.strip()
        if not line or line.lower() in _STRUCTURAL_COLUMN_LABELS:
            continue
        if ';' in line and len(line) > _SEMICOLON_SPLIT_MIN_CHARS:
            lines.extend([item.strip() for item in line.split(';') if item.strip()])
        else:
            lines.append(line)

    if lines:
        return lines

    # Every line was a bare column label ("Title"/"Institution"/etc.) -- the
    # filter above left nothing to bullet. Institution and role are right there
    # in the extracted fields even though this branch only exists because both
    # formatted_text and course_title came back empty, so fall back to those
    # before giving up on the entry (#574).
    text = ', '.join(part for part in (institution, role) if part)
    return [text] if text else []


# What can go wrong opening the source CV or indexing into it. Anything else is
# a code bug and must surface, not read as "this CV has no hierarchy" (§5.4).
_SOURCE_READ_ERRORS = (OSError, KeyError, IndexError, zipfile.BadZipFile, PackageNotFoundError)


def _cell_zero_paragraphs(row: _Row) -> tuple[list[str], list[bool]]:
    """A source row's cell-0 paragraphs as `(texts, has_numPr)`, blanks dropped.

    Text is read with the reader's own `get_paragraph_text`, the function
    `extract_table_metadata` builds a cell's text from, so this side and the
    entry's text cannot disagree over tabs, line breaks, hyperlinks or tracked
    insertions. `w:numPr` is read through `._p` because python-docx has no
    public accessor for a paragraph's numbering; `_apply_list_bullet` writes it
    the same way.
    """
    texts: list[str] = []
    flags: list[bool] = []
    for para in row.cells[0].paragraphs:
        text = get_paragraph_text(para).strip()
        if not text:
            continue
        pPr = para._p.find(qn('w:pPr'))
        texts.append(text)
        flags.append(pPr is not None and pPr.find(qn('w:numPr')) is not None)
    return texts, flags


def _row_matches_lines(source: list[str], lines: list[str]) -> bool:
    """True when `lines` is exactly this source cell, paragraph for paragraph.

    Every line must be its source paragraph verbatim. The one licensed
    difference is the FIRST line: stage 2 attaches a multi-column row's trailing
    columns (date, institution) to cell 0's first paragraph, joined with
    `CELL_SEPARATOR` (#488, `join_row_cells`), so that line may carry a tail and
    no other line may.
    """
    if len(source) != len(lines) or source[1:] != lines[1:]:
        return False
    # `join_row_cells` keeps whatever whitespace sat between the paragraph and
    # the separator, so the tail may follow spaces the stripped source lacks.
    head, tail = lines[0][:len(source[0])], lines[0][len(source[0]):]
    return head == source[0] and (not tail or tail.lstrip().startswith(CELL_SEPARATOR.strip()))


class SourceCellLevels:
    """Sub-bullet levels read back from the source CV's table cells (#423).

    The reader flattens a table cell to newline-joined text, so the
    per-paragraph `w:numPr` that says which lines are children of the entry's
    title reaches no stage artifact. The cell still has it, and both drivers
    already hand stage 6 the source path (`original_doc_path`), so the levels
    are re-read here instead of being threaded through five stages.

    Rows are found by CONTENT, not by `row_index`: that index counts the
    reader's processed rows (row 0 is dropped when it is the table's header,
    merged rows are split), so it is not a position in `table.rows`. A row is
    accepted only when its cell-0 paragraphs reproduce the entry's lines
    (`_row_matches_lines`); several matching rows are accepted only when they
    agree. Every other outcome is None and the caller renders flat, so an
    unusable source is never worse than none.
    """

    def __init__(self, path: str) -> None:
        self._path = path
        self._doc = None
        self._rows_by_table: dict[int, list[tuple[list[str], list[bool]]]] = {}
        self._warned = False

    def _table_rows(self, table_index: int) -> list[tuple[list[str], list[bool]]]:
        if table_index not in self._rows_by_table:
            if self._doc is None:
                self._doc = Document(self._path)
            seen_cells = set()
            rows = []
            for row in self._doc.tables[table_index].rows:
                tc = row.cells[0]._tc
                if tc not in seen_cells:  # a vertically merged cell repeats across rows
                    seen_cells.add(tc)
                    rows.append(_cell_zero_paragraphs(row))
            self._rows_by_table[table_index] = rows
        return self._rows_by_table[table_index]

    def flags_for(self, entry: _TeachingEntry, lines: list[str]) -> list[bool] | None:
        """Per-line "was a sub-bullet in the source cell", or None.

        None when the entry is not a table row, the source is unreadable, no
        source row (or disagreeing rows) reproduces `lines`, or the cell is not
        a title followed by bullets. A cell whose first paragraph is itself a
        list item, or none of whose later ones are, has no title/child
        hierarchy to restore: 104 of 217 source sub-bullet clusters measured
        were whole sections of peer bullets, where flat is the right rendering.
        """
        table_index = entry.get('table_index')
        if entry.get('element_type') != 'table_row' or table_index is None or len(lines) < 2:
            return None
        try:
            rows = self._table_rows(table_index)
        except _SOURCE_READ_ERRORS as e:
            # Once per run, not per entry: a missing source or a stale table
            # index is a whole-run condition, and silence would make the
            # feature no-op read as "this CV has no hierarchy".
            if not self._warned:
                self._warned = True
                logger.warning("Could not read source cell levels from %s, "
                               "section K stays flat: %s", self._path, e)
            return None
        matches = [flags for texts, flags in rows if _row_matches_lines(texts, lines)]
        if not matches or any(m != matches[0] for m in matches):
            return None
        flags = matches[0]
        # Only a title-then-bullets cell carries hierarchy worth restoring.
        if flags[0] or not any(flags[1:]):
            return None
        return flags


class TeachingSection:
    """Section K writers, mixed into `WCMTemplateGenerator`."""

    # Set per render by `generate()`; None means no source CV, so every K
    # bullet stays level 0. A class default so a generator built without
    # `__init__` (several tests do) still renders.
    _source_cell_levels: SourceCellLevels | None = None

    def _fill_teaching(self, entries_by_code: dict[str, list[_TeachingEntry]],
                       original_doc_path: str | None = None) -> None:
        """Fill K. TEACHING ACTIVITIES section.

        Routes K-codes to their appropriate WCM subsections:
        - K1 (didactic) -> "Didactic teaching"
        - K2 (clinical) -> "Clinical teaching"
        - K3 (administrative) -> "Administrative teaching"
        - K4 (CME) -> "Continuing education and professional education"
        - K5 (community) -> "Other education/outreach activities"

        Entries are inserted as a flat chronological list under each K-code section.
        The WCM template provides the structure; Stage 5c handles per-entry formatting.
        Original CV hierarchy labels are not carried over, but a table row whose
        source cell is a title plus bullets renders its bullets one level in
        (#423): `original_doc_path` is where those levels are re-read from, and
        None leaves every bullet at level 0. Built per call, so one CV's rows
        never carry into the next render. When two codes share
        the fallback heading, the module docstring states the order they land in.
        """
        k_section_map = TEACHING_SECTION_HEADERS
        self._source_cell_levels = (
            SourceCellLevels(original_doc_path) if original_doc_path else None)

        # Count total entries
        total_entries = sum(len(entries_by_code.get(code, [])) for code in k_section_map.keys())
        if total_entries == 0:
            return

        logger.info("Filling Teaching (%d entries)...", total_entries)

        # Fill each K-code section separately
        for code, search_texts in k_section_map.items():
            entries = entries_by_code.get(code, [])
            if not entries:
                continue

            # Find the appropriate section header
            section_idx = None
            for search_text in search_texts:
                section_idx = self._find_header_paragraph(search_text)
                if section_idx is not None:
                    break

            if section_idx is None:
                # Fall back to general teaching section. Several codes can land
                # here at once; the module docstring pins the resulting order.
                section_idx = self._find_header_paragraph("EDUCATIONAL CONTRIBUTIONS")
                if section_idx is None:
                    logger.warning("no template heading for teaching code %s; "
                                   "%d entries not rendered", code, len(entries))
                    continue

            # Flat list: sort chronologically and insert without original CV sub-headers.
            # The WCM template's own K-code sections (Didactic, Clinical, Administrative,
            # CME, Community) provide the structure; Stage 5c handles per-entry formatting.
            sorted_entries = sort_entries_reverse_chronological(entries)
            reversed_entries = list(reversed(sorted_entries))

            for i, entry in enumerate(reversed_entries):
                is_first_in_section = (i == len(reversed_entries) - 1)
                self._insert_teaching_entry(section_idx + 1, entry,
                                            is_first_visible=is_first_in_section)

    def _insert_teaching_entry(self, insert_idx: int, entry: _TeachingEntry,
                               is_first_visible: bool = False) -> None:
        """Insert a single teaching entry as bulleted items.

        What text to render is `_teaching_entry_lines`' job (Stage 5c formatted
        text, raw lines, or structured fields). What is left here is the two
        content gates, the warning for an entry that reconstructs to nothing,
        and the insertion order.
        """
        # Skip structural labels from source CV
        if _is_structural_label(entry):
            return

        fields = entry.get('extracted_fields', {}) or {}
        formatted_text = fields.get('formatted_text', '') or ''
        original_text = entry.get('text', '') or ''

        if _is_orphan_fragment(fields, formatted_text, original_text):
            return

        raw_lines = entry_lines(original_text)
        levels = (self._source_cell_levels.flags_for(entry, raw_lines)
                  if self._source_cell_levels else None)
        # A row whose source cell is a title plus bullets is a hierarchy, not one
        # course whose cells wrap (#987): `wrapped_row_text` cannot tell them
        # apart from the text alone, and would weld the bullets into one line.
        lines = _teaching_entry_lines(fields, original_text,
                                      row_text=None if levels else wrapped_row_text(entry),
                                      taxonomy_code=str(entry.get('taxonomy_code') or ''))
        if levels and lines != raw_lines:
            levels = None  # not the raw-lines path (e.g. stage 5c prose); nothing to align
        if not lines:
            logger.warning("teaching entry produced no renderable line (%s): %r",
                           entry.get('taxonomy_code'), original_text[:80])
            return

        # `_insert_bulleted_entry` inserts BEFORE insert_idx, so walking the
        # list backwards leaves it in source order on the page. Exactly one
        # bullet carries `entry` (the last in source order, inserted first), so
        # a multi-line entry is annotated once rather than N times, and the
        # blank line goes above the bullet that ends up visually first.
        for j, line_text in enumerate(reversed(lines)):
            # The loop is reversed, so bullet j is source line len-1-j; indexing
            # `levels` with j would invert the hierarchy.
            is_child = bool(levels and levels[len(lines) - 1 - j])
            self._insert_bulleted_entry(
                insert_idx, line_text, entry if j == 0 else None,
                add_blank_before=is_first_visible and j == len(lines) - 1,
                list_level=1 if is_child else 0
            )
