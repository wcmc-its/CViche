#!/usr/bin/env python3
"""
Create a basic WCM CV template
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path


def create_wcm_template():
    """Create a basic WCM CV template with standard sections."""
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Header - Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('CURRICULUM VITAE')
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()  # Blank line

    # Personal Data Section
    heading = doc.add_paragraph()
    run = heading.add_run('PERSONAL DATA')
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph('Name:')
    doc.add_paragraph('Office address:')
    doc.add_paragraph('Office telephone:')
    doc.add_paragraph('Work email:')
    doc.add_paragraph()  # Blank line

    # Education Section
    heading = doc.add_paragraph()
    run = heading.add_run('EDUCATION')
    run.bold = True
    run.font.size = Pt(12)

    # Create education table
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Degree'
    hdr_cells[1].text = 'Institution'
    hdr_cells[2].text = 'Dates'
    hdr_cells[3].text = 'Year'

    # Make header row bold
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    doc.add_paragraph()  # Blank line

    # Professional Positions Section
    heading = doc.add_paragraph()
    run = heading.add_run('PROFESSIONAL POSITIONS & EMPLOYMENT')
    run.bold = True
    run.font.size = Pt(12)

    # Create positions table
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Title'
    hdr_cells[1].text = 'Institution'
    hdr_cells[2].text = 'Dates'

    # Make header row bold
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    doc.add_paragraph()  # Blank line

    # Grants Section
    heading = doc.add_paragraph()
    run = heading.add_run('RESEARCH SUPPORT')
    run.bold = True
    run.font.size = Pt(12)

    # Create grants table
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Grant Title'
    hdr_cells[1].text = 'Agency'
    hdr_cells[2].text = 'Role'
    hdr_cells[3].text = 'Dates'

    # Make header row bold
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    doc.add_paragraph()  # Blank line

    # Bibliography Section
    heading = doc.add_paragraph()
    run = heading.add_run('BIBLIOGRAPHY')
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph()  # Space for publications

    # Save template
    output_path = Path(__file__).parent / 'cv_template_wcm.docx'
    doc.save(str(output_path))
    print(f"Created WCM template: {output_path}")
    return str(output_path)


if __name__ == '__main__':
    create_wcm_template()
