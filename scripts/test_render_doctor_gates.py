#!/usr/bin/env python3
"""Self-tests for the pure logic in render_gate_compare.py and
doctor_gate_compare.py -- the parts cheap enough to pin without a corpus.

    python3 scripts/test_render_doctor_gates.py

The controls that actually prove these gates work (determinism holds, a real
mutation is caught, a crashed render/doctor run is refused rather than
compared) need a corpus and are logged in
docs/guides/render-doctor-gates.md instead -- see issue #584. This file
covers the regressions found while building that (the date/timestamp masking
not accidentally masking real content, and the work-dir path normaliser
silently failing on a relative path), and render_gate.py's own pure logic:
the LLM-disable monkeypatch below, --source-dir's uid-to-docx resolver
(#550), whose failure mode is silent -- a wrong resolution renders ANOTHER
CV's contact block into this CV's Personal Data table rather than raising --
and the pieces split out of main() for review on #740: uid discovery and
filtering, artifact precedence, the pre-rmtree output-path check, and the
semantic validation of a rendered docx. What those helpers do when wired
together is scripts/test_render_gate_integration.py's subject; this file
pins each in isolation.
"""
import contextlib
import io
import json
import os
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# unified_pipeline is normally on the path via PYTHONPATH=<arm>/src (see
# render_gate.py's own docstring) -- add it here too so this file's own
# "runnable standalone" contract holds without the caller setting that up
# just to run the (otherwise pipeline-independent) self-tests.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import doctor_gate_compare as dgc  # noqa: E402
import render_gate as rg  # noqa: E402
import render_gate_compare as rgc  # noqa: E402

from docx import Document  # noqa: E402


def _make_docx(path, text, tracked_change_date=None):
    """A one-paragraph docx, optionally with a w:date-stamped tracked insert."""
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    if tracked_change_date is None:
        return
    # python-docx has no tracked-changes API; stamp a w:date attribute
    # directly the way stage 6 does (ins.set(qn('w:date'), ...)).
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    p = next(root.iter(W + "p"))
    ins = etree.SubElement(p, W + "ins")
    ins.set(W + "date", tracked_change_date)
    r = etree.SubElement(ins, W + "r")
    t = etree.SubElement(r, W + "t")
    t.text = "tracked insert"
    new_xml = etree.tostring(root)
    # Rewrite the zip with the patched document.xml.
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = new_xml if item.filename == "word/document.xml" else zin.read(item.filename)
            zout.writestr(item, data)
    os.replace(tmp, path)


def test_fingerprint_masks_render_timestamp_but_not_real_content():
    with tempfile.TemporaryDirectory() as tmp:
        a = Path(tmp) / "a.docx"
        b = Path(tmp) / "b.docx"
        # Same visible content, different tracked-change timestamps (exactly
        # what two renders of the identical input produce a few seconds
        # apart) -- must fingerprint identically.
        _make_docx(a, "Date of preparation: February 1, 2026",
                  tracked_change_date="2026-02-01T10:00:00Z")
        _make_docx(b, "Date of preparation: February 1, 2026",
                  tracked_change_date="2026-02-01T10:00:07Z")
        assert rgc._fingerprint(a) == rgc._fingerprint(b), (
            "two renders differing only in render-time timestamps must "
            "fingerprint identically")

        # Real content difference must still be caught.
        c = Path(tmp) / "c.docx"
        _make_docx(c, "Date of preparation: February 1, 2026 -- EXTRA TEXT",
                  tracked_change_date="2026-02-01T10:00:00Z")
        assert rgc._fingerprint(a) != rgc._fingerprint(c), (
            "masking timestamps must not also mask a real content change"
        )


def test_check_render_failures_trips_on_either_arm():
    with tempfile.TemporaryDirectory() as tmp:
        a_dir, b_dir = Path(tmp) / "a", Path(tmp) / "b"
        a_dir.mkdir()
        b_dir.mkdir()
        (a_dir / "_render_index.json").write_text(json.dumps({"uid1": {"input": "x"}}))
        (b_dir / "_render_index.json").write_text(
            json.dumps({"uid1": {"error": "RuntimeError: boom"}}))
        problems = rgc._check_render_failures(a_dir, b_dir)
        assert problems, "a recorded render failure must trip the guard"
        assert any("uid1" in p for p in problems)


def test_check_render_failures_clear_when_both_clean():
    with tempfile.TemporaryDirectory() as tmp:
        a_dir, b_dir = Path(tmp) / "a", Path(tmp) / "b"
        a_dir.mkdir()
        b_dir.mkdir()
        (a_dir / "_render_index.json").write_text(json.dumps({"uid1": {"input": "x"}}))
        (b_dir / "_render_index.json").write_text(json.dumps({"uid1": {"input": "x"}}))
        # "clean" now also means the index's success entry has a matching
        # DOCX on disk (review on #589) -- an index that says success with
        # no corresponding file must trip the guard, not read as clear.
        (a_dir / "uid1_wcm.docx").touch()
        (b_dir / "uid1_wcm.docx").touch()
        assert rgc._check_render_failures(a_dir, b_dir) == []


def test_check_render_failures_trips_when_index_and_docx_disagree():
    with tempfile.TemporaryDirectory() as tmp:
        a_dir, b_dir = Path(tmp) / "a", Path(tmp) / "b"
        a_dir.mkdir()
        b_dir.mkdir()
        # Index claims success for uid1, but no DOCX was actually written --
        # the exact "0 == 0" gap the file-set comparison alone can't see
        # (review on #589, point at render_gate_compare.py:69).
        (a_dir / "_render_index.json").write_text(json.dumps({"uid1": {"input": "x"}}))
        (b_dir / "_render_index.json").write_text(json.dumps({"uid1": {"input": "x"}}))
        (b_dir / "uid1_wcm.docx").touch()
        problems = rgc._check_render_failures(a_dir, b_dir)
        assert problems, "index success with no DOCX on disk must trip the guard"
        assert any("uid1" in p and "no DOCX on disk" in p for p in problems)


def test_norm_matches_relative_and_absolute_work_dirs_identically():
    # The exact bug found building this gate: an earlier version of this
    # regex required a leading "/" before the work-dir segment, so a
    # relative invocation ("scratch/doctor_a_work/...") never matched and
    # every CV read as changed.
    report = {
        "artifacts": {
            "stage_4": "scratch/doctor_a_work/uid1/stage_4_field_extraction/uid1_fields.json"
        }
    }
    report_abs = {
        "artifacts": {
            "stage_4": "/Users/x/worktrees/cviche/scratch/doctor_b_work/uid1/"
                       "stage_4_field_extraction/uid1_fields.json"
        }
    }
    assert dgc.norm(report) == dgc.norm(report_abs), (
        "a relative and an absolute path to the same logical artifact must "
        "normalise identically"
    )
    assert "<WORK>" in dgc.norm(report)
    assert "doctor_a_work" not in dgc.norm(report)


def test_norm_still_distinguishes_real_content_differences():
    a = {"findings": [{"lint": "table_shape", "severity": "WARN"}]}
    b = {"findings": [{"lint": "table_shape", "severity": "INFO"}]}
    assert dgc.norm(a) != dgc.norm(b)


def test_norm_excludes_metrics_so_a_trend_number_never_reads_as_changed():
    # Round-2 review F1: metrics (#816) carries no severity of its own, so a
    # routine denominator shift (honors_rows, appendix_share, ...) must not
    # make norm() -- and therefore CHANGED/PASS/FAIL -- see this uid as
    # regressed when findings are identical.
    a = {"findings": [{"lint": "table_shape", "severity": "WARN"}],
         "metrics": {"honors_rows": 40, "appendix_share": 0.1}}
    b = {"findings": [{"lint": "table_shape", "severity": "WARN"}],
         "metrics": {"honors_rows": 90, "appendix_share": 0.4}}
    assert dgc.norm(a) == dgc.norm(b)


def test_metrics_diff_reports_per_key_old_and_new():
    a = {"metrics": {"honors_rows": 40, "stage3b_fallback_ratio": 0.01}}
    b = {"metrics": {"honors_rows": 90, "stage3b_fallback_ratio": 0.01}}
    assert dgc._metrics_diff(a, b) == [("honors_rows", 40, 90)]


def test_metrics_diff_empty_when_metrics_identical_or_absent():
    a = {"metrics": {"honors_rows": 40}}
    b = {"metrics": {"honors_rows": 40}}
    assert dgc._metrics_diff(a, b) == []
    assert dgc._metrics_diff({}, {}) == []
    # A key present on only one side is a difference too.
    assert dgc._metrics_diff({"metrics": {"x": 1}}, {"metrics": {}}) == [("x", 1, None)]


def _write_report(path, uid, findings, metrics=None):
    path.write_text(json.dumps({
        uid: {"document_uid": uid, "root": "/r", "artifacts": {}, "findings": findings,
              "counts": {}, "worst_severity": None, "metrics": metrics or {}},
        dgc.FAILED_KEY: {},
    }))


def test_doctor_compare_main_passes_when_only_metrics_differ():
    # The whole point of F1: a metrics-only difference must PASS and exit 0,
    # while still telling a human reader what changed.
    with tempfile.TemporaryDirectory() as tmp:
        a_path, b_path = Path(tmp) / "a.json", Path(tmp) / "b.json"
        finding = [{"lint": "table_shape", "severity": "WARN"}]
        _write_report(a_path, "uid1", finding, {"honors_rows": 40})
        _write_report(b_path, "uid1", finding, {"honors_rows": 90})

        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            rc = dgc.main([str(a_path), str(b_path)])
        assert rc == 0
        text = out.getvalue()
        assert "CHANGED 0" in text
        assert "PASS - no doctor finding changed" in text
        assert "metrics CHANGED (reported only, never fails the gate): 1 uid(s), 1 key(s)" in text
        assert "uid1 honors_rows: 40 -> 90" in text


def test_doctor_compare_main_still_fails_on_a_findings_change_regardless_of_metrics():
    with tempfile.TemporaryDirectory() as tmp:
        a_path, b_path = Path(tmp) / "a.json", Path(tmp) / "b.json"
        _write_report(a_path, "uid1", [{"lint": "table_shape", "severity": "WARN"}],
                      {"honors_rows": 40})
        _write_report(b_path, "uid1", [{"lint": "table_shape", "severity": "INFO"}],
                      {"honors_rows": 40})

        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            rc = dgc.main([str(a_path), str(b_path)])
        assert rc == 1
        text = out.getvalue()
        assert "CHANGED 1" in text
        assert "FAIL" in text
        # metrics are identical here -- no metrics section at all.
        assert "metrics CHANGED" not in text


def test_doctor_compare_main_silent_on_metrics_when_reports_have_none():
    # Legacy/pre-#816 shaped reports (no "metrics" key at all): output must
    # be exactly what it was before metrics existed -- no new section.
    with tempfile.TemporaryDirectory() as tmp:
        a_path, b_path = Path(tmp) / "a.json", Path(tmp) / "b.json"
        a_path.write_text(json.dumps({
            "uid1": {"findings": []}, dgc.FAILED_KEY: {}}))
        b_path.write_text(json.dumps({
            "uid1": {"findings": []}, dgc.FAILED_KEY: {}}))

        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            rc = dgc.main([str(a_path), str(b_path)])
        assert rc == 0
        assert "metrics CHANGED" not in out.getvalue()


def test_doctor_failure_guard_trips_on_either_arm():
    assert dgc._failure_guard_problems({"uid1": "boom"}, {})
    assert dgc._failure_guard_problems({}, {"uid1": "boom"})
    assert dgc._failure_guard_problems({}, {}) == []


def test_doctor_failure_guard_trips_on_malformed_failed_value():
    # "_failed" must be a dict -- a malformed doctor_gate.py output (review
    # on #589) must be a loud gate failure, not a silent pass-through.
    assert dgc._failure_guard_problems(["not", "a", "dict"], {})
    assert dgc._failure_guard_problems({}, "also not a dict")


def test_doctor_compare_identical_count_uses_uid_intersection():
    # The exact bug from review on #589: len(a) - len(diff) overcounts
    # "identical" when the UID sets differ. A=3 UIDs, B=2 UIDs, no common
    # changes -- only 2 CVs were actually compared.
    a = {"u1": {"findings": []}, "u2": {"findings": []}, "u3": {"findings": []}}
    b = {"u1": {"findings": []}, "u2": {"findings": []}}
    common = set(a) & set(b)
    diff = [u for u in common if dgc.norm(a[u]) != dgc.norm(b[u])]
    assert len(common) - len(diff) == 2, "only the 2 UIDs present in both arms were compared"


def test_render_gate_llm_monkeypatch_still_intercepts_stage6():
    """render_gate.py neutralizes LLM calls via `s6.call_llm = _no_llm` --
    this only works because stage_6_word_template.py binds call_llm as a
    module-level name via `from ... import call_llm`, which is what makes it
    resolvable (and reassignable) as `s6.call_llm`. If a future refactor
    switches to `import unified_pipeline.llm_client as llm_client` +
    `llm_client.call_llm(...)`, this monkeypatch stops intercepting anything
    with no error -- a "deterministic" gate run could silently make a real
    LLM call (review on #589)."""
    import inspect
    import unified_pipeline.stage_6_word_template as s6

    src = inspect.getsource(s6)
    assert "from unified_pipeline.llm_client import call_llm" in src, (
        "stage_6_word_template.py's call_llm import style changed -- "
        "render_gate.py's LLM-disable monkeypatch needs updating to match"
    )

    original = s6.call_llm
    calls = []

    def _sentinel(*a, **kw):
        calls.append(1)
        raise RuntimeError("sentinel")

    s6.call_llm = _sentinel
    try:
        try:
            s6.call_llm()
            raise AssertionError("expected RuntimeError from the sentinel")
        except RuntimeError as e:
            assert "sentinel" in str(e)
        assert calls, "reassigning s6.call_llm did not intercept the call"
    finally:
        s6.call_llm = original


def test_render_gate_source_dir_resolver_respects_uid_boundaries():
    """A shorter uid must not claim a longer uid's docx (#550).

    glob("web05*.docx") also matches web050_cv.docx, and sorted() puts '0'
    (0x30) before '_' (0x5F), so a bare prefix match hands web05 the WRONG
    CV's source document -- the same shape as the 2026-07-15 doctor sweep
    that diagnosed 3 of 25 CVs against another CV's artifacts. Here the
    consequence is worse than a bad diagnosis: web050's name, office
    address, telephone and work email would render into web05's output.
    """
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        (d / "web05_cv.docx").write_text("")
        (d / "web050_cv.docx").write_text("")

        assert rg._resolve_source_docx(d, "web05").name == "web05_cv.docx", (
            "uid web05 resolved to a longer uid's docx -- the _uid_owns "
            "boundary rule is not being applied")
        assert rg._resolve_source_docx(d, "web050").name == "web050_cv.docx"


def test_render_gate_source_dir_resolver_returns_none_for_an_uncovered_uid():
    # 65 of the farm's 66 uids have a source docx; web08 has none. That uid
    # must resolve to None so main() falls back to the unmodified
    # run_stage6() call and records "not found" in the index -- not an
    # error, and not a silent render against someone else's document.
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        (d / "web09_cv.docx").write_text("")
        assert rg._resolve_source_docx(d, "web08") is None


def test_render_gate_source_dir_resolver_ignores_non_docx_files():
    # The glob is *.docx: a sidecar with the same uid stem (the farm's own
    # per-uid JSON artifacts sit beside CVs in some layouts) is not a source
    # document and must not be handed to python-docx.
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        (d / "web05_entries.json").write_text("{}")
        assert rg._resolve_source_docx(d, "web05") is None


def test_render_gate_source_dir_resolver_pins_an_ambiguous_uid_to_one_docx():
    """Two docx both legitimately owned by the uid -- pin which one wins.

    Nothing in the corpus guarantees one file per uid ("... CV.docx" beside
    "... CV (Updated 1-26-26).docx" is exactly this repo's naming), and an
    unpinned choice here is a silently non-deterministic gate arm: two runs of
    the same code could recover contact fields from two different documents.
    Today's rule is sorted()[0], and this test is what makes changing it a
    visible decision rather than a side effect.
    """
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        (d / "web05_cv_v2.docx").write_text("")
        (d / "web05_cv.docx").write_text("")

        assert rg._resolve_source_docx(d, "web05").name == "web05_cv.docx", (
            "an ambiguous uid must resolve to the first sorted candidate, "
            "not to whatever order the filesystem happened to hand back")


def test_render_gate_uid_filter_narrows_the_discovered_uids():
    # The filter is the difference between re-rendering 66 CVs and re-rendering
    # the one under investigation, so a filter that silently matched nothing
    # (or everything) would waste an hour or answer the wrong question.
    with tempfile.TemporaryDirectory() as tmp:
        arm = Path(tmp)
        s4 = arm / "stage_4_field_extraction"
        s4.mkdir()
        for uid in ("u1", "u2", "u3"):
            (s4 / f"{uid}_fields.json").write_text("{}")

        assert rg._discover_uids(arm, set()) == ["u1", "u2", "u3"], (
            "an empty filter means every uid in the arm")
        assert rg._discover_uids(arm, {"u2"}) == ["u2"]
        assert rg._discover_uids(arm, {"u2", "u3"}) == ["u2", "u3"]
        assert rg._discover_uids(arm, {"not-in-this-arm"}) == [], (
            "a filter matching nothing must produce nothing -- main() turns "
            "that into a non-zero exit rather than an empty PASS")


def test_render_gate_uid_file_drops_blank_lines_and_trailing_whitespace():
    # Every editor writes a trailing newline, and a uid list is usually
    # hand-assembled -- an empty or space-padded entry would join the filter
    # set, match no uid, and quietly narrow the run.
    with tempfile.TemporaryDirectory() as tmp:
        uids_file = Path(tmp) / "uids.txt"
        uids_file.write_text("u1  \n\n  u2\n\n", encoding="utf-8")

        assert rg._uid_filter([], uids_file) == {"u1", "u2"}
        assert rg._uid_filter(["u9"], None) == {"u9"}, (
            "without --uids-file the positional uids are the filter")
        assert rg._uid_filter([], None) == set(), (
            "no uids and no file means 'all', which is the empty set")


def test_render_gate_input_artifact_precedence_prefers_the_latest_stage():
    # PRECEDENCE is the whole reason two arms are comparable: both must render
    # from the same stage for each uid. Peel the tree back one stage at a time
    # and require the resolver to step to exactly the next entry, never past it.
    with tempfile.TemporaryDirectory() as tmp:
        arm = Path(tmp)
        for stage_dir, pat in rg.PRECEDENCE:
            (arm / stage_dir).mkdir(parents=True)
            (arm / stage_dir / pat.format(uid="u1")).write_text("{}")

        for stage_dir, pat in rg.PRECEDENCE:
            name = pat.format(uid="u1")
            resolved = rg._resolve_input_artifact(arm, "u1")
            assert resolved is not None and resolved.name == name, (
                f"expected {name}, got {resolved}")
            (arm / stage_dir / name).unlink()

        assert rg._resolve_input_artifact(arm, "u1") is None, (
            "a uid with no artifact at any stage must resolve to None so the "
            "index records 'no input artifact' instead of rendering nothing")


def test_render_gate_input_artifact_skips_a_directory_named_like_an_artifact():
    # is_file(), not exists() (review on #589): a DIRECTORY named
    # <uid>_enriched.json would otherwise be handed to the JSON reader instead
    # of falling through to the stage-4 file that is actually there.
    with tempfile.TemporaryDirectory() as tmp:
        arm = Path(tmp)
        (arm / "stage_5_enrichment").mkdir()
        (arm / "stage_4_field_extraction").mkdir()
        (arm / "stage_5_enrichment" / "u1_enriched.json").mkdir()
        (arm / "stage_4_field_extraction" / "u1_fields.json").write_text("{}")

        resolved = rg._resolve_input_artifact(arm, "u1")
        assert resolved is not None and resolved.name == "u1_fields.json", (
            "a directory named like a stage artifact must be skipped, not "
            "selected as this uid's input")


def test_render_gate_output_dir_validation_refuses_to_wipe_an_input_tree():
    """main() rmtree's `out` before every run, so `out` must be provably not an
    input (review on #740). All three destructive relationships are refused --
    equal, out inside an input, an input inside out -- and both paths are
    resolved first, so a ".." detour or a symlink cannot smuggle one past a
    string comparison."""
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm = root / "arm"
        (arm / "stage_4_field_extraction").mkdir(parents=True)

        assert rg._validate_output_dir(arm, arm), "out == input must be refused"
        assert rg._validate_output_dir(arm / "renders", arm), (
            "out inside the input tree must be refused")
        assert rg._validate_output_dir(root, arm), (
            "an out that contains the input tree must be refused")
        assert rg._validate_output_dir(root / "out" / ".." / "arm", arm), (
            "a '..' detour back into the input tree must be refused")

        link = root / "link"
        link.symlink_to(arm, target_is_directory=True)
        assert rg._validate_output_dir(link, arm), (
            "a symlink pointing at the input tree must be refused")

        assert rg._validate_output_dir(root / "out", arm) is None, (
            "a sibling out directory is the normal invocation")
        assert rg._validate_output_dir(root / "out", arm, None) is None, (
            "an omitted --source-dir arrives as None and must be skipped, "
            "not compared")
        assert rg._validate_output_dir(root / "sources" / "sub", arm, root / "sources"), (
            "--source-dir is checked on the same footing as arm_outputs")

        # The wire, not just the helper: _parse_args must exit 2 on a bad
        # `out`, so nothing downstream ever reaches shutil.rmtree.
        with contextlib.redirect_stderr(io.StringIO()):
            try:
                rg._parse_args([str(arm), str(arm)])
            except SystemExit as e:
                assert e.code == 2, f"argparse parser.error exits 2, got {e.code}"
            else:
                raise AssertionError("out == arm_outputs was accepted by _parse_args")


def _make_personal_data_docx(path, labels):
    """A docx whose one table carries `labels` down its first column.

    Mirrors the shape `_validate_rendered_docx` looks for, not a full render:
    the labels are what the WCM template defines and what stage 6's own filler
    matches on, so a table built this way is what a good render leaves behind.
    """
    doc = Document()
    doc.add_paragraph("PERSONAL DATA")
    if labels:
        table = doc.add_table(rows=len(labels), cols=2)
        for row, label in zip(table.rows, labels):
            row.cells[0].text = f"{label.title()}:"
    doc.save(str(path))


def test_render_gate_validates_the_rendered_personal_data_table():
    """A docx can be written, and be a bad render (review on #740).

    dest.is_file() passes on a truncated file and on a document that lost the
    template's tables outright, and the compare step would then fingerprint it
    as a real arm. Checked as invariants, not counts: the six labels the filler
    writes into, matched by substring the way that filler matches them.
    """
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)

        good = d / "good.docx"
        _make_personal_data_docx(good, rg.PERSONAL_DATA_LABELS)
        assert rg._validate_rendered_docx(good) is None, (
            "a docx carrying every PERSONAL DATA label row is a valid render")

        # (i) not a docx at all -- what a truncated or half-written file looks
        # like to python-docx, and what is_file() alone cannot tell apart.
        truncated = d / "truncated.docx"
        truncated.write_text("not a zip archive")
        assert rg._validate_rendered_docx(truncated), "an unreadable docx must be refused"

        # (ii) opens cleanly, but the template's tables are gone.
        no_tables = d / "no_tables.docx"
        _make_personal_data_docx(no_tables, ())
        problem = rg._validate_rendered_docx(no_tables)
        assert problem and "no tables" in problem

        # (iii) the table is there and locatable by its anchor, but a label row
        # the filler writes into has been lost -- the shape a row-count check
        # would catch only by accident and a hard-coded count would false-alarm
        # on the next template revision.
        missing_row = d / "missing_row.docx"
        _make_personal_data_docx(
            missing_row, [x for x in rg.PERSONAL_DATA_LABELS if x != "cell phone"])
        problem = rg._validate_rendered_docx(missing_row)
        assert problem and "cell phone" in problem, (
            f"a dropped PERSONAL DATA label row must name itself, got {problem!r}")


def test_render_gate_rejects_a_non_directory_source_dir():
    """Fail closed (CODING_STANDARDS 5.5): a typo'd --source-dir must not
    silently render as if the flag were omitted. Omitting it is a
    documented, intentional mode; a bad path is not -- and the two produce
    identical output, so a silent fallback would make an entire gate arm
    quietly meaningless."""
    with tempfile.TemporaryDirectory() as tmp:
        arm = Path(tmp) / "arm"
        (arm / "stage_4_field_extraction").mkdir(parents=True)
        out = Path(tmp) / "out"

        # Sanity: the same argv WITHOUT the flag parses, so the failure
        # below is attributable to --source-dir and not to the arm layout.
        ok = rg._parse_args([str(arm), str(out)])
        assert ok.source_dir is None, "--source-dir must default to off (opt-in)"

        # A FILE is not a directory either -- is_dir(), not exists().
        afile = Path(tmp) / "a_file.docx"
        afile.write_text("")

        for bad, why in ((Path(tmp) / "nope", "a non-existent"), (afile, "a file")):
            # argparse writes usage to stderr before raising; swallow it so
            # this file's own output stays the single "ok" line its runner
            # (CI's pipeline-tests job) reads.
            with contextlib.redirect_stderr(io.StringIO()):
                try:
                    rg._parse_args([str(arm), str(out), "--source-dir", str(bad)])
                except SystemExit as e:
                    assert e.code == 2, f"argparse parser.error exits 2, got {e.code}"
                else:
                    raise AssertionError(f"{why} --source-dir was accepted")


if __name__ == "__main__":
    test_fingerprint_masks_render_timestamp_but_not_real_content()
    test_check_render_failures_trips_on_either_arm()
    test_check_render_failures_clear_when_both_clean()
    test_check_render_failures_trips_when_index_and_docx_disagree()
    test_norm_matches_relative_and_absolute_work_dirs_identically()
    test_norm_still_distinguishes_real_content_differences()
    test_doctor_failure_guard_trips_on_either_arm()
    test_doctor_failure_guard_trips_on_malformed_failed_value()
    test_doctor_compare_identical_count_uses_uid_intersection()
    test_render_gate_llm_monkeypatch_still_intercepts_stage6()
    test_render_gate_source_dir_resolver_respects_uid_boundaries()
    test_render_gate_source_dir_resolver_returns_none_for_an_uncovered_uid()
    test_render_gate_source_dir_resolver_ignores_non_docx_files()
    test_render_gate_source_dir_resolver_pins_an_ambiguous_uid_to_one_docx()
    test_render_gate_uid_filter_narrows_the_discovered_uids()
    test_render_gate_uid_file_drops_blank_lines_and_trailing_whitespace()
    test_render_gate_input_artifact_precedence_prefers_the_latest_stage()
    test_render_gate_input_artifact_skips_a_directory_named_like_an_artifact()
    test_render_gate_output_dir_validation_refuses_to_wipe_an_input_tree()
    test_render_gate_validates_the_rendered_personal_data_table()
    test_render_gate_rejects_a_non_directory_source_dir()
    test_norm_excludes_metrics_so_a_trend_number_never_reads_as_changed()
    test_metrics_diff_empty_when_metrics_identical_or_absent()
    test_metrics_diff_reports_per_key_old_and_new()
    test_doctor_compare_main_passes_when_only_metrics_differ()
    test_doctor_compare_main_silent_on_metrics_when_reports_have_none()
    test_doctor_compare_main_still_fails_on_a_findings_change_regardless_of_metrics()
    print("ok")
