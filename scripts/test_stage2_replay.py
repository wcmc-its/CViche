#!/usr/bin/env python3
"""Self-tests for stage2_replay.py, on synthetic runs only. No LLM call.

    python3 scripts/test_stage2_replay.py

Pins the prompt-log pairing, the replay (a stored run reproduces itself; a
prompt with no stored reply is counted, never answered; the forced-failure
arm fails every later batch), the gate's criteria, the relocation check and
the census.
"""
import contextlib
import io
import json
import logging
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import stage2_replay as sr  # noqa: E402
from docx import Document  # noqa: E402

stage2 = sr._import_stage2(sr._DEFAULT_SRC)
logging.getLogger(stage2.logger.name).setLevel(logging.ERROR)  # the miss and failure warnings are expected here

UID = "ZZZZZZ"
RECORD = 7  # lines per synthetic record; 63 lines put one hard cut 1 line into record 7


def _phase_reply(prompt: str) -> str:
    """Group every RECORD listed lines from the batch's first, like the model."""
    listed = [int(i) for i in sr._PROMPT_INDEX.findall(prompt)]
    groups = [listed[i:i + RECORD] for i in range(0, len(listed), RECORD)]
    return json.dumps([{"element_idx_start": g[0], "element_idx_end": g[-1], "element_type": "paragraph",
                        "confidence": 0.9} for g in groups])


def _make_run(root: Path, carry_max: int, n_lines: int = 63) -> Path:
    """A run dir in the S3 layout, its prompt logs and entries written by stage 2
    itself (carry set to ``carry_max``) against the phase model."""
    run = root / UID
    for sub in ("input", "outputs", "prompt_logs"):
        (run / sub).mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("MENTORING")
    for i in range(n_lines):
        doc.add_paragraph(f"Field {i % RECORD} of record {i // RECORD}")
    doc.save(run / "input" / f"{UID}.docx")
    hierarchy = run / "outputs" / f"{UID}_hierarchy_mapped.json"
    hierarchy.write_text(json.dumps({
        "document_uid": UID,
        "hierarchy_with_indices": [{"text": "Mentoring", "level": "H1", "element_idx": 0, "children": []}],
        "section_boundaries": [{"hierarchy": ["Mentoring"], "element_idx_start": 0, "element_idx_end": n_lines,
                                "has_children": False}],
    }))
    calls = []

    def logging_model(stage: str, messages: list, response_format: object = None, **_kw: object) -> dict:
        log_id = f"{len(calls):012x}"
        reply = _phase_reply(messages[-1]["content"])
        calls.append(log_id)
        logs = run / "prompt_logs"
        (logs / f"2026-01-01_00-00-00_stage_2_{log_id}.json").write_text(json.dumps(
            {"log_id": log_id, "messages": messages}))
        (logs / f"2026-01-01_00-00-01_stage_2_{log_id}_RESPONSE.json").write_text(json.dumps(
            {"log_id": log_id, "response": {"choices": [{"message": {"content": reply}}],
                                            "usage": {"total_tokens": 10, "prompt_tokens": 8,
                                                      "completion_tokens": 2},
                                            "cost": 0.01}}))
        return {"content": reply, "cost": 0.01, "total_tokens": 10, "prompt_tokens": 8, "completion_tokens": 2}

    real_call, real_om, real_max = stage2.call_llm, stage2.OutputManager, stage2.STAGE2_CARRY_MAX
    stage2.call_llm, stage2.STAGE2_CARRY_MAX = logging_model, carry_max
    stage2.OutputManager = lambda path, base_output_dir=None: real_om(path, base_output_dir=root / "_out")
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            output, _ = stage2.run_stage_2(str(run / "input" / f"{UID}.docx"), str(hierarchy), workers=1)
    finally:
        stage2.call_llm, stage2.OutputManager, stage2.STAGE2_CARRY_MAX = real_call, real_om, real_max
    (run / "outputs" / f"{UID}_entries.json").write_text(json.dumps(output))
    (run / "prompt_logs" / "2026-01-01_00-00-00_stage_3b_aaaaaaaaaaaa.json").write_text("{}")  # not stage 2
    return run


def _replay(run: Path, out: Path, fail_later: bool = False, carry_max: int | None = None) -> dict:
    real_max = stage2.STAGE2_CARRY_MAX
    if carry_max is not None:
        stage2.STAGE2_CARRY_MAX = carry_max
    try:
        out.mkdir(parents=True, exist_ok=True)
        return sr.replay_run(stage2, run, out, fail_later)
    finally:
        stage2.STAGE2_CARRY_MAX = real_max


def test_stored_replies_pair_by_log_id_and_key_by_the_user_prompt():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        replies, paired = sr.stored_replies(run / "prompt_logs")
        assert paired == 2 and len(replies) == 2, (paired, len(replies))
        assert all(prompt.startswith("## Input Data") for prompt in replies)
        assert all(answers[0]["cost"] == 0.01 and answers[0]["total_tokens"] == 10 for answers in replies.values())


def test_a_stored_run_replays_to_its_own_entries_and_a_miss_is_counted():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        base = _replay(run, Path(tmp) / "base", carry_max=0)
        assert base["matches_stored_entries"] and base["served"] == 2 and base["misses"] == 0, base
        assert [len(s["batches"]) for s in base["sections"]] == [2]
        # The carry re-opens batch 2 on record 7: a prompt the stored run never sent.
        new = _replay(run, Path(tmp) / "new")
        assert (new["served"], new["misses"]) == (1, 1), new
        assert new["sections"][0]["batches"][1]["outcome"] == "miss"
        assert new["sections"][0]["batches"][1]["first"] == "50"  # element 50 = record 7's head
        assert new["batch_carry"]["carried"] == 1
        assert (Path(tmp) / "new" / f"{UID}.json").exists()


def test_the_forced_failure_arm_fails_every_later_batch():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        failed = _replay(run, Path(tmp) / "fail", fail_later=True, carry_max=0)
        assert (failed["served"], failed["forced_failures"], failed["misses"]) == (1, 1, 0), failed
        assert [b["outcome"] for b in failed["sections"][0]["batches"]] == ["served", "forced_failure"]


def test_gate_passes_the_carry_and_fails_a_change_outside_the_carried_region():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        arms = {"base": _replay(run, Path(tmp) / "base", carry_max=0),
                "new": _replay(run, Path(tmp) / "new"),
                "base_fail": _replay(run, Path(tmp) / "base_fail", fail_later=True, carry_max=0),
                "new_fail": _replay(run, Path(tmp) / "new_fail", fail_later=True)}
        row = sr.gate_run(arms["base"], arms["new"], arms["base_fail"], arms["new_fail"])
        assert sr.gate_passes(row), row
        assert sr.carried_regions(arms["base"], arms["new"]) == [(50, 63)]
        assert row["fail_entries_identical"] and row["fail_lines_lost"] == 0

        moved = json.loads(json.dumps(arms["new"]))
        first = next(e for e in moved["entries"] if e["element_type"] == "paragraph")
        first["element_idx_end"] += 1  # an entry before the carry changes
        assert sr.gate_run(arms["base"], moved, arms["base_fail"], arms["new_fail"])[
            "changed_entries_outside_carry"] == 2
        assert not sr.gate_passes(sr.gate_run(arms["base"], moved, arms["base_fail"], arms["new_fail"]))

        lossy = json.loads(json.dumps(arms["new_fail"]))
        lossy["entries"] = [e for e in lossy["entries"] if e["element_idx_start"] != 1]
        row = sr.gate_run(arms["base"], arms["new"], arms["base_fail"], lossy)
        assert row["fail_lines_lost"] == RECORD and not sr.gate_passes(row), row

        # A section the base arm could not replay is excluded whole.
        missed = json.loads(json.dumps(arms["base"]))
        missed["sections"][0]["batches"][0]["outcome"] = "miss"
        assert sr.carried_regions(missed, arms["new"]) == [(0, 63)]

        drifted = json.loads(json.dumps(arms["new"]))
        drifted["sections"][0]["batches"][0]["sha"] = "0" * 16
        assert sr.gate_run(arms["base"], drifted, arms["base_fail"], arms["new_fail"])["first_batch_prompt_diffs"] == 1


def test_gate_and_same_cli_exit_codes():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        root = Path(tmp)
        for name, fail_later, carry_max in (("base", False, 0), ("new", False, None),
                                            ("base_fail", True, 0), ("new_fail", True, None),
                                            ("nocarry", False, 0)):
            _replay(run, root / name, fail_later, carry_max)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            assert sr.main(["gate", "--base", str(root / "base"), "--new", str(root / "new"),
                            "--base-fail", str(root / "base_fail"), "--new-fail", str(root / "new_fail")]) == 0
            assert sr.main(["same", "--a", str(root / "base"), "--b", str(root / "nocarry")]) == 0
            assert sr.main(["same", "--a", str(root / "base"), "--b", str(root / "new")]) == 1
            broken = json.loads((root / "new_fail" / f"{UID}.json").read_text())
            broken["entries"] = [e for e in broken["entries"] if e["element_idx_start"] != 1]
            (root / "broken").mkdir()
            (root / "broken" / f"{UID}.json").write_text(json.dumps(broken))
            assert sr.main(["gate", "--base", str(root / "base"), "--new", str(root / "new"),
                            "--base-fail", str(root / "base_fail"), "--new-fail", str(root / "broken")]) == 1
        assert f"{UID}\tTrue" in out.getvalue()
        (root / "empty").mkdir()
        with contextlib.redirect_stderr(io.StringIO()):
            assert sr.main(["gate", "--base", str(root / "empty"), "--new", str(root / "new"),
                            "--base-fail", str(root / "base_fail"), "--new-fail", str(root / "new_fail")]) == 2


def test_entry_lines_cover_paragraph_ranges_and_row_spans():
    assert sr._entry_lines({"element_idx_start": 3, "element_idx_end": 5}) == ["3", "4", "5"]
    assert sr._entry_lines({"element_idx_start": "474.9", "element_idx_end": "474.11"}) == [
        "474.9", "474.10", "474.11"]
    assert sr._entry_lines({"element_idx_start": "474.2", "element_idx_end": "475.5"}) == ["474.2"]
    covered, multi = sr.line_census([
        {"element_type": "paragraph", "element_idx_start": 1, "element_idx_end": 2},
        {"element_type": "paragraph", "element_idx_start": 2, "element_idx_end": 3},
        {"element_type": "break", "element_idx_start": 4, "element_idx_end": 4},
    ])
    assert covered == {"1", "2", "3"} and multi == {"2"}


def test_census_counts_the_edge_and_the_carry_from_stored_entries():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        uid, docx, hierarchy, entries = sr._run_files(run)
        row = sr.census_cv(stage2, docx, hierarchy, entries)
        # Today: batches of 50 and 13. Stored batch 1 ends on record 7's head (element 50).
        assert (row["edges"], row["multi_batch_sections"], row["calls_today"], row["elements_today"]) == (1, 1, 2, 63)
        assert row["carries"] == [1] and row["skipped"] == {}
        assert (row["calls_carry"], row["elements_carry"]) == (2, 64)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            assert sr.run_census(stage2, [(uid, docx, hierarchy, entries)]) == 0
        assert "elements sent 63 -> 64 (+1, 1.6%)" in out.getvalue()
        with contextlib.redirect_stderr(io.StringIO()):
            assert sr.run_census(stage2, []) == 2


def _arms(run: Path, root: Path) -> dict:
    return {"base": _replay(run, root / "base", carry_max=0),
            "new": _replay(run, root / "new"),
            "base_fail": _replay(run, root / "base_fail", fail_later=True, carry_max=0),
            "new_fail": _replay(run, root / "new_fail", fail_later=True)}


def _copy(record: dict) -> dict:
    return json.loads(json.dumps(record))


def _failed_criteria(row: dict) -> list[str]:
    """The gate criteria a row fails, judged here independently of gate_passes."""
    checks = {
        "first_batch": row["first_batch_prompt_diffs"] == 0,
        "single_batch": row["single_batch_prompt_diffs"] == 0,
        "outside_carry": row["changed_entries_outside_carry"] == 0,
        "unaccounted_fail_arms": row["unaccounted_same_fail_arms"],
        "unaccounted_vs_base": row["unaccounted_same_as_base"],
        "lost": row["fail_lines_lost"] == 0,
        "gained": row["fail_lines_gained"] == 0,
        "multi_claimed": row["fail_multi_claimed_base"] == row["fail_multi_claimed_new"],
    }
    return sorted(name for name, ok in checks.items() if not ok)


def _gate(arms: dict, **replace: dict) -> dict:
    merged = {**arms, **replace}
    return sr.gate_run(merged["base"], merged["new"], merged["base_fail"], merged["new_fail"])


def test_a_replay_miss_raises_and_a_repeated_prompt_is_answered_in_log_order():
    answer = {"cost": 0.0, "total_tokens": 0, "prompt_tokens": 0, "completion_tokens": 0}
    replayer = sr._Replayer({"P": [{**answer, "content": "a"}, {**answer, "content": "b"}]}, False)
    replayer.sections.append({"start": 0, "end": 1, "batches": []})
    asked = [{"role": "user", "content": "P"}]
    assert [replayer.call_llm("stage_2", asked)["content"] for _ in range(3)] == ["a", "b", "b"]
    try:
        replayer.call_llm("stage_2", [{"role": "user", "content": "never sent"}])
    except sr.ReplayMiss:
        pass
    else:
        raise AssertionError("a prompt with no stored reply was answered")
    assert (replayer.served, replayer.misses) == (3, 1)
    assert [b["outcome"] for b in replayer.sections[0]["batches"]] == ["served"] * 3 + ["miss"]


def test_a_replay_that_differs_from_the_stored_entries_is_reported():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        stored = run / "outputs" / f"{UID}_entries.json"
        data = json.loads(stored.read_text())
        data["entries"] = data["entries"][1:]
        stored.write_text(json.dumps(data))
        assert not _replay(run, Path(tmp) / "base", carry_max=0)["matches_stored_entries"]


def test_each_gate_criterion_fails_the_gate_on_its_own():
    with tempfile.TemporaryDirectory() as tmp:
        arms = _arms(_make_run(Path(tmp), carry_max=0), Path(tmp))
        assert _failed_criteria(_gate(arms)) == [] and sr.gate_passes(_gate(arms))

        # A section one batch long in the base arm that the new arm sends as two.
        base, new = _copy(arms["base"]), _copy(arms["new"])
        base["sections"][0]["batches"] = base["sections"][0]["batches"][:1]
        new["sections"][0]["batches"][0] = base["sections"][0]["batches"][0]
        cases = {"single_batch": {"base": base, "new": new}}

        gained = _copy(arms["new_fail"])
        gained["entries"].append({"element_type": "paragraph", "element_idx_start": 60, "element_idx_end": 60})
        cases["gained"] = {"new_fail": gained}

        twice = _copy(arms["new_fail"])
        twice["entries"].append(next(e for e in twice["entries"] if e["element_type"] == "paragraph"))
        cases["multi_claimed"] = {"new_fail": twice}

        fail_arm = _copy(arms["base_fail"])
        fail_arm["coverage"]["unaccounted_indices"] = [999]
        cases["unaccounted_fail_arms"] = {"base_fail": fail_arm}

        served = _copy(arms["base"])
        served["coverage"]["unaccounted_indices"] = [999]
        cases["unaccounted_vs_base"] = {"base": served}

        for criterion, replace in cases.items():
            row = _gate(arms, **replace)
            assert _failed_criteria(row) == [criterion], (criterion, row)
            assert not sr.gate_passes(row), criterion


def test_the_gate_reports_served_line_changes_and_restores_without_gating_them():
    with tempfile.TemporaryDirectory() as tmp:
        arms = _arms(_make_run(Path(tmp), carry_max=0), Path(tmp))
        # The served new arm misses batch 2 (a prompt the stored run never sent):
        # lines 51-63 are breaks there, batch 1's held-back record 50 is restored.
        row = _gate(arms)
        assert (row["served_lines_lost"], row["served_lines_gained"]) == (13, 0)
        assert (row["new_restored"], row["new_failed_batches"]) == (1, 1)
        assert sr.gate_passes(row)

        new = _copy(arms["base"])
        new["entries"] = [e for e in new["entries"] if e["element_idx_start"] != 51]  # lines 51-57
        new["entries"].append({"element_type": "paragraph", "element_idx_start": 0, "element_idx_end": 0})
        new["batch_carry"] = {"restored_unclaimed": 2, "restored_after_failure": 1, "failed_batches": 4}
        row = _gate(arms, new=new)
        assert (row["served_lines_lost"], row["served_lines_gained"]) == (RECORD, 1)
        assert (row["new_restored"], row["new_failed_batches"]) == (3, 4)
        new["batch_carry"] = None
        assert (_gate(arms, new=new)["new_restored"], _gate(arms, new=new)["new_failed_batches"]) == (0, 0)


def test_an_entry_on_the_carried_regions_first_line_is_inside_it():
    with tempfile.TemporaryDirectory() as tmp:
        arms = _arms(_make_run(Path(tmp), carry_max=0), Path(tmp))
        assert sr.carried_regions(arms["base"], arms["new"]) == [(50, 63)]
        new = _copy(arms["new"])
        head = next(e for e in new["entries"] if e["element_idx_start"] == 50)
        head["text"] += " (changed)"
        assert _gate(arms, new=new)["changed_entries_outside_carry"] == 0


def test_the_carried_region_starts_at_the_first_batch_whose_prompt_differs():
    def section(shas):
        return {"start": 0, "end": 101, "batches": [{"first": str(20 * k), "sha": sha, "outcome": "served"}
                                                    for k, sha in enumerate(shas)]}
    base, new = {"sections": [section("abc")]}, {"sections": [section("abd")]}
    assert sr.carried_regions(base, new) == [(40, 100)]
    assert sr.carried_regions(base, {"sections": [section("abc")]}) == []
    assert sr.carried_regions(base, {"sections": [section("abcd")]}) == [(60, 100)]


def test_same_exits_1_on_an_entry_only_or_a_coverage_only_difference():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        run = _make_run(root, carry_max=0)
        _replay(run, root / "base", carry_max=0)
        record = json.loads((root / "base" / f"{UID}.json").read_text())
        entries_only, coverage_only = _copy(record), _copy(record)
        entries_only["entries"] = entries_only["entries"][1:]
        coverage_only["coverage"]["unaccounted_indices"] = [999]
        with contextlib.redirect_stdout(io.StringIO()):
            for name, changed in (("entries_only", entries_only), ("coverage_only", coverage_only)):
                (root / name).mkdir()
                (root / name / f"{UID}.json").write_text(json.dumps(changed))
                assert sr.main(["same", "--a", str(root / "base"), "--b", str(root / name)]) == 1, name
        (root / "none").mkdir()
        with contextlib.redirect_stderr(io.StringIO()):
            assert sr.main(["same", "--a", str(root / "none"), "--b", str(root / "base")]) == 2


def test_census_counts_no_edge_for_a_single_batch_section_and_estimates_the_extra_cost():
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0, n_lines=20)
        row = sr.census_cv(stage2, *sr._run_files(run)[1:])
        assert (row["edges"], row["multi_batch_sections"], row["calls_today"], row["carries"]) == (0, 0, 1, [])
    with tempfile.TemporaryDirectory() as tmp:
        run = _make_run(Path(tmp), carry_max=0)
        entries = run / "outputs" / f"{UID}_entries.json"
        data = json.loads(entries.read_text())
        data["total_cost"] = 63.0  # one extra element of 63 sent: an estimated extra $1.00
        entries.write_text(json.dumps(data))
        with contextlib.redirect_stdout(io.StringIO()) as out:
            assert sr.run_census(stage2, [sr._run_files(run)]) == 0
        assert "stage-2 cost $63.00  estimated extra $1.00" in out.getvalue()


def test_percentile_takes_the_value_at_the_quantile_not_the_maximum():
    assert sr._percentile([4, 1, 3, 2], 0.5) == 3
    assert sr._percentile([4, 1, 3, 2], 0.9) == 4
    assert sr._percentile([], 0.5) == 0


if __name__ == "__main__":
    for test in (test_stored_replies_pair_by_log_id_and_key_by_the_user_prompt,
                 test_a_stored_run_replays_to_its_own_entries_and_a_miss_is_counted,
                 test_the_forced_failure_arm_fails_every_later_batch,
                 test_gate_passes_the_carry_and_fails_a_change_outside_the_carried_region,
                 test_gate_and_same_cli_exit_codes,
                 test_entry_lines_cover_paragraph_ranges_and_row_spans,
                 test_census_counts_the_edge_and_the_carry_from_stored_entries,
                 test_a_replay_miss_raises_and_a_repeated_prompt_is_answered_in_log_order,
                 test_a_replay_that_differs_from_the_stored_entries_is_reported,
                 test_each_gate_criterion_fails_the_gate_on_its_own,
                 test_the_gate_reports_served_line_changes_and_restores_without_gating_them,
                 test_an_entry_on_the_carried_regions_first_line_is_inside_it,
                 test_the_carried_region_starts_at_the_first_batch_whose_prompt_differs,
                 test_same_exits_1_on_an_entry_only_or_a_coverage_only_difference,
                 test_census_counts_no_edge_for_a_single_batch_section_and_estimates_the_extra_cost,
                 test_percentile_takes_the_value_at_the_quantile_not_the_maximum):
        test()
        print(f"ok {test.__name__}")
