#!/usr/bin/env python3
"""End-to-end self-tests for render_gate.main(), against a stub stage 6.

    python3 scripts/test_render_gate_integration.py

scripts/test_render_doctor_gates.py pins each helper in isolation; this file
pins what they do wired together, which is where a gate's real failure modes
live -- a crashed render scored as clean, a stale docx from the previous arm
surviving into this one, --source-dir resolving a document and then not
forwarding it. None of those are visible from a helper's own unit test.

The seam is that render_gate imports stage 6 INSIDE main() rather than at
module import, so a stub can stand in for a 9,500-line renderer and every case
below runs in milliseconds with no corpus, no template and no LLM. Replacing
it takes both halves of the binding, measured rather than assumed: `import a.b
as c` consults sys.modules['a.b'] only while the package has no `b` attribute,
and one earlier import of the real module leaves one behind. That happens here
for real: render_gate's own _resolve_source_docx imports
unified_pipeline.run_doctor, which pulls stage 6 in transitively and reads
RENDER_ROUTED_CODES off it. So run_doctor is imported below, before any stub is
installed -- otherwise the --source-dir cases resolve that import against the
stub and die on the missing name.

The stub's WCMTemplateGenerator raises on construction. That is the assertion
behind review point 2 on #740: every render must reach stage 6 through
run_stage6(), the entry point production uses, and not by building a generator
of its own.
"""
import contextlib
import io
import json
import os
import sys
import tempfile
import types
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# Same reason as test_render_doctor_gates.py: unified_pipeline normally
# arrives via PYTHONPATH=<arm>/src, and this file's "runnable standalone"
# contract should not depend on the caller setting that up.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import render_gate as rg  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import unified_pipeline  # noqa: E402
import unified_pipeline.run_doctor  # noqa: F401,E402
import unified_pipeline.stage_6_word_template as _REAL_STAGE6  # noqa: E402

_STAGE6_MODULE = "unified_pipeline.stage_6_word_template"


def _sentinel_call_llm(*a, **kw):
    """Stands in for stage 6's real call_llm, so restoration is observable."""
    return "the original call_llm"


def _make_stub(render):
    """A stand-in stage-6 module whose run_stage6 delegates to `render`.

    `render(stub, input_path, output_path, original_doc_path)` is where each
    test says what this arm's stage 6 does -- write a good docx, write nothing,
    raise, or reach for the LLM. Every call is recorded on `stub.calls` so a
    test can assert on what render_gate actually passed.
    """
    stub = types.ModuleType(_STAGE6_MODULE)
    stub.calls = []
    stub.call_llm = _sentinel_call_llm

    def run_stage6(input_path, output_path=None, verbose=True, original_doc_path=None, **kw):
        stub.calls.append({"input_path": input_path, "output_path": output_path,
                           "original_doc_path": original_doc_path,
                           "discover_original_doc": kw.get("discover_original_doc"),
                           "repair_passed": "repair_protected_data" in kw,
                           "repair_protected_data": kw.get("repair_protected_data"),
                           "subpoints_passed": "supplementary_subpoints" in kw})
        return render(stub, input_path, output_path, original_doc_path)

    def WCMTemplateGenerator(*a, **kw):
        raise AssertionError(
            "render_gate built a WCMTemplateGenerator directly -- every render "
            "must go through run_stage6(), the entry point production uses "
            "(review on #740)")

    stub.run_stage6 = run_stage6
    stub.WCMTemplateGenerator = WCMTemplateGenerator
    return stub


@contextlib.contextmanager
def _stage6(stub):
    """Install `stub` as unified_pipeline.stage_6_word_template, then put the
    real module back -- both in sys.modules and as the package attribute, since
    `import a.b as c` prefers the attribute whenever one exists."""
    sys.modules[_STAGE6_MODULE] = stub
    unified_pipeline.stage_6_word_template = stub
    try:
        yield stub
    finally:
        sys.modules[_STAGE6_MODULE] = _REAL_STAGE6
        unified_pipeline.stage_6_word_template = _REAL_STAGE6


def _write_personal_data_docx(path):
    """A docx shaped like a good render: the PERSONAL DATA table's label rows.

    _validate_rendered_docx reads exactly this, so a stub that writes it counts
    as a successful render and a stub that writes anything else does not.
    """
    doc = Document()
    doc.add_paragraph("PERSONAL DATA")
    table = doc.add_table(rows=len(rg.PERSONAL_DATA_LABELS), cols=2)
    for row, label in zip(table.rows, rg.PERSONAL_DATA_LABELS):
        row.cells[0].text = f"{label.title()}:"
    doc.save(str(path))


def _renders_a_good_docx(stub, input_path, output_path, original_doc_path):
    _write_personal_data_docx(Path(output_path))
    return output_path


def _make_arm(root, uids):
    """A minimal arm_outputs tree: one stage-4 artifact per uid."""
    s4 = root / "arm" / "stage_4_field_extraction"
    s4.mkdir(parents=True)
    for uid in uids:
        (s4 / f"{uid}_fields.json").write_text(
            json.dumps({"document_uid": uid, "entries": []}), encoding="utf-8")
    return root / "arm"


def _run_main(argv):
    """main()'s exit code and its output, kept off this file's own stdout so
    the runner (CI's pipeline-tests job) still reads a single "ok" line."""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
        code = rg.main(argv)
    return code, buf.getvalue()


def _index(out):
    return json.loads((out / "_render_index.json").read_text(encoding="utf-8"))


def test_a_successful_render_exits_zero_and_writes_the_docx():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_renders_a_good_docx)) as stub:
            code, _ = _run_main([str(arm), str(out)])

        assert code == 0, "a clean render must exit 0"
        assert (out / "u1_wcm.docx").is_file()
        assert _index(out)["u1"] == {"input": "u1_fields.json"}
        assert stub.calls[0]["output_path"] == str(out / "u1_wcm.docx")


def test_a_raising_renderer_fails_the_gate():
    def _boom(stub, input_path, output_path, original_doc_path):
        raise ValueError("renderer exploded")

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_boom)):
            code, output = _run_main([str(arm), str(out)])

        assert code == 1, "a crashed render must fail the gate (CODING_STANDARDS 5.5)"
        entry = _index(out)["u1"]
        assert entry["error"] == "ValueError: renderer exploded"
        assert "tb" in entry, "a crash records its traceback tail for triage"
        assert "failed uids: u1" in output


def test_a_renderer_that_writes_nothing_fails_the_gate():
    # "No exception" is not "rendered": without this the compare step sees a
    # missing file rather than a failure, which is defect 1's exact shape.
    def _writes_nothing(stub, input_path, output_path, original_doc_path):
        return output_path

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_writes_nothing)):
            code, _ = _run_main([str(arm), str(out)])

        assert code == 1
        assert _index(out)["u1"]["error"] == "render returned without writing an output file"


def test_a_structurally_broken_docx_fails_the_gate():
    # The file exists and is_file() is happy; it is simply not a WCM document.
    # Without the semantic check this scores as a render and gets compared.
    def _writes_junk(stub, input_path, output_path, original_doc_path):
        Path(output_path).write_text("not a docx at all", encoding="utf-8")
        return output_path

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_writes_junk)):
            code, _ = _run_main([str(arm), str(out)])

        assert code == 1
        assert "not a readable docx" in _index(out)["u1"]["error"]


def test_stale_output_from_a_previous_arm_is_removed():
    # Defect 1 in render_gate.py's docstring: a reused out_dir that keeps the
    # previous arm's files reports PASS on a run that rendered nothing.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        out.mkdir()
        stale = out / "gone_wcm.docx"
        stale.write_text("last arm's output", encoding="utf-8")

        with _stage6(_make_stub(_renders_a_good_docx)):
            code, _ = _run_main([str(arm), str(out)])

        assert code == 0
        assert not stale.exists(), "out_dir must be wiped before every run"
        assert sorted(p.name for p in out.iterdir()) == ["_render_index.json", "u1_wcm.docx"]


def test_the_render_index_records_every_uid():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1", "u2"]), root / "out"
        with _stage6(_make_stub(_renders_a_good_docx)):
            code, _ = _run_main([str(arm), str(out)])

        assert code == 0
        assert _index(out) == {"u1": {"input": "u1_fields.json"},
                               "u2": {"input": "u2_fields.json"}}


def test_the_resolved_source_docx_is_forwarded_as_original_doc_path():
    # The whole point of --source-dir (#550): resolving the document and then
    # not passing it on would leave the flag inert with nothing to show for it.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        sources = root / "sources"
        sources.mkdir()
        (sources / "u1_cv.docx").write_text("", encoding="utf-8")

        with _stage6(_make_stub(_renders_a_good_docx)) as stub:
            code, _ = _run_main([str(arm), str(out), "--source-dir", str(sources)])

        assert code == 0
        assert stub.calls[0]["original_doc_path"] == str(sources / "u1_cv.docx")
        assert _index(out)["u1"]["source_docx"] == "u1_cv.docx"


def test_repair_is_passed_to_run_stage6_only_with_the_flag():
    # #1389: --repair renders as CVICHE_RUN_REPAIR=1 does; without it the
    # keyword is not passed, so an arm whose run_stage6 predates it renders.
    for argv_tail, passed in (([], False), (["--repair"], True)):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            arm, out = _make_arm(root, ["u1"]), root / "out"
            with _stage6(_make_stub(_renders_a_good_docx)) as stub:
                code, _ = _run_main([str(arm), str(out), *argv_tail])
            assert code == 0
            assert stub.calls[0]["repair_passed"] is passed
            assert stub.calls[0]["repair_protected_data"] is (True if passed else None)


def test_subpoints_is_passed_to_run_stage6_only_with_the_flag():
    # #1205: --subpoints renders as CVICHE_SUPPLEMENTARY_SUBPOINTS=1 does; as
    # for --repair, without it the keyword is not passed at all.
    for argv_tail, passed in (([], False), (["--subpoints"], True)):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            arm, out = _make_arm(root, ["u1"]), root / "out"
            with _stage6(_make_stub(_renders_a_good_docx)) as stub:
                code, _ = _run_main([str(arm), str(out), *argv_tail])
            assert code == 0
            assert stub.calls[0]["subpoints_passed"] is passed


def test_a_uid_with_no_source_docx_still_renders_and_says_so():
    # web08 today: 65 of the farm's 66 uids have a source docx. A uid without
    # one is a one-line notice in the index, not an error, and stage 6 must be
    # handed None rather than a guessed path.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        sources = root / "sources"
        sources.mkdir()
        (sources / "someone_else_cv.docx").write_text("", encoding="utf-8")

        with _stage6(_make_stub(_renders_a_good_docx)) as stub:
            code, _ = _run_main([str(arm), str(out), "--source-dir", str(sources)])

        assert code == 0
        assert stub.calls[0]["original_doc_path"] is None
        assert _index(out)["u1"]["source_docx"] == "not found"


def test_one_uid_failing_does_not_stop_the_others():
    # A gate that stopped at the first bad CV would report a green arm as
    # "1 rendered" and hide the other 65 -- and it must still exit non-zero.
    def _fail_u2_only(stub, input_path, output_path, original_doc_path):
        if "u2" in Path(input_path).name:
            raise RuntimeError("u2 is broken")
        return _renders_a_good_docx(stub, input_path, output_path, original_doc_path)

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1", "u2", "u3"]), root / "out"
        with _stage6(_make_stub(_fail_u2_only)):
            code, output = _run_main([str(arm), str(out)])

        assert code == 1
        index = _index(out)
        assert set(index) == {"u1", "u2", "u3"}, "every uid must be attempted"
        assert index["u2"]["error"] == "RuntimeError: u2 is broken"
        assert "error" not in index["u1"] and "error" not in index["u3"]
        assert (out / "u1_wcm.docx").is_file() and (out / "u3_wcm.docx").is_file()
        assert "rendered=2 failed=1" in output


def test_a_uid_whose_output_validation_raises_does_not_kill_the_run():
    """Isolation has to cover the post-render checks, not just the render.

    test_one_uid_failing_does_not_stop_the_others only reaches a failure
    raised by run_stage6 itself, which is inside the per-uid try by
    construction. The checks that run AFTER a render read a document this arm
    just wrote and can raise on a shape no fixture anticipated; if that raise
    escapes the per-uid handler it aborts the whole gate at whichever uid hit
    it and the index is never written, so every later uid is lost and one bad
    CV reads as a crashed run. The validator is replaced here rather than fed
    a shape that happens to raise today, because the guarantee under test is
    the exception scope, not any particular malformed document.
    """
    real_validate = rg._validate_rendered_docx

    def _raises_for_u2(path):
        if "u2" in Path(path).name:
            raise IndexError("tuple index out of range")
        return real_validate(path)

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1", "u2", "u3"]), root / "out"
        rg._validate_rendered_docx = _raises_for_u2
        try:
            with _stage6(_make_stub(_renders_a_good_docx)):
                code, output = _run_main([str(arm), str(out)])
        finally:
            rg._validate_rendered_docx = real_validate

        assert code == 1, "a uid whose validation raises must fail the gate"
        index = _index(out)  # must survive a mid-run validation failure
        assert set(index) == {"u1", "u2", "u3"}, "u3 must still be attempted"
        assert "error" not in index["u1"] and "error" not in index["u3"]
        assert index["u2"]["error"].startswith("IndexError")
        assert "rendered=2 failed=1" in output


def test_a_table_row_with_no_cells_is_tolerated_not_a_validation_failure():
    # A <w:tr> carrying no <w:tc> is legal WordprocessingML. Indexing cells[0]
    # on one raises, so the validator skips such rows rather than reporting a
    # document that is actually fine as broken.
    def _u1_gets_a_cellless_row(stub, input_path, output_path, original_doc_path):
        _renders_a_good_docx(stub, input_path, output_path, original_doc_path)
        doc = Document(str(output_path))
        table = doc.tables[0]
        table._tbl.append(table._tbl.makeelement(qn("w:tr"), {}))
        doc.save(str(output_path))

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_u1_gets_a_cellless_row)):
            code, output = _run_main([str(arm), str(out)])
        assert code == 0, output
        assert "error" not in _index(out)["u1"]


def test_an_unsafe_out_directory_is_rejected_before_anything_is_deleted():
    # The destructive half of this script is one shutil.rmtree; the check that
    # protects it has to run at argument-parse time, or it runs too late. The
    # stub renders nothing at all here, so a guard that let one of these
    # through fails on the raise below rather than by writing a real document
    # into whatever directory it was about to wipe.
    def _must_not_be_reached(stub, input_path, output_path, original_doc_path):
        raise AssertionError("an unsafe out directory reached the render loop")

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm = _make_arm(root, ["u1"])
        artifact = arm / "stage_4_field_extraction" / "u1_fields.json"

        for bad, why in ((arm, "out == arm_outputs"),
                         (arm / "renders", "out inside arm_outputs"),
                         (root, "out containing arm_outputs")):
            with _stage6(_make_stub(_must_not_be_reached)):
                try:
                    _run_main([str(arm), str(bad)])
                except SystemExit as e:
                    assert e.code == 2, f"argparse parser.error exits 2, got {e.code}"
                else:
                    raise AssertionError(f"{why} was accepted")
            assert artifact.is_file(), f"{why} deleted the input tree"


def test_the_llm_is_disabled_during_the_render():
    # Determinism: a gate arm that made a real LLM call would differ from its
    # own re-run, and the compare step would read that as a code change.
    seen = []

    def _calls_the_llm(stub, input_path, output_path, original_doc_path):
        try:
            stub.call_llm("prompt")
        except RuntimeError as exc:
            seen.append(str(exc))
        return _renders_a_good_docx(stub, input_path, output_path, original_doc_path)

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_calls_the_llm)):
            code, _ = _run_main([str(arm), str(out)])

        assert code == 0
        assert seen and "LLM disabled for determinism" in seen[0], (
            "stage 6's call_llm was not neutralised during the render")


def test_the_llm_is_restored_after_main_returns():
    # The replacement is scoped, not permanent (review point 7 on #740): the
    # stage-6 module object outlives this run, and so would a call_llm left
    # raising for whatever imports it next.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm, out = _make_arm(root, ["u1"]), root / "out"
        with _stage6(_make_stub(_renders_a_good_docx)) as stub:
            _run_main([str(arm), str(out)])
            assert stub.call_llm is _sentinel_call_llm, (
                "main() left its LLM-disabling replacement installed")

        # And it is restored even when the run fails part-way through.
        def _boom(s, input_path, output_path, original_doc_path):
            raise ValueError("boom")

        with _stage6(_make_stub(_boom)) as stub:
            _run_main([str(arm), str(root / "out2")])
            assert stub.call_llm is _sentinel_call_llm


def test_a_uids_file_renders_a_trailing_space_uid_and_fails_on_an_unmatched_one():
    # Harvested uids can end in a space ("... CV "); a --uids-file line is
    # stripped, so an exact match dropped that CV and rendered a narrower arm
    # with exit 0. A requested uid the arm does not hold must fail the run.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm = _make_arm(root, ["u1", "u2 "])
        uids_file = root / "uids.txt"
        uids_file.write_text("u2 \n", encoding="utf-8")
        with _stage6(_make_stub(_renders_a_good_docx)):
            code, _ = _run_main([str(arm), str(root / "out"), "--uids-file", str(uids_file)])
        assert code == 0
        assert list(_index(root / "out")) == ["u2 "]

        uids_file.write_text("u2\nnot-in-this-arm\n", encoding="utf-8")
        with _stage6(_make_stub(_renders_a_good_docx)):
            code, output = _run_main([str(arm), str(root / "out2"), "--uids-file", str(uids_file)])
        assert code == 1, "a partial uid match must fail, not render fewer CVs"
        assert "not-in-this-arm" in output


def test_the_gate_never_lets_stage6_guess_a_source_docx():
    # #732: with or without --source-dir (and for a uid it has no docx for),
    # stage 6's SAMPLE_CV_DIR / CWD-relative fallback must be switched off,
    # or the render depends on where the gate was launched.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm = _make_arm(root, ["u1"])
        sources = root / "sources"
        sources.mkdir()
        for label, extra in (("no --source-dir", []),
                             ("--source-dir without a match", ["--source-dir", str(sources)])):
            with _stage6(_make_stub(_renders_a_good_docx)) as stub:
                code, _ = _run_main([str(arm), str(root / "out"), *extra])
            assert code == 0
            assert stub.calls[0]["discover_original_doc"] is False, label


def test_an_arm_with_no_uids_fails_the_gate():
    # An empty arm renders nothing; exit 0 would read as a clean A/B arm.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        arm = _make_arm(root, [])
        with _stage6(_make_stub(_renders_a_good_docx)):
            code, output = _run_main([str(arm), str(root / "out")])
        assert code == 1, "an arm with no uids must fail"
        assert "empty arm" in output


if __name__ == "__main__":
    test_a_successful_render_exits_zero_and_writes_the_docx()
    test_a_raising_renderer_fails_the_gate()
    test_a_renderer_that_writes_nothing_fails_the_gate()
    test_a_structurally_broken_docx_fails_the_gate()
    test_stale_output_from_a_previous_arm_is_removed()
    test_the_render_index_records_every_uid()
    test_the_resolved_source_docx_is_forwarded_as_original_doc_path()
    test_repair_is_passed_to_run_stage6_only_with_the_flag()
    test_a_uid_with_no_source_docx_still_renders_and_says_so()
    test_one_uid_failing_does_not_stop_the_others()
    test_a_uid_whose_output_validation_raises_does_not_kill_the_run()
    test_a_table_row_with_no_cells_is_tolerated_not_a_validation_failure()
    test_an_unsafe_out_directory_is_rejected_before_anything_is_deleted()
    test_the_llm_is_disabled_during_the_render()
    test_the_llm_is_restored_after_main_returns()
    test_a_uids_file_renders_a_trailing_space_uid_and_fails_on_an_unmatched_one()
    test_the_gate_never_lets_stage6_guess_a_source_docx()
    test_an_arm_with_no_uids_fails_the_gate()
    print("ok")
