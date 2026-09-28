#!/usr/bin/env python3
"""Generate template_boilerplate_phrases.json from every WCM faculty CV template
revision in `TEMPLATES`.

Faculty build their CVs on whichever revision they were sent, so a phrase in
ANY revision is boilerplate (#829). For each tracked template DOCX this walks
the document body in order
(paragraphs AND table cells, since most field labels live in tables), normalizes
each text block, and categorizes blocks into:

  * "instructions"     - directive sentences, field/column labels, and
                         subcategory prompt lines. These are DROP-eligible
                         boilerplate that pollutes parsed CV output.
  * "section_headers"  - top-level WCM section names (PERSONAL DATA, EDUCATION,
                         BIBLIOGRAPHY, etc.). Kept SEPARATE and NOT placed in the
                         drop set: dropping a real section header could disturb
                         downstream structure.

Output: src/unified_pipeline/core/template_boilerplate_phrases.json

Usage:
    python3 scripts/gen_template_boilerplate.py
"""

import json
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

REPO_ROOT = Path(__file__).resolve().parent.parent
KEY_FILES = REPO_ROOT / "key_files"
OUTPUT_PATH = (
    REPO_ROOT
    / "src"
    / "unified_pipeline"
    / "core"
    / "template_boilerplate_phrases.json"
)

# Top-level WCM section names. These are kept OUT of the drop set so that an
# entry that is genuinely a section header is never dropped by the filter.
# Listed verbatim (normalized to collapsed whitespace, exact casing) from the
# October 2022 faculty template.
WCM_SECTION_HEADERS = [
    "PERSONAL DATA",
    "EDUCATION",
    "POSTDOCTORAL TRAINING (Include residency/fellowships)",
    "PROFESSIONAL POSITIONS & EMPLOYMENT",
    "EMPLOYMENT STATUS",
    "LICENSURE, BOARD CERTIFICATION",
    "INSTITUTIONAL/HOSPITAL AFFILIATION",
    "HONORS, AWARDS",
    "PROFESSIONAL ORGANIZATIONS AND SOCIETY MEMBERSHIPS",
    "PERCENT EFFORT AND INSTITUTIONAL RESPONSIBILITIES",
    "EDUCATIONAL CONTRIBUTIONS",
    "CLINICAL PRACTICE, INNOVATION, and LEADERSHIP",
    "RESEARCH",
    "MENTORING",
    "INSTITUTIONAL LEADERSHIP ACTIVITIES",
    "INSTITUTIONAL ADMINISTRATIVE ACTIVITIES",
    "EXTRAMURAL PROFESSIONAL RESPONSIBILITIES",
    "INVITATIONS TO SPEAK/PRESENT",
    "BIBLIOGRAPHY",
]

# The same, for the older revisions (#829). A CV built on one of them carries
# these as its real section headers, so they must be protected too.
WCM_2020_SECTION_HEADERS = [
    "GENERAL INFORMATION",
    "EDUCATIONAL BACKGROUND",
    "LICENSURE, BOARD CERTIFICATION, MALPRACTICE",
    "PROFESSIONAL POSITIONS AND EMPLOYMENT",
    "INSTITUTIONAL/HOSPITAL AFFILIATION",
    "PERCENT EFFORT AND INSTITUTIONAL RESPONSIBILITIES",
    "INSTITUTIONAL RESPONSIBILITIES",
    "RESEARCH SUPPORT",
    "EXTRAMURAL PROFESSIONAL RESPONSIBILITIES",
    "PROFESSIONAL MEMBERSHIPS",
    "HONORS AND AWARDS",
    "BIBLIOGRAPHY",
]
WCM_2012_SECTION_HEADERS = [
    "A. GENERAL INFORMATION",
    "B. EDUCATIONAL BACKGROUND",
    "PROFESSIONAL POSITIONS AND EMPLOYMENT",
    "LICENSURE, BOARD CERTIFICATION, MALPRACTICE",
    "F. HONORS AND AWARDS",
    "G. INSTITUTIONAL/HOSPITAL AFFILIATION",
    "H. EMPLOYMENT STATUS",
    "I. CURRENT AND PAST INSTITUTIONAL RESPONSIBILITIES AND",
    "PERCENT EFFORT",
    "K. EXTRAMURAL PROFESSIONAL RESPONSIBILITIES",
    "L. BIBLIOGRAPHY",
]

# Every tracked revision, newest first: (file under key_files/, its section
# headers). The 2012 file is the circulated .doc converted to .docx (macOS
# `textutil -convert docx`), since python-docx cannot read .doc.
TEMPLATES = [
    ("wcm_cv_template_faculty_october_2022_final.docx", WCM_SECTION_HEADERS),
    ("wcm_cv_template_for_website_2020.docx", WCM_2020_SECTION_HEADERS),
    ("curriculum_vitae_format_2012.docx", WCM_2012_SECTION_HEADERS),
]


def _normalize(text: str) -> str:
    """Collapse all whitespace to single spaces and strip ends."""
    return re.sub(r"\s+", " ", text).strip()


def _iter_block_items(parent):
    """Yield Paragraph and Table children of *parent* in document order.

    *parent* is either the Document or a _Cell. Mirrors python-docx internals so
    we walk paragraphs and tables in the exact order they appear in the body.
    """
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    else:
        # _Cell
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _collect_blocks(parent, blocks):
    """Recursively collect non-empty text blocks from paragraphs and table cells."""
    for block in _iter_block_items(parent):
        if isinstance(block, Paragraph):
            if block.text and block.text.strip():
                blocks.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if para.text and para.text.strip():
                            blocks.append(para.text)
                    # Nested tables inside the cell, if any.
                    for nested in cell.tables:
                        _collect_blocks(nested, blocks)


def _is_section_header(normalized: str, section_set_norm: set) -> bool:
    """Return True if *normalized* is a top-level WCM section name.

    Heuristic: lines that exactly match a known WCM top-level section header
    (normalized, case-SENSITIVE). We compare against the curated list rather
    than a pure regex so we never accidentally classify a real CV line as a
    header. Case sensitivity matters: the percent-effort table has a column
    label "Research" which must NOT be treated as the "RESEARCH" section header
    (it stays in the instructions/field-label set instead).
    """
    return normalized in section_set_norm


def main():
    section_set_norm = {
        _normalize(h) for _, headers in TEMPLATES for h in headers
    }

    # Normalize and dedupe across all revisions, preserving order: newest
    # template first, so each older one only appends what it adds.
    seen = set()
    instructions = []
    section_headers = []
    for filename, _ in TEMPLATES:
        path = KEY_FILES / filename
        if not path.exists():
            raise SystemExit(f"Template not found: {path}")
        blocks = []
        _collect_blocks(Document(str(path)), blocks)
        added = 0
        for raw in blocks:
            norm = _normalize(raw)
            if not norm or norm in seen:
                continue
            seen.add(norm)
            added += 1
            if _is_section_header(norm, section_set_norm):
                section_headers.append(norm)
            else:
                instructions.append(norm)
        print(f"{filename}: {len(blocks)} blocks, {added} new")

    output = {
        "source": [filename for filename, _ in TEMPLATES],
        "instructions": instructions,
        "section_headers": section_headers,
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
        f.write("\n")

    print(f"instructions: {len(instructions)}, section_headers: {len(section_headers)}")
    print(f"Wrote: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
