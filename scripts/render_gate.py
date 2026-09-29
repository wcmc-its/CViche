#!/usr/bin/env python3
"""Render every corpus CV through stage 6 from one code arm, into a fresh out_dir.

    PYTHONPATH=<arm>/src python3 scripts/render_gate.py <arm_outputs_dir> <out_dir> [uid ...]
    PYTHONPATH=<arm>/src python3 scripts/render_gate.py <arm_outputs_dir> <out_dir> --uids-file FILE

--uids-file reads one uid per line instead of positional args: this corpus's
uids are derived from real filenames and routinely contain spaces and
parentheses ("038WKA_Jonathan Nahmias CV (Updated 1-26-26)"), which shell
word-splitting on a $(...) uid list silently mangles into unmatched
fragments -- the run then just renders fewer uids than asked, with no error.

<arm_outputs_dir> is a stage-output tree shaped like src/unified_pipeline/outputs/
(stage_4_field_extraction/, stage_5d_citation_formatted/, etc.) -- either that
directory itself, or a farm built by symlinking/copying it from elsewhere. LLM
calls are neutralised so the render is deterministic; both call sites in stage 6
are try/except guarded, so raising is safe.

Companion: render_gate_compare.py, which reads two out_dirs (or one twice, as a
determinism control) and is the thing that actually passes or fails.

--source-dir DIR points this at the directory holding <uid>*.docx (e.g.
data/sample_cvs/word/) so stage 6's personal-data fallback (#550) can recover
contact fields from the original document. Without it -- the default -- this
gate could not see that path at all: SAMPLE_CV_DIR auto-discovery
(stage_6_word_template.py:716-727) resolves for none of the farm uids in a
fresh worktree, so every render took the fallback's "no original doc" branch
and #550 was corpus-unprovable. Omitting the flag, or passing it for a uid
with no matching docx, renders identically to today -- opt-in, and a missing
docx is a one-line notice in the index, not an error. Mirrors
doctor_gate.py's --source-dir: same fail-closed non-directory check, same
_uid_owns boundary rule (a shorter uid must not match a longer uid's file).

With the flag this gate reproduces a live run: run_stage6() takes an
original_doc_path parameter and forwards it to generate(), and both drivers
pass the resolved source path the same way (#550), so a delta this flag
surfaces is a delta a real CV render produces too.

Fixes two defects found in an earlier, uncommitted version of this script
(docs/analysis/HANDOFF-wave1-completion-2026-08-11.md, issue #584):

1. FAILS OPEN ON A CRASHED RENDER. The old version did `OUT.mkdir(exist_ok=True)`
   and never cleared the directory -- re-rendering into a reused out_dir after a
   change that crashes every CV left the *previous* arm's docx files in place,
   and the compare step (seeing files that matched) reported PASS on a render
   that produced zero real output. Fixed: OUT is wiped before every run, and
   this script itself exits non-zero if any render failed (CODING_STANDARDS
   Section 5.5 -- a gate fails closed, "nothing to compare" is a failure).
2. DATE-SENSITIVE. Stage 6 stamps `Date of preparation: {datetime.now()}` (and
   two more datetime.now() call sites reclassify content by calendar date), so
   two arms rendered on different days report differences that are not a
   regression. Not fixable in the renderer -- the operational rule is: render
   both arms on the SAME day, and if a gate run crosses midnight, re-render the
   reference rather than trusting a stale one. render_gate_compare.py's
   fingerprint check masks date text for exactly this reason.
"""
import argparse
import contextlib
import json
import shutil
import sys
import tempfile
import traceback
from pathlib import Path

from docx import Document

PRECEDENCE = [
    ("stage_5d_citation_formatted", "{uid}_citation_formatted.json"),
    ("stage_5c_teaching_formatted", "{uid}_teaching_formatted.json"),
    ("stage_5b_institution_enrichment", "{uid}_institution_enriched.json"),
    ("stage_5_enrichment", "{uid}_enriched.json"),
    ("stage_4_field_extraction", "{uid}_fields.json"),
]

# The cell stage 6 itself locates the PERSONAL DATA table by
# (_find_table_with_cell_text("Work email:"), stage6/sections/personal_data.py:369)
# -- matched here rather than a table index, so the gate and the renderer are
# provably talking about the same table.
PERSONAL_DATA_ANCHOR = "work email:"

# The label rows that table carries, read off the template actually in the repo
# (key_files/wcm_cv_template_faculty_october_2022_final.docx, rows 0-5; rows 6-7
# are the two visa questions no renderer touches). Six labels, not a row count:
# the table gains a row whenever the template does, and asserting eight would
# fail the next template revision for no reason.
PERSONAL_DATA_LABELS = (
    "office address",
    "office telephone",
    "work email",
    "home address",
    "cell phone",
    "personal email",
)


def _validate_output_dir(out: Path, *inputs) -> str | None:
    """Why `out` is unsafe to wipe, or None when it is safe. Deletes nothing.

    Module level, and called from _parse_args rather than from main(), so a bad
    invocation exits 2 before shutil.rmtree is reached at all -- a check sitting
    next to the deletion is worth little, because by then the argument has
    already been believed. Both directions are rejected, not just equality:
    `out` inside an input tree destroys part of that tree, and an input tree
    inside `out` destroys all of it. Every path is resolve()d first -- resolve()
    is non-strict, so it works on the not-yet-created `out`, and it collapses
    the ".." and the symlink a string comparison would wave through.
    """
    out_resolved = Path(out).resolve()
    for other in inputs:
        if other is None:
            continue
        other_resolved = Path(other).resolve()
        if out_resolved == other_resolved:
            relation = "is"
        elif out_resolved.is_relative_to(other_resolved):
            relation = "sits inside"
        elif other_resolved.is_relative_to(out_resolved):
            relation = "contains"
        else:
            continue
        return (f"out directory {out} {relation} {other} -- out is wiped "
                f"(shutil.rmtree) at the start of every run, so it must not be, "
                f"sit inside, or contain an input tree")
    return None


def _parse_args(argv):
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("arm_outputs", type=Path, help="stage-output tree to render from")
    parser.add_argument("out", type=Path, help="output directory (wiped and recreated)")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("uids", nargs="*", default=[], help="uids to render (default: all)")
    group.add_argument("--uids-file", type=Path, default=None,
                        help="file with one uid per line, instead of positional uids")
    parser.add_argument("--source-dir", type=Path, default=None,
                        help="directory holding <uid>*.docx (see docstring)")
    args = parser.parse_args(argv)
    if not (args.arm_outputs / "stage_4_field_extraction").is_dir():
        parser.error(f"no stage_4_field_extraction under {args.arm_outputs} -- point this at "
                     f"a stage-output tree, not a repo root")
    if args.source_dir is not None and not args.source_dir.is_dir():
        # Fail closed (CODING_STANDARDS 5.5), same as doctor_gate.py's
        # identical check: a typo'd --source-dir must not silently render
        # as if the flag were omitted -- omitting it is a documented,
        # intentional mode; a bad path is not.
        parser.error(f"--source-dir is not a directory: {args.source_dir}")
    unsafe_out = _validate_output_dir(args.out, args.arm_outputs, args.source_dir)
    if unsafe_out:
        parser.error(unsafe_out)
    return args


def _atomic_write_json(path: Path, obj) -> None:
    fd, tmp_name = tempfile.mkstemp(dir=path.parent, prefix=path.name + ".", suffix=".tmp")
    try:
        with open(fd, "w", encoding="utf-8") as f:
            json.dump(obj, f, indent=2)
        Path(tmp_name).replace(path)
    except BaseException:
        Path(tmp_name).unlink(missing_ok=True)
        raise


def _uid_filter(uids, uids_file) -> set:
    """The uids the caller asked for; an empty set means "every uid in the arm".

    Blank lines and surrounding whitespace are dropped, so the trailing newline
    every editor writes does not become an empty uid -- which matches nothing
    and would quietly narrow the run rather than fail it.
    """
    if uids_file is None:
        return set(uids)
    return {line.strip()
            for line in uids_file.read_text(encoding="utf-8").splitlines()
            if line.strip()}


def _discover_uids(arm_outputs: Path, only: set) -> list:
    """Every uid this arm holds a stage-4 artifact for, narrowed by `only`."""
    uids = sorted({p.name.replace("_fields.json", "")
                   for p in (arm_outputs / "stage_4_field_extraction").glob("*_fields.json")})
    if only:
        # Stripped, as `_uid_filter` strips each requested uid: harvested
        # filenames can end in a space ("... CV _fields.json"), and an exact
        # match dropped those uids from a --uids-file run without a word.
        uids = [u for u in uids if u.strip() in only]
    return uids


def _resolve_input_artifact(arm_outputs: Path, uid: str):
    """The furthest-along stage artifact this arm holds for `uid`, or None."""
    for stage_dir, pat in PRECEDENCE:
        cand = arm_outputs / stage_dir / pat.format(uid=uid)
        # is_file(), not exists() -- run_stage6 expects a file; exists()
        # would also accept a directory of the same name and hand it to
        # the reader instead of falling through PRECEDENCE (review on #589).
        if cand.is_file():
            return cand
    return None


def _resolve_source_docx(source_dir, uid):
    """The one <uid>*.docx under `source_dir` that belongs to `uid`, or None.

    Module level and not inlined in main() so it can be self-tested without a
    corpus (scripts/test_render_doctor_gates.py) -- the resolver is the only
    part of --source-dir with a rule that can be got subtly wrong, and a
    wrong resolution is silent: it renders another CV's contact block into
    this CV's Personal Data table rather than raising.

    _uid_owns, not a bare prefix match -- glob(f"{uid}*") also matches a
    LONGER uid that starts with this one (the same boundary bug
    doctor_gate.py guards against; see run_doctor.py's docstring on
    _uid_owns for the 2026-07-15 sweep where 3 of 25 CVs were diagnosed
    entirely against another CV's artifacts).
    """
    from unified_pipeline.run_doctor import _uid_owns

    hits = sorted(p for p in source_dir.glob(f"{uid}*.docx") if _uid_owns(p.name, uid))
    return hits[0] if hits else None


def _find_personal_data_table(doc):
    """The table stage 6 fills as PERSONAL DATA, or None.

    Mirrors WCMTemplateGenerator._find_table_with_cell_text
    (stage_6_word_template.py:1339) exactly -- first table with any cell
    containing the anchor, case-insensitively -- so a render that legitimately
    adds or drops some other table cannot shift which table this is.
    """
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if PERSONAL_DATA_ANCHOR in cell.text.lower():
                    return table
    return None


def _validate_rendered_docx(path: Path) -> str | None:
    """Why the rendered document is not worth comparing, or None.

    dest.is_file() accepts a truncated or half-written docx, which is defect 1's
    own shape: a file the compare step will happily fingerprint while proving
    nothing (CODING_STANDARDS 5.5). Stated as invariants rather than counts,
    because both counts here are dynamic -- the document's table count depends
    on which sections had content, and the PERSONAL DATA table's row count
    follows the template. What does not vary is that the six labels the filler
    writes into are all still present, matched by substring on the row's first
    cell, lowercased, the way that filler's own loop matches them
    (stage6/sections/personal_data.py:371-404).
    """
    try:
        doc = Document(str(path))
    except Exception as exc:
        return f"output is not a readable docx: {type(exc).__name__}: {exc}"
    if not doc.tables:
        return "output has no tables -- the WCM template's own tables did not survive"
    table = _find_personal_data_table(doc)
    if table is None:
        return f"no PERSONAL DATA table -- no cell contains {PERSONAL_DATA_ANCHOR!r}"
    # `row.cells` can be empty -- a table row carrying no <w:tc> is legal
    # WordprocessingML, and indexing [0] on one raises rather than reporting
    # a validation failure.
    first_cells = [row.cells[0].text.strip().lower()
                   for row in table.rows if row.cells]
    missing = [label for label in PERSONAL_DATA_LABELS
               if not any(label in cell for cell in first_cells)]
    if missing:
        return f"PERSONAL DATA table lost label row(s): {', '.join(missing)}"
    return None


# Sentinel for "stage 6 has no call_llm at all", which is not the same state as
# "call_llm is None" -- writing None back would leave a name that fails with a
# TypeError where the absence fails with an AttributeError.
_LLM_UNSET = object()


@contextlib.contextmanager
def _llm_disabled(s6):
    """Neutralise stage 6's LLM calls for the duration of the render loop.

    Scoped, not assigned once and left in place: this module is importable (its
    own self-tests import it), and a permanent replacement rides the stage-6
    module object into whatever else the same interpreter runs afterwards --
    hidden global state outliving the run that wanted it. Both call sites in
    stage 6 are try/except guarded, so raising is safe.
    """
    def _no_llm(*a, **kw):
        raise RuntimeError("render_gate: LLM disabled for determinism")

    original = getattr(s6, "call_llm", _LLM_UNSET)
    s6.call_llm = _no_llm
    try:
        yield
    finally:
        if original is _LLM_UNSET:
            delattr(s6, "call_llm")
        else:
            s6.call_llm = original


def _render_uid(s6, src: Path, dest: Path, source_path) -> dict:
    """Render one uid and return its _render_index entry.

    All four outcomes a uid can have are decided here -- crashed, returned
    without writing, wrote something that is not a usable WCM document, or
    succeeded -- because scoring any of the first three as a render is how a
    gate ends up comparing files that prove nothing (CODING_STANDARDS 5.5).
    """
    # One try around the render AND the checks on what it produced. The checks
    # read a document this arm just wrote and can raise on a shape no fixture
    # anticipated; outside the try that would abort the whole gate run at
    # whichever uid hit it, losing every later uid and the index with them --
    # which is the reused-dir failure of defect 1 wearing a different hat. A
    # uid is allowed to fail. A run is not allowed to disappear.
    try:
        s6.run_stage6(input_path=str(src), output_path=str(dest), verbose=False,
                      original_doc_path=str(source_path) if source_path else None)
        # A renderer that returns without raising and without writing the
        # output file is not a successful render -- "no exception" is not
        # "rendered" (review on #589, same fail-closed guarantee this
        # module's docstring claims for defect 1).
        if not dest.is_file():
            return {"error": "render returned without writing an output file"}
        invalid = _validate_rendered_docx(dest)
    except Exception as exc:
        return {"error": f"{type(exc).__name__}: {exc}",
                "tb": traceback.format_exc()[-800:]}
    if invalid:
        return {"error": invalid}
    return {"input": src.name}


def _render_all(s6, arm_outputs: Path, out: Path, source_dir, uids) -> dict:
    """Render every uid in order into the index dict.

    A uid with no resolvable input artifact is recorded and skipped without a
    progress line, which is how this loop has always behaved -- the summary
    and the index still report it.
    """
    results = {}
    with _llm_disabled(s6):
        for i, uid in enumerate(uids, 1):
            src = _resolve_input_artifact(arm_outputs, uid)
            if src is None:
                results[uid] = {"error": "no input artifact"}
                continue
            source_path = _resolve_source_docx(source_dir, uid) if source_dir is not None else None
            results[uid] = _render_uid(s6, src, out / f"{uid}_wcm.docx", source_path)
            if source_dir is not None and "error" not in results[uid]:
                # One-line notice either way -- which docx (if any)
                # backed the fallback write-back for this uid, so a uid
                # with no matching source docx (web08 today) is visible
                # in the index instead of silently rendering as if
                # --source-dir had been omitted (#550).
                results[uid]["source_docx"] = source_path.name if source_path else "not found"
            print(f"[{i}/{len(uids)}] {uid} "
                  f"{'OK' if 'error' not in results[uid] else results[uid]['error']}", flush=True)
    return results


def _write_render_index(out: Path, results: dict) -> None:
    """Persist the per-uid outcome that render_gate_compare.py's guard reads."""
    _atomic_write_json(out / "_render_index.json", results)


def main(argv=None):
    args = _parse_args(argv)
    arm_outputs, out, source_dir = args.arm_outputs, args.out, args.source_dir

    # Read the uid list BEFORE the rmtree below, not after. _validate_output_dir
    # guards `out` against arm_outputs and --source-dir, but a --uids-file can
    # name any path, including one under `out`; reading it second deletes the
    # caller's list and then fails on it.
    only = _uid_filter(args.uids, args.uids_file)

    # Fresh directory every run -- see defect 1 above. _parse_args has already
    # refused an `out` that is, sits inside, or contains an input tree, so this
    # can only reach the scratch directory the caller named for it.
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True)

    import unified_pipeline.stage_6_word_template as s6

    uids = _discover_uids(arm_outputs, only)
    if not uids:
        print("no uids to render -- empty arm, or uid filter matched nothing", file=sys.stderr)
        return 1
    unmatched = sorted(only - {u.strip() for u in uids})
    if unmatched:
        # A partial match is a narrower run than the caller asked for; fail it
        # rather than report a clean arm over fewer CVs.
        print("requested uids not in this arm:", *unmatched, sep="\n  ", file=sys.stderr)
        return 1

    results = _render_all(s6, arm_outputs, out, source_dir, uids)
    _write_render_index(out, results)
    bad = sorted(u for u, r in results.items() if "error" in r)
    print(f"\nrendered={len(results) - len(bad)} failed={len(bad)}")
    if bad:
        print("failed uids:", " ".join(bad))
        return 1  # defect 1: a failed render must never look like a clean arm
    return 0


if __name__ == "__main__":
    sys.exit(main())
